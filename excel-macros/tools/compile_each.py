"""Compiles each procedure of a .bas alone, with the module's declarations, in LibreOffice and names any that does not
compile. A module LibreOffice cannot compile runs nothing and reports nothing, so when every test fails at once, this
says which procedure to look at. Run from excel-macros/: python tools/compile_each.py macros/Hygiene.bas"""
import re
import sys
import tempfile
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tests"))
from harness import MACROS, Office  # noqa: E402
from openpyxl import Workbook, load_workbook
bas = Path(sys.argv[1]).read_text(encoding="utf-8")
i = re.search(r"\n(?:Public |Private )(?:Sub|Function) ", bas).start()
head, body = bas[:i], bas[i:]
procs = [p for p in re.split(r'\n(?=(?:Public |Private )(?:Sub|Function) )', body) if p.strip()]
ping = '\nPublic Sub ZPing()\n    ActiveWorkbook.Worksheets(1).Cells(1, 20).Value = "ping"\nEnd Sub\n'
d = Path(tempfile.mkdtemp(prefix="compile_each_"))
wb = Workbook(); wb.active.title = "use"; wb.save(d / "in.xlsx")
mod = MACROS / "ZCompile.bas"
o = Office()
try:
    for p in procs:
        name = re.match(r'\s*(?:Public|Private) (?:Sub|Function) (\w+)', p).group(1)
        mod.write_text(re.sub(r'Attribute VB_Name = "\w+"', 'Attribute VB_Name = "ZCompile"', head) + ping + "\n" + p)
        o.run(d / "in.xlsx", d / "out.xlsx", [("use", "ZPing")], modules=("ZCompile",))
        if load_workbook(d / "out.xlsx")["use"]["T1"].value != "ping":
            print("DOES NOT COMPILE:", name)
finally:
    o.close(); mod.unlink(missing_ok=True)
print("done")
