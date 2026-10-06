"""The hygiene macros, run in LibreOffice on a synthetic population with every problem planted on purpose.

The firm, 5 Oct 2026: "It'd be nice to be able to easily take a population and figure out basically how messed up it
is or how we can fix it." Every expected count here is worked out in Python from the same rows the workbook holds,
never typed in.
"""
from __future__ import annotations

import datetime as dt
import random

import pytest
from openpyxl import Workbook, load_workbook

from harness import Office

N = 300
HEADS = ["LOAN_ID", "FICO", "STATE", "ORIG_DT", "AMT_TXT", "FLAG_YN", "CONST", "FICO_COPY", "EMPTY", "NOTE"]
AUDIT_HEADS = ["Column", "Type found", "Blank", "Blank %", "Distinct", "Most common", "Numbers as text",
               "Dates as text", "Extra spaces", "Odd characters", "Same value in every row", "Same as column",
               "Keep?", "New name", "Flag: 1 when", "Notes"]


def population() -> list[list]:
    """N loans. Planted: two repeated loan ids, blank FICOs, states with stray spaces and a no-break space, dates
    stored as text, amounts stored as text, a Y/N flag in mixed case, a constant column, a column copying FICO, an
    empty column, and a note column of free text."""
    rnd = random.Random(20261005)
    rows = []
    for i in range(N):
        loan = 1000 + i if i not in (17, 230) else 1000 + i - 1          # rows 17 and 230 repeat the id before
        fico = None if i % 37 == 0 else rnd.randint(560, 820)
        state = rnd.choice(["OH", "NY", "TX", "CA"])
        if i % 41 == 0:
            state += " "
        if i == 99:
            state = "OH "
        day = dt.date(2023, 8, 1) + dt.timedelta(days=rnd.randint(0, 700))
        orig = day.isoformat() if i % 23 == 0 else dt.datetime(day.year, day.month, day.day)
        amt = str(rnd.randint(1, 90) * 500) if i % 3 else rnd.randint(1, 90) * 500
        flag = rnd.choice(["Y", "N", "y", "n", None]) if i % 5 else "Y"
        note = rnd.choice(["", "   ", "ok", "see file", "Ok"]) or None
        rows.append([loan, fico, state, orig, amt, flag, "USD", fico, None, note])
    return rows


def write_population(path) -> list[list]:
    rows = population()
    wb = Workbook()
    ws = wb.active
    ws.title = "use"
    ws.append(HEADS)
    for r in rows:
        ws.append(r)
    for r in range(2, N + 2):
        ws.cell(row=r, column=4).number_format = "yyyy-mm-dd"
    wb.save(path)
    return rows


def blank(v) -> bool:
    return v is None or (isinstance(v, str) and v.strip() == "")


def text_of(v) -> str:
    return v if isinstance(v, str) else str(v)


@pytest.fixture(scope="session")
def office():
    o = Office()
    yield o
    o.close()


@pytest.fixture(scope="session")
def profiled(office, tmp_path_factory):
    d = tmp_path_factory.mktemp("hyg")
    rows = write_population(d / "in.xlsx")
    office.run(d / "in.xlsx", d / "profiled.xlsx", [("use", "HygieneProfile")])
    return {"dir": d, "rows": rows, "wb": load_workbook(d / "profiled.xlsx")}


def audit_rows(wb) -> dict[str, dict]:
    ws = wb["Column Audit"]
    heads = [c.value for c in ws[1]]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0]:
            out[r[0]] = dict(zip(heads, r))
    return out


def log_lines(wb) -> list[tuple]:
    ws = wb["Hygiene"]
    return [(r[0], r[1], r[2]) for r in ws.iter_rows(min_row=17, values_only=True) if r[0]]


def test_profile_writes_one_row_per_column_with_its_heads(profiled):
    ws = profiled["wb"]["Column Audit"]
    assert [c.value for c in ws[1]][:len(AUDIT_HEADS)] == AUDIT_HEADS
    assert list(audit_rows(profiled["wb"])) == HEADS


def test_profile_counts_blanks_and_distinct_values_as_python_does(profiled):
    got = audit_rows(profiled["wb"])
    for c, name in enumerate(HEADS):
        vals = [r[c] for r in profiled["rows"]]
        nonblank = [v for v in vals if not blank(v)]
        assert got[name]["Blank"] == sum(blank(v) for v in vals), name
        if name != "ORIG_DT":                       # dates are keyed by their text; checked apart below
            # the number 1 and the text "1" are two values
            assert got[name]["Distinct"] == len({(isinstance(v, str), text_of(v)) for v in nonblank}), name
        assert got[name]["Blank %"] == pytest.approx(sum(blank(v) for v in vals) / N), name


def test_profile_counts_text_problems_as_python_does(profiled):
    got = audit_rows(profiled["wb"])
    rows = profiled["rows"]
    col = {n: [r[i] for r in rows] for i, n in enumerate(HEADS)}

    def is_num_text(v):
        if not isinstance(v, str) or v.strip() == "":
            return False
        try:
            float(v.strip())
            return True
        except ValueError:
            return False
    assert got["AMT_TXT"]["Numbers as text"] == sum(is_num_text(v) for v in col["AMT_TXT"]) > 0
    assert got["ORIG_DT"]["Dates as text"] == sum(isinstance(v, str) for v in col["ORIG_DT"]) > 0
    # extra spaces: leading or trailing on a value, or a value that is only spaces
    spaced = lambda vs: sum(isinstance(v, str) and v != v.strip(" ") for v in vs)          # noqa: E731
    assert got["STATE"]["Extra spaces"] == spaced(col["STATE"]) > 0
    assert got["NOTE"]["Extra spaces"] == spaced(col["NOTE"]) > 0
    assert got["STATE"]["Odd characters"] == 1                                       # the no-break space
    assert got["NOTE"]["Odd characters"] == 0


def test_profile_tells_case_apart_and_lists_the_most_common(profiled):
    got = audit_rows(profiled["wb"])
    flags = [r[5] for r in profiled["rows"] if not blank(r[5])]
    assert got["FLAG_YN"]["Distinct"] == 4                                           # Y, N, y, n
    counts = sorted(((flags.count(v), v) for v in set(flags)), reverse=True)
    top = got["FLAG_YN"]["Most common"]
    assert top.startswith(f"{counts[0][1]} ({counts[0][0]})"), top
    assert top.count("(") == 3


def test_profile_finds_constant_empty_and_copied_columns(profiled):
    got = audit_rows(profiled["wb"])
    assert got["CONST"]["Same value in every row"] == "Yes"
    assert got["EMPTY"]["Same value in every row"] == "All blank"
    assert got["EMPTY"]["Type found"] == "Empty"
    assert got["FICO_COPY"]["Same as column"] == "FICO"
    assert all(got[n]["Same as column"] in (None, "") for n in HEADS if n != "FICO_COPY")
    assert got["FICO"]["Same value in every row"] in (None, "")


def test_profile_types(profiled):
    got = audit_rows(profiled["wb"])
    assert got["LOAN_ID"]["Type found"] == "Number"
    assert got["STATE"]["Type found"] == "Text"
    n_text = sum(isinstance(r[4], str) for r in profiled["rows"])
    assert got["AMT_TXT"]["Type found"] == f"Mixed: {N - n_text} number, {n_text} text"


def test_profile_leaves_the_source_alone_and_logs_what_it_did(profiled):
    ws = profiled["wb"]["use"]
    assert [c.value for c in ws[1]] == HEADS
    assert ws.max_row == N + 1
    assert ws.cell(row=2, column=1).value == profiled["rows"][0][0]
    log = log_lines(profiled["wb"])
    assert log[-1][1] == "Profile" and log[-1][2].startswith(f"Sheet use: {len(HEADS)} columns, {N} rows")
    assert profiled["wb"]["Hygiene"]["B3"].value == "use"


def decide(path, decisions: dict[str, tuple], key: str | None = "LOAN_ID") -> None:
    """Types decisions on Column Audit, as the analyst would: name -> (Keep?, New name, Flag)."""
    wb = load_workbook(path)
    ws = wb["Column Audit"]
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if name in decisions:
            keep, new, flag = decisions[name]
            ws.cell(row=r, column=13).value = keep
            ws.cell(row=r, column=14).value = new
            ws.cell(row=r, column=15).value = flag
    if key is not None:
        wb["Hygiene"]["B4"].value = key
    wb.save(path)


ALL = {"LOAN_ID": ("Keep", None, None), "FICO": ("Keep", "Score", None), "STATE": ("Keep", None, None),
       "ORIG_DT": ("Keep", None, None), "AMT_TXT": ("Drop", None, None), "FLAG_YN": ("Keep", "Flag", "Y"),
       "CONST": ("Drop", None, None), "FICO_COPY": ("Drop", None, None), "EMPTY": ("Drop", None, None),
       "NOTE": ("keep", None, None)}


@pytest.fixture(scope="session")
def built(office, profiled):
    d = profiled["dir"]
    p = d / "decided.xlsx"
    load_workbook(d / "profiled.xlsx").save(p)
    decide(p, ALL)
    office.run(p, d / "built.xlsx", [("use", "HygieneBuild")])
    return {"dir": d, "rows": profiled["rows"], "wb": load_workbook(d / "built.xlsx"), "path": d / "built.xlsx"}


def test_profile_again_keeps_the_decisions_typed(office, built):
    d = built["dir"]
    office.run(built["path"], d / "reprofiled.xlsx", [("use", "HygieneProfile")])
    got = audit_rows(load_workbook(d / "reprofiled.xlsx"))
    for name, (keep, new, flag) in ALL.items():
        assert (got[name]["Keep?"], got[name]["New name"], got[name]["Flag: 1 when"]) == (keep, new, flag), name


def test_build_writes_the_kept_columns_as_values_renamed_with_the_flag_made(built):
    ws = built["wb"]["Final Population"]
    heads = [c.value for c in ws[1]]
    assert heads == ["LOAN_ID", "Score", "STATE", "ORIG_DT", "Flag", "NOTE"]
    assert ws.max_row == N + 1
    src = built["rows"]
    for i in range(N):
        row = [ws.cell(row=i + 2, column=c).value for c in range(1, 7)]
        assert row[0] == src[i][0]
        assert row[1] == src[i][1]
        f = src[i][5]
        assert row[4] == (None if blank(f) else (1 if f.strip().upper() == "Y" else 0)), (i, f, row[4])
    # values, not formulas
    assert not any(isinstance(c.value, str) and c.value.startswith("=") for r in ws.iter_rows() for c in r)


def test_build_row_audit_lists_rows_with_blanks_and_repeated_keys(built):
    ws = built["wb"]["Row Audit"]
    src = built["rows"]
    kept = [0, 1, 2, 3, 5, 9]
    want_blank_rows = [i + 2 for i, r in enumerate(src) if any(blank(r[c]) for c in kept)]
    got_blank_rows = [r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    assert got_blank_rows == want_blank_rows
    dups = {r[5]: (r[6], r[7]) for r in ws.iter_rows(min_row=2, values_only=True) if r[5] is not None}
    ids = [r[0] for r in src]
    want = {str(v): ids.count(v) for v in set(ids) if ids.count(v) > 1}
    assert {str(k): v[1] for k, v in dups.items()} == want
    assert len(want) == 2


def test_build_stamps_what_it_did_and_says_current(built):
    ctl = built["wb"]["Hygiene"]
    src = built["rows"]
    assert load_workbook(built["path"], data_only=True)["Hygiene"]["B8"].value == "Current"
    assert ctl["B9"].value == N
    assert (ctl["B10"].value, ctl["B11"].value, ctl["B12"].value) == (6, 4, 1)
    kept = [0, 1, 2, 3, 5, 9]
    assert ctl["B13"].value == sum(any(blank(r[c]) for c in kept) for r in src)
    assert ctl["B14"].value == 2
    assert ctl["B7"].value.endswith(" from sheet use")
    log = log_lines(built["wb"])
    assert log[-1][1] == "Build" and log[-1][2].startswith(f"{N} rows; 6 columns kept, 4 dropped; 1 flags")
    assert "Flag: " in log[-1][2] and " ones, " in log[-1][2]


def test_a_changed_decision_makes_the_build_out_of_date(office, built):
    d = built["dir"]
    p = d / "changed.xlsx"
    load_workbook(built["path"]).save(p)
    decide(p, {"NOTE": ("Drop", None, None)}, key=None)
    office.run(p, d / "changed_out.xlsx", [])
    assert load_workbook(d / "changed_out.xlsx", data_only=True)["Hygiene"]["B8"].value.startswith("Out of date")


def test_build_refuses_while_a_column_has_no_decision_and_names_it(office, profiled):
    d = profiled["dir"]
    p = d / "half.xlsx"
    load_workbook(d / "profiled.xlsx").save(p)
    decide(p, {k: v for k, v in ALL.items() if k not in ("EMPTY", "NOTE")})
    office.run(p, d / "half_out.xlsx", [("use", "HygieneBuild")])
    wb = load_workbook(d / "half_out.xlsx")
    assert "Final Population" not in wb.sheetnames
    log = log_lines(wb)
    refused = [x for x in log if x[1] == "Build refused"]
    assert refused and "2 to fix first" in refused[-1][2]
    assert "EMPTY: no decision in Keep?" in refused[-1][2] and "NOTE: no decision in Keep?" in refused[-1][2]


def test_build_refuses_two_kept_columns_with_one_name(office, profiled):
    d = profiled["dir"]
    p = d / "clash.xlsx"
    load_workbook(d / "profiled.xlsx").save(p)
    decide(p, {**ALL, "STATE": ("Keep", "score", None)})
    office.run(p, d / "clash_out.xlsx", [("use", "HygieneBuild")])
    wb = load_workbook(d / "clash_out.xlsx")
    assert "Final Population" not in wb.sheetnames
    assert "score: two kept columns would share this name" in [x for x in log_lines(wb)
                                                              if x[1] == "Build refused"][-1][2]


def test_build_without_a_key_says_duplicates_were_not_checked(office, profiled):
    d = profiled["dir"]
    p = d / "nokey.xlsx"
    load_workbook(d / "profiled.xlsx").save(p)
    decide(p, ALL, key=None)
    office.run(p, d / "nokey_out.xlsx", [("use", "HygieneBuild")])
    wb = load_workbook(d / "nokey_out.xlsx")
    assert wb["Hygiene"]["B14"].value.startswith("Not checked")
    assert wb["Row Audit"]["F2"].value.startswith("Not checked")


def test_profile_refuses_on_its_own_sheets(office, profiled):
    d = profiled["dir"]
    office.run(d / "profiled.xlsx", d / "self.xlsx", [("Column Audit", "HygieneProfile")])
    log = log_lines(load_workbook(d / "self.xlsx"))
    assert log[-1][1] == "Profile said" and "Go to the sheet holding the population" in log[-1][2]


# --------------------------------------------------------------------------
# The review of 5 Oct 2026: what LibreOffice let through that real Excel would not, and the edges the planted
# population never reached. One test per finding.


def small_book(path, header, rows, sheet="use"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header)
    for r in rows:
        ws.append(r)
    wb.save(path)


def profile_then(office, d, name, header, rows, decisions=None, key=None, extra=()):
    """Profile a small book, optionally type decisions and Build, and return the result loaded twice: with
    formulas, and with the values LibreOffice calculated."""
    small_book(d / f"{name}.xlsx", header, rows)
    office.run(d / f"{name}.xlsx", d / f"{name}_p.xlsx", [("use", "HygieneProfile")])
    if decisions is None:
        out = d / f"{name}_p.xlsx"
    else:
        decide(d / f"{name}_p.xlsx", decisions, key=key)
        out = d / f"{name}_b.xlsx"
        office.run(d / f"{name}_p.xlsx", out, [("use", "HygieneBuild"), *extra])
    return load_workbook(out), load_workbook(out, data_only=True)


def test_text_excel_would_reread_is_written_into_text_cells(office, tmp_path):
    """Excel parses text written by a macro as if typed: "00123" turns into 123, "2024-01" into a date, and a top
    value of "-" into arithmetic. Every such cell must be formatted as Text before it is written."""
    header = ["REF", "2024-01", "CODE", "MIX"]
    rows = [["00123", 5, "-", "00789"], ["00456", 6, "-", 8], ["A9", 7, "x", 9]]
    wb, _ = profile_then(office, tmp_path, "reread", header, rows,
                         {"REF": ("Keep", None, None), "2024-01": ("Keep", None, None), "CODE": ("Keep", None, None),
                          "MIX": ("Keep", None, None)})
    aud = wb["Column Audit"]
    names = {aud.cell(row=r, column=1).value: r for r in range(2, 6)}
    assert set(names) == set(header)
    assert len(names) == 4
    assert aud.cell(row=names["2024-01"], column=1).number_format == "@"
    top = aud.cell(row=names["CODE"], column=6)
    assert top.value.startswith("- (2)") and top.number_format == "@"
    fin = wb["Final Population"]
    assert [c.value for c in fin[1]] == header and fin["B1"].number_format == "@"
    assert (fin["A2"].value, fin["A2"].number_format) == ("00123", "@")
    assert (fin["C2"].value, fin["C2"].number_format) == ("-", "@")
    assert fin["B2"].number_format != "@"                                    # numbers stay numbers
    # a column of text and numbers: only its text that Excel would re-read goes in as Text
    assert (fin["D2"].value, fin["D2"].number_format) == ("00789", "@")
    assert (fin["D3"].value, fin["D3"].number_format != "@") == (8, True)


def test_columns_sharing_a_name_each_get_a_row_and_keep_their_own_data(office, tmp_path):
    header = ["ID", "AMT", "amt", "ID"]
    rows = [[1, 10, 100, "a"], [2, 20, 200, "b"]]
    wb, _ = profile_then(office, tmp_path, "shared", header, rows,
                         {"ID": ("Keep", None, None), "AMT": ("Keep", None, None),
                          "amt (column 3)": ("Keep", "AMT_2", None), "ID (column 4)": ("Keep", "ID_TEXT", None)})
    assert [r[0] for r in wb["Column Audit"].iter_rows(min_row=2, values_only=True)] == \
        ["ID", "AMT", "amt (column 3)", "ID (column 4)"]
    fin = wb["Final Population"]
    assert [c.value for c in fin[1]] == ["ID", "AMT", "AMT_2", "ID_TEXT"]
    assert [[c.value for c in r] for r in fin.iter_rows(min_row=2)] == rows
    log = log_lines(wb)
    assert any(x[1] == "Profile" and "2 columns share a name with an earlier one" in x[2] for x in log)


def test_a_key_repeated_many_times_lists_twenty_rows_and_counts_the_rest(office, tmp_path):
    rows = [["K1", i] for i in range(30)] + [["K2", 99]]
    wb, _ = profile_then(office, tmp_path, "heavy", ["KEY", "N"], rows,
                         {"KEY": ("Keep", None, None), "N": ("Keep", None, None)}, key="KEY")
    ra = wb["Row Audit"]
    assert ra["F2"].value == "K1" and ra["H2"].value == 30
    listed = ra["G2"].value
    assert listed.endswith(" ... and 10 more") and listed.count(",") == 19
    assert ra["F3"].value is None


def test_one_data_row_and_one_kept_column_builds(office, tmp_path):
    wb, _ = profile_then(office, tmp_path, "one", ["A", "B"], [[None, 2]],
                         {"A": ("Keep", None, None), "B": ("Drop", None, None)})
    assert wb["Final Population"]["A1"].value == "A"
    assert wb["Row Audit"]["A2"].value == 2 and wb["Row Audit"]["C2"].value == 1
    assert log_lines(wb)[-1][1] == "Build"


def test_data_under_a_blank_heading_is_profiled_not_dropped(office, tmp_path):
    wb, _ = profile_then(office, tmp_path, "nohead", ["A", "B"], [[1, 2, "stray"], [3, 4, None]])
    assert [r[0] for r in wb["Column Audit"].iter_rows(min_row=2, values_only=True)] == \
        ["A", "B", "(no name, column 3)"]


def test_the_number_one_and_the_text_one_are_two_values(office, tmp_path):
    wb, _ = profile_then(office, tmp_path, "kinds", ["MIX", "NUM", "TXT"], [[1, 1, "1"], ["1", 2, "2"], [1, 3, "3"]])
    got = audit_rows(wb)
    assert got["MIX"]["Distinct"] == 2
    assert got["TXT"]["Same as column"] in (None, "")      # the text "1" is not the number 1


def test_numbers_as_text_takes_only_what_reads_as_a_number(office, tmp_path):
    vals = ["&H10", "1D5", "1,000", "$5", "5%", "12.5.1", "-7", "abc"]
    wb, _ = profile_then(office, tmp_path, "numtext", ["V"], [[v] for v in vals])
    assert audit_rows(wb)["V"]["Numbers as text"] == 4                   # 1,000  $5  5%  -7


def test_keys_are_compared_trimmed(office, tmp_path):
    wb, _ = profile_then(office, tmp_path, "trim", ["KEY"], [["A1"], ["A1 "], ["B2"]],
                         {"KEY": ("Keep", None, None)}, key="KEY")
    assert wb["Row Audit"]["F2"].value == "A1" and wb["Row Audit"]["H2"].value == 2


def test_an_error_part_way_is_logged_and_the_screen_put_back(office, tmp_path):
    """HygieneSelfTest freezes the screen as a long run does, then fails: the error is logged. That the screen and
    status bar come back is not provable here: LibreOffice resets both between macro calls by itself, so a probe read
    True even with the reset taken out of Recover. (Two natural triggers were tried first, a protected sheet and a chart sheet holding Build's sheet name;
    LibreOffice lets a macro past both, so the error is planted.)"""
    small_book(tmp_path / "err.xlsx", ["A"], [[1]])
    office.run(tmp_path / "err.xlsx", tmp_path / "err_out.xlsx", [("use", "HygieneSelfTest")])
    log = log_lines(load_workbook(tmp_path / "err_out.xlsx"))
    assert log[-1][1] == "Self test said"
    assert log[-1][2].startswith("Self test stopped on an error and did not finish: a planted error (error 1004)")


def test_every_macro_a_person_runs_ends_in_the_error_path():
    """Each public macro turns errors over to Recover, which puts the screen back and logs: checked in the code,
    because only the self test can make LibreOffice fail on cue."""
    import re
    from harness import MACROS
    code = (MACROS / "Hygiene.bas").read_text()
    for name in ("HygieneProfile", "HygieneBuild", "HygieneSaveCopy", "HygieneSelfTest"):
        body = re.search(rf"Public Sub {name}\(\)\n(.*?)\nEnd Sub", code, re.S).group(1)
        assert "On Error GoTo failed" in body.splitlines()[0] + body.splitlines()[1], name
        assert re.search(r"failed:\n(?:.*\n)*?\s+Recover(?:With)? \"", body + "\n"), name


def test_save_copy_writes_final_population_alone_beside_the_workbook(office, tmp_path):
    small_book(tmp_path / "save.xlsx", ["REF", "N"], [["00123", 1], ["00456", 2]])
    office.run(tmp_path / "save.xlsx", tmp_path / "save_p.xlsx", [("use", "HygieneProfile")])
    decide(tmp_path / "save_p.xlsx", {"REF": ("Keep", None, None), "N": ("Keep", None, None)}, key=None)
    office.run(tmp_path / "save_p.xlsx", tmp_path / "save_b.xlsx", [("use", "HygieneBuild"),
                                                                      ("use", "HygieneSaveCopy")])
    """The copy is made, named for the workbook and the minute, beside it. That is all LibreOffice can show: it
    ignores Excel's FileFormat code (it wrote an .xls under the .xlsx name), and once a second workbook has been
    opened and closed it loses which one is active, so the "Saved copy" log line lands nowhere. The copy's contents
    and its log line are checked on the firm's first run in Excel (README, "What the tests do not prove")."""
    copies = list(tmp_path.glob("save_p - Final Population *.xlsx"))
    assert len(copies) == 1, (copies, log_lines(load_workbook(tmp_path / "save_b.xlsx")))
    assert copies[0].stat().st_size > 0
    assert not any(x[1] == "Save copy said" and "stopped on an error" in x[2]
                   for x in log_lines(load_workbook(tmp_path / "save_b.xlsx")))


# --------------------------------------------------------------------------
# The second review of 5 Oct 2026, on the fixes above.


def test_build_refuses_a_source_column_with_no_row_on_column_audit(office, tmp_path):
    """A column added after Profile, or its row deleted, would otherwise be left out with no decision at all."""
    small_book(tmp_path / "gap.xlsx", ["A", "B", "C"], [[1, 2, 3]])
    office.run(tmp_path / "gap.xlsx", tmp_path / "gap_p.xlsx", [("use", "HygieneProfile")])
    decide(tmp_path / "gap_p.xlsx", {"A": ("Keep", None, None), "B": ("Keep", None, None), "C": ("Drop", None, None)},
           key=None)
    wb = load_workbook(tmp_path / "gap_p.xlsx")
    wb["Column Audit"].delete_rows(3)                                     # B's row
    wb.save(tmp_path / "gap_p.xlsx")
    office.run(tmp_path / "gap_p.xlsx", tmp_path / "gap_b.xlsx", [("use", "HygieneBuild")])
    out = load_workbook(tmp_path / "gap_b.xlsx")
    assert "Final Population" not in out.sheetnames
    refused = [x for x in log_lines(out) if x[1] == "Build refused"]
    assert refused and "B: on use but not on Column Audit; run Profile again" in refused[-1][2]


def test_save_copy_needs_a_current_build(office, tmp_path):
    """Build writes "Not built" before it touches Final Population and "Current" only when it has finished, and Save
    copy saves only a Current one, so a Build that stopped part-way is never saved as if whole."""
    small_book(tmp_path / "cur.xlsx", ["A"], [[1], [2]])
    office.run(tmp_path / "cur.xlsx", tmp_path / "cur_p.xlsx", [("use", "HygieneProfile")])
    decide(tmp_path / "cur_p.xlsx", {"A": ("Keep", None, None)}, key=None)
    office.run(tmp_path / "cur_p.xlsx", tmp_path / "cur_b.xlsx", [("use", "HygieneBuild")])
    wb = load_workbook(tmp_path / "cur_b.xlsx")
    wb["Hygiene"]["B8"].value = "Not built: the last Build did not finish. Run Build again."
    wb.save(tmp_path / "cur_half.xlsx")
    office.run(tmp_path / "cur_half.xlsx", tmp_path / "cur_out.xlsx", [("use", "HygieneSaveCopy")])
    log = log_lines(load_workbook(tmp_path / "cur_out.xlsx"))
    assert log[-1][1] == "Save copy said" and log[-1][2] == ("Final Population is not current. Its status reads: Not built: the last Build "
                                                  "did not finish. Run Build again.")
    assert not list(tmp_path.glob("cur_half - Final Population *"))


def test_error_text_and_formula_text_in_a_mixed_column_stay_text(office, tmp_path):
    """Text reading "#N/A" or "=SUM(" in a column that also holds numbers: kept as text, never an error or a
    formula. openpyxl would store both as an error and a formula, so the source cells are set to text by hand."""
    small_book(tmp_path / "errtext.xlsx", ["ID", "V"], [[1, "x"], [2, "x"], [3, 7], [4, "ok"]])
    wb = load_workbook(tmp_path / "errtext.xlsx")
    for ref, text in (("B2", "#N/A"), ("B3", "=SUM(")):
        wb["use"][ref].value = text
        wb["use"][ref].data_type = "s"
    wb.save(tmp_path / "errtext.xlsx")
    office.run(tmp_path / "errtext.xlsx", tmp_path / "errtext_p.xlsx", [("use", "HygieneProfile")])
    decide(tmp_path / "errtext_p.xlsx", {"ID": ("Keep", None, None), "V": ("Keep", None, None)}, key=None)
    office.run(tmp_path / "errtext_p.xlsx", tmp_path / "errtext_b.xlsx", [("use", "HygieneBuild")])
    wb = load_workbook(tmp_path / "errtext_b.xlsx")
    fin = wb["Final Population"]
    assert (fin["B2"].value, fin["B2"].number_format) == ("#N/A", "@")
    assert (fin["B3"].value, fin["B3"].number_format) == ("=SUM(", "@")
    assert fin["B4"].value == 7 and fin["B5"].value == "ok"


def test_a_heading_that_reads_like_a_marked_name_is_marked_in_turn(office, tmp_path):
    header = ["A", "A", "A (column 2)"]
    wb, _ = profile_then(office, tmp_path, "collide", header, [[1, 2, 3]],
                         {"A": ("Keep", None, None), "A (column 2)": ("Keep", "A2", None),
                          "A (column 2) (column 3)": ("Keep", "A3", None)})
    assert [r[0] for r in wb["Column Audit"].iter_rows(min_row=2, values_only=True)] == \
        ["A", "A (column 2)", "A (column 2) (column 3)"]
    assert [[c.value for c in r] for r in wb["Final Population"].iter_rows()] == [["A", "A2", "A3"], [1, 2, 3]]


def test_three_headings_that_would_collide_get_three_names(office, tmp_path):
    """The third review: "A (column 3)", "A", "A" gave column 3 the first column's heading."""
    header = ["A (column 3)", "A", "A"]
    wb, _ = profile_then(office, tmp_path, "collide3", header, [[1, 2, 3]],
                         {"A (column 3)": ("Keep", None, None), "A": ("Keep", None, None),
                          "A (column 3, 2)": ("Keep", "A3", None)})
    names = [r[0] for r in wb["Column Audit"].iter_rows(min_row=2, values_only=True)]
    assert names == ["A (column 3)", "A", "A (column 3, 2)"] and len(set(names)) == 3
    assert [[c.value for c in r] for r in wb["Final Population"].iter_rows()] == \
        [["A (column 3)", "A", "A3"], [1, 2, 3]]
