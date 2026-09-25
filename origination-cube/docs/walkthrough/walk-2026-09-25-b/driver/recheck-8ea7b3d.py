import sys, shutil, os, csv
sys.path.insert(0, "/tmp/walk2/code2/origination-cube/src")
from pathlib import Path
from openpyxl import load_workbook
from origination_cube import book, meanings, synth
R = Path("/tmp/walk2/re"); mem = R / "mem.yaml"
for p in R.glob("d*"): shutil.rmtree(p)
d = R / "d1"; d.mkdir()
synth.write_extract(str(d), n=5000); (d / "loans.csv").rename(d / "Consumer book Q3.csv")
ex = d / "Consumer book Q3.csv"
o = book.set_up(ex, memory_path=mem); bk = o.book
print("SETUP:", o.lines)
wb = load_workbook(bk); ws = wb["Columns"]
print("COLUMNS HEAD:", [c.value for c in ws[5]][1:])
for r in ws.iter_rows(min_row=6, values_only=True): print("  ", r[1:8])
print("category says:", meanings.catalog()["category"].says, "| unknown:", meanings.catalog()["unknown"].says)
o = book.run(bk, memory_path=mem); print("RUN0:", [l for l in o.lines if "judged" in l or "C18" in l])
def put(path, **cells):
    wb = load_workbook(path)
    for ref, v in cells.items():
        s, c = ref.split("__"); wb[s.replace("_", " ")][c].value = v
    wb.save(path)
put(bk, Control__C6="Every loan", Control__C13=30, Control__C14=10, Control__C16="1% of the book's total losses",
    Control__C18="The rest of its band", Control__C19="1.25 times", Control__C20="0.8 times", Control__C22=0.95,
    Columns__C3="Yes", Odd_values__E5="missing", Odd_values__E6="real")
o = book.run(bk, memory_path=mem); print("RUN 0.95:", o.lines)
put(bk, Control__C22="95%")
o = book.run(bk, memory_path=mem); print("RUN ok:", o.lines[:2])
wb = load_workbook(bk); print("TABS:", wb.sheetnames)
print("START HERE:", [r for r in wb["Start here"].iter_rows(min_row=12, values_only=True)])
print("CHECK materiality/evidence rows:", [r[1:3] for r in wb["Check"].iter_rows(min_row=4, values_only=True) if r[1] and ("ateria" in str(r[1]) or "vidence" in str(r[1]))])
# FICO row edges as a number
fico_row = next(r[0].row for r in wb["Columns"].iter_rows(min_row=6) if r[1].value == "FICO")
put(bk, **{f"Columns__F{fico_row}": 620680740})
o = book.run(bk, memory_path=mem); print("RUN edges 620680740:", o.lines[:3])
put(bk, **{f"Columns__F{fico_row}": "620, 680, 740"})
# Forget
wb = load_workbook(bk); lr = next(r[0].row for r in wb["Learned"].iter_rows(min_row=4) if r[2].value == "CHANNEL")
put(bk, **{f"Learned__A{lr}": "Forget"})
o = book.run(bk, memory_path=mem)
wb = load_workbook(bk); print("AFTER FORGET:", [r[:7] for r in wb["Learned"].iter_rows(min_row=4, values_only=True) if r[2] == "CHANNEL"], o.lines[-2:])
# Set up again
o = book.set_up(ex, memory_path=mem); wb = load_workbook(bk); print("SETUP AGAIN tabs:", wb.sheetnames, [r for r in wb["Start here"].iter_rows(min_row=12, values_only=True)][-1])
# new columns
rows = list(csv.reader(open(ex))); rows = [rows[0] + ["CURR_STATUS", "DTI"]] + [r + ["Current", "0.3"] for r in rows[1:]]
csv.writer(open(ex, "w", newline="")).writerows(rows)
o = book.set_up(ex, memory_path=mem); wb = load_workbook(bk); print("NEW COLS C3:", wb["Columns"]["C3"].value)
o = book.run(bk, memory_path=mem); print("NEW COLS RUN:", o.lines[:1])
# copied folder
d2 = R / "d2"; d2.mkdir(); shutil.copy(bk, d2 / bk.name); synth.write_extract(str(d2), n=3000); (d2 / "loans.csv").rename(d2 / ex.name)
o = book.run(d2 / bk.name, memory_path=mem); print("COPIED FOLDER RUN (3,000-loan extract beside it):", o.lines[:1])
