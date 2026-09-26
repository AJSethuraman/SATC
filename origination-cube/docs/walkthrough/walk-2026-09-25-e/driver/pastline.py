"""Losses vs revenue, read back: pockets drawn past a Control line on the chart but boxed as "the same".

    python3.12 pastline.py BOOK.xlsx...

The chart draws the four Control lines as dashed lines; the box also needs the side's own test
to say "not luck". This counts, per workbook, the pockets whose dot sits past a line on either
axis while the table puts that side at "the same"."""
import re, sys
from openpyxl import load_workbook
for p in sys.argv[1:]:
    ws = load_workbook(p)["Losses vs revenue"]
    sub = ws["B2"].value
    gw, gb = map(float, re.search(r"GCO counts as more at ([\d.]+)x or above and less at ([\d.]+)x", sub).groups())
    rw, rb = map(float, re.search(r"Revenue counts as more at ([\d.]+)x or above and less at ([\d.]+)x", sub).groups())
    n = past_g = past_r = 0
    ex = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not isinstance(r[3], (int, float)) or not r[8] or r[8].startswith("Not tested"):
            continue
        n += 1
        box = r[8]
        gsame = box.startswith("Losing the same") or box == "About the same on both"
        rsame = box.endswith("earning the same") or box == "About the same on both"
        if gsame and (r[4] >= gw or r[4] <= gb):
            past_g += 1; ex.append((r[1], r[2], round(r[4], 2), r[5], "GCO"))
        if rsame and (r[6] >= rw or r[6] <= rb):
            past_r += 1; ex.append((r[1], r[2], round(r[6], 2), r[7], "RANR"))
    print(f"{p.split('/')[-2]}: {n} boxed pockets; past a GCO line but 'the same': {past_g}; "
          f"past a revenue line but 'the same': {past_r}  (GCO {gb}/{gw}, revenue {rb}/{rw})")
    for e in ex[:6]:
        print("   ", e)
