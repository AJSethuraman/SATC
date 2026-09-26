"""The Three-way tab, read back: how many split pockets are flagged worse, by band column and half.

    python3.12 threeway.py BOOK.xlsx [N]

Defect 1: on the no-effect book, which rows still read worse, and on the planted book, how many
come from grids that don't hold the score fixed."""
import sys, collections
from openpyxl import load_workbook
ws = load_workbook(sys.argv[1])["Three-way"]
print(ws["B2"].value)
hdr = None; rows = []
for r in ws.iter_rows(values_only=True):
    if r[1] == "Measure": hdr = list(r); continue
    if hdr and r[1]: rows.append(dict(zip(hdr, r)))
print(len(rows), "rows; headers:", [h for h in hdr if h])
fl = [h for h in hdr if h and h.startswith("Flag")][0]
c = collections.Counter()
for d in rows:
    half = "high" if "high half" in str(d["Segment"]) else "low" if "low half" in str(d["Segment"]) else "value"
    c[(d["Measure"], d["Band column"], half, d[fl])] += 1
for k, v in sorted(c.items(), key=str):
    if k[3] in ("worse",): print(f"{v:3}", k)
for d in rows[:int(sys.argv[2]) if len(sys.argv) > 2 else 5]: print({k: d[k] for k in ("Measure","Band column","Band","Segment","Loans","Excess",fl)})
