"""Losses vs revenue, read back: each row's box set against its flags and its dollars over the book.

    python3.12 boxcheck.py BOOK.xlsx [N]   (prints the first N rows too)

Defects 2 and 3 come from its counts: boxes whose GCO side disagrees with the excess over the book,
and boxes placed on a GCO the tab calls too few losses to test."""
"""Read the Losses vs revenue tab back and set each row's box against its flags and dollars."""
import sys, collections
from openpyxl import load_workbook
ws = load_workbook(sys.argv[1])["Losses vs revenue"]
print("subtitle:", ws["B2"].value)
grid = None; rows = []
for r in ws.iter_rows(min_row=4, values_only=True):
    v = r[1:11]
    if v[0] and v[1] is None and v[2] is None: grid = v[0]; continue
    if v[0] == "Band" or v[2] is None or not isinstance(v[2], (int, float)): continue
    rows.append((grid,) + tuple(v))
print(len(rows), "rows")
c = collections.Counter()
for g, bl, dl, n, gi, gf, ri, rf, box, gx, rx in rows:
    gside = box.split(",")[0] if "," in box else box
    rside = box.split(",")[1].strip() if "," in box else box
    if "more" in gside and gx is not None and gx < 0: c["Losing more but GCO excess over book < 0"] += 1
    if "less" in gside and gx is not None and gx > 0: c["Losing less but GCO excess over book > 0"] += 1
    if "too few" in (gf or "") and "the same" not in gside and "About" not in box: c["box placed on an untested GCO"] += 1
    if "too few" in (rf or "") : c["RANR untested"] += 1
    if ("same" in rside or "About" in box) and rf in ("better", "worse"): c["revenue 'same' but RANR flag significant"] += 1; print("  same-but-significant:", g, bl, dl, n, ri, rf, box)
    if ("more" in rside or "less" in rside) and rf == "in line": c["revenue more/less but RANR flag in line"] += 1
    if ("same" in gside or "About" in box) and gf in ("worse", "better"): c["GCO 'same' but flag significant"] += 1; print("  gco-same-but-significant:", g, bl, dl, n, gi, gf, box)
for k, v in c.items(): print(f"{v:4}  {k}")
if len(sys.argv) > 2:
    for row in rows[: int(sys.argv[2])]: print(row)
print("--- boxes x RANR flag")
bx = collections.Counter((row[8], row[7]) for row in rows)
for k, v in sorted(bx.items()): print(f"{v:4}  {k}")
print("--- boxes x GCO flag (non-'About')")
bx = collections.Counter((row[8], row[5]) for row in rows if row[8] != "About the same on both")
for k, v in sorted(bx.items()): print(f"{v:4}  {k}")
