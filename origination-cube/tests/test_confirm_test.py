"""The confirmatory test (capabilities 4b and 4e): the pre-spec's column in its own groups, against its reference
group, inside its pockets, on the development range and the holdout; and how concentrated the bad loans are on
the holdout.

The goal it finishes (the firm, 25 Sep 2026): "the cube should be ready to test a derived column (income /
sales) against a dated outcome, on a holdout, from a committed pre-spec." It ends when, on the dated synthetic
book with the example pre-spec, a run finds the cliffs on development and confirms them on the holdout, and
Check no longer says "deviates" on the reference group. The first tests here are those two sentences.

Every count is checked against the extract read by hand (csv), never against the cube's own numbers."""

import bisect
import csv
import math
import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import cube, table
from origination_cube import book, confirm_tab, confirmatory, control, engine, kgroups, prespec, stats, synth
from recalc import recalc
from test_book import _answer
from test_book_dates import _check, _columns, _control

needs_git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed on this machine")
EXAMPLE = Path(__file__).resolve().parents[1] / "docs" / "prespec-example.yaml"
BINS = [0.1, 0.25, 0.5, 1.0, 2.0]
PLANTED = {0: 2.0, 5: 3.0}                  # synth.RATIO_ODDS: below 0.1 twice the odds, above 2.0 three times


def _git(repo, *args):
    subprocess.run(["git", "-c", "user.name=Test", "-c", "user.email=test@example.com", "-c", "commit.gpgsign=false",
                    *args], cwd=repo, check=True, capture_output=True)


def _set_up(folder: Path, n: int, spec: Path = EXAMPLE, cut_no=("ORIG_BAL", "REV_DEBT", "ASSET_CLASS", "INCOME",
                                                                 "SALES", "income_to_sales")):
    """The dated synthetic book, income / sales made on Control under the example pre-spec's name, testing from
    that pre-spec, committed beside the workbook."""
    x = synth.write_extract(folder, n=n, ratio=True)
    b = book.set_up(x).book
    _answer(b)
    _control(b, run_kind="Finding and testing a new variable", new_variable_step="Test from a pre-spec",
             **{"derived|1": ("income_to_sales", "INCOME", "SALES")})
    book.set_up(x)
    for c in cut_no:
        _columns(b, c, C_CUT="No")
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    ws = wb[control.SHEET]
    ws.cell(row=control.row_of(ws, control.PRESPEC_KEY), column=control.CHOOSE_COL).value = "prespec.yaml"
    wb.save(b)
    shutil.copy(spec, folder / "prespec.yaml")
    _git(folder, "init", "-q")
    _git(folder, "add", "prespec.yaml")
    _git(folder, "commit", "-q", "-m", "pre-spec")
    return x, b


@pytest.fixture(scope="module")
def route(tmp_path_factory):
    """The goal's run: 20,000 loans (synth's default), the example pre-spec as committed in docs/."""
    folder = tmp_path_factory.mktemp("goal")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("GIT_CEILING_DIRECTORIES", str(folder))
        mp.setenv("CUBE_MEMORY", str(folder / "memory.yaml"))
        x, b = _set_up(folder, 20000)
        seen = {}
        real = confirmatory.state

        def spy(*a, **k):
            seen["st"] = real(*a, **k)
            return seen["st"]

        mp.setattr(confirmatory, "state", spy)
        ran = book.run(b)
    assert ran.ok, ran.lines
    return {"x": x, "b": b, "ran": ran, "st": seen["st"], "t": seen["st"].test}


def _rows(x):
    with open(x, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _by_hand(x, lo: str, hi: str):
    """Loans made lo to hi (both ends in) with a readable ratio and outcome, as (group, bad, gco or None)."""
    out = []
    for r in _rows(x):
        if not lo <= r["ORIG_DATE"] <= hi or r["SALES"] in ("", "0") or r["BAD_FLAG"] not in ("0", "1"):
            continue
        g = bisect.bisect_right(BINS, float(r["INCOME"]) / float(r["SALES"]))
        try:
            gco = float(r["GCO_AMT"])
        except ValueError:
            gco = None
        out.append((g, int(r["BAD_FLAG"]), gco))
    return out


# --------------------------------------------------------------------------
# The goal


def test_the_goal_the_cliffs_are_found_on_development_and_confirmed_on_the_holdout(route):
    t = route["t"]
    assert t.problem is None and t.method == kgroups.CONDITIONAL
    assert t.column == "income_to_sales" and t.groups[t.ref] == "0.25 - 0.49" and t.strata == ("FICO", "CHANNEL")
    for side in (t.development, t.holdout):
        a, f = side.association, side.fit
        assert a.p_general < 0.05 and f.p_block < 0.05 and f.settled
        for k, planted in PLANTED.items():
            # each cliff is found: worse than the reference, significant, and the planted odds inside its range
            lo, hi = f.interval(k, 1.96)
            assert f.odds[k] > 1 and f.p[k] < 0.05 and lo < planted < hi, (side.name, k, f.odds[k], lo, hi)
    # on development, where nothing was planted between the cliffs, nothing between them reads significant
    assert all(t.development.fit.p[k] > 0.05 for k in (1, 3, 4))


def test_the_goal_check_no_longer_says_deviates(route):
    chk = _check(route["b"])
    warns = chk.get("Warning", [])
    warns = warns if isinstance(warns, list) else [warns]
    assert not [w for w in warns if w.startswith(confirmatory.DEVIATES)]
    assert chk["Differs from the pre-spec"] == "nowhere: this run used what it says"
    assert route["st"].deviations == []
    assert any(line.startswith("Follows pre-spec prespec.yaml (commit ") for line in route["ran"].lines)
    assert route["ran"].lines[-1].endswith("start with Confirmatory test.")


# --------------------------------------------------------------------------
# What went in: counted by hand from the extract


def test_the_holdout_holds_only_its_own_loans_counted_by_hand(route):
    t = route["t"]
    for side, (lo, hi) in ((t.holdout, ("2024-01-01", "2024-12-31")), (t.development, ("2022-01-01", "2023-12-31"))):
        mine = _by_hand(route["x"], lo, hi)
        assert side.n == len(mine)
        assert side.loans == [sum(1 for g, _, _ in mine if g == k) for k in range(6)]
        assert side.bad == [sum(y for g, y, _ in mine if g == k) for k in range(6)]
    rows = _rows(route["x"])
    outside = sum(1 for r in rows if not ("2022-01-01" <= r["ORIG_DATE"] <= "2024-12-31"))
    assert t.left_out[confirmatory.OUTSIDE] == outside
    chk = _check(route["b"])["Confirmatory test"]
    assert chk.startswith(f"Development 2022-01-01 to 2023-12-31: {t.development.n:,} loans, "
                          f"{t.development.n_bad:,} bad; holdout 2024-01-01 to 2024-12-31: {t.holdout.n:,} loans")
    assert f"{outside:,} made outside both ranges" in chk


def test_the_pockets_are_the_pre_specs_strata_as_the_grids_cut_them(route):
    """B3 worked out again here from the extract: each loan's FICO band from the edges Check says were used,
    its channel, its group, and nothing else."""
    t = route["t"]
    edges = [float(e) for e in _check(route["b"])["Band edges used: FICO"].split("(")[0].split(";")]
    cells = {}
    for r in _rows(route["x"]):
        if not "2024-01-01" <= r["ORIG_DATE"] <= "2024-12-31" or r["SALES"] in ("", "0") or r["BAD_FLAG"] not in ("0", "1"):
            continue
        f = r["FICO"]
        band = "blank" if f == "" else "missing" if f == "-9999" else bisect.bisect_right(edges, float(f))
        c = cells.setdefault((band, r["CHANNEL"]), [[0] * 6, [0] * 6])
        g = bisect.bisect_right(BINS, float(r["INCOME"]) / float(r["SALES"]))
        c[0][g] += 1
        c[1][g] += int(r["BAD_FLAG"])
    mine = kgroups.association([kgroups.Pocket(*c) for c in cells.values()], range(1, 7))
    assert t.holdout.association.general == pytest.approx(mine.general, rel=1e-12)
    assert t.holdout.association.trend == pytest.approx(mine.trend, rel=1e-12)


# --------------------------------------------------------------------------
# 4e: concentration on the holdout


def test_4e_capture_is_on_the_holdout_only_counted_by_hand(route):
    t = route["t"]
    mine = _by_hand(route["x"], "2024-01-01", "2024-12-31")
    n, nbad = len(mine), sum(y for _, y, _ in mine)
    gco = math.fsum(g for _, _, g in mine if g is not None)
    got = t.concentration()
    ws = load_workbook(route["b"])[confirm_tab.SHEET]
    table_rows = {ws.cell(row=r, column=2).value: r for r in range(1, ws.max_row + 1)}
    for k, name in enumerate(t.groups):
        inside = [x for x in mine if x[0] == k]
        want_flag, want_cap = len(inside) / n, sum(y for _, y, _ in inside) / nbad
        want_gco = math.fsum(g for _, _, g in inside if g is not None) / gco
        want_lift = (sum(y for _, y, _ in inside) / len(inside)) / (nbad / n)
        c = got[k]
        assert (c.flag_rate, c.capture, c.gco_capture) == pytest.approx((want_flag, want_cap, want_gco), rel=1e-12)
        assert c.lift == pytest.approx(want_lift, rel=1e-12)
        # and the tab's own table says the same (its last block: rows named by the group, share columns)
        r = max(rr for label, rr in table_rows.items() if label and str(label).startswith(name))
        assert ws.cell(row=r, column=4).value == pytest.approx(want_flag, rel=1e-12)
        assert ws.cell(row=r, column=6).value == pytest.approx(want_cap, rel=1e-12)
        assert ws.cell(row=r, column=8).value == pytest.approx(want_gco, rel=1e-12)
        assert ws.cell(row=r, column=10).value == pytest.approx(want_lift, rel=1e-12)


# --------------------------------------------------------------------------
# The tab, calculated


@pytest.fixture(scope="module")
def calculated(route, tmp_path_factory):
    return recalc(route["b"], tmp_path_factory.mktemp("calc"))[confirm_tab.SHEET]


def _texts(ws) -> list[str]:
    return [str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def _found(ws, lead: str) -> str:
    return next(s for s in _texts(ws) if s.startswith(lead) and "pockets held fixed" in s)


def test_the_tab_says_what_it_found_in_one_plain_line_per_range(route, calculated):
    t = route["t"]
    texts = _texts(calculated)
    assert not [s for s in texts if s.startswith("#") or "Err:" in s]
    for side, lead in ((t.development, "On development,"), (t.holdout, "On the holdout,")):
        line = _found(calculated, lead)
        f = side.fit
        for k in PLANTED:
            lo, hi = f.interval(k, stats.z_for_confidence(0.95))
            assert (f"{t.groups[k]} goes bad {f.odds[k]:.2f} times as often as 0.25 - 0.49 ({lo:.2f}x to "
                    f"{hi:.2f}x)") in line, line
        for k in range(6):
            if k != t.ref and f.p[k] >= 0.05:
                assert t.groups[k] + " goes" not in line
    first = next(s for s in texts if s.startswith("On the holdout, the bad rate differs across the groups"))
    assert first.endswith("rises steadily as income_to_sales rises.")          # general and trend both at 95%
    assert "Choices made here that no ruling settles yet" in texts
    for term in ("Mantel-Haenszel", "conditional logistic regression", "Odds are bad loans divided by good ones"):
        assert any(term in s for s in texts), term


@needs_git
def test_the_readings_follow_the_confidence_level_on_control(route, tmp_path):
    """OC-40: at 99% sure, a holdout p-value between 1% and 5% stops reading significant and the ranges widen,
    with no Run."""
    t = route["t"]
    mid = [k for k in range(6) if k != t.ref and t.holdout.fit.p[k] is not None and 0.01 < t.holdout.fit.p[k] < 0.05]
    assert mid, "the book has no holdout group between 1% and 5% to show the change"
    assert 0.01 < t.holdout.association.p_trend < 0.05
    copy = tmp_path / "at-99.xlsx"
    wb = load_workbook(route["b"])
    ws = wb[control.SHEET]
    ws.cell(row=control.row_of(ws, "confidence"), column=control.CHOOSE_COL).value = "99% sure"
    wb.save(copy)
    calc = recalc(copy, tmp_path / "calc")[confirm_tab.SHEET]
    line = _found(calc, "On the holdout,")
    for k in mid:
        assert t.groups[k] + " goes" not in line
    f = t.holdout.fit
    lo, hi = f.interval(5, stats.z_for_confidence(0.99))
    assert f"2.00 - {t.groups[5].split(' - ')[1]} goes bad {f.odds[5]:.2f} times as often as 0.25 - 0.49 " \
           f"({lo:.2f}x to {hi:.2f}x)" in line
    texts = _texts(calc)
    assert any(s.startswith("On the holdout, the bad rate differs across the groups, but not in one direction")
               for s in texts)
    assert "Range (99% sure)" in texts
    # the tables follow too: the trend's reading on the holdout, and the range beside each odds ratio
    rows = {}
    for r in range(1, calc.max_row + 1):
        rows.setdefault(calc.cell(row=r, column=2).value, r)
    hold = confirm_tab.FIRST + 7
    assert calc.cell(row=rows["A steady climb or fall (trend)"], column=hold + 3).value == "not significant"
    assert calc.cell(row=rows["Any difference across the groups (general)"], column=hold + 3).value == "significant"
    assert calc.cell(row=rows[t.groups[5]], column=hold + 4).value == f"{lo:.2f}x to {hi:.2f}x"


# --------------------------------------------------------------------------
# The same, on small books built in memory


def _dated(rows_spec):
    """rows_spec: (date, chan, ratio, bad) per loan."""
    return [{"ID": f"L{i}", "SCORE": 600, "CHAN": ch, "BAL": 100, "BAD": bad, "GCO": 50 * bad, "RANR": 3, "R": v,
             "ORIG": d} for i, (d, ch, v, bad) in enumerate(rows_spec)]


SPEC = {"prespec": 1, "written": "2026-09-26", "column": "R", "bins": [1.0, 2.0], "reference": 1,
        "strata": ["CHAN"], "confidence": 0.95, "holdout": {"from": "2024-01-01", "to": "2024-12-31"},
        "development": {"from": "2022-01-01", "to": "2023-12-31"}}


def _run(rows, **spec):
    ps = prespec.parse({**SPEC, **spec})
    res = engine.run(cube(origination_date="ORIG"), table(_dated(rows)))
    ps = prespec.named(ps, *confirmatory.column_range(res, ps.column))
    return res, ps, confirmatory.run_test(res, ps)


def _mixed_book():
    out = []
    for y, n in (("2021-06-01", 7), ("2022-03-01", 40), ("2023-11-30", 40), ("2024-01-01", 30), ("2024-12-31", 30),
                 ("2025-01-01", 9)):
        for i in range(n):
            out.append((y, "AB"[i % 2], (0.5, 1.5, 2.5)[i % 3], 1 if i % 5 == 0 or (i % 3 == 2 and i % 4 == 0) else 0))
    return out


def test_the_holdout_takes_no_development_loan_and_the_range_held_to_is_the_pre_specs():
    rows = _mixed_book()
    res, ps, t = _run(rows)
    assert t.holdout.n == 60 and t.development.n == 80
    assert t.left_out == {confirmatory.OUTSIDE: 16}
    assert t.holdout.range == ps.holdout and t.development.range == ps.development
    used = confirmatory.in_use(res, ps, t)
    assert used["holdout"] is ps.holdout
    # the extract runs from 2021 to 2025, and the test held itself to the holdout: no deviation (final check F5)
    assert prespec.deviations(ps, used, where="in this run") == []


def test_the_reference_group_is_the_pre_specs():
    rows = _mixed_book()
    for ref in (0, 1, 2):
        res, ps, t = _run(rows, reference=ref)
        assert t.ref == ref and confirmatory.in_use(res, ps, t)["reference"] == t.groups[ref]
        pockets = {}
        for d, ch, v, bad in rows:
            if "2024-01-01" <= d <= "2024-12-31":
                c = pockets.setdefault(ch, [[0] * 3, [0] * 3])
                g = bisect.bisect_right([1.0, 2.0], v)
                c[0][g] += 1
                c[1][g] += bad
        want = kgroups.conditional_fit([kgroups.Pocket(*c) for c in pockets.values()], ref)
        assert t.holdout.fit.odds[ref] == 1.0
        assert t.holdout.fit.beta == pytest.approx(want.beta, abs=1e-12)


def test_the_run_fits_conditionally_so_thin_pockets_are_not_biased():
    """Matched pairs through the whole route: every pocket (CHAN, the pre-spec's stratum) is one loan in the
    group and one in the reference, one of them bad. The conditional odds ratio is 60 / 30 = 2.0; a constant per
    pocket would give 4.0 (B5, "Breaks when")."""
    rows = []
    for i in range(90):
        worse = i < 60
        rows.append(("2024-05-01", f"P{i:03d}", 2.5, 1 if worse else 0))
        rows.append(("2024-05-01", f"P{i:03d}", 1.5, 0 if worse else 1))
    rows.append(("2023-05-01", "P000", 0.5, 1))           # development needs a loan; its own business
    res, ps, t = _run(rows, reference=1)
    assert t.holdout.fit.method == kgroups.CONDITIONAL and t.method == kgroups.CONDITIONAL
    assert t.holdout.fit.odds[2] == pytest.approx(2.0, rel=1e-9)
    assert len(t.holdout.pockets) == 90


def test_a_pocket_below_every_floor_still_counts_in_the_pooled_test():
    """The goal's rule: fewest loans (here 1,000) governs per-pocket readings only."""
    rows = _mixed_book() + [("2024-06-01", "Z", 0.5, 1), ("2024-06-01", "Z", 1.5, 0), ("2024-06-01", "Z", 2.5, 0)]
    ps = prespec.parse(SPEC)
    bench = {**cube().raw["benchmark"], "min_units": 1000, "min_events": 50}
    res = engine.run(cube(origination_date="ORIG", benchmark=bench), table(_dated(rows)))
    t = confirmatory.run_test(res, prespec.named(ps, *confirmatory.column_range(res, "R")))
    assert len(t.holdout.pockets) == 3 and t.holdout.association.pockets == 3


@needs_git
def test_a_stratum_the_run_does_not_cut_by_is_refused_by_the_pre_spec_cell(tmp_path, monkeypatch):
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    spec = tmp_path / "spec.yaml"
    spec.write_text(EXAMPLE.read_text(encoding="utf-8").replace("strata: [FICO, CHANNEL]",
                                                                "strata: [FICO, ASSET_CLASS]"), encoding="utf-8")
    x, b = _set_up(tmp_path / "book", 1500, spec)
    ran = book.run(b)
    assert not ran.ok
    cell = f"Control!C{control.read_prespec(load_workbook(b)[control.SHEET])[1].split('!C')[1]}"
    want = (f'{cell}: the pre-spec cuts the pockets by ASSET_CLASS, and Columns doesn\'t cut by it. Set its "Cut by '
            f'it?" to Yes (it becomes a band or a segment). Or fix the pre-spec.')
    assert any(want in line for line in ran.lines), ran.lines
    _columns(b, "ASSET_CLASS", C_CUT="Yes")
    assert book.run(b).ok


@needs_git
def test_a_holdout_with_no_loans_in_the_extract_says_so_and_tests_nothing(tmp_path, monkeypatch):
    """A pre-spec whose holdout is still to come: development is tested, and the holdout's lines say there is
    nothing to test rather than "no difference shows"."""
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    spec = tmp_path / "spec.yaml"
    spec.write_text(EXAMPLE.read_text(encoding="utf-8").replace("holdout: {from: 2024-01-01, to: 2024-12-31}",
                                                                "holdout: {from: 2030-01-01, to: 2030-12-31}"),
                    encoding="utf-8")
    x, b = _set_up(tmp_path / "book", 1500, spec)
    ran = book.run(b)
    assert ran.ok, ran.lines
    ws = recalc(b, tmp_path / "calc")[confirm_tab.SHEET]
    said = _texts(ws)
    lines = [s for s in said if s.startswith("On the holdout, nothing to test: no loan in this test was made "
                                             "2030-01-01 to 2030-12-31.")]
    assert len(lines) == 3                                     # the tests, the odds ratios, the concentration
    assert any(s.startswith("On development,") for s in said)
    assert not [s for s in said if s.startswith("#") or "Err:" in s]
    chk = _check(b)
    assert chk["Holdout"] == "None of this extract's loans were made in the holdout (2030-01-01 to 2030-12-31)."
    assert chk["Differs from the pre-spec"] == "nowhere: this run used what it says"


def test_without_readable_dates_the_test_says_why_and_the_run_goes_on():
    rows = _mixed_book()
    ps = prespec.parse(SPEC)
    res = engine.run(cube(), table(_dated(rows)))                # no column marked Origination date
    t = confirmatory.run_test(res, ps)
    assert t.problem.startswith("the loans can't be split into development and holdout: no column is marked")
    assert confirmatory.in_use(res, ps, t)["reference"] != ps.reference     # it falls back to what the tabs used
