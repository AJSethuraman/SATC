"""The control center: generated from settings.yaml, read back from the
cells a person edits, and refusing rather than defaulting."""

import pytest
from openpyxl import load_workbook

from origination_cube import control


@pytest.fixture
def book(tmp_path):
    return control.build_control_book(tmp_path / "control.xlsx")


def _row(ws, key):
    for r in ws.iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == key:
            return r[0].row
    raise KeyError(key)


JUDGMENT = {"min_age_months", "min_loans", "min_events", "materiality", "compare_to", "worse_at", "better_at",
            "confidence", "revenue_line"}


def test_judgment_settings_recommend_nothing_and_method_settings_recommend_one():
    """The firm: what is material, what is enough loans, what counts as worse
    is the professional's judgment, never the tool's."""
    settings = control.load_settings()
    assert {s.key for s in settings if s.judgment} == JUDGMENT
    for s in settings:
        assert sum(o.recommended for o in s.options) == (0 if s.judgment else 1), s.key


def test_every_setting_has_options_and_short_explanations():
    for s in control.load_settings():
        assert len(s.options) >= 2, s.key
        assert len({o.label for o in s.options}) == len(s.options), s.key
        assert s.takes_effect in ("live", "re-run"), s.key
        for o in s.options:
            # the firm's tell: a sentence past ~25 words was written to be complete, not read
            assert len(o.explains.split()) <= 30, (s.key, o.label)


def _answer_judgment(book, choices=None):
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    for s in control.load_settings():
        if s.judgment:
            pick = (choices or {}).get(s.key, s.options[0].shown)
            ws.cell(row=_row(ws, s.key), column=control.CHOOSE_COL).value = pick
    wb.save(book)


def test_a_fresh_tab_waits_for_every_judgment_and_names_each(book):
    with pytest.raises(control.ControlError) as exc:
        control.read_control(book)
    assert len(exc.value.problems) == len(JUDGMENT)
    assert all("needs an answer" in p for p in exc.value.problems)
    assert not any("judgment" in p.lower() for p in exc.value.problems)


def test_unanswered_cells_are_shaded_by_a_rule_not_labelled(book):
    """The firm: no "your judgment" labels; shade what still needs entering."""
    ws = load_workbook(book)[control.SHEET]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "judgment" not in text.lower() and "whose call" not in text.lower()
    rules = [r for rng in ws.conditional_formatting for r in rng.rules]
    assert len(rules) == len(control.load_settings())
    r = _row(ws, "materiality")
    ranges = [str(rng.sqref) for rng in ws.conditional_formatting]
    assert f"C{r}:D{r}" in ranges
    assert "re-run" not in text and "live" not in text.split()   # results are worked out on Run, not live


def test_method_settings_open_on_their_recommendation(book):
    _answer_judgment(book)
    got = control.read_control(book)
    for s in control.load_settings():
        if not s.judgment:
            assert got[s.key] == s.recommended().value


def test_your_own_value_wins(book):
    _answer_judgment(book)
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    ws.cell(row=_row(ws, "worse_at"), column=control.OWN_COL).value = 1.4
    wb.save(book)
    assert control.read_control(book)["worse_at"] == 1.4


def test_picking_another_option(book):
    _answer_judgment(book)
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    ws.cell(row=_row(ws, "confidence"), column=control.CHOOSE_COL).value = "99%"
    wb.save(book)
    assert control.read_control(book)["confidence"] == 0.99


def test_nothing_chosen_is_refused_by_name(book):
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    ws.cell(row=_row(ws, "band_count"), column=control.CHOOSE_COL).value = None
    wb.save(book)
    _answer_judgment(book)
    with pytest.raises(control.ControlError, match="How many bands"):
        control.read_control(book)


def test_text_where_a_number_is_needed_is_refused(book):
    _answer_judgment(book)
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    ws.cell(row=_row(ws, "materiality"), column=control.OWN_COL).value = "lots"
    wb.save(book)
    with pytest.raises(control.ControlError, match="dollar amount"):
        control.read_control(book)


def test_a_setting_missing_from_the_tab_is_refused(book):
    _answer_judgment(book)
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    ws.cell(row=_row(ws, "many_tests"), column=control.KEY_COL).value = None
    wb.save(book)
    with pytest.raises(control.ControlError, match="missing the setting \"Allowing for testing many pockets"):
        control.read_control(book)


def test_in_use_and_meaning_are_formulas_over_the_options_tab(book):
    ws = load_workbook(book)[control.SHEET]
    r = _row(ws, "confidence")
    assert str(ws.cell(row=r, column=5).value).startswith("=IF(")
    assert "_options" in ws.cell(row=r, column=6).value
    assert ws.data_validations.dataValidation, "every Choose cell carries a dropdown"


@pytest.mark.parametrize("key,bad,why", [("confidence", 95, "from 0.5 to 0.999"), ("worse_at", 0.9, "from 1.01 to 100"),
                                         ("min_loans", 30.5, "from 2 to 100000")])
def test_an_out_of_range_value_is_refused_naming_the_cell(book, key, bad, why):
    """Walkthrough defect 3: 95 for confidence and 0.9 for 'worse' got through."""
    _answer_judgment(book)
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    r = _row(ws, key)
    ws.cell(row=r, column=control.OWN_COL).value = bad
    wb.save(book)
    with pytest.raises(control.ControlError, match=f"Control!D{r}.*{why}"):
        control.read_control(book)


def test_the_tab_checks_typed_values_itself(book):
    ws = load_workbook(book)[control.SHEET]
    r = _row(ws, "confidence")
    dvs = [dv for dv in ws.data_validations.dataValidation if f"D{r}" in str(dv.sqref)]
    assert dvs and dvs[0].type == "decimal" and dvs[0].showErrorMessage
    assert "IFERROR" in ws.cell(row=r, column=5).value and "IFERROR" in ws.cell(row=r, column=6).value
    assert ws.column_dimensions["G"].hidden


def test_read_back_is_in_words(book):
    _answer_judgment(book)
    words = dict(control.describe(control.read_control(book)))
    assert words["Allowing for testing many pockets at once"].startswith("Hold down the share")
    assert "bh" not in words.values()


def test_the_lookups_read_the_key_column_wherever_it_is(book):
    """Found by rendering, 25 Sep 2026: after a column was removed, the key
    column moved from H to G and every lookup still read H, so every "In use"
    cell said "not an option". The tests never evaluate formulas, so this
    checks the reference instead."""
    from openpyxl.utils import get_column_letter
    ws = load_workbook(book)[control.SHEET]
    r = _row(ws, "confidence")
    key = f"${get_column_letter(control.KEY_COL)}{r}"
    assert key in ws.cell(row=r, column=5).value and key in ws.cell(row=r, column=6).value


@pytest.mark.parametrize("stored,want", [("95% sure", 0.95), (0.95, 0.95), ("95%", 0.95), (99, None)])
def test_a_pick_is_read_however_excel_stored_it(book, stored, want):
    """The second walk, defect 3: Excel can store a picked "95%" as 0.95."""
    _answer_judgment(book)
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    ws.cell(row=_row(ws, "confidence"), column=control.CHOOSE_COL).value = stored
    wb.save(book)
    if want is None:
        with pytest.raises(control.ControlError, match="is not an option"):
            control.read_control(book)
    else:
        assert control.read_control(book)["confidence"] == want


def test_no_option_label_looks_like_a_number():
    import re
    for s in control.load_settings():
        for o in s.options:
            assert not re.fullmatch(r"[\d.,]+%?", o.label), (s.key, o.label)


def test_a_row_without_its_own_value_cell_doesnt_point_at_one(book):
    """The second walk, defect 11: "enter your own in column D" on a grey n/a row."""
    with pytest.raises(control.ControlError) as exc:
        control.read_control(book)
    by_row = {p.split(":")[0]: p for p in exc.value.problems}
    ws = load_workbook(book)[control.SHEET]
    compare = by_row[f"Control!C{_row(ws, 'compare_to')}"]
    assert "Pick one from the list." in compare and "column D" not in compare
    assert "column D" in by_row[f"Control!C{_row(ws, 'min_loans')}"]
    r = _row(ws, "compare_to")
    assert "column" not in ws.cell(row=r, column=6).value.split('"That isn')[1]


def test_the_range_is_said_once_and_95_gets_a_hint(book):
    """The second walk, defect 15: "between 0.5 and 0.999, from 0.5 to 0.999; got 95"."""
    _answer_judgment(book)
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    ws.cell(row=_row(ws, "confidence"), column=control.OWN_COL).value = 95
    wb.save(book)
    with pytest.raises(control.ControlError) as exc:
        control.read_control(book)
    said = next(p for p in exc.value.problems if "95" in p)
    assert said.count("0.999") == 1 and "For 95%, type 0.95." in said
