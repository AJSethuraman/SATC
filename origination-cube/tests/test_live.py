"""The judging settings are live in the workbook (OC-40; the firm, 26 Sep 2026: "this is the stuff i want to be
able to adjust in book on the fly ... i know it cannot reband and such").

The loss line, the profit line, materiality, the confidence level and what a pocket is judged against are
Excel formulas over numbers the Run already worked out. These tests calculate the workbook with LibreOffice
(tests/recalc.py) and hold every live reading to the engine:

  - at the settings the Run used, pocket by pocket: the flag, profit's literal wording, the dollars that decide
    and whether the pocket is material, on the hidden _pockets sheet and on the tabs a person reads;
  - after changing each live setting on Control, against the engine run again in Python with the new setting;
  - a setting that takes effect on the next Run changes nothing on the result tabs, and Control says so;
  - a p-value exactly at the bar reads not significant (canon S36's own case: 0.05 at 95%).

openpyxl never calculates a formula, so without LibreOffice these tests skip, saying why."""

from __future__ import annotations

import dataclasses
import re
import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from conftest import cube, row, table
from origination_cube import book, config as cfgmod, control, engine, live, stats, synth
from origination_cube.ingest import read_table
from recalc import SOFFICE, recalc, recalc_file, values_of
from test_book import _answer

pytestmark = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice (soffice) isn't installed, so the "
                                                         "workbook's formulas can't be calculated here")

GCO_WORDS = {engine.WORSE: "losing more", engine.UNSURE_WORSE: "losing more (not significant)",
             engine.BETTER: "losing less", engine.UNSURE_BETTER: "losing less (not significant)",
             engine.IN_LINE: "about the same", engine.FEW: engine.FEW}


# --------------------------------------------------------------------------
# The helper itself


def test_the_recalculation_really_recalculates(tmp_path):
    """A file LibreOffice has already calculated carries its values. Change an input with openpyxl, calculate
    again, and the answer must be the new one, not the value carried in the file."""
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = 2, "=A1*10"
    wb.save(tmp_path / "one.xlsx")
    first = recalc_file(tmp_path / "one.xlsx", tmp_path / "a")
    assert load_workbook(first, data_only=True).active["B1"].value == 20
    again = load_workbook(first)                          # the calculated copy, its value of 20 inside it
    assert again.active["B1"].value == "=A1*10"
    again.active["A1"] = 7
    again.save(tmp_path / "two.xlsx")
    assert load_workbook(recalc_file(tmp_path / "two.xlsx", tmp_path / "b"), data_only=True).active["B1"].value == 70


# --------------------------------------------------------------------------
# The workbook route, calculated


@pytest.fixture(scope="module")
def ran(tmp_path_factory):
    """A whole Run on the synthetic book, split by revolving debt so the Split and Three-way tabs are there
    too, and its calculated values."""
    import os
    from origination_cube import perm
    d = tmp_path_factory.mktemp("live")
    old_mem, old_shuffles = os.environ.get("CUBE_MEMORY"), perm.SHUFFLES
    os.environ["CUBE_MEMORY"] = str(d / "memory.yaml")
    perm.SHUFFLES = 2_000
    try:
        extract = synth.write_extract(d, n=4000)
        b = book.set_up(extract).book
        _answer(b)
        wb = load_workbook(b)
        ws = wb["Columns"]
        for r in ws.iter_rows(min_row=book.COL_FIRST):
            if r[book.C_NAME - 1].value == "REV_DEBT":
                r[book.C_SPLIT - 1].value = "Yes"
        wb.save(b)
        assert book.run(b).ok
        raw, problems, _ = book.read_book(b)
        assert not problems
        cfg = cfgmod.parse(raw)
        table_ = read_table(extract)
        yield {"book": b, "dir": d, "cfg": cfg, "table": table_, "values": recalc(b, d / "run"),
               "res": engine.run(cfg, table_)}
    finally:
        perm.SHUFFLES = old_shuffles
        if old_mem is None:
            os.environ.pop("CUBE_MEMORY", None)
        else:
            os.environ["CUBE_MEMORY"] = old_mem


def _names(res):
    return book._names(res)


def _oracle(res):
    """What the engine says, by (kind, grid name, band, segment, rate): the flag as the tabs print it, the flag
    word, the dollars that decide, material in the tab's words."""
    b = res.config.benchmark
    line = engine.profit_line(b, res.materiality_line.get("gco_rate"))
    names = _names(res)
    out = {}
    for kind, grids in (("grids", res.grids), ("three-way", res.three_way)):
        for g in grids:
            gname = f"{names[g.band]} x {names[g.dimension]}"
            for (bl, dl), c in g.inner():
                for m in res.measures:
                    if not m.is_rate:
                        continue
                    s = c.rates[m.name]
                    out[(kind, gname, bl, dl, m.name)] = {
                        "said": (engine.said(s, line) if m.in_points else s.flag) or "", "flag": s.flag or "",
                        "dollars": s.dollars, "material": {True: "yes", False: "below the line", None: ""}[s.material],
                        "s": s, "m": m}
    return out


def _blank(v):
    return "" if v is None else v


def _pockets(values):
    ws = values[live.POCKETS]
    out = {}
    for r in ws.iter_rows(min_row=live.P_FIRST):
        if r[live.P_KIND - 1].value is None:
            continue
        key = tuple(r[c - 1].value for c in (live.P_KIND, live.P_GRID, live.P_BAND, live.P_SEG, live.P_MEASURE))
        out[key] = {"said": _blank(r[live.P_SAID - 1].value), "flag": _blank(r[live.P_FLAG - 1].value),
                    "dollars": r[live.P_DOLLARS - 1].value, "material": _blank(r[live.P_MATERIAL - 1].value)}
    return out


def _agree(values, res) -> list[str]:
    """Every difference between the calculated workbook and the engine, pocket by pocket: on _pockets, on
    Where it bleeds, Three-way, Losses vs revenue, Check's counts and the Materiality ladder."""
    want, got = _oracle(res), _pockets(values)
    bad = []
    assert set(want) == set(got), (len(want), len(got))
    for k, w in want.items():
        g = got[k]
        for f in ("said", "flag", "material"):
            if g[f] != w[f]:
                bad.append(f"_pockets {k} {f}: {g[f]!r}, the engine {w[f]!r}")
        if (w["dollars"] is None) != (g["dollars"] in (None, "")) or (
                w["dollars"] is not None and g["dollars"] != pytest.approx(w["dollars"], rel=1e-9, abs=1e-6)):
            bad.append(f"_pockets {k} dollars: {g['dollars']!r}, the engine {w['dollars']!r}")
    names = _names(res)
    titles = {m.title: m.name for m in res.measures}
    peers = res.config.benchmark.compare_to == "peers"
    for tab, kind in (("Where it bleeds", "grids"), ("Three-way", "three-way")):
        ws = values[tab]
        seen = 0
        for r in ws.iter_rows(min_row=5):
            title = r[1].value
            if title not in titles:
                continue
            key = (kind, f"{r[2].value} x {r[4].value}", r[3].value, r[5].value, titles[title])
            w = want[key]
            seen += 1
            if _blank(r[17].value) != w["said"]:
                bad.append(f"{tab} row {r[0].row} flag: {r[17].value!r}, the engine {w['said']!r}")
            if _blank(r[12].value) != w["material"]:
                bad.append(f"{tab} row {r[0].row} material: {r[12].value!r}, the engine {w['material']!r}")
            s = w["s"]
            first = s.excess_band if peers else s.excess
            if (first is None) != (r[9].value in (None, "")) or (
                    first is not None and r[9].value != pytest.approx(first, rel=1e-9, abs=1e-6)):
                bad.append(f"{tab} row {r[0].row} first dollars: {r[9].value!r}, the engine {first!r}")
        # every pocket losing more than its share against the comparison that decides now is on the list
        losing = {k for k, w in want.items() if k[0] == kind and w["dollars"] is not None and w["dollars"] > 0}
        listed = {(kind, f"{r[2].value} x {r[4].value}", r[3].value, r[5].value, titles[r[1].value])
                  for r in ws.iter_rows(min_row=5) if r[1].value in titles}
        if losing - listed:
            bad.append(f"{tab}: {len(losing - listed)} pockets losing more than their share aren't listed: {sorted(losing - listed)[:3]}")
        if not seen:
            bad.append(f"{tab}: no rows read")
    # Losses vs revenue: each side's reading, and the two read together
    ws = values["Losses vs revenue"]
    grid = None
    n = 0
    for r in range(4, ws.max_row + 1):
        b_ = ws.cell(row=r, column=2).value
        if isinstance(b_, str) and " x " in b_ and ws.cell(row=r, column=4).value is None:
            grid = b_
            continue
        if not isinstance(ws.cell(row=r, column=4).value, int):
            continue
        n += 1
        k = lambda m: ("grids", grid, b_, ws.cell(row=r, column=3).value, m)          # noqa: E731
        gw, rw, cw = want[k("gco_rate")], want[k("ranr_rate")], want[k("contribution_rate")]
        for col_, w in ((book.LVR_G + 3, GCO_WORDS.get(gw["flag"], "")), (book.LVR_R + 3, rw["said"]),
                        (book.LVR_C + 3, cw["said"])):
            if _blank(ws.cell(row=r, column=col_).value) != w:
                bad.append(f"Losses vs revenue {grid} row {r} col {col_}: {ws.cell(row=r, column=col_).value!r}, "
                           f"the engine {w!r}")
        untested = gw["flag"] in (engine.FEW, engine.THIN)
        pair = {(engine.WORSE, engine.BETTER): "priced for it", (engine.WORSE, engine.WORSE): "net drain",
                (engine.BETTER, engine.WORSE): "safe but idle"}.get((gw["flag"], rw["flag"]), "")
        if _blank(ws.cell(row=r, column=book.LVR_T).value) != ("" if untested else pair):
            bad.append(f"Losses vs revenue {grid} row {r} Together: {ws.cell(row=r, column=book.LVR_T).value!r}")
    if not n:
        bad.append("Losses vs revenue: no rows read")
    # Check counts the pockets that read worse now
    check = {r[1].value: r[2].value for r in values["Check"].iter_rows(min_row=4)}
    for m in res.measures:
        if not m.is_rate:
            continue
        k = sum(1 for key, w in want.items() if key[0] == "grids" and key[4] == m.name and w["flag"] == engine.WORSE)
        j = sum(1 for key, w in want.items() if key[0] == "grids" and key[4] == m.name and w["flag"] == engine.WORSE
                and w["material"] == "yes")
        said = check[f"Worse now: {m.title}"]
        if not said.startswith(f"{k} of ") or f"; {j} of them are material." not in said:
            bad.append(f"Check, {m.title}: {said!r}, the engine {k} worse and {j} material")
    # the Materiality ladder, from the dollars that decide now
    ws = values["Materiality"]
    ladders = {}
    for g in res.grids:
        for m in res.measures:
            if m.is_rate:
                ladders[f"{names[g.band]} x {names[g.dimension]}: {m.title}"] = engine.materiality(g, m, res.total)
    r = 4
    while r <= ws.max_row:
        head = ws.cell(row=r, column=2).value
        if head in ladders:
            for i, want_row in enumerate(ladders[head]):
                got_k, got_s = ws.cell(row=r + 2 + i, column=4).value, ws.cell(row=r + 2 + i, column=5).value
                if got_k != want_row.pockets or got_s != pytest.approx(want_row.captured, abs=1e-9):
                    bad.append(f"Materiality {head} level {i}: {got_k}, {got_s}; the engine {want_row.pockets}, "
                               f"{want_row.captured}")
            r += 2 + len(ladders[head])
        r += 1
    return bad


def test_every_live_reading_equals_the_engine_at_the_runs_settings(ran):
    res = ran["res"]
    # the oracle is the Run: the same numbers the workbook holds
    got = _pockets(ran["values"])
    assert len(got) == sum(1 for g in (*res.grids, *res.three_way) for _ in g.inner()) * sum(
        1 for m in res.measures if m.is_rate)
    assert _agree(ran["values"], res) == []


def _set(path: Path, out: Path, changes: dict) -> Path:
    """A copy of the workbook with Control changed: {key: (pick, own)}; None leaves a cell as it is."""
    wb = load_workbook(path)
    ws = wb[control.SHEET]
    for key, (pick, own) in changes.items():
        r = control.row_of(ws, key)
        if pick is not None:
            ws.cell(row=r, column=control.CHOOSE_COL).value = pick
        if own is not None:
            ws.cell(row=r, column=control.OWN_COL).value = own
    wb.save(out)
    return out


#: Each live setting changed on Control, and the same change to the engine's settings
CHANGES = {
    "worse_at": ({"worse_at": ("2 times", None)}, {"worse_at": 2.0}),
    "better_at": ({"better_at": ("0.5 times", None)}, {"better_at": 0.5}),
    "worse_at_own_value": ({"worse_at": (None, 1.4)}, {"worse_at": 1.4}),
    "profit_points": ({"revenue_line": ("0.5 points either way", None)}, {"revenue_line": 0.5}),
    "profit_materiality": ({"revenue_line": ("A gap as big as the materiality line", None)},
                           {"revenue_line": "materiality"}),
    "confidence": ({"confidence": ("99% sure", None)}, {"confidence": 0.99}),
    "judged_against": ({"compare_to": ("The rest of the book", None)}, {"compare_to": "topline"}),
    "materiality_dollars": ({"materiality": (None, 20_000)}, {"materiality": ("dollars", 20_000.0)}),
}


@pytest.mark.parametrize("change", list(CHANGES))
def test_live_readings_follow_a_change_on_control(ran, change, tmp_path):
    """Change one live setting on Control and calculate: every reading equals the engine run again with that
    setting, pocket by pocket."""
    on_control, on_engine = CHANGES[change]
    b = _set(ran["book"], tmp_path / "changed.xlsx", on_control)
    values = recalc(b, tmp_path / "rc")
    cfg = ran["cfg"]
    bench = dataclasses.replace(cfg.benchmark, **on_engine)
    res = engine.run(dataclasses.replace(cfg, benchmark=bench), ran["table"])
    # the change moved something, or this proves nothing
    before = {k: (w["said"], w["material"]) for k, w in _oracle(ran["res"]).items()}
    after = {k: (w["said"], w["material"]) for k, w in _oracle(res).items()}
    assert before != after, change
    assert _agree(values, res) == []
    # the tab says what it's using now, and Check shows the change beside the last Run's record
    assert values["Where it bleeds"]["B3"].value.startswith("Lines in use now, from Control: ")
    check = values["Check"]
    changed = [r[1].value for r in check.iter_rows(min_row=4) if r[3].value and r[2].value != r[3].value
               and r[1].value.startswith(("The loss line", "The profit line", "Confidence", "Materiality",
                                          "Judged against"))]
    assert changed, change


def test_at_the_runs_settings_check_shows_no_line_changed(ran):
    check = ran["values"]["Check"]
    rows = [r for r in check.iter_rows(min_row=4) if r[3].value and r[1].value.startswith(
        ("The loss line", "The profit line", "Confidence", "Materiality", "Judged against"))]
    assert len(rows) == 6
    assert all(r[2].value == r[3].value for r in rows), [(r[2].value, r[3].value) for r in rows]


#: settings that take effect on the next Run: changed on Control, nothing on a result tab moves
RERUN = {"min_events": ("20 losses", None), "min_loans": ("300 loans", None), "many_tests": ("No allowance", None),
         "power": ("90% of the time", None), "band_count": ("3 bands", None),
         "band_cut": ("The same, snapped to round numbers", None)}
RESULT = ("Where it bleeds", "Losses vs revenue", "Grids", "Split", "Three-way", "Materiality", "Check")


def test_a_setting_for_the_next_run_changes_nothing_on_the_result_tabs(ran, tmp_path):
    b = _set(ran["book"], tmp_path / "rerun.xlsx", RERUN)
    values = recalc(b, tmp_path / "rc")
    for tab in RESULT:
        a, z = ran["values"][tab], values[tab]
        diff = [(c.coordinate, c.value, z[c.coordinate].value) for r in a.iter_rows() for c in r
                if c.value != z[c.coordinate].value]
        assert diff == [], (tab, diff[:5])
    ws = load_workbook(b)[control.SHEET]
    for s in control.load_settings():
        r = control.row_of(ws, s.key)
        said = ws.cell(row=r, column=control.WHEN_COL).value
        if s.key in live.LIVE_KEYS:
            assert said == "Now, on the result tabs", s.key
        elif s.key in control.AT_SET_UP:
            assert said == "At the next Set up", s.key
        else:
            assert said == "At the next Run", s.key
    assert ws.cell(row=4, column=control.WHEN_COL).value == "When a change shows"


def test_the_live_settings_are_the_ones_settings_yaml_marks_live():
    assert {s.key for s in control.load_settings() if s.takes_effect == "live"} == set(live.LIVE_KEYS)


def test_the_tabs_say_what_stays_as_of_the_run(ran):
    v = ran["values"]
    for tab in ("Where it bleeds", "Three-way", "Losses vs revenue"):
        assert "as of the last Run" in v[tab]["B3"].value, tab
        assert "not the order" in v[tab]["B3"].value, tab
    assert "(as of the last Run)" in load_workbook(ran["book"])["Losses vs revenue"]._charts[0].title.tx.rich.p[
        0].r[0].t
    log = load_workbook(ran["book"])["Log"]
    assert "a line changed on Control afterwards shows on the result tabs, not here" in log["A2"].value


# --------------------------------------------------------------------------
# The bar: one cell, rounded once, and a p-value at it is not significant


def _edge_book(tmp_path, p: float, gap: float, measure: str, confidence: float = 0.95):
    """The small book's workbook with one pocket's p-value and gap set by hand, calculated."""
    res = engine.run(cube(benchmark={**_bench(), "confidence": confidence}), table(_rows()))
    wb = Workbook()
    lv = live.ensure(wb, res)
    g = res.grids[0]
    bl, dl = next(k for k, _ in g.inner())
    r = lv.rows[("grids", id(g), bl, dl, measure)]
    ws = wb[live.POCKETS]
    ws.cell(row=r, column=live.P_P_BOOK).value = p
    ws.cell(row=r, column=live.P_VS_BOOK).value = gap
    ws.cell(row=r, column=live.P_FEW).value = False
    return values_of(wb, tmp_path)[live.POCKETS].cell(row=r, column=live.P_READ_BOOK).value


def _bench():
    return {"min_units": 2, "min_events": 1, "worse_at": 1.25, "better_at": 0.8, "confidence": 0.95, "power": 0.8,
            "compare_to": "topline", "many_tests": "none", "materiality": "none", "revenue_line": 0.25,
            "shuffles": 200}


def _rows():
    return [row(i, 600 + 100 * (i % 2), "AB"[i % 3 == 0], 1000, int(i % 7 == 0), 500 * int(i % 7 == 0), 30)
            for i in range(120)]


@pytest.mark.parametrize("confidence, p", [(0.95, 0.05), (0.99, 0.01), (0.9, 0.1)])
def test_a_p_value_exactly_at_the_bar_is_not_significant(tmp_path, confidence, p):
    assert _edge_book(tmp_path, p, 2.0, "gco_rate", confidence) == engine.UNSURE_WORSE
    assert engine.reading_of(2.0, 100, dataclasses.replace(cube(benchmark=_bench()).benchmark,
                                                           confidence=confidence), 0, p) == engine.UNSURE_WORSE


def test_just_under_the_bar_is_significant(tmp_path):
    assert _edge_book(tmp_path, 0.0499, 2.0, "gco_rate") == engine.WORSE
    assert _edge_book(tmp_path / "b", 0.0499, -0.01, "ranr_rate") == engine.WORSE     # 1 point short: past 0.25


def test_profit_at_the_bar_is_not_significant(tmp_path):
    assert _edge_book(tmp_path, 0.05, -0.01, "ranr_rate") == engine.UNSURE_WORSE


def test_every_significance_test_compares_with_the_one_rounded_bar(ran):
    """S36: the bar is worked out once, ROUND(1 - confidence, 12), in one named cell; no formula anywhere in
    the workbook writes 1 - confidence into a comparison. LibreOffice compares within a hair, so a bar left
    unrounded reads the same there: this holds the formula itself, which is what Excel calculates from."""
    wb = load_workbook(ran["book"])
    assert wb.defined_names[live.BAR].attr_text == f"'{live.LIVE_SHEET}'!$C${live.L_BAR}"
    assert wb[live.LIVE_SHEET].cell(row=live.L_BAR, column=3).value == f"=ROUND(1-C{live.L_CONF},12)"
    inline = re.compile(r"1\s*-\s*(confidence|C\d+|\$?[A-Z]+\$?\d+)", re.I)
    uses = 0
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            for c in r:
                v = c.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                uses += live.BAR in v
                for m in inline.finditer(v):
                    # the bar's own cell, and the z for a range (NORMSINV(1-(1-confidence)/2)), aren't comparisons
                    ok = (ws.title == live.LIVE_SHEET and c.row == live.L_BAR) or "NORMSINV(1-(1-confidence)/2)" in v
                    assert ok, (ws.title, c.coordinate, v)
                if "<" in v and "p-value" not in v:
                    for p in re.findall(r"(\$?[A-Z]+\$?\d+)<significance_bar", v):
                        assert f"ISNUMBER({p})" in v
    assert uses > 100


def test_the_live_formulas_use_no_dynamic_arrays(ran):
    """Excel 2016 and LibreOffice: no SORT, FILTER, LET, UNIQUE or XLOOKUP anywhere."""
    wb = load_workbook(ran["book"])
    bad = re.compile(r"\b(SORT|SORTBY|FILTER|LET|UNIQUE|XLOOKUP|SEQUENCE|LAMBDA)\(")
    for ws in wb.worksheets:
        for r in ws.iter_rows():
            for c in r:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert not bad.search(c.value), (ws.title, c.coordinate, c.value)


def test_nothing_on_the_result_tabs_is_an_error(ran):
    errors = ("#N/A", "#VALUE!", "#NAME?", "#REF!", "#DIV/0!", "#NUM!", "Err:")
    for tab in RESULT + (live.POCKETS,):
        for r in ran["values"][tab].iter_rows():
            for c in r:
                assert not (isinstance(c.value, str) and c.value.startswith(errors)), (tab, c.coordinate, c.value)
