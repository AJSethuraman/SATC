"""The launcher's five steps, without the window: every state the redesign draws
(docs/redesign-2026-09-26/README.md, L1 to L5) is a Flow a test can drive, and
every outcome is words."""

import re
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from origination_cube import book, choices, control, house, launcher, synth
from test_book import _answer


@pytest.fixture(autouse=True)
def _nothing_opens(monkeypatch):
    """Open at... opens Excel; here it only notes what it would have opened."""
    opened = []
    monkeypatch.setattr(launcher, "open_file", lambda p: opened.append(Path(p)))
    return opened


def _flow(x=None, **kw) -> launcher.Flow:
    f = launcher.Flow(gate=launcher.AddOns(), **kw)
    if x is not None:
        f.pick(str(x))
    return f


def _read(tmp_path, n=1500, **kw) -> launcher.Flow:
    f = _flow(synth.write_extract(tmp_path, n=n), **kw)
    f.set_up()
    assert f.screen() == "L2", f.message
    return f


def _control_rows(b) -> dict:
    ws = load_workbook(b)[control.SHEET]
    return {r[control.KEY_COL - 1].value: r[control.CHOOSE_COL - 1].value for r in ws.iter_rows(min_row=control.FIRST_ROW)
            if r[control.KEY_COL - 1].value}


# ---- the old buttons, still in words


def test_the_workbook_sits_beside_the_extract_named_pocketbook(tmp_path):
    """The firm, 26 Sep 2026: the product is PocketBook, and so is its workbook."""
    assert launcher.book_for(tmp_path / "sep loans.csv") == tmp_path / "sep loans - PocketBook.xlsx"
    assert book.book_for(tmp_path / "sep loans.csv") == launcher.book_for(tmp_path / "sep loans.csv")


def test_a_workbook_under_either_name_is_refused_as_the_extract(tmp_path):
    for name in ("sep - PocketBook.xlsx", "sep - Origination Cube.xlsx"):
        out = book.set_up(tmp_path / name)
        assert not out.ok and out.lines[0].startswith(f"{name} is the workbook, not the loan file")
        f = _flow(tmp_path / name)
        (tmp_path / name).write_bytes(b"")
        f.set_up()
        assert f.screen() == "L1" and "is the workbook, not the loan file" in f.message[0]


def test_buttons_pressed_out_of_order_say_what_to_do(tmp_path):
    assert "Pick the extract first" in launcher.do_set_up("")[0]
    x = synth.write_extract(tmp_path, n=500)
    assert "Choose the tests and press Next first" in launcher.do_run(str(x))[0]


def test_set_up_then_run_through_the_buttons(tmp_path):
    x = synth.write_extract(tmp_path, n=2000)
    lines = launcher.do_set_up(str(x))
    assert lines[0].startswith("Set up") and launcher.book_for(x).exists()
    lines = launcher.do_run(str(x))
    assert lines[0].startswith("Couldn't run yet")


def test_an_unexpected_failure_is_a_sentence_not_a_traceback(tmp_path, monkeypatch):
    monkeypatch.setattr(launcher.Path, "home", lambda: tmp_path)

    def boom(*a, **k):
        raise ZeroDivisionError("inside")
    monkeypatch.setattr(book, "set_up", boom)
    x = synth.write_extract(tmp_path, n=100)
    lines = launcher.do_set_up(str(x))
    assert len(lines) == 1 and "send that file over" in lines[0] and "Traceback" not in lines[0]
    assert "ZeroDivisionError" in (tmp_path / ".origination-cube" / "last-error.txt").read_text()


def test_run_hands_the_picked_extract_to_the_workbook(tmp_path, monkeypatch):
    """Second walk, defect 1: the window's extract wins over the path stored at Set up."""
    x = synth.write_extract(tmp_path, n=100)
    launcher.book_for(x).write_bytes(b"")
    seen = {}

    def fake(target, extract=None):
        seen["extract"] = extract
        return book.Outcome(True, target, ["ok"])
    monkeypatch.setattr(book, "run", fake)
    assert launcher.do_run(str(x)) == ["ok"] and seen["extract"] == str(x)


def test_the_window_takes_its_colours_from_the_house_module():
    """One place for the colours, and no copy of credit-suite's style file (its conformance test refuses one)."""
    src = (Path(launcher.__file__)).read_text()
    assert not re.search(r"#[0-9A-Fa-f]{6}", src)                 # every colour comes from house.py
    assert house.tk(house.KEY_RED) == "#CC0000" and house.INK == "0A0908"
    assert not (Path(launcher.__file__).parent / "keybank_style.py").exists()


# ---- L1 · Opened


def test_l1_set_up_stays_off_until_an_extract_is_picked(tmp_path):
    f = _flow()
    assert f.screen() == "L1" and f.states()["setup"] == "disabled"
    assert [s.mark for s in f.steps()] == ["current", "todo", "todo", "todo", "todo"]
    assert f.steps()[0].sub == "Pick the loan file"
    f.pick(str(tmp_path / "not there.csv"))
    assert f.states()["setup"] == "disabled"                      # a path to nothing is not a pick
    f.pick(str(synth.write_extract(tmp_path, n=200)))
    assert f.states()["setup"] == "normal" and f.states()["run"] == "disabled"


def _graded(tmp_path, n=1500):
    """The synthetic book with TIER, a number column holding the eight values 1 to 8."""
    import csv
    x = synth.write_extract(tmp_path, n=n)
    with open(x, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    for i, r in enumerate(rows):
        r["TIER"] = str(i % 8 + 1)
    with open(x, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return x


def test_l1_the_column_limits_are_chosen_before_the_workbook_exists(tmp_path):
    """They moved from Control to the launcher: Set up needs them before there is a workbook. TIER has eight
    values: a category at 12 values or fewer, a number column at 6."""
    x = _graded(tmp_path)
    f = _flow(x)
    f.set_up()
    assert {c.name: c.kind for c in f.read.columns}["TIER"] == "cat"
    f = _flow(x, few=6, many=25)
    f.set_up()
    assert {c.name: c.kind for c in f.read.columns}["TIER"] == "num"
    assert not launcher.book_for(f.extract).exists()             # reading the extract writes nothing
    f.next()
    ws = load_workbook(launcher.book_for(f.extract))["Columns"]
    grade = next(r for r in ws.iter_rows(min_row=book.COL_FIRST) if r[book.C_NAME - 1].value == "TIER")
    assert grade[book.C_MEANS - 1].value == "Amount or number"   # Set up read it at 6, as the launcher said
    rows = _control_rows(launcher.book_for(f.extract))
    assert rows["few_values"] == "6 values" and rows["many_values"] == "25 values"
    ws = load_workbook(launcher.book_for(f.extract))[control.SHEET]
    head = next(r[0].row for r in ws.iter_rows(min_row=control.FIRST_ROW) if r[control.KEY_COL - 1].value ==
                f"{choices.KEY}|head")
    assert head < control.row_of(ws, "few_values") < control.row_of(ws, "min_loans")   # in the launcher's block


# ---- L2 · Choose tests


def test_l2_key_and_date_columns_are_greyed_with_nothing_to_tick(tmp_path):
    f = _read(tmp_path)
    rows = {r["name"]: r for r in f.rows()}
    for name in ("LOAN_NBR", "ORIG_DATE"):
        assert rows[name]["grey"] and rows[name]["a"] is rows[name]["b"] is rows[name]["c"] is None
    f.set_mode("new")
    rows = {r["name"]: r for r in f.rows()}
    assert rows["LOAN_NBR"]["grey"] and rows["LOAN_NBR"]["b"] is None
    s = f.steps()
    assert [x.mark for x in s] == ["done", "done", "current", "todo", "todo"]
    assert s[0].sub == "loans.csv · 1,500 loans" and s[1].sub == "10 columns read" and s[2].sub == "Test new variables"


def test_l2_bleed_starts_with_every_number_and_category_column_and_no_split(tmp_path):
    f = _read(tmp_path)
    assert f.heads() == ("Cut into bands", "Segment by", "Split pockets by")
    rows = {r["name"]: r for r in f.rows()}
    assert rows["FICO"]["a"]["on"] and rows["ORIG_BAL"]["a"]["on"] and rows["REV_DEBT"]["a"]["on"]
    assert rows["CHANNEL"]["b"]["on"] and rows["ASSET_CLASS"]["b"]["on"]
    assert not any(r["c"]["on"] for r in rows.values() if r["c"])
    assert rows["BAD_FLAG"].get("every") and rows["GCO_AMT"].get("every") and rows["BAD_FLAG"]["a"] is None


def test_l2_only_one_column_splits_and_it_is_not_also_cut(tmp_path):
    """The radio replaces two Yes cells on Columns, which the run had to refuse ("only one column can split")."""
    f = _read(tmp_path)
    f.click("REV_DEBT", "c")
    f.click("ORIG_BAL", "c")
    got = f.choices()
    assert got.split == "ORIG_BAL" and "ORIG_BAL" not in got.bands and "REV_DEBT" not in got.bands
    f.click("ORIG_BAL", "c")                                      # pressed again: nothing splits
    assert f.choices().split is None


def test_l2_test_it_and_hold_fixed_exclude_each_other_on_a_row(tmp_path):
    f = _read(tmp_path)
    f.set_mode("new")
    assert f.heads() == ("Outcome", "Test it", "Hold fixed")
    rows = {r["name"]: r for r in f.rows()}
    assert rows["BAD_FLAG"]["a"] == {"on": True, "radio": True}   # the outcome, picked for you only if it's the one
    assert rows["GCO_AMT"]["a"] is None
    f.click("CHANNEL", "b")
    f.click("CHANNEL", "c")
    assert f.test == [] and f.hold == ["CHANNEL"]
    f.click("CHANNEL", "b")
    assert f.test == ["CHANNEL"] and f.hold == []


def test_l2_the_summary_says_what_will_run_in_both_modes(tmp_path):
    f = _read(tmp_path)
    f.click("REV_DEBT", "c")
    ok, said = f.summary()
    assert ok and said == ("2 band columns × 2 segment columns = 4 grids, five measures each; split by REV_DEBT "
                           "adds 4 more.")
    f.set_mode("new")
    assert f.summary() == (False, "Tick at least one input to test.")
    for c in ("CHANNEL", "ORIG_BAL", "ASSET_CLASS", "REV_DEBT"):
        f.click(c, "b")
    f.click("FICO", "c")
    assert f.summary() == (True, "4 inputs (CHANNEL, ORIG_BAL, ASSET_CLASS, REV_DEBT) against BAD_FLAG, each with "
                                 "and without FICO held fixed: 8 tests, found on 70% and confirmed on 30%.")


def test_l2_next_is_off_until_there_is_a_grid_to_run(tmp_path):
    f = _read(tmp_path)
    for c in ("CHANNEL", "ASSET_CLASS"):
        f.click(c, "b")
    assert f.states()["next"] == "disabled"
    assert f.summary() == (False, "Tick at least one band column and one segment column.")
    f.click("CHANNEL", "b")
    assert f.states()["next"] == "normal"


def test_l2_next_writes_the_choices_where_control_shows_them_and_the_run_reads_them(tmp_path):
    f = _read(tmp_path)
    f.click("REV_DEBT", "c")
    f.click("ASSET_CLASS", "b")
    f.next()
    assert f.screen() == "L2b" and [s.mark for s in f.steps()] == ["done", "done", "done", "current", "todo"]
    b = launcher.book_for(f.extract)
    rows = _control_rows(b)
    assert rows["run_kind"] == "Where the book bleeds"
    assert rows[f"{choices.KEY}|bands"] == "FICO, ORIG_BAL" and rows[f"{choices.KEY}|segments"] == "CHANNEL"
    assert rows[f"{choices.KEY}|split"] == "REV_DEBT"
    _answer(b)
    raw, problems, _ = book.read_book(b)
    assert not problems
    assert [x["field"] for x in raw["bands"]] == ["FICO", "ORIG_BAL"] and [x["field"] for x in raw["dimensions"]] \
        == ["CHANNEL"] and raw["split"] == {"field": "REV_DEBT", "how": "own_median"}
    # Columns no longer asks what to cut: both old columns are hidden and empty
    ws = load_workbook(b)["Columns"]
    assert ws.column_dimensions["D"].hidden and ws.column_dimensions["H"].hidden
    assert all(ws.cell(row=r, column=c).value is None for r in range(book.COL_FIRST, ws.max_row + 1) for c in (4, 8))


def test_l2_choosing_again_keeps_the_answers_already_given(tmp_path):
    f = _read(tmp_path)
    f.next()
    b = launcher.book_for(f.extract)
    _answer(b)
    f.go(2)
    assert f.screen() == "L2"
    f.click("REV_DEBT", "c")
    f.next()
    rows = _control_rows(b)
    assert rows[f"{choices.KEY}|split"] == "REV_DEBT" and rows["min_loans"] == "30"
    assert load_workbook(b)["Columns"][book.CONFIRM_CELL].value == "Yes"
    again = _flow(f.extract)                                      # a new window: the table starts where it was left
    again.set_up()
    assert again.split == "REV_DEBT" and "ASSET_CLASS" in again.seg


def test_l2_a_saved_shortlist_fills_test_it_and_hold_fixed_from_the_file(tmp_path):
    from test_confirmatory import SPEC
    f = _read(tmp_path)
    spec = tmp_path / "prespec.yaml"
    spec.write_text(yaml.safe_dump(SPEC, sort_keys=False), encoding="utf-8")
    f.set_mode("new")
    f.pick_shortlist(str(spec))
    assert f.test == ["INCOME_TO_SALES"] and f.hold == ["FICO", "CHANNEL"]
    assert all(r["locked"] for r in f.rows())
    f.click("ORIG_BAL", "b")                                      # the file decides: the boxes don't move
    assert f.test == ["INCOME_TO_SALES"]
    ok, said = f.summary()
    assert ok and said == ("the saved shortlist prespec.yaml: INCOME_TO_SALES against BAD_FLAG, with FICO and CHANNEL "
                           "held fixed, confirmed on the loans it held back.")
    got = f.choices()
    assert got.run_kind == choices.NEW_VARIABLE and got.shortlist == str(spec) and got.split == "INCOME_TO_SALES"
    assert got.bands == ("FICO",) and got.segments == ("CHANNEL",)
    f.next()
    rows = _control_rows(launcher.book_for(f.extract))
    assert rows[control.PRESPEC_KEY] == str(spec) and rows["new_variable_step"] == "Test from a pre-spec"
    bad = tmp_path / "bad.yaml"
    bad.write_text("prespec: 2\n", encoding="utf-8")
    f.pick_shortlist(str(bad))
    ok, said = f.summary()
    assert not ok and said.startswith("bad.yaml can't be used:") and f.states()["next"] == "disabled"


def test_the_outcome_the_launcher_tests_against_must_be_the_one_columns_marks(tmp_path):
    from test_book_dates import _choose
    f = _read(tmp_path)
    f.set_mode("new")
    f.click("CHANNEL", "b")
    f.next()
    b = launcher.book_for(f.extract)
    _answer(b)
    _choose(b, run_kind=choices.NEW_VARIABLE, outcome="ORIG_BAL")
    _, problems, _ = book.read_book(b)
    cell = f"Control!C{control.row_of(load_workbook(b)[control.SHEET], 'launcher|outcome')}"
    assert any(p.startswith(f"{cell}: the launcher tests against ORIG_BAL, and Columns!C9 doesn't mark it Outcome "
                            f"(yes/no).") for p in problems), problems
    _choose(b, run_kind=choices.NEW_VARIABLE, outcome="BAD_FLAG")
    assert not any("the launcher tests against" in p for p in book.read_book(b)[1])


# ---- suggestions exist when the workbook is written


def _suggested_cell(b, key):
    ws = load_workbook(b)[control.SHEET]
    return ws.cell(row=control.row_of(ws, key), column=book.SUGGEST_COL).value


def test_the_suggested_value_is_on_control_before_the_first_run(tmp_path):
    """The firm, 26 Sep 2026: choose the cuts first "so that there are suggestions to be made". The value is
    shown beside the setting; the answer cell stays blank for the analyst (ruling OC-13)."""
    f = _read(tmp_path, n=3000)
    f.next()
    b = launcher.book_for(f.extract)
    said = _suggested_cell(b, "min_loans")
    assert re.fullmatch(r"suggested: \d+, from this extract", said), said
    for key in ("worse_at", "better_at"):
        assert re.fullmatch(r"suggested: \d\.\d\dx, from this extract", _suggested_cell(b, key))
    assert _control_rows(b)["min_loans"] is None                  # never chosen for you
    assert "fewest loans" in f.suggested_line()
    # the same number the Run works out when the suggestion is picked
    _answer(b)
    wb = load_workbook(b)
    ws = wb[control.SHEET]
    ws.cell(row=control.row_of(ws, "min_loans"), column=control.CHOOSE_COL).value = \
        "Enough for 5 expected losses (suggested)"
    wb.save(b)
    ran = book.run(b)
    assert ran.ok, ran.lines
    worked = next(x for x in ran.lines if x.startswith("Worked out from this book"))
    assert f"fewest loans {said.split()[1].rstrip(',')}" in worked


def test_run_refreshes_the_suggestion_and_records_what_it_used(tmp_path):
    f = _read(tmp_path, n=3000)
    f.next()
    b = launcher.book_for(f.extract)
    _answer(b)
    assert book.run(b).ok
    ws = load_workbook(b)[control.SHEET]
    r = control.row_of(ws, "min_loans")
    assert ws.cell(row=r, column=book.SUGGEST_COL).value.endswith("from this extract at the last Run")
    assert ws.cell(row=r, column=control.KEY_COL + 1).value == "30 loans"          # Last Run used


# ---- L3 · Run pressed before everything is answered


def test_l3_each_answer_still_needed_is_listed_by_cell_with_its_question(tmp_path):
    f = _read(tmp_path)
    f.next()
    f.run()
    assert f.screen() == "L3" and f.states()["run"] == "normal"
    tags = {n.tag: n.says for n in f.needs}
    b = launcher.book_for(f.extract)
    ws = load_workbook(b)[control.SHEET]
    assert tags[f"CONTROL C{control.row_of(ws, 'min_loans')}"].startswith("Fewest loans in a pocket for the usual test")
    assert tags["COLUMNS C3"] == "Checked every column? Set it to Yes once you have."
    s = f.steps()
    assert s[3].mark == "blocked" and s[3].sub == f"{len(f.needs)} left" and s[4].mark == "todo"


def test_l3_open_at_puts_the_workbook_on_that_cell(tmp_path, _nothing_opens):
    f = _read(tmp_path)
    f.next()
    f.run()
    need = next(n for n in f.needs if n.tag == "COLUMNS C3")
    assert f.open_at(need) == ""
    wb = load_workbook(launcher.book_for(f.extract))
    assert wb.active.title == "Columns"
    assert all(s.activeCell == "C3" for s in wb.active.sheet_view.selection)
    assert _nothing_opens == [launcher.book_for(f.extract)]


def test_l3_while_the_workbook_is_open_the_banner_shows_and_run_stays_off(tmp_path):
    f = _read(tmp_path)
    f.next()
    f.run()
    b = launcher.book_for(f.extract)
    lock = b.with_name(f"~${b.name}")                             # what Excel leaves beside a workbook it has open
    lock.write_text("")
    assert f.refresh() and f.book_open and f.states()["run"] == "disabled"
    assert book.open_at(b, "Columns", "C3") is False              # it can't be moved to the cell: said instead
    assert f.open_at(f.needs[0]).startswith(f"Go to {f.needs[0].sheet} {f.needs[0].cell}")
    lock.unlink()
    assert f.refresh() and f.states()["run"] == "normal"


def test_l3_a_problem_without_a_cell_is_still_listed():
    got = launcher.needs_of(['Columns: Where the book bleeds needs one column marked GCO dollars, and none is.',
                             'Control!C23: "Fewest loans" needs an answer. Pick one.',
                             'Odd values!E6: something', "Couldn't find the extract loans.csv."])
    assert [(n.sheet, n.cell) for n in got] == [("Columns", None), ("Control", "C23"), ("Odd values", "E6"),
                                                (None, None)]
    assert got[1].says == "Fewest loans" and got[1].tag == "CONTROL C23" and got[3].tag == ""


# ---- L4 · An add-on is missing


def test_l4_with_an_add_on_missing_the_rail_fades_and_set_up_and_run_stay_off(tmp_path):
    f = _read(tmp_path)
    f.next()
    f.gate.missing = ["numpy"]
    assert f.screen() == "L4"
    assert [s.mark for s in f.steps()] == ["todo"] * 5
    st = f.states()
    assert st["setup"] == st["run"] == st["next"] == "disabled" and st["open"] == "normal"
    assert st["install"] == "normal"


# ---- L5 · Run finished


def test_l5_a_finished_run_shows_the_three_tiles_and_the_open_question(tmp_path):
    f = _read(tmp_path, n=3000)
    f.next()
    _answer(launcher.book_for(f.extract), odd=False)
    f.run()
    assert f.screen() == "L5", [n.says for n in f.needs]
    h = f.headline()
    assert h["gco"] and h["worse"] >= 1 and h["dollars"] > 0 and h["pockets"] >= h["worse"]
    assert h["tie_outs"] > 0
    assert any(q["says"].startswith("FICO:") or q["says"].startswith("RANR_AMT:") for q in h["open"])
    s = f.steps()
    assert [x.mark for x in s] == ["done"] * 5 and s[3].sub == "All answered"
    assert re.fullmatch(r"\d\d:\d\d · \d+ seconds?", s[4].sub)
    assert f.open_at() == ""
    assert load_workbook(launcher.book_for(f.extract)).active.title == "Start here"



# ---- the window itself, where there is a display


def test_the_window_draws_every_state_and_its_buttons_follow_the_flow(tmp_path, monkeypatch):
    import time
    tk = pytest.importorskip("tkinter")
    try:
        root = tk.Tk()
    except tk.TclError:
        pytest.skip("no display to open a window on")
    monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
    try:
        w = launcher.build(root)
        flow = w["flow"]

        def settle():
            for _ in range(1200):
                root.update()
                if not flow.busy:
                    return
                time.sleep(0.05)
        assert str(w["setup"].cget("state")) == "disabled"                      # L1, nothing picked
        w["extract"].set(str(synth.write_extract(tmp_path, n=1500)))
        root.update()
        assert str(w["setup"].cget("state")) == "normal"
        w["setup"].invoke()
        settle()
        assert flow.screen() == "L2" and str(w["next"].cget("state")) == "normal"
        w["box_REV_DEBT_c"].event_generate("<Button-1>")                        # the split radio, clicked
        root.update()
        assert flow.split == "REV_DEBT" and "split by REV_DEBT" in w["summary"].cget("text")
        w["mode_new"].event_generate("<Button-1>")
        root.update()
        assert flow.mode == "new" and str(w["next"].cget("state")) == "disabled"   # nothing to test yet
        w["mode_bleed"].event_generate("<Button-1>")
        root.update()
        w["next"].invoke()
        settle()
        assert flow.screen() == "L2b" and str(w["run"].cget("state")) == "normal"
        w["run"].invoke()
        settle()
        assert flow.screen() == "L3" and "open_at_0" in w
        flow.gate.missing = ["numpy"]
        w["render"]()
        assert str(w["setup"].cget("state")) == "disabled" and w["headline"].winfo_exists()
    finally:
        root.destroy()
