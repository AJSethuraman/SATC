"""What is learned can be seen and pruned. The firm: "we also need an intuitive
way to go and prune rules that shouldn't have been added as we learn." """

import re

from openpyxl import load_workbook

from origination_cube import cli, config as cfgmod, memory, profile, synth
from origination_cube.ingest import read_table


def _confirmed_file(tmp_path, n=2000):
    _, data = synth.write(tmp_path, n=n)
    path, _, _ = profile.write_cube_file(read_table(data), tmp_path / "cube.yaml")
    s = path.read_text().replace("columns_confirmed: no", "columns_confirmed: yes")
    for k, v in dict(min_units=30, worse_at=1.25, better_at=0.8, confidence=0.95).items():
        s = re.sub(rf'{k}: "\[CONFIRM[^"]*"', f"{k}: {v}", s)
    s = s.replace("{column: FICO, pattern: repeated_value, value: -9999, rows: 40, answer: }",
                  "{column: FICO, pattern: repeated_value, value: -9999, rows: 40, answer: missing}")
    path.write_text(s)
    return path, data


def test_a_confirmed_run_is_remembered_and_an_unconfirmed_one_is_not(tmp_path, capsys):
    path, data = _confirmed_file(tmp_path)
    assert cli.main(["validate", str(path), "--data", str(data)]) == 0
    mem = memory.load()
    assert mem["columns"]["FICO"]["means"] == "fico" and mem["columns"]["GCO_AMT"]["means"] == "gco"
    assert any(v["answer"] == "missing" for v in mem["answers"].values())
    path.write_text(path.read_text().replace("columns_confirmed: yes", "columns_confirmed: no"))
    before = memory.load()
    assert cli.main(["validate", str(path), "--data", str(data)]) == 2
    assert memory.load() == before


def test_the_next_init_uses_what_was_learned(tmp_path):
    path, data = _confirmed_file(tmp_path)
    cli.main(["validate", str(path), "--data", str(data)])
    again, _, _ = profile.write_cube_file(read_table(data), tmp_path / "again.yaml")
    text = again.read_text()
    assert "REMEMBERED - you confirmed this as a FICO score" in text
    assert "answer: missing}   # REMEMBERED - you answered missing" in text


def test_nothing_from_the_data_is_stored(tmp_path):
    path, data = _confirmed_file(tmp_path)
    cli.main(["validate", str(path), "--data", str(data)])
    stored = memory.default_path().read_text()
    for r in read_table(data).rows[:50]:
        assert r["LOAN_NBR"] not in stored


def test_forget_by_name(tmp_path, capsys):
    path, data = _confirmed_file(tmp_path)
    cli.main(["validate", str(path), "--data", str(data)])
    assert cli.main(["memory", "--forget", "FICO"]) == 0
    mem = memory.load()
    assert "FICO" not in mem["columns"] and not any(v["column"] == "FICO" for v in mem["answers"].values())
    assert "GCO_AMT" in mem["columns"]


def test_prune_in_excel(tmp_path):
    path, data = _confirmed_file(tmp_path)
    cli.main(["validate", str(path), "--data", str(data)])
    review = tmp_path / "learned.xlsx"
    assert cli.main(["memory", "--out", str(review)]) == 0
    wb = load_workbook(review)
    ws = wb["Learned"]
    for row in ws.iter_rows(min_row=4):
        if row[2].value == "CHANNEL":
            row[0].value = memory.FORGET
    wb.save(review)
    assert cli.main(["memory", "--read", str(review)]) == 0
    mem = memory.load()
    assert "CHANNEL" not in mem["columns"] and "FICO" in mem["columns"]
