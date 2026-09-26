"""A pocket alone in its band, judged against the rest of its band (the firm,
26 Sep 2026). There is nothing in its band to compare it with, so it is
compared with the rest of the book instead, and its row says so. Before this,
its flag was blank, no reason was given, and "Worst for" skipped it even at 10
times the book's rate (the adversarial pass, reported in its hand-back)."""

from __future__ import annotations

from openpyxl import Workbook

from recalc import calculated
from conftest import cube, row, table
from origination_cube import book, checks, engine

ALONE = "alone in its band: compared with the book"
BANDS = [{"name": "score", "field": "SCORE", "edges": [650, 750]}]


def _bench(compare_to: str) -> dict:
    return {"min_units": 2, "min_events": 1, "worse_at": 1.25, "better_at": 0.8, "confidence": 0.95, "power": 0.8,
            "compare_to": compare_to, "many_tests": "none", "materiality": "none", "shuffles": 200}


def _rows():
    """Two bands of two channels, 2% bad; the top band holds channel A only, 25 bad of 50 (about 10x the book)."""
    out, i = [], 0
    for score, chan, n, bad in ((600, "A", 200, 4), (600, "B", 200, 4), (700, "A", 200, 4), (700, "B", 200, 4),
                                (800, "A", 50, 25)):
        for j in range(n):
            flag = int(j < bad)
            out.append(row(i, score, chan, 1000, flag, 600 * flag, 30 + i % 7))
            i += 1
    return out


def _run(compare_to: str = "peers"):
    return engine.run(cube(bands=BANDS, benchmark=_bench(compare_to)), table(_rows()))


def _cells(res):
    g = res.grids[0]
    lone = g.band_labels[-1]
    return g, lone


def test_a_pocket_alone_in_its_band_is_flagged_against_the_book():
    res, book_res = _run("peers"), _run("topline")
    g, lone = _cells(res)
    s = g.cell(lone, "A").rates["outcome_loans"]
    assert s.rate / res.total.rates["outcome_loans"].rate > 10
    assert s.vs_band is None and s.flag == engine.WORSE
    assert s.alone
    # exactly as under compare_to: book, rate by rate
    for m in res.measures:
        if m.is_rate:
            got, want = g.cell(lone, "A").rates[m.name], book_res.grids[0].cell(lone, "A").rates[m.name]
            assert (got.flag, got.vs_rest, got.p_book, got.test) == (want.flag, want.vs_rest, want.p_book, want.test)
    # and "Worst for" no longer skips it
    worst = [x for x in book._top_lines(res) if x.startswith("Worst for ")]
    assert worst and all(x.endswith(f": SCORE {lone} / CHAN A.") for x in worst), book._top_lines(res)


def test_pockets_with_band_mates_are_still_judged_against_their_band():
    res = _run("peers")
    g, lone = _cells(res)
    for (b, d), c in g.inner():
        if b == lone:
            continue
        for m in res.measures:
            if m.is_rate:
                s = c.rates[m.name]
                assert s.flag == s.reading_band and not s.alone, (b, d, m.name)
    # under compare_to: book nothing is marked alone
    assert not any(c.rates[m.name].alone for _, c in _run("topline").grids[0].inner()
                   for m in res.measures if m.is_rate)


def test_the_families_count_a_lone_pocket_once_against_the_book():
    """Its one test was already in the rest-of-the-book family, and it has none against its band:
    judging it by the book adds no test anywhere."""
    fam = {(f[2], f[3]): f[4] for f in checks.families(_run("peers"))}
    assert fam[("outcome_loans", "the rest of the book")] == 5
    assert fam[("outcome_loans", "the rest of its band")] == 4


def test_a_pocket_alone_in_its_band_says_why_on_its_row():
    res = _run("peers")
    _, lone = _cells(res)
    ws = Workbook().active
    book._bleeds(ws, res)
    ws = calculated(ws)                       # the Test column is a formula (OC-40)
    heads = [c.value for c in ws[4]]
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=5) if r[1].value]
    band, test = heads.index("Band"), heads.index("Test")
    said = [r for r in rows if r[band] == lone]
    assert said and all(ALONE in (r[test] or "") for r in said), said
    assert not any(ALONE in (r[test] or "") for r in rows if r[band] != lone)


def test_check_counts_the_pockets_alone_in_their_band():
    res = _run("peers")
    got = [v for k, v in checks.rows(res) if k == "Alone in its band"]
    assert got == ["1 pocket had no other pocket in its band, so it was compared with the rest of the book."]
    assert not [v for k, v in checks.rows(_run("topline")) if k == "Alone in its band"]


def test_losses_vs_revenue_says_why_in_its_own_column_and_together_stays_the_pair():
    """Found 26 Sep 2026 by the full suite: the note first went into Together, which only ever reads the
    pair (priced for it, net drain, safe but idle)."""
    res = _run("peers")
    _, lone = _cells(res)
    ws = Workbook().active
    book._losses_vs_revenue(ws, res)
    ws = calculated(ws)                       # Compared with and Together are formulas (OC-40)
    head = next(r for r in ws.iter_rows() if any(c.value == "Compared with" for c in r))
    cols = {c.value: c.column for c in head if c.value}
    rows = [r for r in ws.iter_rows(min_row=head[0].row + 1) if r[cols["Band"] - 1].value]
    lone_rows = [r for r in rows if r[cols["Band"] - 1].value == lone]
    assert lone_rows and all(r[cols["Compared with"] - 1].value == ALONE for r in lone_rows)
    assert all(ALONE not in str(r[cols["Together"] - 1].value or "") for r in rows)
    assert not any(r[cols["Compared with"] - 1].value for r in rows if r[cols["Band"] - 1].value != lone)
