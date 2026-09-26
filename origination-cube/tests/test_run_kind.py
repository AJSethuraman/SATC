"""What are you running? (the firm, 26 Sep 2026: "it likely makes sense for the
script to ask which we are doing so it does indeed have the minimum required").

Two answers on Control, and each checks its own minimum:
- Where the book bleeds: the five core columns, and never a date. A pre-spec is
  refused: the run isn't a test of a new variable.
- Finding and testing a new variable: the core columns, a column marked
  Origination date, and then one more answer: scout first (not built, so
  refused) or test from a pre-spec, whose file must be named and whose column
  must be on Columns.
A blank answer is refused by its cell, like every other call on Control: the
tool never picks for the analyst."""

import csv
import shutil

import pytest
from openpyxl import load_workbook

from origination_cube import book, confirmatory, control, synth
from test_book import _answer
from test_book_dates import _check, _columns, _control
from test_confirmatory import _held, _log, _spec, git

BLEED, NEW = "Where the book bleeds", "Finding and testing a new variable"
SCOUT, FROM_SPEC = "Scout first", "Test from a pre-spec"
needs_git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed on this machine")


def _cell(b, key) -> str:
    return f"Control!C{control.row_of(load_workbook(b)[control.SHEET], key)}"


def _without(x, *drop):
    """The extract with some columns taken out."""
    with open(x, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    with open(x, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=[k for k in rows[0] if k not in drop])
        w.writeheader()
        w.writerows({k: v for k, v in r.items() if k not in drop} for r in rows)
    return x


def _refused(ran) -> str:
    assert not ran.ok, ran.lines
    text = "\n".join(ran.lines)
    assert "Traceback" not in text and "`" not in text
    return text


def _new_variable(tmp_path, monkeypatch, n=1500, step=FROM_SPEC):
    """A new-variable run ready but for the pre-spec: INCOME / SALES made on Control and splitting the pockets."""
    from test_confirmatory import _ready
    x, b = _ready(tmp_path, monkeypatch, n=n)
    _control(b, run_kind=NEW, new_variable_step=step)
    return x, b


# --------------------------------------------------------------------------
# The question itself


def test_control_asks_what_you_are_running_first_with_two_answers_and_no_default(tmp_path):
    b = book.set_up(synth.write_extract(tmp_path, n=1500)).book
    s = {x.key: x for x in control.load_settings()}
    assert s["run_kind"].question == "What are you running?"
    assert [o.label for o in s["run_kind"].options] == [BLEED, NEW]
    assert s["new_variable_step"].question == "Scout first, or test from a pre-spec already written?"
    assert [o.label for o in s["new_variable_step"].options] == [SCOUT, FROM_SPEC]
    assert s["run_kind"].judgment and s["new_variable_step"].judgment
    ws = load_workbook(b)[control.SHEET]
    keys = [r[control.KEY_COL - 1].value for r in ws.iter_rows(min_row=control.FIRST_ROW)
            if r[control.KEY_COL - 1].value]
    assert keys[:2] == ["run_kind", "new_variable_step"]                   # the first thing on the tab
    for k in keys[:2]:
        assert ws.cell(row=control.row_of(ws, k), column=control.CHOOSE_COL).value is None     # nothing picked
    # the scouting answer says plainly, on the tab, that it isn't built
    said = {o.label: o.explains for o in s["new_variable_step"].options}
    assert "isn't built yet" in said[SCOUT]
    # Set up asks it, in the launcher's words
    out = book.set_up(synth.write_extract(tmp_path, n=1500))
    assert any('"What are you running?"' in line and BLEED in line and NEW in line for line in out.lines), out.lines


def test_a_blank_answer_is_refused_by_its_cell_and_nothing_is_picked_for_you(tmp_path):
    b = book.set_up(synth.write_extract(tmp_path, n=1500)).book
    _answer(b)
    _control(b, run_kind=None)
    text = _refused(book.run(b))
    assert f'{_cell(b, "run_kind")}: "What are you running?" needs an answer. Pick one from the list.' in text
    # the follow-up isn't asked until the answer says it applies
    assert "Scout first" not in text
    wb = load_workbook(b)
    assert "Couldn't run" in _log(b)[0]
    assert wb[control.SHEET].cell(row=control.row_of(wb[control.SHEET], "run_kind"),
                                  column=control.CHOOSE_COL).value is None


def test_the_follow_up_is_refused_blank_only_when_testing_a_new_variable(tmp_path):
    x = synth.write_extract(tmp_path, n=1500)
    b = book.set_up(x).book
    _answer(b)
    _control(b, run_kind=NEW, new_variable_step=None)
    text = _refused(book.run(b))
    assert (f'{_cell(b, "new_variable_step")}: "Scout first, or test from a pre-spec already written?" needs an '
            f'answer. Pick one from the list.') in text
    _control(b, run_kind=BLEED)
    assert book.run(b).ok


# --------------------------------------------------------------------------
# Where the book bleeds: the five core columns, and no date


def test_a_bleed_run_needs_only_the_five_core_columns_and_never_asks_for_a_date(tmp_path):
    x = _without(synth.write_extract(tmp_path, n=3000), "ORIG_DATE")
    out = book.set_up(x)
    _answer(out.book)
    b = out.book
    ran = book.run(b)
    assert ran.ok, ran.lines
    said = list(out.lines) + list(ran.lines) + _log(b) + [f"{k} {v}" for k, v in _check(b).items()]
    assert not any("date" in str(s).lower() for s in said), [s for s in said if "date" in str(s).lower()]
    chk = _check(b)
    assert chk["What was run"] == BLEED
    assert "Pre-spec" not in chk and "Holdout" not in chk


def test_a_bleed_run_refuses_a_missing_core_column_by_name(tmp_path):
    x = synth.write_extract(tmp_path, n=1500)
    b = book.set_up(x).book
    _answer(b)
    for name in ("GCO_AMT", "RANR_AMT"):
        _columns(b, name, C_MEANS="Not used")
    text = _refused(book.run(b))
    assert f"Columns: {BLEED} needs one column marked GCO dollars, and none is." in text
    assert f"Columns: {BLEED} needs one column marked RANR dollars, and none is." in text
    assert "Origination date" not in text


def test_a_bleed_run_uses_a_marked_origination_date_only_for_checks_line(tmp_path):
    x = synth.write_extract(tmp_path, n=1500)
    b = book.set_up(x).book
    _answer(b)
    assert book.run(b).ok
    chk = _check(b)
    assert chk["Origination dates"].endswith("(1,500 loans; 0 without a readable date)")
    assert "Holdout" not in chk and not any(k.startswith("Pre-spec") for k in chk)


def test_a_pre_spec_under_a_bleed_run_is_refused_with_its_reason(tmp_path, monkeypatch):
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    x = synth.write_extract(tmp_path, n=1500)
    b = book.set_up(x).book
    _answer(b)
    f = _spec(x.parent, commit=False)
    cell = _held(b, str(f))
    text = _refused(book.run(b))
    assert (f'{cell}: Where the book bleeds isn\'t a test of a new variable, so it isn\'t held to a pre-spec. '
            f'Clear the cell, or change "What are you running?" ({_cell(b, "run_kind")}).') in text
    _held(b, None)
    assert book.run(b).ok


# --------------------------------------------------------------------------
# Finding and testing a new variable: an origination date, and the column tested


def test_a_new_variable_run_needs_an_origination_date_and_says_so_by_name(tmp_path, monkeypatch):
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    x = _without(synth.write_extract(tmp_path / "x", n=1500, ratio=True), "ORIG_DATE")
    b = book.set_up(x).book
    _answer(b)
    _control(b, run_kind=NEW, new_variable_step=FROM_SPEC,
             **{"derived|1": ("INCOME_TO_SALES", "INCOME", "SALES")})
    book.set_up(x)
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    _held(b, str(_spec(x.parent, commit=False)))
    text = _refused(book.run(b))
    assert f"Columns: {NEW} needs one column marked Origination date, and none is." in text
    # the same extract runs Where the book bleeds once the pre-spec is cleared
    _held(b, None)
    _control(b, run_kind=BLEED)
    assert book.run(b).ok


def test_a_new_variable_run_refuses_two_origination_dates(tmp_path, monkeypatch):
    x, b = _new_variable(tmp_path, monkeypatch)
    _held(b, str(_spec(x.parent, commit=False)))
    _columns(b, "INCOME", C_MEANS="Origination date")
    text = _refused(book.run(b))
    assert f"Columns: {NEW} needs one column marked Origination date, and 2 are: ORIG_DATE, INCOME." in text


def test_scouting_is_refused_as_not_built(tmp_path, monkeypatch):
    x, b = _new_variable(tmp_path, monkeypatch, step=SCOUT)
    text = _refused(book.run(b))
    assert (f"{_cell(b, 'new_variable_step')}: Scouting isn't built yet. Pick \"{FROM_SPEC}\", or run "
            f"{BLEED}.") in text
    assert "Scouting isn't built yet" in "\n".join(_log(b))


def test_testing_from_a_pre_spec_needs_the_file_named(tmp_path, monkeypatch):
    x, b = _new_variable(tmp_path, monkeypatch)
    cell = _held(b, None)
    text = _refused(book.run(b))
    assert f'{cell}: "{FROM_SPEC}" needs the pre-spec file named here.' in text


def test_the_column_a_pre_spec_tests_must_be_on_columns(tmp_path, monkeypatch):
    x, b = _new_variable(tmp_path, monkeypatch)
    cell = _held(b, str(_spec(x.parent, commit=False, column="DEBT_TO_SALES")))
    text = _refused(book.run(b))
    assert (f"{cell}: the pre-spec tests DEBT_TO_SALES, and no column on Columns has that name. Make it under "
            f"New columns on this tab and press Set up again, or fix the pre-spec.") in text


@needs_git
def test_what_was_run_is_recorded_on_check_the_log_control_and_the_record(tmp_path, monkeypatch):
    x, b = _new_variable(tmp_path, monkeypatch)
    _spec(x.parent)
    _held(b, "prespec.yaml")
    ran = book.run(b)
    assert ran.ok, ran.lines
    said = f"{NEW}: {FROM_SPEC.lower()}"
    assert _check(b)["What was run"] == said
    assert _log(b)[0].endswith(f"What was run: {said}.")
    assert any(line.endswith(f"What was run: {said}.") for line in ran.lines)
    ws = load_workbook(b)[control.SHEET]
    used = {r[control.KEY_COL - 1].value: r[control.KEY_COL].value for r in ws.iter_rows(min_row=control.FIRST_ROW)}
    assert used["run_kind"] == NEW and used["new_variable_step"] == FROM_SPEC
    record = b.with_name(f"{b.stem} - what ran.yaml").read_text(encoding="utf-8")
    assert f"# What was run: {said}\n" in record
    # and a bleed run records itself, with no follow-up
    _held(b, None)
    _control(b, run_kind=BLEED)
    assert book.run(b).ok
    assert _check(b)["What was run"] == BLEED
    assert _log(b)[0].endswith(f"What was run: {BLEED}.")
    used = {r[control.KEY_COL - 1].value: r[control.KEY_COL].value
            for r in load_workbook(b)[control.SHEET].iter_rows(min_row=control.FIRST_ROW)}
    assert used["run_kind"] == BLEED and not used.get("new_variable_step")
