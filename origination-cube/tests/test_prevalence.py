"""The Prevalence tab (NEXT-GOAL 3.12): loans and booked dollars per group per
pocket, with no test attached. docs/statistics.md B8: "A count of the book, not
a finding." The groups are the split column's halves (or values) and each new
column's bands; every count here is made by hand from the rows."""

import csv
from bisect import bisect_right

from openpyxl import load_workbook

from conftest import cube, table
from origination_cube import book, engine, prevalence, synth
from test_book import _answer
from test_book_dates import _columns, _control

BENCH = {"min_units": 2, "min_events": 1, "worse_at": 1.25, "better_at": 0.8, "confidence": 0.95, "power": 0.8,
         "compare_to": "peers", "many_tests": "none", "materiality": "none"}

#: three pockets, small enough to count by eye: (score, channel, X, booked, income, sales)
LOANS = [(600, "A", 1, 100, 20, 100), (600, "A", 2, 100, 60, 100), (600, "A", 3, 200, 150, 100),
         (600, "A", 4, 400, 80, 100),
         (700, "A", 5, 100, 40, 100), (700, "A", "", 50, 30, 0),
         (700, "B", 1, 10, 100, 100), (700, "B", 9, 20, 200, 100), (700, "B", 5, 30, 50, 100)]


def _res(**extra):
    rows = [{"ID": f"L{i}", "SCORE": s, "CHAN": c, "X": x, "BAL": bal, "INC": inc, "SALES": sales,
             "BAD": i % 2, "GCO": bal * (i % 2) / 2, "RANR": 5 + i} for i, (s, c, x, bal, inc, sales) in
            enumerate(LOANS)]
    res = engine.run(cube(benchmark=BENCH, split={"field": "X", "how": "own_median"},
                          derived=[{"name": "R", "top": "INC", "bottom": "SALES"}], **extra), table(rows))
    res.typed_edges = {"R": "0.5; 1"}
    return res


def test_each_pocket_is_counted_by_the_splits_halves():
    res = _res()
    (halves, bands), notes = prevalence.groupings(res)
    assert not notes and halves.kind == "halves" and bands.kind == "bands"
    g = res.grids[0]
    low_band, high_band = g.band_labels                    # 600 - 649 and 650 - 700
    order, per, shown = prevalence.count(res, g, halves)
    assert shown == ["High half", "Low half", "No X"]
    H, L, N = engine.HIGH, engine.LOW, engine.NO_SPLIT_VALUE
    # 600 / A: X is 1, 2, 3, 4, so its median is 2.5; 700 / A: only X = 5, which is its median, so low;
    # 700 / B: 1, 9, 5, median 5, so 9 is the only high one
    assert per[(low_band, "A")] == {H: [2, 600.0], L: [2, 200.0]}
    assert per[(high_band, "A")] == {L: [1, 100.0], N: [1, 50.0]}
    assert per[(high_band, "B")] == {H: [1, 20.0], L: [2, 40.0]}


def test_each_pocket_is_counted_by_a_new_columns_bands():
    res = _res()
    bands = prevalence.groupings(res)[0][1]
    assert bands.edges == (0.5, 1.0) and "R = INC ÷ SALES, by its bands (0.5; 1: the band edges on Columns)" == \
        bands.title
    g = res.grids[0]
    low_band, high_band = g.band_labels
    order, per, shown = prevalence.count(res, g, bands)
    # R is 0.2, 0.6, 1.5, 0.8 | 0.4, blank (sales 0) | 1.0, 2.0, 0.5; named as the grids name bands
    assert shown == ["0.2 - 0.4", "0.5 - 0.9", "1.0 - 2.0", "(blank)"]
    lo, mid, hi, blank = order
    assert per[(low_band, "A")] == {lo: [1, 100.0], mid: [2, 500.0], hi: [1, 200.0]}
    assert per[(high_band, "A")] == {lo: [1, 100.0], blank: [1, 50.0]}
    assert per[(high_band, "B")] == {hi: [2, 30.0], mid: [1, 30.0]}


def test_edges_typed_as_every_are_worked_out_over_the_values():
    res = _res()
    res.typed_edges = {"R": "every 0.5"}
    # R runs from 0.2 to 2.0: the first edge is the first multiple of 0.5 above 0.2
    assert prevalence.edges_of(res, "R") == ((0.5, 1.0, 1.5, 2.0), "every 0.5 on Columns")


def test_a_new_column_with_no_edges_is_said_not_counted():
    res = _res()
    res.typed_edges = {}
    gs, notes = prevalence.groupings(res)
    assert [g.kind for g in gs] == ["halves"]
    assert notes == ["R = INC ÷ SALES has no band edges, so it isn't counted by band here. Type its edges on "
                     "Columns to count it."]


def test_a_count_that_does_not_add_up_to_the_grid_is_not_shown():
    res = _res()
    g = res.grids[0]
    halves = prevalence.groupings(res)[0][0]
    assert prevalence.count(res, g, halves) is not None
    g.cells[next(k for k, _ in g.inner())].rows += 1           # the grid now says one loan more
    assert prevalence.count(res, g, halves) is None


# --------------------------------------------------------------------------
# On the workbook


def test_the_prevalence_tab_counts_the_book_by_the_new_columns_bands(tmp_path):
    x = synth.write_extract(tmp_path, n=3000, ratio=True)
    b = book.set_up(x).book
    _answer(b)
    _control(b, **{"derived|1": ("INCOME_TO_SALES", "INCOME", "SALES")})
    book.set_up(x)
    _columns(b, "INCOME_TO_SALES", C_SPLIT="Yes", C_EDGES="0.1; 0.25; 0.5; 1; 2")
    for c in ("ORIG_BAL", "REV_DEBT", "ASSET_CLASS", "INCOME", "SALES"):
        _columns(b, c, C_CUT="No")
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    ran = book.run(b)
    assert ran.ok, ran.lines
    wb = load_workbook(b)
    names = wb.sheetnames
    assert names.index("Split") + 1 == names.index("Prevalence") == names.index("Three-way") - 1
    ws = wb["Prevalence"]
    assert ws["B1"].value == "Prevalence: a count, not a test"
    assert "Nothing here is tested" in ws["B2"].value
    text = [ws.cell(row=r, column=2).value for r in range(1, ws.max_row + 1)]
    assert "INCOME_TO_SALES, each pocket cut at its own median" in text
    assert "INCOME_TO_SALES = INCOME ÷ SALES, by its bands (0.1; 0.25; 0.5; 1; 2: the band edges on Columns)" in text
    totals = [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "Every pocket"]
    assert len(totals) == 2                                    # one grid, two ways of dividing it

    with open(x, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    edges = [0.1, 0.25, 0.5, 1, 2]
    by_hand = [[0, 0.0] for _ in range(len(edges) + 2)]         # six bands, then blank
    for r in rows:
        bal = float(r["ORIG_BAL"]) if r["ORIG_BAL"] else 0.0
        i = len(edges) + 1 if r["SALES"] in ("", "0") else bisect_right(edges, float(r["INCOME"]) / float(r["SALES"]))
        by_hand[i][0] += 1
        by_hand[i][1] += bal
    halves, bands = totals

    def cells(r):
        return [ws.cell(row=r, column=c).value for c in range(4, ws.max_column + 1)]

    got = cells(bands)
    assert got[0] == 3000
    assert [got[2 + 2 * i] for i in range(len(by_hand))] == [k for k, _ in by_hand]
    assert all(abs(got[3 + 2 * i] - d) < 0.01 for i, (_, d) in enumerate(by_hand))
    got = cells(halves)
    assert got[0] == 3000 and got[2] + got[4] + got[6] == 3000 and got[6] == by_hand[-1][0]   # no ratio: no half
