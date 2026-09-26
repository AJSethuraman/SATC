"""Check's newer lines (NEXT-GOAL 3.16 to 3.18): the pocket budget and how much
of the book sits in testable pockets, how many families of tests the run holds,
and a warning when credit products are mixed in a pocket. Every number is
worked out here by hand from the rows, not read back from the engine.

docs/statistics.md B9: the most pockets an extract can test is its bad loans
divided by 5, and a single red across many families is weak evidence."""

import csv
import random

import pytest
from openpyxl import load_workbook

from conftest import cube, table
from origination_cube import book, checks, engine, meanings, synth
from test_book import _answer

BENCH = {"min_units": 30, "min_events": 51, "worse_at": 1.25, "better_at": 0.8, "confidence": 0.95, "power": 0.8,
         "compare_to": "peers", "many_tests": "bh", "materiality": "none"}


def _test1(bad=(150, 50, 50, 50), n=1000, bal=(300, 100, 100, 100)):
    """docs/statistics.md's Test 1 book: four pockets of 1,000 (score 600 / 700 x Broker / Branch), 50 bad
    in each but Broker 600 with 150. Broker 600's loans are booked at 300, the rest at 100."""
    rows, i = [], 0
    for (score, chan), k, amount in zip(((600, "Broker"), (600, "Branch"), (700, "Broker"), (700, "Branch")), bad, bal):
        for j in range(n):
            flag = 1 if j < k else 0
            rows.append({"ID": f"L{i}", "SCORE": score, "CHAN": chan, "BAL": amount, "BAD": flag,
                         "GCO": amount * 0.5 * flag, "RANR": 10 + (i % 7), "X": (i * 37) % 101})
            i += 1
    return rows


def _rows(res, label):
    return [v for k, v in checks.rows(res) if k == label]


def test_the_pocket_budget_is_the_bad_loans_over_five_for_each_grid():
    res = engine.run(cube(benchmark=BENCH), table(_test1()))
    got = dict(checks.budget_rows(res))
    # 300 bad loans: 300 / 5 = 60 pockets at most, per grid
    assert got["Pocket budget"].startswith("At most 60 pockets per grid can be tested: at the suggested floor a "
                                           "pocket needs to expect 5 bad loans, and the book has 300 (300 ÷ 5 = 60).")
    # fewest losses is 51, so only Broker 600 (150 bad) can be tested: a quarter of the loans, and half the
    # booked dollars (1,000 x 300 of 1,000 x 300 + 3,000 x 100)
    assert got["Pockets: SCORE x CHAN"] == ("4 pockets, against a budget of 60. 1 has at least 51 bad loans and "
                                            "can be tested, holding 25% of the loans and 50% of the booked dollars.")
    assert got["Pockets in all"].startswith("4 pockets across 1 grid. Testable means at least 51 bad loans "
                                            "(fewest losses): that is the only floor")
    assert "Warning" not in got


def test_more_pockets_than_the_book_can_test_is_warned_about():
    # 4 + 3 + 2 + 1 = 10 bad loans: a budget of 2, and the grid has 4 pockets
    res = engine.run(cube(benchmark={**BENCH, "min_events": 1}), table(_test1(bad=(4, 3, 2, 1), n=200)))
    got = checks.budget_rows(res)
    assert ("Warning", "SCORE x CHAN has 4 pockets, more than the 2 this book can test. Most will have too few bad "
                       "loans to say anything: use fewer bands or segments.") in got
    assert dict(got)["Pockets: SCORE x CHAN"].startswith("4 pockets, against a budget of 2. 4 have at least 1 bad "
                                                         "loans and can be tested, holding 100% of the loans")


def test_the_families_are_grids_by_rates_by_comparisons():
    two_grids = [{"name": "score", "field": "SCORE", "edges": [650]}, {"name": "bal", "field": "BAL", "edges": [200]}]
    res = engine.run(cube(benchmark=BENCH, bands=two_grids), table(_test1()))
    got = dict(checks.family_rows(res))
    # 2 grids x 4 rates x 2 comparisons. Tests per rate: the score grid's 4 pockets, each against the book and
    # its band (8); the balance grid's 3 pockets against the book, but only 2 against their band, since Broker
    # 600 is alone in the 300 band (5). 13 x 4 rates = 52
    assert got["Families of tests"].startswith("16, holding 52 tests: 2 grids x 4 rates x 2 comparisons.")
    assert got["Reading a single red"] == ("Each family gets its own allowance for many tests, not one for the whole "
                                           "run. So a single red across 16 families is weak evidence.")
    # with no allowance, every test stands alone
    res = engine.run(cube(benchmark={**BENCH, "many_tests": "none"}, bands=two_grids), table(_test1()))
    assert dict(checks.family_rows(res))["Reading a single red"] == (
        "No allowance for many tests is in use, so each of the 52 tests stands alone: a single red among them is "
        "weak evidence.")


def test_a_rate_nothing_tested_is_no_family():
    """With no shuffles the dollar rates carry no test, so only the share of loans makes families."""
    from dataclasses import replace
    cfg = cube(benchmark=BENCH)
    res = engine.run(replace(cfg, benchmark=replace(cfg.benchmark, shuffles=0)), table(_test1()))
    assert dict(checks.family_rows(res))["Families of tests"].startswith(
        "2, holding 8 tests: 1 grid x 1 rate x 2 comparisons.")


def test_the_split_adds_its_three_way_grid_and_its_halves_as_families():
    res = engine.run(cube(benchmark={**BENCH, "min_events": 1}, split={"field": "X", "how": "own_median"}),
                     table(_test1()))
    fam = checks.families(res)
    kinds = sorted({(f[0], f[3]) for f in fam})
    assert kinds == [("grids", "the rest of its band"), ("grids", "the rest of the book"),
                     ("split halves", "the other half"), ("three-way grids", "the rest of its band"),
                     ("three-way grids", "the rest of the book")]
    # 8 on the grid, 8 on the three-way grid, 4 on the split's halves (one per rate)
    assert dict(checks.family_rows(res))["Families of tests"].startswith(
        "20, holding ")
    assert "1 grid x 4 rates x 2 comparisons; 1 three-way grid x 4 rates x 2 comparisons; 1 grid x 4 rates for " \
           "the split's halves" in dict(checks.family_rows(res))["Families of tests"]
    # the split's halves: every pocket's halves compared, once per rate
    assert sum(f[4] for f in fam if f[0] == "split halves") == 4 * 4


# --------------------------------------------------------------------------
# 3.18 credit products


def _product_cube(cut_product: bool, dims_extra=()):
    from origination_cube import config as cfgmod
    dims = [{"name": "chan", "field": "CHAN"}] + ([{"name": "prod", "field": "LOAN_TYPE"}] if cut_product else [])
    return cfgmod.parse({
        "name": "t", "schema_version": 1, "columns_confirmed": True,
        "columns": {"ID": "key", "BAL": "booked", "BAD": "outcome", "GCO": "gco", "RANR": "ranr", "SCORE": "fico",
                    "CHAN": "category", "LOAN_TYPE": "product", "X": "amount"},
        "bands": [{"name": "score", "field": "SCORE", "edges": [650]}], "dimensions": dims,
        "measures": [{"name": "loans", "mode": "count"}], "benchmark": BENCH, "min_age_months": 0})


def _with_products(products):
    rows = _test1(n=100, bad=(15, 5, 5, 5))
    for i, r in enumerate(rows):
        r["LOAN_TYPE"] = products[i % len(products)]
    return rows


def test_products_mixed_in_a_pocket_are_warned_about():
    res = engine.run(_product_cube(cut_product=False), table(_with_products(["Auto", "Card", "", "HELOC"])))
    assert checks.product_rows(res) == [("Warning", "LOAN_TYPE holds 3 credit products (Auto, Card, HELOC) and "
                                                    "isn't a band or segment, so profit per booked dollar is compared "
                                                    "across products. A pocket can read better or worse just for "
                                                    "holding more of one.")]


def test_one_product_or_products_cut_by_is_not_warned_about():
    assert checks.product_rows(engine.run(_product_cube(cut_product=True),
                                          table(_with_products(["Auto", "Card"])))) == []
    assert checks.product_rows(engine.run(_product_cube(cut_product=False),
                                          table(_with_products(["Auto", ""])))) == []


def test_a_credit_product_column_is_recognised_by_its_name():
    rows = []
    for i in range(300):
        rows.append({"ID": f"L{i}", "LOAN_TYPE": ["Auto", "Card", "HELOC"][i % 3], "PROD_TYPE": f"P{i % 4}",
                     "Product": ["Term", "Line"][i % 2], "OTHER_KIND": ["a", "b"][i % 2]})
    got = {c: s.means for c, s in meanings.suggest(table(rows)).items()}
    assert got["LOAN_TYPE"] == got["PROD_TYPE"] == got["Product"] == "product"
    assert got["OTHER_KIND"] == "category"
    assert meanings.catalog()["product"].label == "Credit product" and meanings.catalog()["product"].cut == "dimension"


# --------------------------------------------------------------------------
# On the workbook: the lines land on Check, counted from the extract by hand


def test_check_carries_the_budget_the_families_and_the_product_mix(tmp_path):
    x = synth.write_extract(tmp_path, n=3000)
    with open(x, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    rng = random.Random(3)
    for r in rows:
        r["LOAN_TYPE"] = rng.choice(["Auto", "Card"])
    with open(x, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    out = book.set_up(x)
    _answer(out.book)
    wb = load_workbook(out.book)
    for r in wb["Columns"].iter_rows(min_row=book.COL_FIRST):
        if r[book.C_NAME - 1].value == "LOAN_TYPE":
            assert r[book.C_MEANS - 1].value == "Credit product" and r[book.C_CUT - 1].value == "Yes"
            r[book.C_CUT - 1].value = "No"
    wb.save(out.book)
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    check = {}
    for r in load_workbook(out.book)["Check"].iter_rows(min_row=4):
        if r[1].value:
            check.setdefault(r[1].value, []).append(r[2].value)
    bad = sum(1 for r in rows if r["BAD_FLAG"] == "1")
    assert check["Pocket budget"][0].startswith(f"At most {bad // 5:,} pockets per grid can be tested")
    assert f"the book has {bad:,} ({bad:,} ÷ 5 = {bad // 5:,})" in check["Pocket budget"][0]
    grids = [k for k in check if k.startswith("Pockets: ")]
    assert "Pockets: FICO x CHANNEL" in grids and "Pockets: ORIG_BAL x ASSET_CLASS" in grids
    assert check["Families of tests"][0].startswith(f"{len(grids) * 8}, holding ")
    assert check["Reading a single red"][0].endswith(f"a single red across {len(grids) * 8} families is weak "
                                                     f"evidence.")
    assert any(w.startswith("LOAN_TYPE holds 2 credit products (Auto, Card) and isn't a band or segment")
               for w in check["Warning"])
