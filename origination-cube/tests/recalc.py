"""Recalculate a workbook through LibreOffice, for the tests that read what a formula gives.

openpyxl writes formulas and never calculates them, so a workbook the cube writes has no values in its live
cells until a spreadsheet program opens it. LibreOffice headless converts a copy to xlsx, calculating every
formula on the way, and the copy is read back with openpyxl's data_only.

It calculates for certain, not only when a file has no cached values: the conversion runs under a profile of
its own whose one setting is "recalculate Excel files on load: always" (OOXMLRecalcMode = 0). A file saved by
LibreOffice carries its values; changing an input with openpyxl and converting again gives the new answer,
which tests/test_live.py proves.

About a second a workbook, so a few end-to-end tests use it rather than every test."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

SOFFICE = shutil.which("soffice") or ("/usr/bin/soffice" if Path("/usr/bin/soffice").exists() else None)

PROFILE_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema"
           xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""

_profile: Path | None = None


def need_soffice():
    if SOFFICE is None:
        pytest.skip("LibreOffice (soffice) isn't installed, so the workbook's formulas can't be calculated here")


def _profile_dir() -> Path:
    global _profile
    if _profile is None:
        _profile = Path(tempfile.mkdtemp(prefix="cube-lo-profile-"))
        user = _profile / "user"
        user.mkdir(parents=True)
        (user / "registrymodifications.xcu").write_text(PROFILE_XCU, encoding="utf-8")
    return _profile


def recalc_file(path: str | Path, out_dir: str | Path | None = None) -> Path:
    """A calculated copy of the workbook at `path`, beside it (or in `out_dir`)."""
    need_soffice()
    path = Path(path)
    out = Path(out_dir) if out_dir else path.parent / "recalculated"
    out.mkdir(parents=True, exist_ok=True)
    got = subprocess.run([SOFFICE, f"-env:UserInstallation={_profile_dir().as_uri()}", "--headless", "--calc",
                          "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                         capture_output=True, text=True, timeout=300)
    done = out / (path.stem + ".xlsx")
    if got.returncode != 0 or not done.exists():
        raise RuntimeError(f"LibreOffice couldn't calculate {path.name}: {got.stdout} {got.stderr}")
    return done


def recalc(path: str | Path, out_dir: str | Path | None = None):
    """The calculated workbook's values (openpyxl, data_only)."""
    return load_workbook(recalc_file(path, out_dir), data_only=True)


def values_of(wb, tmp_path: Path, name: str = "book.xlsx"):
    """Save an in-memory workbook, calculate it, and give back its values."""
    Path(tmp_path).mkdir(parents=True, exist_ok=True)
    p = Path(tmp_path) / name
    wb.save(p)
    return recalc(p)


def calculated_book(path: str | Path):
    """The workbook at `path` as LibreOffice calculates it, each sheet carrying the sheet as written (its
    formulas, number formats and rules) as .formulas."""
    values, written = recalc(path), load_workbook(path)
    for ws in values.worksheets:
        ws.formulas = written[ws.title]
    return values


def calculated(ws):
    """The same sheet, calculated: its workbook saved to a temporary folder and recalculated, for a test that
    writes one tab with a book._* writer and reads what its formulas give."""
    d = Path(tempfile.mkdtemp(prefix="cube-calc-"))
    try:
        return values_of(ws.parent, d)[ws.title]
    finally:
        shutil.rmtree(d, ignore_errors=True)
