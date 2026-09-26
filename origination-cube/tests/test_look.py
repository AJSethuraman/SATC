"""The Look tab (fix 3.8, the audit's item f: band edges were chosen blind).

Set up shows each number column's spread, repeats and codes beside a histogram;
a Run with a number column splitting the pockets adds a scatter of it against
each band column (docs/statistics.md A9). Every number is checked against a
count made by hand from the CSV, not against look.py's own arithmetic."""

import csv
import statistics
from collections import Counter

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, ScatterChart

from origination_cube import book, look, synth
from origination_cube.ingest import read_table
from test_book_results import _ready, _set

NUMBER_COLUMNS = ["FICO", "ORIG_BAL", "GCO_AMT", "RANR_AMT", "REV_DEBT"]


def _csv(path, col) -> list[str]:
    with open(path, newline="", encoding="utf-8") as f:
        return [row[col].strip() for row in csv.DictReader(f)]


def _numbers(raw: list[str]) -> list[float]:
    out = []
    for v in raw:
        try:
            out.append(float(v))
        except ValueError:
            pass
    return out


def _blocks(ws) -> dict[str, dict]:
    """Each block by column: its row, its labelled lines as (count, share), and
    the most-repeated rows as (value, loans)."""
    out: dict[str, dict] = {}
    cur = None
    for r in range(look.FIRST, ws.max_row + 1):
        b = ws.cell(row=r, column=2)
        c, d = ws.cell(row=r, column=3).value, ws.cell(row=r, column=4).value
        if b.value in NUMBER_COLUMNS and b.fill.fgColor.rgb[-6:] == look.INK:
            cur = out[b.value] = {"row": r, "lines": {}, "repeats": [], "fmt": {}}
        elif cur is None or b.value is None:
            continue
        elif isinstance(b.value, str):
            cur["lines"][b.value] = (c, d)
            cur["fmt"][b.value] = ws.cell(row=r, column=3).number_format
            if b.fill.fgColor.rgb[-6:] == look.INK:          # the scatters' own heading: the blocks are over
                cur = None
        elif "Most repeated" in cur["lines"]:
            cur["repeats"].append((b.value, c))
    return out


def _title(chart) -> str:
    return chart.title.tx.rich.p[0].r[0].t


def _series(hs, header: str, first_col: int = 1) -> list[tuple]:
    """A chart's numbers on the hidden sheet: the column pair headed `header`."""
    for col in range(first_col, hs.max_column + 1):
        if hs.cell(row=3, column=col).value == header:
            return [(hs.cell(row=r, column=col).value, hs.cell(row=r, column=col + 1).value)
                    for r in range(4, hs.max_row + 1) if hs.cell(row=r, column=col).value is not None]
    raise KeyError(header)


def test_set_up_writes_a_block_and_a_histogram_for_every_number_column(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=3000))
    wb = load_workbook(out.book)
    assert wb.sheetnames.index("Look") == wb.sheetnames.index("Columns") + 1
    assert wb["_look"].sheet_state == "hidden"
    ws = wb["Look"]
    assert "before choosing its edges on Columns" in ws["B2"].value
    blocks = _blocks(ws)
    # the loan number, the channel, the 0/1 flag and the four asset classes have no edges to choose
    assert list(blocks) == NUMBER_COLUMNS
    bars = [c for c in ws._charts if isinstance(c, BarChart)]
    assert not [c for c in ws._charts if isinstance(c, ScatterChart)]        # no split before a Run
    assert sorted((c.anchor._from.row + 1, _title(c).split(" in steps of ")[0]) for c in bars) == \
        sorted((b["row"], name) for name, b in blocks.items())                # one beside each block
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(t.startswith("None yet.") for t in text)


def test_the_numbers_match_a_count_by_hand(tmp_path):
    x = synth.write_extract(tmp_path, n=3000)
    wb = load_workbook(book.set_up(x).book)
    blocks = _blocks(wb["Look"])
    hs = wb["_look"]

    fico = _csv(x, "FICO")
    blank = sum(1 for v in fico if v == "")
    nums = _numbers(fico)
    rest = [v for v in nums if v != -9999]
    f = blocks["FICO"]["lines"]
    assert f["Loans"][0] == 3000
    assert f["Blank"] == (blank, blank / 3000) and blank == 1
    assert f["At -9999, likely a code"] == (60, 60 / 3000)               # every 50th loan
    assert f[f"The other {len(rest):,} values"] is not None
    assert f["Smallest"][0] == min(rest) and f["Largest"][0] == max(rest)
    assert f["Median"][0] == statistics.median(rest)
    assert min(rest) > 300                                                # the code is out of the spread
    assert blocks["FICO"]["fmt"]["Median"] == "#,##0"                     # a score reads as a whole number
    # the histogram holds every loan in the spread, and none of the code
    assert sum(k for _, k in _series(hs, "FICO")) == len(rest)

    bal = _csv(x, "ORIG_BAL")
    got = blocks["ORIG_BAL"]["lines"]
    vals = _numbers(bal)
    assert got["Blank"] == (1, 1 / 3000) and got["Likely a code"] == ("none found", None)
    assert (got["Smallest"][0], got["Median"][0], got["Largest"][0]) == \
        (min(vals), statistics.median(vals), max(vals))
    assert blocks["ORIG_BAL"]["fmt"]["Median"] == "#,##0"                 # dollars too
    assert blocks["GCO_AMT"]["lines"]["Not a number"] == (1, 1 / 3000)   # the "#N/A"


def test_the_five_most_repeated_values_include_the_code(tmp_path):
    x = synth.write_extract(tmp_path, n=3000)
    blocks = _blocks(load_workbook(book.set_up(x).book)["Look"])
    counts = Counter(_numbers(_csv(x, "FICO")))
    shown = blocks["FICO"]["repeats"]
    assert shown[0] == (-9999, 60)
    assert len(shown) == 5 and all(counts[v] == k for v, k in shown)
    assert [k for _, k in shown] == sorted(counts.values(), reverse=True)[:5]
    # a balance with cents repeats nothing, and says so rather than listing five values seen once
    assert blocks["ORIG_BAL"]["repeats"] == [] and "No value repeats." in blocks["ORIG_BAL"]["lines"]
    gco = Counter(_numbers(_csv(x, "GCO_AMT")))
    assert blocks["GCO_AMT"]["repeats"][0] == (0, gco[0.0])


def test_a_run_split_by_revolving_debt_adds_a_scatter_against_each_band_column(tmp_path):
    b = _ready(tmp_path, n=3000)
    _set(b, "REV_DEBT", book.C_SPLIT, "Yes")
    raw, _, _ = book.read_book(b)
    bands = [x["field"] for x in raw["bands"]]
    assert bands == ["FICO", "ORIG_BAL"]
    assert book.run(b).ok
    wb = load_workbook(b)
    assert wb.sheetnames.index("Look") == wb.sheetnames.index("Columns") + 1        # the Run kept its place
    ws, hs = wb["Look"], wb["_look"]
    assert len([c for c in ws._charts if isinstance(c, BarChart)]) == len(NUMBER_COLUMNS)
    scatters = [c for c in ws._charts if isinstance(c, ScatterChart)]
    assert sorted(_title(c) for c in scatters) == [f"REV_DEBT against {x}" for x in bands]
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("random sample of 2,000" in t for t in text)                 # it says it's a sample
    x = b.with_name("loans.csv")
    rev = _csv(x, "REV_DEBT")
    for band in bands:
        # every dot is a real loan's (band, REV_DEBT), at most 2,000 of them, and never the -9999 code
        dots = _series(hs, band, first_col=2 * len(NUMBER_COLUMNS) + 1)
        loans = set(zip(_numbers_or_none(_csv(x, band)), _numbers_or_none(rev)))
        assert len(dots) == 2000 and all(d in loans for d in dots)
        assert all(d[0] != -9999 for d in dots)


def _numbers_or_none(raw):
    out = []
    for v in raw:
        try:
            out.append(float(v))
        except ValueError:
            out.append(None)
    return out


def test_the_sample_is_the_same_every_time(tmp_path):
    table = read_table(synth.write_extract(tmp_path, n=3000))
    got = []
    for _ in range(2):
        wb = Workbook()
        look.write_look(wb, table, NUMBER_COLUMNS, "REV_DEBT", ["FICO"])
        got.append(_series(wb["_look"], "FICO", first_col=2 * len(NUMBER_COLUMNS) + 1))
    assert got[0] == got[1] and len(got[0]) == look.SAMPLE


def test_set_up_again_rebuilds_the_tab_without_a_copy(tmp_path):
    x = synth.write_extract(tmp_path, n=2000)
    book.set_up(x)
    out = book.set_up(x)
    wb = load_workbook(out.book)
    assert [n for n in wb.sheetnames if "look" in n.lower()] == ["Look", "_look"]
    assert len(wb["Look"]._charts) == len(NUMBER_COLUMNS)
    assert wb.sheetnames.index("Look") == wb.sheetnames.index("Columns") + 1
