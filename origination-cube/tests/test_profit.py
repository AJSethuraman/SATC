"""RANR is profit after losses (NEXT-GOAL 3.1 to 3.5; rulings OC-29, OC-35).

The firm, 25 Sep 2026: "RANR is profit after losses - interest income + fees -
cost of funds - losses - and the cube's outputs, tests and synthetic book
should treat it that way." So profit is compared as a gap in points of booked
dollars, never a multiple (a rate near zero or below it has no stable
multiple); contribution before losses is RANR + GCO; the two are read together
with GCO; and every p-value is called one.

Tests 2 and 4 are the firm's own, from docs/for-test-design.md, built here.
Test 1's book is built from the write-up too, since Test 2 is an edit of it.
"""

import copy
import csv
import random
import re

import pytest
import yaml
from openpyxl import load_workbook

from conftest import TEST_SHUFFLES, cube, table
from origination_cube import book, control, engine, meanings, perm, stats, synth
from origination_cube import config as cfgmod
from origination_cube.ingest import read_table
from test_book import PICK, _answer
from test_book_results import _lvr
from recalc import calculated_book

# --------------------------------------------------------------------------
# Test 1's book (docs/for-test-design.md): 4,000 loans at $10,000, two score bands
# (600, 700) by two channels, 1,000 loans in each; 50 bad in each pocket but Broker
# 600, which has 150; $5,000 of GCO on each bad loan; one good 700 loan moved to 620.

BAND_EDGES = [620]
LOW, HIGH = "600 - 619", "620 - 700"


def _test1(ranr) -> list[dict]:
    rows, i = [], 0
    for score in (600, 700):
        for chan in ("Broker", "Branch"):
            bad = 150 if (score, chan) == (600, "Broker") else 50
            for k in range(1000):
                rows.append({"ID": f"T{i:05d}", "SCORE": score, "CHAN": chan, "BAL": 10_000, "BAD": int(k < bad),
                             "GCO": 5_000 if k < bad else 0, "RANR": ranr(score, chan)})
                i += 1
    rows[-1]["SCORE"] = 620                        # a good 700 loan, exactly on the edge
    return rows


def _test2_ranr(score, chan):
    """Test 2: the 600 band's RANR below zero, -$400 through Broker and -$200 through Branch."""
    return (-400 if chan == "Broker" else -200) if score == 600 else 200


def _cube(rows, **bench):
    b = {"min_units": 30, "min_events": 10, "worse_at": 1.25, "better_at": 0.8, "confidence": 0.95, "power": 0.8,
         "compare_to": "peers", "many_tests": "bh", "materiality": "none"}
    b.update(bench)
    return engine.run(cube(bands=[{"name": "score", "field": "SCORE", "edges": BAND_EDGES}], benchmark=b),
                      table(rows))


def test_test_1_by_hand():
    """The fixture is Test 1's book: every number the write-up works out by hand."""
    res = _cube(_test1(lambda s, c: 200))
    g = res.grids[0]
    assert res.total.rates["outcome_loans"].rate == pytest.approx(0.075)
    s = g.cell(LOW, "Broker").rates
    assert s["outcome_loans"].rate == pytest.approx(0.15) and s["outcome_loans"].vs_band == pytest.approx(3.0)
    assert s["outcome_loans"].excess == pytest.approx(75)
    assert s["gco_rate"].excess == pytest.approx(375_000)
    assert g.cell(HIGH, "Branch").rows == 1000                     # the 620 loan sits in "620 - 700"
    # RANR a flat $200 on every loan: about the same everywhere, a gap of exactly 0 points
    for _, c in g.inner():
        assert c.rates["ranr_rate"].vs_band == 0 and c.rates["ranr_rate"].reading_band == engine.IN_LINE


@pytest.mark.parametrize("line", ["luck", 0.5, "materiality"])
def test_test_2_a_band_below_zero_reads_the_right_way(line):
    """Test 2: RANR below zero in the 600 band. Broker 600 is the worse of two
    losing pockets: it keeps less, with a positive shortfall. Branch 600 is the
    better of the two and must not read keeps less. Before NEXT-GOAL 3.2 the
    multiple read Broker "better" (2.0x) and would read the same today, where
    one negative over another loses the direction."""
    res = _cube(_test1(_test2_ranr), revenue_line=line)
    g = res.grids[0]
    broker, branch = g.cell(LOW, "Broker").rates["ranr_rate"], g.cell(LOW, "Branch").rates["ranr_rate"]
    assert broker.vs_band == pytest.approx(-0.02) and branch.vs_band == pytest.approx(0.02)      # -2.00, +2.00 pts
    assert broker.reading_band == engine.WORSE and broker.flag == engine.WORSE
    assert branch.reading_band == engine.BETTER
    book_rate = res.total.rates["ranr_rate"].rate
    assert book_rate == pytest.approx(-200_000 / 40_000_000)
    assert broker.excess == pytest.approx(350_000) and broker.excess > 0          # the shortfall, in RANR dollars
    assert (-0.04) / (-0.02) == 2.0            # the multiple the flaw rested on: two negatives, "2.0x better"


def _workbook(tmp_path, rows, columns, meaning, answers, edges=None, cut_off=()):
    """Set up, answer Control and Columns as a person would, and run."""
    tmp_path.mkdir(parents=True, exist_ok=True)
    x = tmp_path / "book.csv"
    with x.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=columns)
        w.writeheader()
        w.writerows(rows)
    out = book.set_up(x)
    assert out.ok, out.lines
    wb = load_workbook(out.book)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value in answers:
            r[control.CHOOSE_COL - 1].value = answers[r[control.KEY_COL - 1].value]
    cat = meanings.catalog()
    ws = wb["Columns"]
    for r in ws.iter_rows(min_row=book.COL_FIRST):
        name = r[book.C_NAME - 1].value
        if name in meaning:
            r[book.C_MEANS - 1].value = cat[meaning[name]].label
            r[book.C_CUT - 1].value = "No" if name in cut_off else r[book.C_CUT - 1].value
            if edges and name in edges:
                r[book.C_EDGES - 1].value = edges[name]
    ws[book.CONFIRM_CELL] = "Yes"
    for r in wb["Odd values"].iter_rows(min_row=5):
        if r[1].value:
            r[4].value = "real"
    wb.save(out.book)
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    return calculated_book(out.book)          # the readings are formulas over Control's lines (OC-40)


TEST1_ANSWERS = {**PICK, "min_loans": "30 loans", "materiality": "No floor"}
TEST1_MEANING = {"ID": "key", "SCORE": "fico", "CHAN": "category", "BAL": "booked", "BAD": "outcome", "GCO": "gco",
                 "RANR": "ranr"}


@pytest.mark.parametrize("option", ["Each pocket's own test (suggested)", "0.5 points either way"])
def test_test_2_on_the_workbook(tmp_path, option):
    """Test 2 as the firm would run it, with the same Control answers as Test 1,
    then again with a fixed line: Broker 600 falls short of its band, in red,
    with a positive shortfall; Branch 600 is ahead of it. Since 26 Sep 2026 the
    readings say the gap literally, and the dollars are over the rest of its
    band, the comparison that decides; the book's are beside them."""
    wb = _workbook(tmp_path, _test1(_test2_ranr), list(TEST1_MEANING), TEST1_MEANING,
                   {**TEST1_ANSWERS, "revenue_line": option}, edges={"SCORE": "620"}, cut_off=("BAL",))
    rows = {(x["band"], x["seg"]): x for x in _lvr(wb["Losses vs revenue"])}
    broker, branch = rows[(LOW, "Broker")], rows[(LOW, "Branch")]
    assert broker["r_read"] == "short of its band by 2.00 points ($200,000)" and broker["r_fill"] == book.RED_CELL
    assert broker["r_x"] == pytest.approx(-2.0) and broker["r_over"] == pytest.approx(-200_000)
    assert broker["r_other"] == pytest.approx(-350_000)                         # against the book, for reference
    assert branch["r_read"] == "ahead of its band by 2.00 points ($200,000)" and branch["r_fill"] == book.GREEN_CELL
    # losing more, keeping less: a net drain; it pays more (RANR + GCO: 3.5% against 0.5%)
    assert (broker["g_read"], broker["together"]) == ("losing more", "net drain")
    assert broker["c_read"] == "ahead of its band by 3.00 points ($300,000)" and broker["c_x"] == pytest.approx(3.0)
    ws = wb["Where it bleeds"]
    # the main list: after its heading come the pockets losing more only against the other comparison (OC-40)
    end = next((r for r in range(5, ws.max_row + 1)
                if str(ws.cell(row=r, column=2).value).startswith("Losing more than their share against")),
               ws.max_row + 1)
    below = {(ws.cell(row=r, column=4).value, ws.cell(row=r, column=6).value) for r in range(end, ws.max_row + 1)
             if ws.cell(row=r, column=2).value == "Profit after losses: RANR per booked dollar"}
    assert (LOW, "Branch") in below                    # short of the book, ahead of its band: listed, below
    ranr = [[ws.cell(row=r, column=c).value for c in range(2, 21)] for r in range(5, end)
            if ws.cell(row=r, column=2).value == "Profit after losses: RANR per booked dollar"]
    got = {(x[2], x[4]): x for x in ranr}
    assert got[(LOW, "Broker")][8] == pytest.approx(200_000)                     # a positive shortfall, its band's
    assert got[(LOW, "Broker")][9] == pytest.approx(350_000)                     # and the book's beside it
    assert got[(LOW, "Broker")][14] == pytest.approx(-2.0)
    assert got[(LOW, "Broker")][16] == "short of its band by 2.00 points ($200,000)"
    assert (LOW, "Branch") not in got                                           # ahead of its band: no shortfall


# --------------------------------------------------------------------------
# Test 4 (docs/for-test-design.md): the trade-off pocket, on the synthetic book. Online, scores
# 620 to 679: bad loans roughly doubled, with GCO; a share of balance added to RANR, then the new
# GCO taken out, so RANR stays net of its losses.

def _test4_rows(uplift: float) -> list[dict]:
    rows = synth.make_rows(8000)
    rng = random.Random(4)
    pocket = [r for r in rows if r["CHANNEL"] == "Online" and isinstance(r["FICO"], int) and 620 <= r["FICO"] < 680
              and r["BAD_FLAG"] in (0, 1) and r["ORIG_BAL"] != "" and not isinstance(r["GCO_AMT"], str)]
    bad = sum(r["BAD_FLAG"] for r in pocket)
    for r in pocket:
        r["RANR_AMT"] = round(r["RANR_AMT"] + uplift * r["ORIG_BAL"], 2)
    for r in [r for r in pocket if r["BAD_FLAG"] == 0][:bad]:      # as many new bad loans as it had: doubled
        gco = round(r["ORIG_BAL"] * rng.uniform(0.3, 0.8), 2)
        r["BAD_FLAG"], r["GCO_AMT"], r["RANR_AMT"] = 1, gco, round(r["RANR_AMT"] - gco, 2)
    return rows


def _test4_engine(uplift, line):
    """The pocket's three rates against the rest of its band, on the synthetic
    book's own cube file (edges 620, 680, 740)."""
    raw = yaml.safe_load(synth.CONFIG)
    raw["benchmark"].update(revenue_line=line, shuffles=TEST_SHUFFLES)
    res = engine.run(cfgmod.parse(raw), table(_test4_rows(uplift), synth.COLUMNS))
    return res.grids[0].cell("620 - 679", "Online").rates


@pytest.mark.parametrize("line", ["luck", 0.25, "materiality"])
def test_test_4_with_the_write_ups_3_percent(line):
    """Test 4 as written: 3% of balance added. The honest result is that 3% does
    not pay for doubled losses in this book: GCO loses more (red), and profit is
    0.8 points below the rest of its band, which its own test doesn't call
    significant. So it isn't ahead of its band under any option; under a fixed
    line it reads short of it, not significant. The write-up expected it ahead; the
    synthetic book's RANR is years of interest, and doubled losses cost it
    about 3.8 points."""
    c = _test4_engine(0.03, line)
    assert c["gco_rate"].vs_band == pytest.approx(1.97, abs=0.01) and c["gco_rate"].reading_band == engine.WORSE
    assert c["ranr_rate"].vs_band == pytest.approx(-0.0079, abs=0.0001)
    assert c["ranr_rate"].p_band > 0.05
    assert c["ranr_rate"].reading_band == (engine.IN_LINE if line == "luck" else engine.UNSURE_WORSE)
    assert c["contribution_rate"].reading_band == engine.BETTER               # it does pay more


@pytest.mark.parametrize("line", ["luck", 0.25, "materiality"])
def test_test_4_with_a_bigger_uplift_keeps_more(line):
    """The same pocket with 8% of balance added: it pays for its losses and more,
    so it loses more and keeps more under every option."""
    c = _test4_engine(0.08, line)
    assert c["gco_rate"].reading_band == engine.WORSE
    assert c["ranr_rate"].vs_band > 0.03 and c["ranr_rate"].reading_band == engine.BETTER


@pytest.mark.parametrize("uplift, together, reading", [(0.03, "", r"-0\.\d\d points against its band, not significant"),
                                                      (0.08, "priced for it", r"ahead of its band by \d\.\d\d points "
                                                                              r"\(\$[\d,]+\)")])
def test_test_4_on_the_workbook(tmp_path, uplift, together, reading):
    """Test 4 through the workbook, on the suggested option: GCO red, losing more;
    at 3% profit's gap is given and called not significant, and nothing is read
    together; at 8% it is green, ahead of its band, and the pair reads priced
    for it."""
    rows = _test4_rows(uplift)
    d = tmp_path / f"t4-{uplift}"
    d.mkdir()
    x = d / "loans.csv"
    with x.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=synth.COLUMNS)
        w.writeheader()
        w.writerows(rows)
    out = book.set_up(x)
    _answer(out.book)
    wb = load_workbook(out.book)
    ws = wb["Columns"]
    for r in ws.iter_rows(min_row=book.COL_FIRST):
        if r[book.C_NAME - 1].value == "FICO":
            r[book.C_EDGES - 1].value = "620; 680; 740"
    wb.save(out.book)
    assert book.run(out.book).ok
    got = {(x["band"], x["seg"]): x for x in _lvr(calculated_book(out.book)["Losses vs revenue"])}
    p = got[("620 - 679", "Online")]
    assert p["g_read"] == "losing more" and p["g_fill"] == book.RED_CELL and p["g_over"] > 0
    assert re.fullmatch(reading, p["r_read"]), p["r_read"]
    assert (p["together"] or "") == together
    assert p["r_fill"] == (book.GREEN_CELL if uplift == 0.08 else None)
    if uplift == 0.08:
        assert p["r_over"] > 0                               # over-the-rest dollars positive on both sides


# --------------------------------------------------------------------------
# The synthetic book, as the walks use it: 8,000 loans, the LOB's edges, split by revolving debt


@pytest.fixture(scope="module")
def walk_book(tmp_path_factory):
    d = tmp_path_factory.mktemp("profit")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("CUBE_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
        out = book.set_up(synth.write_extract(d, n=8000))
        _answer(out.book)
        wb = load_workbook(out.book)
        for r in wb["Columns"].iter_rows(min_row=book.COL_FIRST):
            if r[book.C_NAME - 1].value == "FICO":
                r[book.C_EDGES - 1].value = "620; 680; 740"
            if r[book.C_NAME - 1].value == "REV_DEBT":
                r[book.C_SPLIT - 1].value = "Yes"
        wb.save(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return calculated_book(out.book)          # the readings are formulas over Control's lines (OC-40)


def test_the_priced_pocket_reads_priced_for_it(walk_book):
    """The plant (synth.py, NEXT-GOAL 3.5): Online, scores 680 to 739, goes bad
    twice as often as the rest of its band and carries 2 points more interest.
    It pays more, loses more and keeps more: priced for it. The loss plant under
    620 / Broker, priced like its band, is a net drain. Profit and contribution
    say their gap literally (26 Sep 2026)."""
    ws = walk_book["Losses vs revenue"]
    got = {(x["band"], x["seg"]): x for x in _lvr(ws) if x["row"] < 30}          # the first grid, FICO x CHANNEL
    p = got[("680 - 739", "Online")]
    assert (p["g_read"], p["together"]) == ("losing more", "priced for it")
    assert p["c_read"].startswith("ahead of its band by ") and p["r_read"].startswith("ahead of its band by ")
    assert (p["c_fill"], p["g_fill"], p["r_fill"]) == (book.GREEN_CELL, book.RED_CELL, book.GREEN_CELL)
    assert p["r_x"] > 0 and p["g_x"] > 1.25
    drain = got[("496 - 619", "Broker")]
    assert (drain["g_read"], drain["together"]) == ("losing more", "net drain")
    assert drain["r_read"].startswith("short of its band by ")
    # only the three pairs are ever named
    assert {x["together"] for x in _lvr(ws)} <= {None, "priced for it", "net drain", "safe but idle"}


def test_profit_is_a_gap_in_points_on_every_tab(walk_book):
    """NEXT-GOAL 3.2: every profit comparison is pocket - rest in points, never a
    multiple: Where it bleeds, Losses vs revenue, the Grids' heat maps, the Split
    tab's halves and its pooled figure."""
    ws = walk_book["Where it bleeds"]
    profit = [r for r in range(5, ws.max_row + 1)
              if ws.cell(row=r, column=2).value in ("Profit after losses: RANR per booked dollar",
                                                     "Contribution before losses per booked dollar")]
    assert profit
    for r in profit:
        for col in (14, 16):
            assert ws.formulas.cell(row=r, column=col).number_format == book.PTS_FMT
        assert ws.formulas.cell(row=r, column=19).number_format == '0.00" pts or less"'
    lvr = walk_book["Losses vs revenue"].formulas
    assert lvr.cell(row=7, column=book.LVR_R + 2).number_format == book.PTS_FMT
    assert lvr.cell(row=7, column=book.LVR_G + 2).number_format == '0.00"x"'
    split = walk_book["Split"]
    rows = {}
    for r in range(10, 40):                      # the first grid's summary, before its heat maps
        rows.setdefault(split.cell(row=r, column=2).value, r)
    r = rows["Profit after losses: RANR per booked dollar"]
    assert split.formulas.cell(row=r, column=5).number_format == book.PTS_FMT
    assert split.cell(row=r, column=6).value.endswith(" pts") and "x" not in split.cell(row=r, column=6).value
    assert split.cell(row=r, column=5).value < 0          # the high-debt half loses more, at the same price
    texts = [c.value for row in split.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(t.startswith("(") and t.endswith(" pts)") for t in texts)        # not significant: bracketed


def test_the_split_pools_profit_as_a_difference_even_when_the_low_halves_lose_money():
    """The audit, item a: with the low halves losing money in total, the pooled
    ratio O / E came out -2.15 with its range upside down, or vanished when E was
    at or below zero. In points it is (O - E) over the high halves' booked
    dollars, and it is always there."""
    rows, i = [], 0
    for score in (600, 700):
        for chan in ("A", "B"):
            for k in range(200):
                high = k % 2 == 0
                rows.append({"ID": f"L{i}", "SCORE": score, "CHAN": chan, "BAL": 1000, "BAD": 0, "GCO": 0,
                             "RANR": (-10 if high else -30) + (k % 7), "DEBT": 5000 + (1 if high else -1) * k})
                i += 1
    res = engine.run(cube(split={"field": "DEBT", "how": "own_median"},
                          benchmark={**_bench(), "shuffles": 200}), table(rows))
    g = res.grids[0]
    p = g.split_pooled["ranr_rate"]
    assert "ratio" not in p and p["pockets"] == 4
    # each pocket's halves: the high half's rate less the low half's, in points, never a multiple
    for (b, d), got in g.split_compare.items():
        hi_rate = g.split_cells[(b, d, engine.HIGH)].rates["ranr_rate"].rate
        lo_rate = g.split_cells[(b, d, engine.LOW)].rates["ranr_rate"].rate
        assert got["ranr_rate"][0] == pytest.approx(hi_rate - lo_rate, rel=1e-12) and lo_rate < 0
    hi = [r for r in rows if r["DEBT"] >= 5000]
    lo = [r for r in rows if r["DEBT"] < 5000]
    # equal halves of equal loans, so E is each low half's own total, and (O - E) / high dollars is:
    want = (sum(r["RANR"] for r in hi) - sum(r["RANR"] for r in lo)) / (1000 * len(hi))
    assert p["gap"] == pytest.approx(want, rel=1e-9) and p["gap"] > 0
    assert p["gap_lo"] < p["gap"] < p["gap_hi"]


def _bench():
    return {"min_units": 30, "min_events": 1, "worse_at": 1.25, "better_at": 0.8, "confidence": 0.95, "power": 0.8,
            "compare_to": "peers", "many_tests": "none", "materiality": "none"}


def test_two_negatives_and_near_zero_read_by_the_gap():
    """The audit, item a: a pocket 0.19 points ahead of a rest earning 0.01% read
    20x; one 0.10 short read "worse" on 0.02% against 0.12% and "in line" on 1.90%
    against 2.00% under one line. As points, the same gap reads the same way."""
    assert engine.gap_of(0.0020, 0.0001, True) == pytest.approx(0.0019)
    assert engine.gap_of(0.0020, 0.0001, False) == pytest.approx(20.0)             # the multiple it replaces
    line = engine.ProfitLine("points", 0.0025)
    for pocket, rest in ((0.0002, 0.0012), (0.0190, 0.0200)):
        assert engine.reading_gap(pocket - rest, 1e6, line, 0.001, 0.95) == engine.IN_LINE
    assert engine.reading_gap(-0.038 + 0.018, 1e6, line, 0.001, 0.95) == engine.WORSE   # -3.8% against -1.8%


def test_contribution_is_ranr_plus_gco():
    """NEXT-GOAL 3.4, OC-35: contribution before losses = RANR + GCO per booked
    dollar, higher is better, and a loan whose GCO can't be read is left out of
    it and counted against the GCO column."""
    rows = [{"ID": f"L{i}", "SCORE": 600 if i < 4 else 700, "CHAN": "A" if i % 2 else "B", "BAL": 1000,
             "BAD": 1 if i in (0, 5) else 0, "GCO": 400 if i in (0, 5) else 0, "RANR": 50 - (400 if i in (0, 5) else 0)}
            for i in range(8)]
    rows.append({"ID": "L8", "SCORE": 700, "CHAN": "A", "BAL": 1000, "BAD": 0, "GCO": "#N/A", "RANR": 50})
    res = engine.run(cube(), table(rows))
    m = next(x for x in res.measures if x.name == "contribution_rate")
    assert m.higher_is == "better" and m.in_points and m.plus == "GCO" and m.value == "RANR"
    assert m.title == "Contribution before losses per booked dollar"
    c = res.total.rates["contribution_rate"]
    assert c.num == pytest.approx(8 * 50) and c.den == 8000                   # every good loan earned 50
    assert res.left_out["contribution_rate"][("GCO", "not a number")] == 1
    assert res.total.rates["ranr_rate"].num == pytest.approx(9 * 50 - 800)
    ranr = next(x for x in res.measures if x.name == "ranr_rate")
    assert ranr.title == "Profit after losses: RANR per booked dollar"


def test_profit_is_material_at_the_loss_sides_dollar_line():
    """The audit, item a: RANR's materiality was a share of |total RANR|, which
    collapses when the book's profit is near zero. It is the GCO line now: the
    materiality answer applied to GCO, for a share or a dollar amount."""
    rows = []
    for i in range(400):
        bad = i % 10 == 0
        rows.append({"ID": f"L{i}", "SCORE": 600 if i < 200 else 700, "CHAN": "AB"[i % 2], "BAL": 1000,
                     "BAD": int(bad), "GCO": 500 if bad else 0, "RANR": (1 if i % 2 else -1) - (500 if bad else 0)
                     + 13})
    for mat in ("10% of losses", 3000):
        res = engine.run(cube(benchmark={**_bench(), "materiality": mat}), table(rows))
        want = 0.10 * 40 * 500 if mat == "10% of losses" else 3000
        assert res.materiality_line["gco_rate"] == pytest.approx(want)
        assert res.materiality_line["ranr_rate"] == pytest.approx(want)
        assert res.materiality_line["contribution_rate"] == pytest.approx(want)
        # and the Materiality tab's ladder is built on the same GCO dollars
        ladder = engine.materiality(res.grids[0], next(m for m in res.measures if m.name == "ranr_rate"), res.total)
        assert ladder[0].threshold == pytest.approx(0.005 * 40 * 500)


def test_the_smallest_profit_gap_is_in_points_and_ignores_the_rate():
    """The audit, item a: sized as a multiple of a book earning next to nothing,
    the smallest gap was 10,401x and the loans needed 944,772,979. As a
    difference it doesn't depend on the rate at all."""
    pairs = [(v, 1000.0) for v in [50.0, -50.0] * 10_000]                 # a book earning exactly nothing
    ln = stats.loans_needed_difference("ranr", pairs, 0.0025, 0.95, 0.8)
    s_d = ln.s_d
    assert ln.rate == 0 and s_d == pytest.approx(50.0, rel=1e-3)
    gap = stats.smallest_gap_for(ln, 100, 0.95, 0.8)
    z = stats.norm_s_inv(0.975) + stats.norm_s_inv(0.8)
    assert gap == pytest.approx(z * s_d / 1000 * (1 / 100 + 1 / 19_900) ** 0.5, rel=1e-9)
    # loans needed for a quarter point, turned round: that many loans can show about 0.25 points
    assert ln.loans and stats.smallest_gap_for(ln, ln.loans, 0.95, 0.8) <= 0.0025 < \
        stats.smallest_gap_for(ln, ln.loans - 1, 0.95, 0.8)


def test_the_words_are_p_value_and_not_significant(walk_book):
    """NEXT-GOAL 3.1: "Luck alone" is "p-value" on every tab, Control and Check;
    "could be luck" is "not significant"; nothing says "wobble" or calls profit
    revenue or earnings; standard error is defined on Check."""
    text = " ".join(str(c.value) for t in walk_book.sheetnames for row in walk_book[t].iter_rows() for c in row
                    if isinstance(c.value, str))
    low = text.lower()
    for gone in ("luck alone", "could be luck", "wobble", "earning", "earns"):
        assert gone not in low, gone
    assert "p-value" in text and "not significant" in text
    ws = walk_book["Where it bleeds"]
    assert [ws.cell(row=4, column=c).value for c in (15, 17)] == ["p-value", "p-value"]
    check = {r[1].value: r[2].value for r in walk_book["Check"].iter_rows(min_row=4)}
    assert check["Standard error"].startswith("How far a rate worked out from this many loans typically lands")
    assert "Two-sided" in check["p-value"]
    assert check["Contribution before losses"].startswith("RANR + GCO, per booked dollar. This assumes RANR has "
                                                         "gross charge-offs taken out.")
    options = [r[2] for r in walk_book["_options"].iter_rows(min_row=2, values_only=True)]
    assert not any("luck" in str(o).lower() for o in options)


def test_the_shuffle_count_is_in_the_test_column(walk_book):
    """NEXT-GOAL 3.6: a dollar rate's p-value made literal, as N of the shuffles
    that made a gap at least as big against the comparison that decides the flag."""
    ws = walk_book["Where it bleeds"]
    tests = [ws.cell(row=r, column=20).value for r in range(5, ws.max_row + 1)
             if ws.cell(row=r, column=2).value == "GCO per booked dollar"]
    assert tests and all(t.startswith("shuffled: ") and t.endswith(f" of {TEST_SHUFFLES:,}") for t in tests if t)
    first = tests[0]                                   # the planted pocket: no shuffle came close
    assert first == f"shuffled: 0 of {TEST_SHUFFLES:,}"


def test_the_shuffle_count_names_the_flags_comparison():
    """The count on the Test column is the one behind the flag: the rest of its
    band when that decides, the rest of the book otherwise."""
    s = engine.RateStat(test=engine.SHUFFLE_TEST, shuffles=10_000, hits_book=12, hits_band=3)
    assert book.which_test(s, peers=True) == "shuffled: 3 of 10,000"
    assert book.which_test(s, peers=False) == "shuffled: 12 of 10,000"
    assert book.which_test(engine.RateStat(test=engine.Z_TEST), peers=True) == "z test"


def test_old_profit_lines_are_refused():
    """The loss lines no longer lend profit a multiple (NEXT-GOAL 3.2): a cube file
    saying revenue_line: losses is refused with what to write instead, and so is
    a line of points past any real book."""
    with pytest.raises(cfgmod.ConfigError, match="losses is no longer an option"):
        cube(benchmark={**_bench(), "revenue_line": "losses"})
    with pytest.raises(cfgmod.ConfigError, match="number of points"):
        cube(benchmark={**_bench(), "revenue_line": 25})
    assert cube(benchmark={**_bench(), "revenue_line": 0.25}).benchmark.revenue_line == 0.25
    assert cube(benchmark={**_bench(), "revenue_line": "materiality"}).benchmark.revenue_line == "materiality"
