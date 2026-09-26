"""The engine knows no column, no value and no number by name.

The firm, 25 Sep 2026: "it's confusing to me that you talk in terms of specific
datapoints because it worries me we are not working with a generic engine".
Every walkthrough has used one synthetic book (FICO, CHANNEL, ASSET_CLASS,
REV_DEBT, a planted FICO-under-620 / Broker pocket), so every number quoted
back came from it. This file runs the whole route on a different book: other
column names, another product, other sizes, a problem planted somewhere else,
and checks the workbook finds that problem and says nothing about the other
book.
"""

import csv
import random
import re
from pathlib import Path

from openpyxl import load_workbook

from origination_cube import book, control, meanings, synth
from test_book import PICK
from test_book_results import _set

# an auto book, not a card book: none of these names are in the synthetic extract
COLUMNS = ["AcctId", "BureauScore", "Dealer", "FinancedAmt", "ChargedOff", "NetLossDollars",
           "NetRevenueDollars", "Region", "PTI"]
DEALERS = ["North Motors", "Lakeside Auto", "Fleet Direct", "Metro Cars"]
REGIONS = ["Coast", "Valley", "Hills"]


def _auto_book(d: Path, n: int = 7000, seed: int = 11) -> Path:
    """The planted problem: Fleet Direct in the Hills defaults about four times as
    often as the rest, at every score. The revenue column is profit after losses,
    as RANR is (OC-29, OC-35): what the loan paid, less the whole loss. It runs
    lower at low scores and is negative for most loans that charged off, so a
    comparison below zero is exercised too."""
    rnd = random.Random(seed)
    rows = []
    for i in range(n):
        score = int(min(820, max(540, rnd.gauss(690, 55))))
        dealer, region = rnd.choice(DEALERS), rnd.choice(REGIONS)
        amt = round(rnd.uniform(8_000, 45_000), 2)
        p = 0.02 + max(0, 700 - score) / 900
        if dealer == "Fleet Direct" and region == "Hills":
            p *= 4
        bad = rnd.random() < min(p, 0.9)
        loss = round(amt * rnd.uniform(0.3, 0.7), 2) if bad else 0.0
        paid = amt * (0.04 + (score - 640) / 4000)          # contribution before losses
        rev = round(paid - loss, 2)                          # the whole loss comes out; it took out half until 26 Sep 2026
        rows.append({"AcctId": f"A{i:06d}", "BureauScore": score, "Dealer": dealer, "FinancedAmt": amt,
                     "ChargedOff": "Y" if bad else "N", "NetLossDollars": loss, "NetRevenueDollars": rev,
                     "Region": region, "PTI": round(rnd.uniform(0.04, 0.2), 3)})
    d.mkdir(parents=True, exist_ok=True)
    out = d / "auto book.csv"
    with out.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    return out


def _answer_as_a_person_would(b: Path) -> None:
    """Control from the same picks as every other test; Columns by meaning, as a
    person reads them, for the columns whose name gives nothing away."""
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        k = r[control.KEY_COL - 1].value
        if k in PICK:
            r[control.CHOOSE_COL - 1].value = PICK[k]
    for r in wb["Odd values"].iter_rows(min_row=5):
        if r[1].value and not r[4].value:
            r[4].value = "real"                     # revenue below zero is real here
    wb.save(b)
    cat = meanings.catalog()
    for col, code in {"AcctId": "key", "BureauScore": "fico", "Dealer": "category", "FinancedAmt": "booked",
                      "ChargedOff": "outcome", "NetLossDollars": "gco", "NetRevenueDollars": "ranr",
                      "Region": "category", "PTI": "dti"}.items():
        _set(b, col, book.C_MEANS, cat[code].label)
    _set(b, "ChargedOff", book.C_IS, "Y")
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)


def test_another_book_with_other_names_finds_its_own_problem(tmp_path):
    extract = _auto_book(tmp_path)
    out = book.set_up(extract)
    assert out.ok, out.lines
    _answer_as_a_person_would(out.book)
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    wb = load_workbook(out.book)

    # the grids are this book's columns
    grids = [c.value for c in wb["Grids"]["B"] if isinstance(c.value, str) and " x " in c.value]
    assert grids and all(re.match(r"(BureauScore|FinancedAmt|PTI) x (Dealer|Region)", g) for g in grids), grids

    # the planted pocket is found: Fleet Direct is the worst Dealer pocket for the outcome, in every
    # BureauScore band that was tested
    ws = wb["Where it bleeds"]
    rows = [[ws.cell(row=r, column=c).value for c in range(2, 19)] for r in range(5, ws.max_row + 1)]
    flagged = [x for x in rows if x[0] == "Outcome, share of loans" and x[3] == "Dealer" and x[15] == "worse"]
    assert flagged and {x[4] for x in flagged} == {"Fleet Direct"}, [(x[2], x[4], x[15]) for x in flagged]
    assert any(x[3] == "Region" and x[4] == "Hills" and x[15] == "worse" for x in rows)

    # nothing from the other book leaks in: no column, value, or pocket of the synthetic extract
    text = " ".join(str(c.value) for t in wb.sheetnames for row in wb[t].iter_rows() for c in row
                    if isinstance(c.value, str))
    # "FICO score" and "the FICO scale" stay: they're a meaning, chosen here for BureauScore, not a column.
    # Control's explanations used to carry the other book's pocket as their example ("Broker under 620
    # against the rest of under 620"); found by this test
    words = set(re.findall(r"[A-Za-z_]+", text.replace("FICO score", "").replace("FICO scale", "")))
    for name in synth.COLUMNS + synth.CHANNELS:
        assert name not in words, name
    assert "496 - 619" not in text and "under 620" not in text
