"""A test of a new variable builds no bleed analysis (OC-42; the firm, 26 Sep 2026, on BACKLOG 6d: "Yes. Seems
obvious I think. They have entirely different outputs generally").

A Run from a pre-spec runs the confirmatory test and what it needs: the pre-spec's strata cut into the bands the
Columns tab gives them, the development and holdout ranges, and 4e's concentration; then Check and the Log. It
builds no pocket grid, no three-way or split grid, and no shuffle test, and it writes none of the bleed's tabs. A
workbook that held them from an earlier bleed Run has them taken off, and Check says so in one line, so no tab
shows an old Run's pockets beside this Run's test.

Start here's "What the last Run found" and the launcher's last step read the confirmation, not nought pockets of
nought. A bleed Run is unchanged: it still builds its grids and runs one shuffle test for all of them.

Everything counted here is counted from the test's own groups by hand, never from the tiles' numbers."""

import pytest
from openpyxl import load_workbook

from origination_cube import book, confirm_tab, confirmatory, engine, launcher, perm
from recalc import recalc
from test_book_dates import _check, _control
from test_confirmatory import _held, _log, _ready, _spec, needs_git

BLEED, NEW = "Where the book bleeds", "Finding and testing a new variable"
FROM_SPEC = "Test from a pre-spec"
BLEED_TABS = {"Where it bleeds", "Losses vs revenue", "Grids", "Split", "Three-way", "Prevalence"}


def _counting(mp):
    """How many grids the engine builds and how many shuffle tests it runs, from here on."""
    seen = {"grids": 0, "shuffles": 0}
    build, shuffle = engine._build_grid, perm.run

    def built(*a, **k):
        seen["grids"] += 1
        return build(*a, **k)

    def shuffled(*a, **k):
        seen["shuffles"] += 1
        return shuffle(*a, **k)

    mp.setattr(engine, "_build_grid", built)
    mp.setattr(perm, "run", shuffled)
    return seen


@pytest.fixture(scope="module")
def runs(tmp_path_factory):
    """One workbook: a bleed Run first, so it holds the bleed's tabs, then a test of a new variable from a
    committed pre-spec on the same extract."""
    folder = tmp_path_factory.mktemp("new-variable")
    out = {}
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("CUBE_MEMORY", str(folder / "memory.yaml"))
        x, b = _ready(folder, mp, n=20000)
        _spec(x.parent)
        _control(b, run_kind=BLEED)
        with pytest.MonkeyPatch.context() as spy:
            out["bleed_counts"] = _counting(spy)
            out["bleed"] = book.run(b)
        out["bleed_tabs"] = set(load_workbook(b).sheetnames)
        _control(b, run_kind=NEW, new_variable_step=FROM_SPEC)
        _held(b, "prespec.yaml")
        seen = {}
        real = confirmatory.state

        def state(*a, **k):
            seen["st"] = real(*a, **k)
            return seen["st"]

        with pytest.MonkeyPatch.context() as spy:
            out["counts"] = _counting(spy)
            spy.setattr(confirmatory, "state", state)
            out["ran"] = book.run(b)
    out.update(x=x, b=b, st=seen["st"], t=seen["st"].test)
    assert out["bleed"].ok, out["bleed"].lines
    assert out["ran"].ok, out["ran"].lines
    return out


def _worse_by_hand(t, confidence: float = 0.95) -> list[int]:
    """The groups going bad significantly more often than the reference on the holdout."""
    f = t.holdout.fit
    return [k for k in range(len(t.groups)) if k != t.ref and f.odds[k] is not None and f.odds[k] > 1
            and f.p[k] is not None and f.p[k] < round(1 - confidence, 12)]


# --------------------------------------------------------------------------
# What a test of a new variable builds, and doesn't


@needs_git
def test_a_new_variable_run_builds_no_grid_and_runs_no_shuffle_test(runs):
    assert runs["counts"] == {"grids": 0, "shuffles": 0}
    # the counter can see them: the bleed Run on the same workbook built its grids and ran one shuffle test
    assert runs["bleed_counts"]["grids"] > 0 and runs["bleed_counts"]["shuffles"] == 1


@needs_git
def test_a_new_variable_run_writes_no_bleed_tab_and_takes_off_the_ones_a_bleed_run_left(runs):
    assert BLEED_TABS <= runs["bleed_tabs"]
    tabs = set(load_workbook(runs["b"]).sheetnames)
    assert not tabs & BLEED_TABS, tabs & BLEED_TABS
    assert {confirm_tab.SHEET, "Check", "Log", "Start here"} <= tabs


@needs_git
def test_check_says_so_in_one_line_and_says_nothing_of_pockets_or_tie_outs(runs):
    chk = _check(runs["b"])
    assert chk[book.NO_BLEED[0]] == book.NO_BLEED[1]
    for label in ("Tie-out checks", "Pockets tested", "Pocket budget", "Pockets in all", "Families of tests",
                  "Tests", "Decides each pocket", "Split"):
        assert label not in chk, label
    assert not [k for k in chk if str(k).startswith(("Worse now: ", "Pockets: ", "Loans needed for "))]
    # what the confirmatory test rests on is still there: the strata's band edges, and the pre-spec's lines
    assert "Band edges used: FICO" in chk and chk["Differs from the pre-spec"] == "nowhere: this run used what it says"
    assert "tie-out" not in _log(runs["b"])[0] and "tie-out" not in runs["ran"].lines[0]
    assert runs["ran"].lines[-1].endswith("start with Confirmatory test.")


# --------------------------------------------------------------------------
# What the last Run found, read from the confirmation


@needs_git
def test_the_launchers_last_step_reads_the_confirmation(runs):
    t, h = runs["t"], runs["ran"].summary
    worse = _worse_by_hand(t)
    assert worse, "the synthetic book plants two cliffs; the test should find at least one"
    assert h["kind"] == "confirm" and h["reference"] == t.groups[t.ref] and h["column"] == t.column
    assert (h["worse"], h["groups"]) == (len(worse), len(t.groups) - 1)
    assert h["capture"] == pytest.approx(sum(t.holdout.bad[k] for k in worse) / sum(t.holdout.bad), rel=1e-12)
    assert (h["development"], h["holdout"], h["deviations"]) == (t.development.n, t.holdout.n, 0)
    tiles = launcher.confirm_tiles(h)
    assert tiles[0][:2] == (f"Groups worse than {t.groups[t.ref]}", f"{len(worse)} of {len(t.groups) - 1}")
    assert tiles[1][1] == f"{h['capture']:.0%}" and tiles[2][1] == "Yes"
    assert "Pockets worse and material" not in [x[0] for x in tiles]


def test_the_launcher_says_why_when_the_test_could_not_run():
    tiles = launcher.confirm_tiles({"kind": "confirm", "problem": "no column is marked Origination date"})
    assert tiles == (("Confirmatory test", "Not run", "no column is marked Origination date", "KEY_RED", "INK"),)


@needs_git
def test_start_here_shows_each_group_on_the_holdout_and_follows_the_confidence(runs, tmp_path):
    t = runs["t"]
    worse = _worse_by_hand(t)
    ws = recalc(runs["b"], tmp_path / "calc")["Start here"]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    at = {v: k for k, v in cells.items() if isinstance(v, str)}
    label = f"Groups worse than {t.groups[t.ref]}, on the holdout"
    r, c = at[label]
    assert cells[(r + 1, c)] == f"{len(worse)} of {len(t.groups) - 1}"
    share = cells[at["Their share of the holdout's bad loans"][0] + 1, at["Their share of the holdout's bad loans"][1]]
    assert share == pytest.approx(sum(t.holdout.bad[k] for k in worse) / sum(t.holdout.bad), rel=1e-9)
    assert cells[at["Follows the pre-spec"][0] + 1, at["Follows the pre-spec"][1]] == "Yes"
    # one row per group, the reference marked, each worse group saying so
    head = at[f"Group of {t.column}"][0]
    rows = {cells.get((head + 1 + k, 2)): head + 1 + k for k in range(len(t.groups))}
    assert f"{t.groups[t.ref]} (reference)" in rows
    for k, name in enumerate(t.groups):
        rr = rows[name + (" (reference)" if k == t.ref else "")]
        assert cells[(rr, 3)] == t.holdout.loans[k]
        if k != t.ref:
            assert cells[(rr, 7)] == ("Yes, worse" if k in worse else cells[(rr, 7)])
            assert (cells[(rr, 7)] == "Yes, worse") == (k in worse)
    assert "Nothing yet" not in " ".join(v for v in cells.values() if isinstance(v, str))

    # at 99% sure, a group whose p-value sits between 1% and 5% stops counting (OC-40: the lines are live)
    wb = load_workbook(runs["b"])
    from origination_cube import control
    cws = wb[control.SHEET]
    cws.cell(row=control.row_of(cws, "confidence"), column=control.CHOOSE_COL).value = "99%"
    changed = tmp_path / "changed.xlsx"
    wb.save(changed)
    ws = recalc(changed, tmp_path / "calc99")["Start here"]
    got = ws.cell(row=r + 1, column=c).value
    assert len(_worse_by_hand(t, 0.99)) != len(worse)       # this book has a group between the two bars
    assert got == f"{len(_worse_by_hand(t, 0.99))} of {len(t.groups) - 1}"


# --------------------------------------------------------------------------
# The engine, directly


def test_the_engine_builds_no_grid_for_a_new_variable_and_the_same_grids_as_before_for_the_bleed(book):
    from conftest import cube, table
    tb = table(book)
    bled = engine.run(cube(), tb)
    tested = engine.run(cube(run_kind="new_variable"), tb)
    assert bled.bleed and bled.grids and bled.tie_outs > 0
    assert not tested.bleed and tested.grids == [] and tested.three_way == [] and tested.tie_outs == 0
    # what the confirmatory test reads is the same either way: the bands' edges, the loans, the book's totals
    assert tested.band_edges == bled.band_edges and tested.rows == bled.rows
    assert {k: (v.num, v.den) for k, v in tested.total.rates.items()} == \
        {k: (v.num, v.den) for k, v in bled.total.rates.items()}
