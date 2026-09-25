"""`cube init`: every column is classified from its own values, only the clear
cases are decided, and nothing the professional must judge is filled in."""

import re

import pytest

from conftest import cube, row, table
from origination_cube import config as cfgmod
from origination_cube import control, engine, profile, synth
from origination_cube.ingest import read_table


def roles(cols):
    return {c.name: c.role for c in cols}


def test_the_synthetic_extract_is_read_the_way_a_person_would(tmp_path):
    _, data = synth.write(tmp_path, n=5000)
    cols = profile.classify(read_table(data), few_values=12, many_values=50)
    assert roles(cols) == {"LOAN_NBR": "key", "FICO": "band", "CHANNEL": "dimension", "ORIG_BAL": "band",
                           "BAD_FLAG": "dimension", "GCO_AMT": "band", "RANR_AMT": "band"}
    gco = next(c for c in cols if c.name == "GCO_AMT")
    assert "not a number" in gco.why          # the one "#N/A" is noted, not a reason to doubt the column


def test_each_kind_of_column(tmp_path):
    rows = [{"ID": f"A{i}", "ZIP": f"0{1000 + i % 7}", "TERM": [36, 48, 60][i % 3], "ONE": "x", "EMPTY": "",
             "CITY": f"city{i}" if i % 2 else "same", "WHEN": f"2024-01-{1 + i % 28:02d}", "SCORE": 500 + i % 90,
             "APP": 4000000 + i, "SEQ": i + 1}
            for i in range(200)]
    got = roles(profile.classify(table(rows), few_values=12, many_values=50))
    assert got == {"ID": "key", "ZIP": "dimension", "TERM": "dimension", "ONE": "skipped", "EMPTY": "skipped",
                   "CITY": "question", "WHEN": "date", "SCORE": "band",
                   "APP": "key",          # 7 digits on every row
                   "SEQ": "question"}     # 1..200: a loan number or an amount? never cut by until said


def test_odd_values_are_asked_about_only_where_they_make_sense(tmp_path):
    _, data = synth.write(tmp_path, n=5000)
    qs = [q for c in profile.classify(read_table(data), 12, 50) for q in c.questions]
    found = {(q["column"], q["pattern"]) for q in qs}
    assert ("FICO", "repeated_value") in found           # the bureau's -9999
    assert ("RANR_AMT", "negatives") in found            # real, but the tool must ask, not decide
    assert not any(q["column"] == "BAD_FLAG" for q in qs)  # a 0 in a 0/1 flag is not a code
    assert not any(q["column"] == "GCO_AMT" for q in qs)   # a spike of zeros is losses that did not happen


def test_the_required_columns_are_suggested_with_reasons(tmp_path):
    _, data = synth.write(tmp_path, n=4000)
    t = read_table(data)
    sug = profile.suggest_roles(t, profile.classify(t, 12, 50), control.role_hints())
    assert {r: sg.column for r, sg in sug.items()} == {"key": "LOAN_NBR", "booked": "ORIG_BAL",
                                                       "outcome": "BAD_FLAG", "gco": "GCO_AMT", "ranr": "RANR_AMT"}
    assert all(sg.why for sg in sug.values())


def test_nothing_is_used_until_a_person_confirms_it(tmp_path):
    _, data = synth.write(tmp_path, n=2000)
    path, _ = profile.write_cube_file(read_table(data), tmp_path / "cube.yaml")
    with pytest.raises(cfgmod.ConfigError) as exc:
        cfgmod.load(path)
    text = str(exc.value)
    assert "not confirmed yet" in text and "columns_confirmed: yes" in text
    for must in ("min_units", "worse_at", "better_at", "confidence"):     # judgment: never suggested
        assert f"`benchmark.{must}` still reads" in text, must


def test_with_no_name_to_go_on_it_asks_and_lists_candidates(tmp_path):
    rows = [{"A": f"x{i}", "B": 1000 + i % 37, "C": i % 2, "D": 0 if i % 3 else 55.5, "E": (i % 11) - 3}
            for i in range(300)]
    t = table(rows)
    sug = profile.suggest_roles(t, profile.classify(t, 12, 50), control.role_hints())
    assert sug["key"].column == "A"                         # the only column different on every row
    assert sug["outcome"].column == "C"                     # the only 0/1 column
    assert sug["ranr"].column is None and sug["ranr"].candidates   # nothing named RANR: ask, with candidates


def _answer(path, **judgment):
    s = path.read_text()
    s = s.replace("columns_confirmed: no", "columns_confirmed: yes")
    for k, v in judgment.items():
        s = re.sub(rf'{k}: "\[CONFIRM[^"]*"', f"{k}: {v}", s)
    path.write_text(s)


def test_once_confirmed_it_runs_and_the_outcome_is_not_cut_by(tmp_path):
    _, data = synth.write(tmp_path, n=4000)
    path, _ = profile.write_cube_file(read_table(data), tmp_path / "cube.yaml")
    _answer(path, min_units=30, worse_at=1.25, better_at=0.8, confidence=0.95)
    res = engine.run(cfgmod.load(path), read_table(data))
    cut_by = {g.band for g in res.grids} | {g.dimension for g in res.grids}
    assert not cut_by & {"gco_amt", "ranr_amt", "bad_flag"}
    assert any("GCO_AMT` is not cut by" in w for w in res.warnings)
    assert any("open data question" in w for w in res.warnings)


def test_a_wrong_suggestion_is_fixed_by_typing_over_it(tmp_path):
    _, data = synth.write(tmp_path, n=2000)
    path, _ = profile.write_cube_file(read_table(data), tmp_path / "cube.yaml")
    _answer(path, min_units=30, worse_at=1.25, better_at=0.8, confidence=0.95)
    path.write_text(re.sub(r"^booked: ORIG_BAL", "booked: GCO_AMT", path.read_text(), flags=re.M))
    cfg = cfgmod.load(path)
    assert cfg.booked == "GCO_AMT"


def test_a_filled_control_tab_fills_the_judgment(tmp_path):
    _, data = synth.write(tmp_path, n=2000)
    book = control.build_control_book(tmp_path / "control.xlsx")
    from openpyxl import load_workbook
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    for r in ws.iter_rows(min_row=control.FIRST_ROW):
        key = r[control.KEY_COL - 1].value
        s = next((x for x in control.load_settings() if x.key == key), None)
        if s and s.judgment:
            r[control.CHOOSE_COL - 1].value = s.options[1].shown
    wb.save(book)
    path, _ = profile.write_cube_file(read_table(data), tmp_path / "cube.yaml", control.read_control(book))
    text = path.read_text()
    assert "min_units: 100" in text and "worse_at: 1.5" in text and "confidence: 0.95" in text
    assert "[CONFIRM: fewest loans" not in text


def test_a_band_on_the_outcome_is_dropped_and_said_so(book):
    res = engine.run(cube(bands=[{"name": "score", "field": "SCORE", "edges": [650]},
                                 {"name": "g", "field": "GCO", "edges": [10]}]), table(book))
    assert [g.band for g in res.grids] == ["score"]
    assert any("`GCO` is not cut by" in w for w in res.warnings)
