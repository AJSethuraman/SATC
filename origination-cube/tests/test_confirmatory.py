"""A run held to a pre-spec (NEXT-GOAL 3.15): the file is named on Control,
read and refused by that cell, echoed on Check with the git commit it came
from, compared with what the run used line by line, and every run that touched
the holdout is marked in the Log and counted on Check. The holdout's loans are
counted here by hand from the extract."""

import csv
import shutil
import subprocess
from datetime import date, timedelta

import pytest
import yaml
from openpyxl import load_workbook

from conftest import cube, table
from origination_cube import book, confirmatory, control, engine, prespec, prevalence, synth
from test_book import _answer
from test_book_dates import _check, _columns, _control

needs_git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed on this machine")

SPEC = {"prespec": 1, "written": "2026-09-26", "column": "INCOME_TO_SALES", "bins": [0.1, 0.25, 0.5, 1.0, 2.0],
        "reference": "0.25 - 0.49", "strata": ["FICO", "CHANNEL"], "confidence": 0.95,
        "holdout": {"from": "2024-01-01", "to": "2024-12-31"},
        "development": {"from": "2022-01-01", "to": "2023-12-31"}}


def git(repo, *args):
    return subprocess.run(["git", "-c", "user.name=Test", "-c", "user.email=test@example.com",
                           "-c", "commit.gpgsign=false", *args], cwd=repo, check=True, capture_output=True,
                          text=True).stdout.strip()


def _spec(folder, commit=True, **changes):
    """The pre-spec, beside the workbook, committed in a repository of its own."""
    f = folder / "prespec.yaml"
    f.write_text(yaml.safe_dump({**SPEC, **changes}, sort_keys=False), encoding="utf-8")
    if commit:
        if not (folder / ".git").exists():
            git(folder, "init", "-q")
        git(folder, "add", f.name)
        git(folder, "commit", "-q", "-m", "pre-spec")
    return f


def _held(b, path):
    wb = load_workbook(b)
    ws = wb[control.SHEET]
    row = control.row_of(ws, control.PRESPEC_KEY)
    ws.cell(row=row, column=control.CHOOSE_COL).value = path
    wb.save(b)
    return f"Control!C{row}"


def _ready(tmp_path, monkeypatch, n=3000):
    """The synthetic book with INCOME and SALES, INCOME / SALES made on Control and splitting the pockets,
    cut only by FICO and CHANNEL: every setting the pre-spec says except the ones it can't."""
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    x = synth.write_extract(tmp_path / "x", n=n, ratio=True)
    b = book.set_up(x).book
    _answer(b)
    _control(b, run_kind="Finding and testing a new variable", new_variable_step="Test from a pre-spec",
             **{"derived|1": ("INCOME_TO_SALES", "INCOME", "SALES")})
    book.set_up(x)
    _columns(b, "INCOME_TO_SALES", C_SPLIT="Yes", C_EDGES="0.1; 0.25; 0.5; 1; 2")
    for c in ("ORIG_BAL", "REV_DEBT", "ASSET_CLASS", "INCOME", "SALES"):
        _columns(b, c, C_CUT="No")
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    return x, b


def _log(b) -> list[str]:
    ws = load_workbook(b)["Log"]
    return [ws.cell(row=r, column=2).value for r in range(book.LOG_FIRST, ws.max_row + 1)
            if ws.cell(row=r, column=2).value]


def _warnings(chk) -> list[str]:
    w = chk.get("Warning", [])
    return w if isinstance(w, list) else [w]


@needs_git
def test_a_run_held_to_a_committed_pre_spec_echoes_it_says_where_it_differs_and_counts_the_holdout(
        tmp_path, monkeypatch):
    x, b = _ready(tmp_path, monkeypatch)
    f = _spec(x.parent)
    head = git(x.parent, "rev-parse", "HEAD")
    _held(b, "prespec.yaml")                                    # its name alone: it sits beside the workbook
    ran = book.run(b)
    assert ran.ok, ran.lines
    chk = _check(b)
    assert chk["Pre-spec"] == str(f.resolve())
    assert chk["Pre-spec commit"].startswith(f"{head[:12]}, committed ")
    says = chk["What the pre-spec says"].splitlines()
    assert "column: INCOME_TO_SALES" in says and "holdout: 2024-01-01 to 2024-12-31" in says
    assert "reference: 0.25 - 0.49" in says and "strata: FICO, CHANNEL" in says
    assert not any("window" in s for s in says)

    # what the run could follow, it did: the column, its bins (typed on Columns), the strata and the
    # confidence. What it couldn't: its reference is each pocket's low half, and its loans, every one in the
    # extract, run from 2021, not only the holdout's
    with open(x, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    kept = sorted(r["ORIG_DATE"] for r in rows)
    assert chk["Loans run"] == "3,000"
    assert chk["Origination dates"] == f"{kept[0]} to {kept[-1]} (3,000 loans; 0 without a readable date)"
    devs = [w for w in _warnings(chk) if w.startswith(confirmatory.DEVIATES)]
    assert devs == [
        'Deviates from pre-spec: The reference group is "the low half of each pocket", which is not one of its '
        'bins\' groups in this run; the pre-spec says "0.25 - 0.49".',
        f"Deviates from pre-spec: The pre-spec's holdout is 2024-01-01 to 2024-12-31; this run used loans made "
        f"{kept[0]} to {kept[-1]}, not only the holdout."]

    # the holdout, counted by hand: every loan in the extract made in 2024, both ends included
    held = sorted(r["ORIG_DATE"] for r in rows if "2024-01-01" <= r["ORIG_DATE"] <= "2024-12-31")
    assert held
    assert chk["Holdout"] == (f"{len(held):,} of this extract's loans were made in the holdout (2024-01-01 to "
                              f"2024-12-31), the first on {held[0]} and the last on {held[-1]}. This run touched "
                              f"the holdout.")
    assert chk["Runs that touched the holdout"] == "1 in this workbook's Log, this one included"
    log = _log(b)
    assert log[1] == f"Deviates from pre-spec prespec.yaml (commit {head[:12]}): 2 places, listed on Check."
    assert log[2] == f"Touched the holdout: {len(held):,} loans made {held[0]} to {held[-1]}."
    assert log[1] in ran.lines and "Runs that touched the holdout, in this workbook's Log: 1." in ran.lines
    ran_yaml = b.with_name(f"{b.stem} - what ran.yaml").read_text(encoding="utf-8")
    assert f"# pre-spec: {f.resolve()} ({head[:12]}, committed " in ran_yaml
    assert f"# Touched the holdout: {len(held):,} loans made" in ran_yaml

    # Control strays from the pre-spec in one more place: said, and counted; the holdout count goes up
    _control(b, confidence="90%")
    assert book.run(b).ok
    chk = _check(b)
    devs = [w for w in _warnings(chk) if w.startswith(confirmatory.DEVIATES)]
    assert "Deviates from pre-spec: Confidence is 90% in this run; the pre-spec says 95%." in devs and len(devs) == 3
    assert chk["Runs that touched the holdout"] == "2 in this workbook's Log, this one included"
    assert _log(b)[1] == f"Deviates from pre-spec prespec.yaml (commit {head[:12]}): 3 places, listed on Check."

    # the pre-spec edited after its commit: said on Check and in the Log, and the run doesn't count
    f.write_text(f.read_text(encoding="utf-8") + "# an afterthought\n", encoding="utf-8")
    assert book.run(b).ok
    chk = _check(b)
    assert chk["Pre-spec commit"].endswith("; edited since that commit, so the file read here isn't the one "
                                           "committed")
    assert ("This run doesn't count as the pre-specified one until the pre-spec is committed, unchanged."
            in _warnings(chk))
    assert _log(b)[1].startswith(f"Deviates from pre-spec prespec.yaml (commit {head[:12]}, edited since)")
    assert chk["Runs that touched the holdout"] == "3 in this workbook's Log, this one included"


@needs_git
def test_a_pre_spec_the_run_cannot_use_is_refused_by_its_cell(tmp_path, monkeypatch):
    x, b = _ready(tmp_path, monkeypatch, n=1500)
    cell = _held(b, str(tmp_path / "nowhere.yaml"))
    ran = book.run(b)
    assert not ran.ok
    assert any(f"{cell}: there's no pre-spec file at {tmp_path / 'nowhere.yaml'}. Fix the path, or clear the cell"
               in line for line in ran.lines), ran.lines
    _spec(x.parent, bins=[0.5, 0.1], confidence=95)
    _held(b, f'"{x.parent / "prespec.yaml"}"')                # pasted with the quotes Windows adds
    ran = book.run(b)
    assert not ran.ok
    said = [line for line in ran.lines if cell in line]
    assert any('in the pre-spec prespec.yaml, "bins:" must rise' in s for s in said), said
    assert any('in the pre-spec prespec.yaml, "confidence:" must be a share between 0.5 and 1' in s for s in said)
    assert all("`" not in s for s in said)
    assert "Couldn't run" in _log(b)[0] and any(cell in line for line in _log(b))
    # Set up again keeps the path
    book.set_up(x)
    ws = load_workbook(b)[control.SHEET]
    assert control.read_prespec(ws) == (str(x.parent / "prespec.yaml"), cell)
    # a blank cell: refused while testing from a pre-spec (tests/test_run_kind.py), and a run with no pre-spec
    # once the run is Where the book bleeds
    _held(b, None)
    assert not book.run(b).ok
    _control(b, run_kind="Where the book bleeds")
    assert book.run(b).ok
    chk = _check(b)
    assert "Pre-spec" not in chk and "Holdout" not in chk
    assert not any(line.startswith(("Deviates", "Follows", confirmatory.HOLDOUT_MARK)) for line in _log(b))


@needs_git
def test_dates_that_cannot_be_read_leave_the_holdout_said_unchecked(tmp_path, monkeypatch):
    """Testing a new variable needs a column marked Origination date (tests/test_run_kind.py), so the holdout
    goes unchecked only when that column's dates can't be read: here every one reads two ways (01/02/2024,
    January or February?), which is said, never read one way."""
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    x = synth.write_extract(tmp_path / "x", n=1500, ratio=True)
    with open(x, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    for r in rows:
        y, m, d = r["ORIG_DATE"].split("-")
        r["ORIG_DATE"] = f"{m}/{min(int(d), 12):02d}/{y}"
    with open(x, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    b = book.set_up(x).book
    _answer(b)
    _control(b, run_kind="Finding and testing a new variable", new_variable_step="Test from a pre-spec",
             **{"derived|1": ("INCOME_TO_SALES", "INCOME", "SALES")})
    book.set_up(x)
    _columns(b, "ORIG_DATE", C_MEANS="Origination date")
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    f = _spec(x.parent, commit=False)
    _held(b, str(f))
    ran = book.run(b)
    assert ran.ok, ran.lines
    chk = _check(b)
    assert chk["Pre-spec commit"] == "not in a git repository, so not committed: a pre-spec only counts once it " \
                                     "is committed"
    assert chk["Holdout"].startswith('Couldn\'t be checked: the dates in "ORIG_DATE" (when each loan was made) read '
                                     'two ways: ')
    assert chk["Holdout"].endswith(" The holdout is 2024-01-01 to 2024-12-31.")
    assert chk["Runs that touched the holdout"] == "0 in this workbook's Log before this one, which couldn't be checked"
    devs = [w for w in _warnings(chk) if w.startswith(confirmatory.DEVIATES)]
    assert not any("window" in d for d in devs)
    assert ("Deviates from pre-spec: The pre-spec's holdout is 2024-01-01 to 2024-12-31; when this run's loans were "
            "made isn't known." in devs)
    log = _log(b)
    assert log[1].startswith("Deviates from pre-spec prespec.yaml (not committed): ")
    assert log[2].startswith('Holdout not checked: the dates in "ORIG_DATE" (when each loan was made) read two ways')


def _prespec_rows(b) -> list[tuple[str, str]]:
    """Check's pre-spec rows, in order, from "Pre-spec" to the holdout count."""
    out = []
    for r in load_workbook(b)["Check"].iter_rows(min_row=4):
        if r[1].value == "Pre-spec" or out:
            out.append((r[1].value, r[2].value))
        if r[1].value == "Runs that touched the holdout":
            break
    return out


def _prevalence_groups(b, column) -> list[str]:
    """The groups the Prevalence tab shows for a new column, lowest first, from its first grid."""
    ws = load_workbook(b)[prevalence.SHEET]
    top = next(r for r in range(1, ws.max_row + 1)
               if str(ws.cell(row=r, column=2).value or "").startswith(f"{column} = "))
    row = next(r for r in range(top + 1, ws.max_row + 1) if ws.cell(row=r, column=prevalence.GROUP_COL).value)
    return [ws.cell(row=row, column=c).value for c in range(prevalence.GROUP_COL, ws.max_column + 1, 2)
            if ws.cell(row=row, column=c).value]


@needs_git
def test_check_names_the_pre_specs_groups_as_the_tabs_name_them(tmp_path, monkeypatch):
    """Found 26 Sep 2026 (final check, F13): the pre-spec named its lowest group "up to 0.09" from the cut
    points alone, so Check echoed "up to 0.09" while Prevalence named the same group "0.02 - 0.09" from the
    data. The firm: bands read as ranges, never "up to". The file may say it either way; the workbook says it
    one way, the tabs'."""
    x, b = _ready(tmp_path, monkeypatch, n=1500)
    _spec(x.parent, reference="up to 0.09")
    _held(b, "prespec.yaml")
    assert book.run(b).ok
    shown = [g for g in _prevalence_groups(b, "INCOME_TO_SALES") if g not in engine.REASON_LABEL.values()]
    lowest, highest = shown[0], shown[-1]
    assert lowest.endswith(" - 0.09") and highest.startswith("2.00 - ")

    # the end groups' names hold the smallest and largest ratio among the loans run, counted here by hand
    with open(x, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    ratios = [float(r["INCOME"]) / float(r["SALES"]) for r in rows if r["SALES"] not in ("", "0")]
    low, high = float(lowest.split(" - ")[0]), float(highest.split(" - ")[1])
    assert low <= min(ratios) < low + 0.01 and high - 0.01 < max(ratios) <= high

    got = _prespec_rows(b)
    says = dict(got)["What the pre-spec says"].splitlines()
    assert f"reference: {lowest}" in says
    groups = next(s for s in says if s.startswith("bins: ")).split("(groups: ")[1].rstrip(")").split("; ")
    assert [g for g in groups if g in shown] == shown and (groups[0], groups[-1]) == (lowest, highest)
    assert (f'Deviates from pre-spec: The reference group is "the low half of each pocket", which is not one of its '
            f'bins\' groups in this run; the pre-spec says "{lowest}".') in [v for k, v in got if k == "Warning"]
    assert not any("up to" in str(v) for _, v in got)


@needs_git
def test_a_pre_spec_written_after_the_run_is_warned_of_and_the_run_goes_on(tmp_path, monkeypatch):
    """Found 26 Sep 2026 (final check, F13): the example pre-spec was dated 2026-10-01, five days after the run,
    and nothing said so. A pre-spec can't have been written after the run held to it."""
    x, b = _ready(tmp_path, monkeypatch, n=1500)
    later = date.today() + timedelta(days=1)
    _spec(x.parent, written=later.isoformat())
    _held(b, "prespec.yaml")
    ran = book.run(b)
    assert ran.ok, ran.lines
    got = _prespec_rows(b)
    # with the other pre-spec lines: straight after what it says
    said = [k for k, _ in got].index("What the pre-spec says")
    assert got[said + 1] == ("Warning", f"The pre-spec says it was written on {later.isoformat()}, after this run.")
    # written the day of the run: nothing to say
    _spec(x.parent, written=date.today().isoformat())
    assert book.run(b).ok
    assert not any("after this run" in str(v) for _, v in _prespec_rows(b))


# --------------------------------------------------------------------------
# What the run used, worked out from the engine's result


def _dated_rows(dates):
    return [{"ID": f"L{i}", "SCORE": 600 + 50 * (i % 3), "CHAN": "AB"[i % 2], "BAL": 100, "BAD": i % 2,
             "GCO": 50 * (i % 2), "RANR": 3, "R": 0.1 * (i % 20), "ORIG": d} for i, d in enumerate(dates)]


def test_a_run_whose_loans_all_sit_in_the_holdout_used_the_holdout():
    ps = prespec.parse({**SPEC, "column": "R"})
    inside = ["2024-01-01", "2024-06-30", "2024-12-31"] * 4
    res = engine.run(cube(origination_date="ORIG"), table(_dated_rows(inside)))
    assert confirmatory.run_range(res, ps) is ps.holdout
    res = engine.run(cube(origination_date="ORIG"), table(_dated_rows(inside + ["2025-01-01"])))
    assert confirmatory.run_range(res, ps) == prespec.DateRange(date(2024, 1, 1), date(2025, 1, 1))
    assert confirmatory.run_range(engine.run(cube(), table(_dated_rows(inside))), ps) is None


def test_a_band_column_tested_is_left_out_of_the_strata_and_its_edges_are_the_bins():
    ps = prespec.parse({**SPEC, "column": "R", "strata": ["CHAN"]})
    bands = [{"name": "r", "field": "R", "edges": [0.1, 0.25, 0.5, 1.0, 2.0]},
             {"name": "score", "field": "SCORE", "edges": [650]}]
    res = engine.run(cube(bands=bands, origination_date="ORIG"),
                     table(_dated_rows(["2024-03-01"] * 30)))
    used = confirmatory.in_use(res, ps)
    assert used["column"] == "R" and used["bins"] == [0.1, 0.25, 0.5, 1.0, 2.0]
    assert used["strata"] == ["SCORE", "CHAN"] and used["reference"] == "the rest of the book"
    assert "window_months" not in used and used["confidence"] == 0.95 and used["holdout"] is ps.holdout
    assert prespec.deviations(ps, used, where="in this run") == [
        "The reference group is `the rest of the book`, which is not one of its bins' groups in this run; the "
        "pre-spec says `0.25 - 0.49`.",
        "Pockets are cut by SCORE, CHAN in this run; the pre-spec says CHAN."]
