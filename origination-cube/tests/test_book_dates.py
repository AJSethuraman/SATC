"""The workbook route for NEXT-GOAL 3.9 to 3.11: a new column made on Control,
and the period and meaning of an amount on Columns; and the origination date on
Columns, whose range Check gives and which leaves no loan out. Each refusal
names its cell, and every number on Check is checked against the extract
counted by hand."""

import csv
import statistics

import pytest
import yaml
from openpyxl import load_workbook

from origination_cube import book, control, synth
from test_book import _answer


def _control(path, **answers):
    """Answer Control by key: a setting's option label, ("own", value) for column D, or a new-column
    slot as {"derived|1": (name, top, bottom)}."""
    wb = load_workbook(path)
    ws = wb[control.SHEET]
    for r in ws.iter_rows(min_row=control.FIRST_ROW):
        k = r[control.KEY_COL - 1].value
        if k in answers:
            v = answers[k]
            if k.startswith("derived|"):
                for c, x in zip((3, 4, 5), v):
                    r[c - 1].value = x
            elif isinstance(v, tuple):
                r[control.OWN_COL - 1].value = v[1]
            else:
                r[control.CHOOSE_COL - 1].value = v
    wb.save(path)


def _columns(path, name, **cells):
    """Set cells on one Columns row, by constant name: _columns(b, "INCOME", C_PERIOD="per year")."""
    wb = load_workbook(path)
    ws = wb["Columns"]
    for r in ws.iter_rows(min_row=book.COL_FIRST):
        if r[book.C_NAME - 1].value == name:
            for const, v in cells.items():
                ws.cell(row=r[0].row, column=getattr(book, const)).value = v
            wb.save(path)
            return r[0].row
    raise KeyError(name)


def _check(path) -> dict:
    out = {}
    for r in load_workbook(path)["Check"].iter_rows(min_row=4):
        if r[1].value:
            out.setdefault(r[1].value, []).append(r[2].value)
    return {k: v[0] if len(v) == 1 else v for k, v in out.items()}


def _dated(tmp_path, n=3000):
    x = synth.write_extract(tmp_path, n=n, ratio=True)
    out = book.set_up(x)
    _answer(out.book)
    return x, out.book


def _ratios(x):
    with open(x, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    return [float(r["INCOME"]) / float(r["SALES"]) for r in rows if r["SALES"] not in ("", "0")]


def test_check_gives_the_origination_dates_and_no_young_loan_is_left_out(tmp_path):
    """The extract with some loans made days before, one in the future and one with no date: every loan
    runs, and Check gives the range and the count without a date, as a fact that removes nothing."""
    x = synth.write_extract(tmp_path, n=3000, ratio=True)
    with open(x, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for i, d in ((0, "2026-09-25"), (1, "2026-09-26"), (2, "2027-01-01"), (3, ""), (4, "not a date")):
        rows[i]["ORIG_DATE"] = d
    with open(x, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    b = book.set_up(x).book
    _answer(b)
    wb = load_workbook(b)
    means = {r[book.C_NAME - 1].value: r[book.C_MEANS - 1].value
             for r in wb["Columns"].iter_rows(min_row=book.COL_FIRST)}
    assert means["ORIG_DATE"] == "Origination date"
    shown = [str(c.value) for ws in (wb[control.SHEET], wb[control.OPTIONS_SHEET]) for r in ws.iter_rows() for c in r
             if c.value is not None]
    assert not any(w in s for s in shown for w in ("loan age", "Loan age", "window", "as-of", "As-of",
                                                    "Outcome date", "months on book"))
    ran = book.run(b)
    assert ran.ok, ran.lines
    chk = _check(b)
    dates = sorted(r["ORIG_DATE"] for r in rows if r["ORIG_DATE"][:2] == "20")
    assert chk["Loans run"] == f"{len(rows):,}" == "3,000"
    assert chk["Origination dates"] == f"{dates[0]} to {dates[-1]} (3,000 loans; 2 without a readable date)"
    assert dates[-1] == "2027-01-01"
    assert not any(k.startswith(("Left out for loan age", "Left out: under", "As-of", "Outcome window"))
                   for k in chk)


def test_a_new_column_is_made_on_control_listed_on_columns_looked_at_and_run(tmp_path):
    x, b = _dated(tmp_path)
    _control(b, **{"derived|1": ("INCOME_TO_SALES", "INCOME", "SALES")})
    wb = load_workbook(b)
    opts = [c.value for c in wb[control.OPTIONS_SHEET]["J"][1:] if c.value]
    assert "INCOME" in opts and "SALES" in opts and "CHANNEL" not in opts     # number columns only

    ran = book.run(b)
    slot = control.row_of(load_workbook(b)[control.SHEET], "derived|1")
    assert not ran.ok and any(f"Control!C{slot}" in x and "isn't on Columns yet. Press Set up again" in x
                              for x in ran.lines)

    out = book.set_up(x)
    assert any(line.startswith("Made INCOME_TO_SALES = INCOME ÷ SALES on Columns and Look. Blank on 3 where "
                               "SALES is zero; 3 where SALES blank") for line in out.lines), out.lines
    wb = load_workbook(b)
    cols = wb["Columns"]
    row = next(r for r in cols.iter_rows(min_row=book.COL_FIRST) if r[book.C_NAME - 1].value == "INCOME_TO_SALES")
    assert row[book.C_MEANS - 1].value == "Amount or number" and row[book.C_CUT - 1].value == "Yes"
    assert row[book.C_WHY - 1].value == "made on Control: INCOME ÷ SALES"
    assert row[book.C_MADE - 1].value == "INCOME ÷ SALES" and cols.column_dimensions["P"].hidden
    # to four figures (0.3075), not seventeen
    assert all(len(s.strip().lstrip("0.").replace(".", "")) <= 4 for s in row[book.C_SAMPLES - 1].value.split(","))
    assert cols[book.CONFIRM_CELL].value is None and "INCOME_TO_SALES" in cols["D3"].value    # check it first
    # its Look block, against the extract divided by hand
    look = wb["Look"]
    at = next(r for r in range(1, look.max_row + 1) if look.cell(row=r, column=2).value == "INCOME_TO_SALES")
    lines = {look.cell(row=r, column=2).value: look.cell(row=r, column=3).value for r in range(at + 1, at + 10)}
    ratios = _ratios(x)
    assert lines["Loans"] == 3000 and lines["Blank"] == 6
    assert lines["Smallest"] == pytest.approx(min(ratios)) and lines["Largest"] == pytest.approx(max(ratios))
    assert lines["Median"] == pytest.approx(statistics.median(ratios))

    _columns(b, "INCOME_TO_SALES", C_EDGES="0.1; 0.25; 0.5; 1; 2")
    _columns(b, "INCOME", C_CUT="No")
    _columns(b, "SALES", C_CUT="No")
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    ran = book.run(b)
    assert ran.ok, ran.lines
    chk = _check(b)
    assert chk["New column: INCOME_TO_SALES"] == ("INCOME ÷ SALES on each loan. Made on 2,994; blank on 6 "
                                                  "(3 where SALES is zero; 3 where SALES blank).")
    assert chk["Band edges used: INCOME_TO_SALES"] == "0.1; 0.25; 0.5; 1; 2  (6 bands)"
    got = yaml.safe_load(b.with_name(f"{b.stem} - what ran.yaml").read_text(encoding="utf-8"))
    assert got["derived"] == [{"name": "INCOME_TO_SALES", "top": "INCOME", "bottom": "SALES"}]
    wb = load_workbook(b)
    assert any("INCOME_TO_SALES" in str(c.value) for r in wb["Grids"].iter_rows() for c in r)
    # the Run writes Look again, and the new column's block is still there
    assert any(wb["Look"].cell(row=r, column=2).value == "INCOME_TO_SALES" for r in range(1, wb["Look"].max_row + 1))

    _control(b, **{"derived|1": ("INCOME_TO_SALES", "INCOME", "REV_DEBT")})
    ran = book.run(b)
    assert not ran.ok and any(f"Control!C{slot}" in x and "was made as INCOME ÷ SALES and Control now says INCOME "
                              "÷ REV_DEBT. Press Set up again" in x for x in ran.lines)
    _control(b, **{"derived|1": (None, None, None)})
    ran = book.run(b)
    assert not ran.ok and any("was a new column made on Control, and Control doesn't have it now. Press Set up "
                              "again" in x for x in ran.lines)
    _control(b, **{"derived|1": ("INCOME_TO_SALES", None, "SALES")})
    ran = book.run(b)
    assert not ran.ok and any(f"Control!D{slot}: new column 1 needs a top" in x for x in ran.lines)


def test_periods_and_meanings_on_columns_are_recorded_warned_and_split_by(tmp_path):
    x, b = _dated(tmp_path)
    _control(b, **{"derived|1": ("INCOME_TO_SALES", "INCOME", "SALES")})
    x2 = book.set_up(x)
    assert x2.ok
    _columns(b, "INCOME", C_CUT="No", C_PERIOD="per year", C_DEFINE="household income, trailing twelve months")
    _columns(b, "SALES", C_CUT="No", C_PERIOD="per month", C_DEFINE="business sales, one month times twelve")
    _columns(b, "INCOME_TO_SALES", C_SPLIT="Yes")
    row = _columns(b, "CHANNEL", C_PERIOD="per year")
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    ran = book.run(b)
    assert not ran.ok and any(f"Columns!{book._col(book.C_PERIOD)}{row}" in x and "not an amount, so it has no "
                              "period" in x for x in ran.lines)
    _columns(b, "CHANNEL", C_PERIOD=None)
    ran = book.run(b)
    assert ran.ok, ran.lines
    chk = _check(b)
    warnings = chk["Warning"] if isinstance(chk["Warning"], list) else [chk["Warning"]]
    assert any('INCOME_TO_SALES divides "INCOME" (per year) by "SALES" (per month): they aren\'t over the same '
               'period, so it reads 12 times a like-for-like ratio' in w for w in warnings)
    assert chk["What INCOME is"] == "per year; household income, trailing twelve months"
    assert chk["What SALES is"] == "per month; business sales, one month times twelve"
    got = yaml.safe_load(b.with_name(f"{b.stem} - what ran.yaml").read_text(encoding="utf-8"))
    assert got["columns"]["INCOME"] == {"means": "amount", "period": "per_year",
                                        "definition": "household income, trailing twelve months"}
    assert got["split"] == {"field": "INCOME_TO_SALES", "how": "own_median"}
    # the new column splits the pockets like any number column, and Look plots it against each band
    wb = load_workbook(b)
    assert "Split" in wb.sheetnames
    text = [c.value for r in wb["Look"].iter_rows() for c in r if isinstance(c.value, str)]
    assert "INCOME_TO_SALES against FICO" in text


def test_an_old_control_row_for_a_removed_setting_is_refused_until_set_up_again(tmp_path):
    """A workbook set up before the loan age filter was taken out still shows its row, answered. The run would
    ignore it while the tab says it applies, so Run refuses by cell, and Set up again takes the row off."""
    x = synth.write_extract(tmp_path, n=500)
    b = book.set_up(x).book
    _answer(b)
    wb = load_workbook(b)
    ws = wb[control.SHEET]
    r = ws.max_row + 1
    ws.cell(row=r, column=2).value = "Loan age filter: how old a loan must be to count"
    ws.cell(row=r, column=control.CHOOSE_COL).value = "24 months on book or more"
    ws.cell(row=r, column=control.KEY_COL).value = "min_age_months"
    wb.save(b)
    ran = book.run(b)
    assert not ran.ok
    assert any(f'Control!C{r}: "Loan age filter: how old a loan must be to count" is no longer used: every loan in '
               f'the extract is run. Press Set up again to take the row off.' in line for line in ran.lines), ran.lines
    book.set_up(x)
    assert control.row_of(load_workbook(b)[control.SHEET], "min_age_months") is None
    ran = book.run(b)
    assert ran.ok, ran.lines
    assert _check(b)["Loans run"] == "500"


def test_a_remembered_outcome_date_is_not_suggested_and_is_said_to_be_unused(tmp_path):
    """A column confirmed as Outcome date before that meaning was taken out: Set up doesn't suggest it again,
    and the Learned tab says, in words, that nothing uses it, rather than showing the code."""
    from origination_cube import memory
    memory.save({"columns": {"BAD_DATE": {"means": "outcome_date", "first": "2026-09-26", "last": "2026-09-26",
                                          "times": 3}}, "answers": {}})
    x = synth.write_extract(tmp_path, n=500)
    with open(x, newline="", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    for r in rows:
        r["BAD_DATE"] = r["ORIG_DATE"] if r["BAD_FLAG"] == "1" else ""
    with open(x, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    b = book.set_up(x).book
    wb = load_workbook(b)
    means = {r[book.C_NAME - 1].value: r[book.C_MEANS - 1].value
             for r in wb["Columns"].iter_rows(min_row=book.COL_FIRST)}
    assert means["BAD_DATE"] != "Outcome date"
    learned = {r[2]: r[3] for r in wb["Learned"].iter_rows(min_row=4, values_only=True) if r[2]}
    assert learned["BAD_DATE"] == "Outcome date, which the cube no longer uses. Set it to Forget."
