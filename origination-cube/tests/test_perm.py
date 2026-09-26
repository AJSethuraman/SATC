"""The shuffle test (docs/statistics.md B2), ruling OC-34: the worked example,
every labelling of a tiny book counted out by hand, shuffling within the
comparison group, the +1, ties, reproducibility, numpy loaded only when used,
and the production default of 10,000 run end to end."""

import copy
import itertools
import math
import os
import re
import subprocess
import sys
import time
from pathlib import Path

import numpy as np
import pytest
import yaml
from openpyxl import load_workbook

from conftest import PRODUCTION_SHUFFLES, cube, row, table
from origination_cube import book, engine, perm, synth
from origination_cube import config as cfgmod
from origination_cube.ingest import read_table


def _rate(y, x, rows):
    return sum(y[i] for i in rows) / sum(x[i] for i in rows)


def _gap(y, x, pocket, group):
    rest = [i for i in group if i not in pocket]
    return _rate(y, x, pocket) - _rate(y, x, rest)


def _share_as_far(gaps, seen):
    return sum(abs(g) >= abs(seen) - 1e-12 for g in gaps) / len(gaps)


def _within(p, exact, shuffles):
    """A shuffle p-value against the exact one: within four standard errors of Monte Carlo."""
    return abs(p - exact) <= 4 * math.sqrt(max(exact * (1 - exact), 1e-4) / shuffles) + 1 / shuffles


# --------------------------------------------------------------------------
# The worked example


def _b2_example():
    """statistics-examples.py section 6, regenerated exactly: numpy's
    default_rng(7), lognormal balances, 15% and 7% of loans losing 30-80%."""
    rng = np.random.default_rng(7)
    bal_p = rng.lognormal(np.log(20000), 0.5, 40)
    bal_r = rng.lognormal(np.log(20000), 0.5, 400)
    gco_p = np.where(rng.random(40) < 0.15, bal_p * rng.uniform(0.3, 0.8, 40), 0.0)
    gco_r = np.where(rng.random(400) < 0.07, bal_r * rng.uniform(0.3, 0.8, 400), 0.0)
    return gco_p, bal_p, gco_r, bal_r


def test_b2_the_worked_example():
    """A 40-loan pocket at 2.11% GCO per dollar against 400 loans at 5.07%, gap
    -2.96 points: the document's shuffles found 3,497 of 10,000 at least as far
    apart, p 0.3498. The cube's own shuffles are a different draw, so its p must
    land within Monte Carlo error of that: two independent runs of 10,000 differ
    by about sqrt(2 p (1 - p) / 10,000) = 0.0067."""
    gco_p, bal_p, gco_r, bal_r = _b2_example()
    assert round(gco_p.sum() / bal_p.sum(), 4) == 0.0211 and round(gco_r.sum() / bal_r.sum(), 4) == 0.0507
    got, _ = perm.pocket_vs_rest(gco_p, bal_p, gco_r, bal_r, shuffles=10_000)
    assert round(got.gap * 100, 2) == -2.96
    assert got.shuffles == 10_000 and got.p == (got.hits + 1) / 10_001
    assert abs(got.p - 0.3498) < 4 * math.sqrt(2 * 0.35 * 0.65 / 10_000)
    assert got.p > 0.05                                                  # in line


# --------------------------------------------------------------------------
# Correct against counting every labelling

Y = [5.0, 0.0, 3.0, 0.0, 8.0, 1.0, 0.0]
X = [10.0, 20.0, 15.0, 5.0, 30.0, 10.0, 12.0]


def test_every_shuffle_is_a_real_labelling_and_the_p_is_the_share_of_them():
    """Seven loans, a pocket of three: 35 ways to deal the label. Every shuffled
    gap must be one of those 35 gaps, each turning up about 1 time in 35, and the
    p-value must be the share of labellings at least as far apart as the one
    seen, to within Monte Carlo error."""
    k, shuffles = 3, 20_000
    labellings = list(itertools.combinations(range(7), k))
    gaps = [_gap(Y, X, c, range(7)) for c in labellings]
    seen = _gap(Y, X, range(k), range(7))
    exact = _share_as_far(gaps, seen)
    got, st = perm.pocket_vs_rest(Y[:k], X[:k], Y[k:], X[k:], shuffles=shuffles, keep=True)
    assert got.gap == pytest.approx(seen, abs=1e-15)
    draws = np.concatenate(st.draws)[:, 0]
    distinct = np.array(sorted({round(g, 12) for g in gaps}))
    nearest = np.abs(draws[:, None] - distinct[None, :]).argmin(axis=1)
    assert np.abs(draws - distinct[nearest]).max() < 1e-12
    counts = np.bincount(nearest, minlength=len(distinct))
    for j, value in enumerate(distinct):
        share = sum(abs(g - value) < 1e-12 for g in gaps) / len(gaps)
        assert abs(counts[j] / shuffles - share) < 5 * math.sqrt(share * (1 - share) / shuffles), value
    assert _within(got.hits / got.shuffles, exact, shuffles)
    assert got.p == (got.hits + 1) / (got.shuffles + 1)


def test_pockets_sharing_one_shuffle_each_get_their_own_exact_answer():
    """Every pocket of a grid is a slice of the same shuffled order; each must
    still come out as if it were shuffled alone."""
    lay = [0, 0, 1, 1, 1, 2, 2]
    st = perm.RestGap()
    s = perm.Structure("book", None, [lay], {(0, "r"): st})
    perm.run(7, [perm.Column("r", Y, X)], [s], 20_000, perm.seed_of("three pockets"))
    for pocket in (0, 1, 2):
        members = [i for i, p in enumerate(lay) if p == pocket]
        gaps = [_gap(Y, X, c, range(7)) for c in itertools.combinations(range(7), len(members))]
        exact = _share_as_far(gaps, _gap(Y, X, members, range(7)))
        assert _within(st.answers[pocket].hits / 20_000, exact, 20_000), pocket


def test_shuffles_stay_within_the_group():
    """Against the rest of its band a pocket's loans are shuffled with its band's
    loans only. Band 0 loses 10-40%, band 1 about 90%: a gap from mixing them
    would be far outside what band 0's own labellings can make."""
    y = [1.0, 2.0, 3.0, 4.0, 9.0, 9.5, 8.5, 9.0, 9.2, 8.8]
    x = [10.0] * 10
    group = [0, 0, 0, 0, 1, 1, 1, 1, 1, 1]
    lay = [0, 0, 1, 1, 2, 2, 2, 3, 3, 3]
    st = perm.RestGap(keep=True)
    s = perm.Structure("band", group, [lay], {(0, "r"): st})
    perm.run(10, [perm.Column("r", y, x)], [s], 5_000, perm.seed_of("two bands"))
    draws = np.concatenate(st.draws)
    for col, pocket in enumerate(st.pockets):
        members = [i for i, p in enumerate(lay) if p == pocket]
        band = [i for i, g in enumerate(group) if g == group[members[0]]]
        possible = np.array(sorted({_gap(y, x, c, band) for c in itertools.combinations(band, len(members))}))
        assert np.abs(draws[:, col][:, None] - possible[None, :]).min(axis=1).max() < 1e-12, pocket
    # band 0's first pocket is as far apart as its band allows: 2 labellings of 6
    assert _within(st.answers[0].hits / 5_000, 2 / 6, 5_000)


def test_the_engine_shuffles_the_rest_of_its_band_within_the_band():
    """The wiring, counted out: eight loans in two bands. Against the rest of its
    band a pocket's p is the share of its band's 6 labellings; against the rest
    of the book, of the book's 28."""
    loans = [(600, "A", 100, 40), (600, "A", 200, 0), (600, "B", 150, 10), (600, "B", 250, 0),
             (700, "A", 300, 0), (700, "A", 120, 90), (700, "B", 180, 0), (700, "B", 90, 5)]
    rows = [row(i, sc, ch, bal, 1 if gco else 0, gco) for i, (sc, ch, bal, gco) in enumerate(loans)]
    bench = {"min_units": 2, "min_events": 1, "worse_at": 1.25, "better_at": 0.8, "confidence": 0.95, "power": 0.8,
             "compare_to": "peers", "many_tests": "none", "materiality": "none", "shuffles": 20_000}
    res = engine.run(cube(benchmark=bench), table(rows))
    y, x = [ln[3] for ln in loans], [ln[2] for ln in loans]
    for (b, d), c in res.grids[0].inner():
        s = c.rates["gco_rate"]
        members = [i for i, ln in enumerate(loans) if (f"{ln[0]}" in b, ln[1]) == (True, d)]
        band = [i for i, ln in enumerate(loans) if f"{ln[0]}" in b]
        in_band = [_gap(y, x, cc, band) for cc in itertools.combinations(band, 2)]
        in_book = [_gap(y, x, cc, range(8)) for cc in itertools.combinations(range(8), 2)]
        assert s.test == engine.SHUFFLE_TEST and s.shuffles == 20_000
        assert _within(s.hits_band / 20_000, _share_as_far(in_band, _gap(y, x, members, band)), 20_000), (b, d)
        assert _within(s.hits_book / 20_000, _share_as_far(in_book, _gap(y, x, members, range(8))), 20_000), (b, d)


def test_the_split_halves_are_shuffled_within_their_pocket():
    """The split's dollar test (A7 on B2): high and low halves re-dealt inside
    each pocket. Per pocket, rate(high) - rate(low); pooled, the high halves'
    actual total less what their own low half's rate predicts. Two pockets of
    four: 6 labellings each, 36 together, all counted."""
    y = [30.0, 12.0, 0.0, 5.0, 40.0, 0.0, 8.0, 1.0]
    x = [100.0, 60.0, 80.0, 50.0, 90.0, 70.0, 60.0, 40.0]
    group = [0, 0, 0, 0, 1, 1, 1, 1]
    lay = [0, 0, 1, 1, 2, 2, 3, 3]                       # 2g: high half, 2g + 1: low half
    st = perm.HalfGap(pooled={0, 1})
    s = perm.Structure("halves", group, [lay], {(0, "r"): st})
    perm.run(8, [perm.Column("r", y, x)], [s], 20_000, perm.seed_of("halves"))

    def stat(highs):
        per, t = [], 0.0
        for g, hi in zip((0, 1), highs):
            rows = [i for i in range(8) if group[i] == g]
            lo = [i for i in rows if i not in hi]
            per.append(_rate(y, x, hi) - _rate(y, x, lo))
            t += sum(y[i] for i in hi) - _rate(y, x, lo) * sum(x[i] for i in hi)
        return per, t

    seen_per, seen_t = stat(([0, 1], [4, 5]))
    options = [list(itertools.combinations(range(4), 2)), list(itertools.combinations(range(4, 8), 2))]
    joint = [stat((a, b)) for a in options[0] for b in options[1]]
    exact_t = _share_as_far([t for _, t in joint], seen_t)
    assert st.pooled.gap == pytest.approx(seen_t) and _within(st.pooled.hits / 20_000, exact_t, 20_000)
    for g in (0, 1):
        exact = _share_as_far([per[g] for per, _ in joint], seen_per[g])
        assert _within(st.answers[g].hits / 20_000, exact, 20_000), g


# --------------------------------------------------------------------------
# The +1, ties, the seed


def test_the_plus_one_keeps_p_off_zero():
    """B2: no finite number of shuffles can justify p = 0, so (hits + 1) / (B + 1)."""
    y, x = [1.0] * 5 + [0.0] * 50, [1.0] * 55
    got, _ = perm.pocket_vs_rest(y[:5], x[:5], y[5:], x[5:], shuffles=1_000)
    assert got.hits == 0 and got.p == 1 / 1_001


def test_an_exact_tie_counts_however_the_sums_round():
    """Every loan at 10% of its balance: every labelling gives a gap of 0, which
    floating-point sums in a shuffled order land either side of. A tie counts."""
    rng = np.random.default_rng(3)
    x = rng.lognormal(np.log(20000), 0.5, 60)
    y = x * 0.1
    got, _ = perm.pocket_vs_rest(y[:15], x[:15], y[15:], x[15:], shuffles=2_000)
    assert got.hits == 2_000 and got.p == 1.0


def test_the_same_extract_gives_the_same_answer(tmp_path):
    cfg, data = synth.write(tmp_path, n=3000)
    a = engine.run(cfgmod.load(cfg), read_table(data))
    b = engine.run(cfgmod.load(cfg), read_table(data))
    got = [[(k, c.rates[m].hits_book, c.rates[m].hits_band) for k, c in g.inner()] for g in (a.grids[0], b.grids[0])
           for m in ("gco_rate", "ranr_rate", "outcome_booked")]
    assert got[:3] == got[3:] and any(h for rows in got for _, h, _ in rows)
    assert perm.seed_of("x") == perm.seed_of("x") != perm.seed_of("y")


def test_a_pockets_answer_does_not_hang_on_the_other_grids(tmp_path):
    """One shuffled order serves every grid; adding a grid slices it again but
    must not change what an existing grid's pockets get."""
    cfg, data = synth.write(tmp_path, n=3000)
    raw, tbl = cfgmod.load(cfg).raw, read_table(data)
    one = engine.run(cfgmod.parse(raw), tbl).grids[0]
    more = copy.deepcopy(raw)
    more["dimensions"] = raw["dimensions"] + [{"name": "asset", "field": "ASSET_CLASS"}]
    two = engine.run(cfgmod.parse(more), tbl).grids[0]
    for (k, c), (k2, c2) in zip(one.inner(), two.inner()):
        assert k == k2 and c.rates["gco_rate"].hits_book == c2.rates["gco_rate"].hits_book
        assert c.rates["gco_rate"].hits_band == c2.rates["gco_rate"].hits_band


# --------------------------------------------------------------------------
# numpy, and the production default


def test_the_cube_starts_without_numpy_and_says_what_is_missing():
    """OC-34: numpy is a required add-on, and the launcher must be able to start
    and say it is missing, so nothing imports it until a shuffle test runs."""
    code = ("import sys; sys.modules['numpy'] = None\n"            # as if it weren't installed
            "from origination_cube import book, cli, engine, perm\n"
            "from conftest import cube, row, table\n"
            "rows = [row(i, 600 + 100 * (i % 2), 'AB'[i % 3 == 0], 100, int(i % 5 == 0), 50 * (i % 5 == 0))\n"
            "        for i in range(40)]\n"
            "try:\n"
            "    engine.run(cube(), table(rows))\n"
            "except perm.NumpyMissing as exc:\n"
            "    print('refused:', exc)\n")
    here = Path(__file__).parent
    env = dict(os.environ, PYTHONPATH=os.pathsep.join([str(here.parent / "src"), str(here)]))
    got = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True, env=env, cwd=here.parent)
    assert got.returncode == 0, got.stderr
    assert "refused:" in got.stdout and "needs numpy" in got.stdout, got.stdout + got.stderr


def test_the_production_default_runs_end_to_end(tmp_path, monkeypatch):
    """The real 10,000 shuffles through the workbook on the 8,000-loan synthetic
    book, from Set up to the tabs: every other test runs fewer (conftest.py)."""
    assert PRODUCTION_SHUFFLES == 10_000
    monkeypatch.setattr(perm, "SHUFFLES", PRODUCTION_SHUFFLES)
    from test_book import _answer
    out = book.set_up(synth.write_extract(tmp_path, n=8000))
    _answer(out.book)
    start = time.perf_counter()
    ran = book.run(out.book)
    took = time.perf_counter() - start
    assert ran.ok, ran.lines
    wb = load_workbook(out.book)
    check = {r[1].value: r[2].value for r in wb["Check"].iter_rows(min_row=4)}
    assert "shuffled 10,000 times" in check["Tests"] and "no continuity correction" in check["Tests"]
    ran_with = yaml.safe_load(out.book.with_name(f"{out.book.stem} - what ran.yaml").read_text())
    assert ran_with["benchmark"]["shuffles"] == 10_000
    ws = wb["Where it bleeds"]
    assert ws.cell(row=4, column=20).value == "Test"
    gco = [r for r in range(5, ws.max_row + 1) if ws.cell(row=r, column=2).value == "GCO per booked dollar"]
    first = gco[0]                                       # the largest GCO excess: the planted pocket
    assert ws.cell(row=first, column=6).value == "Broker" and ws.cell(row=first, column=18).value == "worse"
    # the shuffle count made literal (NEXT-GOAL 3.6): how many of the 10,000 made a gap as big, against the
    # rest of its band (the comparison that decides the flag)
    tests = [ws.cell(row=r, column=20).value for r in gco]
    assert all(t is None or re.fullmatch(r"shuffled: [\d,]+ of 10,000", t) for t in tests)     # None: untested
    assert ws.cell(row=first, column=20).value == "shuffled: 0 of 10,000"
    outcome = [ws.cell(row=r, column=20).value for r in range(5, ws.max_row + 1)
               if ws.cell(row=r, column=2).value == "Outcome, share of loans"]
    assert set(outcome) <= {"z test", "exact test", None} and "z test" in outcome
    print(f"\n8,000 loans, 6 grids, 10,000 shuffles: {took:.1f} s for the whole Run")
