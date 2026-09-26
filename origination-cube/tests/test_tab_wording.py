"""The final check's wording finding (F9, 26 Sep 2026): four places where a tab
spoke to the repository instead of the credit analyst reading it.

1. Check named the split's test and "no continuity correction" without saying
   what either means (docs/statistics.md A6).
2. Check carried a ruling number and a dated "firm's call".
3. Split answered "Same size in every pocket?" with "yes/no outcome only" for a
   dollar rate, which reads as an answer. Cochran's Q (A8) isn't run there.
4. Columns!D3 still asked for C3 to be set to Yes after a Run had gone through
   with it set.
"""

from __future__ import annotations

import re

import pytest
from openpyxl import load_workbook

from conftest import TEST_SHUFFLES
from origination_cube import book, config as cfgmod, engine, perm, synth
from test_book import _answer


@pytest.fixture(scope="module")
def split_book(tmp_path_factory):
    """The synthetic book split by revolving debt, so Check and Split both have every row."""
    d = tmp_path_factory.mktemp("wording")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("CUBE_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
        out = book.set_up(synth.write_extract(d, n=8000))
        _answer(out.book)
        wb = load_workbook(out.book)
        for r in wb["Columns"].iter_rows(min_row=book.COL_FIRST):
            if r[book.C_NAME - 1].value == "REV_DEBT":
                r[book.C_SPLIT - 1].value = "Yes"
        wb.save(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return load_workbook(out.book)


def _check(wb) -> dict:
    return {r[1].value: r[2].value for r in wb["Check"].iter_rows(min_row=4)}


def test_check_says_what_the_split_test_asks_and_which_form(split_book):
    said = _check(split_book)["Tests"]
    # the name stays, so it can be looked up; what it asks and which form follow in plain words (A6)
    assert "Cochran-Mantel-Haenszel" in said
    assert "asks whether an odds ratio this far from 1 could come from shuffling loans within their pockets" in said
    assert "no continuity correction: nothing is taken off the gap between actual and expected before it is " \
           "squared. That is the modern form." in said
    for s in re.split(r"(?<=\.) ", said):
        if "Cochran" in s or "continuity" in s:
            assert len(s.split()) <= 25, s


def test_no_ruling_number_or_dated_call_on_any_tab(split_book):
    for ws in split_book.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert not re.search(r"\bOC-\d|firm's call", c.value), f"{ws.title}!{c.coordinate}: {c.value}"
    check = _check(split_book)
    # what each sentence told the reader is still there
    assert check["The allowance for many tests covers"] == "each grid and measure on its own, one comparison at a time"
    assert check["Contribution before losses"] == ("RANR + GCO, per booked dollar. This assumes RANR has gross "
                                                   "charge-offs taken out. If RANR nets recoveries instead, "
                                                   "contribution is overstated by the recoveries.")


def test_same_size_is_not_tested_for_a_dollar_rate(split_book):
    ws = split_book["Split"]
    heads = [c for row in ws.iter_rows() for c in row if c.value == "Same size in every pocket?"]
    assert heads
    seen = set()
    for h in heads:
        r = h.row + 1
        while ws.cell(row=r, column=2).value:
            name, said = ws.cell(row=r, column=2).value, ws.cell(row=r, column=h.column).value
            if name == "Outcome, share of loans":
                assert said in ("yes", "no: bigger in some pockets"), said        # tested: an answer
            else:
                assert said == "not tested: dollar rate", (name, said)
            # wider than its column: it wraps there, or the p-value beside it cuts it off on the page
            assert ws.cell(row=r, column=h.column).alignment.wrap_text, name
            seen.add(name)
            r += 1
    assert "Outcome, share of loans" in seen and len(seen) > 1
    # the tab still says what the question is for the row that is tested
    how = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(5, 11)}
    assert "Cochran's Q checks whether the gap is about the same size in every pocket" in how["Pooled across pockets"]


def test_same_size_says_why_the_outcome_was_not_tested():
    outcome = cfgmod.Measure(name="outcome_loans", mode="flagwt", flag="BAD", per=engine.EACH_LOAN)
    gco = cfgmod.Measure(name="gco_rate", mode="sumnum", value="GCO", per="BAL")
    assert book._same_size(outcome, {"pockets": 1}, 0.95) == "not tested: too few pockets"
    # pockets enough, but the high halves had none with the outcome: pooled odds 0, and Q can't be worked out
    assert book._same_size(outcome, {"pockets": 4, "odds": 0.0}, 0.95) == "not tested: couldn't be worked out"
    assert book._same_size(gco, {"pockets": 6}, 0.95) == "not tested: dollar rate"
    assert book._same_size(outcome, {"steady_p": 0.5}, 0.95) == "yes"
    assert book._same_size(outcome, {"steady_p": 0.001}, 0.95) == "no: bigger in some pockets"


def _add_column(x, name, values):
    text = x.read_text().splitlines()
    x.write_text("\n".join([text[0] + f",{name}"] + [t + f",{values(i)}" for i, t in enumerate(text[1:])]) + "\n")


def test_a_run_retires_the_ask_to_check_new_columns(tmp_path):
    x = synth.write_extract(tmp_path, n=2000)
    out = book.set_up(x)
    _answer(out.book)
    _add_column(x, "NEW_LINE", lambda i: i % 7)
    book.set_up(x)
    ws = load_workbook(out.book)["Columns"]
    assert ws[book.CONFIRM_CELL].value is None
    assert "New since the last check: NEW_LINE" in ws["D3"].value and "set C3 to Yes again" in ws["D3"].value

    wb = load_workbook(out.book)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(out.book)
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    ws = load_workbook(out.book)["Columns"]
    assert ws[book.CONFIRM_CELL].value == "Yes"
    assert "set C3 to Yes" not in str(ws["D3"].value or "")        # checked, and the Run went through
    assert "NEW_LINE" not in str(ws["D3"].value or "")

    # Set up again with nothing new says nothing; a column that is new again says so again
    book.set_up(x)
    ws = load_workbook(out.book)["Columns"]
    assert ws[book.CONFIRM_CELL].value == "Yes" and "New since" not in str(ws["D3"].value or "")
    _add_column(x, "NEWER_LINE", lambda i: i % 5)
    book.set_up(x)
    ws = load_workbook(out.book)["Columns"]
    assert ws[book.CONFIRM_CELL].value is None
    assert "New since the last check: NEWER_LINE. Check them, then set C3 to Yes again." in ws["D3"].value
