"""The workbook route (ruling OC-22): set up from an extract, answer in Excel,
run, read the results in Excel. Nobody types a command, and nothing a person
has answered is thrown away."""

from openpyxl import load_workbook

from origination_cube import book, control, memory, synth

PICK = {"run_kind": "Where the book bleeds", "min_loans": "30", "min_events": "10",
        "materiality": "1% of the book's total losses", "compare_to": "The rest of its band",
        "worse_at": "1.25 times", "better_at": "0.8 times", "confidence": "95%",
        "revenue_line": "Each pocket's own test (suggested)"}


def _answer(path, confirm=True, odd=True):
    wb = load_workbook(path)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        k = r[control.KEY_COL - 1].value
        if k in PICK:
            r[control.CHOOSE_COL - 1].value = PICK[k]
    if confirm:
        wb["Columns"][book.CONFIRM_CELL] = "Yes"
    if odd:
        for r in wb["Odd values"].iter_rows(min_row=5):
            if r[1].value == "FICO":
                r[4].value = "missing"
    wb.save(path)


def test_set_up_writes_every_tab_a_person_needs(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=3000))
    assert out.ok and out.book.exists()
    names = load_workbook(out.book).sheetnames
    assert names[:6] == ["Start here", "Control", "Columns", "Look", "Odd values", "Learned"]


def test_run_refuses_until_answered_naming_each_cell(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=3000))
    ran = book.run(out.book)
    assert not ran.ok
    text = "\n".join(ran.lines)
    assert "Control!C" in text and "Columns!C3" in text
    assert "Traceback" not in text and "`" not in text          # words, not code
    log = load_workbook(out.book)["Log"]
    assert log["A1"].value == "Log" and "Couldn't run" in str(log.cell(row=book.LOG_FIRST, column=2).value)


def test_answers_survive_a_second_set_up(tmp_path):
    x = synth.write_extract(tmp_path, n=3000)
    out = book.set_up(x)
    _answer(out.book)
    wb = load_workbook(out.book)
    wb["Columns"]["C8"] = "servicing"                       # CHANNEL, changed by hand
    wb.save(out.book)
    book.set_up(x)
    wb = load_workbook(out.book)
    assert wb["Columns"][book.CONFIRM_CELL].value == "Yes"
    assert wb["Columns"]["C8"].value == "Servicing data"         # shown as its label
    picked = {r[control.KEY_COL - 1].value: r[control.CHOOSE_COL - 1].value
              for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW)}
    assert picked["confidence"] == "95%" and picked["materiality"] == PICK["materiality"]


def test_a_full_run_writes_results_into_the_workbook(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=6000))
    _answer(out.book)
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    assert any("tie-out checks agree" in line for line in ran.lines)
    assert any("Worst for GCO per booked dollar: FICO " in line and "Broker" in line for line in ran.lines)
    wb = load_workbook(out.book)
    for t in ("Where it bleeds", "Grids", "Check", "Log"):
        assert t in wb.sheetnames
    first = [c.value for c in wb["Where it bleeds"][5]]
    assert first[1] == "Outcome, share of loans" and first[5] == "Broker"
    assert out.book.with_name(f"{out.book.stem} - what ran.yaml").exists()      # the record of what ran


def test_the_learned_tab_prunes_on_the_next_run(tmp_path):
    x = synth.write_extract(tmp_path, n=3000)
    out = book.set_up(x)
    _answer(out.book)
    assert book.run(out.book).ok
    assert memory.load()["columns"]["FICO"]["means"] == "fico"
    wb = load_workbook(out.book)
    for r in wb["Learned"].iter_rows(min_row=4):
        if r[2].value == "FICO":
            r[0].value = memory.FORGET
    wb["Columns"]["C7"] = "score"                # it was a custom score all along
    wb.save(out.book)
    assert book.run(out.book).ok
    assert "FICO" not in memory.load()["columns"]          # a Forget is not learned straight back
    assert not book.run(out.book).ok                        # ... nor on the next Run, until a person says yes
    wb = load_workbook(out.book)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(out.book)
    assert book.run(out.book).ok
    assert memory.load()["columns"]["FICO"]["means"] == "score"
    assert "changed_from" not in memory.load()["columns"]["FICO"]      # forgotten first, so learned fresh


def test_taking_away_every_category_is_said_in_words(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=3000))
    _answer(out.book)
    wb = load_workbook(out.book)
    wb["Columns"]["C8"] = "servicing"            # CHANNEL
    wb["Columns"]["C13"] = "servicing"           # ASSET_CLASS: now no category is left
    wb.save(out.book)
    ran = book.run(out.book)
    assert not ran.ok and any("Nothing is left to cut across" in line for line in ran.lines)
    said = next(line for line in ran.lines if "Nothing is left" in line)
    assert "CHANNEL (Columns!C8, now Servicing data)" in said and "ASSET_CLASS (Columns!C13" in said


def test_a_bad_band_edge_is_named_by_cell(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=3000))
    _answer(out.book)
    wb = load_workbook(out.book)
    wb["Columns"]["F7"] = "700, 650"                        # FICO, edges not rising
    wb.save(out.book)
    ran = book.run(out.book)
    assert not ran.ok and any("Columns!F7" in line for line in ran.lines)


def test_own_band_edges_are_used(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=6000))
    _answer(out.book)
    wb = load_workbook(out.book)
    wb["Columns"]["F7"] = "620, 680, 740"
    wb.save(out.book)
    assert book.run(out.book).ok
    check = {r[1].value: r[2].value for r in load_workbook(out.book)["Check"].iter_rows(min_row=4)}
    assert check["Band edges used: FICO"] == "620; 680; 740  (4 bands)"


def test_cut_by_it_chooses_what_goes_into_the_grids(tmp_path):
    """Start with FICO and one segment, then add more: the Cut by it column."""
    out = book.set_up(synth.write_extract(tmp_path, n=4000))
    _answer(out.book)
    wb = load_workbook(out.book)
    assert wb["Columns"]["D9"].value == "Yes"               # ORIG_BAL, a band, cut by default
    wb["Columns"]["D9"] = "No"
    wb.save(out.book)
    assert book.run(out.book).ok
    check = {r[1].value for r in load_workbook(out.book)["Check"].iter_rows(min_row=4)}
    assert "Band edges used: FICO" in check and "Band edges used: ORIG_BAL" not in check
    book.set_up(synth.write_extract(tmp_path, n=4000))       # and the No survives a second set-up
    assert load_workbook(out.book)["Columns"]["D9"].value == "No"


def test_grids_are_heat_maps_against_the_book_and_against_peers(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=4000))
    _answer(out.book)
    assert book.run(out.book).ok
    ws = load_workbook(out.book)["Grids"]
    heads = {c.value for row in ws.iter_rows(max_row=12) for c in row if c.value}
    assert {"Rate", "Vs the book", "Vs the rest of its band"} <= heads
    scales = [r for rng in ws.conditional_formatting for r in rng.rules if r.type == "colorScale"]
    # white at 1.00x for a multiple; for profit, a gap in points, white at 0, even either way, red below
    # (NEXT-GOAL 3.2)
    mids = [float(r.colorScale.cfvo[1].val) for r in scales]
    assert scales and set(mids) == {1.0, 0.0}
    for r in scales:
        if float(r.colorScale.cfvo[1].val) == 0.0:
            lo, hi = float(r.colorScale.cfvo[0].val), float(r.colorScale.cfvo[2].val)
            assert lo == -hi < 0 and r.colorScale.color[0].rgb.endswith(book.RED)
            assert r.colorScale.color[2].rgb.endswith(book.GREEN)
