"""The judging settings, live in the workbook (OC-40; the firm, 26 Sep 2026: "this is the stuff i want to be
able to adjust in book on the fly ... i know it cannot reband and such").

Five settings on Control judge numbers the Run has already worked out, so they are Excel formulas rather than
values: the loss line (worse at, better at), the profit line, materiality, the confidence level, and what a
pocket is judged against. Change one on Control and every reading, verdict word, dollar figure and colour on
the result tabs follows, with no Run.

Two hidden sheets carry it, both rewritten by every Run:

    _live      each live setting as the number the formulas use, read from Control's cells (the pick, or the
               value typed beside it), the one rounded significance bar (canon S36), what the last Run used, and
               each rate's materiality line. Every formula refers to these by name: worse_at, better_at,
               confidence, significance_bar, judged_band, materiality_kind, materiality_share,
               materiality_gco, profit_kind, profit_line.
    _pockets   one row per pocket and rate: what the Run worked out (the gaps, the p-values after the allowance
               for many tests, both dollar figures, alone in its band, too few losses) and, beside them, the
               formulas that read them: each comparison's reading, the flag, the dollars, whether it is
               material, and profit's literal wording.

The result tabs point at _pockets, so one set of formulas decides for every tab, as the engine does for the
command line.

Why this is sound: the p-values do not depend on the confidence level. The shuffle test, the z test, the exact
test and both allowances for many tests (Benjamini-Hochberg's adjusted p-values and Bonferroni's) are worked
out without it; confidence only sets the bar they are compared with. What does depend on a live setting and is
left as of the Run, and said so where it shows: the order of the rows, the charts, the smallest gap a pocket
could show (confidence, power and the lines), a suggested line worked out from the book (its multiple is the
one the last Run worked out), and Check's other counts.

Excel without dynamic arrays: no SORT, FILTER or LET, so Excel 2016 and LibreOffice both calculate it.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from . import control, engine, stats
from .config import PROFIT

LIVE_SHEET = "_live"
POCKETS = "_pockets"
#: the settings whose change shows at once on the result tabs (settings.yaml: takes_effect: live)
LIVE_KEYS = ("worse_at", "better_at", "revenue_line", "materiality", "confidence", "compare_to")
ALONE = "alone in its band: compared with the book"
BAR = "significance_bar"
IN_USE = "in_use_words"

# _pockets: what the Run worked out (values) ...
(P_KIND, P_GRID, P_BAND, P_SEG, P_MEASURE, P_TITLE, P_LOANS, P_BOOKED, P_FEW, P_ALONE, P_VS_BOOK, P_P_BOOK,
 P_VS_BAND, P_P_BAND, P_EX_BOOK, P_EX_BAND, P_LINE, P_TEST_BOOK, P_TEST_BAND, P_REST_BOOK, P_REST_BAND) = range(1, 22)
# ... and the formulas that read them
(P_BY_BAND, P_READ_BOOK, P_READ_BAND, P_FLAG, P_DOLLARS, P_MATERIAL, P_GAP, P_P, P_SAID, P_TEST, P_REST,
 P_OTHER) = range(22, 34)
POCKET_HEADS = {
    P_KIND: "Grids or three-way", P_GRID: "Grid", P_BAND: "Band", P_SEG: "Segment", P_MEASURE: "Rate (key)",
    P_TITLE: "Rate", P_LOANS: "Loans", P_BOOKED: "Booked dollars (the rate's bottom)",
    P_FEW: "Too few losses to test (fewest losses is a Run setting)", P_ALONE: "Alone in its band",
    P_VS_BOOK: "Vs rest of book (a multiple; profit: the gap, 0.01 = 1 point)",
    P_P_BOOK: "p-value vs book, after the allowance", P_VS_BAND: "Vs rest of band",
    P_P_BAND: "p-value vs band, after the allowance", P_EX_BOOK: "Excess over the book ($)",
    P_EX_BAND: "Excess over its band ($)", P_LINE: "Materiality line in use (from _live)",
    P_TEST_BOOK: "Test vs book", P_TEST_BAND: "Test vs band", P_REST_BOOK: "Rest of book's rate",
    P_REST_BAND: "Rest of band's rate",
    P_BY_BAND: "Judged against its band?", P_READ_BOOK: "Reading vs book", P_READ_BAND: "Reading vs band",
    P_FLAG: "Flag", P_DOLLARS: "Dollars that decide", P_MATERIAL: "Material", P_GAP: "Gap that decides",
    P_P: "p-value that decides", P_SAID: "Flag as the tabs print it", P_TEST: "Test, as printed",
    P_REST: "Rest's rate that decides", P_OTHER: "The other dollars, for reference",
}
P_FIRST = 4               # the first pocket row on _pockets

# _live rows (column C: in use now, D: the last Run used, E: in words now, F: in words at the last Run)
L_WORSE, L_BETTER, L_CONF, L_BAR, L_BAND, L_MKIND, L_MSHARE, L_MGCO, L_PKIND, L_PLINE, L_GCO_TOTAL = range(5, 16)
L_SENTENCE = 17
L_LINES = 20              # the first rate's materiality line
OPT_FIRST, OPT_LAST = 5, 400
L_NAMES = {L_WORSE: "worse_at", L_BETTER: "better_at", L_CONF: "confidence", L_BAR: BAR, L_BAND: "judged_band",
           L_MKIND: "materiality_kind", L_MSHARE: "materiality_share", L_MGCO: "materiality_gco",
           L_PKIND: "profit_kind", L_PLINE: "profit_line"}
L_LABELS = {L_WORSE: "Worse at (a multiple)", L_BETTER: "Better at (a multiple)", L_CONF: "Confidence",
            L_BAR: "Significance bar: ROUND(1 - confidence, 12). A p-value below it is significant",
            L_BAND: "Judged against the rest of its band? (FALSE: the rest of the book)",
            L_MKIND: "Materiality: a share, a dollar amount or none",
            L_MSHARE: "Materiality as a share of the book's GCO", L_MGCO: "Materiality line in GCO dollars",
            L_PKIND: "Profit line: each pocket's own test, points or dollars",
            L_PLINE: "Profit line (0.0025 = 0.25 points; dollars for the materiality line)",
            L_GCO_TOTAL: "The book's GCO (from the Run)"}


def col(n: int) -> str:
    return get_column_letter(n)


def q(s) -> str:
    """Words as an Excel string."""
    return '"' + str(s).replace('"', '""') + '"'


def text(*parts) -> str:
    """A formula joining words and expressions: a str part is words, a 1-tuple an Excel expression. Words
    are cut into pieces of 250 characters, under Excel's limit on a string inside a formula."""
    out = []
    for p in parts:
        if isinstance(p, tuple):
            out.append(p[0])
            continue
        s = str(p)
        out += [q(s[i:i + 250]) for i in range(0, len(s), 250)] or ['""']
    return "=" + "&".join(out or ['""'])


def num(v) -> str:
    return repr(float(v))


def sig(p: str) -> str:
    """Significant: a p-value below the one rounded bar (canon S36), never 1 - confidence written inline."""
    return f"AND(ISNUMBER({p}),{p}<{BAR})"


def signed(x: str) -> str:
    """engine._signed as a formula: "+0.12", "-0.80", and "0.00" for no gap."""
    return f'IF(ROUND({x},2)=0,"0.00",TEXT({x},"+0.00;-0.00"))'


def sure_words(c: float) -> str:
    """The confidence level as the tabs say it, the same as the formula in _live: "95% sure"."""
    pct = c * 100
    return f"{pct:.0f}% sure" if abs(round(pct) - pct) < 1e-9 else f"{pct:.1f}% sure"


@dataclass
class Live:
    """Where the live cells are, for the tab writers."""
    rows: dict = field(default_factory=dict)          # (kind, id(grid), band, seg, measure) -> _pockets row
    line_cell: dict = field(default_factory=dict)     # measure -> _live cell of its materiality line
    last: dict = field(default_factory=dict)          # words for what the last Run used, by _live row
    last_sentence: str = ""
    kind_of: dict = field(default_factory=dict)       # id(grid) -> "grids" | "three-way"

    def ref(self, kind: str, grid, band, seg, measure: str, column: int) -> str | None:
        r = self.rows.get((kind, id(grid), band, seg, measure))
        return None if r is None else f"'{POCKETS}'!${col(column)}${r}"

    def row(self, grid, band, seg, measure: str) -> int | None:
        return self.rows.get((self.kind_of.get(id(grid), "grids"), id(grid), band, seg, measure))

    def at(self, grid, band, seg, measure: str, column: int) -> str | None:
        r = self.row(grid, band, seg, measure)
        return None if r is None else f"'{POCKETS}'!${col(column)}${r}"


def ensure(wb, res) -> Live:
    """The two hidden sheets for this run, written once per workbook (every tab writer calls this)."""
    got = getattr(wb, "_cube_live", None)
    if got is not None and got[0] is res:
        return got[1]
    for t in (LIVE_SHEET, POCKETS):
        if t in wb.sheetnames:
            del wb[t]
    lv = Live()
    _write_live(wb, res, lv)
    _write_pockets(wb, res, lv)
    wb._cube_live = (res, lv)
    return lv


# --------------------------------------------------------------------------
# _live: each live setting as the number the formulas use


def _options(res) -> list[tuple[str, str, str, object]]:
    """(key, label, kind, value) for every option of a live setting, both as the list shows it and as its
    plain label. A suggested line is the multiple the last Run worked out, or not a number when this Run
    didn't work one out."""
    sug = getattr(res, "suggested", None) or {}
    out = []
    for s in control.load_settings():
        if s.key not in LIVE_KEYS:
            continue
        for o in s.options:
            v, kind = o.value, ""
            if s.key in ("worse_at", "better_at") and v == "luck":
                v = "=NA()" if sug.get(s.key) is None else float(sug[s.key])
            elif s.key == "materiality":
                kind, v = ("none", 0.0) if v == "none" else ("share", float(str(v).split("%")[0]) / 100)
            elif s.key == "revenue_line":
                kind, v = {"luck": ("test", 0.0), "materiality": ("dollars", 0.0)}.get(v, ("points", v))
            for label in dict.fromkeys((o.shown, o.label)):
                out.append((s.key, label, kind, v))
    return out


def _control_rows(wb) -> dict[str, int]:
    if control.SHEET not in wb.sheetnames:
        return {}
    ws = wb[control.SHEET]
    return {k: control.row_of(ws, k) for k in LIVE_KEYS if control.row_of(ws, k)}


def _write_live(wb, res, lv: Live) -> None:
    ws = wb.create_sheet(LIVE_SHEET)
    ws.sheet_state = "hidden"
    ws["B1"] = "The lines in use: what the formulas on the result tabs read"
    ws["B1"].font = Font(name="Arial", bold=True, size=14)
    ws["B2"] = ("Column C reads Control: your own value where one is typed, otherwise the option picked. Change "
                "Control, not this sheet. Column D is what the last Run used. The result tabs refer to column C "
                "by name (worse_at, better_at, confidence, significance_bar, judged_band, materiality_gco, "
                "profit_kind, profit_line).")
    ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    for i, h in enumerate(("Setting", "In use now", "The last Run used", "In words now", "In words at the last Run"),
                          start=2):
        ws.cell(row=4, column=i, value=h).font = Font(bold=True)
    ws.cell(row=4, column=11, value="Option (setting|label)").font = Font(bold=True)
    ws.cell(row=4, column=12, value="Kind").font = Font(bold=True)
    ws.cell(row=4, column=13, value="Value").font = Font(bold=True)
    for i, (key, label, kind, v) in enumerate(_options(res), start=OPT_FIRST):
        ws.cell(row=i, column=11, value=f"{key}|{label}")
        ws.cell(row=i, column=12, value=kind or None)
        ws.cell(row=i, column=13, value=v)
    b = res.config.benchmark
    total = res.total.rates.get("gco_rate")
    gco_total = abs(total.num) if total is not None else 0.0
    rows = _control_rows(wb)
    opt = lambda key, c, what="M": (f'INDEX(${what}${OPT_FIRST}:${what}${OPT_LAST},'     # noqa: E731
                                    f'MATCH("{key}|"&{c},$K${OPT_FIRST}:$K${OPT_LAST},0))')

    def cells(key):
        r = rows.get(key)
        return (f"{control.SHEET}!$C${r}", f"{control.SHEET}!$D${r}") if r else (None, None)

    def own_set(d):
        return f'AND({d}<>"",{d}<>"n/a")'

    last = _last(res, gco_total)

    def number(key, row):
        """Your own value where one is typed, else the option picked; a number Excel stored in place of
        the option's label ("95%" as 0.95) is that number."""
        c, d = cells(key)
        return f"=IF({own_set(d)},{d},IFERROR({opt(key, c)},IF(ISNUMBER({c}),{c},VALUE({c}))))" if c else last[row]

    formulas: dict[int, object] = {}
    formulas[L_WORSE] = number("worse_at", L_WORSE)
    formulas[L_BETTER] = number("better_at", L_BETTER)
    formulas[L_CONF] = number("confidence", L_CONF)
    # the one bar every significance test on the tabs compares with, rounded once (canon S36)
    formulas[L_BAR] = f"=ROUND(1-C{L_CONF},12)"
    c, _ = cells("compare_to")
    formulas[L_BAND] = f'=IFERROR({opt("compare_to", c)}="peers",NA())' if c else last[L_BAND]
    c, d = cells("materiality")
    if c:
        formulas[L_MKIND] = f'=IF({own_set(d)},"dollars",IFERROR({opt("materiality", c, "L")},NA()))'
        formulas[L_MSHARE] = f'=IF(C{L_MKIND}="share",{opt("materiality", c)},0)'
        formulas[L_MGCO] = f'=IF(C{L_MKIND}="none",0,IF(C{L_MKIND}="share",C{L_MSHARE}*C{L_GCO_TOTAL},{d}))'
    else:
        formulas[L_MKIND], formulas[L_MSHARE] = last[L_MKIND], last[L_MSHARE]
        formulas[L_MGCO] = (f'=IF(C{L_MKIND}="none",0,IF(C{L_MKIND}="share",C{L_MSHARE}*C{L_GCO_TOTAL},'
                            f'{num(last[L_MGCO])}))')
    c, d = cells("revenue_line")
    if c:
        formulas[L_PKIND] = (f'=IF({own_set(d)},"points",IFERROR({opt("revenue_line", c, "L")},'
                             f'IFERROR(IF(VALUE({c})>0,"points",NA()),NA())))')
        formulas[L_PLINE] = (f'=IF(C{L_PKIND}="points",IF({own_set(d)},{d},IFERROR({opt("revenue_line", c)},'
                             f'VALUE({c})))/100,IF(C{L_PKIND}="dollars",C{L_MGCO},0))')
    else:
        formulas[L_PKIND] = last[L_PKIND]
        formulas[L_PLINE] = (f'=IF(C{L_PKIND}="points",{num(last[L_PLINE])},IF(C{L_PKIND}="dollars",C{L_MGCO},0))')
    formulas[L_GCO_TOTAL] = gco_total
    words = {
        L_WORSE: f'"worse at "&TEXT(C{L_WORSE},"0.00")&"x"',
        L_BETTER: f'"better at "&TEXT(C{L_BETTER},"0.00")&"x"',
        L_CONF: (f'IF(ROUND(C{L_CONF}*100,9)=ROUND(C{L_CONF}*100,0),TEXT(C{L_CONF}*100,"0"),'
                 f'TEXT(C{L_CONF}*100,"0.0"))&"% sure"'),
        L_BAND: f'IF(C{L_BAND},"the rest of its band","the rest of the book")',
        L_MKIND: (f'IF(C{L_MKIND}="none","no floor",IF(C{L_MKIND}="share",TEXT(C{L_MSHARE}*100,"0.0")&"% of the '
                  f'book\'s GCO, $"&TEXT(C{L_MGCO},"#,##0"),"$"&TEXT(C{L_MGCO},"#,##0")&" of GCO"))'),
        L_PKIND: (f'IF(C{L_PKIND}="test","each pocket\'s own test",IF(C{L_PKIND}="points",TEXT(C{L_PLINE}*100,'
                  f'"0.00")&" points either way","a gap of $"&TEXT(C{L_PLINE},"#,##0")&" either way (the '
                  f'materiality line)"))'),
    }
    for r, label in L_LABELS.items():
        ws.cell(row=r, column=2, value=label)
        ws.cell(row=r, column=3, value=formulas[r])
        ws.cell(row=r, column=4, value=last.get(r))
        if r in words:
            ws.cell(row=r, column=5, value=f'=IFERROR({words[r]},"no usable answer on Control")')
            ws.cell(row=r, column=6, value=lv_words(last, r))
            lv.last[r] = lv_words(last, r)
        if r in L_NAMES:
            wb.defined_names[L_NAMES[r]] = DefinedName(L_NAMES[r], attr_text=f"'{LIVE_SHEET}'!$C${r}")
    ws.cell(row=L_SENTENCE, column=2, value="In use now, in one sentence (the result tabs show it at the top)")
    ws.cell(row=L_SENTENCE, column=3, value=(
        f'="Lines in use now, from Control: "&E{L_WORSE}&" and "&E{L_BETTER}&" (the loss line); profit line "&'
        f'E{L_PKIND}&"; "&E{L_CONF}&"; materiality "&E{L_MKIND}&"; judged against "&E{L_BAND}&"."'))
    lv.last_sentence = (f"Lines the last Run used: {lv.last[L_WORSE]} and {lv.last[L_BETTER]} (the loss line); "
                        f"profit line {lv.last[L_PKIND]}; {lv.last[L_CONF]}; materiality {lv.last[L_MKIND]}; "
                        f"judged against {lv.last[L_BAND]}.")
    ws.cell(row=L_SENTENCE, column=4, value=lv.last_sentence)
    wb.defined_names[IN_USE] = DefinedName(IN_USE, attr_text=f"'{LIVE_SHEET}'!$C${L_SENTENCE}")
    # each rate's materiality line, as engine._materiality_lines draws it: a share of the book's total for a
    # loss rate, the GCO dollar line for GCO and for profit, and no line for another rate under a dollar amount
    ws.cell(row=L_LINES - 1, column=2, value="Each rate's materiality line").font = Font(bold=True)
    ws.cell(row=L_LINES - 1, column=3, value="In use now").font = Font(bold=True)
    ws.cell(row=L_LINES - 1, column=4, value="The last Run used").font = Font(bold=True)
    ws.cell(row=L_LINES - 1, column=5, value="The rate's book total, for a share").font = Font(bold=True)
    r = L_LINES
    has_gco = "gco_rate" in res.total.rates
    for m in res.measures:
        if not m.is_rate:
            continue
        ws.cell(row=r, column=2, value=m.title)
        ws.cell(row=r, column=5, value=abs(res.total.rates[m.name].num))
        if m.name == "gco_rate" or (m.name in PROFIT and has_gco):
            f = f"=C{L_MGCO}"
        else:
            f = f'=IF(C{L_MKIND}="share",C{L_MSHARE}*E{r},IF(C{L_MKIND}="none",0,""))'
        ws.cell(row=r, column=3, value=f)
        ws.cell(row=r, column=4, value=res.materiality_line.get(m.name))
        lv.line_cell[m.name] = f"'{LIVE_SHEET}'!$C${r}"
        r += 1
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 40
    ws.column_dimensions["K"].width = 50


def _last(res, gco_total: float) -> dict:
    """What the last Run used, as the numbers _live's column C holds."""
    b = res.config.benchmark
    if b is None:
        return {}
    kind, v = b.materiality
    gco_line = res.materiality_line.get("gco_rate", 0.0 if kind == "none" else (v * gco_total if kind == "share"
                                                                              else v))
    line = engine.profit_line(b, gco_line)
    return {L_WORSE: b.worse_at, L_BETTER: b.better_at, L_CONF: b.confidence, L_BAR: stats.bar(b.confidence),
            L_BAND: b.compare_to == "peers", L_MKIND: kind, L_MSHARE: v if kind == "share" else 0.0,
            L_MGCO: gco_line, L_PKIND: line.kind, L_PLINE: line.value, L_GCO_TOTAL: gco_total}


def lv_words(last: dict, r: int) -> str | None:
    """The words _live's column E would give for the last Run's settings: the same formats as its formulas."""
    if not last:
        return None
    if r == L_WORSE:
        return f"worse at {last[L_WORSE]:.2f}x"
    if r == L_BETTER:
        return f"better at {last[L_BETTER]:.2f}x"
    if r == L_CONF:
        return sure_words(last[L_CONF])
    if r == L_BAND:
        return "the rest of its band" if last[L_BAND] else "the rest of the book"
    if r == L_MKIND:
        k = last[L_MKIND]
        if k == "none":
            return "no floor"
        if k == "share":
            return f"{last[L_MSHARE] * 100:.1f}% of the book's GCO, ${last[L_MGCO]:,.0f}"
        return f"${last[L_MGCO]:,.0f} of GCO"
    if r == L_PKIND:
        k = last[L_PKIND]
        if k == "test":
            return "each pocket's own test"
        if k == "points":
            return f"{last[L_PLINE] * 100:.2f} points either way"
        return f"a gap of ${last[L_PLINE]:,.0f} either way (the materiality line)"
    return None


# --------------------------------------------------------------------------
# _pockets: every pocket's numbers, and the formulas that read them


def test_words(s, hits) -> str:
    """book.which_test with the comparison's own shuffle count."""
    if s.test == engine.EXACT_TEST:
        return "exact test"
    if s.test == engine.Z_TEST:
        return "z test"
    if s.test == engine.SHUFFLE_TEST and s.shuffles:
        return f"shuffled: {hits:,} of {s.shuffles:,}" if hits is not None else f"{s.shuffles:,} shuffles"
    return ""


def _rest(s, parent) -> float | None:
    den = parent.den - s.den
    return (parent.num - s.num) / den if den else None


def reading_loss(g: str, p: str, few: str) -> str:
    """engine.reading_of for a pocket against a comparison, as a formula: the multiple against the loss line,
    and significance against the one bar."""
    return (f'IF({g}="","",IF({few},{q(engine.FEW)},IF(AND({g}>better_at,{g}<worse_at),{q(engine.IN_LINE)},'
            f'IF({g}>=worse_at,IF({sig(p)},{q(engine.WORSE)},{q(engine.UNSURE_WORSE)}),'
            f'IF({sig(p)},{q(engine.BETTER)},{q(engine.UNSURE_BETTER)})))))')


def reading_points(g: str, p: str, den: str) -> str:
    """engine.reading_gap for profit, as a formula: each pocket's own test, a line in points, or a line in
    dollars (the gap times the pocket's booked dollars)."""
    worse_or = f'IF({g}<0,{q(engine.WORSE)},{q(engine.BETTER)})'
    unsure_or = f'IF({g}<0,{q(engine.UNSURE_WORSE)},{q(engine.UNSURE_BETTER)})'
    size = f'IF(profit_kind="points",ABS({g}),ABS({g}*{den}))'
    return (f'IF({g}="","",IF(profit_kind="test",IF(OR(NOT({sig(p)}),{g}=0),{q(engine.IN_LINE)},{worse_or}),'
            f'IF(OR({g}=0,{size}<profit_line*(1-1E-9)),{q(engine.IN_LINE)},IF({sig(p)},{worse_or},{unsure_or}))))')


def literal(word: str, gap: str, p: str, dollars: str, by_band: str) -> str:
    """engine.literal as a formula: the gap against the comparison that decides it, in points of booked
    dollars, and the dollars it comes to."""
    against = f'IF({by_band},"its band","the book")'
    inside = (f'IF(profit_kind="test",{signed(gap + "*100")}&" points against "&{against}&", "&IF({p}="",'
              f'"not tested","not significant"),IF(profit_kind="points","within "&TEXT(profit_line*100,"0.00")&'
              f'" points of "&{against}&" ("&{signed(gap + "*100")}&")","within $"&TEXT(profit_line,"#,##0")&'
              f'" of "&{against}&" ("&{signed(gap + "*100")}&" points)"))')
    past = (f'IF(LEFT({word},5)="worse","short of ","ahead of ")&{against}&" by "&TEXT(ABS({gap})*100,"0.00")&'
            f'" points"&IF({dollars}="",""," ($"&TEXT(ABS({dollars}),"#,##0")&")")&'
            f'IF(RIGHT({word},15)="not significant"," (not significant)","")')
    return f'IF({word}="","",IF({word}={q(engine.IN_LINE)},{inside},{past}))'


def _write_pockets(wb, res, lv: Live) -> None:
    ws = wb.create_sheet(POCKETS)
    ws.sheet_state = "hidden"
    ws["A1"] = ("Every pocket's numbers from the last Run (columns A to U), and the formulas that read them with the "
                "lines on Control (V onwards). The result tabs point here, so one set of formulas decides for "
                "every tab.")
    ws["A1"].font = Font(bold=True)
    ws["A2"] = ("A p-value doesn't depend on the confidence level: the lines only decide what it is compared with. "
                "The gaps, p-values and dollars are as of the last Run; change band edges, segments, the floors or "
                "the allowance for many tests and Run again.")
    for c, h in POCKET_HEADS.items():
        cell = ws.cell(row=P_FIRST - 1, column=c, value=h)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if c >= P_BY_BAND:
            cell.fill = PatternFill("solid", fgColor="E4DFD5")
    b = res.config.benchmark
    if b is None:
        return
    names = {x.name: x.field for x in res.config.bands}
    names.update({x.name: x.field for x in res.config.dimensions})
    if res.config.split:
        sf = res.config.split[0]
        names.update({f"{d.name} / {sf}": f"{d.field} / {sf}" for d in res.config.dimensions})
    rates = [m for m in res.measures if m.is_rate]
    r = P_FIRST
    for kind, grids in (("grids", res.grids), ("three-way", res.three_way)):
        for g in grids:
            lv.kind_of[id(g)] = kind
            gname = f"{names.get(g.band, g.band)} x {names.get(g.dimension, g.dimension)}"
            mates: dict = {}
            for (bl, _), _c in g.inner():
                mates[bl] = mates.get(bl, 0) + 1
            for (bl, dl), c in g.inner():
                for m in rates:
                    s = c.rates[m.name]
                    tested = s.reading_topline not in (engine.THIN, engine.FEW)
                    few = (not m.in_points) and s.events < b.min_events
                    vals = {P_KIND: kind, P_GRID: gname, P_BAND: bl, P_SEG: dl, P_MEASURE: m.name, P_TITLE: m.title,
                            P_LOANS: s.units, P_BOOKED: s.den, P_FEW: few, P_ALONE: mates[bl] == 1,
                            P_VS_BOOK: s.vs_rest, P_P_BOOK: s.p_book, P_VS_BAND: s.vs_band, P_P_BAND: s.p_band,
                            P_EX_BOOK: s.excess, P_EX_BAND: s.excess_band,
                            P_TEST_BOOK: test_words(s, s.hits_book) if tested else None,
                            P_TEST_BAND: test_words(s, s.hits_band) if tested else None,
                            P_REST_BOOK: _rest(s, res.total.rates[m.name]),
                            P_REST_BAND: _rest(s, g.cells[(bl, engine.ALL)].rates[m.name])}
                    for cc, v in vals.items():
                        ws.cell(row=r, column=cc, value=v)
                    ws.cell(row=r, column=P_LINE, value=f"={lv.line_cell[m.name]}")
                    R = lambda cc: f"${col(cc)}{r}"         # noqa: E731

                    def pick(band_c, book_c):
                        return (f'IF({R(P_BY_BAND)},IF({R(band_c)}="","",{R(band_c)}),'
                                f'IF({R(book_c)}="","",{R(book_c)}))')

                    ws.cell(row=r, column=P_BY_BAND, value=f"=AND(judged_band,NOT({R(P_ALONE)}))")
                    if m.in_points:
                        rb = reading_points(R(P_VS_BOOK), R(P_P_BOOK), R(P_BOOKED))
                        rn = reading_points(R(P_VS_BAND), R(P_P_BAND), R(P_BOOKED))
                    else:
                        rb = reading_loss(R(P_VS_BOOK), R(P_P_BOOK), R(P_FEW))
                        rn = reading_loss(R(P_VS_BAND), R(P_P_BAND), R(P_FEW))
                    ws.cell(row=r, column=P_READ_BOOK, value=f"={rb}")
                    ws.cell(row=r, column=P_READ_BAND, value=f"={rn}")
                    ws.cell(row=r, column=P_FLAG, value=f"=IF({R(P_BY_BAND)},{R(P_READ_BAND)},{R(P_READ_BOOK)})")
                    ws.cell(row=r, column=P_DOLLARS, value=f"={pick(P_EX_BAND, P_EX_BOOK)}")
                    ws.cell(row=r, column=P_MATERIAL, value=(
                        f'=IF(OR({R(P_LINE)}="",{R(P_DOLLARS)}=""),"",IF(AND({R(P_DOLLARS)}>0,'
                        f'{R(P_DOLLARS)}>={R(P_LINE)}),"yes","below the line"))'))
                    ws.cell(row=r, column=P_GAP, value=f"={pick(P_VS_BAND, P_VS_BOOK)}")
                    ws.cell(row=r, column=P_P, value=f"={pick(P_P_BAND, P_P_BOOK)}")
                    ws.cell(row=r, column=P_SAID, value=(
                        f"={literal(R(P_FLAG), R(P_GAP), R(P_P), R(P_DOLLARS), R(P_BY_BAND))}" if m.in_points
                        else f"={R(P_FLAG)}"))
                    t = pick(P_TEST_BAND, P_TEST_BOOK)
                    ws.cell(row=r, column=P_TEST, value=(
                        f'={t}&IF(AND(judged_band,{R(P_ALONE)}),IF({t}="","","; ")&{q(ALONE)},"")'))
                    ws.cell(row=r, column=P_REST, value=f"={pick(P_REST_BAND, P_REST_BOOK)}")
                    ws.cell(row=r, column=P_OTHER, value=(
                        f'=IF({R(P_BY_BAND)},IF({R(P_EX_BOOK)}="","",{R(P_EX_BOOK)}),'
                        f'IF({R(P_EX_BAND)}="","",{R(P_EX_BAND)}))'))
                    lv.rows[(kind, id(g), bl, dl, m.name)] = r
                    r += 1
    ws.freeze_panes = f"A{P_FIRST}"


def count_formula(measure: str, criteria: list[tuple[int, str]], kind: str = "grids") -> str:
    """COUNTIFS over _pockets: one kind of grid, one rate, and more criteria (column, criterion)."""
    rng = lambda c: f"'{POCKETS}'!${col(c)}${P_FIRST}:${col(c)}$1048576"        # noqa: E731
    parts = [f'{rng(P_KIND)},{q(kind)}', f'{rng(P_MEASURE)},{q(measure)}']
    parts += [f"{rng(c)},{crit}" for c, crit in criteria]
    return f"COUNTIFS({','.join(parts)})"


def pockets_range(c: int) -> str:
    return f"'{POCKETS}'!${col(c)}${P_FIRST}:${col(c)}$1048576"
