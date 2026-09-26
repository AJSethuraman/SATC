"""The Prevalence tab (fix 3.12): loans and booked dollars per group, pocket by
pocket, with no test attached.

docs/statistics.md B8: "Loans and dollars per group per pocket, with no p-value
attached. A count of the book, so it needs no holdout and no confidence level.
It is the first finding whenever the line of business believes a pattern is
rare."

Which groups:
- the column that splits the pockets: its two halves (each pocket cut at its own
  median, as the Split tab cuts it) or its values (a category);
- each new column made on Control: its bands. The edges are the ones the Run cut
  it at when it is a band column, else the ones typed on Columns ("every 0.5"
  worked out over its values the way a band column's would be). A new column
  with no edges anywhere is not counted by band, and the tab says so rather
  than choosing edges for it.

The pockets are the grids' own, counted again from the loans the Run used, and
each pocket's count is tied out to the grid's before anything is written: a
count that does not add up to the grid is not shown.
"""

from __future__ import annotations

import math
import statistics
from dataclasses import dataclass

from openpyxl.styles import Alignment, Font, PatternFill

from . import engine

SHEET = "Prevalence"
INK, SLATE, PAPER, CANVAS = "16130F", "57534B", "FFFFFF", "F4F1EC"       # the workbook's colours (book.py)
FIRST_COL = 2
GROUP_COL = 6            # the first group's column: band, segment, the pocket's loans and dollars before it


@dataclass
class Grouping:
    """One way of dividing every pocket."""
    column: str
    kind: str                      # "halves" | "values" | "bands"
    title: str
    edges: tuple = ()              # bands only
    skip_band: str | None = None   # bands only: a grid cut by this column is its own grouping, so is left out


def rows_run(res) -> list[dict]:
    """The loans the Run counted: every loan in the extract, with any new columns."""
    return res.table.rows if res.table is not None else []


def edges_of(res, column: str) -> tuple[tuple[float, ...] | None, str]:
    """A number column's band edges for this run, and where they came from: the
    edges it was cut at, the edges typed on Columns, or None when there are none."""
    cfg = res.config
    band = next((b for b in cfg.bands if b.field == column), None)
    if band is not None and band.name in res.band_edges:
        return tuple(res.band_edges[band.name]), "the bands it was cut into"
    typed = str((getattr(res, "typed_edges", None) or {}).get(column) or "").strip()
    if not typed:
        return None, ""
    if typed.lower().startswith("every"):
        try:
            w = float(typed.lower().removeprefix("every").split()[0].replace(",", ""))
        except (ValueError, IndexError):
            return None, ""
        rule = cfg.missing.get(column)
        vals = [v for v in (engine.classify_number(r.get(column), rule)[0] for r in rows_run(res)) if v is not None]
        if w <= 0 or not vals:
            return None, ""
        lo, hi = min(vals), max(vals)
        pts, x = [], math.floor(lo / w) * w + w           # as a band column's "every" is cut (book._band_widths)
        while x <= hi and len(pts) < 50:
            pts.append(round(x, 10))
            x += w
        return tuple(pts) or (round(math.floor(lo / w) * w + w, 10),), f"{typed} on Columns"
    try:
        pts = tuple(float(x) for x in typed.replace(";", ",").split(",") if x.strip())
    except ValueError:
        return None, ""
    if not pts or any(b <= a for a, b in zip(pts, pts[1:])):
        return None, ""
    return pts, "the band edges on Columns"


def groupings(res) -> tuple[list[Grouping], list[str]]:
    """Each grouping this run can be counted by, and a note for each new column
    that can't be (no edges)."""
    cfg = res.config
    out, notes = [], []
    if cfg.split:
        sf, how = cfg.split
        out.append(Grouping(sf, "halves", f"{sf}, each pocket cut at its own median") if how == "own_median"
                   else Grouping(sf, "values", f"{sf}, by value"))
    for d in res.derived:
        edges, said = edges_of(res, d.name)
        if edges is None:
            notes.append(f"{d.name} = {d.text()} has no band edges, so it isn't counted by band here. Type its "
                         f"edges on Columns to count it.")
            continue
        shown = "; ".join(engine._fmt(x) for x in edges)
        band = next((b.name for b in cfg.bands if b.field == d.name), None)
        out.append(Grouping(d.name, "bands", f"{d.name} = {d.text()}, by its bands ({shown}: {said})",
                            edges=edges, skip_band=band))
    return out, notes


def _labels_by_band(res, rows) -> dict[str, list[str]]:
    """Each band column's label on every loan, cut exactly as the grids cut it."""
    cfg, out = res.config, {}
    for b in cfg.bands:
        read = [engine.classify_number(r.get(b.field), cfg.missing.get(b.field)) for r in rows]
        edges = tuple(res.band_edges[b.name])
        seen = [v for v, why in read if why is None]
        labels = engine.band_labels(edges, min(seen), max(seen)) if seen else engine.band_labels(edges)
        out[b.name] = [engine.band_of(v, edges, labels) if why is None else engine.REASON_LABEL[why]
                       for v, why in read]
    return out


def _bands_of(values: list[tuple], edges: tuple) -> tuple[list[str], list[str]]:
    """Each loan's band of a column, and the bands in order."""
    seen = [v for v, why in values if why is None]
    labels = engine.band_labels(edges, min(seen), max(seen)) if seen else engine.band_labels(edges)
    got = [engine.band_of(v, edges, labels) if why is None else engine.REASON_LABEL[why] for v, why in values]
    present = set(got)
    return got, [x for x in labels if x in present] + [x for x in engine._order(got) if x not in labels]


def count(res, grid, grouping: Grouping, rows=None, bands=None) -> tuple[list, dict, list[str]] | None:
    """Loans and booked dollars per pocket and group for one grid: (the groups in
    order, {(band, segment): {group: [loans, dollars]}}, the words each group is
    shown as), or None when the count does not tie out to the grid."""
    cfg = res.config
    rows = rows_run(res) if rows is None else rows
    bands = _labels_by_band(res, rows) if bands is None else bands
    bl = bands[grid.band]
    dname = grid.dimension
    dim = next(d for d in cfg.dimensions if d.name == dname)
    dl = [engine.classify_text(r.get(dim.field), cfg.missing.get(dim.field)) for r in rows]
    col, rule = grouping.column, cfg.missing.get(grouping.column)
    if grouping.kind == "halves":
        vals = [engine.classify_number(r.get(col), rule)[0] for r in rows]
        pools: dict[tuple, list[float]] = {}
        for b, d, v in zip(bl, dl, vals):
            if v is not None:
                pools.setdefault((b, d), []).append(v)
        med = {k: statistics.median(v) for k, v in pools.items()}           # as engine._split cuts them
        groups = [engine.NO_SPLIT_VALUE if v is None else engine.HIGH if v > med[(b, d)] else engine.LOW
                  for b, d, v in zip(bl, dl, vals)]
        order = [x for x in (engine.HIGH, engine.LOW, engine.NO_SPLIT_VALUE) if x in set(groups)]
        words = {engine.HIGH: "High half", engine.LOW: "Low half", engine.NO_SPLIT_VALUE: f"No {col}"}
        shown = [words[g] for g in order]
    elif grouping.kind == "values":
        groups = [engine.classify_text(r.get(col), rule) for r in rows]
        order = engine._order(groups)
        shown = list(order)
    else:
        groups, order = _bands_of([engine.classify_number(r.get(col), rule) for r in rows], grouping.edges)
        shown = list(order)
    booked_rule = cfg.missing.get(cfg.booked)
    out: dict[tuple, dict] = {}
    for b, d, g, r in zip(bl, dl, groups, rows):
        cell = out.setdefault((b, d), {})
        got = cell.setdefault(g, [0, 0.0])
        got[0] += 1
        v = engine.classify_number(r.get(cfg.booked), booked_rule)[0]
        if v is not None:
            got[1] += v
    # the tie-out: every pocket's groups add up to the grid's pocket, and the halves to the Split tab's
    inner = dict(grid.inner())
    if set(out) != set(inner):
        return None
    for k, c in inner.items():
        if sum(x[0] for x in out[k].values()) != c.rows:
            return None
        if grouping.kind == "halves" and grid.split_cells:
            for g, x in out[k].items():
                sc = grid.split_cells.get((k[0], k[1], g))
                if sc is None or sc.rows != x[0]:
                    return None
    return order, out, shown


def write(wb, res) -> None:
    """The tab, after Split: one block per grouping and grid. Written only when
    the run has a split column or a new column."""
    gs, notes = groupings(res)
    if not gs and not notes:
        return
    names = {b.name: b.field for b in res.config.bands}
    names.update({d.name: d.field for d in res.config.dimensions})
    rows = rows_run(res)
    bands = _labels_by_band(res, rows)
    # every count first, so the tab is only as wide as its widest block
    blocks = [(g, [(grid, count(res, grid, g, rows, bands)) for grid in res.grids
                   if g.skip_band is None or grid.band != g.skip_band]) for g in gs]
    widest = max((len(got[0]) for _, done in blocks for _, got in done if got is not None), default=1)
    last = max(GROUP_COL + 2 * widest - 1, 9)
    ws = wb.create_sheet(SHEET)
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=last)
    ws.cell(row=1, column=2, value="Prevalence: a count, not a test").font = Font(name="Arial", bold=True, size=16,
                                                                                   color=PAPER)
    for c in range(2, last + 1):
        ws.cell(row=1, column=c).fill = PatternFill("solid", fgColor=INK)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=2, end_row=2, end_column=last)
    sub = ws.cell(row=2, column=2, value="How many loans and booked dollars sit in each group, pocket by pocket. "
                                         "Nothing here is tested, so nothing reads better or worse: it shows how "
                                         "common a group is, not whether it goes bad.")
    sub.font = Font(name="Calibri", size=10, color=SLATE)
    sub.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[2].height = 32
    ws.sheet_view.showGridLines = False
    r = 4
    for n in notes:
        ws.cell(row=r, column=2, value=n).font = Font(name="Calibri", italic=True, color=SLATE)
        r += 1
    if notes:
        r += 1
    for g, done in blocks:
        ws.cell(row=r, column=2, value=g.title).font = Font(name="Calibri", bold=True, size=13)
        r += 2
        for grid, got in done:          # a grid cut by the new column itself is left out: its bands are the grid's
            head = f"{names.get(grid.band, grid.band)} x {names.get(grid.dimension, grid.dimension)}"
            ws.cell(row=r, column=2, value=head).font = Font(name="Calibri", bold=True)
            if got is None:
                ws.cell(row=r + 1, column=2, value="Not shown: the count didn't add up to this grid's loans.")
                r += 3
                continue
            order, per, shown = got
            r = _block(ws, r + 1, grid, names, order, per, shown)
            r += 1
        r += 1
    for col, w in zip("ABCDE", (2, 17, 16, 10, 14)):
        ws.column_dimensions[col].width = w
    from openpyxl.utils import get_column_letter
    for c in range(GROUP_COL, last + 1):
        ws.column_dimensions[get_column_letter(c)].width = 11 if (c - GROUP_COL) % 2 == 0 else 14
    ws.freeze_panes = "D4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _block(ws, r: int, grid, names: dict, order: list, per: dict, shown: list[str]) -> int:
    """One grid: a header, a row per pocket, the whole grid, and each group's
    share of it. Returns the row after it."""
    band, seg = names.get(grid.band, grid.band), names.get(grid.dimension, grid.dimension)
    heads = [band, seg, "Loans", "Booked dollars"]
    for i, h in enumerate(heads):
        _heading(ws, r, FIRST_COL + i, h)
        _heading(ws, r + 1, FIRST_COL + i, None)
    for j, s in enumerate(shown):
        c = GROUP_COL + 2 * j
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        _heading(ws, r, c, s)
        _heading(ws, r, c + 1, None)
        _heading(ws, r + 1, c, "Loans")
        _heading(ws, r + 1, c + 1, "Booked dollars")
    r += 2
    totals = {g: [0, 0.0] for g in order}
    for bl in grid.band_labels:
        for dl in grid.dim_labels:
            got = per.get((bl, dl))
            if got is None:
                continue
            loans = sum(x[0] for x in got.values())
            dollars = math.fsum(x[1] for x in got.values())
            vals = [bl, dl, loans, dollars]
            for g in order:
                x = got.get(g, [0, 0.0])
                vals += [x[0], x[1]]
                totals[g][0] += x[0]
                totals[g][1] += x[1]
            _row(ws, r, vals)
            r += 1
    loans = sum(x[0] for x in totals.values())
    dollars = math.fsum(x[1] for x in totals.values())
    _row(ws, r, ["Every pocket", "", loans, dollars] + [v for g in order for v in totals[g]], bold=True)
    r += 1
    share = ["Share of the grid", "", None, None]
    for g in order:
        share += [totals[g][0] / loans if loans else None, totals[g][1] / dollars if dollars else None]
    _row(ws, r, share, bold=True, pct=True)
    return r + 1


def _heading(ws, r: int, c: int, text) -> None:
    cell = ws.cell(row=r, column=c, value=text)
    cell.font = Font(name="Calibri", bold=True, color=PAPER)
    cell.fill = PatternFill("solid", fgColor=INK)
    cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if c >= FIRST_COL + 2 else None)


def _row(ws, r: int, vals: list, bold: bool = False, pct: bool = False) -> None:
    for i, v in enumerate(vals):
        c = ws.cell(row=r, column=FIRST_COL + i, value=v)
        if i >= 2:
            c.number_format = "0.0%" if pct else "#,##0"
        if bold:
            c.font = Font(name="Calibri", bold=True)
            c.fill = PatternFill("solid", fgColor=CANVAS)
