"""A run held to a pre-spec (fix 3.15): prespec.py read at Run, echoed on
Check, compared with what the run used, and the holdout's use counted in the
Log.

prespec.py reads the file and returns data; this is where a Run meets it.

- The file is named in one cell on Control, under the new columns
  (control.write_prespec). Blank means no pre-spec. A file that is not there,
  or that prespec.load refuses, stops the Run, each problem named by that cell.
- Check echoes what the file says, and the git commit it was read from, or why
  there is none (prespec.provenance). A pre-spec edited since its commit says so,
  and so does one dated after the day of the run. Its groups are named as the
  tabs name them, the lowest and highest from the tested column's smallest and
  largest value among the loans run (prespec.named), so one group has one name.
- Check says, one line each, every place the run differs from the pre-spec
  (prespec.deviations), and the Log labels the run "Deviates from pre-spec".
- The holdout: how many of the extract's loans were made inside the pre-spec's
  holdout range (prespec.holdout_touch). A run that touched it says so in the
  Log, on a line starting "Touched the holdout:", and Check counts those lines,
  so how many runs have looked at the holdout stays visible.

The test itself (capability 4b, and 4e's concentration; `run_test`): the
pre-spec's column cut at its bins, each group against its reference group,
inside pockets cut by its strata, once on the development range and once on the
holdout, as docs/statistics.md B3 to B6 write them (kgroups.py does the
arithmetic). The tab is confirm_tab.py. Loans made outside both ranges, or with
no readable date, tested column or outcome, are left out of this test only and
counted on Check; every other tab still uses every loan.

What the run used, as the pre-spec's keys (`in_use`). When the test ran, it
used the pre-spec's column, bins, reference and strata, and held itself to the
pre-spec's holdout range: `in_use` reports those, as the test used them, so a
run that did what its pre-spec says differs nowhere. When it couldn't run (no
readable dates, say), `in_use` reports what the other tabs used:
- column: the column that splits the pockets; with no split, the pre-spec's
  column if the run cuts it into bands; otherwise not set.
- bins: for a band column, the edges it was cut at; for the split column, the
  edges typed for it on Columns (the Prevalence tab's bands). The Split tab
  itself tests each pocket's halves, not bins.
- reference: what each group was compared with. The split's is the low half of
  each pocket; a band's is the rest of its band, or of the book. None of those
  is one of the pre-spec's groups.
- strata: the band and segment columns, the tested column left out.
- confidence: as Control has it.
- holdout: the origination range of the loans the run used. When every one of
  them was made inside the pre-spec's holdout, the run used the holdout;
  otherwise the range it used is reported, since a confirmatory run on loans
  outside the holdout is not the one pre-specified.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

import bisect

from . import control, engine, kgroups, perm, prespec

HOLDOUT_MARK = "Touched the holdout:"       # how a Log line says a run touched the holdout; Check counts them
DEVIATES = "Deviates from pre-spec"          # how a Log line labels a run that differs from its pre-spec
LOW_HALF = "the low half of each pocket"


@dataclass
class State:
    spec: prespec.PreSpec
    path: Path
    cell: str
    provenance: dict
    deviations: list[str] = field(default_factory=list)
    touched: int | None = None               # extract loans made in the holdout; None when it couldn't be checked
    first: date | None = None                # the first and last origination date among them
    last: date | None = None
    undated: int = 0                         # extract loans with no readable origination date
    unchecked: str | None = None             # why the holdout couldn't be checked
    runs_before: int = 0                     # Log lines already saying a run touched the holdout
    failed: str | None = None                # something went wrong working the state out, in words
    ran_on: date = field(default_factory=date.today)     # the day of the run, which the pre-spec can't postdate
    test: "Test | None" = None               # the confirmatory test itself (4b, 4e)

    @property
    def name(self) -> str:
        return self.path.name

    @property
    def runs(self) -> int:
        """Runs in this workbook's Log that touched the holdout, this one included."""
        return self.runs_before + (1 if self.touched else 0)


# --------------------------------------------------------------------------
# Reading the cell and the file (book.read_book)


def read(wb, book: Path, problems: list[str]) -> dict | None:
    """The pre-spec named on Control, loaded, as {spec, path, cell}; None when
    the cell is blank. Every problem is added to `problems`, named by the cell."""
    if control.SHEET not in wb.sheetnames:
        return None
    text, cell = control.read_prespec(wb[control.SHEET])
    if text is None:
        return None
    p = Path(text).expanduser()
    if not p.is_absolute():
        p = Path(book).resolve().parent / p
    if not p.is_file():
        problems.append(f"{cell}: there's no pre-spec file at {p}. Fix the path, or clear the cell for a run "
                        f"without a pre-spec.")
        return None
    try:
        spec = prespec.load(p)
    except prespec.PreSpecError as exc:
        for x in exc.problems:
            x = x.removeprefix(f"{p}: ")
            problems.append(f"{cell}: in the pre-spec {p.name}, {_plain(x)}")
        return None
    return {"spec": spec, "path": p, "cell": cell}


def _plain(text: str) -> str:
    return text.replace("`", '"')


# --------------------------------------------------------------------------
# After the engine has run (checks.attach)


def state(book, about: dict, res) -> State | None:
    got = (about or {}).get("_prespec")
    if not got:
        return None
    ps = got["spec"]
    st = State(spec=ps, path=Path(got["path"]), cell=got["cell"], provenance=prespec.provenance(got["path"]))
    try:
        ps = st.spec = prespec.named(ps, *column_range(res, ps.column))
        st.test = run_test(res, ps)
        st.deviations = [_plain(x) for x in prespec.deviations(ps, in_use(res, ps, st.test), where="in this run")]
        dates, why = origination_dates(res)
        if dates is None:
            st.unchecked = why
        else:
            st.touched, st.first, st.last = prespec.holdout_touch(ps, dates)
            st.undated = sum(1 for d in dates if d is None)
        st.runs_before = holdout_runs(book)
    except Exception as exc:                     # a line on Check, never a failed run after the engine ran
        st.failed = f"the run couldn't be compared with the pre-spec ({exc})"
    return st


def in_use(res, ps: prespec.PreSpec, test: "Test | None" = None) -> dict[str, Any]:
    """What the run used, as prespec.deviations reads it (the module docstring says how each is found)."""
    from .prevalence import edges_of
    cfg = res.config
    if test is not None and test.problem is None:
        b = cfg.benchmark
        return {"column": test.column, "bins": list(test.bins), "reference": test.groups[test.ref],
                "strata": list(test.strata), "confidence": b.confidence if b is not None else None,
                "holdout": test.holdout.range}
    b = cfg.benchmark
    cut = [x.field for x in cfg.bands] + [x.field for x in cfg.dimensions]
    out: dict[str, Any] = {"column": None, "bins": None, "reference": None}
    if cfg.split:
        sf, how = cfg.split
        out["column"] = sf
        if how == "own_median":
            out["bins"] = list(edges_of(res, sf)[0] or ()) or None
            out["reference"] = LOW_HALF
    elif ps.column in [x.field for x in cfg.bands]:
        out["column"] = ps.column
        out["bins"] = list(edges_of(res, ps.column)[0] or ()) or None
    if out["column"] is not None and out["reference"] is None and b is not None:
        out["reference"] = "the rest of its band" if b.compare_to == "peers" else "the rest of the book"
    out["strata"] = [c for c in cut if c != out["column"]]
    out["confidence"] = b.confidence if b is not None else None
    out["holdout"] = run_range(res, ps)
    return out


def column_range(res, column: str) -> tuple[float | None, float | None]:
    """A number column's smallest and largest value among the loans run, read as
    the tabs read it to name their lowest and highest bands; (None, None) when no
    loan has one."""
    from .prevalence import rows_run
    rule = res.config.missing.get(column)
    seen = [v for v, why in (engine.classify_number(r.get(column), rule) for r in rows_run(res)) if why is None]
    return (min(seen), max(seen)) if seen else (None, None)


def _reader(res):
    """How to read the origination dates, or (None, why not)."""
    col = res.config.origination_date
    table = res.table
    if not col:
        return None, "no column is marked Origination date on Columns"
    if table is None or col not in table.columns:
        return None, f"{col} isn't in the extract"
    try:
        return engine._date_reader(table, col, "when each loan was made"), None
    except engine.NothingToCut as exc:            # dates that read two ways (DataRefused)
        return None, _plain(str(exc))


def _date_or_none(v) -> date | None:
    return v if isinstance(v, date) else None


def origination_dates(res) -> tuple[list[date | None] | None, str | None]:
    """Every extract loan's origination date (None where it can't be read), or
    (None, why) when there is no date to read."""
    read, why = _reader(res)
    if read is None:
        return None, why
    col = res.config.origination_date
    return [_date_or_none(read(r.get(col))) for r in res.table.rows], None


def run_range(res, ps: prespec.PreSpec):
    """The origination range the run used, as a holdout: the pre-spec's own
    holdout when every loan run was made inside it, else the range run; None
    when the run's loans carry no readable origination date."""
    from .prevalence import rows_run
    read, _ = _reader(res)
    if read is None:
        return None
    col = res.config.origination_date
    seen = [d for d in (_date_or_none(read(r.get(col))) for r in rows_run(res)) if d is not None]
    if not seen:
        return None
    first, last = min(seen), max(seen)
    if ps.holdout.holds(first) and ps.holdout.holds(last):
        return ps.holdout
    return prespec.DateRange(first, last)


def holdout_runs(book) -> int:
    """How many runs this workbook's Log already records as touching the holdout."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(book, read_only=True)
    except Exception:
        return 0
    try:
        if "Log" not in wb.sheetnames:
            return 0
        return sum(1 for (v,) in wb["Log"].iter_rows(min_col=2, max_col=2, values_only=True)
                   if isinstance(v, str) and v.startswith(HOLDOUT_MARK))
    finally:
        wb.close()


# --------------------------------------------------------------------------
# The test (4b), and the concentration on the holdout (4e)

DEVELOPMENT, HOLDOUT = "development", "holdout"
#: Why a loan is left out of the test, in the order each is checked; a loan is counted under the first only
NO_DATE, OUTSIDE, NO_VALUE, NO_OUTCOME = ("no readable origination date", "made outside both ranges",
                                          "no readable value of the column tested", "no readable outcome")


@dataclass
class RangeTest:
    """The test on one origination range."""
    name: str                                   # DEVELOPMENT or HOLDOUT
    range: prespec.DateRange                    # the range the loans were held to
    loans: list[int]                            # per group
    bad: list[int]                              # per group
    pockets: list[kgroups.Pocket]               # one per pocket holding a loan of this range
    association: kgroups.Association            # B3 and B4
    fit: kgroups.Fit                            # B5
    mh_gap: float | None                        # the largest gap between B5's odds ratio and Mantel-Haenszel's
    group_of: list[int] = field(default_factory=list, repr=False)     # per loan in this range, for B6
    bad_of: list[int] = field(default_factory=list, repr=False)
    gco_of: list[float | None] = field(default_factory=list, repr=False)

    @property
    def n(self) -> int:
        return sum(self.loans)

    @property
    def n_bad(self) -> int:
        return sum(self.bad)


def concentration(*ranges: RangeTest) -> list[kgroups.Concentration]:
    """B6 per group, on the loans of the ranges given together. 4e gives it the holdout alone."""
    group = [k for r in ranges for k in r.group_of]
    bad = [y for r in ranges for y in r.bad_of]
    gco = [g for r in ranges for g in r.gco_of]
    return [kgroups.concentration(group, bad, gco, k) for k in range(len(ranges[0].loans))]


@dataclass
class Test:
    column: str
    bins: tuple[float, ...]
    groups: tuple[str, ...]                     # named as the tabs name them (prespec.named)
    ref: int
    strata: tuple[str, ...]
    scores: tuple[float, ...]                   # B4's: 1, 2, ... K, lowest group first
    development: RangeTest | None = None
    holdout: RangeTest | None = None
    left_out: dict[str, int] = field(default_factory=dict)
    gco_unread: int = 0                         # holdout loans in the test with no readable GCO (B6's dollars only)
    problem: str | None = None                  # why the test couldn't be run, in words
    method: str = kgroups.CONDITIONAL

    def ranges(self) -> list[RangeTest]:
        return [r for r in (self.development, self.holdout) if r is not None]

    def concentration(self) -> list[kgroups.Concentration]:
        """4e: B6 on the holdout only (statistics.md B6: "on the holdout only")."""
        return concentration(self.holdout)


def outcome_of(raw, m, rules) -> int | None:
    """One loan's outcome as the engine reads it for the share of loans (engine.run, flagwt): 1, 0, or None
    when it can't be read."""
    if m.flag_is is not None:
        if engine.is_blank(raw) or engine.classify_text(raw, rules.get(m.flag)) == engine.MISSING_RULE_LABEL:
            return None
        return 1 if engine.cell_text(raw) == engine.cell_text(m.flag_is) else 0
    t, why = engine.classify_number(raw, rules.get(m.flag))
    if why or t not in (0.0, 1.0):
        return None
    return int(t)


def _stratum_labels(res, strata: tuple[str, ...], rows) -> tuple[list[tuple] | None, str | None]:
    """Each loan's pocket: its label on every strata column, cut exactly as the grids cut it. A number column
    is read in the bands the grids use; a category by its values. (None, why) when a strata column isn't cut."""
    from .prevalence import _labels_by_band
    cfg = res.config
    by_band = _labels_by_band(res, rows)
    cols = []
    for c in strata:
        band = next((b for b in cfg.bands if b.field == c), None)
        dim = next((d for d in cfg.dimensions if d.field == c), None)
        if band is not None:
            cols.append(by_band[band.name])
        elif dim is not None:
            cols.append([engine.classify_text(r.get(c), cfg.missing.get(c)) for r in rows])
        else:
            return None, (f"the pre-spec cuts the pockets by {c}, and this run doesn't cut by it (a band or a "
                          f"segment on Columns)")
    return [tuple(x) for x in zip(*cols)] if cols else [()] * len(rows), None


def _range_test(name, rng, K, ref, scores, items) -> RangeTest:
    """items: (stratum, group, bad, gco) per loan in the range."""
    cells: dict[tuple, list[list[float]]] = {}
    for st, k, y, _ in items:
        c = cells.setdefault(st, [[0.0] * K, [0.0] * K])
        c[0][k] += 1
        c[1][k] += y
    pockets = [kgroups.Pocket(c[0], c[1]) for c in cells.values()]
    loans = [int(sum(c[0][k] for c in cells.values())) for k in range(K)]
    bad = [int(sum(c[1][k] for c in cells.values())) for k in range(K)]
    assoc = kgroups.association(pockets, scores)
    fit = kgroups.conditional_fit(pockets, ref, K)
    gaps = []
    for k in range(K):
        if k == ref or fit.odds[k] is None:
            continue
        mh = kgroups.mh_odds(pockets, k, ref)
        if mh:
            gaps.append(abs(fit.odds[k] / mh - 1))
    return RangeTest(name, rng, loans, bad, pockets, assoc, fit, max(gaps) if gaps else None,
                     [k for _, k, _, _ in items], [y for _, _, y, _ in items], [g for _, _, _, g in items])


def run_test(res, ps: prespec.PreSpec) -> Test:
    """The confirmatory test: B3, B4 and B5 on the development range and on the holdout, and B6 on the holdout.
    Every loan with a readable date in one of the two ranges, a readable value of the column and a readable
    outcome is in it; the rest are counted by why they are not."""
    perm.numpy()                                  # B5 needs numpy (OC-34); refused by name without it
    K = len(ps.bins) + 1
    t = Test(column=ps.column, bins=tuple(ps.bins), groups=tuple(ps.groups), ref=ps.reference_index,
             strata=tuple(ps.strata), scores=tuple(float(k) for k in range(1, K + 1)))
    read, why = _reader(res)
    if read is None:
        t.problem = f"the loans can't be split into development and holdout: {why}"
        return t
    table = res.table
    if table is None or ps.column not in table.columns:
        t.problem = f"{ps.column} isn't among the columns this run read"
        return t
    rows = table.rows
    labels, why = _stratum_labels(res, t.strata, rows)
    if labels is None:
        t.problem = why
        return t
    cfg = res.config
    m = next((x for x in res.measures if x.name == "outcome_loans"), None)
    g = next((x for x in res.measures if x.name == "gco_rate"), None)
    if m is None:
        t.problem = "this run has no yes/no outcome to test"
        return t
    col, rules = cfg.origination_date, cfg.missing
    dev, hold = ps.development, ps.holdout
    left = dict.fromkeys((NO_DATE, OUTSIDE, NO_VALUE, NO_OUTCOME), 0)
    items = {DEVELOPMENT: [], HOLDOUT: []}
    for r, st in zip(rows, labels):
        d = _date_or_none(read(r.get(col)))
        if d is None:
            left[NO_DATE] += 1
            continue
        which = HOLDOUT if hold.holds(d) else DEVELOPMENT if dev.holds(d) else None
        if which is None:
            left[OUTSIDE] += 1
            continue
        v, bad_v = engine.classify_number(r.get(ps.column), rules.get(ps.column))
        if bad_v:
            left[NO_VALUE] += 1
            continue
        y = outcome_of(r.get(m.flag), m, rules)
        if y is None:
            left[NO_OUTCOME] += 1
            continue
        gco = None
        if g is not None:
            gv, gw = engine.classify_number(r.get(g.value), rules.get(g.value))
            gco = None if gw else gv
        items[which].append((st, bisect.bisect_right(ps.bins, v), y, gco))
    t.left_out = {k: v for k, v in left.items() if v}
    t.development = _range_test(DEVELOPMENT, dev, K, t.ref, t.scores, items[DEVELOPMENT])
    t.holdout = _range_test(HOLDOUT, hold, K, t.ref, t.scores, items[HOLDOUT])
    t.gco_unread = sum(1 for x in items[HOLDOUT] if x[3] is None) if g is not None else 0
    return t


def _test_words(t: Test) -> str:
    """Check's line on the test."""
    if t.problem:
        return f"Couldn't be run: {_plain(t.problem)}."
    parts = [f"{r.name} {r.range.text()}: {r.n:,} loans, {r.n_bad:,} bad" for r in t.ranges()]
    said = "; ".join(parts) + ". See the Confirmatory test tab."
    said = said[0].upper() + said[1:]
    if t.left_out:
        said += (" Left out of this test only (the other tabs use every loan): "
                 + "; ".join(f"{v:,} with {k}" if k != OUTSIDE else f"{v:,} {k}" for k, v in t.left_out.items())
                 + ".")
    return said


# --------------------------------------------------------------------------
# What it says: Check, the Log, the launcher and the record of what ran


def _commit_words(st: State) -> str:
    pv = st.provenance
    if not pv.get("commit"):
        return pv.get("reason") or "no commit could be read"
    when = str(pv.get("committed_at") or "").replace("T", " ")
    said = f"{pv['commit'][:12]}, committed {when}" if when else pv["commit"][:12]
    if pv.get("dirty"):
        said += "; edited since that commit, so the file read here isn't the one committed"
    elif pv.get("reason"):
        said += f"; {pv['reason']}"
    return said


def _says(ps: prespec.PreSpec) -> str:
    """The pre-spec's lines, as read."""
    strata = ", ".join(ps.strata) if ps.strata else "none (the whole book is one pocket)"
    return "\n".join([
        f"written: {ps.written.isoformat()}",
        f"column: {ps.column}",
        f"bins: {', '.join(engine._fmt(x) for x in ps.bins)} (groups: {'; '.join(ps.groups)})",
        f"reference: {ps.reference}",
        f"strata: {strata}",
        f"confidence: {ps.confidence:g}",
        f"holdout: {ps.holdout.text()}",
        f"development: {ps.development.text()}",
    ])


def _holdout_words(st: State) -> str:
    rng = st.spec.holdout.text()
    if st.touched is None:
        return f"Couldn't be checked: {st.unchecked}. The holdout is {rng}."
    if st.touched:
        said = (f"{st.touched:,} of this extract's loans were made in the holdout ({rng}), the first on "
                f"{st.first.isoformat()} and the last on {st.last.isoformat()}. This run touched the holdout.")
    else:
        said = f"None of this extract's loans were made in the holdout ({rng})."
    if st.undated:
        said += (f" {st.undated:,} loan{'s have' if st.undated != 1 else ' has'} no readable origination date, so "
                 f"whether {'they are' if st.undated != 1 else 'it is'} in it is unknown.")
    return said


def check_rows(res) -> list[tuple[str, str]]:
    st = getattr(res, "prespec", None)
    if st is None:
        return []
    out = [("Pre-spec", str(st.path)), ("Pre-spec commit", _commit_words(st)),
           ("What the pre-spec says", _says(st.spec))]
    if st.spec.written > st.ran_on:
        out.append(("Warning", f"The pre-spec says it was written on {st.spec.written.isoformat()}, after this run."))
    if st.provenance.get("reason"):
        out.append(("Warning", "This run doesn't count as the pre-specified one until the pre-spec is committed, "
                               "unchanged."))
    if st.failed:
        out.append(("Warning", f"Pre-spec: {st.failed}."))
    elif st.deviations:
        out += [("Warning", f"{DEVIATES}: {d}") for d in st.deviations]
    else:
        out.append(("Differs from the pre-spec", "nowhere: this run used what it says"))
    out.append(("Holdout", _holdout_words(st)))
    if st.test is not None:
        out.append(("Confirmatory test", _test_words(st.test)))
    runs = f"{st.runs:,} in this workbook's Log" + (", this one included" if st.touched else "")
    if st.touched is None:
        runs = f"{st.runs_before:,} in this workbook's Log before this one, which couldn't be checked"
    out.append(("Runs that touched the holdout", runs))
    return out


def log_lines(res) -> list[str]:
    """Lines for this run's Log entry: whether it followed its pre-spec, and whether it touched the holdout."""
    st = getattr(res, "prespec", None)
    if st is None:
        return []
    commit = st.provenance.get("commit")
    of = (f"commit {commit[:12]}" + (", edited since" if st.provenance.get("dirty") else "")) if commit \
        else "not committed"
    if st.failed:
        out = [f"Pre-spec {st.name} ({of}): {st.failed}."]
    elif st.deviations:
        k = len(st.deviations)
        out = [f"{DEVIATES} {st.name} ({of}): {k} {'place' if k == 1 else 'places'}, listed on Check."]
    else:
        out = [f"Follows pre-spec {st.name} ({of})."]
    if st.touched:
        out.append(f"{HOLDOUT_MARK} {st.touched:,} loan{'s' if st.touched != 1 else ''} made "
                   f"{st.first.isoformat()} to {st.last.isoformat()}.")
    elif st.touched is None:
        out.append(f"Holdout not checked: {st.unchecked}.")
    return out


def launcher_lines(res) -> list[str]:
    st = getattr(res, "prespec", None)
    if st is None:
        return []
    return log_lines(res) + [f"Runs that touched the holdout, in this workbook's Log: {st.runs:,}."]


def what_ran(res) -> str:
    """Comment lines for the top of the record of what ran."""
    st = getattr(res, "prespec", None)
    if st is None:
        return ""
    out = f"# pre-spec: {st.path} ({_commit_words(st)})\n"
    for d in st.deviations:
        out += f"# differs from the pre-spec: {d}\n"
    for line in log_lines(res)[1:]:
        out += f"# {line}\n"
    return out
