"""What the tool has learned, kept so it can be seen and pruned.

Two kinds of thing are learned, and only when someone confirms them:
  column   a column name and what it meant: key, booked, fico, score, dti ...
  answer   an answer to an odd-value question: FICO's -9999 is missing,
           RANR's negatives are real

Nothing else is stored: no values, no rows, no file contents. A column name,
a meaning, the dates it was confirmed and how many times. The file lives on
the machine that runs the tool (never in the repository), so what is learned
at work stays at work.

Pruning. The firm, 25 Sep 2026: "we also need an intuitive way to go and prune
rules that shouldn't have been added as we learn." So every entry can be
listed (`cube memory`), dropped by name (`cube memory --forget NAME`), or
reviewed in Excel with a Keep / Forget dropdown per row (`cube memory --out
memory.xlsx`, then `cube memory --read memory.xlsx`).

Where: the path given with --memory, else $CUBE_MEMORY, else
~/.origination-cube/memory.yaml.
"""

from __future__ import annotations

import os
from datetime import date
from pathlib import Path
from typing import Any

import yaml

KEEP, FORGET = "Keep", "Forget"


def default_path() -> Path:
    env = os.environ.get("CUBE_MEMORY")
    return Path(env) if env else Path.home() / ".origination-cube" / "memory.yaml"


def load(path: str | Path | None = None) -> dict[str, Any]:
    p = Path(path) if path else default_path()
    if not p.exists():
        return {"columns": {}, "answers": {}}
    raw = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    return {"columns": dict(raw.get("columns") or {}), "answers": dict(raw.get("answers") or {})}


def save(mem: dict[str, Any], path: str | Path | None = None) -> Path:
    p = Path(path) if path else default_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    head = ("# What origination-cube has learned. Column names and meanings only; no data.\n"
            "# Prune with `cube memory --forget NAME`, or `cube memory --out memory.xlsx`.\n")
    p.write_text(head + yaml.safe_dump(mem, sort_keys=True, allow_unicode=True), encoding="utf-8")
    return p


def _touch(entry: dict | None, today: str, **fields) -> dict:
    e = dict(entry or {})
    changed = any(e.get(k) != v for k, v in fields.items() if k in e)
    if changed:
        e["changed_from"] = {k: e.get(k) for k in fields}
        e["times"] = 0
    e.update(fields)
    e.setdefault("first", today)
    e["last"] = today
    e["times"] = int(e.get("times", 0)) + 1
    return e


def remember(config, path: str | Path | None = None, today: date | None = None) -> tuple[Path, int]:
    """Record a confirmed cube file's column meanings and answered questions.
    Only called once a file has passed with `columns_confirmed: yes`."""
    t = (today or date.today()).isoformat()
    mem = load(path)
    n = 0
    for col, (means, is_value) in (config.columns or {}).items():
        if means in ("unknown", "unused"):
            continue
        fields = {"means": means}
        if is_value is not None:
            fields["is"] = is_value
        mem["columns"][col] = _touch(mem["columns"].get(col), t, **fields)
        n += 1
    for q in config.questions:
        if q.answer:
            key = f"{q.column}|{q.pattern}|{q.value if q.value is not None else ''}"
            mem["answers"][key] = _touch(mem["answers"].get(key), t, column=q.column, pattern=q.pattern,
                                         value=q.value, answer=q.answer)
            n += 1
    return save(mem, path), n


def remember_edges(edges: dict[str, str], path: str | Path | None = None) -> None:
    """Band edges a person typed for a column ("620; 680; 740" or "every 20"),
    kept with that column's meaning so the next set-up fills them in (the
    firm, 25 Sep 2026: "yes, remember them"). Only for columns already
    remembered: an edge without a confirmed meaning has nothing to hang on."""
    mem = load(path)
    changed = False
    for col, text in edges.items():
        e = mem["columns"].get(col)
        if e is None or e.get("edges") == text:
            continue
        if text:
            e["edges"] = text
        else:
            e.pop("edges", None)            # cleared on a confirmed Columns tab: forget the edges too
        changed = True
    if changed:
        save(mem, path)


def answer_for(mem: dict[str, Any], column: str, pattern: str, value: Any) -> dict | None:
    key = f"{column}|{pattern}|{value if value is not None else ''}"
    return mem["answers"].get(key)


def forget(names: list[str], path: str | Path | None = None) -> tuple[Path, list[str]]:
    """Drop entries by column name (a column's meaning and any answers about it)."""
    mem = load(path)
    gone = []
    for n in names:
        if mem["columns"].pop(n, None) is not None:
            gone.append(f"column {n}")
        for k in [k for k, v in mem["answers"].items() if v.get("column") == n]:
            mem["answers"].pop(k)
            gone.append(f"answer {k}")
    return save(mem, path), gone


def rows(mem: dict[str, Any]) -> list[dict]:
    out = []
    for col, e in sorted(mem["columns"].items()):
        what = e["means"] + (f" (yes when {e['is']!r})" if "is" in e else "") + \
            (f"; band edges {e['edges']}" if e.get("edges") else "")
        out.append({"kind": "column", "id": col, "column": col, "learned": what, "first": e.get("first"),
                    "last": e.get("last"), "times": e.get("times", 1)})
    for key, e in sorted(mem["answers"].items()):
        v = e.get("value")
        v = f"{v:g}" if isinstance(v, float) else str(v)
        if e.get("pattern") == "negatives":
            said = "negative values are real" if e.get("answer") == "real" else "negative values mean missing"
        else:
            said = f"{v} is a real value" if e.get("answer") == "real" else f"{v} means missing"
        out.append({"kind": "answer", "id": key, "column": e.get("column"),
                    "learned": said, "first": e.get("first"),
                    "last": e.get("last"), "times": e.get("times", 1)})
    return out


def write_review(path: str | Path, mem_path: str | Path | None = None) -> Path:
    """An Excel sheet of everything learned, a Keep / Forget dropdown per row."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook()
    ws = wb.active
    ws.title = "Learned"
    ws["A1"] = ("Everything the cube has learned. Set a row to Forget if it shouldn't have been learned, "
                "save, then run: cube memory --read this-file.xlsx")
    ws["A1"].font = Font(italic=True, size=10)
    heads = ["Keep?", "Kind", "Column", "What it learned", "First confirmed", "Last confirmed", "Times", "id"]
    ws.append([])
    ws.append(heads)
    for i, _ in enumerate(heads, start=1):
        c = ws.cell(row=3, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="16130F")
    dv = DataValidation(type="list", formula1=f'"{KEEP},{FORGET}"', allow_blank=False)
    ws.add_data_validation(dv)
    for r in rows(load(mem_path)):
        ws.append([KEEP, r["kind"], r["column"], r["learned"], r["first"], r["last"], r["times"], r["id"]])
        dv.add(ws.cell(row=ws.max_row, column=1))
    for col, w in zip("ABCDEFGH", (9, 9, 22, 44, 15, 15, 7, 30)):
        ws.column_dimensions[col].width = w
    ws.column_dimensions["H"].hidden = True
    ws.freeze_panes = "A4"
    for row in ws.iter_rows(min_row=4):
        for c in row:
            c.alignment = Alignment(vertical="top")
    p = Path(path)
    wb.save(p)
    return p


def apply_review(path: str | Path, mem_path: str | Path | None = None) -> tuple[Path, list[str]]:
    """Drop every row the review sheet marks Forget. Nothing else changes."""
    from openpyxl import load_workbook
    ws = load_workbook(path)["Learned"]
    mem = load(mem_path)
    gone = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not row or row[0] != FORGET:
            continue
        kind, ident = row[1], row[7]
        if kind == "column" and mem["columns"].pop(ident, None) is not None:
            gone.append(f"column {ident}")
        elif kind == "answer" and mem["answers"].pop(ident, None) is not None:
            gone.append(f"answer {ident}")
    return save(mem, mem_path), gone
