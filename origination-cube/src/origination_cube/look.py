"""The Look tab: what each number column holds, before its edges are chosen.

Fix 3.8 (docs/NEXT-GOAL.md), from the audit's item f: band edges were typed
blind. The Columns tab shows a blank share and three sample values, and
nothing showed a column's spread, its repeats or its codes.

Set up writes one block per number column, top to bottom, numbers on the
left and a histogram beside them:
    loans; blank; not a number; and the share at a value that looks like a
    code (profile.odd_values, the same detection the Odd values tab asks about)
    smallest, median and largest of the other values
    the five most-repeated exact values, with counts

Run writes the tab again. When a number column splits the pockets, it adds a
scatter of that column against each band column: a correlation near zero
misses a U, and a picture does not (docs/statistics.md A9).

Nothing here changes a number the cube uses. A value that looks like a code is
counted on its own row and left out of the spread and the charts; whether it
really is a code is still answered on Odd values.
"""

from __future__ import annotations

import math
import random
import statistics
from bisect import bisect_right
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.pagebreak import Break

from . import engine, profile
from .ingest import Bad, is_blank, parse_number

LOOK = "Look"
DATA = "_look"            # the charts' own numbers, hidden
AFTER = "Columns"         # where the tab goes when it is new
BINS = 20                 # at most about this many equal bars between the 1st and 99th percentile
SAMPLE = 2000             # dots on a scatter, at most
SEED = 7                  # so the same loans are drawn every Run
FIRST = 5                 # first block's row
BLOCK = 17                # rows per block, the gap under it included
CHART_AT = "F"
PAGE_BLOCKS = 2           # blocks per printed page

INK, SLATE, PAPER = "16130F", "57534B", "FFFFFF"        # the workbook's own colours (book.py)
BAR, END_BAR, DOT = "2F5597", "A6A6A6", "2F5597"


@dataclass
class Shape:
    """One column, counted."""
    name: str
    rows: int
    blank: int
    text: int                                   # not a number
    code: float | None                          # a value that looks like a code, or None
    at_code: int
    values: list[float]                         # the other numbers, sorted
    top: list[tuple[float, int]]                # up to five exact values seen more than once, codes included
    step: float = 0.0
    bars: list[tuple[str, int, bool]] = field(default_factory=list)     # label, loans, an end bar?
    fmt: str = "#,##0"                          # the spread's format
    exact: str = "#,##0"                        # the most-repeated values' format


def number_columns(table, cols, few_values: int = 12) -> list[str]:
    """The columns that get a block, from Set up's profile.classify reading
    (`cols`): numbers to cut into bands, and numbers it asks about because every
    loan's differs (a balance with cents does). A 0/1 flag, a four-value class,
    a fixed-width loan number or a date has no edges to choose."""
    out = []
    for c in cols:
        if c.role == "question":
            vals = [parse_number(r.get(c.name)) for r in table.rows if not is_blank(r.get(c.name))]
            nums = {v for v in vals if not isinstance(v, Bad)}
            if not vals or len(nums) <= few_values or \
                    sum(not isinstance(v, Bad) for v in vals) < profile.NUMERIC_SHARE * len(vals):
                continue
        elif c.role != "band":
            continue
        out.append(c.name)
    return out


def shape_of(table, col: str) -> Shape:
    nums, blank, text = [], 0, 0
    for r in table.rows:
        v = r.get(col)
        if is_blank(v):
            blank += 1
            continue
        p = parse_number(v)
        if isinstance(p, Bad):
            text += 1
        else:
            nums.append(p)
    code = next((q["value"] for q in profile.odd_values(col, nums) if q["pattern"] == "repeated_value"), None)
    values = sorted(x for x in nums if x != code)
    counts = Counter(nums)
    # values seen once aren't "repeated": on a balance with cents, five of them would be picked at random
    top = sorted(((v, k) for v, k in counts.items() if k > 1), key=lambda t: (-t[1], t[0]))[:5]
    s = Shape(col, len(table.rows), blank, text, code, len(nums) - len(values), values, top)
    if values:
        s.fmt, s.exact = _formats(values, [v for v, _ in top])
        s.step, s.bars = _bars(values)
    return s


def _formats(values: list[float], top: list[float]) -> tuple[str, str]:
    """Scores and dollars read as whole numbers; a ratio keeps its decimals.
    The most-repeated values are shown exactly, cents and all."""
    if all(float(x).is_integer() for x in values):
        spread = "#,##0"
    else:
        big = max(abs(values[0]), abs(values[-1]), abs(statistics.median(values)))
        spread = "#,##0" if big >= 100 else "#,##0.0" if big >= 10 else "0.00" if big >= 1 else "0.000"
    places = max((len(f"{x:.4f}".rstrip("0").split(".")[1]) for x in top), default=0)
    return spread, ("#,##0" if places == 0 else "#,##0." + "0" * places)


def _nice(raw: float) -> float:
    """The smallest of 1, 2 or 5 times a power of ten that is at least raw:
    a width a person would type as "every 20"."""
    if raw <= 0:
        return 1.0
    k = math.floor(math.log10(raw))
    for m in (1, 2, 5, 10):
        w = float(f"{m * 10 ** k:.12g}")
        if w >= raw:
            return w
    return w


def _step(lo: float, hi: float, whole: bool) -> float:
    """The bar width that cuts lo to hi into at most BINS bars; a whole
    number for a column of whole numbers, so no bar holds half a score."""
    w = _nice((hi - lo) / BINS)
    return max(w, 1.0) if whole else w


def _bars(values: list[float]) -> tuple[float, list[tuple[str, int, bool]]]:
    """Equal-width bars between the 1st and 99th percentile, plus one end bar
    below and one above for whatever is left, drawn grey.

    Equal widths, not equal counts: bars holding the same number of loans would
    all stand the same height and hide the shape the analyst is looking for.
    Between the percentiles, so one $2m loan or a stray score of 12 doesn't
    squash everything into one bar. The width is a round step and the bars sit
    on its multiples, so each bar is exactly a band that "every 20" on Columns
    would make, and the labels read the way the Grids write bands."""
    n = len(values)
    lo, hi = values[int(0.01 * (n - 1))], values[math.ceil(0.99 * (n - 1))]
    if hi <= lo:
        lo, hi = values[0], values[-1]
    whole = all(float(x).is_integer() for x in values)
    if hi <= lo:                                                  # one value only
        return 0.0, [(_plain(lo), n, False)]
    w = _step(lo, hi, whole)
    start, stop = math.floor(lo / w) * w, (math.floor(hi / w) + 1) * w
    edges = [round(start + i * w, 10) for i in range(round((stop - start) / w) + 1)]
    labels = engine.band_labels(tuple(edges))                     # "up to 559", "560 - 579", ..., "840 and up"
    counts = [0] * len(labels)
    for v in values:
        counts[bisect_right(edges, v)] += 1
    bars = [(lab, k, i in (0, len(labels) - 1)) for i, (lab, k) in enumerate(zip(labels, counts))]
    if not bars[0][1]:
        bars = bars[1:]
    if not bars[-1][1]:
        bars = bars[:-1]
    return w, bars


def _plain(x: float) -> str:
    """A value as the Odd values tab writes it: -9999, 0.35."""
    return str(int(x)) if float(x).is_integer() else f"{x:g}"


# --------------------------------------------------------------------------
# The tab


def write_look(wb, table, columns, split: str | None = None, bands=()) -> None:
    """Write (or write again) the Look tab and its hidden chart data. `columns`
    are the number columns to show, `split` the column that splits the pockets
    (after a Run), and `bands` the band columns it is plotted against."""
    at = wb.sheetnames.index(LOOK) if LOOK in wb.sheetnames else (
        wb.sheetnames.index(AFTER) + 1 if AFTER in wb.sheetnames else None)
    for t in (LOOK, DATA):
        if t in wb.sheetnames:
            del wb[t]
    ws = wb.create_sheet(LOOK, at)
    hs = wb.create_sheet(DATA)
    hs.sheet_state = "hidden"
    columns = [c for c in columns if c in table.columns]
    hs.cell(row=1, column=1, value="columns")                     # read back by refresh()
    for i, c in enumerate(columns, start=2):
        hs.cell(row=1, column=i, value=c)

    _title(ws)
    shapes = {c: shape_of(table, c) for c in columns}
    r = FIRST
    for i, c in enumerate(columns):
        if i and i % PAGE_BLOCKS == 0:
            _page(ws, r)
        _block(ws, hs, r, shapes[c], 2 * i + 1)
        r += BLOCK
    end = _scatters(ws, hs, r, table, shapes, split, [b for b in bands if b != split], 2 * len(columns) + 1)

    for col, w in zip("ABCDE", (2, 30, 11, 9, 3)):
        ws.column_dimensions[col].width = w
    ws.print_area = f"A1:O{end}"                                   # to the foot of the last chart
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.sheet_view.showGridLines = False


def refresh(book: str | Path, table, split: str | None = None, bands=()) -> None:
    """At Run: the same columns Set up showed, and the split's scatters. A
    workbook set up before the Look tab existed gets it at its next Set up;
    working out its columns here would repeat Set up's slowest step."""
    wb = load_workbook(book)
    if DATA not in wb.sheetnames:
        return
    columns = [str(c.value) for c in wb[DATA][1][1:] if c.value]
    write_look(wb, table, columns, split, bands)
    wb.save(book)


def _page(ws, r: int) -> None:
    """A new printed page before the block at row r. The break sits in the gap
    row above, not on the chart's top edge, which would print a stray line."""
    ws.row_breaks.append(Break(id=r - 2))


def _axis(ax, vals: list[float], fmt: str) -> None:
    """A scatter axis on round numbers, set rather than left to the program:
    Excel starts a FICO axis at 0 and squashes every dot to the right."""
    lo, hi = min(vals), max(vals)
    unit = _nice((hi - lo) / 8)                                     # about eight gridlines
    ax.scaling.min, ax.scaling.max = math.floor(lo / unit) * unit, math.ceil(hi / unit) * unit
    if ax.scaling.max <= ax.scaling.min:
        ax.scaling.max = ax.scaling.min + unit
    ax.majorUnit = unit
    ax.number_format = fmt
    ax.delete = False


def _title(ws) -> None:
    ws.merge_cells("B1:O1")
    ws["B1"] = "Look"
    ws["B1"].font = Font(name="Arial", bold=True, size=16, color=PAPER)
    for c in "BCDEFGHIJKLMNO":
        ws[f"{c}1"].fill = PatternFill("solid", fgColor=INK)
    ws.row_dimensions[1].height = 28
    ws["B2"] = "Look at each number column before choosing its edges on Columns."
    ws["B2"].font = Font(name="Calibri", bold=True, size=11)
    ws["B3"] = ('A chart "in steps of 20" shows the bands that "every 20" on Columns would cut. '
                "The grey bar at each end holds the last 1% or less.")
    ws["B3"].font = Font(name="Calibri", size=10, color=SLATE)


def _bar_row(ws, r: int, text: str) -> None:
    for col in (2, 3, 4):
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=INK)
    c = ws.cell(row=r, column=2, value=text)
    c.font = Font(name="Calibri", bold=True, size=12, color=PAPER)


def _line(ws, r: int, label: str, value=None, share=None, fmt: str = "#,##0", indent: int = 0) -> None:
    a = ws.cell(row=r, column=2, value=label)
    a.alignment = Alignment(indent=indent)
    a.font = Font(name="Calibri")
    if value is not None:
        b = ws.cell(row=r, column=3, value=value)
        b.number_format = fmt
        b.font = Font(name="Calibri")
    if share is not None:
        d = ws.cell(row=r, column=4, value=share)
        d.number_format = "0.0%"
        d.font = Font(name="Calibri", color=SLATE)


def _sub(ws, r: int, left: str, *heads: str) -> None:
    ws.cell(row=r, column=2, value=left).font = Font(name="Calibri", bold=True)
    for i, h in enumerate(heads, start=3):
        c = ws.cell(row=r, column=i, value=h)
        c.font = Font(name="Calibri", bold=True, color=SLATE)
        c.alignment = Alignment(horizontal="right")


def _block(ws, hs, r: int, s: Shape, hcol: int) -> None:
    n = s.rows or 1
    _bar_row(ws, r, s.name)
    _line(ws, r + 1, "Loans", s.rows)
    _line(ws, r + 2, "Blank", s.blank, s.blank / n)
    _line(ws, r + 3, "Not a number", s.text, s.text / n)
    if s.code is not None:
        _line(ws, r + 4, f"At {_plain(s.code)}, likely a code", s.at_code, s.at_code / n)
    else:
        _line(ws, r + 4, "Likely a code", "none found")
        ws.cell(row=r + 4, column=3).alignment = Alignment(horizontal="right")
    left_out = s.blank + s.text + s.at_code
    _sub(ws, r + 5, f"The other {len(s.values):,} values" if left_out else f"All {len(s.values):,} values")
    if s.values:
        _line(ws, r + 6, "Smallest", s.values[0], fmt=s.fmt, indent=1)
        _line(ws, r + 7, "Median", statistics.median(s.values), fmt=s.fmt, indent=1)
        _line(ws, r + 8, "Largest", s.values[-1], fmt=s.fmt, indent=1)
    else:
        _line(ws, r + 6, "No numbers to show.", indent=1)
    _sub(ws, r + 9, "Most repeated", "Loans", "Share")
    for k, (v, cnt) in enumerate(s.top):
        _line(ws, r + 10 + k, None, cnt, cnt / n, indent=1)
        c = ws.cell(row=r + 10 + k, column=2, value=v)
        # a code as the row above and the Odd values tab write it: -9999, not -9,999
        c.number_format = "0" if v == s.code and float(v).is_integer() else s.exact
        c.alignment = Alignment(horizontal="left", indent=1)
    if len(s.top) < 5:
        _line(ws, r + 10 + len(s.top), "No other value repeats." if s.top else "No value repeats.", indent=1)
        ws.cell(row=r + 10 + len(s.top), column=2).font = Font(name="Calibri", color=SLATE)
    if not s.bars:
        return
    hs.cell(row=3, column=hcol, value=s.name)
    hs.cell(row=3, column=hcol + 1, value="Loans")
    for k, (lab, cnt, _) in enumerate(s.bars, start=4):
        hs.cell(row=k, column=hcol, value=lab)
        hs.cell(row=k, column=hcol + 1, value=cnt)
    last = 3 + len(s.bars)
    ch = BarChart()
    ch.type = "col"
    ch.gapWidth = 15
    ch.add_data(Reference(hs, min_col=hcol + 1, min_row=3, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(hs, min_col=hcol, min_row=4, max_row=last))
    ser = ch.series[0]
    ser.graphicalProperties.solidFill = BAR
    ser.graphicalProperties.line.solidFill = BAR
    for k, (_, _, end) in enumerate(s.bars):
        if end:
            pt = DataPoint(idx=k)
            pt.graphicalProperties.solidFill = END_BAR
            pt.graphicalProperties.line.solidFill = END_BAR
            ser.dPt.append(pt)
    ch.title = f"{s.name} in steps of {s.step:,g}" if s.step else s.name
    ch.legend = None
    ch.y_axis.title = "Loans"
    ch.y_axis.number_format = "#,##0"
    ch.y_axis.scaling.min = 0
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.width, ch.height = 18, 7.8
    ws.add_chart(ch, f"{CHART_AT}{r}")




def _note(ws, r: int, text: str) -> None:
    c = ws.cell(row=r, column=2, value=text)
    c.font = Font(name="Calibri", color=SLATE)


def _scatters(ws, hs, r: int, table, shapes: dict[str, Shape], split: str | None, bands: list[str],
              hcol: int) -> int:
    """The split column against each band column, one scatter each, on a page
    of their own. Returns the row under the last one."""
    if not split:
        _bar_row(ws, r, "Scatters")
        _note(ws, r + 1, "None yet. Set a number column to split the pockets on Columns, then Run. A scatter of "
                         "it against each band column appears here.")
        return r + 2
    if split not in shapes:
        _bar_row(ws, r, f"{split} splits the pockets")
        _note(ws, r + 1, f"{split} isn't a number column, so it has no scatter. The Three-way tab shows it value "
                         f"by value.")
        return r + 2
    if r > FIRST:
        _page(ws, r)
    _bar_row(ws, r, f"{split} against each band column")
    _note(ws, r + 1, f"Each dot is one loan. If the dots slope or bend, {split} moves with that column, and part "
                     f"of its gap on Split may be that column's.")
    r += 3
    ys = _readable(table, split, shapes[split].code)
    for j, band in enumerate(b for b in bands if b in table.columns):
        sb = shapes.get(band) or shape_of(table, band)
        pairs = [(x, y) for x, y in zip(_readable(table, band, sb.code), ys) if x is not None and y is not None]
        shown = pairs
        if len(pairs) > SAMPLE:
            # a sample keeps the workbook light; drawn with a fixed seed, so every Run shows the same loans
            shown = [pairs[i] for i in sorted(random.Random(SEED).sample(range(len(pairs)), SAMPLE))]
        if j and j % PAGE_BLOCKS == 0:
            _page(ws, r)
        _bar_row(ws, r, f"{split} against {band}")
        _line(ws, r + 1, "Loans with both values", len(pairs))
        _line(ws, r + 2, "Dots shown", len(shown))
        if len(shown) < len(pairs):
            _note(ws, r + 3, f"A random sample of {len(shown):,}: the same loans every Run.")
        elif not pairs:
            _note(ws, r + 3, "No loan has both, so there is nothing to plot.")
        if pairs:
            c = hcol + 2 * j
            hs.cell(row=3, column=c, value=band)
            hs.cell(row=3, column=c + 1, value=split)
            for k, (x, y) in enumerate(shown, start=4):
                hs.cell(row=k, column=c, value=x)
                hs.cell(row=k, column=c + 1, value=y)
            ch = ScatterChart()
            ch.title = f"{split} against {band}"
            ch.style = 13
            ser = Series(Reference(hs, min_col=c + 1, min_row=4, max_row=3 + len(shown)),
                         Reference(hs, min_col=c, min_row=4, max_row=3 + len(shown)), title=split)
            ser.marker.symbol = "circle"
            ser.marker.size = 3
            ser.marker.graphicalProperties.solidFill = DOT
            ser.marker.graphicalProperties.line.solidFill = DOT
            ser.graphicalProperties.line.noFill = True
            ch.series.append(ser)
            ch.legend = None
            ch.x_axis.title, ch.y_axis.title = band, split
            _axis(ch.x_axis, [x for x, _ in shown], sb.fmt)
            _axis(ch.y_axis, [y for _, y in shown], shapes[split].fmt)
            ch.width, ch.height = 18, 7.8
            ws.add_chart(ch, f"{CHART_AT}{r}")
        r += BLOCK
    return r


def _readable(table, col: str, code: float | None) -> list[float | None]:
    """Each loan's value, None where it is blank, not a number or the code."""
    out: list[float | None] = []
    for row in table.rows:
        p = parse_number(row.get(col))
        out.append(p if isinstance(p, float) and p != code else None)
    return out
