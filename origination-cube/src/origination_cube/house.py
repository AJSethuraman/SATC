"""The house colours, in one place: the launcher and the workbook both read them.

The firm's redesign (docs/redesign-2026-09-26/README.md, "Design tokens") uses
the bank's workbook palette. The values are written here once, as the cube's
own module: `credit-suite` owns the palette's shared code, and a copy of that
file anywhere else in the repository fails its conformance test. Nothing here
imports openpyxl or Tk, so the launcher can read the colours before it has
checked that the add-ons are installed (ruling OC-34).

Hex is written without the '#', the way openpyxl takes it; `tk()` adds it for
the window.
"""

from __future__ import annotations

INK = "0A0908"          # title bands, column headers, a done step
ONYX = "16130F"         # section bands on the record; data text
INK_TEXT = ONYX
KEY_RED = "CC0000"      # the accent rule, the current step, the primary button
CRIMSON = "960019"      # alert text: "4 left", a refusal's tag
PAPER = "FFFFFF"        # the page
CANVAS = "F4F1EC"       # tiles, method notes, the step rail
MIST = "E4DFD5"         # dividers, the read-only band, a disabled button
STONE = "B9B4AC"        # a step still to do, faint rules, a greyed row
SLATE = "57534B"        # secondary text
ALERT_FG = "F7DEDE"     # a cell that needs an answer, a worse pocket, the workbook-open banner
POSITIVE = "1E7A47"     # tie-outs, holds up
POSITIVE_BG = "EAF6EE"
DISABLED_TEXT = "9A958C"
ROW_RULE = "EFEBE4"     # the one-pixel line between rows
HEAT_GOOD, HEAT_MID, HEAT_BAD, HEAT_BAD2 = "BBD3BD", "F4F1EC", "E0A6A6", "CF7777"

PALETTE = {"INK": INK, "ONYX": ONYX, "KEY_RED": KEY_RED, "CRIMSON": CRIMSON, "PAPER": PAPER, "CANVAS": CANVAS,
           "MIST": MIST, "STONE": STONE, "SLATE": SLATE, "ALERT_FG": ALERT_FG, "POSITIVE": POSITIVE,
           "POSITIVE_BG": POSITIVE_BG, "HEAT_GOOD": HEAT_GOOD, "HEAT_MID": HEAT_MID, "HEAT_BAD": HEAT_BAD,
           "HEAT_BAD2": HEAT_BAD2}


def tk(hex_: str) -> str:
    """A colour as Tk takes it: '#CC0000'."""
    return f"#{hex_}"


def fill(hex_: str):
    """A solid cell fill (openpyxl, imported only when a workbook is being written)."""
    from openpyxl.styles import PatternFill
    return PatternFill("solid", fgColor=hex_, bgColor=hex_)


def band(ws, row: int, first: int, last: int, text: str, hex_: str = INK, color: str = PAPER,
         rule: str | None = KEY_RED, key_col: int | None = None, key: str | None = None) -> None:
    """A section's header band across columns first..last: a fill, bold white text, and a rule under it."""
    from openpyxl.styles import Border, Font, Side
    for col in range(first, last + 1):
        c = ws.cell(row=row, column=col)
        c.fill = fill(hex_)
        if rule:
            c.border = Border(bottom=Side(style="medium", color=rule))
    head = ws.cell(row=row, column=first, value=text)
    head.font = Font(name="Arial", bold=True, size=10, color=color)
    if key_col and key:
        ws.cell(row=row, column=key_col, value=key)
