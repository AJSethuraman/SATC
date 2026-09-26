"""The control center: every setting of a run on one Excel tab.

Each setting has a dropdown of options and a cell for your own value. Two
more cells are Excel formulas: the value in use, and what the chosen option
means. The options and their explanations come from `settings.yaml` and
nowhere else, so the tab, the run and the docs cannot drift apart.

Reading back does not depend on Excel having recalculated. Python reads the
dropdown and the override cells directly: your own value wins when it is
filled in. A setting with neither is refused and named.

Two kinds of setting (settings.yaml says which). A judgment setting - is it
material, is it enough loans, how much worse counts - opens blank and the
run refuses until someone answers it: the tool never decides what matters.
A method setting - how bands are cut, which multiple-test allowance - opens
on its recommended option, which is printed on every output.

The tab does not label which is which. The firm, 25 Sep 2026: "i don't want
things to be labeled 'your judgment' that is very AI coded ... conditionally
format the workbook to indicate where we must enter things and provide
instructions that don't sound robotic." So a cell that still needs an answer
is shaded by conditional formatting, and the shading goes away once it is
answered; the instructions at the top are written the way the firm talks.
"""

from __future__ import annotations

from dataclasses import dataclass
from importlib import resources
from pathlib import Path
import math
from typing import Any

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SHEET = "Control"
OPTIONS_SHEET = "_options"
RECOMMENDED = " (recommended)"
FIRST_ROW = 5
KEY_COL, CHOOSE_COL, OWN_COL = 7, 3, 4            # G, C, D
NEEDS = "FCE4C4"          # the shade on a cell that still needs an answer
INSTRUCTIONS = (
    "Fill in the shaded cells. The rest have starting values; change them if the population calls for it. "
    "Pick from the list, or type a number in the next column to override it. Changes apply on the next Run."
)

INK, CANVAS, MIST, SLATE, PAPER, KEY_RED = "16130F", "F4F1EC", "E4DFD5", "57534B", "FFFFFF", "CC0000"


class ControlError(Exception):
    def __init__(self, problems: list[str]):
        self.problems = problems
        super().__init__("\n".join(problems))


@dataclass(frozen=True)
class Option:
    value: Any
    label: str
    explains: str
    recommended: bool = False

    @property
    def shown(self) -> str:
        return self.label + (RECOMMENDED if self.recommended else "")


@dataclass(frozen=True)
class Setting:
    key: str
    question: str
    takes_effect: str
    override: str | None
    options: tuple[Option, ...]
    group: str
    judgment: bool = False
    valid: dict | None = None        # {min, max, whole} or {date: true}: what a typed value may be
    needed_when: str | None = None   # asked for only then (NEEDED_WHEN); blank is no answer, not a refusal, otherwise
    def recommended(self) -> Option | None:
        return next((o for o in self.options if o.recommended), None)

    @property
    def is_date(self) -> bool:
        return bool((self.valid or {}).get("date"))


#: When a setting that isn't always needed is needed, in the words its row shows while blank. The run
#: (book.read_book) refuses a blank one only then, naming its cell.
NEEDED_WHEN = {
    "outcome_date": "Needs an answer when a column on Columns is marked Outcome date.",
    "age_or_window": "Needs an answer for a loan age filter or an outcome window, unless a column is marked "
                     "As-of date.",
}
DATE_FORMAT = "yyyy-mm-dd"
#: The new-columns block under the settings (fix 3.9): a heading, a note, a header and this many rows.
DERIVED_KEY = "derived"
DERIVED_ROWS = 3
DERIVED_NOTE = ("Name a new column and pick the two columns it divides. Press Set up again: it then shows on "
                "Columns and Look like any number column. Where the bottom is zero or blank, it is blank.")


def load_settings(path: str | Path | None = None) -> list[Setting]:
    text = (Path(path).read_text(encoding="utf-8") if path
            else resources.files("origination_cube").joinpath("settings.yaml").read_text(encoding="utf-8"))
    raw = yaml.safe_load(text)
    out = []
    for g in raw["groups"]:
        for s in g["settings"]:
            opts = tuple(Option(value=o["value"], label=str(o["label"]), explains=" ".join(str(o["explains"]).split()),
                                recommended=bool(o.get("recommended", False))) for o in s["options"])
            out.append(Setting(key=s["key"], question=s["question"], takes_effect=s["takes_effect"],
                               override=s.get("override"), options=opts, group=g["title"],
                               judgment=bool(s.get("judgment", False)), valid=s.get("valid"),
                               needed_when=s.get("needed_when")))
    return out


# --------------------------------------------------------------------------


def _by_value(key: str, v: Any) -> str:
    """An option's value as Excel writes a number joined to text: 0.95, 30, 1.25."""
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return ""
    t = f"{float(v):.10f}".rstrip("0").rstrip(".")
    return f"{key}|{t}"


def write_control(wb: Workbook, settings: list[Setting], with_columns: bool = False) -> None:
    """`with_columns`: the workbook has a Columns tab (Set up), so a setting
    needed only when a column is marked a certain way is shaded only then."""
    ws = wb.create_sheet(SHEET, 0) if SHEET not in wb.sheetnames else wb[SHEET]
    opt = wb.create_sheet(OPTIONS_SHEET)
    opt.append(["lookup", "setting", "option", "value", "what it means", "your own value accepts", "plain",
                "lookup by value"])
    ranges: dict[str, tuple[int, int]] = {}
    for s in settings:
        first = opt.max_row + 1
        for o in s.options:
            opt.append([f"{s.key}|{o.shown}", s.key, o.shown, o.value, o.explains, s.override or "", o.label,
                        _by_value(s.key, o.value)])
        ranges[s.key] = (first, opt.max_row)
    opt.sheet_state = "hidden"

    thin = Side(style="thin", color=MIST)
    ws.sheet_view.showGridLines = False
    widths = {"A": 2, "B": 34, "C": 44, "D": 19, "E": 22, "F": 70, "G": 12}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.merge_cells("B1:F1")
    ws["B1"] = "Control center"
    ws["B1"].font = Font(name="Arial", bold=True, size=16, color=PAPER)
    for c in "BCDEF":
        ws[f"{c}1"].fill = PatternFill("solid", fgColor=INK)
    ws.row_dimensions[1].height = 28
    ws.merge_cells("B2:F2")
    ws["B2"] = INSTRUCTIONS
    ws["B2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws["B2"].font = Font(name="Calibri", size=10, color=SLATE)
    ws.row_dimensions[2].height = 44
    ws["B3"] = "Shaded = still needs an answer"
    ws["B3"].fill = PatternFill("solid", fgColor=NEEDS)
    ws["B3"].font = Font(name="Calibri", size=9, color=INK)
    heads = ["Setting", "Choose", "Or enter your own", "In use", "What it means", "key"]
    for i, h in enumerate(heads):
        c = ws.cell(row=4, column=2 + i, value=h)
        c.font = Font(name="Calibri", bold=True, color=PAPER)
        c.fill = PatternFill("solid", fgColor=INK)
    ws.freeze_panes = "C5"

    r = FIRST_ROW
    group = None
    row_of: dict[str, int] = {}
    for s in settings:
        if s.group != group:
            group = s.group
            ws.cell(row=r, column=2, value=group).font = Font(name="Calibri", bold=True, color=INK)
            for col in range(2, 7):
                ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=CANVAS)
            r += 1
        first, last = ranges[s.key]
        row_of[s.key] = r
        ws.cell(row=r, column=2, value=s.question)
        rec = s.recommended()
        choose = ws.cell(row=r, column=CHOOSE_COL, value=None if s.judgment or rec is None else rec.shown)
        dv = DataValidation(type="list", formula1=f"='{OPTIONS_SHEET}'!$C${first}:$C${last}", allow_blank=True,
                            showDropDown=False, showErrorMessage=True)
        dv.errorTitle = "Pick from the list"
        dv.error = "Pick one of the listed options. To use your own number, type it in the next column."
        ws.add_data_validation(dv)
        dv.add(choose)
        own = ws.cell(row=r, column=OWN_COL)
        if s.is_date:
            # a typed date, shown as one; the run reads it whether Excel kept it as a date or as text
            dvd = DataValidation(type="date", operator="between", formula1="32874", formula2="73415",
                                 allow_blank=True, showErrorMessage=True)
            dvd.errorTitle = "Not a date"
            dvd.error = "Type a date between 1990 and 2100, such as 2026-06-30."
            ws.add_data_validation(dvd)
            dvd.add(own)
            own.number_format = DATE_FORMAT
            ws.cell(row=r, column=5).number_format = DATE_FORMAT
        elif s.override is not None and s.valid:
            v = s.valid
            dvo = DataValidation(type="whole" if v.get("whole") else "decimal", operator="between",
                                 formula1=str(v["min"]), formula2=str(v["max"]), allow_blank=True,
                                 showErrorMessage=True)
            dvo.errorTitle = "Out of range"
            dvo.error = f"Enter {_range_words(s)}."
            ws.add_data_validation(dvo)
            dvo.add(own)
        if s.override is None:
            own.value = "n/a"
            own.font = Font(name="Calibri", italic=True, color=SLATE)
            own.fill = PatternFill("solid", fgColor=MIST)
        from openpyxl.utils import get_column_letter
        C, D, H = f"C{r}", f"D{r}", f"${get_column_letter(KEY_COL)}{r}"   # H: the key column, wherever it is
        # shaded while unanswered, on any row: a method setting someone clears needs an answer too.
        # A row with no own-value cell shades only its dropdown (the grey n/a cell is not an answer).
        if s.needed_when is None:
            ws.conditional_formatting.add(
                f"C{r}:D{r}" if s.override is not None else f"C{r}",
                FormulaRule(formula=[f'AND($C{r}="",OR($D{r}="",$D{r}="n/a"))'],
                            fill=PatternFill("solid", fgColor=NEEDS, bgColor=NEEDS)))
        own_set = f'AND({D}<>"",{D}<>"n/a")'
        # the label, or a number equal to an option's value, as the reader takes it: Excel turns "95%"
        # into 0.95 (the second walk, defect 3), and the tab mustn't say "not an option" to a value
        # the run then uses (found on the render, 25 Sep 2026)
        lookup = (f"IFERROR(MATCH({H}&\"|\"&{C},{OPTIONS_SHEET}!$A:$A,0),"
                  f"MATCH({H}&\"|\"&IFERROR(VALUE({C}),{C}),{OPTIONS_SHEET}!$H:$H,0))")
        ws.cell(row=r, column=5, value=(f'=IF({own_set},{D},IF({C}="","",'
                                        f'IFERROR(INDEX({OPTIONS_SHEET}!$G:$G,{lookup}),"not an option")))'))
        note = f"Your own value, in place of the options ({s.override})." if s.override else ""
        pick = "Pick one, or enter your own." if s.override is not None else "Pick one from the list."
        instead = (f" Pick from the list, or put your {'date' if s.is_date else 'number'} in the next column."
                   if s.override is not None else " Pick from the list.")
        blank = NEEDED_WHEN[s.needed_when] if s.needed_when else "Needs an answer before we run."
        ws.cell(row=r, column=6, value=(f'=IF({own_set},"{note}",IF({C}="","{blank} '
                                        f'{pick}",'
                                        f'IFERROR(INDEX({OPTIONS_SHEET}!$E:$E,{lookup}),"That isn\'t one of the '
                                        f'options.{instead}")))'))
        k = ws.cell(row=r, column=KEY_COL, value=s.key)
        k.font = Font(name="Consolas", size=8, color=SLATE)
        for col in range(2, 8):
            cell = ws.cell(row=r, column=col)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col != OWN_COL and col != KEY_COL and cell.font.color is None:
                cell.font = Font(name="Calibri", size=10, color=INK)
        ws.cell(row=r, column=5).font = Font(name="Calibri", bold=True, color=INK)
        r += 1
    _shade_when_needed(ws, settings, row_of, with_columns)
    ws.column_dimensions["G"].hidden = True
    ws.print_area = f"B1:F{r - 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def _in_use(r: int, s: Setting) -> str:
    """An Excel test that a whole-months setting is answered with more than 0."""
    zero = next((o.shown for o in s.options if o.value == 0), "")
    return f'OR(AND(ISNUMBER($D{r}),$D{r}>0),AND($C{r}<>"",$C{r}<>"{zero}"))'


def _shade_when_needed(ws, settings: list[Setting], row_of: dict[str, int], with_columns: bool) -> None:
    """A setting needed only sometimes is shaded, while blank, only when it is
    needed: the window when a column is marked Outcome date on Columns, the
    as-of date when a loan age filter or a window is in use and no column is
    marked As-of date. The run refuses on the same conditions (book.read_book)."""
    by_key = {s.key: s for s in settings}
    for s in settings:
        if not s.needed_when:
            continue
        r = row_of[s.key]
        if s.needed_when == "outcome_date":
            when = 'COUNTIF(Columns!$C:$C,"Outcome date")>0' if with_columns else "FALSE"
        else:
            uses = [_in_use(row_of[k], by_key[k]) for k in ("min_age_months", "window_months") if k in row_of]
            when = f"OR({','.join(uses)})" if uses else "FALSE"
            if with_columns:
                when = f'AND(COUNTIF(Columns!$C:$C,"As-of date")=0,{when})'
        ws.conditional_formatting.add(
            f"C{r}:D{r}" if s.override is not None else f"C{r}",
            FormulaRule(formula=[f'AND($C{r}="",OR($D{r}="",$D{r}="n/a"),{when})'],
                        fill=PatternFill("solid", fgColor=NEEDS, bgColor=NEEDS)))


def build_control_book(out: str | Path, settings: list[Setting] | None = None) -> Path:
    wb = Workbook()
    wb.remove(wb.active)
    write_control(wb, settings or load_settings())
    p = Path(out)
    wb.save(p)
    return p


# --------------------------------------------------------------------------


def read_control(path: str | Path, settings: list[Setting] | None = None) -> dict[str, Any]:
    """The settings in use, read from the cells a person edits. Refuses, all
    at once, any setting left with nothing chosen, an unknown option, or an
    own value the setting cannot take."""
    settings = settings or load_settings()
    by_key = {s.key: s for s in settings}
    ws = load_workbook(path)[SHEET]
    found: dict[str, Any] = {}
    seen: set[str] = set()
    problems: list[str] = []
    for row in ws.iter_rows(min_row=FIRST_ROW):
        key = row[KEY_COL - 1].value
        if key not in by_key:
            continue
        s = by_key[key]
        seen.add(key)
        chosen, own = row[CHOOSE_COL - 1].value, row[OWN_COL - 1].value
        where = f"{SHEET}!C{row[0].row}"
        if own not in (None, "", "n/a"):
            if s.override is None:
                problems.append(f'{where}: "{s.question}" takes one of the listed options only.')
            elif s.is_date:
                got = as_date(own)
                if got is None:
                    problems.append(f'{SHEET}!D{row[0].row}: "{s.question}" needs {s.override}; got {own!r}.')
                else:
                    found[key] = got
            elif not isinstance(own, (int, float)) or isinstance(own, bool):
                problems.append(f'{SHEET}!D{row[0].row}: "{s.question}" needs {s.override}; got {own!r}.')
            elif s.valid and not (s.valid["min"] <= own <= s.valid["max"]) or \
                    (s.valid and s.valid.get("whole") and not float(own).is_integer()):
                problems.append(f'{SHEET}!D{row[0].row}: "{s.question}" needs {_range_words(s)}; got {own!r}.'
                                + _percent_hint(s, own))
            else:
                found[key] = int(own) if (s.valid or {}).get("whole") else own
            continue
        if chosen in (None, ""):
            if s.needed_when:
                continue                # asked for only when it matters; book.read_book says when it does
            how = "Pick one from the list." if s.override is None else "Pick one, or enter your own in column D."
            problems.append(f'{where}: "{s.question}" needs an answer. {how}')
            continue
        match = _matching(s, chosen)
        if not match:
            problems.append(f'{where}: {chosen!r} is not an option for "{s.question}".')
            continue
        found[key] = match[0].value
    for k in by_key:
        if k not in seen:
            problems.append(f'The {SHEET} tab is missing the setting "{by_key[k].question}". Press Set up again.')
    if problems:
        raise ControlError(problems)
    return found


def describe(found: dict[str, Any], settings: list[Setting] | None = None) -> list[tuple[str, str]]:
    """The settings in use as a person reads them: the question and the chosen
    option's words, or the typed value (walkthrough defect 13 echoed codes)."""
    out = []
    for s in settings or load_settings():
        if s.key not in found:
            continue
        v = found[s.key]
        o = next((o for o in s.options if o.value == v), None)
        label = o.label.replace(" (suggested)", "") if o else None
        out.append((s.question, label if o else f"{v:g}" if isinstance(v, float) else str(v)))
    return out


def _as_number(v: Any) -> float | None:
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    t = str(v).strip().replace(",", "")
    try:
        return float(t[:-1]) / 100 if t.endswith("%") else float(t)
    except ValueError:
        return None


def _matching(s: Setting, chosen: Any) -> list[Option]:
    """The option a Choose cell means, however Excel stored it. A pick from the
    list is its label; but Excel can turn a label that looks like a number into
    that number ("95%" becomes 0.95), so a number that equals an option's value
    means that option too (the second walkthrough, defect 3)."""
    text = str(chosen).strip()
    hit = [o for o in s.options if text in (o.shown, o.label)]
    if hit:
        return hit
    n = _as_number(chosen)
    if n is None:
        return []
    return [o for o in s.options if isinstance(o.value, (int, float)) and not isinstance(o.value, bool)
            and abs(float(o.value) - n) < 1e-9]


def answer_of(key: str, choose: Any, own: Any) -> Any:
    """One setting's answer from its two cells, or None when there isn't a usable
    one: a valid own value first, then the option picked. For Set up, which reads
    Control before the Run's full check does."""
    s = next((x for x in load_settings() if x.key == key), None)
    if s is None:
        return None
    n = _as_number(own) if own not in (None, "", "n/a") else None
    v = s.valid or {}
    if n is not None and v.get("min", -math.inf) <= n <= v.get("max", math.inf):
        return int(n) if v.get("integer") or float(n).is_integer() else n
    hit = _matching(s, choose) if choose not in (None, "") else []
    return hit[0].value if hit else None


def _range_words(s: Setting) -> str:
    """What the own-value cell takes, with its range said once (second walk,
    defect 15: "a share between 0.5 and 0.999, from 0.5 to 0.999")."""
    v = s.valid or {}
    if "min" not in v:
        return s.override
    return f"{s.override}, from {v['min']:g} to {v['max']:g}"


def _percent_hint(s: Setting, own: Any) -> str:
    """95 typed where 0.95 is meant: say so."""
    v = s.valid or {}
    if v.get("max", 2) < 1 and isinstance(own, (int, float)) and v["max"] < own <= 1:
        return f" The most it takes is {v['max']:g}."
    if v.get("max", 2) < 1 and isinstance(own, (int, float)) and 1 < own <= 100:
        if v.get("min", 0) <= own / 100 <= v["max"]:
            return f" For {own:g}%, type {own / 100:g}."
        return f" Type a share between {v['min']:g} and {v['max']:g}: for 15%, type 0.15."
    return ""


def as_date(v: Any):
    """A typed date as a date: an Excel date cell, or text written 2026-06-30. None otherwise."""
    from datetime import date, datetime
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        try:
            return date.fromisoformat(v.strip())
        except ValueError:
            return None
    return None


def row_of(ws, key: str) -> int | None:
    """The row a setting (or a new-column slot) sits on, by its key."""
    for r in ws.iter_rows(min_row=FIRST_ROW):
        if r[KEY_COL - 1].value == key:
            return r[0].row
    return None


# --------------------------------------------------------------------------
# New columns (fix 3.9): a small block under the settings. Each row names a
# column made by dividing one of the extract's columns by another. Set up
# writes the block with the extract's number columns in the dropdowns and puts
# every answer back; the run reads it and refuses a half-filled row by cell.

NUMBER_LIST_COL = 10              # J on the hidden options tab: the dropdowns' column names


def write_derived(wb: Workbook, number_columns: list[str], kept: dict[int, tuple] | None = None) -> None:
    """The block, under the last setting. `kept` maps a slot (1 to DERIVED_ROWS)
    to the (name, top, bottom) already typed there."""
    ws, opt = wb[SHEET], wb[OPTIONS_SHEET]
    kept = kept or {}
    opt.cell(row=1, column=NUMBER_LIST_COL, value="number columns")
    for i, c in enumerate(number_columns, start=2):
        opt.cell(row=i, column=NUMBER_LIST_COL, value=c)
    last = max(r[0].row for r in ws.iter_rows(min_row=FIRST_ROW) if r[KEY_COL - 1].value)
    r = last + 2
    ws.cell(row=r, column=2, value="New columns: one column divided by another").font = \
        Font(name="Calibri", bold=True, color=INK)
    for col in range(2, 7):
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=CANVAS)
    ws.cell(row=r, column=KEY_COL, value=f"{DERIVED_KEY}|head")
    ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 1, end_column=6)
    note = ws.cell(row=r + 1, column=2, value=DERIVED_NOTE)
    note.font = Font(name="Calibri", size=10, color=SLATE)
    note.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[r + 1].height = 30
    ws.cell(row=r + 1, column=KEY_COL, value=f"{DERIVED_KEY}|note")
    for i, h in enumerate(("", "Name", "Top (divided)", "Bottom (divided by)", "What it makes"), start=2):
        c = ws.cell(row=r + 2, column=i, value=h or None)
        c.font = Font(name="Calibri", bold=True, color=PAPER)
        c.fill = PatternFill("solid", fgColor=INK)
    ws.cell(row=r + 2, column=KEY_COL, value=f"{DERIVED_KEY}|cols")
    dv = None
    if number_columns:
        dv = DataValidation(type="list", formula1=f"='{OPTIONS_SHEET}'!$J$2:$J${len(number_columns) + 1}",
                            allow_blank=True, showErrorMessage=True)
        dv.errorTitle = "Pick a column"
        dv.error = "Pick one of the extract's number columns from the list."
        ws.add_data_validation(dv)
    thin = Side(style="thin", color=MIST)
    for slot in range(1, DERIVED_ROWS + 1):
        row = r + 2 + slot
        ws.cell(row=row, column=2, value=f"New column {slot}")
        name, top, bottom = kept.get(slot, (None, None, None))
        ws.cell(row=row, column=3, value=name)
        ws.cell(row=row, column=4, value=top)
        ws.cell(row=row, column=5, value=bottom)
        if dv is not None:
            dv.add(ws.cell(row=row, column=4))
            dv.add(ws.cell(row=row, column=5))
        ws.cell(row=row, column=6, value=(
            f'=IF(COUNTA(C{row}:E{row})=0,"",IF(COUNTA(C{row}:E{row})<3,"Needs a name, a top and a bottom.",'
            f'C{row}&" = "&D{row}&" ÷ "&E{row}&" on each loan."))'))
        ws.cell(row=row, column=KEY_COL, value=f"{DERIVED_KEY}|{slot}")
        for col in range(2, 8):
            cell = ws.cell(row=row, column=col)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if col != KEY_COL:
                cell.font = Font(name="Calibri", size=10, color=INK)
        ws.cell(row=row, column=KEY_COL).font = Font(name="Consolas", size=8, color=SLATE)
        # a half-filled row is shaded where it still needs something; an empty one is left alone
        ws.conditional_formatting.add(f"C{row}:E{row}", FormulaRule(
            formula=[f'AND(C{row}="",COUNTA($C{row}:$E{row})>0)'],
            fill=PatternFill("solid", fgColor=NEEDS, bgColor=NEEDS)))
    ws.print_area = f"B1:F{r + 2 + DERIVED_ROWS}"


def read_derived(ws) -> tuple[list[dict], list[str]]:
    """Each filled new-column row as {slot, row, name, top, bottom}, and a
    problem for each half-filled one, named by cell."""
    out, problems = [], []
    for r in ws.iter_rows(min_row=FIRST_ROW):
        key = r[KEY_COL - 1].value
        if not (isinstance(key, str) and key.startswith(f"{DERIVED_KEY}|") and key.split("|")[1].isdigit()):
            continue
        row = r[0].row
        vals = [r[c - 1].value for c in (3, 4, 5)]
        vals = [str(v).strip() if v not in (None, "") and str(v).strip() else None for v in vals]
        if not any(vals):
            continue
        if not all(vals):
            lacking = [w for w, v in zip(("a name", "a top", "a bottom"), vals) if not v]
            col = "CDE"[[v is None for v in vals].index(True)]
            problems.append(f"{SHEET}!{col}{row}: new column {key.split('|')[1]} needs "
                            f"{' and '.join(lacking)}. Fill it in, or clear the row.")
            continue
        out.append({"slot": int(key.split("|")[1]), "row": row, "name": vals[0], "top": vals[1], "bottom": vals[2]})
    return out, problems
