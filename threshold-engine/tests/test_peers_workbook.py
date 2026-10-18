"""The peer workbook: one dashboard tab (Peers) over a hidden Workings tab.

Excel's live composite, band, best and worst, book-size band, lagged pack
total and normalised pack composite must equal peers.books and
peers.analyses (themselves tied out to the earlier build_peers.py), as built
and after the Yes/No cells, the product, the charge-off rate and the
delinquency pickers are changed. Recalculated with the `formulas` engine, so
no copy of Excel is needed; a skip is reported, never counted as a pass. The
roster pull is tests/data/peers/roster_2006.json, with the earlier build's
settings. tests/test_peers_dashboard.py holds every view against the old
one-tab-per-view builder."""
import datetime as dt
import json
import os
import re
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from conftest import DATA, as_earlier_build, ROOT
from threshold_engine import reporting_changes as RC
from threshold_engine.peers import roster as R
from threshold_engine.peers import delinquency as D
from threshold_engine.peers import peers_workbook as PW
from threshold_engine.peers.analyses import lagged_rate, normalisation, organic_index
from threshold_engine.peers.books import band, bank_book, best_worst, composite, industry
from threshold_engine.peers.glossary import PEER_TERMS
from threshold_engine.peers.peers_workbook import build_peers_workbook
from threshold_engine.profile import percentile

formulas = pytest.importorskip("formulas", reason="the formulas engine recalculates the workbook")

PEERS = DATA / "peers"
CHANGES = RC.load(ROOT / "config" / "reporting-changes.json")
FOUR = ("Credit card", "Mortgage", "Home equity", "Other consumer")
TWO = ("Credit card", "Mortgage")
KEYBANK, JPMORGAN, FIFTH_THIRD = "17534", "628", "6672"
#: The earlier build's settings: 25% jump, $1bn ($000), KeyBank exempt, five banks.
JUMP, MIN_BOOK, MIN_BANKS = 0.25, 1_000_000, 5
RANK = dict(rank_by="net", last_n=20, min_obs=12)
EXTRAS = {FIFTH_THIRD: ["2022Q2"]}
#: No acquisition brought in the share in these tests: every product, none.
NONE_ACQUIRED = {p: {} for p in ("Credit card", "Mortgage", "Home equity", "Other consumer", "Auto")}
DQ_RULES = dict(jump=0.25, min_book=1_000_000, exempt_self=True, acquired=NONE_ACQUIRED)
#: The pickers and settings on Peers (fixed by peers_workbook.PICKERS and SETTINGS).
PRODUCT, RATE, PEER_A, PEER_B, WINDOW, DQ_MEASURE, DQ_BASIS = "B5", "C5", "D5", "E5", "F5", "J5", "K5"


def recalc(path):
    """Every cell's value after a full recalculation, keyed 'SHEET!A1'."""
    model = formulas.ExcelModel().loads(str(path)).finish()
    sol = model.calculate()
    out = {}
    for key, val in sol.items():
        ref = key.split("]")[-1].upper()
        v = getattr(val, "value", val)
        try:
            v = v.item(0) if hasattr(v, "item") else v
        except Exception:                          # noqa: BLE001
            pass
        out[ref.replace("'", "")] = v
    return out


def num(v):
    """A recalculated number, or None for #N/A, blank or text."""
    if isinstance(v, bool) or not isinstance(v, (int, float)):
        return None
    return float(v)


@pytest.fixture(scope="module")
def data():
    return as_earlier_build(json.loads((PEERS / "roster_2006.json").read_text()))


@pytest.fixture(scope="module")
def roster():
    return R.load(ROOT / "config" / "peers.json")


@pytest.fixture(scope="module")
def books(data, roster):
    return {p: {b["cert"]: bank_book(data, b["cert"], p, JUMP, MIN_BOOK,
                                     exempt_min_book=b["cert"] == KEYBANK, acquired=())
                for b in roster} for p in FOUR}


def MERGERS():
    return json.loads((PEERS / "mergers_roster.json").read_text())


def make(path, data, roster, pack, products=FOUR, base="2019Q4", **over):
    kw = dict(exempt_self=True, min_banks=MIN_BANKS, index_min_book=500_000, lagged_years=(2019, 2025),
              lag_years=2, normalisation_years=(2019, 2025), mix_base_year=2019, industry=None,
              delinquency=None, acquired=NONE_ACQUIRED, combined={}, changes=CHANGES, **RANK)
    kw.update(over)
    return build_peers_workbook(path, data, roster, products, JUMP, MIN_BOOK, pack, MERGERS(), EXTRAS, base, **kw)


def typed(src, dst, cells):
    """A copy of ``src`` with ``cells`` ({'A1': value} on Peers) typed in, recalculated."""
    wb = load_workbook(src)
    for ref, v in cells.items():
        wb["Peers"][ref] = v
    wb.save(dst)
    return recalc(dst)


@pytest.fixture(scope="module")
def recent(data):
    """2017 on: enough for the 2019 base and the 2019 to 2025 views, and quick to build."""
    return {q: rows for q, rows in data.items() if q >= "20170331"}


@pytest.fixture(scope="module")
def all_peers(tmp_path_factory, data, roster):
    p = tmp_path_factory.mktemp("peers") / "all.xlsx"
    info = make(p, data, roster, "all")
    return p, info, recalc(p)


def inc(info, short):
    return info["layout"]["include"][short]


def wcol(vals, info, view, name, offset=0):
    """A live helper column on Workings, one value per row of the view."""
    c = get_column_letter(column_index_from_string(info[view]["helpers"][name]) + offset)
    n = len(info[view]["quarters"])
    return [vals.get("WORKINGS!%s%d" % (c, 2 + k)) for k in range(n)]


def scalar(vals, info, key):
    return vals.get("WORKINGS!%s" % info["scalars"][key].replace("$", ""))


def same(got, want, places=9):
    """Equal within rounding, a missing value (#N/A) matching None."""
    if want is None:
        return num(got) is None
    return num(got) is not None and abs(num(got) - want) < 10 ** -places


def assert_series(got, want, what):
    bad = [(i, g, w) for i, (g, w) in enumerate(zip(got, want)) if not same(g, w)]
    assert len(got) == len(want)
    assert not bad, (what, bad[:5])


def assert_composite(vals, info, books_p, certs, kind="net"):
    want = composite(books_p, certs, min_banks=MIN_BANKS)
    assert [r["quarter"] for r in want] == info["charge_offs"]["quarters"]
    assert [num(v) for v in wcol(vals, info, "charge_offs", "banks")] == [r["banks"] for r in want]
    assert_series(wcol(vals, info, "charge_offs", "c"), [r[kind + "_rate"] for r in want], "composite " + kind)
    b_ = band(books_p, certs, kind, min_banks=MIN_BANKS)
    for pct in (10, 25, 75, 90):
        assert_series(wcol(vals, info, "charge_offs", "p%d" % pct), [r["p%d" % pct] for r in b_],
                      "%s p%d" % (kind, pct))
    return want


def short_of(roster, cert):
    return next(b["short"] for b in roster if b["cert"] == cert)


def peers_cell(vals, ref):
    return vals.get("PEERS!" + ref)


def run_rows(wb):
    """The standalone Run tab's rows from column B, where its sections start."""
    return [r[1:] for r in wb["Run"].iter_rows(values_only=True)]


def glossary_rows(wb):
    """The standalone Glossary tab's (term, meaning, calculation) rows, from column B."""
    return [r[1:4] for r in wb["Glossary"].iter_rows(values_only=True)]


def assert_rank_table(vals, info, books_p, certs, roster, kind="net"):
    """Where each bank stands on charge-offs: the banks in ``certs`` kept in the
    latest quarter, and KeyBank, by rate, a peer level with KeyBank after it;
    each row's rate, its distance from the composite and the In pack mark."""
    order = [b["cert"] for b in roster]
    short = {b["cert"]: b["short"] for b in roster}
    last = {c: (books_p[c][-1][kind + "_rate"] if books_p[c][-1]["kept"] else None) for c in order}
    ranked = sorted((c for c in order if (c in certs or c == KEYBANK) and last[c] is not None),
                    key=lambda c: (last[c], c != KEYBANK, order.index(c)))
    comp = composite(books_p, certs, MIN_BANKS)[-1][kind + "_rate"]
    head = info["layout"]["rank_head"]
    for k, c in enumerate(ranked):
        r = head + 1 + k
        assert num(peers_cell(vals, "B%d" % r)) == k + 1 and peers_cell(vals, "C%d" % r) == short[c], (k, c)
        assert num(peers_cell(vals, "D%d" % r)) == pytest.approx(last[c], abs=1e-12)
        assert num(peers_cell(vals, "E%d" % r)) == pytest.approx(last[c] - comp, abs=1e-9)
        assert (peers_cell(vals, "H%d" % r) == "—") if c == KEYBANK else \
            peers_cell(vals, "H%d" % r).startswith("Yes")
    assert all(peers_cell(vals, "%s%d" % (c_, head + 1 + k)) in ("", None)
               for k in range(len(ranked), len(roster)) for c_ in "BCD")
    me = ranked.index(KEYBANK) + 1
    assert peers_cell(vals, "F%d" % (info["layout"]["tiles"] + 1)) == "%d of %d" % (me, len(ranked))


# --- the book-size view ---------------------------------------------------------

def test_book_size_is_a_full_view_of_every_bank_switched_on(all_peers, books, roster):
    """Book size reads like charge-offs: the pack's percentiles of the balances
    above zero as bands and a median, KeyBank, Peer A and B as lines on a
    logarithmic axis (nothing at or below zero is drawn, or Excel warns in a
    dialog), the strip, and every bank switched on, and KeyBank, ranked by
    balance beside it, largest first."""
    p, info, vals = all_peers
    grid = info["quarters"]
    bal = {c: {r["quarter"]: r["avg_balance"] for r in rows} for c, rows in books["Credit card"].items()}
    peers = [b["cert"] for b in R.peers(roster)]
    for pct in (10, 25, 50, 75, 90):
        want = []
        for q in grid:
            v = [bal[c][q] / 1e6 for c in peers if bal[c].get(q)]
            want.append(percentile(v, pct) if len(v) >= MIN_BANKS else None)
        assert_series(wcol(vals, info, "charge_offs", "s%d" % pct), want, "size p%d" % pct)
    assert_series(wcol(vals, info, "charge_offs", "bk_k"),
                  [bal[KEYBANK][q] / 1e6 if bal[KEYBANK].get(q) else None for q in grid], "KeyBank balance")
    for name in ("s10", "s50", "s90", "bk_k", "bk_a", "bk_b", "smark"):
        assert all(num(v) is None or num(v) > 0 for v in wcol(vals, info, "charge_offs", name)), name

    order = [b["cert"] for b in roster]
    last, before = grid[-1], grid[-5]
    ranked = sorted((c for c in order if bal[c].get(last) is not None),
                    key=lambda c: (-bal[c][last], c != KEYBANK, order.index(c)))
    total = sum(bal[c][last] for c in ranked)
    lay = info["layout"]
    for k, c in enumerate(ranked):
        r = lay["size_head"] + 1 + k
        assert num(peers_cell(vals, "J%d" % r)) == k + 1 and peers_cell(vals, "K%d" % r) == short_of(roster, c)
        assert num(peers_cell(vals, "L%d" % r)) == pytest.approx(bal[c][last] / 1e6)
        assert num(peers_cell(vals, "M%d" % r)) == pytest.approx(bal[c][last] / total)
        if bal[c].get(before):
            assert num(peers_cell(vals, "N%d" % r)) == pytest.approx(bal[c][last] / bal[c][before] - 1)
    me = ranked.index(KEYBANK) + 1
    assert peers_cell(vals, "F%d" % (lay["size_tiles"] + 1)) == "%d of %d" % (me, len(ranked))
    assert num(peers_cell(vals, "B%d" % (lay["size_tiles"] + 1))) == pytest.approx(bal[KEYBANK][last] / 1e6)
    with zipfile.ZipFile(p) as z:
        xml = next(x for x in (z.read(n).decode("utf-8") for n in z.namelist()
                               if re.match(r"xl/charts/chart\d+\.xml$", n)) if "Peers!co_s10" in x)
    assert '<logBase val="10"' in xml
    assert re.findall(r'<legendEntry><idx val="(\d+)" /><delete val="1" />', xml) == ["0", "1", "2", "3", "8"]


# --- the charge-off view --------------------------------------------------------

def test_all_peers_composite_and_band_equal_books(all_peers, books, roster):
    _, info, vals = all_peers
    certs = [b["cert"] for b in R.peers(roster)]
    assert info["pack"] == certs and len(certs) == 15
    want = assert_composite(vals, info, books["Credit card"], certs)
    # 2007Q1 to 2010Q3 have fewer than five kept card books, so #N/A there
    # is checked as well as the 63 quarters with a composite.
    assert sum(r["net_rate"] is not None for r in want) == 63
    assert_series(wcol(vals, info, "charge_offs", "k"), [r["net_rate"] for r in books["Credit card"][KEYBANK]], "k")
    # No industry was supplied: the line is #N/A throughout and the tile says so.
    assert all(num(v) is None for v in wcol(vals, info, "charge_offs", "ind"))
    t = info["layout"]["tiles"]
    assert peers_cell(vals, "H%d" % (t + 1)) == "none" and peers_cell(vals, "H%d" % (t + 2)) == "not supplied"


def test_the_headline_numbers_are_the_latest_quarter(all_peers, books, roster):
    _, info, vals = all_peers
    t = info["layout"]["tiles"]
    card = books["Credit card"]
    certs = [b["cert"] for b in R.peers(roster)]
    me, pack = card[KEYBANK][-1]["net_rate"], composite(card, certs, MIN_BANKS)[-1]
    assert num(peers_cell(vals, "B%d" % (t + 1))) == pytest.approx(me, abs=1e-12)
    assert num(peers_cell(vals, "D%d" % (t + 1))) == pytest.approx(pack["net_rate"], abs=1e-12)
    rates = [card[c][-1]["net_rate"] for c in certs if card[c][-1]["kept"]]
    assert peers_cell(vals, "F%d" % (t + 1)) == "%d of %d" % (sum(r < me for r in rates) + 1, len(rates) + 1)
    assert peers_cell(vals, "D%d" % (t + 2)) == "%d banks in, %s" % (pack["banks"], info["quarters"][-1])


def test_best_and_worst_and_the_ranking_window_equal_books(all_peers, books, roster):
    _, info, vals = all_peers
    peers = R.peers(roster)
    certs = [b["cert"] for b in peers]
    best, worst = best_worst(books["Credit card"], certs, "net", 20, 12)
    assert (scalar(vals, info, "best"), scalar(vals, info, "worst")) == (short_of(roster, best),
                                                                          short_of(roster, worst))
    assert scalar(vals, info, "a_name") == short_of(roster, worst)      # Peer A starts as Worst in pack
    assert scalar(vals, info, "b_name") == short_of(roster, best)
    for name, cert in (("a", worst), ("b", best)):
        assert_series(wcol(vals, info, "charge_offs", name),
                      [r["net_rate"] for r in books["Credit card"][cert]], name)
    for j, b in enumerate(peers):
        rates = [r["net_rate"] for r in books["Credit card"][b["cert"]][-20:] if r["net_rate"] is not None]
        assert num(vals["WORKINGS!N%d" % (2 + j)]) == len(rates), b["short"]
        mean = vals["WORKINGS!O%d" % (2 + j)]
        if len(rates) >= 12:
            assert num(mean) == pytest.approx(sum(rates) / len(rates), abs=1e-9), b["short"]
        else:
            assert num(mean) is None, b["short"]
    assert "Best in pack: %s. Worst in pack: %s." % (short_of(roster, best), short_of(roster, worst)) in \
        peers_cell(vals, "B%d" % info["layout"]["best_worst"])


# --- switching a bank off ---------------------------------------------------------

@pytest.fixture(scope="module")
def jpm_off(all_peers, tmp_path_factory):
    p, info, _ = all_peers
    return typed(p, tmp_path_factory.mktemp("off") / "off.xlsx", {inc(info, "JPMorgan"): "No"})


def moved(before, after):
    return sum(1 for a, b in zip(before, after) if num(a) is not None and num(b) is not None
               and abs(num(a) - num(b)) > 1e-9)


def test_switching_jpmorgan_off_moves_every_pack_view(all_peers, jpm_off, books, data, roster):
    """JPMorgan to No: the charge-off composite and band, best and worst, the
    book-size band and median, the lagged pack total, the normalised pack
    composite and the headline numbers all leave it out, each equal to the
    module over the other fourteen; and each of them moved."""
    _, info, before = all_peers
    after = jpm_off
    peers = R.peers(roster)
    without = [b["cert"] for b in peers if b["cert"] != JPMORGAN]
    card = books["Credit card"]
    assert_composite(after, info, card, without)
    assert moved(wcol(before, info, "charge_offs", "c"), wcol(after, info, "charge_offs", "c")) > 40
    best, worst = best_worst(card, without, "net", 20, 12)
    assert (scalar(after, info, "best"), scalar(after, info, "worst")) == (short_of(roster, best),
                                                                         short_of(roster, worst))
    # The book size: the band and median of the indexes of the banks still in.
    quarters = info["book"]["quarters"]
    mergers = MERGERS()
    idx = {}
    for b in peers:
        acq = list(mergers.get(b["cert"], [])) + EXTRAS.get(b["cert"], [])
        try:
            rows = organic_index(data, b["cert"], "Credit card", "2019Q4", acq)
        except ValueError:
            continue
        if next(x["balance"] for x in rows if x["date"] == "2019-12-31") >= 500_000:
            idx[b["cert"]] = {x["quarter"]: x["organic"] for x in rows}
    for name, p in (("p10", 10), ("p50", 50), ("p90", 90)):
        for certs, vals in ((list(idx), before), ([c for c in idx if c != JPMORGAN], after)):
            want = []
            for q in quarters:
                v = [idx[c][q] for c in certs if idx[c].get(q) is not None]
                want.append(percentile(v, p) if len(v) >= MIN_BANKS else None)
            assert_series(wcol(vals, info, "book", name), want, "book " + name)
    assert moved(wcol(before, info, "book", "p50"), wcol(after, info, "book", "p50")) > 10
    # The lagged pack total, and the normalised pack composite.
    lay = info["layout"]
    for j, y in enumerate((2019, 2025)):
        cell = "%s%d" % ("DE"[j], lay["lagged"] + 4)
        for certs, vals in (([b["cert"] for b in peers], before), (without, after)):
            assert num(peers_cell(vals, cell)) == pytest.approx(
                lagged_rate(data, certs, "Credit card", y, 2)["rate"], abs=1e-9), (y, cell)
        assert num(peers_cell(before, cell)) != pytest.approx(num(peers_cell(after, cell)), abs=1e-6)
    assert peers_cell(after, "B%d" % (lay["lagged"] + 4)) == "Pack total (14 banks switched on)"
    r = lay["norm_table"] + 4
    for certs, vals in (([b["cert"] for b in peers], before), (without, after)):
        comp = composite(card, certs, MIN_BANKS)
        want = normalisation([(x["quarter"], x["net_rate"]) for x in comp], 2019, 2025)
        for c, key in (("L", "base"), ("M", "end"), ("N", "ratio")):
            assert num(peers_cell(vals, "%s%d" % (c, r))) == pytest.approx(want[key], abs=1e-9), key
    assert num(peers_cell(before, "N%d" % r)) != pytest.approx(num(peers_cell(after, "N%d" % r)), abs=1e-6)
    # The headline pack number moves; KeyBank's own does not.
    t = lay["tiles"] + 1
    assert num(peers_cell(before, "D%d" % t)) != pytest.approx(num(peers_cell(after, "D%d" % t)), abs=1e-6)
    assert num(peers_cell(before, "B%d" % t)) == num(peers_cell(after, "B%d" % t))
    # The rank table leaves JPMorgan out, and is the module's order over the other fourteen and KeyBank.
    assert_rank_table(after, info, card, without, roster)
    assert_rank_table(before, info, card, [b["cert"] for b in peers], roster)


# --- another product, the gross rate, the regionals --------------------------------

@pytest.fixture(scope="module")
def regionals_gross(all_peers, tmp_path_factory, roster):
    """Other consumer (the fourth product, so the picker's offset is exercised),
    the gross rate, and only the regionals switched on."""
    p, info, _ = all_peers
    cells = {PRODUCT: "Other consumer", RATE: "Gross"}
    for b in R.peers(roster):
        cells[inc(info, b["short"])] = "Yes" if b["regional"] else "No"
    return typed(p, tmp_path_factory.mktemp("reg") / "reg.xlsx", cells)


def test_the_product_rate_and_yes_no_cells_move_the_lines(all_peers, regionals_gross, books, roster):
    _, info, _ = all_peers
    vals = regionals_gross
    certs = [b["cert"] for b in R.regionals(roster)]
    assert len(certs) == 10
    want = assert_composite(vals, info, books["Other consumer"], certs, kind="gross")
    assert any(r["gross_rate"] is not None for r in want)
    # Best and worst are still ranked on the net rate (the ranking setting), among the regionals.
    best, worst = best_worst(books["Other consumer"], certs, "net", 20, 12)
    assert scalar(vals, info, "a_name") == short_of(roster, worst)
    for name, cert in (("k", KEYBANK), ("a", worst), ("b", best)):
        assert_series(wcol(vals, info, "charge_offs", name),
                      [r["gross_rate"] for r in books["Other consumer"][cert]], name)
    assert peers_cell(vals, "B%d" % info["layout"]["v1"]) == \
        "Charge-offs: other consumer, gross rate, trailing four quarters (%)"


def test_a_named_peer_is_drawn_switched_off_and_min_obs_is_met_at_exactly(all_peers, books, roster, tmp_path):
    """Peer A named Fifth Third and Fifth Third switched off: it is drawn, but
    not ranked. Capital One has exactly 10 kept card quarters in the last 20,
    so with 10 needed it is ranked, and best and worst follow books."""
    p, info, _ = all_peers
    peers = R.peers(roster)
    certs = [b["cert"] for b in peers]
    rates = {c: [r["net_rate"] for r in books["Credit card"][c][-20:] if r["net_rate"] is not None] for c in certs}
    edge = [c for c in certs if 0 < len(rates[c]) < 20 and c != FIFTH_THIRD]
    assert edge, "no peer sits between 0 and 20 kept quarters, so the boundary is not tested"
    need = len(rates[edge[0]])
    vals = typed(p, tmp_path / "pick.xlsx", {PEER_A: "Fifth Third", inc(info, "Fifth Third"): "No",
                                             info["layout"]["settings_cells"]["min_obs"]: need})
    assert scalar(vals, info, "a_name") == "Fifth Third"
    assert_series(wcol(vals, info, "charge_offs", "a"), [r["net_rate"] for r in books["Credit card"][FIFTH_THIRD]],
                  "A")
    j = certs.index(edge[0])
    assert num(vals["WORKINGS!O%d" % (2 + j)]) == pytest.approx(sum(rates[edge[0]]) / need, abs=1e-9)
    on = [c for c in certs if c != FIFTH_THIRD]
    best, worst = best_worst(books["Credit card"], on, "net", 20, need)
    assert (scalar(vals, info, "best"), scalar(vals, info, "worst")) == (short_of(roster, best),
                                                                       short_of(roster, worst))


# --- lagged, book size, normalisation, mix ------------------------------------------

def lagged_values(vals_or_ws, info, roster, product_i, year):
    """Each bank's lagged rate for one product and year, from Workings' values block."""
    c = column_index_from_string(info["lagged"]["helpers"]["v_%d_%d" % (product_i, year)])
    return {b["short"]: vals_or_ws.cell(2 + i, c).value for i, b in enumerate(roster)}


def test_lagged_values_carry_the_acceptance_numbers_and_equal_analyses(all_peers, data, roster):
    p, info, vals = all_peers
    ws = load_workbook(p)["Workings"]
    # docs/PEERS-SPEC.md: card, 2025 net over the 2023 book; 2019 over 2017.
    card19, card25 = (lagged_values(ws, info, roster, 0, y) for y in (2019, 2025))
    for name, y2019, y2025 in (("Fifth Third", 6.18, 3.47), ("KeyBank", 3.52, 3.82), ("Huntington", 7.89, 4.92)):
        assert (round(card19[name], 2), round(card25[name], 2)) == (y2019, y2025)
    for pi, product in enumerate(FOUR):
        for y in (2019, 2025):
            got = lagged_values(ws, info, roster, pi, y)
            for b in roster:
                want = lagged_rate(data, b["cert"], product, y, 2)["rate"]
                assert (got[b["short"]] is None) if want is None else got[b["short"]] == pytest.approx(want)
    r = info["layout"]["lagged"] + 1
    assert peers_cell(vals, "B%d" % r) == "KeyBank" and round(num(peers_cell(vals, "D%d" % r)), 2) == 3.52
    assert peers_cell(vals, "B%d" % (r + 3)) == "Pack total (15 banks switched on)"


def test_the_book_index_is_on_the_data_tab_and_carries_the_acceptance_numbers(all_peers, data, roster):
    """Data's two index columns follow organic_index from the base, for every
    bank indexed and every product; a bank under the minimum is not indexed
    and the bank table says why; the acceptance numbers at 2026Q2."""
    p, info, vals = all_peers
    ws = load_workbook(p)["Data"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    by = {(r[0], r[1], r[3]): r for r in rows}
    for name, rep, org in (("M&T", 156, 145), ("PNC", 100, 80), ("KeyBank", 82, 82), ("Fifth Third", 66, 67)):
        r = by["Credit card", name, "2026Q2"]
        assert (round(r[11]), round(r[12])) == (rep, org)
    mergers = MERGERS()
    for product in FOUR:
        for b in roster:
            acq = list(mergers.get(b["cert"], [])) + EXTRAS.get(b["cert"], [])
            try:
                idx = organic_index(data, b["cert"], product, "2019Q4", acq)
            except ValueError:
                idx = None
            base = next((x["balance"] for x in idx if x["date"] == "2019-12-31"), None) if idx else None
            ours = [r for r in rows if r[0] == product and r[1] == b["short"]]
            if idx is None or base < 500_000:
                assert all(r[11] is None and r[12] is None for r in ours), (product, b["short"])
                continue
            want = {x["quarter"]: x for x in idx if x["date"] >= "2019-12-31"}
            for r in ours:
                x = want.get(r[3])
                if x is None:
                    assert r[11] is None and r[12] is None
                else:
                    for got, w in ((r[11], x["reported"]), (r[12], x["organic"])):
                        assert (got is None) if w is None else got == pytest.approx(w), (product, b["short"], r[3])
    # Zions' card book is under the smallest indexed at the base: the book index says so for it.
    wk = load_workbook(p)["Workings"]
    notes = column_index_from_string(info["banks"]["helpers"]["notes"])
    zions = [b["short"] for b in roster].index("Zions")
    assert wk.cell(2 + zions, notes).value == "too small"
    bt = info["layout"]["book_table"]
    assert peers_cell(vals, "J%d" % (bt + 1)) == "KeyBank"
    assert num(peers_cell(vals, "M%d" % (bt + 1))) == pytest.approx(82, abs=0.5)     # as reported, as on Data
    run = run_rows(load_workbook(p))
    assert any(r[0] == "Fifth Third" and "2022Q2 (declared)" in str(r[1]) for r in run)


def test_normalisation_rows_equal_analyses(all_peers, books, roster):
    _, info, vals = all_peers
    r = info["layout"]["norm_table"] + 1
    card = books["Credit card"]
    who = {"KeyBank": KEYBANK, "a": None, "b": None}
    certs = [b["cert"] for b in R.peers(roster)]
    best, worst = best_worst(card, certs, "net", 20, 12)
    for i, cert in enumerate((KEYBANK, worst, best)):
        want = normalisation([(x["quarter"], x["net_rate"]) for x in card[cert]], 2019, 2025)
        for c, key in (("L", "base"), ("M", "end"), ("N", "ratio")):
            got = peers_cell(vals, "%s%d" % (c, r + i))
            assert (num(got) is None) if want[key] is None else num(got) == pytest.approx(want[key]), (cert, key)
    assert round(num(peers_cell(vals, "L%d" % r)), 2) == 3.29 and round(num(peers_cell(vals, "M%d" % r)), 2) == 4.30
    del who


def test_mix_says_when_no_industry_was_supplied(all_peers):
    _, info, vals = all_peers
    assert peers_cell(vals, "B%d" % info["layout"]["mix_peers"]).startswith(
        "No industry series was supplied for Credit card")
    assert all(peers_cell(vals, "G%d" % (info["layout"]["mix"] + 1 + k)) in ("", None)
               for k in range(len(info["mix"]["years"])))


def test_mix_values_with_an_industry(tmp_path, recent, roster):
    """The values the mix picks from, by product, are analyses' for KeyBank and
    every bank's own factor (read without a recalculation)."""
    from threshold_engine.peers.analyses import calendar_year, mix_factor
    ind = {p: industry(recent, p) for p in TWO}      # the roster's aggregate, standing in for a test
    p = tmp_path / "mix.xlsx"
    info = make(p, recent, roster, "all", products=TWO, industry=ind)
    ws = load_workbook(p)["Workings"]
    years = info["mix"]["years"]
    assert years == list(range(2019, 2027)) or years == list(range(2019, 2026))
    for pi, product in enumerate(TWO):
        ind_rate = {int(x["quarter"][:4]): x["net_rate"] for x in ind[product] if x["quarter"].endswith("Q4")}
        for i, b in enumerate(roster):
            rate = {y: lagged_rate(recent, b["cert"], product, y, 0)["rate"] for y in years}
            bal = {y: calendar_year(recent, b["cert"], product, y)["balance"] for y in years}
            rows = mix_factor(rate, {y: ind_rate.get(y) for y in years}, bal, 2019)
            c = column_index_from_string(info["mix"]["helpers"]["own_%d" % pi]) + i
            for x in rows:
                got = ws.cell(2 + years.index(x["year"]), c).value
                assert (got is None) if x["own_factor"] is None else got == pytest.approx(x["own_factor"])
            if b["cert"] == KEYBANK:
                c0 = column_index_from_string(info["mix"]["helpers"]["m_%d" % pi])
                for x in rows:
                    got = [ws.cell(2 + years.index(x["year"]), c0 + j).value for j in range(6)]
                    want = [rate[x["year"]], ind_rate.get(x["year"]), x["bank_ratio"], x["industry_ratio"],
                            x["own_factor"], x["balance_factor"]]
                    for g, w in zip(got, want):
                        assert (g is None) if w is None else g == pytest.approx(w)
    status = column_index_from_string(info["mix"]["helpers"]["status"])
    assert [ws.cell(2 + i, status).value for i in range(2)] == ["OK", "OK"]


# --- the glossary, the data, the structure ------------------------------------------

def test_every_glossary_term_is_on_the_glossary_tab(all_peers):
    p, _, _ = all_peers
    on_tab = {r[0] for r in glossary_rows(load_workbook(p))}
    terms = [t for t, _, _ in PEER_TERMS]
    assert set(terms) <= on_tab
    for must in ("Structural move", "Pack", "In the pack", "Chart shows", "Composite", "Band", "Lagged rate",
                 "Organic index", "Own factor", "Book index pack median", "Pack total (lagged)", "Normalization",
                 "Rank tables", "Book size", "Charge-off rate \u00f7 pack composite", "Consumer book by product"):
        assert must in terms
    assert all(m and c for _, m, c in PEER_TERMS)
    assert not any("Pack tab" in m + c or "Include" in m + c for _, m, c in PEER_TERMS)


def test_data_tab_holds_every_bank_product_and_quarter(all_peers, books, roster):
    p, info, _ = all_peers
    ws = load_workbook(p)["Data"]
    assert info["data_rows"] == 4 * 16 * len(info["quarters"]) == ws.max_row - 1
    want = books["Credit card"][KEYBANK][0]
    first = [c.value for c in ws[2 + 8 * len(info["quarters"])]]        # KeyBank is ninth on the roster
    assert first[:4] == ["Credit card", "KeyBank", KEYBANK, want["quarter"]]
    assert first[9] is want["kept"] and first[10] == want["why"]


def test_one_dashboard_a_hidden_workings_and_charts_over_named_ranges(all_peers):
    """Peers is the one tab the bank looks at; Workings is hidden; the tabs
    carry the design's colours and Peers its column widths. Each chart reads
    names on Peers (OFFSET over Workings, from the first row drawn, as many
    rows as are drawn); the main chart's band is up-down bars between
    invisible lines, then the lines in the design's colours and widths, with
    KeyBank last and a 6 pt circle on its last point; the legend is at the
    top with the helper series' keys deleted; #N/A is a gap; the hidden tab
    is plotted."""
    p, info, _ = all_peers
    wb = load_workbook(p)
    assert wb.sheetnames == list(PW.TABS)
    assert wb["Workings"].sheet_state == "hidden" and wb["Peers"].sheet_state == "visible"
    assert wb.active.title == "Peers"
    colours = {t: wb[t].sheet_properties.tabColor.rgb[-6:] for t in PW.TABS if t != "Workings"}
    assert colours == {"Peers": "0A0908", "Glossary": "B9B4AC", "Run": "B9B4AC", "Data": "E4DFD5",
                       "DelinquencyData": "E4DFD5"}
    pe = wb["Peers"]
    px = {c: round(pe.column_dimensions[c].width * 7) for c in "ABCDEFGHIJKLMNOPQ"}
    assert px == dict({c: 130 for c in "BCDEFGHJKLMNOP"}, A=16, I=24, Q=8)
    names = pe.defined_names
    f = names["co_c"].attr_text
    assert f.startswith("OFFSET(Workings!$") and "Workings!%s-1" % info["scalars"]["co_first"] in f
    assert f.endswith(",Workings!%s,1)" % info["scalars"]["co_count"])
    with zipfile.ZipFile(p) as z:
        charts = [z.read(n).decode("utf-8") for n in sorted(z.namelist())
                  if re.match(r"xl/charts/chart\d+\.xml$", n)]
    # Charge-offs, book size, the rate over the composite, the mix; delinquency was not built.
    assert len(charts) == 4
    for xml in charts:
        assert 'dispNaAsBlank val="1"' in xml and '<dispBlanksAs val="gap"' in xml and '<plotVisOnly val="0"' in xml
        assert '<legendPos val="t"' in xml and "<barChart>" not in xml
        assert all(c == "Peers!co_q" for c in re.findall(r"<cat><strRef><f>([^<]*)</f>", xml))
        # Every series carries a value cache, or Excel drops the legend deletions.
        assert xml.count("<numRef>") == xml.count("<numCache>") > 0
    main = next(x for x in charts if "Peers!co_p10" in x)
    groups = re.findall(r"<lineChart>.*?</lineChart>", main)
    assert len(groups) == 3 and "<areaChart>" not in main
    for g, (lo, hi), fill in zip(groups[:2], ((10, 90), (25, 75)), ("E4DFD5", "CFC8BC")):
        assert "<upDownBars>" in g and fill in g
        assert re.findall(r"<val><numRef><f>Peers!([a-z]+)_p(\d+)</f>", g) == [("co", str(lo)), ("co", str(hi))]
        assert all("<a:ln><a:noFill" in x for x in re.findall(r"<ser>.*?</ser>", g))
    assert "<upDownBars>" not in groups[2]
    sers = re.findall(r"<ser>.*?</ser>", groups[2])
    lines = [re.search(r"<val><numRef><f>Peers!([a-z_0-9]+)</f>", s_).group(1) for s_ in sers]
    # Then the reporting change marker (16 Oct 2026): a 7 pt INK diamond, no line, kept in the legend.
    assert lines == ["co_c", "co_ind", "co_a", "co_b", "co_k", "co_mark", "co_chg"]
    assert '<symbol val="diamond" /><size val="7" />' in sers[6] and "<a:ln><a:noFill" in sers[6]
    assert "Workings!$%s$1" % info["charge_offs"]["helpers"]["chg"] in sers[6]
    style = {n: (re.search(r'<a:ln w="(\d+)"', s_), re.search(r'<a:srgbClr val="([0-9A-F]{6})"', s_),
                 re.search(r'<a:prstDash val="(\w+)"', s_)) for n, s_ in zip(lines, sers)}
    got = {n: (int(w.group(1)) if w else None, c.group(1) if c else None, d.group(1) if d else None)
           for n, (w, c, d) in style.items()}
    assert got["co_c"] == (25400, "0A0908", "solid") and got["co_k"] == (31750, "CC0000", "solid")
    assert got["co_ind"] == (19050, "8C877E", "solid")
    assert got["co_a"] == (19050, "57534B", "dash") and got["co_b"] == (19050, "57534B", "sysDot")
    assert '<symbol val="circle" /><size val="6" />' in sers[5]
    assert re.findall(r'<legendEntry><idx val="(\d+)" /><delete val="1" />', main) == ["0", "1", "2", "3", "9"]
    assert '<a:ln w="6350"><a:solidFill><a:srgbClr val="E4DFD5"' in main          # Y gridlines, MIST 0.5 pt
    small = {tuple(re.findall(r"<val><numRef><f>Peers!([a-z_0-9]+)</f>", x)): x for x in charts if x is not main}
    size = ("co_s10", "co_s90", "co_s25", "co_s75", "co_s50", "co_bk_a", "co_bk_b", "co_bk_k", "co_smark")
    assert set(small) == {size, ("co_rk", "co_ra", "co_rb", "co_one"),
                          tuple("co_mix%d" % j for j in range(1, len(FOUR) + 1))}
    ratio = small["co_rk", "co_ra", "co_rb", "co_one"]
    assert re.findall(r'<legendEntry><idx val="(\d+)" /><delete val="1" />', ratio) == ["3"]
    mix = small[tuple("co_mix%d" % j for j in range(1, len(FOUR) + 1))]
    assert '<areaChart><grouping val="percentStacked" />' in mix and "CC0000" in mix


def formulas_reaching_above_row_1(wb):
    """Formulas that do, or could at recalculation, refer above row 1: a
    literal row below 1 (what ``"A%d" % (r - 1)`` writes when r is 1), or an
    OFFSET or INDIRECT, whose reach a reading of the formula cannot see. No
    cell uses either; the charts' OFFSET names start at row 1 and move down."""
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    text = re.sub(r'"[^"]*"', "", c.value)            # not inside text
                    if re.search(r"\b(OFFSET|INDIRECT)\(", text, re.I) or any(
                            int(m.group(1)) < 1
                            for m in re.finditer(r"(?<![A-Z0-9_.])\$?[A-Z]{1,3}\$?(\d+)\b", text)):
                        bad.append((ws.title, c.coordinate, c.value))
    return bad


def test_no_formula_refers_to_a_row_above_row_1(all_peers):
    p, _, _ = all_peers
    wb = load_workbook(p)
    bad = formulas_reaching_above_row_1(wb)
    assert not bad, bad[:3]
    for name, dn in wb["Peers"].defined_names.items():
        assert re.fullmatch(r"OFFSET\(Workings!\$[A-Z]+\$1,Workings!\$B\$\d+-1,0,Workings!\$B\$\d+,1\)",
                            dn.attr_text), (name, dn.attr_text)


def test_the_row_check_can_fail():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "=SUM(B1:B3)"
    ws["A2"] = '="A0 in text"&B2'
    assert formulas_reaching_above_row_1(wb) == []
    for f in ("=B0+1", "=SUM($C$0:C4)", "=OFFSET(A1,-1,0)", '=INDIRECT("A"&ROW()-1)'):
        ws["A3"] = f
        assert formulas_reaching_above_row_1(wb), f


def test_the_yellow_cells_refuse_what_the_build_would(all_peers):
    """The counts refuse anything but a whole number of at least 1, as the
    build does; every picker is a list; a bank's cell takes Pack (follow
    Banks in the pack), Yes or No, and starts as Pack; Banks in the pack takes
    one of its four picks and starts as the build's pack."""
    p, info, _ = all_peers
    ws = load_workbook(p)["Peers"]
    dvs = ws.data_validations.dataValidation
    whole = [d for d in dvs if d.type == "whole"]
    assert len(whole) == 1
    d = whole[0]
    assert d.operator == "greaterThanOrEqual" and d.formula1 == "1" and not d.allow_blank and d.showErrorMessage
    cells = info["layout"]["settings_cells"]
    assert {str(c) for c in d.sqref.ranges} == {cells["last_n"], cells["min_obs"], cells["min_banks"]}
    lists = {str(r): dv.formula1 for dv in dvs if dv.type == "list" for r in dv.sqref.ranges}
    assert lists[PRODUCT].startswith("=Workings!") and lists[RATE] == '"Net,Gross"'
    assert lists[WINDOW].startswith("=Workings!") and lists[cells["rank_by"]] == '"gross,net"'
    assert all(lists[c] == '"Pack,Yes,No"' and ws[c].value == "Pack" for c in info["layout"]["include"].values())
    pack = info["layout"]["pickers"]["pack"]
    assert lists[pack] == '"%s"' % ",".join(PW.PACK_PICKS) and ws[pack].value == "Every peer"
    assert ws[DQ_MEASURE].value == "not built" and DQ_MEASURE not in lists
    assert ws[WINDOW].value == "Last 48 quarters"


def test_run_tab_states_the_rules_as_built(all_peers):
    p, _, _ = all_peers
    rows = {r[0]: r for r in run_rows(load_workbook(p)) if r and r[0]}
    assert rows["Jump"][1] == JUMP
    assert rows["Minimum book ($000)"][1] == MIN_BOOK and rows["Minimum book ($000)"][2] == "KeyBank exempt"
    assert rows["Banks needed for a composite or band (starting)"][1] == MIN_BANKS
    assert rows["Pack (starting)"][1] == "all"
    assert rows["Charts show (starting)"][1] == "Last 48 quarters"
    assert rows["Pulled"][1].startswith("not recorded")          # no cache folder was given
    mergers = MERGERS()
    fp = rows["Acquisitions record (SHA-256 of it as canonical JSON)"]
    assert fp[2] == "%d banks, %d acquisition quarters" % (len(mergers), sum(len(v) for v in mergers.values()))
    assert re.fullmatch(r"[0-9a-f]{64}", fp[1])


def test_run_tab_records_when_each_quarter_and_the_mergers_were_pulled(tmp_path, recent, roster):
    cache = tmp_path / "cache"
    cache.mkdir()
    pulled = dt.datetime(2026, 10, 1, 9, 30, 0)
    stamp = pulled.timestamp()
    for q, rows in recent.items():
        f = cache / ("q%s.json" % q)
        f.write_text(json.dumps(rows), encoding="utf-8")
        os.utime(f, (stamp, stamp))
    late = cache / ("q%s.json" % max(recent))
    later = dt.datetime(2026, 10, 2, 8, 0, 0)
    os.utime(late, (later.timestamp(), later.timestamp()))
    mfile = tmp_path / "mergers.json"
    mfile.write_text(json.dumps(MERGERS()), encoding="utf-8")
    os.utime(mfile, (stamp, stamp))
    p = tmp_path / "run.xlsx"
    make(p, recent, roster, "all", products=("Credit card",), cache_dir=cache, mergers_file=mfile)
    rows = run_rows(load_workbook(p))
    by = {r[0]: r for r in rows if r and r[0]}
    assert by["Pulled (cache files written)"][1] == "2026-10-01T09:30:00 to 2026-10-02T08:00:00"
    files = {r[0]: r for r in rows if r and isinstance(r[0], str) and r[0].startswith("q")}
    assert len(files) == len(recent)
    assert files[late.name][1] == "2026-10-02T08:00:00"
    import hashlib
    first = cache / ("q%s.json" % min(recent))
    assert files[first.name][2] == hashlib.sha256(first.read_bytes()).hexdigest()
    m = by["Acquisitions file"]
    assert m[1] == str(mfile) and "2026-10-01T09:30:00" in m[2]


@pytest.mark.parametrize("over, why", [
    (dict(products=("Credit card", "Credit card")), "distinct"),
    (dict(products=("Credit card", "Cars")), "distinct"),
    (dict(normalisation_years=(2019,)), "normalisation_years"),
    (dict(min_banks=0), "min_banks"),
    (dict(min_banks=2.5), "min_banks"),
    (dict(min_banks=True), "min_banks"),
    (dict(last_n=0), "last_n"),
    (dict(min_obs=0), "min_obs"),
    (dict(product="Mortgage", products=("Credit card",)), "starting product"),
    (dict(base="2031Q4"), "base quarter 2031Q4 is not in the data"),
    (dict(base="2019Q5"), "not a quarter"),
    (dict(base="2017Q2"), "base quarter 2017Q2 is not one of the TTM quarters .2018Q1 to 2026Q2."),
    (dict(products=("Credit card", "Auto")), "LNAUTO"),
])
def test_a_bad_setting_is_refused_with_its_reason(tmp_path, recent, roster, over, why):
    over = dict(over)
    kw = {k: over.pop(k) for k in ("products", "base") if k in over}
    with pytest.raises(ValueError, match=why):
        make(tmp_path / "x.xlsx", recent, roster, "all", **kw, **over)
    assert not (tmp_path / "x.xlsx").exists()


def test_data_too_short_for_a_ttm_quarter_is_refused(tmp_path, data, roster):
    four = {q: data[q] for q in ("20190331", "20190630", "20190930", "20191231")}
    with pytest.raises(ValueError, match="no TTM quarter"):
        make(tmp_path / "x.xlsx", four, roster, "all", products=("Credit card",))


def test_a_roster_without_peers_or_without_regionals(tmp_path, recent, roster):
    me = [b for b in roster if b["group"] == "self"]
    with pytest.raises(ValueError, match="no peers"):
        make(tmp_path / "x.xlsx", recent, me, "all", products=("Credit card",))
    flat = [dict(b, regional=False) for b in roster]
    with pytest.raises(ValueError, match="regional"):
        make(tmp_path / "x.xlsx", recent, flat, "regionals", products=("Credit card",))
    make(tmp_path / "flat.xlsx", recent, flat, "all", products=("Credit card",))


def test_mix_says_so_when_the_data_lacks_a_quarter_of_the_base_year(tmp_path, recent, roster):
    ind = {"Credit card": industry(recent, "Credit card")}
    gap = {q: rows for q, rows in recent.items() if q != "20190630"}
    p = tmp_path / "gap.xlsx"
    info = make(p, gap, roster, "all", products=("Credit card",), industry=ind)
    ws = load_workbook(p)["Workings"]
    status = ws.cell(2, column_index_from_string(info["mix"]["helpers"]["status"])).value
    assert "does not hold all four quarters of 2019" in status


def test_judgements_are_required_and_bad_ones_refused(tmp_path, data, roster):
    with pytest.raises(TypeError):
        build_peers_workbook(tmp_path / "x.xlsx", data, roster, FOUR, JUMP, MIN_BOOK, "all", {}, {}, "2019Q4")
    with pytest.raises(ValueError, match="pack"):
        make(tmp_path / "x.xlsx", data, roster, "everyone")
    with pytest.raises(ValueError, match="rank_by"):
        make(tmp_path / "x.xlsx", data, roster, "all", rank_by="mean")
    with pytest.raises(ValueError, match="peer pick"):
        make(tmp_path / "x.xlsx", data, roster, "all", peer_a="KeyBank")


# --- delinquency and nonaccrual (docs/V1-ROUND2-SPEC.md, section A) ------------
#
# Built over tests/data/peers/delinquency_roster.json (the 16 roster banks,
# 2017Q1 to 2026Q2) with two delinquency products and card charge-offs only, so
# the product picker's offset is exercised, and picking mortgage shows the
# charge-off view saying it was not built. The live lines must equal
# peers.delinquency, which is itself checked against arithmetic on the filing.

@pytest.fixture(scope="module")
def dl_data():
    return json.loads((PEERS / "delinquency_roster.json").read_text())


def dq_spec(dl_data, **over):
    spec = dict(data=dl_data, products=TWO, measure="90+ accruing", basis="quarter-end", **DQ_RULES)
    spec.update(over)
    return spec


def dl_books(dl_data, roster, product, measure):
    return {b["cert"]: D.bank_delinquency(dl_data, b["cert"], product, measure, DQ_RULES["jump"],
                                          DQ_RULES["min_book"], DQ_RULES["exempt_self"] and b["cert"] == KEYBANK,
                                          acquired=())
            for b in roster}


@pytest.fixture(scope="module")
def dlq(tmp_path_factory, recent, dl_data, roster):
    """The all-peers pack with delinquency, as built: card, 90+ accruing, quarter-end."""
    p = tmp_path_factory.mktemp("delinquency") / "dlq.xlsx"
    info = make(p, recent, roster, "all", products=("Credit card",), delinquency=dq_spec(dl_data))
    return p, info, recalc(p)


def assert_delinquency(vals, info, books, certs, basis, what, min_banks=MIN_BANKS):
    """Composite, both bank counts and the band on ``basis`` equal peers.delinquency."""
    comp = D.composite_delinquency(books, certs, min_banks)
    assert [r["q"] for r in comp] == info["delinquency"]["quarters"]
    assert [num(v) for v in wcol(vals, info, "delinquency", "banks")] == [r["banks"] for r in comp], what
    key = D.basis_key(basis)
    assert_series(wcol(vals, info, "delinquency", "c"), [r[key] for r in comp], what + " composite")
    want = D.band_delinquency(books, certs, basis, min_banks=min_banks)
    assert [num(v) for v in wcol(vals, info, "delinquency", "bandn")] == [r["banks"] for r in want], what
    for pct in (10, 25, 75, 90):
        assert_series(wcol(vals, info, "delinquency", "p%d" % pct), [r["p%d" % pct] for r in want],
                      "%s p%d" % (what, pct))
    return comp, want


def assert_lines(vals, info, books, roster, basis):
    key = D.basis_key(basis)
    cert = {b["short"]: b["cert"] for b in roster}
    for name, who in (("k", KEYBANK), ("a", cert[scalar(vals, info, "a_name")]),
                      ("bb", cert[scalar(vals, info, "b_name")])):
        assert_series(wcol(vals, info, "delinquency", name), [r[key] for r in books[who]], name)


def test_delinquency_composite_and_band_equal_the_module(dlq, dl_data, roster):
    _, info, vals = dlq
    books = dl_books(dl_data, roster, "Credit card", "90+ accruing")
    certs = [b["cert"] for b in R.peers(roster)]
    comp, band_ = assert_delinquency(vals, info, books, certs, "quarter-end", "card 90+")
    assert len(info["delinquency"]["quarters"]) == 38
    assert sum(r["rate"] is not None for r in comp) == 38 and all(r["p10"] is not None for r in band_)
    assert_lines(vals, info, books, roster, "quarter-end")
    t = info["layout"]["tiles"] + 1
    assert num(peers_cell(vals, "J%d" % t)) == pytest.approx(books[KEYBANK][-1]["rate"], abs=1e-12)
    assert num(peers_cell(vals, "L%d" % t)) == pytest.approx(comp[-1]["rate"], abs=1e-12)
    assert num(peers_cell(vals, "P%d" % t)) == band_[-1]["banks"]


def test_delinquency_follows_jpmorgan_switched_off(dlq, dl_data, roster, tmp_path):
    p, info, before = dlq
    after = typed(p, tmp_path / "flip.xlsx", {inc(info, "JPMorgan"): "No"})
    books = dl_books(dl_data, roster, "Credit card", "90+ accruing")
    without = [b["cert"] for b in R.peers(roster) if b["cert"] != JPMORGAN]
    assert_delinquency(after, info, books, without, "quarter-end", "after the switch")
    assert moved(wcol(before, info, "delinquency", "c"), wcol(after, info, "delinquency", "c")) == 38


def test_delinquency_pickers_move_product_measure_and_basis(dlq, dl_data, roster, tmp_path):
    """Mortgage (the second delinquency product, not a charge-off product
    here), 30+ incl. nonaccrual (the last measure), four-quarter average: the
    composite is the mean of four composite quarter-ends, the band is over the
    banks' four-quarter averages; and the charge-off view says, in words, that
    mortgage charge-offs were not built rather than drawing anything."""
    p, info, _ = dlq
    vals = typed(p, tmp_path / "pick.xlsx", {PRODUCT: "Mortgage", DQ_MEASURE: "30+ incl. nonaccrual",
                                             DQ_BASIS: "four-quarter average"})
    books = dl_books(dl_data, roster, "Mortgage", "30+ incl. nonaccrual")
    certs = [b["cert"] for b in R.peers(roster)]
    comp, band_ = assert_delinquency(vals, info, books, certs, "four-quarter average", "mortgage 30+ avg")
    assert [r["avg4"] is None for r in comp[:4]] == [True, True, True, False]
    assert sum(r["avg4"] is not None for r in comp) == 35
    # Best and worst are ranked on the product's charge-offs, which mortgage
    # does not have here: Peer A and B resolve to no one, and say so.
    assert scalar(vals, info, "a_name") == scalar(vals, info, "b_name") == "(none qualifies)"
    assert peers_cell(vals, "D6") == "Drawn: (none qualifies)"
    assert all(num(v) is None for n_ in ("a", "bb") for v in wcol(vals, info, "delinquency", n_))
    assert_series(wcol(vals, info, "delinquency", "k"), [r["avg4"] for r in books[KEYBANK]], "k")
    differ = sum(num(a) != num(b) for a, b in zip(wcol(vals, info, "delinquency", "banks"),
                                                  wcol(vals, info, "delinquency", "bandn")))
    assert differ >= 10, differ
    # The charge-off view: refused in words, every line blank, the tiles 'none'.
    lay = info["layout"]
    assert scalar(vals, info, "co_p") == 0
    assert peers_cell(vals, "B%d" % lay["note1"]).startswith("No charge-offs were built for this product")
    assert all(num(v) is None for name in ("c", "k", "a", "b", "p10") for v in wcol(vals, info, "charge_offs", name))
    assert peers_cell(vals, "B%d" % (lay["tiles"] + 1)) == "none"
    assert peers_cell(vals, "B%d" % lay["mix_peers"]).startswith("No charge-offs were built for Mortgage")
    assert peers_cell(vals, "B%d" % 6) == "No charge-offs built for this product."
    assert peers_cell(vals, "J%d" % lay["v1"]) == "Delinquency: mortgage, 30+ incl. nonaccrual, four-quarter average (%)"


def test_delinquency_banks_needed_is_met_at_exactly_that_many(dlq, dl_data, roster, tmp_path):
    p, info, _ = dlq
    vals = typed(p, tmp_path / "fifteen.xlsx", {info["layout"]["settings_cells"]["min_banks"]: 15,
                                                PRODUCT: "Mortgage", DQ_MEASURE: "30-89",
                                                DQ_BASIS: "four-quarter average"})
    books = dl_books(dl_data, roster, "Mortgage", "30-89")
    certs = [b["cert"] for b in R.peers(roster)]
    comp, band_ = assert_delinquency(vals, info, books, certs, "four-quarter average", "mortgage at 15",
                                     min_banks=15)
    assert max(r["banks"] for r in comp) == 15
    assert sum(r["avg4"] is not None for r in comp) == 2
    assert [r["banks"] for r in band_ if r["p10"] is not None] == [15, 15]


def test_delinquency_average_needs_four_consecutive_quarter_ends(tmp_path, recent, roster, dl_data):
    gap = {q: rows for q, rows in dl_data.items() if q != "20200630"}
    p = tmp_path / "gap.xlsx"
    info = make(p, recent, roster, "all", products=("Credit card",),
                delinquency=dq_spec(gap, products=("Credit card",), basis="four-quarter average"))
    vals = recalc(p)
    books = dl_books(gap, roster, "Credit card", "90+ accruing")
    certs = [b["cert"] for b in R.peers(roster)]
    comp, band_ = assert_delinquency(vals, info, books, certs, "four-quarter average", "card, 2020Q2 gone")
    assert len(comp) == 37 and "2020Q2" not in [r["q"] for r in comp]
    blank = [r["q"] for r in comp[3:] if r["avg4"] is None]
    assert blank == ["2020Q3", "2020Q4", "2021Q1"]
    assert [r["q"] for r in band_[3:] if r["p10"] is None] == blank
    assert_lines(vals, info, books, roster, "four-quarter average")


def test_delinquency_data_tab_holds_the_module_rows(dlq, dl_data, roster):
    p, info, _ = dlq
    ws = load_workbook(p)["DelinquencyData"]
    nq, nb, nm = 38, len(roster), len(D.MEASURES)
    assert info["delinquency"]["data_rows"] == len(TWO) * nm * nb * nq == ws.max_row - 1
    rows = list(ws.iter_rows(values_only=True))
    for pi, product in enumerate(TWO):
        for mi, measure in enumerate(D.MEASURES):
            books = dl_books(dl_data, roster, product, measure)
            for bi, b in enumerate(roster):
                at = 1 + ((pi * nm + mi) * nb + bi) * nq
                for k, want in enumerate(books[b["cert"]]):
                    got = list(rows[at + k])
                    assert got[:5] == [product, measure, b["short"], b["cert"], want["q"]]
                    # The last column is whether the bank is in the pack now: a live switch for a
                    # peer (what Decide's peer group sums over), FALSE for KeyBank.
                    assert got[-1] is False if b["short"] == "KeyBank" else got[-1].endswith('="Yes"')
                    got = got[:-1]
                    assert got[5:] == pytest.approx([want["numerator"], want["balance"], want["rate"],
                                                     want["kept"], want["why"], want["avg4"],
                                                     want["avg4"] is not None, want["why_avg4"],
                                                     (want["why_avg4"] or "").startswith(D.WHY_JUMP_PREFIX)],
                                                    rel=1e-14, abs=1e-14)


def test_the_delinquency_chart_and_pickers(dlq):
    p, info, _ = dlq
    with zipfile.ZipFile(p) as z:
        charts = [z.read(n).decode("utf-8") for n in sorted(z.namelist())
                  if re.match(r"xl/charts/chart\d+\.xml$", n) and "Peers!dl_" in z.read(n).decode("utf-8")]
    assert len(charts) == 1
    xml = charts[0]
    groups = re.findall(r"<lineChart>.*?</lineChart>", xml)
    assert len(groups) == 3 and all("<upDownBars>" in g for g in groups[:2])
    assert re.findall(r"<val><numRef><f>Peers!(dl_[a-z0-9]+)</f>", groups[2]) == ["dl_c", "dl_a", "dl_bb", "dl_k",
                                                                                  "dl_mark", "dl_chg"]
    ws = load_workbook(p)["Peers"]
    lists = {str(r) for dv in ws.data_validations.dataValidation if dv.type == "list" for r in dv.sqref.ranges}
    assert {DQ_MEASURE, DQ_BASIS, PRODUCT} <= lists
    assert (ws[DQ_MEASURE].value, ws[DQ_BASIS].value) == ("90+ accruing", "quarter-end")
    assert not formulas_reaching_above_row_1(load_workbook(p))


def test_not_built_the_dashboard_and_the_tab_say_so(all_peers):
    p, info, vals = all_peers
    wb = load_workbook(p)
    assert info["delinquency"] is None
    assert "Delinquency was not built" in wb["DelinquencyData"]["A1"].value
    assert peers_cell(vals, "J%d" % info["layout"]["tiles"]) == PW.NOT_BUILT
    run = {r[0]: r for r in run_rows(wb) if r and r[0]}
    assert run["Delinquency"][1] == "not built"


def test_delinquency_is_a_required_argument(tmp_path, recent, roster):
    kw = dict(exempt_self=True, min_banks=MIN_BANKS, index_min_book=500_000, lagged_years=(2019, 2025),
              lag_years=2, normalisation_years=(2019, 2025), mix_base_year=2019, industry=None,
              acquired=NONE_ACQUIRED, combined={}, **RANK)
    with pytest.raises(TypeError, match="delinquency"):
        build_peers_workbook(tmp_path / "x.xlsx", recent, roster, ("Credit card",), JUMP, MIN_BOOK, "all",
                             MERGERS(), EXTRAS, "2019Q4", **kw)


def test_delinquency_run_tab_and_glossary(dlq):
    p, _, _ = dlq
    wb = load_workbook(p)
    run = {r[0]: r for r in run_rows(wb) if r and r[0]}
    assert list(run["Delinquency"][1:3]) == ["Credit card, Mortgage", "quarter-ends 2017Q1 to 2026Q2"]
    assert run["Delinquency jump"][1] == 0.25
    assert list(run["Delinquency minimum book ($000)"][1:3]) == [1_000_000, "KeyBank exempt"]
    terms = [t for t, _, _ in PEER_TERMS]
    for must in ["Delinquency rate", "Measure", "Basis", "Quarter-end rate", "Four-quarter average",
                 "Not filed", "Delinquency composite", "Delinquency band", "Banks in the quarter-end composite",
                 "Banks in the band"] + list(D.MEASURES):
        assert must.lower() in [t.lower() for t in terms], must
    assert set(terms) <= {r[0] for r in glossary_rows(wb)}


def test_glossary_lists_everything_left_out_and_why(dlq):
    p, _, _ = dlq
    rows = [r for r in glossary_rows(load_workbook(p)) if r and r[0]]
    left = {r[0]: r for r in rows if r[0].startswith("Left out:")}
    assert sorted(left) == sorted([
        "Left out: an acquisition (the jump)", "Left out: a small book (the minimum book)",
        "Left out: not filed", "Left out: fewer than four quarters", "Left out: not included",
        "Left out: too few banks", "Left out: before a gap (series files)",
        "Left out: a stated acquisition"])                 # the firm, 17 Oct 2026, question 14
    terms = [r[0] for r in rows]
    first = terms.index("Left out, and why")
    assert set(terms[first + 1:first + 1 + len(left)]) == set(left)
    assert "counted as 0" in rows[first][1]
    for term, where in (("an acquisition (the jump)", "Why not kept"), ("not filed", "Why not kept"),
                        ("fewer than four quarters", "Why no four-quarter average"),
                        ("not included", "Switched off (No)"), ("too few banks", "never drawn from fewer"),
                        ("a stated acquisition", "Why not kept"),
                        ("before a gap (series files)", "note names")):
        assert where in left["Left out: " + term][2], term


@pytest.mark.parametrize("over, why", [
    (dict(measure=None), "missing: measure"),
    (dict(measure="60+"), "measure must be one of"),
    (dict(basis="annual"), "basis must be one of"),
    (dict(products=("Credit card", "Credit card")), "distinct"),
    (dict(products=("Cars",)), "distinct"),
    (dict(products=()), "distinct"),
    (dict(exempt_self="yes"), "exempt_self must be True or False"),
    (dict(jump=-0.1), "jump"),
    (dict(min_book=-1), "min_book"),
    (dict(data={}), "no quarter-ends"),
    (dict(cutoff=3), "not known: cutoff"),
    (dict(jump="0.25"), "jump must be a number"),
    (dict(jump=True), "jump must be a number"),
    (dict(data=[{}]), r"data must be \{REPDTE: \[rows\]\}.*not a list"),
])
def test_a_bad_delinquency_setting_is_refused(tmp_path, recent, roster, dl_data, over, why):
    spec = dq_spec(dl_data)
    for k, v in over.items():
        if v is None:
            spec.pop(k)
        else:
            spec[k] = v
    with pytest.raises(ValueError, match=why):
        make(tmp_path / "x.xlsx", recent, roster, "all", products=("Credit card",), delinquency=spec)
    assert not (tmp_path / "x.xlsx").exists()


def test_delinquency_refuses_a_field_that_was_not_pulled(tmp_path, recent, roster, dl_data):
    gap = {q: [{k: v for k, v in r.items() if k != "P9RERSFM"} if q == "20200630" else r for r in rows]
           for q, rows in dl_data.items()}
    with pytest.raises(ValueError, match="P9RERSFM.*1 quarter.*20200630"):
        make(tmp_path / "x.xlsx", recent, roster, "all", products=("Credit card",), delinquency=dq_spec(gap))
    with pytest.raises(ValueError, match="None .not built. or a dict"):
        make(tmp_path / "x.xlsx", recent, roster, "all", products=("Credit card",), delinquency=[dl_data])
    assert not (tmp_path / "x.xlsx").exists()


# --- reporting changes on the charts (the firm, 16 Oct 2026: "Definitely mark reporting changes.") ----------

def test_the_charge_off_chart_marks_each_quarter_a_reporting_change_took_effect(all_peers, regionals_gross):
    """Card, net (as built): a diamond on the pack composite in 2008Q1, 2010Q1
    and 2020Q1 (the changes touching card net charge-offs in the registry,
    config/reporting-changes.json), #N/A in every other quarter, and none
    where the composite has no rate. Other consumer, gross (typed): the
    2011Q1 auto split and the three quarters of the 2024Q4 to 2025Q2
    revision as well. Each flag column is the registry's, per product and rate."""
    _, info, vals = all_peers
    grid = info["charge_offs"]["quarters"]

    def marked(vs):
        comp, chg = wcol(vs, info, "charge_offs", "c"), wcol(vs, info, "charge_offs", "chg")
        for q, c_, m_ in zip(grid, comp, chg):
            assert num(m_) is None or num(m_) == num(c_), q
        return [q for q, m_ in zip(grid, chg) if num(m_) is not None]

    def want(product, measure):
        took = PW.change_quarters(CHANGES, product, measure)
        comp = wcol(vals_, info, "charge_offs", "c")
        return [q for q, c_ in zip(grid, comp) if q in took and num(c_) is not None]

    vals_ = vals
    assert PW.change_quarters(CHANGES, "Credit card", "net charge-offs") == {"2008Q1", "2010Q1", "2020Q1"}
    assert marked(vals) == want("Credit card", "net charge-offs") and "2020Q1" in marked(vals)
    vals_ = regionals_gross
    oc = want("Other consumer", "gross charge-offs")
    assert {"2024Q4", "2025Q1", "2025Q2", "2010Q1"} <= set(oc)   # FAS 166 and 167 touch other consumer too
    assert marked(regionals_gross) == oc
    # The 2011Q1 auto split touches other consumer, but this pack has no
    # composite rate that quarter, so there is nothing to put a diamond on.
    assert "2011Q1" in PW.change_quarters(CHANGES, "Other consumer", "gross charge-offs") and "2011Q1" not in oc
    assert num(wcol(regionals_gross, info, "charge_offs", "c")[grid.index("2011Q1")]) is None


def test_the_delinquency_chart_marks_the_measure_picked(dlq, tmp_path):
    """Card, 90+ accruing (as built): 2020Q1 (CECL) and 2020Q2 (COVID-19
    deferrals not reported as past due); nonaccrual (typed): 2020Q1 only,
    since the deferral relief is about past due, not nonaccrual."""
    p, info, vals = dlq
    grid = info["delinquency"]["quarters"]

    def marked(vs):
        comp, chg = wcol(vs, info, "delinquency", "c"), wcol(vs, info, "delinquency", "chg")
        for q, c_, m_ in zip(grid, comp, chg):
            assert num(m_) is None or num(m_) == num(c_), q
        return [q for q, m_ in zip(grid, chg) if num(m_) is not None]

    assert marked(vals) == ["2020Q1", "2020Q2"]
    assert marked(typed(p, tmp_path / "na.xlsx", {DQ_MEASURE: "nonaccrual"})) == ["2020Q1"]


def test_no_reporting_changes_stated_marks_nothing(tmp_path, recent, roster):
    """changes=[] states there are none: the marker column is #N/A throughout."""
    p = tmp_path / "none.xlsx"
    info = make(p, recent, roster, "all", products=("Credit card",), changes=[])
    vals = recalc(p)
    assert all(num(v) is None for v in wcol(vals, info, "charge_offs", "chg"))
    with pytest.raises(TypeError, match="changes"):
        kw = dict(exempt_self=True, min_banks=MIN_BANKS, index_min_book=500_000, lagged_years=(2019, 2025),
                  lag_years=2, normalisation_years=(2019, 2025), mix_base_year=2019, industry=None,
                  delinquency=None, acquired=NONE_ACQUIRED, combined={}, **RANK)
        build_peers_workbook(tmp_path / "x.xlsx", recent, roster, ("Credit card",), JUMP, MIN_BOOK, "all",
                             MERGERS(), EXTRAS, "2019Q4", **kw)
