"""Stated acquisitions (peers/stated.py, config/stated-acquisitions.json):
the acquisitions the FDIC record cannot see, each with its filing, read by
`peers build` like the FDIC record.

The firm, 17 Oct 2026, on the commercial proposal's question 14 ("agreed
with the recommendation" on the docket form): add a stated list of the
acquisitions the FDIC record cannot see (JPMorgan and Washington Mutual,
M&T and Wilmington Trust, asset purchases such as Wells Fargo and GE
Capital), each with its 8-K, read by the build like the FDIC record, so the
windows holding them are left out with a reason.
"""
import copy
import json
from pathlib import Path

import pytest

from conftest import ROOT
from threshold_engine.peers import books, delinquency as D, roster as R, run, stated as ST

FILE = ROOT / "config" / "stated-acquisitions.json"
ROSTER = [b["cert"] for b in R.load(ROOT / "config" / "peers.json")]


# ---- the file in the repository ------------------------------------------------------------------------

def test_the_file_holds_the_four_deals_the_firm_named_and_says_why_macquarie_is_not_there():
    """Only the deals docs/commercial-2026-10/research-breaks.md (on main),
    section B, marks not in the record, and only those the answer names.
    Huntington and Macquarie Equipment Finance is left out: its close date is
    recorded there as not checked, so its quarter is not known."""
    entries = ST.load(FILE)
    assert [(e["acquirer"], e["quarter"]) for e in entries] == [
        ("628", "2008Q3"), ("588", "2011Q2"), ("3511", "2016Q1"), ("4297", "2015Q4")]
    for e in entries:
        assert e["url"].startswith("https://www.sec.gov/Archives/edgar/data/"), e
        assert e["acquirer"] in ROSTER
    what = json.loads(FILE.read_text(encoding="utf-8"))["_what"]
    assert "Macquarie" in what and "close date as not checked" in what
    assert all("Macquarie" not in e["what"] for e in entries)


# ---- the refusals: by entry, every problem at once -----------------------------------------------------

def good():
    return {"_what": "A test list.",
            "acquisitions": [{"acquirer": "628", "quarter": "2008Q3", "what": "A bought B.",
                              "citation": "A's 8-K", "url": "https://example.org/8k", "products": ["C&I"],
                              "products_source": "A's 8-K names its loans.", "products_unsettled": []}]}


def refused(raw, tmp_path):
    p = tmp_path / "s.json"
    p.write_text(json.dumps(raw), encoding="utf-8")
    with pytest.raises(ST.StatedError) as e:
        ST.load(p)
    return str(e.value)


@pytest.mark.parametrize("over, says", [
    ({"acquirer": "JPM"}, "acquirer must be the FDIC certificate written as digits"),
    ({"acquirer": 628}, "acquirer must be the FDIC certificate"),
    ({"quarter": "2008-09-30"}, "quarter must be written like 2008Q3"),
    ({"quarter": "2008Q5"}, "quarter must be written like 2008Q3"),
    ({"what": ""}, "what must be one plain sentence"),
    ({"what": "two\nlines"}, "what must be one plain sentence"),
    ({"citation": " "}, "citation must be the filing"),
    ({"url": "http://example.org"}, "url must be the filing's address, starting https://"),
    ({"url": None}, "url must be the filing's address"),
])
def test_a_malformed_entry_is_refused_by_name(tmp_path, over, says):
    raw = good()
    raw["acquisitions"][0].update(over)
    msg = refused(raw, tmp_path)
    assert "acquisition 1" in msg and says in msg, msg


def test_every_problem_is_listed_at_once_and_a_missing_or_extra_key_is_named(tmp_path):
    raw = good()
    second = copy.deepcopy(raw["acquisitions"][0])
    del second["url"]
    second["product"] = ["C&I"]
    third = copy.deepcopy(raw["acquisitions"][0])             # the same deal again
    raw["acquisitions"] += [second, third]
    raw["acquisitions"][0]["quarter"] = "2008"
    msg = refused(raw, tmp_path)
    assert "acquisition 1 (certificate 628, 2008): quarter must be written like 2008Q3" in msg
    assert "acquisition 2 (certificate 628, 2008Q3) is missing 'url' (every key is required" in msg
    assert "acquisition 2 (certificate 628, 2008Q3) has 'product', which is not a key it takes" in msg
    # 1 was refused, so 3 is the first of its kind; it is a repeat of nothing kept
    assert "already listed" not in msg
    ok = good()
    ok["acquisitions"].append(copy.deepcopy(ok["acquisitions"][0]))
    assert ("acquisition 2 (certificate 628, 2008Q3): certificate 628 in 2008Q3 is already listed "
            "(acquisition 1)") in refused(ok, tmp_path)


def test_the_same_acquirer_and_quarter_is_one_acquisition_however_worded(tmp_path):
    """The build reads the acquirer and the quarter, never the wording, so
    the same pair twice is a duplicate even when the sentence differs."""
    raw = good()
    again = dict(raw["acquisitions"][0], what="A acquired the whole of B.", citation="A's 10-Q",
                 url="https://example.org/10q")
    raw["acquisitions"].append(again)
    msg = refused(raw, tmp_path)
    assert "acquisition 2 (certificate 628, 2008Q3): certificate 628 in 2008Q3 is already listed" in msg
    assert "two deals in one quarter go in one entry" in msg
    # Another quarter, or another acquirer, is another acquisition.
    raw["acquisitions"][1]["quarter"] = "2008Q4"
    assert len(ST.check(raw)) == 2


def test_a_key_given_twice_in_one_entry_is_refused(tmp_path):
    """JSON keeps the last of a key given twice, silently; the file is refused."""
    p = tmp_path / "s.json"
    p.write_text('{"_what": "A test list.", "acquisitions": [{"acquirer": "628", "quarter": "2008Q3", '
                 '"quarter": "2009Q1", "what": "A bought B.", "citation": "A\'s 8-K", '
                 '"url": "https://example.org/8k"}]}', encoding="utf-8")
    with pytest.raises(ST.StatedError, match=r"s\.json: the key 'quarter' is given twice in one object "
                                             r"\(certificate 628\); give it once"):
        ST.load(p)
    p.write_text('{"_what": "A.", "_what": "B.", "acquisitions": []}', encoding="utf-8")
    with pytest.raises(ST.StatedError, match="the key '_what' is given twice"):
        ST.load(p)


@pytest.mark.parametrize("raw, says", [
    ({"acquisitions": []}, "is missing '_what'"),
    ({"_what": "x"}, "is missing 'acquisitions'"),
    ({"_what": "x", "acquisitions": {}}, "acquisitions must be a list"),
    ({"_what": " ", "acquisitions": []}, "_what must say what the file is"),
])
def test_the_file_itself_is_checked(tmp_path, raw, says):
    assert says in refused(raw, tmp_path)


def test_an_empty_list_states_none_and_a_missing_or_broken_file_is_refused(tmp_path):
    p = tmp_path / "none.json"
    p.write_text(json.dumps({"_what": "None stated.", "acquisitions": []}), encoding="utf-8")
    assert ST.load(p) == []
    with pytest.raises(ST.StatedError, match="does not exist"):
        ST.load(tmp_path / "nope.json")
    (tmp_path / "bad.json").write_text("{", encoding="utf-8")
    with pytest.raises(ST.StatedError, match="not valid JSON"):
        ST.load(tmp_path / "bad.json")


def test_by_cert_refuses_a_certificate_not_on_the_roster_and_reports_quarters_outside_the_data():
    entries = ST.load(FILE)
    quarters = ["20150930", "20151231", "20160331"]
    got, outside = ST.by_cert(entries, ROSTER, quarters)
    assert got == {"3511": ["20160331"], "4297": ["20151231"]}
    assert [(e["acquirer"], e["quarter"]) for e in outside] == [("628", "2008Q3"), ("588", "2011Q2")]
    with pytest.raises(ST.StatedError, match=r"acquisition 1 \(certificate 628, 2008Q3\)"):
        ST.by_cert(entries, [c for c in ROSTER if c != "628"], quarters)


# ---- what a stated acquisition does to a bank's windows ------------------------------------------------

def quarters_(first_year, n):
    out = []
    y, q = first_year, 1
    for _ in range(n):
        out.append("%d%s" % (y, ("0331", "0630", "0930", "1231")[q - 1]))
        y, q = (y + 1, 1) if q == 4 else (y, q + 1)
    return out


def flat_book(cert="588", n=12, balance=1000.0, ytd=4.0):
    """A bank with a flat C&I book (no jump) and 4 a quarter charged off."""
    data = {}
    for k, q in enumerate(quarters_(2010, n)):
        qn = (int(q[4:6]) - 1) // 3 + 1
        data[q] = [{"CERT": cert, "REPDTE": q, "LNCI": balance, "DRCI": ytd * qn, "CRCI": 0.0,
                    "P3CI": 10.0, "P9CI": 5.0, "NACI": 2.0, "RSSDHCR": "1"}]
    return data


def test_a_stated_quarter_leaves_out_every_window_holding_it_with_its_own_reason():
    """M&T and Wilmington Trust (2011Q2): +13% in the book, under the 25% jump
    and invisible to the 5% share, so before the list the windows were kept,
    two books mixed."""
    data = flat_book()
    before = books.bank_book(data, "588", "C&I", 0.25, 0, acquired=[], stated=())
    assert all(r["kept"] for r in before)
    rows = books.bank_book(data, "588", "C&I", 0.25, 0, acquired=[], stated=["20110630"])
    out = {r["quarter"]: r["why"] for r in rows if not r["kept"]}
    assert out == {q: books.WHY_STATED for q in ("2011Q2", "2011Q3", "2011Q4", "2012Q1")}
    # A recorded acquisition in the same window comes first, as before.
    both = books.bank_book(data, "588", "C&I", 0.25, 0, acquired=["20110630"], stated=["20110630"])
    assert {r["why"] for r in both if not r["kept"]} == {books.WHY_ACQUIRED}
    with pytest.raises(ValueError, match="stated is a list"):
        books.bank_book(data, "588", "C&I", 0.25, 0, acquired=[], stated="20110630")


def test_a_stated_quarter_blanks_the_delinquency_average_never_a_quarter_end_rate():
    data = flat_book()
    rows = D.bank_delinquency(data, "588", "C&I", "30-89", 0.25, 0, False, acquired=[], stated=["20110630"])
    assert all(r["rate"] is not None for r in rows)
    blank = {r["q"]: r["why_avg4"] for r in rows if r["avg4"] is None and r["q"] >= "2011Q2"}
    assert blank == {q: D.WHY_STATED % "2011Q2" for q in ("2011Q2", "2011Q3", "2011Q4")}
    assert D.WHY_STATED_PREFIX in D.LEFT_OUT



def test_the_stated_acquisitions_are_required_wherever_the_recorded_ones_are():
    """Every setting is required: left out, a stated acquisition's windows
    would be kept with two books mixed and nothing said. ``acquired`` is
    keyword-only with no default on the same functions, so ``stated`` is
    too; an empty list (``{}`` for the workbook) states none."""
    import inspect
    from threshold_engine.peers import peers_workbook as PW
    data = flat_book()
    with pytest.raises(TypeError, match="stated"):
        books.bank_book(data, "588", "C&I", 0.25, 0, acquired=[])
    with pytest.raises(TypeError, match="stated"):
        D.bank_delinquency(data, "588", "C&I", "30-89", 0.25, 0, False, acquired=[])
    from threshold_engine.peers import flows as FL
    for f in (books.bank_book, D.bank_delinquency, FL.bank_flow, PW.build_peers_workbook, PW._check_delinquency):
        params = inspect.signature(f).parameters
        assert "stated" in params, f.__name__
        for name in [n for n in ("acquired", "stated") if n in params]:
            par = params[name]
            assert par.default is inspect.Parameter.empty, (f.__name__, name)
            assert par.kind is inspect.Parameter.KEYWORD_ONLY, (f.__name__, name)
    with pytest.raises(ValueError, match="stated must be"):
        PW._check_delinquency(None, [], {}, stated=None)

# ---- the build reads it like the FDIC record -----------------------------------------------------------

def cache_of(tmp_path, data):
    cache = tmp_path / "cache"
    cache.mkdir()
    for q, rows in data.items():
        (cache / ("q%s.json" % q)).write_text(json.dumps(rows), encoding="utf-8")
    return cache


def build_kw(tmp_path, cache, **over):
    m = tmp_path / "mergers.json"
    m.write_text("{}", encoding="utf-8")
    a = tmp_path / "absorptions.json"
    a.write_text("[]", encoding="utf-8")
    kw = dict(cache=cache, roster=ROOT / "config" / "peers.json", series_dir=tmp_path / "series",
              first="2010Q1", last="2012Q4", products=["C&I"], mergers=m, absorptions=a,
              extra_acquisitions=[], industry_combined=[], jump=0.25, acquisition_share=0.05, min_book=0,
              exempt_self=True, pro_forma=False, pack="all", min_banks=1, rank_by="net", last_n=4, min_obs=2,
              base_quarter="2011Q4", index_min_book=0, lagged_years=[2011], lag_years=1,
              normalisation_years=[2011, 2012], mix_base_year=2011, delinquency=False,
              reporting_changes=ROOT / "config" / "reporting-changes.json",
              stated_acquisitions=FILE, say=lambda line: None)
    kw.update(over)
    return kw


def test_the_build_leaves_out_ms_windows_and_says_so_on_its_series_and_the_run_tab(tmp_path):
    from openpyxl import load_workbook
    cache = cache_of(tmp_path, flat_book())
    out = tmp_path / "peers.xlsx"
    run.build(out, **build_kw(tmp_path, cache))
    mt = tmp_path / "series" / "banks" / "c-i_m-t_net.csv"
    text = mt.read_text(encoding="utf-8")
    # 2011Q2 to 2012Q1 hold the stated quarter, so the series starts after them.
    assert "date,value\n2012-06-30," in text
    assert "Written from 2012Q2, the first quarter after the last break" in text
    assert "# Stated acquisitions, the ones the FDIC record cannot see (stated-acquisitions.json): 1 of the 1 " \
           "in these quarters brought C&I, so each bank's windows holding one are left out: M&T 2011Q2 (M&T Bank " \
           "Corporation, Form 8-K, Item 2.01" in text
    assert "; products: A whole bank: the filing quoted merges Wilmington Trust Corporation into M&T" in text
    run_tab = load_workbook(out)["Run"]
    said = " ".join(str(c.value) for row in run_tab.iter_rows() for c in row if c.value)
    assert "Stated acquisitions: %s, 4 entries, SHA-256" % FILE in said
    assert "in these quarters: M&T 2011Q2 (brought Credit card, Mortgage, Home equity, Other consumer, Auto, " \
           "C&I, Construction, Multifamily, Nonfarm nonresidential, Owner-occupied nonfarm nonresidential, Other " \
           "nonfarm nonresidential); outside them, not read: JPMorgan 2008Q3; Wells Fargo 2016Q1; " \
           "Capital One 2015Q4" in said
    # The book index reads it as declared, like an extra acquisition.
    assert "2011Q2 (declared)" in said


def test_none_states_that_no_list_is_read_and_nothing_is_left_out(tmp_path):
    cache = cache_of(tmp_path, flat_book())
    run.build(tmp_path / "peers.xlsx", **build_kw(tmp_path, cache, stated_acquisitions=None))
    text = (tmp_path / "series" / "banks" / "c-i_m-t_net.csv").read_text(encoding="utf-8")
    assert "date,value\n2011-03-31," in text
    assert "(none read): none named, so none is read." in text


@pytest.mark.parametrize("value, says", [
    (True, "--stated-acquisitions must name the file"),
    (5, "--stated-acquisitions must name the file"),
    ("nope.json", "--stated-acquisitions: the stated acquisitions file nope.json does not exist"),
])
def test_the_build_refuses_a_bad_list_before_reading_anything(tmp_path, value, says):
    with pytest.raises(ValueError, match=says.replace("(", r"\(")):
        run.build(tmp_path / "p.xlsx", **build_kw(tmp_path, tmp_path / "no-cache", stated_acquisitions=value))
    assert not (tmp_path / "p.xlsx").exists()


def test_a_listed_certificate_not_on_the_roster_is_refused(tmp_path):
    p = tmp_path / "s.json"
    p.write_text(json.dumps({"_what": "x", "acquisitions": [
        {"acquirer": "99999", "quarter": "2011Q2", "what": "Someone bought something.",
         "citation": "An 8-K", "url": "https://example.org/8k", "products": ["C&I"],
         "products_source": "The 8-K.", "products_unsettled": []}]}), encoding="utf-8")
    with pytest.raises(ValueError, match=r"--stated-acquisitions: .*acquisition 1 \(certificate 99999, 2011Q2\)"):
        run.build(tmp_path / "p.xlsx", **build_kw(tmp_path, cache_of(tmp_path, flat_book()), stated_acquisitions=p))


def test_the_cli_takes_a_file_or_the_word_none(tmp_path, monkeypatch):
    from threshold_engine import cli
    seen = {}
    monkeypatch.setattr(run, "build", lambda out, **kw: seen.update(kw))
    args = ["peers", "build", str(tmp_path / "p.xlsx"), "--cache", "c", "--roster", "r", "--series-dir", "s",
            "--first", "2010Q1", "--last", "2012Q4", "--products", "C&I", "--mergers", "m", "--absorptions", "a",
            "--extra-acquisitions", "--industry-combined", "--delinquency", "no", "--reporting-changes", "x",
            "--jump", "0.25", "--acquisition-share", "0.05", "--min-book", "0", "--exempt-self", "yes",
            "--pro-forma", "no", "--pack", "all", "--min-banks", "1", "--rank-by", "net", "--last-n", "4",
            "--min-obs", "2", "--base-quarter", "2011Q4", "--index-min-book", "0", "--lagged-years", "2011",
            "--lag-years", "1", "--normalisation-years", "2011", "2012", "--mix-base-year", "2011"]
    assert cli.main(args + ["--stated-acquisitions", "none"]) == 0 and seen["stated_acquisitions"] is None
    assert cli.main(args + ["--stated-acquisitions", str(FILE)]) == 0 and seen["stated_acquisitions"] == str(FILE)
    with pytest.raises(SystemExit):
        cli.main(args)                                     # required: no default


def test_the_consumer_book_states_null_and_says_why():
    raw = json.loads((ROOT / "books" / "consumer-public.json").read_text(encoding="utf-8"))
    assert raw["peers"]["stated_acquisitions"] is None
    assert any(n.startswith("Stated acquisitions: none is read for this book") for n in raw["notes"])


def test_the_delinquency_files_name_the_list_and_the_pack_average_names_the_bank_it_blanks(tmp_path):
    """The four-quarter delinquency average is blanked like a recorded
    acquisition, and the files the rules apply to say which list was read."""
    data = flat_book()
    for q, rows in flat_book(cert="628").items():           # a second bank keeps the composite going
        data[q] = data[q] + rows
    cache = cache_of(tmp_path, data)
    run.build(tmp_path / "peers.xlsx", **build_kw(tmp_path, cache, delinquency=True))
    series = tmp_path / "series"
    comp = (series / "c-i_pack-composite_delinquency-30-89_four-quarter-average.csv").read_text(encoding="utf-8")
    assert "# Stated acquisitions, the ones the FDIC record cannot see (stated-acquisitions.json): 1 of the 1 in " \
           "these quarters brought C&I" in comp
    assert "blanked by a stated acquisition (an acquisition the FDIC record cannot see, listed with its " \
           "filing): M&T (588) 2011Q2" in comp
    # 30+ to charge-off (flows, merged 18 Oct 2026; flows.bank_flow is tested on its own): the files the
    # rules apply to name the list and say the rule.
    pack = (series / "c-i_pack-composite_30-plus-to-charge-off_quarter-end.csv").read_text(encoding="utf-8")
    assert "# Stated acquisitions, the ones the FDIC record cannot see (stated-acquisitions.json): 1 of the 1 in " \
           "these quarters brought C&I" in pack and "and the stated acquisitions, applied to the two quarters" in pack


def test_left_out_note_words_each_reason_apart():
    tail = [("2011Q2", 1.0), ("2011Q3", 1.0)]
    note = run.left_out_note({"2011Q2": [("588", "a stated acquisition")], "2011Q3": [("588", "the jump")]},
                             tail, {"588": "M&T"})
    assert "blanked by the jump (the book moved by more than the jump): M&T 2011Q3" in note
    assert "blanked by a stated acquisition (an acquisition the FDIC record cannot see, listed with its " \
           "filing): M&T 2011Q2" in note
    assert "acquisition share" not in note
    assert "a stated acquisition" in run.left_out_note({}, tail, {})


# ---- the book file: the key is required, null or a file ------------------------------------------------

def test_the_book_requires_the_key_and_takes_null_or_a_file(tmp_path):
    from threshold_engine import book as BK
    from test_book import write_book
    peers = {k: "x" for k in BK.PEERS_KEYS}
    del peers["stated_acquisitions"]
    with pytest.raises(BK.BookError, match="peers is missing 'stated_acquisitions'"):
        BK.load(write_book(tmp_path, peers=peers))
    for bad in (5, "", True):
        peers["stated_acquisitions"] = bad
        with pytest.raises(BK.BookError, match="peers: stated_acquisitions must be null or a file"):
            BK.load(write_book(tmp_path, peers=peers))
    peers["stated_acquisitions"] = None                 # null: none is read, and nothing is said against it
    with pytest.raises(BK.BookError) as e:              # the other "x" settings are still refused
        BK.load(write_book(tmp_path, peers=peers))
    assert "stated_acquisitions" not in str(e.value)
