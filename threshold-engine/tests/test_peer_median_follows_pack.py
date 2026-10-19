"""The peer median follows the pack switches on Peers, as the peer total does.

The firm, 18 Oct 2026, question 27 of the round 2 commercial docket, agreed
with the recommendation: *"The peer median follows the pack switches"*.
Before, ``book._live_peer_group`` made only the pack composite (the peer
total) live; the pack median stayed as built with all 15 banks.

A small workbook laid out as the book's is (Data, Peers, Peers data, Peers
delinquency data, Lines by source, and the switches on Peers workings) is
made live and evaluated with the formulas engine: with every bank in, each
median equals the middle of the kept rates (the file's, as built); switching
a bank out moves it, and below the banks needed it is blank.
"""
import statistics
from pathlib import Path

import pytest
from openpyxl import Workbook

from test_workbook import recalc
from threshold_engine import book as BK
from threshold_engine.peers import delinquency as D

QUARTERS = ["2020Q1", "2020Q2", "2020Q3"]
PRODUCTS = ["C&I", "Construction"]
BANKS = ["KeyBank", "P1", "P2", "P3", "P4"]            # KeyBank is never in the pack
MIN_BANKS = 2
#: Each bank's net rate by product and quarter; None: not kept that quarter.
NET = {"C&I": {"KeyBank": [9.0, 9.0, 9.0], "P1": [0.1, 0.4, 0.3], "P2": [0.5, None, 0.9], "P3": [0.2, 0.6, None],
               "P4": [0.8, 0.7, None]},
       "Construction": {"KeyBank": [5.0, 5.0, 5.0], "P1": [1.0, 1.1, 1.2], "P2": [2.0, 2.5, 3.0],
                        "P3": [0.5, 0.25, 0.75], "P4": [None, None, 4.0]}}
#: Quarter-end nonaccrual rates, C&I only (the delinquency product here).
NONACCRUAL = {"KeyBank": [7.0, 7.0, 7.0], "P1": [1.0, 2.0, 3.0], "P2": [4.0, None, 6.0], "P3": [0.5, 0.5, None],
              "P4": [3.0, 1.0, 2.0]}
SERIES = [("C&I - Net - Peer median", "c-i_pack-median_net"),
          ("Construction - Net - Peer median", "construction_pack-median_net"),
          ("Construction - Gross - Peer median", "construction_pack-median_gross"),
          ("C&I - Net - Peer total", "c-i_pack-composite_net"),
          ("C&I - Nonaccrual - Peer median", "c-i_pack-median_delinquency-nonaccrual_quarter-end"),
          ("C&I - Net - Industry", "c-i_industry_net")]


def gross(v):
    return None if v is None else v + 1.0          # a gross rate a point above the net, so the two differ


def expected(table, k, out=(), f=lambda v: v):
    vals = [f(table[b][k]) for b in BANKS[1:] if b not in out and table[b][k] is not None]
    return statistics.median(vals) if len(vals) >= MIN_BANKS else ""


def make(path: Path, off=()):
    """The workbook, made live; ``off`` names the banks switched out on Workings."""
    wb = Workbook()
    data = wb.active
    data.title = "Data"
    data.append(["Quarter"] + [lbl for lbl, _ in SERIES])
    for k, q in enumerate(QUARTERS):
        # As built: each series' value with every bank in (what the file holds).
        data.append([q, expected(NET["C&I"], k), expected(NET["Construction"], k),
                     expected(NET["Construction"], k, f=gross), 0.123, expected(NONACCRUAL, k), 0.5])
    wk = wb.create_sheet("Peers workings")
    for i, b in enumerate(BANKS[1:], start=1):
        wk["B%d" % i] = "No" if b in off else "Yes"
    pe = wb.create_sheet("Peers")
    pe["B3"] = MIN_BANKS
    da = wb.create_sheet(BK.PEER_TAB_NAMES["Data"])
    da.append(["Product", "Bank", "Certificate", "Quarter", "Gross", "Recoveries", "Average balance", "Gross rate",
               "Net rate", "Kept", "Why", "", "", "", "", "In the pack now"])
    for p in PRODUCTS:
        for bi, b in enumerate(BANKS):
            for k, q in enumerate(QUARTERS):
                n = NET[p][b][k]
                kept = n is not None
                da.append([p, b, str(100 + bi), q, 10.0, 1.0, 1000.0 * bi, gross(n), n, kept, None if kept else "missing",
                           None, None, None, None,
                           False if bi == 0 else "='Peers workings'!$B$%d=\"Yes\"" % bi])
    dd = wb.create_sheet(BK.PEER_TAB_NAMES["DelinquencyData"])
    dd.append(["Product", "Measure", "Bank", "Certificate", "Quarter-end", "Numerator", "Balance", "Rate", "Kept",
               "Why", "", "", "", "", "In the pack now"])
    for m in D.MEASURES:
        for bi, b in enumerate(BANKS):
            for k, q in enumerate(QUARTERS):
                # Every measure but nonaccrual a constant, so a block read in the wrong place shows.
                v = NONACCRUAL[b][k] if m == "nonaccrual" else 50.0
                dd.append(["C&I", m, b, str(100 + bi), q, 1.0, 100.0, v, v is not None, None, None, None, None, None,
                           False if bi == 0 else "='Peers workings'!$B$%d=\"Yes\"" % bi])
    lbs = wb.create_sheet("Lines by source")
    for i, (lbl, _) in enumerate(SERIES):
        lbs["A%d" % (5 + i)] = lbl
    info = {"layout": {"settings_cells": {"min_banks": "B3"}}, "products": PRODUCTS, "quarters": QUARTERS,
            "data_rows": len(PRODUCTS) * len(BANKS) * len(QUARTERS),
            "delinquency": {"products": ["C&I"], "quarters": QUARTERS}}
    said = []
    made = BK._live_peer_group(wb, [(lbl, Path("series/%s.csv" % stem)) for lbl, stem in SERIES], info, said.append)
    wb.save(path)
    return made, said, wb


def values(path):
    v = recalc(path)
    return {lbl: [v["DATA!%s%d" % (chr(ord("B") + j), 2 + k)] for k in range(len(QUARTERS))]
            for j, (lbl, _) in enumerate(SERIES)}, v


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    d = tmp_path_factory.mktemp("median")
    made, said, wb = make(d / "all.xlsx")
    _, _, _ = make(d / "p2-off.xlsx", off=("P2",))
    _, _, _ = make(d / "p2-p3-off.xlsx", off=("P2", "P3"))
    return {"made": made, "said": said, "wb": wb, "all": values(d / "all.xlsx"),
            "p2": values(d / "p2-off.xlsx"), "p23": values(d / "p2-p3-off.xlsx")}


def close(got, want):
    return all((g == "" and w == "") or (w != "" and g == pytest.approx(w)) for g, w in zip(got, want))


def test_every_median_and_the_total_are_made_live(built):
    assert built["made"] == 5                                      # four medians and the total
    assert built["said"] == ["peer group: 5 series follow the banks in the pack on Peers"]
    data = built["wb"]["Data"]
    for j, (lbl, stem) in enumerate(SERIES):
        cell = data.cell(2, 2 + j).value
        if "pack-median" in stem:
            assert str(cell).startswith("=IF(COUNT(") and "MEDIAN(" in cell, lbl
        elif "pack-composite" in stem:
            assert str(cell).startswith("=IF(COUNTIFS("), lbl
        else:
            assert cell == 0.5                                     # the industry is not a peer group series


def test_with_every_bank_in_each_median_is_the_files(built):
    got, _ = built["all"]
    assert close(got["C&I - Net - Peer median"], [expected(NET["C&I"], k) for k in range(3)])
    assert close(got["Construction - Net - Peer median"], [expected(NET["Construction"], k) for k in range(3)])
    assert close(got["Construction - Gross - Peer median"],
                 [expected(NET["Construction"], k, f=gross) for k in range(3)])
    assert close(got["C&I - Nonaccrual - Peer median"], [expected(NONACCRUAL, k) for k in range(3)])
    # KeyBank's 9.0 is never counted: every median is below it.
    assert all(x == "" or x < 5 for x in got["C&I - Net - Peer median"] + got["Construction - Net - Peer median"])


def test_switching_a_bank_out_moves_the_median(built):
    """P2 out: each median is the middle of the banks left, and differs from
    the median with P2 in where P2 was counted."""
    all_, _ = built["all"]
    off, _ = built["p2"]
    for lbl, table, f in (("C&I - Net - Peer median", NET["C&I"], lambda v: v),
                          ("Construction - Net - Peer median", NET["Construction"], lambda v: v),
                          ("Construction - Gross - Peer median", NET["Construction"], gross),
                          ("C&I - Nonaccrual - Peer median", NONACCRUAL, lambda v: v)):
        want = [expected(table, k, out=("P2",), f=f) for k in range(3)]
        assert close(off[lbl], want), (lbl, off[lbl], want)
    assert off["C&I - Net - Peer median"][0] != all_["C&I - Net - Peer median"][0]
    assert off["Construction - Net - Peer median"][2] != all_["Construction - Net - Peer median"][2]
    assert off["C&I - Nonaccrual - Peer median"][2] != all_["C&I - Nonaccrual - Peer median"][2]
    # The peer total moves too, as it already did.
    assert off["C&I - Net - Peer total"][0] != all_["C&I - Net - Peer total"][0]


def test_below_the_banks_needed_the_median_is_blank(built):
    """P2 and P3 out: C&I in 2020Q3 has only P1 kept (P4 is not), below the
    two banks needed, so blank, never a median of one."""
    got, _ = built["p23"]
    assert got["C&I - Net - Peer median"] == [expected(NET["C&I"], k, out=("P2", "P3")) for k in range(3)]
    assert got["C&I - Net - Peer median"][2] == ""
    assert got["C&I - Net - Peer median"][0] == pytest.approx(statistics.median([0.1, 0.8]))


def test_lines_by_source_reads_the_live_latest_median(built):
    _, v_all = built["all"]
    _, v_off = built["p2"]
    r = 5 + [lbl for lbl, _ in SERIES].index("Construction - Net - Peer median")
    assert v_all["LINES BY SOURCE!M%d" % r] == pytest.approx(expected(NET["Construction"], 2))
    assert v_off["LINES BY SOURCE!M%d" % r] == pytest.approx(expected(NET["Construction"], 2, out=("P2",)))
    assert v_all["LINES BY SOURCE!L%d" % r] == "2020Q3"


def test_a_median_in_a_quarter_the_peer_tabs_do_not_hold_is_refused(tmp_path):
    """The book's Data says 2020Q4 and the peer tabs end at 2020Q3: refused
    rather than read from the wrong row."""
    wb = Workbook()
    wb.active.title = "Data"
    wb["Data"].append(["Quarter", "C&I - Net - Peer median"])
    for q in ["2020Q1", "2020Q2", "2020Q4"]:
        wb["Data"].append([q, 0.3])
    da = wb.create_sheet(BK.PEER_TAB_NAMES["Data"])
    da.append(["h"] * 16)
    for b in BANKS:
        for q in QUARTERS:
            da.append(["C&I", b, "1", q, 1, 0, 1, 1, 1, True, None, None, None, None, None, True])
    wb.create_sheet("Peers")["B3"] = 2
    info = {"layout": {"settings_cells": {"min_banks": "B3"}}, "products": ["C&I"], "quarters": QUARTERS,
            "data_rows": len(BANKS) * len(QUARTERS), "delinquency": None}
    with pytest.raises(ValueError, match="has a value in 2020Q4, a quarter Peers data does not hold"):
        BK._live_peer_group(wb, [("C&I - Net - Peer median", Path("s/c-i_pack-median_net.csv"))], info,
                            lambda s: None)
