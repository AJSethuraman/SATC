"""The Synthesis tab: every candidate line, chosen line, rating and Check the
workbook works out live must equal synthesis.synthesize on the same inputs,
as built, after the firm changes a choice on the tab, and after Your settings
on Assess change. Recalculated with the `formulas` engine; a skip is
reported, never counted as a pass. The series are made up (test_synthesis's)."""
import hashlib
import re
import statistics

import pytest
from openpyxl import load_workbook

from test_synthesis import FILED, INDUSTRY, INTERNAL, PCT, PEERS, SETTINGS, series
from threshold_engine.cli import main
from threshold_engine.glossary import TERMS
from threshold_engine.scale import scale
from threshold_engine.series import Point, read_notes, smooth
from threshold_engine.synthesis import LINES, synthesize
from threshold_engine.synthesis_tab import HELP, from_flags
from threshold_engine.workbook import SET_TOP, assess_layout, build, quarter

formulas = pytest.importorskip("formulas", reason="the formulas engine recalculates the workbook")

LEADS = ["internal", "internal", "internal", "industry"]
# One peer and KeyBank's filed series begin four years before the internal
# series, far from where they sit after: their medians must be taken over the
# internal quarters only, and these quarters would move them if they were not.
PEERS = dict(PEERS, p4=series(2010, [9.0] * 16 + [p.value for p in PEERS["p4"]]))
FILED = series(2010, [9.0] * 16 + [p.value for p in FILED])
OK = "OK: every line is taken from a source that has it, and the lines run in order"
NOTES = "# measure: {m}\n# as-of: 2025Q4\n# basis: made up for the tests\n# source: test_synthesis.py\n"


def write(path, pts, measure):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(NOTES.format(m=measure) + "date,value\n" + "".join("%s,%r\n" % (p.date, p.value) for p in pts),
                    encoding="utf-8")
    return path


def flip(pts):
    """The same shapes as a measure where lower is worse."""
    return [Point(p.date, 10 - p.value) for p in pts]


def inputs(d, internal=INTERNAL, industry=INDUSTRY, filed=FILED, peers=PEERS):
    files = dict(internal=write(d / "internal.csv", internal, "internal, made up"),
                 industry=write(d / "industry.csv", industry, "industry, made up"),
                 filed=write(d / "filed.csv", filed, "KeyBank filed, made up") if filed else None,
                 peers=[(k, write(d / "peers" / (k + ".csv"), v, "peer %s, made up" % k)) for k, v in peers.items()])
    return files


def make(path, files, leads=LEADS, translation="median ratio", pct=PCT, direction="higher_is_worse",
         floor_at_zero=True, smoothing=1, extra=()):
    syn = {"Credit card": dict(industry=files["industry"], filed=files["filed"], peers=files["peers"],
                               translation=translation, peer_percentiles=pct if files["peers"] else None,
                               leads=leads)}
    return build(path, [("Credit card", files["internal"])] + list(extra), "NCO", "%", direction, "quarterly",
                 smoothing, floor_at_zero, SETTINGS["halfwidth"], SETTINGS["low_step"], SETTINGS["high_z"], "worse",
                 (), (50, 75, 90, 95), 4, 20, synthesis=syn)


def recalc(path):
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for key, val in sol.items():
        v = getattr(val, "value", val)
        try:
            v = v.item(0) if hasattr(v, "item") else v
        except Exception:                          # noqa: BLE001
            pass
        out[key.split("]")[-1].upper().replace("'", "")] = v
    return out


def edited(path, tmp_path, name, synthesis=None, assess=None):
    """Type into the Synthesis and Assess tabs, by their own cell names, and recalculate."""
    wb = load_workbook(path)
    for sheet, cells in (("Synthesis", synthesis or {}), ("Assess", assess or {})):
        for ref, v in cells.items():
            wb[sheet][ref] = v
    q = tmp_path / name
    wb.save(q)
    return recalc(q)


def reference(direction="higher_is_worse", floor_at_zero=True, internal=INTERNAL, industry=INDUSTRY, filed=FILED,
              peers=PEERS, **over):
    args = dict(SETTINGS, direction=direction, floor_at_zero=floor_at_zero, translation="median ratio",
                peer_percentiles=PCT if peers else None, leads=LEADS)
    args.update(over)
    return synthesize(internal, industry, filed, peers, **args)


def cell(vals, ref):
    return vals["SYNTHESIS!%s" % ref]


def helper(vals, at, name):
    from openpyxl.utils import get_column_letter
    return vals["SYNTHESIS!%s%d" % (get_column_letter(HELP[name]), at["helpers"])]


def same(got, want):
    if want is None:
        return got == ""
    return isinstance(got, (int, float)) and got == pytest.approx(want, rel=1e-9, abs=1e-12)


def assert_equals_reference(vals, at, r):
    """Every candidate, every chosen line, the rating and the competitive check."""
    for k, row in enumerate(at["lines"]):
        for src, col in (("internal", "B"), ("industry", "C"), ("peers", "E")):
            want = (r["candidates"][src]["lines"] or [None] * 4)[k]
            got = cell(vals, "%s%d" % (col, row))
            assert same(got, want), (LINES[k], src, got, want)
        assert cell(vals, "F%d" % row) == r["leads"][k]
        assert cell(vals, "G%d" % row) == pytest.approx(r["chosen"][k])
    assert cell(vals, "B%d" % at["check"]) == OK
    assert cell(vals, "B%d" % at["latest"]) == r["latest"][0]
    assert cell(vals, "C%d" % at["latest"]) == pytest.approx(r["latest"][1])
    assert cell(vals, "D%d" % at["latest"]) == r["rating"]
    ch = r["check"]
    if ch is None:
        assert cell(vals, "B%d" % at["competitive"]).startswith("No competitive check")
    else:
        assert helper(vals, at, "kb_med") == pytest.approx(ch["filed_median"])
        assert helper(vals, at, "worse_than") == ch["worse_than"]
        assert helper(vals, at, "held") == ch["peers"]
        # synthesis.report's sentence; the median to three decimals (Excel's
        # TEXT and Python's %.3f can round a final 5 different ways).
        said = cell(vals, "B%d" % at["competitive"])
        m = re.fullmatch(r"KeyBank's filed median (\d+\.\d{3})% over the internal window; it is worse than "
                         r"(\d+) of (\d+) peers", said)
        assert m, said
        assert float(m.group(1)) == pytest.approx(ch["filed_median"], abs=0.00051)
        assert (int(m.group(2)), int(m.group(3))) == (ch["worse_than"], ch["peers"])


def refusal(vals, at):
    return cell(vals, "B%d" % at["check"])


def refused_as(**kw):
    with pytest.raises(ValueError) as exc:
        reference(**kw)
    return "Refused: %s" % exc.value


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    d = tmp_path_factory.mktemp("syn")
    files = inputs(d)
    info = make(d / "s.xlsx", files)
    return d / "s.xlsx", info["synthesis"]["Credit card"], recalc(d / "s.xlsx"), files


# ---------------------------------------------------------------- equal to synthesize

def test_internal_leads_with_high_from_the_industry_equal_the_reference(built):
    _, at, vals, _ = built
    r = reference()
    assert r["candidates"]["internal"]["lines"][3] is None          # the calm history cannot place High
    assert_equals_reference(vals, at, r)


def test_a_line_led_by_peers_follows_the_lead_typed_on_the_tab(built, tmp_path):
    p, at, _, _ = built
    leads = ["internal", "internal", "peers", "peers"]
    vals = edited(p, tmp_path, "peers.xlsx", synthesis={"F%d" % at["lines"][2]: "peers",
                                                         "F%d" % at["lines"][3]: "peers"})
    assert_equals_reference(vals, at, reference(leads=leads))


def test_translation_none_takes_the_industry_lines_as_they_are(built, tmp_path):
    p, at, _, _ = built
    vals = edited(p, tmp_path, "none.xlsx", synthesis={"B%d" % at["translation"]: "none"})
    r = reference(translation="none")
    assert_equals_reference(vals, at, r)
    ind = scale([x.value for x in INDUSTRY], "higher_is_worse", 0.5, 1.0, 2.5, True)["bounds"]
    assert [cell(vals, "C%d" % row) for row in at["lines"]] == pytest.approx(ind)


def test_the_peers_percentiles_typed_on_the_tab_are_live(built, tmp_path):
    p, at, _, _ = built
    pct = [10, 30, 60, 90]
    vals = edited(p, tmp_path, "pct.xlsx", synthesis={"D%d" % row: v for row, v in zip(at["lines"], pct)})
    assert_equals_reference(vals, at, reference(peer_percentiles=pct))


def test_your_settings_on_assess_move_every_source(built, tmp_path):
    p, at, _, _ = built
    vals = edited(p, tmp_path, "set.xlsx", assess={"O%d" % SET_TOP: 0.25, "P%d" % SET_TOP: 0.75,
                                                    "Q%d" % SET_TOP: 2.0})
    assert_equals_reference(vals, at, reference(halfwidth=0.25, low_step=0.75, high_z=2.0))


def test_a_period_left_out_on_assess_moves_the_internal_lines_only(built, tmp_path):
    """The documented choice: periods left out shape the internal lines (they
    are Thresholds' lines); the translations and the peer window keep every
    quarter the internal series holds, as synthesis.py does."""
    p, at, _, _ = built
    lo = assess_layout(1)["lo_top"]
    vals = edited(p, tmp_path, "lo.xlsx", assess={"N%d" % lo: "Credit card", "O%d" % lo: "2020Q1",
                                                   "P%d" % lo: "2021Q4", "Q%d" % lo: "a test"})
    kept = [None if "2020Q1" <= quarter(x.date) <= "2021Q4" else x.value for x in INTERNAL]
    own = scale(kept, "higher_is_worse", 0.5, 1.0, 2.5, True)
    med, move = own["median"], own["typical_yearly_move"]
    assert [cell(vals, "B%d" % row) for row in at["lines"][:3]] == pytest.approx(
        [med - move, med - 0.5 * move, med + 0.5 * move])
    r = reference()
    for src, col in (("industry", "C"), ("peers", "E")):
        assert [cell(vals, "%s%d" % (col, row)) for row in at["lines"]] == pytest.approx(r["candidates"][src]["lines"])


def test_a_measure_where_lower_is_worse_equals_the_reference(tmp_path):
    files = inputs(tmp_path, flip(INTERNAL), flip(INDUSTRY), flip(FILED), {k: flip(v) for k, v in PEERS.items()})
    leads = ["peers", "internal", "internal", "industry"]
    info = make(tmp_path / "low.xlsx", files, leads=leads, direction="lower_is_worse", floor_at_zero=False)
    vals = recalc(tmp_path / "low.xlsx")
    r = reference(direction="lower_is_worse", floor_at_zero=False, internal=flip(INTERNAL), industry=flip(INDUSTRY),
                  filed=flip(FILED), peers={k: flip(v) for k, v in PEERS.items()}, leads=leads)
    assert_equals_reference(vals, info["synthesis"]["Credit card"], r)


def test_smoothing_applies_to_every_source_alike(tmp_path):
    files = inputs(tmp_path)
    info = make(tmp_path / "sm.xlsx", files, smoothing=2)
    vals = recalc(tmp_path / "sm.xlsx")
    sm = lambda pts: smooth(pts, 2)                                     # noqa: E731
    r = reference(internal=sm(INTERNAL), industry=sm(INDUSTRY), filed=sm(FILED),
                  peers={k: sm(v) for k, v in PEERS.items()})
    assert_equals_reference(vals, info["synthesis"]["Credit card"], r)


def test_a_latest_value_exactly_on_a_chosen_line_follows_on_a_line_rate_it(tmp_path):
    """The latest internal quarter sits exactly on line 3, taken from peers:
    one peer is the internal series itself, filed equals internal (so the
    carry is exactly 1) and the 25th percentile of five medians is the second
    exactly. 'worse' rates it 3, 'better' 2, in the workbook as in Python."""
    vals_ = [p.value for p in INTERNAL]
    m = statistics.median(vals_)
    internal = INTERNAL[:-1] + [Point(INTERNAL[-1].date, m)]
    assert statistics.median(p.value for p in internal) == m
    peers = dict(PEERS, p1=internal)
    pct, leads = [10, 25, 60, 95], ["internal", "peers", "internal", "industry"]
    files = inputs(tmp_path, internal=internal, filed=internal, peers=peers)
    info = make(tmp_path / "on.xlsx", files, leads=leads, pct=pct)
    at = info["synthesis"]["Credit card"]
    kw = dict(internal=internal, filed=internal, peers=peers, peer_percentiles=pct, leads=leads)
    worse, better = reference(**kw), reference(on_the_line="better", **kw)
    assert worse["chosen"][1] == m == worse["latest"][1] and (worse["rating"], better["rating"]) == (3, 2)
    assert_equals_reference(recalc(tmp_path / "on.xlsx"), at, worse)
    vals = edited(tmp_path / "on.xlsx", tmp_path, "on-better.xlsx", assess={"R%d" % SET_TOP: "better"})
    assert_equals_reference(vals, at, better)


# ---------------------------------------------------------------- refusals, in synthesis.py's words

def test_a_lead_from_a_source_without_that_line_is_refused(built, tmp_path):
    p, at, _, _ = built
    vals = edited(p, tmp_path, "nohigh.xlsx", synthesis={"F%d" % at["lines"][3]: "internal"})
    want = refused_as(leads=["internal"] * 4)
    assert want.startswith("Refused: line 5 High is led by internal, which has no such line: line 5:")
    assert refusal(vals, at) == want
    assert cell(vals, "D%d" % at["latest"]) == ""                     # no rating on refused lines


def test_lines_out_of_order_are_refused_naming_the_pair(built, tmp_path):
    p, at, _, _ = built
    leads = ["internal", "internal", "industry", "peers"]
    cells = {"D%d" % row: v for row, v in zip(at["lines"], [1, 2, 3, 4])}
    cells.update({"F%d" % row: v for row, v in zip(at["lines"], leads)})
    vals = edited(p, tmp_path, "order.xlsx", synthesis=cells)
    want = refused_as(leads=leads, peer_percentiles=[1, 2, 3, 4])
    assert "5 High (peers, " in want and "does not begin beyond 4 Moderate-High (industry, " in want
    assert refusal(vals, at) == want


@pytest.mark.parametrize("pct, words", [([20, 40, 40, 95], "peer percentiles must rise from line 2 to line 5"),
                                        ([20, 40, 75, 101], "peer percentiles are four numbers from 0 to 100, "
                                                            "one per line"),
                                        ([20, "forty", 75, 95], "peer percentiles are four numbers from 0 to 100, "
                                                                "one per line")])
def test_peer_percentiles_that_do_not_rise_or_are_not_percentiles_are_refused(built, tmp_path, pct, words):
    p, at, _, _ = built
    vals = edited(p, tmp_path, "pct%s.xlsx" % pct[3], synthesis={"D%d" % row: v for row, v in zip(at["lines"], pct)})
    assert refusal(vals, at) == "Refused: " + words
    if not isinstance(pct[1], str):
        with pytest.raises(ValueError, match=words):
            reference(peer_percentiles=pct)


@pytest.mark.parametrize("ref, typed, words", [
    ("F", "fed", "leads names the source of each of the four lines (2, 3, 4, 5), each one of internal, industry, "
                 "peers"),
    ("B", "half", "translation must be one of median ratio, none")])
def test_a_lead_or_translation_typed_outside_its_list_is_refused(built, tmp_path, ref, typed, words):
    p, at, _, _ = built
    row = at["lines"][0] if ref == "F" else at["translation"]
    vals = edited(p, tmp_path, "typed%s.xlsx" % ref, synthesis={"%s%d" % (ref, row): typed})
    assert refusal(vals, at) == "Refused: " + words
    with pytest.raises(ValueError) as exc:
        reference(**({"leads": [typed] + LEADS[1:]} if ref == "F" else {"translation": typed}))
    assert str(exc.value).startswith(words)


def test_peers_without_keybanks_filed_series_cannot_lead(tmp_path):
    files = inputs(tmp_path, filed=None)
    leads = ["peers", "internal", "internal", "industry"]
    info = make(tmp_path / "nofiled.xlsx", files, leads=leads)
    at = info["synthesis"]["Credit card"]
    vals = recalc(tmp_path / "nofiled.xlsx")
    want = refused_as(filed=None, leads=leads)
    assert "led by peers, which has no such line: KeyBank's filed series shares no quarter" in want
    assert refusal(vals, at) == want
    assert cell(vals, "B%d" % at["competitive"]).startswith("No competitive check")


def test_no_peers_and_no_shared_quarter_are_refused_and_none_clears_the_second(tmp_path):
    """Built with no peers and an industry series that ends before the
    internal one begins: High from the industry cannot be carried by a median
    ratio; from peers, nothing was built; taken as it is, it stands."""
    far = series(1990, [1.0 + 0.2 * ((i % 8) - 3.5) / 3.5 + (2.0 if 8 <= i < 14 else 0.0) for i in range(40)])
    files = inputs(tmp_path, industry=far, filed=None, peers={})
    info = make(tmp_path / "far.xlsx", files)
    at = info["synthesis"]["Credit card"]
    p = tmp_path / "far.xlsx"
    vals = recalc(p)
    assert refusal(vals, at) == refused_as(industry=far, filed=None, peers={})
    assert "the internal and industry series share no quarter" in refusal(vals, at)
    vals = edited(p, tmp_path, "far-peers.xlsx", synthesis={"F%d" % at["lines"][3]: "peers"})
    assert refusal(vals, at) == refused_as(industry=far, filed=None, peers={},
                                           leads=LEADS[:3] + ["peers"])
    assert refusal(vals, at).endswith("not built: no peer series were given")
    vals = edited(p, tmp_path, "far-none.xlsx", synthesis={"B%d" % at["translation"]: "none"})
    assert_equals_reference(vals, at, reference(industry=far, filed=None, peers={}, translation="none"))


# ---------------------------------------------------------------- around the tab

def test_assess_shows_the_synthesis_for_the_series_showing(built):
    p, at, vals, _ = built
    r = reference()
    lay = assess_layout(1)
    top = lay["expl_top"] + 6
    ws = load_workbook(p)["Assess"]
    assert ws.cell(top, 21).hyperlink.location == "Synthesis!A1" or "#Synthesis!A1" in str(ws.cell(top, 21).hyperlink.target)
    assert top + 7 < lay["dial_title"]
    assert vals["ASSESS!N%d" % top] == "Where each colour starts · Credit card"
    for k in range(4):
        row = top + 2 + k
        got = [vals["ASSESS!%s%d" % (c, row)] for c in "OPQRS"]
        want = [(r["candidates"][s]["lines"] or [None] * 4)[k] for s in ("internal", "industry", "peers")]
        for g, w in zip(got[:3], want):
            assert same(g, w)
        assert got[3] == pytest.approx(r["chosen"][k]) and got[4] == LEADS[k]
    assert vals["ASSESS!O%d" % (top + 6)] == "%s: %d %s" % (r["latest"][0], r["rating"],
                                                         ("Low", "Moderate-Low", "Moderate", "Moderate-High",
                                                          "High")[r["rating"] - 1])
    assert vals["ASSESS!N%d" % (top + 7)] == "Synthesis check: OK"


def test_a_series_without_a_synthesis_says_so_on_assess(tmp_path):
    files = inputs(tmp_path)
    other = write(tmp_path / "other.csv", INTERNAL, "another, made up")
    make(tmp_path / "two.xlsx", files, extra=[("Mortgage", other)])
    p = tmp_path / "two.xlsx"
    wb = load_workbook(p)
    wb["Assess"]["B5"] = "Mortgage"
    wb.save(tmp_path / "two-m.xlsx")
    vals = recalc(tmp_path / "two-m.xlsx")
    top = assess_layout(2)["expl_top"] + 6
    assert vals["ASSESS!N%d" % (top + 7)] == "No synthesis for Mortgage: the workbook was built without one for it"
    assert all(vals["ASSESS!%s%d" % (c, top + 2 + k)] == "" for c in "OPQRS" for k in range(4))


def test_without_synthesis_inputs_the_workbook_is_as_before(tmp_path):
    files = inputs(tmp_path)
    info = build(tmp_path / "plain.xlsx", [("Credit card", files["internal"])], "NCO", "%", "higher_is_worse",
                 "quarterly", 1, True, 0.5, 1.0, 2.5, "worse", (), (50, 75, 90, 95), 4, 20)
    assert info["synthesis"] == {}
    assert load_workbook(tmp_path / "plain.xlsx").sheetnames == [
        "Assess", "Thresholds", "Chart", "Evidence", "Backtest", "Sources", "Glossary", "Data", "Run"]


def test_the_synthesis_tabs_sit_beside_thresholds_and_data(built):
    p, _, _, _ = built
    assert load_workbook(p).sheetnames == ["Assess", "Thresholds", "Synthesis", "Chart", "Evidence", "Backtest",
                                           "Sources", "Glossary", "Data", "SynthesisData", "Run"]


def test_sources_and_run_carry_every_synthesis_file(built):
    p, _, _, files = built
    wb = load_workbook(p)
    src = {(wb["Sources"].cell(r, 1).value, wb["Sources"].cell(r, 2).value) for r in range(1, wb["Sources"].max_row + 1)}
    titles = {wb["Sources"].cell(r, 1).value for r in range(1, wb["Sources"].max_row + 1)}
    named = [("Credit card · industry", files["industry"]), ("Credit card · KeyBank filed", files["filed"])] + \
            [("Credit card · peer %s" % k, f) for k, f in files["peers"]]
    for title, path in named:
        assert title in titles
        assert ("Measure", dict(read_notes(path))["measure"]) in src
    run = {r[0].value: r[5].value for r in wb["Run"].iter_rows(min_row=4) if len(str(r[5].value or "")) == 64}
    for title, path in named:
        assert run[title] == hashlib.sha256(path.read_bytes()).hexdigest()


def test_every_synthesis_heading_is_defined_in_the_glossary(built):
    p, at, _, _ = built
    ws = load_workbook(p)["Synthesis"]
    terms = {t for t, _, _ in TERMS}
    plain = {"Rating begins", "Peer", "Median over the internal quarters", "Check"}
    heads = [c.value for c in ws[at["top"] + 1][:7] if c.value]
    heads += [ws.cell(at["translation"], 1).value, ws.cell(at["competitive"], 1).value]
    for h in heads:
        assert h in plain or h in terms, "no glossary entry for %r" % h


def test_no_synthesis_formula_points_above_row_one_or_runs_past_excels_limits(built):
    import re
    p, _, _, _ = built
    wb = load_workbook(p)
    for name in ("Synthesis", "SynthesisData", "Assess", "Run"):
        for row in wb[name].iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert len(c.value) < 8192, (name, c.coordinate)
                    for m in re.finditer(r"(?<![A-Za-z])\$?[A-Z]{1,3}\$?(-?\d+)(?![\d(])", c.value):
                        assert int(m.group(1)) >= 1, (name, c.coordinate, c.value)


def test_no_cell_the_synthesis_writes_is_an_error(built):
    _, _, vals, _ = built
    bad = {k: v for k, v in vals.items() if k.split("!")[0] in ("SYNTHESIS", "SYNTHESISDATA", "RUN")
           and type(v).__name__ == "XlError"}
    assert bad == {}


def test_the_starting_choices_are_the_ones_stated(built):
    p, at, _, _ = built
    ws = load_workbook(p)["Synthesis"]
    assert [ws["F%d" % r].value for r in at["lines"]] == LEADS
    assert [ws["D%d" % r].value for r in at["lines"]] == PCT
    assert ws["B%d" % at["translation"]].value == "median ratio"


# ---------------------------------------------------------------- the command line

def _cli(tmp_path, files, extra):
    return ["workbook", str(tmp_path / "c.xlsx"), "--series", "Credit card=%s" % files["internal"],
            "--name", "NCO", "--unit", "%", "--direction", "higher_is_worse", "--frequency", "quarterly",
            "--smoothing", "1", "--floor-at-zero", "yes", "--moderate-halfwidth", "0.5", "--low-step", "1.0",
            "--high-z", "2.5", "--on-the-line", "worse", "--percentiles", "50", "75", "90", "95",
            "--horizon", "4", "--min-history", "20"] + extra


def _spec(files):
    return "Credit card=industry:%s;filed:%s;peers:%s" % (files["industry"], files["filed"],
                                                           files["peers"][0][1].parent)


def test_the_command_builds_the_synthesis_from_its_flags(tmp_path, capsys):
    files = inputs(tmp_path)
    assert main(_cli(tmp_path, files, ["--synthesis", _spec(files), "--translation", "Credit card=median ratio",
                                       "--peer-percentiles", "Credit card=20,40,75,95",
                                       "--leads", "Credit card=internal,internal,internal,industry"])) == 0
    assert "synthesis for Credit card" in capsys.readouterr().out
    ws = load_workbook(tmp_path / "c.xlsx")["Synthesis"]
    assert ws["A4"].value == "Credit card" and ws["F9"].value == "industry"
    peers = [ws.cell(r, 1).value for r in range(18, 23)]
    assert peers == sorted(PEERS)                                       # every .csv in the folder, by name


@pytest.mark.parametrize("extra, words", [
    (["--translation", "Credit card=none", "--peer-percentiles", "Credit card=20,40,75,95"],
     "needs --leads"),
    (["--leads", "Credit card=internal,internal,internal,industry", "--peer-percentiles", "Credit card=20,40,75,95"],
     "needs --translation"),
    (["--translation", "Credit card=none", "--leads", "Credit card=internal,internal,internal,industry"],
     "has peers, so it needs --peer-percentiles"),
    (["--translation", "Credit card=none", "--leads", "Credit card=internal,internal,internal",
      "--peer-percentiles", "Credit card=20,40,75,95"], "leads names the source"),
    (["--translation", "Credit card=sideways", "--leads", "Credit card=internal,internal,internal,industry",
      "--peer-percentiles", "Credit card=20,40,75,95"], "translation must be one of"),
    (["--translation", "Credit card=none", "--leads", "Credit card=internal,internal,internal,industry",
      "--peer-percentiles", "Credit card=20,40,40,95"], "peer percentiles must rise"),
    (["--translation", "Mortgage=none", "--leads", "Credit card=internal,internal,internal,industry",
      "--peer-percentiles", "Credit card=20,40,75,95"], "--translation names 'Mortgage', which has no --synthesis"),
])
def test_the_command_refuses_a_synthesis_with_a_choice_missing(tmp_path, capsys, extra, words):
    files = inputs(tmp_path)
    assert main(_cli(tmp_path, files, ["--synthesis", _spec(files)] + extra)) == 2
    assert words in capsys.readouterr().err
    assert not (tmp_path / "c.xlsx").exists()


@pytest.mark.parametrize("args, words", [
    (["--leads", "Credit card=internal,internal,internal,industry"], "--leads is for the synthesis: give --synthesis"),
    (["--synthesis", "Mortgage=industry:x.csv"], "not a --series label"),
    (["--synthesis", "Credit card=filed:x.csv"], "needs industry:FILE"),
    (["--synthesis", "Credit card=industry:x.csv;fed:y.csv"], "takes industry:FILE"),
])
def test_the_command_refuses_synthesis_flags_that_do_not_fit(tmp_path, capsys, args, words):
    files = inputs(tmp_path)
    assert main(_cli(tmp_path, files, args)) == 2
    assert words in capsys.readouterr().err


def test_percentiles_without_peers_are_refused(tmp_path):
    files = inputs(tmp_path, peers={})
    with pytest.raises(ValueError, match="peer percentiles and peer series go together"):
        from_flags(["Credit card=industry:%s" % files["industry"]], ["Credit card=none"],
                   ["Credit card=20,40,75,95"], ["Credit card=internal,internal,internal,industry"], ["Credit card"])


def test_a_peer_pattern_takes_the_files_it_matches(tmp_path):
    files = inputs(tmp_path)
    got = from_flags(["Credit card=industry:%s;filed:%s;peers:%s" % (files["industry"], files["filed"],
                                                                      tmp_path / "peers" / "p[13].csv")],
                     ["Credit card=none"], ["Credit card=20,40,75,95"],
                     ["Credit card=internal,internal,internal,industry"], ["Credit card"])
    assert [nm for nm, _ in got["Credit card"]["peers"]] == ["p1", "p3"]
    with pytest.raises(ValueError, match="matches no .csv file"):
        from_flags(["Credit card=industry:x.csv;peers:%s" % (tmp_path / "nothing_*.csv")], [], [], [], ["Credit card"])


def test_the_reference_competitive_check_is_counted_from_the_medians():
    """A guard on the guard: the counts the workbook is held to are worked here
    from the peers' medians directly, not only taken from synthesis.py."""
    r = reference()
    window = {p.date for p in INTERNAL}
    kb = statistics.median(p.value for p in FILED if p.date in window)
    meds = [statistics.median(p.value for p in v if p.date in window) for v in PEERS.values()]
    assert r["check"]["filed_median"] == pytest.approx(kb)
    assert r["check"]["worse_than"] == sum(1 for m in meds if m < kb)
    assert kb != pytest.approx(statistics.median(p.value for p in FILED))     # the early quarters would move it
