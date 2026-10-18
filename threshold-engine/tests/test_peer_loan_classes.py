"""The Peer loan classes tab (peers/classes.py, peers/loan_classes_tab.py):
the SEC pull, reading a filing's classes, the stated file, and the tab. Offline:
a made-up filing stands in for EDGAR."""
import json

import pytest
from openpyxl import Workbook

from conftest import ROOT
from threshold_engine.peers import classes as C
from threshold_engine.peers import loan_classes_tab as LC

PERIOD = "2026-06-30"
INSTANCE = """<?xml version="1.0" encoding="utf-8"?>
<xbrli:xbrl xmlns:xbrli="http://www.xbrl.org/2003/instance" xmlns:xbrldi="http://xbrl.org/2006/xbrldi"
  xmlns:us-gaap="http://fasb.org/us-gaap/2025" xmlns:bk="http://bank.example/2026">
 <xbrli:context id="tot"><xbrli:entity><xbrli:identifier scheme="x">1</xbrli:identifier></xbrli:entity>
  <xbrli:period><xbrli:instant>2026-06-30</xbrli:instant></xbrli:period></xbrli:context>
 <xbrli:context id="stu"><xbrli:entity><xbrli:identifier scheme="x">1</xbrli:identifier>
  <xbrli:segment><xbrldi:explicitMember dimension="us-gaap:FinancingReceivablePortfolioSegmentAxis">us-gaap:ConsumerPortfolioSegmentMember</xbrldi:explicitMember>
  <xbrldi:explicitMember dimension="us-gaap:FinancingReceivableRecordedInvestmentByClassOfFinancingReceivableAxis">us-gaap:StudentLoanMember</xbrldi:explicitMember></xbrli:segment></xbrli:entity>
  <xbrli:period><xbrli:instant>2026-06-30</xbrli:instant></xbrli:period></xbrli:context>
 <xbrli:context id="oth"><xbrli:entity><xbrli:identifier scheme="x">1</xbrli:identifier>
  <xbrli:segment><xbrldi:explicitMember dimension="us-gaap:FinancingReceivableRecordedInvestmentByClassOfFinancingReceivableAxis">bk:OtherRetailMember</xbrldi:explicitMember>
  <xbrldi:explicitMember dimension="us-gaap:FinancingReceivablePortfolioSegmentAxis">us-gaap:ConsumerPortfolioSegmentMember</xbrldi:explicitMember></xbrli:segment></xbrli:entity>
  <xbrli:period><xbrli:instant>2026-06-30</xbrli:instant></xbrli:period></xbrli:context>
 <xbrli:context id="late"><xbrli:entity><xbrli:identifier scheme="x">1</xbrli:identifier>
  <xbrli:segment><xbrldi:explicitMember dimension="us-gaap:FinancingReceivablesPeriodPastDueAxis">us-gaap:FinancingReceivables30To59DaysPastDueMember</xbrldi:explicitMember></xbrli:segment></xbrli:entity>
  <xbrli:period><xbrli:instant>2026-06-30</xbrli:instant></xbrli:period></xbrli:context>
 <xbrli:context id="before"><xbrli:entity><xbrli:identifier scheme="x">1</xbrli:identifier>
  <xbrli:segment><xbrldi:explicitMember dimension="us-gaap:FinancingReceivableRecordedInvestmentByClassOfFinancingReceivableAxis">us-gaap:StudentLoanMember</xbrldi:explicitMember></xbrli:segment></xbrli:entity>
  <xbrli:period><xbrli:instant>2025-12-31</xbrli:instant></xbrli:period></xbrli:context>
 <us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss contextRef="tot" unitRef="usd" decimals="-6">100000000000</us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss>
 <us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss contextRef="stu" unitRef="usd" decimals="-6">8044000000</us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss>
 <us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss contextRef="oth" unitRef="usd" decimals="-6">4031000000</us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss>
 <us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss contextRef="late" unitRef="usd" decimals="-6">5000000</us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss>
 <us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss contextRef="before" unitRef="usd" decimals="-6">9000000000</us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss>
</xbrli:xbrl>"""
LABELS = """<?xml version="1.0" encoding="utf-8"?>
<link:linkbase xmlns:link="http://www.xbrl.org/2003/linkbase" xmlns:xlink="http://www.w3.org/1999/xlink">
 <link:labelLink xlink:type="extended" xlink:role="http://www.xbrl.org/2003/role/link">
  <link:loc xlink:type="locator" xlink:href="bk-20260630.xsd#bk_OtherRetailMember" xlink:label="loc_o"/>
  <link:label xlink:type="resource" xlink:label="lab_o" xlink:role="http://www.xbrl.org/2003/role/label">Other Retail [Member]</link:label>
  <link:labelArc xlink:type="arc" xlink:from="loc_o" xlink:to="lab_o"/>
 </link:labelLink>
</link:linkbase>"""
SUBMISSIONS = {"name": "MADE-UP BANCORP", "filings": {"recent": {
    "form": ["10-Q", "8-K", "10-Q", "10-Q/A", "10-K"],
    "reportDate": ["2026-06-30", "2026-07-01", "2026-03-31", "2026-06-30", "2025-12-31"],
    "filingDate": ["2026-08-03", "2026-07-02", "2026-05-04", "2026-08-20", "2026-02-20"],
    "accessionNumber": ["0000000001-26-000001", "x", "0000000001-26-000000", "0000000001-26-000009",
                        "0000000001-26-000002"],
    "primaryDocument": ["bk-20260630.htm", "x.htm", "bk-20260331.htm", "bka-20260630.htm", "bk-20251231.htm"]}}}
UA = "A Firm someone@example.com"


def fake_edgar(seen):
    def fetch(url):
        seen.append(url)
        if url.endswith("CIK0000000001.json"):
            return json.dumps(SUBMISSIONS).encode()
        if url.endswith("index.json"):
            return json.dumps({"directory": {"item": [{"name": "bk-20260630.htm"}, {"name": "bk-20260630_htm.xml"},
                                                      {"name": "bk-20260630_lab.xml"}]}}).encode()
        if url.endswith("_htm.xml"):
            return INSTANCE.encode()
        if url.endswith("_lab.xml"):
            return LABELS.encode()
        raise AssertionError(url)
    return fetch


def spec(**over):
    bank = {"short": "Made-up", "cik": 1, "holding_company": "Made-up Bancorp", "filing_url": "https://x/bk.htm",
            "classes": [{"members": ["us-gaap:ConsumerPortfolioSegmentMember", "us-gaap:StudentLoanMember"],
                         "name": "Education", "is": "Student"},
                        {"members": ["bk:OtherRetailMember", "us-gaap:ConsumerPortfolioSegmentMember"],
                         "is": "Mixed other consumer"},
                        {"name": "Unsecured", "amount_musd": 626, "where": "Table 44", "is": "Unsecured or personal"}],
            "findings": {"student": {"answer": "Yes", "short": "Own class", "quote": "five categories ... education", "where": "MD&A"},
                         "unsecured": {"answer": "Yes", "short": "Own line", "quote": "Unsecured | 626", "where": "Table 44"},
                         "secured_non_auto": {"answer": "None stated in the 10-Q", "short": "Not stated", "quote": "not stated",
                                              "where": "searched RV, marine"}},
            "comparison": {"answer": "Mixed", "why": "Education is twice other retail."}}
    bank.update(over)
    return {"quarter": "2026Q2", "banks": [bank]}


def test_the_contact_is_required_and_never_stored(tmp_path):
    for bad in ("", "A Firm", "a@b"):
        with pytest.raises(C.Refused, match="SEC requires a contact email"):
            C.pull(tmp_path, [{"short": "Made-up", "cik": 1}], "2026Q2", bad, fetch=lambda u: b"")
    seen = []
    got = C.pull(tmp_path, [{"short": "Made-up", "cik": 1}], "2026Q2", UA, fetch=fake_edgar(seen), pause=0)
    # The latest filed for the quarter: the amendment, not the original or another quarter.
    assert got["1"]["accession"] == "0000000001-26-000009" and got["1"]["form"] == "10-Q/A"
    assert seen[0] == "https://data.sec.gov/submissions/CIK0000000001.json"
    assert all("someone@example.com" not in p.read_text(encoding="utf-8") for p in tmp_path.rglob("*") if p.is_file())
    said = []
    C.pull(tmp_path, [{"short": "Made-up", "cik": 1}], "2026Q2", UA, fetch=fake_edgar([]), say=said.append, pause=0)
    assert said == ["Made-up: 0000000001-26-000009 already held"]


def test_a_december_quarter_reads_the_10k_and_a_missing_filing_is_said():
    assert C.filing_for(SUBMISSIONS, "2025-12-31")["form"] == "10-K"
    assert C.filing_for(SUBMISSIONS, "2026-09-30") is None
    with pytest.raises(C.Refused, match="YYYYQn"):
        C.quarter_end("2026-06")


def test_facts_keep_only_segment_and_class_cells_at_the_quarter_end(tmp_path):
    (tmp_path / "i.xml").write_text(INSTANCE, encoding="utf-8")
    got = C.facts(tmp_path / "i.xml", PERIOD)
    assert got == {(): 100e9, ("us-gaap:ConsumerPortfolioSegmentMember", "us-gaap:StudentLoanMember"): 8044e6,
                   ("bk:OtherRetailMember", "us-gaap:ConsumerPortfolioSegmentMember"): 4031e6}
    (tmp_path / "two.xml").write_text(INSTANCE.replace(
        '<us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss contextRef="tot" unitRef="usd" '
        'decimals="-6">100000000000',
        '<us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss contextRef="stu" unitRef="usd" '
        'decimals="-6">1</us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss>'
        '<us-gaap:FinancingReceivableExcludingAccruedInterestBeforeAllowanceForCreditLoss contextRef="tot" unitRef="usd" '
        'decimals="-6">100000000000'), encoding="utf-8")
    with pytest.raises(C.Refused, match="two values"):
        C.facts(tmp_path / "two.xml", PERIOD)


def test_the_file_states_every_judgement_with_the_filings_words(tmp_path):
    p = tmp_path / "c.json"
    p.write_text(json.dumps(spec()), encoding="utf-8")
    assert LC.load(p)["banks"][0]["short"] == "Made-up"
    bad = spec(classes=[{"members": ["us-gaap:StudentLoanMember"], "is": "Education"}],
               findings={"student": {"answer": "Yes", "quote": "", "where": "MD&A"}},
               comparison={"answer": "Fine", "why": ""})
    p.write_text(json.dumps(bad), encoding="utf-8")
    with pytest.raises(C.Refused) as e:
        LC.load(p)
    msg = str(e.value)
    assert "counted as 'Education'" in msg and "the student finding needs" in msg
    assert "the unsecured finding needs" in msg and "comparison needs an answer" in msg


def test_the_tab_reads_amounts_from_the_filing_and_sums_them_live(tmp_path):
    C.pull(tmp_path / "sec", [{"short": "Made-up", "cik": 1}], "2026Q2", UA, fetch=fake_edgar([]), pause=0)
    wb = Workbook()
    ws = wb.active
    info = LC.write(ws, spec(), tmp_path / "sec")
    rows = {ws.cell(r, 2).value: (ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 6).value)
            for r in range(info["first"], info["last"] + 1)}
    assert rows == {"Education": (8044.0, "Student", "Read from the filing's XBRL"),
                    "Other Retail": (4031.0, "Mixed other consumer", "Read from the filing's XBRL"),
                    "Unsecured": (626.0, "Unsecured or personal", "As printed (not tagged by class)")}
    r = info["summary"]["Made-up"]
    assert ws.cell(r, 2).value.startswith('=IF(COUNTIFS(') and '"Student"' in ws.cell(r, 2).value
    # Where no class is counted as a kind, the finding's short form shows; its full answer is in the words table.
    assert '"Not stated"' in ws.cell(r, 4).value
    assert ws.cell(info["first"], 5).hyperlink.target == "https://www.sec.gov/Archives/edgar/data/1/" \
                                                          "000000000126000009/bka-20260630.htm"
    missing = spec(classes=[{"members": ["us-gaap:AutomobileLoanMember"], "is": "Other"}])
    with pytest.raises(C.Refused, match="has no amortized cost"):
        LC.write(Workbook().active, missing, tmp_path / "sec")


def test_the_repositorys_file_loads_and_cites_every_filing():
    s = LC.load(ROOT / "config" / "peer-loan-classes.json")
    assert s["quarter"] == "2026Q2" and len(s["banks"]) == 17
    assert all(b["filing_url"].startswith("https://www.sec.gov/Archives/edgar/data/%d/" % b["cik"]) for b in s["banks"])
    assert all("amount_musd" not in c for b in s["banks"] for c in b["classes"] if "members" in c)
