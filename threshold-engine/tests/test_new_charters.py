"""A new charter's first filing counts as that quarter's activity.

The firm, 18 Oct 2026, question 21 of the round 2 commercial docket ("agreed
with the recommendation"): "A new charter's first filing counts as that
quarter's activity, where the FDIC record shows it opened that quarter".

A bank added pro forma whose first filing falls partway through a year has
no filing the quarter before, so its flow that quarter (a year to date less
the quarter before's) was not known and every window holding it was missing.
Now its first year to date counts as that quarter's activity when the FDIC
institutions record (api.fdic.gov/banks/institutions, ESTYMD) shows it opened
in that quarter; otherwise that quarter stays missing. The record is pulled
and cached with the merger record ('peers pull --mergers'); the build reads
it from the cache, never the network, and refuses, naming the bank, when the
cache cannot answer.

Real data: tests/data/peers/keybank_usa_1994.json is KeyBank (17534) and Key
Bank USA (34055), 1994Q1 to 1997Q4, from the commercial cache; Key Bank USA
first filed in 1995Q3 and KeyBank absorbed it in 2004Q4 (FDIC history).
tests/data/peers/institutions_keybank.json is six institutions rows as the
API returned them on 18 Oct 2026: Key Bank USA opened 5 Sep 1995. The unit
tests on ``books`` use made-up banks "1" and "2", as test_peers_books.py does.
"""
import json
import urllib.parse

import pytest

from conftest import ROOT
from threshold_engine import cli
from threshold_engine.peers import books, fdic, run
from threshold_engine.series import read_csv

PEERS = ROOT / "tests" / "data" / "peers"
KEYBANK, KEY_USA = "17534", "34055"


def row(cert, q, ln, dr, cr=0.0, parent="9"):
    return {"CERT": cert, "REPDTE": q, "LNCRCD": ln, "DRCRCD": dr, "CRCRCD": cr, books.PARENT: parent}


def quarters(first_year, n):
    return [q for y in range(first_year, first_year + 3) for q in
            ("%d0331" % y, "%d0630" % y, "%d0930" % y, "%d1231" % y)][:n]


def made_up():
    """Bank 1 files from 2020Q1; bank 2 first files in 2020Q3 (year to date
    7 by then, 10 by Q4) and is absorbed by bank 1 in 2021Q2."""
    qs = quarters(2020, 8)
    data = {q: [row(1, q, 1000.0, 4.0 * (1 + i % 4))] for i, q in enumerate(qs)}
    data["20200930"].append(row(2, "20200930", 200.0, 7.0))
    data["20201231"].append(row(2, "20201231", 210.0, 10.0))
    data["20210331"].append(row(2, "20210331", 220.0, 3.0))
    return data


# ---------------------------------------------------------------- books

def test_a_members_first_filing_partway_through_a_year_is_found():
    data = made_up()
    assert books.first_filings(data, [("2", "20210630")]) == {"2": "20200930"}
    # not after its until, not in the data's first quarter, not in a Q1
    assert books.first_filings(data, [("2", "20200630")]) == {}
    assert books.first_filings(data, [("1", "20210630")]) == {}
    q1 = {q: rows for q, rows in data.items() if q >= "20201231"}
    assert books.first_filings(q1, [("2", "20210630")]) == {}
    first_q1 = dict(data)
    first_q1["20200930"] = [r for r in data["20200930"] if r["CERT"] != 2]
    first_q1["20201231"] = [r for r in data["20201231"] if r["CERT"] != 2]
    assert books.first_filings(first_q1, [("2", "20210630")]) == {}


def test_it_counts_only_when_the_record_shows_it_opened_that_quarter():
    firsts = {"2": "20200930"}
    assert books.opened_in_first_quarter(firsts, {"2": "2020-08-03"}) == {"2": "20200930"}
    assert books.opened_in_first_quarter(firsts, {"2": "2020-07-01"}) == {"2": "20200930"}
    assert books.opened_in_first_quarter(firsts, {"2": "2020-06-30"}) == {}
    assert books.opened_in_first_quarter(firsts, {"2": "1985-01-01"}) == {}
    assert books.opened_in_first_quarter(firsts, {}) == {}


def test_its_first_year_to_date_is_that_quarters_flow_and_otherwise_none():
    data = made_up()
    members = [("2", "20210630")]
    before = {q: (g, b) for q, g, _, b in books.combined_flows(data, 1, "Credit card", members, {})}
    after = {q: (g, b) for q, g, _, b in
             books.combined_flows(data, 1, "Credit card", members, {"2": "20200930"})}
    # bank 1's own flow in 2020Q3 is 12 - 8 = 4; bank 2's year to date is 7
    assert before["20200930"] == (None, 1200.0)
    assert after["20200930"] == (4.0 + 7.0, 1200.0)
    # the next quarter is a year to date less the one before, as ever
    assert after["20201231"] == before["20201231"] == (4.0 + 3.0, 1210.0)
    # a 2021Q1 flow is the year to date
    assert after["20210331"] == before["20210331"] == (4.0 + 3.0, 1220.0)
    # the bank's own first filing is never read this way
    own = {q: g for q, g, _, _ in books.combined_flows(data, 2, "Credit card", [("1", "20211231")], {"2": "20200930"})}
    assert own["20200930"] is None


def test_the_window_holding_it_is_kept_and_without_the_rule_missing():
    data = made_up()
    members = [("2", "20210630")]
    without = {r["quarter"]: r for r in books.bank_book(data, 1, "Credit card", 10.0, 0, acquired=(), stated=(),
                                                         members=members, opened={})}
    withit = {r["quarter"]: r for r in books.bank_book(data, 1, "Credit card", 10.0, 0, acquired=(), stated=(),
                                                        members=members, opened={"2": "20200930"})}
    assert without["2021Q1"]["why"] == books.WHY_MISSING
    assert withit["2021Q1"]["kept"]
    # 2020Q2 to 2021Q1: bank 1's 4 a quarter; bank 2's 7, then 3 and 3
    assert withit["2021Q1"]["gross"] == pytest.approx(4 + (4 + 7) + (4 + 3) + (4 + 3))


def test_opened_is_required_with_members():
    data = made_up()
    with pytest.raises(ValueError, match="opened must be"):
        books.combined_flows(data, 1, "Credit card", [("2", "20210630")])
    with pytest.raises(ValueError, match="opened must be"):
        books.bank_book(data, 1, "Credit card", 0.25, 0, acquired=(), stated=(), members=[("2", "20210630")])
    # without members, nothing is needed
    assert books.combined_flows(data, 1, "Credit card") == books.quarterly_flows(data, 1, "Credit card")


def test_the_record_reads_the_fdics_dates_and_keeps_an_unreadable_one():
    rows = json.loads((PEERS / "institutions_keybank.json").read_text(encoding="utf-8"))
    dated, undated = fdic.established(rows + [{"CERT": 1, "NAME": "x", "ESTYMD": "9/5/1995"},
                                              {"CERT": 2, "NAME": "y", "ESTYMD": None},
                                              {"CERT": 3, "NAME": "z", "ESTYMD": "02/30/1995"}])
    assert dated[KEY_USA] == {"name": "Key Bank USA, National Association", "established": "1995-09-05"}
    assert dated[KEYBANK]["established"] == "1849-01-01" and len(dated) == 6
    assert undated == {"1": {"name": "x", "established": "9/5/1995"}, "2": {"name": "y", "established": None},
                       "3": {"name": "z", "established": "02/30/1995"}}


# ---------------------------------------------------------------- pull

class Fake:
    """The institutions endpoint over the fixture, paged as the API pages."""

    def __init__(self, rows):
        self.rows, self.urls = rows, []

    def __call__(self, url):
        self.urls.append(url)
        q = urllib.parse.parse_qs(urllib.parse.urlsplit(url).query)
        offset, limit = int(q["offset"][0]), int(q["limit"][0])
        assert url.startswith(fdic.INSTITUTIONS_URL)
        return {"data": [{"data": r} for r in self.rows[offset:offset + limit]], "meta": {"total": len(self.rows)}}


def institutions_rows():
    return json.loads((PEERS / "institutions_keybank.json").read_text(encoding="utf-8"))


def test_the_pull_writes_the_record_and_refuses_an_empty_one(tmp_path):
    fake = Fake(institutions_rows())
    path = run.write_institutions(tmp_path, fake, say=lambda s: None)
    got = run.load_institutions(path)
    assert got["rows"] == 6 and got["source"] == fdic.INSTITUTIONS_URL and got["undated"] == {}
    assert got["established"][KEY_USA]["established"] == "1995-09-05"
    assert "ESTYMD" in urllib.parse.unquote_plus(fake.urls[0])
    with pytest.raises(RuntimeError, match="no institutions"):
        run.write_institutions(tmp_path / "empty", Fake([]), say=lambda s: None)
    assert not (tmp_path / "empty" / run.INSTITUTIONS).exists()


def test_a_malformed_record_is_refused_by_name(tmp_path):
    path = tmp_path / run.INSTITUTIONS
    path.write_text(json.dumps({"34055": "1995-09-05"}), encoding="utf-8")
    with pytest.raises(ValueError, match="pull the merger record again"):
        run.load_institutions(path)
    good = {"source": "s", "fields": "f", "pulled": "2026-10-18", "rows": 1, "undated": {},
            "established": {"34055": {"name": "n", "established": "09/05/1995"}}}
    path.write_text(json.dumps(good), encoding="utf-8")
    with pytest.raises(ValueError, match="not a date written like"):
        run.load_institutions(path)


# ---------------------------------------------------------------- build

def fixture():
    return json.loads((PEERS / "keybank_usa_1994.json").read_text(encoding="utf-8"))


def lay_out(tmp, institutions=True, drop=(), pulled=None):
    """A cache from the fixture, the merger record of the one absorption, and
    (unless ``institutions`` is False) the institutions record written by
    the pull from the fixture's rows, less ``drop``, with ``pulled`` set."""
    cache = tmp / "cache"
    cache.mkdir()
    for q, rows in fixture().items():
        (cache / ("q%s.json" % q)).write_text(json.dumps(rows), encoding="utf-8")
    (cache / "mergers.json").write_text(json.dumps({KEYBANK: ["2004-12-31"]}), encoding="utf-8")
    (cache / "absorptions.json").write_text(json.dumps(
        [{"acquirer": KEYBANK, "absorbed": KEY_USA, "quarter": "2004-12-31"}]), encoding="utf-8")
    if institutions:
        path = run.write_institutions(cache, Fake([r for r in institutions_rows() if str(r["CERT"]) not in drop]),
                                      say=lambda s: None)
        if pulled:
            rec = json.loads(path.read_text(encoding="utf-8"))
            rec["pulled"] = pulled
            path.write_text(json.dumps(rec), encoding="utf-8")
    return cache


def build(tmp, cache, monkeypatch, pro_forma="yes"):
    monkeypatch.setattr(fdic, "urllib_fetch", lambda url: (_ for _ in ()).throw(AssertionError(url)))
    return cli.main([
        "peers", "build", str(tmp / "peers.xlsx"), "--cache", str(cache),
        "--roster", str(ROOT / "config" / "peers.json"), "--series-dir", str(tmp / "series"),
        "--first", "1994Q1", "--last", "1997Q4", "--products", "C&I",
        "--mergers", str(cache / "mergers.json"), "--absorptions", str(cache / "absorptions.json"),
        "--extra-acquisitions", "--industry-combined", "--delinquency", "no",
        "--reporting-changes", str(ROOT / "config" / "reporting-changes.json"),
        "--stated-acquisitions", "none",
        "--jump", "0.25", "--acquisition-share", "0.05", "--min-book", "0", "--exempt-self", "yes",
        "--pro-forma", pro_forma, "--pack", "all", "--min-banks", "1", "--rank-by", "net",
        "--last-n", "4", "--min-obs", "2", "--base-quarter", "1995Q4", "--index-min-book", "0",
        "--lagged-years", "1995", "--lag-years", "1", "--normalisation-years", "1995", "1996",
        "--mix-base-year", "1995"])


def run_lines(tmp):
    from openpyxl import load_workbook
    ws = load_workbook(tmp / "peers.xlsx")["Run"]
    return [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)]


def test_the_fixture_is_what_the_docstring_says():
    data = fixture()
    first = min(q for q, rows in data.items() if any(str(r["CERT"]) == KEY_USA for r in rows))
    assert first == "19950930" and len(data) == 16 and sum(len(v) for v in data.values()) == 26
    assert {str(r["CERT"]) for rows in data.values() for r in rows} == {KEYBANK, KEY_USA}


def test_keybanks_ci_runs_across_key_bank_usas_first_filing(tmp_path, monkeypatch):
    cache = lay_out(tmp_path)
    assert build(tmp_path, cache, monkeypatch) == 0
    path = tmp_path / "series" / "c-i_keybank_net.csv"
    pts = read_csv(path)
    # From the first window (1995Q1) to 1997Q1; KeyBank's own C&I more than
    # doubled in 1997Q2 (5,538,986 to 11,594,374), so the jump ends it there.
    assert (pts[0].date, pts[-1].date, len(pts)) == ("1995-03-31", "1997-03-31", 9)
    # 1995Q3 by hand: KeyBank's own flow is its year to date less 1995Q2's,
    # and Key Bank USA's first year to date is its flow (0: it filed no C&I).
    data = fixture()

    def kb(q, f):
        return next(r[f] for r in data[q] if str(r["CERT"]) == KEYBANK)

    def usa(q, f):
        return next((r[f] for r in data[q] if str(r["CERT"]) == KEY_USA), 0)
    net = {"19950331": kb("19950331", "DRCI") - kb("19950331", "CRCI")}
    for q, p in (("19950630", "19950331"), ("19950930", "19950630"), ("19951231", "19950930")):
        net[q] = (kb(q, "DRCI") - kb(p, "DRCI")) - (kb(q, "CRCI") - kb(p, "CRCI"))
    net["19950930"] += usa("19950930", "DRCI") - usa("19950930", "CRCI")
    net["19951231"] += (usa("19951231", "DRCI") - usa("19950930", "DRCI")) \
        - (usa("19951231", "CRCI") - usa("19950930", "CRCI"))
    avg = sum(kb(q, "LNCI") + usa(q, "LNCI") for q in net) / 4
    q4 = next(p for p in pts if p.date == "1995-12-31")
    assert q4.value == pytest.approx(100 * sum(net.values()) / avg)
    # The workbook's Data tab reads the same rule: the 1995Q4 row is kept.
    from openpyxl import load_workbook
    da = load_workbook(tmp_path / "peers.xlsx")["Data"]
    head = [c.value for c in da[1]]
    rows = [dict(zip(head, (c.value for c in r))) for r in da.iter_rows(min_row=2)]
    kept = {r["Quarter"]: r["Kept"] for r in rows if r["Bank"] == "KeyBank" and r["Product"] == "C&I"}
    assert kept["1995Q3"] is True and kept["1995Q4"] is True and kept["1996Q2"] is True
    lines = "\n".join(run_lines(tmp_path))
    assert ("First filing counted, KeyBank: 1 bank: Key Bank USA, National Association (certificate 34055), "
            "first filed 1995Q3, opened 1995-09-05.") in lines
    assert "institutions.json, pulled " in lines


def test_a_bank_the_record_does_not_hold_stays_missing_and_is_named(tmp_path, monkeypatch):
    # The record pulled whole, after 1995Q3, and with no row for 34055: it
    # does not show the bank opened then, so 1995Q3 has no flow, as before.
    cache = lay_out(tmp_path, drop=(KEY_USA,))
    assert build(tmp_path, cache, monkeypatch) == 0
    path = tmp_path / "series" / "c-i_keybank_net.csv"
    pts = read_csv(path)
    # 1995Q3 to 1996Q2 hold it, so the series starts at 1996Q3
    assert (pts[0].date, pts[-1].date, len(pts)) == ("1996-09-30", "1997-03-31", 3)
    assert ("First filing not counted, not in the record, KeyBank: 1 bank: certificate 34055, first filed "
            "1995Q3.") in run_lines(tmp_path)


def test_a_cache_without_the_record_is_refused_naming_the_bank(tmp_path, monkeypatch, capsys):
    cache = lay_out(tmp_path, institutions=False)
    assert build(tmp_path, cache, monkeypatch) == 2
    err = capsys.readouterr().err
    assert "certificate 34055 (first filed 1995Q3, in KeyBank's history)" in err
    assert "institutions.json is not there" in err and "--mergers SINCE" in err
    assert not (tmp_path / "peers.xlsx").exists() and not (tmp_path / "series").exists()


def test_a_record_pulled_before_the_bank_first_filed_is_refused(tmp_path, monkeypatch, capsys):
    # A record pulled on 1995-09-30 could not hold a bank first filing for
    # the quarter ending that day, so its absence says nothing.
    cache = lay_out(tmp_path, drop=(KEY_USA,), pulled="1995-09-30")
    assert build(tmp_path, cache, monkeypatch) == 2
    err = capsys.readouterr().err
    assert "pulled on 1995-09-30" in err and "certificate 34055 (first filed 1995Q3, in KeyBank's history)" in err
    # pulled the day after, its absence is an answer: built, the quarter missing
    later = tmp_path / "later"
    later.mkdir()
    assert build(later, lay_out(later, drop=(KEY_USA,), pulled="1995-10-01"), monkeypatch) == 0
    assert read_csv(later / "series" / "c-i_keybank_net.csv")[0].date == "1996-09-30"


def test_without_pro_forma_no_record_is_needed(tmp_path, monkeypatch):
    cache = lay_out(tmp_path, institutions=False)
    assert build(tmp_path, cache, monkeypatch, pro_forma="no") == 0
    assert not [x for x in run_lines(tmp_path) if "First filing" in x]
