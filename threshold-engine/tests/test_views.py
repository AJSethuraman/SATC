"""Saved views (views_tab): Signals' picks kept by number with a note, and the
tab following a view when its Saved view box holds that number."""

import pytest
from openpyxl import load_workbook

from threshold_engine import views_tab as VT
from test_signals import GRID, tiny_book
from test_workbook import recalc

formulas = pytest.importorskip("formulas")


def made(tmp_path):
    wb, info = tiny_book(tmp_path)
    out = VT.build(wb, "=Data!$A$2:$A$%d" % (1 + len(GRID)))
    p = tmp_path / "v.xlsx"
    wb.save(p)
    return p, info, out


def test_the_tab_is_laid_out_with_every_signals_pick_and_the_box_on_signals(tmp_path):
    p, info, out = made(tmp_path)
    wb = load_workbook(p)
    ws, sg = wb[VT.TITLE], wb["Signals"]
    ps = out["picks"]["Signals"]
    assert [x["cell"] for x in ps] == ["B5", "E5", "I5", "K5", "M5", "B20", "C20", "D20", "E20", "F20"]
    assert [x["label"] for x in ps][:5] == ["Compare against: product", "Measure", "Source", "Shift the signal by",
                                           "Window for the normal"]
    assert ps[5]["label"] == "On the charts: Credit card"
    assert all(x["reads"] > 0 for x in ps)
    ly = out["layout"]["Signals"]
    assert [ws.cell(ly["head"], c).value for c in (2, 3, 4, 5)] == ["View", "Note", "Data through",
                                                                   "Compare against: product"]
    assert [ws.cell(r, 2).value for r in range(ly["first"], ly["last"] + 1)] == list(range(1, VT.VIEWS + 1))
    assert ws.cell(ly["now"], 5).value == "='Signals'!$B$5"
    assert sg["U4"].value == "Saved view" and sg["U5"].value is None
    # nothing on Signals reads its own yellow cells any more but the in-effect cells
    for row in sg.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                assert "'Signals'!$B$5" not in c.value and "'Signals'!$E$5" not in c.value, c.coordinate


def test_a_saved_view_replaces_the_picks_it_holds_and_keeps_the_rest(tmp_path):
    p, info, out = made(tmp_path)
    wb = load_workbook(p)
    ly = out["layout"]["Signals"]
    ws = wb[VT.TITLE]
    r = ly["first"] + 2                                     # view 3
    ws.cell(r, 3, "card against the pack")
    ws.cell(r, 5, "Credit card")
    wb["Signals"]["U5"] = 3
    wb.save(tmp_path / "v3.xlsx")
    v = recalc(tmp_path / "v3.xlsx")
    assert v["SIGNALS!P5"].startswith("Credit card · 30+ days past due incl. nonaccrual · KeyBank")
    assert v["SIGNALS!U6"] == "Saved view 3: card against the pack"
    # an empty view says so, and the tab keeps its own picks
    wb["Signals"]["U5"] = 4
    wb.save(tmp_path / "v4.xlsx")
    v = recalc(tmp_path / "v4.xlsx")
    assert v["SIGNALS!P5"].startswith("Other consumer · ")
    assert v["SIGNALS!U6"] == "Saved view 4 has no picks"
    # no view: blank note, the tab's picks
    wb["Signals"]["U5"] = None
    wb.save(tmp_path / "v0.xlsx")
    v = recalc(tmp_path / "v0.xlsx")
    assert v["SIGNALS!P5"].startswith("Other consumer · ") and v["SIGNALS!U6"] == ""


def test_a_pick_no_formula_reads_is_refused(tmp_path):
    wb, info = tiny_book(tmp_path)
    wb["Signals"]["B5"].value  # the pick exists
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and "'Signals'!$M$5" in c.value:
                    c.value = c.value.replace("'Signals'!$M$5", '"All quarters"')
    with pytest.raises(ValueError, match="no formula reads Signals!M5"):
        VT.build(wb, "=Data!$A$2:$A$%d" % (1 + len(GRID)))
