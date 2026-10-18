"""Table slicers written into an openpyxl workbook by hand (slicers.py).

Excel itself is not here; tools/check_slicers_in_excel.ps1 opens the output
in Excel 16, proves it opens without a repair, and clicks a button. These
hold the package to what Excel wrote when it added the same slicers.
"""
import posixpath
import re
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.table import Table

from threshold_engine import slicers as S
from threshold_engine.slicers import add_table_slicers

M = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
HEADER_ROW = 4
ROWS = [
    ("Credit card", "Net charge-offs", "Industry", 1.5),       # 5
    ("Credit card", "Delinquency", "Industry", 2.5),           # 6
    ("Cards & loans", "Net charge-offs", "Peers <10>", 3.5),   # 7
    ('Say "hi"', "Net charge-offs", "Industry", 4.5),          # 8
    ("Auto+lease", "Delinquency", "Peers <10>", 5.5),          # 9
    ("Credit card", "Net charge-offs", "Peers <10>", 6.5),     # 10
    ("Cards & loans", "Delinquency", "Industry", 7.5),         # 11
]


def spec(product="Credit card", measure="Net charge-offs", source="Industry"):
    return [
        {"column": "Product", "caption": "Product", "anchor": (1, 0, 3, 3), "selected": product,
         "name": "PickProduct"},
        {"column": "Measure", "caption": "Measure & kind", "anchor": (3, 0, 5, 3),
         "selected": measure},
        {"column": "Source", "caption": 'Source "of"', "anchor": (5, 0, 7, 3), "selected": source},
    ]


def build(path, chart=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Pick"
    for j, h in enumerate(("Product", "Measure", "Source", "Value")):
        ws.cell(HEADER_ROW, 2 + j, h)
    for i, row in enumerate(ROWS):
        for j, v in enumerate(row):
            ws.cell(HEADER_ROW + 1 + i, 2 + j, v)
    last = HEADER_ROW + len(ROWS)
    ws.add_table(Table(displayName="Picks", ref="B%d:E%d" % (HEADER_ROW, last)))
    if chart:
        c = BarChart()
        c.add_data(Reference(ws, min_col=5, min_row=HEADER_ROW, max_row=last), titles_from_data=True)
        ws.add_chart(c, "H2")
    wb.create_sheet("Other")["A1"] = "x"
    wb.save(path)
    return path


@pytest.fixture
def book(tmp_path):
    return build(tmp_path / "pick.xlsx")


def parts(path):
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist()}


def xml(path, name):
    return ET.fromstring(parts(path)[name])


def rels(path, part):
    d, f = posixpath.split(part)
    root = xml(path, posixpath.join(d, "_rels", f + ".rels"))
    out = {}
    for r in root:
        t = r.get("Target")
        t = t[1:] if t.startswith("/") else posixpath.normpath(posixpath.join(d, t))
        out[r.get("Id")] = (r.get("Type").rsplit("/", 1)[1], t)
    return out


def hidden_rows(path):
    sheet = xml(path, "xl/worksheets/sheet1.xml")
    return {int(r.get("r")) for r in sheet.iter(M + "row") if r.get("hidden") == "1"}


def test_every_part_excel_writes_is_there_and_linked(book):
    add_table_slicers(book, "Pick", "Picks", spec())
    p = parts(book)
    for n, body in p.items():
        if n.endswith(".xml") or n.endswith(".rels"):
            ET.fromstring(body)                       # every part is well-formed XML
    caches = sorted(n for n in p if n.startswith("xl/slicerCaches/"))
    assert caches == ["xl/slicerCaches/slicerCache%d.xml" % i for i in (1, 2, 3)]
    assert [n for n in p if n.startswith("xl/slicers/")] == ["xl/slicers/slicer1.xml"]

    ct = p["[Content_Types].xml"].decode()
    for n in caches:
        assert 'PartName="/%s" ContentType="application/vnd.ms-excel.slicerCache+xml"' % n in ct
    assert 'PartName="/xl/slicers/slicer1.xml" ContentType="application/vnd.ms-excel.slicer+xml"' in ct

    # the workbook names each cache, and lists them in the x15 ext by relationship
    wb_rels = rels(book, "xl/workbook.xml")
    wb = p["xl/workbook.xml"].decode()
    names = re.findall(r'<definedName name="([^"]+)">#N/A</definedName>', wb)
    assert names == ["Slicer_Product", "Slicer_Measure", "Slicer_Source"]
    ext = re.search(r'<ext uri="\{46BE6895-7355-4a93-B00E-2C351335B9C9\}".*?</ext>', wb).group(0)
    rids = re.findall(r'r:id="([^"]+)"', ext)
    assert [wb_rels[r] for r in rids] == [("slicerCache", n) for n in caches]

    # each cache: its defined name, its column, and the table it filters
    table = xml(book, "xl/tables/table1.xml")
    col_ids = {c.get("name"): c.get("id") for c in table.iter(M + "tableColumn")}
    for n, name, col in zip(caches, names, ("Product", "Measure", "Source")):
        c = xml(book, n)
        assert (c.get("name"), c.get("sourceName")) == (name, col)
        tsc = c.find(".//{http://schemas.microsoft.com/office/spreadsheetml/2010/11/main}tableSlicerCache")
        assert (tsc.get("tableId"), tsc.get("column")) == (table.get("id"), col_ids[col])

    # the sheet points at its slicers part and at one drawing holding the frames
    sh_rels = rels(book, "xl/worksheets/sheet1.xml")
    sheet = p["xl/worksheets/sheet1.xml"].decode()
    rid = re.search(r'<x14:slicerList[^>]*><x14:slicer [^>]*r:id="([^"]+)"', sheet).group(1)
    assert sh_rels[rid] == ("slicer", "xl/slicers/slicer1.xml")
    drawings = [t for k, t in sh_rels.values() if k == "drawing"]
    assert len(drawings) == 1
    assert sheet.index("<drawing") < sheet.index("<tableParts") < sheet.index("<extLst>")
    frames = re.findall(r'<sle:slicer [^>]*name="([^"]+)"', p[drawings[0]].decode())
    sl = xml(book, "xl/slicers/slicer1.xml")
    assert [(s.get("name"), s.get("cache"), s.get("caption")) for s in sl] == [
        ("PickProduct", "Slicer_Product", "Product"),
        ("Measure", "Slicer_Measure", "Measure & kind"),
        ("Source", "Slicer_Source", 'Source "of"')]
    assert frames == ["PickProduct", "Measure", "Source"]
    # placed where asked, floating
    anchors = re.findall(r'<xdr:twoCellAnchor editAs="absolute".*?<xdr:col>(\d+)</xdr:col>.*?'
                         r'<xdr:row>(\d+)</xdr:row>.*?<xdr:col>(\d+)</xdr:col>.*?<xdr:row>(\d+)</xdr:row>',
                         p[drawings[0]].decode())
    assert [tuple(map(int, a)) for a in anchors] == [(1, 0, 3, 3), (3, 0, 5, 3), (5, 0, 7, 3)]


def test_autofilter_selects_exactly_the_chosen_values(book):
    add_table_slicers(book, "Pick", "Picks", spec())
    af = xml(book, "xl/tables/table1.xml").find(M + "autoFilter")
    assert af.get("ref") == "B4:E11"
    got = {fc.get("colId"): [f.get("val") for f in fc.iter(M + "filter")]
           for fc in af.findall(M + "filterColumn")}
    assert got == {"0": ["Credit card"], "1": ["Net charge-offs"], "2": ["Industry"]}


def test_hidden_rows_are_exactly_the_ones_the_filter_hides(book):
    add_table_slicers(book, "Pick", "Picks", spec())
    assert hidden_rows(book) == {6, 7, 8, 9, 10, 11}          # only row 5 matches all three


def test_names_with_markup_characters_are_escaped_and_round_trip(book):
    add_table_slicers(book, "Pick", "Picks", spec("Cards & loans", "Net charge-offs", "Peers <10>"))
    af = xml(book, "xl/tables/table1.xml").find(M + "autoFilter")
    assert [f.get("val") for f in af.iter(M + "filter")] == ["Cards & loans", "Net charge-offs", "Peers <10>"]
    assert hidden_rows(book) == {5, 6, 8, 9, 10, 11}
    raw = parts(book)["xl/tables/table1.xml"].decode()
    assert 'val="Cards &amp; loans"' in raw and 'val="Peers &lt;10&gt;"' in raw


@pytest.mark.parametrize("product, shown", [('Say "hi"', 8), ("Auto+lease", None)])
def test_quotes_and_plus_round_trip(tmp_path, product, shown):
    book = build(tmp_path / "q.xlsx")
    measure = "Net charge-offs" if shown else "Delinquency"
    source = "Industry" if shown else "Peers <10>"
    add_table_slicers(book, "Pick", "Picks", spec(product, measure, source))
    af = xml(book, "xl/tables/table1.xml").find(M + "autoFilter")
    assert [f.get("val") for f in af.iter(M + "filter")][0] == product
    visible = set(range(5, 12)) - hidden_rows(book)
    assert visible == {shown or 9}


def test_openpyxl_still_loads_it_and_the_cells_are_untouched(book):
    add_table_slicers(book, "Pick", "Picks", spec())
    with pytest.warns(UserWarning):                  # openpyxl drops the extension it cannot read
        wb = openpyxl.load_workbook(book)
    ws = wb["Pick"]
    got = [tuple(c.value for c in r) for r in ws.iter_rows(min_row=5, max_row=11, min_col=2, max_col=5)]
    assert got == ROWS
    assert "Picks" in ws.tables
    assert ws.row_dimensions[6].hidden and not ws.row_dimensions[5].hidden


def test_slicer_frames_join_the_drawing_a_chart_already_has(tmp_path):
    book = build(tmp_path / "chart.xlsx", chart=True)
    before = rels(book, "xl/worksheets/sheet1.xml")
    add_table_slicers(book, "Pick", "Picks", spec())
    after = rels(book, "xl/worksheets/sheet1.xml")
    drawings = [t for k, t in after.values() if k == "drawing"]
    assert drawings == [t for k, t in before.values() if k == "drawing"]
    d = parts(book)[drawings[0]].decode()
    assert "<c:chart" in d and d.count("<sle:slicer ") == 3
    ids = re.findall(r'cNvPr id="(\d+)"', d)
    ids = [i for i in ids if i != "0"]                # the fallback shapes are id 0, as Excel writes
    assert len(ids) == len(set(ids)) == 4
    assert parts(book)["xl/worksheets/sheet1.xml"].decode().count("<drawing") == 1


def test_inline_strings_are_read_too():
    sheet = (b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>'
             b'<row r="2"><c r="B2" t="inlineStr"><is><t>A &amp; B</t></is></c></row>'
             b'<row r="3"><c r="B3" t="str"><f>"x"</f><v>x</v></c></row>'
             b'<row r="4"><c r="B4" t="s"><v>0</v></c></row>'
             b'</sheetData></worksheet>')
    assert S._column_text(sheet, ["shared"], 2, 2, 5, "P") == {2: "A & B", 3: "x", 4: "shared", 5: None}


def test_refuses_what_it_cannot_do_faithfully(book):
    with pytest.raises(ValueError, match="no value 'Mortgage'"):
        add_table_slicers(book, "Pick", "Picks", spec(product="Mortgage"))
    bad = spec()
    bad[0]["column"] = "Value"
    bad[0]["selected"] = "1.5"
    with pytest.raises(ValueError, match="must hold text"):
        add_table_slicers(book, "Pick", "Picks", bad)
    with pytest.raises(ValueError, match="no column"):
        add_table_slicers(book, "Pick", "Picks", [dict(spec()[0], column="Nope")])
    with pytest.raises(ValueError, match="no table"):
        add_table_slicers(book, "Pick", "Nope", spec())
    with pytest.raises(ValueError, match="anchor"):
        add_table_slicers(book, "Pick", "Picks", [dict(spec()[0], anchor=(3, 3, 1, 1))])
    untouched = parts(book)
    add_table_slicers(book, "Pick", "Picks", spec())
    assert parts(book) != untouched
    with pytest.raises(ValueError, match="already has slicers"):
        add_table_slicers(book, "Pick", "Picks", spec())


def test_a_refusal_leaves_the_file_as_it_was(book):
    before = parts(book)
    with pytest.raises(ValueError):
        add_table_slicers(book, "Pick", "Picks", spec(source="Nowhere"))
    assert parts(book) == before


def test_button_columns_style_and_height_are_written_as_excel_writes_them(book):
    """Excel 16, asked for 3 columns, SlicerStyleOther2 and 0.25 in buttons
    (Slicer.NumberOfColumns, .Style, .RowHeight = 18 pt), saved:
    <slicer name=.. cache=.. caption=.. columnCount="3" style="SlicerStyleOther2" rowHeight="228600"/>."""
    s = spec()
    for x, n in zip(s, (3, 4, 1)):
        x.update(columns=n, style="SlicerStyleOther2", button_height=0.25)
    add_table_slicers(book, "Pick", "Picks", s)
    sl = list(xml(book, "xl/slicers/slicer1.xml"))
    assert [(x.get("columnCount"), x.get("style"), x.get("rowHeight")) for x in sl] == [
        ("3", "SlicerStyleOther2", "228600"), ("4", "SlicerStyleOther2", "228600"),
        (None, "SlicerStyleOther2", "228600")]                # one column is Excel's default, not written
    raw = parts(book)["xl/slicers/slicer1.xml"].decode()
    assert re.search(r'caption="Product" columnCount="3" style="SlicerStyleOther2" rowHeight="228600"', raw)


def test_without_the_options_the_slicer_is_as_before(book):
    add_table_slicers(book, "Pick", "Picks", spec())
    sl = list(xml(book, "xl/slicers/slicer1.xml"))
    assert {(x.get("columnCount"), x.get("style"), x.get("rowHeight")) for x in sl} == {
        (None, S.DEFAULT_STYLE, str(S.ROW_HEIGHT_EMU))}
    assert 'hiddenButton' not in parts(book)["xl/tables/table1.xml"].decode()


def test_hidden_filter_buttons_cover_every_column_and_keep_the_selection(book):
    """Excel 16 with ListObject.ShowAutoFilterDropDown = False wrote a
    filterColumn with hiddenButton="1" for every column, an empty one for a
    column nothing filters."""
    add_table_slicers(book, "Pick", "Picks", spec(), hide_buttons=True)
    af = xml(book, "xl/tables/table1.xml").find(M + "autoFilter")
    fcs = af.findall(M + "filterColumn")
    assert [(fc.get("colId"), fc.get("hiddenButton")) for fc in fcs] == [(str(i), "1") for i in range(4)]
    got = {fc.get("colId"): [f.get("val") for f in fc.iter(M + "filter")] for fc in fcs}
    assert got == {"0": ["Credit card"], "1": ["Net charge-offs"], "2": ["Industry"], "3": []}
    assert hidden_rows(book) == {6, 7, 8, 9, 10, 11}


@pytest.mark.parametrize("key, value, said", [("columns", 0, "columns must be"), ("columns", 2.5, "columns must be"),
                                              ("button_height", -1, "button_height"),
                                              ("button_height", "0.25", "button_height")])
def test_bad_layout_options_are_refused(book, key, value, said):
    s = spec()
    s[0][key] = value
    before = parts(book)
    with pytest.raises(ValueError, match=said):
        add_table_slicers(book, "Pick", "Picks", s)
    assert parts(book) == before


def test_cache_names_are_valid_defined_names_and_unique(tmp_path):
    taken = {"slicer_product"}
    assert S._defined_name("Product", taken) == "Slicer_Product1"
    assert S._defined_name("Width of Moderate", taken) == "Slicer_Width_of_Moderate"
    assert S._defined_name("How rare (z)", taken) == "Slicer_How_rare__z_"
