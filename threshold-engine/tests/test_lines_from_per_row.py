"""Each row of Decide names where its lines start (``lines_from.start`` as a list).

The firm, 18 Oct 2026, question 24 of the round 2 commercial docket, agreed
with the recommendation: *"Each row chooses where its lines start; CRE
charge-off rows start on Industry"*. Question 2 (17 Oct 2026) had already
agreed that CRE net charge-offs are rated on lines set from the industry's
history; the book could only start every row the same way.

So ``start`` is either one value for every row (null or a source), as the
consumer book keeps it, or a list naming every rated row once with its own
start, null or a source that row has. Nothing is defaulted: a rated row the
list leaves out is refused by name, as is a start the row has no series for.
"""
import json
import re

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

from conftest import ROOT
from test_book import write_book
from test_decide import NAMES, SER, SETTINGS
from test_workbook import recalc
from threshold_engine import book as BK
from threshold_engine import decide_tab as DT
from threshold_engine.workbook import build

LABELS = [lbl for lbl, _ in SER]
PAIRS = [("Card", "Net"), ("Mortgage", "Net"), ("Card", "Gross")]          # Decide's rows, in order
CRE = {"Owner-occupied CRE", "Investor CRE (non-owner-occupied)", "Multifamily", "Construction"}


def spec(lines_from):
    return {"names": NAMES, "parts": {lbl: tuple(lbl.split(" - ")) for lbl in LABELS}, "absent": [],
            "compare": None, "notes": [], "lines_from": lines_from, "kind": "consumer", "shown_not_rated": {}}


def row(p, m, start):
    return {"Product": p, "Measure": m, "start": start}


def lines_from_column(info):
    """Decide's Lines set from, row by row, as written."""
    d = info["decide"]
    ws = load_workbook(info["path"])["Decide"]
    c = L(d["columns"][DT.LINES_FROM])
    return [ws["%s%d" % (c, r)].value for r in range(d["first"], d["last"] + 1)]


# ---------------------------------------------------------------- the book

def test_each_rows_start_is_kept(tmp_path):
    bk = BK.load(write_book(tmp_path, lines_from={"part": "Source", "start": [
        row("Card", "Net", "Bank"), row("Mortgage", "Net", None)]}))
    assert bk["decide"]["lines_from"] == {("Card", "Net"): "Bank", ("Mortgage", "Net"): None}


@pytest.mark.parametrize("entries, said", [
    ([row("Card", "Net", "Bank")], "names no start for Mortgage - Net"),
    ([], "names no start for Card - Net, Mortgage - Net"),
    ([{"Product": "Card", "Measure": "Net"}, row("Mortgage", "Net", None)], "lines_from start 1 is missing 'start'"),
    ([dict(row("Card", "Net", None), Source="Bank"), row("Mortgage", "Net", None)], "'Source', which is not a key"),
    ([row("Card", "Net", "Bank"), row("Card", "Net", None), row("Mortgage", "Net", None)],
     "lines_from start 2 names Card - Net a second time"),
    ([row("Card", "Net", None), row("Mortgage", "Net", None), row("Card", "Gross", None)],
     "names Card - Gross, which is no product and measure of the book's series"),
    ([row("Card", "Net", None), row("Mortgage", "Net", "Bank")],
     "Mortgage - Net starts as 'Bank', which is not one of its sources"),
    ([row("Card", "Net", None), row("Mortgage", "Net", 0)],
     "Mortgage - Net starts as 0, which is not one of its sources"),
])
def test_each_entry_is_checked_and_refused_by_name(tmp_path, entries, said):
    with pytest.raises(BK.BookError, match=said):
        BK.load(write_book(tmp_path, lines_from={"part": "Source", "start": entries}))


@pytest.mark.parametrize("bad", [["Industry"], {"x": 1}])
def test_a_start_that_is_a_list_or_an_object_is_refused_by_name_not_crashed(tmp_path, bad):
    """A start must be a source's name or null. A list or an object used to
    reach ``st in set`` and crash with a TypeError (review, 19 Oct 2026); it
    is refused naming the entry, in both forms of ``start``."""
    with pytest.raises(BK.BookError, match=r"lines_from start 1: Card - Net starts as %s, which is not one of its "
                       "sources" % re.escape(repr(bad))):
        BK.load(write_book(tmp_path, lines_from={"part": "Source", "start": [
            row("Card", "Net", bad), row("Mortgage", "Net", None)]}))
    if isinstance(bad, dict):                   # a list is the per-row form, tried above
        with pytest.raises(BK.BookError, match=r"lines_from: start must be null .* not %s" % re.escape(repr(bad))):
            BK.load(write_book(tmp_path, lines_from={"part": "Source", "start": bad}))


def test_a_row_shown_and_not_rated_has_no_start(tmp_path):
    """Question 23's rows have no row on Decide, so no Lines set from: naming
    one is refused, and leaving it out is right."""
    unrated = [{"Product": "Mortgage", "Measure": "Net", "why": "Shown only (a test)."}]
    with pytest.raises(BK.BookError, match="names Mortgage - Net, which the book shows and does not rate"):
        BK.load(write_book(tmp_path, shown_not_rated=unrated, lines_from={"part": "Source", "start": [
            row("Card", "Net", None), row("Mortgage", "Net", None)]}))
    bk = BK.load(write_book(tmp_path, shown_not_rated=unrated, lines_from={"part": "Source", "start": [
        row("Card", "Net", "Industry")]}))
    assert bk["decide"]["lines_from"] == {("Card", "Net"): "Industry"}


def test_the_commercial_book_starts_its_cre_charge_off_rows_on_industry():
    """Question 24: the four CRE net charge-off rows start on Industry, every
    other rated row blank; Nonfarm nonresidential CRE, shown and not rated,
    is not named."""
    bk = BK.load(ROOT / "books" / "commercial-public.json")
    starts = bk["decide"]["lines_from"]
    assert len(starts) == 10
    assert {pr for pr, s in starts.items() if s is not None} == {(p, "Net charge-offs") for p in CRE}
    assert {s for s in starts.values()} == {"Industry", None}
    assert not any(p == "Nonfarm nonresidential CRE" for p, _ in starts)
    raw = json.loads((ROOT / "books" / "commercial-public.json").read_text(encoding="utf-8"))
    assert any("question 24" in n and "start on Industry" in n for n in raw["notes"])


def test_the_consumer_book_keeps_one_start_for_every_row():
    """The consumer book is not changed: one start, null, for every row."""
    raw = json.loads((ROOT / "books" / "consumer-public.json").read_text(encoding="utf-8"))
    assert raw["lines_from"] == {"part": "Source", "start": None}
    assert BK.load(ROOT / "books" / "consumer-public.json")["decide"]["lines_from"] is None


# ---------------------------------------------------------------- the workbook

@pytest.fixture(scope="module")
def per_row(tmp_path_factory):
    d = tmp_path_factory.mktemp("per_row")
    info = build(d / "p.xlsx", SER, decide=spec({("Card", "Net"): "Bank", ("Mortgage", "Net"): None,
                                                ("Card", "Gross"): "Industry"}), **SETTINGS)
    return info, recalc(info["path"])


def test_each_row_of_decide_starts_on_its_own_source(per_row):
    info, _ = per_row
    assert info["decide"]["pairs"] == PAIRS
    assert lines_from_column(info) == ["Bank", None, "Industry"]


def test_thresholds_and_summary_follow_each_rows_start(per_row):
    """The lines each row is rated on are its own start's: Thresholds reads
    each row's Lines set from, and Summary counts the one row with none."""
    info, v = per_row
    t = info["thresholds"]
    from threshold_engine import thresholds_tab as TT
    got = [v.get("THRESHOLDS!%s%d" % (TT.FIX["lines_from"], r)) for r in range(t["first"], t["last"] + 1)]
    assert got == ["Bank", "", "Industry"]
    s = info["summary"]
    head = v["SUMMARY!" + s["count"]]
    assert head.endswith("1 with no source picked"), head


@pytest.mark.parametrize("start", [None, "Industry"])
def test_one_start_for_every_row_is_the_same_as_naming_it_on_each(tmp_path, start):
    """The book-wide form, which the consumer book uses, is unchanged: a
    workbook built with one start is cell for cell the one built with that
    start named on every row."""
    a = build(tmp_path / "a.xlsx", SER, decide=spec(start), **SETTINGS)
    b = build(tmp_path / "b.xlsx", SER, decide=spec({pr: start for pr in PAIRS}), **SETTINGS)
    wa, wb_ = load_workbook(a["path"]), load_workbook(b["path"])
    assert wa.sheetnames == wb_.sheetnames
    def cells(ws, own):
        """Every cell's value; the Run tab's time of the build and the file's own name aside."""
        out = {}
        for r_ in ws.iter_rows():
            for c in r_:
                v = c.value
                if isinstance(v, str):
                    v = re.sub(r"^\d{4}-\d\d-\d\dT\d\d:\d\d:\d\d$", "<time>", v).replace(own, "<out>")
                if v is not None:
                    out[c.coordinate] = v
        return out
    for name in wa.sheetnames:
        assert cells(wa[name], "a.xlsx") == cells(wb_[name], "b.xlsx"), name
    assert lines_from_column(a) == [start] * 3


@pytest.mark.parametrize("starts, said", [
    ({("Card", "Net"): "Bank", ("Mortgage", "Net"): None}, "gives no start for Card - Gross"),
    ({("Card", "Net"): None, ("Mortgage", "Net"): None, ("Card", "Gross"): None, ("Auto", "Net"): None},
     "names Auto - Net, which no row of Decide is"),
    ({("Card", "Net"): None, ("Mortgage", "Net"): "Bank", ("Card", "Gross"): None},
     r"no series for: Mortgage - Net \('Bank'\)"),
])
def test_the_workbook_refuses_a_start_it_cannot_place(tmp_path, starts, said):
    with pytest.raises(ValueError, match=said):
        build(tmp_path / "x.xlsx", SER, decide=spec(starts), **SETTINGS)


@pytest.mark.parametrize("bad", [["Bank"], {"x": 1}])
def test_check_start_refuses_a_start_that_is_not_a_name_rather_than_crash(bad):
    """Review, 19 Oct 2026: an unhashable start crashed ``check_start`` with a
    TypeError; it is refused naming the row, per row and as one start."""
    parts = {lbl: tuple(lbl.split(" - ")) for lbl in LABELS}
    starts = {pr: None for pr in PAIRS}
    starts[("Card", "Net")] = bad
    with pytest.raises(ValueError, match=r"no series for: Card - Net \(%s\)" % re.escape(repr(bad))):
        DT.check_start(starts, LABELS, parts)
    one = bad if isinstance(bad, list) else [bad]          # one start for every row; a dict is the per-row form
    with pytest.raises(ValueError, match="lines_from starts as %s, which is not a source" % re.escape(repr(one))):
        DT.check_start(one, LABELS, parts)


def test_a_book_with_each_rows_start_builds_with_it(tmp_path):
    """Through the book command, as the commercial book is built."""
    out = tmp_path / "out.xlsx"
    info = BK.build(write_book(tmp_path, lines_from={"part": "Source", "start": [
        row("Card", "Net", "Bank"), row("Mortgage", "Net", None)]}), out, say=lambda s: None)
    assert lines_from_column(dict(info, path=out)) == ["Bank", None]
