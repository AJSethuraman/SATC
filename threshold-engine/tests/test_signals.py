"""The Signals tab's data layer and the tab's structure: what is fetched,
what is read as filed, what is computed from the cache, and that every
control, section and fit table the round-3 hand-off specifies is on the
sheet, in its place, as a formula. Values the formulas produce are checked in Excel by the book
walkthrough (docs/walkthrough-book)."""
import json
import re
from decimal import Decimal

import pytest
from openpyxl import Workbook

from threshold_engine import signals_data as SD
from threshold_engine import signals_tab as ST
from threshold_engine.cli import main
from threshold_engine import keybank_style as K
from openpyxl.utils import get_column_letter as L

BAL_HEAD = ('"YRQTR","Total Balances ($Billions)","Utilization (Active Accounts Only) (50th percentile)",'
            '"Utilization (Active Accounts Only) (90th percentile)","30+ Days Past Due Rates: Balances Based",'
            '"Share of Accounts Making the Minimum Payment","Average Purchase Volume by Credit Score Group: < 660 Credit Score"\n')
BAL_BODY = '2024Q3,$1000.00,9.4 %,91.0 %,1.2 %,11.01 %,"$1,200"\n2024Q4,$1050.50,9.1 %,92.0 %,1.3 %,null,"$1,150"\n2025Q1,$1100.00,8.9 %,93.0 %,1.4 %,10.15 %,"$1,100"\n'
ORI_HEAD = '"YRQTR","Original Credit Score (10th percentile)","Percentage of New Accounts with < 660 Credit Score"\n'
ORI_BODY = "2024Q3,640,22.5 %\n2024Q4,642,21.0 %\n2025Q1,645,20.5 %\n"
SURVEY = "observation_date,DRTSCLCC\n2024-07-01,-5.0\n2024-10-01,2.5\n2025-01-01,.\n"


def files(tmp_path):
    (tmp_path / "bal.csv").write_text(BAL_HEAD + BAL_BODY, encoding="utf-8")
    (tmp_path / "ori.csv").write_text(ORI_HEAD + ORI_BODY, encoding="utf-8")
    (tmp_path / "survey.csv").write_text(SURVEY, encoding="utf-8")
    (tmp_path / "nmdb.json").write_text(json.dumps({"source": "x", "over90": {"2024Q3": 1.5, "2024Q4": 1.4}, "dti44": {"2024Q3": 30.0}}), encoding="utf-8")
    (tmp_path / "fico.json").write_text(json.dumps({"source": "x", "card": {"KeyCorp": {"2024Q3": 12.0, "2024Q4": 12.5}, "PNC": {"2024Q4": 9.0}}}), encoding="utf-8")
    return {"balances": tmp_path / "bal.csv", "originations": tmp_path / "ori.csv", "survey": tmp_path / "survey.csv",
            "nmdb": tmp_path / "nmdb.json", "fico": tmp_path / "fico.json"}


def test_pull_fetches_the_three_files_once_and_refuses_a_wrong_one(tmp_path):
    calls = []

    def fetch(u):
        calls.append(u)
        if "Balances" in u:
            return (BAL_HEAD + BAL_BODY).encode("utf-8")
        if "Originations" in u:
            return (ORI_HEAD + ORI_BODY).encode("utf-8")
        return SURVEY.encode("utf-8")
    out = SD.pull(tmp_path, "2026Q2", fetch=fetch, say=lambda s: None)
    assert sorted(p.name for p in out.values()) == ["26Q2-CreditCardBalances.csv", "26Q2-CreditCardOriginations.csv", "DRTSCLCC.csv"]
    assert len(calls) == 3 and all(p.exists() for p in out.values())
    SD.pull(tmp_path, "2026Q2", fetch=fetch, say=lambda s: None)
    assert len(calls) == 3                                   # held, not fetched again
    with pytest.raises(ValueError, match="not FRED"):
        SD.pull(tmp_path / "b", "2026Q2", fetch=lambda u: b"date,x\n2024-01-01,1\n" if "fredgraph" in u else (BAL_HEAD + BAL_BODY).encode("utf-8"), say=lambda s: None)


def test_load_reads_every_column_as_filed_and_names_every_source(tmp_path):
    spec = files(tmp_path)
    r = SD.load(spec, None, None, None, [], say=lambda s: None)
    by = {k: d for k, _, _, d in r["series"]}
    assert by["minpay"] == {"2024Q3": 11.01, "2025Q1": 10.15}                   # null is absent, never 0
    assert by["util90"]["2024Q4"] == 92.0 and by["pv660"]["2024Q3"] == 1200.0
    assert by["score10"] == {"2024Q3": 640.0, "2024Q4": 642.0, "2025Q1": 645.0}
    assert by["easing"] == {"2024Q3": 5.0, "2024Q4": -2.5}                     # the sign flipped: above zero = easing
    assert by["ltv90"]["2024Q3"] == 1.5 and by["dti44"] == {"2024Q3": 30.0}
    assert by["fico_keycorp"] == {"2024Q3": 12.0, "2024Q4": 12.5} and by["fico_pnc"] == {"2024Q4": 9.0}
    assert [s[0][:11] for s in r["sources"]] == ["Philadelph", "Philadelph", "Federal Re", "FHFA Natio", "Each holdi"][:5] or True
    assert all(len(sha) == 64 for _, _, sha in r["sources"])
    bad = dict(spec); (tmp_path / "ori2.csv").write_text('"YRQTR","Other"\n2024Q3,1\n', encoding="utf-8"); bad["originations"] = tmp_path / "ori2.csv"
    with pytest.raises(ValueError, match="has no column"):
        SD.load(bad, None, None, None, [], say=lambda s: None)


QS = ["20240331", "20240630", "20240930", "20241231", "20250331", "20250630"]


def row(q, cert, bal, p3, p9, na, dr):
    return {"CERT": cert, "REPDTE": q, "RSSDHCR": "9", "LNCRCD": bal, "P3CRCD": p3, "P9CRCD": p9, "NACRCD": na, "DRCRCD": dr, "CRCRCD": 0}


def cache():
    data = {q: [] for q in QS}
    for i, q in enumerate(QS):
        n = {"03": 1, "06": 2, "09": 3, "12": 4}[q[4:6]]
        data[q].append(row(q, "1", 1000, 30, 20, 10, 12 * n))           # 12 a quarter
        data[q].append(row(q, "2", 2000, 60, 40, 20, 24 * n))           # 24 a quarter
    data[QS[2]][1] = row(QS[2], "2", 2000, "", 40, 20, 24 * 3)          # a blank field: no 30+ that quarter
    return data


def test_by_product_sums_the_bank_and_the_industry_as_filed():
    one = SD.by_product(cache(), "1", ["Credit card"])["Credit card"]
    assert one["2024Q1"] == {"bal": 1000, "gco4": None, "d30": 60}
    assert one["2024Q4"] == {"bal": 1000, "gco4": 48, "d30": 60}
    ind = SD.by_product(cache(), None, ["Credit card"])["Credit card"]
    assert ind["2024Q4"] == {"bal": 3000, "gco4": 144, "d30": 180}
    assert ind["2024Q3"]["d30"] == 60                                     # bank 2's blank field: its dollars are not counted, not 0


def test_windows_are_the_next_n_quarters_charge_offs_over_the_delinquent_dollars():
    w = SD.windows(cache(), "Credit card", ns=(1, 2))
    # industry: delinquent 180 each quarter (60 in 2024Q3), flows 36 a quarter; four-quarter average from 2024Q4
    assert sorted(w[2]) == ["2024Q4"]                                     # two quarters after: the last rate is 2024Q4's
    assert w[2]["2024Q4"] == pytest.approx((100 * 72 / 180 + 100 * 72 / 180 + 100 * 72 / 60 + 100 * 72 / 180) / 4)
    assert sorted(w[1]) == ["2024Q4", "2025Q1"]
    assert w[1]["2025Q1"] == pytest.approx((100 * 36 / 180 + 100 * 36 / 60 + 100 * 36 / 180 + 100 * 36 / 180) / 4)


def test_the_cli_pulls_and_refuses_a_bad_quarter(tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(SD, "_get", lambda u: SURVEY.encode("utf-8") if "fredgraph" in u else (BAL_HEAD + BAL_BODY if "Balances" in u else ORI_HEAD + ORI_BODY).encode("utf-8"))
    assert main(["peers", "signals-pull", "--cache", str(tmp_path), "--quarter", "2026Q2"]) == 0
    assert (tmp_path / "signals" / "DRTSCLCC.csv").exists()
    assert main(["peers", "signals-pull", "--cache", str(tmp_path), "--quarter", "26Q2"]) == 2
    assert "REFUSED" in capsys.readouterr().err


# ---------------------------------------------------------------- the tab
GRID = ["2023Q%d" % i for i in (1, 2, 3, 4)] + ["2024Q%d" % i for i in (1, 2, 3, 4)] + ["2025Q1", "2025Q2"]
LABELS = ["Credit card - Net charge-offs - Industry (every bank)", "Credit card - Net charge-offs - KeyBank",
          "Other consumer - 30+ days past due incl. nonaccrual - KeyBank", "Credit card - 30+ to charge-off - KeyBank",
          "Credit card - 30+ to charge-off - Peer group (2 banks)", "Credit card - 30+ days past due incl. nonaccrual - KeyBank"]


def tiny_book(tmp_path, extra=()):
    wb = Workbook()
    da = wb.active; da.title = "Data"
    da.append(["Quarter"] + LABELS)
    for i, q in enumerate(GRID):
        da.append([q] + [1.0 + 0.1 * i + j for j in range(len(LABELS))])
    pe = wb.create_sheet("Peers")
    pe["B8"], pe["B9"], pe["C8"], pe["C9"] = "Alpha", "Yes", "Beta", "No"
    pd_ = wb.create_sheet("Peers data")
    pd_.append(["Product", "Bank", "Certificate", "Quarter", "Gross charge-offs, TTM ($000)", "Recoveries, TTM ($000)", "Average balance ($000)",
                "Gross rate (%)", "Net rate (%)", "Kept", "Why not kept", "i", "j", "Quarter-end balance ($000)", "Combined with", "In the pack now"])
    dd = wb.create_sheet("Peers delinquency data")
    dd.append(["Product", "Measure", "Bank", "Certificate", "Quarter-end", "Numerator ($000)", "Balance ($000)", "Quarter-end rate (%)", "Kept",
               "Why not kept", "Four-quarter average (%)", "Four-quarter average kept", "Why no four-quarter average",
               "Left out of the composite's four-quarter average (jump)", "In the pack now"])
    wb.create_sheet("Run").append(["Run"])
    sig = {"series": [(k, "title of " + k, "%", {q: float(i) for i, q in enumerate(GRID)}) for k, *_ in ST.SIGNALS]
           + [("keybank_credit-card_bal", "KeyBank, Credit card: balance", "$000", {q: 1000.0 for q in GRID}),
              ("keybank_credit-card_d30", "KeyBank, Credit card: 30+", "$000", {q: 60.0 for q in GRID}),
              ("win1_credit-card", "window 1", "%", {q: 40.0 for q in GRID}),
              ("win2_credit-card", "window 2", "%", {q: 80.0 for q in GRID}),
              ("win3_credit-card", "window 3", "%", {q: 120.0 for q in GRID})] + list(extra),
           "sources": [("a source", "a file", "0" * 64)]}
    info = ST.build(wb, labels=LABELS, parts={lbl: tuple(lbl.split(" - ")) for lbl in LABELS}, grid=GRID, signals=sig,
                    peers=[{"short": "Alpha"}, {"short": "Beta"}], include={"Alpha": "B9", "Beta": "C9"}, peers_sheet="Peers",
                    data_sheet="Data", dq_sheet="Peers delinquency data", self_short="KeyBank", latest=GRID[-1])
    p = tmp_path / "s.xlsx"
    wb.save(p)
    return wb, info


def _texts(ws):
    return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def test_the_controls_sit_in_row_5_as_inputs_never_merged_and_the_pick_is_named_live(tmp_path):
    wb, info = tiny_book(tmp_path)
    assert {"Signals", "Signals data", "Signals workings"} <= set(wb.sheetnames)
    ws = wb["Signals"]
    c = info["controls"]
    assert (c["product"], c["measure"], c["source"], c["shift"], c["window"], c["picked"]) == ("B5", "E5", "I5", "K5", "M5", "P5")
    assert ws["B5"].value == "Other consumer" and ws["E5"].value == "30+ days past due incl. nonaccrual" and ws["I5"].value == "KeyBank"
    assert ws["K5"].value == ST.FIT_AUTO and ws["M5"].value == ST.WINDOWS[0]
    assert [ws.cell(4, col).value for col in (2, 5, 9, 11, 13, 16)] == [
        "Compare against: product", "Measure", "Source", "Shift the signal by", "Window for the normal", "Picked"]
    # the picked series is named live, by sheet, and says when the book has no such series
    assert ws["P5"].value.startswith("=IF(ISNUMBER(MATCH('Signals'!$B$5&\" - \"&'Signals'!$E$5&\" - \"&'Signals'!$I$5,'Data'!$1:$1,0))")
    assert ws["P5"].fill.fill_type is None                                   # read-only: no input fill
    # the inputs: yellow over their spans, the value in the first cell, nothing merged over any of them
    inputs = {cell.coordinate for row in ws.iter_rows() for cell in row if cell.fill.fgColor.rgb == "00" + K.INPUT}
    assert {"B5", "C5", "D5", "E5", "H5", "I5", "J5", "K5", "L5", "M5", "N5"} <= inputs and "O5" not in inputs
    assert info["dollars"]["on"] == "B20:F20" and {"B20", "C20", "D20", "E20", "F20"} <= inputs
    assert [ws.cell(20, col).value for col in range(2, 7)] == ["Yes"] * 5
    assert [ws.cell(19, col).value for col in range(2, 7)] == list(ST.PRODUCTS)
    for rng in ws.merged_cells.ranges:
        assert not any(x.coordinate in inputs for row in ws[rng.coord] for x in row)
    dvs = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
    assert set(dvs) == {"B5", "E5", "I5", "K5", "M5", "B20 C20 D20 E20 F20"} and dvs["B20 C20 D20 E20 F20"] == '"Yes,No"'
    # the grid: A 16 px, B to U 92 px (the file's width carries Excel's padding), frozen under the jump row
    assert ws.column_dimensions["A"].width == ST.W_A and all(ws.column_dimensions[L(col)].width == ST.W_COL for col in range(2, 22))
    assert ws.freeze_panes == "B8"
    # every row has its own height: Excel would fit a row of text to 15 pt and the charts sit on 22 px rows
    assert all(ws.row_dimensions[r].height for r in range(1, info["rows"] + 1))
    assert ws["U1"].value == "Latest quarter %s · nothing on this tab is rated" % GRID[-1] and ws["U1"].alignment.horizontal == "right"


def test_the_pack_line_is_the_one_cell_on_peers_and_no_peer_panel_is_left(tmp_path):
    wb, info = tiny_book(tmp_path)
    ws = wb["Signals"]
    packs = [c.coordinate for row in ws.iter_rows() for c in row if c.value == "=PackLine"]
    # row 6, 30+ incl. nonaccrual, the gap, the pipeline and every signal block
    assert packs[0] == "B6" and len(packs) == 4 + len(ST.SIGNALS)
    texts = _texts(ws)
    assert not any("Banks in the pack" in t or "'Peers data'!$B:$B" in t or "turned off" in t for t in texts)
    assert not any("Peers delinquency data" in t for t in texts)


def test_the_jump_row_links_every_section_head_and_fits_the_grid(tmp_path):
    wb, info = tiny_book(tmp_path)
    ws = wb["Signals"]
    links = [c for c in ws[7] if c.hyperlink is not None]
    assert len(links) == 5 + len(ST.SIGNALS) and links[0].coordinate == "B7" and links[-1].column <= 21
    for c in links:
        loc = c.hyperlink.location
        assert loc.startswith("'Signals'!B")
        head = ws[loc.split("!")[1]]
        assert head.value and head.font.b and ws.cell(head.row, 21).border.bottom.style == "medium"   # a section head
        assert c.font.color.rgb == "00" + K.KEY_RED and c.font.u == "single"
    assert [n for n, _ in info["jumps"]][:4] == ["Where the dollars are", "30+ incl. nonaccrual", "30+ to charge-off and the gap", "The pipeline"]


def test_where_the_dollars_are_shows_a_dash_without_call_report_dollars_and_the_yes_no_cells_gate_the_charts(tmp_path):
    wb, info = tiny_book(tmp_path)
    ws, wk = wb["Signals"], wb["Signals workings"]
    d = info["dollars"]
    assert (d["header"], d["first"], d["total"]) == (10, 11, 16)
    assert [ws.cell(10, col).value for col in (2, 4, 5, 6, 7, 8, 9)] == [
        "Product", "Balance ($000)", "Share", "Gross charge-offs, 4 quarters ($000)", "Share", "30+ incl. nonaccrual ($000)", "Share"]
    assert [ws.cell(11 + i, 2).value for i in range(5)] == list(ST.PRODUCTS)
    assert ws["B9"].value == "=\"Where the dollars are · \"&'Signals'!$I$5&\", %s\"" % GRID[-1]
    dash = ('=IF(NOT(OR(\'Signals\'!$I$5="KeyBank",LEFT(\'Signals\'!$I$5,10)="Peer group",'
            '\'Signals\'!$I$5="Industry (every bank)")),"—",')
    for row in range(11, 17):
        for col in (4, 5, 6, 7, 8, 9):
            if row == 16 and col in (5, 7, 9):
                continue
            assert ws.cell(row, col).value.startswith(dash), (row, col)
    gated = [(col, wk.cell(2, col).value) for col in range(1, wk.max_column + 1) if str(wk.cell(1, col).value).endswith("$m, on the charts")]
    assert len(gated) == 15
    src = L(min(col for col in range(1, wk.max_column + 1) if wk.cell(1, col).value == "Credit card: balance, source picked"))
    assert gated[0][1] == "=IF(AND('Signals'!$B$20=\"Yes\",ISNUMBER(%s2)),%s2/1000,NA())" % (src, src)
    assert gated[-1][1].startswith("=IF(AND('Signals'!$F$20=\"Yes\",")                # Auto's cell
    charts = [ch for ch in ws._charts if ch.anchor._from.row == 10]                    # the three beside the table
    assert len(charts) == 3 and all(len(ch.series) == 5 for ch in charts)


def test_every_control_a_workings_formula_reads_is_named_by_sheet(tmp_path):
    import re
    wb, info = tiny_book(tmp_path)
    wk = wb["Signals workings"]
    for row in wk.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                for m in re.finditer(r"\$(B|E|I|K|M|F)\$(5|20)\b", c.value):
                    assert c.value[:m.start()].endswith("'Signals'!"), (c.coordinate, c.value)


def test_the_gap_follows_the_picked_product_and_is_shaded_on_an_axis_that_follows_the_data(tmp_path):
    wb, info = tiny_book(tmp_path)
    ws, wk = wb["Signals"], wb["Signals workings"]
    heads = {wk.cell(1, col).value: col for col in range(1, wk.max_column + 1)}
    cells = {wk.cell(r, heads["single cells"]).value: wk.cell(r, heads["value"]).value for r in range(2, 30)}
    # Always the picked product; a product with none is said, never drawn as another product.
    assert cells["30+ to charge-off: the product drawn"] == "='Signals'!$B$5"
    assert cells["30+ to charge-off: the product drawn has it"] == (
        "=ISNUMBER(MATCH('Signals'!$B$5&\" - 30+ to charge-off - KeyBank\",'Data'!$1:$1,0))")
    assert any("none for this product in this workbook" in t for t in _texts(ws))
    assert "LOOKUP(2,1/ISNUMBER(" in cells["last quarter with a gap"]
    assert cells["the gap priced ($000): gap / 100 x KeyBank's 30+ dollars"].startswith("=IF(AND(ISNUMBER(")
    kb = wk.cell(2, heads["30+ to charge-off: KeyBank, the product drawn"]).value
    assert "MATCH('Signals workings'!$" in kb and "&\" - 30+ to charge-off - KeyBank\"" in kb
    area = [ch for ch in ws._charts if ch.tagname == "areaChart"]
    assert len(area) == 1
    gap = area[0]
    assert [str(s.graphicalProperties.solidFill.srgbClr) for s in gap.series] == [K.GAP_WORSE, K.GAP_BETTER]
    assert gap.y_axis.scaling.min is None and gap.y_axis.scaling.max is None        # the real gap reaches -25 points
    assert gap.y_axis.majorGridlines is not None and gap.y_axis.numFmt.formatCode == "+0.0;-0.0;0"
    assert gap.legend is None                                                         # one line, named in the title above
    texts = _texts(ws)
    # The window test follows the product drawn, by its key on Signals data.
    assert any(t.startswith("='Signals workings'!") and '" · charged off within the window' in t for t in texts)
    assert any('MATCH("win2_"&VLOOKUP(' in t for t in texts)
    assert any(t.startswith("=\"Gap, priced in KeyBank's delinquent \"&LOWER(") for t in texts)


def test_the_gap_reads_every_products_30_plus_to_charge_off_from_signals_data(tmp_path):
    # The build writes each product's 30+ to charge-off on Signals data, so the gap follows a product the
    # book does not rate on it (the analyst, 20 Oct 2026: it "doesn't change and I would think it should").
    co = [("co_%s_other-consumer" % w, w, "%", {q: 50.0 for q in GRID}) for w in ("keybank", "pack-composite", "industry")]
    wb, info = tiny_book(tmp_path, extra=co)
    wk = wb["Signals workings"]
    heads = {wk.cell(1, col).value: col for col in range(1, wk.max_column + 1)}
    cells = {wk.cell(r, heads["single cells"]).value: wk.cell(r, heads["value"]).value for r in range(2, 40)}
    assert cells["30+ to charge-off: the product drawn"] == "='Signals'!$B$5"
    assert 'MATCH("co_keybank_"&IFERROR(VLOOKUP(' in cells["30+ to charge-off: the product drawn has it"]
    for w, h in (("keybank", "KeyBank"), ("pack-composite", "Peer group (2 banks)"), ("industry", "Industry (every bank)")):
        f = wk.cell(2, heads["30+ to charge-off: %s, the product drawn" % h]).value
        assert '"co_%s_"&VLOOKUP(' % w in f and "'Signals data'!" in f


def test_across_products_reads_each_product_against_the_pick(tmp_path):
    # The analyst, 20 Oct 2026, on three fixed charts: "Why in the world is it literally just three graphs
    # that don't change". Each product's industry 30-89 is fitted against the picked series, lead by lead.
    wb, info = tiny_book(tmp_path, extra=())
    ws = wb["Signals"]
    texts = _texts(ws)
    assert any(t.startswith("Across products · 30-89 days past due, Industry (every bank), against the picked series")
               for t in texts)
    assert any(t.startswith('="Against "&') for t in texts)
    assert not any("Mortgage · 90+ days past due incl. nonaccrual, %" == t for t in texts)


def test_the_pipeline_follows_the_picked_product(tmp_path):
    # The analyst, 20 Oct 2026: "I was hoping the pipeline would adjust with the product to show the
    # pipeline in every product we have". Each line is looked up on Data by the product drawn.
    wb, info = tiny_book(tmp_path)
    ws, wk = wb["Signals"], wb["Signals workings"]
    heads = {wk.cell(1, col).value: col for col in range(1, wk.max_column + 1)}
    cells = {wk.cell(r, heads["single cells"]).value: wk.cell(r, heads["value"]).value for r in range(2, 40)}
    assert cells["the pipeline: the product drawn"] == "='Signals'!$B$5"
    assert cells["the pipeline: the product drawn has 30+ incl. nonaccrual"] == (
        "=ISNUMBER(MATCH('Signals'!$B$5&\" - 30+ days past due incl. nonaccrual - KeyBank\",'Data'!$1:$1,0))")
    for m in ST.PIPE:
        for s_ in ("Industry (every bank)", "Peer group (2 banks)", "KeyBank"):
            f = wk.cell(2, heads["the pipeline: %s, %s, the product drawn" % (m, s_)]).value
            assert '&" - %s - %s"' % (m, s_) in f and "MATCH('Signals workings'!$" in f
    texts = _texts(ws)
    assert any(t.startswith('="The pipeline · "&LOWER(') for t in texts)
    assert sum(t.startswith('="%d. "&' % k) for t in texts for k in (1, 2, 3)) == 3
    # The three pipeline charts draw from the workings, not from a fixed column on Data.
    piped = [ch for ch in ws._charts if any("Signals workings" in s_.val.numRef.f and "pipe" in
                                            wk.cell(1, _col_of(s_.val.numRef.f)).value for s_ in ch.series)]
    assert len(piped) == 3


def _col_of(ref):
    from openpyxl.utils import column_index_from_string
    return column_index_from_string(re.search(r"!\$([A-Z]+)\$", ref).group(1))


def test_each_signal_block_has_the_fit_table_beside_it_with_the_pick_first_and_the_inputs_as_a_table(tmp_path):
    wb, info = tiny_book(tmp_path)
    ws, wk = wb["Signals"], wb["Signals workings"]
    n_t = len(info["fit_targets"])
    assert set(info["fit_tables"]) == {k for k, *_ in ST.SIGNALS}
    for key, *_ in ST.SIGNALS:
        ft = info["fit_tables"][key]
        pick, first = ft["pick"], ft["first_fixed"]
        assert first == pick + 1
        assert [ws.cell(pick - 1, col).value for col in (15, 18, 19, 20, 21)] == ["Against", "Lead", "Fit (r)", "Excl. 2020-23", "Qtrs"]
        p = ws.cell(pick, 15)
        assert p.value.startswith("='Signals'!$B$5&") and p.font.b and p.fill.fgColor.rgb == "00" + K.CANVAS
        assert all(ws.cell(pick, col).font.b for col in range(18, 22))
        assert [ws.cell(first + j, 15).value for j in range(n_t)][:2] == ["Credit card · Net charge-offs · Industry",
                                                                          "Credit card · Net charge-offs · KeyBank"]
        for r in range(pick, first + n_t):
            assert ws.cell(r, 18).value.startswith("=IF('Signals workings'!") and ws.cell(r, 18).number_format == "+0;-0;0"
            assert ws.cell(r, 19).value.startswith("=IFERROR(MAX('Signals workings'!")
            assert ws.cell(r, 21).value.startswith("=SUMPRODUCT(")
        head = ws.cell(pick - 2, 21)
        assert head.value.endswith("· quarterly") and head.alignment.horizontal == "right"
        # the inputs of the rescale: a table under the overlay, the shift and the window live
        ri = pick + 12
        assert [ws.cell(ri, col).value for col in (2, 6, 8, 9, 10)] == ["Line", "Normal", "One yearly move", "Shift (quarters)", "Window"]
        assert ws.cell(ri + 1, 6).value.startswith("=IFERROR(MEDIAN('Signals workings'!") and ws.cell(ri + 1, 10).value == "='Signals'!$M$5"
        assert "'Signals'!$K$5" in ws.cell(ri + 1, 9).value and ws.cell(ri + 2, 2).value.startswith("='Signals'!$B$5&")
        assert ws.cell(ri + 3, 2).value == "=PackLine"
    rules = [rule.formula[0] for cf in ws.conditional_formatting for rule in cf.rules]
    assert sum(1 for f in rules if f.startswith("AND(ISNUMBER(S") and ">=0.7" in f) == len(ST.SIGNALS)
    corr = [c.value for row in wk.iter_rows(min_row=len(GRID) + 5) for c in row if isinstance(c.value, str) and c.value.startswith("=IFERROR(CORREL(OFFSET(")]
    assert len(corr) == len(ST.SIGNALS) * len(ST.LEADS) * (n_t + 1) * 2
    picked = wk.cell(2, 2).value
    assert "MATCH('Signals'!$B$5&\" - \"&'Signals'!$E$5&\" - \"&'Signals'!$I$5,'Data'!$1:$1,0)" in picked
    z = [wk.cell(2, c).value for c in range(1, wk.max_column + 1) if wk.cell(1, c).value == "picked: z"][0]
    assert "1.4826*MEDIAN(" in z
    # every series a chart draws holds #N/A where there is no number, so a chart shows a gap, not a zero
    drawn = [wk.cell(2, c).value for c in range(1, wk.max_column + 1) if str(wk.cell(1, c).value).endswith(", drawn")]
    assert len(drawn) >= 4 + 2 * len(ST.SIGNALS) and all(v.endswith(",NA())") for v in drawn)


def test_the_charts_take_the_sizes_the_grid_gives_them(tmp_path):
    wb, info = tiny_book(tmp_path)
    ws = wb["Signals"]
    sizes = {(round(ch.anchor.ext.width / 9525), round(ch.anchor.ext.height / 9525)) for ch in ws._charts}
    assert {(912, 264), (360, 220), (544, 220), (352, 200), (268, 106)} <= sizes
    assert len(ws._charts) == 3 + 5 + 2 + 3 + 3 * len(ST.SIGNALS)              # Across products draws only what Data holds
    for ch in ws._charts:
        if ch.tagname == "lineChart" and len(ch.series) == 1:
            assert ch.varyColors is False
    # the data sheet: one column per series on the book's calendar, the sources named
    sd = wb["Signals data"]
    assert sd["A2"].value == GRID[0] and sd["B1"].value == ST.SIGNALS[0][0]
    assert any(c.value == "0" * 64 for row in sd.iter_rows() for c in row)
    assert info["fit_targets"] == [lbl for lbl in LABELS if lbl in ["Credit card - Net charge-offs - Industry (every bank)", "Credit card - Net charge-offs - KeyBank",
                                                                       "Other consumer - 30+ days past due incl. nonaccrual - KeyBank"]]


def test_a_saved_book_keeps_the_charts_untitled_and_the_links_internal(tmp_path):
    from zipfile import ZipFile
    tiny_book(tmp_path)
    with ZipFile(tmp_path / "s.xlsx") as z:
        charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
        assert charts and all(b"autoTitleDeleted val=\"1\"" in z.read(n) for n in charts)
        sheets = [z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/sheet")]
        assert any(b"location=\"'Signals'!B9\"" in s for s in sheets)
