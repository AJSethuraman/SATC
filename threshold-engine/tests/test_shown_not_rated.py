"""A row shown for its history and not rated (the book's ``shown_not_rated``).

The firm, 18 Oct 2026, question 23 of the round 2 commercial docket, agreed
with the recommendation: *"Nonfarm CRE: rate owner-occupied and investor CRE
separately from 2007; show nonfarm as a whole, unrated"* (10 rated rows).

So a book names each product and measure it shows and does not rate, with
why. Decide does not offer it for a decision and Thresholds has no row for
it; Summary shows its latest values and says "Shown, not rated", with the
book's why under the panels; Lines by source says the same where its rating
would be. The key is required in every book ([] for none), as every key is.
"""
import json

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as L

from conftest import ROOT
from test_book import write_book
from test_decide import NAMES, SER, SETTINGS
from test_workbook import recalc
from threshold_engine import book as BK
from threshold_engine import decide_tab as DT
from threshold_engine.series import read_csv
from threshold_engine.workbook import build

WHY = "Shown for its history only (a test)."
LABELS = [lbl for lbl, _ in SER]


def spec(unrated, lines_from="Industry", compare=None):
    return {"names": NAMES, "parts": {lbl: tuple(lbl.split(" - ")) for lbl in LABELS}, "absent": [],
            "compare": compare, "notes": [], "lines_from": lines_from, "kind": "consumer",
            "shown_not_rated": unrated}


@pytest.fixture(scope="module")
def shown(tmp_path_factory):
    """Mortgage - Net shown, not rated; Card - Net and Card - Gross rated,
    their lines set from the industry from the start."""
    d = tmp_path_factory.mktemp("shown")
    info = build(d / "s.xlsx", SER, decide=spec({("Mortgage", "Net"): WHY}), **SETTINGS)
    return info, load_workbook(info["path"]), recalc(info["path"])


# ---------------------------------------------------------------- the book

@pytest.mark.parametrize("entries, said", [
    ("Mortgage", "shown_not_rated must be a list"),
    ([{"Product": "Mortgage", "Measure": "Net"}], "shown_not_rated 1 is missing 'why'"),
    ([{"Product": "Mortgage", "Measure": "Net", "Source": "Industry", "why": "x"}], "'Source', which is not a key"),
    ([{"Product": "Mortgage", "Measure": "Gross", "why": "x"}], "names Mortgage - Gross, which is no product and "
                                                                "measure of the book's series"),
    ([{"Product": "Mortgage", "Measure": "Net", "why": " "}], "why must say why"),
    ([{"Product": "Mortgage", "Measure": "Net", "why": "x"}, {"Product": "Mortgage", "Measure": "Net", "why": "y"}],
     "names Mortgage - Net a second time"),
    ([{"Product": "Mortgage", "Measure": "Net", "why": "x"}, {"Product": "Card", "Measure": "Net", "why": "y"}],
     "names every row of the book's series, so nothing would be rated"),
])
def test_each_entry_is_checked_and_refused_by_name(tmp_path, entries, said):
    with pytest.raises(BK.BookError, match=said.replace("(", r"\(")):
        BK.load(write_book(tmp_path, shown_not_rated=entries))


def test_the_key_is_required_and_an_entry_is_kept_with_its_why(tmp_path):
    p = write_book(tmp_path)
    raw = json.loads(p.read_text())
    del raw["shown_not_rated"]
    p.write_text(json.dumps(raw))
    with pytest.raises(BK.BookError, match="missing 'shown_not_rated'"):
        BK.load(p)
    bk = BK.load(write_book(tmp_path, shown_not_rated=[{"Product": "Mortgage", "Measure": "Net", "why": WHY}]))
    assert bk["decide"]["shown_not_rated"] == {("Mortgage", "Net"): WHY}
    assert BK.load(write_book(tmp_path))["decide"]["shown_not_rated"] == {}


def test_the_workbook_refuses_a_decide_spec_without_it(tmp_path):
    s = spec({})
    del s["shown_not_rated"]
    with pytest.raises(ValueError, match="decide needs shown_not_rated"):
        build(tmp_path / "x.xlsx", SER, decide=s, **SETTINGS)
    with pytest.raises(ValueError, match="shown_not_rated names Mortgage - Gross, which no series has"):
        build(tmp_path / "y.xlsx", SER, decide=spec({("Mortgage", "Gross"): WHY}), **SETTINGS)
    with pytest.raises(ValueError, match="names every row"):
        build(tmp_path / "z.xlsx", SER, decide=spec({("Mortgage", "Net"): WHY, ("Card", "Net"): WHY,
                                                     ("Card", "Gross"): WHY}), **SETTINGS)


def test_the_commercial_book_rates_ten_rows_and_shows_nonfarm_as_a_whole():
    """Question 23: owner-occupied and investor CRE rated separately, with
    C&I, multifamily and construction (5 portfolios, 2 measures: 10 rows);
    Nonfarm nonresidential CRE shown for its history and not rated."""
    bk = BK.load(ROOT / "books" / "commercial-public.json")
    unrated = bk["decide"]["shown_not_rated"]
    assert set(unrated) == {("Nonfarm nonresidential CRE", "Nonaccrual"),
                            ("Nonfarm nonresidential CRE", "Net charge-offs")}
    assert all("question 23" in why for why in unrated.values())
    rows = {v[:2] for v in bk["decide"]["parts"].values()}
    rated = rows - set(unrated)
    assert len(rated) == 10
    assert {p for p, _ in rated} == {"C&I", "Owner-occupied CRE", "Investor CRE (non-owner-occupied)",
                                     "Multifamily", "Construction"}
    # Its series stay in the book, every source of both rows, for the history.
    assert sum(1 for v in bk["decide"]["parts"].values() if v[0] == "Nonfarm nonresidential CRE") == 8
    assert BK.load(ROOT / "books" / "consumer-public.json")["decide"]["shown_not_rated"] == {}


# ---------------------------------------------------------------- the tabs

def test_decide_does_not_offer_the_row(shown):
    info, wb, _ = shown
    d = info["decide"]
    assert d["pairs"] == [("Card", "Net"), ("Card", "Gross")]
    ws = wb["Decide"]
    b = L(d["columns"]["Product"])
    assert [ws["%s%d" % (b, r)].value for r in range(d["first"], d["last"] + 1)] == ["Card", "Card"]
    assert "of 2 ·" in ws["L1"].value
    assert not any(c.value == "Mortgage" for row in ws.iter_rows(min_row=d["first"], max_row=d["last"]) for c in row)


def test_thresholds_has_no_row_for_it(shown):
    info, wb, _ = shown
    th = wb["Thresholds"]
    t = info["thresholds"]
    keys = [th["A%d" % r].value for r in range(t["first"], t["last"] + 1)]
    assert keys == ["Card - Net", "Card - Gross"]


def test_lines_by_source_says_shown_not_rated_and_gives_no_decision(shown):
    info, wb, v = shown
    ws = wb[DT.LINES_TAB]
    T = DT.TH_BOOK
    r = 5 + LABELS.index("Mortgage - Net - Industry")
    assert ws["%s%d" % (T["rating_shown"], r)].value == DT.NOT_RATED
    assert ws["%s%d" % (T["decision"], r)].value in (None, "")          # no formula reading Decide
    assert ws["%s%d" % (T["check"], r)].value.startswith(DT.NOT_RATED)
    # A rated row keeps its live rating and decision.
    r0 = 5 + LABELS.index("Card - Net - Industry")
    assert ws["%s%d" % (T["rating_shown"], r0)].value.startswith("=")
    assert ws["%s%d" % (T["decision"], r0)].value.startswith("=IF(Decide!")


def test_summary_shows_its_latest_value_and_says_shown_not_rated(shown):
    info, wb, v = shown
    s = info["summary"]
    ws = wb["Summary"]
    r, base = s["placed"][("Mortgage", "Net")]
    lat, rat = s["columns"]["Industry"]
    off = base - 2
    assert ws.cell(r, rat + off).value == DT.NOT_RATED
    # Its latest value, read live from its own row of Lines by source: the
    # file's last quarter (0.0 for this fixture, so the cell is named too).
    row = 5 + LABELS.index("Mortgage - Net - Industry")
    want = read_csv(dict(SER)["Mortgage - Net - Industry"])[-1].value
    assert "'%s'!$%s$%d" % (DT.LINES_TAB, DT.TH_BOOK["value"], row) in ws.cell(r, lat + off).value
    assert v["%s!%s%d" % (DT.LINES_TAB.upper(), DT.TH_BOOK["value"], row)] == pytest.approx(want)
    assert v["SUMMARY!%s%d" % (L(lat + off), r)] == pytest.approx(want)
    notes = [ws.cell(r_, 2).value for r_ in range(s["notes"], ws.max_row + 1)]
    assert "%s: Mortgage and net. %s" % (DT.NOT_RATED, WHY) in notes


def test_summary_counts_and_lists_only_rated_rows(shown):
    """Industry leads (no compare pair): the count of its series is the two
    rated rows, never the row shown and not rated."""
    info, wb, v = shown
    s = info["summary"]
    a = s["attention"]
    cell = "SUMMARY!%s%d" % (L(s["panels"][1]), a["refused"])
    assert v[cell] == "0 of 2 Industry series"
    helpers = [wb["Summary"].cell(r, s["helpers"]["name"]).value for r in range(s["first"], s["first"] + 3)]
    assert "Mortgage · Net" not in helpers


def test_its_settings_stay_the_books_on_assess(shown):
    """No row on Decide, so nothing there feeds its settings: Assess keeps
    the book's numbers for it, and formulas for the rated series."""
    info, wb, _ = shown
    from threshold_engine.workbook import assess_layout
    lay = assess_layout(len(LABELS))
    asx = wb["Assess"]
    i = LABELS.index("Mortgage - Net - Industry")
    assert asx["O%d" % (lay["set_top"] + i)].value == SETTINGS["moderate_halfwidth"]
    assert str(asx["O%d" % lay["set_top"]].value).startswith("=IF(Decide!")


def test_with_nothing_named_the_workbook_is_unchanged(tmp_path):
    """[] in the book ({} here): every row on Decide, as before."""
    info = build(tmp_path / "all.xlsx", SER, decide=spec({}, lines_from=None), **SETTINGS)
    assert info["decide"]["pairs"] == [("Card", "Net"), ("Mortgage", "Net"), ("Card", "Gross")]
    ws = load_workbook(info["path"])[DT.LINES_TAB]
    assert DT.NOT_RATED not in [c.value for c in ws[DT.TH_BOOK["rating_shown"]]]


# ------------------------------------------------- no lines and no rating anywhere (review, 19 Oct 2026)
#
# The first build blanked only the visible Rating, Decision and Check of the
# row: its lines (2 Moderate-Low to 5 High) and its hidden rating number and
# name were still written, so Assess, picked on the series, showed "1 Low"
# and five bands, Chart drew its bands, and Backtest rated it.

UNRATED_ROW = 5 + LABELS.index("Mortgage - Net - Industry")
LINE_KEYS = ("l2", "l3", "l4", "l5")


def test_lines_by_source_gives_it_no_lines_and_no_rating(shown):
    _, _, v = shown
    T = DT.TH_BOOK
    got = {k: v["%s!%s%d" % (DT.LINES_TAB.upper(), T[k], UNRATED_ROW)] for k in LINE_KEYS + ("rating", "rating_name")}
    assert all(x == "" for x in got.values()), got
    assert v["%s!%s%d" % (DT.LINES_TAB.upper(), T["refusal"], UNRATED_ROW)] == DT.NOT_RATED_SAYS
    # A rated row keeps its lines and its rating.
    r0 = 5 + LABELS.index("Card - Net - Industry")
    assert all(isinstance(v["%s!%s%d" % (DT.LINES_TAB.upper(), T[k], r0)], float) for k in LINE_KEYS)
    assert v["%s!%s%d" % (DT.LINES_TAB.upper(), T["rating"], r0)] in (1, 2, 3, 4, 5)


@pytest.fixture(scope="module")
def picked(shown, tmp_path_factory):
    """The same workbook with Assess and Chart both picked on the unrated
    series, recalculated."""
    info, _, _ = shown
    wb = load_workbook(info["path"])
    wb["Assess"]["B5"] = "Mortgage - Net - Industry"
    n = len(LABELS)
    wb["Chart"].cell(5, 3 + n).value = "Mortgage - Net - Industry"
    p = tmp_path_factory.mktemp("picked") / "p.xlsx"
    wb.save(p)
    return recalc(p)


def test_assess_picked_on_it_shows_no_rating_no_bands_and_says_why(picked, shown):
    from threshold_engine.workbook import assess_layout
    lay = assess_layout(len(LABELS))
    v = picked
    assert v["ASSESS!H5"] == "No rating: see Check"
    assert v["ASSESS!B6"] == "Check: %s." % DT.NOT_RATED_SAYS
    # The score table's lines: none.
    froms = [v["ASSESS!%s%d" % (c, lay["score_top"] + i)] for i in range(5) for c in "OP"
             if "ASSESS!%s%d" % (c, lay["score_top"] + i) in v]
    assert froms and all(x == "" for x in froms), froms
    # Keeping every quarter, and every dial: no lines either.
    assert [v["ASSESS!%s%d" % (c, lay["scen_all"])] for c in "CDEF"] == [""] * 4
    for top in lay["dial_blocks"]:
        for r in range(top + 2, top + 7):
            assert [v["ASSESS!%s%d" % (c, r)] for c in "CDEF"] == [""] * 4, (top, r)
            assert v["ASSESS!L%d" % r] == DT.NOT_RATED.lower()
    # Your settings' Check names it.
    i = LABELS.index("Mortgage - Net - Industry")
    assert v["ASSESS!S%d" % (lay["set_top"] + i)] == DT.NOT_RATED_SAYS
    # Picked on a rated series (as built), Assess rates it.
    _, _, v0 = shown
    assert v0["ASSESS!H5"][0] in "12345"


def test_chart_draws_no_bands_for_it(picked, shown):
    n = len(LABELS)
    drawable = "CHART!%s4" % L(max(40, n + 10) + 1 + n + 6)
    assert picked[drawable] == 0
    assert shown[2][drawable] == 1                      # as built, on a rated series


def test_backtest_does_not_test_it(shown):
    info, wb, _ = shown
    bts = wb["Backtest"]
    rows = [[c.value for c in r_] for r_ in bts.iter_rows(min_row=5)]
    mine = [r_ for r_ in rows if r_[1:4] == ["Mortgage", "Net", "Industry"]]
    assert len(mine) == 1, mine
    assert mine[0][4] == "Not tested: shown, not rated: the book shows this row and does not rate it"
    rated = [r_ for r_ in rows if r_[1:4] == ["Card", "Net", "Industry"]]
    assert rated and not any(DT.NOT_RATED.lower() in str(r_[4]) for r_ in rated)


def test_the_internal_tab_rates_nothing_on_it(shown):
    """The Internal tab (consumer books) rates the bank's figures on outside
    series' lines; the unrated row has none, so it is not among them."""
    _, wb, _ = shown
    T = DT.TH_BOOK
    it = wb["Internal"]
    import re
    at = re.compile(r"(?:'%s'|Thresholds)!\$(?:%s)\$%d(?!\d)" % (
        DT.LINES_TAB, "|".join(T[k] for k in LINE_KEYS + ("rating", "refusal")), UNRATED_ROW))
    refs = [c.coordinate for row_ in it.iter_rows() for c in row_ if isinstance(c.value, str) and at.search(c.value)]
    assert not refs, refs
    # The pattern finds a rated row's lines, which the tab does read.
    r0 = 5 + LABELS.index("Card - Net - Industry")
    at0 = re.compile(at.pattern.replace(r"\$%d(" % UNRATED_ROW, r"\$%d(" % r0))
    assert any(isinstance(c.value, str) and at0.search(c.value) for row_ in it.iter_rows() for c in row_)
