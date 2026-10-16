"""The template command: `template list`, `template write KEY OUT.csv` and
`template check KEY FILE.csv`, through the command line only. A file the
bank fills passes check and then builds through the `workbook` command as an
ordinary series; every problem is reported at once with exit 2, each refusal
with its exact message. Every value here is synthetic (a sine wave), never a
bank's number."""
import math
import re

import pytest
from openpyxl import load_workbook

from threshold_engine.cli import main
from threshold_engine.templates import SCALE_REFUSAL, TEMPLATES

ENDS = ("03-31", "06-30", "09-30", "12-31")

MEASURE = ["--frequency", "quarterly", "--smoothing", "1", "--floor-at-zero", "yes",
           "--moderate-halfwidth", "0.5", "--low-step", "1.0", "--high-z", "2.5",
           "--on-the-line", "worse", "--percentiles", "50", "75", "90", "95",
           "--horizon", "4", "--min-history", "8"]


def dates(n, first_year=2015):
    return ["%d-%s" % (first_year + i // 4, ENDS[i % 4]) for i in range(n)]


def synthetic(unit, n):
    """A made-up cycle: a rate between about 0.5 and 3.5, or a score near 700."""
    if unit == "%":
        return [round(2.0 + 1.2 * math.sin(i / 3.0) + 0.3 * math.sin(i), 4) for i in range(n)]
    return [round(700 - 25 * math.sin(i / 3.0) - 5 * math.sin(i), 2) for i in range(n)]


def write(tmp_path, key, capsys, name="t.csv"):
    out = tmp_path / name
    assert main(["template", "write", key, str(out)]) == 0
    capsys.readouterr()
    return out


def fill(path, key, n=24, rows=None, header="date,value"):
    """Answer every [STATE: ...] note of the written template with a synthetic
    answer (a scale of 300-850 for a score) and add the rows."""
    lines = []
    for line in path.read_text(encoding="utf-8").splitlines():
        if line.startswith("#") and "[STATE:" in line:
            k = line[1:].split(":", 1)[0].strip()
            lines.append("# %s: %s" % (k, "300-850" if k == "scale" else "synthetic answer for " + k))
        elif line == "date,value":
            lines.append(header)
        else:
            lines.append(line)
    if rows is None:
        rows = ["%s,%s" % r for r in zip(dates(n), synthetic(TEMPLATES[key].unit, n))]
    path.write_text("\n".join(lines + rows) + "\n", encoding="utf-8")
    return path


def check(key, path, capsys):
    rc = main(["template", "check", key, str(path)])
    printed = capsys.readouterr()
    return rc, printed.out, printed.err


def test_list_names_every_template_with_its_unit_and_direction(capsys):
    assert main(["template", "list"]) == 0
    lines = capsys.readouterr().out.splitlines()
    assert len(lines) == len(TEMPLATES) == 7
    for line, t in zip(lines, TEMPLATES.values()):
        assert line.startswith(t.key + " ") and t.title in line
        assert "(%s; %s)" % (t.unit, t.direction.replace("_", " ")) in line


@pytest.mark.parametrize("key", list(TEMPLATES))
def test_each_template_written_filled_checked_and_built(tmp_path, capsys, key):
    """Spec B's acceptance, through the command line: write, fill with
    synthetic values, check OK, then the workbook command builds from it."""
    out = tmp_path / "t.csv"
    assert main(["template", "write", key, str(out)]) == 0
    said = capsys.readouterr().out
    assert out.exists() and (tmp_path / "t.checklist.md").exists()
    assert "template check %s %s" % (key, out) in said
    rc, _, err = check(key, out, capsys)                 # the empty template is not a series
    assert rc == 2 and "REFUSED" in err and "still holds its placeholder" in err

    fill(out, key)
    rc, said, err = check(key, out, capsys)
    assert rc == 0 and err == "" and said.startswith("OK: %s passes template %s" % (out, key))

    t = TEMPLATES[key]
    book = tmp_path / "w.xlsx"
    rc = main(["workbook", str(book), "--series", "Synthetic=%s" % out, "--name", t.title,
               "--unit", t.unit, "--direction", t.direction] + MEASURE)
    assert rc == 0, capsys.readouterr().err
    assert "Synthetic, 24 quarters" in capsys.readouterr().out
    assert load_workbook(book).sheetnames[:2] == ["Assess", "Thresholds"]


def test_write_will_not_overwrite_a_filled_template(tmp_path, capsys):
    out = fill(write(tmp_path, "wapd", capsys), "wapd")
    before = out.read_text(encoding="utf-8")
    assert main(["template", "write", "wapd", str(out)]) == 2
    assert ("REFUSED: %s already exists; it is not written over, so a filled template cannot be "
            "lost. Choose another name or move it." % out) in capsys.readouterr().err
    assert out.read_text(encoding="utf-8") == before


def test_an_unknown_template_is_refused_and_the_known_ones_named(tmp_path, capsys):
    assert main(["template", "write", "pd", str(tmp_path / "x.csv")]) == 2
    assert ("REFUSED: no template 'pd'; the templates are: %s" % ", ".join(TEMPLATES)
            in capsys.readouterr().err)
    assert not (tmp_path / "x.csv").exists()
    rc, _, err = check("pd", tmp_path / "x.csv", capsys)
    assert rc == 2 and "no template 'pd'" in err


def test_a_file_that_is_not_there_is_refused_not_a_traceback(tmp_path, capsys):
    rc, _, err = check("epd", tmp_path / "nowhere.csv", capsys)
    assert rc == 2 and err.startswith("REFUSED:") and "nowhere.csv" in err


def test_a_third_column_is_refused_and_says_why(tmp_path, capsys):
    """The no-PII rule: an `account` column could carry account-level data."""
    out = write(tmp_path, "epd", capsys)
    fill(out, "epd", rows=["%s,%s,A-%d" % (d, v, i) for i, (d, v) in
                           enumerate(zip(dates(12), synthetic("%", 12)))], header="date,value,account")
    rc, said, err = check("epd", out, capsys)
    assert rc == 2 and said == ""
    lines = err.splitlines()
    assert lines[0] == "REFUSED: %s does not pass template epd; 1 problem:" % out
    assert lines[1] == ("  - line %d: column(s) 'account' refused: a template holds date and value and "
                        "nothing else, so no account-level data or personal information can travel in it"
                        % header_line(out))


def header_line(path):
    """The 1-based line of the date,value header: a row's line is this plus its position."""
    lines = path.read_text(encoding="utf-8").splitlines()
    return next(i for i, ln in enumerate(lines, start=1) if not ln.startswith("#"))


@pytest.mark.parametrize("rows,says", [
    (["2015-03-31,1.0", "2015-03-31,1.1"] + ["%s,1.0" % d for d in dates(9)[1:]],
     "line {2}: 2015-03-31 is repeated"),
    (["2015-03-31,1.0", "2015-09-30,1.1"] + ["%s,1.0" % d for d in dates(10)[3:]],
     "line {2}: 2015-09-30 follows 2015-03-31; the quarters between are missing, and a gap is "
     "refused, not filled in"),
    (["2015-02-28,1.0"] + ["%s,1.0" % d for d in dates(10)[1:]],
     "line {1}: date '2015-02-28' is not a quarter-end (YYYY-03-31, YYYY-06-30, YYYY-09-30 or YYYY-12-31)"),
    (["%s,1.0" % d for d in dates(8)] + ["2017-03-31,nan"],
     "line {9}: value 'nan' is not a finite number"),
    (["%s,1.0" % d for d in dates(8)] + ["2017-03-31,n/a"],
     "line {9}: value 'n/a' is not a number"),
    (["%s,1.0" % d for d in dates(8)] + ["2017-03-31,140"],
     "line {9}: 140 is outside 0-100, and this measure is a %"),
    (["%s,1.0" % d for d in dates(8)],
     "8 quarter(s) of history; at least 9 are needed. scale.py refuses with '%s', and nine "
     "quarters hold two four-quarter changes end to end" % SCALE_REFUSAL),
])
def test_each_refusal_gives_its_exact_message(tmp_path, capsys, rows, says):
    """``{n}`` in the message is the line of the n-th row."""
    out = write(tmp_path, "card-rewards", capsys)
    fill(out, "card-rewards", rows=rows)
    h = header_line(out)
    says = re.sub(r"\{(\d+)\}", lambda m: str(h + int(m.group(1))), says)
    rc, said, err = check("card-rewards", out, capsys)
    assert rc == 2 and said == "" and err.startswith("REFUSED: ")
    assert "  - " + says in err.splitlines()


def test_a_score_outside_its_stated_scale_is_refused(tmp_path, capsys):
    out = write(tmp_path, "refreshed-score", capsys)
    fill(out, "refreshed-score", rows=["%s,700" % d for d in dates(9)] + ["2017-06-30,900"])
    rc, _, err = check("refreshed-score", out, capsys)
    assert rc == 2 and "is outside the stated scale 300-850" in err


def test_every_problem_is_reported_at_once(tmp_path, capsys):
    """An unanswered note, a repeated date and a short history: three problems, one run."""
    out = write(tmp_path, "wapd", capsys)
    fill(out, "wapd", rows=["2015-03-31,1.0", "2015-03-31,1.2"])
    text = out.read_text(encoding="utf-8").replace("# source: synthetic answer for source",
                                                   "# source: [STATE: the system]")
    out.write_text(text, encoding="utf-8")
    rc, _, err = check("wapd", out, capsys)
    problems = [ln for ln in err.splitlines() if ln.startswith("  - ")]
    assert rc == 2 and len(problems) == 3
    assert err.splitlines()[0].endswith("3 problems:")
    assert any("note 'source' still holds its placeholder" in p for p in problems)
    assert any("is repeated" in p for p in problems)
    assert any("1 quarter(s) of history" in p for p in problems)


def test_ok_gives_the_quarters_and_what_the_workbook_needs_of_them(tmp_path, capsys):
    """A file at the check's floor (9 quarters) passes check, and the OK line
    says what the workbook command needs of its own settings to build from
    it. Both sides of that line hold: at 9 it builds, at 10 it is refused.
    The OK line used to promise a build the check cannot know about."""
    out = fill(write(tmp_path, "wapd", capsys), "wapd", n=9)
    rc, said, _ = check("wapd", out, capsys)
    assert rc == 0
    assert said.strip().endswith(
        "(Weighted average probability of default): 9 quarters. It is an ordinary series for the "
        "workbook command's --series, which builds from it only when --min-history + --horizon + "
        "--smoothing - 1 is 9 or less.")

    def build(min_history, horizon, smoothing):
        args = [a for a in MEASURE]
        for flag, v in (("--min-history", min_history), ("--horizon", horizon), ("--smoothing", smoothing)):
            args[args.index(flag) + 1] = str(v)
        return main(["workbook", str(tmp_path / "w.xlsx"), "--series", "S=%s" % out, "--name", "x",
                     "--unit", "%", "--direction", "higher_is_worse"] + args)
    assert build(4, 4, 2) == 0                        # 4 + 4 + 2 - 1 = 9
    capsys.readouterr()
    assert build(5, 4, 2) == 2                        # 10
    assert "8 observations after smoothing; at least min_history + horizon = 9" in capsys.readouterr().err
