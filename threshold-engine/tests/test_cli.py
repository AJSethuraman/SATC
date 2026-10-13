"""The command line: each command runs, refuses plainly, and will not run on
a judgement nobody stated."""
import pytest

from conftest import DATA
from threshold_engine.cli import main

MEASURE = ["--name", "Card net charge-offs", "--unit", "%", "--direction", "higher_is_worse",
           "--frequency", "quarterly", "--smoothing", "1"]
SCALE = ["--floor-at-zero", "yes", "--moderate-halfwidth", "0.5", "--low-step", "1.0",
         "--high-z", "2.5", "--on-the-line", "worse"]


def test_evidence_reports_on_the_workbooks_scale(capsys):
    assert main(["evidence", str(DATA / "cards_nco_ttm.csv")] + MEASURE + SCALE) == 0
    out = capsys.readouterr().out
    assert "on the workbook's scale" in out and "High z 2.5" in out
    assert "lines 3.066 / 3.393 / 4.047 / 5.972" in out       # scenario C: 2008 run left out


def test_a_refusal_exits_non_zero_and_says_why(tmp_path, capsys):
    bad = tmp_path / "gap.csv"
    bad.write_text("date,value\n2020-03-31,1\n2020-09-30,2\n", encoding="utf-8")
    assert main(["evidence", str(bad)] + MEASURE + SCALE) == 2
    assert "REFUSED" in capsys.readouterr().err


@pytest.mark.parametrize("drop", ["--high-z", "--low-step", "--on-the-line"])
def test_every_scale_setting_must_be_stated(drop):
    i = SCALE.index(drop)
    with pytest.raises(SystemExit):
        main(["evidence", str(DATA / "cards_nco_ttm.csv")] + MEASURE + SCALE[:i] + SCALE[i + 2:])


def test_the_old_cutoffs_command_is_gone():
    with pytest.raises(SystemExit):
        main(["cutoffs", str(DATA / "cards_nco_ttm.csv")] + MEASURE)


def test_the_workbook_command_writes_a_workbook(tmp_path, capsys):
    out = tmp_path / "t.xlsx"
    rc = main(["workbook", str(out), "--series", "Credit card=%s" % (DATA / "cards_nco_ttm.csv")]
              + MEASURE + SCALE + ["--percentiles", "50", "75", "90", "95", "--horizon", "4",
                                   "--min-history", "41"])
    assert rc == 0 and out.exists() and "163 quarters" in capsys.readouterr().out


def test_a_badly_formed_series_argument_is_refused(tmp_path, capsys):
    rc = main(["workbook", str(tmp_path / "t.xlsx"), "--series", "no-equals-sign"]
              + MEASURE + SCALE + ["--percentiles", "50", "--horizon", "4", "--min-history", "41"])
    assert rc == 2 and "LABEL=FILE" in capsys.readouterr().err


PROFILE = ["--name", "Mortgage net charge-offs", "--unit", "%", "--direction",
           "higher_is_worse", "--frequency", "quarterly", "--smoothing", "1"]


def test_a_profile_takes_no_judgement_settings(capsys):
    assert main(["profile", str(DATA / "mortgage_nco_ttm.csv")] + PROFILE) == 0
    out = capsys.readouterr().out
    assert "SPELLS WORSE THAN THE MEDIAN (0.122%)" in out
    assert "PERCENTILES WITHOUT THE LARGEST SPELL" in out


def test_a_profile_refuses_a_judgement_setting():
    with pytest.raises(SystemExit):
        main(["profile", str(DATA / "mortgage_nco_ttm.csv")] + PROFILE + ["--high-z", "2.5"])


@pytest.mark.parametrize("value", ["nan", "inf", "-inf"])
def test_a_setting_that_is_not_a_number_is_refused(value):
    # argparse's float takes 'nan' and 'inf'; nothing is above NaN, so a NaN
    # High z would switch a check off rather than fail it.
    args = list(SCALE)
    args[args.index("--high-z") + 1] = value
    with pytest.raises(SystemExit):
        main(["evidence", str(DATA / "cards_nco_ttm.csv")] + MEASURE + args)


# ---- Found by the Codex review of PR #426 (12 Oct 2026): each a rule the
# window enforced and the command line did not, or a value no front door caught.

WB = SCALE + ["--percentiles", "50", "75", "90", "95", "--horizon", "4", "--min-history", "41"]
BT = ["--floor-at-zero", "yes", "--percentiles", "50", "75", "90", "95", "--horizon", "4",
      "--min-history", "41"]
CARDS = DATA / "cards_nco_ttm.csv"


@pytest.mark.parametrize("labels, why", [
    (("Cards", "cards"), "Excel ignores capitals"),
    (("Cards, US", "Mortgage"), "comma"),
    (('Cards "US"', "Mortgage"), "double quote"),
    ((" ", "Mortgage"), "needs a product label"),
])
def test_the_command_line_refuses_a_label_excel_cannot_hold(tmp_path, capsys, labels, why):
    out = tmp_path / "t.xlsx"
    args = []
    for lbl in labels:
        args += ["--series", "%s=%s" % (lbl, CARDS)]
    assert main(["workbook", str(out)] + args + MEASURE + WB) == 2
    assert why in capsys.readouterr().err and not out.exists()


def test_the_build_itself_refuses_a_label_that_excel_would_read_as_a_formula(tmp_path):
    """The command line's LABEL=FILE parse already stops '=Cards=...' (an empty
    label); ``build`` refuses it too, for every other caller."""
    from threshold_engine.workbook import build
    out = tmp_path / "t.xlsx"
    with pytest.raises(ValueError, match="begin with '='"):
        build(out, [("=Cards", CARDS)], name="NCO", unit="%", direction="higher_is_worse",
              frequency="quarterly", smoothing=1, floor_at_zero=True, moderate_halfwidth=0.5,
              low_step=1.0, high_z=2.5, on_the_line="worse", half_lives=(), percentiles=(50,),
              horizon=4, min_history=41)
    assert not out.exists()


@pytest.mark.parametrize("bad", ["nan", "inf", "-inf", "NaN"])
@pytest.mark.parametrize("command", ["profile", "evidence", "backtest", "workbook"])
def test_a_value_that_is_not_a_number_is_refused_at_every_front_door(tmp_path, capsys, bad, command):
    rows = CARDS.read_text(encoding="utf-8").splitlines()
    i = next(k for k, r in enumerate(rows) if r[:1].isdigit())
    date = rows[i + 10].split(",")[0]
    rows[i + 10] = "%s,%s" % (date, bad)
    p = tmp_path / "with_nan.csv"
    p.write_text("\n".join(rows) + "\n", encoding="utf-8")
    out = tmp_path / "t.xlsx"
    if command == "profile":
        argv = ["profile", str(p)] + MEASURE
    elif command == "workbook":
        argv = ["workbook", str(out), "--series", "Credit card=%s" % p] + MEASURE + WB
    elif command == "backtest":
        argv = ["backtest", str(p)] + MEASURE + BT
    else:
        argv = ["evidence", str(p)] + MEASURE + SCALE
    assert main(argv) == 2
    err = capsys.readouterr().err
    assert "with_nan.csv, %s" % date in err and "not a number" in err
    assert not out.exists()


@pytest.mark.parametrize("life", ["0", "-5"])
@pytest.mark.parametrize("command", ["evidence", "backtest", "workbook"])
def test_a_half_life_of_zero_or_below_is_refused(tmp_path, capsys, life, command):
    out = tmp_path / "t.xlsx"
    if command == "workbook":
        argv = ["workbook", str(out), "--series", "Credit card=%s" % CARDS] + MEASURE + WB
    elif command == "backtest":
        argv = ["backtest", str(CARDS)] + MEASURE + BT
    else:
        argv = ["evidence", str(CARDS)] + MEASURE + SCALE
    assert main(argv + ["--half-lives", "10", life]) == 2
    assert "half-life must be a number of years above 0" in capsys.readouterr().err
    assert not out.exists()


def test_the_package_has_one_version_the_workbooks_write():
    """pyproject.toml takes its version from threshold_engine.VERSION, the one
    the Run tabs write, so an installed package and its workbooks agree."""
    tomllib = pytest.importorskip("tomllib")
    from pathlib import Path
    from threshold_engine import VERSION
    meta = tomllib.loads((Path(__file__).parents[1] / "pyproject.toml").read_text(encoding="utf-8"))
    assert "version" not in meta["project"] and "version" in meta["project"]["dynamic"]
    assert meta["tool"]["setuptools"]["dynamic"]["version"] == {"attr": "threshold_engine.VERSION"}
    assert VERSION.count(".") == 2


@pytest.mark.parametrize("command", ["profile", "evidence", "backtest", "workbook"])
def test_a_file_without_the_date_value_header_is_refused_not_a_traceback(tmp_path, capsys, command):
    p = tmp_path / "amount.csv"
    p.write_text("date,amount\n2020-03-31,1\n", encoding="utf-8")
    out = tmp_path / "t.xlsx"
    argv = {"profile": ["profile", str(p)] + MEASURE,
            "evidence": ["evidence", str(p)] + MEASURE + SCALE,
            "backtest": ["backtest", str(p)] + MEASURE + BT,
            "workbook": ["workbook", str(out), "--series", "Cards=%s" % p] + MEASURE + WB}[command]
    assert main(argv) == 2
    assert "amount.csv must begin with the header line date,value (it has: date,amount)" in capsys.readouterr().err


@pytest.mark.parametrize("command", ["profile", "evidence", "backtest", "workbook"])
def test_smoothing_longer_than_the_series_is_refused_not_a_traceback(tmp_path, capsys, command):
    p = tmp_path / "two.csv"
    p.write_text("date,value\n2020-03-31,1\n2020-06-30,2\n", encoding="utf-8")
    out = tmp_path / "t.xlsx"
    m = MEASURE[:-1] + ["3"]                                           # --smoothing 3
    argv = {"profile": ["profile", str(p)] + m,
            "evidence": ["evidence", str(p)] + m + SCALE,
            "backtest": ["backtest", str(p)] + m + BT,
            "workbook": ["workbook", str(out), "--series", "Cards=%s" % p] + m + WB}[command]
    assert main(argv) == 2
    assert "smoothing over 3 periods needs at least 3 observations; the series has 2" in capsys.readouterr().err


def test_the_backtest_tab_has_one_change_column_per_score(tmp_path):
    """Three percentiles make four scores: four change columns, then Quarters
    tested under its own heading (it used to sit under a fixed fifth)."""
    from openpyxl import load_workbook
    from threshold_engine.backtest import backtest
    from threshold_engine.series import read_csv
    out = tmp_path / "t.xlsx"
    assert main(["workbook", str(out), "--series", "Credit card=%s" % CARDS] + MEASURE
                + SCALE + ["--percentiles", "50", "75", "90", "--horizon", "4", "--min-history", "41"]) == 0
    rows = [list(r) for r in load_workbook(out)["Backtest"].iter_rows(values_only=True)]
    i = next(k for k, r in enumerate(rows) if r[0] == "Weighting")
    head = [c for c in rows[i] if c is not None]
    assert head == ["Weighting", "Spearman: score vs level ahead", "Spearman: score vs change ahead",
                    "Change after 1", "after 2", "after 3", "after 4", "Quarters tested"]
    want = backtest(read_csv(CARDS), "c", "%", "higher_is_worse", "quarterly", 1, (50, 75, 90), 4, 41, True)
    assert rows[i + 1][len(head) - 1] == want["runs"][0]["periods_tested"]
    assert rows[i + 1][3:7] == pytest.approx([b["mean_change_ahead"] for b in want["runs"][0]["by_score"]])
