"""The Peers dashboard against the builder it replaced.

tests/data/peers/peers_workbook_642ff55f.py is the one-tab-per-view builder
(Pack, Peers, Delinquency, Balances, Lagged, Normalisation, Mix) copied
unchanged from commit 642ff55f, loaded inside the package beside today's. The
same inputs go to both. With every bank switched on, every view on the
dashboard must equal the old builder's value for the same quantity: the
charge-off composite, band, KeyBank, Peer A and B (net, and gross once the
rate picker says so), best and worst and the ranking window, delinquency, the
book index, the lagged rates and the pack total, how far each rate moved, and
the mix. Then "Chart shows" is changed, and nothing but what is drawn may move.
Since the round 2 layout (15 Oct 2026) the cells sit elsewhere on Peers; the
values compared are the same.

Both workbooks are recalculated with the `formulas` engine.
"""
import importlib.util
import json

import pytest
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from conftest import DATA, ROOT
from threshold_engine import reporting_changes as RC
from threshold_engine.peers import peers_workbook as PW
from threshold_engine.peers import roster as R
from threshold_engine.peers.books import industry

from test_peers_workbook import num, recalc, typed

PEERS = DATA / "peers"
CHANGES = RC.load(ROOT / "config" / "reporting-changes.json")
TWO = ("Credit card", "Mortgage")
NONE_ACQUIRED = {p: {} for p in ("Credit card", "Mortgage", "Home equity", "Other consumer", "Auto")}
#: No stated acquisition in any product (question 18: one {cert: [quarter-ends]} per product).
NONE_STATED = {p: {} for p in NONE_ACQUIRED}


def _old():
    spec = importlib.util.spec_from_file_location("threshold_engine.peers._peers_workbook_642ff55f",
                                                  PEERS / "peers_workbook_642ff55f.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


OLD = _old()


def build(module, out, recent, dl, roster):
    return module.build_peers_workbook(
        out, recent, roster, TWO, 0.25, 1_000_000, "all", json.loads((PEERS / "mergers_roster.json").read_text())["acquirers"],
        {"6672": ["2022Q2"]}, "2019Q4", exempt_self=True, min_banks=5, index_min_book=500_000,
        lagged_years=(2019, 2025), lag_years=2, normalisation_years=(2019, 2025), mix_base_year=2019,
        industry={p: industry(recent, p) for p in TWO},
        delinquency=dict(data=dl, products=TWO, jump=0.25, min_book=1_000_000, exempt_self=True,
                         acquired=NONE_ACQUIRED, measure="90+ accruing", basis="quarter-end"),
        acquired=NONE_ACQUIRED, rank_by="net", last_n=20, min_obs=12,
        **({} if module is OLD else {"combined": {}, "opened": {}, "changes": CHANGES, "stated": NONE_STATED}))


@pytest.fixture(scope="module")
def roster():
    return R.load(ROOT / "config" / "peers.json")


@pytest.fixture(scope="module")
def both(tmp_path_factory, roster):
    data = json.loads((PEERS / "roster_2006.json").read_text())
    recent = {q: rows for q, rows in data.items() if q >= "20170331"}
    dl = json.loads((PEERS / "delinquency_roster.json").read_text())
    base = tmp_path_factory.mktemp("dashboard")
    old = build(OLD, base / "old.xlsx", recent, dl, roster)
    new = build(PW, base / "new.xlsx", recent, dl, roster)
    return dict(old=old, new=new, oldv=recalc(old["path"]), newv=recalc(new["path"]), base=base)


def old_col(vals, info, tab, helpers, first, name, n, offset=0):
    c = get_column_letter(column_index_from_string(helpers[name]) + offset)
    return [vals.get("%s!%s%d" % (tab, c, first + k)) for k in range(n)]


def new_col(vals, info, view, name, offset=0, n=None):
    c = get_column_letter(column_index_from_string(info[view]["helpers"][name]) + offset)
    n = len(info[view]["quarters"]) if n is None else n
    return [vals.get("WORKINGS!%s%d" % (c, 2 + k)) for k in range(n)]


def assert_equal(got, want, what):
    """Equal within 1e-9; #N/A, a blank and text all read as None."""
    assert len(got) == len(want), what
    bad = [(i, g, w) for i, (g, w) in enumerate(zip(got, want))
           if not ((num(g) is None and num(w) is None) or
                   (num(g) is not None and num(w) is not None and abs(num(g) - num(w)) < 1e-9))]
    assert not bad, (what, bad[:5])
    return sum(num(w) is not None for w in want)


def charge_off_pairs(kind):
    k = kind[0]
    return [("banks", "banks"), ("c", "c" + k), ("p10", k + "10"), ("p25", k + "25"), ("p75", k + "75"),
            ("p90", k + "90"), ("k", "k" + k), ("a", "a" + k), ("b", "b" + k)]


def test_the_net_charge_off_view_equals_the_old_peers_tab(both):
    old, new, ov, nv = both["old"], both["new"], both["oldv"], both["newv"]
    assert new["quarters"] == old["quarters"]
    n = len(old["quarters"])
    seen = 0
    for mine, theirs in charge_off_pairs("net"):
        seen += assert_equal(new_col(nv, new, "charge_offs", mine),
                             old_col(ov, old, "PEERS", old["peers_helpers"], old["peers_first_row"], theirs, n),
                             mine)
    assert seen > 200
    t = new["layout"]["tiles"] + 1
    last = old["peers_first_row"] + n - 1
    for cell, theirs in (("B%d" % t, "kn"), ("D%d" % t, "cn")):
        c = old["peers_helpers"][theirs]
        assert num(nv["PEERS!" + cell]) == pytest.approx(num(ov["PEERS!%s%d" % (c, last)]), abs=1e-12)


def test_the_gross_rate_equals_the_old_gross_chart(both):
    old, new, ov = both["old"], both["new"], both["oldv"]
    gv = typed(new["path"], both["base"] / "gross.xlsx", {"C5": "Gross"})
    n = len(old["quarters"])
    seen = 0
    for mine, theirs in charge_off_pairs("gross"):
        seen += assert_equal(new_col(gv, new, "charge_offs", mine),
                             old_col(ov, old, "PEERS", old["peers_helpers"], old["peers_first_row"], theirs, n),
                             mine)
    assert seen > 200


def test_best_worst_and_the_ranking_window_equal_the_old_pack_tab(both, roster):
    ov, nv, new = both["oldv"], both["newv"], both["new"]
    s = new["scalars"]
    assert nv["WORKINGS!" + s["best"].replace("$", "")] == ov["PACK!B11"]
    assert nv["WORKINGS!" + s["worst"].replace("$", "")] == ov["PACK!B12"]
    assert nv["WORKINGS!" + s["a_name"].replace("$", "")] == ov["PACK!C5"]
    assert nv["WORKINGS!" + s["b_name"].replace("$", "")] == ov["PACK!C6"]
    for j in range(len(R.peers(roster))):
        assert_equal([nv["WORKINGS!N%d" % (2 + j)], nv["WORKINGS!O%d" % (2 + j)]],
                     [ov["PACK!F%d" % (15 + j)], ov["PACK!E%d" % (15 + j)]], "ranking row %d" % j)


def test_the_delinquency_view_equals_the_old_delinquency_tab(both):
    old, new, ov, nv = both["old"], both["new"], both["oldv"], both["newv"]
    d_old, d_new = old["delinquency"], new["delinquency"]
    assert d_new["quarters"] == d_old["quarters"]
    n = len(d_old["quarters"])
    seen = 0
    for name in ("banks", "c", "bandn", "p10", "p25", "p75", "p90", "k", "a", "bb"):
        seen += assert_equal(new_col(nv, new, "delinquency", name),
                             old_col(ov, old, "DELINQUENCY", d_old["helpers"], d_old["first_row"], name, n), name)
    assert seen > 300


def rows_of(ws):
    return [[c.value for c in r] for r in ws.iter_rows()]


def section(ws, title, stop=TWO):
    rows = rows_of(ws)
    i = next(k for k, r in enumerate(rows) if r[0] == title)
    out = []
    for r in rows[i + 2:]:
        if r[0] is None or r[0] in stop:
            break
        out.append(r)
    return out


def test_the_book_index_equals_the_old_balances_tab(both, roster):
    """Each indexed bank's quarterly index on the old Balances tab (without
    acquisitions, then as reported) is the live one the dashboard draws from,
    and KeyBank's, Peer A's and Peer B's latest are the old latest."""
    old, new, nv = both["old"], both["new"], both["newv"]
    ws = load_workbook(old["path"])["Balances"]
    title = "Credit card: book index, 2019Q4 = 100"
    latest = {r[0]: r for r in section(ws, title)}
    allr = rows_of(ws)
    start = next(i for i, r in enumerate(allr) if r[0] == title)
    head = next(i for i in range(start, len(allr)) if allr[i][0] == "Quarter")
    names = [x for x in allr[head][1:] if x]
    n_ix = len(names) // 2
    quarters = new["book"]["quarters"]
    order = [b["short"] for b in roster]
    seen = 0
    for j, short in enumerate(names[:n_ix]):
        i = order.index(short)
        org = [allr[head + 1 + k][1 + j] for k in range(len(quarters))]
        rep = [allr[head + 1 + k][2 + n_ix + j] for k in range(len(quarters))]
        assert [allr[head + 1 + k][0] for k in range(len(quarters))] == quarters
        seen += assert_equal(new_col(nv, new, "book", "bo", i), org, short + " without acquisitions")
        seen += assert_equal(new_col(nv, new, "book", "br", i), rep, short + " as reported")
    assert seen > 300
    bt = new["layout"]["book_table"]
    a, b = nv["WORKINGS!" + new["scalars"]["a_name"].replace("$", "")], \
        nv["WORKINGS!" + new["scalars"]["b_name"].replace("$", "")]
    for r, who in ((bt + 1, "KeyBank"), (bt + 2, a), (bt + 3, b)):
        assert_equal([nv["PEERS!M%d" % r], nv["PEERS!O%d" % r]], [latest[who][2], latest[who][3]], who)


def test_lagged_rates_and_the_pack_total_equal_the_old_lagged_tab(both, roster):
    old, new, nv = both["old"], both["new"], both["newv"]
    rows = {r[0]: r for r in section(load_workbook(old["path"])["Lagged"], "Credit card")}
    for j, y in enumerate((2019, 2025)):
        c = get_column_letter(column_index_from_string(new["lagged"]["helpers"]["s_%d" % y]))
        got = [nv["WORKINGS!%s%d" % (c, 2 + i)] for i in range(len(roster))]
        assert_equal(got, [rows[b["short"]][1 + j] for b in roster], "lagged %d" % y)
        cell = "PEERS!%s%d" % ("DE"[j], new["layout"]["lagged"] + 4)
        assert num(nv[cell]) == pytest.approx(rows["All peers, total (15 banks)"][1 + j], abs=1e-12)


def test_how_far_each_rate_moved_equals_the_old_normalisation_tab(both):
    old, new, nv = both["old"], both["new"], both["newv"]
    rows = {r[0]: r for r in section(load_workbook(old["path"])["Normalisation"], "Credit card")}
    a = nv["WORKINGS!" + new["scalars"]["a_name"].replace("$", "")]
    b = nv["WORKINGS!" + new["scalars"]["b_name"].replace("$", "")]
    r = new["layout"]["norm_table"] + 1
    for i, who in enumerate(("KeyBank", a, b, "All-peer composite (15 banks)")):
        assert_equal([nv["PEERS!%s%d" % (c, r + i)] for c in "LMN"], rows[who][1:4], who)


def test_the_mix_equals_the_old_mix_tab(both, roster):
    old, new, nv = both["old"], both["new"], both["newv"]
    ws = load_workbook(old["path"])["Mix"]
    rows = section(ws, "Credit card")
    years = new["mix"]["years"]
    m = new["layout"]["mix"]
    for k, y in enumerate(years):
        assert rows[k][0] == y and nv["PEERS!B%d" % (m + 1 + k)] == y
        assert_equal([nv["PEERS!%s%d" % (c, m + 1 + k)] for c in "CDEFGH"], rows[k][1:7], y)
    allr = rows_of(ws)
    at = next(i for i, r in enumerate(allr) if r[0] == "Own factor, every bank")
    table = {r[0]: r for r in allr[at + 2:at + 2 + len(roster)]}
    own = get_column_letter(column_index_from_string(new["banks"]["helpers"]["own"]))
    assert_equal([nv["WORKINGS!%s%d" % (own, 2 + i)] for i in range(len(roster))],
                 [table[b["short"]][len(years)] for b in roster], "own factor, %d" % years[-1])


def test_chart_shows_changes_only_what_is_drawn(both):
    """Every value on every tab is the same at "Last 20 quarters" as at the
    starting 48, but the window picker itself, how many quarters it allows
    and each chart's first row and row count, the cells the named ranges read."""
    new, nv = both["new"], both["newv"]
    wv = typed(new["path"], both["base"] / "window.xlsx", {"F5": "Last 20 quarters"})
    s = new["scalars"]
    drawn = {"WORKINGS!" + s[k].replace("$", "") for k in s if k.endswith(("_first", "_count")) or k == "span"}
    assert len(drawn) == 5
    differ = {k for k in nv if k.split("!")[0] in ("PEERS", "WORKINGS") and
              str(nv[k]) != str(wv.get(k)) and not (num(nv[k]) is not None and num(wv.get(k)) is not None
                                                     and abs(num(nv[k]) - num(wv[k])) < 1e-12)}
    assert differ <= drawn | {"PEERS!F5"}, sorted(differ - drawn)[:10]
    assert "PEERS!F5" in differ and drawn & differ
    for key, view in (("co", "charge_offs"), ("dl", "delinquency")):
        n = len(new[view]["quarters"])
        first, count = (num(wv["WORKINGS!" + s[key + x].replace("$", "")]) for x in ("_first", "_count"))
        assert count == min(20, n) and first + count - 1 == n + 1, key       # the last 20, ending at the last
        first, count = (num(nv["WORKINGS!" + s[key + x].replace("$", "")]) for x in ("_first", "_count"))
        assert count == n and first == 2, key                                 # 48 holds them all here
    names = load_workbook(new["path"])["Peers"].defined_names
    assert names["co_c"].attr_text == "OFFSET(Workings!$%s$1,Workings!%s-1,0,Workings!%s,1)" % (
        new["charge_offs"]["helpers"]["c"], s["co_first"], s["co_count"])
