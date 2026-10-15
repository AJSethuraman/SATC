"""The peer workbook: KeyBank against its peers, from public FDIC Call Reports.

docs/PEERS-SPEC.md, "peers_workbook.py", is the contract. Python works out
every bank's books once (``books.bank_book``) and writes them to the Data tab
as values. Everything the bank looks at is on one tab, Peers, a dashboard: the
filters at the top, then a strip of headline numbers, then each view as a
compact chart or a short table, every one of them Excel formulas over the
data tabs, so the lines move as the filters are changed. The firm, 14 Oct
2026: *"it is annoying to flip between stuff for peer analysis - can't they
all exist on one tab with filters like a dashboard at the top?"*

    Peers           (red) the dashboard. Yellow cells: product, charge-off
                    rate (net or gross), Peer A and Peer B (any peer, best or
                    worst in pack), what the charts show (all quarters or the
                    last 80, 48 or 20), the delinquency measure and basis,
                    a Yes/No per peer for whether it is in the pack, and how
                    best and worst are ranked and how many banks a composite
                    or band needs. Under them: the headline numbers, the
                    charge-off chart (band, pack composite, KeyBank, Peer A
                    and B, industry), the delinquency chart, the book-size
                    chart (each bank's book indexed, the pack's band and
                    median), lagged rates, how far each rate moved, the mix
                    against the industry, and every bank's row.
    Workings        (hidden) the live helpers every view reads, one block per
                    view, and the per-year values the lagged and mix views
                    pick from by product. Headers in row 1.
    Glossary        every term, and how it is calculated
    Data            per bank, product and quarter: the TTM flows, average
                    balance, rates, kept and why, and the book index from the
                    base quarter (as reported and without acquisitions)
    DelinquencyData per product, measure, bank and quarter-end: numerator,
                    balance, quarter-end rate, four-quarter average, kept and
                    why (``peers.delinquency``, docs/V1-ROUND2-SPEC.md, A)
    Run             sources, when each cache file was pulled, fingerprints,
                    the acquisitions record, roster and the rules as stated

Every view with a pack in it follows the Yes/No cells: the charge-off and
delinquency composites and bands, the book-size band and median, the lagged
pack total and the normalised pack composite, and the headline numbers. The
old tabs drew Balances, Lagged and Normalisation as values worked out once
for every peer (or the regionals), so switching a bank off moved none of them
(the firm, 14 Oct 2026: *"peers balances do not follow peer pack stuff so i
cant turn off like JPM"*). Mix compares KeyBank with the industry, which has
no pack in it, and follows the product and Peer A and B.

The charts read named ranges (OFFSET over the Workings columns), as the
Decide tab's does: an axis cannot follow a cell, but a named range can, so a
chart draws only the quarters "Chart shows" picks, from the first quarter any
of its lines has, and its axis fits what is drawn. Only the picture: no line
depends on the window.

Lessons from ``workbook.py`` kept here: a chart combining lines and bands
has a LineChart as its base (with an AreaChart base Excel 16 drew no tick
labels); helper cells are on a tab of their own and the charts plot them
anyway (``visible_cells_only = False``); no formula refers to a row above row
1. Excel joins a line straight across #N/A, which would draw a false bridge
over quarters not kept, so every chart here is also told to show #N/A as a
gap (Excel 2016+'s ``dispNaAsBlank``; see :func:`_na_as_gap`). The band is
drawn as up-down bars between two invisible lines (10th to 90th, then 25th to
75th), not as a stacked area: an area cannot leave a gap, so Excel drew it
falling to 0 at every quarter without a band (seen in Excel, 12 Oct 2026).

The tabs can also be written into another workbook (``into``), under titles
of the caller's choosing (``names``), so the thresholds workbook and the peer
tabs are one file built by Python alone. Every formula and every defined name
names a tab through :class:`_Tabs`, never by a literal, and once the tabs are
written every reference on them is checked to name one of them
(:func:`foreign_refs`): a peer formula left saying ``Data!`` in a host with a
Data tab of its own would otherwise read that tab, silently.

Public data only. Every setting that is a judgement is a required argument.
"""
from __future__ import annotations

import datetime as dt
import functools
import hashlib
import json
import math
import re
import zipfile
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.chart import LineChart
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.data_source import AxDataSource, NumDataSource, NumRef, StrRef
from openpyxl.chart.legend import LegendEntry
from openpyxl.chart.marker import Marker
from openpyxl.chart.series import Series, SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.chart.updown_bars import UpDownBars
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties, RichTextProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from .. import VERSION
from .. import keybank_style as K
from . import roster as R
from .analyses import calendar_year, lagged_rate, mix_factor, organic_index
from . import delinquency as D
from .books import PRODUCTS, _fields, bank_book, label, quarter_end, raw_fields
from .glossary import PEER_TERMS

INPUT = PatternFill("solid", fgColor=K.INPUT, bgColor=K.INPUT)
SELF_FILL = PatternFill("solid", fgColor=K.ALERT_FG, bgColor=K.ALERT_FG)
NOTE = Font(name="Calibri", size=9, color=K.SLATE)
WORST, BEST = "Worst in pack", "Best in pack"
PACKS = ("all", "regionals")
SOURCES = ("FDIC BankFind financials: https://api.fdic.gov/banks/financials",
           "FDIC BankFind structure changes: https://api.fdic.gov/banks/history")
YES, NO = "Yes", "No"
#: The charge-off rate picker. Lagged and Mix are net by definition and say so.
RATES = ("Net", "Gross")
#: What the charts show; no line depends on it (Decide's choices, the same words).
WINDOWS = ("All quarters", "Last 80 quarters", "Last 48 quarters", "Last 20 quarters")
#: The firm, 14 Oct 2026: "it can focus - like maybe the last 48 quarters".
START_WINDOW = WINDOWS[2]
_SPANS = (100000, 80, 48, 20)

#: What ``delinquency`` must name when the tabs are built; nothing is defaulted.
D_KEYS = ("data", "products", "jump", "min_book", "exempt_self", "acquired", "measure", "basis")
NOT_BUILT = ("Delinquency was not built for this workbook: the build was told not to "
             "(delinquency=None, or --delinquency no).")

#: Every tab the builder makes, in the order it makes them. ``names`` maps
#: each of these to the title it is written under.
TABS = ("Peers", "Workings", "Glossary", "Data", "DelinquencyData", "Run")

# ---- the dashboard's grid (the Peers tab) ----
#: Columns B to Q; the left views over B:I, the right over J:Q.
FIRST_COL, LAST_COL = 2, 17
LEFT, RIGHT = (2, 9), (10, 17)
COL_WIDTH = 9.7
#: Each picker: (first column, last column) of its merged yellow cell, on R_PICK.
PICKERS = (("product", "Product", 2, 4), ("rate", "Charge-off rate", 5, 6), ("peer_a", "Peer A", 7, 8),
           ("peer_b", "Peer B", 9, 10), ("window", "Chart shows", 11, 12),
           ("dq_measure", "Delinquency measure", 13, 15), ("dq_basis", "Delinquency basis", 16, 17))
R_LABEL, R_PICK, R_RESOLVED = 4, 5, 6
R_BANKS = 8                       # bank names; their Yes/No on the row below; 15 banks a pair of rows
PER_ROW = 15                      # banks across C:Q
#: The ranking settings' yellow cells, on the row after the banks: (label span, cell column).
SETTINGS = (("rank_by", "Best and worst: ranked on the", (2, 4), 5),
            ("last_n", "rate's mean over the last", (6, 8), 9),
            ("min_obs", "quarters, if kept in at least", (10, 12), 13),
            ("min_banks", "Banks a composite or band needs", (14, 16), 17))
CHART_ROWS = 15                   # rows each top chart covers
BOOK_ROWS = 13                    # rows the book-size chart covers


def _width_cm(cols: int, width: float = COL_WIDTH) -> float:
    """``cols`` columns of ``width`` characters in cm, as Excel draws them at 100%
    (Calibri 11, a 7-pixel digit), less a margin so a chart stays inside them."""
    px = cols * math.floor((256 * width + math.floor(128 / 7)) / 256 * 7)
    return round(px / 96 * 2.54 - 0.25, 2)

_BAD_TITLE = re.compile(r"[\[\]:*?/\\]")
_STRING = re.compile(r'"(?:[^"]|"")*"')
_SHEET_REF = re.compile(r"'((?:[^']|'')+)'!|(?<![\w.'])([A-Za-z_\\][\w.]*)!")


def _quoted(title: str) -> str:
    """A tab's title as a formula names it: bare when it is letters (and _ or .)
    only, so the standalone workbook's formulas are as they always were, and
    quoted otherwise -- a space, a digit (``A1`` would read as a cell), or a
    name Excel could read as something else."""
    if re.fullmatch(r"[A-Za-z_][A-Za-z_.]*", title) and title.upper() not in ("R", "C", "TRUE", "FALSE"):
        return title
    return "'%s'" % title.replace("'", "''")


class _Tabs:
    """The title each tab is written under, and how a formula or a line of
    prose names it. ``ref("Data")`` is the prefix a formula reading the Data
    tab uses (``Data!``, or ``'Peers data'!``); ``words`` re-points the prose
    that names a tab ("the Pack tab", "DelinquencyData says why") and is the
    identity when no tab is renamed."""

    def __init__(self, names: Dict[str, str]):
        self.names = names
        moved = [t for t in TABS if names[t] != t]
        self._prose = (re.compile(r"\b(%s)\b(?=(?: and [A-Z]\w*)? tabs?\b)|\b(DelinquencyData)\b"
                                  % "|".join(sorted(moved, key=len, reverse=True)))
                       if moved else None)

    def title(self, tab: str) -> str:
        return self.names[tab]

    def ref(self, tab: str) -> str:
        return _quoted(self.names[tab]) + "!"

    def words(self, text: str) -> str:
        if self._prose is None:
            return text
        return self._prose.sub(lambda m: self.names[m.group(1) or m.group(2)], text)


def _check_names(names, into) -> Dict[str, str]:
    """``names`` checked: None for every tab under its own title, or a title
    for exactly every tab in :data:`TABS`, each a title Excel accepts, no two
    the same (Excel ignores case), and none already in ``into`` -- openpyxl
    would quietly call a second "Data" "Data1", and the formulas would still
    say Data."""
    if names is None:
        names = {t: t for t in TABS}
    elif not isinstance(names, dict):
        raise ValueError("names must be None or {tab: title} for every tab of %s, not %r" % (", ".join(TABS), names))
    else:
        missing = [t for t in TABS if t not in names]
        unknown = sorted(str(t) for t in names if t not in TABS)
        if missing or unknown:
            raise ValueError("names needs a title for exactly %s; %s" % (", ".join(TABS), "; ".join(
                x for x in ("missing: " + ", ".join(missing) if missing else "",
                            "not a tab: " + ", ".join(unknown) if unknown else "") if x)))
        names = {t: names[t] for t in TABS}
    for t, title in names.items():
        if not isinstance(title, str) or not title.strip() or len(title) > 31 or _BAD_TITLE.search(title) \
                or title.startswith("'") or title.endswith("'") or title.lower() == "history":
            raise ValueError("the %s tab's title must be 1 to 31 characters, none of []:*?/\\, not starting "
                             "or ending with ' and not History, not %r" % (t, title))
    folded = [n.lower() for n in names.values()]
    twice = sorted({n for n in names.values() if folded.count(n.lower()) > 1})
    if twice:
        raise ValueError("two tabs would have the same title (Excel ignores case): %s" % ", ".join(twice))
    if into is not None:
        if not isinstance(into, Workbook):
            raise ValueError("into must be an openpyxl Workbook, not a %s" % type(into).__name__)
        if into.write_only:
            raise ValueError("into is a write-only workbook; the peer tabs need one they can write cells into")
        wanted = {n.lower() for n in names.values()}
        clash = [s for s in into.sheetnames if s.lower() in wanted]
        if clash:
            raise ValueError("the workbook already has a tab called %s; give the peer tabs other titles "
                             "(names=...), or a formula would read the wrong one" % ", ".join(clash))
    return names


def _sheets_named(formula: str) -> List[str]:
    """Every tab a formula names, as written (unquoted), its strings left out."""
    text = _STRING.sub('""', formula)
    return [(q.replace("''", "'") if q else bare) for q, bare in _SHEET_REF.findall(text)]


def foreign_refs(wb, titles: Sequence[str]) -> List[str]:
    """Every reference on the tabs ``titles`` to a tab that is not one of them:
    cell formulas, data-validation and conditional-formatting formulas,
    hyperlinks, sheet-level defined names and chart series (values,
    categories and series names). Empty when every one stays among the peer
    tabs. Excel ignores case in a tab's name, so this does too."""
    own = {t.lower() for t in titles}
    bad = []

    def check(where, formula):
        if not isinstance(formula, str):
            return
        for name in _sheets_named(formula):
            if name.lower() not in own:
                bad.append("%s names %s: %s" % (where, name, formula[:120]))

    for title in titles:
        ws = wb[title]
        # The cells written, read in place: iter_rows would add an empty cell
        # at every gap it walks over, and an empty row to the file with it.
        for c in list(ws._cells.values()):
            if c.data_type == "f" or (isinstance(c.value, str) and c.value.startswith("=")):
                check("%s!%s" % (title, c.coordinate), c.value)
            if c.hyperlink is not None and c.hyperlink.location:
                check("%s!%s (hyperlink)" % (title, c.coordinate), c.hyperlink.location)
        for dv in ws.data_validations.dataValidation:
            for f in (dv.formula1, dv.formula2):
                check("%s data validation %s" % (title, dv.sqref), f)
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                for f in rule.formula or ():
                    check("%s conditional format %s" % (title, cf.sqref), f)
        for name, dn in getattr(ws, "defined_names", {}).items():
            check("%s defined name %s" % (title, name), dn.attr_text)
        for chart in ws._charts:
            for part in chart._charts:
                for s_ in part.series:
                    for what, src in (("values", s_.val), ("categories", s_.cat)):
                        for r_ in (getattr(src, "numRef", None), getattr(src, "strRef", None)):
                            if r_ is not None:
                                check("%s chart %s" % (title, what), r_.f)
                    if s_.tx is not None and s_.tx.strRef is not None:
                        check("%s chart series name" % title, s_.tx.strRef.f)
    return bad


def _wrap(ws, row, height=34):
    """Header labels wrap in their column instead of being cut off."""
    for cell in ws[row]:
        if cell.value is not None:
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal=cell.alignment.horizontal or "left")
    ws.row_dimensions[row].height = height


def _fingerprint(data: Dict[str, list]) -> str:
    """SHA-256 of the pull as canonical JSON: the same rows give the same print."""
    blob = json.dumps(data, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(blob).hexdigest()


_REL = "{http://schemas.openxmlformats.org/package/2006/relationships}Relationship"
_MAIN = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
_R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"


def _charts_on(z: zipfile.ZipFile, sheets: Sequence[str]) -> set:
    """The chart parts (``xl/charts/chartN.xml``) drawn on the tabs named,
    followed from the workbook to each sheet, its drawing and its charts."""
    import posixpath
    import xml.etree.ElementTree as ET

    def rels(part):
        """{Id: target part} for a part, its targets resolved against its folder."""
        folder, name = posixpath.split(part)
        rel = posixpath.join(folder, "_rels", name + ".rels")
        if rel not in z.namelist():
            return {}
        out = {}
        for r in ET.fromstring(z.read(rel)).iter(_REL):
            t = r.get("Target")
            out[r.get("Id")] = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(folder, t))
        return out

    book = rels("xl/workbook.xml")
    by_name = {s.get("name"): book[s.get(_R_ID)]
               for s in ET.fromstring(z.read("xl/workbook.xml")).iter(_MAIN + "sheet")}
    missing = [s for s in sheets if s not in by_name]
    if missing:
        raise ValueError("the workbook has no tab called %s" % ", ".join(missing))
    charts = set()
    for s in sheets:
        for drawing in rels(by_name[s]).values():
            if re.match(r"xl/drawings/drawing\d+\.xml$", drawing):
                charts.update(c for c in rels(drawing).values() if re.match(r"xl/charts/chart\d+\.xml$", c))
    return charts


def na_as_gap(path, sheets: Optional[Sequence[str]] = None) -> None:
    """Tell the charts in the file to show #N/A as a gap, not to join across it:
    every chart when ``sheets`` is None (the standalone peer workbook, which
    is all peer tabs), else only those on the tabs named -- for a workbook the
    peer tabs were written into, so its own charts are left as they were.

    openpyxl cannot write the option, so it is added to each chart's XML as
    Excel 2016+ writes it (``c16r3:dataDisplayOptions16/dispNaAsBlank`` in the
    extLst that closes ``<c:chart>``), with ``dispBlanksAs="gap"`` already set
    by openpyxl. Older Excel ignores it."""
    _na_as_gap(Path(path), sheets)


def _na_as_gap(path: Path, sheets: Optional[Sequence[str]] = None) -> None:
    """:func:`na_as_gap`, as the standalone build has always called it."""
    if sheets is not None:
        with zipfile.ZipFile(path) as z:
            only = _charts_on(z, list(sheets))
    ext = ('<extLst><ext uri="{56B9EC1D-385E-4148-901F-78D8002777C0}" '
           'xmlns:c16r3="http://schemas.microsoft.com/office/drawing/2017/03/chart">'
           '<c16r3:dataDisplayOptions16><c16r3:dispNaAsBlank val="1"/>'
           '</c16r3:dataDisplayOptions16></ext></extLst>')
    chart_ns = 'xmlns="http://schemas.openxmlformats.org/drawingml/2006/chart"'
    tmp = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            body = zin.read(item.filename)
            if re.match(r"xl/charts/chart\d+\.xml$", item.filename) and (sheets is None or item.filename in only):
                text = body.decode("utf-8")
                if "dispNaAsBlank" not in text:
                    # openpyxl writes the chart namespace as the default one;
                    # Excel keeps this option in the extLst closing <chart>.
                    if chart_ns not in text or text.count("</chart>") != 1 or "</extLst></chart>" in text:
                        raise RuntimeError("unexpected chart XML in %s" % item.filename)
                    text = text.replace("</chart>", ext + "</chart>")
                body = text.encode("utf-8")
            zout.writestr(item, body)
    tmp.replace(path)



def _written(path: Path) -> str:
    """When a cache file was last written, i.e. pulled: its modification time, local, to the second."""
    return dt.datetime.fromtimestamp(path.stat().st_mtime).isoformat(timespec="seconds")


def _whole_at_least_1(name: str, value) -> None:
    """A count setting must be a whole number of at least 1 (a bool is not one)."""
    if isinstance(value, bool) or not isinstance(value, int) or value < 1:
        raise ValueError("%s must be a whole number of at least 1, not %r" % (name, value))


def _acquisitions(mergers: Dict[str, Sequence[str]], extras: Dict[str, Sequence[str]], cert: str):
    """The bank's acquisition quarters as (quarter, declared?) in date order."""
    found = {quarter_end(q): False for q in mergers.get(cert, [])}
    for q in extras.get(cert, []):
        found.setdefault(quarter_end(q), True)
    return [(label(d.replace("-", "")), found[d]) for d in sorted(found)]


def _check_delinquency(spec: Optional[dict], roster: List[dict], me: dict) -> Optional[dict]:
    """None for no delinquency tabs; otherwise the stated settings, checked,
    with every bank's rows worked out once (``delinquency.bank_delinquency``,
    per product, measure and bank) under ``books`` and the quarter-ends under
    ``grid``.

    Refused (ValueError), each with its reason, whatever the type of the
    value given: anything but None or a dict; a key missing or not known
    (every one is a judgement or the data, so none is filled in); products
    that are not distinct names from ``delinquency.FIELDS``; a measure or
    basis that is not one of the names the module builds; ``exempt_self``
    that is not True or False; a jump or minimum book that is not a number
    (a bool or a string is not one) or is below 0 (refused by the module);
    data that is not ``{REPDTE: [rows]}``, or has no quarter; a quarter with
    no rows (nobody filed, which is not a field missing); and a field of a
    product absent from every row of a quarter. That last is a field that was
    not pulled, which is not the same as a field not filed (a blank value,
    which the module reads as not filed): ``fdic.pull`` refuses the same."""
    if spec is None:
        return None
    if not isinstance(spec, dict):
        raise ValueError("delinquency must be None (not built) or a dict of %s, not %r" % (", ".join(D_KEYS), spec))
    missing = [k for k in D_KEYS if k not in spec]
    unknown = sorted(k for k in spec if k not in D_KEYS)
    if missing or unknown:
        raise ValueError("delinquency needs exactly %s; %s" % (", ".join(D_KEYS), "; ".join(
            x for x in ("missing: " + ", ".join(missing) if missing else "",
                        "not known: " + ", ".join(unknown) if unknown else "") if x)))
    products = list(spec["products"]) if not isinstance(spec["products"], str) else [spec["products"]]
    if not products or any(p not in D.FIELDS for p in products) or len(set(products)) != len(products):
        raise ValueError("delinquency products must be distinct names from %s, not %r"
                         % (list(D.FIELDS), spec["products"]))
    if not isinstance(spec["measure"], str) or spec["measure"] not in D.MEASURES:
        raise ValueError("delinquency measure must be one of %s, not %r" % (list(D.MEASURES), spec["measure"]))
    if not isinstance(spec["basis"], str) or spec["basis"] not in D.BASES:
        raise ValueError("delinquency basis must be one of %s, not %r" % (list(D.BASES), spec["basis"]))
    if not isinstance(spec["exempt_self"], bool):
        raise ValueError("delinquency exempt_self must be True or False, not %r" % (spec["exempt_self"],))
    for name in ("jump", "min_book"):
        if isinstance(spec[name], bool) or not isinstance(spec[name], (int, float)):
            raise ValueError("delinquency %s must be a number, not %r" % (name, spec[name]))
    if not isinstance(spec["acquired"], dict) or any(p not in spec["acquired"] for p in products):
        raise ValueError("delinquency acquired must be {product: {cert: [quarter-ends]}} for every "
                         "product, not %r" % (spec["acquired"],))
    data = spec["data"]
    if not isinstance(data, dict):
        raise ValueError("delinquency data must be {REPDTE: [rows]} as fdic.pull returns it, not a %s"
                         % type(data).__name__)
    if not data:
        raise ValueError("delinquency data has no quarter-ends, so nothing can be drawn")
    empty = [q for q in sorted(data) if not data[q]]
    if empty:
        raise ValueError("delinquency data has no rows (no filer) in %d quarter(s), first %s: a quarter-end "
                         "nobody filed cannot be drawn; pull it again or leave it out" % (len(empty), empty[0]))
    for p in products:
        for f in D.fields_for([p]):
            lacking = [q for q in sorted(data) if not any(f in row for row in data[q])]
            if lacking:
                raise ValueError("delinquency for %s needs %s, which the data does not carry in %d quarter(s), "
                                 "first %s: pull it with the delinquency fields" % (p, f, len(lacking), lacking[0]))
    books = {p: {m: {b["cert"]: D.bank_delinquency(data, b["cert"], p, m, spec["jump"], spec["min_book"],
                                                   spec["exempt_self"] and b is me,
                                                   acquired=spec["acquired"][p].get(b["cert"], []))
                     for b in roster} for m in D.MEASURES} for p in products}
    grid = [r["q"] for r in books[products[0]][spec["measure"]][me["cert"]]]
    return dict(spec, products=products, books=books, grid=grid)





#: Rows each note keeps: a note is one line per row, never a wrapped cell, so
#: no note on one half of the tab makes a row of the other half taller.
NOTE_ROWS = dict(note1=3, book_note=5, lagged_note=3, norm_note=3, mix_note=4, banks_note=4)
NOTE_CHARS = 88                   # characters a line of a note holds across eight columns


def layout(n_peers: int, n_banks: int, n_years: int) -> dict:
    """Where everything sits on the Peers tab for ``n_peers`` peers,
    ``n_banks`` banks on the roster and ``n_years`` mix years: one place for
    the builder, the tests and anyone photographing it."""
    chunks = max(1, math.ceil(n_peers / PER_ROW))
    r_set = R_BANKS + 2 * chunks
    v1 = r_set + 3
    note1 = v1 + 4 + CHART_ROWS
    v2 = note1 + NOTE_ROWS["note1"] + 1
    book_table = v2 + 1 + BOOK_ROWS
    book_note = book_table + 5
    norm = v2 + 2 + 4 + NOTE_ROWS["lagged_note"] + 1
    norm_note = norm + 6
    end2 = max(book_note + NOTE_ROWS["book_note"], norm_note + NOTE_ROWS["norm_note"])
    v3 = end2 + 1
    return dict(chunks=chunks, settings=r_set, best_worst=r_set + 1,
                v1=v1, tiles=v1 + 1, chart1=v1 + 4, note1=note1,
                v2=v2, chart2=v2 + 1, book_table=book_table, book_note=book_note,
                lagged=v2 + 1, lagged_note=v2 + 6, norm=norm, norm_table=norm + 1, norm_note=norm_note,
                v3=v3, mix=v3 + 1, mix_peers=v3 + 2 + n_years, mix_note=v3 + 3 + n_years,
                banks=v3 + 1, banks_note=v3 + 2 + n_banks,
                end=max(v3 + 3 + n_years + NOTE_ROWS["mix_note"], v3 + 2 + n_banks + NOTE_ROWS["banks_note"]) - 1)


def include_cell(j: int) -> str:
    """The Yes/No cell of the j-th peer (roster order, peers only) on Peers."""
    chunk, at = divmod(j, PER_ROW)
    return "%s%d" % (L(3 + at), R_BANKS + 2 * chunk + 1)


def picker_cell(name: str) -> str:
    """A picker's yellow cell on Peers (the top-left of its merged range)."""
    col = next(first for key, _, first, _ in PICKERS if key == name)
    return "%s%d" % (L(col), R_PICK)


def setting_cell(name: str, n_peers: int) -> str:
    col = next(c for key, _, _, c in SETTINGS if key == name)
    return "%s%d" % (L(col), layout(n_peers, 1, 0)["settings"])


class _Block:
    """Columns on Workings, allocated left to right: a header in row 1, the
    values from row 2."""

    def __init__(self, ws, start: int):
        self.ws, self.next, self.col, self.width = ws, start, {}, {}

    def add(self, name: str, width: int = 1, heads: Optional[Sequence[str]] = None) -> int:
        self.col[name], self.width[name] = self.next, width
        for j, h in enumerate(heads if heads is not None else [name]):
            self.ws.cell(1, self.next + j, h)
        self.next += width
        return self.col[name]

    def at(self, name: str, r: int, j: int = 0) -> str:
        return "%s%d" % (L(self.col[name] + j), r)

    def fixed(self, name: str, r: int, j: int = 0) -> str:
        return "$%s$%d" % (L(self.col[name] + j), r)

    def across(self, name: str, r: int) -> str:
        """The block's cells on row r: one per bank or peer."""
        return "$%s%d:$%s%d" % (L(self.col[name]), r, L(self.col[name] + self.width[name] - 1), r)

    def down(self, name: str, r1: int, r2: int, j: int = 0) -> str:
        c = L(self.col[name] + j)
        return "$%s$%d:$%s$%d" % (c, r1, c, r2)


def _small_text(size=800, bold=False):
    """Chart text at ``size`` hundredths of a point."""
    cp = CharacterProperties(sz=size, b=bold, solidFill=K.SLATE)
    return RichText(bodyPr=RichTextProperties(), p=[Paragraph(pPr=ParagraphProperties(defRPr=cp),
                                                             endParaRPr=cp)])


class _Chart:
    """A chart over named ranges (OFFSET over Workings), the band as up-down
    bars between invisible lines, then the lines. ``name(col)`` defines a
    name on the Peers tab for a Workings column and returns how a series
    refers to it."""

    def __init__(self, cats: str):
        self.cats, self.idx, self.groups, self.hide = cats, 0, [], []

    def _series(self, f: str, label) -> Series:
        s_ = Series(val=NumDataSource(numRef=NumRef(f=f)), tx=label)
        s_.cat = AxDataSource(strRef=StrRef(f=self.cats))
        s_.idx = s_.order = self.idx
        self.idx += 1
        s_.smooth = False
        s_.marker = Marker(symbol="none")
        return s_

    def band(self, lo: str, hi: str, fill: str):
        lc = LineChart()
        for f in (lo, hi):
            s_ = self._series(f, SeriesLabel(v=" "))
            s_.graphicalProperties.line.noFill = True
            self.hide.append(s_.idx)
            lc.series.append(s_)
        lc.upDownBars = UpDownBars(gapWidth=0)
        for which in ("upBars", "downBars"):
            gp = GraphicalProperties(solidFill=fill)
            gp.line.noFill = True
            setattr(lc.upDownBars, which, ChartLines(spPr=gp))
        self.groups.append(lc)

    def lines(self, specs):
        """``specs``: (f, label, colour, width EMU, dash or None)."""
        lc = LineChart()
        for f, lbl, colour, width, dash in specs:
            s_ = self._series(f, lbl)
            s_.graphicalProperties.line.solidFill = colour
            s_.graphicalProperties.line.width = width
            if dash:
                s_.graphicalProperties.line.dashStyle = dash
            lc.series.append(s_)
        self.groups.append(lc)

    def finish(self, number_format: str, skip: int, height: float, width: float) -> LineChart:
        base = self.groups[0]
        for g in self.groups[1:]:
            base += g
        base.height, base.width = height, width
        base.x_axis.delete = False
        base.y_axis.delete = False
        base.y_axis.number_format = number_format
        base.y_axis.majorGridlines = ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill="E7E2D8")))
        base.x_axis.tickLblSkip = skip
        base.x_axis.tickMarkSkip = skip
        base.x_axis.tickLblPos = "low"
        base.x_axis.txPr = _small_text()
        base.y_axis.txPr = _small_text()
        base.legend.position = "b"
        base.legend.overlay = False
        base.legend.txPr = _small_text()
        # The band's invisible ends would be blank keys: the note names the shading.
        base.legend.legendEntry = [LegendEntry(idx=i, delete=True) for i in self.hide]
        base.display_blanks = "gap"
        base.visible_cells_only = False          # its data is on the hidden Workings tab
        return base


#: Line styles: pack (ink), KeyBank (red), Peer A (slate dash), Peer B (blue dot), industry (green).
STYLE = {"pack": (K.INK, 28575, None), "self": (K.KEY_RED, 28575, None), "a": (K.SLATE, 15875, "dash"),
         "b": ("1F5F8B", 15875, "sysDot"), "industry": ("2E7D4F", 19050, None)}
BAND_LIGHT, BAND_DARK = K.MIST, "CFC6B6"

F_TITLE = Font(name="Arial", size=16, bold=True, color=K.INK)
F_SUB = Font(name="Calibri", size=10, color=K.SLATE)
F_LAB = Font(name="Arial", size=9, bold=True, color=K.SLATE)
F_BODY = Font(name="Calibri", size=10, color=K.INK_TEXT)
F_BOLD = Font(name="Calibri", size=10, bold=True, color=K.INK_TEXT)
F_SELF = Font(name="Calibri", size=10, bold=True, color=K.CRIMSON)
F_TILE = Font(name="Arial", size=16, bold=True, color=K.INK)
F_PICK = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
EDGE = Side(style="thin", color=K.INPUT_EDGE)
BOX = Border(EDGE, EDGE, EDGE, EDGE)
HAIR = Border(bottom=Side(style="thin", color="E7E2D8"))


def _yellow(c, merge_to=None):
    c.fill = INPUT
    c.font = F_PICK
    c.border = BOX
    c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    if merge_to is not None:
        ws = c.parent
        for col in range(c.column + 1, merge_to + 1):
            ws.cell(c.row, col).border = BOX
            ws.cell(c.row, col).fill = INPUT
        ws.merge_cells(start_row=c.row, start_column=c.column, end_row=c.row, end_column=merge_to)


def _merged(ws, row, first, last, value, font=F_BODY, height=None, wrap=True, align="left"):
    if last > first:
        ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    c = ws.cell(row, first, value)
    c.font = font
    c.alignment = Alignment(wrap_text=wrap, vertical="top" if wrap else "center", horizontal=align)
    if height:
        ws.row_dimensions[row].height = height
    return c


def _note(ws, row, span, room, text, cond=None, alt=None, words=None):
    """A note as lines, one per row from ``row``, none wrapped: ``room`` rows
    at most (more is a RuntimeError, so a longer note cannot quietly run into
    what is under it). With ``cond`` (a formula), each line shows ``alt``'s
    line when it is TRUE. ``words`` (``_Tabs.words``) renames the tabs a line
    names after the lines are cut, so a note breaks in the same places
    whatever its tabs are called, and the cells are the same in both."""
    import textwrap
    main = [(words or str)(x) for x in textwrap.wrap(text, NOTE_CHARS)]
    other = textwrap.wrap(alt, NOTE_CHARS) if alt else []
    if max(len(main), len(other)) > room:
        raise RuntimeError("a note needs %d lines and %d are kept for it: %s" % (max(len(main), len(other)), room,
                                                                               text[:60]))
    for i in range(room):
        a = main[i] if i < len(main) else ""
        if cond:
            b = other[i] if i < len(other) else ""
            v = '=IF(%s,"%s","%s")' % (cond, b.replace('"', '""'), a.replace('"', '""')) if (a or b) else None
        else:
            v = a or None
        _merged(ws, row + i, span[0], span[1], v, NOTE, wrap=False)


def _section(ws, row, first, last, value):
    """A view's title: bold, over a medium ink rule (keybank_style.quiet_section)."""
    ws.merge_cells(start_row=row, start_column=first, end_row=row, end_column=last)
    c = ws.cell(row, first, value)
    c.font = Font(name="Arial", size=11, bold=True, color=K.INK)
    c.alignment = Alignment(vertical="bottom")
    for col in range(first, last + 1):
        ws.cell(row, col).border = Border(bottom=Side(style="medium", color=K.INK))
    ws.row_dimensions[row].height = 18


def _table_head(ws, row, first, heads, right_from=1, height=26):
    """canvas column headers (keybank_style.canvas_header), wrapped."""
    for i, h in enumerate(heads):
        c = ws.cell(row, first + i, h)
        c.fill = K.solid(K.CANVAS)
        c.font = Font(name="Arial", bold=True, size=8, color=K.SLATE)
        c.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="right" if i >= right_from else "left")
        c.border = Border(bottom=Side(style="thin", color=K.STONE))
    ws.row_dimensions[row].height = height


def _tile(ws, row, first, label_, value, caption, fmt):
    """A headline number: label, number, caption, on canvas under a red rule."""
    last = first + 1
    for r, v, font in ((row, label_, F_LAB), (row + 1, value, F_TILE), (row + 2, caption, NOTE)):
        ws.merge_cells(start_row=r, start_column=first, end_row=r, end_column=last)
        c = ws.cell(r, first, v)
        c.font = font
        c.fill = K.solid(K.CANVAS)
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=r == row + 2)
        ws.cell(r, last).fill = K.solid(K.CANVAS)
    ws.cell(row + 1, first).number_format = fmt
    for col in (first, last):
        ws.cell(row, col).border = Border(top=Side(style="medium", color=K.KEY_RED))



def _delinquency_data(dd, dq: dict, roster: List[dict]) -> int:
    """DelinquencyData: one row per product, measure (``delinquency.MEASURES``
    order), bank (roster order) and quarter-end, in that order: the live
    formulas find a row by position, so the order is load-bearing. Each row is
    ``bank_delinquency``'s, plus whether the four-quarter average was kept and
    whether the jump blanked it, so no formula has to read a blank cell (Excel
    reads one as 0). Returns the number of rows."""
    products, books = dq["products"], dq["books"]
    K.header_row(dd, 1, ["Product", "Measure", "Bank", "Certificate", "Quarter-end", "Numerator ($000)",
                         "Balance ($000)", "Quarter-end rate (%)", "Kept", "Why not kept",
                         "Four-quarter average (%)", "Four-quarter average kept", "Why no four-quarter average",
                         "Left out of the composite's four-quarter average (jump)"],
                 right_from=5)
    _wrap(dd, 1, 46)
    for p in products:
        for m in D.MEASURES:
            for b in roster:
                for r in books[p][m][b["cert"]]:
                    dd.append([p, m, b["short"], b["cert"], r["q"], r["numerator"], r["balance"], r["rate"],
                               r["kept"], r["why"], r["avg4"], r["avg4"] is not None, r["why_avg4"],
                               (r["why_avg4"] or "").startswith(D.LEFT_OUT)])
    n = len(products) * len(D.MEASURES) * len(roster) * len(dq["grid"])
    for row in dd.iter_rows(min_row=2, max_row=n + 1, min_col=6, max_col=11):
        for c in row:
            if c.column in (6, 7):
                c.number_format = "#,##0"
            elif c.column in (8, 11):
                c.number_format = "0.0000"
    for c_, w in zip("ABCDEFGHIJKLMN", (14, 20, 16, 13, 12, 14, 16, 14, 8, 40, 16, 16, 60, 18)):
        dd.column_dimensions[c_].width = w
    K.freeze_below(dd, 1)
    return n


def _delinquency_block(wk, dq: dict, roster: List[dict], me: dict, peers: List[dict], start: int, x: dict):
    """The delinquency view's live helpers on Workings, one row per quarter-end
    from row 2, driven by the Peers tab's product, delinquency measure and
    basis, Yes/No cells, Peer A and B and banks needed (``x`` holds those
    references):

    - per bank: kept on the quarter-end basis, and its value on the basis
      picked (#N/A when there is none, or the product was not built for
      delinquency);
    - per peer switched on: the numerator and balance when kept (for the
      composite) and the value on the basis (for the band);
    - the composite: summed numerators / summed balances of the peers switched
      on and kept, #N/A below the banks needed; its four-quarter average at i
      is the mean of the four composite quarter-end rates i-3..i, each
      recomputed without the peers whose own four-quarter average at i the
      jump blanked, only when all four exist on consecutive quarter-ends and
      each has the banks needed (``composite_delinquency``);
    - the band: PERCENTILE of the values of the peers switched on, #N/A below
      the banks needed (``band_delinquency``);
    - KeyBank, Peer A and Peer B: their values on the basis."""
    products, grid = dq["products"], dq["grid"]
    NB, NP, DQ, NM = len(roster), len(peers), len(grid), len(D.MEASURES)
    bidx = {b["cert"]: i for i, b in enumerate(roster)}
    DD = x["DD"]
    sel, prod, basis = x["dq_sel"], x["dq_p"], x["dq_b"]

    def at(pi, mi, bi, k):
        return 2 + ((pi * NM + mi) * NB + bi) * DQ + k

    def pick(letter, bi, k):
        """The DelinquencyData cell in this column for the picked product and measure."""
        cells = ",".join("%s$%s$%d" % (DD, letter, at(pi, mi, bi, k)) for pi in range(len(products))
                         for mi in range(NM))
        return "CHOOSE(%s,%s)" % (sel, cells)

    def pick_balance(bi, k):
        """The balance for the picked product: the same for every measure, so the first's."""
        cells = ",".join("%s$G$%d" % (DD, at(pi, 0, bi, k)) for pi in range(len(products)))
        return "CHOOSE(%s,%s)" % (prod, cells)

    bl = _Block(wk, start)
    bl.add("q", heads=["Quarter-end (delinquency)"])
    bl.add("four", heads=["Four consecutive quarter-ends"])
    for name, banks in (("uk", roster), ("uv", roster), ("ok", peers), ("n", peers), ("b", peers), ("v", peers),
                        ("s", peers)):
        bl.add(name, len(banks), ["%s %s" % (b["short"], name) for b in banks])
    for name, h in (("banks", "Banks in the quarter-end composite"), ("cq", "Composite, quarter-end"),
                    ("ca", "Composite, four-quarter average"), ("c", "Pack composite"),
                    ("bandn", "Banks in the band"), ("p10", "10th"), ("p25", "25th"), ("p75", "75th"),
                    ("p90", "90th"), ("k", me["short"])):
        bl.add(name, heads=[h])
    bl.add("a", heads=['="Peer A: "&%s' % x["a_name"]])
    bl.add("bb", heads=['="Peer B: "&%s' % x["b_name"]])
    bl.add("drawn", heads=["Row drawn"])
    mb = x["mb"]
    follows = D._consecutive(grid)
    me_i = bidx[me["cert"]]
    for k, q in enumerate(grid):
        r = 2 + k
        four = k >= 3 and all(follows[i] for i in range(k - 2, k + 1))
        wk.cell(r, bl.col["q"], q)
        wk.cell(r, bl.col["four"], four)
        for j in range(NB):
            uk = bl.at("uk", r, j)
            wk.cell(r, bl.col["uk"] + j, "=IF(%s>0,%s,FALSE)" % (prod, pick("I", j, k)))
            wk.cell(r, bl.col["uv"] + j, "=IF({p}=0,NA(),IF({b}=1,IF({uk},{rate},NA()),IF({k4},{avg},NA())))"
                    .format(p=prod, b=basis, uk=uk, rate=pick("H", j, k), k4=pick("L", j, k), avg=pick("K", j, k)))
        for j, b in enumerate(peers):
            i = bidx[b["cert"]]
            inc = x["inc"][j]
            ok = bl.at("ok", r, j)
            uv = bl.at("uv", r, i)
            wk.cell(r, bl.col["ok"] + j, '=AND(%s="%s",%s)' % (inc, YES, bl.at("uk", r, i)))
            wk.cell(r, bl.col["n"] + j, '=IF(%s,%s,"")' % (ok, pick("F", i, k)))
            wk.cell(r, bl.col["b"] + j, '=IF(%s,%s,"")' % (ok, pick_balance(i, k)))
            wk.cell(r, bl.col["v"] + j, '=IF(AND(%s="%s",ISNUMBER(%s)),%s,"")' % (inc, YES, uv, uv))
            # 1 when this peer stays in the composite's four-quarter average
            # at this quarter-end, 0 when the jump blanked its own.
            wk.cell(r, bl.col["s"] + j, "=IF(%s=0,1,IF(%s,0,1))" % (prod, pick("N", i, k)))
        banks, bandn = bl.at("banks", r), bl.at("bandn", r)
        wk.cell(r, bl.col["banks"], "=COUNT(%s)" % bl.across("n", r))
        wk.cell(r, bl.col["cq"], "=IF({n}>={mb},100*SUM({x})/SUM({b}),NA())"
                .format(n=banks, mb=mb, x=bl.across("n", r), b=bl.across("b", r)))
        if four:
            stay = bl.across("s", r)
            cnt = ["COUNTIFS(%s,1,%s,TRUE)" % (stay, bl.across("ok", w)) for w in range(r - 3, r + 1)]
            rate = ["100*SUMIFS(%s,%s,1)/SUMIFS(%s,%s,1)" % (bl.across("n", w), stay, bl.across("b", w), stay)
                    for w in range(r - 3, r + 1)]
            wk.cell(r, bl.col["ca"], "=IF(AND(%s),AVERAGE(%s),NA())"
                    % (",".join("%s>=%s" % (c, mb) for c in cnt), ",".join(rate)))
        else:
            wk.cell(r, bl.col["ca"], "=NA()")
        wk.cell(r, bl.col["c"], "=IF(%s=1,%s,%s)" % (basis, bl.at("cq", r), bl.at("ca", r)))
        wk.cell(r, bl.col["bandn"], "=COUNT(%s)" % bl.across("v", r))
        for p in (10, 25, 75, 90):
            wk.cell(r, bl.col["p%d" % p], "=IF({n}>={mb},PERCENTILE({v},{p}),NA())"
                    .format(n=bandn, mb=mb, v=bl.across("v", r), p=p / 100))
        wk.cell(r, bl.col["k"], "=" + bl.at("uv", r, me_i))
        for name, num in (("a", x["a_no"]), ("bb", x["b_no"])):
            wk.cell(r, bl.col[name], "=IFERROR(INDEX({row},{n}),NA())".format(row=bl.across("uv", r), n=num))
        wk.cell(r, bl.col["drawn"], '=IF(OR(%s),%d,"")' % (",".join(
            "ISNUMBER(%s)" % bl.at(n_, r) for n_ in ("c", "k", "a", "bb", "p10")), r))
    return bl


def _rolled_back(build):
    """``build``, which on any failure takes back out of ``into`` every tab it
    added, so a refused or failed build leaves the caller's workbook as it was."""
    @functools.wraps(build)
    def wrapper(*args, **kwargs):
        into = kwargs.get("into")
        had = list(into.sheetnames) if isinstance(into, Workbook) else None
        try:
            return build(*args, **kwargs)
        except BaseException:
            if had is not None:
                for title in [t for t in into.sheetnames if t not in had]:
                    into.remove(into[title])
            raise
    return wrapper



@_rolled_back
def build_peers_workbook(out, data: Dict[str, list], roster: List[dict], products: Sequence[str],
                         jump: float, min_book: float, pack: str,
                         mergers: Dict[str, Sequence[str]], extra_acquisitions: Dict[str, Sequence[str]],
                         base_quarter: str, *,
                         exempt_self: bool, min_banks: int, rank_by: str, last_n: int, min_obs: int,
                         index_min_book: float, lagged_years: Sequence[int], lag_years: int,
                         normalisation_years: Sequence[int], mix_base_year: int,
                         industry: Optional[Dict[str, List[dict]]],
                         delinquency: Optional[dict],
                         acquired: Dict[str, Dict[str, Sequence[str]]],
                         sources: Iterable[str] = (), cache_dir=None, mergers_file=None,
                         peer_a: str = WORST, peer_b: str = BEST,
                         product: Optional[str] = None,
                         into: Optional[Workbook] = None,
                         names: Optional[Dict[str, str]] = None) -> dict:
    """Write the peer workbook to ``out`` and return what went in.

    ``data`` is ``{REPDTE: [rows]}`` as ``fdic.pull`` returns it; ``roster`` the
    banks as ``roster.load`` returns them. The judgements, none defaulted:

    - ``jump``, ``min_book`` ($000) and ``exempt_self``: the rules of
      ``books.bank_book``; ``exempt_self`` lifts the minimum for KeyBank;
    - ``acquired``: ``{product: {cert: [YYYYMMDD]}}``, the quarters whose
      recorded acquisitions brought in at least the acquisition share
      (``books.acquired_quarters``), for every product built; the delinquency
      spec carries the same under its ``acquired`` key;
    - ``pack``: 'all' or 'regionals', the peers whose Yes/No starts Yes;
    - ``min_banks``: how many kept banks a composite or band needs (a live
      cell on Peers, starting here);
    - ``rank_by`` ('gross' or 'net'), ``last_n``, ``min_obs``: how best and
      worst in pack are found (``books.best_worst``), live cells on Peers;
      the charge-off rate picker starts on ``rank_by``'s rate;
    - ``mergers`` (``{cert: [quarter-end dates]}``, as ``fdic.classify``
      returns) and ``extra_acquisitions`` (declared purchases the FDIC record
      does not show), ``base_quarter`` and ``index_min_book`` ($000, the
      smallest book at the base that is indexed): the book-size view. The
      base must be one of the TTM quarters, as the view is drawn on them;
    - ``lagged_years`` and ``lag_years``: the lagged view;
    - ``normalisation_years``: (base year, end year) for how far each rate moved;
    - ``mix_base_year`` and ``industry`` (``{product: books.industry(...)}``
      over every filer, or None, in which case the industry line and the mix
      say it was not supplied rather than standing the roster in for the
      industry);
    - ``delinquency``: None, and the delinquency view and DelinquencyData say
      it was not built; or a dict of exactly these keys, none defaulted (see
      :func:`_check_delinquency`): ``data`` (``{REPDTE: [rows]}`` carrying
      ``delinquency.FIELDS`` for its products, usually the same pull as
      ``data``), ``products`` (names from ``delinquency.FIELDS``), ``jump``,
      ``min_book`` ($000) and ``exempt_self`` (the rules of
      ``delinquency.bank_delinquency``, stated for delinquency rather than
      borrowed from the charge-off ones), ``acquired``, and ``measure`` and
      ``basis``, the delinquency pickers' starting positions.

    The product picker lists the charge-off products, then any delinquency
    product not among them; a view whose product was not built says so in
    words. ``peer_a``, ``peer_b`` and ``product`` are only the pickers'
    starting positions, as "Chart shows" starts at the last 48 quarters (the
    firm's own); nothing is computed from them that the bank cannot change.

    Provenance for the Run tab: ``sources`` (lines of text), ``cache_dir``
    (the folder ``fdic.pull`` read; each file's name, when it was written,
    which is when it was pulled, and its SHA-256) and ``mergers_file`` (the
    acquisitions file ``mergers`` came from, named with its date and SHA-256).
    The ``mergers`` given are fingerprinted whatever their source.

    Where the tabs go: by default a workbook of their own, saved to ``out``.
    ``into`` (an openpyxl Workbook) writes them into that workbook instead,
    after its own tabs, and saves nothing (``out`` must then be None); the
    caller saves it and then runs :func:`na_as_gap` on the file, over the tabs
    returned under ``sheets``, as the standalone build does to its own. In
    that mode the workbook's own settings are left as they are: which tab is
    active, its calculation and document properties. ``names`` (``{tab:
    title}`` for exactly every tab in :data:`TABS`) writes each tab under the
    title given; every formula, defined name and line that names a tab names
    it by that title. A title the workbook already has is refused, as is the
    same title twice. If the build fails after it has begun adding tabs, it
    takes them out again before the error reaches the caller. Before
    returning, it checks every reference on the peer tabs names a peer tab
    (:func:`foreign_refs`), and raises RuntimeError, saving nothing, if one
    does not.

    Refused (ValueError), each with its reason: a setting outside its range, a
    base quarter that is not in the data or is not a TTM quarter, a product
    whose fields the data does not carry, a roster with no peers, a regionals
    pack with no regionals, and banks or products that do not share one
    quarter calendar; and for where the tabs go, an ``into`` that is not a
    Workbook (or is write-only), an ``out`` given with it or missing without
    it, and ``names`` as above.
    """
    if into is None and out is None:
        raise ValueError("out is needed: the peer workbook is written there unless into names a workbook")
    if into is not None and out is not None:
        raise ValueError("into writes the peer tabs into that workbook and saves nothing, so out must be None, "
                         "not %r; save the workbook yourself" % (out,))
    t = _Tabs(_check_names(names, into))
    if pack not in PACKS:
        raise ValueError("pack must be 'all' or 'regionals', not %r" % (pack,))
    if rank_by not in ("gross", "net"):
        raise ValueError("rank_by must be 'gross' or 'net', not %r" % (rank_by,))
    products = list(products)
    if not products or any(p not in PRODUCTS for p in products) or len(set(products)) != len(products):
        raise ValueError("products must be distinct names from %s, not %r" % (sorted(PRODUCTS), products))
    if len(normalisation_years) != 2:
        raise ValueError("normalisation_years is (base year, end year)")
    for name, value in (("min_banks", min_banks), ("last_n", last_n), ("min_obs", min_obs)):
        _whole_at_least_1(name, value)
    product = product or products[0]
    if product not in products:
        raise ValueError("the starting product %r is not one of %r" % (product, products))
    base_repdte = quarter_end(base_quarter).replace("-", "")       # refuses '2019Q5'
    if base_repdte not in data:
        qs = sorted(data)
        raise ValueError("the base quarter %s is not in the data (%s), so no book can be indexed to it"
                         % (base_quarter, "%s to %s" % (label(qs[0]), label(qs[-1])) if qs else "no quarters"))
    # A field absent from a quarter's every row would be read as 0 for every
    # bank: an empty book, not a missing one. fdic.pull refuses the same.
    for p in products:
        for f in raw_fields(_fields(p)):
            lacking = [q for q in sorted(data) if not any(f in row for row in data[q])]
            if lacking:
                raise ValueError("%s needs %s, which the data does not carry in %d quarter(s), first %s: "
                                 "pull it with that field" % (p, f, len(lacking), lacking[0]))
    me = R.self_bank(roster)
    peers = R.peers(roster)
    if not peers:
        raise ValueError("the roster has no peers, so there is nothing to compare %s with" % me["short"])
    if pack == "regionals" and not R.regionals(roster):
        raise ValueError("pack 'regionals' was asked for, but no peer on the roster is marked regional")
    shorts = [b["short"] for b in peers]
    for pick in (peer_a, peer_b):
        if pick not in shorts + [WORST, BEST]:
            raise ValueError("a peer pick must be a peer's short name, %r or %r, not %r" % (WORST, BEST, pick))
    in_pack = {b["cert"] for b in (R.peers(roster) if pack == "all" else R.regionals(roster))}
    dq = _check_delinquency(delinquency, roster, me)

    books = {p: {b["cert"]: bank_book(data, b["cert"], p, jump, min_book,
                                      exempt_min_book=exempt_self and b is me,
                                      acquired=acquired[p].get(b["cert"], []))
                 for b in roster} for p in products}
    grid = [r["quarter"] for r in books[products[0]][me["cert"]]]
    for p in products:
        for b in roster:
            if [r["quarter"] for r in books[p][b["cert"]]] != grid:
                raise ValueError("%s %s does not share the quarter calendar" % (b["short"], p))
    NB, NP, Q = len(roster), len(peers), len(grid)
    if Q < 1:
        raise ValueError("the data has no TTM quarter: it needs at least five quarter-ends")
    base_lbl = label(base_repdte)
    if base_lbl not in grid:
        raise ValueError("the base quarter %s is not one of the TTM quarters (%s to %s): the book index is drawn "
                         "on the quarters the rates are, so the base must be one of them"
                         % (base_lbl, grid[0], grid[-1]))
    base_k = grid.index(base_lbl)
    bidx = {b["cert"]: i for i, b in enumerate(roster)}
    me_i = bidx[me["cert"]]
    peer_i = [bidx[b["cert"]] for b in peers]
    all_products = products + [p for p in (dq["products"] if dq else []) if p not in products]
    y0, y1 = normalisation_years
    years = sorted({int(q[:4]) for q in data if all("%s%s" % (q[:4], md) in data
                                                    for md in ("0331", "0630", "0930", "1231"))})
    years = [y for y in years if y >= mix_base_year]
    NY = len(years)
    lay = layout(NP, NB, NY)

    if into is None:
        wb = Workbook()
        pe = wb.active
        pe.title = t.title("Peers")
    else:
        wb = into
        pe = wb.create_sheet(t.title("Peers"))
    wk, gl, da, dd, run = (wb.create_sheet(t.title(x)) for x in TABS[1:])
    wk.sheet_state = "hidden"
    PE, WK, DA, DD = t.ref("Peers"), t.ref("Workings"), t.ref("Data"), t.ref("DelinquencyData")
    pe.sheet_properties.tabColor = K.KEY_RED
    for ws in (pe, gl, run):
        K.hide_gridlines(ws)

    # ---- the book index: per product and bank, from the base on ----
    acq_of = {b["cert"]: _acquisitions(mergers, extra_acquisitions, b["cert"]) for b in roster}
    indexed, index_note = {}, {}
    for p in products:
        for b in roster:
            try:
                idx = organic_index(data, b["cert"], p, base_quarter, [q for q, _ in acq_of[b["cert"]]])
            except ValueError:
                # The base quarter is in the data (checked above), so this is
                # the bank having no filing or a zero book there.
                idx = None
            at_base = next((x["balance"] for x in idx if x["date"] == quarter_end(base_quarter)), None) \
                if idx else None
            if idx is None:
                index_note[p, b["cert"]] = "no book"
            elif at_base < index_min_book:
                index_note[p, b["cert"]] = "too small"
            else:
                index_note[p, b["cert"]] = "indexed"
                indexed[p, b["cert"]] = {x["quarter"]: x for x in idx if x["date"] >= quarter_end(base_quarter)}

    # ---- Data: one row per product, bank (roster order) and quarter ----
    # The live formulas find a row by position: product block, then bank
    # block, then quarter, so the order here is load-bearing.
    K.header_row(da, 1, ["Product", "Bank", "Certificate", "Quarter", "Gross charge-offs, TTM ($000)",
                         "Recoveries, TTM ($000)", "Average balance ($000)", "Gross rate (%)",
                         "Net rate (%)", "Kept", "Why not kept", "Book index, as reported (base = 100)",
                         "Book index, without acquisitions (base = 100)"], right_from=4)
    _wrap(da, 1, 46)
    for p in products:
        for b in roster:
            ix = indexed.get((p, b["cert"]), {})
            for r in books[p][b["cert"]]:
                x = ix.get(r["quarter"])
                da.append([p, b["short"], b["cert"], r["quarter"], r["gross"], r["recoveries"],
                           r["avg_balance"], r["gross_rate"], r["net_rate"], r["kept"], r["why"],
                           x["reported"] if x else None, x["organic"] if x else None])
    nrow = 1 + len(products) * NB * Q
    for row in da.iter_rows(min_row=2, max_row=nrow, min_col=5, max_col=13):
        for c in row:
            if c.column in (10, 11):
                continue
            c.number_format = "#,##0" if c.column <= 7 else "0.0000" if c.column <= 9 else "0.0"
    for col, w in zip("ABCDEFGHIJKLM", (16, 16, 11, 9, 14, 14, 16, 10, 10, 8, 44, 14, 14)):
        da.column_dimensions[col].width = w
    K.freeze_below(da, 1)

    def drow(pi, bi, k):
        return 2 + pi * NB * Q + bi * Q + k

    # ---- Workings: lists, then the settings every block reads ----
    lists, lists_pe = {}, {}
    for col, key, head, items in (
            (3, "banks", "All banks", [b["short"] for b in roster]),
            (4, "picks", "Peer picks", shorts + [WORST, BEST]),
            (5, "products", "Products", all_products),
            (6, "co_products", "Charge-off products", products),
            (7, "dq_products", "Delinquency products", dq["products"] if dq else []),
            (8, "windows", "Chart shows", list(WINDOWS)),
            (9, "dq_measures", "Delinquency measures", list(D.MEASURES) if dq else []),
            (10, "bases", "Delinquency bases", list(D.BASES) if dq else [])):
        wk.cell(1, col, head)
        for i, v in enumerate(items):
            wk.cell(2 + i, col, v)
        n_ = max(1, len(items))
        lists[key] = "$%s$2:$%s$%d" % (L(col), L(col), 1 + n_)
        lists_pe[key] = WK + lists[key]

    def pe_abs(cell):
        m = re.fullmatch(r"([A-Z]+)(\d+)", cell)
        return "%s$%s$%s" % (PE, m.group(1), m.group(2))

    PICK = {key: pe_abs(picker_cell(key)) for key, _, _, _ in PICKERS}
    SET = {key: pe_abs(setting_cell(key, NP)) for key, _, _, _ in SETTINGS}
    INC = [pe_abs(include_cell(j)) for j in range(NP)]
    mb = SET["min_banks"]

    S, S_PE = {}, {}

    def scalar(key, text, formula):
        r = len(S) + 1
        wk.cell(r, 1, text)
        wk.cell(r, 2, formula)
        S[key] = "$B$%d" % r
        S_PE[key] = WK + S[key]
        return S[key]

    scalar("co_p", "Charge-off product number (0: not built)",
           "=IFERROR(MATCH(%s,%s,0),0)" % (PICK["product"], lists["co_products"]))
    scalar("net", "Net picked", '=%s="Net"' % PICK["rate"])
    # Best and worst, from the ranking block below (its columns are fixed here).
    RK = 12                                     # ranking block: L (peer), M (on), N (kept), O (mean)
    means, rnames = "$O$2:$O$%d" % (1 + NP), "$L$2:$L$%d" % (1 + NP)
    scalar("best", "Best in pack", '=IF(COUNT({m})=0,"(none qualifies)",INDEX({n},MATCH(MIN({m}),{m},0)))'
           .format(m=means, n=rnames))
    scalar("worst", "Worst in pack", '=IF(COUNT({m})=0,"(none qualifies)",INDEX({n},MATCH(MAX({m}),{m},0)))'
           .format(m=means, n=rnames))
    for who in ("a", "b"):
        p_ = PICK["peer_" + who]
        scalar(who + "_name", "Peer %s is" % who.upper(), '=IF({p}="{w}",{wr},IF({p}="{b}",{br},{p}))'
               .format(p=p_, w=WORST, b=BEST, wr=S["worst"], br=S["best"]))
        scalar(who + "_no", "Peer %s, bank number" % who.upper(),
               "=IFERROR(MATCH(%s,%s,0),NA())" % (S[who + "_name"], lists["banks"]))
    scalar("span", "Quarters the charts show at most",
           "=CHOOSE(MATCH(%s,%s,0),%s)" % (PICK["window"], lists["windows"], ",".join(str(s) for s in _SPANS)))
    if dq:
        NM = len(D.MEASURES)
        scalar("dq_p", "Delinquency product number (0: not built)",
               "=IFERROR(MATCH(%s,%s,0),0)" % (PICK["product"], lists["dq_products"]))
        scalar("dq_m", "Delinquency measure number", "=MATCH(%s,%s,0)" % (PICK["dq_measure"], lists["dq_measures"]))
        scalar("dq_b", "Delinquency basis number", "=MATCH(%s,%s,0)" % (PICK["dq_basis"], lists["bases"]))
        scalar("dq_sel", "Delinquency product and measure", "=(%s-1)*%d+%s" % (S["dq_p"], NM, S["dq_m"]))
    co_p, net = S["co_p"], S["net"]

    def pick(dcol, bi, k):
        """The Data cell in this column for the picked product, bank (roster position) and quarter."""
        return "CHOOSE(%s,%s)" % (co_p, ",".join("%s$%s$%d" % (DA, dcol, drow(pi, bi, k))
                                                for pi in range(len(products))))

    def pick_number(cells, missing='""'):
        """One of ``cells`` (one per charge-off product) for the picked product,
        or ``missing`` when it is blank or the product was not built: a blank
        read through a formula would be 0."""
        return "IF({p}=0,{m},IF(CHOOSE({p},{isn}),CHOOSE({p},{c}),{m}))".format(
            p=co_p, m=missing, isn=",".join("ISNUMBER(%s)" % c for c in cells), c=",".join(cells))

    # ---- Workings: the ranking block (L:O, one row per peer) ----
    for c_, h in zip(range(RK, RK + 4), ("Peer", "Switched on", "Kept quarters in the window",
                                         "Mean rate in the window")):
        wk.cell(1, c_, h)

    # ---- Workings: the industry's rates, as values, per product ----
    blk = _Block(wk, 18)
    ind_cols = {}
    for pi, p in enumerate(products):
        rows_ = {x["quarter"]: x for x in ((industry or {}).get(p) or [])}
        for kind in ("gross", "net"):
            c_ = blk.add("ind_%s_%d" % (kind, pi), heads=["Industry %s, %s" % (kind, p)])
            ind_cols[kind, pi] = c_
            for k, q in enumerate(grid):
                v = rows_.get(q, {}).get(kind + "_rate")
                wk.cell(2 + k, c_, v if v is not None else "=NA()")

    # ---- Workings: the charge-off view, one row per TTM quarter ----
    co = blk
    co.add("q", heads=["Quarter"])
    co.add("inwin", heads=["In the ranking window"])
    co.add("uk", NB, ["%s kept" % b["short"] for b in roster])
    co.add("ur", NB, ["%s rate" % b["short"] for b in roster])
    for name, what in (("ok", "in"), ("x", "charge-offs"), ("bal", "balance"), ("r", "rate"),
                       ("win", "ranking rate")):
        co.add(name, NP, ["%s %s" % (b["short"], what) for b in peers])
    for name, h in (("banks", "Banks in"), ("c", "Pack composite"), ("p10", "10th"), ("p25", "25th"),
                    ("p75", "75th"), ("p90", "90th"), ("k", me["short"])):
        co.add(name, heads=[h])
    co.add("a", heads=['="Peer A: "&%s' % S["a_name"]])
    co.add("b", heads=['="Peer B: "&%s' % S["b_name"]])
    co.add("ind", heads=["Industry"])
    co.add("drawn", heads=["Row drawn"])
    for k, q in enumerate(grid):
        r = 2 + k
        wk.cell(r, co.col["q"], q)
        wk.cell(r, co.col["inwin"], "=%d>=%d-%s" % (k, Q, SET["last_n"]))
        for j in range(NB):
            wk.cell(r, co.col["uk"] + j, "=IF(%s>0,%s,FALSE)" % (co_p, pick("J", j, k)))
            wk.cell(r, co.col["ur"] + j, "=IF({uk},IF({n},{nr},{gr}),NA())".format(
                uk=co.at("uk", r, j), n=net, nr=pick("I", j, k), gr=pick("H", j, k)))
        for j in range(NP):
            i = peer_i[j]
            ok = co.at("ok", r, j)
            wk.cell(r, co.col["ok"] + j, '=AND(%s="%s",%s)' % (INC[j], YES, co.at("uk", r, i)))
            wk.cell(r, co.col["x"] + j, '=IF({ok},{g}-IF({n},{rc},0),"")'.format(
                ok=ok, g=pick("E", i, k), n=net, rc=pick("F", i, k)))
            wk.cell(r, co.col["bal"] + j, '=IF(%s,%s,"")' % (ok, pick("G", i, k)))
            wk.cell(r, co.col["r"] + j, '=IF(%s,%s,"")' % (ok, co.at("ur", r, i)))
            wk.cell(r, co.col["win"] + j, '=IF(AND({w},{ok}),IF({rk}="gross",{g},{n}),"")'.format(
                w=co.at("inwin", r), ok=ok, rk=SET["rank_by"], g=pick("H", i, k), n=pick("I", i, k)))
        banks = co.at("banks", r)
        wk.cell(r, co.col["banks"], "=COUNT(%s)" % co.across("r", r))
        wk.cell(r, co.col["c"], "=IF({n}>={mb},100*SUM({x})/SUM({b}),NA())".format(
            n=banks, mb=mb, x=co.across("x", r), b=co.across("bal", r)))
        for p in (10, 25, 75, 90):
            wk.cell(r, co.col["p%d" % p], "=IF({n}>={mb},PERCENTILE({v},{p}),NA())".format(
                n=banks, mb=mb, v=co.across("r", r), p=p / 100))
        wk.cell(r, co.col["k"], "=" + co.at("ur", r, me_i))
        for name, num in (("a", S["a_no"]), ("b", S["b_no"])):
            wk.cell(r, co.col[name], "=IFERROR(INDEX(%s,%s),NA())" % (co.across("ur", r), num))
        wk.cell(r, co.col["ind"], "=IF({p}>0,IF({n},CHOOSE({p},{nc}),CHOOSE({p},{gc})),NA())".format(
            p=co_p, n=net, nc=",".join("%s%d" % (L(ind_cols["net", pi]), r) for pi in range(len(products))),
            gc=",".join("%s%d" % (L(ind_cols["gross", pi]), r) for pi in range(len(products)))))
        wk.cell(r, co.col["drawn"], '=IF(OR(%s),%d,"")' % (",".join(
            "ISNUMBER(%s)" % co.at(n_, r) for n_ in ("c", "k", "a", "b", "ind", "p10")), r))
    LR = Q + 1                                   # the latest quarter's row

    # The ranking block, now the window columns are placed.
    for j, b in enumerate(peers):
        r = 2 + j
        w = co.down("win", 2, LR, j)
        wk.cell(r, RK, b["short"])
        wk.cell(r, RK + 1, '=%s="%s"' % (INC[j], YES))
        wk.cell(r, RK + 2, "=COUNT(%s)" % w)
        wk.cell(r, RK + 3, '=IF(%s>=%s,AVERAGE(%s),"")' % ("$N$%d" % r, SET["min_obs"], w))

    # ---- Workings: the book-size view, one row per TTM quarter from the base ----
    bk = _Block(wk, co.next + 1)
    bk.add("q", heads=["Quarter (book index)"])
    bk.add("bo", NB, ["%s without acquisitions" % b["short"] for b in roster])
    bk.add("br", NB, ["%s as reported" % b["short"] for b in roster])
    bk.add("v", NP, ["%s in, without acquisitions" % b["short"] for b in peers])
    bk.add("vr", NP, ["%s in, as reported" % b["short"] for b in peers])
    for name, h in (("n", "Banks in"), ("p10", "10th"), ("p25", "25th"), ("p50", "Pack median"), ("p75", "75th"),
                    ("p90", "90th"), ("k", me["short"])):
        bk.add(name, heads=[h])
    bk.add("a", heads=['="Peer A: "&%s' % S["a_name"]])
    bk.add("b", heads=['="Peer B: "&%s' % S["b_name"]])
    bk.add("rn", heads=["Banks in, as reported"])
    bk.add("rmed", heads=["Pack median, as reported"])
    bk.add("drawn", heads=["Row drawn"])
    BQ = Q - base_k
    for j_, k in enumerate(range(base_k, Q)):
        r = 2 + j_
        wk.cell(r, bk.col["q"], grid[k])
        for i in range(NB):
            for name, dcol in (("bo", "M"), ("br", "L")):
                cells = ["%s$%s$%d" % (DA, dcol, drow(pi, i, k)) for pi in range(len(products))]
                wk.cell(r, bk.col[name] + i, "=" + pick_number(cells, "NA()"))
        for j, i in enumerate(peer_i):
            for name, src in (("v", "bo"), ("vr", "br")):
                u = bk.at(src, r, i)
                wk.cell(r, bk.col[name] + j, '=IF(AND(%s="%s",ISNUMBER(%s)),%s,"")' % (INC[j], YES, u, u))
        n_ = bk.at("n", r)
        wk.cell(r, bk.col["n"], "=COUNT(%s)" % bk.across("v", r))
        for p in (10, 25, 50, 75, 90):
            wk.cell(r, bk.col["p%d" % p], "=IF({n}>={mb},PERCENTILE({v},{p}),NA())".format(
                n=n_, mb=mb, v=bk.across("v", r), p=p / 100))
        wk.cell(r, bk.col["k"], "=" + bk.at("bo", r, me_i))
        for name, num in (("a", S["a_no"]), ("b", S["b_no"])):
            wk.cell(r, bk.col[name], "=IFERROR(INDEX(%s,%s),NA())" % (bk.across("bo", r), num))
        wk.cell(r, bk.col["rn"], "=COUNT(%s)" % bk.across("vr", r))
        wk.cell(r, bk.col["rmed"], "=IF({n}>={mb},PERCENTILE({v},0.5),NA())".format(
            n=bk.at("rn", r), mb=mb, v=bk.across("vr", r)))
        wk.cell(r, bk.col["drawn"], '=IF(OR(%s),%d,"")' % (",".join(
            "ISNUMBER(%s)" % bk.at(n_, r) for n_ in ("p50", "k", "a", "b", "p10")), r))
    BR = BQ + 1                                  # the book index's latest row

    # ---- Workings: lagged rates, per bank (roster order) ----
    lg = _Block(wk, bk.next + 1)
    lg.add("bank", heads=["Bank (lagged)"])
    for pi, p in enumerate(products):
        for y in lagged_years:
            lg.add("v_%d_%d" % (pi, y), 3, ["%s %d rate" % (p, y), "%s %d net ($000)" % (p, y),
                                             "%s %d-year-earlier balance ($000)" % (p, y)])
    lg.add("on", heads=["In the pack"])
    for y in lagged_years:
        lg.add("s_%d" % y, 4, ["%d rate" % y, "%d net" % y, "%d balance" % y, "%d missing" % y])
    for i, b in enumerate(roster):
        r = 2 + i
        wk.cell(r, lg.col["bank"], b["short"])
        for pi, p in enumerate(products):
            for y in lagged_years:
                lr = lagged_rate(data, b["cert"], p, y, lag_years)
                for j, v in enumerate((lr["rate"], lr["net"], lr["balance"])):
                    wk.cell(r, lg.col["v_%d_%d" % (pi, y)] + j, v)
        wk.cell(r, lg.col["on"], 0 if b is me else '=IF(%s="%s",1,0)' % (INC[peer_i.index(i)], YES))
        for y in lagged_years:
            for j in range(3):
                cells = [lg.at("v_%d_%d" % (pi, y), r, j) for pi in range(len(products))]
                wk.cell(r, lg.col["s_%d" % y] + j, "=" + pick_number(cells))
            wk.cell(r, lg.col["s_%d" % y] + 3, "=IF(AND({on}=1,OR(NOT(ISNUMBER({n})),NOT(ISNUMBER({b})))),1,0)"
                    .format(on=lg.at("on", r), n=lg.at("s_%d" % y, r, 1), b=lg.at("s_%d" % y, r, 2)))

    # ---- Workings: the mix, per year, per product, as values ----
    mx = _Block(wk, lg.next + 1)
    mx.add("year", heads=["Year (mix)"])
    MIX_KEYS = ("bank_rate", "industry_rate", "bank_ratio", "industry_ratio", "own_factor", "balance_factor")
    status = {}
    for pi, p in enumerate(products):
        mx.add("m_%d" % pi, len(MIX_KEYS), ["%s %s" % (p, k_) for k_ in MIX_KEYS])
        mx.add("own_%d" % pi, NB, ["%s own factor, %s" % (p, b["short"]) for b in roster])
        ind = (industry or {}).get(p)
        if not ind:
            status[p] = ("No industry series was supplied for %s, so no mix factor is worked out. The roster "
                         "banks are not the industry and do not stand in for it." % p)
            continue
        ind_rate = {int(x["quarter"][:4]): x["net_rate"] for x in ind if x["quarter"].endswith("Q4")}
        if mix_base_year not in ind_rate:
            status[p] = "The industry series for %s has no %d to measure from." % (p, mix_base_year)
            continue
        if mix_base_year not in years:
            status[p] = ("The data does not hold all four quarters of %d, so the banks have no %d to measure "
                         "from." % (mix_base_year, mix_base_year))
            continue
        status[p] = "OK"
        for i, b in enumerate(roster):
            cy = {y: calendar_year(data, b["cert"], p, y) for y in years}
            rate = {y: lagged_rate(data, b["cert"], p, y, 0)["rate"] for y in years}
            rows_ = mix_factor(rate, {y: ind_rate.get(y) for y in years}, {y: cy[y]["balance"] for y in years},
                               mix_base_year)
            for x_ in rows_:
                r = 2 + years.index(x_["year"])
                wk.cell(r, mx.col["own_%d" % pi] + i, x_["own_factor"])
                if b is me:
                    vals = (rate[x_["year"]], ind_rate.get(x_["year"]), x_["bank_ratio"], x_["industry_ratio"],
                            x_["own_factor"], x_["balance_factor"])
                    for j, v in enumerate(vals):
                        wk.cell(r, mx.col["m_%d" % pi] + j, v)
    for y in years:
        wk.cell(2 + years.index(y), mx.col["year"], y)
    mx.add("status", heads=["Mix, per product"])
    for pi, p in enumerate(products):
        wk.cell(2 + pi, mx.col["status"], status[p])
    mix_status = scalar("mix", "Mix for the product picked",
                        '=IF({p}=0,"No charge-offs were built for "&{pp}&", so there is no mix.",CHOOSE({p},{s}))'
                        .format(p=co_p, pp=PICK["product"], s=",".join(mx.fixed("status", 2 + pi)
                                                                       for pi in range(len(products)))))
    mx.add("sel", len(MIX_KEYS), ["Picked: %s" % k_ for k_ in MIX_KEYS])
    for y in years:
        r = 2 + years.index(y)
        for j in range(len(MIX_KEYS)):
            cells = [mx.at("m_%d" % pi, r, j) for pi in range(len(products))]
            wk.cell(r, mx.col["sel"] + j, "=" + pick_number(cells))

    # ---- Workings: each bank's own row for the bank table (roster order) ----
    pb = _Block(wk, mx.next + 1)
    pb.add("bank", heads=["Bank (table)"])
    pb.add("latest", heads=["Latest rate"])
    pb.add("index", heads=["Book index, latest"])
    pb.add("notes", len(products), ["Book index note, %s" % p for p in products])
    pb.add("own", heads=["Own factor, latest year"])
    for i, b in enumerate(roster):
        r = 2 + i
        wk.cell(r, pb.col["bank"], b["short"])
        wk.cell(r, pb.col["latest"], "=" + co.at("ur", LR, i))
        wk.cell(r, pb.col["index"], "=" + bk.at("bo", BR, i))
        for pi, p in enumerate(products):
            wk.cell(r, pb.col["notes"] + pi, index_note[p, b["cert"]])
        if NY:
            cells = [mx.at("own_%d" % pi, 1 + NY, i) for pi in range(len(products))]
            wk.cell(r, pb.col["own"], "=" + pick_number(cells, "NA()"))
        else:
            wk.cell(r, pb.col["own"], "=NA()")

    # ---- Workings: the delinquency view ----
    if dq:
        dl = _delinquency_block(wk, dq, roster, me, peers, pb.next + 1, dict(
            DD=DD, dq_sel=S["dq_sel"], dq_p=S["dq_p"], dq_b=S["dq_b"], inc=INC, a_name=S["a_name"],
            b_name=S["b_name"], a_no=S["a_no"], b_no=S["b_no"], mb=mb))
        DLR = len(dq["grid"]) + 1
        d_rows = _delinquency_data(dd, dq, roster)
    else:
        dl = None
        for ws in (dd,):
            ws.cell(1, 1, NOT_BUILT).font = Font(bold=True)
            ws.column_dimensions["A"].width = 100

    # Where each chart starts and how many rows it draws: from the first
    # quarter any of its lines has, at most "Chart shows" quarters back from
    # the last. An axis cannot follow a cell; a named range can.
    def fit(key, block, n_rows):
        drawn = block.down("drawn", 2, 1 + n_rows)
        first = scalar(key + "_first", "First row drawn (%s)" % key,
                       "=IF(COUNT({a})=0,2,MAX(MIN({a}),MAX({a})-{s}+1))".format(a=drawn, s=S["span"]))
        scalar(key + "_count", "Rows drawn (%s)" % key,
               "=IF(COUNT({a})=0,{n},MAX({a})-{f}+1)".format(a=drawn, n=n_rows, f=first))

    fit("co", co, Q)
    fit("bk", bk, BQ)
    if dl is not None:
        fit("dl", dl, len(dq["grid"]))
    wk.column_dimensions["A"].width = 40
    wk.column_dimensions["B"].width = 16

    def named(key, block, name, j=0, nm=None):
        """A name on the Peers tab for one Workings column, as drawn."""
        nm = nm or "%s_%s" % (key, name)
        pe.defined_names[nm] = DefinedName(nm, attr_text="OFFSET(%s$%s$1,%s-1,0,%s,1)" % (
            WK, L(block.col[name] + j), S_PE[key + "_first"], S_PE[key + "_count"]))
        return PE + nm

    def label_ref(block, name):
        return SeriesLabel(strRef=StrRef(f="%s$%s$1" % (WK, L(block.col[name]))))

    # ======================= the Peers tab: the dashboard =======================
    pe.column_dimensions["A"].width = 1.7
    for c_ in range(FIRST_COL, LAST_COL + 1):
        pe.column_dimensions[L(c_)].width = COL_WIDTH
    K.light_title(pe, 1, LAST_COL - 1, "Peers: KeyBank against its pack",
                  t.words("Pick in the yellow cells; every number and chart on this tab follows them. From "
                          "public FDIC Call Reports; the Glossary tab says how each figure is worked out."),
                  live=True, first_col=FIRST_COL)
    pe.row_dimensions[3].height = 6
    starts = {"product": product, "rate": rank_by.capitalize(), "peer_a": peer_a, "peer_b": peer_b,
              "window": START_WINDOW, "dq_measure": dq["measure"] if dq else "not built",
              "dq_basis": dq["basis"] if dq else "not built"}
    choices = {"product": "=" + lists_pe["products"], "rate": '"%s"' % ",".join(RATES),
               "peer_a": "=" + lists_pe["picks"], "peer_b": "=" + lists_pe["picks"],
               "window": "=" + lists_pe["windows"],
               "dq_measure": "=" + lists_pe["dq_measures"] if dq else None,
               "dq_basis": "=" + lists_pe["bases"] if dq else None}
    for key, text, first, last in PICKERS:
        pe.cell(R_LABEL, first, text).font = F_LAB
        c = pe.cell(R_PICK, first, starts[key])
        _yellow(c, last)
        if choices[key]:
            dv = DataValidation(type="list", formula1=choices[key], allow_blank=False, showErrorMessage=True,
                                errorTitle=text, error="Pick one from the list.")
            pe.add_data_validation(dv)
            dv.add(c.coordinate)
        else:
            c.fill = K.solid(K.CANVAS)
            c.font = Font(name="Calibri", size=11, color=K.SLATE)
    pe.row_dimensions[R_PICK].height = 20
    for who in ("a", "b"):
        first = next(f for key, _, f, _ in PICKERS if key == "peer_" + who)
        _merged(pe, R_RESOLVED, first, first + 1, '="Drawn: "&%s' % S_PE[who + "_name"], NOTE, wrap=False)
    dq_p = S_PE.get("dq_p")
    _merged(pe, R_RESOLVED, 2, 6, '=IF({co}=0,"No charge-offs built for this product.",{dq})'.format(
        co=S_PE["co_p"], dq='IF(%s=0,"No delinquency built for this product.","")' % dq_p if dq
        else '"Delinquency was not built for this workbook."'), Font(name="Calibri", size=9, bold=True,
                                                                     color=K.CRIMSON), wrap=False)
    pe.row_dimensions[7].height = 8

    # Which banks are in the pack: a Yes/No under each peer's name.
    yes_no = DataValidation(type="list", formula1='"%s,%s"' % (YES, NO), allow_blank=False, showErrorMessage=True,
                            errorTitle="In the pack", error="Yes or No.")
    pe.add_data_validation(yes_no)
    for ch in range(lay["chunks"]):
        r = R_BANKS + 2 * ch
        pe.cell(r, 2, "Bank").font = F_LAB
        pe.cell(r + 1, 2, "In the pack").font = F_LAB
        pe.row_dimensions[r].height = 26
        pe.row_dimensions[r + 1].height = 18
    for j, b in enumerate(peers):
        cell = include_cell(j)
        m = re.fullmatch(r"([A-Z]+)(\d+)", cell)
        c = pe[cell]
        name = pe.cell(int(m.group(2)) - 1, c.column, b["short"])
        name.font = Font(name="Calibri", size=9, bold=True, color=K.INK_TEXT)
        name.alignment = Alignment(wrap_text=True, horizontal="center", vertical="bottom")
        c.value = YES if b["cert"] in in_pack else NO
        c.fill = INPUT
        c.border = BOX
        c.font = Font(name="Calibri", size=10, bold=True, color=K.INK_TEXT)
        c.alignment = Alignment(horizontal="center", vertical="center")
        yes_no.add(cell)
        # A bank switched off is greyed, so who is out reads at a glance.
        pe.conditional_formatting.add(cell, FormulaRule(formula=['%s="%s"' % (cell, NO)],
                                                        fill=K.solid(K.MIST), font=Font(color=K.SLATE)))
    rs = lay["settings"]
    counts = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=False,
                            showErrorMessage=True, errorTitle="A whole number of at least 1",
                            error="Type a whole number of 1 or more.")
    pe.add_data_validation(counts)
    kinds = DataValidation(type="list", formula1='"gross,net"', allow_blank=False, showErrorMessage=True,
                           errorTitle="Ranked by", error="gross or net.")
    pe.add_data_validation(kinds)
    for key, text, (a_, z_), col in SETTINGS:
        _merged(pe, rs, a_, z_, text, F_LAB, wrap=False, align="right")
        c = pe.cell(rs, col, {"rank_by": rank_by, "last_n": last_n, "min_obs": min_obs,
                              "min_banks": min_banks}[key])
        _yellow(c)
        c.alignment = Alignment(horizontal="center", vertical="center")
        (kinds if key == "rank_by" else counts).add(c.coordinate)
    pe.row_dimensions[rs].height = 20
    _merged(pe, lay["best_worst"], 2, LAST_COL, (
        '="Best in pack: "&{b}&". Worst in pack: "&{w}&". Only banks switched on are ranked; Peer A and B are '
        'drawn either way. The pack started as {start}."').format(
            b=S_PE["best"], w=S_PE["worst"], start="every peer" if pack == "all" else "the regionals only"),
        NOTE, wrap=False)

    charts = []                                  # (chart, format, label skip, columns, first row, rows)
    # ---- row 1 of views: charge-offs (left), delinquency (right) ----
    v1, tl = lay["v1"], lay["tiles"]
    _section(pe, v1, *LEFT, '="Charge-offs: "&{p}&", "&LOWER({r})&" rate, trailing four quarters (%)"'
             .format(p=PICK["product"], r=PICK["rate"]))
    last_q = grid[-1]

    def w_(block, name, r, j=0):
        return WK + block.fixed(name, r, j)

    k_, c_l, ind_l = w_(co, "k", LR), w_(co, "c", LR), w_(co, "ind", LR)
    r_row = WK + co.across("r", LR)
    _tile(pe, tl, 2, me["short"].upper(), '=IF(ISNUMBER({x}),{x},"none")'.format(x=k_),
          '=IF(ISNUMBER({x}),"in {q}","not kept in {q}")'.format(x=k_, q=last_q), "0.00")
    _tile(pe, tl, 4, "PACK COMPOSITE", '=IF(ISNUMBER({x}),{x},"none")'.format(x=c_l),
          '=IF(ISNUMBER({x}),{n}&" banks in, {q}","under "&{mb}&" banks kept")'
          .format(x=c_l, n=w_(co, "banks", LR), mb=mb, q=last_q), "0.00")
    _tile(pe, tl, 6, "%s'S RANK" % me["short"].upper(),
          '=IF(ISNUMBER({k}),(COUNTIF({r},"<"&{k})+1)&" of "&(COUNT({r})+1),"none")'.format(k=k_, r=r_row),
          "1 = the lowest rate", "General")
    _tile(pe, tl, 8, "INDUSTRY", '=IF(ISNUMBER({x}),{x},"none")'.format(x=ind_l),
          '=IF(ISNUMBER({x}),"every filer, {q}","not supplied")'.format(x=ind_l, q=last_q),
          "0.00")
    ch = _Chart(named("co", co, "q"))
    ch.band(named("co", co, "p10"), named("co", co, "p90"), BAND_LIGHT)
    ch.band(named("co", co, "p25"), named("co", co, "p75"), BAND_DARK)
    ch.lines([(named("co", co, "c"), label_ref(co, "c")) + STYLE["pack"],
              (named("co", co, "k"), label_ref(co, "k")) + STYLE["self"],
              (named("co", co, "a"), label_ref(co, "a")) + STYLE["a"],
              (named("co", co, "b"), label_ref(co, "b")) + STYLE["b"],
              (named("co", co, "ind"), label_ref(co, "ind")) + STYLE["industry"]])
    charts.append((ch, "0.0", 8, LEFT, lay["chart1"], CHART_ROWS))
    _note(pe, lay["note1"], LEFT, NOTE_ROWS["note1"], words=t.words, text=(
        "Shaded: the middle 80% (light) and middle 50% (dark) of the banks switched on and kept; none where "
        "fewer than the banks needed are kept. A gap: a quarter not kept (the Data tab says why)."),
        cond="%s=0" % S_PE["co_p"], alt="No charge-offs were built for this product. The products built are: %s."
        % ", ".join(products))

    if dl is not None:
        _section(pe, v1, *RIGHT, '="Delinquency: "&{p}&", "&{m}&", "&{b}&" (%)"'.format(
            p=PICK["product"], m=PICK["dq_measure"], b=PICK["dq_basis"]))
        dlast = dq["grid"][-1]
        dk, dc = w_(dl, "k", DLR), w_(dl, "c", DLR)
        v_row = WK + dl.across("v", DLR)
        _tile(pe, tl, 10, me["short"].upper(), '=IF(ISNUMBER({x}),{x},"none")'.format(x=dk),
              '=IF(ISNUMBER({x}),"at {q}","no rate at {q}")'.format(x=dk, q=dlast), "0.00")
        _tile(pe, tl, 12, "PACK COMPOSITE", '=IF(ISNUMBER({x}),{x},"none")'.format(x=dc),
              '=IF(ISNUMBER({x}),{n}&" banks in, {q}","under "&{mb}&" banks kept")'
              .format(x=dc, n=w_(dl, "banks", DLR), q=dlast, mb=mb), "0.00")
        _tile(pe, tl, 14, "%s'S RANK" % me["short"].upper(),
              '=IF(ISNUMBER({k}),(COUNTIF({r},"<"&{k})+1)&" of "&(COUNT({r})+1),"none")'.format(k=dk, r=v_row),
              "1 = the lowest rate", "General")
        _tile(pe, tl, 16, "BANKS WITH A RATE", "=" + w_(dl, "bandn", DLR),
              "switched on, %s" % dlast, "0")
        ch = _Chart(named("dl", dl, "q"))
        ch.band(named("dl", dl, "p10"), named("dl", dl, "p90"), BAND_LIGHT)
        ch.band(named("dl", dl, "p25"), named("dl", dl, "p75"), BAND_DARK)
        ch.lines([(named("dl", dl, "c"), label_ref(dl, "c")) + STYLE["pack"],
                  (named("dl", dl, "k"), label_ref(dl, "k")) + STYLE["self"],
                  (named("dl", dl, "a"), label_ref(dl, "a")) + STYLE["a"],
                  (named("dl", dl, "bb"), label_ref(dl, "bb")) + STYLE["b"]])
        charts.append((ch, "0.00", 8, RIGHT, lay["chart1"], CHART_ROWS))
        _note(pe, lay["note1"], RIGHT, NOTE_ROWS["note1"], words=t.words, text=(
            "Shaded as on the left. On the four-quarter average a bank mid-acquisition is left out of the "
            "composite, as of the band. A gap: no rate that quarter-end (DelinquencyData says why)."),
            cond="%s=0" % S_PE["dq_p"], alt="No delinquency was built for this product. The products built are: "
            "%s." % ", ".join(dq["products"]))
    else:
        _section(pe, v1, *RIGHT, "Delinquency")
        _merged(pe, tl, *RIGHT, NOT_BUILT, F_BOLD, 30)

    # ---- row 2: book size (left); lagged and how far rates moved (right) ----
    v2 = lay["v2"]
    _section(pe, v2, *LEFT, '="Book size: each bank\'s "&LOWER({p})&" loans, {b} = 100, without acquisitions"'
             .format(p=PICK["product"], b=base_lbl))
    ch = _Chart(named("bk", bk, "q"))
    ch.band(named("bk", bk, "p10"), named("bk", bk, "p90"), BAND_LIGHT)
    ch.band(named("bk", bk, "p25"), named("bk", bk, "p75"), BAND_DARK)
    ch.lines([(named("bk", bk, "p50"), label_ref(bk, "p50")) + STYLE["pack"],
              (named("bk", bk, "k"), label_ref(bk, "k")) + STYLE["self"],
              (named("bk", bk, "a"), label_ref(bk, "a")) + STYLE["a"],
              (named("bk", bk, "b"), label_ref(bk, "b")) + STYLE["b"]])
    charts.append((ch, "0", 4, LEFT, lay["chart2"], BOOK_ROWS))
    bt = lay["book_table"]
    _table_head(pe, bt, 2, ["Latest, %s (%s = 100)" % (last_q, base_lbl), "", "", "As reported", "",
                            "Without acquisitions", ""], height=18)
    for a_, z_ in ((2, 4), (5, 6), (7, 8)):
        pe.merge_cells(start_row=bt, start_column=a_, end_row=bt, end_column=z_)
    book_rows = (
        (me["short"], w_(bk, "br", BR, me_i), w_(bk, "bo", BR, me_i), True),
        ('="Peer A: "&%s' % S_PE["a_name"], "INDEX(%s,%s)" % (WK + bk.across("br", BR), S_PE["a_no"]),
         "INDEX(%s,%s)" % (WK + bk.across("bo", BR), S_PE["a_no"]), False),
        ('="Peer B: "&%s' % S_PE["b_name"], "INDEX(%s,%s)" % (WK + bk.across("br", BR), S_PE["b_no"]),
         "INDEX(%s,%s)" % (WK + bk.across("bo", BR), S_PE["b_no"]), False),
        ('="Pack median ("&%s&" banks)"' % w_(bk, "n", BR), w_(bk, "rmed", BR),
         w_(bk, "p50", BR), False))
    for i, (who, rep, org, mine) in enumerate(book_rows):
        r = bt + 1 + i
        _merged(pe, r, 2, 4, who, F_SELF if mine else F_BODY, wrap=False)
        for c_, f in ((5, rep), (7, org)):
            pe.merge_cells(start_row=r, start_column=c_, end_row=r, end_column=c_ + 1)
            cell = pe.cell(r, c_, '=IFERROR(%s+0,"")' % f)
            cell.number_format = "0"
            cell.font = F_BODY
            cell.alignment = Alignment(horizontal="right")
        for c_ in range(2, 9):
            pe.cell(r, c_).border = HAIR
    _note(pe, lay["book_note"], LEFT, NOTE_ROWS["book_note"], words=t.words, text=(
        "Each bank's book at a quarter-end / its book at %s x 100. Without acquisitions, growth in a quarter "
        "the bank bought a book counts as none (its acquisition quarters are on the Run tab). Shaded: the middle "
        "80%% and 50%% of the banks switched on and indexed; the black line is their median. A bank with less "
        "than $%sm at the base is not indexed." % (base_lbl, "{:,.0f}".format(index_min_book / 1000))))

    _section(pe, v2, *RIGHT, '="Lagged rate: "&LOWER({p})&" net losses over the book {n} year{s} '
             'before (%)"'.format(p=PICK["product"], n=lag_years, s="" if lag_years == 1 else "s"))
    lr0 = lay["lagged"]
    _table_head(pe, lr0, 10, ["", "", ""] + ["%d over %d" % (y, y - lag_years) for y in lagged_years])
    pe.merge_cells(start_row=lr0, start_column=10, end_row=lr0, end_column=12)
    on_rng = WK + lg.down("on", 2, 1 + NB)
    lag_rows = [(me["short"], lambda y: w_(lg, "s_%d" % y, 2 + me_i), True),
                ('="Peer A: "&%s' % S_PE["a_name"],
                 lambda y: "INDEX(%s,%s)" % (WK + lg.down("s_%d" % y, 2, 1 + NB), S_PE["a_no"]), False),
                ('="Peer B: "&%s' % S_PE["b_name"],
                 lambda y: "INDEX(%s,%s)" % (WK + lg.down("s_%d" % y, 2, 1 + NB), S_PE["b_no"]), False)]
    for i, (who, ref, mine) in enumerate(lag_rows):
        r = lr0 + 1 + i
        _merged(pe, r, 10, 12, who, F_SELF if mine else F_BODY, wrap=False)
        for j, y in enumerate(lagged_years):
            cell = pe.cell(r, 13 + j, '=IFERROR(%s+0,"")' % ref(y))
            cell.number_format = "0.00"
            cell.font = F_BODY
    r = lr0 + 4
    _merged(pe, r, 10, 12, '="Pack total ("&SUM(%s)&" banks switched on)"' % on_rng, F_BOLD, wrap=False)
    for j, y in enumerate(lagged_years):
        n_, b_, miss = (WK + lg.down("s_%d" % y, 2, 1 + NB, x_) for x_ in (1, 2, 3))
        cell = pe.cell(r, 13 + j, '=IF(OR(SUM({on})=0,SUM({m})>0),"",IF(SUMIF({on},1,{b})>0,100*SUMIF({on},1,{n})'
                                  '/SUMIF({on},1,{b}),""))'.format(on=on_rng, m=miss, b=b_, n=n_))
        cell.number_format = "0.00"
        cell.font = F_BOLD
    for rr in range(lr0 + 1, lr0 + 5):
        for c_ in range(10, 13 + len(lagged_years)):
            pe.cell(rr, c_).border = HAIR
    _note(pe, lay["lagged_note"], RIGHT, NOTE_ROWS["lagged_note"],
          "100 x the year's net charge-offs (Q4 year-to-date, as filed) / the mean balance %d year%s earlier. "
          "Blank: a bank did not file every quarter of a year needed. The pack total sums the dollars, and is "
          "blank if any bank switched on is." % (lag_years, "" if lag_years == 1 else "s"))

    _section(pe, lay["norm"], *RIGHT, '="How far each "&LOWER({r})&" rate moved, {a} to {b}"'.format(
        r=PICK["rate"], a=y0, b=y1))
    nt = lay["norm_table"]
    _table_head(pe, nt, 10, ["", "", "", "%d (%%)" % y0, "%d (%%)" % y1, "%d / %d" % (y1, y0)])
    pe.merge_cells(start_row=nt, start_column=10, end_row=nt, end_column=12)

    def year_rows(y):
        ks = [grid.index("%dQ%d" % (y, n_)) for n_ in range(1, 5) if "%dQ%d" % (y, n_) in grid]
        return (2 + ks[0], 2 + ks[-1]) if len(ks) == 4 and ks[-1] - ks[0] == 3 else None

    yr = {y: year_rows(y) for y in (y0, y1)}
    norm_rows = ((me["short"], ("ur", me_i), True), ('="Peer A: "&%s' % S_PE["a_name"], ("a", 0), False),
                 ('="Peer B: "&%s' % S_PE["b_name"], ("b", 0), False),
                 ("Pack composite", ("c", 0), False))
    for i, (who, (name, j), mine) in enumerate(norm_rows):
        r = nt + 1 + i
        _merged(pe, r, 10, 12, who, F_SELF if mine else (F_BOLD if name == "c" else F_BODY), wrap=False)
        for c_, y in ((13, y0), (14, y1)):
            if yr[y] is None:
                f = '=""'
            else:
                rng = WK + co.down(name, yr[y][0], yr[y][1], j)
                f = '=IF(COUNT({r})=4,AVERAGE({r}),"")'.format(r=rng)
            pe.cell(r, c_, f).number_format = "0.00"
        pe.cell(r, 15, '=IF(AND(ISNUMBER(M{r}),ISNUMBER(N{r})),IF(M{r}>0,N{r}/M{r},""),"")'.format(r=r)
                ).number_format = "0.00"
        for c_ in range(10, 16):
            pe.cell(r, c_).border = HAIR
            if c_ >= 13:
                pe.cell(r, c_).font = F_BOLD if name == "c" else F_BODY
    _note(pe, lay["norm_note"], RIGHT, NOTE_ROWS["norm_note"],
          "Each year: the mean of its four quarterly trailing-four-quarter rates, blank unless all four were "
          "kept. The ratio is blank when the earlier year is 0 or below (net recoveries): a ratio to it would "
          "read the wrong way.")

    # ---- row 3: against the industry (left); every bank (right) ----
    v3 = lay["v3"]
    _section(pe, v3, *LEFT, '="Against the industry: "&{me}&"\'s "&LOWER({p})&" net rate since {y}"'.format(
        me='"%s"' % me["short"], p=PICK["product"], y=mix_base_year))
    mt = lay["mix"]
    _table_head(pe, mt, 2, ["Year", "%s net rate (%%)" % me["short"], "Industry net rate (%)",
                            "%s / %d" % (me["short"], mix_base_year), "Industry / %d" % mix_base_year,
                            "Own factor", "Balance factor"])
    for y in years:
        r = mt + 1 + years.index(y)
        pe.cell(r, 2, y).font = F_BODY
        pe.cell(r, 2).alignment = Alignment(horizontal="left")
        for j in range(len(MIX_KEYS)):
            cell = pe.cell(r, 3 + j, '=IF({s}<>"OK","",{v})'.format(
                s=S_PE["mix"], v=WK + mx.fixed("sel", 2 + years.index(y), j)))
            cell.number_format = "0.00"
            cell.font = F_BOLD if MIX_KEYS[j] == "own_factor" else F_BODY
        for c_ in range(2, 9):
            pe.cell(r, c_).border = HAIR
    own = WK + pb.down("own", 2, 1 + NB)
    last_year = years[-1] if years else mix_base_year
    _merged(pe, lay["mix_peers"], *LEFT, (
        '=IF({s}<>"OK",{s},"Own factor in {y}: Peer A ("&{a}&") "&IFERROR(TEXT(INDEX({o},{an}),"0.00"),"none")'
        '&"; Peer B ("&{b}&") "&IFERROR(TEXT(INDEX({o},{bn}),"0.00"),"none")&".")').format(
            s=S_PE["mix"], y=last_year, a=S_PE["a_name"], b=S_PE["b_name"], o=own, an=S_PE["a_no"],
            bn=S_PE["b_no"]), F_BOLD, wrap=False)
    _note(pe, lay["mix_note"], LEFT, NOTE_ROWS["mix_note"],
          "Calendar-year net rate: Q4 year-to-date net charge-offs as filed / the mean of the year's four "
          "quarter-end balances; the industry's is its trailing-four-quarter net rate at Q4. Own factor above "
          "1: %s worsened more than the industry since %d. Balance factor: the %d book / this year's."
          % (me["short"], mix_base_year, mix_base_year))

    _section(pe, v3, *RIGHT, '="Every bank: "&{p}&", "&LOWER({r})&" rate"'.format(
        p=PICK["product"], r=PICK["rate"]))
    bt3 = lay["banks"]
    _table_head(pe, bt3, 10, ["Bank", "", "In the pack", "Latest, %s" % last_q, "Ranking mean",
                              "Kept in the window", "Book index, %s" % last_q, "Own factor, %d" % last_year],
                right_from=3, height=36)
    pe.merge_cells(start_row=bt3, start_column=10, end_row=bt3, end_column=11)
    for i, b in enumerate(roster):
        r = bt3 + 1 + i
        mine = b is me
        _merged(pe, r, 10, 11, b["short"], F_SELF if mine else F_BODY, wrap=False)
        if mine:
            pe.cell(r, 12, "itself")
        else:
            j = peer_i.index(i)
            pe.cell(r, 12, "=" + INC[j])
            pe.cell(r, 14, '=%s%s' % (WK, "$O$%d" % (2 + j))).number_format = "0.00"
            pe.cell(r, 15, '=%s%s' % (WK, "$N$%d" % (2 + j))).number_format = "0"
        pe.cell(r, 13, '=IFERROR(%s+0,"")' % w_(pb, "latest", 2 + i)).number_format = "0.00"
        pe.cell(r, 16, '=IF(ISNUMBER({x}),{x},IF({p}=0,"",CHOOSE({p},{n})))'.format(
            x=w_(pb, "index", 2 + i), p=S_PE["co_p"],
            n=",".join(w_(pb, "notes", 2 + i, pi) for pi in range(len(products))))).number_format = "0"
        pe.cell(r, 17, '=IFERROR(%s+0,"")' % w_(pb, "own", 2 + i)).number_format = "0.00"
        for c_ in range(10, 18):
            cell = pe.cell(r, c_)
            cell.border = HAIR
            if c_ >= 12:
                cell.alignment = Alignment(horizontal="center" if c_ == 12 else "right", vertical="center")
                cell.font = F_SELF if mine else F_BODY
            if mine:
                cell.fill = SELF_FILL
    _note(pe, lay["banks_note"], RIGHT, NOTE_ROWS["banks_note"], words=t.words, text=(
        "Ranking mean: the rate's mean over the window set above, for banks switched on and kept in enough of "
        "it; best and worst in pack are the lowest and highest. Book index: too small, under the minimum at the "
        "base; no book, none there. Every term is on the Glossary tab."))
    # Each chart is as tall as the rows it covers, whatever height a note on
    # the other half gave them, so nothing under it is pushed down or covered.
    for ch, fmt, skip, (a_, z_), r0, n_ in charts:
        tall = sum(pe.row_dimensions[r].height or 15 for r in range(r0, r0 + n_)) * 2.54 / 72
        pe.add_chart(ch.finish(fmt, skip, round(tall - 0.1, 2), _width_cm(z_ - a_ + 1)), "%s%d" % (L(a_), r0))
    pe.freeze_panes = "A%d" % (R_RESOLVED + 1)
    pe.sheet_view.zoomScale = 100

    # ---- Glossary ----
    K.brand_banner(gl, 1, 3, "Glossary", "What each term means, and exactly how the workbook calculates it.")
    K.header_row(gl, 4, ["Term", "What it means", "How it is calculated"])
    for term, meaning, calc in PEER_TERMS:
        gl.append([term, t.words(meaning), t.words(calc)])
        for c_ in range(1, 4):
            gl.cell(gl.max_row, c_).alignment = Alignment(wrap_text=True, vertical="top")
        gl.cell(gl.max_row, 1).font = Font(bold=True)
    for c_, w in zip("ABC", (30, 60, 80)):
        gl.column_dimensions[c_].width = w
    K.freeze_below(gl, 4)

    # ---- Run ----
    run.append(["threshold-engine", VERSION])
    run.append(["Built", dt.datetime.now().isoformat(timespec="seconds")])
    run.append(["Sources"] + list(SOURCES))
    for s in sources:
        run.append(["", s])
    qs = sorted(data)
    run.append(["Quarter-ends in the pull", len(qs), "%s to %s" % (qs[0], qs[-1]) if qs else ""])
    run.append(["Pull fingerprint (SHA-256 of the rows as canonical JSON)", _fingerprint(data)])
    # Pull dates: a cache file is written when its quarter is fetched, so its
    # modification time is when that quarter was pulled.
    if cache_dir is not None:
        files = sorted(Path(cache_dir).glob("q*.json"))
        stamps = sorted(_written(f) for f in files)
        run.append(["Cache folder", str(cache_dir), "%d files" % len(files)])
        run.append(["Pulled (cache files written)", "%s to %s" % (stamps[0], stamps[-1]) if stamps else "no files"])
        run.append(["Cache file", "Pulled (written)", "SHA-256"])
        for f in files:
            run.append([f.name, _written(f), hashlib.sha256(f.read_bytes()).hexdigest()])
    else:
        run.append(["Pulled", "not recorded: no cache folder was given, so the rows came straight from the caller"])
    run.append(["Acquisitions record (SHA-256 of it as canonical JSON)", _fingerprint(mergers),
                "%d banks, %d acquisition quarters" % (len(mergers), sum(len(v) for v in mergers.values()))])
    if mergers_file is not None:
        mf = Path(mergers_file)
        run.append(["Acquisitions file", str(mf), "pulled (written) %s, SHA-256 %s"
                    % (_written(mf), hashlib.sha256(mf.read_bytes()).hexdigest())])
    else:
        run.append(["Acquisitions file", "not named by the caller"])
    run.append([])
    run.append(["Rules, as stated"])
    run.append(["Jump", jump, "a quarter-on-quarter balance move above this is structural"])
    run.append(["Minimum book ($000)", min_book, "KeyBank exempt" if exempt_self else "KeyBank not exempt"])
    run.append(["Banks needed for a composite or band (starting)", min_banks])
    run.append(["Pack (starting)", pack])
    run.append(["Best and worst (starting)", "%s rate, last %d quarters, at least %d kept"
                % (rank_by, last_n, min_obs)])
    run.append(["Balance index", "base %s" % base_lbl, "books of at least %s ($000) at the base" % index_min_book])
    run.append(["Declared acquisitions", "; ".join("%s: %s" % (c, ", ".join(v))
                                                   for c, v in sorted(extra_acquisitions.items())) or "none"])
    run.append(["Lagged", "years %s" % ", ".join(str(y) for y in lagged_years), "lag %d" % lag_years])
    run.append(["Normalisation", "%d to %d" % (y0, y1)])
    run.append(["Mix base year", mix_base_year, "industry supplied" if industry else "no industry supplied"])
    run.append(["Charts show (starting)", START_WINDOW, "only the picture: no figure depends on it"])
    if dq:
        dqs = sorted(dq["data"])
        run.append(["Delinquency", ", ".join(dq["products"]),
                    "quarter-ends %s to %s" % (label(dqs[0]), label(dqs[-1]))])
        run.append(["Delinquency pull fingerprint (SHA-256 of the rows as canonical JSON)", _fingerprint(dq["data"])])
        run.append(["Delinquency jump", dq["jump"], "blanks a four-quarter average; never a quarter-end rate"])
        run.append(["Delinquency minimum book ($000)", dq["min_book"],
                    "KeyBank exempt" if dq["exempt_self"] else "KeyBank not exempt"])
        run.append(["Delinquency pickers (starting)", "%s, %s" % (dq["measure"], dq["basis"])])
    else:
        run.append(["Delinquency", "not built"])
    run.append([])
    run.append(["Acquisition quarters, per bank", "taken out of the book index without acquisitions (declared: "
                "named at the build, not in the FDIC record)"])
    for b in roster:
        run.append([b["short"], ", ".join(q + (" (declared)" if d else "") for q, d in acq_of[b["cert"]])
                    or "none"])
        run.cell(run.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    run.append([])
    run.append(["Roster"])
    run.append(["Certificate", "Name", "Short", "Group", "Regional"])
    for b in roster:
        run.append([b["cert"], b["name"], b["short"], b["group"], "yes" if b["regional"] else "no"])
    run.column_dimensions["A"].width = 44
    run.column_dimensions["B"].width = 60
    run.column_dimensions["C"].width = 66

    sheets = [t.title(x) for x in TABS]
    stray = foreign_refs(wb, sheets)
    if stray:
        raise RuntimeError("a peer tab refers outside the peer tabs, so it would read the wrong one; nothing "
                           "is saved: %s%s" % ("; ".join(stray[:5]), " (and %d more)" % (len(stray) - 5)
                                               if len(stray) > 5 else ""))
    if into is None:
        out = Path(out)
        wb.save(out)
        _na_as_gap(out)

    def cols(block):
        return {name: L(c) for name, c in block.col.items()}

    return {"path": out, "sheets": sheets, "products": products, "quarters": grid,
            "pack": [b["cert"] for b in peers if b["cert"] in in_pack],
            "data_rows": nrow - 1,
            # Where things are: the dashboard's rows and cells, and each
            # view's helper columns on Workings (values from row 2).
            "layout": dict(lay, pickers={key: picker_cell(key) for key, _, _, _ in PICKERS},
                           settings_cells={key: setting_cell(key, NP) for key, _, _, _ in SETTINGS},
                           include={b["short"]: include_cell(j) for j, b in enumerate(peers)}),
            "scalars": dict(S),
            "charge_offs": {"helpers": cols(co), "first_row": 2, "quarters": grid},
            "book": {"helpers": cols(bk), "first_row": 2, "quarters": grid[base_k:]},
            "lagged": {"helpers": cols(lg), "first_row": 2},
            "mix": {"helpers": cols(mx), "first_row": 2, "years": years},
            "banks": {"helpers": cols(pb), "first_row": 2},
            "ranking": {"first_row": 2, "columns": dict(peer="L", on="M", kept="N", mean="O")},
            # None when not built; else its quarters, helper columns and first row.
            "delinquency": None if dl is None else {"quarters": dq["grid"], "products": dq["products"],
                                                    "helpers": cols(dl), "first_row": 2, "data_rows": d_rows}}
