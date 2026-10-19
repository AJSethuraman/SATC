"""The peer tabs written into another workbook (``into``, ``names``), so the
thresholds workbook and the peers are one file built by Python alone.

Five things are proved here:

1. The series CSVs ``peers build`` writes (they feed the thresholds workbook)
   are byte for byte what they were before the round 2 layout: run.build is
   run twice over the same cache, once with today's builder and once with the
   dashboard builder it replaced (tests/data/peers/peers_workbook_96ee5f2f.py,
   copied unchanged from commit 96ee5f2f). The pin here before, against the
   one-tab-per-view builder (peers_workbook_642ff55f.py), is retired: the
   standalone output was changed on purpose (15 Oct 2026, Claude Design's
   round 2), and the builder each change replaces is the one frozen.
2. Written into a workbook that already has its own Data, Glossary and Run
   tabs, each holding a value no peer cell could produce, every peer tab
   recalculates (the ``formulas`` engine) to exactly the standalone build's
   values, cell for cell, under the titles the book gives them
   (book.PEER_TAB_NAMES); the peer glossary and run record are second
   sections of the host's own Glossary and Run, after what they held, with
   the standalone tabs' content; and no formula, validation, conditional
   format, hyperlink, defined name or chart series on a peer tab names a tab
   that is not a peer tab.
3. In the book's order (the peers first, the host's Glossary and Run after,
   then the save) the sections land after the host's own content, once.
4. ``foreign_refs``, the check the builder runs before it returns, finds a
   stray reference of each of those kinds when one is planted.
5. Bad titles, and ``out`` with ``into``, are refused in words.
"""
import datetime as dt
import importlib.util
import json
import re
import types
import zipfile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from conftest import DATA, ROOT
from threshold_engine import reporting_changes as RC
from threshold_engine.book import PEER_TAB_NAMES
from threshold_engine.peers import peers_workbook as PW
from threshold_engine.peers.glossary import PEER_TERMS
from threshold_engine.peers import roster as R
from threshold_engine.peers import run

PEERS = DATA / "peers"
CHANGES = RC.load(ROOT / "config" / "reporting-changes.json")
TWO = ("Credit card", "Mortgage")
JUMP, MIN_BOOK = 0.25, 1_000_000
NAMES = dict(PEER_TAB_NAMES)
OWN = [t for t in PW.TABS if t not in PW.FOLDED]
NONE_ACQUIRED = {p: {} for p in ("Credit card", "Mortgage", "Home equity", "Other consumer", "Auto")}
#: No stated acquisition in any product (question 18: one {cert: [quarter-ends]} per product).
NONE_STATED = {p: {} for p in NONE_ACQUIRED}
#: What the host's own Data, Glossary and Run hold: a peer formula pointed at
#: one of them would read one of these.
HOST = {"Data": ("HOST DATA A1", 987654321.0), "Glossary": ("HOST GLOSSARY A1", 876543219.0),
        "Run": ("HOST RUN A1", 765432198.0)}


class _Frozen(dt.datetime):
    @classmethod
    def now(cls, tz=None):
        return cls(2026, 10, 14, 9, 30, 0)


FROZEN = types.SimpleNamespace(datetime=_Frozen, date=dt.date)


def _old():
    """The dashboard builder this layout replaced (96ee5f2f), loaded inside the
    package so its relative imports find today's books, roster and glossary."""
    spec = importlib.util.spec_from_file_location("threshold_engine.peers._peers_workbook_96ee5f2f",
                                                  PEERS / "peers_workbook_96ee5f2f.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


OLD = _old()


@pytest.fixture(scope="module")
def recent():
    data = json.loads((PEERS / "roster_2006.json").read_text())
    return {q: rows for q, rows in data.items() if q >= "20170331"}


@pytest.fixture(scope="module")
def dl_data():
    return json.loads((PEERS / "delinquency_roster.json").read_text())


@pytest.fixture(scope="module")
def roster():
    return R.load(ROOT / "config" / "peers.json")


def settings(**over):
    kw = dict(exempt_self=True, min_banks=5, index_min_book=500_000, lagged_years=(2019, 2025), lag_years=2,
              normalisation_years=(2019, 2025), mix_base_year=2019, industry=None, delinquency=None,
              acquired=NONE_ACQUIRED, rank_by="net", last_n=20, min_obs=12)
    kw.update(over)
    return kw


def build(module, out, data, roster, products, pack="all", **over):
    """``module``'s build_peers_workbook with the tests' settings, its clock frozen."""
    mp = pytest.MonkeyPatch()
    mp.setattr(module, "dt", FROZEN)
    try:
        return module.build_peers_workbook(out, data, roster, products, JUMP, MIN_BOOK, pack,
                                           json.loads((PEERS / "mergers_roster.json").read_text())["acquirers"],
                                           {"6672": ["2022Q2"]}, "2019Q4", **settings(**over),
                                           **({} if module is OLD else {"combined": {}, "opened": {}, "changes": CHANGES, "stated": NONE_STATED}))
    finally:
        mp.undo()


def dq(dl_data, **over):
    spec = dict(data=dl_data, products=TWO, measure="90+ accruing", basis="quarter-end",
                jump=0.25, min_book=1_000_000, exempt_self=True, acquired=NONE_ACQUIRED)
    spec.update(over)
    return spec


# ---- 1. the series CSVs are byte for byte as they were ----

def dq_cache(base, dl_data):
    """delinquency_roster.json with each bank's charge-off fields, laid out as
    the cache 'peers pull --delinquency yes' writes (as test_peers_cli)."""
    co = [json.loads((PEERS / n).read_text()) for n in ("roster_2006.json", "roster_1992.json")]
    cache = base / "cache"
    cache.mkdir()
    for q, rows in dl_data.items():
        merged = []
        for r in rows:
            row = dict(r)
            for d in co:
                row.update(next(x for x in d[q] if x["CERT"] == r["CERT"]))
            merged.append(row)
        (cache / ("q%s.json" % q)).write_text(json.dumps(merged), encoding="utf-8")
    return cache


def run_build(out, cache, series, **over):
    kw = dict(cache=cache, roster=ROOT / "config" / "peers.json", series_dir=series, first="2017Q1",
              last="2026Q2", products=list(TWO), mergers=PEERS / "mergers_roster.json",
              absorptions=PEERS / "absorptions_none.json", extra_acquisitions=["6672=2022Q2"],
              industry_combined=[], jump=0.25, acquisition_share=0.05, min_book=1_000_000.0, exempt_self=True,
              pro_forma=False, pack="all", min_banks=5, rank_by="net", last_n=20, min_obs=12, base_quarter="2019Q4",
              index_min_book=500_000.0, lagged_years=[2019, 2025], lag_years=2, normalisation_years=[2019, 2025],
              mix_base_year=2019, delinquency=True, reporting_changes=ROOT / "config" / "reporting-changes.json",
              stated_acquisitions=None,
              say=lambda s: None)
    kw.update(over)
    return run.build(out, **kw)


def tree(folder):
    return {str(p.relative_to(folder)): p.read_bytes() for p in sorted(folder.rglob("*")) if p.is_file()}


def test_peers_build_writes_the_same_series_as_before_the_round_2_layout(tmp_path, dl_data, monkeypatch):
    """Through peers.run.build (a cache folder, a mergers file, the industry,
    delinquency, each peer's own book): every series CSV, byte for byte, as
    the builder it replaced wrote it. The workbooks differ, on purpose; the
    files the thresholds workbook reads do not."""
    cache = dq_cache(tmp_path, dl_data)
    new = run_build(tmp_path / "after.xlsx", cache, tmp_path / "after")
    assert load_workbook(new["out"]).sheetnames == list(PW.TABS)
    monkeypatch.setattr(PW, "build_peers_workbook",
                        lambda *a, into=None, names=None, combined=None, opened=None, changes=None, stated=None,
                        book_kind=None, **kw:
                        OLD.build_peers_workbook(*a, **kw))
    old = run_build(tmp_path / "before.xlsx", cache, tmp_path / "before")
    assert load_workbook(old["out"]).sheetnames == ["Peers", "Workings", "Glossary", "Data", "DelinquencyData",
                                                    "Run"]
    before, after = tree(tmp_path / "before"), tree(tmp_path / "after")
    assert len(after) > 100 and any(k.startswith("banks") for k in after)
    assert sorted(before) == sorted(after)
    assert [k for k in before if before[k] != after[k]] == []
    assert [(p.relative_to(tmp_path / "after"), n) for p, n, *_ in new["series"]] == \
        [(p.relative_to(tmp_path / "before"), n) for p, n, *_ in old["series"]]


# ---- 2. into a workbook with its own Data, Glossary and Run ----

def host():
    wb = Workbook()
    wb.active.title = "Data"
    wb.create_sheet("Glossary")
    wb.create_sheet("Run")
    for title, (a1, b2) in HOST.items():
        wb[title]["A1"], wb[title]["B2"] = a1, b2
    return wb


def recalc(path):
    """Every cell's value after a full recalculation, keyed 'SHEET!A1', as
    test_peers_workbook does it."""
    formulas = pytest.importorskip("formulas", reason="the formulas engine recalculates the workbook")
    sol = formulas.ExcelModel().loads(str(path)).finish().calculate()
    out = {}
    for key, val in sol.items():
        v = getattr(val, "value", val)
        try:
            v = v.item(0) if hasattr(v, "item") else v
        except Exception:                          # noqa: BLE001
            pass
        out[key.split("]")[-1].upper().replace("'", "")] = v
    return out


def same_value(v):
    """A value as compared: numbers as floats (a bool stays a bool), anything
    else (text, an error such as #N/A) as its text."""
    if isinstance(v, bool) or v is None:
        return v
    if isinstance(v, (int, float)):
        return float(v)
    return str(v)


@pytest.fixture(scope="module")
def into_and_alone(tmp_path_factory, recent, dl_data, roster):
    """The same build twice, frozen clock: standalone, and into host() under
    the book's titles, saved and given na_as_gap over the peer tabs."""
    base = tmp_path_factory.mktemp("into")
    kw = dict(product="Mortgage", delinquency=dq(dl_data))
    alone = build(PW, base / "alone.xlsx", recent, roster, TWO, **kw)
    wb = host()
    info = build(PW, None, recent, roster, TWO, into=wb, names=NAMES, **kw)
    path = base / "into.xlsx"
    wb.save(path)
    PW.na_as_gap(path, info["sheets"])
    return dict(alone=alone, info=info, wb=wb, path=path)


def section_rows(ws, first_row, first_col, n_cols):
    """The values of a block of rows from ``first_row``, ``n_cols`` wide from ``first_col``, blank rows left out."""
    out = []
    for r in range(first_row, ws.max_row + 1):
        vals = tuple(ws.cell(r, c).value for c in range(first_col, first_col + n_cols))
        if any(v is not None for v in vals):
            out.append(vals)
    return out


def test_into_writes_the_peer_tabs_after_the_hosts_and_saves_nothing(into_and_alone):
    info, wb = into_and_alone["info"], into_and_alone["wb"]
    assert info["path"] is None
    assert info["sheets"] == [NAMES[t] for t in OWN]
    assert info["folded"] == {"Glossary": "Glossary", "Run": "Run"} == {t: NAMES[t] for t in PW.FOLDED}
    assert wb.sheetnames == ["Data", "Glossary", "Run"] + info["sheets"]
    assert wb.active.title == "Data"                       # the host's active tab is left as it was
    assert wb[NAMES["Workings"]].sheet_state == "hidden"
    ws = wb["Data"]
    assert (ws["A1"].value, ws["B2"].value, ws.max_row, ws.max_column) == HOST["Data"] + (2, 2)
    for title in PW.FOLDED:                                # the host's own content is where it was
        ws = wb[title]
        assert (ws["A1"].value, ws["B2"].value) == HOST[title]
        assert [ws.cell(3, c).value for c in range(1, 5)] == [None] * 4
    alone = into_and_alone["alone"]
    assert {k: v for k, v in info.items() if k not in ("path", "sheets", "folded")} == \
        {k: v for k, v in alone.items() if k not in ("path", "sheets", "folded")}


def test_the_glossary_and_run_are_second_sections_of_the_hosts(into_and_alone):
    """The host's Glossary carries "Peers terms" after its own content: the
    standalone Glossary's rows, word for word (each naming the renamed tabs);
    its Run carries the standalone Run's rows under "Peers: " headings."""
    wb, alone = into_and_alone["wb"], load_workbook(into_and_alone["alone"]["path"])
    words = PW._Tabs(NAMES).words
    g = wb["Glossary"]
    assert g["A4"].value == "Peers terms" and [g.cell(5, c).value for c in (1, 2, 3)] == \
        ["Term", "What it means", "How it is calculated"]
    mine = section_rows(alone["Glossary"], 5, 2, 3)
    theirs = section_rows(g, 6, 1, 3)
    assert len(mine) == len(PEER_TERMS) == len(theirs)
    assert theirs == [tuple(words(v) for v in r) for r in mine]
    assert any("Peers data tab" in v for r in theirs for v in r)
    r_ = wb["Run"]
    assert r_["A4"].value == "Peers: this build"
    heads = [r_.cell(r, 1).value for r in range(4, r_.max_row + 1) if r_.cell(r, 1).font.b and
             r_.cell(r, 1).font.sz == 11]
    assert heads == ["Peers: this build", "Peers: rules, as stated", "Peers: acquisition quarters, per bank",
                     "Peers: roster"]
    body = [r for r in section_rows(r_, 4, 1, 5) if r[0] not in heads]
    want = [r for r in section_rows(alone["Run"], 4, 2, 5)
            if r[0] not in ("This build", "Rules, as stated", "Acquisition quarters, per bank", "Roster")]
    assert body == want and any(r[0] == "Jump" for r in body)


@pytest.fixture(scope="module")
def recalculated(into_and_alone):
    return recalc(into_and_alone["alone"]["path"]), recalc(into_and_alone["path"])


def test_every_peer_cell_recalculates_to_the_standalone_value(into_and_alone, recalculated):
    """Cell for cell, every value on the peer tabs equal. The one difference
    allowed is the prose naming a tab, which names the renamed tab (the 'Data
    tab' is the 'Peers data tab'): the standalone text, renamed, must equal
    it, and those cells are counted so the allowance cannot grow unseen."""
    alone, into = recalculated
    words = PW._Tabs(NAMES).words
    host_values = {same_value(v) for a1, b2 in HOST.values() for v in (a1, b2)}
    compared, renamed, live = 0, 0, 0
    wb = load_workbook(into_and_alone["alone"]["path"])
    for tab in OWN:
        mine = {k.split("!", 1)[1]: v for k, v in alone.items() if k.split("!", 1)[0] == tab.upper()}
        theirs = {k.split("!", 1)[1]: v for k, v in into.items() if k.split("!", 1)[0] == NAMES[tab].upper()}
        assert mine, tab
        assert sorted(mine) == sorted(theirs), tab
        bad = []
        for c in mine:
            want, got = same_value(mine[c]), same_value(theirs[c])
            if want != got:
                if isinstance(mine[c], str) and words(mine[c]) == theirs[c]:
                    renamed += 1
                else:
                    bad.append((c, mine[c], theirs[c]))
        assert not bad, (tab, bad[:5])
        compared += len(mine)
        live += sum(1 for row in wb[tab].iter_rows() for c in row if c.data_type == "f")
        assert not host_values & {same_value(v) for v in theirs.values()}, tab
    assert compared > 10_000 and live > 5_000
    # The notes on the dashboard that name the Data or DelinquencyData tab.
    assert 1 <= renamed <= 4, renamed


def test_the_recalculated_dashboard_is_live_not_blank(into_and_alone, recalculated):
    """The comparison above is over real numbers: best and worst resolve and
    the charge-off composite has values, in both."""
    alone, into = recalculated
    s = into_and_alone["info"]["scalars"]
    best = s["best"].replace("$", "")
    assert alone["WORKINGS!" + best] == into["PEERS WORKINGS!" + best] and isinstance(alone["WORKINGS!" + best], str)
    nums = [v for k, v in into.items() if k.startswith("PEERS WORKINGS!") and isinstance(v, float)]
    assert len(nums) > 1000
    assert alone["PEERS!B5"] == into["PEERS!B5"] == "Mortgage"


def test_no_peer_tab_refers_outside_the_peer_tabs(into_and_alone):
    """In memory and as saved, every formula, validation, conditional format,
    hyperlink, defined name and chart series on a peer tab names peer tabs
    only; the charts read the names on Peers, which read Peers workings."""
    sheets = into_and_alone["info"]["sheets"]
    for wb in (into_and_alone["wb"], load_workbook(into_and_alone["path"])):
        assert PW.foreign_refs(wb, sheets) == []
    wb = load_workbook(into_and_alone["path"])
    names = wb["Peers"].defined_names
    assert names and all(dn.attr_text.startswith("OFFSET('Peers workings'!") for dn in names.values())
    with zipfile.ZipFile(into_and_alone["path"]) as z:
        charts = [z.read(n).decode() for n in z.namelist() if n.startswith("xl/charts/chart")]
    refs = [f for c in charts for f in re.findall(r"<f>([^<]*)</f>", c)]
    assert len(charts) == 5 and refs                     # two panels, book size, the ratio, the mix
    assert all(f.startswith(("Peers!", "'Peers workings'!")) for f in refs), refs[:5]
    assert all("dispNaAsBlank" in c for c in charts)


def test_the_prose_names_the_renamed_tabs(into_and_alone):
    """A line telling the reader to look at 'the Data tab' would send them to
    the host's: every such line names the peer tab's title. The Glossary and
    Run are the host's own, which now hold the peer sections."""
    wb = into_and_alone["wb"]
    moved = [t for t in PW.TABS if NAMES[t] != t]
    assert moved == ["Data", "DelinquencyData", "Workings"]
    stale = re.compile(r"\b(%s)\b(?=(?: and [A-Z]\w*)? tabs?\b)|\bDelinquencyData\b" % "|".join(moved))
    seen = []
    for title in into_and_alone["info"]["sheets"] + list(PW.FOLDED):
        for row in wb[title].iter_rows():
            for c in row:
                if isinstance(c.value, str) and " tab" in c.value:
                    seen.append(c.value)
                    assert not stale.search(c.value), (title, c.coordinate, c.value)
    assert any("Peers data tab" in v for v in seen) and any("Peers delinquency data tab" in v for v in seen)
    assert any("the Glossary tab" in v for v in seen) and any("the Run tab" in v for v in seen)


# ---- 3. in the book's order: the peers first, the host's tabs after ----

def test_the_sections_go_after_the_hosts_content_when_its_tabs_come_later(tmp_path, recent, roster):
    """book.build writes the peers into an empty workbook, then its own
    Glossary and Run, then saves: the save writes each peer section after the
    host's own rows, in the column its title starts in, and a second save
    writes nothing more. A host that never makes the tab gets one."""
    wb = Workbook()
    info = build(PW, None, recent, roster, ("Credit card",), into=wb, names=NAMES)
    assert info["sheets"] == [NAMES[t] for t in OWN] and "Glossary" not in wb.sheetnames
    g = wb.create_sheet("Glossary")
    g["B1"], g["B4"], g["B12"] = "Glossary", "Term", "The host's last term"
    path = tmp_path / "book.xlsx"
    wb.save(path)
    wb.save(path)
    saved = load_workbook(path)
    g = saved["Glossary"]
    assert g["B14"].value == "Peers terms" and g["B15"].value == "Term" and g["B16"].value == "Pack"
    assert [r for r in range(1, g.max_row + 1) if g.cell(r, 2).value == "Peers terms"] == [14]
    assert g.max_row == 15 + len(PEER_TERMS)
    r_ = saved["Run"]                                      # nothing made a Run: the save did
    assert r_["A1"].value == "Peers: this build" and saved.sheetnames[-1] == "Run"
    assert "save" not in wb.__dict__                       # the workbook saves as any other from here


def test_na_as_gap_over_the_peer_tabs_leaves_the_hosts_charts_alone(tmp_path, recent, roster):
    wb = host()
    for r in range(1, 6):
        wb["Data"].cell(r, 3, r)
    lc = LineChart()
    lc.add_data(Reference(wb["Data"], min_col=3, min_row=1, max_row=5))
    wb["Data"].add_chart(lc, "E2")
    info = build(PW, None, recent, roster, ("Credit card",), into=wb, names=NAMES)
    wb.save(tmp_path / "x.xlsx")
    PW.na_as_gap(tmp_path / "x.xlsx", info["sheets"])
    with zipfile.ZipFile(tmp_path / "x.xlsx") as z:
        charts = {n: z.read(n).decode() for n in z.namelist() if n.startswith("xl/charts/chart")}
    host_charts = [n for n, c in charts.items() if "'Data'!" in c]
    assert len(host_charts) == 1 and "dispNaAsBlank" not in charts[host_charts[0]]
    assert len(charts) > 1 and all("dispNaAsBlank" in c for n, c in charts.items() if n not in host_charts)
    with pytest.raises(ValueError, match="no tab called Nope"):
        PW.na_as_gap(tmp_path / "x.xlsx", ["Nope"])


def test_peers_build_into_a_workbook_writes_the_series_and_no_peer_file(tmp_path, dl_data):
    cache = dq_cache(tmp_path, dl_data)
    wb = host()
    got = run_build(None, cache, tmp_path / "series", into=wb, names=NAMES)
    assert got["out"] is None and got["info"]["sheets"] == [NAMES[t] for t in OWN]
    assert wb.sheetnames[3:] == got["info"]["sheets"]
    written = [p for p, n, *_ in got["series"] if n]
    assert len(written) > 12 and all(p.exists() for p in written)
    assert sorted(f.name for f in tmp_path.iterdir()) == ["cache", "series"]
    with pytest.raises(ValueError, match="out must be None"):
        run_build(tmp_path / "x.xlsx", cache, tmp_path / "s2", into=host(), names=NAMES)


# ---- 5. refusals ----

@pytest.mark.parametrize("names, why", [
    ({k: v for k, v in NAMES.items() if k != "Run"}, "missing: Run"),
    (dict(NAMES, Extra="Peers extra"), "not a tab: Extra"),
    (dict(NAMES, Pack="Peers pack"), "not a tab: Pack"),
    (dict(NAMES, Run="peers DATA"), "same title"),
    (dict(NAMES, Run="Peers: run"), "must be 1 to 31"),
    (dict(NAMES, Run="P" * 32), "must be 1 to 31"),
    (dict(NAMES, Run="History"), "not History"),
    (["Peers pack"], "names must be None"),
])
def test_bad_names_are_refused(names, why, recent, roster):
    with pytest.raises(ValueError, match=re.escape(why)):
        build(PW, None, recent, roster, ("Credit card",), into=host(), names=names)


def test_a_title_the_host_already_has_is_refused(recent, roster):
    wb = host()
    with pytest.raises(ValueError, match="already has a tab called Data;"):
        build(PW, None, recent, roster, ("Credit card",), into=wb)
    wb2 = host()
    wb2.create_sheet("PEERS WORKINGS")
    with pytest.raises(ValueError, match="already has a tab called PEERS WORKINGS"):
        build(PW, None, recent, roster, ("Credit card",), into=wb2, names=NAMES)
    assert wb.sheetnames == ["Data", "Glossary", "Run"] and wb2.sheetnames[-1] == "PEERS WORKINGS"


def test_out_and_into_are_one_or_the_other(tmp_path, recent, roster):
    with pytest.raises(ValueError, match="out must be None"):
        build(PW, tmp_path / "x.xlsx", recent, roster, ("Credit card",), into=host(), names=NAMES)
    with pytest.raises(ValueError, match="out is needed"):
        build(PW, None, recent, roster, ("Credit card",))
    with pytest.raises(ValueError, match="openpyxl Workbook"):
        build(PW, None, recent, roster, ("Credit card",), into="book.xlsx", names=NAMES)
    with pytest.raises(ValueError, match="write-only"):
        build(PW, None, recent, roster, ("Credit card",), into=Workbook(write_only=True), names=NAMES)
    assert not list(tmp_path.iterdir())


def test_a_title_that_needs_quoting_is_quoted():
    assert PW._quoted("Peers") == "Peers" and PW._quoted("Peers workings") == "'Peers workings'"
    assert PW._quoted("A1") == "'A1'" and PW._quoted("R") == "'R'" and PW._quoted("O'Neil") == "'O''Neil'"
    assert PW._quoted("Q4data") == "'Q4data'"


# ---- 4. the check can fail ----

def planted():
    """A peer tab beside a host Data tab, and one reference of each kind on
    it: the builder must never leave any of these."""
    wb = Workbook()
    wb.active.title = "Data"
    ws = wb.create_sheet("Peers workings")
    ws["A1"] = "='Peers workings'!B1+1"                     # its own tab: fine
    ws["A2"] = '="Data!A1 in a string is text"'              # a string: fine
    return wb, ws


@pytest.mark.parametrize("kind", ["formula", "lowercase formula", "quoted formula", "validation",
                                  "conditional format", "hyperlink", "defined name", "chart"])
def test_foreign_refs_finds_each_kind_of_stray_reference(kind):
    wb, ws = planted()
    assert PW.foreign_refs(wb, ["Peers workings"]) == []
    if kind == "formula":
        ws["B3"] = "=SUM(Data!$A$1:$A$4)"
    elif kind == "lowercase formula":
        ws["B3"] = "=data!A1"
    elif kind == "quoted formula":
        ws["B3"] = "='Data'!A1"
    elif kind == "validation":
        dv = DataValidation(type="list", formula1="=Data!$A$1:$A$3")
        ws.add_data_validation(dv)
        dv.add("B3")
    elif kind == "conditional format":
        ws.conditional_formatting.add("B3", FormulaRule(formula=["Data!$A$1>0"]))
    elif kind == "hyperlink":
        ws["B3"] = "see"
        ws["B3"].hyperlink = "#Data!A1"
        ws["B3"].hyperlink.location = "Data!A1"
    elif kind == "defined name":
        ws.defined_names["stray"] = DefinedName("stray", attr_text="OFFSET(Data!$A$1,0,0,3,1)")
    elif kind == "chart":
        lc = LineChart()
        lc.add_data(Reference(wb["Data"], min_col=1, min_row=1, max_row=4))
        ws.add_chart(lc, "D2")
    found = PW.foreign_refs(wb, ["Peers workings"])
    assert len(found) == 1 and re.search("names data", found[0], re.I), found


def test_the_builder_refuses_to_return_with_a_stray_reference(recent, roster, monkeypatch):
    """Were a literal ``Data!`` to creep back into the builder, the build
    stops rather than hand back a workbook that reads the host's Data."""
    real = PW._Tabs.ref
    monkeypatch.setattr(PW._Tabs, "ref", lambda self, tab: "Data!" if tab == "Data" else real(self, tab))
    wb = host()
    with pytest.raises(RuntimeError, match="refers outside the peer tabs.*names Data"):
        build(PW, None, recent, roster, ("Credit card",), into=wb, names=NAMES)
    assert wb.sheetnames == ["Data", "Glossary", "Run"]
    assert [(wb[t]["A1"].value, wb[t]["B2"].value) for t in HOST] == list(HOST.values())
