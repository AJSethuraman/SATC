"""The Internal and Internal data tabs (internal_tab.py).

Every figure here is made up by the test: four outside sources built so that
one agrees with the internal composite, one does not, one has too few
quarters and one has lines that are refused; internal rows for two segments
of a card book and an auto book with CRR grades. Nothing internal is ever in
the repository.

The proof is a Python reference. The composite comes from internal.py
(``composite_series``), the lines from the recalculated Thresholds tab, and
every figure in the reliability table, the verdict included, is worked out
again here and compared with what the formulas give.
"""
import math
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from threshold_engine import internal as I
from threshold_engine import internal_tab as IT
from threshold_engine.glossary import TERMS
from threshold_engine.workbook import build
from test_workbook import recalc

NAMES = ["Product", "Measure", "Source"]
SETTINGS = dict(name="N", unit="%", direction="higher_is_worse", frequency="quarterly", smoothing=1,
                floor_at_zero=True, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5, on_the_line="worse",
                half_lives=(), percentiles=(50,), horizon=4, min_history=41)
NET = "Net charge-offs"
QS = ["%dQ%d" % (2005 + i // 4, i % 4 + 1) for i in range(60)]          # 2005Q1-2019Q4
INT_QS = QS[12:]                                                           # 2008Q1-2019Q4


def path_rate(i):
    """A made-up annual loss rate, %, with one stress stretch."""
    return 3 + 1.2 * math.sin(i / 5.0) + (5.0 if 16 <= i <= 23 else 0)


def internal_rows():
    """Card: two segments, no All row. Auto: All rows with a CRR grade. In
    $ thousands, made up."""
    rows = []
    for i, q in enumerate(INT_QS):
        x = path_rate(i + 12)
        for seg, bal, tilt in (("Bank card", 600000 + 2000 * i, 1.1), ("Private label", 400000 - 1000 * i, 0.85)):
            nco = round(bal * x * tilt / 400)
            rec = round(nco * 0.12)
            rows.append(I.Row(0, q, "Credit card", seg, Decimal(bal), Decimal(nco + rec), Decimal(nco), Decimal(rec),
                              None, Decimal(round(bal * 0.03)), Decimal(round(bal * 0.02)), Decimal(round(bal * 0.01)),
                              ()))
        bal = 250000 + 500 * i
        nco = round(bal * (0.8 + 0.3 * math.cos(i / 4.0) + (1.5 if 10 <= i <= 15 else 0)) / 400)
        rows.append(I.Row(0, q, "Auto", "All", Decimal(bal), Decimal(nco + 50), Decimal(nco), Decimal(50), None,
                          Decimal(round(bal * 0.02)), Decimal(round(bal * 0.01)), Decimal(round(bal * 0.004)),
                          (("crr_grade", Decimal(1 + (i % 9))),)))
    return rows


def the_internal():
    rows = tuple(internal_rows())
    return I.Internal("synthetic", {}, {"Bank card": "x", "Private label": "x", "All": "x"}, (), rows)


def write_series(path, pairs):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("date,value\n")
        for q, v in pairs:
            fh.write("%s,%r\n" % (I.books.quarter_end(q), v))
    return path


@pytest.fixture(scope="module")
def book(tmp_path_factory):
    """Sources made from the internal composite: Agrees tracks it, Disagrees
    has its stress years later, Short has ten quarters, Refused sits below zero."""
    d = tmp_path_factory.mktemp("internal_tab")
    comp = dict((q, v) for q, v in I.composite_series(the_internal(), "Credit card", "net-charge-off-rate")
                if v is not None)
    auto = dict((q, v) for q, v in I.composite_series(the_internal(), "Auto", "net-charge-off-rate")
                if v is not None)
    early = [(q, path_rate(i)) for i, q in enumerate(QS) if q not in comp]
    agree = early + [(q, comp[q] + 0.002 * ((i % 3) - 1)) for i, q in enumerate(QS) if q in comp]
    other = [(q, path_rate((i + 24) % 60)) for i, q in enumerate(QS)]        # its stress comes years later
    short = [(q, comp[q]) for q in QS[-10:]]
    refused = [(q, v - 10) for q, v in agree]
    auto_s = [(q, 1.0) for q in QS if q not in auto] + [(q, auto[q]) for q in QS if q in auto]
    ser = [("Credit card - %s - Agrees" % NET, write_series(d / "a.csv", agree)),
           ("Credit card - %s - Disagrees" % NET, write_series(d / "b.csv", other)),
           ("Credit card - %s - Short" % NET, write_series(d / "c.csv", short)),
           ("Credit card - %s - Refused" % NET, write_series(d / "e.csv", refused)),
           ("Auto - %s - Agrees" % NET, write_series(d / "f.csv", auto_s))]
    dec = {"names": NAMES, "parts": {lbl: tuple(lbl.split(" - ")) for lbl, _ in ser}, "absent": [],
           "compare": None, "notes": [], "internal": {"capacity": 160}, "lines_from": None}
    info = build(d / "book.xlsx", ser, decide=dec, **SETTINGS)
    return d, info, ser


def paste(src, dst, rows, extra=()):
    """The slot filled as an analyst would paste it: the header row at A4 in
    another order, products under the bank's own names, some quarters as
    dates and one row giving only gross charge-offs and recoveries."""
    import datetime as dt
    wb = load_workbook(src)
    ds = wb[IT.DATA_TITLE]
    head = ["segment", "quarter", "product", "balance", "net_charge_offs", "gross_charge_offs", "recoveries",
            "dpd_90_plus", "dpd_60_plus", "dpd_30_plus", "crr_grade", "further_1", "further_2", "further_3",
            "further_4", "further_5"]
    for c, h in enumerate(head, start=1):
        ds.cell(4, c, h)
    names = {"Credit card": "Cards", "Auto": "Auto loans"}
    for k, r_ in enumerate(list(rows) + list(extra)):
        r = 5 + k
        if isinstance(r_, dict):
            for c, h in enumerate(head, start=1):
                if h in r_:
                    ds.cell(r, c, r_[h])
            continue
        y, qq = int(r_.quarter[:4]), int(r_.quarter[5])
        quarter = dt.datetime(y, 3 * qq, 28) if k % 5 == 0 else r_.quarter
        grade = dict(r_.extra).get("crr_grade")
        vals = {"segment": r_.segment, "quarter": quarter, "product": names[r_.product],
                "balance": float(r_.balance), "net_charge_offs": float(r_.net_charge_offs),
                "gross_charge_offs": float(r_.gross_charge_offs), "recoveries": float(r_.recoveries),
                "dpd_90_plus": float(r_.dpd_90_plus), "dpd_60_plus": float(r_.dpd_60_plus),
                "dpd_30_plus": float(r_.dpd_30_plus), "crr_grade": float(grade) if grade is not None else None}
        if k == 3:
            vals["net_charge_offs"] = None              # worked out: gross less recoveries
        for c, h in enumerate(head, start=1):
            if vals.get(h) is not None:
                ds.cell(r, c, vals[h])
    wb.save(dst)
    return recalc(dst)


BAD = {"segment": "All", "quarter": "2019Q4", "product": "Boats", "balance": 10, "net_charge_offs": 1,
       "gross_charge_offs": 1, "recoveries": 0, "dpd_90_plus": 0, "dpd_60_plus": 0, "dpd_30_plus": 0}


@pytest.fixture(scope="module")
def filled(book):
    d, info, ser = book
    vals = paste(d / "book.xlsx", d / "filled.xlsx", internal_rows(), extra=[BAD])
    return d, info, vals


def v(vals, ref):
    return vals.get(ref.upper())


def iv(vals, col, row):
    return v(vals, "INTERNAL!%s%d" % (col, row))


# ---------------------------------------------------------------- the reference

def rating(x, lines, worse, floor=True):
    if floor and x <= 0:
        return 1
    return 1 + sum(1 for l_ in lines if (x > l_) or (x == l_ and worse))


def pearson(xs, ys):
    n = len(xs)
    if n < 2:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    sxy = sum((a - mx) * (b - my) for a, b in zip(xs, ys))
    sxx = sum((a - mx) ** 2 for a in xs)
    syy = sum((b - my) ** 2 for b in ys)
    if sxx == 0 or syy == 0:
        return None
    return sxy / math.sqrt(sxx * syy)


def reference(internal_by_q, source_by_q, lines, worse, refused, mins=(12, 0.7, 0.6)):
    qs = [q for q in QS if q in internal_by_q and q in source_by_q]
    n = len(qs)
    xs, ys = [internal_by_q[q] for q in qs], [source_by_q[q] for q in qs]
    chg = [(internal_by_q[q] - internal_by_q[p], source_by_q[q] - source_by_q[p])
           for q, p in zip(QS[4:], QS) if all(k in internal_by_q and k in source_by_q for k in (q, p))]
    level = pearson(xs, ys)
    change = pearson([a for a, _ in chg], [b for _, b in chg])
    agree = within = None
    if n and not refused:
        ri = [rating(x, lines, worse) for x in xs]
        rs = [rating(y, lines, worse) for y in ys]
        agree = sum(a == b for a, b in zip(ri, rs)) / n
        within = sum(abs(a - b) <= 1 for a, b in zip(ri, rs)) / n
    min_q, min_a, min_c = mins
    if n < min_q:
        verdict = IT.VERDICTS[0]
    elif refused:
        verdict = IT.VERDICTS[1]
    elif agree is not None and change is not None:
        verdict = (IT.VERDICTS[2] if agree >= min_a and change >= min_c else
                   IT.VERDICTS[3] if agree >= 0.75 * min_a else IT.VERDICTS[4])
    else:
        verdict = IT.VERDICTS[3] if agree is not None and agree >= 0.75 * min_a else IT.VERDICTS[4]
    return {"overlap": n, "level": level, "change": change, "gap": (sum(xs) - sum(ys)) / n if n else None,
            "agree": agree, "within": within, "verdict": verdict}


def lines_of(vals, info, label):
    r = IT.thresholds_refs(load_workbook(info["path"]), [label], {label: tuple(label.split(" - "))})[label]
    ls = [v(vals, x.replace("$", "").replace("'", "")) for x in r["lines"]]
    check = v(vals, r["refused"].replace("$", "").replace("'", ""))
    otl = v(vals, r["otl"].replace("$", "").replace("'", ""))
    return ls, otl == "worse", bool(check)


def source_values(ser, label):
    from threshold_engine.series import read_csv
    path = dict(ser)[label]
    from threshold_engine.workbook import quarter
    return {quarter(p.date): p.value for p in read_csv(Path(path))}


# ---------------------------------------------------------------- tests

def test_the_slot_ships_empty_and_says_so(book):
    d, info, _ = book
    wb = load_workbook(d / "book.xlsx")
    ds = wb[IT.DATA_TITLE]
    assert IT.TABLE in ds.tables
    assert ds.tables[IT.TABLE].ref == "A4:P%d" % (4 + 160)
    assert all(ds.cell(r, c).value is None for r in range(5, 165) for c in range(1, 17))
    assert [ds.cell(4, c).value for c in range(1, 12)] == list(IT.SLOT[:11])
    vals = recalc(d / "book.xlsx")
    assert iv(vals, "J", 9) == 0 and iv(vals, "J", 10) == 0
    chart = info["decide"]["internal"]
    # The chart's title cell says there is nothing yet.
    titles = [val for key, val in vals.items() if key.startswith("INTERNAL!") and val == "No internal data yet"]
    assert titles
    assert all(iv(vals, c, chart["verdict_row"]) == IT.VERDICTS[0] for c in "MNOP")


def test_the_workbook_lists_every_internal_tab_heading_in_the_glossary(book):
    d, info, _ = book
    terms = {t for t, _, _ in TERMS}
    for t in ("Internal data", "Mapped", "Composite", "Segment", "Quarters both have", "Level correlation",
              "Change correlation", "Average gap", "Ratings that agree",
              "Within one rating", "Verdict", "CRR grade", "Lowest CRR grade", "Like for like", "Rows not used",
              "Call Report source", "Why not used", "Share of the product", "Balance by segment ($ millions)",
              "What counts as reliable"):
        assert t in terms, t
    ws = load_workbook(d / "book.xlsx")[IT.TITLE]
    rel = info["decide"]["internal"]["rel_top"]
    for r in range(rel + 1, rel + 8):
        assert ws.cell(r, 12).value in terms


def test_names_are_mapped_by_the_guess_rules_and_rows_read(filled):
    d, info, vals = filled
    got = {iv(vals, "B", r): iv(vals, "C", r) for r in range(5, 13) if iv(vals, "B", r)}
    assert got == {"Cards": "Credit card", "Auto loans": "Auto", "Boats": ""}
    assert iv(vals, "J", 9) == len(internal_rows()) + 1
    assert iv(vals, "J", 10) == 1
    # The unmapped row says why; a row with net charge-offs blank had it worked out.
    whys = [v(vals, "INTERNAL DATA!AH%d" % r) for r in range(5, 5 + len(internal_rows()) + 1)]
    assert whys.count("The product is not mapped on the Internal tab.") == 1
    assert all(w in ("", "The product is not mapped on the Internal tab.") for w in whys)
    r4 = internal_rows()[3]
    assert v(vals, "INTERNAL DATA!AE8") == float(r4.net_charge_offs)
    # Quarters read from a date and from text alike.
    assert v(vals, "INTERNAL DATA!R5") == INT_QS[0]


def test_coverage_counts_each_product(filled):
    d, info, vals = filled
    prods = info["decide"]["internal"]["products"]
    col = {p: "MNOPQ"[i] for i, p in enumerate(prods)}
    card, auto = col["Credit card"], col["Auto"]
    assert iv(vals, card, 6) == 48 and iv(vals, auto, 6) == 48
    assert iv(vals, card, 7) == "2008Q1" and iv(vals, card, 8) == "2019Q4"
    assert iv(vals, card, 9) == 2 and iv(vals, auto, 9) == 0
    assert iv(vals, card, 10) == "Summed" and iv(vals, auto, 10) == "All rows"
    last = [r for r in internal_rows() if r.quarter == "2019Q4" and r.product == "Credit card"]
    assert iv(vals, card, 11) == pytest.approx(sum(float(r.balance) for r in last) / 1000)


def test_the_composite_matches_internal_py(filled):
    """The formulas and internal.py give the same composite, every quarter
    and every measure."""
    d, info, vals = filled
    t = info["decide"]["internal"]
    w0 = t["w0"]
    the = the_internal()
    for p in ("Credit card", "Auto"):
        b = t["pblock"][p]
        for key, col in zip(("net-charge-off-rate", "gross-charge-off-rate", "dpd-30-plus-quarter-end",
                             "dpd-60-plus-quarter-end", "dpd-90-plus-quarter-end"), IT_RATE_KEYS):
            ref = dict(I.composite_series(the, p, key))
            from openpyxl.utils import get_column_letter as L
            for i, q in enumerate(t["calendar"]):
                got = v(vals, "INTERNAL!%s%d" % (L(b[col]), w0 + i))
                want = ref.get(q)
                if want is None:
                    assert got in ("", None), (p, key, q, got)
                else:
                    assert got == pytest.approx(want, rel=1e-12), (p, key, q)


IT_RATE_KEYS = ("r_net", "r_gross", "r_d30", "r_d60", "r_d90")


def test_every_figure_in_the_reliability_table_matches_the_reference(filled, book):
    d, info, vals = filled
    _, _, ser = book
    t = info["decide"]["internal"]
    comp = {q: x for q, x in I.composite_series(the_internal(), "Credit card", "net-charge-off-rate")
            if x is not None}
    rel = t["rel_top"]
    want_verdicts = {"Agrees": IT.VERDICTS[2], "Disagrees": IT.VERDICTS[4], "Short": IT.VERDICTS[0],
                     "Refused": IT.VERDICTS[1]}
    for k, s in enumerate(t["outside"]):
        label = "Credit card - %s - %s" % (NET, s)
        lines, worse, refused = lines_of(vals, info, label)
        ref = reference(comp, source_values(ser, label), lines, worse, refused)
        col = "MNOPQ"[k]
        got = {key: iv(vals, col, rel + 1 + i) for i, key in
               enumerate(("overlap", "level", "change", "gap", "agree", "within", "verdict"))}
        assert got["verdict"] == ref["verdict"] == want_verdicts[s], s
        assert got["overlap"] == ref["overlap"], s
        for key in ("level", "change", "gap", "agree", "within"):
            if ref[key] is None:
                assert got[key] in ("", None), (s, key, got[key])
            else:
                assert got[key] == pytest.approx(ref[key], abs=1e-9), (s, key)
        # The same figures in the every-series statistics.
        st = t["stats"][label]
        assert v(vals, "INTERNAL!" + st["verdict"].replace("$", "")) == ref["verdict"]


def test_every_product_and_measure_names_the_best_source(filled):
    d, info, vals = filled
    t = info["decide"]["internal"]
    a, b_ = t["every_rows"]
    cols = t["every_cols"]
    from openpyxl.utils import get_column_letter as L
    rows = {(iv(vals, "B", r), iv(vals, "C", r)): r for r in range(a, b_ + 1)}
    r = rows[("Credit card", "Net charge-off rate")]
    assert iv(vals, L(cols["Best source"]), r) == "Agrees"
    assert iv(vals, L(cols["Verdict"]), r) == IT.VERDICTS[2]
    assert iv(vals, L(cols["Workbook measure"]), r) == NET
    assert iv(vals, L(cols["Quarters"]), r) == 45
    # A measure not mapped to the workbook has no outside series.
    r = rows[("Credit card", "30+ days past due")]
    assert iv(vals, L(cols["Workbook measure"]), r) == IT.NOT_MAPPED
    assert iv(vals, L(cols["Verdict"]), r) == IT.NO_SERIES


def test_the_segment_table_and_shares(filled):
    d, info, vals = filled
    t = info["decide"]["internal"]
    comp_r, tog_r = t["seg_rows"]
    last = {r.segment: r for r in internal_rows() if r.quarter == "2019Q4" and r.product == "Credit card"}
    total = sum(float(r.balance) for r in last.values())
    assert iv(vals, "B", comp_r) == IT.COMPOSITE
    assert iv(vals, "C", comp_r) == pytest.approx(total / 1000)
    segs = {iv(vals, "B", r): r for r in range(comp_r + 1, tog_r)}
    assert set(k for k in segs if k) == {"Bank card", "Private label"}
    r = segs["Bank card"]
    assert iv(vals, "C", r) == pytest.approx(float(last["Bank card"].balance) / 1000)
    assert iv(vals, "E", r) == pytest.approx(float(last["Bank card"].balance) / total)
    one = I.Internal("x", {}, {"Bank card": "x"}, (), tuple(internal_rows()))
    want = dict(I.series(one, "Credit card", "Bank card", "net-charge-off-rate"))["2019Q4"]
    assert iv(vals, "F", r) == pytest.approx(want, rel=1e-12)
    assert iv(vals, "E", tog_r) == pytest.approx(1.0)


def test_crr_grades_are_rated_and_compared(book):
    d, info, ser = book
    wb = load_workbook(d / "book.xlsx")
    wb[IT.TITLE]["B%d" % IT.PICK_ROW] = "Auto"
    wb.save(d / "auto.xlsx")
    vals = paste(d / "auto.xlsx", d / "auto_filled.xlsx", internal_rows())
    t = info["decide"]["internal"]
    sentence = iv(vals, "L", t["crr_row"])
    label = "Auto - %s - Agrees" % NET
    lines, worse, _ = lines_of(vals, info, label)
    auto = {q: x for q, x in I.composite_series(the_internal(), "Auto", "net-charge-off-rate") if x is not None}
    grades = {r.quarter: float(dict(r.extra)["crr_grade"]) for r in internal_rows() if r.product == "Auto"}
    lows = (1, 3, 5, 7, 9)
    both = [q for q in auto if q in grades]
    match = sum(sum(1 for lo in lows if lo <= grades[q]) == rating(auto[q], lines, worse) for q in both)
    assert sentence == ("Your CRR grades agree with the rating on the Agrees lines in %d%% of %d quarters."
                        % (round(100 * match / len(both)), len(both)))


def test_like_for_like_reads_the_call_report_series(filled, book):
    d, info, vals = filled
    t = info["decide"]["internal"]
    rows = {iv(vals, "B", r): r for r in t["like_rows"]}
    auto = rows["Auto"]
    q = "2019Q4"
    the = the_internal()
    win = [r for r in the.rows if r.product == "Auto" and I.qindex(r.quarter) > I.qindex(q) - 4]
    avg = sum(float(r.balance) for r in win) / 4
    assert iv(vals, "H", auto) == pytest.approx(avg / 1000)
    assert iv(vals, "L", auto) == pytest.approx(100 * sum(float(r.net_charge_offs) for r in win) / avg)
    pub = source_values(book[2], "Auto - %s - Agrees" % NET)[q]
    assert iv(vals, "M", auto) == pytest.approx(pub)
    assert iv(vals, "N", auto) == pytest.approx(iv(vals, "L", auto) - pub)
    # Card's composite is its segments summed, set against the same line.
    card = rows["Credit card"]
    win = [r for r in the.rows if r.product == "Credit card" and I.qindex(r.quarter) > I.qindex(q) - 4]
    assert iv(vals, "H", card) == pytest.approx(sum(float(r.balance) for r in win) / 4000)
    # No public figure for a line no product of the workbook holds.
    assert iv(vals, "F", rows["Mortgage and home equity loans"]) == (
        "No product in this workbook holds DRRERSF2 or LNRERSF2.")


def test_a_product_with_a_bad_row_has_no_composite(book):
    d, info, _ = book
    rows = internal_rows()
    bad = dict(BAD, product="Cards", segment="Bank card", quarter="2010Q1")       # a repeated quarter
    vals = paste(d / "book.xlsx", d / "bad.xlsx", rows, extra=[bad])
    t = info["decide"]["internal"]
    col = "MNOPQ"[t["products"].index("Credit card")]
    assert iv(vals, col, 12) == 2                         # the repeat and the row it repeats
    from openpyxl.utils import get_column_letter as L
    b = t["pblock"]["Credit card"]
    assert all(v(vals, "INTERNAL!%s%d" % (L(b["r_net"]), t["w0"] + i)) in ("", None) for i in range(60))


# ---------------------------------------------------------------- the book's internal section

def write_internal_csv(path):
    """internal_rows() as the import file internal.py reads, notes and all."""
    lines = ["# entity: Synthetic Bancorp (made up for the tests)",
             "# as-of: balances at quarter-end; charge-offs for the quarter",
             "# basis: synthetic loans, $ thousands", "# source: invented in tests/test_internal_tab.py",
             "# segment-All: the synthetic product total", "# segment-Bank card: synthetic bank cards",
             "# segment-Private label: synthetic private-label cards",
             ",".join(I.COLUMNS + ("crr_grade",))]
    for r in internal_rows():
        g = dict(r.extra).get("crr_grade")
        lines.append(",".join([r.quarter, r.product, r.segment] + [str(getattr(r, c)) for c in I.COLUMNS[3:]]
                              + ["" if g is None else str(g)]))
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return path


def write_book(folder, ser, internal):
    import json
    book = {"workbook": {k: (list(x) if isinstance(x, tuple) else x) for k, x in SETTINGS.items()},
            "parts": NAMES, "lines_from": {"part": NAMES[2], "start": None},
            "series": [dict(zip(NAMES, lbl.split(" - ")), file=str(p)) for lbl, p in ser],
            "not_in_book": [], "compare": None, "notes": [], "peers": None, "internal": internal,
            "loan_classes": None, "reporting_changes": None, "floor_off": [], "large_banks": None}
    p = folder / "book.json"
    p.write_text(json.dumps(book), encoding="utf-8")
    return p


def test_a_book_with_an_internal_file_adds_its_series_and_fills_the_slot(book, tmp_path):
    from threshold_engine import book as BK
    from threshold_engine.series import read_csv
    from threshold_engine.workbook import quarter
    _, _, ser = book
    f = write_internal_csv(tmp_path / "internal.csv")
    spec = {"file": str(f), "source": "Bank internal", "segments": True,
            "measures": {"net-charge-off-rate": NET}, "call_report_source": "Agrees"}
    out = tmp_path / "out" / "book.xlsx"
    out.parent.mkdir()
    info = BK.build(write_book(tmp_path, ser, spec), out, say=lambda *_: None)
    added = [lbl for lbl in info["products"] if "Bank internal" in lbl]
    assert sorted(added) == sorted(["Credit card - %s - Bank internal" % NET, "Auto - %s - Bank internal" % NET,
                                    "Credit card - %s - Bank internal: Bank card" % NET,
                                    "Credit card - %s - Bank internal: Private label" % NET])
    # The composite written is internal.py's; the series sit beside the workbook.
    folder = out.parent / "book internal series"
    got = {quarter(p.date): p.value
           for p in read_csv(folder / "internal_credit-card_composite_net-charge-off-rate.csv")}
    want = {q: x for q, x in I.composite_series(the_internal(), "Credit card", "net-charge-off-rate")
            if x is not None}
    assert got == pytest.approx(want)
    wb = load_workbook(out)
    ds = wb[IT.DATA_TITLE]
    assert ds["A5"].value == internal_rows()[0].quarter and ds["B5"].value == "Credit card"
    assert sum(1 for r in range(5, ds.max_row + 1) if ds.cell(r, 1).value) == len(internal_rows())
    # The internal source is not an outside source on the Internal tab, and the
    # past-due measures the book did not map stay on the tab only.
    t = info["decide"]["internal"]
    assert t["outside"] == ["Agrees", "Disagrees", "Short", "Refused"]
    ws = wb[IT.TITLE]
    assert [ws.cell(r, 6).value for r in range(5, 10)] == [NET] + [IT.NOT_MAPPED] * 4
    assert ws["J11"].value == "Agrees"
    assert "Credit card - %s - Bank internal" % NET in [c.value for c in wb["Lines by source"]["A"]]


@pytest.mark.parametrize("change, said", [
    ({"measures": {"dpd-30-plus": NET}}, "not an internal measure"),
    ({"measures": {"dpd-30-plus-quarter-end": "30+ incl. nonaccrual"}}, "not a measure of the book's series"),
    ({"source": "Agrees"}, "already a source"),
    ({"call_report_source": "Nobody"}, "call_report_source must be"),
    ({"segments": "yes"}, "segments must be true"),
])
def test_the_internal_section_is_refused_by_name(book, tmp_path, change, said):
    from threshold_engine import book as BK
    _, _, ser = book
    spec = {"file": str(write_internal_csv(tmp_path / "internal.csv")), "source": "Bank internal",
            "segments": False, "measures": {"net-charge-off-rate": NET}, "call_report_source": "Agrees"}
    spec.update(change)
    with pytest.raises(BK.BookError, match=said):
        BK.load(write_book(tmp_path, ser, spec))


def test_an_internal_file_that_fails_its_check_is_refused_with_its_problems(book, tmp_path):
    from threshold_engine import book as BK
    _, _, ser = book
    f = write_internal_csv(tmp_path / "internal.csv")
    f.write_text(f.read_text().replace("# entity: Synthetic Bancorp (made up for the tests)\n", ""))
    spec = {"file": str(f), "source": "Bank internal", "segments": False, "measures": {},
            "call_report_source": "Agrees"}
    with pytest.raises(BK.BookError, match="internal file: note 'entity' is not stated"):
        BK.load(write_book(tmp_path, ser, spec))


def test_a_book_must_say_whether_it_has_an_internal_file(book, tmp_path):
    import json
    from threshold_engine import book as BK
    _, _, ser = book
    p = write_book(tmp_path, ser, None)
    raw = json.loads(p.read_text())
    del raw["internal"]
    p.write_text(json.dumps(raw))
    with pytest.raises(BK.BookError, match="missing 'internal'"):
        BK.load(p)


def test_the_composite_takes_the_all_row_quarter_by_quarter():
    """Codex on #442: a product with All rows for part of its history and
    segment rows for the rest keeps every quarter; the All row is used where
    given and the segments are summed where not, under one segment name."""
    rows = []
    for i, q in enumerate(INT_QS):
        if i < 6:
            rows.append(I.Row(0, q, "Auto", "All", Decimal(1000 + i), Decimal(10), Decimal(8), Decimal(2), None,
                              Decimal(30), Decimal(20), Decimal(10), ()))
        else:
            for seg, bal in (("New", 600 + i), ("Used", 400)):
                rows.append(I.Row(0, q, "Auto", seg, Decimal(bal), Decimal(5), Decimal(4), Decimal(1), None,
                                  Decimal(15), Decimal(10), Decimal(5), ()))
    one = I.Internal("synthetic", {}, {"All": "x", "New": "x", "Used": "x"}, (), tuple(rows))
    comp = I.composite_rows(one, "Auto")
    assert [r.quarter for r in comp] == list(INT_QS)
    assert {r.segment for r in comp} == {I.COMPOSITE} and I.composite_segment(one, "Auto") == I.COMPOSITE
    assert comp[0].balance == Decimal(1000) and comp[0].net_charge_offs == Decimal(8)
    assert comp[6].balance == Decimal(1006) and comp[6].net_charge_offs == Decimal(8)
    # Every quarter with an All row: the composite is the All rows, under All.
    alls = tuple(r for r in rows if r.segment == "All")
    only = I.Internal("synthetic", {}, {"All": "x"}, (), alls)
    assert I.composite_segment(only, "Auto") == I.ALL
    assert [r.segment for r in I.composite_rows(only, "Auto")] == [I.ALL] * 6
