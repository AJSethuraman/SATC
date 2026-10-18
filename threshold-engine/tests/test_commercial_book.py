"""A commercial book: the words its reader sees, stated by the book.

The firm, 17 Oct 2026: *"names in this should be changed to be more
commercial ... product sounds consumer"*. A book states its ``kind``; a
commercial one says Portfolio wherever a consumer one says Product, shows the
portfolios under the firm's names (NDFI in full), and keeps the Call Report's
own name beside each on the Glossary. A consumer book keeps Product exactly
as before. Also here: the display problems the example commercial build found
on 17 Oct 2026 (the Peers headings, the share panel's title, the industry's
caption, the Internal tab).

The commercial book is built end to end on tests/data/peers/commercial_2007.json
(five filers, 2006Q3 to 2008Q2, public FDIC data; its README says how it was
pulled), with at least one bank in a composite so the peer series are
written from the one peer the fixture holds.
"""
import json
import re

import pytest
from openpyxl import load_workbook

from conftest import DATA, ROOT
from test_book import write_book
from threshold_engine import book as BK
from threshold_engine import kinds as KD
from threshold_engine.peers import books, fdic

PEERS = DATA / "peers"
WORD = re.compile(r"\bproducts?\b", re.I)
BARE_NDFI = re.compile(r"(?<!\()\bNDFI\b(?!\))")
#: The fixture's series that the peers build writes (measured: C&I's net
#: charge-offs hold 4 quarters, the nonaccrual rates 8, or 6 for the parts).
SERIES = [("C&I", "Net charge-offs", "c-i", "net"), ("C&I", "Nonaccrual", "c-i", "delinquency-nonaccrual_quarter-end"),
          ("Nonfarm nonresidential CRE", "Nonaccrual", "nonfarm-nonresidential", "delinquency-nonaccrual_quarter-end"),
          ("Owner-occupied CRE", "Nonaccrual", "owner-occupied-nonfarm-nonresidential",
           "delinquency-nonaccrual_quarter-end"),
          ("Investor CRE (non-owner-occupied)", "Nonaccrual", "other-nonfarm-nonresidential",
           "delinquency-nonaccrual_quarter-end")]
SOURCES = [("Industry", "industry"), ("Peer median", "pack-median"), ("Peer total", "pack-composite"),
           ("KeyBank", "keybank")]


# ---------------------------------------------------------------- the words

def test_the_commercial_words_and_names():
    t = KD.commercial_text
    assert t("Pick a product and a measure. Products: Owner-occupied nonfarm nonresidential, "
             "Other nonfarm nonresidential, Nonfarm nonresidential.") == (
        "Pick a portfolio and a measure. Portfolios: Owner-occupied CRE, Investor CRE (non-owner-occupied), "
        "Nonfarm nonresidential CRE.")
    # Said once, never twice; the Call Report's own words are left alone.
    assert t("Nonfarm nonresidential CRE") == "Nonfarm nonresidential CRE"
    assert t("Loans secured by owner-occupied nonfarm nonresidential properties") == \
        "Loans secured by owner-occupied nonfarm nonresidential properties"
    assert t("NDFI and margin-loan revision") == "Loans to nonbank financials (NDFI) and margin-loan revision"
    assert t(KD.NDFI) == KD.NDFI
    assert t("Product") == "Portfolio" and t("product's") == "portfolio's"


def test_a_formula_has_only_its_quoted_text_reworded():
    f = '=IF(SUMPRODUCT((A1:A3="x")*1)>0,"Refused: this product is not whole",PRODUCT(B1:B2))'
    assert KD._formula(f) == '=IF(SUMPRODUCT((A1:A3="x")*1)>0,"Refused: this portfolio is not whole",PRODUCT(B1:B2))'
    # A quoted string grown past Excel's 255 characters is split, joined by &.
    long_ = '="%s"' % ("NDFI " * 50).strip()
    out = KD._formula(long_)
    pieces = re.findall(r'"((?:[^"]|"")*)"', out)
    assert all(len(p) <= KD.LITERAL_MAX for p in pieces) and len(pieces) > 1
    assert "".join(pieces) == KD.commercial_text(("NDFI " * 50).strip())


# ---------------------------------------------------------------- the book states its kind

def test_a_book_without_a_kind_is_refused_by_name(tmp_path):
    p = write_book(tmp_path)
    raw = json.loads(p.read_text())
    del raw["kind"]
    p.write_text(json.dumps(raw))
    with pytest.raises(BK.BookError, match="missing 'kind'"):
        BK.load(p)


@pytest.mark.parametrize("over, said", [
    (dict(kind="retail"), "kind must be one of consumer, commercial"),
    (dict(kind=None), "kind must be one of consumer, commercial"),
    (dict(kind="consumer", parts=["Portfolio", "Measure", "Source"]),
     "a consumer book calls its first part 'Product', not 'Portfolio'"),
    (dict(kind="commercial"), "a commercial book calls its first part 'Portfolio', not 'Product'"),
])
def test_the_kind_is_checked_and_decides_the_first_part(tmp_path, over, said):
    with pytest.raises(BK.BookError, match=re.escape(said)):
        BK.load(write_book(tmp_path, **over))


def commercial_rows(names):
    return [{"Portfolio": n, "Measure": "Net", "Source": "Industry", "file": "files/cards_nco_ttm.csv"} for n in names]


def test_a_commercial_book_names_its_portfolios_as_the_firm_did(tmp_path):
    over = dict(kind="commercial", parts=["Portfolio", "Measure", "Source"], compare=None, not_in_book=[])
    # The peers build's name is the Call Report's, not the one shown.
    with pytest.raises(BK.BookError, match="not 'Owner-occupied nonfarm nonresidential'"):
        BK.load(write_book(tmp_path, series=commercial_rows(["Owner-occupied nonfarm nonresidential"]), **over))
    bk = BK.load(write_book(tmp_path, series=commercial_rows(["Owner-occupied CRE", "C&I"]), **over))
    assert bk["kind"] == "commercial" and bk["decide"]["kind"] == "commercial"


def test_a_commercial_book_has_no_internal_or_loan_classes_section(tmp_path):
    over = dict(kind="commercial", parts=["Portfolio", "Measure", "Source"], compare=None, not_in_book=[],
                series=commercial_rows(["C&I"]))
    with pytest.raises(BK.BookError, match="internal: a commercial book has no Internal tab"):
        BK.load(write_book(tmp_path, internal={k: "x" for k in BK.INTERNAL_KEYS}, **over))
    with pytest.raises(BK.BookError, match="loan_classes: the classes file records each peer's consumer book"):
        BK.load(write_book(tmp_path, loan_classes={"file": "x", "cache": "y"}, **over))


def test_the_books_in_the_repository_state_their_kinds():
    assert json.loads((ROOT / "books" / "consumer-public.json").read_text())["kind"] == "consumer"
    bk = BK.load(ROOT / "books" / "commercial-public.json")
    assert bk["kind"] == "commercial"
    assert {v[0] for v in bk["decide"]["parts"].values()} == set(KD.COMMERCIAL_NAMES.values())
    assert bk["decide"]["compare"] == ("KeyBank", "Peer median (15 banks)")
    assert bk["peers"]["first"] == "1991Q1" and bk["peers"]["min_book"] == 0
    # No extra acquisition: the consumer book's (Fifth Third, 6672=2022Q2) is a
    # consumer deal, and no commercial one is recorded. The acquisitions the
    # FDIC record cannot see come from the stated list (question 14).
    assert bk["peers"]["extra_acquisitions"] == []
    assert bk["peers"]["stated_acquisitions"] is not None
    assert bk["internal"] is None and bk["loan_classes"] is None
    # Question 3: nonaccrual leads every CRE portfolio.
    first = {}
    for p, m, _ in bk["decide"]["parts"].values():
        first.setdefault(p, m)
    assert {p: m for p, m in first.items() if p != "C&I"} == {p: "Nonaccrual" for p in first if p != "C&I"}
    # Question 4: each portfolio and measure has the median and the total beside it.
    srcs = {s for _, _, s in bk["decide"]["parts"].values()}
    assert srcs == {"Industry", "Peer median (15 banks)", "Peer total (15 banks)", "KeyBank"}


# ---------------------------------------------------------------- built end to end

def _book(base, data, kind):
    cache = base / "cache"
    cache.mkdir()
    for q, rows in data.items():
        (cache / ("q%s.json" % q)).write_text(json.dumps(rows), encoding="utf-8")
    (base / "mergers.json").write_text("{}", encoding="utf-8")
    (base / "absorptions.json").write_text("[]", encoding="utf-8")
    word = KD.KINDS[kind]["word"]
    shown = (lambda p: p) if kind == "commercial" else (
        lambda p: next(k for k, v in KD.COMMERCIAL_NAMES.items() if v == p))
    series = [{word: shown(p), "Measure": m, "Source": s, "file": "series/%s_%s_%s.csv" % (slug, f, ms)}
              for p, m, slug, ms in SERIES for s, f in SOURCES]
    book = {"kind": kind, "workbook": {"name": "Commercial", "unit": "%", "direction": "higher_is_worse",
                                        "frequency": "quarterly", "smoothing": 1, "floor_at_zero": True,
                                        "moderate_halfwidth": 0.5, "low_step": 1.0, "high_z": 2.5,
                                        "on_the_line": "worse", "half_lives": [], "percentiles": [50, 75, 90, 95],
                                        "horizon": 4, "min_history": 41},
            "parts": [word, "Measure", "Source"], "lines_from": {"part": "Source", "start": None},
            "series": series, "not_in_book": [], "compare": ["KeyBank", "Peer median"],
            "notes": ["A test book on the commercial fixture."],
            "peers": {"cache": str(cache), "roster": str(ROOT / "config" / "peers.json"), "series_dir": "series",
                      "first": "2006Q3", "last": "2008Q2", "products": list(books.COMMERCIAL),
                      "mergers": str(base / "mergers.json"), "absorptions": str(base / "absorptions.json"),
                      "extra_acquisitions": [], "industry_combined": [], "jump": 0.25, "acquisition_share": 0.05,
                      "min_book": 0, "exempt_self": True, "pro_forma": True, "pack": "all", "min_banks": 1,
                      "rank_by": "net", "last_n": 8, "min_obs": 6, "base_quarter": "2007Q4",
                      "index_min_book": 1000000, "lagged_years": [2007, 2007], "lag_years": 1,
                      "normalisation_years": [2007, 2007], "mix_base_year": 2007, "delinquency": True,
                      "stated_acquisitions": None},
            "internal": None, "loan_classes": None,
            "reporting_changes": str(ROOT / "config" / "reporting-changes.json")}
    p = base / "book.json"
    p.write_text(json.dumps(book), encoding="utf-8")
    return p


@pytest.fixture(scope="module")
def data():
    return json.loads((PEERS / "commercial_2007.json").read_text())


@pytest.fixture(scope="module")
def built(data, tmp_path_factory):
    out = {}
    mp = pytest.MonkeyPatch()
    mp.setattr(fdic, "urllib_fetch", lambda url: (_ for _ in ()).throw(AssertionError(url)))
    try:
        for kind in ("commercial", "consumer"):
            base = tmp_path_factory.mktemp(kind)
            said = []
            BK.build(_book(base, data, kind), base / "book.xlsx", say=said.append)
            out[kind] = (base / "book.xlsx", said)
    finally:
        mp.undo()
    return out


def texts(wb):
    """Every piece of text a reader can see: cells (a formula's quoted text
    only), dropdown prompts and titles, by tab."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str):
                    for t in (re.findall(r'"((?:[^"]|"")*)"', v) if v.startswith("=") else [v]):
                        yield ws.title, c.coordinate, t
        for dv in ws.data_validations.dataValidation:
            for t in (dv.prompt, dv.promptTitle, dv.error, dv.errorTitle):
                if isinstance(t, str):
                    yield ws.title, "validation %s" % dv.sqref, t


def test_a_commercial_book_says_portfolio_on_every_tab_and_never_product(built):
    path, said = built["commercial"]
    wb = load_workbook(path)
    hits = [(t, c, x[:80]) for t, c, x in texts(wb) if WORD.search(x)]
    assert hits == []
    # The part's own headings and the pickers.
    assert wb["Decide"]["B8"].value == "Portfolio"
    assert wb["Summary"]["B5"].value == "Portfolio"
    assert wb["Peers"]["B4"].value == "Portfolio"
    assert any(r[0].value == "Portfolio" for r in wb["Glossary"].iter_rows(min_col=2, max_col=2))
    assert any(s.startswith("commercial names: ") for s in said)


def test_a_consumer_book_keeps_product_exactly(built):
    path, _ = built["consumer"]
    wb = load_workbook(path)
    # The one "portfolio" is a source's own words, quoted on Reporting changes.
    hits = [x for t, c, x in texts(wb) if re.search(r"\bportfolios?\b", x, re.I)]
    assert len(hits) == 1 and hits[0].startswith(
        "FDIC Quarterly Banking Profile, First Quarter 2020: C&I loan portfolio increased by $339.4 billion")
    assert wb["Decide"]["B8"].value == "Product"
    assert wb["Summary"]["B5"].value == "Product"
    assert wb["Peers"]["B4"].value == "Product"
    gl = {r[1].value: r[2].value for r in wb["Glossary"].iter_rows()}
    assert gl["Product"] == "The first of the three parts that name a series, such as Credit card."
    assert "Internal" in wb.sheetnames and "Internal data" in wb.sheetnames
    # The Peers tab lists the peers build's own names, as before.
    wk = wb["Peers workings"]
    assert [wk.cell(r, 5).value for r in range(2, 8)] == list(books.COMMERCIAL)


def test_a_commercial_book_shows_the_firms_names_and_the_call_reports_beside_them(built, data):
    path, _ = built["commercial"]
    wb = load_workbook(path)
    wk = wb["Peers workings"]
    assert [wk.cell(r, 5).value for r in range(2, 8)] == [KD.COMMERCIAL_NAMES[p] for p in books.COMMERCIAL]
    assert {wb["Peers data"].cell(r, 1).value for r in range(2, wb["Peers data"].max_row + 1)} == \
        set(KD.COMMERCIAL_NAMES.values())
    # NDFI is never bare.
    assert [(t, c) for t, c, x in texts(wb) if BARE_NDFI.search(x)] == []
    assert any(KD.NDFI in x for t, _, x in texts(wb) if t == "Reporting changes")
    # The Glossary's last section: each name as shown beside the Call Report's.
    gl = wb["Glossary"]
    rows = [tuple(c.value for c in r[1:4]) for r in gl.iter_rows()]
    at = rows.index(("Shown as", "The Call Report's name", "Where it is filed"))
    assert rows[at + 1:at + 1 + len(KD.CALL_REPORT)] == list(KD.CALL_REPORT)
    assert "Internal" not in wb.sheetnames and "Internal data" not in wb.sheetnames
    gterms = {r[0] for r in rows}
    assert not (gterms & {"Internal data", "Competitive check", "CRR grade"})
    assert gl.cell(5, 3).value == "The first of the three parts that name a series, such as C&I."


def test_peers_headings_show_the_name_as_built_never_lower_cased(built):
    for kind in ("commercial", "consumer"):
        pe = load_workbook(built[kind][0])["Peers"]
        fs = [c.value for r in pe.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("=")]
        heads = [f for f in fs if re.match(r'="(Charge-offs|Delinquency|Book size|Lagged rate|Book index): ', f)
                 or f.startswith('="Against the industry')]
        assert len(heads) == 6
        assert not [f for f in heads if "LOWER('Peers'!$B$5)" in f or "LOWER(Peers!$B$5)" in f]


def test_the_share_panel_names_the_books_own_kind(built):
    for kind, want in (("commercial", "KeyBank's commercial book: share by portfolio"),
                       ("consumer", "KeyBank's consumer book: share by product")):
        pe = load_workbook(built[kind][0])["Peers"]
        assert any(isinstance(c.value, str) and want in c.value for r in pe.iter_rows() for c in r), kind
        assert not any(isinstance(c.value, str) and "consumer book" in c.value for r in pe.iter_rows()
                       for c in r) or kind == "consumer"


def test_the_glossary_names_the_share_panel_as_the_book_kind_says(built):
    """The panel says "KeyBank's commercial book: share by portfolio"; its
    Glossary term said "Consumer book by portfolio", "KeyBank's consumer book".
    A consumer book keeps its term exactly."""
    for kind, term, meaning in (
            ("commercial", "Commercial book by portfolio",
             "The share of KeyBank's commercial book in each portfolio built, the portfolio picked drawn first "
             "in red."),
            ("consumer", "Consumer book by product",
             "The share of KeyBank's consumer book in each product built, the product picked drawn first in red.")):
        gl = load_workbook(built[kind][0])["Glossary"]
        rows = {r[1].value: r[2].value for r in gl.iter_rows() if isinstance(r[1].value, str)}
        assert rows.get(term) == meaning, (kind, {t: m for t, m in rows.items() if "book by" in t})
        if kind == "commercial":
            assert not [t for t, m in rows.items() if t.startswith("Consumer book")
                        or "KeyBank's consumer book" in (m or "")]

def test_the_industry_caption_states_how_many_filers_it_holds(built, data):
    last = sorted(data)[-1]
    n = len(data[last])
    want = []
    for p in books.COMMERCIAL:
        k = books.industry_filers(data, [p], last)
        want.append("all %d filers" % n if k == n else "%d of %d filers" % (k, n))
    # The fixture's Bank of China files no charge-off field, so no commercial
    # industry series holds every filer.
    assert all(" of %d filers" % n in w for w in want)
    for kind in ("commercial", "consumer"):
        wb = load_workbook(built[kind][0])
        wk = wb["Peers workings"]
        assert wk.cell(1, 11).value == "Industry filers, last quarter"
        assert [wk.cell(2 + i, 11).value for i in range(len(want))] == want
        f = [c.value for r in wb["Peers"].iter_rows() for c in r
             if isinstance(c.value, str) and "Industry filers" not in c.value and "$K$2:$K$7" in c.value]
        assert len(f) == 1 and "every filer" not in f[0] and '"not supplied"' in f[0]


def test_summary_is_wide_enough_for_the_longest_portfolio_name(built):
    ws = load_workbook(built["commercial"][0])["Summary"]
    longest = max((n for n, *_ in SERIES), key=len)
    cols = [c.column_letter for c in ws[5] if c.value == "Portfolio"]
    assert len(cols) == 2                        # the two panels
    for col in cols:
        assert ws.column_dimensions[col].width >= len(longest)
    assert longest in [ws["%s%d" % (col, r)].value for col in cols for r in range(6, 20)]
    wb = load_workbook(built["commercial"][0])
    for tab in ("Thresholds", "Lines by source"):
        assert wb[tab].column_dimensions["B"].width >= len(longest), tab
    # A consumer book's names keep the widths they always had.
    from threshold_engine.summary_tab import W_PRODUCT, product_width
    consumer = ["Credit card", "Mortgage", "Home equity", "Other consumer", "Auto"]
    assert product_width(consumer) == W_PRODUCT and product_width(consumer, 16.43) == 16.43
