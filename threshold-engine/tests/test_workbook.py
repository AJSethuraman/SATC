"""The workbook: Excel's live lines must equal Python's reference scale
(threshold_engine.scale), as built and after the bank changes a setting.
Recalculated with the `formulas` engine, so no copy of Excel is needed; a skip
is reported, never counted as a pass."""
import hashlib

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from conftest import DATA, load
from threshold_engine.glossary import TERMS
from threshold_engine.scale import scale, score
from threshold_engine.workbook import LO_ROWS, SET_TOP, assess_layout, build, quarter

# Where things sit on Assess since Claude Design's hand-off (13 Oct 2026):
# scores in N:U from row 10, Your settings N:R, Periods left out N:Q (Status
# T), and the chart's hidden helpers from Z (Z quarter, AC every quarter,
# AD:AH the bands, AI top/floor/height, AJ the base, AK the strip, AM its base).
LAY = assess_layout(4)
LO_TOP = LAY["lo_top"]

formulas = pytest.importorskip("formulas", reason="the formulas engine recalculates the workbook")

SERIES = [("Credit card", DATA / "cards_nco_ttm.csv"), ("Mortgage", DATA / "mortgage_nco_ttm.csv"),
          ("Home equity", DATA / "home_equity_nco_ttm.csv"),
          ("Other consumer", DATA / "other_consumer_as_filed_nco_ttm.csv")]
CARD_CUTS = (("2007Q4", "2012Q4"), ("2021Q3", "2023Q1"))
SET = dict(halfwidth=0.5, low_step=1.0, high_z=2.5)


def make(path, **over):
    kw = dict(name="NCO", unit="%", direction="higher_is_worse", frequency="quarterly", smoothing=1,
              floor_at_zero=True, moderate_halfwidth=SET["halfwidth"], low_step=SET["low_step"],
              high_z=SET["high_z"], on_the_line="worse", half_lives=(10,),
              percentiles=(50, 75, 90, 95), horizon=4, min_history=41)
    kw.update(over)
    return build(path, SERIES, **kw)


def recalc(path):
    """Every cell's value after a full recalculation, keyed 'SHEET!A1'."""
    model = formulas.ExcelModel().loads(str(path)).finish()
    sol = model.calculate()
    out = {}
    for key, val in sol.items():
        ref = key.split("]")[-1].upper()
        v = getattr(val, "value", val)
        try:
            v = v.item(0) if hasattr(v, "item") else v
        except Exception:                          # noqa: BLE001
            pass
        out[ref.replace("'", "")] = v
    return out


def row(vals, r):
    """Median, typical yearly move, worst, the four lines, score."""
    return [vals.get("THRESHOLDS!%s%d" % (c, r)) for c in "BCDEFGHK"]


def reference(file, cuts=(), **over):
    pts = load(file)
    kept = [None if any(a <= quarter(p.date) <= b for a, b in cuts) else p.value for p in pts]
    kw = dict(SET, **over)
    ref = scale(kept, "higher_is_worse", kw["halfwidth"], kw["low_step"], kw["high_z"], True)
    return pts, kept, ref


def top_of(vals, first):
    return vals["ASSESS!AI%d" % first]


# The tests below still say what they change in the Settings tab's old cells
# (row 6 Credit card ... row 9; B-D the first leave-out, E-G the second, H-K
# the three distances and the on-the-line rule), which reads well. Since 13 Oct
# 2026 the inputs live on Assess, and this puts each one where it now is: a
# leave-out becomes a row of Periods left out naming the product, a setting a
# cell of Your settings.
def assess_cells(cells):
    out = {}
    for ref, v in cells.items():
        col, row = ref[0], int(ref[1:])
        k = row - 6
        if col in "BCDEFG":
            slot = 2 * k + (0 if col in "BCD" else 1)
            r = LO_TOP + slot
            out["N%d" % r] = SERIES[k][0]
            out["%s%d" % ("OPQ"["BCDEFG".index(col) % 3], r)] = v
        else:
            out["%s%d" % ("OPQR"["HIJK".index(col)], SET_TOP + k)] = v
    return out


def with_settings(p, tmp_path, name, cells):
    wb = load_workbook(p)
    for ref, v in assess_cells(cells).items():
        wb["Assess"][ref] = v
    q = tmp_path / name
    wb.save(q)
    return recalc(q)


def with_assess(p, tmp_path, name, cells):
    """Type straight into Assess, by its own cell names."""
    wb = load_workbook(p)
    for ref, v in cells.items():
        wb["Assess"][ref] = v
    q = tmp_path / name
    wb.save(q)
    return recalc(q)


def dial_top(p):
    """The first dial's band row on Assess."""
    ws = load_workbook(p)["Assess"]
    return next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("Dial:"))


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    p = tmp_path_factory.mktemp("wb") / "t.xlsx"
    info = make(p)
    return p, info, recalc(p)


def test_as_built_every_product_equals_the_reference_scale(built):
    _, _, vals = built
    for r, (label, path) in enumerate(SERIES, start=5):
        pts, _, ref = reference(path.name)
        got = row(vals, r)
        assert got[0] == pytest.approx(ref["median"])
        assert got[1] == pytest.approx(ref["typical_yearly_move"])
        assert got[2] == pytest.approx(ref["worst"])
        assert got[3:7] == pytest.approx(ref["bounds"])
        assert got[7] == score(pts[-1].value, ref["bounds"], "higher_is_worse", "worse", True)


def test_cards_with_both_periods_left_out_equal_the_reference(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "cuts.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "Financial crisis",
        "E6": "2021Q3", "F6": "2023Q1", "G6": "Pandemic forbearance"})
    _, _, ref = reference("cards_nco_ttm.csv", CARD_CUTS)
    got = row(vals, 5)
    assert got[1] == pytest.approx(ref["typical_yearly_move"])
    assert got[3:7] == pytest.approx(ref["bounds"])
    assert [round(x, 3) for x in got[3:7]] == [3.171, 3.464, 4.050, 5.981]
    assert got[7] == 3                                 # today's 3.98% is Moderate


def test_each_multiplier_on_settings_is_live(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "mult.xlsx", {"H6": 0.25, "I6": 1.5, "J6": 3.0})
    _, _, ref = reference("cards_nco_ttm.csv", halfwidth=0.25, low_step=1.5, high_z=3.0)
    assert row(vals, 5)[3:7] == pytest.approx(ref["bounds"])


def test_the_on_the_line_rule_is_live(built, tmp_path):
    # Half-width 0 puts Moderate and Moderate-High both on the median, 3.978,
    # which is also cards' latest value: "worse" scores it 4, "better" 2.
    p, _, _ = built
    worse = with_settings(p, tmp_path, "w.xlsx", {"H6": 0})
    better = with_settings(p, tmp_path, "b.xlsx", {"H6": 0, "K6": "better"})
    assert row(worse, 5)[0] == pytest.approx(3.978)
    assert (row(worse, 5)[7], row(better, 5)[7]) == (4, 2)


def test_a_median_at_or_below_zero_is_refused(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "z.xlsx", {"B7": "1991Q4", "C7": "2017Q4", "D7": "test"})
    assert vals["THRESHOLDS!B6"] <= 0
    assert vals["THRESHOLDS!M6"] == "Refused: the median is at or below zero"
    assert vals["THRESHOLDS!K6"] == ""


def test_moderate_low_must_begin_below_moderate(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "order.xlsx", {"H6": 1.0, "I6": 0.5})
    assert vals["THRESHOLDS!M5"] == "Refused: Moderate-Low must begin further from the median than Moderate"


def test_every_series_is_on_one_quarterly_calendar(built):
    p, info, _ = built
    da = load_workbook(p)["Data"]
    qs = [c.value for c in da["A"][1:]]
    assert qs == sorted(qs) and len(qs) == info["quarters"] and qs[0] == "1985Q4"
    assert quarter("2010-04-01") == quarter("2010-06-30") == "2010Q2"


def test_the_run_tab_fingerprints_every_file(built):
    p, _, _ = built
    run = load_workbook(p)["Run"]
    rows = {r[1].value: r for r in run.iter_rows(min_row=4)
            if r[1].value and len(str(r[6].value or "")) == 64}         # the hash goes last, in G
    for label, path in SERIES:
        assert rows[label][6].value == hashlib.sha256(path.read_bytes()).hexdigest()


def test_mixed_duplicate_or_impossible_inputs_are_refused(tmp_path):
    with pytest.raises(ValueError):
        build(tmp_path / "x.xlsx", SERIES[:1] * 2, "NCO", "%", "higher_is_worse", "quarterly", 1, True,
              0.5, 1.0, 2.5, "worse", (), (50, 75, 90, 95), 4, 41)
    nan, inf = float("nan"), float("inf")
    for bad in (dict(on_the_line="sometimes"), dict(low_step=0.5), dict(high_z=0),
                dict(high_z=nan), dict(high_z=inf), dict(moderate_halfwidth=nan), dict(low_step=inf)):
        with pytest.raises(ValueError):
            make(tmp_path / "y.xlsx", **bad)


def test_two_leave_out_windows_carry_their_reasons_to_the_run_tab(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "two.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "Financial crisis",
                                                   "E6": "2021Q3", "F6": "2023Q1", "G6": "Pandemic forbearance"})
    run = {v for k, v in vals.items() if k.startswith("RUN!")}
    assert {"2007Q4", "2012Q4", "Financial crisis", "2021Q3", "2023Q1", "Pandemic forbearance",
            "Left out: 21 quarters", "Left out: 7 quarters"} <= run


def test_assess_says_what_the_evidence_found_for_the_product_picked(built):
    p, _, vals = built
    ws = load_workbook(p)["Assess"]
    top = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith(
        '="What the evidence says about "'))
    assert [ws.cell(top + i, 2).value for i in (1, 2, 3)] == ["Suggested", "Departures", "Why"]
    said, deps, why = (vals["ASSESS!C%d" % (top + i)] for i in (1, 2, 3))
    assert vals["ASSESS!B%d" % top] == "What the evidence says about Credit card"
    assert said == "Leave out 2007Q4 to 2012Q4"
    assert "2021Q3 to 2023Q1 better than trend" in deps
    assert "2008" not in (why or "")


def test_assess_counts_the_quarters_in_each_score(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "assess.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "Financial crisis",
        "E6": "2021Q3", "F6": "2023Q1", "G6": "Pandemic forbearance"})
    pts, kept, ref = reference("cards_nco_ttm.csv", CARD_CUTS)
    sc = lambda v: score(v, ref["bounds"], "higher_is_worse", "worse", True)      # noqa: E731
    want_kept = [sum(1 for v in kept if v is not None and sc(v) == k) for k in range(1, 6)]
    want_all = [sum(1 for x in pts if sc(x.value) == k) for k in range(1, 6)]
    assert [vals["ASSESS!Q%d" % r] for r in range(10, 15)] == want_kept
    assert [vals["ASSESS!S%d" % r] for r in range(10, 15)] == want_all
    assert sum(want_kept) == 135


def test_assess_follows_the_on_the_line_rule(built, tmp_path):
    p, _, _ = built
    worse = with_settings(p, tmp_path, "aw.xlsx", {"H6": 0})
    better = with_settings(p, tmp_path, "ab.xlsx", {"H6": 0, "K6": "better"})
    on_line = sum(1 for x in load("cards_nco_ttm.csv") if x.value == 3.978)
    assert on_line >= 1
    # With Moderate empty, a quarter on the median is Moderate-High under
    # "worse" and Moderate-Low under "better".
    assert worse["ASSESS!S13"] - better["ASSESS!S13"] == on_line
    assert better["ASSESS!S11"] - worse["ASSESS!S11"] == on_line


def test_the_strip_marks_exactly_the_quarters_left_out_and_the_caption_says_why(built, tmp_path):
    p, info, _ = built
    vals = with_settings(p, tmp_path, "strip.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "Financial crisis",
                                                     "E6": "2021Q3", "F6": "2023Q1", "G6": "Pandemic forbearance"})
    qcol, scol = get_column_letter(26), get_column_letter(37)
    qs = [vals["ASSESS!%s%d" % (qcol, r)] for r in range(2, info["quarters"] + 2)]
    marked = [q for r, q in zip(range(2, info["quarters"] + 2), qs) if vals["ASSESS!%s%d" % (scol, r)] != 0]
    assert marked == [q for q in qs if any(a <= q <= b for a, b in CARD_CUTS)] and len(marked) == 28
    assert vals["ASSESS!T%d" % LO_TOP] == "Left out: 21 quarters"
    assert vals["ASSESS!T%d" % (LO_TOP + 1)] == "Left out: 7 quarters"


def test_with_nothing_left_out_every_status_is_blank(built):
    _, _, vals = built
    assert all(vals["ASSESS!T%d" % r] == "" for r in range(LO_TOP, LO_TOP + LO_ROWS))


# Every heading a reader meets on Thresholds and Assess's inputs, and the glossary
# entry that explains it. Headings that need no definition are listed as such,
# so a new heading fails here until someone decides which it is.
PLAIN = {"Series", "Reason", "From", "To"}
DEFINED_BY = {
    "Median kept": "Median kept", "Typical yearly move": "Typical yearly move", "Worst kept": "Worst kept",
    "2 Moderate-Low from": "2 Moderate-Low from, 3 Moderate from, 4 Moderate-High from, 5 High from",
    "3 Moderate from": "2 Moderate-Low from, 3 Moderate from, 4 Moderate-High from, 5 High from",
    "4 Moderate-High from": "2 Moderate-Low from, 3 Moderate from, 4 Moderate-High from, 5 High from",
    "5 High from": "2 Moderate-Low from, 3 Moderate from, 4 Moderate-High from, 5 High from",
    "Latest quarter": "Latest", "Latest value": "Latest", "Rating": "Rating", "Rating name": "Rating name",
    "Check": "Check",
    "Width of Moderate (yearly moves)": "Width of Moderate (yearly moves)",
    "Moderate-Low begins (yearly moves)": "Moderate-Low begins (yearly moves)",
    "How rare High is (z)": "How rare High is (z)", "On a line, rate it": "On a line, rate it",
    "Status": "Status",
}


def test_every_heading_is_defined_in_the_glossary(built):
    p, _, _ = built
    wb = load_workbook(p)
    terms = {t for t, _, _ in TERMS}
    headings = ([c.value for c in wb["Thresholds"][4] if c.value]
                + [c.value for c in wb["Assess"][LO_TOP - 1][13:21] if c.value]     # N-U: the inputs' headers
                + [c.value for c in wb["Assess"][SET_TOP - 1][13:21] if c.value])
    for h in headings:
        assert h in PLAIN or DEFINED_BY.get(h) in terms, "no glossary entry for %r" % h


def test_the_glossary_tab_carries_every_term(built):
    p, _, _ = built
    gl = load_workbook(p)["Glossary"]
    on_tab = [r[1].value for r in gl.iter_rows(min_row=5) if r[1].value]          # A is the margin
    assert on_tab == [t for t, _, _ in TERMS]
    assert all(m and c for _, m, c in TERMS)


def test_no_formula_refers_to_a_row_that_does_not_exist(built):
    # The four-quarter change starts on the fifth quarter: a formula on an
    # earlier row would point above row 1, which Excel reports as a damaged
    # file even where the formulas engine shrugs.
    import re
    p, _, _ = built
    wb = load_workbook(p)
    bad = []
    for ws in wb.worksheets:
        for row_ in ws.iter_rows():
            for c in row_:
                if isinstance(c.value, str) and c.value.startswith("="):
                    for m in re.finditer(r"(?<![A-Za-z])\$?[A-Z]{1,3}\$?(-?\d+)(?![\d(])", c.value):
                        if int(m.group(1)) < 1:
                            bad.append((ws.title, c.coordinate, c.value))
    assert bad == []



def test_high_beyond_everything_kept_is_refused(built, tmp_path):
    # Cards with the crisis left out: z 3.5 puts High above the worst quarter kept.
    p, _, _ = built
    vals = with_settings(p, tmp_path, "z35.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "Financial crisis",
        "E6": "2021Q3", "F6": "2023Q1", "G6": "Pandemic forbearance",
                                                   "J6": 3.5})
    assert vals["THRESHOLDS!M5"] == "Refused: High would begin beyond the worst quarter kept"
    _, _, ref = reference("cards_nco_ttm.csv", CARD_CUTS, high_z=3.5)
    assert ref["bounds"] is None and "beyond the worst" in ref["refused"]


def test_every_dial_row_equals_the_reference_scale(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "dial.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "Financial crisis",
        "E6": "2021Q3", "F6": "2023Q1", "G6": "Pandemic forbearance"})
    pts, kept, _ = reference("cards_nco_ttm.csv", CARD_CUTS)
    rows = {}
    r = dial_top(p) + 2
    for name, values in (("high_z", (2.0, 2.5, 3.0, 3.5)), ("halfwidth", (0.25, 0.5, 0.75, 1.0)),
                         ("low_step", (0.75, 1.0, 1.5, 2.0))):
        for v in list(values) + [None]:
            rows[(name, v)] = r
            r += 1
        r += 3
    for (name, v), r in rows.items():
        kw = dict(SET)
        if v is not None:
            kw[name] = v
        ref = scale(kept, "higher_is_worse", kw["halfwidth"], kw["low_step"], kw["high_z"], True)
        note = vals["ASSESS!L%d" % r]
        if ref["bounds"] is None:
            assert str(note).startswith("refused"), (name, v, note)
            continue
        assert [vals["ASSESS!%s%d" % (c, r)] for c in "CDEF"] == pytest.approx(ref["bounds"]), (name, v)
        want = [sum(1 for x in kept if x is not None and score(x, ref["bounds"], "higher_is_worse", "worse", True) == k)
                for k in range(1, 6)]
        assert [vals["ASSESS!%s%d" % (c, r)] for c in "GHIJK"] == want, (name, v)
        assert note == ("◀ yours" if v is None else ""), (name, v, note)


@pytest.mark.parametrize("rule", ["worse", "better"])
def test_the_dials_current_row_counts_exactly_as_the_main_table(built, tmp_path, rule):
    # Half-width 0 puts two lines on the median, 3.978, a value cards actually
    # took: counts then depend on the on-the-line rule, in the dial as above.
    p, _, _ = built
    vals = with_settings(p, tmp_path, "dl-%s.xlsx" % rule, {"H6": 0, "K6": rule})
    main = [vals["ASSESS!Q%d" % r] for r in range(10, 15)]
    top = dial_top(p)
    for now_row in (top + 6, top + 14, top + 22):
        assert [vals["ASSESS!%s%d" % (c, now_row)] for c in "GHIJK"] == main
        assert vals["ASSESS!L%d" % now_row] == "◀ yours"



def test_too_little_history_kept_is_refused_in_the_workbook(built, tmp_path):
    # Leave out everything but the last four quarters of cards: no four-quarter
    # change has both ends kept, so there is no typical yearly move.
    p, _, _ = built
    vals = with_settings(p, tmp_path, "short.xlsx", {"B6": "1985Q4", "C6": "2025Q2", "D6": "test"})
    assert vals["THRESHOLDS!M5"] == "Refused: too little history kept for a typical yearly move"
    assert vals["THRESHOLDS!K5"] == ""
    # A period was left out by hand, so putting one back is the fix.
    assert vals["THRESHOLDS!U5"] == "Put back a period, so that more than a year of quarters is kept."


def test_a_series_too_short_for_a_yearly_move_shows_blanks_and_names_the_real_cause(tmp_path):
    """The example commercial build (17 Oct 2026): KeyBank's C&I charge-off
    series held too few quarters for a four-quarter change, so its typical
    yearly move and three of its lines showed #NUM! on Lines by source, and
    the fix said to put back a period when none had been left out. Now
    those cells are blank and the fix says why."""
    lines = (DATA / "cards_nco_ttm.csv").read_text(encoding="utf-8").splitlines()
    short = tmp_path / "short.csv"
    rows = [ln for ln in lines if not ln.startswith("#")]
    short.write_text("\n".join([ln for ln in lines if ln.startswith("#")] + rows[:1] + rows[-4:]) + "\n",
                     encoding="utf-8")
    p = tmp_path / "w.xlsx"
    kw = dict(name="NCO", unit="%", direction="higher_is_worse", frequency="quarterly", smoothing=1,
              floor_at_zero=True, moderate_halfwidth=SET["halfwidth"], low_step=SET["low_step"],
              high_z=SET["high_z"], on_the_line="worse", half_lives=(), percentiles=(50, 75, 90, 95), horizon=4,
              min_history=41)
    build(p, SERIES[:1] + [("Short", short)], **kw)
    vals = recalc(p)
    errors = {k: v for k, v in vals.items() if k.startswith("THRESHOLDS!") and isinstance(v, str)
              and v.startswith("#")}
    assert errors == {}
    # Row 6 is the short series: no move, no lines 2 to 4, and the reason.
    assert [vals["THRESHOLDS!%s6" % c] for c in "CEFG"] == ["", "", "", ""]
    assert vals["THRESHOLDS!M6"] == "Refused: too little history kept for a typical yearly move"
    assert vals["THRESHOLDS!U6"] == (
        "Only 4 quarters, and nothing was left out by hand: a typical yearly move needs two quarters four "
        "apart, so only a longer series clears this. Its notes on Sources say where it starts and why.")
    # The series beside it is untouched.
    assert vals["THRESHOLDS!M5"] == "" and isinstance(vals["THRESHOLDS!C5"], float)


def test_a_smoothed_workbook_builds_its_live_lines_from_the_smoothed_series(tmp_path):
    """The Codex review (13 Oct 2026): with smoothing above 1 the evidence and
    backtest used the smoothed series while Data, Thresholds and the score
    used the raw file, so the workbook disagreed with itself. Now every line
    is built from the smoothed values, and Excel's equal Python's on them."""
    from threshold_engine.series import check, smooth
    p = tmp_path / "s.xlsx"
    make(p, smoothing=4)
    vals = recalc(p)
    da = load_workbook(p)["Data"]
    for r, (label, path) in enumerate(SERIES, start=5):
        sm = smooth(check(load(path.name), "quarterly"), 4)
        k = [c.value for c in da[1]].index(label) + 1
        col = {da.cell(i, 1).value: da.cell(i, k).value for i in range(2, da.max_row + 1)}
        assert [col[quarter(x.date)] for x in sm] == pytest.approx([x.value for x in sm])
        assert sum(v is not None for v in col.values()) == len(sm)
        ref = scale([x.value for x in sm], "higher_is_worse", SET["halfwidth"], SET["low_step"], SET["high_z"], True)
        got = row(vals, r)
        assert got[0] == pytest.approx(ref["median"]) and got[3:7] == pytest.approx(ref["bounds"])
        assert vals["THRESHOLDS!J%d" % r] == pytest.approx(sm[-1].value)
        assert got[7] == score(sm[-1].value, ref["bounds"], "higher_is_worse", "worse", True)
    run = [r for r in load_workbook(p)["Run"].iter_rows(values_only=True) if r and r[1] == "Smoothing"]
    assert run and "trailing mean over 4 quarters" in run[0][2]


def test_a_lower_is_worse_workbook_shades_five_bands_of_positive_width(tmp_path):
    """The Codex review (13 Oct 2026): for a lower-is-worse measure the lines
    fall, and the Assess chart's stacked bands came out with negative widths.
    The stack now starts at score 5 and every band is the gap between two
    lines, in the measure's own order."""
    src = tmp_path / "score.csv"
    rows = ["date,value"] + ["%s,%r" % (p.date, 800 - 40 * p.value) for p in load("cards_nco_ttm.csv")]
    src.write_text("\n".join(rows) + "\n", encoding="utf-8")
    p = tmp_path / "low.xlsx"
    build(p, [("Score", src)], name="Score", unit="", direction="lower_is_worse", frequency="quarterly",
          smoothing=1, floor_at_zero=False, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5,
          on_the_line="worse", half_lives=(), percentiles=(50, 75, 90, 95), horizon=4, min_history=41)
    vals = recalc(p)
    lines = [vals["THRESHOLDS!%s5" % c] for c in "EFGH"]            # 2..5 begin, falling
    assert lines == sorted(lines, reverse=True) and vals["THRESHOLDS!M5"] == ""
    heads = [vals.get("ASSESS!%s%d" % (L, 15)) for L in ("P", "Q", "R", "S", "T")]
    ws = load_workbook(p)["Assess"]
    first = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 26).value == "Quarter") + 1
    names = [ws.cell(first - 1, c).value for c in range(30, 35)]
    assert names == ["5 High", "4 Moderate-High", "3 Moderate", "2 Moderate-Low", "1 Low"]
    widths = [vals["ASSESS!%s%d" % (get_column_letter(c), first)] for c in range(30, 35)]
    floor = vals["ASSESS!AI%d" % (first + 1)]
    mn = min(min(lines), min(800 - 40 * x.value for x in load("cards_nco_ttm.csv")))
    assert 0 < floor < mn                     # not floored at zero: the axis starts just below the data
    # Nothing the chart plots sits at zero, or Excel stretches the axis down to it (the walkthrough, 13 Oct
    # 2026): the strip stands on an invisible base under the floor and is 0 high where nothing is left out.
    h = vals["ASSESS!AI%d" % (first + 2)]
    rows = range(first, first + len(load("cards_nco_ttm.csv")))
    assert all(vals["ASSESS!AM%d" % r] == pytest.approx(floor - h) for r in rows) and floor - h > 0
    assert all(vals["ASSESS!AK%d" % r] == 0 for r in rows)                   # nothing left out here
    assert all(vals["ASSESS!AJ%d" % r] == pytest.approx(floor) for r in rows)
    # and the axis itself starts at a round number just under that, not at Excel's automatic zero
    y_min = load_workbook(p)["Assess"]._charts[0].y_axis.scaling.min
    assert y_min is not None and 0 < y_min <= floor - h and floor - y_min < 0.2 * (top_of(vals, first) - floor)
    assert floor + widths[0] == pytest.approx(lines[3])                 # up to where High begins
    assert widths[1:4] == pytest.approx([lines[2] - lines[3], lines[1] - lines[2], lines[0] - lines[1]])
    assert all(x >= 0 for x in widths)


@pytest.mark.parametrize("rule", ["worse", "better"])
def test_a_lower_is_worse_workbook_equals_the_reference_scale(tmp_path, rule):
    """NEXT-GOAL section 2: a lower-is-worse measure (a credit score) was
    tested in scale.py and never through the workbook's lines and score. A
    mirrored cards series, built, recalculated, and held to scale() and
    score() with the direction reversed, under each on-the-line rule."""
    src = tmp_path / "score.csv"
    pts = load("cards_nco_ttm.csv")
    src.write_text("date,value\n" + "".join("%s,%r\n" % (x.date, 800 - 40 * x.value) for x in pts),
                   encoding="utf-8")
    p = tmp_path / ("low-%s.xlsx" % rule)
    build(p, [("Score", src)], name="Score", unit="", direction="lower_is_worse", frequency="quarterly",
          smoothing=1, floor_at_zero=False, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5,
          on_the_line=rule, half_lives=(), percentiles=(50, 75, 90, 95), horizon=4, min_history=41)
    vals = recalc(p)
    mirrored = [800 - 40 * x.value for x in pts]
    ref = scale(mirrored, "lower_is_worse", 0.5, 1.0, 2.5, False)
    got = row(vals, 5)
    assert vals["THRESHOLDS!M5"] == "" and ref["refused"] is None
    assert got[0] == pytest.approx(ref["median"]) and got[3:7] == pytest.approx(ref["bounds"])
    assert ref["bounds"] == sorted(ref["bounds"], reverse=True)          # lower is worse: lines fall
    assert vals["THRESHOLDS!J5"] == pytest.approx(mirrored[-1])
    assert got[7] == score(mirrored[-1], ref["bounds"], "lower_is_worse", rule, False)
    want = [sum(1 for x in mirrored if score(x, ref["bounds"], "lower_is_worse", rule, False) == k)
            for k in range(1, 6)]
    assert [vals["ASSESS!Q%d" % r] for r in range(10, 15)] == want


LO_REFUSED = ("Refused: a period left out for this product is not whole; its Status on Assess says what is "
              "missing")


@pytest.mark.parametrize("cells, status", [
    ({"B6": "2008Q1"}, "Missing: To, reason (Credit card refused)"),               # a start and nothing else
    ({"B6": "2012Q4", "C6": "2008Q1", "D6": "x"}, "From is after To (Credit card refused)"),
    ({"B6": "2008Q1", "C6": "2012Q4"}, "Missing: reason (Credit card refused)"),   # no reason
    ({"B6": "2008-Q1", "C6": "2012Q4", "D6": "x"}, "Missing: From (Credit card refused)"),   # not on the list
    ({"E6": "2020Q1", "F6": "2021Q1"}, "Missing: reason (Credit card refused)"),   # the second row
])
def test_a_leave_out_that_is_not_whole_is_refused_live(built, tmp_path, cells, status):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "lo.xlsx", cells)
    assert vals["THRESHOLDS!M5"] == LO_REFUSED
    assert vals["THRESHOLDS!K5"] == ""
    row = LO_TOP + (1 if "E6" in cells else 0)
    assert vals["ASSESS!T%d" % row] == status
    assert vals["THRESHOLDS!M6"] == "" and vals["THRESHOLDS!K6"] != ""          # no other product is touched


def test_a_row_without_a_product_asks_for_one_and_refuses_nothing(built, tmp_path):
    p, _, _ = built
    vals = with_assess(p, tmp_path, "noprod.xlsx", {"O%d" % LO_TOP: "2008Q1", "P%d" % LO_TOP: "2012Q4",
                                                    "Q%d" % LO_TOP: "x"})
    assert vals["ASSESS!T%d" % LO_TOP] == "Pick a series"
    assert all(vals["THRESHOLDS!M%d" % r] == "" for r in range(5, 9))


def test_any_number_of_periods_for_a_product_each_in_its_own_row(built, tmp_path):
    """The bank, 13 Oct 2026: leave-outs picked on Assess from lists, not two
    fixed pairs per product. Three periods for cards, in rows apart, with a
    period for mortgage between: each leaves out its own product's quarters."""
    p, _, _ = built
    rows = [("Credit card", "1994Q1", "1995Q4", "a"), ("Mortgage", "2008Q1", "2012Q4", "b"),
            ("Credit card", "2007Q4", "2012Q4", "c"), ("Credit card", "2021Q3", "2023Q1", "d")]
    cells = {}
    for i, (prod, a, b, why) in enumerate(rows):
        r = LO_TOP + 2 * i
        cells.update({"N%d" % r: prod, "O%d" % r: a, "P%d" % r: b, "Q%d" % r: why})
    vals = with_assess(p, tmp_path, "many.xlsx", cells)
    cuts = (("1994Q1", "1995Q4"), ("2007Q4", "2012Q4"), ("2021Q3", "2023Q1"))
    _, kept, ref = reference("cards_nco_ttm.csv", cuts)
    assert ref["refused"] is None and vals["THRESHOLDS!M5"] == ""
    assert row(vals, 5)[3:7] == pytest.approx(ref["bounds"])
    _, _, mref = reference("mortgage_nco_ttm.csv", (("2008Q1", "2012Q4"),))
    assert vals["THRESHOLDS!B6"] == pytest.approx(mref["median"])
    assert vals["THRESHOLDS!B7"] == pytest.approx(reference("home_equity_nco_ttm.csv")[2]["median"])


def test_a_whole_leave_out_is_not_refused(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "ok.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "the 2008 crisis"})
    assert vals["THRESHOLDS!M5"] == ""


@pytest.mark.parametrize("cells, why", [({"H6": -0.5}, "Moderate's half-width must be 0 or more"),
                                        ({"J6": 0}, "High's z must be above 0")])
def test_the_workbook_refuses_a_negative_half_width_or_z_live(built, tmp_path, cells, why):
    p, _, _ = built
    assert with_settings(p, tmp_path, "neg.xlsx", cells)["THRESHOLDS!M5"] == "Refused: " + why


def test_the_workbook_is_never_written_over_its_own_input(tmp_path):
    src = tmp_path / "cards.csv"
    src.write_bytes((DATA / "cards_nco_ttm.csv").read_bytes())
    before = src.read_bytes()
    for out in (src, tmp_path / "x" / ".." / "cards.csv"):
        with pytest.raises(ValueError, match="must be written to a .xlsx file"):
            make(out)
    alias = tmp_path / "cards.xlsx"
    alias.write_bytes(before)
    with pytest.raises(ValueError, match="would be written over cards.xlsx"):
        build(alias, [("Cards", alias)], name="NCO", unit="%", direction="higher_is_worse",
              frequency="quarterly", smoothing=1, floor_at_zero=True, moderate_halfwidth=0.5, low_step=1.0,
              high_z=2.5, on_the_line="worse", half_lives=(), percentiles=(50,), horizon=4, min_history=41)
    assert src.read_bytes() == before and alias.read_bytes() == before


def test_the_assess_bands_stack_from_the_floor_when_values_go_below_zero(tmp_path):
    """The Codex review (13 Oct 2026): with cut-offs below zero the first band
    was the cut-off itself, a negative width. The stack now starts at an
    invisible base, the chart's floor, and every band is a gap >= 0."""
    src = tmp_path / "neg.csv"
    src.write_text("date,value\n" + "".join("%s,%r\n" % (p.date, p.value - 6) for p in load("cards_nco_ttm.csv")),
                   encoding="utf-8")
    p = tmp_path / "neg.xlsx"
    build(p, [("Net", src)], name="Net", unit="%", direction="higher_is_worse", frequency="quarterly",
          smoothing=1, floor_at_zero=False, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5,
          on_the_line="worse", half_lives=(), percentiles=(50, 75, 90, 95), horizon=4, min_history=41)
    vals = recalc(p)
    lines = [vals["THRESHOLDS!%s5" % c] for c in "EFGH"]
    assert vals["THRESHOLDS!M5"] == "" and lines[0] < 0
    ws = load_workbook(p)["Assess"]
    first = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 26).value == "Quarter") + 1
    floor = vals["ASSESS!AI%d" % (first + 1)]
    top = vals["ASSESS!AI%d" % first]
    data = [v for v in (vals.get("ASSESS!AC%d" % r) for r in range(first, ws.max_row + 1)) if isinstance(v, float)]
    mn, mx = min(min(data), min(lines)), max(max(data), max(lines))
    assert floor == pytest.approx(mn - 0.08 * (mx - mn)) and top > mx        # not floored: 8% below the lowest
    assert vals["ASSESS!AJ%d" % first] == pytest.approx(floor)                  # the base
    widths = [vals["ASSESS!%s%d" % (get_column_letter(c), first)] for c in range(30, 35)]
    assert all(w >= 0 for w in widths)
    assert floor + sum(widths) == pytest.approx(top)
    assert floor + widths[0] == pytest.approx(lines[0])


def test_the_assess_axis_labels_negative_numbers_only_when_they_are_real(built, tmp_path):
    """Mortgage and home equity net charge-offs go below zero (net recoveries),
    and the axis used to hide every negative label; now it hides them only
    when nothing real can be negative (floored, and no value below zero)."""
    fmt = lambda p: load_workbook(p)["Assess"]._charts[0].y_axis.number_format.formatCode     # noqa: E731
    assert fmt(built[0]) == "0.0"                          # mortgage and home equity dip below zero
    cards = tmp_path / "cards.xlsx"
    build(cards, SERIES[:1], name="NCO", unit="%", direction="higher_is_worse", frequency="quarterly",
          smoothing=1, floor_at_zero=True, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5,
          on_the_line="worse", half_lives=(), percentiles=(50,), horizon=4, min_history=41)
    assert fmt(cards) == "0.0;;0.0"                        # nothing below zero: the strip is unlabelled
    assert load_workbook(cards)["Assess"]._charts[0].y_axis.scaling.min is None      # floored: Excel's own
    src = tmp_path / "neg.csv"
    src.write_text("date,value\n" + "".join("%s,%r\n" % (p.date, p.value - 6) for p in load("cards_nco_ttm.csv")),
                   encoding="utf-8")
    p = tmp_path / "neg.xlsx"
    build(p, [("Net", src)], name="Net", unit="%", direction="higher_is_worse", frequency="quarterly",
          smoothing=1, floor_at_zero=False, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5,
          on_the_line="worse", half_lives=(), percentiles=(50,), horizon=4, min_history=41)
    assert fmt(p) == "0.0"


def test_text_in_the_moderate_low_step_is_refused_not_an_error(built, tmp_path):
    p, _, _ = built
    vals = with_settings(p, tmp_path, "ls.xlsx", {"I6": "one"})
    assert vals["THRESHOLDS!M5"] == "Refused: Moderate-Low's step must be a number"
    assert vals["THRESHOLDS!K5"] == "" and vals["THRESHOLDS!E5"] == ""


@pytest.mark.parametrize("cell", ["H6", "I6", "J6"])
def test_text_in_a_setting_leaves_no_error_on_the_assess_dials(built, tmp_path, cell):
    """The Codex review (13 Oct 2026, on #426 after merge): the Thresholds tab
    refused text in the low step, but the Assess dials multiplied it and
    showed #VALUE!. Every dial row is now blank, and says why, when a setting
    it uses is not a number; the same holds for the half-width and z."""
    p, _, _ = built
    vals = with_settings(p, tmp_path, "t%s.xlsx" % cell, {cell: "one"})
    ws = load_workbook(p)["Assess"]
    rows = [r for r in range(dial_top(p), ws.max_row + 1) if isinstance(ws.cell(r, 3).value, str)
            and ws.cell(r, 3).value.startswith("=") and ws.cell(r, 12).value]
    assert rows
    for r in rows:
        for c in "BCDEFGHIJKL":
            v = vals.get("ASSESS!%s%d" % (c, r))
            assert not (isinstance(v, str) and v.startswith("#")), (c, r, v)
    refused = [vals["ASSESS!L%d" % r] for r in rows]
    assert any(v == "refused: a setting is not a number" for v in refused)


@pytest.mark.parametrize("product", ["Mortgage", "Home equity", "Other consumer"])
def test_assess_draws_every_band_for_a_product_that_starts_after_the_calendar(built, tmp_path, product):
    """The walkthrough (13 Oct 2026), in Excel: for a product whose data starts
    after the first quarter on the calendar, the chart's top and floor were
    MAX and MIN over a column holding #N/A, so both were #N/A, and Low and
    High vanished while the other bands were drawn from zero."""
    p = built[0]
    wb = load_workbook(p)
    wb["Assess"]["B5"] = product
    q = tmp_path / ("assess-%s.xlsx" % product.replace(" ", "-"))
    wb.save(q)
    vals = recalc(q)
    ws = wb["Assess"]
    first = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 26).value == "Quarter") + 1
    top, floor = vals["ASSESS!AI%d" % first], vals["ASSESS!AI%d" % (first + 1)]
    assert all(isinstance(x, (int, float)) and not isinstance(x, bool) for x in (top, floor)), (top, floor)
    lines = [vals["ASSESS!O%d" % r] for r in (11, 12, 13, 14)]
    data = [v for v in (vals.get("ASSESS!AC%d" % r) for r in range(first, ws.max_row + 1)) if isinstance(v, float)]
    assert floor == pytest.approx(min(0, min(data), min(lines))) and top > max(max(data), max(lines))
    widths = [vals["ASSESS!%s%d" % (get_column_letter(c), first)] for c in range(30, 35)]
    assert all(isinstance(w, (int, float)) and w >= 0 for w in widths), widths
    assert floor + widths[0] == pytest.approx(lines[0]) and floor + sum(widths) == pytest.approx(top)


def test_the_chart_tabs_key_is_the_swatch_over_each_switch(built):
    """The walkthrough (13 Oct 2026) found the legend over the cutoff lines;
    Claude Design's layout puts the switches in one row, each under its
    product's name with a swatch in its series colour, and the chart needs no
    legend (Excel 16 would also list the bands' invisible base)."""
    ws = load_workbook(built[0])["Chart"]
    assert ws._charts[0].legend is None
    for k, (label, _) in enumerate(SERIES):
        head = ws.cell(4, 2 + k)
        assert head.value == label and head.border.bottom.style == "thick"
        assert head.border.bottom.color.rgb.endswith(("0A0908", "CC0000", "57534B", "9C7A4A")[k])
        assert ws.cell(5, 2 + k).value in (True, False)


def test_the_backtest_glossary_entry_says_one_thing_about_the_history_used():
    """The Codex review (13 Oct 2026, on #426 after merge): the meaning said
    the scored quarter is included and the calculation said 'before it'."""
    from threshold_engine.glossary import TERMS
    entry = next(t for t in TERMS if t[0] == "Backtest")
    assert "the quarter rated and those before it" in entry[1]
    assert "that quarter and the quarters before it" in entry[2]


def test_the_pickers_read_their_choices_from_cells_so_long_labels_survive(tmp_path):
    """The Codex review (13 Oct 2026): an inline dropdown list over 255
    characters is one Excel drops or repairs. The choices are cells now."""
    labels = ["%s, a long product label to push the list well past Excel's limit" % k
              for k in ("Cards", "Mortgage", "Home equity", "Other consumer")]
    labels = [l.replace(",", "") for l in labels]
    assert len(",".join(labels)) > 255
    p = tmp_path / "long.xlsx"
    build(p, list(zip(labels, [s for _, s in SERIES])), name="NCO", unit="%", direction="higher_is_worse",
          frequency="quarterly", smoothing=1, floor_at_zero=True, moderate_halfwidth=0.5, low_step=1.0,
          high_z=2.5, on_the_line="worse", half_lives=(), percentiles=(50,), horizon=4, min_history=41)
    wb = load_workbook(p)
    for sheet in ("Chart", "Assess"):
        for dv in wb[sheet].data_validations.dataValidation:
            if dv.type != "list":
                continue                                  # an input message, not a list
            assert len(dv.formula1) < 255 and not dv.formula1.startswith('"' + labels[0][:5]), (sheet, dv.formula1)
    chart = wb["Chart"]
    col = next(c for c in range(1, chart.max_column + 1) if chart.cell(1, c).value == "(none)")
    assert [chart.cell(1 + i, col).value for i in range(5)] == ["(none)"] + labels
    assert chart.column_dimensions[get_column_letter(col)].hidden


def test_the_backtest_tab_says_the_scored_quarter_is_in_its_own_history(built):
    banner = " ".join(str(c.value) for row in load_workbook(built[0])["Backtest"].iter_rows(max_row=3) for c in row if c.value)
    assert "up to and including it" in banner and "history before it" not in banner


def _scen_rows(p):
    ws = load_workbook(p)["Assess"]
    top = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, 2).value or "").startswith("What happens if"))
    return top + 2, top + 3


def test_the_scenarios_set_your_lines_beside_keeping_every_quarter(built, tmp_path):
    """The bank, 13 Oct 2026: views on Assess that show what changes as
    quarters are left out. The every-quarter row equals the reference scale on
    the whole series, under the same settings; the other row is the table above."""
    p, _, _ = built
    vals = with_settings(p, tmp_path, "scen.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "Financial crisis",
                                                    "E6": "2021Q3", "F6": "2023Q1", "G6": "Pandemic forbearance"})
    ra, rb = _scen_rows(p)
    _, kept, ref = reference("cards_nco_ttm.csv", CARD_CUTS)
    pts, _, full = reference("cards_nco_ttm.csv")
    assert [vals["ASSESS!%s%d" % (c, ra)] for c in "CDEF"] == pytest.approx(ref["bounds"])
    assert [vals["ASSESS!%s%d" % (c, rb)] for c in "CDEF"] == pytest.approx(full["bounds"])
    sc = lambda v: score(v, full["bounds"], "higher_is_worse", "worse", True)      # noqa: E731
    assert [vals["ASSESS!%s%d" % (c, rb)] for c in "GHIJK"] == [sum(1 for x in pts if sc(x.value) == k)
                                                                for k in range(1, 6)]
    assert [vals["ASSESS!%s%d" % (c, ra)] for c in "GHIJK"] == [vals["ASSESS!Q%d" % r] for r in range(10, 15)]
    charts = load_workbook(p)["Assess"]._charts
    assert len(charts) == 5                      # the main chart, the comparison, a bar chart per setting


def test_a_refused_product_still_shows_where_its_lines_fall(built, tmp_path):
    """The bank, 13 Oct 2026, of other consumer refused at z 2.5 with the
    crisis left out: "what happened to the colors". The bands and counts stay
    while the lines are numbers in order; only the Check says refused."""
    p, _, _ = built
    vals = with_settings(p, tmp_path, "ref.xlsx", {"B6": "2007Q4", "C6": "2012Q4", "D6": "x", "J6": 3.5})
    assert vals["THRESHOLDS!M5"] == "Refused: High would begin beyond the worst quarter kept"
    assert vals["ASSESS!AI5"] == 1
    assert all(isinstance(vals["ASSESS!Q%d" % r], (int, float)) for r in range(10, 15))
    assert vals["ASSESS!Q14"] == 0                                  # nothing kept reaches High: the reason
    first = 2
    assert all(isinstance(vals["ASSESS!%s%d" % (get_column_letter(c), first)], (int, float)) for c in range(30, 35))


def test_assess_shows_the_latest_quarter_and_its_score(built):
    _, _, vals = built
    assert (vals["ASSESS!D5"], vals["ASSESS!F5"], vals["ASSESS!H5"]) == ("2026Q2", pytest.approx(3.978),
                                                                       "3 Moderate")
    assert vals["ASSESS!U12"].endswith("2026Q2") and vals["ASSESS!U11"] == ""      # the marker, by score 3


@pytest.mark.parametrize("cells, refusal", [
    ({"B6": "2008Q1"}, "not whole"),
    ({"H6": -0.1}, "half-width must be 0 or more"),
    ({"J6": 0}, "High's z must be above 0"),
    ({"I6": "one"}, "step must be a number"),
    ({"B6": "1985Q4", "C6": "2025Q2", "D6": "t"}, "too little history"),
    ({"I6": 0.4}, "Moderate-Low must begin further"),
    ({"J6": 0.3}, "High must begin beyond"),
    ({"J6": 8}, "beyond the worst quarter kept"),
    ({"I6": 6}, "Moderate-Low would begin at or below zero"),
])
def test_every_refusal_says_what_to_fix(built, tmp_path, cells, refusal):
    """Claude Design's hand-off (13 Oct 2026): beside every refusal, a "To fix"
    in the same order, its numbers worked out live; the Check on Assess
    carries both."""
    p, _, _ = built
    vals = with_settings(p, tmp_path, "fix.xlsx", cells)
    assert refusal in vals["THRESHOLDS!M5"]
    fix = vals["THRESHOLDS!U5"]
    assert fix and fix.endswith(".")
    assert vals["ASSESS!B6"] == vals["THRESHOLDS!M5"] + ". To fix: " + fix


@pytest.mark.parametrize("cells, setting", [({"J6": 8}, "J6"), ({"J6": 0.3}, "J6"), ({"I6": 6}, "I6")])
def test_following_the_to_fix_clears_the_refusal(built, tmp_path, cells, setting):
    """The number a To fix gives is one that works: type it and the lines stand."""
    import re
    p, _, _ = built
    vals = with_settings(p, tmp_path, "fix1.xlsx", cells)
    nums = re.findall(r"\d+\.\d\d", vals["THRESHOLDS!U5"])
    assert nums
    ok = with_settings(p, tmp_path, "fix2.xlsx", {setting: float(nums[0]) + (0.01 if "Raise" in
                                                                             vals["THRESHOLDS!U5"] else 0)})
    assert ok["THRESHOLDS!M5"] == "", (vals["THRESHOLDS!U5"], ok["THRESHOLDS!M5"])


def test_the_sources_tab_prints_what_each_file_says_about_itself(built):
    """The bank, 13 Oct 2026: define the data and how it is used, as a tab.
    Generic by construction: the tab prints each file's own description lines
    (the keys the internal templates ask for) and adds nothing of its own."""
    from threshold_engine.series import read_notes
    p, _, _ = built
    ws = load_workbook(p)["Sources"]
    cells = {(ws.cell(r, 2).value, ws.cell(r, 3).value) for r in range(1, ws.max_row + 1)}
    for label, path in SERIES:
        notes = dict(read_notes(path))
        assert ("Measure", notes["measure"]) in cells and ("Source", notes["source"]) in cells
    assert not any(v == "Not stated in the file" for _, v in cells)          # every public file says it all
    assert load_workbook(p).sheetnames.index("Sources") == load_workbook(p).sheetnames.index("Glossary") - 1


def test_a_file_that_says_nothing_about_itself_is_marked_not_filled_in(tmp_path):
    """A file without description lines gets every core key marked "Not
    stated in the file": nothing is inferred to fill the gap."""
    from threshold_engine.series import CORE_NOTES
    bare = tmp_path / "bare.csv"
    bare.write_text("\n".join(l for l in (DATA / "cards_nco_ttm.csv").read_text().splitlines()
                              if not l.startswith("#")) + "\n", encoding="utf-8")
    p = tmp_path / "bare.xlsx"
    build(p, [("Bare", bare)], name="N", unit="%", direction="higher_is_worse", frequency="quarterly", smoothing=1,
          floor_at_zero=True, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5, on_the_line="worse", half_lives=(),
          percentiles=(50,), horizon=4, min_history=41)
    ws = load_workbook(p)["Sources"]
    marked = [ws.cell(r, 2).value for r in range(1, ws.max_row + 1) if ws.cell(r, 3).value == "Not stated in the file"]
    assert len(marked) == len(CORE_NOTES)


def test_the_fixed_axis_holds_every_line_the_dials_offer(tmp_path):
    """Codex on #432: a measure not floored at zero gets a fixed axis minimum
    at build, which cannot follow later typing, so a larger Moderate-Low step
    could push score 2's line off the bottom. The minimum now sits under every
    line the dials try. A noisy series makes the point: at step 2 its line is
    below every value it ever took."""
    import random
    rnd = random.Random(7)
    src = tmp_path / "noisy.csv"
    pts = load("cards_nco_ttm.csv")
    src.write_text("date,value\n" + "".join("%s,%r\n" % (x.date, round(100 + 2 * rnd.random(), 3)) for x in pts),
                   encoding="utf-8")
    p = tmp_path / "noisy.xlsx"
    build(p, [("Noisy", src)], name="N", unit="", direction="higher_is_worse", frequency="quarterly", smoothing=1,
          floor_at_zero=False, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5, on_the_line="worse", half_lives=(),
          percentiles=(50,), horizon=4, min_history=41)
    vals = [float(l.split(",")[1]) for l in src.read_text().splitlines()[1:]]
    ref = scale(vals, "higher_is_worse", 0.5, 2.0, 2.5, False)
    low_line = ref["median"] - 2.0 * ref["typical_yearly_move"]
    assert low_line < min(vals)                                   # the case that was clipped
    y_min = load_workbook(p)["Assess"]._charts[0].y_axis.scaling.min
    assert y_min is not None and y_min < low_line


def test_a_line_below_the_fixed_axis_is_said_not_clipped_in_silence(tmp_path):
    """Codex on #432: a fixed axis cannot follow every leave-out or setting.
    When the floor or the left-out strip falls below it, the row over the
    chart says so; as built it says nothing."""
    import random
    rnd = random.Random(7)
    src = tmp_path / "noisy.csv"
    pts = load("cards_nco_ttm.csv")
    src.write_text("date,value\n" + "".join("%s,%r\n" % (x.date, round(100 + 2 * rnd.random(), 3)) for x in pts),
                   encoding="utf-8")
    p = tmp_path / "noisy.xlsx"
    build(p, [("Noisy", src)], name="N", unit="", direction="higher_is_worse", frequency="quarterly", smoothing=1,
          floor_at_zero=False, moderate_halfwidth=0.5, low_step=1.0, high_z=2.5, on_the_line="worse", half_lives=(),
          percentiles=(50,), horizon=4, min_history=41)
    assert recalc(p)["ASSESS!B7"] == ""
    vals = with_settings(p, tmp_path, "far.xlsx", {"I6": 6})
    assert "below the axis of this chart" in vals["ASSESS!B7"]


def test_many_series_build_with_the_chart_helpers_clear_of_the_title(tmp_path):
    """14 Oct 2026, one workbook holding every product, measure and source (95
    series): with the Chart tab's hidden helpers fixed at column AN, the title
    over 36 or more series ran into them and the build failed. Each helper
    column now holds what it says, the switch row stays visible."""
    n = 40
    p = tmp_path / "many.xlsx"
    build(p, [("Series %d" % k, DATA / "cards_nco_ttm.csv") for k in range(n)], name="N", unit="%",
          direction="higher_is_worse", frequency="quarterly", smoothing=1, floor_at_zero=True,
          moderate_halfwidth=0.5, low_step=1.0, high_z=2.5, on_the_line="worse", half_lives=(),
          percentiles=(50,), horizon=4, min_history=41)
    ch = load_workbook(p)["Chart"]
    q = next(c for c in range(1, ch.max_column + 1) if ch.cell(1, c).value == "Quarter")
    assert q > 2 + n + 6                                          # past the switches, the picker and the chart
    assert [ch.cell(1, q + 1 + k).value for k in range(n)] == ["Series %d" % k for k in range(n)]
    assert [ch.cell(4, 2 + k).value for k in range(n)] == ["Series %d" % k for k in range(n)]
    assert not any(ch.column_dimensions[get_column_letter(2 + k)].hidden for k in range(n))
