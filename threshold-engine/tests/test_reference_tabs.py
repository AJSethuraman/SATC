"""The reference and data tabs on round 3's one template (reference_tabs.py):
Evidence as two Excel tables, Glossary with a Tab column, Sources in three
sections with Reporting changes folded in, Run with the folder once and file
names only, and Data with the quarters before any value folded away.

Built with Decide (the book's tabs) from the test series, a left-out file and
the repository's registry of reporting changes.
"""
import re

import pytest
from openpyxl import Workbook, load_workbook

from conftest import ROOT
from threshold_engine import reference_tabs as RT
from threshold_engine import reporting_changes as RC
from threshold_engine.glossary import TERMS, by_tab
from threshold_engine.series import CORE_NOTES, read_notes
from threshold_engine.workbook import build
from test_decide import SER, SETTINGS, decide_spec

REGISTRY = ROOT / "config" / "reporting-changes.json"
CHANGES = RC.load(REGISTRY)
LEFT_OUT = ("# Quarters each series left out as built.\n"
            "Product,Measure,Source,first,last,why\n"
            "Card,Net,Bank,2007Q1,2013Q3,\"before the last break: 2013Q3 has no value\"\n")


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    d = tmp_path_factory.mktemp("reference")
    lo = d / "left-out.csv"
    lo.write_text(LEFT_OUT, encoding="utf-8")
    info = build(d / "ref.xlsx", SER, decide=decide_spec(), left_out_file=lo,
                 reporting=RC.for_sources(CHANGES, REGISTRY), **SETTINGS)
    return d, info, load_workbook(d / "ref.xlsx")


def head(ws, row, first=2):
    out = []
    c = first
    while ws.cell(row, c).value is not None or ws.cell(row, c).fill.fgColor.rgb in ("00E4DFD5", "FFE4DFD5"):
        out.append(ws.cell(row, c).value)
        c += 1
    return out


def find(ws, text, col=2):
    return next(r for r in range(1, ws.max_row + 1) if str(ws.cell(r, col).value or "").startswith(text))


# ---------------------------------------------------------------- the template

def test_each_tab_has_the_reference_title_and_a_thin_stone_rule(built):
    _, _, wb = built
    for name in ("Evidence", "Glossary", "Sources", "Run"):
        ws = wb[name]
        t = ws["B1"]
        assert (t.font.name, t.font.sz, t.font.b) == ("Arial", 14, True), name
        assert ws["B2"].border.bottom.style == "thin" and ws["B2"].border.bottom.color.rgb.endswith("B9B4AC"), name
        assert ws.row_dimensions[1].height == 25.5


# ---------------------------------------------------------------- Evidence

def test_evidence_is_two_excel_tables_named_on_every_row(built):
    _, info, wb = built
    ws = wb["Evidence"]
    assert set(ws.tables) == {"Spells", "Departures"}
    sp, dp = ws.tables["Spells"].ref, ws.tables["Departures"].ref
    assert re.fullmatch(r"B\d+:L\d+", sp) and re.fullmatch(r"B\d+:K\d+", dp)
    h = int(sp.split(":")[0][1:])
    assert [ws.cell(h, c).value for c in range(2, 13)] == ["Product", "Measure", "Source"] + list(RT.SPELL_HEAD)
    hd = int(dp.split(":")[0][1:])
    assert [ws.cell(hd, c).value for c in range(2, 12)] == ["Product", "Measure", "Source"] + list(RT.DEPART_HEAD)
    # every row names its series; one row per spell and per departure
    last = int(sp.split(":")[1][1:])
    rows = [[ws.cell(r, c).value for c in range(2, 5)] for r in range(h + 1, last + 1)]
    assert all(all(x) for x in rows)
    n_spells = sum(len(e["spells"]) for e in info["evidence"].values())
    assert len(rows) == n_spells
    # the robust z format, and the filter buttons on Product, Measure and Source only
    assert ws.cell(h + 1, 9).number_format == "+0.0;-0.0;0.0"
    hidden = {f.colId for f in ws.tables["Spells"].autoFilter.filterColumn if f.hiddenButton}
    assert hidden == set(range(3, 11))


def test_evidence_hides_repeated_names_and_bolds_an_unusual_z(built):
    _, _, wb = built
    ws = wb["Evidence"]
    rules = {str(rng): [r.formula[0] for r in rs] for rng, rs in ws.conditional_formatting._cf_rules.items()}
    sp = ws.tables["Spells"].ref
    first = int(sp.split(":")[0][1:]) + 1
    assert any(f == "AND($B%d=$B%d)" % (first, first - 1) for fs in rules.values() for f in fs)
    assert any(f == "AND($B%d=$B%d,$C%d=$C%d,$D%d=$D%d)" % ((first, first - 1) * 3) for fs in rules.values()
               for f in fs)
    assert any("ABS($I%d)>=3.5" % first in f for fs in rules.values() for f in fs)


# ---------------------------------------------------------------- Glossary

def test_the_glossary_has_a_tab_column_and_every_term(built):
    _, _, wb = built
    gl = wb["Glossary"]
    assert [gl.cell(4, c).value for c in range(2, 6)] == list(RT.GLOSSARY_HEAD)
    rows = [tuple(gl.cell(r, c).value for c in range(2, 6)) for r in range(5, gl.max_row + 1)
            if gl.cell(r, 3).value]
    assert rows == by_tab() and {r[1] for r in rows} == {t for t, _, _ in TERMS}
    assert gl.auto_filter.ref == "B4:E%d" % (4 + len(rows))
    assert {f.colId for f in gl.auto_filter.filterColumn if f.hiddenButton} == {1, 2, 3}
    assert gl.cell(5, 3).font.b and gl.cell(5, 2).font.sz == 10


def test_the_new_signals_terms_are_defined():
    terms = {t: (m, c) for t, m, c in TERMS}
    for t in ("Pack line", "Rescaled", "Lead", "Fit (r)", "Window for the normal", "Shift", "Jump row"):
        assert t in terms and all(terms[t]), t
    assert "Peer group in this chart" not in terms                      # the pack line replaced the panel


def test_peer_terms_can_be_added_under_their_tab(built):
    d, _, wb = built
    gl = wb["Glossary"]
    last = RT.append_terms(gl, "Peers", [("A peer term", "What it is.", "How.")])
    assert [gl.cell(last, c).value for c in range(2, 6)] == ["Peers", "A peer term", "What it is.", "How."]
    assert gl.auto_filter.ref.endswith("E%d" % last)


# ---------------------------------------------------------------- Sources

def test_sources_prints_shared_lines_once_and_each_file_on_one_row(built):
    _, _, wb = built
    ws = wb["Sources"]
    h = find(ws, "Product")
    assert [ws.cell(h, c).value for c in range(2, 12)][:7] == ["Product", "Measure", "Source", "File", "Rows",
                                                               "First", "Last"]
    rows = {ws.cell(r, 5).value: r for r in range(h + 1, h + 1 + len(SER))}
    assert set(rows) == {p.name for _, p in SER}
    shared = {ws.cell(r, 2).value: r for r in range(5, h) if re.fullmatch(r"S\d+", str(ws.cell(r, 2).value))}
    text_of = {k: ((ws.cell(r, 3).value + ": ") if ws.cell(r, 3).value else "") + ws.cell(r, 4).value
               for k, r in shared.items()}
    for lbl, path in SER:
        r = rows[path.name]
        keys = (ws.cell(r, 9).value or "").split()
        own = ws.cell(r, 11).value
        own = str(own) if own is not None else ""
        for k, t in read_notes(path):
            ln = t if k == "note" else "%s: %s" % (k, t)
            if k == RC.NOTE_KEY:
                continue
            assert ln in own or any(text_of[x] == RT.short_paths(ln) for x in keys), (path.name, ln)
        assert ws.cell(r, 6).value == len([x for x in path.read_text().splitlines()
                                           if x and not x.startswith("#")]) - 1


def test_sources_marks_a_required_line_a_file_does_not_give(tmp_path):
    bare = tmp_path / "bare.csv"
    src = SER[0][1].read_text().splitlines()
    bare.write_text("\n".join(x for x in src if not x.startswith("#")) + "\n", encoding="utf-8")
    p = tmp_path / "bare.xlsx"
    build(p, [("Bare", bare)], **SETTINGS)
    ws = load_workbook(p, rich_text=True)["Sources"]
    h = find(ws, "Series")
    cell = str(ws.cell(h + 1, 2 + 7).value)
    assert "Not stated: Measure, As of, Basis, Source." in cell
    assert len(CORE_NOTES) == 4


def test_reporting_changes_are_the_last_section_of_sources_with_linked_documents(built):
    _, info, wb = built
    ws = wb["Sources"]
    r = find(ws, "Reporting changes")
    assert ws.cell(r, 2).value == "Reporting changes \u00b7 12 rows, from reporting-changes.json"
    assert ws.cell(r + 1, 2).value == RC.what(REGISTRY)
    h = r + 2
    heads = [ws.cell(h, c).value for c in range(2, 12) if ws.cell(h, c).value]
    assert heads == list(RT.RC_HEAD)
    names = [ws.cell(x, 2).value for x in range(h + 1, ws.max_row + 1) if ws.cell(x, 2).value]
    assert names == [c["name"] for c in CHANGES]
    links = [ws.cell(x, 11) for x in range(h + 1, ws.max_row + 1) if ws.cell(x, 11).hyperlink]
    assert [c.hyperlink.target for c in links] == [s["url"] for c in CHANGES for s in c["sources"]]
    assert [c.value for c in links] == [s["citation"] for c in CHANGES for s in c["sources"]]


def test_no_full_local_path_is_printed_outside_runs_folder_row(built):
    d, _, wb = built
    folder = str(d)
    for name in ("Sources", "Evidence", "Glossary"):
        for row in wb[name].iter_rows(values_only=True):
            assert not any(folder in str(v) for v in row if v), name
    run = wb["Run"]
    hits = [(r, c) for r in range(1, run.max_row + 1) for c in range(1, run.max_column + 1)
            if str(SER[0][1].parent) in str(run.cell(r, c).value or "")]
    assert len(hits) == 1 and run.cell(hits[0][0], 2).value == "Series folder"


def test_a_long_path_in_a_note_is_cut_to_its_last_two_folders():
    assert RT.short_paths("from the cache C:\\Users\\ajish\\Temp\\scratchpad\\ind_cache.") == \
        "from the cache \u2026\\scratchpad\\ind_cache."
    assert RT.short_paths("in /home/me/work/cache/q.json now") == "in \u2026/cache/q.json now"
    url = "see https://www.federalreserve.gov/releases/chargeoff/default.htm (also FRED)"
    assert RT.short_paths(url) == url


# ---------------------------------------------------------------- Run

def test_run_prints_the_folder_once_and_each_file_by_name(built):
    d, _, wb = built
    ws = wb["Run"]
    r = find(ws, "Files")
    h = r + 1
    names = {ws.cell(h, c).value: c for c in range(2, 10) if ws.cell(h, c).value}
    assert list(names) == ["Product", "Measure", "Source", "First", "Last", "Rows", "File", "SHA-256"]
    import hashlib
    for i, (lbl, path) in enumerate(SER):
        row = h + 1 + i
        assert [ws.cell(row, c).value for c in range(2, 5)] == lbl.split(" - ")
        assert ws.cell(row, names["File"]).value == path.name
        sha = ws.cell(row, names["SHA-256"])
        assert sha.value == hashlib.sha256(path.read_bytes()).hexdigest()
        assert (sha.font.name, sha.font.sz, sha.font.color.rgb[-6:]) == ("Consolas", 9, "57534B")


def test_run_shows_left_out_as_built_after_this_build(built):
    d, _, wb = built
    ws = wb["Run"]
    lo = find(ws, "Left out as built")
    assert lo < find(ws, "Files")
    names = [ws.cell(lo + 2, c).value for c in range(2, 10)]
    assert names == ["Product", "Measure", "Source", "From", "To", None, "Why", None]
    assert [ws.cell(lo + 3, c).value for c in (2, 3, 4, 5, 6, 8)] == [
        "Card", "Net", "Bank", "2007Q1", "2013Q3", "before the last break: 2013Q3 has no value"]
    assert ws.cell(lo + 1, 8).value == "left-out.csv"


# ---------------------------------------------------------------- Data

def test_data_groups_only_the_quarters_before_any_value():
    wb = Workbook()
    ws = wb.active
    ws.append(["Quarter", "a", "b"])
    for i in range(10):
        ws.append(["Q%d" % i, None if i < 4 else 1.0, "=A%d" % (i + 2)])
    assert RT.data_layout(ws, 1, 2, 11, 3, "0.0000") == (2, 5)
    assert all(ws.row_dimensions[r].hidden and ws.row_dimensions[r].outline_level == 1 for r in range(2, 6))
    assert not ws.row_dimensions[6].hidden and ws.freeze_panes == "B2"
    assert ws.row_dimensions[1].height == 45 and ws["B1"].fill.fgColor.rgb.endswith("E4DFD5")
    ws2 = Workbook().active
    ws2.append(["Quarter", "a"])
    ws2.append(["Q0", 2.0])
    assert RT.data_layout(ws2, 1, 2, 2, 2, "0.0000") == (0, 0)


def test_the_books_data_tab_starts_at_its_first_value(built):
    _, _, wb = built
    da = wb["Data"]
    assert da.freeze_panes == "B2" and da.row_dimensions[1].height == 45
    assert not any(da.row_dimensions[r].hidden for r in range(2, da.max_row + 1))     # the first quarter has one
