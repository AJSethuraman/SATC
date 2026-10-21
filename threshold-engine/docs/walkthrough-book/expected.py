"""The book walkthrough's expected answers, from the source files, not the workbook.

Run from threshold-engine/, after the build (WALKTHROUGH.md, step 1):
    python docs/walkthrough-book/expected.py
Writes steps.json (what to click and type, where, and what each cell must then
say) and EXPECTED.md (the same, for a person). Every number comes from the
series CSVs the build read and scale.py, the reference the workbook's formulas
are held to; walk.ps1 does each step in Excel 16 and reads the answer back.

One set of lines per product and measure (the firm, 15 Oct 2026: "we don't
need to set thresholds on all 3, we need to use all 3 to set appropriate
thresholds"): Decide's row picks the source whose history sets the lines, its
settings apply to every source's own lines (Lines by source), its period left
out to the source picked, and every source's latest quarter is rated on the
lines picked (Thresholds, Summary).

The steps are cumulative, as a person's session is: what step 3 types is
still there in step 4, and step 10 saves, closes and reopens the file to show
the decisions were kept.
"""
import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "src"))
from threshold_engine import book as BK                              # noqa: E402
from threshold_engine import decide_tab as DT                        # noqa: E402
from threshold_engine import summary_tab as SU                       # noqa: E402
from threshold_engine import thresholds_tab as TT                    # noqa: E402
from threshold_engine.scale import scale, score                      # noqa: E402
from threshold_engine.series import check, read_csv                  # noqa: E402
from threshold_engine.workbook import assess_layout, quarter         # noqa: E402
from openpyxl.utils import get_column_letter as L                    # noqa: E402

HERE = Path(__file__).parent
BOOK = Path(sys.argv[1]).resolve() if len(sys.argv) > 1 else ROOT / "books" / "consumer-public.json"
RATINGS = ["Low", "Moderate-Low", "Moderate", "Moderate-High", "High"]
LBS = DT.LINES_TAB

bk = BK.load(BOOK)
W = bk["workbook"]
SER = bk["series"]
LABELS = [lbl for lbl, _ in SER]
N = len(SER)
PARTS = bk["decide"]["parts"]
NAMES = bk["decide"]["names"]
COMPARE = bk["decide"]["compare"]
PAIRS, BY = DT.pairs_of(LABELS, PARTS)
THIRD = DT.sources_of(LABELS, PARTS)
LEAD = COMPARE[0] if COMPARE else THIRD[0]
COLS = DT.columns(NAMES)
from threshold_engine.peers import roster as _RO                    # noqa: E402
N_PEERS = len(_RO.peers(_RO.load(bk["peers"]["roster"]))) if bk["peers"] else 0
DL = DT.layout(len(PAIRS), len(THIRD), COMPARE is not None, N_PEERS)
AL = assess_layout(N)
SL = SU.layout(LABELS, PARTS, COMPARE)
TL = TT.layout(SL["sources"], SL["dropped"])
PTS = {lbl: check(read_csv(p), "quarterly") for lbl, p in SER}
FLOOR_OFF = set(bk["decide"]["floor_off"])


def floor_of(lbl):
    """The series' floor at zero: the workbook's, unless its measure has none (floor_off)."""
    return bool(W["floor_at_zero"]) and PARTS[lbl][1] not in FLOOR_OFF


def dcell(name, pr):
    return "%s%d" % (L(COLS[name]), DL["first"] + PAIRS.index(pr))


def refusal_text(msg):
    if msg.startswith("too little history"):
        return "Refused: too little history kept for a typical yearly move"
    if msg.startswith("High would begin beyond the worst"):
        return "Refused: High would begin beyond the worst quarter kept"
    return "Refused: " + msg


def own(lbl, window=None, hw=None, ls=None, z=None, rule=None):
    """A series' own lines, as Lines by source holds them, keyed by the plain
    layout's letters, plus "bounds" (None when refused)."""
    hw = W["moderate_halfwidth"] if hw is None else hw
    ls = W["low_step"] if ls is None else ls
    z = W["high_z"] if z is None else z
    rule = W["on_the_line"] if rule is None else rule
    kept = [None if window and window[0] <= quarter(p.date) <= window[1] else p.value for p in PTS[lbl]]
    ref = scale(kept, W["direction"], hw, ls, z, floor_of(lbl))
    latest = PTS[lbl][-1]
    row = {"B": ref["median"], "C": ref["typical_yearly_move"], "D": max(v for v in kept if v is not None),
           "I": quarter(latest.date), "J": latest.value, "bounds": None, "rule": rule}
    if ref["refused"]:
        row.update({"K": "", "M": refusal_text(ref["refused"])})
        return row
    row.update(zip("EFGH", ref["bounds"]))
    row["bounds"] = ref["bounds"]
    sc = score(latest.value, ref["bounds"], W["direction"], rule, floor_of(lbl))
    row.update({"K": sc, "L": RATINGS[sc - 1], "M": ""})
    return row


def fits_reads(lbl, window=None, hw=None, ls=None, z=None):
    """What fits the history the lines are set from, on Decide, beside each
    setting: the ranges from the series' own median, yearly move, spread and
    worst kept (scale.py), worded and rounded as the cells word them."""
    from decimal import ROUND_HALF_UP, Decimal
    hw = W["moderate_halfwidth"] if hw is None else hw
    ls = W["low_step"] if ls is None else ls
    z = W["high_z"] if z is None else z
    kept = [None if window and window[0] <= quarter(p.date) <= window[1] else p.value for p in PTS[lbl]]
    ref = scale(kept, W["direction"], hw, ls, z, floor_of(lbl))
    med, mv, sp, wo = ref["median"], ref["typical_yearly_move"], ref["spread_of_levels"], ref["worst"]
    t = lambda x: str(Decimal(repr(x)).quantize(Decimal("0.01"), ROUND_HALF_UP))      # noqa: E731
    say = lambda ok: "It fits." if ok else "It does not fit."                        # noqa: E731
    floor = floor_of(lbl) and W["direction"] == "higher_is_worse"
    sg = 1 if W["direction"] == "higher_is_worse" else -1
    zmin, zmax, lsmax = hw * mv / sp, sg * (wo - med) / sp, med / mv
    texts = ["Fits below %s, where Moderate-Low begins. %s" % (t(ls), say(hw < ls)),
             "Fits above %s%s. %s" % (t(hw), ", and below %s so line 2 stays above zero" % t(lsmax) if floor else "",
                                     say(ls > hw and (not floor or ls < lsmax))),
             "Fits above %s (High beyond Moderate-High) and up to %s (some quarter kept can rate High). %s"
             % (t(zmin), t(zmax), say(zmin < z <= zmax)),
             "Either fits."]
    return [{"sheet": "Decide", "cell": "L%d" % (DL["settings"] + 1 + i), "expect": x} for i, x in enumerate(texts)]


#: own() keys its answers by the plain layout's letters; Lines by source lays
#: the table out as decide_tab.TH_BOOK, so each is read by meaning.
PLAIN_TO_BOOK = {DT.TH_PLAIN[k]: DT.TH_BOOK[k] for k in DT.TH_PLAIN}


def lbs_reads(lbl, row):
    r = 5 + LABELS.index(lbl)
    return [{"sheet": LBS, "cell": "%s%d" % (PLAIN_TO_BOOK[c], r), "expect": v}
            for c, v in row.items() if c in PLAIN_TO_BOOK]


def on_lines(lbl, lines):
    """A source's latest value rated on the lines picked: (number, text)."""
    if lines["bounds"] is None:
        return "", "Refused"
    sc = score(PTS[lbl][-1].value, lines["bounds"], W["direction"], lines["rule"], floor_of(lbl))
    return sc, "%d %s" % (sc, RATINGS[sc - 1])


def row_reads(pr, src, lines):
    """Thresholds' row for a product and measure, and Decide's two tables,
    with the lines set from ``src`` (``lines``, from own())."""
    r = 5 + PAIRS.index(pr)
    c, per = TL["cols"], TL["per"]
    out = [{"sheet": "Thresholds", "cell": "%s%d" % (c["lines_from"], r), "expect": src}]
    # Round 3: what the lines rest on (median, yearly move, worst kept) beside them.
    out += [{"sheet": "Thresholds", "cell": "%s%d" % (c[k], r), "expect": lines[v]}
            for k, v in (("median", "B"), ("move", "C"), ("worst", "D"))]
    for k, v in zip(("l2", "l3", "l4", "l5"), "EFGH"):
        out.append({"sheet": "Thresholds", "cell": "%s%d" % (c[k], r), "expect": lines[v] if lines["bounds"] else ""})
    for s, lbl in BY[pr].items():
        _, text = on_lines(lbl, lines)
        out.append({"sheet": "Thresholds", "cell": "%s%d" % (per[s]["rating"], r), "expect": text})
        out.append({"sheet": "Thresholds", "cell": "%s%d" % (per[s]["latest"], r), "expect": PTS[lbl][-1].value})
    return out


def decide_reads(pr, lines, others):
    """Decide's screen for the row shown: every source on the lines picked,
    and what each source's own history gives (``others``: source -> own())."""
    out = []
    for i, s in enumerate(THIRD):
        lbl = BY[pr].get(s)
        if lbl is None:
            continue
        k, text = on_lines(lbl, lines)
        out += [{"sheet": "Decide", "cell": "M%d" % (DL["compare_first"] + i),
                 "expect": text if lines["bounds"] else "No rating"},
                {"sheet": "Decide", "cell": "N%d" % (DL["compare_first"] + i), "expect": PTS[lbl][-1].value}]
        o = others[s]
        g = DL["gives_first"] + i
        if o["bounds"]:
            out += [{"sheet": "Decide", "cell": "%s%d" % (c, g), "expect": o[v]} for c, v in zip("MNOP", "EFGH")]
        else:
            out.append({"sheet": "Decide", "cell": "M%d" % g, "expect": "Refused"})
        out.append({"sheet": "Decide", "cell": "Q%d" % g, "expect": o["D"]})
    return out


def chip(pr, lines):
    k, text = on_lines(BY[pr][LEAD], lines)
    return "No rating: see the Check" if k == "" else "%s: %s" % (LEAD, text)


def headline(pr, src):
    tail = ": lines set from %s" % src if src else ": no %s picked to set the lines yet" % NAMES[2].lower()
    return "%s, %s%s" % (pr[0], pr[1].lower(), tail)


def decided(n_acc, n_rej):
    return "Decided: %d of %d · accepted %d · rejected %d" % (n_acc + n_rej, len(PAIRS), n_acc, n_rej)


def summary_cell(lbl):
    """The row and column of a series' latest value on Summary; its rating is
    the next column (summary_tab.layout, the builder's own)."""
    r, latest, _ = SU.layout(LABELS, PARTS, COMPARE)["cells"][lbl]
    return r, latest


def slicers(**picks):
    return [{"type": "slicer", "caption": k, "value": v} for k, v in picks.items()]


def typed(pr, name, value):
    return {"type": "cell", "sheet": "Decide", "cell": dcell(name, pr), "value": value}


#: Periods left out on Decide: row k's column for each field.
LO_FIELD = dict(zip(("product", "measure", "from", "to", "why"), "BCDEF"))


def lo_cell(k, field):
    return "%s%d" % (LO_FIELD[field], DL["lo_first"] + k)


def lo_typed(k, field, value):
    return {"type": "cell", "sheet": "Decide", "cell": lo_cell(k, field), "value": value}


def every_own(pr, **kw):
    return {s: own(lbl, **kw) for s, lbl in BY[pr].items()}


SHEETS = ["Summary", "Decide", "Signals", "Peers", "Saved views", "Thresholds", "Internal", LBS, "Peer books",
          "Evidence", "Glossary", "Sources", "Run", "Data", "Signals data", "Internal data", "Peers data", "Peers delinquency data", "Assess",
          "Chart", "Peers workings", "Signals workings"]
HIDDEN = (["Internal"] if bk["internal"] is None else []) + (["Internal data"] if bk["internal"] is None else []) + [
    "Peers data", "Peers delinquency data", "Assess", "Chart", "Peers workings", "Signals workings"]
IND, PEER = [s for s in THIRD if s != LEAD][:2] if len(THIRD) >= 3 else (THIRD[0], THIRD[-1])
OK = "Check: OK, your settings fit the history the lines are set from."
steps = []
# ---- Round 3's Summary: the strip for a source left out with one series, the
# rows only such a source has, the attention block, the counts top right ----
SU_LAST = SL["last_col"]
STRIPS = [(2 + k, lbl) for k, lbl in enumerate(
    next(lbl for lbl in LABELS if PARTS[lbl][2] == s_) for s_ in SL["dropped"]
    if sum(1 for lbl in LABELS if PARTS[lbl][2] == s_) == 1)]
ATT_HEAD = SL["ends"][0] + 2
DOT = "·"


def strip_reads(rating_text=None):
    """Each strip, one cell under the counts: the series' latest, its quarter and the rating, or the dash."""
    out = []
    for r_, lbl in STRIPS:
        out.append({"sheet": "Summary", "cell": "%s%d" % (L(SU_LAST), r_),
                    "expect": "%s %s %s %s %s %s %.2f%% %s %s %s rating %s" % (
                        PARTS[lbl][2], DOT, PARTS[lbl][0], DOT, PARTS[lbl][1], DOT, PTS[lbl][-1].value, DOT,
                        quarter(PTS[lbl][-1].date), DOT, rating_text or SU.NO_LINES)})
    return out


def only_reads():
    """A row only a source left out has: "<source> only" where the lead's rating goes."""
    out = []
    for pm, srcs in SL["only"].items():
        r_, base = SL["placed"][pm]
        out.append({"sheet": "Summary", "cell": "%s%d" % (L(base + SL["offsets"][LEAD][1]), r_),
                    "expect": "%s only" % " and ".join(SU.short(x) for x in srcs)})
    return out


def counts(picked, hot, refused):
    t = "Lines picked: %d of %d %s %s rated 4 or 5: %d %s refused: %d" % (picked, len(PAIRS), DOT, LEAD, hot, DOT, refused)
    return [{"sheet": "Summary", "cell": "%s1" % L(SU_LAST), "expect": t},
            {"sheet": "Thresholds", "cell": "%s1" % TL["cols"]["check"], "expect": t}]


def undecided_reads(pr):
    """A row with no source picked on Thresholds: blank from Lines set from to Check, but each
    source's latest quarter."""
    r_ = 5 + PAIRS.index(pr)
    c_, per = TL["cols"], TL["per"]
    out = [{"sheet": "Thresholds", "cell": "%s%d" % (c_[k], r_), "expect": ""}
           for k in ("lines_from", "median", "move", "worst", "l2", "l3", "l4", "l5", "decision", "check")]
    for s_, lbl in BY[pr].items():
        out += [{"sheet": "Thresholds", "cell": "%s%d" % (per[s_]["rating"], r_), "expect": ""},
                {"sheet": "Thresholds", "cell": "%s%d" % (per[s_]["latest"], r_), "expect": PTS[lbl][-1].value}]
    return out


# ---- B1: open ----
first = PAIRS[0]
steps.append({"id": "B1", "what": "Open the workbook as built",
              "shows": "It opens on Summary; the tabs in order; two slicers, each with one button picked; Decide "
                       "shows the first product and measure, no source picked and nothing decided",
              "actions": [], "reads": [
                  {"special": "active", "expect": "Summary"},
                  # Round 3: work, then reference, then data, then the hidden tabs; Backtest folded into
                  # Lines by source, Reporting changes the last section of Sources.
                  {"special": "sheets", "expect": SHEETS},
                  {"special": "hidden", "expect": HIDDEN},
                  {"special": "slicers", "expect": {NAMES[0]: first[0], NAMES[1]: first[1]}},
                  {"special": "visible_rows", "sheet": "Decide", "expect": [DL["first"]]},
                  {"sheet": "Decide", "cell": "L1", "expect": decided(0, 0)},
                  {"sheet": "Decide", "cell": "B%d" % DL["headline"], "expect": headline(first, None)},
                  {"sheet": "Decide", "cell": "I%d" % DL["chip"], "expect": "No rating: see the Check"},
                  {"sheet": "Decide", "cell": "B%d" % DL["check"],
                   "expect": "Check: pick the %s whose history sets the lines, under %s."
                             % (NAMES[2].lower(), DT.LINES_FROM)},
                  {"special": "links", "expect": []},
                  {"special": "errors", "expect": 0},
                  {"sheet": "Summary", "cell": "B%d" % ATT_HEAD, "expect": "Needs your attention %s %s rated 4 or 5" % (DOT, LEAD)},
                  {"sheet": "Summary", "cell": "B%d" % (ATT_HEAD + 1), "expect": SU.NONE_HOT}]
                 + counts(0, 0, 0) + strip_reads() + only_reads() + undecided_reads(first),
              "picture": [{"sheet": "Decide", "range": "A1:Q%d" % (DL["notes"] + 2), "file": "b1-decide.png"},
                          {"sheet": "Summary", "range": "A1:%s%d" % (L(SU_LAST + 1), ATT_HEAD + 10), "file": "b1-summary.png"}]})

# ---- B2: every Summary cell, before a source is picked ----
reads = []
for lbl in LABELS:
    if lbl not in SL["cells"]:          # a source left out of Summary's columns: its strip, or Lines by source
        continue
    r, c, rc = SL["cells"][lbl]
    reads.append({"sheet": "Summary", "cell": "%s%d" % (L(c), r), "expect": PTS[lbl][-1].value})
    reads.append({"sheet": "Summary", "cell": "%s%d" % (L(rc), r), "expect": SU.NO_LINES})
reads += [own_read for lbl in LABELS for own_read in lbs_reads(lbl, own(lbl))]
# Round 3: Backtest folded into Lines by source. Each series' unweighted run (backtest.py, the
# reference) past the hidden workings: Spearman against the value ahead and the change ahead, the
# quarters tested, then the average change ahead by rating; a series too short has 0 quarters tested,
# no Spearman, and its reason under the table.
from threshold_engine.backtest import backtest as _backtest          # noqa: E402
from openpyxl.utils import column_index_from_string as _CI           # noqa: E402
FOLD0 = _CI(DT.TH_HIDDEN[1]) + 1


def _bt(lbl):
    try:
        return _backtest(read_csv(dict(SER)[lbl]), lbl, W["unit"], W["direction"], W["frequency"], W["smoothing"],
                         W["percentiles"], W["horizon"], W["min_history"], floor_of(lbl), W["half_lives"]), None
    except ValueError as exc:
        if "at least min_history" not in str(exc):
            raise
        return None, str(exc)


UNTESTED = [lbl for lbl in LABELS if _bt(lbl)[0] is None]


def fold_reads(lbl):
    r_ = 5 + LABELS.index(lbl)
    bt, why = _bt(lbl)
    if bt is None:
        out = [{"sheet": LBS, "cell": "%s%d" % (L(FOLD0 + k), r_), "expect": v} for k, v in enumerate(("", "", 0))]
        out.append({"sheet": LBS, "cell": "B%d" % (4 + N + 3 + 1 + UNTESTED.index(lbl)),
                    "startswith": "%s: Not tested: %s" % (lbl, why)})
        return out
    run = bt["runs"][0]
    vals = [run["score_vs_level_ahead"], run["score_vs_change_ahead"], run["periods_tested"]] + [
        b["mean_change_ahead"] for b in run["by_score"]]
    return [{"sheet": LBS, "cell": "%s%d" % (L(FOLD0 + k), r_), "expect": "" if v is None else v}
            for k, v in enumerate(vals)]


FOLD_HEADS = [{"sheet": LBS, "cell": "%s4" % L(FOLD0 + k), "expect": h}
              for k, h in enumerate(("Spearman, value", "Spearman, change", "Quarters tested"))]
reads += FOLD_HEADS + fold_reads("Credit card - Net charge-offs - Industry (every bank)") + fold_reads(
    "Auto - Net charge-offs - KeyBank")
steps.append({"id": "B2", "what": "Read Summary, and every source's own lines on Lines by source",
              "shows": "Every series' latest value, and a dash where its rating goes: %d series, each with "
                       "its own lines on Lines by source" % N,
              "actions": [], "reads": reads, "picture": []})

# ---- B3: click Credit card, Nonaccrual; set the lines from KeyBank ----
CN = ("Credit card", "Nonaccrual")
kb = BY[CN][LEAD]
lines3 = own(kb)
others3 = every_own(CN)
steps.append({"id": "B3", "what": "On Decide, click Credit card and Nonaccrual, and pick %s under %s"
                                  % (LEAD, DT.LINES_FROM),
              "shows": "One row shows; the lines are %s's own; every source is rated on them, beside what its own "
                       "history would give" % LEAD,
              "actions": slicers(**{NAMES[0]: CN[0], NAMES[1]: CN[1]}) + [typed(CN, DT.LINES_FROM, LEAD)],
              "reads": [{"special": "visible_rows", "sheet": "Decide", "expect": [DL["first"] + PAIRS.index(CN)]},
                        {"sheet": "Decide", "cell": "B%d" % DL["headline"], "expect": headline(CN, LEAD)},
                        {"sheet": "Decide", "cell": "I%d" % DL["chip"], "expect": chip(CN, lines3)},
                        {"sheet": "Decide", "cell": "B%d" % DL["check"], "expect": OK}]
                       + [{"sheet": "Decide", "cell": "M%d" % (DL["ratings_first"] + 1 + j), "expect": lines3[c]}
                          for j, c in enumerate("EFGH")]
                       + decide_reads(CN, lines3, others3) + row_reads(CN, LEAD, lines3) + lbs_reads(kb, lines3)
                       + fits_reads(kb),
              "picture": [{"sheet": "Decide", "range": "A1:Q%d" % (DL["notes"] + 2), "file": "b3-decide.png"},
                          {"sheet": "Thresholds", "range": "A1:%s%d" % (TL["cols"]["check"], 5 + len(PAIRS)),
                           "file": "b3-thresholds.png"}]})

# ---- B4: raise How rare High is to 3.5 on that row: every source's own lines move ----
lines4 = own(kb, z=3.5)
r, c = summary_cell(kb)
steps.append({"id": "B4", "what": "Type 3.5 under How rare High is (z), in the yellow row",
              "shows": "High's line moves for every source of the row; Thresholds and Summary follow",
              "actions": [typed(CN, "How rare High is (z)", 3.5)],
              "reads": row_reads(CN, LEAD, lines4) + [lr for s, lbl in BY[CN].items()
                                                     for lr in lbs_reads(lbl, own(lbl, z=3.5))]
                       + [{"sheet": "Summary", "cell": "%s%d" % (L(c + 1), r), "expect": on_lines(kb, lines4)[1]}]
                       + fits_reads(kb, z=3.5),
              "picture": []})

# ---- B5: leave out what the evidence suggests, from the history picked ----
WINDOW = ("2012Q3", "2013Q4")                  # Decide's Suggested for KeyBank's series in the build of 14 Oct 2026
lines5 = own(kb, window=WINDOW, z=3.5)
ind5 = BY[CN][IND]
steps.append({"id": "B5", "what": "Leave out %s to %s, the period Decide suggests, with a reason" % WINDOW,
              "shows": "%s's lines move as the reference says; the other sources' own lines do not; Assess lists "
                       "the period under %s's series, its Status whole" % (LEAD, LEAD),
              "actions": [lo_typed(0, "product", CN[0]), lo_typed(0, "measure", CN[1]), lo_typed(0, "from", WINDOW[0]),
                          lo_typed(0, "to", WINDOW[1]), lo_typed(0, "why", "Suggested by the evidence: unusual stretch")],
              "reads": [{"sheet": "Decide", "cell": "D%d" % DL["evidence"], "expect": "Leave out %s to %s" % WINDOW},
                        {"sheet": "Decide", "cell": "L%d" % DL["lo_first"], "expect": "Left out: 6 quarters"},
                        {"sheet": "Decide", "cell": dcell("Periods left out", CN), "expect": 1},
                        {"sheet": "Decide", "cell": dcell("Quarters left out", CN), "expect": 6},
                        {"sheet": "Decide", "cell": "D%d" % DL["left_now"], "expect": "Left out: %s to %s, Suggested by the "
                         "evidence: unusual stretch. The chart draws these quarters in gray; the lines are set without them."
                         % WINDOW},
                        {"sheet": "Assess", "cell": "N%d" % AL["lo_top"], "expect": kb},
                        {"sheet": "Assess", "cell": "T%d" % AL["lo_top"], "expect": "Left out: 6 quarters"}]
                       + lbs_reads(kb, lines5) + lbs_reads(ind5, own(ind5, z=3.5)) + row_reads(CN, LEAD, lines5),
              "picture": [{"sheet": "Decide", "range": "A%d:Q%d" % (DL["history"], DL["lo_last"]),
                           "file": "b5-periods-left-out.png"}]})

# ---- B6: half a period refuses, in words; then put it back ----
steps.append({"id": "B6", "what": "Clear the reason", "shows": "The lines are refused, saying why and what to do",
              "actions": [lo_typed(0, "why", "")],
              "reads": [{"sheet": LBS, "cell": "%s%d" % (DT.TH_BOOK["refusal"], 5 + LABELS.index(kb)),
                         "expect": "Refused: a period left out for this product is not whole; its Status on Assess "
                                   "says what is missing"},
                        {"sheet": "Thresholds", "cell": "%s%d" % (TL["cols"]["check"], 5 + PAIRS.index(CN)),
                         "expect": "To fix: Finish its row under Periods left out on Decide (product, measure, From, "
                                   "To and Why), or clear all five of its cells."},
                        {"sheet": "Decide", "cell": dcell("Check", CN), "expect": "Refused: see the fix below"}],
              "picture": []})
steps.append({"id": "B6b", "what": "Type the reason back", "shows": "The lines are back as in B5",
              "actions": [lo_typed(0, "why", "Suggested by the evidence: unusual stretch")],
              "reads": lbs_reads(kb, lines5) + row_reads(CN, LEAD, lines5), "picture": []})

# ---- B7: the decision ----
steps.append({"id": "B7", "what": "Pick Accept these lines and write a note",
              "shows": "The count at the top; the decision on Thresholds, and on Lines by source beside the source "
                       "the lines are set from",
              "actions": [typed(CN, "Decision", "Accept these lines"),
                          typed(CN, "Note", "Walkthrough: accepted with z 3.5")],
              "reads": [{"sheet": "Decide", "cell": "L1", "expect": decided(1, 0)},
                        {"sheet": "Thresholds", "cell": "%s%d" % (TL["cols"]["decision"], 5 + PAIRS.index(CN)),
                         "expect": "Accept these lines"},
                        {"sheet": LBS, "cell": "%s%d" % (DT.TH_BOOK["decision"], 5 + LABELS.index(kb)),
                         "expect": "Accept these lines"},
                        {"sheet": LBS, "cell": "%s%d" % (DT.TH_BOOK["decision"], 5 + LABELS.index(ind5)),
                         "expect": ""}], "picture": []})

# ---- B8: KeyBank's short history rated on the industry's lines ----
AN = ("Auto", "Net charge-offs")
au, aui = BY[AN][LEAD], BY[AN][IND]
lines8 = own(aui)
r8, c8 = summary_cell(au)
steps.append({"id": "B8", "what": "Click Auto and Net charge-offs, and pick %s under %s" % (IND, DT.LINES_FROM),
              "shows": "%s's life-to-date history (%d quarters) rated on the lines the industry's longer history "
                       "gives; Summary agrees" % (LEAD, len(PTS[au])),
              "actions": slicers(**{NAMES[0]: AN[0], NAMES[1]: AN[1]}) + [typed(AN, DT.LINES_FROM, IND)],
              "reads": [{"special": "visible_rows", "sheet": "Decide", "expect": [DL["first"] + PAIRS.index(AN)]},
                        {"sheet": "Decide", "cell": "B%d" % DL["headline"], "expect": headline(AN, IND)},
                        {"sheet": "Decide", "cell": "I%d" % DL["chip"], "expect": chip(AN, lines8)},
                        {"sheet": "Summary", "cell": "%s%d" % (L(c8 + 1), r8), "expect": on_lines(au, lines8)[1]}]
                       + ([{"sheet": "Decide", "cell": "B%d" % DL["check"], "expect": OK}] if lines8["bounds"] else
                          [{"sheet": "Decide", "cell": "B%d" % DL["check"], "expect": lines8["M"]}])
                       + decide_reads(AN, lines8, every_own(AN)) + row_reads(AN, IND, lines8),
              "picture": [{"sheet": "Decide", "range": "A1:Q%d" % (DL["notes"] + 2), "file": "b8-life-to-date.png"}]})

# ---- B9: lines that are refused, and the To fix ----
MN = ("Mortgage", "Net charge-offs")
mg = BY[MN][IND]
lines9 = own(mg)
steps.append({"id": "B9", "what": "Click Mortgage and Net charge-offs, and pick %s" % IND,
              "shows": "Refused, with what to type" if lines9["M"] else "The lines stand",
              "actions": slicers(**{NAMES[0]: MN[0], NAMES[1]: MN[1]}) + [typed(MN, DT.LINES_FROM, IND)],
              "reads": ([{"sheet": "Decide", "cell": "B%d" % DL["check"], "expect": lines9["M"]},
                         {"sheet": "Decide", "cell": "I%d" % DL["check"], "startswith": "To fix: "},
                         {"sheet": "Decide", "cell": "I%d" % DL["chip"], "expect": "No rating: see the Check"},
                         {"sheet": "Thresholds", "cell": "%s%d" % (TL["cols"]["check"], 5 + PAIRS.index(MN)),
                          "startswith": "To fix: "}] if lines9["M"] else
                        [{"sheet": "Decide", "cell": "B%d" % DL["check"], "expect": OK}])
                       + row_reads(MN, IND, lines9) + lbs_reads(mg, lines9),
              "picture": [{"sheet": "Decide", "range": "A1:Q%d" % (DL["notes"] + 2), "file": "b9-refused.png"}]})

# The three rows picked so far, as Summary's counts and Needs your attention read them.
PICKED_ROWS = [(CN, LEAD, lines5), (AN, IND, lines8), (MN, IND, lines9)]


def attention():
    hot = []
    for pr, src, lines in PICKED_ROWS:
        k, text = on_lines(BY[pr][LEAD], lines)
        if k != "" and k >= 4:
            hot.append((k, PTS[BY[pr][LEAD]][-1].value, pr, text, src))
    hot.sort(key=lambda h: (-h[0], -h[1]))
    out = counts(len(PICKED_ROWS), len(hot), sum(1 for _, _, ln in PICKED_ROWS if ln["bounds"] is None))
    if not hot:
        return out + [{"sheet": "Summary", "cell": "B%d" % (ATT_HEAD + 1), "expect": SU.NONE_HOT}]
    for i, (k, v, pr, text, src) in enumerate(hot[:SU.ATTENTION_ROWS]):
        r_ = ATT_HEAD + 1 + i
        out += [{"sheet": "Summary", "cell": "B%d" % r_, "expect": pr[1]},
                {"sheet": "Summary", "cell": "C%d" % r_, "expect": v},
                {"sheet": "Summary", "cell": "D%d" % r_, "expect": text},
                {"sheet": "Summary", "cell": "E%d" % r_, "expect": "%s %s lines from %s" % (pr[0], DOT, SU.short(src))}]
    return out


# ---- B10: save, close, reopen ----
steps.append({"id": "B10", "what": "Save (Ctrl+S), close Excel, open the file again",
              "shows": "Every decision, pick and setting typed is still there, and the lines with them",
              "actions": [{"type": "reopen"}],
              "reads": [{"sheet": "Decide", "cell": dcell(DT.LINES_FROM, CN), "expect": LEAD},
                        {"sheet": "Decide", "cell": dcell("How rare High is (z)", CN), "expect": 3.5},
                        {"sheet": "Decide", "cell": lo_cell(0, "from"), "expect": WINDOW[0]},
                        {"sheet": "Decide", "cell": dcell("Decision", CN), "expect": "Accept these lines"},
                        {"sheet": "Decide", "cell": dcell("Note", CN), "expect": "Walkthrough: accepted with z 3.5"},
                        {"sheet": "Decide", "cell": "L1", "expect": decided(1, 0)},
                        {"special": "slicers", "expect": {NAMES[0]: MN[0], NAMES[1]: MN[1]}}]
                       + lbs_reads(kb, lines5) + row_reads(CN, LEAD, lines5) + attention(),
              "picture": [{"sheet": "Summary", "range": "A1:%s%d" % (L(SU_LAST + 1), ATT_HEAD + 10), "file": "b10-summary.png"}]})

# ---- B11 to B13: the peer group follows the banks in the pack set on Decide ----
# (the firm, 15 Oct 2026: "this tab needs to adjust with peers and i need a
# faster way to select and de-select peers"). The expected composites are
# summed here, in Python, from the per-bank rows the build wrote to Peers data
# and Peers delinquency data (facts per bank and quarter, values, not
# formulas), over the banks the step leaves in the pack; the workbook's live
# formulas must land on the same numbers.
from openpyxl import load_workbook                                   # noqa: E402
from threshold_engine.peers import peers_workbook as PW              # noqa: E402
from threshold_engine.peers import roster as RO                      # noqa: E402

BUILT = HERE / "out" / "consumer-credit.xlsx"
roster = RO.load(bk["peers"]["roster"])
peers_ = RO.peers(roster)
MB = bk["peers"]["min_banks"]
_wb = load_workbook(BUILT, read_only=True)
CO_ROWS = [r for r in _wb["Peers data"].iter_rows(min_row=2, values_only=True)]
DQ_ROWS = [r for r in _wb["Peers delinquency data"].iter_rows(min_row=2, values_only=True)]
LIVE = [("Credit card", "Net charge-offs", "net"), ("Other consumer", "Net charge-offs", "net"),
        ("Credit card", "90+ days past due incl. nonaccrual", "90+ incl. nonaccrual")]


def peer_label(product, measure):
    return next(lbl for lbl in LABELS if PARTS[lbl][0] == product and PARTS[lbl][1] == measure
                and PARTS[lbl][2] == PEER_SRC)


PEER_SRC = next(s_ for s_ in THIRD if s_.startswith("Peer group"))


def composite(product, kind, q, certs):
    if kind in ("net", "gross"):
        rows = [r for r in CO_ROWS if r[0] == product and r[3] == q and r[2] in certs and r[9] is True]
        if len(rows) < MB:
            return None
        g, rc, b = (sum(r[i] for r in rows) for i in (4, 5, 6))
        return 100 * ((g - rc) if kind == "net" else g) / b
    rows = [r for r in DQ_ROWS if r[0] == product and r[1] == kind and r[4] == q and r[3] in certs and r[8] is True]
    if len(rows) < MB:
        return None
    return 100 * sum(r[5] for r in rows) / sum(r[6] for r in rows)


def latest_live(product, measure, kind, certs):
    """The latest quarter the series had a value in, with a composite over
    ``certs``, and that value: what Lines by source must then show."""
    for pt in reversed(PTS[peer_label(product, measure)]):
        v = composite(product, kind, quarter(pt.date), certs)
        if v is not None:
            return quarter(pt.date), v
    return "", ""


def live_reads(certs):
    out = []
    for product, measure, kind in LIVE:
        q, v = latest_live(product, measure, kind, certs)
        r = 5 + LABELS.index(peer_label(product, measure))
        out += [{"sheet": LBS, "cell": "L%d" % r, "expect": q}, {"sheet": LBS, "cell": "M%d" % r, "expect": v}]
    return out


every = {b["cert"] for b in peers_}
regionals = {b["cert"] for b in peers_ if b["regional"]}
jpm = next(j for j, b in enumerate(peers_) if b["short"] == "JPMorgan")
_half = -(-len(peers_) // 2)


def switch_cell(j):
    """The j-th peer's switch on Decide (decide_tab.peer_switches): two lists side by side."""
    return "%s%d" % (L(DT.PEER_LISTS[j // _half][1]), DL["peers_first"] + j % _half)


def name_cell(j):
    return "%s%d" % (L(DT.PEER_LISTS[j // _half][0][0]), DL["peers_first"] + j % _half)


PACK_CELL, JPM_CELL = "D%d" % DL["peers_pick"], switch_cell(jpm)

# Round 3's pack line (the workbook name PackLine, on Peers, Decide and Signals): each peer is in
# (switched on, and a four-quarter average of 30+ incl. nonaccrual at the latest quarter, for the
# product picked on Peers), out by the book's rules (switched on, and that average blanked or the bank
# not kept), or switched off on Decide. Worked out here from Peers delinquency data's flags.
_dq_head = [c for c in next(_wb["Peers delinquency data"].iter_rows(min_row=1, max_row=1, values_only=True))]
PEERS_PRODUCT = _wb["Peers"][PW.picker_cell("product")].value
DQ_LAST = max(r[4] for r in DQ_ROWS)
_DQ_AT = {r[3]: r for r in DQ_ROWS if r[0] == PEERS_PRODUCT and r[1] == PW.DQ_PACK_MEASURE and r[4] == DQ_LAST}
PL = PW.layout(len(peers_), len(roster), 0)


def _out_by_rule(b):
    r = _DQ_AT[b["cert"]]
    return r[13] is True or r[11] is not True


# The rule each is out by, read from the data's own reason (Why no four-quarter average), and the
# first sentence of Why out on Decide said from it with the book's jump and acquisition share.
_JUMP, _SHARE = bk["peers"]["jump"], bk["peers"]["acquisition_share"]


def _rule(b):
    why = _DQ_AT[b["cert"]][12] or ""
    for prefix, rule in (("an acquisition brought in", "acquisition"), ("book moved more than the jump", "jump"),
                         ("a quarter-end rate in the window is None", "no rate"),
                         ("fewer than four consecutive", "too few quarters")):
        if why.startswith(prefix):
            return rule
    return ""


def _first_sentence(b):
    why = _DQ_AT[b["cert"]][12]
    rest = why.split(": ", 1)[1]
    q = rest.split(" ")[0]
    pct = lambda x: "%g%%" % round(100 * x, 6)                                # noqa: E731
    return {"acquisition": lambda: "An acquisition in %s brought in at least %s of its book, or a book not filed." % (q, pct(_SHARE)),
            "jump": lambda: ("Its book went from 0 to above 0 in %s." % q if rest.endswith("from 0 to above 0") else
                             "Its book moved %s in %s, more than the %s jump." % (rest.split(" ")[1], q, pct(_JUMP))),
            "no rate": lambda: "No rate in %s: %s." % (q, rest.split(" ", 1)[1][1:-1]),
            "too few quarters": lambda: "Fewer than four consecutive quarter-ends in the data."}[_rule(b)]()


def PACK_TEXT(included):
    outs = [b["short"] + (" (%s)" % _rule(b) if _rule(b) else "") for b in peers_
            if b["cert"] in included and _out_by_rule(b)]
    offs = [b["short"] for b in peers_ if b["cert"] not in included]
    t = "Peer group %s %d of %d switched on" % (DOT, len(included), len(peers_))
    if outs:
        t += " %s out by the book's rules on %s, %s: %s" % (DOT, PW.DQ_PACK_MEASURE, DQ_LAST, ", ".join(outs))
    if offs:
        t += " %s switched off on Decide: %s" % (DOT, ", ".join(offs))
    return t


def pack_reads(included):
    t = PACK_TEXT(included)
    return [{"sheet": "Peers", "cell": "B%d" % PL["pack"], "expect": t},
            {"sheet": "Peers", "cell": "B%d" % PL["pack_charts"], "expect": t},
            {"sheet": "Peers", "cell": "J%d" % PL["pack_charts"], "expect": t},
            {"sheet": "Decide", "cell": "B%d" % DL["pack_line"], "expect": t},
            {"sheet": "Signals", "cell": "B6", "expect": t}]


# The roster as built: each peer's switch on Decide under its name, reading Pack, and Peers reading Decide;
# each bank out says its rule under its name, and Why out lists them in roster order with the reason.
_outs = [b for b in peers_ if _out_by_rule(b)]
steps[0]["reads"] += [{"sheet": "Peers", "cell": PW.state_cell(j), "expect": "out: " + _rule(b) if _rule(b) else
                       PW.STATE_OUT} for j, b in enumerate(peers_) if _out_by_rule(b)] + [
    {"sheet": "Decide", "cell": "B%d" % (DL["why_first"] + k), "expect": b["short"]} for k, b in enumerate(_outs[:_half])] + [
    {"sheet": "Decide", "cell": "D%d" % (DL["why_first"] + k), "startswith": _first_sentence(b)}
    for k, b in enumerate(_outs[:_half]) if _rule(b)] + [
    {"sheet": "Decide", "cell": "B%d" % DL["why_head"], "expect": "Why out %s %s, %s, %s" % (
        DOT, PW.DQ_PACK_MEASURE, PEERS_PRODUCT.lower(), DQ_LAST)}]
steps[0]["reads"] += pack_reads(every) + [
    {"sheet": "Decide", "cell": switch_cell(j), "expect": PW.FOLLOW} for j in (jpm, len(peers_) - 1)] + [
    {"sheet": "Decide", "cell": name_cell(j), "expect": peers_[j]["short"]} for j in (jpm, len(peers_) - 1)] + [
    {"sheet": "Peers", "cell": PW.include_cell(j), "expect": PW.FOLLOW} for j in (jpm, len(peers_) - 1)] + [
    {"sheet": "Decide", "cell": PACK_CELL, "expect": PW.PACK_PICKS[0]},
    {"sheet": "Peers", "cell": PW.picker_cell("pack"), "expect": PW.PACK_PICKS[0]}]
steps.append({"id": "B11", "what": "On Decide, set Banks in the pack to Regionals only",
              "shows": "Every peer group series on Decide, Lines by source, Thresholds and Summary is summed over "
                       "the %d regionals: the latest values below are their composite" % len(regionals),
              "actions": [{"type": "cell", "sheet": "Decide", "cell": PACK_CELL, "value": PW.PACK_PICKS[1]}],
              "reads": live_reads(regionals) + pack_reads(regionals), "picture": []})
steps.append({"id": "B12", "what": "Set Banks in the pack back to Every peer, and JPMorgan's cell to No",
              "shows": "The peer group is every peer but JPMorgan",
              "actions": [{"type": "cell", "sheet": "Decide", "cell": PACK_CELL, "value": PW.PACK_PICKS[0]},
                          {"type": "cell", "sheet": "Decide", "cell": JPM_CELL, "value": PW.NO}],
              "reads": live_reads(every - {peers_[jpm]["cert"]}) + pack_reads(every - {peers_[jpm]["cert"]}),
              "picture": [{"sheet": "Peers", "range": "A1:Q%d" % (PL["pack_charts"] + 1), "file": "b12-peers.png"},
                          {"sheet": "Decide", "range": "A%d:Q%d" % (DL["pack_line"], DL["why_first"] + _half),
                           "file": "b12-decide-peers.png"}]})
steps.append({"id": "B13", "what": "Set JPMorgan's cell back to Pack",
              "shows": "The peer group is every peer again: each latest value is the series file's own",
              "actions": [{"type": "cell", "sheet": "Decide", "cell": JPM_CELL, "value": PW.FOLLOW}],
              "reads": live_reads(every) + [
                  {"sheet": LBS, "cell": "M%d" % (5 + LABELS.index(peer_label(pr, me))),
                   "expect": PTS[peer_label(pr, me)][-1].value} for pr, me, _ in LIVE] + pack_reads(every),
              "picture": []})

# Round 3: Reporting changes is the last section of Sources (no tab of its own). Its heading, the
# registry's first change and when it took effect, from the registry file the book names.
from threshold_engine import reporting_changes as RCH                # noqa: E402
_rc = RCH.load(bk["reporting_changes"])
_src_rows = list(_wb["Sources"].iter_rows(min_row=1, values_only=True))
_heads = [i for i, r in enumerate(_src_rows, start=1) if len(r) > 1 and isinstance(r[1], str)
          and r[1].startswith("Reporting changes %s " % DOT)]
assert len(_heads) == 1, _heads
RC_ROW = _heads[0]
steps[0]["reads"] += [
    {"sheet": "Sources", "cell": "B%d" % RC_ROW, "startswith": "Reporting changes %s %d rows" % (DOT, len(_rc))},
    {"sheet": "Sources", "cell": "B%d" % (RC_ROW + 2), "expect": "Change"},
    {"sheet": "Sources", "cell": "B%d" % (RC_ROW + 3), "expect": _rc[0]["name"]}]

# ---------------------------------------------------------------- B14: the Signals tab, as built
# What the tab's formulas should show with the pick as built (other consumer, 30+ incl.
# nonaccrual, KeyBank; the window all quarters; every peer in the pack): the controls, the
# where-the-dollars-are table from the series the build wrote on Signals data, the peer
# panel's states from Peers delinquency data, and two fit tables worked out here from the
# same yearly changes (CORREL over the quarters both series have, at each lead).
from threshold_engine import signals_tab as ST                       # noqa: E402
import math as _m                                                    # noqa: E402

_sd = _wb["Signals data"]
SD_HEAD = [c for c in next(_sd.iter_rows(min_row=1, max_row=1, values_only=True))]
SD_ROWS = [r for r in _sd.iter_rows(min_row=2, values_only=True) if r[0]]
GRID_Q = [r[0] for r in SD_ROWS]
SIGNAL = {k: {r[0]: r[i] for r in SD_ROWS if r[i] is not None} for i, k in enumerate(SD_HEAD) if i and k}
_sig = _wb["Signals"]
SIG_ROWS = [r for r in _sig.iter_rows(min_row=1, values_only=True)]


def _pearson(xs, ys):
    n = len(xs)
    if n < 2:
        return None
    mx, my = sum(xs) / n, sum(ys) / n
    sx = _m.sqrt(sum((x - mx) ** 2 for x in xs)); sy = _m.sqrt(sum((y - my) ** 2 for y in ys))
    return sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / (sx * sy) if sx and sy else None


def _change(series, pct=False):
    out = {}
    for i, q in enumerate(GRID_Q):
        if i >= 4 and q in series and GRID_Q[i - 4] in series:
            a, b_ = series[q], series[GRID_Q[i - 4]]
            out[q] = (100 * (a / b_ - 1) if b_ else None) if pct else a - b_
    return {q: v for q, v in out.items() if v is not None}


def _fit(sig_change, tgt_change, excl=False):
    """(best lead, r at it, r at it without 2020-23, overlapping quarters) as the tab's formulas read them."""
    n_ = len(GRID_Q)
    out_all, out_ex = {}, {}
    for k in ST.LEADS:
        xs, ys, xe, ye = [], [], [], []
        for i in range(max(0, -k), n_ - max(0, k)):
            qa, qb = GRID_Q[i], GRID_Q[i + k]
            if qa in sig_change and qb in tgt_change:
                xs.append(sig_change[qa]); ys.append(tgt_change[qb])
                if not (ST.OUT_FROM <= qa <= ST.OUT_TO) and not (ST.OUT_FROM <= qb <= ST.OUT_TO):
                    xe.append(sig_change[qa]); ye.append(tgt_change[qb])
        r_ = _pearson(xs, ys); re_ = _pearson(xe, ye)
        if r_ is not None:
            out_all[k] = r_; out_ex[k] = re_
    if not out_all:
        return None
    k = max(out_all, key=lambda k_: out_all[k_])
    # the quarters the fit at that lead rests on
    return k, out_all[k], out_ex.get(k), sum(1 for i in range(max(0, -k), n_ - max(0, k))
                                             if GRID_Q[i] in sig_change and GRID_Q[i + k] in tgt_change)


DATA_SER = {lbl: {quarter(p_.date): p_.value for p_ in PTS[lbl]} for lbl in LABELS}
PICKED_LBL = "Other consumer - 30+ days past due incl. nonaccrual - KeyBank"
TARGET_LBLS = [" - ".join(t) for t in ST.FIT_TARGETS if " - ".join(t) in DATA_SER] + [PICKED_LBL]
sig_reads = [{"sheet": "Signals", "cell": "B5", "expect": "Other consumer"},
             {"sheet": "Signals", "cell": "E5", "expect": "30+ days past due incl. nonaccrual"},
             {"sheet": "Signals", "cell": "I5", "expect": "KeyBank"},
             {"sheet": "Signals", "cell": "K5", "expect": ST.FIT_AUTO},
             {"sheet": "Signals", "cell": "M5", "expect": ST.WINDOWS[0]},
             # Picked, in words: the series is in the book
             {"sheet": "Signals", "cell": "P5", "expect": (" %s " % DOT).join(PICKED_LBL.split(" - "))},
             {"sheet": "Signals", "cell": "B6", "expect": PACK_TEXT(every)}]
# where the dollars are: KeyBank's row for each product, from Signals data, and the shares; the total row
_last = GRID_Q[-1]
assert SIG_ROWS[ST.A_TH - 1][1] == "Product" and SIG_ROWS[ST.A_TH - 1][3] == "Balance ($000)", SIG_ROWS[ST.A_TH - 1]
tots = {k: sum(SIGNAL.get("keybank_%s_%s" % (p.lower().replace(" ", "-"), k), {}).get(_last, 0) for p in ST.PRODUCTS)
        for k in ("bal", "gco4", "d30")}
for i, p_ in enumerate(ST.PRODUCTS):
    slug_ = p_.lower().replace(" ", "-")
    sig_reads.append({"sheet": "Signals", "cell": "B%d" % (ST.A_FIRST + i), "expect": p_})
    for j, k in enumerate(("bal", "gco4", "d30")):
        v = SIGNAL.get("keybank_%s_%s" % (slug_, k), {}).get(_last)
        sig_reads.append({"sheet": "Signals", "cell": "%s%d" % (L(4 + 2 * j), ST.A_FIRST + i), "expect": "" if v is None else v})
        sig_reads.append({"sheet": "Signals", "cell": "%s%d" % (L(5 + 2 * j), ST.A_FIRST + i),
                          "expect": "" if v is None or not tots[k] else v / tots[k]})
for j, k in enumerate(("bal", "gco4", "d30")):
    sig_reads.append({"sheet": "Signals", "cell": "%s%d" % (L(4 + 2 * j), ST.A_TOTAL), "expect": tots[k]})
# the Yes/No under the table: every product on the charts as built
sig_reads += [{"sheet": "Signals", "cell": "%s%d" % (L(2 + i), ST.A_ON), "expect": ST.YES} for i in range(len(ST.PRODUCTS))]
# two fit tables: the first signal's and the minimum-payment share's. Each: a header row (Against in
# column O), the pick first (merged O:Q), then the fixed series; Lead a number, Fit, Excl. 2020-23, Qtrs.
_fit_rows = [i for i, r in enumerate(SIG_ROWS, start=1) if len(r) > 14 and r[14] == "Against"]
for si, (key, title, what, kind, pkind, fmt) in enumerate(ST.SIGNALS):
    if key not in ("fico_keycorp", "minpay"):
        continue
    head_ = _fit_rows[si]
    x = SIGNAL[key] if kind == "change" else _change(SIGNAL[key], pct=(key == "pv660"))
    for j, lbl in enumerate(TARGET_LBLS):
        is_pick = j == len(TARGET_LBLS) - 1          # the pick, last here, first on the tab; the same series may be fixed too
        rr = head_ + 1 if is_pick else head_ + 2 + j
        if is_pick:
            pr_, me_, so_ = lbl.split(" - ")
            sig_reads.append({"sheet": "Signals", "cell": "O%d" % rr, "expect": (" %s " % DOT).join(
                (pr_, me_.replace(" days past due incl.", " incl."), so_.replace(" (every bank)", "")))})
        f = _fit(x, _change(DATA_SER[lbl]))
        if f is None:
            continue
        k, r_all, r_ex, n_ = f
        sig_reads += [{"sheet": "Signals", "cell": "R%d" % rr, "expect": k},
                      {"sheet": "Signals", "cell": "S%d" % rr, "expect": r_all},
                      {"sheet": "Signals", "cell": "U%d" % rr, "expect": n_}]
        if r_ex is not None:
            sig_reads.append({"sheet": "Signals", "cell": "T%d" % rr, "expect": r_ex})
# Across products: each product's industry 30-89 against the pick, worked out here as the fit tables are.
_xp_head = next(c.row for row in _sig.iter_rows(min_col=13, max_col=13) for c in row if c.value == "Product")
for _i, _p in enumerate([p_ for p_ in ST.PRODUCTS + ST.COMBINED if " - ".join((p_, ST.XP_MEASURE, ST.INDUSTRY)) in DATA_SER]):
    _rr = _xp_head + 1 + _i
    sig_reads.append({"sheet": "Signals", "cell": "M%d" % _rr, "expect": _p})
    _f = _fit(_change(DATA_SER[" - ".join((_p, ST.XP_MEASURE, ST.INDUSTRY))]), _change(DATA_SER[PICKED_LBL]))
    if _f is None:
        continue
    _k, _r_all, _r_ex, _n = _f
    sig_reads += [{"sheet": "Signals", "cell": "R%d" % _rr, "expect": _k},
                  {"sheet": "Signals", "cell": "S%d" % _rr, "expect": _r_all},
                  {"sheet": "Signals", "cell": "U%d" % _rr, "expect": _n}]
    if _r_ex is not None:
        sig_reads.append({"sheet": "Signals", "cell": "T%d" % _rr, "expect": _r_ex})

# The pipeline and the window test follow the product drawn (the pick as built, other consumer;
# the analyst, 20 Oct 2026: "show the pipeline in every product we have").
_PICK_P = _sig["B5"].value
_at = {c.value: c.coordinate for row in _sig.iter_rows() for c in row if isinstance(c.value, str)}
_pipe_cell = next(v for k, v in _at.items() if k.startswith('="The pipeline'))
sig_reads.append({"sheet": "Signals", "cell": _pipe_cell, "expect": "The pipeline %s %s: 30+ incl. nonaccrual, then 90+ incl. "
                  "nonaccrual, then charge-offs" % (DOT, _PICK_P.lower())})
_win_key = "win%d_" + re.sub(r"[^a-z0-9]+", "-", _PICK_P.lower()).strip("-")
_qw = max(q for q in SIGNAL[_win_key % 3])
_win_cell = next(v for k, v in _at.items() if '" · charged off within the window' in k)
sig_reads.append({"sheet": "Signals", "cell": _win_cell, "expect": "%s %s charged off within the window, %% of delinquent dollars, %s"
                  % (_PICK_P, DOT, _qw)})
_wr = int(re.sub(r"[A-Z]+", "", _win_cell))
for _n in (1, 2, 3):
    if _qw in SIGNAL[_win_key % _n]:
        sig_reads.append({"sheet": "Signals", "cell": "L%d" % (_wr + 1 + _n), "expect": SIGNAL[_win_key % _n][_qw]})
steps.append({"id": "B14", "what": "Open Signals",
              "shows": "The pick as built (other consumer, 30+ incl. nonaccrual, KeyBank) is in the book; where the dollars are reads "
                       "KeyBank's Call Report by product; the pack line names the peers in and out; the pipeline and the "
                       "window test are other consumer's; two fit tables agree with the same correlations worked out "
                       "from the series, lead by lead",
              "actions": [], "reads": sig_reads, "picture": [{"sheet": "Signals", "range": "A1:U110", "file": "b14-signals.png"}, {"sheet": "Signals", "range": "A%d:U%d" % (_xp_head - 3, _xp_head + 14), "file": "b14-across.png"}]})

# Saved views (views_tab): a view for each tab, called up by its number, then put away.
from threshold_engine import views_tab as VT                          # noqa: E402
_sv = _wb[VT.TITLE]
_sec = {_sv.cell(r, 2).value: r for r in range(1, 60) if _sv.cell(r, 2).value in VT.TABS}
_v1 = {t: _sec[t] + 3 for t in _sec}                                  # each tab's view 1
_SV_NOTE = {"Signals": "Card against the pack", "Peers": "Home equity, gross rate"}
_sv_typed = lambda t, col, v: {"type": "cell", "sheet": VT.TITLE, "cell": "%s%d" % (col, _v1[t]), "value": v}  # noqa: E731
_box = {t: VT.TABS[t]["box"] for t in _sec}
_note = {t: VT.TABS[t]["note"] for t in _sec}
_PICKED_NOW = "%s %s %s %s KeyBank" % (_PICK_P, DOT, _sig["E5"].value, DOT)
steps.append({"id": "B15", "what": "On Saved views, fill view 1 for Signals (credit card) and for Peers (home equity, "
                                   "gross); type 1 under Saved view on each tab",
              "shows": "Each tab follows its view: the picks the view holds replace the tab's, which are struck "
                       "through, the rest are kept, and the note shows under the box",
              "actions": [_sv_typed("Signals", "C", _SV_NOTE["Signals"]), _sv_typed("Signals", "D", GRID_Q[-1]),
                          _sv_typed("Signals", "E", "Credit card"),
                          _sv_typed("Peers", "C", _SV_NOTE["Peers"]), _sv_typed("Peers", "D", GRID_Q[-1]),
                          _sv_typed("Peers", "E", "Home equity"), _sv_typed("Peers", "F", "Gross"),
                          {"type": "cell", "sheet": "Signals", "cell": _box["Signals"], "value": 1},
                          {"type": "cell", "sheet": "Peers", "cell": _box["Peers"], "value": 1}],
              "reads": [{"sheet": "Signals", "cell": "P5", "expect": _PICKED_NOW.replace(_PICK_P, "Credit card", 1)},
                        {"sheet": "Signals", "cell": _note["Signals"], "expect": "Saved view 1: " + _SV_NOTE["Signals"]},
                        {"sheet": "Peers", "cell": _note["Peers"], "expect": "Saved view 1: " + _SV_NOTE["Peers"]},
                        {"sheet": "Signals", "cell": "B5", "expect": _PICK_P},
                        {"sheet": "Peers", "cell": "B5", "expect": "Credit card"}],
              "picture": [{"sheet": VT.TITLE, "range": "A1:N%d" % (_sec["Peers"] + 3 + VT.VIEWS), "file": "b15-saved-views.png"},
                          {"sheet": "Signals", "range": "A1:U7", "file": "b15-signals-view.png"},
                          {"sheet": "Peers", "range": "A1:Q36", "file": "b15-peers-view.png"}]})
steps.append({"id": "B16", "what": "Clear Saved view on Signals and on Peers",
              "shows": "Each tab is back on its own picks; the view stays on Saved views",
              "actions": [{"type": "cell", "sheet": "Signals", "cell": _box["Signals"], "value": ""},
                          {"type": "cell", "sheet": "Peers", "cell": _box["Peers"], "value": ""}],
              "reads": [{"sheet": "Signals", "cell": "P5", "expect": _PICKED_NOW},
                        {"sheet": "Signals", "cell": _note["Signals"], "expect": ""},
                        {"sheet": "Peers", "cell": _note["Peers"], "expect": ""},
                        {"sheet": VT.TITLE, "cell": "E%d" % _v1["Peers"], "expect": "Home equity"}]})

(HERE / "steps.json").write_text(json.dumps({"book": str(BOOK), "series": N, "rows": len(PAIRS), "steps": steps},
                                            indent=1, default=str), encoding="utf-8")
lines = ["# Expected answers: the book walkthrough", "",
         "From `%s` and the series files the build wrote, through `scale.py`. %d series, %d products and measures."
         % (BOOK.name, N, len(PAIRS)), ""]
for s in steps:
    lines += ["## %s · %s" % (s["id"], s["what"]), "", s["shows"] + ".", "",
              "%d reads." % len(s["reads"]), ""]
(HERE / "EXPECTED.md").write_text("\n".join(lines), encoding="utf-8")
print("%d steps, %d reads" % (len(steps), sum(len(s["reads"]) for s in steps)))
