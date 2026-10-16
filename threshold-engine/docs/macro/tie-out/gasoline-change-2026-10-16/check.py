"""Tie the workbook's figures to the publishers' own data. Imports nothing from threshold_engine.

    python check.py WORKBOOK.xlsx PULL_FOLDER EIA_CSV BLS_HTML SAHM_CSV FHFA_ZIP SETTINGS_JSON [OUT.json]

WORKBOOK      the economic conditions workbook the teams open (rate --workbook)
PULL_FOLDER   the dated pull it was written from (its Sources tab names the files)
EIA_CSV       EIA's own weekly history, EMM_EPMR_PTE_NUS_DPGw.xls saved as CSV
BLS_HTML      data.bls.gov/timeseries/LNS14000000, the data view
SAHM_CSV      fred.stlouisfed.org/graph/fredgraph.csv?id=SAHMCURRENT
FHFA_ZIP      FHFA's National Mortgage Database file, downloaded fresh
SETTINGS_JSON docs/macro/settings.json (the firm's cut-offs)

Every "ours" value is read out of the workbook's cells. Every "source" value is
read from the publisher's own file. The 52-week change, its percentile and its
rating are worked out again here, from EIA's file, by code written for this
check: the tool's own code is not used.
"""
import csv
import datetime
import hashlib
import html
import io
import json
import os
import re
import sys
import zipfile

import openpyxl


def sha(path):
    with open(path, "rb") as f:
        return hashlib.sha256(f.read()).hexdigest()


def ratings_row(wb, title):
    ws = wb["Ratings"]
    for row in ws.iter_rows():
        if row[1].value == title:
            return {c.column_letter: (c.value, c.number_format) for c in row}
    raise SystemExit("no row %r on the Ratings tab" % title)


def day(text):
    return datetime.datetime.strptime(text, "%m/%d/%Y").date()


def main(argv):
    wbp, pull, eia_csv, bls_html, sahm_csv, fhfa_zip, settings_json = argv[:7]
    out = {"files": {os.path.basename(p): sha(p) for p in (wbp, eia_csv, bls_html, sahm_csv, fhfa_zip, settings_json)}}
    wb = openpyxl.load_workbook(wbp)

    # -- ours: out of the workbook's cells -----------------------------------
    gas = ratings_row(wb, "Retail gasoline price, regular")
    une = ratings_row(wb, "Unemployment rate")
    gov = ratings_row(wb, "Mortgages 90+ days late, government-backed")
    sahm_cell = next(c for row in wb["Ratings"].iter_rows() for c in row
                     if isinstance(c.value, str) and "Sahm's published trigger" in c.value)
    sources_tab = {r[0].value: r[8].value for r in wb["Sources"].iter_rows() if r[0].value}
    tab = wb["Gasoline"]                     # the series' own tab: Date in N, Value in O
    by_date = {str(r[13].value)[:10]: (r[13].coordinate, r[14].value) for r in tab.iter_rows() if r[13].value}
    out["ours"] = {
        "gasoline_latest": gas["D"], "gasoline_latest_date": str(gas["C"][0])[:10],
        "gasoline_change": gas["H"], "gasoline_rating": gas["I"][0], "gasoline_percentile": gas["J"],
        "gasoline_tab_2025_10_13": by_date.get("2025-10-13"), "gasoline_tab_2026_10_12": by_date.get("2026-10-12"),
        "unemployment": une["D"], "unemployment_date": str(une["C"][0])[:10],
        "gov_90_level": gov["D"], "gov_90_change": gov["H"], "gov_90_date": str(gov["C"][0])[:10],
        "sahm_text": "%s!%s: %s" % ("Ratings", sahm_cell.coordinate, sahm_cell.value),
        "sources_tab_gasreg_sha256": sources_tab.get("Retail gasoline price, regular"),
        "run": {k: wb["Run"].cell(r, 2).value for r, k in ((4, "built"), (5, "pull"))},
    }

    # -- EIA's own weekly history ---------------------------------------------
    rows = list(csv.reader(open(eia_csv, encoding="utf-8")))
    eia = {}
    for r in rows:
        if len(r) >= 2 and re.fullmatch(r"\d\d/\d\d/\d{4}", r[0]) and r[1].strip():
            eia[day(r[0])] = float(r[1])
    weeks = sorted(eia)
    latest = weeks[-1]
    then = latest - datetime.timedelta(weeks=52)

    def a_year_before(d):
        """The reading nearest the date 52 weeks before, at most three days away; on a tie, the earlier."""
        target = d - datetime.timedelta(weeks=52)
        near = [(abs((w - target).days), w) for w in (target + datetime.timedelta(days=k) for k in range(-3, 4))
                if w in eia]
        return min(near)[1] if near else None

    changes = {}
    for d in weeks:
        b = a_year_before(d)
        if b is not None:
            changes[d] = (eia[d] / eia[b] - 1) * 100
    x = changes[latest]
    at_or_below = sum(1 for v in changes.values() if v <= x)
    pct = 100.0 * at_or_below / len(changes)
    s = json.load(open(settings_json))
    cut = s["cutoffs"]["answer"]
    on_line = s["on_the_line"]["answer"]
    rating = 1 + sum(1 for c in cut if (pct >= c if on_line == "worse" else pct > c))
    # The same again the way a reader can in a spreadsheet: EIA's sheet has a row
    # for every week, blank where no price was published, so 52 rows up is 52
    # weeks earlier.
    sheet = [(r[0], r[1].strip()) for r in rows if len(r) >= 2 and re.fullmatch(r"\d\d/\d\d/\d{4}", r[0])]
    by_row = [float(b) / float(a) - 1 for (_, a), (_, b) in zip(sheet, sheet[52:]) if a and b]
    row_pct = 100.0 * sum(1 for v in by_row if v <= by_row[-1]) / len(by_row)
    out["eia_rows"] = {"rows": len(sheet), "blank_rows": sum(1 for _, v in sheet if not v), "changes": len(by_row),
                       "percentile": row_pct, "first_row_date": sheet[0][0], "last_row_date": sheet[-1][0],
                       "weeks_reading_3.061": [d for d, v in sheet if v == "3.061"]}
    out["eia"] = {"weeks": len(weeks), "first": str(weeks[0]), "latest": str(latest), "latest_value": eia[latest],
                  "year_before": str(a_year_before(latest)), "year_before_value": eia[a_year_before(latest)],
                  "change_pct": x, "changes": len(changes), "first_change": str(min(changes)),
                  "at_or_below": at_or_below, "percentile": pct, "cutoffs": cut, "on_the_line": on_line,
                  "rating": rating, "exact_52_weeks": str(then)}

    # -- the file the workbook read, against EIA's own, week by week ----------
    fred_path = "%s/fred-GASREGW.csv" % pull
    fred = {}
    for r in list(csv.reader(open(fred_path, encoding="utf-8")))[1:]:
        if r and r[1] not in ("", "."):
            fred[datetime.date.fromisoformat(r[0])] = float(r[1])
    both = set(fred) & set(eia)
    out["fred_vs_eia"] = {"fred_file_sha256": sha(fred_path), "fred_weeks": len(fred), "eia_weeks": len(eia),
                          "in_both": len(both), "differ": sorted(str(d) for d in both if fred[d] != eia[d])[:10],
                          "only_fred": sorted(str(d) for d in set(fred) - set(eia))[:10],
                          "only_eia": sorted(str(d) for d in set(eia) - set(fred))[:10]}

    # -- BLS's own table: 2026, September --------------------------------------
    t = open(bls_html, encoding="utf-8", errors="replace").read()
    row = re.search(r'<TH scope="row">2026</TH>(.*?)</TR>', t, re.S).group(1)
    cells = [html.unescape(re.sub(r"<[^>]+>", "", c)).strip() for c in re.findall(r"<TD[^>]*>(.*?)</TD>", row, re.S)]
    out["bls"] = {"series": "LNS14000000" if "LNS14000000" in t else None, "2026_sep": cells[8]}

    # -- FRED's own Sahm indicator ---------------------------------------------
    sahm = dict(r[:2] for r in list(csv.reader(open(sahm_csv, encoding="utf-8")))[1:] if r)
    out["sahm"] = {"2026-09-01": sahm.get("2026-09-01"), "latest": max(sahm)}

    # -- FHFA's own file --------------------------------------------------------
    z = zipfile.ZipFile(fhfa_zip)
    name = next(n for n in z.namelist() if n.endswith(".csv"))
    fh = {}
    for r in csv.DictReader(io.StringIO(z.read(name).decode("utf-8-sig"))):
        if r["GEOID"] == "USA" and r["MARKET"] == "Government / Non-Conventional" and r["SERIESID"] == "P90DL":
            fh[r["PERIOD"]] = (r["VALUE1"], r["SUPPRESSED"])
    out["fhfa"] = {"csv": name, "2025Q2": fh["2025Q2"], "2026Q2": fh["2026Q2"],
                   "change_points": round(float(fh["2026Q2"][0]) - float(fh["2025Q2"][0]), 10)}

    text = json.dumps(out, indent=1, default=str)
    print(text)
    if len(argv) > 7:
        open(argv[7], "w").write(text + "\n")


if __name__ == "__main__":
    main(sys.argv[1:])
