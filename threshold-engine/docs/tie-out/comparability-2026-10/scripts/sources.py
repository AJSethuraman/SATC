"""Readers for the sources fetch_sources.py downloads, and for the tool's files.

Each reader returns numbers exactly as the source states them. Nothing here
computes a figure the tie-out reports; collect.py does that, in the open.
"""
import csv
import glob
import html
import json
import os
import re
from pathlib import Path

# ------------------------------------------------------------ Call Reports

_FACT = re.compile(r'<[A-Za-z]+:([A-Z]{4}[A-Z0-9]{4})\b[^>]*contextRef="([^"]+)"[^>]*unitRef="USD"[^>]*>(-?\d+)<')


def xbrl(path):
    """{code: value in $000} from a CDR Call Report XBRL instance.

    The instance states dollars (unitRef USD, decimals 0); the PDF pages and
    the FDIC API state thousands. Dollars are divided by 1,000, and a value
    that is not whole thousands is refused rather than rounded."""
    text = Path(path).read_text("utf-8", "replace")
    out = {}
    for code, _ctx, raw in _FACT.findall(text):
        v = int(raw)
        if v % 1000:
            raise ValueError("%s: %s = %s is not whole thousands" % (path, code, raw))
        out[code] = v // 1000
    return out


def line(facts, code4):
    """A 4-character line: consolidated (RCFD) first, then domestic (RCON),
    which is how an 041 filer (Zions) and the domestic-only 1-4 family lines
    are filed. Returns (value, full code) or (None, None)."""
    for p in ("RCFD", "RCON"):
        if p + code4 in facts:
            return facts[p + code4], p + code4
    return None, None


def filing(sources, cert, mmddyyyy):
    return xbrl(Path(sources) / ("call-%s-%s.xbrl" % (cert, mmddyyyy)))

# ------------------------------------------------------------ FDIC API


def api_rows(path):
    """The rows of a saved BankFind response, and its meta block."""
    d = json.loads(Path(path).read_text("utf-8"))
    return [x["data"] for x in d["data"]], d.get("meta", {})


def api_by_cert(path):
    rows, meta = api_rows(path)
    return {str(r["CERT"]): r for r in rows}, meta

# ------------------------------------------------------------ the tool's files


def cache(cache_dir):
    """{REPDTE: {cert: row}} from the tool's own cache (peers pull)."""
    out = {}
    for fn in sorted(glob.glob(os.path.join(cache_dir, "q*.json"))):
        out[os.path.basename(fn)[1:9]] = {str(r["CERT"]): r for r in json.load(open(fn, encoding="utf-8"))}
    return out


def series(path):
    """{'2026Q2': value} from a date,value CSV the tool wrote (peers build)."""
    out = {}
    with open(path, encoding="utf-8") as f:
        for row in csv.reader(l for l in f if not l.startswith("#")):
            if row[0] == "date":
                continue
            d = row[0]
            out["%sQ%d" % (d[:4], (int(d[5:7]) - 1) // 3 + 1)] = float(row[1])
    return out

# ------------------------------------------------------------ QBP spreadsheet

QBP_LABELS = {"30-89 days past due", "90 days or more past due", "In nonaccrual status",
              "Total outstanding", "30-89 day past due rate", "Noncurrent rate",
              "Total charge-offs", "Total recoveries", "Net charge-offs",
              "Average outstanding", "Net charge-off rate"}


def qbp_blocks(xlsx, titles):
    """{title: {'excel_row': n, 'rows': {label: {'2026Q2': v}}}} from the QBP
    time-series workbook, sheet 'Loan Performance'. A block ends at the next
    title; a label appearing twice inside one block is refused (that is how a
    first reading of this file went wrong: the next block's rows overwrote the
    GNMA rows)."""
    import openpyxl
    ws = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)["Loan Performance"]
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
    hdr = rows[4]
    cols = {hdr[j]: j for j in range(len(hdr)) if isinstance(hdr[j], str) and re.fullmatch(r"\d{4}Q[1-4]", hdr[j])}
    out = {}
    for i, r in enumerate(rows):
        t = (r[1] or "").strip() if len(r) > 1 and isinstance(r[1], str) else ""
        if t in titles and t not in out:      # first occurrence: the composition table repeats titles
            blk, j = {}, i + 1
            while j < len(rows):
                lab = (rows[j][1] or "").strip() if isinstance(rows[j][1], str) else ""
                if lab and lab not in QBP_LABELS:
                    break
                if lab:
                    if lab in blk:
                        raise ValueError("label %r twice in QBP block %r" % (lab, t))
                    blk[lab] = {q: rows[j][c] for q, c in cols.items()}
                j += 1
            out[t] = {"excel_row": i + 1, "rows": blk}
    missing = set(titles) - set(out)
    if missing:
        raise ValueError("QBP blocks not found: %s" % sorted(missing))
    return out

# ------------------------------------------------------------ SEC filings


def htm_text(path):
    """The visible text of an EDGAR inline-XBRL document, whitespace folded."""
    t = Path(path).read_text("utf-8", "replace")
    t = re.sub(r"<ix:header>.*?</ix:header>", " ", t, flags=re.S)
    t = re.sub(r"<[^>]+>", " ", t)
    t = html.unescape(t)
    return re.sub(r"\s+", " ", t)


def find(text, pattern, what):
    """One regex match, refused when absent or ambiguous."""
    m = list(re.finditer(pattern, text))
    if len(m) != 1:
        raise ValueError("%s: expected one match for %r, found %d" % (what, pattern, len(m)))
    return m[0]


def num(s):
    s = s.strip()
    if s in ("—", "-", "–"):
        return 0
    neg = s.startswith("(") and s.endswith(")")
    v = float(s.strip("()").replace(",", ""))
    return -v if neg else v
