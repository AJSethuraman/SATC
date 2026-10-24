"""The proposal's section 3 Q1 stress table, recomputed from what `peers build`
wrote, so the tool's own rules are applied (pro forma, the jump, the
acquisition share, the repeated year-to-date and push-down fixes, a blank
read as not filed), and set beside the as-filed table.

    python stress_tool_rules.py SERIES_DIR PEERS.xlsx STATS-OUTPUT.txt

SERIES_DIR and PEERS.xlsx are what `peers build` wrote; STATS-OUTPUT.txt is
the as-filed table (docs/commercial-2026-10/stats-output.txt, on main).
Nothing is fetched and nothing is computed from the Call Report here: every
rate is read from the build's output.

Per segment, measure and window (2006Q4 on, 2014Q1 on):
- n, median, spread of levels (scale.level_spread, the method's own), the
  worst quarter, and the largest z a history allows: (worst - median) / spread.
- Industry: the industry series file (every filer, no jump or minimum rule).
- KeyBank, two ways: the series file `peers build` wrote, which is the latest
  unbroken run the workbook reads (a quarter not kept breaks it), and every
  kept quarter on the peer workbook's Data and DelinquencyData tabs, gaps
  left as gaps.
"""
import csv
import statistics
import sys
from pathlib import Path

HERE = Path(__file__).resolve()
sys.path.insert(0, str(HERE.parents[2] / "src"))
from threshold_engine.scale import level_spread  # noqa: E402

SEGMENTS = (("C&I", "c-i", "C&I"), ("Nonfarm nonresidential", "nonfarm-nonresidential", "Nonfarm nonres"),
            ("Multifamily", "multifamily", "Multifamily"), ("Construction", "construction", "Construction"))
MEASURES = (("nco", "TTM net charge-off", "net", "Net rate (%)"),
            ("na", "Quarter-end nonaccrual", "delinquency-nonaccrual_quarter-end", "nonaccrual"))
WINDOWS = ("2006Q4", "2014Q1")
CERT = "17534"


def label(date):
    return "%sQ%d" % (date[:4], (int(date[5:7]) - 1) // 3 + 1)


def read_series(path):
    """(quarter, value) pairs; none when the build wrote no file (no value)."""
    if not Path(path).exists():
        return []
    with open(path, encoding="utf-8") as fh:
        rows = csv.DictReader(ln for ln in fh if not ln.startswith("#"))
        return [(label(r["date"]), float(r["value"])) for r in rows]


def stats(points, start):
    s = [(q, v) for q, v in points if q >= start]
    if len(s) < 2:
        return None
    v = [x for _, x in s]
    med, sp = statistics.median(v), level_spread(v)
    i = max(range(len(v)), key=lambda k: v[k])
    return dict(n=len(v), first=s[0][0], last=s[-1][0], median=med, spread=sp, worst=v[i], worstq=s[i][0],
                z=(v[i] - med) / sp if sp else None)


def kept_from_workbook(xlsx):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    out = {}
    rows = wb["Data"].iter_rows(values_only=True)
    head = list(next(rows))
    ix = {h: head.index(h) for h in ("Product", "Certificate", "Quarter", "Net rate (%)", "Kept")}
    for r in rows:
        if str(r[ix["Certificate"]]) == CERT and r[ix["Kept"]] is True:
            out.setdefault((r[ix["Product"]], "nco"), []).append((r[ix["Quarter"]], r[ix["Net rate (%)"]]))
    rows = wb["DelinquencyData"].iter_rows(values_only=True)
    head = list(next(rows))
    ix = {h: head.index(h) for h in ("Product", "Measure", "Certificate", "Quarter-end", "Quarter-end rate (%)",
                                     "Kept")}
    for r in rows:
        if str(r[ix["Certificate"]]) == CERT and r[ix["Measure"]] == "nonaccrual" and r[ix["Kept"]] is True:
            out.setdefault((r[ix["Product"]], "na"), []).append((r[ix["Quarter-end"]], r[ix["Quarter-end rate (%)"]]))
    return {k: sorted(v) for k, v in out.items()}


def as_filed(path):
    out = {}
    for ln in Path(path).read_text(encoding="utf-8").splitlines()[1:]:
        f = ln.split("|")
        if len(f) == 13:
            out[(f[0], f[1], f[2], f[3])] = dict(n=int(f[4]), median=float(f[6]), spread=float(f[8]),
                                                 worst=float(f[9]), worstq=f[10],
                                                 z=None if f[12] == "None" else float(f[12]))
    return out


def fmt(st, digits=3):
    if st is None:
        return "| — | — | — | — | — |"
    z = "—" if st["z"] is None else "%.1f" % st["z"]
    return "| %d | %.*f | %.*f | %.*f (%s) | %s |" % (st["n"], digits, st["median"], digits, st["spread"], digits,
                                                    st["worst"], st["worstq"], z)


def main(series_dir, xlsx, filed_path):
    series_dir = Path(series_dir)
    filed = as_filed(filed_path)
    kept = kept_from_workbook(xlsx)
    print("| Segment | Measure | Who | Window | n | Median | Spread | Worst (quarter) | Largest z | "
          "As filed: n | median | spread | worst (quarter) | largest z |")
    print("|---|---|---|---|---|---|---|---|---|---|---|---|---|---|")
    for product, slug, filed_name in SEGMENTS:
        for key, words, suffix, _ in MEASURES:
            ind = read_series(series_dir / ("%s_industry_%s.csv" % (slug, suffix)))
            own = read_series(series_dir / ("%s_keybank_%s.csv" % (slug, suffix)))
            every = kept.get((product, key), [])
            for who, pts, fname in (("Industry", ind, "Industry"), ("KeyBank, series written", own, "KeyBank"),
                                    ("KeyBank, every kept quarter", every, None)):
                for w in WINDOWS:
                    a = filed.get((fname, filed_name, key, w)) if fname else None
                    fa = ("| %d | %.3f | %.3f | %.3f (%s) | %s |" % (
                        a["n"], a["median"], a["spread"], a["worst"], a["worstq"],
                        "—" if a["z"] is None else "%.1f" % a["z"])) if a else "| | | | | |"
                    print("| %s | %s | %s | %s on " % (product, words, who, w) + fmt(stats(pts, w)) + fa[1:])
    print()
    print("Quarters held: " + "; ".join(
        "%s %s: industry %s to %s, KeyBank written %s to %s, KeyBank kept %d" % (
            p, k, *(lambda i, o, e: (i[0][0], i[-1][0], o[0][0] if o else "none", o[-1][0] if o else "none",
                                     len(e)))(
                read_series(series_dir / ("%s_industry_%s.csv" % (s, suf))),
                read_series(series_dir / ("%s_keybank_%s.csv" % (s, suf))), kept.get((p, k), [])))
        for p, s, _ in SEGMENTS for k, _, suf, _ in MEASURES))


if __name__ == "__main__":
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    main(*sys.argv[1:])
