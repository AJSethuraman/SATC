"""Write a small workbook with a table and three slicers, for the Excel check.

    python tools/slicer_sample.py OUT.xlsx

Then: powershell -File tools/check_slicers_in_excel.ps1 -Path OUT.xlsx ...
(see that script). The item names carry &, <, quotes and + on purpose: they
are the characters a hand-written XML part gets wrong.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))
from threshold_engine.slicers import add_table_slicers  # noqa: E402

HEADER = ("Product", "Measure", "Source", "Value")
ROWS = [
    ("Credit card", "Net charge-offs", "Industry", 1.5),
    ("Credit card", "Delinquency", "Industry", 2.5),
    ("Cards & loans", "Net charge-offs", "Peers <10>", 3.5),
    ('Say "hi"', "Net charge-offs", "Industry", 4.5),
    ("Auto+lease", "Delinquency", "Peers <10>", 5.5),
    ("Credit card", "Net charge-offs", "Peers <10>", 6.5),
    ("Cards & loans", "Delinquency", "Industry", 7.5),
]
#: Where the table's header sits (1-based); the slicers go above it.
HEADER_ROW = 12
SLICERS = [
    {"column": "Product", "caption": "Product", "anchor": (1, 1, 3, 10), "selected": "Credit card",
     "name": "PickProduct"},
    {"column": "Measure", "caption": "Measure & kind", "anchor": (3, 1, 5, 10),
     "selected": "Net charge-offs", "name": "PickMeasure"},
    {"column": "Source", "caption": "Source", "anchor": (5, 1, 7, 10), "selected": "Industry",
     "name": "PickSource"},
]


def build(path, rows=ROWS, slicers=SLICERS, chart=False):
    """The sample table (sheet Pick, table Picks, header row 12, from B), saved."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pick"
    ws["B1"] = "Pick a row"
    for j, h in enumerate(HEADER):
        ws.cell(HEADER_ROW, 2 + j, h)
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            ws.cell(HEADER_ROW + 1 + i, 2 + j, v)
    for col, width in zip("BCDE", (16, 18, 14, 10)):
        ws.column_dimensions[col].width = width
    last = HEADER_ROW + len(rows)
    ws.add_table(Table(displayName="Picks", ref="B%d:E%d" % (HEADER_ROW, last)))
    if chart:
        from openpyxl.chart import BarChart, Reference
        c = BarChart()
        c.add_data(Reference(ws, min_col=5, min_row=HEADER_ROW, max_row=last), titles_from_data=True)
        ws.add_chart(c, "H2")
    wb.create_sheet("Other")["A1"] = "another sheet"
    wb.save(path)
    if slicers:
        add_table_slicers(path, "Pick", "Picks", slicers)
    return path


if __name__ == "__main__":
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    print(build(sys.argv[1]))
