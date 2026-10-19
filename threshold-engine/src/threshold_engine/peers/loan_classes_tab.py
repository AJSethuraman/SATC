"""The Peer books tab (Peer loan classes until design round 3): what each
peer's consumer book holds, from its own 10-Q, so an "other consumer" rate
can be read for what is in it. A reference tab, laid out to the round 3
reference template (docs/design/round-3/handoff/README.md, "Peer books").

The file it is built from (``config/peer-loan-classes.json``) states, per
bank, which of the filing's consumer classes counts as what (student,
unsecured, secured non-auto, mixed), with the filing's own words quoted and
where they are. Amounts are never typed in it: each class names the XBRL
members the filing tags it with, and the amount is read from the cached
filing (:mod:`peers.classes`). A class the filing prints but does not tag by
class may be given as printed, with its page; the tab marks it so, and the
tie-out checks it against the page.

Three tables, top to bottom:

- **Who holds what**: one row per bank. Student loans, unsecured loans and
  lines, secured non-auto and mixed other consumer, each the sum of the
  classes counted as that (live SUMIFS over the next table), or the finding
  in words where no class is; and whether the bank is a clean comparison for
  unsecured lending, with why.
- **Consumer classes as filed**: one row per class: the filer's own label,
  $ millions, what it is counted as, and the filing. An Excel table, filtered
  by bank or by what a class is counted as.
- **In the filing's words**: the quotes behind every finding."""
import json
from pathlib import Path
from typing import List

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.filters import AutoFilter, FilterColumn
from openpyxl.worksheet.table import Table, TableStyleInfo

from .. import keybank_style as K
from . import classes as C

TITLE = "Peer books"
#: What a class may be counted as, in the order the summary shows them.
KINDS = ("Student", "Unsecured or personal", "Secured non-auto", "Mixed other consumer", "Other")
SHOWN = KINDS[:4]
FINDINGS = ("student", "unsecured", "secured_non_auto")
COMPARISON = ("Clean", "Mixed", "Not comparable")
BANK_KEYS = ("short", "cik", "holding_company", "filing_url", "classes", "findings", "comparison")
CLASS_KEYS_TAGGED = ("members", "is")
CLASS_KEYS_PRINTED = ("name", "amount_musd", "where", "is")


def load(path) -> dict:
    """The loan classes file, checked: every key present, every kind and
    answer one of the allowed, every finding quoted with where it is."""
    spec = json.loads(Path(path).read_text(encoding="utf-8"))
    problems: List[str] = []
    for k in ("quarter", "banks"):
        if k not in spec:
            problems.append("the loan classes file has no %r" % k)
    if problems:
        raise C.Refused("; ".join(problems))
    C.quarter_end(spec["quarter"])
    seen = set()
    for b in spec["banks"]:
        who = b.get("short", "?")
        for k in BANK_KEYS:
            if k not in b:
                problems.append("%s: no %r" % (who, k))
        if who in seen:
            problems.append("%s is listed twice" % who)
        seen.add(who)
        for c in b.get("classes", []):
            keys = CLASS_KEYS_TAGGED if "members" in c else CLASS_KEYS_PRINTED
            missing = [k for k in keys if k not in c]
            if missing:
                problems.append("%s: a class lacks %s" % (who, ", ".join(missing)))
            if c.get("is") not in KINDS:
                problems.append("%s: a class is counted as %r, not one of %s" % (who, c.get("is"), ", ".join(KINDS)))
        for f in FINDINGS:
            x = b.get("findings", {}).get(f)
            if not isinstance(x, dict) or not all(x.get(k) for k in ("answer", "short", "quote", "where")):
                problems.append("%s: the %s finding needs an answer, a short form, the filing's words and where "
                                "they are" % (who, f.replace("_", " ")))
            elif len(x["short"]) > 24:
                problems.append("%s: the %s finding's short form is over 24 characters" % (who, f.replace("_", " ")))
        cmp_ = b.get("comparison", {})
        if cmp_.get("answer") not in COMPARISON or not cmp_.get("why"):
            problems.append("%s: comparison needs an answer (%s) and why" % (who, ", ".join(COMPARISON)))
    if problems:
        raise C.Refused("\n".join(problems))
    return spec


def rows(spec: dict, cache) -> List[dict]:
    """Every class with its amount: from the cached filing where tagged
    (refused if the members are not found), as printed otherwise."""
    period = C.quarter_end(spec["quarter"])
    out = []
    for b in spec["banks"]:
        record, facts, labs = C.read(cache, b["cik"], period)
        for c in b["classes"]:
            if "members" in c:
                key = tuple(sorted(c["members"]))
                if key not in facts:
                    raise C.Refused("%s: the filing %s has no amortized cost at %s for %s"
                                    % (b["short"], record["accession"], period, ", ".join(key)))
                # The name as the filing prints it where given: a filer's XBRL
                # label can differ (JPMorgan tags "Auto and other" with the
                # standard "Automobile Loan" member).
                name = c.get("name") or " / ".join(C.member_label(m, labs) for m in c["members"]
                                                   if "PortfolioSegment" not in m) or C.member_label(c["members"][0], labs)
                out.append(dict(bank=b["short"], name=name, musd=facts[key] / 1e6, kind=c["is"], how="tagged",
                                where=", ".join(c["members"]), record=record))
            else:
                out.append(dict(bank=b["short"], name=c["name"], musd=float(c["amount_musd"]), kind=c["is"],
                                how="printed", where=c["where"], record=record))
    return out


#: Column widths in pixels from B (A is a 16 px margin): the README's Why
#: column (876 px) last; the rest share the two tables' columns, which a
#: column can only have one width for.
COL_PX = (16, 120, 200, 150, 150, 170, 170, 876)
FIRST = 2                                          # the first column used: B
NCOLS = 7                                          # B:H
HAIRLINE = "E7E2D8"
CLASSES_TABLE = "PeerBookClasses"
#: The classes table's columns with a filter button (the reference template:
#: filters only on Bank and Counted as).
FILTERED = ("Bank", "Counted as")
#: Characters a wrapped line holds: the Why column at Calibri 10, a class's
#: label at Calibri 11, a quote across D:G at Calibri 11.
WHY_CHARS, LABEL_CHARS, QUOTE_CHARS, WHERE_CHARS, WORDS_CHARS = 150, 25, 95, 150, 22


def _lines(text: str, per: int) -> int:
    return max(1, -(-len(text or "") // per))


def _height(lines: int) -> float:
    """A row's height for ``lines`` wrapped lines of Calibri 11: 15 pt each."""
    return 15 * lines


def _head(ws, row, heads, right=()):
    """A reference table's header: MIST, Arial 9 bold INK, wrapped, 27 pt, a STONE rule."""
    for i, h in enumerate(heads):
        c = ws.cell(row, FIRST + i, h)
        c.fill = K.solid(K.MIST)
        c.font = Font(name="Arial", size=9, bold=True, color=K.INK)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="right" if i in right else "left")
        c.border = Border(bottom=Side(style="thin", color=K.STONE))
    ws.row_dimensions[row].height = 27


def write(ws, spec: dict, cache) -> dict:
    """Fill ``ws`` and return where things are (for the tests and the tie-out)."""
    body = Font(name="Calibri", size=11, color=K.INK_TEXT)
    bold = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
    small = Font(name="Calibri", size=10, color=K.SLATE)
    slate = Font(name="Calibri", size=11, color=K.SLATE)
    code = Font(name="Consolas", size=9, color=K.SLATE)
    top = Alignment(vertical="top")
    wrap = Alignment(wrap_text=True, vertical="top")
    hair = Border(bottom=Side(style="thin", color=HAIRLINE))
    cls = rows(spec, cache)
    period = C.quarter_end(spec["quarter"])
    form = "10-K" if spec["quarter"].endswith("Q4") else "10-Q"
    last_col = FIRST + NCOLS - 1
    ws.sheet_properties.tabColor = K.TAB_REFERENCE
    c = ws.cell(1, FIRST, "%s · what each bank's consumer book holds" % TITLE)
    c.font = Font(name="Arial", size=14, bold=True, color=K.INK)
    c.alignment = Alignment(vertical="bottom")
    c = ws.cell(2, FIRST, "From each holding company's %s for the quarter ended %s, at amortized cost, $ millions. "
                          "The Call Report's other consumer line mixes these classes." % (form, period))
    c.font = Font(name="Calibri", size=11, color=K.SLATE)
    c.alignment = Alignment(vertical="center")
    for col in range(FIRST, last_col + 1):
        ws.cell(2, col).border = Border(bottom=Side(style="thin", color=K.STONE))
    ws.row_dimensions[1].height = 25.5
    ws.row_dimensions[2].height = 15
    ws.row_dimensions[3].height = 6

    banks = spec["banks"]
    s1 = 4
    t2 = s1 + 2 + len(banks) + 1                      # the classes table's section row
    first, last = t2 + 2, t2 + 1 + len(cls)
    col_of = {"bank": L(FIRST), "name": L(FIRST + 1), "musd": L(FIRST + 2), "kind": L(FIRST + 3)}
    rng = lambda key: "$%s$%d:$%s$%d" % (col_of[key], first, col_of[key], last)    # noqa: E731

    K.quiet_section(ws, s1, "Who holds what", NCOLS, FIRST)
    _head(ws, s1 + 1, ["Bank", "Student loans", "Unsecured or personal loans and lines", "Secured non-auto",
                       "Other consumer, mixed", "For an unsecured comparison", "Why"], right=(1, 2, 3, 4))
    summary = {}
    for i, b in enumerate(banks):
        r = s1 + 2 + i
        summary[b["short"]] = r
        bank = "$%s%d" % (L(FIRST), r)
        ws.cell(r, FIRST, b["short"]).font = bold
        longest = 1
        for j, kind in enumerate(SHOWN):
            # Where no class is counted as this, the finding's short form; the
            # full answer and the filing's words are in the last table.
            words = {"Student": b["findings"]["student"]["short"],
                     "Unsecured or personal": b["findings"]["unsecured"]["short"],
                     "Secured non-auto": b["findings"]["secured_non_auto"]["short"],
                     "Mixed other consumer": "None"}[kind]
            c = ws.cell(r, FIRST + 1 + j, '=IF(COUNTIFS({b},{bank},{k},"{kind}")=0,"{w}",SUMIFS({m},{b},{bank},'
                                          '{k},"{kind}"))'.format(b=rng("bank"), k=rng("kind"), m=rng("musd"),
                                                                  bank=bank, kind=kind, w=words.replace('"', '""')))
            c.number_format = "#,##0"
            c.font = body
            c.alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
            longest = max(longest, _lines(words, WORDS_CHARS))
        verdict = b["comparison"]["answer"]
        c = ws.cell(r, FIRST + 5, verdict)
        c.font = Font(name="Calibri", size=11, bold=True, color={"Not comparable": K.CRIMSON,
                                                                 "Clean": K.INTERNAL_TEXT}.get(verdict, K.INK))
        c = ws.cell(r, FIRST + 6, b["comparison"]["why"])
        c.font = small
        ws.row_dimensions[r].height = _height(max(longest, _lines(b["comparison"]["why"], WHY_CHARS)))
        for col in range(FIRST, last_col + 1):
            ws.cell(r, col).border = hair
            if not FIRST + 1 <= col <= FIRST + 4:              # the amounts are aligned above
                ws.cell(r, col).alignment = wrap if col == last_col else top
    # Words in an amount column (no class counted as that kind) read in SLATE.
    a1 = "%s%d" % (L(FIRST + 1), s1 + 2)
    ws.conditional_formatting.add("%s:%s%d" % (a1, L(FIRST + 4), s1 + 1 + len(banks)), FormulaRule(
        formula=["ISTEXT(%s)" % a1], font=Font(color=K.SLATE)))

    K.quiet_section(ws, t2, "Consumer classes as filed", NCOLS, FIRST)
    heads = ["Bank", "Class, as the filing labels it", "$ millions", "Counted as", "Filing",
             "How the amount was read", "Tagged with (XBRL members) or printed where"]
    _head(ws, t2 + 1, heads, right=(2,))
    where = {}
    for k, c in enumerate(cls):
        r = first + k
        rec = c["record"]
        vals = [c["bank"], c["name"], c["musd"], c["kind"], "%s, %s" % (rec["form"], rec["filed"]),
                "From the filing's XBRL" if c["how"] == "tagged" else "As printed, not tagged", c["where"]]
        for j, v in enumerate(vals):
            cell = ws.cell(r, FIRST + j, v)
            cell.font = body if j < 4 else (code if j == 6 and c["how"] == "tagged" else slate)
            cell.border = hair
            cell.alignment = Alignment(vertical="top", wrap_text=j in (1, 6), horizontal="right" if j == 2 else "left")
        ws.cell(r, FIRST + 2).number_format = "#,##0"
        ws.cell(r, FIRST + 4).hyperlink = rec["url"]
        ws.row_dimensions[r].height = _height(max(_lines(c["name"], LABEL_CHARS), _lines(c["where"], WHERE_CHARS)))
        where.setdefault(c["bank"], []).append((r, c))
    # The classes as an Excel table, filtered by bank or by what a class is counted as.
    tbl = Table(displayName=CLASSES_TABLE, ref="%s%d:%s%d" % (L(FIRST), t2 + 1, L(last_col), max(last, t2 + 2)))
    tbl.tableStyleInfo = TableStyleInfo(name=None, showRowStripes=False)
    tbl.autoFilter = AutoFilter(ref=tbl.ref, filterColumn=[FilterColumn(colId=i, hiddenButton=True)
                                                           for i, h in enumerate(heads) if h not in FILTERED])
    ws.add_table(tbl)

    t3 = last + 3
    K.quiet_section(ws, t3, "In the filing's words", NCOLS, FIRST)
    _head(ws, t3 + 1, ["Bank", "About", "The filing says", "", "", "", "Where"])
    r = t3 + 2
    quotes = []
    for b in banks:
        for f, about in (("student", "Student loans"), ("unsecured", "Unsecured or personal loans and lines"),
                         ("secured_non_auto", "Secured non-auto")):
            x = b["findings"][f]
            for col, v in ((FIRST, b["short"]), (FIRST + 1, about)):
                ws.cell(r, col, v).font = body
                ws.cell(r, col).alignment = wrap
            ws.merge_cells(start_row=r, start_column=FIRST + 2, end_row=r, end_column=FIRST + 5)
            said = "Not stated in the 10-Q." if x["quote"] == "not stated" else "“%s”" % x["quote"]
            text = said if x["answer"] == "None stated in the 10-Q" else "%s. %s" % (x["answer"].rstrip("."), said)
            q = ws.cell(r, FIRST + 2, text)
            q.font = body
            q.alignment = wrap
            ws.cell(r, last_col, x["where"]).font = small
            ws.cell(r, last_col).alignment = wrap
            for col in range(FIRST, last_col + 1):
                ws.cell(r, col).border = hair
            ws.row_dimensions[r].height = _height(max(_lines(text, QUOTE_CHARS), _lines(x["where"], WHERE_CHARS),
                                                      _lines(about, LABEL_CHARS)))
            quotes.append(r)
            r += 1
    ws.cell(r + 1, FIRST, "The amounts are the filings' own: each is read from the filing's XBRL by the members "
                          "named, or, where the filing does not tag a class, as printed on the page named. What each "
                          "class is counted as is a judgment, stated with the filing's words; the Glossary tab "
                          "defines each heading.").font = small
    for i, px in enumerate(COL_PX):
        ws.column_dimensions[L(1 + i)].width = round(px / 7, 2)
    ws.freeze_panes = "C4"
    K.hide_gridlines(ws)
    return {"summary": summary, "classes": where, "first": first, "last": last, "quotes": quotes,
            "table": CLASSES_TABLE}
