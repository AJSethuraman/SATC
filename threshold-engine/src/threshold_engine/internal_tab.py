"""The Internal and Internal data tabs: the bank's own figures beside the outside series.

Claude Design's round 2 hand-off (docs/design/round-2/README.md, "Internal")
as the firm then widened it (15 Oct 2026): the internal file is granular, one
row per quarter, product and segment, amounts in $ thousands, in the layout
``internal.py`` already reads (docs/INTERNAL-IMPORT.md). Nothing on these tabs
is computed in Python. Every figure is a formula over the ``InternalData``
table, so a file pasted into it on site is read the moment it lands.

    Internal data   the paste slot: the table ``InternalData`` (header at A4,
                    columns found by their header names, in any order), and
                    helper columns that read each row the way ``internal.py``
                    does and say why a row is not used
    Internal        the name mapping (products, and which workbook measure
                    each internal measure is); the coverage by product; the
                    picked series on a chart over the zones of a chosen
                    source's lines; how reliable each outside source is for
                    it; the CRR grade check; the composite and each segment
                    with balances; the like-for-like check against the Call
                    Report; and every product and measure in one table

The arithmetic follows ``internal.py`` exactly, and the tests prove the two
agree: a charge-off rate is 100 x four quarters' charge-offs over the mean
of the four quarter-end balances; a past-due rate is the amount over the
balance at the same quarter-end. The composite is the product's ``All`` rows
where the file has them, otherwise every segment summed
(``internal.composite_rows``).

A source's ratings are its own: the four lines on Lines by source, the
same On a line rule, and a value at or below zero rated 1 where Decide does
so. Nothing internal is ever shipped: the table is empty unless the book
names an internal file, which is read and checked by ``internal.read`` first.
"""
from __future__ import annotations

import re
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl.chart import AreaChart, LineChart
from openpyxl.chart.data_source import AxDataSource, NumDataSource, NumRef, StrRef
from openpyxl.chart.series import Series, SeriesLabel
from openpyxl.chart.text import Text
from openpyxl.chart.title import Title
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from . import assess_tab as AT
from . import internal as I
from . import keybank_style as K
from .peers import books

TITLE = "Internal"
DATA_TITLE = "Internal data"
TABLE = "InternalData"
DT = "'Internal data'!"
IN = "Internal!"

#: The slot's columns as the table starts out; a pasted header replaces them.
SLOT = ("quarter", "product", "segment", "balance", "gross_charge_offs", "net_charge_offs", "recoveries",
        "dpd_30_plus", "dpd_60_plus", "dpd_90_plus", "crr_grade",
        "further_1", "further_2", "further_3", "further_4", "further_5")
HDR, FIRST = 4, 5
#: Rows the slot holds; a book's internal file widens it.
CAPACITY = 2000

#: The internal measures the tab works with: (internal.MEASURES key, name).
MEASURES = (("net-charge-off-rate", "Net charge-off rate"),
            ("gross-charge-off-rate", "Gross charge-off rate"),
            ("dpd-30-plus-quarter-end", "30+ days past due"),
            ("dpd-60-plus-quarter-end", "60+ days past due"),
            ("dpd-90-plus-quarter-end", "90+ days past due"))
NOT_MAPPED = "Internal tab only"
COMPOSITE = "Composite"
N_NAMES = 8                       # product names the mapping holds (rows 5-12)
N_SEGS = 8                        # segments shown for the picked product
WINDOW = 48                       # quarters on the charts

VERDICTS = ("Too few quarters", "Its lines are refused", "Reliable for setting lines", "Use with care",
            "Not reliable")
NO_SERIES = "No outside series"

#: The helper columns on Internal data, from R.
H = dict(q="R", qn="S", prod="T", seg="U", bal="V", gco_in="W", nco_in="X", rec_in="Y", d30="Z", d60="AA",
         d90="AB", grade="AC", gco="AD", nco="AE", rec="AF", total="AG", why="AH", used="AI", name="AJ",
         name_n="AK", seg_n="AL", has="AM", has_grade="AN", latest="AO", seg_first="AP")
#: Which slot field each helper column reads, by its header name.
READS = dict(q="quarter", prod="product", seg="segment", bal="balance", gco_in="gross_charge_offs",
             nco_in="net_charge_offs", rec_in="recoveries", d30="dpd_30_plus", d60="dpd_60_plus",
             d90="dpd_90_plus", grade="crr_grade")
HELPER_HEADS = dict(q="Quarter, read as", qn="Quarter number", prod="Product, mapped", seg="Segment",
                    bal="Balance", gco_in="Gross charge-offs, as given", nco_in="Net charge-offs, as given",
                    rec_in="Recoveries, as given", d30="30+ past due", d60="60+ past due", d90="90+ past due",
                    grade="CRR grade", gco="Gross charge-offs", nco="Net charge-offs", rec="Recoveries",
                    total="Product total row", why="Why not used", used="Used", name="Product, as written",
                    name_n="Product order", seg_n="Segment order", has="Has data", has_grade="Has a grade",
                    latest="Quarter number, picked product", seg_first="First row of its segment")

#: Visible layout of the Internal tab (README: columns and rows).
WIDTHS = {"A": 16, "B": 150, "C": 190, "D": 20, "E": 150, "F": 250, "G": 20, "H": 110, "I": 110, "J": 110,
          "K": 20, "L": 170, "M": 100, "N": 100, "O": 100, "P": 100, "Q": 152}
PICK_ROW = 15
REL_TOP = 15                      # the reliability table's header row; its seven rows follow
CRR_TOP = 24
SEG_TOP = 33
F_TITLE = Font(name="Arial", size=16, bold=True, color=K.INK)
F_SUB = Font(name="Calibri", size=11, color=K.SLATE)
F_SECTION = Font(name="Arial", size=11, bold=True, color=K.INK)
F_LAB = Font(name="Arial", size=9, bold=True, color=K.SLATE)
F_HEAD = Font(name="Arial", size=9, bold=True, color=K.PAPER)
F_BODY = Font(name="Calibri", size=11, color=K.INK_TEXT)
F_BOLD = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
F_NOTE = Font(name="Calibri", size=10, color=K.SLATE)
F_HELP = Font(name="Calibri", size=10, color=K.SLATE)
EDGE = Side(style="thin", color=K.INPUT_EDGE)
BOX = Border(EDGE, EDGE, EDGE, EDGE)
RULE = Side(style="medium", color=K.INK)
HAIR = Side(style="thin", color=K.MIST)


def fill(hex_):
    return PatternFill("solid", fgColor=hex_, bgColor=hex_)


def lit(text: str) -> str:
    """An Excel string literal."""
    return '"%s"' % str(text).replace('"', '""')


def crit(text: str) -> str:
    """An exact-match criterion for COUNTIFS and SUMIFS, as a literal: the
    wildcards * ? ~ in a name are escaped so they match only themselves."""
    return lit("=" + str(text).replace("~", "~~").replace("*", "~*").replace("?", "~?"))


def crit_of(expr: str) -> str:
    """The same, for a name held in a cell."""
    return '"="&SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(%s,"~","~~"),"*","~*"),"?","~?")' % expr


def _width(px: int) -> float:
    return round((px - 5) / 7, 2)


def _input(cell, value=None, bold=True):
    if value is not None:
        cell.value = value
    cell.fill = fill(K.INPUT)
    cell.border = BOX
    cell.font = Font(name="Calibri", size=11, bold=bold, color=K.INK_TEXT)


def _section(ws, row, c1, c2, text):
    c = ws.cell(row, c1, text)
    c.font = F_SECTION
    c.alignment = Alignment(vertical="bottom")
    for col in range(c1, c2 + 1):
        ws.cell(row, col).border = Border(bottom=RULE)


def _head(ws, row, c1, labels, right_from=1):
    for i, t in enumerate(labels):
        c = ws.cell(row, c1 + i, t)
        c.font = F_HEAD
        c.fill = fill(K.INK)
        c.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="right" if i >= right_from and t else "left")


def _ln(hex_, width_pt, dash=None):
    return AT._ln(hex_, width_pt, dash)


# ------------------------------------------------------------------ discovery

_OTL = re.compile(r"((?:'[^']+'|[A-Za-z_][\w.]*)!\$?[A-Z]{1,3}\$?\d+)\s*=\s*\"worse\"")


def thresholds_refs(wb, labels: Sequence[str], parts: Dict[str, Tuple[str, str, str]]) -> Dict[str, dict]:
    """Where each series' four lines, refusal and On a line setting are.
    Thresholds' columns are read by meaning from ``decide_tab.TH_BOOK``, the
    one place the book's layout is stated; each series' row is found by its
    label in the series column, and the On a line cell is the one its rating
    formula compares with "worse". Anything not found is refused by name,
    never guessed. The per-series table is titled Thresholds while the book
    is built and Lines by source once Decide has renamed it."""
    from .decide_tab import LINES_TAB, TH_BOOK as T, _sheet_ref
    title = LINES_TAB if LINES_TAB in wb.sheetnames else "Thresholds"
    ws = wb[title]
    prefix = _sheet_ref(title)
    rows = {}
    for r in range(1, ws.max_row + 1):
        v = ws["%s%d" % (T["series"], r)].value
        if isinstance(v, str) and v in labels and v not in rows:
            rows[v] = r
    missing = [lab for lab in labels if lab not in rows]
    if missing:
        raise ValueError("Thresholds has no row for %s in its series column (%s)"
                         % (", ".join(missing[:3]), T["series"]))
    out = {}
    for lab, r in rows.items():
        rating = ws["%s%d" % (T["rating"], r)].value
        m = _OTL.search(rating) if isinstance(rating, str) else None
        if m is None:
            raise ValueError("the rating formula for %s on Thresholds %s%d does not say where its On a line "
                             "setting is" % (lab, T["rating"], r))
        ref = lambda key: "%s$%s$%d" % (prefix, T[key], r)        # noqa: E731
        out[lab] = {"lines": [ref(k) for k in ("l2", "l3", "l4", "l5")], "refused": ref("refusal"),
                    "otl": m.group(1), "row": r}
    return out


def decompose(comparison) -> Tuple[Optional[List[str]], List[str]]:
    """The book products whose Call Report lines add up to a comparison's
    (``internal.COMPARISONS``), or None and the fields no product holds."""
    bal, gross = set(books.parts(comparison.balance)), set(books.parts(comparison.gross))
    picked = [p for p, (b, g, _) in books.PRODUCTS.items()
              if set(books.parts(b)) <= bal and set(books.parts(g)) <= gross]
    got_b = [f for p in picked for f in books.parts(books.PRODUCTS[p][0])]
    got_g = [f for p in picked for f in books.parts(books.PRODUCTS[p][1])]
    if sorted(got_b) == sorted(bal) and sorted(got_g) == sorted(gross) and picked:
        return picked, []
    return None, sorted((bal - set(got_b)) | (gross - set(got_g)))


def guess_product(products: Sequence[str]) -> List[Tuple[Tuple[str, ...], str]]:
    """The README's guess rules, kept to products this workbook has."""
    rules = ((("card",), "Credit card"), (("mtg", "mortgage"), "Mortgage"), (("heloc", "home eq"), "Home equity"),
             (("auto",), "Auto"), (("oth", "other"), "Other consumer"))
    have = {p.casefold(): p for p in products}
    return [(keys, have[t.casefold()]) for keys, t in rules if t.casefold() in have]


def guess_measures(measures: Sequence[str]) -> Dict[str, str]:
    """Net and gross charge-offs map by definition, to the workbook measure
    whose name says so; the past-due buckets are never assumed to match."""
    out = {}
    for key, words in (("net-charge-off-rate", ("net", "nco")), ("gross-charge-off-rate", ("gross", "gco"))):
        hit = [m for m in measures if any(re.search(r"\b%s\b" % w, m.casefold()) for w in words)
               and ("charge" in m.casefold() or any(w in m.casefold() for w in ("nco", "gco")))]
        if hit:
            out[key] = hit[0]
    return out


# ------------------------------------------------------------------ the build

class _Cols:
    """Hidden columns, handed out left to right."""

    def __init__(self, start: int):
        self.next = start

    def take(self, n: int = 1) -> int:
        c = self.next
        self.next += n
        return c


def build(wb, *, labels: Sequence[str], parts: Dict[str, Tuple[str, str, str]], part_names: Sequence[str],
          grid: Sequence[str], not_rated: Sequence[str], higher: bool, floor_at_zero: bool,
          internal: Optional[dict] = None,
          capacity: Optional[int] = None) -> dict:
    """Add Internal and Internal data to a workbook that has Decide.

    ``internal`` is the book's internal section as ``book.load`` returns it,
    or None: ``{"rows": [internal.Row], "source": name of its series' source,
    "measures": {internal measure key: workbook measure}, "call_report":
    the source holding the bank's Call Report series, "peers_self": the
    bank's short name on Peers data, or None}``. Without it the slot is empty
    and the starting values are the guesses the README states.

    ``not_rated``: the series of rows the book shows and does not rate
    (question 23), [] for none. They have no lines, so nothing here is rated
    on them: they are left out of the outside series compared, as a series
    the workbook does not have."""
    internal = internal or {}
    rows_in = list(internal.get("rows") or [])
    # Room for the rows given and 500 more; with none, the standing size.
    cap = capacity or internal.get("capacity") or (max(CAPACITY, len(rows_in) + 500) if rows_in else CAPACITY)
    last_slot = FIRST + cap - 1
    n = len(labels)
    products = list(dict.fromkeys(parts[lab][0] for lab in labels))
    measures = list(dict.fromkeys(parts[lab][1] for lab in labels))
    own = internal.get("source")
    outside = [s for s in dict.fromkeys(parts[lab][2] for lab in labels)
               if not (own and (s == own or s.startswith(own + ":")))]
    if not outside:
        raise ValueError("the Internal tab compares the internal figures with outside sources, and the book has "
                         "none besides the internal one")
    call_report = internal.get("call_report") or next(
        (s for s in outside if not re.search(r"industry|peer", s, re.I)), outside[0])
    mapped = dict(guess_measures(measures))
    if internal.get("measures"):
        mapped = {k: v for k, v in internal["measures"].items()}
    # The calendar: the Data tab's quarters, then twelve more for internal
    # figures newer than the public ones.
    qn = [I.qindex(q) for q in grid]
    if qn != list(range(qn[0], qn[0] + len(qn))):
        raise ValueError("the Internal tab needs the Data tab's quarters to run one after another")
    G = len(grid)
    cal = list(grid) + [I._quarter_label(qn[-1] + k) for k in range(1, 13)]
    CAL = len(cal)
    refs = thresholds_refs(wb, [lab for lab in labels if lab not in not_rated], parts)
    d_last = L(1 + n)
    data_hdr = "Data!$B$1:$%s$1" % d_last
    data_col = {lab: L(2 + i) for i, lab in enumerate(labels)}

    ws = wb.create_sheet(TITLE, wb.sheetnames.index("Thresholds") + 1)
    ds = wb.create_sheet(DATA_TITLE, wb.sheetnames.index("Data") + 1 if "Data" in wb.sheetnames else len(wb._sheets))
    ws.sheet_properties.tabColor = K.TAB_WORK
    ds.sheet_properties.tabColor = K.TAB_DATA
    ws.sheet_view.showGridLines = False
    ds.sheet_view.showGridLines = False

    slot = lambda c: "%s$%s$%d:$%s$%d" % (DT, c, FIRST, c, last_slot)     # noqa: E731
    S = {k: slot(v) for k, v in H.items()}

    # ================================================================ Internal data
    ds["A1"] = "Internal data"
    ds["A1"].font = Font(name="Arial", size=14, bold=True, color=K.INK)
    ds["A2"] = ("The bank's own figures, pasted at A4 with their header row: one row per quarter, product and "
                "segment, amounts in $ thousands. Columns are found by their header names, in any order. The "
                "gray columns to the right read each row; Why not used names the first problem in a row.")
    ds["A2"].font = F_SUB
    ds.merge_cells("A2:P2")
    ds["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ds.row_dimensions[1].height = 25.5
    ds.row_dimensions[2].height = 30
    for col in range(1, 17):
        ds.cell(2, col).border = Border(bottom=Side(style="medium", color=K.KEY_RED))
    for i, h in enumerate(SLOT, start=1):
        ds.cell(HDR, i, h)
    for r_i, row in enumerate(rows_in):
        r = FIRST + r_i
        extra = dict(row.extra)
        vals = [row.quarter, row.product, row.segment, float(row.balance), float(row.gross_charge_offs),
                float(row.net_charge_offs), float(row.recoveries), float(row.dpd_30_plus), float(row.dpd_60_plus),
                float(row.dpd_90_plus)]
        grade = extra.get("crr_grade")
        vals.append(float(grade) if grade is not None else None)
        for c, v in enumerate(vals, start=1):
            if v is not None:
                ds.cell(r, c, v)
    tab = Table(displayName=TABLE, ref="A%d:P%d" % (HDR, last_slot))
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ds.add_table(tab)
    in_fill = fill(K.INPUT)
    for r in range(FIRST, last_slot + 1):
        for c in range(1, 17):
            ds.cell(r, c).fill = in_fill
    # Where each field is: its column within A:P, by header name (0: not there).
    for k, field in READS.items():
        c = ds["%s3" % H[k]]
        c.value = '=IFERROR(MATCH(%s,$A$4:$P$4,0),0)' % lit(field)
        c.font = Font(name="Calibri", size=9, color=K.SLATE)
    ds["Q3"] = "Column"
    ds["Q3"].font = Font(name="Calibri", size=9, color=K.SLATE)
    for k, h in HELPER_HEADS.items():
        c = ds["%s%d" % (H[k], HDR)]
        c.value = h
        c.font = Font(name="Arial", size=9, bold=True, color=K.SLATE)
        c.fill = fill(K.CANVAS)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ds.row_dimensions[HDR].height = 34.5
    esc = lambda x: 'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"~","~~"),"*","~*"),"?","~?")'.format(x=x)  # noqa: E731
    pick_p = "Internal!$B$%d" % PICK_ROW
    for r in range(FIRST, last_slot + 1):
        h = {k: "%s%d" % (v, r) for k, v in H.items()}
        row = "$A{r}:$P{r}".format(r=r)

        def g(k):
            return 'IF({c}$3=0,"",IF(INDEX({row},{c}$3)="","",INDEX({row},{c}$3)))'.format(c=H[k], row=row)
        f = {}
        f["has"] = "=IF(COUNTA(%s)>0,1,0)" % row
        f["q"] = ('=IF({has}=0,"",IF(ISNUMBER({g}),YEAR({g})&"Q"&ROUNDUP(MONTH({g})/3,0),'
                  'UPPER(SUBSTITUTE(TRIM({g}&"")," ",""))))').format(has=h["has"], g=g("q"))
        f["qn"] = ('=IFERROR(IF(AND(LEN({q})=6,MID({q},5,1)="Q",VALUE(RIGHT({q},1))>=1,VALUE(RIGHT({q},1))<=4),'
                   'VALUE(LEFT({q},4))*4+VALUE(RIGHT({q},1))-1,""),"")').format(q=h["q"])
        f["name"] = '=IF({has}=0,"",TRIM({g}&""))'.format(has=h["has"], g=g("prod"))
        f["prod"] = ('=IF({n}="","",IFERROR(INDEX(Internal!$C$5:$C${e},MATCH({n},Internal!$B$5:$B${e},0))&"",""))'
                     .format(n=h["name"], e=4 + N_NAMES))
        f["seg"] = '=IF({has}=0,"",TRIM({g}&""))'.format(has=h["has"], g=g("seg"))
        for k in ("bal", "gco_in", "nco_in", "rec_in", "d30", "d60", "d90", "grade"):
            f[k] = '=IF(ISNUMBER({g}),{g},IF({g}="","","not a number"))'.format(g=g(k))
        W_, X_, Y_ = h["gco_in"], h["nco_in"], h["rec_in"]
        f["gco"] = '=IF(ISNUMBER({w}),{w},IF(AND(ISNUMBER({x}),ISNUMBER({y})),{x}+{y},""))'.format(w=W_, x=X_, y=Y_)
        f["nco"] = '=IF(ISNUMBER({x}),{x},IF(AND(ISNUMBER({w}),ISNUMBER({y})),{w}-{y},""))'.format(w=W_, x=X_, y=Y_)
        f["rec"] = '=IF(ISNUMBER({y}),{y},IF(AND(ISNUMBER({w}),ISNUMBER({x})),{w}-{x},""))'.format(w=W_, x=X_, y=Y_)
        f["total"] = '=IF(EXACT({s},"All"),1,0)'.format(s=h["seg"])
        same = ('{qn},{qv},{p},"="&{ep},{s},"="&{es}'
                .format(qn=S["qn"], qv="%s", p=S["prod"], ep=esc(h["prod"]), s=S["seg"], es=esc(h["seg"])))
        dup = "COUNTIFS(%s)>1" % (same % h["qn"])
        gap = 'AND(COUNTIFS({a})=0,COUNTIFS({b})>0)'.format(a=same % ("%s-1" % h["qn"]),
                                                             b=same % ('"<"&%s' % h["qn"]))
        f["why"] = ('=IF({has}=0,"",IF({qn}="","The quarter is not written YYYYQn or as a date.",'
                    'IF({p}="","The product is not mapped on the Internal tab.",'
                    'IF({s}="","The segment is blank.",'
                    'IF(OR(NOT(ISNUMBER({bal})),NOT(ISNUMBER({d30})),NOT(ISNUMBER({d60})),NOT(ISNUMBER({d90}))),'
                    '"A balance or past-due amount is blank or not a number.",'
                    'IF(OR({w}="not a number",{x}="not a number",{y}="not a number",{gr}="not a number"),'
                    '"A charge-off amount or the CRR grade is not a number.",'
                    'IF(ISNUMBER({w})*1+ISNUMBER({x})*1+ISNUMBER({y})*1<2,'
                    '"Fewer than two of gross charge-offs, net charge-offs and recoveries are given.",'
                    'IF(IF(ISNUMBER({w})*ISNUMBER({x})*ISNUMBER({y})=1,ROUND({w}-{y}-{x},6)<>0,FALSE),'
                    '"Gross charge-offs less recoveries does not equal net charge-offs.",'
                    'IF(OR({bal}<0,{d30}<0,{d60}<0,{d90}<0,{g_}<0,{r_}<0),'
                    '"An amount other than net charge-offs is below zero.",'
                    'IF(OR({d60}>{d30},{d90}>{d60},{d30}>{bal}),'
                    '"The past-due amounts are not cumulative: 90+ is at most 60+, 60+ at most 30+, and 30+ at '
                    'most the balance.",'
                    'IF({dup},"The same quarter, product and segment appear on more than one row.",'
                    'IF({gap},"The quarter before this one is missing for the same product and segment.",'
                    '""))))))))))))').format(has=h["has"], qn=h["qn"], p=h["prod"], s=h["seg"], bal=h["bal"],
                                              d30=h["d30"], d60=h["d60"], d90=h["d90"], w=W_, x=X_, y=Y_,
                                              gr=h["grade"], g_=h["gco"], r_=h["rec"], dup=dup, gap=gap)
        f["used"] = '=IF(AND({has}=1,{why}=""),1,0)'.format(has=h["has"], why=h["why"])
        f["name_n"] = ('=IF({n}="","",IF(COUNTIF($%s$%d:%s,"="&%s)=1,MAX($%s$%d:%s%d)+1,""))'
                       % (H["name"], FIRST, h["name"], esc(h["name"]), H["name_n"], HDR, H["name_n"], r - 1)
                       ).replace("{n}", h["name"])
        f["seg_first"] = ('=IF(AND({u}=1,{t}=0),IF(COUNTIFS($%s$%d:%s,"="&%s,$%s$%d:%s,"="&%s,$%s$%d:%s,1)=1,1,0),0)'
                          % (H["prod"], FIRST, h["prod"], esc(h["prod"]), H["seg"], FIRST, h["seg"], esc(h["seg"]),
                             H["used"], FIRST, h["used"])).format(u=h["used"], t=h["total"])
        f["seg_n"] = ('=IF(AND({f}=1,{p}={pick}),MAX($%s$%d:%s%d)+1,"")' % (H["seg_n"], HDR, H["seg_n"], r - 1)
                      ).format(f=h["seg_first"], p=h["prod"], pick=pick_p)
        f["has_grade"] = '=IF(AND({u}=1,ISNUMBER({g})),1,0)'.format(u=h["used"], g=h["grade"])
        f["latest"] = '=IF(AND({u}=1,{p}={pick}),{qn},"")'.format(u=h["used"], p=h["prod"], pick=pick_p, qn=h["qn"])
        for k, v in f.items():
            c = ds["%s%d" % (H[k], r)]
            c.value = v
            c.font = F_HELP
    ds.freeze_panes = "A5"
    for col in range(1, 17):
        ds.column_dimensions[L(col)].width = 13
    ds.column_dimensions["Q"].width = 7
    for k, v in H.items():
        ds.column_dimensions[v].width = 34 if k == "why" else 12
    for c in (H["qn"], H["total"], H["has"], H["name_n"], H["seg_n"], H["has_grade"], H["latest"], H["gco_in"],
              H["nco_in"], H["rec_in"], H["seg_first"]):
        ds.column_dimensions[c].hidden = True
    n_below = 'COUNTA(%s$A$%d:$P$%d)' % (DT, last_slot + 1, last_slot + 5000)

    # ================================================================ Internal: visible
    for col, px in WIDTHS.items():
        ws.column_dimensions[col].width = _width(px)
    heights = {1: 25.5, 2: 15, 3: 6, 4: 16.5, 5: 15, 13: 7.5, 32: 7.5}
    for r in range(6, 13):
        heights[r] = 16.5
    for r in range(14, 32):
        heights[r] = 16.5
    for r, hgt in heights.items():
        ws.row_dimensions[r].height = hgt
    ws["B1"] = "Internal: the bank's own figures beside the Call Report, peers and industry"
    ws["B1"].font = F_TITLE
    ws["B1"].alignment = Alignment(vertical="bottom")
    ws["B2"] = ("Built from the rows on Internal data. The yellow cells under What counts as reliable are "
                "starting values to set.")
    ws["B2"].font = F_SUB
    for col in range(2, 18):
        ws.cell(2, col).border = Border(bottom=Side(style="medium", color=K.KEY_RED))

    hc = _Cols(max(27, 13 + len(products) + 3))
    # Lists the dropdowns read, in hidden columns.
    lists = {}

    def a_list(key, items, head):
        col = hc.take()
        ws.cell(1, col, head)
        for i, v in enumerate(items):
            ws.cell(2 + i, col, v)
        lists[key] = "$%s$2:$%s$%d" % (L(col), L(col), 1 + max(1, len(items)))
        return col
    a_list("products", products, "Products")
    a_list("measures", measures + [NOT_MAPPED], "Measures")
    a_list("sources", outside, "Sources")
    a_list("imeasures", [m for _, m in MEASURES], "Internal measures")
    seg_list_col = hc.take()
    ws.cell(1, seg_list_col, "Segments")
    ws.cell(2, seg_list_col, COMPOSITE)
    for k in range(N_SEGS):
        ws.cell(3 + k, seg_list_col, '=IFERROR(INDEX(%s,MATCH(%d,%s,0)),"")' % (S["seg"], k + 1, S["seg_n"]))
    lists["segments"] = "$%s$2:$%s$%d" % (L(seg_list_col), L(seg_list_col), 2 + N_SEGS)
    SEG_NAME = lambda k: "$%s$%d" % (L(seg_list_col), 3 + k)                # noqa: E731

    def dv(ref, key, title):
        v = DataValidation(type="list", formula1="=%s" % lists[key], showErrorMessage=True, errorTitle=title,
                           error="Pick one from the list.")
        ws.add_data_validation(v)
        v.add(ref)

    # ---- 1. Your products are (B4:C12)
    _section(ws, 4, 2, 3, "Your products are")
    rules = guess_product(products)
    for i in range(N_NAMES):
        r = 5 + i
        ws.cell(r, 2, '=IFERROR(INDEX(%s,MATCH(%d,%s,0)),"")' % (S["name"], i + 1, S["name_n"])).font = F_BODY
        guess = '""'
        for keys, target in reversed(rules):
            guess = 'IF(OR(%s),%s,%s)' % (",".join('ISNUMBER(SEARCH(%s,B%d))' % (lit(k), r) for k in keys),
                                          lit(target), guess)
        _input(ws.cell(r, 3), '=IF(B{r}="","",IFERROR(INDEX({p},MATCH(B{r},{p},0)),{g}))'
               .format(r=r, p=lists["products"], g=guess))
    dv("C5:C%d" % (4 + N_NAMES), "products", "Product")

    # ---- 2. Your measures are (E4:F12)
    _section(ws, 4, 5, 6, "Your measures are")
    for i, (key, name) in enumerate(MEASURES):
        r = 5 + i
        ws.cell(r, 5, name).font = F_BODY
        _input(ws.cell(r, 6), mapped.get(key, NOT_MAPPED))
    dv("F5:F%d" % (4 + len(MEASURES)), "measures", "Workbook measure")
    M_NAMES = "$E$5:$E$%d" % (4 + len(MEASURES))
    M_MAP = "$F$5:$F$%d" % (4 + len(MEASURES))
    ws.cell(11, 5, "CRR grade column").font = F_LAB
    ws.cell(11, 6, '=IF(%s$AC$3=0,"Not in the file",IF(SUM(%s)=0,"In the table, with no grades",'
                   'SUM(%s)&" rows have a grade"))' % (DT, S["has_grade"], S["has_grade"])).font = F_BODY

    # ---- 3. What counts as reliable (H4:J12)
    _section(ws, 4, 8, 10, "What counts as reliable")
    for r, lab, v, fmt in ((5, "Quarters both have, at least", 12, "0"),
                           (6, "Ratings that agree, at least", 0.7, "0%"),
                           (7, "Change correlation, at least", 0.6, "0.00")):
        ws.cell(r, 8, lab).font = F_LAB
        _input(ws.cell(r, 10), v)
        ws.cell(r, 10).number_format = fmt
    MIN_Q, MIN_A, MIN_C = "$J$5", "$J$6", "$J$7"
    ws.cell(9, 8, "Rows on Internal data").font = F_LAB
    ws.cell(9, 10, "=SUM(%s)+%s" % (S["has"], n_below)).font = F_BODY
    ws.cell(10, 8, "Rows not used").font = F_LAB
    ws.cell(10, 10, "=SUM(%s)-SUM(%s)+%s" % (S["has"], S["used"], n_below)).font = F_BOLD
    ws.conditional_formatting.add("J10", FormulaRule(formula=["$J$10>0"], fill=fill(K.CANVAS),
                                                     font=Font(bold=True, color=K.CRIMSON)))
    ws.cell(11, 8, "Call Report source").font = F_LAB
    _input(ws.cell(11, 10), call_report)
    dv("J11", "sources", "Call Report source")
    ws.cell(12, 8, "Rows the slot holds").font = F_LAB
    ws.cell(12, 10, cap).font = F_BODY
    for r in (9, 10, 12):
        ws.cell(r, 10).number_format = "#,##0"
    CR = "$J$11"

    # ---- hidden: calendar and one block per product (composite)
    W0 = 5
    WE = W0 + CAL - 1
    c_q, c_qn = hc.take(), hc.take()
    ws.cell(4, c_q, "Quarter")
    ws.cell(4, c_qn, "Quarter number")
    for i, q in enumerate(cal):
        ws.cell(W0 + i, c_q, q)
        ws.cell(W0 + i, c_qn, I.qindex(q))
    QN = lambda r: "$%s%d" % (L(c_qn), r)                                 # noqa: E731
    CALQ = "$%s$%d:$%s$%d" % (L(c_q), W0, L(c_q), WE)
    CALN = "$%s$%d:$%s$%d" % (L(c_qn), W0, L(c_qn), WE)
    FIELDS = ("bal", "gco", "nco", "d30", "d60", "d90")
    pblock = {}
    for pi, p in enumerate(products):
        b = {k: hc.take() for k in ("cnt", "hq") + FIELDS + ("r_net", "r_gross", "r_d30", "r_d60", "r_d90", "at")}
        pblock[p] = b
        ws.cell(1, b["cnt"], p)
        hasall = "$%s$2" % L(b["cnt"])
        bad = "$%s$3" % L(b["cnt"])
        # hq, per quarter: 1 when the file gives the product's All row for that
        # quarter (used there), else 0 and the segments are summed. hasall
        # reads them back: 0 no All row, 1 All in every quarter with rows, 2
        # All in some (Codex on #442: one global choice dropped the quarters
        # that had segments only).
        HQ = "$%s$%d:$%s$%d" % (L(b["hq"]), W0, L(b["hq"]), WE)
        CNT = "$%s$%d:$%s$%d" % (L(b["cnt"]), W0, L(b["cnt"]), WE)
        ws[hasall[1:].replace("$", "")] = ('=IF(SUM({hq})=0,0,IF(SUM({hq})=COUNTIF({cnt},">0"),1,2))'
                                           .format(hq=HQ, cnt=CNT))
        ws[bad[1:].replace("$", "")] = '=COUNTIFS({p},{c},{u},0)'.format(p=S["prod"], c=crit(p), u=S["used"])
        b["hasall"], b["bad"] = hasall, bad
        for k in b:
            if isinstance(b[k], int):
                ws.cell(4, b[k], k)
        for i in range(CAL):
            r = W0 + i
            ws.cell(r, b["hq"], '=IF(COUNTIFS({qn},{q},{p},{c},{u},1,{t},1)>0,1,0)'
                    .format(qn=S["qn"], q=QN(r), p=S["prod"], c=crit(p), u=S["used"], t=S["total"]))
            base = '{p},{c},{u},1,{t},{h}'.format(p=S["prod"], c=crit(p), u=S["used"], t=S["total"],
                                                  h="%s%d" % (L(b["hq"]), r))
            ws.cell(r, b["cnt"], "=COUNTIFS(%s,%s,%s)" % (S["qn"], QN(r), base))
            ws.cell(r, b["at"], '=IF(%s%d>0,%d,"")' % (L(b["cnt"]), r, i + 1))
            for k in FIELDS:
                ws.cell(r, b[k], "=SUMIFS(%s,%s,%s,%s)" % (S[k], S["qn"], QN(r), base))
            cn = lambda k_, rr: "%s%d" % (L(b[k_]), rr)                     # noqa: E731
            for rk, fk in (("r_net", "nco"), ("r_gross", "gco")):
                if i < 3:
                    ws.cell(r, b[rk], '=""')
                    continue
                win = lambda k_: "SUM(%s:%s)" % (cn(k_, r - 3), cn(k_, r))  # noqa: E731
                ws.cell(r, b[rk], '=IF({bad}>0,"",IF(AND({c0}>0,{c1}>0,{c2}>0,{c3}>0),IF({sb}>0,100*{sn}/({sb}/4),'
                                  '""),""))'.format(bad=bad, c0=cn("cnt", r), c1=cn("cnt", r - 1),
                                                    c2=cn("cnt", r - 2), c3=cn("cnt", r - 3), sb=win("bal"),
                                                    sn=win(fk)))
            for rk, fk in (("r_d30", "d30"), ("r_d60", "d60"), ("r_d90", "d90")):
                ws.cell(r, b[rk], '=IF(OR({bad}>0,{c}=0),"",IF({b}>0,100*{d}/{b},""))'
                        .format(bad=bad, c=cn("cnt", r), b=cn("bal", r), d=cn(fk, r)))
    RATE_KEYS = ("r_net", "r_gross", "r_d30", "r_d60", "r_d90")

    # ---- hidden: (product, workbook measure) blocks, aligned with Data
    out_labels = [lab for lab in labels if parts[lab][2] in outside and lab not in not_rated]
    pw = {}
    for lab in out_labels:
        p, w, _ = parts[lab]
        if (p, w) in pw:
            continue
        v, fl = hc.take(), hc.take()
        pw[(p, w)] = (v, fl)
        kcell = "$%s$3" % L(v)
        ws.cell(1, v, "%s - %s" % (p, w))
        ws.cell(3, v, '=IFERROR(MATCH(%s,%s,0),0)' % (lit(w), M_MAP))
        b = pblock[p]
        for i in range(G):
            r = W0 + i
            ch = "CHOOSE(%s,%s)" % (kcell, ",".join("%s%d" % (L(b[k]), r) for k in RATE_KEYS))
            ws.cell(r, v, "=IF(%s=0,0,IF(ISNUMBER(%s),%s,0))" % (kcell, ch, ch))
            ws.cell(r, fl, "=IF(%s=0,0,IF(ISNUMBER(%s),1,0))" % (kcell, ch))

    def rng(col, a, b_):
        return "$%s$%d:$%s$%d" % (L(col), a, L(col), b_)

    def rating(v, lines, wc):
        op = ">" if higher else "<"
        terms = "+".join("(({v}{op}{l})+({v}={l})*{w})".format(v=v, op=op, l=l_, w=wc) for l_ in lines)
        if floor_at_zero and higher:
            return "(1+({v}>0)*({t}))".format(v=v, t=terms)
        return "(1+%s)" % terms

    def stats(row, c0, I_, F_, S_, SF_, Ic, Ip, Fc, Fp, Sc, Sp, SFc, SFp, lines, wc, refused, found=None):
        """One row of statistics from column c0: sums, then the seven
        figures and the verdict. Returns the cells of each figure."""
        M = "(%s*%s)" % (F_, SF_)
        Mc = "(%s*%s*%s*%s)" % (Fc, Fp, SFc, SFp)
        dx, dy = "(%s-%s)" % (Ic, Ip), "(%s-%s)" % (Sc, Sp)
        sums = [("n", "SUMPRODUCT(%s*1)" % M), ("sx", "SUMPRODUCT(%s*%s)" % (M, I_)),
                ("sy", "SUMPRODUCT(%s*%s)" % (M, S_)), ("sxx", "SUMPRODUCT(%s*%s*%s)" % (M, I_, I_)),
                ("syy", "SUMPRODUCT(%s*%s*%s)" % (M, S_, S_)), ("sxy", "SUMPRODUCT(%s*%s*%s)" % (M, I_, S_)),
                ("dn", "SUMPRODUCT(%s*1)" % Mc), ("dsx", "SUMPRODUCT(%s*%s)" % (Mc, dx)),
                ("dsy", "SUMPRODUCT(%s*%s)" % (Mc, dy)), ("dsxx", "SUMPRODUCT(%s*%s*%s)" % (Mc, dx, dx)),
                ("dsyy", "SUMPRODUCT(%s*%s*%s)" % (Mc, dy, dy)), ("dsxy", "SUMPRODUCT(%s*%s*%s)" % (Mc, dx, dy)),
                ("agn", "IFERROR(SUMPRODUCT(%s*(%s=%s)),\"\")" % (M, rating(I_, lines, wc), rating(S_, lines, wc))),
                ("win", "IFERROR(SUMPRODUCT(%s*(ABS(%s-%s)<=1)),\"\")"
                 % (M, rating(I_, lines, wc), rating(S_, lines, wc)))]
        at = {}
        c = c0
        for k, fm in sums:
            ws.cell(row, c, "=" + fm)
            at[k] = "$%s$%d" % (L(c), row)
            c += 1

        def corr(n_, sx, sy, sxx, syy, sxy):
            return ('IFERROR(({n}*{sxy}-{sx}*{sy})/SQRT(({n}*{sxx}-{sx}^2)*({n}*{syy}-{sy}^2)),"")'
                    .format(n=at[n_], sx=at[sx], sy=at[sy], sxx=at[sxx], syy=at[syy], sxy=at[sxy]))
        figs = [("overlap", "%s" % at["n"]), ("level", corr("n", "sx", "sy", "sxx", "syy", "sxy")),
                ("change", corr("dn", "dsx", "dsy", "dsxx", "dsyy", "dsxy")),
                ("gap", 'IF({n}>0,({sx}-{sy})/{n},"")'.format(**at)),
                # No rating is given on lines that are refused, so neither is agreement.
                ("agree", 'IF(AND({n}>0,{rf}<>1,ISNUMBER({agn})),{agn}/{n},"")'.format(rf=refused, **at)),
                ("within", 'IF(AND({n}>0,{rf}<>1,ISNUMBER({win})),{win}/{n},"")'.format(rf=refused, **at))]
        for k, fm in figs:
            ws.cell(row, c, "=" + fm)
            at[k] = "$%s$%d" % (L(c), row)
            c += 1
        verdict = ('IF({n}<{mq},"{v0}",IF({ref}=1,"{v1}",IF(AND(ISNUMBER({a}),ISNUMBER({ch})),'
                   'IF(AND({a}>={ma},{ch}>={mc}),"{v2}",IF({a}>=0.75*{ma},"{v3}","{v4}")),'
                   'IF(AND(ISNUMBER({a}),{a}>=0.75*{ma}),"{v3}","{v4}"))))'
                   .format(n=at["n"], mq=MIN_Q, ref=refused, a=at["agree"], ch=at["change"], ma=MIN_A, mc=MIN_C,
                           v0=VERDICTS[0], v1=VERDICTS[1], v2=VERDICTS[2], v3=VERDICTS[3], v4=VERDICTS[4]))
        if found is not None:
            verdict = 'IF(%s,%s,"%s")' % (found, verdict, NO_SERIES)
        ws.cell(row, c, "=" + verdict)
        at["verdict"] = "$%s$%d" % (L(c), row)
        c += 1
        # Ranks the sources for the best one: agreement, where the verdict allows it.
        ws.cell(row, c, '=IF(OR({v}="{a}",{v}="{b}",{v}="{c}"),{ag},-1)'.format(
            v=at["verdict"], a=VERDICTS[2], b=VERDICTS[3], c=VERDICTS[4], ag=at["agree"]))
        at["score"] = "$%s$%d" % (L(c), row)
        return at

    # ---- hidden: one row of statistics per outside series
    st_lab = hc.take()
    st_lines = [hc.take() for _ in range(4)]
    st_w, st_ref = hc.take(), hc.take()
    st0 = hc.take(22)
    ws.cell(4, st_lab, "Series")
    for k in range(4):
        ws.cell(4, st_lines[k], "Line %d" % (k + 2))
    ws.cell(4, st_w, "On a line is worse")
    ws.cell(4, st_ref, "Refused")
    st_at = {}
    STAT0 = W0
    I_rows = (W0, W0 + G - 1)
    for j, lab in enumerate(out_labels):
        r = STAT0 + j
        p, w, s = parts[lab]
        rf = refs[lab]
        ws.cell(r, st_lab, lab)
        for k in range(4):
            ws.cell(r, st_lines[k], "=%s" % rf["lines"][k])
        ws.cell(r, st_w, '=IF(%s="worse",1,0)' % rf["otl"])
        ws.cell(r, st_ref, '=IF(%s<>"",1,0)' % rf["refused"])
        v, fl = pw[(p, w)]
        dc = data_col[lab]
        Srng = "Data!$%s$2:$%s$%d" % (dc, dc, G + 1)
        st_at[lab] = stats(
            r, st0, rng(v, *I_rows), rng(fl, *I_rows), Srng, "ISNUMBER(%s)" % Srng,
            rng(v, W0 + 4, W0 + G - 1), rng(v, W0, W0 + G - 5), rng(fl, W0 + 4, W0 + G - 1), rng(fl, W0, W0 + G - 5),
            "Data!$%s$6:$%s$%d" % (dc, dc, G + 1), "Data!$%s$2:$%s$%d" % (dc, dc, G - 3),
            "ISNUMBER(Data!$%s$6:$%s$%d)" % (dc, dc, G + 1), "ISNUMBER(Data!$%s$2:$%s$%d)" % (dc, dc, G - 3),
            ["$%s$%d" % (L(c_), r) for c_ in st_lines], "$%s$%d" % (L(st_w), r), "$%s$%d" % (L(st_ref), r))
    ST_END = STAT0 + max(1, len(out_labels)) - 1
    ST_LABELS = rng(st_lab, STAT0, ST_END)
    def st_range(key):
        col = re.match(r"\$([A-Z]+)\$", st_at[out_labels[0]][key]).group(1)
        return "$%s$%d:$%s$%d" % (col, STAT0, col, ST_END)

    # ---- the pickers (row 15)
    for col, t in ((2, "Product"), (3, "Measure"), (5, "Segment"), (6, "Shade with lines from")):
        ws.cell(14, col, t).font = F_LAB
        ws.cell(14, col).alignment = Alignment(vertical="bottom")
    _input(ws.cell(PICK_ROW, 2), products[0])
    _input(ws.cell(PICK_ROW, 3), MEASURES[0][1])
    _input(ws.cell(PICK_ROW, 5), COMPOSITE)
    _input(ws.cell(PICK_ROW, 6), call_report)
    dv("B%d" % PICK_ROW, "products", "Product")
    dv("C%d" % PICK_ROW, "imeasures", "Measure")
    dv("E%d" % PICK_ROW, "segments", "Segment")
    dv("F%d" % PICK_ROW, "sources", "Shade with lines from")
    PB, PC, PE, PF = ("$%s$%d" % (c_, PICK_ROW) for c_ in "BCEF")

    # ---- hidden: the picked series, every calendar quarter
    pk = {k: hc.take() for k in ("cnt",) + FIELDS + ("rate", "val", "flag", "last", "grade", "gcount")}
    for k, c_ in pk.items():
        ws.cell(4, c_, "picked " + k)
    c0 = pk["cnt"]
    KM = "$%s$1" % L(c0)
    ws.cell(1, c0, "=IFERROR(MATCH(%s,%s,0),0)" % (PC, M_NAMES))
    HASALL = "$%s$2" % L(c0)
    ws.cell(2, c0, '=IF(COUNTIFS(%s,%s,%s,1,%s,1)>0,1,0)' % (S["prod"], crit_of(PB), S["total"], S["used"]))
    BAD = "$%s$3" % L(c0)
    ws.cell(3, c0, '=COUNTIFS(%s,%s,%s,0)' % (S["prod"], crit_of(PB), S["used"]))
    COMP = '(%s="%s")' % (PE, COMPOSITE)
    WPK = "$%s$1" % L(pk["bal"])                                    # the workbook measure the pick maps to
    ws.cell(1, pk["bal"], '=IF(%s=0,"",INDEX(%s,%s))' % (KM, M_MAP, KM))
    base_c = '{p},{c},{u},1,{t},{h}'.format(p=S["prod"], c=crit_of(PB), u=S["used"], t=S["total"], h=HASALL)
    base_s = '{p},{c},{u},1,{s},{e}'.format(p=S["prod"], c=crit_of(PB), u=S["used"], s=S["seg"], e=crit_of(PE))
    for i in range(CAL):
        r = W0 + i
        ws.cell(r, pk["cnt"], "=IF(%s,COUNTIFS(%s,%s,%s),COUNTIFS(%s,%s,%s))"
                % (COMP, S["qn"], QN(r), base_c, S["qn"], QN(r), base_s))
        for k in FIELDS:
            ws.cell(r, pk[k], "=IF(%s,SUMIFS(%s,%s,%s,%s),SUMIFS(%s,%s,%s,%s))"
                    % (COMP, S[k], S["qn"], QN(r), base_c, S[k], S["qn"], QN(r), base_s))
        cn = lambda k_, rr: "%s%d" % (L(pk[k_]), rr)                       # noqa: E731
        if i >= 3:
            have = "AND(%s>0,%s>0,%s>0,%s>0)" % tuple(cn("cnt", r - k) for k in range(4))
            win = lambda k_: "SUM(%s:%s)" % (cn(k_, r - 3), cn(k_, r))      # noqa: E731
            ttm = lambda k_: 'IF(%s,IF(%s>0,100*%s/(%s/4),""),"")' % (have, win("bal"), win(k_), win("bal"))  # noqa
        else:
            ttm = lambda k_: '""'                                           # noqa: E731
        qe = lambda k_: 'IF(AND(%s>0,%s>0),100*%s/%s,"")' % (cn("cnt", r), cn("bal", r), cn(k_, r), cn("bal", r))  # noqa
        ws.cell(r, pk["rate"], '=IF(OR(%s>0,%s=0),"",CHOOSE(%s,%s,%s,%s,%s,%s))'
                % (BAD, KM, KM, ttm("nco"), ttm("gco"), qe("d30"), qe("d60"), qe("d90")))
        ws.cell(r, pk["val"], '=IF(ISNUMBER(%s),%s,0)' % (cn("rate", r), cn("rate", r)))
        ws.cell(r, pk["flag"], '=IF(ISNUMBER(%s),1,0)' % cn("rate", r))
        ws.cell(r, pk["last"], '=IF(%s=1,%d,%s)' % (cn("flag", r), i + 1, cn("last", r - 1) if i else 0))
        gq = "%s,%s,%s,1" % (S["qn"], QN(r), S["has_grade"])
        ws.cell(r, pk["gcount"], "=IF(%s,IF(%s=1,COUNTIFS(%s,%s,%s,1),0),COUNTIFS(%s,%s,%s))"
                % (COMP, HASALL, gq, S["prod"] + "," + crit_of(PB), S["total"], gq, S["prod"] + "," + crit_of(PB),
                   S["seg"] + "," + crit_of(PE)))
        ws.cell(r, pk["grade"], '=IF(%s=1,IF(%s,SUMIFS(%s,%s,%s,%s,1),SUMIFS(%s,%s,%s,%s)),"")'
                % (cn("gcount", r), COMP, S["grade"], gq, S["prod"] + "," + crit_of(PB), S["total"],
                   S["grade"], gq, S["prod"] + "," + crit_of(PB), S["seg"] + "," + crit_of(PE)))
    LAST = "$%s$%d" % (L(pk["last"]), WE)

    # ---- hidden: each outside source for the pick, and its statistics
    psrc = {}
    for k, s in enumerate(outside):
        v, fl = hc.take(), hc.take()
        psrc[s] = (v, fl)
        lab_c, col_c, idx_c = "$%s$1" % L(v), "$%s$2" % L(v), "$%s$3" % L(v)
        ws.cell(1, v, '=%s&" - "&%s&" - "&%s' % (PB, WPK, lit(s)))
        ws.cell(2, v, '=IFERROR(MATCH(%s,%s,0),0)' % (lab_c, data_hdr))
        ws.cell(3, v, '=IFERROR(MATCH(%s,%s,0),0)' % (lab_c, ST_LABELS))
        ws.cell(4, v, s)
        for i in range(CAL):
            r = W0 + i
            if i < G:
                cell = "INDEX(Data!$B${d}:${e}${d},{c})".format(d=2 + i, e=d_last, c=col_c)
                ws.cell(r, v, '=IF({c}=0,0,IF({x}="",0,{x}))'.format(c=col_c, x=cell))
                ws.cell(r, fl, '=IF({c}=0,0,IF({x}="",0,1))'.format(c=col_c, x=cell))
            else:
                ws.cell(r, v, 0)
                ws.cell(r, fl, 0)
    pst_lines = [hc.take() for _ in range(4)]
    pst_w, pst_ref, pst_found = hc.take(), hc.take(), hc.take()
    pst0 = hc.take(22)
    PST0 = ST_END + 3
    pst_at = {}
    for k, s in enumerate(outside):
        r = PST0 + k
        v, fl = psrc[s]
        idx = "$%s$3" % L(v)
        for j in range(4):
            ws.cell(r, pst_lines[j], '=IF(%s=0,"",INDEX(%s,%s))' % (idx, rng(st_lines[j], STAT0, ST_END), idx))
        ws.cell(r, pst_w, '=IF(%s=0,0,INDEX(%s,%s))' % (idx, rng(st_w, STAT0, ST_END), idx))
        ws.cell(r, pst_ref, '=IF(%s=0,0,INDEX(%s,%s))' % (idx, rng(st_ref, STAT0, ST_END), idx))
        ws.cell(r, pst_found, '=IF(%s>0,TRUE,FALSE)' % idx)
        pst_at[s] = stats(
            r, pst0, rng(pk["val"], *I_rows), rng(pk["flag"], *I_rows), rng(v, *I_rows), rng(fl, *I_rows),
            rng(pk["val"], W0 + 4, W0 + G - 1), rng(pk["val"], W0, W0 + G - 5),
            rng(pk["flag"], W0 + 4, W0 + G - 1), rng(pk["flag"], W0, W0 + G - 5),
            rng(v, W0 + 4, W0 + G - 1), rng(v, W0, W0 + G - 5), rng(fl, W0 + 4, W0 + G - 1), rng(fl, W0, W0 + G - 5),
            ["$%s$%d" % (L(pst_lines[j]), r) for j in range(4)], "$%s$%d" % (L(pst_w), r),
            "$%s$%d" % (L(pst_ref), r), found="$%s$%d" % (L(pst_found), r))
    # The shading source's lines.
    sh = hc.take()
    SH_IDX = "$%s$1" % L(sh)
    ws.cell(1, sh, '=IFERROR(MATCH(%s&" - "&%s&" - "&%s,%s,0),0)' % (PB, WPK, PF, ST_LABELS))
    SH = []
    for j in range(4):
        ws.cell(2 + j, sh, '=IF(%s=0,"",INDEX(%s,%s))' % (SH_IDX, rng(st_lines[j], STAT0, ST_END), SH_IDX))
        SH.append("$%s$%d" % (L(sh), 2 + j))
    SH_W = "$%s$6" % L(sh)
    ws.cell(6, sh, '=IF(%s=0,0,INDEX(%s,%s))' % (SH_IDX, rng(st_w, STAT0, ST_END), SH_IDX))
    SH_REF = "$%s$7" % L(sh)
    ws.cell(7, sh, '=IF(%s=0,1,INDEX(%s,%s))' % (SH_IDX, rng(st_ref, STAT0, ST_END), SH_IDX))
    SH_OK = "$%s$8" % L(sh)
    ws.cell(8, sh, '=AND(%s=0,%s)' % (SH_REF, ",".join("ISNUMBER(%s)" % x for x in SH)))
    misc = hc.take()                     # the latest quarters: the picked product's, and the one compared
    ws.cell(1, misc, "latest quarters")

    # ---- 4. Quarters loaded, per product (L4:Q12)
    _section(ws, 4, 12, 12 + len(products), "Quarters loaded, per product")
    for j, p in enumerate(products):
        c = ws.cell(5, 13 + j, p)
        c.font = F_LAB
        c.alignment = Alignment(horizontal="right")
    ws.row_dimensions[5].height = 15
    cov_rows = ("Quarters", "First quarter", "Latest quarter", "Segments", "Composite from",
                "Balance, latest ($ millions)", "Rows not used")
    for i, t in enumerate(cov_rows):
        ws.cell(6 + i, 12, t).font = F_LAB
    for j, p in enumerate(products):
        b = pblock[p]
        col = 13 + j
        at = rng(b["at"], W0, WE)
        ws.cell(6, col, "=COUNT(%s)" % at)
        ws.cell(7, col, '=IF(%s6=0,"",INDEX(%s,MIN(%s)))' % (L(col), CALQ, at))
        ws.cell(8, col, '=IF(%s6=0,"",INDEX(%s,MAX(%s)))' % (L(col), CALQ, at))
        ws.cell(9, col, '=COUNTIFS(%s,%s,%s,1)' % (S["prod"], crit(p), S["seg_first"]))
        ws.cell(10, col, '=IF({c}6=0,"",IF({h}=1,"All rows",IF({h}=2,"All rows where given, else summed",'
                         '"Summed")))'.format(c=L(col), h=b["hasall"]))
        ws.cell(11, col, '=IF({c}6=0,"",INDEX({b},MATCH({c}8,{q},0))/1000)'
                .format(c=L(col), b=rng(b["bal"], W0, WE), q=CALQ))
        ws.cell(12, col, "=%s" % b["bad"])
        for r in range(6, 13):
            cell = ws.cell(r, col)
            cell.font = F_BODY
            cell.alignment = Alignment(horizontal="right")
        ws.cell(6, col).number_format = "0"
        ws.cell(11, col).number_format = "#,##0.0"
        ws.cell(12, col).number_format = "0"
    if products:
        grid_rng = "M6:%s12" % L(12 + len(products))
        ws.conditional_formatting.add(grid_rng, FormulaRule(formula=['M$5=%s' % PB], fill=fill(K.INTERNAL),
                                                            font=Font(color=K.PAPER, bold=True)))
        ws.conditional_formatting.add("M6:%s6" % L(12 + len(products)),
                                      FormulaRule(formula=["M6>=%s" % MIN_Q], fill=fill(K.INTERNAL_TINT),
                                                  font=Font(color=K.INTERNAL_TEXT)))
        ws.conditional_formatting.add("M12:%s12" % L(12 + len(products)),
                                      FormulaRule(formula=["M12>0"], fill=fill(K.CANVAS),
                                                  font=Font(bold=True, color=K.CRIMSON)))

    # ---- 7. How reliable is each outside source here? (L14:Q31)
    _section(ws, 14, 12, 17, "How reliable is each outside source here?")
    _head(ws, REL_TOP, 12, [""] + outside)
    ws.row_dimensions[REL_TOP].height = 27
    rel_rows = (("Quarters both have", "overlap", "0"), ("Level correlation", "level", "0.00"),
                ("Change correlation", "change", "0.00"),
                ("Average gap", "gap", "+0.00;-0.00;0.00"),
                ("Ratings that agree", "agree", "0%"), ("Within one rating", "within", "0%"),
                ("Verdict", "verdict", None))
    for i, (t, key, fmt) in enumerate(rel_rows):
        r = REL_TOP + 1 + i
        ws.cell(r, 12, t).font = F_LAB
        for k, s in enumerate(outside):
            c = ws.cell(r, 13 + k)
            src = pst_at[s][key]
            if key == "verdict":
                c.value = "=%s" % src
            else:
                c.value = '=IF(%s="%s","",%s)' % (pst_at[s]["verdict"], NO_SERIES, src)
            c.font = F_BOLD if key == "verdict" else F_BODY
            c.alignment = Alignment(horizontal="right" if key != "verdict" else "center", wrap_text=key == "verdict")
            if fmt:
                c.number_format = fmt
        if key == "verdict":
            ws.row_dimensions[r].height = 30
    V_ROW = REL_TOP + len(rel_rows)
    _verdict_formats(ws, "M%d:%s%d" % (V_ROW, L(12 + len(outside)), V_ROW), "M%d" % V_ROW)
    ws.cell(V_ROW + 1, 12, '="Lines: each source\'s own, from Lines by source, for "&%s&" - "&IF(%s="","(not mapped)",%s)&"."'
            % (PB, WPK, WPK)).font = F_NOTE

    # ---- CRR grades (L24:Q30)
    _section(ws, CRR_TOP, 12, 17, "If your file has CRR grades")
    ws.cell(CRR_TOP + 1, 12, "Rating").font = F_LAB
    ws.cell(CRR_TOP + 1, 14, "Lowest CRR grade").font = F_LAB
    for k, (nm, low) in enumerate(zip(("1 Low", "2 Moderate-Low", "3 Moderate", "4 Moderate-High", "5 High"),
                                      (1, 3, 5, 7, 9))):
        r = CRR_TOP + 2 + k
        c = ws.cell(r, 12, nm)
        c.font = F_BODY
        c.fill = fill(K.RATING_FILL[k])
        _input(ws.cell(r, 14), low)
    LOWS = "$N$%d:$N$%d" % (CRR_TOP + 2, CRR_TOP + 6)
    # Per quarter: the grade's rating, and the internal value's rating on the shading source's lines.
    gr_c, ir_c = hc.take(), hc.take()
    ws.cell(4, gr_c, "grade rating")
    ws.cell(4, ir_c, "internal rating on the shading lines")
    for i in range(CAL):
        r = W0 + i
        g_ = "%s%d" % (L(pk["grade"]), r)
        ws.cell(r, gr_c, '=IF(ISNUMBER(%s),COUNTIF(%s,"<="&%s),0)' % (g_, LOWS, g_))
        v_ = "%s%d" % (L(pk["val"]), r)
        ws.cell(r, ir_c, '=IF(AND(%s%d=1,%s),%s,0)' % (L(pk["flag"]), r, SH_OK, rating(v_, SH, SH_W)))
    both = "(%s>0)*(%s>0)" % (rng(gr_c, W0, WE), rng(ir_c, W0, WE))
    crr_n = "SUMPRODUCT(%s*1)" % both
    crr_a = "SUMPRODUCT(%s*(%s=%s))" % (both, rng(gr_c, W0, WE), rng(ir_c, W0, WE))
    asc = "AND(%s)" % ",".join("$N$%d<$N$%d" % (CRR_TOP + 2 + k, CRR_TOP + 3 + k) for k in range(4))
    r = CRR_TOP + 7
    c = ws.cell(r, 12, ('=IF(NOT({asc}),"The lowest grades must rise from rating 1 to rating 5.",'
                        'IF({ix}=0,"The workbook has no "&{f}&" series for this product and measure.",'
                        'IF({rf}=1,"The "&{f}&" lines for this series are refused on Lines by source, so no quarter is '
                        'rated.",IF({n}=0,"No quarter of the picked series has both a CRR grade and an internal '
                        'value.","Your CRR grades agree with the rating on the "&{f}&" lines in "&TEXT({a}/{n},"0%")'
                        '&" of "&{n}&" quarters."))))').format(asc=asc, n=crr_n, a=crr_a, f=PF, ix=SH_IDX, rf=SH_REF))
    c.font = F_BOLD
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=r, start_column=12, end_row=r, end_column=17)
    ws.row_dimensions[r].height = 30

    # ---- the chart block: the last 48 quarters of the picked series
    cb = {k: hc.take() for k in ("idx", "q", "int", "num_lo", "num_hi", "base", "z1", "z2", "z3", "z4", "z5")}
    cb_src = {s: hc.take() for s in outside}
    for k, c_ in cb.items():
        ws.cell(4, c_, "chart " + k)
    START = "$%s$1" % L(cb["idx"])
    ws.cell(1, cb["idx"], "=IF(%s=0,MAX(1,%d),MAX(1,%s-%d))" % (LAST, G - WINDOW + 1, LAST, WINDOW - 1))
    FLOOR, TOP = "$%s$2" % L(cb["idx"]), "$%s$3" % L(cb["idx"])
    C0, CE = W0, W0 + WINDOW - 1
    lo_r, hi_r = rng(cb["num_lo"], C0, CE), rng(cb["num_hi"], C0, CE)
    fz = "TRUE" if (floor_at_zero and higher) else "FALSE"
    ws.cell(2, cb["idx"], "=IF(COUNT({lo})=0,0,IF({fz},MIN(0,MIN({lo})),MIN({lo})-0.08*(MAX({hi})-MIN({lo}))))"
            .format(lo=lo_r, hi=hi_r, fz=fz))
    ws.cell(3, cb["idx"], "=IF(COUNT({hi})=0,1,MAX({hi})+0.08*(MAX({hi})-{f})+IF(MAX({hi})={f},1,0))"
            .format(hi=hi_r, f=FLOOR))
    ws.cell(1, cb["q"], '=IF(%s=0,"No internal data yet",%s&", "&%s&IF(%s,"",", "&%s)&": the last 48 quarters")'
            % (LAST, PB, PC, COMP, PE))
    order = [0, 1, 2, 3] if higher else [3, 2, 1, 0]
    zone_fills = list(K.ZONE_FILL) if higher else list(K.ZONE_FILL[::-1])
    for j in range(WINDOW):
        r = C0 + j
        idx = "%s%d" % (L(cb["idx"]), r)
        ws.cell(r, cb["idx"], "=MIN(%s+%d,%d)" % (START, j, CAL))
        ws.cell(r, cb["q"], "=INDEX(%s,%s)" % (CALQ, idx))
        ws.cell(r, cb["int"], "=IF(INDEX({f},{i})=1,INDEX({v},{i}),NA())".format(
            f=rng(pk["flag"], W0, WE), v=rng(pk["val"], W0, WE), i=idx))
        drawn = ["%s%d" % (L(cb["int"]), r)]
        for s in outside:
            v, fl = psrc[s]
            ws.cell(r, cb_src[s], "=IF(INDEX({f},{i})=1,INDEX({v},{i}),NA())".format(
                f=rng(fl, W0, WE), v=rng(v, W0, WE), i=idx))
            drawn.append("%s%d" % (L(cb_src[s]), r))
        ws.cell(r, cb["num_lo"], '=IF(COUNT({c})=0,"",MIN({m}))'.format(
            c=",".join(drawn), m=",".join("IF(ISNUMBER(%s),%s,1E+300)" % (x, x) for x in drawn)))
        ws.cell(r, cb["num_hi"], '=IF(COUNT({c})=0,"",MAX({m}))'.format(
            c=",".join(drawn), m=",".join("IF(ISNUMBER(%s),%s,-1E+300)" % (x, x) for x in drawn)))
        clip = lambda e: "MIN(MAX(%s,%s),%s)" % (e, FLOOR, TOP)              # noqa: E731
        edges = [FLOOR] + [clip(SH[k]) for k in order] + [TOP]
        ws.cell(r, cb["base"], "=IF(%s,%s,NA())" % (SH_OK, FLOOR))
        for k in range(5):
            ws.cell(r, cb["z%d" % (k + 1)], "=IF(%s,MAX(%s-%s,0),NA())" % (SH_OK, edges[k + 1], edges[k]))
    TITLE_CELL = "'%s'!$%s$1" % (TITLE, L(cb["q"]))
    cats = "'%s'!$%s$%d:$%s$%d" % (TITLE, L(cb["q"]), C0, L(cb["q"]), CE)

    def ser(col, title, idx):
        s_ = Series(val=NumDataSource(numRef=NumRef(f="'%s'!$%s$%d:$%s$%d" % (TITLE, L(col), C0, L(col), CE))),
                    tx=SeriesLabel(v=title))
        s_.cat = AxDataSource(strRef=StrRef(f=cats))
        s_.idx = s_.order = idx
        return s_
    line = LineChart()
    main = ser(cb["int"], "Internal", 0)
    main.smooth = False
    main.graphicalProperties.line = _ln(K.INTERNAL, 2.5)
    main.marker.symbol = "none"
    line.series.append(main)
    looks = line_styles(outside, call_report)
    for k, s in enumerate(outside):
        s_ = ser(cb_src[s], s, 1 + k)
        s_.smooth = False
        s_.marker.symbol = "none"
        s_.graphicalProperties.line = _ln(looks[s][0], 1.5, looks[s][1])
        line.series.append(s_)
    area = AreaChart()
    area.grouping = "stacked"
    n_lines = 1 + len(outside)
    for k, nm in enumerate(["base", "z1", "z2", "z3", "z4", "z5"]):
        s_ = ser(cb[nm], nm, n_lines + k)
        if k == 0:
            s_.graphicalProperties.noFill = True
        else:
            s_.graphicalProperties.solidFill = zone_fills[k - 1]
        s_.graphicalProperties.line.noFill = True
        area.series.append(s_)
    AT._quiet_axes(line, 4)
    line.y_axis.number_format = "0.00"
    line.x_axis.tickLblPos = "low"
    # The key is the row under the chart: Excel 16 ignores a deleted legend
    # entry in a file openpyxl writes, so a legend would list the zones' helpers.
    line.legend = None
    line.title = Title(tx=Text(strRef=StrRef(f=TITLE_CELL)), overlay=False)
    line.title.txPr = AT._text(10, K.INK, True)
    line.visible_cells_only = False
    line += area
    AT._square(line)
    AT._place(ws, line, 2, 16, 11, 32)
    _key(ws, 32, outside, looks)

    # ---- 5. The composite and each segment (B33:J45) and the balance chart (L33:Q47)
    _section(ws, SEG_TOP, 2, 10, None)
    LQ = "$%s$4" % L(misc)                                      # the picked product's latest quarter number
    ws.cell(4, misc, "=IF(COUNT(%s)=0,0,MAX(%s))" % (S["latest"], S["latest"]))
    ws.cell(SEG_TOP, 2, '=%s&": the composite and each segment"&IF(%s=0,"",", "&INDEX(%s,MATCH(%s,%s,0)))'
            % (PB, LQ, CALQ, LQ, CALN)).font = F_SECTION
    seg_head = ["Segment", "Balance ($ millions)", None, "Share of the product", "Net charge-off rate", None,
                "Gross charge-off rate", "30+ days past due", "90+ days past due"]
    for i, t in enumerate(seg_head):
        if t:
            c = ws.cell(SEG_TOP + 1, 2 + i, t)
            c.font = F_HEAD
            c.fill = fill(K.INK)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if i == 0 else "right")
    ws.row_dimensions[SEG_TOP + 1].height = 27

    def seg_formulas(r, crit_pair):
        """Balance at the latest quarter and the four rates, for rows matching
        crit_pair (an extra SUMIFS criterion pair on top of product and used)."""
        base = "%s,%s,%s,1,%s" % (S["prod"], crit_of(PB), S["used"], crit_pair)
        cnt = lambda q: "COUNTIFS(%s,%s,%s)" % (S["qn"], q, base)                       # noqa: E731
        sm = lambda k, q: "SUMIFS(%s,%s,%s,%s)" % (S[k], S["qn"], q, base)              # noqa: E731
        span = lambda k: 'SUMIFS({c},{q},">="&({l}-3),{q},"<="&{l},{b})'.format(c=S[k], q=S["qn"], l=LQ, b=base)  # noqa
        have = "AND(%s)" % ",".join("%s>0" % cnt("%s-%d" % (LQ, k) if k else LQ) for k in range(4))
        ttm = lambda k: 'IF({h},IF({b}>0,100*{n}/({b}/4),""),"")'.format(h=have, b=span("bal"), n=span(k))  # noqa
        qe = lambda k: 'IF(AND({c}>0,{b}>0),100*{d}/{b},"")'.format(c=cnt(LQ), b=sm("bal", LQ), d=sm(k, LQ))  # noqa
        ok = "AND(%s=0,%s>0)" % (BAD, LQ)
        return {3: '=IF(AND(%s,%s>0),%s/1000,"")' % (ok, cnt(LQ), sm("bal", LQ)),
                6: '=IF(%s,%s,"")' % (ok, ttm("nco")), 8: '=IF(%s,%s,"")' % (ok, ttm("gco")),
                9: '=IF(%s,%s,"")' % (ok, qe("d30")), 10: '=IF(%s,%s,"")' % (ok, qe("d90"))}
    comp_r = SEG_TOP + 2
    rows_seg = [(comp_r, COMPOSITE, "%s,%s" % (S["total"], HASALL))]
    for k in range(N_SEGS):
        rows_seg.append((comp_r + 1 + k, "=%s" % SEG_NAME(k), "%s,%s" % (S["seg"], crit_of(SEG_NAME(k)))))
    for r, name, pair in rows_seg:
        ws.cell(r, 2, name).font = F_BOLD if r == comp_r else F_BODY
        if r != comp_r:
            # The segment's swatch, shown only when the row has a segment.
            ws.conditional_formatting.add("A%d" % r, FormulaRule(formula=['$B%d<>""' % r],
                                                                 fill=fill(SEG_COLORS[r - comp_r - 1])))
        for col, fm in seg_formulas(r, pair).items():
            if r != comp_r:
                fm = '=IF(%s="","",%s)' % (SEG_NAME(r - comp_r - 1), fm[1:])
            c = ws.cell(r, col, fm)
            c.font = F_BOLD if r == comp_r else F_BODY
            c.number_format = "#,##0.0" if col == 3 else "0.00"
            c.alignment = Alignment(horizontal="right")
        sh_c = ws.cell(r, 5, '=IF(AND(ISNUMBER(C{r}),ISNUMBER($C${c}),$C${c}>0),C{r}/$C${c},"")'.format(r=r, c=comp_r))
        sh_c.number_format = "0%"
        sh_c.font = F_BOLD if r == comp_r else F_BODY
        sh_c.alignment = Alignment(horizontal="right")
        for col in range(2, 11):
            ws.cell(r, col).border = Border(bottom=HAIR)
    tog_r = comp_r + 1 + N_SEGS
    ws.cell(tog_r, 2, "Segments together").font = F_LAB
    c = ws.cell(tog_r, 3, '=IF(COUNT(C%d:C%d)=0,"",SUM(C%d:C%d))' % (comp_r + 1, tog_r - 1, comp_r + 1, tog_r - 1))
    c.number_format = "#,##0.0"
    c.alignment = Alignment(horizontal="right")
    c = ws.cell(tog_r, 5, '=IF(AND(ISNUMBER(C{r}),ISNUMBER($C${c}),$C${c}>0),C{r}/$C${c},"")'.format(r=tog_r, c=comp_r))
    c.number_format = "0%"
    c.alignment = Alignment(horizontal="right")
    for r in range(SEG_TOP + 2, tog_r + 1):
        ws.row_dimensions[r].height = 15

    # Balance by segment, the chart window: each segment's balance in $ millions.
    bseg = [hc.take() for _ in range(N_SEGS)]
    bcomp = hc.take()
    for k in range(N_SEGS):
        ws.cell(4, bseg[k], "=%s" % SEG_NAME(k))
    ws.cell(4, bcomp, "Composite")
    for j in range(WINDOW):
        r = C0 + j
        qn_ = "INDEX(%s,%s%d)" % (CALN, L(cb["idx"]), r)
        for k in range(N_SEGS):
            base = "%s,%s,%s,1,%s,%s,%s,%s" % (S["prod"], crit_of(PB), S["used"], S["seg"], crit_of(SEG_NAME(k)),
                                              S["qn"], qn_)
            ws.cell(r, bseg[k], '=IF(OR({n}="",{b}>0,COUNTIFS({x})=0),NA(),SUMIFS({v},{x})/1000)'
                    .format(n=SEG_NAME(k), b=BAD, x=base, v=S["bal"]))
        base = "%s,%s,%s,1,%s,%s,%s,%s" % (S["prod"], crit_of(PB), S["used"], S["total"], HASALL, S["qn"], qn_)
        ws.cell(r, bcomp, '=IF(OR({b}>0,COUNTIFS({x})=0),NA(),SUMIFS({v},{x})/1000)'.format(b=BAD, x=base, v=S["bal"]))
    _section(ws, SEG_TOP, 12, 17, "Balance by segment ($ millions)")
    bal_area = AreaChart()
    bal_area.grouping = "stacked"
    for k in range(N_SEGS):
        s_ = Series(val=NumDataSource(numRef=NumRef(f="'%s'!$%s$%d:$%s$%d" % (TITLE, L(bseg[k]), C0, L(bseg[k]), CE))),
                    tx=SeriesLabel(strRef=StrRef(f="'%s'!$%s$4" % (TITLE, L(bseg[k])))))
        s_.cat = AxDataSource(strRef=StrRef(f=cats))
        s_.idx = s_.order = 1 + k
        s_.graphicalProperties.solidFill = SEG_COLORS[k]
        s_.graphicalProperties.line.noFill = True
        bal_area.series.append(s_)
    bal_line = LineChart()
    s_ = Series(val=NumDataSource(numRef=NumRef(f="'%s'!$%s$%d:$%s$%d" % (TITLE, L(bcomp), C0, L(bcomp), CE))),
                tx=SeriesLabel(v="Composite"))
    s_.cat = AxDataSource(strRef=StrRef(f=cats))
    s_.idx = s_.order = 0
    s_.smooth = False
    s_.marker.symbol = "none"
    s_.graphicalProperties.line = _ln(K.INK, 1.5)
    bal_line.series.append(s_)
    AT._quiet_axes(bal_line, 8)
    bal_line.y_axis.number_format = "#,##0"
    bal_line.x_axis.tickLblPos = "low"
    bal_line.legend = None                  # the swatches in column A of the segment table are its key
    bal_line.visible_cells_only = False
    bal_line += bal_area
    AT._square(bal_line)
    AT._place(ws, bal_line, 12, SEG_TOP + 1, 18, tog_r + 4)

    # ---- 6. Against the Call Report, like for like
    LR = tog_r + 5
    _section(ws, LR, 2, 17, None)
    ws.cell(LR, 2, '="Against "&%s&"\'s Call Report, like for like, "&IF(%s=0,"(no quarter yet)",%s)'
            % (CR, "$%s$5" % L(misc), "$%s$6" % L(misc))).font = F_SECTION
    # The quarter compared: the latest the internal file and the Data tab both hold.
    all_at = ",".join(rng(pblock[p]["at"], W0, WE) for p in products)
    ws.cell(5, misc, '=IF(COUNT({a})=0,0,MIN(INDEX({n},MAX({a})),{last}))'.format(a=all_at, n=CALN, last=qn[-1]))
    QSTAR = "$%s$5" % L(misc)
    ws.cell(6, misc, '=IF(%s=0,"",INDEX(%s,MATCH(%s,%s,0)))' % (QSTAR, CALQ, QSTAR, CALN))
    lr_head = ["Comparison", "Internal figure", None, "Call Report lines", "Read from", None,
               "Balance, internal ($ millions)", "Balance, Call Report", "Balance gap", None,
               "Net charge-off rate, internal", "Net charge-off rate, Call Report", "Rate gap (pts)"]
    for i, t in enumerate(lr_head):
        if t:
            c = ws.cell(LR + 1, 2 + i, t)
            c.font = F_HEAD
            c.fill = fill(K.INK)
            c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="left" if i < 5 else "right")
    ws.row_dimensions[LR + 1].height = 39
    peers_ok = "Peers data" in wb.sheetnames and internal.get("peers_self")
    net_w = "INDEX(%s,1)" % M_MAP                                     # the workbook measure net maps to
    lr_rows = []
    for k, comp in enumerate(I.COMPARISONS):
        r = LR + 2 + k
        lr_rows.append(r)
        ws.cell(r, 2, comp.name).font = F_BOLD
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="center")
        ws.cell(r, 3, " + ".join("%s, composite" % ps[0] if ps[1] == I.ALL else "%s / %s" % ps
                                 for ps in comp.internal)).font = F_BODY
        ws.cell(r, 5, comp.items).font = F_BODY
        for col in (3, 5):
            ws.cell(r, col).alignment = Alignment(wrap_text=True, vertical="center")
        # The internal side: four quarters ending at the quarter compared, every part present.
        hcs = hc.take(4)
        for j, (what, k_) in enumerate((("cnt", None), ("bal", "bal"), ("gco", "gco"), ("nco", "nco"))):
            terms = []
            for ps in comp.internal:
                if ps[1] == I.ALL:
                    # The product's composite: its All rows, or its segments summed.
                    b = "%s,%s,%s,IF(COUNTIFS(%s,%s,%s,1,%s,1)>0,1,0),%s,1" % (
                        S["prod"], crit(ps[0]), S["total"], S["prod"], crit(ps[0]), S["total"], S["used"], S["used"])
                else:
                    b = "%s,%s,%s,%s,%s,1" % (S["prod"], crit(ps[0]), S["seg"], crit(ps[1]), S["used"])
                if what == "cnt":
                    terms.append("AND(%s)" % ",".join(["COUNTIFS(%s,%s-%d,%s)>0" % (S["qn"], QSTAR, d, b)
                                                        for d in range(4)]
                                                       + ["COUNTIFS(%s,%s,%s,0)=0" % (S["prod"], crit(ps[0]),
                                                                                      S["used"])]))
                else:
                    terms.append('SUMIFS({v},{q},">="&({s}-3),{q},"<="&{s},{b})'
                                 .format(v=S[k_], q=S["qn"], s=QSTAR, b=b))
            ws.cell(r, hcs + j, "=AND(%s)" % ",".join(terms) if what == "cnt" else "=" + "+".join(terms))
        H_OK, H_BAL, H_GCO, H_NCO = ("$%s$%d" % (L(hcs + j), r) for j in range(4))
        ok_int = "AND(%s>0,%s,%s>0)" % (QSTAR, H_OK, H_BAL)
        # The public side.
        prods, missing = decompose(comp)
        pub_bal = pub_rate = None
        if prods is None:
            how = "No product in this workbook holds %s." % " or ".join(missing)
        elif peers_ok:
            me = internal["peers_self"]
            pc = "'Peers data'!"
            pq = 'INDEX(%s,MATCH(%s,%s,0))' % (CALQ, QSTAR, CALN)

            def pd(col):
                return "+".join("SUMIFS({pc}${c}:${c},{pc}$A:$A,{p},{pc}$B:$B,{me},{pc}$D:$D,{q})"
                                .format(pc=pc, c=col, p=crit(p), me=crit(me), q=pq) for p in prods)
            have = "AND(%s)" % ",".join("COUNTIFS({pc}$A:$A,{p},{pc}$B:$B,{me},{pc}$D:$D,{q})=1"
                                        .format(pc=pc, p=crit(p), me=crit(me), q=pq) for p in prods)
            pub_bal = 'IF(AND(%s>0,%s),(%s)/1000,"")' % (QSTAR, have, pd("G"))
            pub_rate = 'IF(AND(%s>0,%s,(%s)>0),100*((%s)-(%s))/(%s),"")' % (QSTAR, have, pd("G"), pd("E"), pd("F"),
                                                                            pd("G"))
            how = "Peers data, %s, %s" % (me, " + ".join(prods))
        elif len(prods) == 1:
            lab = '%s&" - "&%s&" - "&%s' % (lit(prods[0]), net_w, CR)
            pub_rate = ('IFERROR(IF(INDEX(Data!$B$2:${e}${g},MATCH({q},Data!$A$2:$A${g},0),MATCH({lab},{h},0))="","",'
                        'INDEX(Data!$B$2:${e}${g},MATCH({q},Data!$A$2:$A${g},0),MATCH({lab},{h},0))),"")'
                        .format(e=d_last, g=G + 1, q='INDEX(%s,MATCH(%s,%s,0))' % (CALQ, QSTAR, CALN), lab=lab,
                                h=data_hdr))
            how = ('=IF(ISNUMBER(MATCH(%s,%s,0)),"The rate is the Data tab\'s "&%s&". Balances need Peers '
                   'data, which this workbook does not have.","This workbook has no series "&%s&".")'
                   % (lab, data_hdr, lab, lab))
        else:
            how = ("Needs the %s amounts on Peers data, which this workbook does not have."
                   % " and ".join(prods))
        c = ws.cell(r, 6, how)
        c.font = F_NOTE
        c.alignment = Alignment(wrap_text=True, vertical="center")
        cells = {8: '=IF(%s,(%s/4)/1000,"")' % (ok_int, H_BAL),
                 9: "=%s" % pub_bal if pub_bal else '=""',
                 10: '=IF(AND(ISNUMBER(H{r}),ISNUMBER(I{r})),H{r}-I{r},"")'.format(r=r),
                 12: '=IF({ok},100*{n}/({b}/4),"")'.format(ok=ok_int, n=H_NCO, b=H_BAL),
                 13: "=%s" % pub_rate if pub_rate else '=""',
                 14: '=IF(AND(ISNUMBER(L{r}),ISNUMBER(M{r})),L{r}-M{r},"")'.format(r=r)}
        for col, fm in cells.items():
            c = ws.cell(r, col, fm)
            c.font = F_BODY
            c.alignment = Alignment(horizontal="right", vertical="center")
            c.number_format = "#,##0.0" if col in (8, 9) else "+#,##0.0;-#,##0.0;0.0" if col == 10 else \
                "+0.00;-0.00;0.00" if col == 14 else "0.00"
        ws.row_dimensions[r].height = 45
        for col in range(2, 18):
            ws.cell(r, col).border = Border(bottom=HAIR)
    note_r = lr_rows[-1] + 1
    ws.cell(note_r, 2, ("Balances are the mean of four quarter-end balances, in $ millions; the net charge-off rate "
                        "is four quarters' net charge-offs over that mean. Named exclusions are not taken off "
                        "here; the internal reconcile command applies them.")).font = F_NOTE
    ws.merge_cells(start_row=note_r, start_column=2, end_row=note_r, end_column=17)
    ws.row_dimensions[note_r].height = 28
    ws.cell(note_r, 2).alignment = Alignment(wrap_text=True, vertical="top")

    # ---- 8. Every product and measure you loaded
    EV = note_r + 2
    _section(ws, EV, 2, 17, "Every product and measure you loaded")
    agree_cols = [8, 9, 10, 12][:len(outside)]
    cols_ev = {"Product": 2, "Measure": 3, "Workbook measure": 5, "Quarters": 6}
    for s_, c_ in zip(outside, agree_cols):
        cols_ev["Agree: %s" % s_] = c_
    cols_ev["Best source"] = 13
    cols_ev["Verdict"] = 15
    for t, col in cols_ev.items():
        c = ws.cell(EV + 1, col, t)
        c.font = F_HEAD
        c.fill = fill(K.INK)
        c.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="left" if col < 6 or t in ("Best source", "Verdict") else "right")
    for col in range(2, 18):
        ws.cell(EV + 1, col).fill = fill(K.INK)
    ws.row_dimensions[EV + 1].height = 27
    r = EV + 2
    best_c, verdict_c = cols_ev.get("Best source"), cols_ev.get("Verdict")
    for p in products:
        b = pblock[p]
        for mi, (key, name) in enumerate(MEASURES):
            ws.cell(r, 2, p).font = F_BODY
            ws.cell(r, 3, name).font = F_BODY
            mp = "$F$%d" % (5 + mi)
            ws.cell(r, 5, "=%s" % mp).font = F_BODY
            ws.cell(r, 6, "=COUNT(%s)" % rng(b[RATE_KEYS[mi]], W0, WE)).font = F_BODY
            ws.cell(r, 6).number_format = "0"
            scores, verdicts = [], []
            for s in outside:
                idx = 'IFERROR(MATCH(%s&" - "&%s&" - "&%s,%s,0),0)' % (lit(p), mp, lit(s), ST_LABELS)
                agree = 'IF({i}=0,"",INDEX({a},{i}))'.format(i=idx, a=st_range("agree"))
                scores.append('IF({i}=0,-1,INDEX({a},{i}))'.format(i=idx, a=st_range("score")))
                verdicts.append('IF({i}=0,"",INDEX({a},{i}))'.format(i=idx, a=st_range("verdict")))
                t = "Agree: %s" % s
                if t in cols_ev:
                    c = ws.cell(r, cols_ev[t], "=" + agree)
                    c.number_format = "0%"
                    c.font = F_BODY
                    c.alignment = Alignment(horizontal="right")
            hidden_s = hc.take(len(outside) + 1)
            for k in range(len(outside)):
                ws.cell(r, hidden_s + k, "=" + scores[k])
            sc = "$%s$%d:$%s$%d" % (L(hidden_s), r, L(hidden_s + len(outside) - 1), r)
            vv = hidden_s + len(outside)
            # Verdicts of every source, for the row when no source qualifies.
            ws.cell(r, vv, "=" + "&\"|\"&".join(verdicts))
            if best_c:
                ws.cell(r, best_c, '=IF(MAX({s})<0,"",INDEX({o},MATCH(MAX({s}),{s},0)))'
                        .format(s=sc, o=lists["sources"])).font = F_BOLD
                ws.merge_cells(start_row=r, start_column=best_c, end_row=r, end_column=best_c + 1)
            if verdict_c:
                c = ws.cell(r, verdict_c, ('=IF(MAX({s})>=0,CHOOSE(MATCH(MAX({s}),{s},0),{vs}),IF(ISNUMBER(SEARCH("{v1}",'
                                           '{vv})),"{v1}",IF(ISNUMBER(SEARCH("{v0}",{vv})),"{v0}","{ns}")))')
                           .format(s=sc, vs=",".join(verdicts), vv="$%s$%d" % (L(vv), r), v0=VERDICTS[0],
                                   v1=VERDICTS[1], ns=NO_SERIES))
                c.font = F_BOLD
                c.alignment = Alignment(horizontal="center")
                ws.merge_cells(start_row=r, start_column=verdict_c, end_row=r, end_column=verdict_c + 2)
            for c_ in range(2, 18):
                ws.cell(r, c_).border = Border(bottom=HAIR)
            ws.row_dimensions[r].height = 15
            r += 1
    EV_END = r - 1
    if verdict_c:
        _verdict_formats(ws, "%s%d:%s%d" % (L(verdict_c), EV + 2, L(verdict_c), EV_END), "%s%d" % (L(verdict_c), EV + 2))

    # ---- notes, the method once
    notes = [
        "Composite: a product's All rows where the file has them; otherwise its segments summed, quarter by quarter.",
        "Charge-off rates are four quarters' charge-offs over the mean of the four quarter-end balances, times 100. "
        "Past-due rates are the amount over the balance at the same quarter-end, times 100.",
        "A product with any row not used has no composite until the row is fixed; Why not used on Internal data "
        "names the problem.",
        "Each source is rated on its own lines from Lines by source with the On a line rule set on Decide. The verdict "
        "rules are on the Glossary tab.",
    ]
    for i, t in enumerate(notes):
        c = ws.cell(EV_END + 2 + i, 2, t)
        c.font = F_NOTE
    for col in range(hc_start(products), hc.next):
        ws.column_dimensions[L(col)].hidden = True
    ws.freeze_panes = None
    ws.sheet_view.zoomScale = 100
    return {"sheet": TITLE, "data_sheet": DATA_TITLE, "table": TABLE, "capacity": cap, "first": FIRST,
            "last": last_slot, "outside": outside, "call_report": call_report, "products": products,
            "pick_row": PICK_ROW, "rel_top": REL_TOP, "verdict_row": V_ROW, "crr_row": CRR_TOP + 7,
            "seg_rows": (comp_r, tog_r), "like_rows": lr_rows, "every_rows": (EV + 2, EV_END),
            "every_cols": cols_ev, "stats": st_at, "stat_rows": {lab: STAT0 + j for j, lab in enumerate(out_labels)},
            "pblock": pblock, "w0": W0, "calendar": cal, "mapping_rows": (5, 4 + N_NAMES),
            "measure_rows": (5, 4 + len(MEASURES)), "helpers": H, "pick_stats": pst_at}


#: Segment fills on the balance chart, darkest at the bottom of the stack.
SEG_COLORS = ("57534B", "8C877E", "B9B4AC", "CFC8BC", "2F5D50", "7A9A8F", "A9C1B8", "E1EAE6")


def line_styles(outside, call_report):
    """Each outside source's line: the Call Report source in KEY_RED, dashed;
    a peer group INK sysDash; an industry SLATE sysDot (the README's)."""
    out, spare = {}, [(K.MEDIAN, "dash"), (K.STONE, "sysDash"), (K.SLATE, "dash")]
    for s in outside:
        if s == call_report:
            out[s] = (K.KEY_RED, "dash")
        elif re.search(r"peer", s, re.I):
            out[s] = (K.INK, "sysDash")
        elif re.search(r"industry", s, re.I):
            out[s] = (K.SLATE, "sysDot")
        else:
            out[s] = spare[len([x for x in out if out[x] in spare]) % len(spare)]
    return out


def _key(ws, row, outside, looks):
    """The chart's key under it: each line's name in its color, one cell
    each (B, C, E, F, then H, I), and what the shading is."""
    mark = {None: "\u2014", "dash": "- -", "sysDash": "- - -", "sysDot": "\u00b7\u00b7\u00b7"}
    items = [("%s Internal" % mark[None], K.INTERNAL)] + [("%s %s" % (mark[looks[s][1]], s), looks[s][0])
                                                          for s in outside]
    cols = [2, 3, 5, 6, 8, 9, 10]
    for (text, color), col in zip(items, cols):
        c = ws.cell(row, col, text)
        c.font = Font(name="Calibri", size=9, bold=True, color=color)
    nxt = cols[min(len(items), len(cols) - 1)]
    ws.cell(row, nxt, "Shading: ratings on the lines picked under Shade with lines from.").font = \
        Font(name="Calibri", size=9, color=K.SLATE)
    ws.row_dimensions[row].height = 15


def hc_start(products):
    return max(27, 13 + len(products) + 3)


def _verdict_formats(ws, rng_, first):
    for text, fill_, font in ((VERDICTS[2], K.INTERNAL_TINT, K.INTERNAL_TEXT), (VERDICTS[3], K.RATING_FILL[2], K.INK),
                              (VERDICTS[4], K.CANVAS, K.CRIMSON), (VERDICTS[1], K.CANVAS, K.CRIMSON),
                              (VERDICTS[0], K.PAPER, K.SLATE), (NO_SERIES, K.PAPER, K.SLATE)):
        ws.conditional_formatting.add(rng_, FormulaRule(formula=['%s="%s"' % (first.replace("$", ""), text)],
                                                        fill=fill(fill_), font=Font(bold=True, color=font)))
