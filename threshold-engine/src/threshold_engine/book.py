"""The book: one file that says everything a workbook is built from.

``threshold-engine book BOOK.json OUT.xlsx`` reads it and writes one workbook:
Summary and Decide first, then Thresholds, Assess and the reference tabs, and,
when the book has a ``peers`` section, the peer tabs built into the same file.
No Excel is needed to build it; the Decide tab's slicers are written into the
file (``slicers.py``).

Every key is required and none has a default (the tool's rule: a missing
judgement is refused, not filled in). Relative paths are read from the book's
own folder, so a book and its files can be copied to another machine whole.

    {
      "workbook": {"name", "unit", "direction", "frequency", "smoothing",
                   "floor_at_zero", "moderate_halfwidth", "low_step", "high_z",
                   "on_the_line", "half_lives", "percentiles", "horizon",
                   "min_history"},
      "parts": ["Product", "Measure", "Source"],
      "lines_from": {"part": "Source", "start": null or one of its values},
      "series": [{"Product": ..., "Measure": ..., "Source": ..., "file": ...}]
                or "a CSV with those columns",
      "not_in_book": [{"Product": ..., "Measure": ..., "Source": ..., "why": ...}],
      "compare": ["KeyBank", "Peer group"] or null,
      "notes": ["..."],
      "peers": null or {every setting 'peers build' takes, by its keyword name},
      "internal": null or {"file": the bank's internal file (docs/INTERNAL-IMPORT.md),
                           "source": the source its series go under, e.g. "KeyBank internal",
                           "segments": true to add each segment as a series as well,
                           "measures": {internal measure: the book's measure it is},
                           "call_report_source": the source holding the bank's Call Report series},
      "loan_classes": null or {"file": the loan classes file (config/peer-loan-classes.json),
                               "cache": the folder 'peers classes-pull' fetched the filings into},
      "reporting_changes": null or the registry of Call Report changes (config/reporting-changes.json),
      "floor_off": [measures whose series have no floor at zero (a gap can be negative); [] for none],
      "large_banks": null or {"file": the Philadelphia Fed credit card balances file 'peers large-banks-pull' fetched},
      "signals": null or {"balances", "originations", "survey": the files 'peers signals-pull' fetched;
                          "nmdb", "fico": the committed files under data/signals}
    }

``signals`` (18 Oct 2026): the Signals tab, every public signal found in the
early-warning work laid beside the book's series (``signals_tab.py``,
``signals_data.py``). Each of its five files is named; a file that is not
there stops the build and says which pull or folder gives it. Null builds
the book without the tab.

``floor_off`` (17 Oct 2026, the gap measures): the workbook's ``floor_at_zero``
says a value at or below zero is always Low, which is true of a rate and
false of a difference between two rates. Each measure named here is rated
without the floor, on Decide, Thresholds and Lines by source alike, and its
lines are set without it. A name that is not a measure of the book's series
is refused.

``large_banks``: the build writes each of its measures (``peers.large_banks``)
as a series file beside the peer series, and the series list names it under
the source "Large banks (Philadelphia Fed)".

``reporting_changes`` (the firm, 16 Oct 2026: "Definitely mark reporting
changes."): the registry is checked as ``reporting_changes.load`` checks it,
given to the peers build (every series file it writes notes each change
touching it, and the Peers charts mark the quarters), and written out whole
on the Reporting changes tab. Decide's chart marks the quarters a series'
own notes name. Null states the book's series are not Call Report series; a
book with a peers section reads the Call Report and must name it.

With an internal section, run on site only: the file is checked as the
``internal check`` command checks it; each product's composite (and, when
asked, each segment) is written beside the workbook as a series for every
measure the book maps; and the file's rows fill the Internal data tab. A
measure the book does not map stays on the Internal tab: whether an internal
30+ includes nonaccrual loans is the book's to say, and is never assumed.

A part value of ``*`` in ``not_in_book`` stands for every value of that part.

``lines_from`` says which part is the source and what each row's Lines set
from starts as. Decide and Thresholds have one row per value of the other two
parts (a product and a measure), and its lines are set from the history of
the source picked there (the firm, 15 Oct 2026: "we don't need to set
thresholds on all 3, we need to use all 3 to set appropriate thresholds").
``start`` null leaves every row's pick blank, which is the bank's to make.
The source part is the third of ``parts``.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl.styles import Font
from typing import Dict, List, Optional, Tuple

from .series import LABEL_RULE, label_ok

WORKBOOK_KEYS = ("name", "unit", "direction", "frequency", "smoothing", "floor_at_zero", "moderate_halfwidth",
                 "low_step", "high_z", "on_the_line", "half_lives", "percentiles", "horizon", "min_history")
TOP_KEYS = ("workbook", "parts", "lines_from", "series", "not_in_book", "compare", "notes", "peers", "internal",
            "floor_off", "large_banks", "signals",
            "loan_classes", "reporting_changes")
#: The Peer books tab: the stated classes file and the SEC filings' cache.
LOAN_CLASSES_KEYS = ("file", "cache")
LINES_FROM_KEYS = ("part", "start")
#: The internal section: the file (docs/INTERNAL-IMPORT.md) and the judgments made about it.
INTERNAL_KEYS = ("file", "source", "segments", "measures", "call_report_source")
#: 'peers build' settings, by the keyword ``peers.run.build`` takes; paths among them. The one
#: it takes that is not here, reporting_changes, is the book's own key, given to it.
PEERS_KEYS = ("cache", "roster", "series_dir", "first", "last", "products", "mergers", "absorptions",
              "extra_acquisitions", "industry_combined", "jump", "acquisition_share", "min_book", "exempt_self",
              "pro_forma", "pack", "min_banks", "rank_by", "last_n", "min_obs", "base_quarter", "index_min_book",
              "lagged_years", "lag_years", "normalisation_years", "mix_base_year", "delinquency")
PEERS_PATHS = ("cache", "roster", "series_dir", "mergers", "absorptions")
#: The peer tabs' titles inside the book's workbook, which has its own Data, Glossary and Run.
#: One dashboard (Peers) and its hidden Workings since 14 Oct 2026; before, one tab per view. Since
#: 15 Oct 2026 the peer glossary and run record are sections of the book's own Glossary and Run.
PEER_TAB_NAMES = {"Peers": "Peers", "Workings": "Peers workings", "Glossary": "Glossary",
                  "Data": "Peers data", "DelinquencyData": "Peers delinquency data", "Run": "Run"}


def _live_peer_group(wb, series, info, say) -> int:
    """Make the peer group series follow the banks in the pack on the Peers
    tab (the analyst, 15 Oct 2026: "this tab needs to adjust with peers").

    A peer group series is a pack composite the peers build wrote at build
    time with the pack as it started. Here its Data column becomes live: each
    quarter it had a value, the same composite is summed from Peers data (or
    Peers delinquency data) over the banks kept that quarter and in the pack
    now, needing at least the banks the Peers tab asks for, and blank below
    that. With the pack as built, every value equals the file's. Its latest
    value and quarter on Lines by source become live too. Which composite a
    series is comes from its file's name, as the peers build wrote it.
    Returns the number of series made live."""
    import re
    from openpyxl.utils import get_column_letter as L
    from .peers import delinquency as D
    from .peers.run import slug
    from .peers.peers_workbook import _quoted

    sheets = dict(zip(("Peers", "Data", "DelinquencyData"), (PEER_TAB_NAMES[k] for k in
                                                              ("Peers", "Data", "DelinquencyData"))))
    pe, pd_, dd = (_quoted(sheets[k]) + "!" for k in ("Peers", "Data", "DelinquencyData"))
    mb = pe + "$%s$%s" % re.fullmatch(r"([A-Z]+)(\d+)", info["layout"]["settings_cells"]["min_banks"]).groups()
    products, Q = info["products"], len(info["quarters"])
    NB = info["data_rows"] // (len(products) * Q)
    blocks = {("co", slug(p_)): (2 + i * NB * Q, 1 + (i + 1) * NB * Q) for i, p_ in enumerate(products)}
    dq = info.get("delinquency")
    if dq:
        NM, Qd = len(D.MEASURES), len(dq["quarters"])
        for i, p_ in enumerate(dq["products"]):
            for j, m in enumerate(D.MEASURES):
                a = 2 + (i * NM + j) * NB * Qd
                blocks[("dq", slug(p_), slug(m))] = (a, a + NB * Qd - 1)
    data = wb["Data"]
    head = {data.cell(1, c).value: c for c in range(1, data.max_column + 1)}
    last = data.max_row
    while last > 1 and data.cell(last, 1).value is None:
        last -= 1
    lbs = wb["Lines by source"] if "Lines by source" in wb.sheetnames else None
    made = 0
    for label, path in series:
        m = re.fullmatch(r"(.+)_pack-composite_(net|gross|delinquency-(.+)_quarter-end)", Path(path).stem)
        if not m or label not in head:
            continue
        col = head[label]
        if m.group(3) is None:
            a, z = blocks[("co", m.group(1))]
            rng = lambda c_: "%s$%s$%d:$%s$%d" % (pd_, c_, a, c_, z)                       # noqa: E731
            crit = lambda r: "%s,$A%d,%s,TRUE,%s,TRUE" % (rng("D"), r, rng("J"), rng("P"))  # noqa: E731
            num = (lambda r: "SUMIFS(%s,%s)-SUMIFS(%s,%s)" % (rng("E"), crit(r), rng("F"), crit(r))) \
                if m.group(2) == "net" else (lambda r: "SUMIFS(%s,%s)" % (rng("E"), crit(r)))
            den = lambda r: "SUMIFS(%s,%s)" % (rng("G"), crit(r))                           # noqa: E731
        else:
            key = ("dq", m.group(1), m.group(3))
            if key not in blocks:
                continue
            a, z = blocks[key]
            rng = lambda c_: "%s$%s$%d:$%s$%d" % (dd, c_, a, c_, z)                       # noqa: E731
            crit = lambda r: "%s,$A%d,%s,TRUE,%s,TRUE" % (rng("E"), r, rng("I"), rng("O"))  # noqa: E731
            num = lambda r: "SUMIFS(%s,%s)" % (rng("F"), crit(r))                           # noqa: E731
            den = lambda r: "SUMIFS(%s,%s)" % (rng("G"), crit(r))                           # noqa: E731
        for r in range(2, last + 1):
            if isinstance(data.cell(r, col).value, (int, float)):
                data.cell(r, col).value = '=IF(COUNTIFS(%s)>=%s,100*(%s)/%s,"")' % (crit(r), mb, num(r), den(r))
        if lbs is not None:
            c_, rows_ = L(col), "Data!$%s$2:$%s$%d" % (L(col), L(col), last)
            hit = "MATCH(9.99E+307,%s)" % rows_
            for r in range(1, lbs.max_row + 1):
                if lbs.cell(r, 1).value == label:
                    lbs["L%d" % r] = '=IFERROR(INDEX(Data!$A$2:$A$%d,%s),"")' % (last, hit)
                    lbs["M%d" % r] = '=IFERROR(INDEX(%s,%s),"")' % (rows_, hit)
        made += 1
    say("peer group: %d series follow the banks in the pack on Peers" % made)
    return made


#: The file of quarters left out as built, written beside the series files
#: (the agent building Book 2, 17 Oct 2026, asked for it; the analyst: "sure").
LEFT_OUT_FILE = "left-out.csv"
LEFT_OUT_HEAD = ("Product", "Measure", "Source", "first", "last", "why")


def write_left_out(path: Path, series, parts: Dict[str, tuple], spans: Dict) -> int:
    """Write ``path``: one row per span a book series left out as built
    (``peers.run.left_out_spans``), in the series list's order, named by the
    book's Product, Measure and Source. A series with nothing left out has no
    row. The periods left out on Decide are the analyst's and are not here.
    Returns the number of rows."""
    by_path = {str(Path(p_)).lower(): v for p_, v in spans.items()}
    rows = []
    for lbl, p_ in series:
        for first, last, why in by_path.get(str(Path(p_)).lower(), []):
            rows.append(list(parts[lbl]) + [first, last, why])
    with open(path, "w", newline="", encoding="utf-8") as fh:
        fh.write("# Quarters each series left out as built, before its last break or after its last value, "
                 "with why. A series with nothing left out has no row.\n")
        fh.write("# The periods left out on Decide are the analyst's, typed in the saved workbook, and are not "
                 "here. Quarters are YYYYQn.\n")
        w = csv.writer(fh)
        w.writerow(LEFT_OUT_HEAD)
        w.writerows(rows)
    return len(rows)


def _signals_on_run(wb, signals):
    """A section on Run naming every source behind the Signals tab, with its file and SHA-256."""
    from . import reference_tabs as RT
    RT.run_signals(wb["Run"], signals["sources"])


class BookError(ValueError):
    """A book that cannot be built from, with every problem found."""


def _keys(where: str, got: dict, want, problems: List[str]) -> None:
    if not isinstance(got, dict):
        problems.append("%s must be an object with the keys %s" % (where, ", ".join(want)))
        return
    for k in want:
        if k not in got:
            problems.append("%s is missing %r (every key is required; none has a default)" % (where, k))
    for k in got:
        if k not in want:
            problems.append("%s has %r, which is not a key it takes (check the spelling)" % (where, k))


def load(path) -> dict:
    """Read and check a book. Returns ``{"workbook": settings, "series":
    [(label, Path)], "decide": spec for workbook.build, "peers": settings or
    None, "folder": Path}``; raises BookError listing every problem at once."""
    path = Path(path)
    folder = path.resolve().parent
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise BookError("%s is not valid JSON: %s" % (path.name, e)) from None
    problems: List[str] = []
    _keys("the book", raw, TOP_KEYS, problems)
    if problems:
        raise BookError("\n".join(problems))
    _keys("workbook", raw["workbook"], WORKBOOK_KEYS, problems)
    parts = raw["parts"]
    if not (isinstance(parts, list) and len(parts) == 3 and all(isinstance(p, str) and p.strip() for p in parts)):
        problems.append("parts must be a list of three names, for example [\"Product\", \"Measure\", \"Source\"]")
        raise BookError("\n".join(problems))

    def resolve(p) -> Path:
        q = Path(p)
        return q if q.is_absolute() else folder / q

    rows = raw["series"]
    if isinstance(rows, str):
        src = resolve(rows)
        if not src.exists():
            problems.append("series: the file %s does not exist" % src)
            rows = []
        else:
            with src.open(newline="", encoding="utf-8-sig") as fh:
                rows = list(csv.DictReader(fh))
            # A CSV's file paths are read from the CSV's own folder.
            for r in rows:
                if r.get("file") and not Path(r["file"]).is_absolute():
                    r["file"] = str(src.parent / r["file"])
    if not isinstance(rows, list) or not rows:
        problems.append("series must list at least one series")
        rows = []
    series, decide_parts, seen = [], {}, set()
    for i, r in enumerate(rows, start=1):
        where = "series %d" % i
        if not isinstance(r, dict):
            problems.append("%s must be an object" % where)
            continue
        _keys(where, r, list(parts) + ["file"], problems)
        if any(k not in r for k in list(parts) + ["file"]):
            continue
        vals = tuple(str(r[p]).strip() for p in parts)
        if not all(vals):
            problems.append("%s: every part needs a value" % where)
            continue
        label = " - ".join(vals)
        if any(" - " in v for v in vals):
            problems.append("%s: a part value may not hold ' - ', which joins the parts into the label" % where)
            continue
        if not label_ok(label):
            problems.append("%s: %r can't be used: %s" % (where, label, LABEL_RULE))
            continue
        if label.casefold() in seen:
            problems.append("%s: %r is listed twice" % (where, label))
            continue
        seen.add(label.casefold())
        series.append((label, resolve(r["file"])))
        decide_parts[label] = vals
    absent = raw["not_in_book"]
    if not isinstance(absent, list):
        problems.append("not_in_book must be a list (empty to say there is nothing missing to explain)")
        absent = []
    absent_out = []
    for i, a in enumerate(absent, start=1):
        where = "not_in_book %d" % i
        _keys(where, a, list(parts) + ["why"], problems)
        if isinstance(a, dict) and all(k in a for k in list(parts) + ["why"]):
            if not str(a["why"]).strip():
                problems.append("%s: why must say why" % where)
            absent_out.append({"parts": {p: a[p] for p in parts}, "why": str(a["why"]).strip()})
    compare = raw["compare"]
    third = {v[2] for v in decide_parts.values()}
    if compare is not None:
        if not (isinstance(compare, list) and len(compare) == 2 and all(c in third for c in compare)
                and compare[0] != compare[1]):
            problems.append("compare must be null or two different values of %s that the series use (%s)"
                            % (parts[2], ", ".join(sorted(third))))
            compare = None
        else:
            compare = tuple(compare)
    notes = raw["notes"]
    if not (isinstance(notes, list) and all(isinstance(t, str) for t in notes)):
        problems.append("notes must be a list of sentences (empty for none)")
        notes = []
    peers = raw["peers"]
    if peers is not None:
        _keys("peers", peers, PEERS_KEYS, problems)
        if isinstance(peers, dict) and all(k in peers for k in PEERS_KEYS):
            peers = dict(peers)
            for k in PEERS_PATHS:
                peers[k] = resolve(peers[k])
    # A series file that is not there is only a problem if no peers build will write it.
    if peers is None:
        for label, p in series:
            if not p.exists():
                problems.append("the file for %r does not exist: %s" % (label, p))
    internal = _internal(raw["internal"], resolve, decide_parts, problems)
    loan_classes = raw["loan_classes"]
    if loan_classes is not None:
        _keys("loan_classes", loan_classes, LOAN_CLASSES_KEYS, problems)
        if isinstance(loan_classes, dict) and all(k in loan_classes for k in LOAN_CLASSES_KEYS):
            loan_classes = {k: resolve(loan_classes[k]) for k in LOAN_CLASSES_KEYS}
            if not loan_classes["file"].exists():
                problems.append("loan_classes: the file %s does not exist" % loan_classes["file"])
    lines_from = _lines_from(raw["lines_from"], parts, third, internal, problems)
    floor_off = raw["floor_off"]
    measures_ = {v[1] for v in decide_parts.values()}
    if not (isinstance(floor_off, list) and all(isinstance(m, str) for m in floor_off)):
        problems.append("floor_off must be a list of the measures that have no floor at zero ([] for none)")
        floor_off = []
    else:
        for m in floor_off:
            if m not in measures_:
                problems.append("floor_off names %r, which is not a %s of the book's series" % (m, parts[1]))
        if floor_off and not raw["workbook"].get("floor_at_zero", True):
            problems.append("floor_off names measures, but the workbook has no floor at zero to switch off")
    large = raw["large_banks"]
    if large is not None:
        _keys("large_banks", large, ("file",), problems)
        if isinstance(large, dict) and "file" in large:
            large = {"file": resolve(large["file"])}      # a cache file: its presence is checked at build
    signals = raw["signals"]
    if signals is not None:
        from . import signals_data as SD
        _keys("signals", signals, SD.KEYS, problems)
        if isinstance(signals, dict) and all(k in signals for k in SD.KEYS):
            signals = {k: resolve(signals[k]) for k in SD.KEYS}
            for k in ("nmdb", "fico"):
                if not signals[k].exists():
                    problems.append("signals: the file %s does not exist (data/signals holds the committed ones)" % signals[k])
    reporting = raw["reporting_changes"]
    changes = None
    if reporting is None:
        if peers is not None:
            problems.append("reporting_changes: the peers build reads the Call Report, so name the registry of its "
                            "changes (config/reporting-changes.json); null is for a book with no Call Report series")
    elif not isinstance(reporting, str) or not reporting.strip():
        problems.append("reporting_changes must be null or the registry's file, not %r" % (reporting,))
        reporting = None
    else:
        from . import reporting_changes as RC
        reporting = resolve(reporting)
        try:
            changes = RC.load(reporting)
        except RC.RegistryError as e:
            problems.extend("reporting_changes: %s" % p for p in str(e).splitlines())
    if problems:
        raise BookError("\n".join(problems))
    wbk = dict(raw["workbook"])
    wbk["half_lives"] = tuple(wbk["half_lives"])
    wbk["percentiles"] = tuple(wbk["percentiles"])
    return {"workbook": wbk, "series": series, "peers": peers, "folder": folder, "internal": internal,
            "loan_classes": loan_classes, "reporting_changes": reporting, "changes": changes,
            "large_banks": large, "signals": signals,
            "decide": {"names": list(parts), "parts": decide_parts, "absent": absent_out, "compare": compare,
                       "notes": notes, "lines_from": lines_from, "floor_off": list(floor_off)}}


def _lines_from(spec, parts, sources, internal, problems) -> Optional[str]:
    """The book's lines_from, checked: which part is the source (the third),
    and the source every row starts as, or None for blank."""
    _keys("lines_from", spec, LINES_FROM_KEYS, problems)
    if not isinstance(spec, dict) or any(k not in spec for k in LINES_FROM_KEYS):
        return None
    if spec["part"] != parts[2]:
        problems.append("lines_from: part must be %r, the third of parts: the source is the part each row picks "
                        "its lines from, and it is listed third" % parts[2])
    start = spec["start"]
    allowed = set(sources) | ({internal["source"]} if internal else set())
    if start is not None and start not in allowed:
        problems.append("lines_from: start must be null (every row blank, for the bank to pick) or one of %s, "
                        "not %r" % (", ".join(sorted(allowed)), start))
        return None
    return start


def _internal(spec, resolve, decide_parts, problems) -> Optional[dict]:
    """The book's internal section, checked: None, or the file read by
    ``internal.read`` (every one of its problems listed here) and the
    choices, each stated. Which workbook measure an internal measure is, is a
    judgment the book states; one it does not state stays on the Internal
    tab only."""
    if spec is None:
        return None
    _keys("internal", spec, INTERNAL_KEYS, problems)
    if not isinstance(spec, dict) or any(k not in spec for k in INTERNAL_KEYS):
        return None
    from . import internal as I
    products = {v[0] for v in decide_parts.values()}
    measures = {v[1] for v in decide_parts.values()}
    sources = {v[2] for v in decide_parts.values()}
    out = {"file": resolve(spec["file"]), "source": spec["source"], "segments": spec["segments"],
           "measures": spec["measures"], "call_report_source": spec["call_report_source"], "data": None}
    src = spec["source"]
    if not (isinstance(src, str) and src.strip() and label_ok(src) and " - " not in src):
        problems.append("internal: source must name the internal series' source, for example \"KeyBank "
                        "internal\" (%s, and no ' - ')" % LABEL_RULE)
    elif src in sources:
        problems.append("internal: source %r is already a source of the book's series; name the internal one "
                        "differently" % src)
    if not isinstance(spec["segments"], bool):
        problems.append("internal: segments must be true (each segment also becomes a series) or false")
    ms = spec["measures"]
    if not isinstance(ms, dict):
        problems.append("internal: measures must be an object from an internal measure (%s) to the workbook "
                        "measure it is; {} puts none on the views" % ", ".join(I.MEASURES))
    else:
        for k, v in ms.items():
            if k not in I.MEASURES:
                problems.append("internal: measures has %r, which is not an internal measure (%s)"
                                % (k, ", ".join(I.MEASURES)))
            elif v not in measures:
                problems.append("internal: measures maps %s to %r, which is not a measure of the book's series "
                                "(%s)" % (k, v, ", ".join(sorted(measures))))
        if len(set(ms.values())) != len(ms):
            problems.append("internal: measures maps two internal measures to one workbook measure")
    if spec["call_report_source"] not in sources:
        problems.append("internal: call_report_source must be the source holding the bank's Call Report series, "
                        "one of %s" % ", ".join(sorted(sources)))
    if not out["file"].exists():
        problems.append("internal: the file %s does not exist" % out["file"])
        return out
    try:
        out["data"] = I.read(out["file"])
    except I.Refused as e:
        problems.extend("internal file: " + p for p in e.problems)
        return out
    if spec["segments"] is True:
        for p, seg in out["data"].keys():
            if seg != I.ALL and not (label_ok(seg) and " - " not in seg):
                problems.append("internal: segment %r of %s can't name a source (%s, and no ' - ')"
                                % (seg, p, LABEL_RULE))
    if not any(p in products for p, _ in out["data"].keys()):
        problems.append("internal: the file's products (%s) are none of the book's (%s)"
                        % (", ".join(sorted({p for p, _ in out["data"].keys()})), ", ".join(sorted(products))))
    return out


def internal_series(bk: dict, out_dir: Path) -> List[Tuple[str, Path, Tuple[str, str, str]]]:
    """Write the book's internal series to ``out_dir`` and return ``[(label,
    path, parts)]``: the composite of every product the book has, for every
    measure the book maps, under the internal source; and each segment too,
    under "SOURCE: SEGMENT", when the book says so. By the public method
    (``internal.composite_series`` and ``internal.series``), each from its
    latest unbroken run, with notes saying what it is."""
    from . import internal as I
    from .peers.books import write_series
    from .peers.run import series_note, slug, unbroken_tail
    spec = bk["internal"]
    data = spec["data"]
    products = {v[0] for v in bk["decide"]["parts"].values()}
    out_dir.mkdir(parents=True, exist_ok=True)
    made = []
    for p in dict.fromkeys(pr for pr, _ in data.keys()):
        if p not in products:
            continue
        segs = [s for pr, s in data.keys() if pr == p and s != I.ALL] if spec["segments"] else []
        for key, measure in spec["measures"].items():
            m = I.MEASURES[key]
            for seg in [None] + segs:
                if seg is None:
                    pairs = I.composite_series(data, p, key)
                    src, what = spec["source"], ("the composite: %s" % (
                        "the All rows" if I.composite_segment(data, p) == I.ALL else "every segment summed"))
                else:
                    pairs = I.series(data, p, seg, key)
                    src, what = "%s: %s" % (spec["source"], seg), "segment %s: %s" % (seg, data.segments[seg])
                tail, before, after = unbroken_tail(pairs)
                if not tail:
                    continue
                path = out_dir / ("internal_%s_%s_%s.csv" % (slug(p), slug(seg or "composite"), key))
                note = ["measure: %s. %s, %s, internal" % (m.title, p, what), "numerator: %s" % m.numerator,
                        "denominator: %s" % m.denominator,
                        "basis: %s. The file's own basis: %s" % (m.basis, data.notes["basis"]),
                        "as-of: %s" % data.notes["as-of"], "source: %s" % data.notes["source"],
                        "entity: %s" % data.notes["entity"],
                        "Written by 'book' from %s. Internal data: run on site, never committed."
                        % Path(data.path).name]
                write_series(tail, path, series_note(note, tail, before, after))
                made.append(("%s - %s - %s" % (p, measure, src), path, (p, measure, src)))
    return made


def build(book_path, out, say=print) -> dict:
    """Build the book's workbook at ``out``. With a peers section, the peers
    build runs first, into the same workbook, and writes the series CSVs the
    book then reads."""
    from openpyxl import Workbook

    from . import workbook as WB

    bk = load(book_path)
    out = Path(out)
    host = Workbook()
    blank = host.active
    peer_tabs = peer_info = left_out = None
    signals = None
    if bk["peers"] is not None:
        from .peers import run as peers_run
        built = peers_run.build(None, into=host, names=PEER_TAB_NAMES, say=say,
                                reporting_changes=bk["reporting_changes"], **bk["peers"])
        peer_info = built["info"]
        peer_tabs = peer_info["sheets"]
        if bk["large_banks"] is not None:
            # The Philadelphia Fed's measures, written beside the peer series from the file fetched.
            from .peers import large_banks as LB
            if not bk["large_banks"]["file"].exists():
                raise BookError("large_banks: the file %s does not exist; run 'peers large-banks-pull' first"
                                % bk["large_banks"]["file"])
            for parts_, path in LB.write(bk["large_banks"]["file"], bk["peers"]["series_dir"], say=say):
                lbl = " - ".join(parts_)
                if lbl not in dict(bk["series"]):
                    raise BookError("large_banks: the book's series list does not list %r (file %s)" % (lbl, path))
        missing = [(lbl, p) for lbl, p in bk["series"] if not p.exists()]
        if missing:
            raise BookError("\n".join("the file for %r was not written by the peers build or the large banks file either: %s" % m
                                      for m in missing))
        if bk["signals"] is not None:
            # The public series behind the Signals tab: the files named, and the
            # by-product dollars from the same cache the peers build read.
            from . import signals_data as SD
            from .peers import books as PB
            from .peers import delinquency as PD
            from .peers import fdic as PF
            from .peers import roster as PR
            from .peers.run import charge_off_fields
            for k in ("balances", "originations", "survey"):
                if not bk["signals"][k].exists():
                    raise BookError("signals: the file %s does not exist; run 'peers signals-pull' first" % bk["signals"][k])
            pe_ = bk["peers"]
            five = [p for p in PB.CONSUMER if p in pe_["products"]]
            qs_ = PF.quarter_ends(pe_["first"], pe_["last"])
            fields_ = charge_off_fields(five) + [PB.PARENT] + [f for f in PD.fields_for(five)]
            data_ = PF.pull(qs_, list(dict.fromkeys(fields_)), pe_["cache"], fetch=lambda u: (_ for _ in ()).throw(ValueError("build never fetches")))
            signals = SD.load(bk["signals"], pe_["cache"], data_, PR.self_bank(PR.load(pe_["roster"]))["cert"], five, say=say,
                              window_products=list(pe_["products"]), flows_dir=pe_["series_dir"])
        lo_path = Path(bk["peers"]["series_dir"]) / LEFT_OUT_FILE
        left_out = (lo_path, write_left_out(lo_path, bk["series"], bk["decide"]["parts"], built["left_out"]))
        say("left out as built: %d spans, in %s" % (left_out[1], lo_path))

    def tidy(wb):
        if blank.title in wb.sheetnames and blank.max_row == 1 and blank["A1"].value is None:
            wb.remove(blank)
        if signals is not None:
            from . import signals_tab as ST
            from .peers import roster as R
            ro = R.load(bk["peers"]["roster"])
            grid_ = [wb["Data"].cell(r, 1).value for r in range(2, wb["Data"].max_row + 1) if wb["Data"].cell(r, 1).value]
            ST.build(wb, labels=[lbl for lbl, _ in series], parts=decide["parts"], grid=grid_, signals=signals,
                     peers=R.peers(ro), include=peer_info["layout"]["include"], peers_sheet=PEER_TAB_NAMES["Peers"],
                     data_sheet="Data", dq_sheet=PEER_TAB_NAMES["DelinquencyData"], self_short=R.self_bank(ro)["short"],
                     latest=grid_[-1])
            _signals_on_run(wb, signals)
            say("signals: the Signals tab, %d sections, %d series on Signals data" % (len(ST.SIGNALS), len(signals["series"])))
        # Round 3: Backtest's columns join Lines by source, one row per series on both.
        if "Backtest" in wb.sheetnames and "Lines by source" in wb.sheetnames:
            WB.fold_backtest(wb)
        # Work (ink), then reference (stone), then data (mist); the peers' filed
        # figures and the workings hidden, read only by formulas.
        work = ["Summary", "Decide", "Signals", "Peers", "Thresholds", "Internal", "Synthesis"]
        reference = ["Lines by source", "Peer books", "Evidence", "Glossary", "Sources", "Run"]
        data = ["Data", "Signals data", "Internal data", "SynthesisData"]
        hidden = ["Peers data", "Peers delinquency data", "Assess", "Chart", "Peers workings", "Signals workings"]
        order = work + reference + data + hidden
        known = [wb[t] for t in order if t in wb.sheetnames]
        wb._sheets = known + [ws for ws in wb._sheets if ws not in known]
        for names, colour in ((work, WB.K.TAB_WORK), (reference, WB.K.TAB_REFERENCE), (data, WB.K.TAB_DATA)):
            for t in names:
                if t in wb.sheetnames:
                    wb[t].sheet_properties.tabColor = colour
        # Without an internal file Internal and Internal data are empty: hidden,
        # not removed, since Decide's formulas still name them.
        for t in ["Peers data", "Peers delinquency data"] + (["Internal", "Internal data"] if bk["internal"] is None else []):
            if t in wb.sheetnames:
                wb[t].sheet_state = "hidden"
        wb.active = wb.sheetnames.index("Summary")
        if peer_info is not None:
            _live_peer_group(wb, series, peer_info, say)
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = ws.title == "Summary"

    series, decide = list(bk["series"]), dict(bk["decide"])
    if peer_info is not None:
        # Each peer's switch is set on Decide, under the chart (decide_tab.peer_switches).
        pl_ = peer_info["layout"]
        decide["peers"] = {"names": list(pl_["include"]), "sheet": PEER_TAB_NAMES["Peers"], "include": pl_["include"],
                           "states": pl_["states"], "pack": pl_["pickers"]["pack"]}
    signals = signals if bk["signals"] is not None and bk["peers"] is not None else None
    if bk["internal"] is not None:
        # The bank's own figures, on site: their series sit beside the
        # workbook, never in the repository, and the rows fill Internal data.
        from .peers import roster as R
        spec = bk["internal"]
        made = internal_series(bk, out.parent / (out.stem + " internal series"))
        taken = {lbl.casefold() for lbl, _ in series}
        for lbl, path, parts_ in made:
            if lbl.casefold() in taken:
                raise BookError("internal: the series %r is already in the book" % lbl)
            series.append((lbl, path))
        decide["parts"] = dict(decide["parts"], **{lbl: parts_ for lbl, _, parts_ in made})
        tab_keys = ("net-charge-off-rate", "gross-charge-off-rate", "dpd-30-plus-quarter-end",
                    "dpd-60-plus-quarter-end", "dpd-90-plus-quarter-end")
        decide["internal"] = {
            "rows": list(spec["data"].rows), "source": spec["source"],
            "measures": {k: spec["measures"][k] for k in tab_keys if k in spec["measures"]},
            "call_report": spec["call_report_source"],
            "peers_self": R.self_bank(R.load(bk["peers"]["roster"]))["short"] if bk["peers"] else None}
        say("internal: %d rows from %s; %d series added under %s"
            % (len(spec["data"].rows), Path(spec["data"].path).name, len(made), spec["source"]))

    if bk["loan_classes"] is not None:
        from .peers import loan_classes_tab as LC
        spec = LC.load(bk["loan_classes"]["file"])
        ws = host.create_sheet(LC.TITLE)
        ws.sheet_properties.tabColor = "0A0908"
        LC.write(ws, spec, bk["loan_classes"]["cache"])
        say("loan classes: %d banks, from their %s filings" % (len(spec["banks"]), spec["quarter"]))

    reporting = None
    if bk["changes"] is not None:
        # The Call Report changes are the last section of Sources.
        from . import reporting_changes as RC
        reporting = RC.for_sources(bk["changes"], bk["reporting_changes"])
        say("reporting changes: %d, from %s, on Sources" % (len(bk["changes"]), Path(bk["reporting_changes"]).name))

    info = WB.build(out, series, decide=decide, into=host, before_save=tidy, floor_off=bk["decide"]["floor_off"],
                    left_out_file=left_out[0] if left_out else None, reporting=reporting, **bk["workbook"])
    if peer_tabs:
        # The peer charts draw a missing quarter as a gap, which is set on the saved file.
        from .peers.peers_workbook import na_as_gap
        na_as_gap(out, peer_tabs)
    return info
