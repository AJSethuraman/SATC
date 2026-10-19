"""The Assess tab: where the bank works, laid out to Claude Design's hand-off.

The design (13 Oct 2026, `docs/design/HANDOFF.md`): on one screen, without
scrolling, the product, its latest score, the Check, the chart with its five
bands, and every input that moves it. Two columns: the chart and the evidence
for it on the left (B:L); the scores, Your settings and Periods left out on the
right (N:U). Below the fold: What happens if, what the lines rest on, what each
setting does, and a dial for each setting with its chart beside it.

Every number is a formula. The cells are placed by `layout()`, which the
builder, the tests and the walkthrough all read.
"""
from __future__ import annotations

from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from . import keybank_style as K

RATINGS = ("Low", "Moderate-Low", "Moderate", "Moderate-High", "High")
SHORT = ("1 Low", "2 Mod-Low", "3 Moderate", "4 Mod-High", "5 High")
SET_TOP = 19                 # Your settings: one row per product from here
LO_ROWS = 12                 # Periods left out: this many rows
HC = 26                      # hidden helper columns start at Z

# Tokens from the hand-off that keybank_style does not have yet.
INPUT = "FFF2CC"
INPUT_EDGE = "D6B656"
PICKED = "F6D9A8"
BAND = ("F4F1EC", "E4DFD5", "EFD7CF", "E0A6A6", "CC8080")
LINE = ("8E877B", "B98577", "B35A5A", "960019")            # 2, 3, 4, 5 from
HAIR = "E7E2D8"


def fill(hex_):
    return PatternFill("solid", fgColor=hex_, bgColor=hex_)


EDGE = Side(style="thin", color=INPUT_EDGE)
BOX = Border(left=EDGE, right=EDGE, top=EDGE, bottom=EDGE)
HAIRLINE = Border(bottom=Side(style="thin", color=HAIR))
F_TITLE = Font(name="Arial", bold=True, size=16, color=K.INK)
F_SUB = Font(name="Calibri", size=11, color=K.SLATE)
F_SECTION = Font(name="Arial", bold=True, size=11, color=K.INK)
F_HEAD = Font(name="Arial", bold=True, size=9, color=K.SLATE)
F_LABEL = Font(name="Arial", bold=True, size=9, color=K.SLATE)
F_DATA = Font(name="Calibri", size=11, color=K.INK_TEXT)
F_INPUT = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
F_NOTE = Font(name="Calibri", size=10, color=K.SLATE)
F_BIG = Font(name="Calibri", size=14, bold=True, color=K.INK_TEXT)
WRAP = Alignment(wrap_text=True, vertical="center")


def layout(n: int) -> dict:
    """Where everything sits on Assess for n products."""
    set_end = SET_TOP + n - 1
    lo_top = set_end + 4
    lo_end = lo_top + LO_ROWS - 1
    low = max(40, lo_end + 3)                  # below the fold
    dial_title = low + 23
    blocks = [dial_title + 1 + 8 * i for i in range(3)]
    return dict(picker="B5", latest_row=5, check_row=6, score_head=9, score_top=10,
                set_title=SET_TOP - 2, set_head=SET_TOP - 1, set_top=SET_TOP, set_end=set_end,
                lo_title=lo_top - 2, lo_head=lo_top - 1, lo_top=lo_top, lo_end=lo_end, lo_note=lo_end + 1,
                ev_top=33, scen_title=low, scen_head=low + 1, scen_yours=low + 2, scen_all=low + 3,
                scen_chart=low + 5, rest_title=low, rest_head=low + 1, rest_top=low + 2,
                expl_title=low + 7, expl_head=low + 8, expl_top=low + 9, dial_title=dial_title,
                dial_blocks=blocks,
                dials={"high_z": blocks[0] + 2, "halfwidth": blocks[1] + 2, "low_step": blocks[2] + 2})


def _ln(hex_, width_pt=None, dash=None):
    lp = LineProperties(solidFill=hex_)
    if width_pt is not None:
        lp.w = int(width_pt * 12700)
    if dash:
        lp.prstDash = dash
    return lp


def _text(size=9, color=K.SLATE, bold=False):
    cp = CharacterProperties(sz=int(size * 100), b=bold, solidFill=color, latin=None)
    return RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=cp), endParaRPr=cp)])


def _quiet_axes(ch, tick_skip=None):
    for ax in (ch.x_axis, ch.y_axis):
        ax.delete = False
        ax.majorGridlines = None
        ax.spPr = GraphicalProperties(ln=_ln(K.STONE, 0.75))
        ax.txPr = _text()
    if tick_skip:
        ch.x_axis.tickLblSkip = tick_skip
        ch.x_axis.tickMarkSkip = tick_skip


def _square(ch):
    ch.roundedCorners = False
    ch.graphical_properties = GraphicalProperties(ln=LineProperties(noFill=True))


def _place(ws, chart, col1, row1, col2, row2):
    """Anchor a chart to cells, top-left of col1/row1 to top-left of
    col2/row2 (1-based), so it fills exactly those columns whatever the
    screen: sized in centimetres it ran past column L into the inputs."""
    a = TwoCellAnchor()
    a._from = AnchorMarker(col=col1 - 1, row=row1 - 1)
    a.to = AnchorMarker(col=col2 - 1, row=row2 - 1)
    chart.anchor = a
    ws.add_chart(chart)


def _section(ws, row, c1, c2, text):
    c = ws.cell(row, c1, text)
    c.font = F_SECTION
    c.alignment = Alignment(vertical="bottom")
    for col in range(c1, c2 + 1):
        ws.cell(row, col).border = Border(bottom=Side(style="medium", color=K.INK))
    ws.row_dimensions[row].height = 16.5


def _header(ws, row, c1, labels, right_from=None, fills=None):
    for i, t in enumerate(labels):
        c = ws.cell(row, c1 + i, t)
        c.font = F_HEAD
        c.fill = fill(K.CANVAS) if not fills or not fills[i] else fill(fills[i])
        c.alignment = Alignment(wrap_text=True, vertical="center",
                                horizontal="right" if right_from is not None and i >= right_from else "left")
        c.border = Border(bottom=Side(style="thin", color=K.STONE))


def _input(cell, prompt=None, dv=None):
    cell.fill = fill(INPUT)
    cell.border = BOX
    cell.font = F_INPUT


def _prompt(ws, ref, title, text):
    """An input message: Excel shows it when the cell is selected."""
    dv = DataValidation(showInputMessage=True, promptTitle=title[:32], prompt=text[:255], allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(ref)
    return dv


def build(asx, *, loaded, grid, first, last, n, higher, floor_at_zero, name, unit, moderate_halfwidth,
          low_step, high_z, on_the_line, y_min, quarter, not_rated, nf="0.000", th_cols=None):
    # Thresholds' columns (decide_tab.TH_PLAIN or TH_BOOK); the letters below
    # are the plain layout's, translated through it.
    # ``not_rated``: the series of rows the book shows and does not rate
    # (question 23), [] for none. Their lines and rating are blank on Lines
    # by source; picked here, the Check says so and the dials and "keeping
    # every quarter" show no lines either. With none, the tab is as before.
    from .decide_tab import TH_PLAIN, NOT_RATED
    tc = th_cols or TH_PLAIN
    plain = {v: k for k, v in TH_PLAIN.items()}
    lay = layout(n)
    sgn = "1" if higher else "-1"
    # A year label every so many years, so a short series shows its years
    # too (40 quarters showed two labels at a fixed 20).
    skip = 4 * max(1, round(len(grid) / 32))
    K.hide_gridlines(asx)
    asx.sheet_properties.tabColor = K.KEY_RED
    E = 4 + n                                             # last product row on Thresholds
    tr = "MATCH($B$5,Thresholds!$A$5:$A${e},0)".format(e=E)
    th = lambda c_: "INDEX(Thresholds!${c}$5:${c}${e},{m})".format(c=tc[plain[c_]], e=E, m=tr)   # noqa: E731
    st = lambda c_: "INDEX(${c}${a}:${c}${b},{m})".format(c=c_, a=lay["set_top"], b=lay["set_end"], m=tr)   # noqa
    online = st("R")
    HW, LS, HZ = st("O"), st("P"), st("Q")
    # True when the series showing is one the book does not rate.
    UNRATED = 'LEFT({m},{k})="{t}"'.format(m=th("M"), k=len(NOT_RATED), t=NOT_RATED) if not_rated else None

    # ---- grid ----
    widths = {"A": 2.71, "B": 21, "L": 47.29, "M": 2.71, "N": 21, "U": 14}
    for c_ in "CDEFGHIJK":
        widths[c_] = 9
    for c_ in "OPQRST":
        widths[c_] = 10
    for c_, w in widths.items():
        asx.column_dimensions[c_].width = w
    heights = {1: 22.5, 2: 13.5, 3: 4.5, 4: 12, 5: 22.5, 6: 25.5, 7: 7.5, 16: 7.5, 23: 7.5, 9: 22.5,
               lay["set_head"]: 36, lay["lo_head"]: 15, 15: 13.5, lay["lo_note"]: 13.5, 39: 10.5}
    for r, h in heights.items():
        asx.row_dimensions[r].height = h

    # ---- top strip ----
    asx.merge_cells("B1:U1")
    t = asx["B1"]
    t.value = "Assess · set the thresholds"
    t.font = F_TITLE
    t.alignment = Alignment(vertical="bottom")
    asx.merge_cells("B2:U2")
    s = asx["B2"]
    s.value = ("Choose what to show in B5. On the right, leave out periods that do not represent the risk and "
               "adjust Your settings. Type only in yellow cells; everything else updates as you type.")
    s.font = F_SUB
    for col in range(2, 22):
        asx.cell(2, col).border = Border(bottom=Side(style="medium", color=K.KEY_RED))
    for col, lbl in ((2, "Showing"), (4, "Latest quarter"), (6, "Latest value"), (8, "Latest rating"),
                     (12, "Measure")):
        c = asx.cell(4, col, lbl)
        c.font = F_LABEL
        if col == 6:
            c.alignment = Alignment(horizontal="right")
    pick = asx["B5"]
    pick.value = loaded[0]["label"]
    _input(pick)
    pick.font = Font(name="Calibri", size=12, bold=True, color=K.INK_TEXT)
    pv = DataValidation(type="list", formula1="=Thresholds!$A$5:$A$%d" % (4 + n), showInputMessage=True,
                        promptTitle="Showing", prompt="Pick what to show. Everything on this tab follows it.")
    asx.add_data_validation(pv)
    pv.add("B5")
    asx.merge_cells("D5:E5")
    asx["D5"] = "=%s" % th("I")
    asx["D5"].font = F_BIG
    asx.merge_cells("F5:G5")
    asx["F5"] = "=%s" % th("J")
    asx["F5"].font = F_BIG
    asx["F5"].number_format = nf
    asx["F5"].alignment = Alignment(horizontal="right")
    asx.merge_cells("H5:J5")
    asx["H5"] = '=IF({k}="","No rating: see Check",{k}&" "&{lb})'.format(k=th("K"), lb=th("L"))
    asx["H5"].font = Font(name="Arial", size=11, bold=True, color=K.INK_TEXT)
    asx["H5"].alignment = Alignment(horizontal="center", vertical="center")
    for k in range(5):
        asx.conditional_formatting.add("H5:J5", FormulaRule(formula=['LEFT($H$5,1)="%d"' % (k + 1)],
                                                             fill=fill(BAND[k])))
    asx.conditional_formatting.add("H5:J5", FormulaRule(formula=['ISNUMBER(SEARCH("No rating",$H$5))'],
                                                         fill=fill(K.ALERT_FG), font=Font(color=K.CRIMSON)))
    asx["L5"] = name if not unit or unit.lower() in name.lower() else "%s (%s)" % (name, unit)
    asx["L5"].font = Font(name="Calibri", size=11, color=K.SLATE)
    # The Check: one sentence, and a "To fix" when it refuses.
    asx.merge_cells("B6:U6")
    c = asx["B6"]
    c.value = ('=IF({m}="","Check: OK. Every setting is valid and every period left out is complete.",'
               '{m}&". To fix: "&{f})'.format(m=th("M"), f=th("U")))
    if UNRATED:
        c.value = '=IF(%s,"Check: "&%s&".",%s)' % (UNRATED, th("M"), c.value[1:])
    c.font = Font(name="Arial", size=11, bold=True, color=K.INK_TEXT)
    c.alignment = Alignment(vertical="center", indent=1, wrap_text=True)
    c.fill = fill(K.CANVAS)
    asx.conditional_formatting.add("B6:U6", FormulaRule(formula=['LEFT($B$6,8)="Refused:"'],
                                                         fill=fill(K.ALERT_FG), font=Font(color=K.CRIMSON)))

    # ---- left: section title over the chart ----
    first_q, last_q = grid[0], grid[-1]
    measure = name if not unit or unit.lower() in name.lower() else "%s (%s)" % (name, unit)
    asx["B8"] = '=$B$5&": {nm}, {a} to {b}, with the five rating bands"'.format(nm=measure, a=first_q, b=last_q)
    _section(asx, 8, 2, 12, asx["B8"].value)

    # ---- right: Scores for {product} ----
    _section(asx, 8, 14, 21, '="Ratings for "&$B$5')
    _header(asx, lay["score_head"], 14, ["Rating", "From", "To", "Quarters kept", "Share kept", "All quarters",
                                         "Share of all", "Latest"], right_from=1)
    hc = HC
    rng = lambda c_: "${c}${a}:${c}${b}".format(c=L(c_), a=first, b=last)          # noqa: E731
    kept_r, raw_r = rng(hc + 1), rng(hc + 3)
    bnd = [th(c_) for c_ in "EFGH"]
    if higher:
        ge, lt = ('">="', '"<"'), ('">"', '"<="')
    else:
        ge, lt = ('"<="', '">"'), ('"<"', '">="')
    draw = "$%s$5" % L(hc + 9)                           # 1 when the four lines are numbers in order
    st0 = lay["score_top"]
    for i in range(5):
        r = st0 + i
        c = asx.cell(r, 14, "%d %s" % (i + 1, RATINGS[i]))
        c.fill = fill(BAND[i])
        c.font = Font(name="Arial", size=9, bold=True, color=K.INK_TEXT)
        lo = "" if i == 0 else bnd[i - 1]
        hi = "" if i == 4 else bnd[i]
        asx.cell(r, 15, '=IFERROR(%s+0,"")' % lo if lo else "")
        asx.cell(r, 16, '=IFERROR(%s+0,"")' % hi if hi else "")
        for col, data in ((17, kept_r), (19, raw_r)):
            parts = []
            if lo:
                parts.append('{d},IF({o}="worse",{w},{b})&{ref}'.format(d=data, o=online, w=ge[0], b=lt[0], ref=lo))
            if hi:
                parts.append('{d},IF({o}="worse",{w},{b})&{ref}'.format(d=data, o=online, w=ge[1], b=lt[1], ref=hi))
            asx.cell(r, col, '=IF(%s<>1,"",COUNTIFS(%s))' % (draw, ",".join(parts)))
        asx.cell(r, 18, '=IF(Q{r}="","",Q{r}/SUM($Q${a}:$Q${b}))'.format(r=r, a=st0, b=st0 + 4))
        asx.cell(r, 20, '=IF(S{r}="","",S{r}/SUM($S${a}:$S${b}))'.format(r=r, a=st0, b=st0 + 4))
        asx.cell(r, 21, '=IF({k}={i},"◀ "&{q},"")'.format(k=th("K"), i=i + 1, q=th("I")))
        asx.cell(r, 21).font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
        for col in (15, 16):
            asx.cell(r, col).number_format = nf
        for col in (18, 20):
            asx.cell(r, col).number_format = "0%"
        for col in range(15, 22):
            if col != 21:
                asx.cell(r, col).font = F_DATA
            asx.cell(r, col).border = HAIRLINE
    asx.cell(st0, 14).border = Border(left=Side(style="thin", color=K.STONE), right=Side(style="thin", color=K.STONE),
                                      top=Side(style="thin", color=K.STONE), bottom=Side(style="thin", color=K.STONE))
    floor_note = " A value at or below zero is always 1 Low." if floor_at_zero and higher else ""
    asx.cell(st0 + 5, 14, '="A value exactly on a line gets the "&{o}&" rating (Your settings).{f}"'
             .format(o=online, f=floor_note)).font = F_NOTE

    # ---- right: Your settings ----
    _section(asx, lay["set_title"], 14, 21, "Your settings · one row each; the one showing is highlighted")
    _header(asx, lay["set_head"], 14, ["Series", "Width of Moderate (yearly moves)", "Moderate-Low begins (yearly moves)",
                                       "How rare High is (z)", "On a line, rate it", "Check"], right_from=1)
    asx.cell(lay["set_head"], 18).alignment = Alignment(wrap_text=True, vertical="center")
    asx.cell(lay["set_head"], 19).alignment = Alignment(wrap_text=True, vertical="center")
    rule_list = DataValidation(type="list", formula1='"worse,better"', allow_blank=False, showInputMessage=True,
                               promptTitle="On a line, rate it",
                               prompt="worse or better: the rating a value exactly on a line gets.")
    asx.add_data_validation(rule_list)
    good, bad = ("below", "above") if higher else ("above", "below")
    prompts = {
        15: ("Width of Moderate", "How wide Moderate is each side of the median, in typical yearly moves. 0.5 is "
                                  "half a year's movement each way. Raise it and more quarters rate Moderate."),
        16: ("Moderate-Low begins", "How far %s the median score 2 starts, in typical yearly moves. Must be more "
                                    "than Width of Moderate. Raise it and fewer quarters rate Low." % good),
        17: ("How rare High is (z)", "High begins this many spreads %s the median (a robust z). 2.5 is rare; 3.5 is "
                                     "the usual line for an outlier. Raise it and fewer quarters rate High." % bad),
    }
    rows_ = "%d:%d" % (lay["set_top"], lay["set_end"])
    for col, (tt, pp) in prompts.items():
        _prompt(asx, "%s%d:%s%d" % (L(col), lay["set_top"], L(col), lay["set_end"]), tt, pp)
    for k, s_ in enumerate(loaded):
        rr = lay["set_top"] + k
        c = asx.cell(rr, 14, s_["label"])
        c.font = F_DATA
        for col, v in ((15, moderate_halfwidth), (16, low_step), (17, high_z), (18, on_the_line)):
            _input(asx.cell(rr, col, v))
        rule_list.add("R%d" % rr)
        asx.cell(rr, 19, '=IF(Thresholds!${m}${t}="","OK",Thresholds!${m}${t})'.format(t=5 + k, m=tc["refusal"]))
        asx.cell(rr, 19).font = Font(name="Calibri", size=11, color=K.SLATE)
    asx.conditional_formatting.add("N{a}:N{b}".format(a=lay["set_top"], b=lay["set_end"]),
                                   FormulaRule(formula=["$N{a}=$B$5".format(a=lay["set_top"])], fill=fill(PICKED),
                                               font=Font(bold=True)))
    asx.conditional_formatting.add("S{a}:S{b}".format(a=lay["set_top"], b=lay["set_end"]),
                                   FormulaRule(formula=['LEFT($S{a},8)="Refused:"'.format(a=lay["set_top"])],
                                               font=Font(bold=True, color=K.CRIMSON)))

    # ---- right: Periods left out ----
    _section(asx, lay["lo_title"], 14, 21, "Periods left out · which series, first and last quarter, and why")
    _header(asx, lay["lo_head"], 14, ["Series", "From", "To", "Reason", "", "", "Status", ""])
    cal = "Data!$A${a}:$A${b}".format(a=first, b=last)
    prod_list = DataValidation(type="list", formula1="=Thresholds!$A$5:$A$%d" % (4 + n), showInputMessage=True,
                               promptTitle="Series", prompt="Pick the series this period is left out of.")
    q_list = DataValidation(type="list", formula1="=" + cal, showInputMessage=True, promptTitle="Quarter",
                            prompt="Pick a quarter from the list of quarters in the data.")
    for dv_ in (prod_list, q_list):
        asx.add_data_validation(dv_)
    _prompt(asx, "Q%d:Q%d" % (lay["lo_top"], lay["lo_end"]), "Reason",
            "Say why this period does not represent the risk. It is recorded on Run.")
    names = "Thresholds!$A$5:$A${e}".format(e=E)
    for r in range(lay["lo_top"], lay["lo_end"] + 1):
        for col in (14, 15, 16, 17):
            _input(asx.cell(r, col))
        # The Reason box runs on through R and S (nothing reads them): room
        # for the text without a merge, which would stop Delete clearing a row.
        for col in (18, 19):
            asx.cell(r, col).fill = fill(INPUT)
            asx.cell(r, col).border = Border(top=EDGE, bottom=EDGE, right=EDGE if col == 19 else None)
        asx.cell(r, 17).border = Border(top=EDGE, bottom=EDGE, left=EDGE)
        prod_list.add("N%d" % r)
        q_list.add("O%d:P%d" % (r, r))
        z, aa, ab, ac = (L(hc + 14 + i) for i in range(4))
        asx.cell(r, hc + 14, '=IFERROR(MATCH(O{r},{cal},0),"")'.format(r=r, cal=cal))
        asx.cell(r, hc + 15, '=IFERROR(MATCH(P{r},{cal},0),"")'.format(r=r, cal=cal))
        asx.cell(r, hc + 16, '=IF(AND(ISNUMBER(MATCH(N{r},{nm},0)),ISNUMBER({z}{r}),ISNUMBER({aa}{r}),'
                             '{z}{r}<={aa}{r},TRIM(Q{r})<>""),1,0)'.format(r=r, nm=names, z=z, aa=aa))
        asx.cell(r, hc + 17, '=IF(OR(N{r}<>"",O{r}<>"",P{r}<>"",TRIM(Q{r})<>""),1,0)'.format(r=r))
        missing = ('SUBSTITUTE(TRIM(IF(NOT(ISNUMBER({z}{r})),"From ","")&IF(NOT(ISNUMBER({aa}{r})),"To ","")'
                   '&IF(TRIM(Q{r})="","a_reason ",""))," ",", ")'.format(z=z, aa=aa, r=r))
        asx.cell(r, 20, '=IF({ac}{r}=0,"",IF({ab}{r}=1,"Left out: "&({aa}{r}-{z}{r}+1)&" quarters",'
                        'IF(NOT(ISNUMBER(MATCH(N{r},{nm},0))),"Pick a series",'
                        'IF(AND(ISNUMBER({z}{r}),ISNUMBER({aa}{r}),TRIM(Q{r})<>""),"From is after To ("&N{r}&" refused)",'
                        '"Missing: "&SUBSTITUTE({mi},"a_reason","reason")&" ("&N{r}&" refused)"))))'
                 .format(r=r, nm=names, z=z, aa=aa, ab=ab, ac=ac, mi=missing))
        asx.cell(r, 20).font = Font(name="Calibri", size=11, color=K.SLATE)
    asx.conditional_formatting.add("T{a}:U{b}".format(a=lay["lo_top"], b=lay["lo_end"]),
                                   FormulaRule(formula=['ISNUMBER(SEARCH("refused",$T{a}))'.format(a=lay["lo_top"])],
                                               fill=fill(K.ALERT_FG), font=Font(bold=True, color=K.CRIMSON)))
    asx.cell(lay["lo_note"], 14, "Left-out quarters show as gray dots over a dark strip. To put a period back, "
                                 "clear its yellow cells.").font = F_NOTE

    # ---- left: the evidence, level with the periods it argues for ----
    srow0, send = lay["set_top"], lay["set_end"]
    ev_row = "MATCH($B$5,$N${a}:$N${b},0)".format(a=srow0, b=send)
    asx.cell(lay["ev_top"], 2, '="What the evidence says about "&$B$5')
    _section(asx, lay["ev_top"], 2, 12, asx.cell(lay["ev_top"], 2).value)
    for i, (lbl, col) in enumerate((("Suggested", hc + 18), ("Departures", hc + 19), ("Why", hc + 20))):
        asx.cell(lay["ev_top"] + 1 + i, 2, lbl).font = F_LABEL
        asx.cell(lay["ev_top"] + 1 + i, 3, '=IFERROR(INDEX(${c}${a}:${c}${b},{m}),"")'
                 .format(c=L(col), a=srow0, b=send, m=ev_row)).font = F_DATA
    for k, s_ in enumerate(loaded):
        rr = srow0 + k
        lg, nx = s_["ev"]["largest"], s_["ev"]["next_unusual"]
        if lg and nx:
            end_ = lg["end"]
            last_q2 = grid[grid.index(quarter(end_)) - 1] if end_ else grid[-1]
            say = "Leave out %s to %s" % (quarter(lg["start"]), last_q2)
            why = ("Another stretch is also unusual (from %s, z %.1f), so the worst is not the only record of stress"
                   % (quarter(nx["start"]), nx["robust_z"]))
        else:
            say = "Keep every period"
            why = ("The worst stretch (peak %s, z %.1f) is the only unusual one, so it is the only record of stress"
                   % (quarter(lg["peak_date"]), lg["robust_z"]) if lg else "No stretch worse than the median")
        deps = "; ".join("%s to %s %s (z %+.1f)" % (quarter(d["first"]), quarter(d["last"]),
                                                    d["direction"], d["robust_z"])
                         for d in s_["ev"]["departures"]) or "None found"
        for col, v in ((hc + 18, say), (hc + 19, deps), (hc + 20, why)):
            asx.cell(rr, col, v)

    # ---- the main chart's helpers (hidden), one row per quarter ----
    order = list(range(5)) if higher else list(range(4, -1, -1))
    heads = ["Quarter", "Kept", "Quarter left out", "Every quarter"] + ["%d %s" % (i + 1, RATINGS[i]) for i in order]
    for j, h in enumerate(heads):
        asx.cell(first - 1, hc + j, h)
    asx.cell(first - 1, hc + 9, "Top")
    asx.cell(first - 1, hc + 10, " ")
    pk = "MATCH($B$5,Data!$B$1:${c}$1,0)".format(c=L(1 + n))
    sc = lambda r_: "$%s$%d" % (L(15), r_)                # noqa: E731  the "From" column of the score table
    b_ = [sc(st0 + 1), sc(st0 + 2), sc(st0 + 3), sc(st0 + 4)]
    if not higher:
        b_ = b_[::-1]
    base = "${c}${r0}".format(c=L(hc + 9), r0=first + 1)
    top = "${c}${r0}".format(c=L(hc + 9), r0=first)
    latest_q = th("I")
    for r in range(first, last + 1):
        asx.cell(r, hc, "=Data!A%d" % r)
        raw = "INDEX(Data!$B{r}:${c}{r},{m})".format(r=r, c=L(1 + n), m=pk)
        kept = "INDEX(Data!${a}{r}:${b}{r},{m})".format(a=L(2 + n), b=L(1 + 2 * n), r=r, m=pk)
        asx.cell(r, hc + 1, '=IF(ISNUMBER({k}),{k},NA())'.format(k=kept))
        asx.cell(r, hc + 2, '=IF(AND(ISNUMBER({raw}),NOT(ISNUMBER({k}))),{raw},NA())'.format(raw=raw, k=kept))
        asx.cell(r, hc + 3, '=IF(ISNUMBER({raw}),{raw},NA())'.format(raw=raw))
        asx.cell(r, hc + 12, '=IF(ISNUMBER({raw}),{raw},"")'.format(raw=raw))
        asx.cell(r, hc + 10, '=IF({d}<>1,NA(),{lo})'.format(d=draw, lo=base))
        asx.cell(r, hc + 4, '=IF({d}<>1,NA(),{b0}-{lo})'.format(d=draw, b0=b_[0], lo=base))
        for i in range(1, 4):
            asx.cell(r, hc + 4 + i, '=IF({d}<>1,NA(),{b1}-{b0})'.format(d=draw, b1=b_[i], b0=b_[i - 1]))
        asx.cell(r, hc + 8, '=IF({d}<>1,NA(),MAX({t}-{b3},0))'.format(d=draw, t=top, b3=b_[3]))
        asx.cell(r, hc + 29, '=IF({q}=Data!A{r},{raw},NA())'.format(q=latest_q, r=r, raw=raw))
    cuts = "${c}${a}:${c}${b}".format(c=L(15), a=st0 + 1, b=st0 + 4)
    axis_r = rng(hc + 12)
    asx.cell(first - 1, hc + 12, "For the axis")
    asx.cell(first, hc + 9, "=MAX(MAX({r}),MAX({c}))+0.08*(MAX(MAX({r}),MAX({c}))-{lo})".format(r=axis_r, c=cuts,
                                                                                            lo=base))
    if floor_at_zero:
        asx.cell(first + 1, hc + 9, "=MIN(0,MIN(%s),MIN(%s))" % (axis_r, cuts))
    else:
        mn, mx = "MIN(MIN(%s),MIN(%s))" % (axis_r, cuts), "MAX(MAX(%s),MAX(%s))" % (axis_r, cuts)
        asx.cell(first + 1, hc + 9, "=%s-0.08*(%s-%s)" % (mn, mx, mn))
    h_c = "${c}${r}".format(c=L(hc + 9), r=first + 2)
    asx.cell(first + 2, hc + 9, "=({t}-{lo})*0.06".format(t=top, lo=base))
    asx.cell(5, hc + 9, "=IF(AND(ISNUMBER({a}),ISNUMBER({b}),ISNUMBER({c}),ISNUMBER({d}),"
                        "{s}*({b}-{a})>=0,{s}*({c}-{b})>=0,{s}*({d}-{c})>=0),1,0)"
             .format(a=sc(st0 + 1), b=sc(st0 + 2), c=sc(st0 + 3), d=sc(st0 + 4), s=sgn))
    asx.cell(first - 1, hc + 11, "Period left out")
    asx.cell(first - 1, hc + 13, " ")
    asx.cell(first - 1, hc + 29, "Latest")
    for r in range(first, last + 1):
        asx.cell(r, hc + 13, '={lo}-{h}'.format(lo=base, h=h_c))
        asx.cell(r, hc + 11, '=IF(ISNUMBER({g}{r}),{h},0)'.format(g=L(hc + 2), r=r, h=h_c))

    # ---- the main chart ----
    line = LineChart()
    line.add_data(Reference(asx, min_col=hc + 3, max_col=hc + 3, min_row=first - 1, max_row=last), titles_from_data=True)
    line.add_data(Reference(asx, min_col=hc + 2, max_col=hc + 2, min_row=first - 1, max_row=last), titles_from_data=True)
    line.add_data(Reference(asx, min_col=hc + 29, max_col=hc + 29, min_row=first - 1, max_row=last),
                  titles_from_data=True)
    line.set_categories(Reference(asx, min_col=hc, min_row=first, max_row=last))
    every, gone, latest = line.series
    every.graphicalProperties.line = _ln(K.INK, 2.25)
    every.marker.symbol = "none"
    gone.graphicalProperties.line.noFill = True
    gone.marker.symbol = "circle"
    gone.marker.size = 5
    gone.marker.graphicalProperties = GraphicalProperties(solidFill=K.STONE, ln=_ln(K.SLATE, 0.75))
    latest.graphicalProperties.line.noFill = True
    latest.marker.symbol = "circle"
    latest.marker.size = 7
    latest.marker.graphicalProperties = GraphicalProperties(solidFill="FFFFFF", ln=_ln(K.INK, 1.5))
    latest.dLbls = DataLabelList()
    latest.dLbls.showCatName = True
    latest.dLbls.showVal = True
    latest.dLbls.showSerName = False
    latest.dLbls.showLegendKey = False
    latest.dLbls.position = "l"
    latest.dLbls.txPr = _text(9, K.INK_TEXT, True)
    for s_ in line.series:
        s_.smooth = False
    area = AreaChart()
    area.grouping = "stacked"
    area.add_data(Reference(asx, min_col=hc + 10, max_col=hc + 10, min_row=first - 1, max_row=last),
                  titles_from_data=True)
    area.add_data(Reference(asx, min_col=hc + 4, max_col=hc + 8, min_row=first - 1, max_row=last),
                  titles_from_data=True)
    area.series[0].graphicalProperties.noFill = True
    area.series[0].graphicalProperties.line.noFill = True
    for i, s_ in enumerate(area.series[1:]):
        s_.graphicalProperties.solidFill = BAND[order[i]]
        s_.graphicalProperties.line.noFill = True
    strip = AreaChart()
    strip.grouping = "stacked"
    strip.add_data(Reference(asx, min_col=hc + 13, max_col=hc + 13, min_row=first - 1, max_row=last),
                   titles_from_data=True)
    strip.add_data(Reference(asx, min_col=hc + 11, max_col=hc + 11, min_row=first - 1, max_row=last),
                   titles_from_data=True)
    strip.series[0].graphicalProperties.noFill = True
    strip.series[0].graphicalProperties.line.noFill = True
    strip.series[1].graphicalProperties.solidFill = K.SLATE
    strip.series[1].graphicalProperties.line.noFill = True
    _quiet_axes(line, skip)
    below_zero = not floor_at_zero or any(p.value < 0 for s_ in loaded for p in s_["points"])
    line.y_axis.number_format = "0.0" if below_zero else "0.0;;0.0"
    line.x_axis.tickLblPos = "low"
    if y_min is not None:
        line.y_axis.scaling.min = y_min
        # No fixed axis can follow every leave-out (Codex on #432): leaving
        # out a calm stretch can widen the spread and push a line below it.
        # Say so, in the row over the chart, rather than clip in silence.
        asx.merge_cells("B7:L7")
        warn = asx["B7"]
        warn.value = ('=IF({f}-{h}<{m},"Part of a line or the left-out strip is below the axis of this chart '
                      '(fixed at build at {m}): the numbers are right, the picture is cut off.","")'
                      .format(f=base, h=h_c, m=y_min))
        warn.font = Font(name="Arial", size=10, bold=True, color=K.CRIMSON)
    line.visible_cells_only = False
    line.height, line.width = 12.3, 29.1
    # No legend: the design deletes the bands' and helpers' entries, and
    # Excel 16 ignores a deleted legend entry written by openpyxl (tested on a
    # two-line chart, 13 Oct 2026). The score table is the bands' key and the
    # row under the chart names the rest.
    line.legend = None
    line += area
    line += strip
    _square(line)
    _place(asx, line, 2, 9, 13, 32)                         # B9 to L31
    asx.cell(32, 2, "━ History     ● Left out     ○ Latest     ▬ Period left out     "
                    "Shaded bands: the five ratings, as in the table on the right").font = F_NOTE

    # ---- below the fold, left: What happens if ----
    _section(asx, lay["scen_title"], 2, 12, "What happens if · your lines against keeping every quarter")
    _header(asx, lay["scen_head"], 2, ["", "2 from", "3 from", "4 from", "5 from"] + list(SHORT) + [""],
            right_from=1, fills=[None] * 5 + list(BAND) + [None])
    th_all = lambda c_: th(c_)                                     # noqa: E731
    med_a, move_a, spread_a = th_all("Q"), th_all("S"), th_all("T")
    all_lines = ["{m}-{s}*{ls}*{mv}".format(m=med_a, s=sgn, ls=LS, mv=move_a),
                 "{m}-{s}*{hw}*{mv}".format(m=med_a, s=sgn, hw=HW, mv=move_a),
                 "{m}+{s}*{hw}*{mv}".format(m=med_a, s=sgn, hw=HW, mv=move_a),
                 "{m}+{s}*{hz}*{sp}".format(m=med_a, s=sgn, hz=HZ, sp=spread_a)]
    ra, rb = lay["scen_yours"], lay["scen_all"]
    asx.cell(ra, 2, "As you have it").font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
    asx.cell(rb, 2, "Keeping every quarter").font = F_DATA
    for i in range(4):
        asx.cell(ra, 3 + i, '=IFERROR(%s+0,"")' % sc(st0 + 1 + i))
        asx.cell(rb, 3 + i, '=IFERROR(%s,"")' % (all_lines[i] if not UNRATED else
                                                  'IF(%s,"",%s)' % (UNRATED, all_lines[i])))
    for i in range(5):
        asx.cell(ra, 7 + i, "=Q%d" % (st0 + i))
    bb = ["%s%d" % (c_, rb) for c_ in "CDEF"]
    for k in range(5):
        parts = []
        if k > 0:
            parts.append('{d},IF({o}="worse",{w},{bt})&{ref}'.format(d=raw_r, o=online, w=ge[0], bt=lt[0], ref=bb[k - 1]))
        if k < 4:
            parts.append('{d},IF({o}="worse",{w},{bt})&{ref}'.format(d=raw_r, o=online, w=ge[1], bt=lt[1], ref=bb[k]))
        asx.cell(rb, 7 + k, '=IF(COUNT(C{r}:F{r})<4,"",COUNTIFS({p}))'.format(r=rb, p=",".join(parts)))
    asx.cell(ra, 12, "Solid lines. Counts: quarters kept").font = F_NOTE
    asx.cell(rb, 12, "Dashed lines. Same settings, nothing left out").font = F_NOTE
    for rr in (ra, rb):
        for c_ in range(3, 7):
            asx.cell(rr, c_).number_format = nf
        for c_ in range(2, 13):
            asx.cell(rr, c_).border = HAIRLINE

    sc0 = hc + 21
    names_ = ["2 from", "3 from", "4 from", "5 from"]
    for i in range(4):
        asx.cell(first - 1, sc0 + i, "%s (yours)" % names_[i])
        asx.cell(first - 1, sc0 + 4 + i, "%s (all kept)" % names_[i])
    for r in range(first, last + 1):
        for i in range(4):
            asx.cell(r, sc0 + i, '=IF(ISNUMBER($%s$%d),$%s$%d,NA())' % (L(3 + i), ra, L(3 + i), ra))
            asx.cell(r, sc0 + 4 + i, '=IF(ISNUMBER($%s$%d),$%s$%d,NA())' % (L(3 + i), rb, L(3 + i), rb))
    cmp_ = LineChart()
    cmp_.add_data(Reference(asx, min_col=hc + 3, max_col=hc + 3, min_row=first - 1, max_row=last),
                  titles_from_data=True)
    cmp_.add_data(Reference(asx, min_col=sc0, max_col=sc0 + 7, min_row=first - 1, max_row=last), titles_from_data=True)
    cmp_.set_categories(Reference(asx, min_col=hc, min_row=first, max_row=last))
    cmp_.series[0].graphicalProperties.line = _ln(K.INK, 1.75)
    for i, s_ in enumerate(cmp_.series[1:]):
        s_.graphicalProperties.line = _ln(LINE[i % 4], 1.5, "sysDash" if i >= 4 else None)
    for s_ in cmp_.series:
        s_.smooth = False
        s_.marker.symbol = "none"
    _quiet_axes(cmp_, skip)
    cmp_.y_axis.number_format = "0.0"
    cmp_.x_axis.tickLblPos = "low"
    if y_min is not None:
        cmp_.y_axis.scaling.min = y_min
    cmp_.height, cmp_.width = 9.0, 29.1
    cmp_.legend = None
    cmp_.visible_cells_only = False
    _square(cmp_)
    _place(asx, cmp_, 2, lay["scen_chart"], 13, lay["scen_chart"] + 17)
    asx.cell(lay["scen_chart"] + 17, 2, "━ History     Solid: where ratings 2, 3, 4 and 5 begin as you have it     "
                                       "Dashed: the same, keeping every quarter     Grey, tan, rose, crimson: "
                                       "2, 3, 4, 5").font = F_NOTE

    # ---- below the fold, right: what the lines rest on, and what each setting does ----
    _section(asx, lay["rest_title"], 14, 21, "What the lines rest on")
    _header(asx, lay["rest_head"], 14, ["", "Quarters kept", "Every quarter"], right_from=1)
    for i, (lbl, kc, ac) in enumerate((("Median", "B", "Q"), ("Typical yearly move", "C", "S"),
                                       ("Spread", "P", "T"), ("Worst", "D", None))):
        r = lay["rest_top"] + i
        asx.cell(r, 14, lbl).font = F_LABEL
        asx.cell(r, 15, "=" + th(kc)).font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
        if ac:
            asx.cell(r, 16, "=" + th(ac)).font = Font(name="Calibri", size=11, color=K.SLATE)
        else:
            asx.cell(r, 16, "=%s(%s)" % ("MAX" if higher else "MIN", rng(hc + 12))).font = Font(
                name="Calibri", size=11, color=K.SLATE)
        for c_ in (15, 16):
            asx.cell(r, c_).number_format = nf
            asx.cell(r, c_).border = HAIRLINE
    _section(asx, lay["expl_title"], 14, 21, "What each setting does")
    _header(asx, lay["expl_head"], 14, ["Setting", "Raise it and", "", "", "", "", "", ""])
    expl = [("Width of Moderate", "more quarters rate Moderate."),
            ("Moderate-Low begins", "fewer quarters rate Low. Keep it above Width of Moderate."),
            ("How rare High is", "fewer quarters rate High. 2.5 is rare; 3.5 is the usual line for an outlier."),
            ("On a line, rate it", "worse: a value exactly on a line gets the worse rating; better: the better."),
            ("A period left out", "is not used for the median, the yearly move or the spread.")]
    for i, (a, b) in enumerate(expl):
        r = lay["expl_top"] + i
        asx.cell(r, 14, a).font = F_LABEL
        asx.cell(r, 15, b).font = F_DATA
        for c_ in range(14, 22):
            asx.cell(r, c_).border = HAIRLINE

    # ---- dials, each with its chart beside it ----
    _section(asx, lay["dial_title"], 2, 21,
             "Dials · each setting at other values, the others as in Your settings. The shaded row is yours")
    WORST, MED, MOVE, SPREAD = th("D"), th("B"), th("C"), th("P")
    dials = [("High's z", HZ, (2.0, 2.5, 3.0, 3.5), "How rare High is (z)"),
             ("Width of Moderate", HW, (0.25, 0.5, 0.75, 1.0), "Width of Moderate"),
             ("Moderate-Low begins", LS, (0.75, 1.0, 1.5, 2.0), "Moderate-Low begins")]
    for (title, now, values, label), top_ in zip(dials, lay["dial_blocks"]):
        asx.cell(top_, 2, "Dial: " + label).font = Font(name="Arial", bold=True, size=10, color=K.INK)
        _header(asx, top_ + 1, 2, ["Value", "2 from", "3 from", "4 from", "5 from"] + list(SHORT) + [""],
                right_from=0, fills=[None] * 5 + list(BAND) + [None])
        asx.row_dimensions[top_ + 1].height = 22.5
        asx.row_dimensions[top_ + 7].height = 22.5
        r = top_ + 2
        for v in list(values) + [None]:
            val = now if v is None else repr(v)
            asx.cell(r, 2, "=" + now if v is None else v)
            hw_, ls_, hz_ = (val if now == x else x for x in (HW, LS, HZ))
            num = "AND(ISNUMBER({}),ISNUMBER({}),ISNUMBER({}))".format(hw_, ls_, hz_)
            if UNRATED:
                num = "AND(NOT(%s),%s)" % (UNRATED, num[4:-1])
            asx.cell(r, 3, '=IF({n},{m}-{s}*{ls}*{mv},"")'.format(n=num, m=MED, s=sgn, ls=ls_, mv=MOVE))
            asx.cell(r, 4, '=IF({n},{m}-{s}*{hw}*{mv},"")'.format(n=num, m=MED, s=sgn, hw=hw_, mv=MOVE))
            asx.cell(r, 5, '=IF({n},{m}+{s}*{hw}*{mv},"")'.format(n=num, m=MED, s=sgn, hw=hw_, mv=MOVE))
            asx.cell(r, 6, '=IF({n},{m}+{s}*{hz}*{sp},"")'.format(n=num, m=MED, s=sgn, hz=hz_, sp=SPREAD))
            b = ["%s%d" % (c_, r) for c_ in "CDEF"]
            for k in range(5):
                parts = []
                if k > 0:
                    parts.append('{d},IF({o}="worse",{w},{bt})&{ref}'.format(d=kept_r, o=online, w=ge[0], bt=lt[0],
                                                                              ref=b[k - 1]))
                if k < 4:
                    parts.append('{d},IF({o}="worse",{w},{bt})&{ref}'.format(d=kept_r, o=online, w=ge[1], bt=lt[1],
                                                                              ref=b[k]))
                asx.cell(r, 7 + k, '=IF(NOT(%s),"",COUNTIFS(%s))' % (num, ",".join(parts)))
            asx.cell(r, 12, ('=IF(NOT({n}),"refused: a setting is not a number",'
                             'IF({ls}<={hw},"refused: Moderate-Low inside Moderate",'
                             'IF({s}*(F{r}-E{r})<=0,"refused: High inside Moderate-High",'
                             'IF({s}*({w}-F{r})<0,"refused: High beyond the worst kept",{now}))))')
                     .format(n=num, ls=ls_, hw=hw_, s=sgn, r=r, w=WORST, now='"◀ yours"' if v is None else '""'))
            if UNRATED:
                asx.cell(r, 12).value = '=IF(%s,"%s",%s)' % (UNRATED, NOT_RATED.lower(), asx.cell(r, 12).value[1:])
            for c_ in range(3, 7):
                asx.cell(r, c_).number_format = nf
            for c_ in range(2, 13):
                asx.cell(r, c_).border = HAIRLINE
            r += 1
        first_val = top_ + 2
        asx.conditional_formatting.add("B{a}:L{b}".format(a=first_val, b=first_val + 4),
                                       FormulaRule(formula=['ISNUMBER(SEARCH("yours",$L{a}))'.format(a=first_val)],
                                                   fill=fill(K.MIST), font=Font(bold=True)))
        asx.conditional_formatting.add("L{a}:L{b}".format(a=first_val, b=first_val + 4),
                                       FormulaRule(formula=['ISNUMBER(SEARCH("refused",$L{a}))'.format(a=first_val)],
                                                   font=Font(color=K.CRIMSON)))
        bc = BarChart()
        bc.type = "col"
        bc.grouping = "stacked"
        bc.overlap = 100
        bc.gapWidth = 60
        bc.title = "Quarters in each rating, by %s (last bar: yours)" % label.replace(" (z)", "")
        bc.title.tx.rich.p[0].pPr = ParagraphProperties(defRPr=CharacterProperties(sz=1000, b=True,
                                                                                   solidFill=K.INK_TEXT))
        bc.add_data(Reference(asx, min_col=7, max_col=11, min_row=first_val - 1, max_row=first_val + 4),
                    titles_from_data=True)
        bc.set_categories(Reference(asx, min_col=2, min_row=first_val, max_row=first_val + 4))
        for i, s_ in enumerate(bc.series):
            s_.graphicalProperties.solidFill = BAND[i]
            s_.graphicalProperties.line = _ln(K.STONE, 0.5)
        _quiet_axes(bc)
        bc.legend = None
        bc.height, bc.width = 4.2, 16.6
        _square(bc)
        _place(asx, bc, 14, top_, 22, top_ + 8)
    asx.cell(lay["dial_blocks"][-1] + 8, 2,
             "Each dial changes one setting and keeps the others as in Your settings. A refused value still "
             "shows where its lines fall, so you can see why it is refused.").font = F_NOTE

    for col in range(hc, hc + 30):
        asx.column_dimensions[L(col)].hidden = True
    asx.freeze_panes = None
    return lay
