"""Thresholds in the book: one set of lines per product and measure.

The firm, 15 Oct 2026: *"we don't need to set thresholds on all 3, we need
to use all 3 to set appropriate thresholds. industry, competitor, and key
reported"*. Each row is a product and measure: the source whose history sets
its lines (picked on Decide), what those lines rest on, where each rating
begins, and every source's latest quarter rated on them. The scoring file
reads the four lines here.

Nothing here computes a line. Each row takes the lines of the source picked
from the Lines by source tab, which computes every source's own lines with
the row's settings; the only arithmetic is each source's latest value rated
on the lines picked, the way the lines rate a value everywhere: at or below
zero is 1 Low when the measure has a floor, otherwise one more than the
number of lines at or past the value, On a line deciding a value exactly on
one.
"""
from __future__ import annotations

from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.filters import FilterColumn

from . import decide_tab as DT
from . import keybank_style as K

TITLE = "Thresholds"
FIRST = 5
#: The fixed columns, by meaning; the sources' (Latest, Rating) pairs follow
#: from L in Summary's order, then Decision and Check, then hidden workings.
FIX = dict(key="A", product="B", measure="C", lines_from="D", median="E", move="F", worst="G", l2="H", l3="I",
           l4="J", l5="K")
NOT_PICKED = "Pick a source"
#: Widths in Excel units, from round 3's pixels ((px - 5) / 7): B 150, C 232,
#: D 146, E-K 68, each source's latest 60 and rating 116, Decision 124, Check 184.
WIDTHS = dict(A=1.57, B=20.71, C=32.43, D=20.14, E=9.0, F=9.0, G=9.0, H=9.0, I=9.0, J=9.0, K=9.0)
W_LATEST, W_RATING, W_DECISION, W_CHECK = 7.86, 15.86, 17.0, 25.57
#: Row heights: a body row; a product name on two lines (Calibri 11 bold);
#: a To fix on three lines; the header, three lines of Arial 9 so
#: "4 Moderate-High" and "Peer group latest" read whole in their columns.
H_ROW, H_WRAP2, H_WRAP3, H_HEAD = 15, 30, 34.5, 39


def _q(s: str) -> str:
    return s.replace('"', '""')


def layout(sources: Sequence[str], dropped: Sequence[str] = ()) -> dict:
    """Every column, by meaning: FIX, then per source shown (latest, rating),
    then decision and check, then the hidden workings: the row on Lines by
    source, the refusal, the fix, the On a line rule, per source shown its
    rating as a number and its quarter, and per source ``dropped`` (one
    Summary shows no columns for) all four of its cells."""
    cols = dict(FIX)
    c = 12
    per = {}
    for s in sources:
        per[s] = {"latest": L(c), "rating": L(c + 1)}
        c += 2
    cols["decision"], cols["check"] = L(c), L(c + 1)
    c += 2
    hidden_from = c
    for key in ("row", "refusal", "fix", "rule", "picked", "refused"):
        cols[key] = L(c)
        c += 1
    for s in sources:
        per[s]["number"], per[s]["quarter"] = L(c), L(c + 1)
        c += 2
    for s in dropped:
        per[s] = {"latest": L(c), "rating": L(c + 1), "number": L(c + 2), "quarter": L(c + 3)}
        c += 4
    return {"cols": cols, "per": per, "hidden": (L(hidden_from), L(c - 1)), "last_shown": cols["check"],
            "shown": list(sources), "dropped": list(dropped)}


def _short(s: str) -> str:
    return s.split(" (")[0]


def build(wb, *, pairs: Sequence[Tuple[str, str]], by_pair: Dict[Tuple[str, str], Dict[str, str]],
          sources: Sequence[str], part_names: Sequence[str], decide: dict, n_series: int, name: str,
          higher: bool, floor_at_zero: bool, number_format: str = "0.00", index: Optional[int] = None,
          floor_off: Sequence[str] = (), dropped: Sequence[str] = (), refused_at_build: Sequence[str] = (),
          start: Optional[str] = None) -> dict:
    """Add Thresholds: one row per product and measure (``pairs``, in Decide's
    order). ``by_pair``: each row's label per source; ``sources``: the order
    of the source columns (Summary's); ``dropped``: sources with hidden cells
    only; ``decide``: what ``decide_tab.build`` returned. The per-series tab
    must already be titled Lines by source.

    A row's height cannot follow a formula, so it is set for the workbook as
    built: a row whose lines come, as built, from a series in
    ``refused_at_build`` (labels) is three lines high, so its To fix reads
    whole. ``start`` is the source every row's lines start from (Decide's
    Lines set from as built), or None."""
    lay = layout(sources, dropped)
    cols, per = lay["cols"], lay["per"]
    ws = wb.create_sheet(TITLE, index)
    ws.sheet_properties.tabColor = K.TAB_WORK
    ws.sheet_view.showGridLines = False
    T = DT.TH_BOOK
    E = 4 + n_series
    lbs = DT._sheet_ref(DT.LINES_TAB)
    lbs_col = lambda key: "%s$%s$5:$%s$%d" % (lbs, T[key], T[key], E)           # noqa: E731
    lbs_a = lbs_col("series")
    C = decide["columns"]
    p0, p1, p2 = part_names
    floor_off = tuple(floor_off or ())
    op_w, op_b = ("<=", "<") if higher else (">=", ">")
    last = FIRST + len(pairs) - 1
    last_col = lay["last_shown"]

    ws["B1"] = "Thresholds · one set of lines per %s and %s" % (p0.lower(), p1.lower())
    ws["B1"].font = Font(name="Arial", bold=True, size=16, color=K.INK)
    ws["B1"].alignment = Alignment(vertical="bottom")
    ws["B2"] = ("Each row's lines come from the %s under %s on Decide. The header arrows filter by %s or %s."
                % (p2.lower(), DT.LINES_FROM, p0.lower(), p1.lower()))
    ws["B2"].font = Font(name="Calibri", size=11, color=K.SLATE)
    from openpyxl.utils import column_index_from_string as CI
    n_last = CI(last_col)
    for c_ in range(2, n_last + 1):
        ws.cell(2, c_).border = Border(bottom=Side(style="medium", color=K.KEY_RED))
    for r, h in ((1, 25.5), (2, 15), (3, 16.5), (4, H_HEAD)):
        ws.row_dimensions[r].height = h
    first_src, last_src = per[sources[0]]["latest"], per[sources[-1]]["rating"]
    groups = [("B", "D", "%s and %s" % (p0, p1.lower()), None), ("E", "G", "What the lines rest on", K.SLATE),
              ("H", "K", "Each rating begins at", K.SLATE),
              (first_src, last_src, "Every %s's latest quarter on these lines" % p2.lower(), K.KEY_RED),
              (cols["decision"], cols["check"], "Your decision", K.SLATE)]
    for a, b, lbl, rule in groups:
        cell = ws["%s3" % a]
        cell.value = lbl
        cell.font = Font(name="Arial", bold=True, size=9, color=K.INK)
        cell.alignment = Alignment(vertical="bottom")
        if rule:
            for c_ in range(CI(a), CI(b) + 1):
                ws.cell(3, c_).border = Border(bottom=Side(style="medium", color=rule))
    heads = [("product", p0), ("measure", p1), ("lines_from", DT.LINES_FROM), ("median", "Median kept"),
             ("move", "Yearly move"), ("worst", "Worst kept"), ("l2", "2 Moderate-Low"),
             ("l3", "3 Moderate"), ("l4", "4 Moderate-High"), ("l5", "5 High"), ("decision", "Decision"),
             ("check", "Check")]
    right = {"median", "move", "worst", "l2", "l3", "l4", "l5"}

    def head(cell, text, align="left"):
        cell.value = text
        cell.font = Font(name="Arial", bold=True, size=9, color=K.PAPER)
        cell.fill = K.solid(K.INK)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal=align)

    for key, t_ in heads:
        head(ws["%s4" % cols[key]], t_, "right" if key in right else "left")
    for s in sources:
        head(ws["%s4" % per[s]["latest"]], "%s latest" % _short(s), "right")
        head(ws["%s4" % per[s]["rating"]], "%s rating" % _short(s))
    for key, t_ in (("row", "Row on Lines by source"), ("refusal", "Refused because"), ("fix", "To fix"),
                    ("picked", "Lines picked (1 or 0)"), ("refused", "Refused (1 or 0)"),
                    ("rule", "On a line, rate it")):
        ws["%s4" % cols[key]].value = t_
    for s in list(sources) + list(dropped):
        ws["%s4" % per[s]["number"]].value = "%s rating number" % s
        ws["%s4" % per[s]["quarter"]].value = "%s quarter" % s
    for s in dropped:
        ws["%s4" % per[s]["latest"]].value = "%s latest" % s
        ws["%s4" % per[s]["rating"]].value = "%s rating" % s
    for c_, w in WIDTHS.items():
        ws.column_dimensions[c_].width = w
    for s in sources:
        ws.column_dimensions[per[s]["latest"]].width = W_LATEST
        ws.column_dimensions[per[s]["rating"]].width = W_RATING
    ws.column_dimensions[cols["decision"]].width = W_DECISION
    ws.column_dimensions[cols["check"]].width = W_CHECK
    h1, h2 = lay["hidden"]
    for c_ in range(CI(h1), CI(h2) + 1):
        ws.column_dimensions[L(c_)].hidden = True

    body = Font(name="Calibri", size=11, color=K.INK_TEXT)
    for i, pr in enumerate(pairs):
        r = FIRST + i
        dr = decide["row_of"][pr]
        at = lambda key: "$%s%d" % (cols[key], r)                                  # noqa: E731
        put = lambda key, v: ws["%s%d" % (cols[key], r)].__setattr__("value", v)  # noqa: E731
        put("key", "%s - %s" % pr)
        put("product", pr[0])
        put("measure", pr[1])
        put("lines_from", '=IF(Decide!{c}{r}="","",Decide!{c}{r})'.format(c=L(C[DT.LINES_FROM]), r=dr))
        put("rule", '=Decide!{c}{r}&""'.format(c=L(C["On a line, rate it"]), r=dr))
        # 1 or 0, for the counts top right here and on Summary.
        put("picked", '=IF({d}="",0,1)'.format(d=at("lines_from")))
        put("refused", '=IF({rf}="",0,1)'.format(rf=at("refusal")))
        put("row", '=IF({d}="","",IFERROR(MATCH({b}&" - "&{c}&" - "&{d},{a},0),""))'.format(
            d=at("lines_from"), b=at("product"), c=at("measure"), a=lbs_a))
        put("refusal", '=IF({w}="","",INDEX({x},{w}))'.format(w=at("row"), x=lbs_col("refusal")))
        put("fix", '=IF({w}="","",INDEX({x},{w}))'.format(w=at("row"), x=lbs_col("fix")))
        for key in ("median", "move", "worst"):
            put(key, '=IF({w}="","",IFERROR(INDEX({x},{w})+0,""))'.format(w=at("row"), x=lbs_col(key)))
        # The lines stand only when nothing refuses them; a refused row has none.
        for key in ("l2", "l3", "l4", "l5"):
            put(key, '=IF(OR({w}="",{rf}<>""),"",IFERROR(INDEX({x},{w})+0,""))'.format(
                w=at("row"), rf=at("refusal"), x=lbs_col(key)))
        lines = "$%s%d:$%s%d" % (cols["l2"], r, cols["l5"], r)
        for s in list(sources) + list(dropped):
            pc = per[s]
            cell_ = lambda k: "$%s%d" % (pc[k], r)                               # noqa: E731
            if s not in by_pair[pr]:
                ws["%s%d" % (pc["rating"], r)] = None
                ws["%s%d" % (pc["latest"], r)] = None
                ws["%s%d" % (pc["number"], r)] = ""
                ws["%s%d" % (pc["quarter"], r)] = ""
                continue
            lab = _q(by_pair[pr][s])
            m = 'MATCH("%s",%s,0)' % (lab, lbs_a)
            ws["%s%d" % (pc["latest"], r)] = '=IFERROR(INDEX({v},{m})+0,"")'.format(v=lbs_col("value"), m=m)
            ws["%s%d" % (pc["quarter"], r)] = '=IFERROR(INDEX({q},{m})&"","")'.format(q=lbs_col("quarter"), m=m)
            x = cell_("latest")
            ws["%s%d" % (pc["number"], r)] = (
                '=IF(OR(COUNT({ln})<4,NOT(ISNUMBER({x}))),"",IF(AND({fz},{x}<=0),1,1+IF({rule}="worse",'
                'COUNTIF({ln},"{w}"&{x}),COUNTIF({ln},"{b}"&{x}))))'
                .format(ln=lines, x=x, fz="TRUE" if (floor_at_zero and higher and pr[1] not in floor_off) else "FALSE",
                        rule=at("rule"), w=op_w, b=op_b))
            k = cell_("number")
            ws["%s%d" % (pc["rating"], r)] = (
                '=IF(OR({d}="",{w}=""),"",IF({rf}<>"","Refused",IF({k}="","",'
                '{k}&" "&CHOOSE({k},"{r0}","{r1}","{r2}","{r3}","{r4}"))))'
                .format(d=at("lines_from"), w=at("row"), rf=at("refusal"), k=k, r0=DT.RATINGS[0],
                        r1=DT.RATINGS[1], r2=DT.RATINGS[2], r3=DT.RATINGS[3], r4=DT.RATINGS[4]))
        # A row with no source picked is blank from Lines set from to Check,
        # but for each source's latest quarter and a decision typed on Decide.
        put("decision", '=IF(OR(Decide!{c}{r}="",AND({d}="",Decide!{c}{r}="{nd}")),"",Decide!{c}{r})'.format(
            d=at("lines_from"), c=L(C["Decision"]), r=dr, nd=DT.DECISIONS[0]))
        put("check", '=IF({d}="","",IF({w}="","This workbook has no "&{d}&'
                     '" series for this row: pick another source on Decide.",IF({rf}="","OK","To fix: "&{fix})))'
            .format(d=at("lines_from"), w=at("row"), rf=at("refusal"), fix=DT.HERE.format(m=at("fix"))))
        for key in ("median", "move", "worst", "l2", "l3", "l4", "l5"):
            ws["%s%d" % (cols[key], r)].number_format = number_format
        for s in list(sources) + list(dropped):
            ws["%s%d" % (per[s]["latest"], r)].number_format = number_format
        for c_ in range(2, n_last + 1):
            ws.cell(r, c_).font = body
            ws.cell(r, c_).alignment = Alignment(vertical="center")
        ws["%s%d" % (cols["product"], r)].font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
        if sources:
            ws["%s%d" % (per[sources[0]]["latest"], r)].font = Font(name="Calibri", size=11, bold=True,
                                                                     color=K.INK_TEXT)
        ws["%s%d" % (cols["check"], r)].font = Font(name="Calibri", size=10, color=K.SLATE)
        # A To fix that arrives after the build, on a row one line high, shows its first line.
        ws["%s%d" % (cols["check"], r)].alignment = Alignment(vertical="top", wrap_text=True)
        prev = pairs[i - 1] if i else None
        # The height as built: a combined product's name wraps on its first
        # row; a row refused as built shows its To fix whole.
        height = H_ROW
        if (prev is None or pr[0] != prev[0]) and len(pr[0]) > WIDTHS["B"] * 1.1:
            height = H_WRAP2
            ws["%s%d" % (cols["product"], r)].alignment = Alignment(vertical="center", wrap_text=True)
            ws["%s%d" % (cols["check"], r)].alignment = Alignment(vertical="center", wrap_text=True)
        if start is not None and by_pair[pr].get(start) in set(refused_at_build):
            height = H_WRAP3
            ws["%s%d" % (cols["check"], r)].alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = height
        side = None
        if prev is None or pr[0] != prev[0]:
            side = Side(style="thin", color=K.INK)
        elif pr[1] != prev[1]:
            side = Side(style="thin", color=K.MIST)
        if side:
            for c_ in range(2, n_last + 1):
                ws.cell(r, c_).border = Border(top=side)
        if i == len(pairs) - 1:
            for c_ in range(2, n_last + 1):
                ws.cell(r, c_).border = Border(top=ws.cell(r, c_).border.top, bottom=Side(style="thin", color=K.INK))
        ws["A%d" % r].number_format = ";;;"

    rng = lambda c_: "%s%d:%s%d" % (c_, FIRST, c_, last)                         # noqa: E731
    p_ = cols["product"]
    ws.conditional_formatting.add(rng(p_), FormulaRule(formula=["$%s5=$%s4" % (p_, p_)], font=Font(color=K.PAPER)))
    for s in sources:
        rc, nc = per[s]["rating"], per[s]["number"]
        for k in range(5):
            ws.conditional_formatting.add(rng(rc), FormulaRule(formula=["$%s5=%d" % (nc, k + 1)],
                                                               fill=K.solid(K.RATING_FILL[k]), stopIfTrue=True))
        ws.conditional_formatting.add(rng(rc), FormulaRule(formula=['$%s5="Refused"' % rc], fill=K.solid(K.CANVAS),
                                                           font=Font(bold=True, color=K.CRIMSON), stopIfTrue=True))
    d_ = cols["decision"]
    ws.conditional_formatting.add(rng(d_), FormulaRule(formula=['$%s5="%s"' % (d_, DT.DECISIONS[1])],
                                                       font=Font(bold=True, color=K.INK)))
    ws.conditional_formatting.add(rng(d_), FormulaRule(formula=['$%s5="%s"' % (d_, DT.DECISIONS[2])],
                                                       font=Font(bold=True, color=K.SLATE)))
    ck = cols["check"]
    ws.conditional_formatting.add(rng(ck), FormulaRule(formula=['LEFT($%s5,7)="To fix:"' % ck],
                                                       font=Font(bold=True, color=K.CRIMSON)))
    lf = cols["lines_from"]
    for s, colour in DT.source_colors(list(sources) + list(dropped)).items():
        if colour == K.KEY_RED:
            ws.conditional_formatting.add(rng(lf), FormulaRule(formula=['$%s5="%s"' % (lf, _q(s))],
                                                               font=Font(color=K.KEY_RED)))
    # The count, top right, as on Summary.
    lead = sources[0] if sources else None
    span = lambda c_: "$%s$%d:$%s$%d" % (c_, FIRST, c_, last)                  # noqa: E731
    count = '="Lines picked: "&SUM(%s)&" of %d' % (span(cols["picked"]), len(pairs))
    if lead:
        count += ' · %s rated 4 or 5: "&COUNTIF(%s,">=4")&"' % (_q(lead), span(per[lead]["number"]))
    count += ' · refused: "&SUM(%s)' % span(cols["refused"])
    ws["%s1" % last_col] = count
    ws["%s1" % last_col].font = Font(name="Arial", size=12, bold=True, color=K.INK)
    ws["%s1" % last_col].alignment = Alignment(horizontal="right", vertical="bottom")
    ws.auto_filter.ref = "B4:%s%d" % (last_col, last)
    # Arrows on the first two parts only: on the narrow number columns an
    # arrow covers the heading (hiddenButton, as Excel writes it).
    for i in range(2, n_last - 1):
        ws.auto_filter.filterColumn.append(FilterColumn(colId=i, hiddenButton=True, showButton=None))
    ws.freeze_panes = "D5"
    floored = floor_at_zero and higher
    floor_words = ("" if not floored else " A value at or below zero is always Low" + (
        "" if not floor_off else ", except on %s, which %s no floor at zero"
        % (", ".join(floor_off), "has" if len(floor_off) == 1 else "have")) + ".")
    ws.cell(last + 2, 2, "Ratings: 1 Low, 2 Moderate-Low, 3 Moderate, 4 Moderate-High, 5 High.%s Each %s's own "
                         "lines, with the same settings, are on the Lines by source tab. Every term is defined on "
                         "the Glossary tab." % (floor_words, p2.lower())).font = Font(name="Calibri", size=10,
                                                                                       color=K.SLATE)
    rows = {pr: FIRST + i for i, pr in enumerate(pairs)}
    return {"sheet": TITLE, "first": FIRST, "last": last, "cols": cols, "per": per, "rows": rows,
            "hidden": lay["hidden"], "shown": list(sources), "dropped": list(dropped),
            "count": "%s1" % last_col}


def refs(info: dict, by_pair: Dict[Tuple[str, str], Dict[str, str]]) -> Dict[str, Dict[str, str]]:
    """For Summary: each series label's cells on Thresholds, by meaning:
    its latest value and quarter, its rating as a number and as words, the
    row's refusal and the source the row's lines are set from."""
    out = {}
    cols, per = info["cols"], info["per"]
    for pr, srcs in by_pair.items():
        r = info["rows"][pr]
        for s, lab in srcs.items():
            c = lambda k: "Thresholds!$%s$%d" % (k, r)                          # noqa: E731
            out[lab] = {"value": c(per[s]["latest"]), "quarter": c(per[s]["quarter"]),
                        "rating": c(per[s]["number"]), "rating_text": c(per[s]["rating"]),
                        "refusal": c(cols["refusal"]), "lines_from": c(cols["lines_from"]),
                        "picked": c(cols["picked"]), "refused": c(cols["refused"])}
    return out
