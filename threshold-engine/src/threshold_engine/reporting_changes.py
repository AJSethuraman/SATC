"""Reporting changes: a stated registry of the Call Report's definition and
form changes, and the labels it puts on every series a change touches.

The firm, 16 Oct 2026, on the commercial proposal's question 13: *"Definitely
mark reporting changes."* Every change is labelled on every series it
touches, in the quarter it took effect, a change spread over several quarters
included, the consumer book too.

The registry is a file (``config/reporting-changes.json``), never code, and
every key of every entry is required; nothing has a default:

    {"_what": "what the file is and what it leaves out",
     "events": [entries of the same shape, [] for none],
     "changes": [{
        "name": "a short name",
        "first": "YYYYQn", "last": "YYYYQn",   # the quarters it took effect in
        "products": [the tool's products it touches, books.PRODUCTS; [] for none],
        "measures": [the measures it touches, MEASURES; [] exactly when products is []],
        "call_report_items": ["RC-C 4: C&I", ...; the items it touches, for commercial later],
        "what_changed": "one plain sentence",
        "sources": [{"citation": "...", "url": "https://..."}, ...],
        "already_labelled": {product: "words of the product's own definition note
                             that already describe this change"} or {}
     }]}

``events`` (the firm, 17 Oct 2026, on the commercial proposal's question 15,
"agreed with the recommendation") holds what moved a series without being a
change to the report: the 2020 revolving-line drawdowns and the PPP loans in
C&I. Each event is checked as a change is, and labelled the same way on every
series it touches, under its own key (``# event: ...``, "Event" on Sources),
so it is never called a reporting change. Events are listed on the
Reporting changes tab after the changes, under their own heading; Decide's
and the Peers charts' diamonds mark reporting changes only.

``first`` to ``last`` is when the change took effect: one quarter, or several
where banks were allowed to report it over more than one (the December 2024
NDFI and margin-loan revision, best efforts until 30 June 2025). A change that
a product's own definition note already describes (the securitized cards
before 2010, auto inside other consumer until 2011) is linked to that note,
not said twice: ``already_labelled`` names the words, and the registry is
refused if that note does not hold them.

What reads it:

- ``load`` checks it and refuses, naming each entry and what is wrong with it,
  every problem at once (``RegistryError``), and returns the changes;
  ``load_events`` checks the same file and returns its events;
- ``event_lines`` gives the ``# event:`` lines for one series, as
  ``note_lines`` gives the reporting-change lines;
- ``labelled_quarters`` gives every quarter the entries given name for a
  product and measure. Given the changes, a window the jump leaves out only
  for a move in one of those quarters does not break the series written (the
  firm, 17 Oct 2026, question 16; ``peers.run.bridged``); the events are
  never given for that, as they are not reporting changes;
- ``note_lines`` gives the ``# reporting-change:`` lines for one series (its
  products, its measure and its quarters), which ``peers build`` writes into
  every series file it makes, so the Sources tab shows them;
- ``quarters_in_notes`` reads those lines back out of a file's notes, which is
  how Decide's chart knows the quarters to mark;
- ``write_tab`` writes the Reporting changes tab of the book's workbook.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Iterable, List, Sequence, Set, Tuple

TOP_KEYS = ("_what", "changes", "events")
KEYS = ("name", "first", "last", "products", "measures", "call_report_items", "what_changed", "sources",
        "already_labelled")
SOURCE_KEYS = ("citation", "url")
#: The charge-off measures, as the registry names them; the delinquency
#: measures are ``peers.delinquency.MEASURES``' keys.
CHARGE_OFFS = ("net charge-offs", "gross charge-offs")
#: The key a series file's note line carries (``# reporting-change: ...``):
#: the Sources tab shows it as "Reporting change".
NOTE_KEY = "reporting-change"
#: The key an event's note line carries (``# event: ...``): "Event" on Sources.
EVENT_KEY = "event"
TAB = "Reporting changes"
#: The heading the events take on the tab, and what it says about them.
EVENTS_HEADING = "Events"
EVENTS_ABOUT = ("What moved a series without being a change to the report's definitions or forms (the firm, "
                "17 Oct 2026). Each is labelled on every series it touches as an Event, not a Reporting change, "
                "and is not marked on the charts.")
HEADINGS = ("Change", "Took effect", "Products touched", "Measures touched", "Call Report items", "What changed",
            "Source document")
_QUARTER = re.compile(r"\d{4}Q[1-4]")
_WHEN = re.compile(r"(\d{4}Q[1-4])(?: to (\d{4}Q[1-4]))?, ")


class RegistryError(ValueError):
    """A registry that cannot be read, with every problem found."""


def products() -> Tuple[str, ...]:
    from .peers.books import PRODUCTS
    return tuple(PRODUCTS)


def measures() -> Tuple[str, ...]:
    from .peers.delinquency import MEASURES
    return CHARGE_OFFS + tuple(MEASURES)


def charge_off_measure(kind: str) -> str:
    """'net' or 'gross' (the peers build's words) -> the registry's measure."""
    if kind not in ("net", "gross"):
        raise ValueError("a charge-off rate is net or gross, not %r" % (kind,))
    return "%s charge-offs" % kind


def _definitions(product: str) -> List[str]:
    """The definition notes a product's series carry: charge-offs' and delinquency's."""
    from .peers import books, delinquency
    return [d[product] for d in (books.DEFINITIONS, delinquency.DEFINITIONS) if product in d]


def _n(q: str) -> int:
    return int(q[:4]) * 4 + int(q[5]) - 1


def quarters(first: str, last: str) -> List[str]:
    """'2024Q4', '2025Q2' -> ['2024Q4', '2025Q1', '2025Q2']."""
    return ["%dQ%d" % (n // 4, n % 4 + 1) for n in range(_n(first), _n(last) + 1)]


def _keys(where: str, got, want, problems: List[str]) -> bool:
    """The book's rule: every key, and no other. True when all are there."""
    if not isinstance(got, dict):
        problems.append("%s must be an object with the keys %s" % (where, ", ".join(want)))
        return False
    for k in want:
        if k not in got:
            problems.append("%s is missing %r (every key is required; none has a default)" % (where, k))
    for k in got:
        if k not in want:
            problems.append("%s has %r, which is not a key it takes (check the spelling)" % (where, k))
    return all(k in got for k in want)


def _names(where, key, got, allowed, problems) -> List[str]:
    if not (isinstance(got, list) and all(isinstance(x, str) for x in got)):
        problems.append("%s: %s must be a list of names (empty for none), not %r" % (where, key, got))
        return []
    bad = [x for x in got if x not in allowed]
    if bad:
        problems.append("%s: %s has %s, not one of %s" % (where, key, ", ".join(repr(x) for x in bad),
                                                        ", ".join(allowed)))
    if len(set(got)) != len(got):
        problems.append("%s: %s names one twice" % (where, key))
    return got


def _one_line(text) -> bool:
    """True when ``text`` holds no line break of any kind (\\n, \\r, U+2028 ...):
    a series note is split into lines, so a break would cut the entry in two."""
    return isinstance(text, str) and text.splitlines() == [text]


def check(raw, name: str = "the registry") -> List[dict]:
    """The registry's changes, checked; ``RegistryError`` listing every
    problem, each naming its entry, when anything is wrong, in the changes or
    in the events."""
    return check_all(raw, name)[0]


def check_all(raw, name: str = "the registry") -> Tuple[List[dict], List[dict]]:
    """``(changes, events)``, both checked by the same rules; a name is used
    once across both lists."""
    problems: List[str] = []
    if not _keys(name, raw, TOP_KEYS, problems):
        raise RegistryError("\n".join(problems))
    if not (isinstance(raw["_what"], str) and raw["_what"].strip()):
        problems.append("%s: _what must say what the file is" % name)
    changes, events = raw["changes"], raw["events"]
    if not (isinstance(changes, list) and changes):
        problems.append("%s: changes must list at least one change" % name)
    if not isinstance(events, list):
        problems.append("%s: events must be a list of events ([] for none)" % name)
    if problems:
        raise RegistryError("\n".join(problems))
    seen: set = set()
    out_c = _entries(changes, "change", seen, problems)
    out_e = _entries(events, "event", seen, problems)
    if problems:
        raise RegistryError("\n".join(problems))
    return out_c, out_e


def _entries(entries, kind: str, seen: set, problems: List[str]) -> List[dict]:
    """The entries of one list (``kind`` 'change' or 'event'), each checked;
    every problem is added to ``problems`` under the entry's name."""
    allowed_p, allowed_m = products(), measures()
    out = []
    for i, c in enumerate(entries, start=1):
        label = c.get("name") if isinstance(c, dict) else None
        where = "%s %d (%s)" % (kind, i, label) if isinstance(label, str) and label.strip() else "%s %d" % (kind, i)
        if not _keys(where, c, KEYS, problems):
            continue
        before = len(problems)
        if not (isinstance(c["name"], str) and c["name"].strip()) or not _one_line(c["name"]):
            problems.append("%s: name must be a short name on one line" % where)
        elif c["name"].strip().casefold() in seen:
            problems.append("%s: the name is used twice" % where)
        else:
            seen.add(c["name"].strip().casefold())
        for k in ("first", "last"):
            if not (isinstance(c[k], str) and _QUARTER.fullmatch(c[k])):
                problems.append("%s: %s must be a quarter written like 2024Q4, not %r" % (where, k, c[k]))
        if len(problems) == before and _n(c["last"]) < _n(c["first"]):
            problems.append("%s: last, %s, comes before first, %s" % (where, c["last"], c["first"]))
        ps = _names(where, "products", c["products"], allowed_p, problems)
        ms = _names(where, "measures", c["measures"], allowed_m, problems)
        if bool(ps) != bool(ms):
            problems.append("%s: a change that touches a product names the measures it touches, and one that "
                            "touches none names none (products %s, measures %s)"
                            % (where, ps or "none", ms or "none"))
        items = c["call_report_items"]
        if not (isinstance(items, list) and all(isinstance(x, str) and x.strip() for x in items)):
            problems.append("%s: call_report_items must be a list of the items it touches (empty for none)"
                            % where)
            items = []
        if not ps and not items:
            problems.append("%s: it touches no product and no Call Report item, so it would label nothing"
                            % where)
        w = c["what_changed"]
        if not (isinstance(w, str) and w.strip()) or not _one_line(w):
            problems.append("%s: what_changed must be one plain sentence on one line" % where)
        srcs = c["sources"]
        if not (isinstance(srcs, list) and srcs):
            problems.append("%s: sources must list at least one source, each a citation and its URL" % where)
        else:
            for j, s in enumerate(srcs, start=1):
                sw = "%s, source %d" % (where, j)
                if not _keys(sw, s, SOURCE_KEYS, problems):
                    continue
                if not (isinstance(s["citation"], str) and s["citation"].strip()) or not _one_line(s["citation"]):
                    problems.append("%s: citation must name the document, on one line" % sw)
                if not (isinstance(s["url"], str) and re.fullmatch(r"https://\S+", s["url"])):
                    problems.append("%s: url must be the document's address, starting https://, not %r"
                                    % (sw, s["url"]))
        al = c["already_labelled"]
        if not isinstance(al, dict):
            problems.append("%s: already_labelled must be an object from a product to the words of its own "
                            "note that describe this change ({} for none)" % where)
        else:
            for p, words in al.items():
                if p not in ps:
                    problems.append("%s: already_labelled names %r, which is not one of its products" % (where, p))
                elif not (isinstance(words, str) and words.strip()):
                    problems.append("%s: already_labelled for %s must quote the note's words" % (where, p))
                elif not any(words in d for d in _definitions(p)):
                    problems.append("%s: already_labelled for %s quotes %r, but no definition note of %s says it, "
                                    "so the link would point at nothing" % (where, p, words, p))
        if len(problems) == before:
            out.append({k: c[k] for k in KEYS})
    return out


def load(path) -> List[dict]:
    """Read and check the registry file at ``path``; its changes."""
    return load_all(path)[0]


def load_events(path) -> List[dict]:
    """Read and check the registry file at ``path``; its events."""
    return load_all(path)[1]


def load_all(path) -> Tuple[List[dict], List[dict]]:
    """Read and check the registry file at ``path``: ``(changes, events)``."""
    path = Path(path)
    try:
        raw = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError:
        raise RegistryError("the reporting changes file %s does not exist" % path) from None
    except json.JSONDecodeError as e:
        raise RegistryError("%s is not valid JSON: %s" % (path.name, e)) from None
    except UnicodeDecodeError:
        raise RegistryError("%s is not a UTF-8 text file" % path) from None
    except OSError as e:
        raise RegistryError("the reporting changes file %s cannot be read: %s" % (path, e.strerror or e)) from None
    return check_all(raw, path.name)


def what(path) -> str:
    """The registry's own description of itself (its ``_what``)."""
    return json.loads(Path(path).read_text(encoding="utf-8"))["_what"]


def touching(changes: Sequence[dict], products_: Iterable[str], measure: str, first: str, last: str) -> List[dict]:
    """The changes that touch any of ``products_`` on ``measure`` and took
    effect in a quarter from ``first`` to ``last``."""
    ps = set(products_)
    lo, hi = _n(first), _n(last)
    return [c for c in changes if ps & set(c["products"]) and measure in c["measures"]
            and _n(c["first"]) <= hi and _n(c["last"]) >= lo]


def when(c: dict) -> str:
    return c["first"] if c["first"] == c["last"] else "%s to %s" % (c["first"], c["last"])


def cited(c: dict) -> str:
    return "; ".join("%s, %s" % (s["citation"], s["url"]) for s in c["sources"])


def note_lines(changes: Sequence[dict], products_: Sequence[str], measure: str, first: str, last: str,
               notes: Sequence[str]) -> List[str]:
    """The ``reporting-change:`` note lines for a series of ``products_`` and
    ``measure`` running ``first`` to ``last``, whose other note lines are
    ``notes``. A change one of those lines already describes (its
    ``already_labelled`` words are in it) is linked to that line, not said
    again."""
    return _lines(changes, NOTE_KEY, products_, measure, first, last, notes)


def event_lines(events: Sequence[dict], products_: Sequence[str], measure: str, first: str, last: str,
                notes: Sequence[str]) -> List[str]:
    """The ``event:`` note lines for a series, as ``note_lines`` gives the
    reporting-change lines: an event is labelled the same way, under its own
    key, so it is never called a reporting change."""
    return _lines(events, EVENT_KEY, products_, measure, first, last, notes)


def labelled_quarters(entries: Sequence[dict], product: str, measure: str) -> Set[str]:
    """Every quarter ('YYYYQn') an entry (a change or an event) in ``entries``
    names for ``product`` on ``measure``: each quarter from its first to its
    last."""
    return {q for c in entries if product in c["products"] and measure in c["measures"]
            for q in quarters(c["first"], c["last"])}


def _lines(changes, key, products_, measure, first, last, notes) -> List[str]:
    out = []
    for c in touching(changes, products_, measure, first, last):
        linked = [w for p, w in c["already_labelled"].items() if p in products_ and any(w in n for n in notes)]
        body = ('The note above that says "%s" describes this change.' % linked[0] if linked
                else c["what_changed"])
        out.append("%s: %s, %s. %s Source: %s." % (key, when(c), c["name"], body, cited(c)))
    return out


def quarters_in_notes(notes: Sequence[Tuple[str, str]]) -> Set[str]:
    """Every quarter a series' ``reporting-change`` notes say a change took
    effect in, from ``series.read_notes``' (key, text) pairs. A
    reporting-change line that does not begin with its quarters is refused."""
    out: Set[str] = set()
    for key, text in notes:
        if key != NOTE_KEY:
            continue
        m = _WHEN.match(text)
        if not m:
            raise ValueError("a reporting-change note must begin with the quarter it took effect, like "
                             "'2024Q4 to 2025Q2, ...', not %r" % text[:60])
        out.update(quarters(m.group(1), m.group(2) or m.group(1)))
    return out


def write_tab(ws, changes: Sequence[dict], about: str, source_name: str, events: Sequence[dict]) -> int:
    """The Reporting changes tab, on the reference tabs' template: the
    registry's own description, then one row per change, then, under their
    own heading, one row per event (``events``: what moved a series without
    being a change to the report; ``[]`` for none, and then no heading).
    Returns the last row written."""
    from openpyxl.styles import Alignment, Font

    from . import keybank_style as K
    from .workbook import REF_EDGE, REF_LABEL, _lines, _ref_header, _ref_row, _ref_title

    last_col = 1 + len(HEADINGS)
    _ref_title(ws, TAB, "Every change to the Call Report's definitions or forms this workbook labels, from %s. "
                        "Each series it touches says so on Sources." % source_name, last_col)
    ws.cell(4, 2, "About").font = REF_LABEL
    c = ws.cell(4, 3, about)
    c.font = Font(name="Calibri", size=11, color=K.INK_TEXT)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=4, start_column=3, end_row=4, end_column=last_col)
    widths = (30, 16, 22, 26, 36, 60, 60)
    ws.row_dimensions[4].height = 15 * _lines(about, sum(widths[1:]))
    r = _ref_header(ws, 6, list(HEADINGS), 2)

    def rows(entries, r):
        for ch in entries:
            vals = (ch["name"], when(ch), ", ".join(ch["products"]) or "None of the products",
                    ", ".join(ch["measures"]) or "None", "; ".join(ch["call_report_items"]) or "None stated",
                    ch["what_changed"], "\n".join("%s, %s" % (s["citation"], s["url"]) for s in ch["sources"]))
            for col, v in enumerate(vals, start=2):
                cell = ws.cell(r, col, v)
                cell.font = Font(name="Calibri", size=11, color=K.INK_TEXT, bold=col == 2)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            _ref_row(ws, r, 2, last_col, lines=max(_lines(v, w) + str(v).count("\n") for v, w in zip(vals, widths)))
            r += 1
        return r

    r = rows(changes, r)
    if events:
        r += 1
        ws.cell(r, 2, EVENTS_HEADING).font = REF_LABEL
        c = ws.cell(r, 3, EVENTS_ABOUT)
        c.font = Font(name="Calibri", size=11, color=K.INK_TEXT)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=last_col)
        ws.row_dimensions[r].height = 15 * _lines(EVENTS_ABOUT, sum(widths[1:]))
        r = _ref_header(ws, r + 1, list(HEADINGS), 2)
        r = rows(events, r)
    ws.column_dimensions["A"].width = REF_EDGE
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(ord("B") + i)].width = w
    ws.freeze_panes = "B7"
    return r - 1
