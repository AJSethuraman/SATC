"""Where each series comes from, and how its file is read.

One reader per kind of source. Each turns the publisher's file into
(date, value) pairs, oldest first, with ``None`` wherever the publisher
printed no number: FRED's "." or blank, BLS's dash. A missing value is never
turned into zero, because a zero would read as the most extreme reading in
the series' history (the October 2025 unemployment rate is the case in point).

Dates are the first day of the period: 2026-07-01 is July 2026, or 2026Q3 for
a quarterly series.
"""
import csv
import html
import io
import re
import zipfile

FRED = "https://fred.stlouisfed.org/graph/fredgraph.csv?id={id}"
BLS = ("https://data.bls.gov/timeseries/{id}?from_year={start}&to_year={end}"
       "&output_view=data&include_graphs=false")
FED_EBP = "https://www.federalreserve.gov/econres/notes/feds-notes/ebp_csv.csv"
FHFA_NMDB = ("https://www.fhfa.gov/document/d/nmdb/"
             "nmdb-mortgage-performance-statistics-national-census-areas-quarterly.zip")
NYFED_SCE = "https://www.newyorkfed.org/medialibrary/interactives/sce/sce/downloads/data/frbny-sce-data.xlsx"

NUMBER = re.compile(r"-?\d+(\.\d+)?")


def url_for(series, today_year):
    """The address one series is downloaded from. No key is used."""
    if series.source == "fred":
        return FRED.format(id=series.source_id)
    if series.source == "bls":
        return BLS.format(id=series.source_id, start=series.starts[:4], end=today_year)
    if series.source == "fed_ebp":
        return FED_EBP
    if series.source == "fhfa_nmdb":
        return FHFA_NMDB
    if series.source == "nyfed_sce":
        return NYFED_SCE
    raise ValueError("no source of kind %r" % series.source)


def _value(text):
    text = text.strip()
    return float(text) if NUMBER.fullmatch(text) else None


def read_fred(raw, source_id):
    """FRED's keyless CSV: a header row, then date,value."""
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig"))))
    if not rows or len(rows[0]) != 2 or rows[0][1] != source_id:
        raise ValueError("not a FRED file for %s: header %r" % (source_id, rows[0] if rows else None))
    return [(d, _value(v)) for d, v in rows[1:] if d]


def read_bls(raw, source_id):
    """The data table on data.bls.gov/timeseries: one row a year, Jan to Dec.

    A cell may carry a footnote mark, as in "4.3(12)", or no number at all, as
    in "-(9)" for October 2025. Every month from the first number to the last
    is kept, and one with no number is missing, whether BLS printed a dash or
    left it blank (motor vehicle insurance, November 2025, is blank). Blank
    months after the last number are months not published yet, and are left
    out.
    """
    text = raw.decode("utf-8", "replace")
    if "Series Id:" not in text or source_id not in text:
        raise ValueError("not the BLS page for %s" % source_id)
    table = text[text.find('id="table0"'):]
    table = table[:table.find("</TABLE>") if "</TABLE>" in table else table.find("</table>")]
    out = []
    for year, row in re.findall(r'<TH scope="row">(\d{4})</TH>(.*?)</TR>', table, re.S):
        cells = [html.unescape(re.sub(r"<[^>]+>", "", c)).strip() for c in re.findall(r"<TD[^>]*>(.*?)</TD>", row, re.S)]
        for month in range(1, 13):
            cell = cells[month - 1] if month <= len(cells) else ""
            cell = re.sub(r"\(\w+\)", "", cell).replace("\xa0", "").strip()
            out.append(("%s-%02d-01" % (year, month), _value(cell) if cell else None))
    present = [i for i, (_, v) in enumerate(out) if v is not None]
    if not present:
        raise ValueError("no numbers in the BLS table for %s" % source_id)
    return out[present[0]:present[-1] + 1]


def read_fed_ebp(raw, column):
    """The Fed's excess bond premium file: date,gz_spread,ebp,est_prob."""
    rows = list(csv.DictReader(io.StringIO(raw.decode("utf-8-sig"))))
    if not rows or column not in rows[0]:
        raise ValueError("the Fed's file has no column %r" % column)
    return [(r["date"], _value(r[column])) for r in rows]


def read_fhfa_nmdb(raw, selector):
    """FHFA's National Mortgage Database file; selector is GEOID|MARKET|SERIESID.

    A suppressed cell is kept as missing.
    """
    geoid, market, seriesid = selector.split("|")
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        name = next(n for n in z.namelist() if n.endswith(".csv"))
        rows = list(csv.DictReader(io.StringIO(z.read(name).decode("utf-8-sig"))))
    out = []
    for r in rows:
        if r["GEOID"] == geoid and r["MARKET"] == market and r["SERIESID"] == seriesid:
            month = 3 * int(r["QUARTER"]) - 2
            v = None if r.get("SUPPRESSED", "0") not in ("0", "") else _value(r["VALUE1"])
            out.append(("%s-%02d-01" % (r["YEAR"], month), v))
    if not out:
        raise ValueError("FHFA file has no rows for %s" % selector)
    return sorted(out)


#: What the SCE file's delinquency sheet says its column holds, checked so another sheet is never read as it.
SCE_MISSED = "Mean probability of missing minimum debt payment over the next three months"


def read_nyfed_sce(raw, sheet):
    """The NY Fed's Survey of Consumer Expectations workbook: one sheet per question, the month as 201306 in
    column A and the mean in column B, under four rows of source and title. A month with no number is kept as
    missing."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        raise ValueError("not the NY Fed's SCE workbook: %s" % e)
    if sheet not in wb.sheetnames:
        raise ValueError("the SCE workbook has no sheet %r" % sheet)
    rows = list(wb[sheet].iter_rows(values_only=True))
    heads = [str(c) for r in rows[:6] for c in r if isinstance(c, str)]
    if SCE_MISSED not in heads:
        raise ValueError("the SCE sheet %r does not hold %r" % (sheet, SCE_MISSED))
    out = []
    for r in rows:
        month = r[0] if r else None
        if isinstance(month, (int, float)) and not isinstance(month, bool) and 190001 <= month <= 299912:
            month = int(month)
            if not 1 <= month % 100 <= 12:
                raise ValueError("the SCE sheet %r has %s where a month belongs" % (sheet, r[0]))
            v = r[1] if len(r) > 1 else None
            out.append(("%04d-%02d-01" % (month // 100, month % 100),
                        float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None))
    if not out:
        raise ValueError("no months in the SCE sheet %r" % sheet)
    return out


READERS = {"fred": read_fred, "bls": read_bls, "fed_ebp": read_fed_ebp, "fhfa_nmdb": read_fhfa_nmdb,
           "nyfed_sce": read_nyfed_sce}


def read(series, raw):
    return READERS[series.source](raw, series.source_id)
