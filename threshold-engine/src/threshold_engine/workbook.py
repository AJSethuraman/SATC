"""One run, one workbook: every product's facts, evidence, backtest and live cutoffs.

Python does the statistics once -- spells, robust z-scores, scenarios, recency,
the backtest -- and writes their results as values. The cutoffs are Excel
formulas over the data, driven by the red Assess tab, so the bank's
judgements are made in the workbook and the thresholds move as they are made:

    Assess (red)    pick a product and see where its history falls; pick the
                    periods to leave out from lists of real quarters, and set
                    the scale's three distances and the on-the-line rule
    Thresholds      median and worst of the periods kept, the four cutoffs,
                    the latest value's score and rating -- all live
    Synthesis       only with synthesis inputs (synthesis_tab.py): internal,
                    industry and peers side by side for each line, which one
                    each line is taken from, the chosen lines -- live, with
                    the series it reads on SynthesisData
    Chart          any combination of products on one scale, with one
                    product's cutoffs drawn, switched on the tab itself
    Evidence        spells, robust z, scenarios, recency (values)
    Backtest        how predictive the scale would have been (values)
    Data            every series on one quarterly calendar, and the periods kept
    Run             what went in: files, fingerprints, rows, settings

Every product must share one measure (name, unit, direction); a credit score
and a loss rate do not belong on one scale or one chart.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import math
from pathlib import Path
from typing import Dict, Optional, Sequence, Tuple

from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from . import assess_tab as AT
from . import keybank_style as K
from . import decide_tab as DT
from . import reporting_changes as RC
from . import reference_tabs as RT
from . import summary_tab as SU
from . import synthesis_tab as ST
from .backtest import backtest
from . import VERSION
from .evidence import evidence
from .glossary import TERMS
from .scale import scale as _scale
from .series import CORE_NOTES, LABEL_RULE, check, label_ok, read_csv, read_notes, smooth

RATINGS = ("Low", "Moderate-Low", "Moderate", "Moderate-High", "High")
INPUT = PatternFill("solid", fgColor="FFF2CC")
# A conditional format's fill is read from its background colour, so both
# are set (with fgColor alone Excel drew the highlighted row white).
PICKED = PatternFill("solid", fgColor="F6D9A8", bgColor="F6D9A8")
# The Assess tab's inputs, laid out by assess_tab (Claude Design's hand-off,
# 13 Oct 2026): Your settings N:R from SET_TOP, one row per product; Periods
# left out N:Q below it, LO_ROWS rows; hidden helpers from column Z (HC).
SET_TOP, LO_ROWS, HC = AT.SET_TOP, AT.LO_ROWS, AT.HC
NOTE = Font(name="Arial", size=9, italic=True, color=K.SLATE)


def _wrap(ws, row, height=34):
    """Header labels wrap in their column instead of being cut off."""
    for cell in ws[row]:
        if cell.value is not None:
            cell.alignment = Alignment(wrap_text=True, vertical="center",
                                       horizontal=cell.alignment.horizontal or "left")
    ws.row_dimensions[row].height = height


def _axis_floor(loaded, direction, halfwidth, low_step, high_z):
    """A round number under every product's lowest value and line, and under
    the strip that marks quarters left out, or None if that is not above zero
    (then the axis starts at zero anyway). Rounded down to a half of the
    range's order of magnitude.

    The axis cannot follow a cell, so its minimum has to hold for the
    settings the bank may still type, not only the starting ones (Codex on
    #432: raising Moderate-Low's step pushed score 2's line off the bottom).
    It covers every value the dials offer as well as the starting values:
    a width up to 1, a Moderate-Low step up to 2, a z from 2 to 3.5."""
    s_ = 1.0 if direction == "higher_is_worse" else -1.0
    lows, highs = [], []
    for s in loaded:
        vals = [p.value for p in s["points"]]
        ref = _scale(vals, direction, halfwidth, low_step, high_z, False)
        med, move, spread = ref["median"], ref["typical_yearly_move"] or 0.0, ref["spread_of_levels"]
        lines = []
        for hw in {halfwidth, 0.25, 1.0}:
            for ls in {low_step, 0.75, 2.0}:
                for z in {high_z, 2.0, 3.5}:
                    lines += [med - s_ * ls * move, med - s_ * hw * move, med + s_ * hw * move, med + s_ * z * spread]
        both = vals + lines
        lows.append(min(both))
        highs.append(max(both))
    lo, hi = min(lows), max(highs)
    # As the sheet works them out: the floor, the top, and the strip under the
    # floor for quarters left out, which must stay on the chart.
    floor = lo - 0.08 * (hi - lo)
    top = hi + 0.08 * (hi - floor)
    target = floor - 0.06 * (top - floor)
    if hi <= lo or target <= 0:
        return None
    unit = 10 ** math.floor(math.log10(hi - lo)) / 2
    return math.floor(target / unit) * unit


def assess_layout(n: int) -> dict:
    """Where things sit on Assess for n products: one place for the builder,
    the tests and the walkthrough (see assess_tab.layout)."""
    return AT.layout(n)


def quarter(date: str) -> str:
    """'2010-04-01' and '2010-06-30' are both 2010Q2."""
    return "%sQ%d" % (date[:4], (int(date[5:7]) - 1) // 3 + 1)


#: One reader for every front door (``series.read_csv``).
read_series = read_csv


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ---- The reference tabs' one template (Claude Design's round 2,
# docs/design/round-2/README.md, "Reference tabs") ----
REF_EDGE = 1.57                                       # column A, a 16 px margin
REF_BODY = Font(name="Calibri", size=11, color=K.INK_TEXT)
REF_LABEL = Font(name="Arial", size=9, bold=True, color=K.SLATE)


def _lines(text, width):
    """Lines a text takes when wrapped in a column ``width`` units wide."""
    return max(1, -(-len(str(text or "")) // max(1, int(width))))


def _ref_title(ws, title, purpose, last_col):
    """Column A a margin; row 1 the title (Arial 14 bold), row 2 the purpose
    (Calibri 11 SLATE) over a medium KEY_RED rule to ``last_col``."""
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = REF_EDGE
    t = ws.cell(1, 2, title)
    t.font = Font(name="Arial", size=14, bold=True, color=K.INK)
    t.alignment = Alignment(vertical="bottom")
    p = ws.cell(2, 2, purpose)
    p.font = Font(name="Calibri", size=11, color=K.SLATE)
    p.alignment = Alignment(vertical="center")
    for col in range(2, last_col + 1):
        ws.cell(2, col).border = Border(bottom=Side(style="medium", color=K.KEY_RED))
    for r_, h in ((1, 25.5), (2, 15), (3, 9)):
        ws.row_dimensions[r_].height = h


def _ref_header(ws, row, names, first_col, numeric=(), height=27):
    """A header row: MIST fill, Arial 9 bold INK, wrapped, a STONE rule;
    the ``numeric`` positions right-aligned. Returns the next row."""
    for i, nm in enumerate(names):
        c = ws.cell(row, first_col + i, nm)
        c.fill = K.solid(K.MIST)
        c.font = Font(name="Arial", size=9, bold=True, color=K.INK)
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="right" if i in numeric else "left")
        c.border = Border(bottom=Side(style="thin", color=K.STONE))
    ws.row_dimensions[row].height = height
    return row + 1


def _ref_block(ws, row, text, first_col, last_col):
    """A series block's heading: Arial 11 bold over a thin INK rule."""
    c = ws.cell(row, first_col, text)
    c.font = Font(name="Arial", size=11, bold=True, color=K.INK)
    c.alignment = Alignment(vertical="bottom")
    for col in range(first_col, last_col + 1):
        ws.cell(row, col).border = Border(bottom=Side(style="thin", color=K.INK))
    ws.row_dimensions[row].height = 18
    return row + 1


def _ref_row(ws, row, first_col, last_col, lines=1):
    """A body row: 15 pt a line, a MIST hairline under it."""
    for col in range(first_col, last_col + 1):
        ws.cell(row, col).border = Border(bottom=Side(style="thin", color=K.MIST))
    ws.row_dimensions[row].height = 15 * lines


def _thresholds_plain(th, name, n, loaded, floored):
    """Thresholds without Decide: the layout it has always had (TH_PLAIN)."""
    th.sheet_properties.tabColor = K.KEY_RED
    K.light_title(th, 1, 13, "%s · candidate thresholds" % name,
                  "Live: follows every choice on Assess. Candidates for the bank to accept, adjust or reject.")
    for c1, c2, lbl in ((2, 4, "What the lines rest on"), (5, 8, "Each rating begins at"), (9, 12, "Latest quarter"),
                        (13, 13, "Check")):
        th.cell(3, c1, lbl).font = Font(name="Arial", bold=True, size=9, color=K.INK)
        for c_ in range(c1, c2 + 1):
            th.cell(3, c_).border = Border(bottom=Side(style="thin", color=K.INK))
    th.row_dimensions[3].height = 15
    K.canvas_header(th, 4, ["Series", "Median kept", "Typical yearly move", "Worst kept", "2 Moderate-Low from",
                            "3 Moderate from", "4 Moderate-High from", "5 High from", "Latest quarter",
                            "Latest value", "Rating", "Rating name", "Check"], right_from=1,
                    fills=[None] * 4 + list(K.BANDS[1:]) + [None] * 5)
    th.column_dimensions["O"].hidden = True
    for c_ in "PQRSTU":
        th.column_dimensions[c_].hidden = True
    for c_, w in zip("ABCDEFGHIJKLM", (15, 11, 11, 11, 11, 11, 11, 11, 10, 10, 8, 15, 70)):
        th.column_dimensions[c_].width = w
    # A label must be read whole: "Credit card - Nonaccrual - KeyBank" and its
    # industry and peer twins differ only at the end (the tie-out, 14 Oct 2026).
    th.column_dimensions["A"].width = max(th.column_dimensions["A"].width or 0,
                                          min(60, 2 + max(len(s["label"]) for s in loaded)))
    for r in range(5, 5 + n):
        th.row_dimensions[r].height = 15
        K.hairlines(th, r, 1, 13)
        for c_ in range(1, 14):
            th.cell(r, c_).font = Font(name="Calibri", size=11, color=K.INK_TEXT)
    for k in range(5):
        th.conditional_formatting.add("K5:K%d" % (4 + n), FormulaRule(formula=["$K5=%d" % (k + 1)],
                                                                      fill=K.solid(K.BANDS[k]), font=Font(bold=True)))
    th.conditional_formatting.add("M5:M%d" % (4 + n), FormulaRule(formula=['LEFT($M5,8)="Refused:"'],
                                                                  fill=K.solid(K.ALERT_FG),
                                                                  font=Font(bold=True, color=K.CRIMSON)))
    th.freeze_panes = "B5"
    th.cell(6 + n, 1, "Ratings: 1 Low, 2 Moderate-Low, 3 Moderate, 4 Moderate-High, 5 High. A value at or below "
                      "zero is always Low. Every term is defined on the Glossary tab." if floored
                      else "Ratings: 1 Low, 2 Moderate-Low, 3 Moderate, 4 Moderate-High, 5 High. Every term is "
                      "defined on the Glossary tab.").font = NOTE


#: Lines by source in the book (round 3): px -> units, as (px - 5) / 7: B 150,
#: C 216, D 150, E-K 72, L 64, M 64, N 120, O 110, P 260.
TH_BOOK_WIDTHS = dict(A=1.57, B=20.71, C=30.14, D=20.71, E=9.57, F=9.57, G=9.57, H=9.57, I=9.57, J=9.57,
                      K=9.57, L=8.43, M=8.43, N=16.43, O=15.0, P=36.43)
#: Backtest folded into Lines by source (fold_backtest): its three summary
#: columns, then the average change ahead by rating in a collapsed group,
#: past the hidden workings (decide_tab.TH_HIDDEN), so no letter moves.
FOLD_SHOWN = ("spearman_value", "spearman_change", "tested")
FOLD_WIDTHS = (9.57, 9.57, 8.43)                     # 72, 72, 64 px; the group's columns 64 px
FOLD_HEADS = ("Spearman, value", "Spearman, change", "Quarters tested")
#: A line of wrapped text at Calibri 10 on Check (P), and a refused row's height.
LBS_CHECK_CHARS, H_REFUSED = 44, 34.5


def _thresholds_book(th, name, n, labels, decide, floored, refused_at_build=()):
    """The per-series table with Decide, retitled Lines by source once
    Decide is built (decide_tab.rename_tab): every source's own lines, with
    its product and measure's settings. Claude Design's round 3 (README,
    "Lines by source"): a reference tab, Product, Measure and Source in B:D,
    frozen with the header; repeated names hidden by formatting, so the
    header filters still work. A row's height cannot follow a formula, so a
    series refused as built (``refused_at_build``) is three lines high and
    its To fix reads whole."""
    T = DT.TH_BOOK
    last = 4 + n
    th.sheet_properties.tabColor = K.TAB_REFERENCE
    th.sheet_view.showGridLines = False
    th["B1"] = "Lines by source · what each %s's own history gives" % decide["names"][2].lower()
    th["B1"].font = Font(name="Arial", bold=True, size=14, color=K.INK)
    th["B1"].alignment = Alignment(vertical="bottom")
    th["B2"] = ("Each %s's own lines, with its %s and %s's settings on Decide; Thresholds uses one of them per row. "
                "The header arrows filter by %s, %s or %s."
                % ((decide["names"][2].lower(), decide["names"][0].lower(), decide["names"][1].lower())
                   + tuple(p.lower() for p in decide["names"])))
    th["B2"].font = Font(name="Calibri", size=11, color=K.SLATE)
    th["B2"].alignment = Alignment(vertical="center")
    for c_ in range(2, 17):
        th.cell(2, c_).border = Border(bottom=Side(style="thin", color=K.STONE))
    # The header is three lines of Arial 9, so "4 Moderate-High" reads whole in 72 px.
    for r, h in ((1, 25.5), (2, 15), (3, 16.5), (4, 39)):
        th.row_dimensions[r].height = h
    groups = (("B", "D", "Series"), ("E", "G", "What the lines rest on"), ("H", "K", "Each rating begins at"),
              ("L", "N", "Latest quarter, on its own lines"), ("O", "P", "Your decision, where its lines are used"))
    from openpyxl.utils import column_index_from_string as CI
    for a, b, lbl in groups:
        _ref_group(th, a, b, lbl)
    heads = [("product", decide["names"][0]), ("measure", decide["names"][1]), ("source", decide["names"][2]),
             ("median", "Median kept"), ("move", "Yearly move"), ("worst", "Worst kept"),
             ("l2", "2 Moderate-Low"), ("l3", "3 Moderate"), ("l4", "4 Moderate-High"), ("l5", "5 High"),
             ("quarter", "Quarter"), ("value", "Value"), ("rating_shown", "Rating"), ("decision", "Decision"),
             ("check", "Check")]
    right = {"median", "move", "worst", "l2", "l3", "l4", "l5", "value"}
    for key, t_ in heads:
        _ref_head(th["%s4" % T[key]], t_, key in right)
    # The hidden workings keep a header, so a column found by its heading stays findable.
    for key, t_ in (("rating", "Rating number"), ("rating_name", "Rating name"), ("refusal", "Refused because"),
                    ("fix", "To fix"), ("med_change", "Median change"), ("spread", "Spread of levels"),
                    ("med_all", "Median, every quarter"), ("med_change_all", "Median change, every quarter"),
                    ("move_all", "Typical yearly move, every quarter"),
                    ("spread_all", "Spread of levels, every quarter"), ("floor", "Floor at zero")):
        th["%s4" % T[key]].value = t_
    a1, a2 = DT.TH_HIDDEN
    for c_ in range(CI(a1), CI(a2) + 1):
        th.column_dimensions[L(c_)].hidden = True
    for c_, w in TH_BOOK_WIDTHS.items():
        th.column_dimensions[c_].width = w
    body = Font(name="Calibri", size=11, color=K.INK_TEXT)
    refused_at_build = set(refused_at_build)
    for r in range(5, last + 1):
        th.row_dimensions[r].height = H_REFUSED if labels[r - 5] in refused_at_build else 15
        for c_ in range(2, 17):
            th.cell(r, c_).font = body
            th.cell(r, c_).alignment = Alignment(vertical="center")
        for key in ("median", "move", "worst", "l2", "l3", "l4", "l5", "value"):
            th["%s%d" % (T[key], r)].number_format = "0.00"
        th["%s%d" % (T["product"], r)].font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
        th["%s%d" % (T["quarter"], r)].font = Font(name="Calibri", size=11, color=K.SLATE)
        th["%s%d" % (T["value"], r)].font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
        th["%s%d" % (T["check"], r)].font = Font(name="Calibri", size=10, color=K.SLATE)
        # A To fix that arrives after the build, on a row one line high, shows its first line.
        th["%s%d" % (T["check"], r)].alignment = Alignment(
            vertical="center" if labels[r - 5] in refused_at_build else "top", wrap_text=True)
        src = th["%s%d" % (T["source"], r)]
        if "keybank" in str(src.value).casefold():
            src.font = Font(name="Calibri", size=11, color=K.KEY_RED)
        # A thin INK rule starts each product, a MIST one each measure.
        parts, prev = decide["parts"][labels[r - 5]], decide["parts"][labels[r - 6]] if r > 5 else None
        side = None
        if prev is None or parts[0] != prev[0]:
            side = Side(style="thin", color=K.INK)
        elif parts[1] != prev[1]:
            side = Side(style="thin", color=K.MIST)
        if side:
            for c_ in range(2, 17):
                th.cell(r, c_).border = Border(top=side)
        # The label lives in A for every lookup; the 16 px gutter shows none of it.
        th["A%d" % r].number_format = ";;;"
    rng = lambda key: "%s5:%s%d" % (T[key], T[key], last)        # noqa: E731
    # Repeated names are written in white, not deleted, so filters and lookups are untouched.
    p_, m_ = T["product"], T["measure"]
    th.conditional_formatting.add(rng("product"), FormulaRule(formula=["$%s5=$%s4" % (p_, p_)],
                                                              font=Font(color=K.PAPER)))
    th.conditional_formatting.add(rng("measure"), FormulaRule(formula=["AND($%s5=$%s4,$%s5=$%s4)" % (p_, p_, m_, m_)],
                                                              font=Font(color=K.PAPER)))
    q_ = T["rating"]
    for k in range(5):
        th.conditional_formatting.add(rng("rating_shown"), FormulaRule(formula=["$%s5=%d" % (q_, k + 1)],
                                                                       fill=K.solid(K.RATING_FILL[k]),
                                                                       stopIfTrue=True))
    th.conditional_formatting.add(rng("rating_shown"), FormulaRule(formula=['$%s5="Refused"' % T["rating_shown"]],
                                                                   fill=K.solid(K.CANVAS),
                                                                   font=Font(bold=True, color=K.CRIMSON)))
    d_ = T["decision"]
    th.conditional_formatting.add(rng("decision"), FormulaRule(formula=['$%s5="%s"' % (d_, DT.DECISIONS[1])],
                                                               font=Font(bold=True, color=K.INK)))
    th.conditional_formatting.add(rng("decision"), FormulaRule(formula=['$%s5="%s"' % (d_, DT.DECISIONS[2])],
                                                               font=Font(bold=True, color=K.SLATE)))
    ck = T["check"]
    th.conditional_formatting.add(rng("check"), FormulaRule(formula=['LEFT($%s5,7)="To fix:"' % ck],
                                                            font=Font(bold=True, color=K.CRIMSON)))
    th.auto_filter.ref = "B4:%s%d" % (T["check"], last)
    # Arrows on Product, Measure and Source only: on the narrow number
    # columns an arrow covers the heading (hiddenButton, as Excel writes it).
    from openpyxl.worksheet.filters import FilterColumn
    for i in range(3, CI(T["check"]) - 1):
        th.auto_filter.filterColumn.append(FilterColumn(colId=i, hiddenButton=True, showButton=None))
    th.freeze_panes = "E5"
    th.cell(6 + n, 2, "Ratings: 1 Low, 2 Moderate-Low, 3 Moderate, 4 Moderate-High, 5 High.%s Every term is "
                      "defined on the Glossary tab." % floored).font = Font(name="Calibri", size=10, color=K.SLATE)


def _ref_group(ws, a, b, label, row=3):
    """A group label over columns ``a`` to ``b``: Arial 9 bold INK over a thin
    STONE rule (the reference template)."""
    from openpyxl.utils import column_index_from_string as CI
    cell = ws["%s%d" % (a, row)]
    cell.value = label
    cell.font = Font(name="Arial", bold=True, size=9, color=K.INK)
    cell.alignment = Alignment(vertical="bottom")
    for c_ in range(CI(a), CI(b) + 1):
        ws.cell(row, c_).border = Border(bottom=Side(style="thin", color=K.STONE))


def _ref_head(cell, text, right=False):
    """A reference table's header cell: Arial 9 bold INK on MIST, wrapped, a
    STONE rule under it."""
    cell.value = text
    cell.font = Font(name="Arial", bold=True, size=9, color=K.INK)
    cell.fill = K.solid(K.MIST)
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="right" if right else "left")
    cell.border = Border(bottom=Side(style="thin", color=K.STONE))


def fold_backtest(wb, backtest: str = "Backtest") -> dict:
    """Fold the Backtest tab into Lines by source and remove it (the book
    only; Claude Design's round 3: "Backtest folded in").

    Each series' unweighted run gives three columns past the hidden workings
    (Spearman against the value ahead, against the change ahead, quarters
    tested), headed "How predictive, N quarters ahead"; the average change
    ahead by rating follows in a collapsed column group. A series too short
    to test has 0 quarters tested and no Spearman. Its notes go under the
    table. Nothing in the workbook refers to the Backtest tab by formula, so
    nothing else moves. Returns the columns, by meaning."""
    import re
    from openpyxl.utils import column_index_from_string as CI
    bts, ws = wb[backtest], wb[DT.LINES_TAB]
    head = {}
    c = 2
    while bts.cell(4, c).value not in (None, ""):
        head[str(bts.cell(4, c).value)] = c
        c += 1
    if "Weighting" not in head or "Quarters tested" not in head:
        raise ValueError("the Backtest tab has no Weighting or Quarters tested column; nothing was folded")
    wt = head["Weighting"]
    id_cols = list(range(2, wt))
    ratings = [k for k in head if k.isdigit()]
    m = re.search(r"value (\d+) quarters? later", str(bts["B2"].value or ""))
    if not m:
        raise ValueError("the Backtest tab does not say its horizon; nothing was folded")
    horizon = int(m.group(1))
    tested, notes, untested = {}, [], []
    r = 5
    while bts.cell(r, 2).value not in (None, ""):
        key = tuple(bts.cell(r, c_).value for c_ in id_cols)
        w = bts.cell(r, wt).value
        if w == "none":
            tested[key] = (bts.cell(r, head["Value ahead"]).value, bts.cell(r, head["Change ahead"]).value,
                           bts.cell(r, head["Quarters tested"]).value,
                           [bts.cell(r, head[k]).value for k in ratings])
        elif isinstance(w, str) and w.startswith("Not tested"):
            tested.setdefault(key, (None, None, 0, [None] * len(ratings)))
            untested.append("%s: %s" % (" - ".join(map(str, key)), w))
        r += 1
    r += 1
    while r <= bts.max_row:
        if bts.cell(r, 2).value:
            notes.append(bts.cell(r, 2).value)
        r += 1
    first = CI(DT.TH_HIDDEN[1]) + 1
    cols = {k: first + i for i, k in enumerate(FOLD_SHOWN)}
    group = [first + len(FOLD_SHOWN) + i for i in range(len(ratings))]
    _ref_group(ws, L(first), L(first + len(FOLD_SHOWN) - 1), "How predictive, %d quarters ahead" % horizon)
    _ref_group(ws, L(group[0]), L(group[-1]), "Average change ahead, by rating")
    for k, t_ in zip(FOLD_SHOWN, FOLD_HEADS):
        _ref_head(ws.cell(4, cols[k]), t_, True)
    for c_, k in zip(group, ratings):
        _ref_head(ws.cell(4, c_), "%s %s" % (k, RATINGS[int(k) - 1]), True)
    T = DT.TH_BOOK
    last, r = 4, 5
    while ws["%s%d" % (T["product"], r)].value not in (None, ""):
        key = tuple(ws["%s%d" % (T[k], r)].value for k in ("product", "measure", "source"))
        got = tested.get(key)
        if got is None:
            raise ValueError("Backtest has no row for %s; nothing was folded" % " - ".join(map(str, key)))
        vals = [got[0], got[1], got[2]] + list(got[3])
        top = ws["%s%d" % (T["check"], r)].border.top
        for c_, v in zip([cols[k] for k in FOLD_SHOWN] + group, vals):
            cell = ws.cell(r, c_, v)
            cell.font = Font(name="Calibri", size=11, color=K.INK_TEXT)
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "0" if c_ == cols["tested"] else (
                "0.00" if c_ in (cols["spearman_value"], cols["spearman_change"]) else "+0.00;-0.00;0.00")
            cell.border = Border(top=top)
        last = r
        r += 1
    for c_, w in zip([cols[k] for k in FOLD_SHOWN], FOLD_WIDTHS):
        ws.column_dimensions[L(c_)].width = w
    for c_ in group:
        ws.column_dimensions[L(c_)].width = 8.43
    ws.column_dimensions.group(L(group[0]), L(group[-1]), outline_level=1, hidden=True)
    for c_ in range(first, group[-1] + 1):
        ws.cell(2, c_).border = Border(bottom=Side(style="thin", color=K.STONE))
    # Backtest's own words, under the table's note.
    nr = last + 3
    lines = [str(bts["B2"].value)] + untested + notes
    for i, t in enumerate(lines):
        ws.cell(nr + i, 2, t).font = Font(name="Calibri", size=10, color=K.SLATE)
    wb.remove(bts)
    return {"columns": {k: L(v) for k, v in cols.items()}, "by_rating": [L(c_) for c_ in group],
            "horizon": horizon, "rows": (5, last), "notes": nr}


def build(out: Path, series: Sequence[Tuple[str, Path]], name: str, unit: str,
          direction: str, frequency: str, smoothing: int, floor_at_zero: bool,
          moderate_halfwidth: float, low_step: float, high_z: float, on_the_line: str,
          half_lives: Sequence[float],
          percentiles: Sequence[float], horizon: int, min_history: int,
          synthesis: Optional[Dict[str, dict]] = None, decide: Optional[dict] = None,
          into: Optional[Workbook] = None, before_save=None, floor_off: Sequence[str] = (),
          left_out_file: Optional[Path] = None, reporting: Optional[dict] = None) -> dict:
    """``decide``, when given, adds the Summary and Decide tabs and makes Decide
    where the settings, periods left out and decisions are typed (decide_tab.py,
    summary_tab.py): ``{"names": [three part names], "parts": {label: (three
    values)}, "absent": [{"parts": {...}, "why": text}], "compare": (a, b) or
    None, "notes": [text]}``, every key stated. ``into`` builds into an
    existing workbook (the book command puts the peer tabs there first);
    ``before_save(wb)`` runs last, before saving.

    ``left_out_file``, the left-out file the book wrote (book.LEFT_OUT_FILE), is
    printed on Run beside This build. ``reporting``, ``{"changes": the
    registry's entries, "about": its own description, "source_name": its
    file's name}``, is the last section of Sources (the book's registry of
    reporting changes; round 3 folded its tab into Sources).

    ``synthesis``, when given, adds the Synthesis tab (synthesis_tab.py):
    ``{label: {industry: PATH, filed: PATH or None, peers: [(NAME, PATH)],
    translation, peer_percentiles (four, or None without peers), leads
    (four)}}``, every choice stated; ``synthesis_tab.from_flags`` reads the
    command line's flags into it. Without it the workbook is as before."""
    if frequency != "quarterly":
        raise ValueError("the workbook's calendar is quarterly; monthly series are not built yet")
    out = Path(out)
    if out.suffix.lower() != ".xlsx":
        raise ValueError("the workbook must be written to a .xlsx file, not %s" % out.name)
    for lbl, path in series:
        # Saving over an input would destroy the data it was built from.
        if out.resolve() == Path(path).resolve():
            raise ValueError("the workbook would be written over %s, the series for %s; write it elsewhere"
                             % (Path(path).name, lbl))
    if not all(math.isfinite(v) for v in
               (moderate_halfwidth, low_step, high_z, *half_lives, *percentiles)):
        raise ValueError("every setting must be a number, not nan or infinity")
    if moderate_halfwidth < 0 or low_step <= moderate_halfwidth or high_z <= 0:
        raise ValueError("need 0 <= moderate_halfwidth < low_step, and high_z above 0")
    if on_the_line not in ("worse", "better"):
        raise ValueError("on_the_line must be 'worse' or 'better'")
    for lbl, _ in series:
        if not lbl.strip():
            raise ValueError("every series needs a product label")
        if not label_ok(lbl):
            raise ValueError("the label %r can't be used: %s" % (lbl, LABEL_RULE))
    if len({lbl.casefold() for lbl, _ in series}) != len(series):
        raise ValueError("two series share a label (Excel ignores capitals, so 'Cards' and 'cards' are one)")
    if synthesis:
        ST.validate(synthesis, [lbl for lbl, _ in series], direction, on_the_line)
        for lbl, s_ in synthesis.items():
            for what, path in ([("the industry series", s_["industry"])]
                               + ([("KeyBank's filed series", s_["filed"])] if s_.get("filed") else [])
                               + [("the peer %s" % nm, p_) for nm, p_ in s_.get("peers") or []]):
                if out.resolve() == Path(path).resolve():
                    raise ValueError("the workbook would be written over %s, %s for %s; write it elsewhere"
                                     % (Path(path).name, what, lbl))
    higher = direction == "higher_is_worse"
    floor_off = tuple(floor_off or ())
    if floor_off and not decide:
        raise ValueError("floor_off names measures, which only a book with Decide has")

    def floored(label: str) -> bool:
        """Whether a value at or below zero is always Low for this series: the
        workbook's floor, unless its measure is in floor_off (a difference
        between two rates has no floor)."""
        return bool(floor_at_zero and higher and not (decide and decide["parts"][label][1] in floor_off))

    floor_words = ("" if not (floor_at_zero and higher) else
                   " A value at or below zero is always Low" + (
                       ("" if not floor_off else ", except on %s, which %s no floor at zero"
                        % (", ".join(floor_off), "has" if len(floor_off) == 1 else "have")) + "."))
    # Thresholds' columns: the book's layout with Decide (round 2), the plain
    # one without. Every formula below takes its letters from here.
    TC = DT.TH_BOOK if decide else DT.TH_PLAIN

    loaded = []
    for label, path in series:
        pts = read_series(Path(path))
        ev = evidence(pts, label, unit, direction, frequency, smoothing, None, floored(label))
        # Life to date (the firm, 14 Oct 2026): a series shorter than the
        # backtest needs keeps its lines; only its backtest says it was not run.
        try:
            bt = backtest(pts, label, unit, direction, frequency, smoothing, percentiles,
                          horizon, min_history, floored(label), half_lives)
        except ValueError as exc:
            if "at least min_history" not in str(exc):
                raise
            bt = {"not_tested": "Not tested: %s, so there is no quarter with %d before it and %d after it"
                                % (exc, min_history, horizon), "runs": []}
        # Smoothing comes before anything else, as the window says: the live
        # lines, the latest value and the score are built from the series the
        # evidence and backtest use, never from the raw file beside it.
        used = smooth(check(pts, frequency), smoothing)
        by_q = {quarter(p.date): p.value for p in used}
        if len(by_q) != len(used):
            raise ValueError("%s has two observations in one quarter" % label)
        loaded.append(dict(label=label, path=Path(path), raw=pts, points=used, ev=ev, bt=bt, by_q=by_q))
    syn_loaded = ST.load(synthesis, smoothing) if synthesis else None
    grid = sorted(set().union(*(s["by_q"] for s in loaded)))
    # Decimals to suit the values: 3.978 (a rate) keeps three, 714.5 (a score)
    # one; three decimals on a credit score read as false precision.
    biggest = max(abs(p.value) for s in loaded for p in s["points"])
    NF = "0.0" if biggest >= 100 else "0.00" if biggest >= 10 else "0.000"
    n, first, last = len(loaded), 2, len(grid) + 1

    if decide:
        DT.validate(decide["names"])
        if "lines_from" not in decide:
            raise ValueError("decide needs lines_from: the source every row's Lines set from starts as, or None "
                             "to start blank (none has a default)")
        missing = [lbl for lbl, _ in series if lbl not in decide["parts"]]
        if missing:
            raise ValueError("no parts given for %s" % ", ".join(missing))
        for lbl, _ in series:
            if lbl != " - ".join(decide["parts"][lbl]):
                raise ValueError("the series %r should be labelled %r, its parts joined by ' - '"
                                 % (lbl, " - ".join(decide["parts"][lbl])))
    if into is not None:
        wb = into
        asx = wb.create_sheet("Assess", 0)
    else:
        wb = Workbook()
        asx = wb.active
        asx.title = "Assess"                 # first, and where it opens: where the bank works
    th, ch, evs, bts, srcs, gl, da, run = (wb.create_sheet(t) for t in
                                           ("Thresholds", "Chart", "Evidence", "Backtest", "Sources", "Glossary",
                                            "Data", "Run"))
    sy = sd = None
    if synthesis:
        # Beside Thresholds, the other live tab; its data beside Data.
        sy = wb.create_sheet("Synthesis", wb.sheetnames.index("Thresholds") + 1)
        sd = wb.create_sheet("SynthesisData", wb.sheetnames.index("Data") + 1)
        K.hide_gridlines(sy)
    for ws in (th, asx, ch, evs, bts, srcs, gl, run):
        K.hide_gridlines(ws)
    lay = assess_layout(n)
    LO_TOP, lo_end, set_end = lay["lo_top"], lay["lo_end"], lay["set_end"]
    # The leave-out list, as ranges: the product typed, the calendar positions
    # of From and To (hidden helpers), whether the row is whole, and whether it
    # is in use at all.
    lo_rng = lambda c_: "Assess!${c}${a}:${c}${b}".format(c=c_, a=LO_TOP, b=lo_end)     # noqa: E731
    LO_PROD, LO_FROM, LO_TO, LO_OK, LO_USED = (lo_rng(c_) for c_ in ("N",) + tuple(L(HC + 14 + i) for i in range(4)))

    # ---- Data: one calendar, raw values, and the values each product keeps ----
    da.append(["Quarter"] + [s["label"] for s in loaded] + ["%s, kept" % s["label"] for s in loaded]
              + ["%s, change over 4 quarters kept" % s["label"] for s in loaded]
              + ["%s, distance from the median change" % s["label"] for s in loaded]
              + ["%s, distance from the median kept" % s["label"] for s in loaded]
              + ["%s, change over 4 quarters, every quarter" % s["label"] for s in loaded]
              + ["%s, distance from the median change, every quarter" % s["label"] for s in loaded]
              + ["%s, distance from the median, every quarter" % s["label"] for s in loaded])
    srow = {s["label"]: SET_TOP + i for i, s in enumerate(loaded)}      # Your settings row per product
    for r, q in enumerate(grid, start=2):
        da.cell(r, 1, q)
        for k, s in enumerate(loaded):
            raw = L(2 + k)
            if q in s["by_q"]:
                da.cell(r, 2 + k, s["by_q"][q])
            # Left out when a whole row of the list names this product and its
            # From..To takes in this quarter's place on the calendar. An exact
            # match on the product (SUMPRODUCT, not COUNTIFS, which would read
            # * and ? in a label as wildcards).
            left_out = ('SUMPRODUCT(({p}=Thresholds!$A${t})*({f}<={pos})*({to}>={pos})*({ok}=1))>0'
                        .format(p=LO_PROD, f=LO_FROM, to=LO_TO, ok=LO_OK, t=5 + k, pos=r - 1))
            da.cell(r, 2 + n + k,
                    '=IF({raw}{r}="","",IF({lo},"",{raw}{r}))'.format(raw=raw, r=r, lo=left_out))
            kc = L(2 + n + k)
            if r >= first + 4:
                da.cell(r, 2 + 2 * n + k, '=IF(AND(ISNUMBER({c}{r}),ISNUMBER({c}{p})),{c}{r}-{c}{p},"")'
                        .format(c=kc, r=r, p=r - 4))
            else:
                da.cell(r, 2 + 2 * n + k, "")
            yc = L(2 + 2 * n + k)
            da.cell(r, 2 + 3 * n + k, '=IF({y}{r}="","",ABS({y}{r}-Thresholds!${o}${tr}))'
                    .format(y=yc, r=r, tr=5 + k, o=TC["med_change"]))
            da.cell(r, 2 + 4 * n + k, '=IF({c}{r}="","",ABS({c}{r}-Thresholds!${b}${tr}))'
                    .format(c=kc, r=r, tr=5 + k, b=TC["median"]))
            # The same three, with nothing left out: the Scenarios on Assess
            # set them beside the lines as the bank has them.
            if r >= first + 4:
                da.cell(r, 2 + 5 * n + k, '=IF(AND(ISNUMBER({c}{r}),ISNUMBER({c}{p})),{c}{r}-{c}{p},"")'
                        .format(c=raw, r=r, p=r - 4))
            ac = L(2 + 5 * n + k)
            da.cell(r, 2 + 6 * n + k, '=IF({a}{r}="","",ABS({a}{r}-Thresholds!${x}${tr}))'.format(a=ac, r=r, tr=5 + k,
                                                                                     x=TC["med_change_all"]))
            da.cell(r, 2 + 7 * n + k, '=IF({c}{r}="","",ABS({c}{r}-Thresholds!${x}${tr}))'.format(c=raw, r=r, tr=5 + k,
                                                                                     x=TC["med_all"]))
    K.freeze_below(da, 1)

    # ---- Thresholds: live. Its columns come from TC: the plain layout, or
    # with Decide the book's (round 2): Product, Measure and Source split,
    # the lines, the latest quarter, the decision and the Check in A..P, the
    # series label kept in A for every lookup, and the workings hidden. ----
    c = TC
    if decide:
        d_pairs, _ = DT.pairs_of([s["label"] for s in loaded], decide["parts"])
    sg = 1 if higher else -1
    for k, s in enumerate(loaded):
        r = 5 + k
        sr = srow[s["label"]]
        fz = "TRUE" if floored(s["label"]) else "FALSE"
        col = lambda c_: "Data!$%s$%d:$%s$%d" % (L(c_), first, L(c_), last)        # noqa: E731
        kept, yoy, dev = col(2 + n + k), col(2 + 2 * n + k), col(2 + 3 * n + k)
        levdev = col(2 + 4 * n + k)
        at = lambda key: "%s%d" % (c[key], r)                                      # noqa: E731
        put = lambda key, v: th["%s%d" % (c[key], r)].__setattr__("value", v)      # noqa: E731
        put("series", s["label"])
        put("median", "=MEDIAN(%s)" % kept)
        put("med_change", "=MEDIAN(%s)" % yoy)                       # helper: the median four-quarter change
        put("move", "=1.4826*MEDIAN(%s)" % dev)
        put("worst", "=%s(%s)" % ("MAX" if higher else "MIN", kept))
        put("spread", "=1.4826*MEDIAN(%s)" % levdev)                  # helper: the spread of levels
        # Helpers for the Scenarios: the same, keeping every quarter.
        raw_all, chg_all = col(2 + k), col(2 + 5 * n + k)
        put("med_all", "=MEDIAN(%s)" % raw_all)
        put("med_change_all", "=MEDIAN(%s)" % chg_all)
        put("move_all", "=1.4826*MEDIAN(%s)" % col(2 + 6 * n + k))
        put("spread_all", "=1.4826*MEDIAN(%s)" % col(2 + 7 * n + k))
        hw, ls, hz = ("Assess!$%s$%d" % (c_, sr) for c_ in "OPQ")
        B, C_, D, E, G, H_, P = (at(x) for x in ("median", "move", "worst", "l2", "l4", "l5", "spread"))
        # A period left out for this product must be whole: two quarters from
        # the calendar, in order, and a reason. Half of one would silently keep
        # every quarter, so the product is refused until the row is fixed.
        bad_lo = ('SUMPRODUCT(({p}=${a}{r})*({u}=1)*({ok}=0))>0'
                  .format(p=LO_PROD, u=LO_USED, ok=LO_OK, a=c["series"], r=r))
        put("refusal", ('=IF({lo},"Refused: a period left out for this product is not whole; its Status on Assess '
                        'says what is missing",'
                        'IF(OR(NOT(ISNUMBER({hw})),{hw}<0),"Refused: Moderate\'s half-width must be 0 or more",'
                        'IF(OR(NOT(ISNUMBER({hz})),{hz}<=0),"Refused: High\'s z must be above 0",'
                        'IF(NOT(ISNUMBER({ls})),"Refused: Moderate-Low\'s step must be a number",'
                        'IF(ISERROR({C}),"Refused: too little history kept for a typical yearly move",'
                        'IF(AND({fz},{B}<=0),"Refused: the median is at or below zero",'
                        'IF({ls}<={hw},"Refused: Moderate-Low must begin further from the median than Moderate",'
                        'IF({sg}*({H}-{G})<=0,"Refused: High must begin beyond Moderate-High\'s line",'
                        'IF({sg}*({D}-{H})<0,"Refused: High would begin beyond the worst quarter kept",'
                        'IF(AND({fz},{E}<=0),"Refused: Moderate-Low would begin at or below zero",""))))))))))')
            .format(fz=fz, ls=ls, hw=hw, hz=hz, sg=sg, lo=bad_lo, B=B, C=C_, D=D, E=E, G=G, H=H_))
        # What to do about the refusal, in the same order, with the numbers
        # that would clear it worked out live.
        went = "higher" if higher else "lower"
        put("fix", ('=IF({lo},"Finish that row under Periods left out (product, From, To and a reason), or clear '
                    'all four of its cells.",'
                    'IF(OR(NOT(ISNUMBER({hw})),{hw}<0),"Type 0 or more in Width of Moderate.",'
                    'IF(OR(NOT(ISNUMBER({hz})),{hz}<=0),"Type a number above 0 in How rare High is.",'
                    'IF(NOT(ISNUMBER({ls})),"Type a number in Moderate-Low begins.",'
                    'IF(ISERROR({C}),"Put back a period, so that more than a year of quarters is kept.",'
                    'IF(AND({fz},{B}<=0),"Put back periods so the quarters kept center above zero.",'
                    'IF({ls}<={hw},"Raise Moderate-Low begins above "&{hw}&", or lower Width of Moderate below "'
                    '&{ls}&".",'
                    'IF({sg}*({H}-{G})<=0,"Raise How rare High is above "&TEXT(ROUNDUP({hw}*{C}/{P},2),'
                    '"0.00")&".",'
                    'IF({sg}*({D}-{H})<0,"Lower How rare High is to "&TEXT(ROUNDDOWN({sg}*({D}-{B})/{P},2),'
                    '"0.00")&" or less, or put back a period that went {went}.",'
                    'IF(AND({fz},{E}<=0),"Lower Moderate-Low begins to "&TEXT(ROUNDDOWN({B}/{C},2),"0.00")'
                    '&" or less.",""))))))))))')
            .format(fz=fz, ls=ls, hw=hw, hz=hz, sg=sg, lo=bad_lo, went=went, B=B, C=C_, D=D, E=E, G=G, H=H_,
                    P=P))
        mb, mv = "$%s%d" % (c["median"], r), "$%s%d" % (c["move"], r)
        put("l4", "={b}+{sg}*{hw}*{c}".format(b=mb, sg=sg, hw=hw, c=mv))
        put("l3", "={b}-{sg}*{hw}*{c}".format(b=mb, sg=sg, hw=hw, c=mv))
        put("l2", "=IF(ISNUMBER({ls}),{b}-{sg}*{ls}*{c},\"\")".format(b=mb, sg=sg, ls=ls, c=mv))
        put("l5", "={b}+{sg}*{hz}*${p}{r}".format(b=mb, sg=sg, hz=hz, p=c["spread"], r=r))
        lp = s["points"][-1]
        put("quarter", quarter(lp.date))
        put("value", lp.value)
        op_worse, op_better = ("<=", "<") if higher else (">=", ">")
        bounds = "${a}{r}:${b}{r}".format(a=c["l2"], b=c["l5"], r=r)
        J = at("value")
        floor = "AND({j}<=0,{f})".format(j=J, f=fz)
        put("rating", '=IF(${m}{r}<>"","",IF({floor},1,1+IF(Assess!$R${sr}="worse",'
                      'COUNTIF({b},"{w}"&{j}),COUNTIF({b},"{bt}"&{j}))))'
            .format(m=c["refusal"], r=r, floor=floor, sr=sr, b=bounds, w=op_worse, bt=op_better, j=J))
        put("rating_name", '=IF({k}="","",CHOOSE({k},"{0}","{1}","{2}","{3}","{4}"))'.format(*RATINGS, k=at("rating")))
        for key in ("median", "move", "worst", "l2", "l3", "l4", "l5", "value"):
            th[at(key)].number_format = NF
        if decide:
            # The book's own columns: the parts, the rating as a person reads
            # it, the decision typed on Decide, and the Check in one cell.
            put("floor", floored(s["label"]))                # read by Decide's rating formulas
            p_ = decide["parts"][s["label"]]
            for key, v in zip(("product", "measure", "source"), p_):
                put(key, v)
            put("rating_shown", '=IF({m}<>"","Refused",{k}&" "&{l})'.format(m=at("refusal"), k=at("rating"),
                                                                             l=at("rating_name")))
            # The decision shows on the source whose lines its row uses.
            dc = DT.columns(decide["names"])
            dr = DT.HDR + 1 + d_pairs.index(tuple(p_[:2]))
            put("decision", '=IF(Decide!{f}{r}="{s}",Decide!{d}{r},"")'.format(
                f=L(dc[DT.LINES_FROM]), d=L(dc["Decision"]), r=dr, s=p_[2].replace('"', '""')))
            put("check", '=IF({m}="","OK","To fix: "&{f})'.format(m=at("refusal"), f=DT.HERE.format(m=at("fix"))))
    if decide:
        # Which series the lines refuse as built (the starting settings, no
        # period left out): their rows are set three lines high, so the To fix
        # reads whole. The same rules as the Check (scale.scale).
        refused_at_build = [s["label"] for s in loaded if _scale(
            [p.value for p in s["points"]], direction, moderate_halfwidth, low_step, high_z,
            floored(s["label"]))["refused"]]
        _thresholds_book(th, name, n, [s["label"] for s in loaded], decide, floor_words, refused_at_build)
    else:
        _thresholds_plain(th, name, n, loaded, floor_at_zero and higher)

    # ---- Chart: every product on one scale, switched on and off in one row,
    # with one product's five score bands behind them (the hand-off) ----
    ch.sheet_properties.tabColor = K.INK
    K.light_title(ch, 1, 2 + n + 2, "%s · every series on one scale" % name,
                  "TRUE in row 5 shows a series; FALSE hides it. The shaded bands are the five ratings of the "
                  "series picked under Bands for.", live=False)
    ch.column_dimensions["A"].width = 2.71
    for k, s in enumerate(loaded):
        col = 2 + k
        c = ch.cell(4, col, s["label"])
        c.font = Font(name="Arial", bold=True, size=9, color=K.SLATE)
        c.alignment = Alignment(wrap_text=True, vertical="bottom")
        c.border = Border(bottom=Side(style="thick", color=K.SERIES[k % len(K.SERIES)]))
        sw_c = ch.cell(5, col, k < 2)
        sw_c.fill = K.solid(K.INPUT)
        sw_c.border = Border(*(Side(style="thin", color=K.INPUT_EDGE),) * 4)
        sw_c.font = Font(name="Calibri", size=11, bold=True)
        ch.column_dimensions[L(col)].width = 15
    ch.row_dimensions[4].height = 25.5
    sw = DataValidation(type="list", formula1='"TRUE,FALSE"', showInputMessage=True, promptTitle="Show",
                        prompt="TRUE draws this product; FALSE hides it.")
    ch.add_data_validation(sw)
    sw.add("B5:%s5" % L(1 + n))
    pc = 3 + n                                            # the Bands for picker, after a gap
    ch.cell(4, pc, "Bands for").font = Font(name="Arial", bold=True, size=9, color=K.SLATE)
    ch.cell(4, pc).alignment = Alignment(vertical="bottom")
    pick = ch.cell(5, pc, loaded[0]["label"])
    pick.fill = K.solid(K.INPUT)
    pick.border = Border(*(Side(style="thin", color=K.INPUT_EDGE),) * 4)
    pick.font = Font(name="Calibri", size=11, bold=True)
    ch.column_dimensions[L(pc)].width = 18
    # The pickers read their choices from cells: an inline list longer than
    # 255 characters is one Excel drops or repairs on open.
    # Hidden helpers from column AN, or past the last series' switch, the
    # picker and the chart when there are more than 30 series: at a fixed AN
    # the title over 36 series ran into them and the build failed.
    h0 = max(40, n + 10)
    lists_col = h0 + n + 16
    for i, v in enumerate(["(none)"] + [s["label"] for s in loaded]):
        ch.cell(1 + i, lists_col, v)
    pv = DataValidation(type="list", formula1="=$%s$1:$%s$%d" % (L(lists_col), L(lists_col), 1 + n))
    ch.add_data_validation(pv)
    pv.add(pick.coordinate)
    ch.cell(1, h0, "Quarter")
    for k, s in enumerate(loaded):
        ch.cell(1, h0 + 1 + k, s["label"])
    band_names = ["1 Low", "2 Moderate-Low", "3 Moderate", "4 Moderate-High", "5 High"]
    bc0 = h0 + 1 + n                                      # base, five bands, top/floor
    ch.cell(1, bc0, " ")
    order_c = list(range(5)) if higher else list(range(4, -1, -1))
    for j, i in enumerate(order_c):
        ch.cell(1, bc0 + 1 + j, band_names[i])
    tf = bc0 + 6                                          # top in row 2, floor in row 3, drawable in row 4
    cutsc = ["IFERROR(INDEX(Thresholds!${c}$5:${c}${e},MATCH(${pc}$5,Thresholds!$A$5:$A${e},0))+0,NA())"
             .format(c=TC["l%d" % (i + 2)], e=4 + n, pc=L(pc)) for i in range(4)]
    if not higher:
        cutsc = cutsc[::-1]
    raw_all = "Data!$B${a}:${e}${b}".format(a=first, b=last, e=L(1 + n))
    ch.cell(2, tf, "=MAX(MAX({r}),{c})*1.08".format(r=raw_all, c=",".join(cutsc)))
    if floor_at_zero:
        ch.cell(3, tf, "=MIN(0,MIN({r}),{c})".format(r=raw_all, c=",".join(cutsc)))
    else:
        ch.cell(3, tf, "=MIN(MIN({r}),{c})-0.08*(MAX(MAX({r}),{c})-MIN(MIN({r}),{c}))".format(r=raw_all,
                                                                                         c=",".join(cutsc)))
    ch.cell(4, tf, "=IF(ISNUMBER(SUM({c})),1,0)".format(c=",".join(cutsc)))
    T_, F_, D_ = ("${c}${r}".format(c=L(tf), r=r_) for r_ in (2, 3, 4))
    for r in range(first, last + 1):
        ch.cell(r, h0, "=Data!A%d" % r)
        for k in range(n):
            col = L(2 + k)
            ch.cell(r, h0 + 1 + k, '=IF(AND(${c}$5,ISNUMBER(Data!{d}{r})),Data!{d}{r},NA())'
                    .format(c=L(2 + k), d=col, r=r))
        ch.cell(r, bc0, "=IF({d}<>1,NA(),{f})".format(d=D_, f=F_))
        ch.cell(r, bc0 + 1, "=IF({d}<>1,NA(),{a}-{f})".format(d=D_, a=cutsc[0], f=F_))
        for j in range(1, 4):
            ch.cell(r, bc0 + 1 + j, "=IF({d}<>1,NA(),{b}-{a})".format(d=D_, a=cutsc[j - 1], b=cutsc[j]))
        ch.cell(r, bc0 + 5, "=IF({d}<>1,NA(),MAX({t}-{a},0))".format(d=D_, t=T_, a=cutsc[3]))
    lc = LineChart()
    lc.add_data(Reference(ch, min_col=h0 + 1, max_col=h0 + n, min_row=1, max_row=last), titles_from_data=True)
    lc.set_categories(Reference(ch, min_col=h0, min_row=2, max_row=last))
    for k, s_ in enumerate(lc.series):
        s_.smooth = False
        s_.marker.symbol = "none"
        s_.graphicalProperties.line = AT._ln(K.SERIES[k % len(K.SERIES)], 2)
    bands_ch = AreaChart()
    bands_ch.grouping = "stacked"
    bands_ch.add_data(Reference(ch, min_col=bc0, max_col=bc0 + 5, min_row=1, max_row=last), titles_from_data=True)
    bands_ch.series[0].graphicalProperties.noFill = True
    bands_ch.series[0].graphicalProperties.line.noFill = True
    for j, s_ in enumerate(bands_ch.series[1:]):
        s_.graphicalProperties.solidFill = K.BANDS[order_c[j]]
        s_.graphicalProperties.line.noFill = True
    AT._quiet_axes(lc, 20)
    lc.y_axis.number_format = "0.0"
    lc.x_axis.tickLblPos = "low"
    # The swatches over row 5 are the products' key; the bands' key is the
    # note under the chart (Excel 16 ignores deleted legend entries, so a
    # legend would also list the bands' invisible base).
    lc.legend = None
    lc.visible_cells_only = False
    lc += bands_ch
    AT._square(lc)
    AT._place(ch, lc, 1, 7, 2 + n + 6, 34)
    for col in range(h0, lists_col + 1):
        ch.column_dimensions[L(col)].hidden = True

    # ---- Assess: the bank's tab, built by assess_tab ----
    AT.build(asx, loaded=loaded, grid=grid, first=first, last=last, n=n, higher=higher,
             floor_at_zero=floor_at_zero, name=name, unit=unit, moderate_halfwidth=moderate_halfwidth,
             low_step=low_step, high_z=high_z, on_the_line=on_the_line, quarter=quarter,
             y_min=None if floor_at_zero else _axis_floor(loaded, direction, moderate_halfwidth, low_step, high_z),
             nf=NF, th_cols=TC)

    # ---- Synthesis: internal, industry and peers, where each colour starts ----
    syn_at = {}
    if synthesis:
        syn_at = ST.build(wb, sy, sd, asx, loaded=loaded, syn=synthesis, syn_loaded=syn_loaded, grid=grid, n=n,
                          higher=higher, floor_at_zero=floor_at_zero, unit=unit, nf=NF, srow=srow,
                          quarter=quarter, assess_lay=lay)
        sd.sheet_properties.tabColor = K.STONE

    # ---- The reference tabs: one quiet template (Claude Design's round 2,
    # docs/design/round-2/README.md, "Reference tabs"). Column A is a margin;
    # row 1 the title, row 2 the purpose over a KEY_RED rule, the header in
    # row 4 on MIST; Calibri 11 on 15 pt rows with MIST hairlines and no
    # fills. Series blocks start with a heading over a thin INK rule. ----
    for ws in (evs, bts, srcs, gl, run):
        ws.sheet_properties.tabColor = K.TAB_REFERENCE
    da.sheet_properties.tabColor = K.TAB_DATA
    DATA_F = REF_BODY
    # A series is named by its parts where the book gives them.
    named = (lambda lbl: " · ".join(decide["parts"][lbl])) if decide else (lambda lbl: lbl)    # noqa: E731
    # Product, Measure and Source where the book gives them; a file it does
    # not name (a synthesis file) under its own name in the first column.
    ids_of = (lambda lbl: list(decide["parts"][lbl]) if decide and lbl in decide["parts"]          # noqa: E731
              else [lbl] + [None] * (len(decide["names"]) - 1 if decide else 0))

    # ---- Evidence: two Excel tables, Spells and Departures, every row named
    # by its series' parts (round 3; reference_tabs.evidence) ----
    id_names = list(decide["names"]) if decide else ["Series"]
    RT.evidence(evs, [dict(ids=ids_of(s["label"]), ev=s["ev"]) for s in loaded], id_names, quarter)

    # ---- Backtest: a row per series and weighting; the series' parts in
    # their own columns where the book gives them ----
    n_scores = len(percentiles) + 1
    id_head = list(decide["names"]) if decide else ["Series"]
    ni = len(id_head)
    WT = 2 + ni                                         # Weighting; the numbers follow it
    bt_last = WT + 3 + n_scores
    _ref_title(bts, "Backtest: how predictive the scale would have been",
               "Each quarter is rated on lines at percentiles %s of the history up to and including it, then "
               "compared with the value %d quarters later." % ("/".join("%g" % p for p in percentiles), horizon),
               bt_last)
    for c1, c2, lbl in ((WT + 1, WT + 2, "Spearman: rating against"),
                        (WT + 3, WT + 2 + n_scores, "Average change ahead, by rating")):
        bts.cell(3, c1, lbl).font = Font(name="Arial", bold=True, size=9, color=K.INK)
        for c_ in range(c1, c2 + 1):
            bts.cell(3, c_).border = Border(bottom=Side(style="thin", color=K.INK))
    bts.row_dimensions[3].height = 15
    _ref_header(bts, 4, id_head + ["Weighting", "Value ahead", "Change ahead"]
                + ["%d" % k for k in range(1, n_scores + 1)] + ["Quarters tested"], 2,
                numeric=range(ni + 1, ni + 4 + n_scores))
    r = 5
    for s in loaded:
        ids = list(decide["parts"][s["label"]]) if decide else [s["label"]]
        if s["bt"].get("not_tested"):
            for col, v in enumerate(ids + [s["bt"]["not_tested"]], start=2):
                bts.cell(r, col, v).font = DATA_F
            _ref_row(bts, r, 2, bt_last)
            r += 1
        for run_ in s["bt"]["runs"]:
            vals_ = ids + ["none" if run_["half_life_years"] is None else "%g-yr half-life"
                           % run_["half_life_years"], run_["score_vs_level_ahead"], run_["score_vs_change_ahead"],
                           *[b["mean_change_ahead"] for b in run_["by_score"]], run_["periods_tested"]]
            for col, v in enumerate(vals_, start=2):
                bts.cell(r, col, v).font = DATA_F
            for col in range(WT + 1, bt_last):
                bts.cell(r, col).number_format = "0.00"
            bts.cell(r, bt_last).number_format = "0"
            _ref_row(bts, r, 2, bt_last)
            r += 1
    for i, t in enumerate(("Spearman is the rank correlation between the rating and what followed: 1 means a "
                           "higher rating always preceded a higher value, and 0 means no relationship.",
                           "Average change ahead is in the measure's own units.",
                           "Neighboring quarters overlap, so the quarters tested are not independent "
                           "observations.")):
        bts.cell(r + 1 + i, 2, t).font = Font(name="Calibri", size=10, color=K.SLATE)
    bts.column_dimensions["A"].width = REF_EDGE
    for i, w in enumerate((17.0, 31.5, 22.0)[:ni] if decide else (max(18, min(60, 2 + max(
            len(s["label"]) for s in loaded))),)):
        bts.column_dimensions[L(2 + i)].width = w
    bts.column_dimensions[L(WT)].width = 15
    for col in range(WT + 1, bt_last + 1):
        bts.column_dimensions[L(col)].width = 11
    bts.freeze_panes = bts.cell(5, WT + 1).coordinate

    # ---- Sources: what each file says about itself, in three sections
    # (round 3; reference_tabs.sources): the lines many files share, printed
    # once; one row per file with its own lines verbatim and the required
    # lines it does not give marked, never filled in; and the reporting
    # changes, when the book names their registry. ----
    # The synthesis' files (industry, KeyBank filed, each peer) follow the
    # series, each under its own name, by the same rule.
    described = [dict(label=s["label"], title=named(s["label"]), path=s["path"], raw=s["raw"]) for s in loaded]
    described += [dict(label=t_, title=t_, path=p_, raw=raw_)
                  for t_, p_, raw_ in (ST.files(syn_loaded) if synthesis else [])]
    files_ = [dict(ids=ids_of(s["label"]), path=Path(s["path"]), rows=len(s["raw"]), first=s["raw"][0].date,
                   last=s["raw"][-1].date, notes=read_notes(s["path"])) for s in described]
    RT.sources(srcs, files_, id_names, CORE_NOTES, reporting=reporting, change_key=RC.NOTE_KEY,
               whens=RC.whens_in_notes)

    # ---- Glossary: with Decide (the book) a Tab column first ----
    if decide:
        from .glossary import by_tab
        RT.glossary(gl, by_tab(), with_tab=True)
    else:
        RT.glossary(gl, TERMS, with_tab=False)

    # ---- Data: quarter in A, a series per column; the quarters before the
    # first value folded away. Rates to four decimals; a measure in the
    # hundreds (a score) keeps its own. ----
    groups = 8                                       # raw, kept, change, distances x2, the same keeping every quarter
    RT.data_layout(da, 1, first, last, 1 + groups * n, NF if biggest >= 100 else "0.0000")

    # ---- Run: This build (the series folder once), Left out as built when
    # the book wrote one, the files by name with their SHA-256, then the live
    # sections; every table on the same columns ----
    grid_r = RT.run_grid(id_names, [f["ids"] for f in files_], [f["path"].name for f in files_])
    LAST_R = grid_r["last_col"]
    folders = list(dict.fromkeys(str(Path(s["path"]).resolve().parent) for s in described))
    build_rows = [("threshold-engine", VERSION), ("Built", dt.datetime.now().isoformat(timespec="seconds"))]
    if smoothing > 1:
        build_rows.append(("Smoothing", "The Data tab holds each file's values as a trailing mean over %d quarters, "
                           "dated at the last; the first %d quarters of each file have no full window and are not "
                           "used. Every line, score and test is built from those values." % (smoothing, smoothing - 1)))
    build_rows += [("Measure", "%s (%s); %s; %s; %s" % (
                        name, unit, "higher is worse" if higher else "lower is worse",
                        "no smoothing" if smoothing == 1 else "smoothed over %d quarters" % smoothing,
                        ("a value at or below zero is always Low" + (
                            "" if not floor_off else ", except on %s" % ", ".join(floor_off)))
                        if floor_at_zero else "no floor at zero")),
                   ("Starting settings", "Width of Moderate %g; Moderate-Low begins %g; How rare High is %g; on a "
                    "line, rate it %s" % (moderate_halfwidth, low_step, high_z, on_the_line)),
                   ("Backtest half-lives", ", ".join("%g" % h for h in half_lives) or "none"),
                   ("Backtest settings", "percentiles %s; horizon %d; min history %d"
                    % ("/".join("%g" % p for p in percentiles), horizon, min_history))]
    r = RT.run_build(run, grid_r, build_rows, folders)
    if left_out_file is not None:
        r = RT.run_left_out(run, r + 1, grid_r, left_out_file)
    else:
        r += 1
    r = RT.run_files(run, r, grid_r, id_names,
                     [dict(ids=f["ids"], name=f["path"].name, rows=f["rows"], first=f["first"], last=f["last"],
                           sha=_sha(f["path"])) for f in files_], "Files \u00b7 %d" % len(files_))

    def live(r_, heading, names_at):
        """A live section's heading and header, its names at Run's columns."""
        last_ = max(LAST_R, max(names_at))
        r_ = RT.section(run, r_, heading, 2, last_)
        names = [None] * (last_ - 1)
        for c_, nm in names_at.items():
            names[c_ - 2] = nm
        return RT.header(run, r_, names, 2)

    G_ = grid_r
    r = live(r, "Periods left out (live)", {2: "Series", G_["first"]: "From", G_["last"]: "To", G_["file"]: "Reason",
                                             G_["sha"]: "Status"})
    for rr in range(LO_TOP, lo_end + 1):
        for col, c_ in zip((2, G_["first"], G_["last"], G_["file"]), "NOPQ"):
            run.cell(r, col, '=IF(Assess!{c}{r}="","",Assess!{c}{r})'.format(c=c_, r=rr)).font = DATA_F
        run.cell(r, G_["sha"], "=Assess!T%d" % rr).font = DATA_F
        RT.body(run, r, 2, LAST_R)
        r += 1
    r = live(r + 1, "Settings in use (live)", {2: "Series", G_["first"]: "Width of Moderate",
                                                G_["last"]: "Moderate-Low begins", G_["file"]: "How rare High is (z)",
                                                G_["sha"]: "On a line, rate it"})
    for s in loaded:
        run.cell(r, 2, s["label"]).font = DATA_F
        for col, c_ in zip((G_["first"], G_["last"], G_["file"], G_["sha"]), "OPQR"):
            run.cell(r, col, "=Assess!%s%d" % (c_, srow[s["label"]])).font = DATA_F
            run.cell(r, col).alignment = Alignment(horizontal="left")
        RT.body(run, r, 2, LAST_R)
        r += 1
    if synthesis:
        r = live(r + 1, "Synthesis choices, as on Synthesis (live)",
                 {2: "Series", G_["first"]: "Industry carried by", G_["file"]: "Taken from: lines 2, 3, 4, 5",
                  G_["sha"]: "Peers at percentile: lines 2, 3, 4, 5", LAST_R + 1: "Check"})
        run.column_dimensions[L(LAST_R + 1)].width = RT.px(300)
        for lbl, at in syn_at.items():
            lr = at["lines"]
            run.cell(r, 2, lbl).font = DATA_F
            run.cell(r, G_["first"], "=Synthesis!B%d&\"\"" % at["translation"]).font = DATA_F
            run.merge_cells(start_row=r, start_column=G_["first"], end_row=r, end_column=G_["rows"])
            run.cell(r, G_["file"], "=" + '&", "&'.join("Synthesis!F%d" % x for x in lr)).font = DATA_F
            run.cell(r, G_["sha"], ("=" + '&", "&'.join("Synthesis!D%d" % x for x in lr))
                     if synthesis[lbl].get("peers") else "no peers").font = DATA_F
            run.cell(r, LAST_R + 1, "=Synthesis!B%d" % at["check"]).font = DATA_F
            RT.body(run, r, 2, LAST_R + 1)
            r += 1

    dec = su = tt = None
    if decide:
        labels = [s["label"] for s in loaded]
        settings = {lbl: (moderate_halfwidth, low_step, high_z, on_the_line) for lbl in labels}
        dec = DT.build(wb, labels=labels, parts=decide["parts"], part_names=decide["names"], settings=settings,
                       lay=lay, n_quarters=len(grid), absent=decide["absent"], compare=decide["compare"],
                       lines_from=decide["lines_from"], unit=unit, higher=higher, floor_at_zero=floor_at_zero,
                       changes={s["label"]: RC.quarters_in_notes(read_notes(s["path"])) for s in loaded},
                       quarters=grid, floor_off=floor_off, peers=decide.get("peers"))
        # Decide is where the bank types and looks, so Assess (whose cells
        # Decide feeds) and Chart (every series on one scale, 95 switches)
        # are kept for their workings but hidden (the firm, 14 Oct 2026: "doesnt
        # this just look ridiculous"). Unhide them from any tab's right-click menu.
        wb["Assess"].sheet_state = "hidden"
        wb["Chart"].sheet_state = "hidden"

        # ---- Internal and Internal data (internal_tab.py): the bank's own
        # figures beside the outside series. Empty unless the book names an
        # internal file (decide["internal"]). ----
        from . import internal_tab as IT
        dec["internal"] = IT.build(wb, labels=labels, parts=decide["parts"], part_names=decide["names"], grid=grid,
                                   higher=higher, floor_at_zero=floor_at_zero, internal=decide.get("internal"))

        # ---- One set of lines per product and measure (the firm, 15 Oct
        # 2026: "we need to use all 3 to set appropriate thresholds"): the
        # per-series table becomes Lines by source, and Thresholds holds the
        # lines of the source each row picks, with every source rated on them.
        DT.rename_tab(wb, "Thresholds", DT.LINES_TAB)
        from . import thresholds_tab as TT
        su_lay = SU.layout(labels, decide["parts"], decide["compare"])
        tt = TT.build(wb, pairs=dec["pairs"], by_pair=dec["by_pair"], sources=su_lay["sources"],
                      dropped=su_lay["dropped"], part_names=decide["names"], decide=dec, n_series=n, name=name,
                      higher=higher, floor_at_zero=floor_at_zero, number_format="0.00", floor_off=floor_off,
                      refused_at_build=refused_at_build, start=decide["lines_from"])
        su = SU.build(wb, name=name, unit=unit, labels=labels, parts=decide["parts"], part_names=decide["names"],
                      absent=decide["absent"], compare=decide["compare"], notes=decide["notes"],
                      refs=TT.refs(tt, dec["by_pair"]))
        dec["thresholds"] = tt
        order = ["Summary", "Decide", "Thresholds", DT.LINES_TAB, "Assess", "Synthesis"]
        rest = [ws for ws in wb._sheets if ws.title not in order]
        wb._sheets = [wb[t] for t in order if t in wb.sheetnames] + rest
        wb.active = 0
        for ws in wb.worksheets:
            ws.sheet_view.tabSelected = ws.title == "Summary"
    if before_save is not None:
        before_save(wb)
    out = Path(out)
    wb.save(out)
    if decide:
        # The buttons over the Decisions table, written into the saved file
        # (openpyxl cannot write slicers); it opens on the first series.
        from . import slicers as SL
        first_parts = decide["parts"][loaded[0]["label"]]
        SL.add_table_slicers(out, DT.TITLE, dec["table"], [
            {"column": nm, "caption": nm, "anchor": anchor, "selected": v, "name": "Pick" + str(i + 1),
             "columns": cols, "style": DT.SLICER_STYLE, "button_height": DT.SLICER_BUTTON_IN}
            for i, (nm, v, anchor, cols) in enumerate(zip(decide["names"][:2], first_parts, DT.SLICER_ANCHORS,
                                                          DT.SLICER_COLUMNS))], hide_buttons=True)
    return {"path": out, "products": [s["label"] for s in loaded], "quarters": len(grid),
            "evidence": {s["label"]: s["ev"] for s in loaded}, "synthesis": syn_at, "decide": dec,
            "summary": su, "thresholds": tt}
