"""The Summary tab: every series' latest quarter and rating on one page.

Each product and measure has one set of lines (Thresholds), set from the
history of the source picked on Decide; every source's latest quarter is
rated on those lines.

Claude Design's round 3 (docs/design/round-3/handoff/README.md, "Summary"):
two panels side by side, each a header row per product (the first part) with
its measures (the second part) under it, and a pair of columns (latest,
rating) per value of the third part (source). The products are split between
the panels where the taller panel is shortest; "Needs your attention" goes
under the left panel.

Generic by construction. Nothing here names a product, a measure or a bank.
The source columns come in the book's order, except that a ``compare`` pair
leads: ``compare[0]`` (the bank) first, at the larger size and in bold, then
the points between it and ``compare[1]`` (the peer group), then
``compare[1]``, then every other source. Without a compare pair the first
source the book uses leads and there is no points column. "Needs your
attention" is about the leading source.

A source gets a column only when it has a series on at least a quarter of
the rows (``shown_sources``); the compare pair always does. A source with
fewer has its lines and latest quarters on Lines by source; one with a single
series is also printed as a strip in row 2, read live from Thresholds.

Every number is a live reference to its product and measure's row on
Thresholds (thresholds_tab.py), so it follows what is typed on Decide;
nothing is computed here but the points between the compare pair and the
attention block's ordering. The attention block uses INDEX/MATCH over hidden
helper columns rather than FILTER/SORTBY: it looks the same, it needs no
dynamic-array metadata in the file, and the test engine (``formulas``) can
recalculate it.
"""
from __future__ import annotations

import re
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L

from . import keybank_style as K
from .thresholds_tab import NOT_PICKED  # noqa: F401  (kept for callers that read it from here)

TITLE = "Summary"
GROUP, HEAD, FIRST = 4, 5, 6                 # group headers, column labels, first body row
ATTENTION_ROWS = 5                           # the 4s and 5s listed by name; the rest are counted
NO_LINES = "—"                               # a rating with no source picked yet
NONE_HOT = "None rated 4 or 5."
INK, SLATE = K.INK_TEXT, K.SLATE
#: The Thresholds cells Summary reads for each series, by meaning
#: (thresholds_tab.refs): the latest quarter and value, the rating as a number
#: and in words, the row's refusal ("Refused: ...", blank when the lines
#: stand) and the source its lines are set from (blank until one is picked).
TH_READ = ("quarter", "value", "rating", "rating_text", "refusal", "lines_from", "picked", "refused")
#: A source shown with columns has a series on at least this share of the rows.
SHOWN_SHARE = 0.25

# Widths in Excel units, from the hand-off's pixels ((px - 5) / 7).
W_MEASURE = 36.43                            # 260 px
W_LEAD_LATEST, W_LEAD_RATING, W_POINTS = 8.43, 17.0, 7.86       # 64, 124, 60 px
W_LATEST, W_RATING = 7.29, 15.29             # 56, 112 px
W_GUTTER, W_EDGE = 2.71, 1.57                # 24, 16 px
H_PRODUCT, H_BODY = 16.5, 15.75


def _order(values):
    out = []
    for v in values:
        if v not in out:
            out.append(v)
    return out


def _side(style, color):
    return Side(style=style, color=color)


def short(s: str) -> str:
    """A source's name up to its bracket: 'Peer group (15 banks)' is 'Peer
    group', 'Large banks (Philadelphia Fed)' is 'Large banks'."""
    return s.split(" (")[0]


def sources_in_order(cols: Sequence[str], compare: Optional[Tuple[str, str]]) -> List[str]:
    """The compare pair first (the bank, then its peer group), then the rest
    in the book's order."""
    if not compare:
        return list(cols)
    return [compare[0], compare[1]] + [c for c in cols if c not in compare]


def shown_sources(labels: Sequence[str], parts: Dict[str, Tuple[str, str, str]],
                  compare: Optional[Tuple[str, str]]) -> Tuple[List[str], List[str]]:
    """(shown, left to Lines by source): the compare pair and every source
    with a series on at least a quarter of the rows get columns; the rest do
    not. Both in the order Summary uses."""
    rows = _order((parts[lab][0], parts[lab][1]) for lab in labels)
    srcs = sources_in_order(_order(parts[lab][2] for lab in labels), compare)
    have = {s: sum(1 for lab in labels if parts[lab][2] == s) for s in srcs}
    keep = [s for s in srcs if (compare and s in compare) or have[s] >= SHOWN_SHARE * len(rows)]
    if not keep:
        keep = list(srcs)
    return keep, [s for s in srcs if s not in keep]


def split_panels(blocks: Sequence[Tuple[str, int]]) -> Tuple[List[str], List[str]]:
    """Split whole product blocks, in order, between the left and right
    panels where the taller panel is shortest (a block is its header row plus
    its rows; the left panel takes the earlier split on a tie). ``blocks`` is
    (product, rows) in order."""
    sizes = [n + 1 for _, n in blocks]
    best = None
    for k in range(1, len(blocks) + 1):
        tall = max(sum(sizes[:k]), sum(sizes[k:]))
        if best is None or tall < best[0]:
            best = (tall, k)
    k = best[1] if best else 0
    return [p for p, _ in blocks[:k]], [p for p, _ in blocks[k:]]


def layout(labels: Sequence[str], parts: Dict[str, Tuple[str, str, str]],
           compare: Optional[Tuple[str, str]]) -> dict:
    """Where everything sits on Summary, for the builder, the tests and the
    walkthrough: the sources shown in order (``sources``) and those left to
    Lines by source (``dropped``), each source's (latest, rating) offsets
    within a panel (the measure is offset 0), the two panels' first columns,
    each product's header row (``heads``), and for each label of a shown
    source the row and columns of its latest value and rating (``cells``)."""
    rows = _order((parts[lab][0], parts[lab][1]) for lab in labels)
    srcs, dropped = shown_sources(labels, parts, compare)
    products = _order(a for a, _ in rows)
    left, right = split_panels([(p, sum(1 for a, _ in rows if a == p)) for p in products])
    # Within a panel: measure, then latest and rating per source, with the
    # points after the leading source's pair (0-based offsets).
    offs, c = {}, 1
    for i, s in enumerate(srcs):
        offs[s] = (c, c + 1)
        c += 2
        if i == 0 and compare:
            offs["points"] = c
            c += 1
    width = c
    panels = [2, 2 + width + 1]                     # B, then past a one-column gutter
    placed, heads, ends = {}, {}, []
    for base, prods in zip(panels, (left, right)):
        r = FIRST
        for p in prods:
            heads[p] = (r, base)
            r += 1
            for pm in (x for x in rows if x[0] == p):
                placed[pm] = (r, base)
                r += 1
        ends.append(r - 1)
    cells = {}
    for lab in labels:
        if parts[lab][2] not in offs:
            continue
        r, base = placed[tuple(parts[lab][:2])]
        a, b = offs[parts[lab][2]]
        cells[lab] = (r, base + a, base + b)
    # A row no shown source has a series for, but a source left out does.
    only = {}
    for pm in rows:
        here = [parts[lab][2] for lab in labels if tuple(parts[lab][:2]) == pm]
        if here and not any(s in offs for s in here):
            only[pm] = here
    return {"rows": rows, "sources": srcs, "dropped": dropped, "offsets": offs, "width": width, "panels": panels,
            "left": left, "right": right, "last_col": panels[1] + width - 1, "placed": placed, "heads": heads,
            "ends": ends, "cells": cells, "only": only}


def _range_of(refs: Dict[str, Dict[str, str]], key: str) -> str:
    """The Thresholds column ``key`` over every row Summary reads, as one
    absolute range (``Thresholds!$D$5:$D$60``)."""
    got = [re.match(r"(.*!)\$([A-Z]+)\$(\d+)$", m[key]).groups() for m in refs.values()]
    sheet, col = got[0][0], got[0][1]
    rows_ = [int(g[2]) for g in got]
    return "%s$%s$%d:$%s$%d" % (sheet, col, min(rows_), col, max(rows_))


def build(wb, *, name: str, unit: str, labels: Sequence[str], parts: Dict[str, Tuple[str, str, str]],
          part_names: Sequence[str], absent: Sequence[dict], compare: Optional[Tuple[str, str]],
          notes: Sequence[str], refs: Dict[str, Dict[str, str]]) -> dict:
    """Add the Summary tab. ``refs`` gives each label's cells on Thresholds by
    meaning (``thresholds_tab.refs``), for the sources shown and left out
    alike. ``absent`` and ``notes`` (the book's own words) are not printed
    here: the footnotes are the three lines the tab's own words need."""
    ws = wb.create_sheet(TITLE)
    missing = [lab for lab in labels if any(k not in refs.get(lab, {}) for k in TH_READ)]
    if missing:
        raise ValueError("Summary needs the Thresholds cells of %s" % ", ".join(missing[:3]))

    def th(col, lab):
        return refs[lab][col]

    p0, p1, p2 = part_names
    lay = layout(labels, parts, compare)
    rows, srcs, offs, width = lay["rows"], lay["sources"], lay["offsets"], lay["width"]
    panel_left, left, right, last_col = lay["panels"], lay["left"], lay["right"], lay["last_col"]
    lead = srcs[0]
    by = {parts[lab]: lab for lab in labels}
    lines_from = {pm: refs[by[next(k for k in by if k[:2] == pm)]]["lines_from"] for pm in rows}

    ws.sheet_properties.tabColor = K.TAB_WORK
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 100

    lead_rows = [(a, b) for a, b in rows if (a, b, lead) in by]
    first_lead = by[lead_rows[0] + (lead,)] if lead_rows else labels[0]
    ws["B1"] = '="%s · where each series stands, "&%s' % (name.replace('"', '""'), th("quarter", first_lead))
    ws["B1"].font = Font(name="Arial", size=16, bold=True, color=K.INK)
    ws["B1"].alignment = Alignment(vertical="bottom")
    ws["B2"] = ("Each rating places the latest quarter on its %s and %s's one set of lines, from 1 Low to 5 High, "
                "set on Decide." % (p0.lower(), p1.lower()))
    ws["B2"].font = Font(name="Calibri", size=11, color=SLATE)
    ws["B2"].alignment = Alignment(vertical="center")
    for col in range(2, last_col + 1):
        ws.cell(2, col).border = Border(bottom=_side("medium", K.KEY_RED))
    for r_, h in ((1, 25.5), (2, 15), (3, 6), (GROUP, 16.5), (HEAD, 16.5)):
        ws.row_dimensions[r_].height = h

    lab_font = Font(name="Arial", size=9, bold=True, color=SLATE)
    for base in panel_left:
        col = lambda off: base + off                # noqa: E731
        for i, s in enumerate(srcs):
            a, b = offs[s]
            lead_ = i == 0
            span_end = offs["points"] if lead_ and compare else b
            fill_ = K.solid(K.INK if lead_ else K.SLATE)
            top = _side("medium", K.KEY_RED if lead_ else K.SLATE)
            ws.merge_cells(start_row=GROUP, start_column=col(a), end_row=GROUP, end_column=col(b))
            for cc in range(a, span_end + 1):
                cell = ws.cell(GROUP, col(cc))
                cell.fill = fill_
                cell.border = Border(top=top, left=_side("thin", "7A766F") if (cc == a and i > 1) else None)
            h = ws.cell(GROUP, col(a), s)
            h.font = Font(name="Arial", size=10, bold=True, color=K.PAPER)
            h.alignment = Alignment(vertical="center", indent=0)
            if lead_ and compare:
                v = ws.cell(GROUP, col(offs["points"]), "vs %s" % short(compare[1]).split(" ")[0])
                v.font = Font(name="Arial", size=9, bold=True, color=K.PAPER)
                v.alignment = Alignment(horizontal="right", vertical="center")
        heads = {0: p1}
        for s in srcs:
            heads[offs[s][0]] = "Latest"
            heads[offs[s][1]] = "Rating"
        if compare:
            heads[offs["points"]] = "Points"
        for off in range(width):
            cell = ws.cell(HEAD, col(off), heads.get(off))
            cell.font = lab_font
            cell.border = Border(bottom=_side("medium", K.INK))
            numeric = off >= 1 and (off == offs.get("points") or any(off == offs[s][0] for s in srcs))
            cell.alignment = Alignment(horizontal="right" if numeric else "left", vertical="center")

    hair = _side("thin", "EFEBE4")
    ink = _side("thin", K.INK)
    placed, ends = lay["placed"], lay["ends"]
    small = Font(name="Calibri", size=10, color=SLATE)
    for base, prods in zip(panel_left, (left, right)):
        for p in prods:
            hr, _ = lay["heads"][p]
            for off in range(width):
                ws.cell(hr, base + off).border = Border(top=ink)
            if compare:
                ws.cell(hr, base + offs["points"]).border = Border(top=ink, right=_side("thin", K.STONE))
            pr = ws.cell(hr, base, p)
            pr.font = Font(name="Calibri", size=12, bold=True, color=INK)
            pr.alignment = Alignment(vertical="center")
            ws.row_dimensions[hr].height = H_PRODUCT
            for k, (a_, b_) in enumerate(x for x in rows if x[0] == p):
                r = placed[(a_, b_)][0]
                top = None if k == 0 else hair
                for off in range(width):
                    ws.cell(r, base + off).border = Border(top=top)
                me = ws.cell(r, base, b_)
                me.font = Font(name="Calibri", size=12, color=INK)
                me.alignment = Alignment(vertical="center")
                lf = lines_from[(a_, b_)]
                for i, s in enumerate(srcs):
                    a, b = offs[s]
                    lab = by.get((a_, b_, s))
                    big = i == 0
                    val = ws.cell(r, base + a, '=IF(ISNUMBER({v}),{v},"")'.format(v=th("value", lab)) if lab else None)
                    val.number_format = "0.00"
                    val.font = Font(name="Calibri", size=12, bold=True, color=INK) if big else small
                    val.alignment = Alignment(horizontal="right", vertical="center")
                    rt = ws.cell(r, base + b, '=IF({d}="","{n}",{t}&"")'.format(d=lf, n=NO_LINES,
                                                                                 t=th("rating_text", lab))
                                 if lab else None)
                    rt.font = Font(name="Calibri", size=12, color=SLATE) if big else small
                    rt.alignment = Alignment(horizontal="left", vertical="center", indent=0)
                if (a_, b_) in lay["only"]:
                    o = ws.cell(r, base + offs[lead][1], "%s only" % " and ".join(short(x)
                                                                               for x in lay["only"][(a_, b_)]))
                    o.font = Font(name="Calibri", size=12, italic=True, color=SLATE)
                    o.alignment = Alignment(horizontal="left", vertical="center")
                if compare:
                    x, y = (L(base + offs[s][0]) for s in compare)
                    d = ws.cell(r, base + offs["points"], '=IF(AND(ISNUMBER({x}{r}),ISNUMBER({y}{r})),{x}{r}-{y}{r},"")'
                                .format(x=x, y=y, r=r))
                    d.number_format = "+0.00;-0.00;0.00"
                    d.font = small
                    d.alignment = Alignment(horizontal="right", vertical="center")
                    d.border = Border(top=top, right=_side("thin", K.STONE))
                ws.row_dimensions[r].height = H_BODY
    for base, e in zip(panel_left, ends):
        if e >= FIRST:
            for off in range(width):
                cell = ws.cell(e, base + off)
                cell.border = Border(top=cell.border.top, right=cell.border.right, bottom=_side("thin", K.INK))

    # Rating cells: the rating's fill; Refused in crimson on canvas; the
    # rest (a dash, a source's absence) stays as written.
    rating_cols = [L(base + offs[s][1]) for base in panel_left for s in srcs]
    end = max(ends)
    rng = " ".join("%s%d:%s%d" % (c_, FIRST, c_, end) for c_ in rating_cols)
    anchor = "%s%d" % (rating_cols[0], FIRST)
    _rating_rules(ws, rng, anchor)

    # ---- Hidden helpers: one row per leading-source series, for the block ----
    hc = last_col + 3                               # past the right panel and a spare column
    hcols = {k: hc + i for i, k in enumerate(("product", "measure", "rating", "latest", "lines_from", "order",
                                                         "row"))}
    hr0 = FIRST
    hr1 = hr0 + max(len(lead_rows), 1) - 1
    rng_h = lambda k: "$%s$%d:$%s$%d" % (L(hcols[k]), hr0, L(hcols[k]), hr1)    # noqa: E731
    for i, (a_, b_) in enumerate(lead_rows):
        r = hr0 + i
        lab = by[(a_, b_, lead)]
        ws.cell(r, hcols["product"], a_)
        ws.cell(r, hcols["measure"], b_)
        ws.cell(r, hcols["rating"], '=IF(ISNUMBER({k}),{k},0)'.format(k=th("rating", lab)))
        ws.cell(r, hcols["latest"], '=IF(ISNUMBER({v}),{v},0)'.format(v=th("value", lab)))
        ws.cell(r, hcols["lines_from"], "=%s&\"\"" % th("lines_from", lab))
        # Its place in the list: rated 4 or 5, the higher rating first, then
        # the higher latest value, then the order on this tab.
        ws.cell(r, hcols["order"], ('=IF({R}{r}>=4,COUNTIFS({RR},">"&{R}{r})+COUNTIFS({RR},{R}{r},{LL},">"&{V}{r})'
                                    '+COUNTIFS({RR},{R}{r},{LL},{V}{r},{OO},"<"&{r})+1,"")')
                .format(R=L(hcols["rating"]), V=L(hcols["latest"]), r=r, RR=rng_h("rating"), LL=rng_h("latest"),
                        OO=rng_h("row")))
        ws.cell(r, hcols["row"], r)
    for k in range(len(hcols)):
        ws.column_dimensions[L(hc + k)].hidden = True
    hot = 'COUNTIF(%s,">=4")' % rng_h("rating")
    picked = 'SUM(%s)' % _range_of(refs, "picked")
    refused = 'SUM(%s)' % _range_of(refs, "refused")

    ws.cell(1, last_col, '="Lines picked: "&%s&" of %d · %s rated 4 or 5: "&%s&" · refused: "&%s'
            % (picked, len(rows), lead.replace('"', '""'), hot, refused))
    ws.cell(1, last_col).font = Font(name="Arial", size=12, bold=True, color=K.INK)
    ws.cell(1, last_col).alignment = Alignment(horizontal="right", vertical="bottom")

    # ---- Row 2, right: each source left out with a single series, as a strip ----
    strips = []
    one = [s for s in lay["dropped"] if sum(1 for lab in labels if parts[lab][2] == s) == 1]
    pct = unit == "%"
    for k, s in enumerate(one):
        r = 2 + k
        lab = next(lab for lab in labels if parts[lab][2] == s)
        a_, b_ = parts[lab][:2]
        # One cell, right-aligned under the counts, so no column's width
        # opens a gap between the value and the words after it.
        t = ws.cell(r, last_col, '="{w} · "&IF(ISNUMBER({v}),TEXT({v},"0.00"){u}&" · ","")&{q}&" · rating "&'
                                 'IF(OR({d}="",{t}=""),"{n}",{t})'.format(
                                     w=("%s · %s · %s" % (s, a_, b_)).replace('"', '""'), v=th("value", lab),
                                     u='&"%"' if pct else "", q=th("quarter", lab), d=th("lines_from", lab),
                                     t=th("rating_text", lab), n=NO_LINES))
        t.font = Font(name="Calibri", size=11, color=SLATE)
        t.alignment = Alignment(horizontal="right", vertical="center")
        if r > 2:
            ws.row_dimensions[r].height = 15
        strips.append({"source": s, "label": lab, "row": r, "text": t.coordinate})

    # ---- Needs your attention, under the left panel ----
    base = panel_left[0]
    # Measure, latest and rating under the panel's Measure and KeyBank columns;
    # the product and the source after them, where the panel's other sources are.
    c_meas, c_val, c_chip, c_prod = base, base + 1, base + 2, base + 3
    r = ends[0] + 2
    head = ws.cell(r, base, "Needs your attention · %s rated 4 or 5" % lead)
    head.font = Font(name="Arial", size=11, bold=True, color=K.INK)
    head.alignment = Alignment(vertical="bottom")
    for off in range(width):
        ws.cell(r, base + off).border = Border(bottom=_side("medium", K.INK))
    ws.row_dimensions[r].height = 16.5
    attention = {"head": r, "rows": [], "columns": {"measure": c_meas, "latest": c_val, "rating": c_chip,
                                                    "product": c_prod}}
    body = Font(name="Calibri", size=11, color=INK)
    for k in range(1, ATTENTION_ROWS + 1):
        r += 1
        at = "MATCH(%d,%s,0)" % (k, rng_h("order"))
        pick = lambda key: "INDEX(%s,%s)" % (rng_h(key), at)              # noqa: E731
        f_meas = 'IFERROR(%s,"")' % pick("measure")
        if k == 1:
            f_meas = 'IF(%s=0,"%s",%s)' % (hot, NONE_HOT, f_meas)
        ws.cell(r, c_meas, "=" + f_meas).font = body
        v = ws.cell(r, c_val, '=IFERROR(%s,"")' % pick("latest"))
        v.font = Font(name="Calibri", size=12, bold=True, color=INK)
        v.number_format = "0.00"
        v.alignment = Alignment(horizontal="right", vertical="center")
        chip = ws.cell(r, c_chip, '=IFERROR({k}&" "&CHOOSE({k},"Low","Moderate-Low","Moderate","Moderate-High",'
                                  '"High"),"")'.format(k=pick("rating")))
        chip.font = Font(name="Calibri", size=11, color=INK)
        chip.alignment = Alignment(vertical="center")
        # The source's name up to its bracket, so the words end inside the panel.
        ws.cell(r, c_prod, '=IFERROR({p}&" · lines from "&IFERROR(LEFT({s},FIND(" (",{s})-1),{s}),"")'.format(
            p=pick("product"), s=pick("lines_from"))).font = small
        for c_ in (c_meas, c_prod):
            ws.cell(r, c_).alignment = Alignment(vertical="center")
        attention["rows"].append(r)
        ws.row_dimensions[r].height = H_BODY
        for off in range(width):
            ws.cell(r, base + off).border = Border(bottom=_side("thin", "EFEBE4"))
    # The first row says "None rated 4 or 5." in slate when the list is empty.
    f0 = "%s%d" % (L(c_meas), attention["rows"][0])
    ws.conditional_formatting.add(f0, FormulaRule(formula=['%s="%s"' % (f0, NONE_HOT)], font=Font(color=SLATE)))
    _rating_rules(ws, "%s%d:%s%d" % (L(c_chip), attention["rows"][0], L(c_chip), r),
                  "%s%d" % (L(c_chip), attention["rows"][0]))

    # ---- Three footnotes, under both panels ----
    r = max(ends + [r]) + 2
    told = ["Rating: where the latest quarter falls on its %s and %s's lines, set on Decide from the history of "
            "one %s. %s : no %s picked yet. Refused: the settings do not fit that history; its Check on Decide "
            "says the fix." % (p0.lower(), p1.lower(), p2.lower(), NO_LINES, p2.lower(), )]
    second = "Blank: the %s has no series for that row." % p2.lower()
    if compare:
        second = ("Points: %s's latest value minus the latest value of %s, in the measure's units. "
                  % (compare[0], compare[1]) + second)
    told.append(second)
    if lay["dropped"]:
        d_ = lay["dropped"]
        told.append("%s: %s on Lines by source." % (
            " and ".join(d_), "its latest quarter and its own lines are" if len(d_) == 1
            else "their latest quarters and their own lines are"))
    for i, t in enumerate(told):
        cell = ws.cell(r + i, 2, t)
        cell.font = Font(name="Calibri", size=10, color=SLATE)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[r + i].height = 15
    notes_at = r

    ws.column_dimensions["A"].width = W_EDGE
    for base in panel_left:
        ws.column_dimensions[L(base)].width = W_MEASURE
        for i, s in enumerate(srcs):
            a, b = offs[s]
            ws.column_dimensions[L(base + a)].width = W_LEAD_LATEST if i == 0 else W_LATEST
            ws.column_dimensions[L(base + b)].width = W_LEAD_RATING if i == 0 else W_RATING
        if compare:
            ws.column_dimensions[L(base + offs["points"])].width = W_POINTS
    ws.column_dimensions[L(panel_left[1] - 1)].width = W_GUTTER
    ws.freeze_panes = "A%d" % FIRST
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    # The left panel's columns (the right panel's are these plus width + 1).
    columns = {s: (2 + offs[s][0], 2 + offs[s][1]) for s in srcs}
    if compare:
        columns["points"] = 2 + offs["points"]
    return {"sheet": TITLE, "first": FIRST, "last": end, "sources": srcs, "dropped": lay["dropped"],
            "panels": panel_left, "columns": columns, "placed": placed, "heads": lay["heads"], "rows": rows,
            "attention": attention, "notes": notes_at, "cells": lay["cells"], "helpers": hcols,
            "count": ws.cell(1, last_col).coordinate, "strips": strips, "only": lay["only"]}


def _rating_rules(ws, rng, anchor):
    for k in range(5):
        ws.conditional_formatting.add(rng, FormulaRule(formula=['LEFT(%s,1)="%d"' % (anchor, k + 1)],
                                                       fill=K.solid(K.RATING_FILL[k]), font=Font(color=K.INK)))
    ws.conditional_formatting.add(rng, FormulaRule(formula=['%s="Refused"' % anchor], fill=K.solid(K.CANVAS),
                                                   font=Font(bold=True, color=K.CRIMSON)))
