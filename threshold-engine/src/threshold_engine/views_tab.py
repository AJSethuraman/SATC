"""Saved views: Signals' and Peers' picks kept by number, with a note, and
called back on either tab with its Saved view box.

The analyst, 20 Oct 2026: "the factors on decide are easy to show as proof of
threshold work. What about other tabs? Options to essentially 'save a view
with notes'?" Excel's Custom Views are off in any workbook with a table, and no
formula can freeze a value, so a view is its picks: the data in a built
workbook does not move, so the same picks show the same view.

Built last (book.tidy), after every tab that reads a pick: each formula that
names a tab's pick is pointed at that pick's in-effect cell here, which reads
the saved view's cell when the tab's Saved view box holds a number and that
cell is not blank, and the tab's own yellow cell otherwise. A pick no formula
names is refused, never skipped.
"""

import re
from typing import Dict, List

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

from . import keybank_style as K

TITLE = "Saved views"
VIEWS = 12                                            # rows per tab
#: The tabs a view is saved for, each with its Saved view box (label above, input) and its note's cell.
TABS = {"Signals": {"box": "U5", "note": "U6", "align": "right"},
        "Peers": {"box": "H5", "note": "H6", "align": "left"}}
#: Peers' picks are its row 5; the rest of its yellow cells are the peer rules, not a view.
PEERS_ROW = 5
FIRST = 5                                             # the first section's title row
COL_VIEW, COL_NOTE, COL_THROUGH, COL_PICK = 2, 3, 4, 5
HELP = 30                                             # hidden: each pick's in-effect value, from column AD

INK, SLATE = K.INK_TEXT, K.SLATE
F_BODY = Font(name="Calibri", size=11, color=INK)
F_BOLD = Font(name="Calibri", size=11, bold=True, color=INK)
F_SLATE = Font(name="Calibri", size=11, color=SLATE)
F_NOTE = Font(name="Calibri", size=10, color=SLATE)
F_LAB = Font(name="Arial", size=9, bold=True, color=SLATE)
EDGE = Side(style="thin", color=K.INPUT_EDGE)
BOX = Border(EDGE, EDGE, EDGE, EDGE)
INPUT = PatternFill("solid", fgColor=K.INPUT, bgColor=K.INPUT)


def _q(title: str) -> str:
    return "'%s'!" % title.replace("'", "''")


def _abs(coord: str) -> str:
    m = re.fullmatch(r"([A-Z]+)(\d+)", coord)
    return "$%s$%s" % m.groups()


def picks_of(wb, tab: str) -> List[dict]:
    """The tab's picks, in sheet order: each yellow cell a list sets, and the
    words over it (Signals' Yes/No per product reads 'On the charts: <product>')."""
    ws = wb[tab]
    out = []
    for dv in ws.data_validations.dataValidation:
        for rng in dv.sqref.ranges:
            for row in ws[rng.coord] if ":" in rng.coord else [[ws[rng.coord]]]:
                for c in row:
                    if tab == "Peers" and c.row != PEERS_ROW:
                        continue
                    above = ws.cell(c.row - 1, c.column).value
                    label = above if c.row == 5 else "On the charts: %s" % above
                    out.append({"cell": c.coordinate, "label": str(label), "list": dv.formula1, "start": c.value})
    out.sort(key=lambda p: (int(re.sub(r"[A-Z]+", "", p["cell"])), len(re.sub(r"\d+", "", p["cell"])),
                            re.sub(r"\d+", "", p["cell"])))
    return out


def layout(n_picks: Dict[str, int]) -> dict:
    """Rows of each section: title, header, Picks now, then views 1..VIEWS."""
    out, r = {}, FIRST
    for tab in TABS:
        if tab not in n_picks:
            continue
        out[tab] = dict(title=r, head=r + 1, now=r + 2, first=r + 3, last=r + 2 + VIEWS)
        r += 3 + VIEWS + 2
    return out


def _repoint(wb, tab: str, coord: str, to: str) -> int:
    """Point every formula that names the tab's pick at ``to``; the count."""
    col, row = re.fullmatch(r"([A-Z]+)(\d+)", coord).groups()
    q = re.escape(tab)
    named = re.compile(r"(?:'%s'|%s)!\$?%s\$?%s(?![0-9:A-Za-z(])" % (q.replace("'", "''"), q, col, row))
    bare = re.compile(r"(?<![!:A-Za-z0-9$_.'])\$?%s\$?%s(?![0-9:A-Za-z(!])" % (col, row))
    n = 0

    def sub(text, own):
        nonlocal n
        text, k = named.subn(to, text)
        n += k
        if own:
            text, k = bare.subn(to, text)
            n += k
        return text

    for ws in wb.worksheets:
        if ws.title == TITLE:
            continue
        own = ws.title == tab
        for line in ws.iter_rows():
            for c in line:
                if isinstance(c.value, str) and c.value.startswith("="):
                    c.value = sub(c.value, own)
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                rule.formula = [sub(f, own) for f in rule.formula]
    for name, dn in list(wb.defined_names.items()):
        if dn.attr_text:
            dn.attr_text = sub(dn.attr_text, False)
    return n


def build(wb, quarters: str) -> dict:
    """The Saved views tab, and Signals and Peers made to follow it. ``quarters``
    is the list Data through picks from, e.g. "=Data!$A$2:$A$82"."""
    picks = {t: picks_of(wb, t) for t in TABS if t in wb.sheetnames}
    picks = {t: p for t, p in picks.items() if p}
    if not picks:
        return {}
    lay = layout({t: len(p) for t, p in picks.items()})
    ws = wb.create_sheet(TITLE)
    K.hide_gridlines(ws)
    width = COL_PICK + max(len(p) for p in picks.values()) - 1
    K.light_title(ws, 1, width, TITLE,
                  "Type a view's number under Saved view on its tab: its picks replace the tab's. A blank cell "
                  "keeps the tab's own pick.", first_col=2)
    ws.column_dimensions["A"].width = 2.3
    for c, w in ((COL_VIEW, 8), (COL_NOTE, 46), (COL_THROUGH, 13)):
        ws.column_dimensions[L(c)].width = w
    # each pick's column as wide as the longest pick it starts with, so Picks now reads whole
    for j in range(width - COL_PICK + 1):
        longest = max(len(str(ps[j]["start"] or "")) for ps in picks.values() if j < len(ps))
        ws.column_dimensions[L(COL_PICK + j)].width = max(16, min(34, longest + 3))
    numbers = '"%s"' % ",".join(str(i) for i in range(1, VIEWS + 1))
    dv_q = DataValidation(type="list", formula1=quarters, allow_blank=True, showErrorMessage=True,
                          errorTitle="Data through", error="Pick a quarter from the list.")
    ws.add_data_validation(dv_q)
    sh = _q(TITLE)
    out = {"layout": lay, "picks": {}, "boxes": {}}
    for tab, ps in picks.items():
        ly, src, cfg = lay[tab], wb[tab], TABS[tab]
        K.quiet_section(ws, ly["title"], tab, width - 1, first_col=2)
        K.canvas_header(ws, ly["head"], ["View", "Note", "Data through"] + [p["label"] for p in ps], first_col=2,
                        height=27)
        rows = range(ly["first"], ly["last"] + 1)
        # Picks now: the tab's own yellow cells, to copy onto a view's row (paste values).
        ws.cell(ly["now"], COL_VIEW, "Picks now").font = F_LAB
        ws.cell(ly["now"], COL_VIEW).alignment = Alignment(vertical="center")
        for j, p in enumerate(ps):
            c = ws.cell(ly["now"], COL_PICK + j, "=%s%s" % (_q(tab), _abs(p["cell"])))
            c.font, c.alignment = F_SLATE, Alignment(vertical="center")
        for col in range(COL_VIEW, width + 1):
            ws.cell(ly["now"], col).border = Border(bottom=Side(style="thin", color=K.STONE))
        for i, r in enumerate(rows):
            ws.cell(r, COL_VIEW, i + 1).font = F_BOLD
            ws.cell(r, COL_VIEW).alignment = Alignment(horizontal="center", vertical="center")
            for col in range(COL_NOTE, COL_PICK + len(ps)):
                c = ws.cell(r, col)
                c.fill, c.border = INPUT, BOX
                c.font = F_BODY if col == COL_NOTE else F_BOLD
                c.alignment = Alignment(vertical="center", wrap_text=col == COL_NOTE)
            ws.row_dimensions[r].height = 18
        dv_q.add("%s%d:%s%d" % (L(COL_THROUGH), ly["first"], L(COL_THROUGH), ly["last"]))
        for j, p in enumerate(ps):
            dv = DataValidation(type="list", formula1=p["list"], allow_blank=True, showErrorMessage=True,
                                errorTitle=p["label"], error="Pick one from the list, or clear the cell.")
            ws.add_data_validation(dv)
            dv.add("%s%d:%s%d" % (L(COL_PICK + j), ly["first"], L(COL_PICK + j), ly["last"]))
        # The tab's Saved view box, under its label, and the note it calls up under it.
        box, note = cfg["box"], cfg["note"]
        bc = src[box]
        src.cell(bc.row - 1, bc.column, "Saved view").font = F_LAB
        src.cell(bc.row - 1, bc.column).alignment = Alignment(horizontal=cfg["align"], vertical="bottom")
        assert bc.value is None and src[note].value is None, "the Saved view box on %s sits on a cell in use" % tab
        bc.fill, bc.border, bc.font = INPUT, BOX, F_BOLD
        bc.alignment = Alignment(horizontal="center", vertical="center")
        src.add_data_validation(DataValidation(type="list", formula1=numbers, allow_blank=True,
                                               showErrorMessage=True, errorTitle="Saved view",
                                               error="Pick a view's number from the list, or clear the cell."))
        src.data_validations.dataValidation[-1].add(box)
        v = _q(tab) + _abs(box)
        rng = lambda col: "%s$%s$%d:$%s$%d" % (sh, L(col), ly["first"], L(col), ly["last"])      # noqa: E731
        filled = "+".join('(INDEX(%s,%s)<>"")' % (rng(COL_PICK + j), v) for j in range(len(ps)))
        nc = src[note]
        nc.value = ('=IF(NOT(ISNUMBER({v})),"",IF(({f})=0,"Saved view "&{v}&" has no picks",'
                    '"Saved view "&{v}&IF(INDEX({n},{v})="","",": "&INDEX({n},{v}))))'
                    .format(v=v, f=filled, n=rng(COL_NOTE)))
        nc.font, nc.alignment = F_NOTE, Alignment(horizontal=cfg["align"], vertical="center")
        out["boxes"][tab] = box
        # Each pick in effect, hidden on this tab: the view's cell when the box holds a
        # number and that cell is not blank, the tab's own otherwise; then every formula
        # that read the tab's cell reads this one, and the tab's cell, overridden, is struck.
        out["picks"][tab] = []
        for j, p in enumerate(ps):
            hc = ws.cell(ly["title"], HELP + j)
            to = "%s$%s$%d" % (sh, L(HELP + j), ly["title"])
            n = _repoint(wb, tab, p["cell"], to)
            if not n:
                raise ValueError("Saved views: no formula reads %s!%s (%s); a view of it would do nothing"
                                 % (tab, p["cell"], p["label"]))
            over = 'AND(ISNUMBER({v}),INDEX({c},MAX(1,{v}))<>"")'.format(v=v, c=rng(COL_PICK + j))
            hc.value = '=IF({o},INDEX({c},{v}),{t})'.format(o=over, c=rng(COL_PICK + j), v=v,
                                                            t=_q(tab) + _abs(p["cell"]))
            src.conditional_formatting.add(p["cell"], FormulaRule(
                formula=[over], font=Font(color=K.STONE, strike=True)))
            out["picks"][tab].append(dict(p, in_effect=to, column=L(COL_PICK + j), reads=n))
    for col in range(HELP, HELP + max(len(p) for p in picks.values())):
        ws.column_dimensions[L(col)].hidden = True
    K.freeze_below(ws, 2)
    return out
