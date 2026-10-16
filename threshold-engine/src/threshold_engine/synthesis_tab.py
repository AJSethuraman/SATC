"""The Synthesis tab: where each colour starts, from internal, industry and peers, live.

The firm, 13 Oct 2026: *"i want to get an industry perspective, competitive
perspective, and internal. all of this to somehow synthesize reasonable
thresholds"*; and, of where it goes, the synthesis on Assess: *where each
colour starts*. ``synthesis.py`` is the arithmetic and the reference; this
writes the same arithmetic as Excel formulas, so the three candidates, the
chosen lines, the latest rating and the competitive check move when Your
settings on Assess move, or when the firm changes a choice here.

For each product given synthesis inputs, a block on the Synthesis tab:

    the four lines      internal, industry and peers side by side; which source
                        each line is taken from (yours, a yellow cell); the
                        chosen line
    internal            Thresholds' lines for the product: Your settings and the
                        periods left out on Assess. A calm history places lines
                        2-4 and cannot place line 5, as ``synthesis._lines`` says
    industry            the same scale on the industry's series, times the
                        median of internal / industry over the quarters both
                        hold, or as it is (yours: median ratio or none)
    peers               the percentile across the peers' medians over the
                        internal quarters (yours, one per line), times the
                        median of internal / KeyBank filed
    latest, check       the latest internal quarter rated on the chosen lines;
                        KeyBank's filed median and how many peers it is worse
                        than (facts, no alarm level)
    Check               OK, or why it refuses, in synthesis.py's words

and on Assess, for the product showing in B5, the same four lines and the
Check, read from here. The series the formulas read are on SynthesisData, one
continuous quarterly calendar.

Every choice starts from what the command stated; nothing has a default. The
industry, filed and peer series are smoothed as the measure says, like the
internal ones. Periods left out on Assess change the internal lines only: the
translations and the peer window use every quarter the internal series holds,
as synthesis.py does.
"""
from __future__ import annotations

import glob
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

from . import assess_tab as AT
from . import keybank_style as K
from .series import check, read_csv, smooth
from .synthesis import LINES, SOURCES, TRANSLATIONS, _check

#: Thresholds' wording for the refusal that leaves lines 2-4 standing, and
#: synthesis.py's for the same thing (scale.py's text after "line 5: ").
TH_HIGH_BEYOND = "Refused: High would begin beyond the worst quarter kept"
HIGH_BEYOND = "line 5: High would begin beyond the worst quarter kept: nothing kept could rate High"
TH_TOO_LITTLE = "Refused: too little history kept for a typical yearly move"
TOO_LITTLE = "too little history kept for a typical yearly move (two four-quarter changes needed)"
NO_PEERS = "not built: no peer series were given"
NO_FILED = ("KeyBank's filed series shares no quarter with the internal one, so peers cannot be carried to the "
            "internal level")
FEW_PEERS = "fewer than two peers hold any quarter of the internal window"
NO_SHARED = "the internal and industry series share no quarter, so the industry cannot be translated"
GO_TOGETHER = ("peer percentiles and peer series go together: give both, or neither (then no line can be led by "
               "peers)")

# Where things sit in a product's block (rows from its top) and in its row
# of hidden helpers (columns). One place for the builder and the tests.
COL = dict(label=1, internal=2, industry=3, pct=4, peers=5, lead=6, chosen=7, note=8)
ROW = dict(title=0, head=1, lines=2, translation=6, latest=7, competitive=8, check=9,
           int_rests=10, ind_rests=11, peer_rests=12, peer_head=13, peer_top=14)
_H = ("ind_med", "ind_mchg", "ind_move", "ind_spread", "ind_worst", "ind_ratio", "ind_ratio_n",
      "ib2", "ib3", "ib4", "ib5", "ind_scale_why", "ind_why", "f_ratio", "f_ratio_n", "held", "kb_med",
      "below", "worse_than", "peer_why", "int_why", "hw", "ls", "hz", "online", "set_ok", "lead_ok",
      "pct_ok", "pct_rise", "trans_ok", "online_ok", "ind_n", "rating", "rating_name", "check")
HELP = {name: 12 + i for i, name in enumerate(_H)}          # hidden, from column L
SUMMARY_COL = 60                                            # hidden: one row per product, read by Assess


# ------------------------------------------------------------------ inputs

def _split(item: str, flag: str) -> Tuple[str, str]:
    label, sep, rest = item.partition("=")
    if not sep or not label.strip() or not rest.strip():
        raise ValueError("%s takes LABEL=VALUE, not %r" % (flag, item))
    return label, rest.strip()


def _peer_files(where: str) -> List[Tuple[str, Path]]:
    """A folder means every .csv in it; anything else is a file pattern
    (``banks/credit-card_*_net.csv``). Each peer is named by its file."""
    p = Path(where)
    files = sorted(p.glob("*.csv")) if p.is_dir() else sorted(Path(f) for f in glob.glob(where))
    if not files:
        raise ValueError("peers: %s %s no .csv file" % (where, "holds" if p.is_dir() else "matches"))
    return [(f.stem, f) for f in files]


def from_flags(specs: Sequence[str], translations: Sequence[str], percentiles: Sequence[str],
               leads: Sequence[str], labels: Sequence[str]) -> Optional[Dict[str, dict]]:
    """The command line's synthesis flags as ``build``'s ``synthesis``.

    ``--synthesis "LABEL=industry:FILE;filed:FILE;peers:FOLDER-OR-PATTERN"``
    (industry required, filed and peers optional), and per product
    ``--translation "LABEL=median ratio"``, ``--leads "LABEL=SOURCE,SOURCE,
    SOURCE,SOURCE"`` and, when it has peers, ``--peer-percentiles
    "LABEL=P2,P3,P4,P5"``. None when no synthesis flag is given; a refusal
    (ValueError) naming what is missing otherwise. Every choice is stated per
    product: there is no set that applies to all."""
    if not (specs or translations or percentiles or leads):
        return None
    if not specs:
        given = [f for f, v in (("--translation", translations), ("--peer-percentiles", percentiles),
                                ("--leads", leads)) if v]
        raise ValueError("%s %s for the synthesis: give --synthesis too, or leave %s out"
                         % (" and ".join(given), "is" if len(given) == 1 else "are",
                            "it" if len(given) == 1 else "them"))
    out: Dict[str, dict] = {}
    for item in specs:
        label, rest = _split(item, "--synthesis")
        if label not in labels:
            raise ValueError("--synthesis names %r, which is not a --series label (%s)" % (label, ", ".join(labels)))
        if label in out:
            raise ValueError("--synthesis names %s twice" % label)
        parts: Dict[str, str] = {}
        for piece in rest.split(";"):
            key, sep, val = piece.strip().partition(":")
            if not sep or key not in ("industry", "filed", "peers") or not val.strip():
                raise ValueError("--synthesis for %s takes industry:FILE, and optionally filed:FILE and "
                                 "peers:FOLDER, separated by ';', not %r" % (label, piece.strip()))
            if key in parts:
                raise ValueError("--synthesis for %s gives %s twice" % (label, key))
            parts[key] = val.strip()
        if "industry" not in parts:
            raise ValueError("--synthesis for %s needs industry:FILE" % label)
        out[label] = dict(industry=Path(parts["industry"]), filed=Path(parts["filed"]) if "filed" in parts else None,
                          peers=_peer_files(parts["peers"]) if "peers" in parts else [],
                          translation=None, peer_percentiles=None, leads=None)
    for flag, items, key in (("--translation", translations, "translation"),
                             ("--peer-percentiles", percentiles, "peer_percentiles"), ("--leads", leads, "leads")):
        seen = set()
        for item in items:
            label, val = _split(item, flag)
            if label not in out:
                raise ValueError("%s names %r, which has no --synthesis" % (flag, label))
            if label in seen:
                raise ValueError("%s names %s twice" % (flag, label))
            seen.add(label)
            if key == "translation":
                out[label][key] = val
            elif key == "leads":
                out[label][key] = [x.strip() for x in val.split(",")]
            else:
                try:
                    out[label][key] = [float(x) for x in val.split(",")]
                except ValueError:
                    raise ValueError("%s for %s takes four numbers separated by commas, not %r"
                                     % (flag, label, val)) from None
    for label, s in out.items():
        if s["translation"] is None:
            raise ValueError('the synthesis for %s needs --translation "%s=median ratio" or "%s=none"'
                             % (label, label, label))
        if s["leads"] is None:
            raise ValueError('the synthesis for %s needs --leads "%s=SOURCE,SOURCE,SOURCE,SOURCE": internal, '
                             'industry or peers for lines 2, 3, 4 and 5' % (label, label))
        if s["peers"] and s["peer_percentiles"] is None:
            raise ValueError('the synthesis for %s has peers, so it needs --peer-percentiles "%s=P2,P3,P4,P5"'
                             % (label, label))
        if not s["peers"] and s["peer_percentiles"] is not None:
            raise ValueError("the synthesis for %s: %s" % (label, GO_TOGETHER))
    return out


def validate(synthesis: Dict[str, dict], labels: Sequence[str], direction: str, on_the_line: str) -> None:
    """The starting choices, refused as synthesis.py refuses them."""
    for label, s in synthesis.items():
        if label not in labels:
            raise ValueError("the synthesis names %r, which is not one of the series" % label)
        if not s.get("industry"):
            raise ValueError("the synthesis for %s needs an industry series" % label)
        peers = s.get("peers") or []
        pct = s.get("peer_percentiles")
        if bool(peers) != (pct is not None):
            raise ValueError("the synthesis for %s: %s" % (label, GO_TOGETHER))
        if peers and not s.get("filed"):            # as synthesis.synthesize refuses it
            raise ValueError("the synthesis for %s: peers are carried to the internal level by KeyBank's "
                             "filed series: give it (filed:FILE), or give no peers" % label)
        try:
            _check(direction, on_the_line, s.get("translation"), pct, s.get("leads"))
        except ValueError as exc:
            raise ValueError("the synthesis for %s: %s" % (label, exc)) from None
        names = [nm for nm, _ in peers]
        if len(set(names)) != len(names):
            raise ValueError("the synthesis for %s names two peers alike" % label)


def load(synthesis: Dict[str, dict], smoothing: int) -> Dict[str, dict]:
    """Each file read, checked and smoothed as the measure says, like the
    internal series: ``{label: {"industry": (path, raw, points), "filed":
    ..., "peers": [(name, path, raw, points)]}}``."""
    def one(path, what):
        raw = read_csv(Path(path))
        try:
            return Path(path), raw, smooth(check(raw, "quarterly"), smoothing)
        except ValueError as exc:
            raise ValueError("%s (%s): %s" % (what, Path(path).name, exc)) from None
    out = {}
    for label, s in synthesis.items():
        out[label] = dict(
            industry=one(s["industry"], "the industry series for %s" % label),
            filed=one(s["filed"], "KeyBank's filed series for %s" % label) if s.get("filed") else None,
            peers=[(nm,) + one(p, "the peer %s for %s" % (nm, label)) for nm, p in s.get("peers") or []])
    return out


def files(synthesis_loaded: Dict[str, dict]) -> List[Tuple[str, Path, list]]:
    """(title, path, raw rows) for each synthesis file, in the order the
    Sources and Run tabs list them."""
    out = []
    for label, d in synthesis_loaded.items():
        out.append(("%s · industry" % label, d["industry"][0], d["industry"][1]))
        if d["filed"]:
            out.append(("%s · KeyBank filed" % label, d["filed"][0], d["filed"][1]))
        for nm, path, raw, _ in d["peers"]:
            out.append(("%s · peer %s" % (label, nm), path, raw))
    return out


def block_height(n_peers: int) -> int:
    return ROW["peer_rests"] + 1 + (2 + n_peers + 1 if n_peers else 0) + 1


# ------------------------------------------------------------------ the tab

def _q_index(q: str) -> int:
    return int(q[:4]) * 4 + int(q[5]) - 1


def _q_name(i: int) -> str:
    return "%dQ%d" % (i // 4, i % 4 + 1)


def build(wb, sy, sd, asx, *, loaded, syn, syn_loaded, grid, n, higher, floor_at_zero, unit, nf, srow,
          quarter, assess_lay) -> Dict[str, dict]:
    """Write the Synthesis and SynthesisData tabs and the panel on Assess.
    Returns, per product, where its block sits (for the tests)."""
    s_ = 1 if higher else -1
    sg = str(s_)
    fz = "TRUE" if floor_at_zero and higher else "FALSE"
    labels = [s["label"] for s in loaded]
    order = [lb for lb in labels if lb in syn]                  # the series' order
    tr = {lb: 5 + labels.index(lb) for lb in labels}            # Thresholds row
    data_row = {q: 2 + i for i, q in enumerate(grid)}            # Data row

    # ---- SynthesisData: one continuous calendar ----
    qs = set()
    for lb in order:
        qs |= set(next(s for s in loaded if s["label"] == lb)["by_q"])
        d = syn_loaded[lb]
        for pts in [d["industry"][2]] + ([d["filed"][2]] if d["filed"] else []) + [p[3] for p in d["peers"]]:
            qs |= {quarter(p.date) for p in pts}
    lo_i, hi_i = min(map(_q_index, qs)), max(map(_q_index, qs))
    cal = [_q_name(i) for i in range(lo_i, hi_i + 1)]
    r0, r1 = 2, len(cal) + 1
    sd.cell(1, 1, "Quarter")
    for r, q in enumerate(cal, start=r0):
        sd.cell(r, 1, q)
    col = 2
    cols: Dict[str, dict] = {}
    rng = lambda c_: "SynthesisData!$%s$%d:$%s$%d" % (L(c_), r0, L(c_), r1)      # noqa: E731

    def by_q(pts):
        return {quarter(p.date): p.value for p in pts}

    top = 4
    tops = {}
    for lb in order:
        tops[lb] = top
        top += block_height(len(syn_loaded[lb]["peers"]))
    hr = {lb: tops[lb] + ROW["lines"] for lb in order}          # the helpers' row
    H = lambda lb, name: "Synthesis!$%s$%d" % (L(HELP[name]), hr[lb])          # noqa: E731

    for lb in order:
        k = labels.index(lb)
        d = syn_loaded[lb]
        c = dict(int=col, ind=col + 1, chg=col + 2, dchg=col + 3, dlev=col + 4, ratio=col + 5)
        col += 6
        if d["filed"]:
            c.update(filed=col, fratio=col + 1, fwin=col + 2)
            col += 3
        c["peers"] = []
        for _ in d["peers"]:
            c["peers"].append((col, col + 1))
            col += 2
        cols[lb] = c
        heads = {c["int"]: "%s, internal" % lb, c["ind"]: "%s, industry" % lb,
                 c["chg"]: "%s, industry, change over 4 quarters" % lb,
                 c["dchg"]: "%s, industry, distance from the median change" % lb,
                 c["dlev"]: "%s, industry, distance from the median" % lb,
                 c["ratio"]: "%s, internal / industry" % lb}
        if d["filed"]:
            heads.update({c["filed"]: "%s, KeyBank filed" % lb, c["fratio"]: "%s, internal / KeyBank filed" % lb,
                          c["fwin"]: "%s, KeyBank filed, internal quarters" % lb})
        for (nm, *_), (vc, wc) in zip(d["peers"], c["peers"]):
            heads.update({vc: "%s, peer %s" % (lb, nm), wc: "%s, peer %s, internal quarters" % (lb, nm)})
        for cc, h in heads.items():
            sd.cell(1, cc, h)
        ind = by_q(d["industry"][2])
        filed = by_q(d["filed"][2]) if d["filed"] else {}
        peers = [by_q(p[3]) for p in d["peers"]]
        dcol = L(2 + k)
        own = next(s for s in loaded if s["label"] == lb)["by_q"]
        for r, q in enumerate(cal, start=r0):
            I, X = "%s%d" % (L(c["int"]), r), "%s%d" % (L(c["ind"]), r)
            if q in own:
                sd.cell(r, c["int"], "=Data!%s%d" % (dcol, data_row[q]))
            if q in ind:
                sd.cell(r, c["ind"], ind[q])
            if r >= r0 + 4:
                sd.cell(r, c["chg"], '=IF(AND(ISNUMBER({x}),ISNUMBER({c}{p})),{x}-{c}{p},"")'
                        .format(x=X, c=L(c["ind"]), p=r - 4))
            sd.cell(r, c["dchg"], '=IF(ISNUMBER({g}{r}),ABS({g}{r}-{m}),"")'
                    .format(g=L(c["chg"]), r=r, m=H(lb, "ind_mchg")))
            sd.cell(r, c["dlev"], '=IF(ISNUMBER({x}),ABS({x}-{m}),"")'.format(x=X, m=H(lb, "ind_med")))
            sd.cell(r, c["ratio"], '=IF(AND(ISNUMBER({i}),ISNUMBER({x})),IF({x}>0,{i}/{x},""),"")'.format(i=I, x=X))
            if d["filed"]:
                F = "%s%d" % (L(c["filed"]), r)
                if q in filed:
                    sd.cell(r, c["filed"], filed[q])
                sd.cell(r, c["fratio"], '=IF(AND(ISNUMBER({i}),ISNUMBER({f})),IF({f}>0,{i}/{f},""),"")'
                        .format(i=I, f=F))
                sd.cell(r, c["fwin"], '=IF(AND(ISNUMBER({i}),ISNUMBER({f})),{f},"")'.format(i=I, f=F))
            for pv, (vc, wc) in zip(peers, c["peers"]):
                if q in pv:
                    sd.cell(r, vc, pv[q])
                sd.cell(r, wc, '=IF(AND(ISNUMBER({i}),ISNUMBER({p}{r})),{p}{r},"")'.format(i=I, p=L(vc), r=r))
    for cc in range(2, col):
        h = sd.cell(1, cc)
        h.fill = K.solid(K.CANVAS)
        h.font = Font(name="Arial", bold=True, size=9, color=K.SLATE)
        h.alignment = Alignment(wrap_text=True, vertical="center")
        sd.column_dimensions[L(cc)].width = 13
        for r in range(r0, r1 + 1):
            sd.cell(r, cc).number_format = nf
    sd.cell(1, 1).font = Font(name="Arial", bold=True, size=9, color=K.SLATE)
    sd.cell(1, 1).fill = K.solid(K.CANVAS)
    sd.row_dimensions[1].height = 46
    sd.freeze_panes = "B2"

    # ---- Synthesis: one block per product ----
    sy.sheet_properties.tabColor = K.KEY_RED
    K.light_title(sy, 1, 8, "Synthesis · where each colour starts",
                  "Live: internal, industry and peers side by side for each rating line; you choose where each line "
                  "is taken from. Follows Your settings on Assess. Type only in yellow cells.")
    lead_dv = DataValidation(type="list", formula1='"internal,industry,peers"', showInputMessage=True,
                             promptTitle="Taken from", prompt="Where this line is taken from: internal, industry "
                             "or peers. The chosen lines must still run in order.")
    trans_dv = DataValidation(type="list", formula1='"median ratio,none"', showInputMessage=True,
                              promptTitle="Industry carried by",
                              prompt="median ratio: the industry's lines times the median of internal / industry "
                                     "over the quarters both hold. none: the industry's lines as they are.")
    pct_dv = DataValidation(showInputMessage=True, allow_blank=True, promptTitle="Peers at percentile",
                            prompt="The percentile across the peers' medians for this line, 0 to 100, rising from "
                                   "line 2 to line 5. For a measure where lower is worse, 100 minus it is taken.")
    for dv in (lead_dv, trans_dv, pct_dv):
        sy.add_data_validation(dv)
    DATA_F = Font(name="Calibri", size=11, color=K.INK_TEXT)
    LABEL_F = Font(name="Arial", bold=True, size=9, color=K.SLATE)
    out: Dict[str, dict] = {}
    for i_s, lb in enumerate(order):
        b = tops[lb]
        t = tr[lb]
        sr = srow[lb]
        d = syn_loaded[lb]
        s = syn[lb]
        c = cols[lb]
        m = len(d["peers"])
        h = lambda name: "$%s$%d" % (L(HELP[name]), hr[lb])                    # noqa: E731
        lr = [b + ROW["lines"] + k for k in range(4)]
        cell = lambda name, r: "$%s$%d" % (L(COL[name]), r)                     # noqa: E731
        LEAD = [cell("lead", r) for r in lr]
        PCT = [cell("pct", r) for r in lr]
        CH = [cell("chosen", r) for r in lr]
        TRANS = cell("internal", b + ROW["translation"])
        meds = ("$B$%d:$B$%d" % (b + ROW["peer_top"], b + ROW["peer_top"] + m - 1)) if m else None

        K.quiet_section(sy, b, lb, 8)
        K.canvas_header(sy, b + 1, ["Rating begins", "Internal", "Industry", "Peers at percentile (yours)", "Peers",
                                    "Taken from (yours)", "Chosen line", ""], right_from=1)
        # hidden helpers: names over them, values on the first line's row
        for name, cc in HELP.items():
            sy.cell(b + 1, cc, name)
        own_n = len(next(x for x in loaded if x["label"] == lb)["points"])
        ind_pts = d["industry"][2]
        hw, ls, hz, onl = (h(x) for x in ("hw", "ls", "hz", "online"))
        f = {
            "hw": "=Assess!$O$%d" % sr, "ls": "=Assess!$P$%d" % sr, "hz": "=Assess!$Q$%d" % sr,
            "online": "=Assess!$R$%d" % sr,
            "set_ok": "=AND(ISNUMBER({hw}),ISNUMBER({ls}),ISNUMBER({hz}))".format(hw=hw, ls=ls, hz=hz),
            "online_ok": '=OR(EXACT({o},"worse"),EXACT({o},"better"))'.format(o=onl),
            "trans_ok": '=OR(EXACT({t},"median ratio"),EXACT({t},"none"))'.format(t=TRANS),
            "lead_ok": "=AND(%s)" % ",".join('OR(EXACT({x},"internal"),EXACT({x},"industry"),EXACT({x},"peers"))'
                                             .format(x=x) for x in LEAD),
            "pct_ok": "=TRUE" if not m else
                      "=IF(AND(%s),AND(%s),FALSE)" % (",".join("ISNUMBER(%s)" % p for p in PCT),
                                                      ",".join("%s>=0,%s<=100" % (p, p) for p in PCT)),
            "pct_rise": "=IF({ok},AND({a}),FALSE)".format(ok=h("pct_ok"), a=",".join(
                "%s>%s" % (PCT[k + 1], PCT[k]) for k in range(3))),
            "ind_n": "=COUNT(%s)" % rng(c["ind"]),
            "ind_med": "=IF(COUNT({x})=0,\"\",MEDIAN({x}))".format(x=rng(c["ind"])),
            "ind_mchg": "=IF(COUNT({x})=0,\"\",MEDIAN({x}))".format(x=rng(c["chg"])),
            "ind_move": "=IF(COUNT({x})<2,\"\",1.4826*MEDIAN({d}))".format(x=rng(c["chg"]), d=rng(c["dchg"])),
            "ind_spread": "=IF(COUNT({x})=0,\"\",1.4826*MEDIAN({d}))".format(x=rng(c["ind"]), d=rng(c["dlev"])),
            "ind_worst": "=%s(%s)" % ("MAX" if higher else "MIN", rng(c["ind"])),
            "ind_ratio": '=IF(EXACT({t},"none"),1,IF(COUNT({r})=0,"",MEDIAN({r})))'.format(t=TRANS, r=rng(c["ratio"])),
            "ind_ratio_n": "=COUNT(%s)" % rng(c["ratio"]),
            "f_ratio": ('=IF(COUNT({r})=0,"",MEDIAN({r}))'.format(r=rng(c["fratio"])) if d["filed"] else '=""'),
            "f_ratio_n": ("=COUNT(%s)" % rng(c["fratio"])) if d["filed"] else "=0",
            "held": ("=COUNT(%s)" % meds) if m else "=0",
            "kb_med": ('=IF(COUNT({r})=0,"",MEDIAN({r}))'.format(r=rng(c["fwin"]))) if d["filed"] else '=""',
        }
        ok = "AND({s},ISNUMBER({mv}))".format(s=h("set_ok"), mv=h("ind_move"))
        med, mv, sp = h("ind_med"), h("ind_move"), h("ind_spread")
        f["ib2"] = '=IF({ok},{m}-{s}*{ls}*{mv},"")'.format(ok=ok, m=med, s=sg, ls=ls, mv=mv)
        f["ib3"] = '=IF({ok},{m}-{s}*{hw}*{mv},"")'.format(ok=ok, m=med, s=sg, hw=hw, mv=mv)
        f["ib4"] = '=IF({ok},{m}+{s}*{hw}*{mv},"")'.format(ok=ok, m=med, s=sg, hw=hw, mv=mv)
        f["ib5"] = '=IF(AND({so},ISNUMBER({sp})),{m}+{s}*{hz}*{sp},"")'.format(so=h("set_ok"), sp=sp, m=med, s=sg,
                                                                                 hz=hz)
        # scale.py's refusals, in its order and its words.
        f["ind_scale_why"] = (
            '=IF(NOT({so}),"a setting in Your settings on Assess is not a number",'
            'IF(NOT(ISNUMBER({mv})),"{tl}",'
            'IF({hw}<0,"Moderate\'s half-width must be 0 or more",'
            'IF({hz}<=0,"High\'s z must be above 0",'
            'IF(AND({fz},{m}<=0),"the median is at or below zero",'
            'IF({ls}<={hw},"Moderate-Low must begin further from the median than Moderate",'
            'IF({s}*({b5}-{b4})<=0,"High must begin beyond Moderate-High\'s line",'
            # the floor before High beyond the worst: lines 2-4 stand without
            # High only when line 2 is above zero (synthesis._lines)
            'IF(AND({fz},{b2}<=0),"Moderate-Low would begin at or below zero",'
            'IF({s}*({w}-{b5})<0,"{hb}","")))))))))'
            .format(so=h("set_ok"), mv=mv, tl=TOO_LITTLE, hw=hw, hz=hz, fz=fz, m=med, ls=ls, s=sg,
                    b5=h("ib5"), b4=h("ib4"), w=h("ind_worst"), hb=HIGH_BEYOND, b2=h("ib2")))
        f["ind_why"] = '=IF(AND(NOT(EXACT({t},"none")),{n}=0),"{ns}",{w})'.format(
            t=TRANS, n=h("ind_ratio_n"), ns=NO_SHARED, w=h("ind_scale_why"))
        if m:
            f["below"] = '=IF({k}="","",SUMPRODUCT(ISNUMBER({r})*({r}<{k})))'.format(k=h("kb_med"), r=meds)
            f["worse_than"] = ('=IF({k}="","",{b})'.format(k=h("kb_med"), b=h("below")) if higher else
                               '=IF({k}="","",{n}-{b}-SUMPRODUCT(ISNUMBER({r})*({r}={k})))'
                               .format(k=h("kb_med"), n=h("held"), b=h("below"), r=meds))
            f["peer_why"] = '=IF({f}="","{nf}",IF({n}<2,"{fp}",""))'.format(f=h("f_ratio"), nf=NO_FILED,
                                                                          n=h("held"), fp=FEW_PEERS)
        else:
            f["below"] = f["worse_than"] = '=""'
            f["peer_why"] = '="%s"' % NO_PEERS
        mth = "Thresholds!$M$%d" % t
        # Thresholds refuses High beyond the worst before the floor; lines 2-4
        # stand only when line 2 is above zero under the floor (synthesis._lines)
        f["int_why"] = ('=IF({m}="","",IF({m}="{thb}",IF(AND({fz},Thresholds!$E${t}<=0),'
                        '"Moderate-Low would begin at or below zero","{hb}"),'
                        'IF({m}="{ttl}","{tl}",MID({m},10,500))))'
                        .format(m=mth, thb=TH_HIGH_BEYOND, hb=HIGH_BEYOND, ttl=TH_TOO_LITTLE, tl=TOO_LITTLE,
                                fz=fz, t=t))
        # The latest internal quarter, rated on the chosen lines when the Check is OK.
        J = "Thresholds!$J$%d" % t
        chosen_rng = "$G$%d:$G$%d" % (lr[0], lr[3])
        hit = ("SUMPRODUCT(({c}<={v})*1)", "SUMPRODUCT(({c}<{v})*1)") if higher else \
              ("SUMPRODUCT(({c}>={v})*1)", "SUMPRODUCT(({c}>{v})*1)")
        f["rating"] = ('=IF({ck}<>"","",IF(AND({fz},{v}<=0),1,1+IF(EXACT({o},"worse"),{w},{bt})))'
                       .format(ck=h("check"), fz=fz, v=J, o=onl, w=hit[0].format(c=chosen_rng, v=J),
                               bt=hit[1].format(c=chosen_rng, v=J)))
        f["rating_name"] = '=IF({r}="","",CHOOSE({r},"Low","Moderate-Low","Moderate","Moderate-High","High"))'.format(
            r=h("rating"))
        for name, formula in f.items():
            sy.cell(hr[lb], HELP[name], formula)

        # the four lines
        why = lambda x: ('IF(EXACT({x},"internal"),IF({a}="","not built",{a}),IF(EXACT({x},"industry"),'    # noqa
                         'IF({b}="","not built",{b}),IF({c}="","not built",{c})))'
                         .format(x=x, a=h("int_why"), b=h("ind_why"), c=h("peer_why")))
        for k, r in enumerate(lr):
            a = sy.cell(r, 1, "%s from" % LINES[k])
            a.fill = AT.fill(AT.BAND[k + 1])
            a.font = Font(name="Arial", bold=True, size=9, color=K.INK_TEXT)
            part = "" if k == 3 else ',AND(LEFT({w},7)="line 5:")'
            sy.cell(r, COL["internal"], '=IF(OR({w}=""{p}),Thresholds!${c}${t},"")'.format(
                w=h("int_why"), p=part.format(w=h("int_why")), c="EFGH"[k], t=t))
            sy.cell(r, COL["industry"], '=IF(AND(ISNUMBER({q}),OR({w}=""{p})),{b}*{q},"")'.format(
                q=h("ind_ratio"), w=h("ind_why"), p=part.format(w=h("ind_why")), b=h("ib%d" % (k + 2))))
            if m:
                p = PCT[k]
                pp = "%s/100" % p if higher else "(100-%s)/100" % p
                sy.cell(r, COL["pct"], s["peer_percentiles"][k])
                AT._input(sy.cell(r, COL["pct"]))
                pct_dv.add(sy.cell(r, COL["pct"]).coordinate)
                sy.cell(r, COL["peers"], '=IF(AND({w}="",ISNUMBER({p}),{p}>=0,{p}<=100),PERCENTILE({r},{pp})*{f},"")'
                        .format(w=h("peer_why"), p=p, r=meds, pp=pp, f=h("f_ratio")))
            else:
                sy.cell(r, COL["pct"], "no peers").font = Font(name="Calibri", size=10, italic=True, color=K.SLATE)
                sy.cell(r, COL["peers"], '=""')
            sy.cell(r, COL["lead"], s["leads"][k])
            AT._input(sy.cell(r, COL["lead"]))
            lead_dv.add(sy.cell(r, COL["lead"]).coordinate)
            sy.cell(r, COL["chosen"], '=IF(EXACT({l},"internal"),{a},IF(EXACT({l},"industry"),{b},'
                                      'IF(EXACT({l},"peers"),{c},"")))'
                    .format(l=LEAD[k], a=cell("internal", r), b=cell("industry", r), c=cell("peers", r)))
            sy.cell(r, COL["chosen"]).font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
            for cc in (COL["internal"], COL["industry"], COL["peers"], COL["chosen"]):
                sy.cell(r, cc).number_format = nf
            for cc in range(2, 9):
                if cc not in (COL["lead"],) and not (m and cc == COL["pct"]):
                    sy.cell(r, cc).border = Border(bottom=Side(style="thin", color=AT.HAIR))

        # the industry's carrier, the latest quarter, the competitive check, the Check
        rt = b + ROW["translation"]
        sy.cell(rt, 1, "Industry carried by (yours)").font = LABEL_F
        sy.cell(rt, 2, s["translation"])
        AT._input(sy.cell(rt, 2))
        trans_dv.add(sy.cell(rt, 2).coordinate)
        sy.merge_cells(start_row=rt, start_column=3, end_row=rt, end_column=8)
        sy.cell(rt, 3, "median ratio: times the median of internal / industry over the quarters both hold. "
                       "none: the industry's lines as they are.").font = AT.F_NOTE

        rl = b + ROW["latest"]
        sy.cell(rl, 1, "Latest quarter, on the chosen lines").font = LABEL_F
        sy.cell(rl, 2, "=Thresholds!$I$%d" % t).font = DATA_F
        sy.cell(rl, 3, "=Thresholds!$J$%d" % t).font = DATA_F
        sy.cell(rl, 3).number_format = nf
        sy.cell(rl, 4, "=%s" % h("rating")).font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
        sy.cell(rl, 5, "=%s" % h("rating_name")).font = DATA_F
        for kk in range(5):
            sy.conditional_formatting.add("D%d" % rl, FormulaRule(formula=["$D$%d=%d" % (rl, kk + 1)],
                                                                  fill=AT.fill(AT.BAND[kk])))

        rc = b + ROW["competitive"]
        sy.cell(rc, 1, "Competitive check").font = LABEL_F
        sy.merge_cells(start_row=rc, start_column=2, end_row=rc, end_column=8)
        if m:
            sy.cell(rc, 2, ('=IF({k}="","No competitive check: KeyBank\'s filed series holds no quarter of the '
                            'internal window","KeyBank\'s filed median "&TEXT({k},"0.000")&"{u} over the internal '
                            'window; it is worse than "&{w}&" of "&{n}&" peers")')
                    .format(k=h("kb_med"), u=unit, w=h("worse_than"), n=h("held")))
        else:
            sy.cell(rc, 2, "No competitive check: no peer series were given")
        sy.cell(rc, 2).font = DATA_F

        ck = b + ROW["check"]
        sy.cell(ck, 1, "Check").font = LABEL_F
        sy.merge_cells(start_row=ck, start_column=2, end_row=ck, end_column=8)
        g = lambda x: 'IFERROR(ROUND({x},3-INT(LOG10(ABS({x})))),{x})'.format(x=x)              # noqa: E731
        tail = '""'
        for k in (2, 1, 0):                                         # order, last pair innermost
            tail = ('IF({s}*({b}-{a})<=0,"Refused: the chosen lines are out of order: {nb} ("&{lb}&", "&{gb}&") '
                    'does not begin beyond {na} ("&{la}&", "&{ga}&"). Lead one of them from another source",{t})'
                    .format(s=sg, b=CH[k + 1], a=CH[k], nb=LINES[k + 1], lb=LEAD[k + 1], gb=g(CH[k + 1]),
                            na=LINES[k], la=LEAD[k], ga=g(CH[k]), t=tail))
        for k in (3, 2, 1, 0):
            tail = ('IF({c}="","Refused: line {nm} is led by "&{l}&", which has no such line: "&{w},{t})'
                    .format(c=CH[k], nm=LINES[k], l=LEAD[k], w=why(LEAD[k]), t=tail))
        if m:
            tail = ('IF(NOT({ok}),"Refused: peer percentiles are four numbers from 0 to 100, one per line",'
                    'IF(NOT({ri}),"Refused: peer percentiles must rise from line 2 to line 5",{t}))'
                    .format(ok=h("pct_ok"), ri=h("pct_rise"), t=tail))
        tail = ('IF(NOT({so}),"Refused: a setting in Your settings on Assess is not a number; its Check there says '
                'which",IF(NOT({oo}),"Refused: on the line must be worse or better",'
                'IF(NOT({to}),"Refused: translation must be one of {tr}",'
                'IF(NOT({lo}),"Refused: leads names the source of each of the four lines (2, 3, 4, 5), each one of '
                '{src}",{t}))))'
                .format(so=h("set_ok"), oo=h("online_ok"), to=h("trans_ok"), tr=", ".join(TRANSLATIONS),
                        lo=h("lead_ok"), src=", ".join(SOURCES), t=tail))
        # The refusal in a hidden helper ("" when there is none), and the Check
        # cell saying OK or the refusal: written once, read twice.
        sy.cell(hr[lb], HELP["check"], "=" + tail)
        sy.cell(ck, 2, '=IF({c}="","OK: every line is taken from a source that has it, and the lines run in order",'
                       '{c})'.format(c=h("check")))
        sy.cell(ck, 2).font = Font(name="Arial", size=10, bold=True, color=K.INK_TEXT)
        sy.cell(ck, 2).alignment = Alignment(wrap_text=True, vertical="center")
        sy.row_dimensions[ck].height = 30
        sy.conditional_formatting.add("B%d:H%d" % (ck, ck), FormulaRule(
            formula=['LEFT($B$%d,8)="Refused:"' % ck], fill=K.solid(K.ALERT_FG), font=Font(bold=True, color=K.CRIMSON)))
        out_ck = ck

        # what each source rests on
        ind_q = sorted(by_q(ind_pts))
        own_q = sorted(next(x for x in loaded if x["label"] == lb)["by_q"])
        rows_ = [
            (ROW["int_rests"], "Internal rests on",
             '="Your lines on Thresholds, from the quarters kept on Assess (the series holds {n} quarters, {a} to '
             '{z})."&IF({w}="","",IF(LEFT({w},7)="line 5:"," It places lines 2 to 4 and not line 5: "&MID({w},9,500),'
             '" Refused: "&{w}))'.format(n=own_n, a=own_q[0], z=own_q[-1], w=h("int_why"))),
            (ROW["ind_rests"], "Industry rests on",
             '="The same scale on the industry\'s {n} quarters, {a} to {z}, "&IF(EXACT({t},"none"),"taken as they '
             'are.",IF(ISNUMBER({q}),"times "&TEXT({q},"0.000")&", the median of internal / industry over "&{qn}&'
             '" quarters.","not carried."))&IF({w}="","",IF(LEFT({w},7)="line 5:"," It places lines 2 to 4 and not '
             'line 5: "&MID({w},9,500)," Refused: "&{w}))'
             .format(n=len(ind_q), a=ind_q[0], z=ind_q[-1], t=TRANS, q=h("ind_ratio"), qn=h("ind_ratio_n"),
                     w=h("ind_why"))),
            (ROW["peer_rests"], "Peers rest on",
             ('="Each peer\'s median over the internal quarters ({m} peers given, "&{held}&" hold one), at the '
              'percentiles you state, times "&IF(ISNUMBER({f}),TEXT({f},"0.000")&", the median of internal / '
              'KeyBank filed over "&{fn}&" quarters.","the median of internal / KeyBank filed, which cannot be '
              'worked out.")&IF({w}="",""," Refused: "&{w})'
              .format(m=m, held=h("held"), f=h("f_ratio"), fn=h("f_ratio_n"), w=h("peer_why"))) if m else
             "Not built: no peer series were given."),
        ]
        for off, lbl_, formula in rows_:
            rr = b + off
            sy.cell(rr, 1, lbl_).font = LABEL_F
            sy.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=8)
            sy.cell(rr, 2, formula).font = Font(name="Calibri", size=10, color=K.SLATE)
            sy.cell(rr, 2).alignment = Alignment(wrap_text=True, vertical="top")
            sy.row_dimensions[rr].height = 27
        if m:
            ph = b + ROW["peer_head"]
            K.canvas_header(sy, ph, ["Peer", "Median over the internal quarters"], right_from=1)
            for j, ((nm, *_), (vc, wc)) in enumerate(zip(d["peers"], c["peers"])):
                rr = b + ROW["peer_top"] + j
                sy.cell(rr, 1, nm).font = DATA_F
                sy.cell(rr, 2, '=IF(COUNT({r})=0,"",MEDIAN({r}))'.format(r=rng(wc))).font = DATA_F
                sy.cell(rr, 2).number_format = nf
                K.hairlines(sy, rr, 1, 2)
            rr = b + ROW["peer_top"] + m
            sy.cell(rr, 1, "KeyBank, filed").font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
            sy.cell(rr, 2, "=%s" % h("kb_med")).font = Font(name="Calibri", size=11, bold=True, color=K.INK_TEXT)
            sy.cell(rr, 2).number_format = nf

        # the summary row Assess reads
        sum_r = 4 + i_s
        vals = ([lb] + ["=" + cell("internal", r) for r in lr] + ["=" + cell("industry", r) for r in lr]
                + ["=" + cell("peers", r) for r in lr] + ["=" + cell("lead", r) for r in lr]
                + ["=" + cell("chosen", r) for r in lr]
                + ['=IF({r}="","",{q}&": "&{r}&" "&{nm})'.format(r=h("rating"), q="Thresholds!$I$%d" % t,
                                                                  nm=h("rating_name")),
                   '=IF({c}="","OK",{c})'.format(c=h("check"))])
        for j, v in enumerate(vals):
            sy.cell(sum_r, SUMMARY_COL + j, v)
        out[lb] = dict(top=b, lines=lr, check=out_ck, latest=rl, competitive=rc, helpers=hr[lb],
                       translation=rt, peer_top=b + ROW["peer_top"] if m else None, summary=sum_r)

    for c_, w in zip("ABCDEFGH", (34, 12, 12, 13, 12, 13, 12, 30)):
        sy.column_dimensions[c_].width = w
    for cc in list(HELP.values()) + list(range(SUMMARY_COL, SUMMARY_COL + 23)):
        sy.column_dimensions[L(cc)].hidden = True
    for j, nm in enumerate(["Series"] + ["internal %d" % k for k in range(2, 6)] + ["industry %d" % k for k in range(2, 6)]
                           + ["peers %d" % k for k in range(2, 6)] + ["taken from %d" % k for k in range(2, 6)]
                           + ["chosen %d" % k for k in range(2, 6)] + ["latest", "check"]):
        sy.cell(3, SUMMARY_COL + j, nm)
    sy.freeze_panes = "B4"

    # ---- Assess: the product showing, where each colour starts ----
    _assess_panel(asx, assess_lay, len(order), nf)
    return out


def _assess_panel(asx, lay, ns, nf):
    """Below What each setting does, on the right: the four lines of the
    product showing in B5 from the three sources, the chosen ones, and the
    Synthesis Check. Read-only; the choices are made on the Synthesis tab."""
    p0 = lay["expl_top"] + 6
    last = 3 + ns
    sm = lambda j: "Synthesis!${c}$4:${c}${e}".format(c=L(SUMMARY_COL + j), e=last)   # noqa: E731
    mt = "MATCH($B$5,%s,0)" % sm(0)
    AT._section(asx, p0, 14, 21, '="Where each colour starts · "&$B$5')
    link = asx.cell(p0, 21, "Synthesis ›")
    link.hyperlink = Hyperlink(ref=link.coordinate, location="Synthesis!A1", display="Synthesis ›")
    link.font = Font(name="Arial", size=9, bold=True, color=K.KEY_RED, underline="single")
    link.alignment = Alignment(horizontal="right", vertical="bottom")
    AT._header(asx, p0 + 1, 14, ["Rating begins", "Internal", "Industry", "Peers", "Chosen line", "Taken from", "", ""],
               right_from=1)
    for k in range(4):
        r = p0 + 2 + k
        a = asx.cell(r, 14, "%s from" % LINES[k])
        a.fill = AT.fill(AT.BAND[k + 1])
        a.font = Font(name="Arial", size=9, bold=True, color=K.INK_TEXT)
        for j, base in enumerate((1, 5, 9, 17, 13)):
            c = asx.cell(r, 15 + j, '=IFERROR(INDEX({r},{m}),"")'.format(r=sm(base + k), m=mt))
            c.font = AT.F_INPUT if j == 3 else AT.F_DATA
            c.number_format = nf
            c.border = AT.HAIRLINE
    r = p0 + 6
    asx.cell(r, 14, "Latest, on the chosen lines").font = AT.F_LABEL
    asx.merge_cells(start_row=r, start_column=15, end_row=r, end_column=21)
    asx.cell(r, 15, '=IFERROR(INDEX({r},{m}),"")'.format(r=sm(21), m=mt)).font = AT.F_DATA
    r = p0 + 7
    asx.merge_cells(start_row=r, start_column=14, end_row=r, end_column=21)
    c = asx.cell(r, 14, '=IFERROR("Synthesis check: "&INDEX({r},{m}),"No synthesis for "&$B$5&": the workbook '
                        'was built without one for it")'.format(r=sm(22), m=mt))
    c.font = Font(name="Arial", size=9, bold=True, color=K.INK_TEXT)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    asx.row_dimensions[r].height = 27
    asx.conditional_formatting.add("N%d:U%d" % (r, r), FormulaRule(
        formula=['ISNUMBER(SEARCH("Refused:",$N$%d))' % r], fill=AT.fill(K.ALERT_FG), font=Font(color=K.CRIMSON)))
    return p0


def assess_panel_top(lay) -> int:
    return lay["expl_top"] + 6
