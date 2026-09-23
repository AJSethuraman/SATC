#!/usr/bin/env python3
"""The desk test, walked: twenty-seven screens from the emailed file to a read workbook.

    python tools/walk_desk_test.py --bundle build_pack.py --desk DIR --out DIR [--pylibs DIR]

`--desk` is an empty folder standing in for the person's; the bundle is copied
into it and every command is run there exactly as `docs/PROCEDURE-desk-test.md`
says to type it. `--out` receives one PNG per step, named as the procedure
references them. Run it again after a change and compare the folder with the
one in `docs/walkthrough/`: a step whose picture differs is either a defect or
an out-of-date procedure, and finding out which is the job (canon `walk`).

The edits a person would make by hand in Excel (step 16, a knob on `_config`;
steps 23-25, the picks in the form) are made here by writing the cells; the
procedure says so.
"""

from __future__ import annotations

import argparse
import os
from datetime import datetime
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import walkshot as w  # noqa: E402

PROMPT = "desk> "
TITLE = "Command Prompt"


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--bundle", required=True, type=Path)
    ap.add_argument("--desk", required=True, type=Path)
    ap.add_argument("--out", required=True, type=Path)
    ap.add_argument("--pylibs", type=Path, help="a folder holding openpyxl and PyYAML, when they are not installed")
    ap.add_argument("--work", type=Path, help="where LibreOffice renders go (default: beside the desk)")
    a = ap.parse_args(argv)
    desk, out = a.desk.resolve(), a.out.resolve()
    work = (a.work or desk.parent / (desk.name + "-render")).resolve()
    if desk.exists() and any(desk.iterdir()):
        print(f"{desk} is not empty; the walk starts from an empty folder", file=sys.stderr)
        return 1
    desk.mkdir(parents=True, exist_ok=True)
    out.mkdir(parents=True, exist_ok=True)
    shutil.copy(a.bundle, desk / "build_pack.py")
    env = dict(os.environ)
    if a.pylibs:
        env["PYTHONPATH"] = str(a.pylibs.resolve())
    import openpyxl  # noqa: F401  (needed for step 16; the desk's own Python may lack it)

    def term(n: int, name: str, cmd: list[str]) -> tuple[int, str, str]:
        rc, so, se = w.terminal(cmd, desk, out / f"step-{n:02d}-{name}.png", title=TITLE, prompt=PROMPT, env=env)
        print(f"step {n:02d} {name}: rc={rc}")
        return rc, so, se

    def shot(wb: w.Workbook, n: int, name: str, page: int, **kw) -> None:
        w.sheet(wb, page, out / f"step-{n:02d}-{name}.png", **kw)
        print(f"step {n:02d} {name}: page {page}")

    w.route([
        ("The emailed file", "Save build_pack.py into an empty folder. Nothing else is needed but Python.", "before step 1"),
        ("Command Prompt", "Check Python and the two libraries it needs.", "steps 1–2"),
        ("Command Prompt", "Make a book with a known answer, look at the folder, build the pack.", "steps 3–5"),
        ("Excel", "Read the pack: the cover, the seven tabs, the knobs; move a knob and watch the cover notice.", "steps 6–17"),
        ("Notepad", "Open the planted answer and compare it with the model's interval.", "step 18"),
        ("Command Prompt + Excel", "Make a book with nothing in it, build it, read a cover that says so.", "steps 19–21"),
        ("Command Prompt", "Your own extract: one command writes a form beside it.", "step 22"),
        ("Excel", "Pick a role beside each column from its dropdown, two answers on the second tab, save.", "steps 23–25"),
        ("Command Prompt", "The same command again reads the form, writes the question file and builds the pack.", "step 26"),
        ("Excel", "The cover of the pack built from your picks.", "step 27"),
    ], out / "step-00-route.png")

    # Part A · is the machine ready?
    term(1, "python-version", ["python", "--version"])
    term(2, "libraries", ["python", "-c", "import openpyxl, yaml; print('ready: openpyxl', openpyxl.__version__, 'and PyYAML', yaml.__version__)"])
    # Part B · a book with a known answer
    term(3, "make-book", ["python", "build_pack.py", "--synth", "demo"])
    term(4, "the-folder", ["ls", "-1", ".", "demo"])
    term(5, "build", ["python", "build_pack.py", "--data", "demo/loans.csv", "--asof", "2026-06-30", "--config", "demo/config.yaml", "-o", "demo/pack.xlsx"])
    # Part C · read the workbook
    wb = w.Workbook(desk / "demo/pack.xlsx", work / "demo")
    shot(wb, 6, "cover", 1, ring=wb.span(1, "The answer, in three lines", "211 events.", full_width=True),
         zoom=wb.find(1, "formula checks agree", pad=0.01), zoom_caption="the last line of the cover", zoom_width=700, stack=True)
    q4 = wb.span(3, "Quarter", "2019Q4", full_width=True)
    shot(wb, 7, "capture", 3, ring=q4, zoom=q4, zoom_caption="the first four quarters, every column", zoom_width=1000, stack=True)
    q4 = wb.span(5, "Quarter", "2019Q4", full_width=True)
    shot(wb, 8, "prevalence", 5, ring=q4, zoom=q4, zoom_caption="the first four quarters: capture rate and flag rate, each with its interval", zoom_width=1000, stack=True)
    top = (0.06, 0.11, 0.9, 0.32)
    shot(wb, 9, "gradient", 6, ring=(0.06, 0.20, 0.40, 0.16), zoom=(0.06, 0.20, 0.40, 0.16), zoom_caption="the bucket table and the two lines under it",
         zoom_width=1000, stack=True, page_crop=top)
    shot(wb, 10, "gradient-chart", 6, ring=(0.69, 0.21, 0.22, 0.13), zoom=(0.69, 0.21, 0.22, 0.13), zoom_caption="the chart to the right of the table",
         zoom_width=600, stack=True, page_crop=top)
    v = (0.06, 0.2835, 0.30, 0.044)
    shot(wb, 11, "stratified", 7, ring=v, zoom=v, zoom_caption="under the first block: crude, pooled, share kept, and the word",
         zoom_width=1000, stack=True, page_crop=(0.06, 0.11, 0.9, 0.40))
    d = wb.span(9, "by category_1", "central", full_width=True)
    shot(wb, 12, "decomposition", 9, ring=d, zoom=d, zoom_caption="the first dimension, rows sorted by flagged events", zoom_width=1000, stack=True)
    m = wb.span(10, "M1: flag + controls", "intercept", full_width=True)
    shot(wb, 13, "model", 10, ring=wb.find(10, "3.729", pad=0.008), zoom=m,
         zoom_caption="M1: the flag row is the first term; its odds ratio and interval are the answer", zoom_width=1000, stack=True, page_crop=(0.06, 0.11, 0.9, 0.62))
    c = wb.span(12, "Seasoned loans with both fields present", "Nothing in the process reacts to it.", full_width=True)
    shot(wb, 14, "control", 12, ring=c, zoom=c, zoom_caption="the counts and the sentence built from them", zoom_width=900, stack=True, page_crop=(0.06, 0.11, 0.9, 0.35))
    pc = wb.page_with("Live knobs")
    k = wb.span(pc, "Live knobs", "not a knob.", full_width=True)
    shot(wb, 15, "config", pc, ring=k, zoom=(0.06, k[1], 0.83, k[3]), zoom_caption="the live knobs", zoom_width=1000, stack=True, page_crop=(0.06, 0.11, 0.9, 0.38))
    # step 16: the knob, written by openpyxl in place of a hand
    book = openpyxl.load_workbook(desk / "demo/pack.xlsx")
    cell = book.defined_names["METHOD"].attr_text.split("!")[1].replace("$", "")
    book["_config"][cell].value = "Clopper-Pearson"
    book.save(desk / "demo/pack-knob.xlsx")
    wbk = w.Workbook(desk / "demo/pack-knob.xlsx", work / "knob")
    pc = wbk.page_with("Live knobs")
    k = wbk.span(pc, "Live knobs", "not a knob.", full_width=True)
    shot(wbk, 16, "knob-changed", pc, ring=wbk.find(pc, "Clopper-Pearson", pad=0.01), zoom=(0.06, k[1], 0.83, k[3]),
         zoom_caption="Interval method now reads Clopper-Pearson", zoom_width=1000, stack=True, page_crop=(0.06, 0.11, 0.9, 0.38))
    line = wbk.span(1, "The live knobs have been moved", "set the knobs back to compare.", full_width=True)
    said = [ln for ln in wbk.text(1).splitlines() if "have been moved" in ln][0]
    shot(wbk, 17, "knob-cover", 1, ring=line, zoom=(line[0], line[1] - 0.004, 0.48, line[3] + 0.008),
         zoom_caption="the first half of the cover's last line, which now reads: " + said, zoom_width=1000, stack=True, page_crop=(0.0, 0.0, 1.0, 0.45))
    w.textfile(desk / "demo/planted.json", out / "step-18-planted.png", highlight=[9], title="demo\\planted.json — Notepad")
    # Part D · a book with nothing in it
    term(19, "make-null-book", ["python", "build_pack.py", "--synth", "null", "--null"])
    term(20, "build-null", ["python", "build_pack.py", "--data", "null/loans.csv", "--asof", "2026-06-30", "--config", "null/config.yaml", "-o", "null/pack.xlsx"])
    wbn = w.Workbook(desk / "null/pack.xlsx", work / "null")
    ans = wbn.span(1, "The answer, in three lines", "events.", full_width=True)
    shot(wbn, 21, "null-cover", 1, ring=ans, zoom=ans, zoom_caption="the three answer lines on the book with nothing planted", zoom_width=1000, stack=True, page_crop=(0.06, 0.11, 0.9, 0.8))
    # Part E · your own extract: the form. One command writes it, the person picks
    # in Excel (openpyxl stands in for the hand here), the same command builds.
    cmd = ["python", "build_pack.py", "--setup", "demo/loans.csv"]
    term(22, "setup-form-written", cmd)
    wbf = w.Workbook(desk / "demo/question.xlsx", work / "form-blank")
    role = wbf.find(1, "pick from the list", pad=0.008)
    rows = wbf.span(1, "loan_id", "measure_b", full_width=True)
    table = (0.06, role[1] - 0.012, 0.86, rows[1] + rows[3] - role[1] + 0.024)     # the whole table, header to last row
    sheet_crop = (0.05, 0.12, 0.90, 0.34)                                          # the sheet on its landscape page
    shot(wbf, 23, "form-blank", 1, ring=(role[0], role[1], role[2], rows[1] + rows[3] - role[1]), zoom=table,
         zoom_caption="one row per column of the extract; the shaded role cell beside each offers a list when clicked",
         zoom_width=1000, stack=True, page_crop=sheet_crop)
    book = openpyxl.load_workbook(desk / "demo/question.xlsx")
    ws, wa = book["Columns"], book["Answers"]
    roles = {"loan_id": "loan number", "origination_date": "origination date", "field_a": "rule top",
             "field_b": "rule bottom + group", "outcome_date": "went bad: date", "amount": "group", "category_2": "group"}
    edges = {"field_b": "100000, 200000, 400000, 800000, 1600000", "amount": "60000, 250000"}
    r = 4
    while ws.cell(r, 2).value:
        col = ws.cell(r, 2).value
        ws.cell(r, 6).value = roles.get(col)
        ws.cell(r, 8).value = edges.get(col)
        r += 1
    r = 4
    while wa.cell(r, 1).value:
        q = wa.cell(r, 1).value
        if q.startswith("The word every tab"):
            wa.cell(r, 2).value = "event"
        if q.startswith("The as-of date"):
            wa.cell(r, 2).value = datetime(2026, 6, 30)
        r += 1
    book.save(desk / "demo/question.xlsx")
    wbd = w.Workbook(desk / "demo/question.xlsx", work / "form-filled")
    picked = wbd.span(1, "loan number", "went bad: date", full_width=True)
    shot(wbd, 24, "form-columns", 1, ring=picked, zoom=table,
         zoom_caption="the roles picked beside the seven columns in use, and edges typed for two of the groups",
         zoom_width=1000, stack=True, page_crop=sheet_crop)
    ans = wbd.span(2, "The word every tab", "The as-of date", full_width=True)
    whole = wbd.span(2, "Compare the two rule", "or its band edges", full_width=True)
    shot(wbd, 25, "form-answers", 2, ring=ans, zoom=(0.06, whole[1] - 0.02, 0.86, whole[3] + 0.03),
         zoom_caption="the Answers tab: the two required answers filled (the word for the event, the as-of date); the rest as written",
         zoom_width=1000, stack=True, page_crop=(0.05, 0.12, 0.90, 0.30))
    term(26, "setup-built", cmd)
    wbm = w.Workbook(desk / "demo/field_a_vs_field_b.xlsx", work / "mine")
    ans = wbm.span(1, "The answer, in three lines", "events.", full_width=True)
    shot(wbm, 27, "picked-cover", 1, ring=ans, zoom=ans, zoom_caption="the three answer lines of the pack built from your picks",
         zoom_width=1000, stack=True, page_crop=(0.06, 0.11, 0.9, 0.8))
    print(f"{len(list(out.glob('step-*.png')))} screens in {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
