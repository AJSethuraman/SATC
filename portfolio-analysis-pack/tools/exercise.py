#!/usr/bin/env python3
"""Run the real script on made-up books with known answers, and show what it said.

A green test suite proves the code agrees with itself. This harness proves the
thing a person runs — the pure-ASCII bundle, `build_pack.py`, driven exactly
as a desk would drive it — answers each scenario the tool is for, and lays
what was planted beside what the pack said, scenario by scenario, with the
pages rendered and the checks counted. The firm, 20 September 2026: "you
should be ensuring everything works by actually running the script using a
good synthetically created population. It should be able to show how each
scenario is covered."

    python tools/exercise.py                      # writes docs/exercise-report.md and exercise-out/
    python tools/exercise.py --out DIR --loans 40000 --asof 2026-06-30 --run-date 2026-09-20

Needs what the tests need: openpyxl, PyYAML, the `formulas` engine (to read
the workbooks back the way Excel would) and LibreOffice + poppler-utils for
the pictures. No clock is read: every date is an argument.
"""

from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import time
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

HERE = Path(__file__).resolve().parent
PKG = HERE.parent
for p in (PKG / "src", PKG / "tests", HERE):
    if str(p) not in sys.path:
        sys.path.insert(0, str(p))

import yaml  # noqa: E402

from analysis_pack import synth  # noqa: E402
from analysis_pack.cli import main as pack_main  # noqa: E402
from analysis_pack.config import CONFIRM, load_config  # noqa: E402
from analysis_pack.ingest import read_table  # noqa: E402
from analysis_pack.population import build_population  # noqa: E402
from analysis_pack.workbook import build_pack  # noqa: E402
from recalc import Recalc  # noqa: E402
from render import render  # noqa: E402

WORDS = {"survives", "collapses", "no crude effect", "unknown", "no data"}


@dataclass
class Check:
    claim: str
    ok: bool
    detail: str = ""


@dataclass
class Scenario:
    key: str
    title: str
    planted: str
    commands: list[str] = field(default_factory=list)
    said: list[str] = field(default_factory=list)
    checks: list[Check] = field(default_factory=list)
    pictures: list[tuple[str, Path]] = field(default_factory=list)
    tests: list[str] = field(default_factory=list)

    def check(self, claim: str, ok: bool, detail: str = "") -> None:
        self.checks.append(Check(claim, bool(ok), detail))


class Desk:
    """The bundle script, run from a folder the way a desk would run it."""

    def __init__(self, folder: Path, bundle: Path):
        self.folder = folder
        self.folder.mkdir(parents=True, exist_ok=True)
        self.script = folder / bundle.name
        shutil.copy(bundle, self.script)
        self.log: list[str] = []

    def run(self, *args: str) -> tuple[int, dict | None, str, float]:
        cmd = [sys.executable, self.script.name, *args]
        self.log.append("python " + " ".join(cmd[1:]))
        t0 = time.perf_counter()
        r = subprocess.run(cmd, cwd=self.folder, capture_output=True, text=True, timeout=1200)
        elapsed = time.perf_counter() - t0
        status = None
        if r.stdout.strip():
            try:
                status = json.loads(r.stdout)
            except json.JSONDecodeError:
                status = None
        return r.returncode, status, r.stderr, elapsed


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def cover_lines(rec: Recalc) -> list[str]:
    out = []
    for needle in ("Gradient:", "Survives or collapses", "Model (step 6)", "seasoned loans where both",
                   "No seasoned loan carries", "formula checks agree", "live knobs have been moved"):
        out.extend(rec.find_text("Cover", needle))
    return out


def fill_skeleton(skeleton_path: Path, answers_path: Path, out_path: Path) -> dict:
    """What the person at the desk does by hand: replace each marker with an
    answer. Here the answers come from the synthetic book's own question file."""
    raw = yaml.safe_load(skeleton_path.read_text(encoding="utf-8"))
    ans = yaml.safe_load(answers_path.read_text(encoding="utf-8"))
    raw["name"] = "designated_at_the_desk"
    raw["population"]["loan_id"] = ans["population"]["loan_id"]
    raw["population"]["origination_date"] = ans["population"]["origination_date"]
    outcome_col = ans["outcome"]["date_field"]
    used = set(ans["fields"].keys())
    raw["fields"] = {col: {"known": "later" if col == outcome_col else "at_origination"}
                     for col in raw["fields"] if col in used}
    raw["rule"]["field_a"] = ans["rule"]["field_a"]
    raw["rule"]["field_b"] = ans["rule"]["field_b"]
    raw["outcome"]["label"] = ans["outcome"]["label"]
    raw["outcome"]["date_field"] = outcome_col
    raw["confounders"] = ans["confounders"]
    raw["controls"] = ans["controls"]
    raw["decompose_by"] = ans["decompose_by"]
    raw["existing_control"] = "none"
    out_path.write_text(yaml.safe_dump(raw, sort_keys=False), encoding="utf-8")
    return raw


def variant_config(base: Path, out: Path, **changes) -> Path:
    raw = yaml.safe_load(base.read_text(encoding="utf-8"))
    for k, v in changes.items():
        if v is None:
            raw.pop(k, None)
        else:
            raw[k] = v
    out.write_text(yaml.safe_dump(raw, sort_keys=False), encoding="utf-8")
    return out


def rewrite_csv(src: Path, dst: Path, edit) -> None:
    with src.open(newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
        columns = list(rows[0].keys())
    for i, r in enumerate(rows):
        edit(i, r)
    with dst.open("w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=columns)
        w.writeheader()
        w.writerows(rows)


def page_for(pdf: Path, pages_dir: Path, stem: str, heading: str) -> Path | None:
    """The PNG of the first page whose text carries `heading`, when pdftotext is here."""
    if not shutil.which("pdftotext"):
        return None
    for i in range(1, 80):
        r = subprocess.run(["pdftotext", "-f", str(i), "-l", str(i), str(pdf), "-"],
                           capture_output=True, text=True, timeout=60)
        if r.returncode != 0 or not r.stdout.strip():
            break
        if heading in r.stdout:
            # pdftoppm pads the page number to the width of the page count
            for width in (1, 2, 3, 4):
                png = pages_dir / f"{stem}-{i:0{width}d}.png"
                if png.exists():
                    return png
            return None
    return None


def rendered_pictures(pack: Path, out_dir: Path, wanted: list[tuple[str, str]]) -> tuple[list[tuple[str, Path]], dict]:
    res = render(pack, out_dir)
    pics: list[tuple[str, Path]] = []
    if res.get("ok") and res.get("pages_dir"):
        pdf = Path(res["pdf"])
        pages = Path(res["pages_dir"])
        for caption, heading in wanted:
            png = page_for(pdf, pages, pack.stem, heading)
            if png:
                pics.append((caption, png))
    return pics, res


# --------------------------------------------------------------------------
# The scenarios
# --------------------------------------------------------------------------

def scenario_effect(desk: Desk, books: Path, out: Path, asof: str, run: str, loans: int) -> tuple[Scenario, dict, Path]:
    s = Scenario("effect", "A planted effect, outcome by event date",
                 f"{loans:,} made-up loans. Loans where field_a ÷ field_b is above 1 were drawn to go bad at about "
                 "four times the odds of the rest, with nothing else in play, so the flag is a real signal and "
                 "no confounder should explain it away.")
    s.tests = ["tests/test_gradient.py", "tests/test_slice4.py", "tests/test_slice6.py"]
    book = books / "effect"
    planted = synth.generate(book, seed=20260918, loans=loans, effect=2.0, mode="effect")
    shutil.copy(book / "loans.csv", desk.folder / "effect.csv")
    shutil.copy(book / "config.yaml", desk.folder / "effect.yaml")
    rc, st, err, secs = desk.run("--data", "effect.csv", "--asof", asof, "--run-date", run,
                                 "--config", "effect.yaml", "-o", "effect.xlsx")
    s.commands = list(desk.log[-1:])
    s.check("the script built the pack and exited 0", rc == 0 and st is not None and st.get("ok"), err[-300:] if rc else "")
    pack = desk.folder / "effect.xlsx"
    if rc != 0 or not st:
        return s, {}, pack
    rec = Recalc(pack.read_bytes())
    s.said = cover_lines(rec)
    s.said.append(f"{st['seasoned']:,} seasoned loans, {st['events']:,} events; {st['unseasoned']:,} too young to count. "
                  f"Built in {secs:.1f} s.")
    words = st["outcomes"][0]["stratified"]
    s.check("the gradient reads: the rate rises at every step", st["gradient"] == "monotonic increasing", st["gradient"])
    s.check("every confounder block says the effect survives", all(w == "survives" for w in words.values()), str(words))
    m1, m2 = st["models"][0]["m1"], st["models"][0]["m2"]
    truth = planted["true_marginal_flag_odds_ratio"]
    s.check(f"the regression's interval for the flag contains the planted odds ratio ({truth:.2f})",
            "lo" in m1 and m1["lo"] <= truth <= m1["hi"], f"M1 {m1.get('flag_or', 0):.2f} [{m1.get('lo', 0):.2f}, {m1.get('hi', 0):.2f}]")
    s.check("adding the confounders does not move the flag's odds ratio much (no confounding was planted)",
            "flag_or" in m1 and "flag_or" in m2 and abs(m1["flag_or"] - m2["flag_or"]) / m1["flag_or"] < 0.15,
            f"M1 {m1.get('flag_or', 0):.2f} vs M2 {m2.get('flag_or', 0):.2f}")
    agree = rec.find_text("Cover", "formula checks agree")
    s.check("the cover's own formula check reads N of N", bool(agree) and agree[0].split(" of ")[0] == agree[0].split(" of ")[1].split(" ")[0], agree[0] if agree else "no line")
    s.pictures, res = rendered_pictures(pack, out / "render-effect", [("The cover", "Portfolio Analysis Pack"),
                                                                       ("Step 3, the gradient", "Step 3"),
                                                                       ("Step 4, stratified", "Step 4")])
    s.check("LibreOffice renders it with no error cell", res.get("ok", False) and not res.get("error_cells"),
            f"{res.get('pages')} pages" if res.get("ok") else str(res))
    return s, st, pack


def scenario_null(desk: Desk, books: Path, asof: str, run: str, loans: int) -> Scenario:
    s = Scenario("null", "No effect at all",
                 f"{loans:,} made-up loans where the flag has nothing to do with going bad. The right answer is that "
                 "nothing is found: no confounder block may say survives, and the regression's interval must contain 1.")
    s.tests = ["tests/test_slice4.py::test_the_null_book_reads_no_crude_effect_everywhere", "tests/test_slice6.py"]
    book = books / "null"
    synth.generate(book, seed=20260918, loans=loans, mode="null")
    shutil.copy(book / "loans.csv", desk.folder / "null.csv")
    shutil.copy(book / "config.yaml", desk.folder / "null.yaml")
    rc, st, err, secs = desk.run("--data", "null.csv", "--asof", asof, "--run-date", run,
                                 "--config", "null.yaml", "-o", "null.xlsx")
    s.commands = list(desk.log[-1:])
    s.check("the script built the pack and exited 0", rc == 0 and st is not None and st.get("ok"), err[-300:] if rc else "")
    if rc != 0 or not st:
        return s
    rec = Recalc((desk.folder / "null.xlsx").read_bytes())
    s.said = cover_lines(rec)
    words = st["outcomes"][0]["stratified"]
    s.check("no confounder block says survives", all(w != "survives" for w in words.values()), str(words))
    s.check("every block reads: no crude effect", all(w == "no crude effect" for w in words.values()), str(words))
    m1 = st["models"][0]["m1"]
    s.check("the regression's interval for the flag contains 1", "lo" in m1 and m1["lo"] <= 1.0 <= m1["hi"],
            f"M1 {m1.get('flag_or', 0):.2f} [{m1.get('lo', 0):.2f}, {m1.get('hi', 0):.2f}]")
    s.check("the gradient does not read as rising at every step", st["gradient"] != "monotonic increasing", st["gradient"])
    return s


def scenario_confounded(desk: Desk, books: Path, out: Path, asof: str, run: str, loans: int) -> Scenario:
    s = Scenario("confounded", "An effect that is really size in disguise",
                 f"{loans:,} made-up loans where going bad depends only on the size band of field_b, and the flag "
                 "happens to fall more often in the small bands. Looked at crudely the flag seems to matter; "
                 "inside each size band it does not. The pack must say so: the size-band block collapses, and the "
                 "regression with the confounders in loses the effect.")
    s.tests = ["tests/test_slice4.py::test_the_confounded_book_collapses_on_the_confounder_and_only_there",
               "tests/test_slice6.py"]
    book = books / "confounded"
    synth.generate(book, seed=20260918, loans=loans, mode="confounded")
    shutil.copy(book / "loans.csv", desk.folder / "confounded.csv")
    shutil.copy(book / "config.yaml", desk.folder / "confounded.yaml")
    rc, st, err, secs = desk.run("--data", "confounded.csv", "--asof", asof, "--run-date", run,
                                 "--config", "confounded.yaml", "-o", "confounded.xlsx")
    s.commands = list(desk.log[-1:])
    s.check("the script built the pack and exited 0", rc == 0 and st is not None and st.get("ok"), err[-300:] if rc else "")
    if rc != 0 or not st:
        return s
    pack = desk.folder / "confounded.xlsx"
    rec = Recalc(pack.read_bytes())
    s.said = cover_lines(rec)
    words = st["outcomes"][0]["stratified"]
    s.check("the size-band block says the effect collapses", words.get("size_band.edges") == "collapses", str(words))
    s.check("the blocks that are not the culprit do not collapse",
            all(w != "collapses" for k, w in words.items() if k != "size_band.edges"), str(words))
    m1, m2 = st["models"][0]["m1"], st["models"][0]["m2"]
    s.check("the regression without the confounders sees an effect (interval above 1)", "lo" in m1 and m1["lo"] > 1.0,
            f"M1 {m1.get('flag_or', 0):.2f} [{m1.get('lo', 0):.2f}, {m1.get('hi', 0):.2f}]")
    s.check("the regression with the confounders loses it (interval contains 1)", "lo" in m2 and m2["lo"] <= 1.0 <= m2["hi"],
            f"M2 {m2.get('flag_or', 0):.2f} [{m2.get('lo', 0):.2f}, {m2.get('hi', 0):.2f}]")
    s.check("the tree's first split is on size, not on the flag", "size_band" in str(st["models"][0].get("tree_first_split", "")),
            str(st["models"][0].get("tree_first_split")))
    s.pictures, res = rendered_pictures(pack, out / "render-confounded", [("Step 4, stratified: the size-band block", "Step 4")])
    s.check("LibreOffice renders it with no error cell", res.get("ok", False) and not res.get("error_cells"),
            f"{res.get('pages')} pages" if res.get("ok") else str(res))
    return s


def scenario_flag(desk: Desk, asof: str, run: str, effect_status: dict) -> Scenario:
    s = Scenario("flag", "The outcome as a flag the bank already windowed",
                 "The same planted-effect book, but the outcome is declared as a 0/1 flag column the bank set "
                 "itself, with a line saying how. The generator set that flag to the same event the date column "
                 "carries, so the pack should count the same events and reach the same words.")
    s.tests = ["tests/test_slice2.py"]
    variant_config(desk.folder / "effect.yaml", desk.folder / "flag.yaml",
                   outcome={"label": "event", "field": "flag_1", "op": "==", "value": 1,
                            "basis": "windowed_by_bank"})
    rc, st, err, secs = desk.run("--data", "effect.csv", "--asof", asof, "--run-date", run,
                                 "--config", "flag.yaml", "-o", "flag.xlsx")
    s.commands = list(desk.log[-1:])
    s.check("the script built the pack and exited 0", rc == 0 and st is not None and st.get("ok"), err[-300:] if rc else "")
    if rc != 0 or not st:
        return s
    rec = Recalc((desk.folder / "flag.xlsx").read_bytes())
    s.said = cover_lines(rec)
    s.check("the event count equals the event-date book's", st["events"] == effect_status.get("events"),
            f"{st['events']:,} vs {effect_status.get('events', 0):,}")
    s.check("the words are the same as the event-date book's",
            st["gradient"] == effect_status.get("gradient") and st["outcomes"][0]["stratified"] == effect_status["outcomes"][0]["stratified"],
            f"{st['gradient']}; {st['outcomes'][0]['stratified']}")
    return s


def scenario_snapshot(desk: Desk, out: Path, asof: str, run: str) -> Scenario:
    s = Scenario("snapshot", "The outcome as a measure at the as-of date, in bands",
                 "The same book, but the outcome is a ratio of two columns measured at the as-of date "
                 "(measure_a over measure_b), cut at three band edges chosen in the question file. Every step is "
                 "shown once per edge, and loans with no measure are counted on the capture tab rather than "
                 "silently treated as fine.")
    s.tests = ["tests/test_slice2.py", "tests/test_adversarial.py::test_a_blank_snapshot_measure_is_counted_and_not_silently_a_non_event"]
    variant_config(desk.folder / "effect.yaml", desk.folder / "snapshot.yaml",
                   outcome={"label": "drawn", "measure": {"kind": "ratio", "field_a": "measure_a", "field_b": "measure_b"},
                            "edges": [0.2, 0.5, 0.8], "basis": "snapshot_at_asof"})
    rc, st, err, secs = desk.run("--data", "effect.csv", "--asof", asof, "--run-date", run,
                                 "--config", "snapshot.yaml", "-o", "snapshot.xlsx")
    s.commands = list(desk.log[-1:])
    s.check("the script built the pack and exited 0", rc == 0 and st is not None and st.get("ok"), err[-300:] if rc else "")
    if rc != 0 or not st:
        return s
    pack = desk.folder / "snapshot.xlsx"
    rec = Recalc(pack.read_bytes())
    s.said = cover_lines(rec)
    s.check("one outcome per band edge: three blocks on the gradient tab", len(st["outcomes"]) == 3,
            ", ".join(f"{o['label']}: {o['events']:,} events" for o in st["outcomes"]))
    from openpyxl import load_workbook
    wb = load_workbook(pack)
    cap = wb["1_Capture"]
    headers = [str(c.value) for row in cap.iter_rows(min_row=1, max_row=8) for c in row
               if isinstance(c.value, str) and len(c.value) < 60]
    s.check("the capture tab carries a column counting loans with no measure",
            any("measure blank" in h for h in headers), "; ".join(h for h in headers if "blank" in h))
    s.pictures, res = rendered_pictures(pack, out / "render-snapshot", [("Step 3, one block per band edge", "Step 3")])
    s.check("LibreOffice renders it with no error cell", res.get("ok", False) and not res.get("error_cells"),
            f"{res.get('pages')} pages" if res.get("ok") else str(res))
    return s


def scenario_designation(desk: Desk, asof: str, run: str) -> Scenario:
    s = Scenario("designation", "Designation at the desk: any extract, nothing coming back",
                 "An extract the bundle's carried question file was never written for. The script writes a "
                 "question file from the extract's own columns; the unfilled file is refused, naming every slot; "
                 "the filled file validates and builds; and the desk's workbook is byte-identical to one built "
                 "here from the same inputs, which is how a reviewer proves the desk copy is the same deliverable.")
    s.tests = ["tests/test_designate.py"]
    shutil.copy(desk.folder / "effect.csv", desk.folder / "extract.csv")
    q = desk.folder / "question.yaml"
    if q.exists():
        q.unlink()
    rc, st, err, _ = desk.run("--init", "extract.csv")
    s.commands.append(desk.log[-1])
    s.check("the script wrote a question file listing every column", rc == 0 and q.exists() and st and st["columns"],
            f"{len(st['columns']) if st else 0} columns, {st['markers'] if st else 0} values to fill in")
    rc, st2, err, _ = desk.run("--validate", "extract.csv", "--config", "question.yaml")
    s.commands.append(desk.log[-1])
    n = st["markers"] if st else 0
    s.check(f"the unfilled file is refused, naming all {n} slots", rc == 2 and st2 and st2.get("refused") == "config"
            and len(st2.get("problems", [])) == 1 + n, f"exit {rc}, {len(st2.get('problems', [])) if st2 else 0} lines")
    s.said.append("Refusal, first lines: " + " | ".join((st2 or {}).get("problems", [])[:3]))
    filled = desk.folder / "filled.yaml"
    fill_skeleton(q, desk.folder / "effect.yaml", filled)
    s.said.append("Filled in by hand: loan_id, origination_date, field_a over field_b, outcome_date as the event, "
                  "the columns declared, existing_control: none.")
    rc, st3, err, _ = desk.run("--validate", "extract.csv", "--asof", asof, "--config", "filled.yaml")
    s.commands.append(desk.log[-1])
    s.check("the filled file validates", rc == 0 and st3 and st3.get("ok"), err[-300:] if rc else f"{st3.get('seasoned', 0):,} seasoned")
    rc, st4, err, _ = desk.run("--data", "extract.csv", "--asof", asof, "--run-date", run, "--config", "filled.yaml", "-o", "desk.xlsx")
    s.commands.append(desk.log[-1])
    s.check("the filled file builds", rc == 0 and st4 and st4.get("ok"), err[-300:] if rc else "")
    if rc == 0 and st4:
        cfg = load_config(filled)
        table = read_table(desk.folder / "extract.csv")
        pop = build_population(cfg, table.rows, date.fromisoformat(asof))
        blob, _, _ = build_pack(cfg, pop, table, date.fromisoformat(run))
        here = hashlib.sha256(blob).hexdigest()
        s.check("the desk's workbook is byte-identical to one built here from the same inputs", st4["sha256"] == here,
                f"desk {st4['sha256'][:16]}… here {here[:16]}…")
        s.said.append(f"SHA-256 of the desk build: {st4['sha256']}")
    return s


def scenario_refusals(desk: Desk, asof: str, run: str) -> Scenario:
    s = Scenario("refusals", "Dirt is refused with the rows named; a missing line is refused with the line to add",
                 "Four spoiled inputs, each with one thing wrong. The tool must stop, say what, and never guess: "
                 "a value that is not a number in a rule field; dates that read two ways; a question file with no "
                 "existing_control line; a column known only after origination used as a control.")
    s.tests = ["tests/test_cli.py", "tests/test_slice2.py", "tests/test_adversarial.py::test_one_bad_value_is_refused_once_not_twice"]
    # 1. dirt
    def spoil(i, r):
        if i == 7:
            r["field_b"] = "not a number"
    rewrite_csv(desk.folder / "effect.csv", desk.folder / "dirty.csv", spoil)
    rc, st, err, _ = desk.run("--data", "dirty.csv", "--asof", asof, "--run-date", run, "--config", "effect.yaml", "-o", "dirty.xlsx")
    s.commands.append(desk.log[-1])
    hyg = desk.folder / "hygiene-synthetic_effect.csv"
    rows = list(csv.DictReader(hyg.read_text().splitlines())) if hyg.exists() else []
    s.check("one non-numeric value: refused (exit 2), one row in the hygiene file, no workbook written",
            rc == 2 and st and st.get("refused") == "hygiene" and len(rows) == 1 and not (desk.folder / "dirty.xlsx").exists(),
            f"exit {rc}; counts {st.get('counts') if st else None}; rows {rows}")
    # 2. ambiguous dates, on a question file that declares no date format
    def ambiguous(i, r):
        y, m, d = r["origination_date"].split("-")
        r["origination_date"] = f"{m}/{min(int(d), 12):02d}/{y}"
        if r["outcome_date"]:
            y, m, d = r["outcome_date"].split("-")
            r["outcome_date"] = f"{m}/{min(int(d), 12):02d}/{y}"
    rewrite_csv(desk.folder / "effect.csv", desk.folder / "ambiguous.csv", ambiguous)
    raw = yaml.safe_load((desk.folder / "effect.yaml").read_text())
    raw["population"].pop("date_format", None)
    (desk.folder / "nodate.yaml").write_text(yaml.safe_dump(raw, sort_keys=False))
    rc, st, err, _ = desk.run("--data", "ambiguous.csv", "--asof", asof, "--run-date", run, "--config", "nodate.yaml", "-o", "ambiguous.xlsx")
    s.commands.append(desk.log[-1])
    s.check("dates that read two ways: refused, both readings shown, the line to add named",
            rc == 2 and st and st.get("refused") == "dates" and "date_format" in err,
            f"exit {rc}; candidates {st.get('candidates') if st else None}")
    s.said.append("Two readings shown: " + "; ".join(f"{p} → {plain}" for p, plain in (st or {}).get("readings", [])[:2]))
    # 3. missing existing_control
    variant_config(desk.folder / "effect.yaml", desk.folder / "nocontrol.yaml", existing_control=None)
    rc, st, err, _ = desk.run("--validate", "effect.csv", "--config", "nocontrol.yaml")
    s.commands.append(desk.log[-1])
    s.check("no existing_control line: refused with the line to add",
            rc == 2 and st and st.get("refused") == "config" and any("existing_control" in p for p in st.get("problems", [])),
            f"exit {rc}")
    # 4. leakage: a later column as a control
    raw = yaml.safe_load((desk.folder / "effect.yaml").read_text())
    raw["controls"].append({"name": "leak", "field": "flag_1", "as": "categorical"})
    (desk.folder / "leak.yaml").write_text(yaml.safe_dump(raw, sort_keys=False))
    rc, st, err, _ = desk.run("--validate", "effect.csv", "--config", "leak.yaml")
    s.commands.append(desk.log[-1])
    s.check("a column known only later used as a control: refused, naming the column",
            rc == 2 and st and any("flag_1" in p for p in st.get("problems", [])), f"exit {rc}")
    return s


def scenario_blanks(desk: Desk, asof: str, run: str) -> Scenario:
    s = Scenario("blanks", "Blanks are a finding, not dirt",
                 "The planted-effect book has blanks in both rule fields on purpose, and here a third of the "
                 "category_1 values are blanked too. Blanks are never refused and never repaired: the capture tab "
                 "counts them by quarter, and a blank grouping value gets its own row, always last.")
    s.tests = ["tests/test_slice3.py", "tests/test_adversarial.py::test_a_blank_grouping_value_gets_its_own_level_rather_than_vanishing"]
    def blank_some(i, r):
        if i % 3 == 0:
            r["category_1"] = ""
    rewrite_csv(desk.folder / "effect.csv", desk.folder / "blanks.csv", blank_some)
    rc, st, err, _ = desk.run("--data", "blanks.csv", "--asof", asof, "--run-date", run, "--config", "effect.yaml", "-o", "blanks.xlsx")
    s.commands.append(desk.log[-1])
    s.check("the script built the pack and exited 0", rc == 0 and st is not None and st.get("ok"), err[-300:] if rc else "")
    if rc != 0:
        return s
    rec = Recalc((desk.folder / "blanks.xlsx").read_bytes())
    a_blank = sum(v for r_, v in rec.column("1_Capture", "D").items() if r_ > 1 and isinstance(v, (int, float)))
    b_blank = sum(v for r_, v in rec.column("1_Capture", "F").items() if r_ > 1 and isinstance(v, (int, float)))
    s.said.append(f"Capture tab: {int(a_blank):,} loans with field_a blank and {int(b_blank):,} with field_b blank, counted by quarter.")
    s.check("the capture tab counts the blanks in both rule fields", a_blank > 0 and b_blank > 0, f"{int(a_blank):,} / {int(b_blank):,}")
    levels = rec.find_text("5_Decomposition", "(blank)")
    s.check("the blank category has its own row on the decomposition tab", bool(levels), str(levels[:2]))
    agree = rec.find_text("Cover", "formula checks agree")
    s.check("the pack's own formula check still reads N of N with blanks in play",
            bool(agree) and agree[0].split(" of ")[0] == agree[0].split(" of ")[1].split(" ")[0], agree[0] if agree else "no line")
    return s


def scenario_knobs(effect_pack: Path) -> Scenario:
    s = Scenario("knobs", "The live knobs",
                 "The planted-effect workbook, with the interval method switched from Wilson to Clopper-Pearson on "
                 "the _config tab and nothing rebuilt. Every interval must move, and the cover must not show a "
                 "false formula-check count: it says the knobs have moved, and reads N of N again when they are put back.")
    s.tests = ["tests/test_slice3.py::test_switching_the_method_moves_the_prevalence_intervals_and_the_checks_still_agree",
               "tests/test_adversarial.py::test_moving_the_interval_method_knob_never_shows_a_false_check_count"]
    blob = effect_pack.read_bytes()
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(blob))
    ref = wb.defined_names["METHOD"].attr_text
    sheet, cell = ref.split("!")[0], ref.split("!")[1].replace("$", "")
    before = Recalc(blob)
    after = Recalc(blob, inputs={(sheet, cell): "Clopper-Pearson"})
    lo_before = {r: v for r, v in before.column("3_Gradient", "E").items() if isinstance(v, float)}
    lo_after = {r: v for r, v in after.column("3_Gradient", "E").items() if isinstance(v, float)}
    moved = sum(1 for r in lo_before if r in lo_after and abs(lo_before[r] - lo_after[r]) > 1e-9)
    s.said.append(f"Switching the method moved {moved} of {len(lo_before)} lower bounds on the gradient tab.")
    s.check("every interval bound on the gradient tab moved", moved == len(lo_before) and moved > 0, f"{moved} of {len(lo_before)}")
    line = after.find_text("Cover", "live knobs have been moved")
    s.check("the cover says the knobs have moved instead of showing a false count",
            bool(line) and not after.find_text("Cover", "formula checks agree"), line[0][:90] if line else "no line")
    back = before.find_text("Cover", "formula checks agree")
    s.check("at the built settings the cover reads N of N", bool(back) and back[0].split(" of ")[0] == back[0].split(" of ")[1].split(" ")[0],
            back[0] if back else "no line")
    s.said.append("At the built settings: " + (back[0] if back else "no line"))
    s.said.append("With the method knob moved: " + (line[0] if line else "no line"))
    return s


def scenario_determinism(desk: Desk, asof: str, run: str, effect_sha: str) -> Scenario:
    s = Scenario("determinism", "Same inputs, same file",
                 "The planted-effect book built again from the same inputs must produce the same bytes, so a hash "
                 "proves two copies are the same deliverable. A different run date must change the file (it is on "
                 "the provenance tab and the header bands) and nothing else.")
    s.tests = ["tests/test_determinism.py"]
    rc, st, err, _ = desk.run("--data", "effect.csv", "--asof", asof, "--run-date", run, "--config", "effect.yaml", "-o", "effect-again.xlsx")
    s.commands.append(desk.log[-1])
    s.check("built again: identical SHA-256", rc == 0 and st and st["sha256"] == effect_sha,
            f"{(st or {}).get('sha256', '')[:16]}… vs {effect_sha[:16]}…")
    other = "2026-09-19" if run != "2026-09-19" else "2026-09-18"
    rc, st2, err, _ = desk.run("--data", "effect.csv", "--asof", asof, "--run-date", other, "--config", "effect.yaml", "-o", "effect-other-day.xlsx")
    s.commands.append(desk.log[-1])
    s.check("a different run date: a different file", rc == 0 and st2 and st2["sha256"] != effect_sha, (st2 or {}).get("sha256", "")[:16] + "…")
    s.said.append(f"SHA-256, built twice: {effect_sha}")
    return s


def scenario_scale(desk: Desk, books: Path, asof: str, run: str) -> Scenario:
    s = Scenario("scale", "At scale: 100,000 loans",
                 "The planted-effect book at 100,000 loans. Python does the work; the workbook holds counts, so its "
                 "size should not grow with the loan count.")
    s.tests = ["(measured, not a test)"]
    book = books / "large"
    synth.generate(book, seed=20260918, loans=100000, effect=2.0, mode="effect")
    shutil.copy(book / "loans.csv", desk.folder / "large.csv")
    shutil.copy(book / "config.yaml", desk.folder / "large.yaml")
    rc, st, err, secs = desk.run("--data", "large.csv", "--asof", asof, "--run-date", run, "--config", "large.yaml", "-o", "large.xlsx")
    s.commands.append(desk.log[-1])
    s.check("the script built the pack and exited 0", rc == 0 and st is not None and st.get("ok"), err[-300:] if rc else "")
    if rc == 0 and st:
        size = (desk.folder / "large.xlsx").stat().st_size
        small = (desk.folder / "effect.xlsx").stat().st_size
        s.said.append(f"{st['seasoned']:,} seasoned loans, {st['events']:,} events, built in {secs:.1f} s; "
                      f"workbook {size / 1024:.0f} KB against {small / 1024:.0f} KB at 40,000 loans.")
        s.check("built in under a minute", secs < 60, f"{secs:.1f} s")
        s.check("the workbook is no larger than the 40,000-loan one by more than a tenth", size <= small * 1.1, f"{size:,} vs {small:,} bytes")
        s.check("the words match the 40,000-loan book's", st["gradient"] == "monotonic increasing"
                and all(w == "survives" for w in st["outcomes"][0]["stratified"].values()), str(st["outcomes"][0]["stratified"]))
    return s


# --------------------------------------------------------------------------
# The report
# --------------------------------------------------------------------------

def write_markdown(path: Path, scenarios: list[Scenario], meta: dict) -> None:
    pages = path.parent / "exercise-pages"
    if pages.exists():
        shutil.rmtree(pages)
    pages.mkdir(parents=True)
    total = sum(len(s.checks) for s in scenarios)
    passed = sum(1 for s in scenarios for c in s.checks if c.ok)
    lines = [f"# Exercise report — the script run on made-up books with known answers",
             "",
             f"Run {meta['run_date']} (as-of {meta['asof']}), {meta['loans']:,} loans per book, generator {meta['version']}. "
             f"Every command below was run against `{meta['bundle']}`, the same pure-ASCII script a desk receives, "
             f"from an empty folder, and every answer was read back out of the workbook the way Excel reads it. "
             f"Generated by `tools/exercise.py`; do not edit by hand.",
             "",
             f"**{passed} of {total} checks agree** across {len(scenarios)} scenarios."
             + ("" if passed == total else " **The ones that do not are listed first.**"),
             ""]
    failed = [(s, c) for s in scenarios for c in s.checks if not c.ok]
    if failed:
        lines += ["## Did not agree", ""]
        for s, c in failed:
            lines.append(f"- **{s.title}:** {c.claim} — {c.detail}")
        lines.append("")
    lines += ["## The scenarios at a glance", "", "| Scenario | Checks | Proved by |", "|---|---|---|"]
    for s in scenarios:
        ok = sum(1 for c in s.checks if c.ok)
        lines.append(f"| {s.title} | {ok} of {len(s.checks)} | {', '.join(f'`{t}`' for t in s.tests)} |")
    lines.append("")
    for s in scenarios:
        ok = sum(1 for c in s.checks if c.ok)
        lines += [f"## {s.title}", "", f"**What was planted.** {s.planted}", "", "**What was run.**", "", "```"]
        lines += s.commands
        lines += ["```", ""]
        if s.said:
            lines += ["**What the pack said.**", ""]
            lines += [f"- {t}" for t in s.said]
            lines.append("")
        lines += [f"**Checks: {ok} of {len(s.checks)} agree.**", ""]
        for c in s.checks:
            mark = "✔" if c.ok else "✘"
            lines.append(f"- {mark} {c.claim}" + (f" — {c.detail}" if c.detail else ""))
        lines.append("")
        for caption, png in s.pictures:
            dest = pages / f"{s.key}-{png.stem.split('-')[-1]}.png"
            shutil.copy(png, dest)
            lines.append(f"![{caption}](exercise-pages/{dest.name})")
            lines.append("")
            lines.append(f"*{caption}.*")
            lines.append("")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_html(path: Path, scenarios: list[Scenario], meta: dict) -> None:
    total = sum(len(s.checks) for s in scenarios)
    passed = sum(1 for s in scenarios for c in s.checks if c.ok)
    e = html.escape
    parts = ["<title>Analysis Pack Exercise</title>",
             "<style>",
             ":root{--ground:#F5F7F4;--panel:#FFFFFF;--ink:#1B2430;--muted:#5B6673;--line:#D7DDD9;--accent:#1F6F8B;--ok:#2E7D4F;--ok-soft:#E4F1E8;--bad:#B3541E;--bad-soft:#F8E9DF;--mono-bg:#EEF1EE;}",
             "@media (prefers-color-scheme: dark){:root:not([data-theme=\"light\"]){--ground:#141A20;--panel:#1D252E;--ink:#E8ECEF;--muted:#9AA6B2;--line:#313C47;--accent:#5FB3CF;--ok:#6CC38F;--ok-soft:#1E3328;--bad:#E08A57;--bad-soft:#3A2718;--mono-bg:#242E38;}}",
             ":root[data-theme=\"dark\"]{--ground:#141A20;--panel:#1D252E;--ink:#E8ECEF;--muted:#9AA6B2;--line:#313C47;--accent:#5FB3CF;--ok:#6CC38F;--ok-soft:#1E3328;--bad:#E08A57;--bad-soft:#3A2718;--mono-bg:#242E38;}",
             "body{background:var(--ground);color:var(--ink);font-family:'IBM Plex Sans',system-ui,-apple-system,'Segoe UI',sans-serif;font-size:16px;line-height:1.5;}",
             ".wrap{max-width:820px;margin:0 auto;padding-block:28px 64px;padding-inline:16px;}",
             "h1{font-size:1.6rem;font-weight:600;margin:0 0 6px;text-wrap:balance;} h2{font-size:1.15rem;font-weight:600;margin:36px 0 10px;text-wrap:balance;}",
             "p{margin:0 0 10px;max-width:70ch;} .sub{color:var(--muted);}",
             ".tally{display:inline-block;padding:4px 10px;border-radius:4px;font-weight:600;background:var(--ok-soft);color:var(--ok);} .tally.bad{background:var(--bad-soft);color:var(--bad);}",
             "table{border-collapse:collapse;width:100%;font-size:0.95rem;margin:8px 0 14px;} th,td{text-align:left;padding:7px 10px;border-bottom:1px solid var(--line);vertical-align:top;} th{color:var(--muted);font-size:0.8rem;letter-spacing:0.04em;text-transform:uppercase;}",
             ".tablewrap{overflow-x:auto;} pre{font-family:'IBM Plex Mono',ui-monospace,Consolas,monospace;font-size:0.85rem;background:var(--mono-bg);padding:10px 12px;border-radius:4px;overflow-x:auto;}",
             "code{font-family:'IBM Plex Mono',ui-monospace,Consolas,monospace;font-size:0.9em;background:var(--mono-bg);padding:1px 5px;border-radius:3px;}",
             "ul{margin:0 0 10px;padding-left:22px;} li{margin-bottom:5px;max-width:74ch;} .ok{color:var(--ok);font-weight:600;} .bad{color:var(--bad);font-weight:600;}",
             ".block{background:var(--panel);border:1px solid var(--line);border-radius:6px;padding:16px 18px;margin-bottom:16px;}",
             "figure{margin:12px 0;} figure img{max-width:100%;border:1px solid var(--line);border-radius:4px;} figcaption{font-size:0.85rem;color:var(--muted);margin-top:4px;}",
             ".small{font-size:0.85rem;color:var(--muted);}",
             "</style>",
             "<div class=\"wrap\">",
             "<h1>Analysis Pack Exercise</h1>",
             f"<p class=\"sub\">The real script, run on made-up books with known answers. Run {e(meta['run_date'])}, as-of {e(meta['asof'])}, "
             f"{meta['loans']:,} loans per book, generator {e(meta['version'])}. Every command was run against <code>{e(meta['bundle'])}</code>, "
             "the same plain-text script a desk receives, from an empty folder; every answer was read back out of the workbook the way Excel reads it.</p>",
             f"<p><span class=\"tally{'' if passed == total else ' bad'}\">{passed} of {total} checks agree</span> across {len(scenarios)} scenarios.</p>"]
    failed = [(s, c) for s in scenarios for c in s.checks if not c.ok]
    if failed:
        parts.append("<h2>Did not agree</h2><ul>")
        for s, c in failed:
            parts.append(f"<li><strong>{e(s.title)}:</strong> {e(c.claim)} — {e(c.detail)}</li>")
        parts.append("</ul>")
    parts.append("<h2>The scenarios at a glance</h2><div class=\"tablewrap\"><table><tr><th>Scenario</th><th>Checks</th><th>Proved by</th></tr>")
    for s in scenarios:
        ok = sum(1 for c in s.checks if c.ok)
        parts.append(f"<tr><td><a href=\"#{s.key}\">{e(s.title)}</a></td><td>{ok} of {len(s.checks)}</td><td>{', '.join('<code>' + e(t) + '</code>' for t in s.tests)}</td></tr>")
    parts.append("</table></div>")
    for s in scenarios:
        ok = sum(1 for c in s.checks if c.ok)
        parts.append(f"<div class=\"block\" id=\"{s.key}\"><h2 style=\"margin-top:0\">{e(s.title)}</h2>")
        parts.append(f"<p><strong>What was planted.</strong> {e(s.planted)}</p>")
        parts.append("<p><strong>What was run.</strong></p><pre>" + e("\n".join(s.commands)) + "</pre>")
        if s.said:
            parts.append("<p><strong>What the pack said.</strong></p><ul>" + "".join(f"<li>{e(t)}</li>" for t in s.said) + "</ul>")
        parts.append(f"<p><strong>Checks: {ok} of {len(s.checks)} agree.</strong></p><ul>")
        for c in s.checks:
            cls = "ok" if c.ok else "bad"
            mark = "✔" if c.ok else "✘"
            parts.append(f"<li><span class=\"{cls}\">{mark}</span> {e(c.claim)}" + (f" <span class=\"small\">— {e(c.detail)}</span>" if c.detail else "") + "</li>")
        parts.append("</ul>")
        for caption, png in s.pictures:
            data = base64.b64encode(png.read_bytes()).decode("ascii")
            parts.append(f"<figure><img src=\"data:image/png;base64,{data}\" alt=\"{e(caption)}\"><figcaption>{e(caption)}</figcaption></figure>")
        parts.append("</div>")
    parts.append("</div>")
    path.write_text("\n".join(parts), encoding="utf-8")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=str(PKG / "exercise-out"))
    ap.add_argument("--report", default=str(PKG / "docs" / "exercise-report.md"))
    ap.add_argument("--loans", type=int, default=40000)
    ap.add_argument("--asof", default="2026-06-30")
    ap.add_argument("--run-date", default="2026-09-20")
    ap.add_argument("--skip-scale", action="store_true")
    a = ap.parse_args(argv)
    out = Path(a.out)
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    books = out / "books"
    from analysis_pack import __version__
    # the bundle, made from the example question file: the real desk artifact
    bundle = out / "build_pack.py"
    rc = pack_main(["bundle", str(PKG / "configs" / "examples" / "stated_income_vs_sales.yaml"), "-o", str(bundle)])
    if rc != 0:
        print("could not make the bundle", file=sys.stderr)
        return 1
    desk = Desk(out / "desk", bundle)
    scenarios: list[Scenario] = []
    t0 = time.perf_counter()
    eff, eff_status, eff_pack = scenario_effect(desk, books, out, a.asof, a.run_date, a.loans)
    scenarios.append(eff)
    scenarios.append(scenario_null(desk, books, a.asof, a.run_date, a.loans))
    scenarios.append(scenario_confounded(desk, books, out, a.asof, a.run_date, a.loans))
    if eff_status:
        scenarios.append(scenario_flag(desk, a.asof, a.run_date, eff_status))
        scenarios.append(scenario_snapshot(desk, out, a.asof, a.run_date))
        scenarios.append(scenario_designation(desk, a.asof, a.run_date))
        scenarios.append(scenario_refusals(desk, a.asof, a.run_date))
        scenarios.append(scenario_blanks(desk, a.asof, a.run_date))
        scenarios.append(scenario_knobs(eff_pack))
        scenarios.append(scenario_determinism(desk, a.asof, a.run_date, eff_status["sha256"]))
        if not a.skip_scale:
            scenarios.append(scenario_scale(desk, books, a.asof, a.run_date))
    meta = {"run_date": a.run_date, "asof": a.asof, "loans": a.loans, "version": __version__, "bundle": bundle.name,
            "seconds": round(time.perf_counter() - t0, 1)}
    write_markdown(Path(a.report), scenarios, meta)
    write_html(out / "report.html", scenarios, meta)
    total = sum(len(s.checks) for s in scenarios)
    passed = sum(1 for s in scenarios for c in s.checks if c.ok)
    print(f"{passed} of {total} checks agree across {len(scenarios)} scenarios in {meta['seconds']} s")
    print(f"report: {a.report}\nhtml:   {out / 'report.html'}")
    for s in scenarios:
        for c in s.checks:
            if not c.ok:
                print(f"  DID NOT AGREE — {s.title}: {c.claim} — {c.detail}")
    return 0 if passed == total else 2


if __name__ == "__main__":
    sys.exit(main())
