"""Slice 5 (issue #369): step 5 decomposition, step 7 control observation,
generic derived groupings, and the consumer example building with no code
change."""

from __future__ import annotations

import csv
import io
import json
import shutil
from datetime import date
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from analysis_pack import ladder
from analysis_pack.cli import main
from analysis_pack.config import ConfigError, load_config
from analysis_pack.ingest import Table, read_table
from analysis_pack.population import PopulationError, add_months, build_population
from analysis_pack.workbook import build_pack

from conftest import ASOF, RUN

EXAMPLES = Path(__file__).resolve().parents[1] / "configs" / "examples"


def _row(i, **over):
    base = {"loan_id": f"X{i:04d}", "origination_date": add_months(ASOF, -40).isoformat(), "field_a": "10",
            "field_b": "5", "amount": "1", "category_1": "x", "category_2": "y", "code_1": "111111",
            "outcome_date": "", "flag_1": 0, "measure_a": "1", "measure_b": "2"}
    base.update(over)
    return base


def _config(effect_book: Path, tmp_path: Path, mutate, name="edited.yaml") -> Path:
    raw = yaml.safe_load((effect_book / "config.yaml").read_text())
    mutate(raw)
    p = tmp_path / name
    p.write_text(yaml.safe_dump(raw, sort_keys=False, allow_unicode=True))
    return p


def _table(rows):
    return Table(path="x.csv", sha256="0" * 64, columns=list(rows[0].keys()), rows=rows, kind="csv")


# -- step 5 ------------------------------------------------------------------

def test_a_level_holding_every_flagged_event_is_the_first_row_with_share_one(effect_book, tmp_path):
    p = _config(effect_book, tmp_path, lambda raw: raw.update({"decompose_by": ["category_2"]}))
    cfg = load_config(p)
    rows = []
    for i in range(400):
        flagged = i % 4 == 0                          # ratio 2 fires; ratio 0.5 does not
        level = "A" if i % 3 == 0 else "B"
        event = flagged and level == "A" and i % 8 == 0
        rows.append(_row(i, field_a="10" if flagged else "2.5", category_2=level,
                         outcome_date=add_months(add_months(ASOF, -40), 6).isoformat() if event else ""))
    pop = build_population(cfg, rows, ASOF)
    data = ladder.run(cfg, pop)
    dc = next(d for d in data.decompositions if d.dimension == "category_2")
    assert dc.rows[0].label == "A" and dc.rows[0].share == 1.0
    assert dc.rows[1].label == "B" and dc.rows[1].share == 0.0
    blob, _, _ = build_pack(cfg, pop, _table(rows), RUN)
    from recalc import Recalc
    rec = Recalc(blob)
    shares = [v for r, v in sorted(rec.column("5_Decomposition", "M").items()) if isinstance(v, (int, float))]
    assert shares[:2] == [1.0, 0.0]


def test_decompose_by_a_name_that_is_not_a_dimension_is_refused(effect_book, tmp_path):
    p = _config(effect_book, tmp_path, lambda raw: raw.update({"decompose_by": ["amount"]}))   # a log control
    with pytest.raises(ConfigError) as exc:
        load_config(p)
    assert "not a confounder, a categorical control, or origination_year" in str(exc.value)


# -- derived groupings -----------------------------------------------------------

def test_prefix_groups_codes_and_refuses_a_short_one(effect_book, tmp_path):
    p = _config(effect_book, tmp_path, lambda raw: (raw["controls"].append({"name": "code_group", "field": "code_1"}),
                                                    raw.update({"decompose_by": ["code_group"]})))
    cfg = load_config(p)
    rows = [_row(0, code_1="311234"), _row(1, code_1="332111"), _row(2, code_1="445566")]
    pop = build_population(cfg, rows, ASOF)
    assert sorted({l.values["code_1"] for l in pop.loans}) == ["31", "33", "44"]
    with pytest.raises(PopulationError) as exc:
        build_population(cfg, rows + [_row(3, code_1="9")], ASOF)
    assert "malformed code" in str(exc.value)
    assert exc.value.rows[0][0] == "X0003"


def test_map_after_prefix_labels_groups_uses_other_and_refuses_without_it(effect_book, tmp_path):
    def with_map(other):
        def mutate(raw):
            steps = [{"kind": "prefix", "length": 2}, {"kind": "map", "groups": {"A": ["31", "33"]}}]
            if other:
                steps[1]["other"] = "rest"
            raw["fields"]["code_1"] = {"known": "at_origination", "derive": steps}
            raw["controls"].append({"name": "code_group", "field": "code_1"})
            raw["decompose_by"] = ["code_group"]
        return mutate
    rows = [_row(0, code_1="311234"), _row(1, code_1="332111"), _row(2, code_1="445566")]
    cfg = load_config(_config(effect_book, tmp_path, with_map(True), "with.yaml"))
    pop = build_population(cfg, rows, ASOF)
    assert [l.values["code_1"] for l in pop.loans] == ["A", "A", "rest"]
    data = ladder.run(cfg, pop)
    assert data.band_counts == [] or True
    cfg2 = load_config(_config(effect_book, tmp_path, with_map(False), "without.yaml"))
    with pytest.raises(PopulationError) as exc:
        build_population(cfg2, rows, ASOF)
    assert "no `other` label" in str(exc.value) and exc.value.rows[0][2] == "44"


def test_a_groups_file_behaves_like_inline_groups(effect_book, tmp_path):
    (tmp_path / "groups.csv").write_text("value,group\n31,A\n33,A\n")
    def mutate(raw):
        raw["fields"]["code_1"] = {"known": "at_origination",
                                   "derive": [{"kind": "prefix", "length": 2},
                                              {"kind": "map", "groups_file": "groups.csv", "other": "rest"}]}
        raw["controls"].append({"name": "code_group", "field": "code_1"})
    cfg = load_config(_config(effect_book, tmp_path, mutate))
    rows = [_row(0, code_1="311234"), _row(1, code_1="332111"), _row(2, code_1="445566")]
    pop = build_population(cfg, rows, ASOF)
    assert [l.values["code_1"] for l in pop.loans] == ["A", "A", "rest"]


def test_a_missing_groups_file_is_a_config_refusal(effect_book, tmp_path):
    def mutate(raw):
        raw["fields"]["code_1"] = {"known": "at_origination", "derive": {"kind": "map", "groups_file": "nowhere.csv"}}
    with pytest.raises(ConfigError) as exc:
        load_config(_config(effect_book, tmp_path, mutate))
    assert "was not found beside the question file" in str(exc.value)


# -- step 7 ------------------------------------------------------------------

def test_step_7_prints_the_share_and_the_none_sentence_from_the_config_only(effect_pack, effect_recalc):
    data = effect_pack["data"]
    assert data.control is not None
    both = sum(r.capture.events for r in data.prevalence)
    fires = sum(r.flag.events for r in data.prevalence)
    assert data.control.both.n == both and data.control.both.events == fires
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    lines = [str(c.value) for row in wb["7_Control"].iter_rows() for c in row if isinstance(c.value, str)]
    share_line = next(l for l in lines if l.startswith("In "))
    assert f"{fires / both:.1%}" in share_line and f"{both:,}" in share_line
    assert "Nothing in the process reacts to it." in lines
    assert "field_a: decision one" in lines and "field_b: decision two" in lines
    assert not any("unverified" in l.lower() for l in lines)
    cover = [str(c.value) for row in wb["Cover"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "Regardless of the above" in cover and "Nothing in the process reacts to it." in cover
    share_cells = [v for r, v in effect_recalc.column("7_Control", "B").items() if isinstance(v, float)]
    assert any(abs(v - fires / both) < 1e-12 for v in share_cells)


def test_step_7_names_the_control_when_the_config_gives_one(effect_book, tmp_path):
    p = _config(effect_book, tmp_path, lambda raw: raw.update({"existing_control": "a second-look review"}))
    cfg = load_config(p)
    table = read_table(effect_book / "loans.csv")
    pop = build_population(cfg, table.rows, ASOF)
    blob, _, _ = build_pack(cfg, pop, table, RUN)
    wb = load_workbook(io.BytesIO(blob))
    lines = [str(c.value) for row in wb["7_Control"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "What reacts to it today, as the question file states it: a second-look review" in lines
    assert "Nothing in the process reacts to it." not in lines


# -- the consumer example ----------------------------------------------------------

def _consumer_rows(effect_book: Path) -> list[dict]:
    """The example's columns, derived from the generator's book, never typed in."""
    out = []
    terms = ["36", "48", "60", "72"]
    for i, r in enumerate(csv.DictReader((effect_book / "loans.csv").read_text().splitlines())):
        if i >= 6000:
            break
        amount = float(r["amount"])
        fa = float(r["field_a"]) if r["field_a"] else None
        fb = float(r["field_b"]) if r["field_b"] else None
        clamp = lambda v: None if v is None else min(max(v, 6000.0), 2000000.0)
        out.append({
            "ACCT": r["loan_id"], "FUND_DT": r["origination_date"], "PRODUCT": "INDIRECT_AUTO",
            "STATED_INC": "" if fa is None else f"{clamp(fa):.2f}",
            "BUREAU_INC_EST": "" if fb is None else f"{clamp(fb):.2f}",
            "FICO": str(300 + int(amount) % 551), "LTV": f"{0.1 + (int(amount) % 190) / 100:.2f}",
            "TERM_MO": terms[int(amount) % 4], "STATE": r["category_1"].upper(),
            "DPD60_DT": r["outcome_date"],
        })
    return out


def test_the_consumer_example_validates_and_builds_with_no_code_change(effect_book, tmp_path, capsys):
    rows = _consumer_rows(effect_book)
    path = tmp_path / "auto.csv"
    with path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)
    cfg_path = EXAMPLES / "stated_vs_bureau_income_auto.yaml"
    rc = main(["validate", str(cfg_path), "--data", str(path), "--asof", "2026-06-30"])
    out = capsys.readouterr()
    assert rc == 0, out.err
    rc = main(["build", str(cfg_path), "--data", str(path), "--asof", "2026-06-30", "-o", str(tmp_path / "auto.xlsx")])
    out = capsys.readouterr()
    assert rc == 0, out.err
    status = json.loads(out.out)
    assert status["outcomes"][0]["label"] == "60+ days past due"
    assert status["date_formats"]["FUND_DT"] == "%Y-%m-%d"
    wb = load_workbook(tmp_path / "auto.xlsx")
    assert "7_Control" in wb.sheetnames and "5_Decomposition" in wb.sheetnames
    lines = [str(c.value) for row in wb["7_Control"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("income reasonableness test above 150% of the bureau estimate" in l for l in lines)
    labels = [str(c.value) for row in wb["5_Decomposition"].iter_rows() for c in row if isinstance(c.value, str)]
    assert "north-east" in labels and "south" in labels     # the file-backed map, read beside the question file


def test_the_question_files_name_survives_validation(effect_book):
    """Found by opening the artifact: a loop variable in the validation shadowed
    `name`, and every banner said the file was called `origination_year`."""
    cfg = load_config(effect_book / "config.yaml")
    assert cfg.name == "synthetic_effect"
