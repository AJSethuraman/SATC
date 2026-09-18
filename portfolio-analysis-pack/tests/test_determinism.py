"""Same input, same bytes. A different run date moves only the provenance."""

from __future__ import annotations

import hashlib
import io
import zipfile
from datetime import date

from openpyxl import load_workbook

from analysis_pack.workbook import build_pack

from conftest import RUN


def test_two_builds_from_identical_input_are_byte_identical(effect_pack):
    cfg, pop, table = effect_pack["cfg"], effect_pack["pop"], effect_pack["table"]
    a, _, _ = build_pack(cfg, pop, table, RUN)
    b, _, _ = build_pack(cfg, pop, table, RUN)
    assert hashlib.sha256(a).hexdigest() == hashlib.sha256(b).hexdigest()
    assert a == effect_pack["bytes"]


def test_a_different_run_date_changes_only_provenance_and_header_bands(effect_pack):
    cfg, pop, table = effect_pack["cfg"], effect_pack["pop"], effect_pack["table"]
    other, _, _ = build_pack(cfg, pop, table, date(2026, 9, 19))
    assert other != effect_pack["bytes"]
    wa = load_workbook(io.BytesIO(effect_pack["bytes"]))
    wb = load_workbook(io.BytesIO(other))
    diffs = []
    for name in wa.sheetnames:
        for row_a, row_b in zip(wa[name].iter_rows(), wb[name].iter_rows()):
            for ca, cb in zip(row_a, row_b):
                if ca.value != cb.value:
                    diffs.append((name, ca.coordinate, ca.value, cb.value))
    assert diffs, "the run date did not reach the workbook"
    for name, coord, va, vb in diffs:
        assert "2026-09-18" in str(va) and "2026-09-19" in str(vb), (name, coord, va, vb)
        assert name in ("Cover", "3_Gradient", "_provenance"), (name, coord)
    with zipfile.ZipFile(io.BytesIO(other)) as z:
        core = z.read("docProps/core.xml").decode()
        assert "2026-09-19T00:00:00Z" in core
        assert all(i.date_time == (1980, 1, 1, 0, 0, 0) for i in z.infolist())
