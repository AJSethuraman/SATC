"""Slice 9 (issue #372): the guards the mutation tool relies on that no other
slice pinned, and the tool's own shape."""

from __future__ import annotations

import io
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

TOOLS = Path(__file__).resolve().parents[1] / "tools"


def test_every_check_row_tolerance_is_tight(effect_pack):
    """A check tab whose tolerance is wide enough to swallow any error says
    'agree' about nothing. Counts are exact; rates and ratios to a part in a
    billion (scaled for odds ratios); interval bounds to a millionth."""
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    tols = [row[0].value for row in wb["_check"].iter_rows(min_row=2, min_col=6, max_col=6) if row[0].value is not None]
    assert tols, "no tolerances on the check tab"
    assert max(tols) <= 1e-3, max(tols)
    assert any(t == 0.0 for t in tols)          # counts are compared exactly


def test_the_mutation_tool_lists_nine_mutations_each_naming_tests_that_exist():
    r = subprocess.run([sys.executable, str(TOOLS / "mutation_check.py"), "--list"], capture_output=True, text=True)
    assert r.returncode == 0
    ids = [line.split()[0] for line in r.stdout.splitlines() if line.strip()]
    assert len(ids) == 9 and len(set(ids)) == 9, ids
    sys.path.insert(0, str(TOOLS))
    import mutation_check as mc
    tests_dir = Path(__file__).resolve().parent
    for m in mc.MUTATIONS:
        assert m.path.exists() and m.path.read_text(encoding="utf-8").count(m.old) == 1, m.id
        for t in m.must_fail:
            f, name = t.split("::")
            assert f"def {name}(" in (tests_dir.parent / f).read_text(encoding="utf-8"), t
