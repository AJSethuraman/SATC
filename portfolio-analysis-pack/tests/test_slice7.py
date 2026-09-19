"""Slice 7 (issue #371): one chart per gradient block and one per stratified
block, with interval bars read from twinned helper cells."""

from __future__ import annotations

import io
import re
import zipfile

from openpyxl import load_workbook

from analysis_pack.workbook import build_pack

from conftest import RUN

REF = re.compile(r"<(?:c:)?f>([^<]+)</(?:c:)?f>")


def _charts(blob: bytes) -> list[str]:
    z = zipfile.ZipFile(io.BytesIO(blob))
    names = sorted(n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
    return [z.read(n).decode("utf-8") for n in names]


def test_one_chart_per_gradient_block_and_per_stratified_block(effect_pack):
    data = effect_pack["data"]
    charts = _charts(effect_pack["bytes"])
    assert len(charts) == len(data.per_outcome) + len(data.strata), len(charts)
    on_gradient = [c for c in charts if "'3_Gradient'!" in c]
    on_strat = [c for c in charts if "'4_Stratified'!" in c]
    assert len(on_gradient) == len(data.per_outcome) and len(on_strat) == len(data.strata)


def test_each_charts_series_and_error_bars_point_at_its_own_blocks_cells(effect_pack):
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    grad = wb["3_Gradient"]
    charts = _charts(effect_pack["bytes"])
    g = [c for c in charts if "'3_Gradient'!" in c][0]
    refs = REF.findall(g)
    values = [r for r in refs if re.search(r"'3_Gradient'!\$?D\$?\d+:\$?D\$?\d+", r)]
    assert len(values) == 1, refs
    first, last = (int(x) for x in re.findall(r"\d+", values[0].split("!")[1]))
    # the rows the chart plots are the bucket rows: their labels are the bucket labels
    labels = [grad.cell(r, 1).value for r in range(first, last + 1)]
    assert labels == effect_pack["pop"].bucket_labels, labels
    plus = [r for r in refs if "$K$" in r]
    minus = [r for r in refs if "$L$" in r]
    assert plus == [f"'3_Gradient'!$K${first}:$K${last}"] and minus == [f"'3_Gradient'!$L${first}:$L${last}"]
    assert g.count("errValType") == 1 and 'val="cust"' in g


def test_the_stratified_chart_carries_two_series_with_their_own_error_bars(effect_pack):
    charts = [c for c in _charts(effect_pack["bytes"]) if "'4_Stratified'!" in c]
    for c in charts:
        assert len(re.findall(r"<(?:c:)?ser>", c)) == 2
        assert len(re.findall(r"<(?:c:)?errBars>", c)) == 2
        refs = REF.findall(c)
        assert any("$Q$" in r for r in refs) and any("$R$" in r for r in refs)
        assert any("$S$" in r for r in refs) and any("$T$" in r for r in refs)


def test_the_helper_cells_behind_the_error_bars_are_twinned_and_agree(effect_recalc, effect_pack):
    helpers = [c for c in effect_pack["checks"]
               if (c.sheet == "3_Gradient" and c.cell[0] in "KL") or (c.sheet == "4_Stratified" and c.cell[0] in "QRST")]
    assert helpers, "no helper cells were twinned"
    verdicts = {r: v for r, v in effect_recalc.column("_check", "G").items() if r >= 2}
    assert all(v == "OK" for v in verdicts.values())
    # a helper equals upper minus rate: check one against the twins directly
    g = effect_pack["data"].per_outcome[0][1]
    ks = [c for c in helpers if c.sheet == "3_Gradient" and c.cell.startswith("K")]
    assert any(abs(c.python - (gr.hi - gr.rate)) < 1e-12 for c in ks for gr in g.rows if gr.hi is not None)


def test_charts_keep_the_build_deterministic(effect_pack):
    cfg, pop, table = effect_pack["cfg"], effect_pack["pop"], effect_pack["table"]
    again, _, _ = build_pack(cfg, pop, table, RUN)
    assert again == effect_pack["bytes"]


def test_every_tab_prints_one_page_wide(effect_pack):
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    for ws in wb.worksheets:
        assert ws.sheet_properties.pageSetUpPr.fitToPage is True, ws.title
        assert ws.page_setup.fitToWidth == 1 and ws.page_setup.fitToHeight == 0, ws.title
