"""The eleven defects the desk walk of 22 Sep 2026 found on the screens
(docs/WALKTHROUGH-DEFECTS.md), each pinned by a test so it stays fixed. The
suite was 99 tests, 9 of 9 mutations and 48 of 48 harness checks, and caught
none of them: every one was a thing a screen said. Defect 11 (the skeleton
quotes real values) is not patched — the firm declined a PII guard on 18 Sep.
"""

from __future__ import annotations

import io
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from datetime import date
from pathlib import Path

import yaml
from openpyxl import load_workbook

from analysis_pack import workbook_style as ks
from analysis_pack.cli import main
from analysis_pack.config import CONFIRM, ConfigError, load_config
from analysis_pack.ingest import read_table
from analysis_pack.population import build_population
from analysis_pack.workbook import build_pack

from conftest import ASOF


def _texts(ws) -> list[str]:
    return [str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


# -- defect 1: a question file with no confounders --------------------------

def _no_confounder_config(effect_book: Path, tmp_path: Path) -> Path:
    raw = yaml.safe_load((effect_book / "config.yaml").read_text(encoding="utf-8"))
    were = {c["name"] for c in raw["confounders"]}
    raw["confounders"] = []
    raw["decompose_by"] = [d for d in raw["decompose_by"] if d not in were]
    raw["name"] = "no_confounders"
    p = tmp_path / "none.yaml"
    p.write_text(yaml.safe_dump(raw, sort_keys=False), encoding="utf-8")
    return p


def test_a_question_file_with_no_confounders_is_told_so_on_every_screen(effect_book, tmp_path, capsys):
    cfg_path = _no_confounder_config(effect_book, tmp_path)
    out = tmp_path / "none.xlsx"
    rc = main(["build", str(cfg_path), "--data", str(effect_book / "loans.csv"), "--asof", ASOF.isoformat(), "-o", str(out)])
    captured = capsys.readouterr()
    assert rc == 0, captured.err
    assert "step 4: nothing to compare within — the question file lists no confounders" in captured.err
    wb = load_workbook(out)
    cover = "\n".join(_texts(wb["Cover"]))
    assert "not built in this version" not in cover
    assert "the question file lists nothing to compare within" in cover
    assert "There is no M2: the question file lists no confounders to add" in cover
    assert "with the confounders added" not in cover
    model = "\n".join(_texts(wb["6_Model"]))
    assert "M2: M1 again — the question file lists no confounders to add" in model
    # validate says it too, before anything is built
    rc = main(["validate", str(cfg_path), "--data", str(effect_book / "loans.csv"), "--asof", ASOF.isoformat()])
    captured = capsys.readouterr()
    assert rc == 0 and "note: no confounders listed: step 4 will have nothing to compare within" in captured.err
    assert json.loads(captured.out)["note"].startswith("no confounders listed")


def test_the_skeleton_marks_confounders_as_a_value_to_fill_in(effect_book, tmp_path, capsys):
    out = tmp_path / "q.yaml"
    assert main(["init", str(effect_book / "loans.csv"), "-o", str(out)]) == 0
    capsys.readouterr()
    raw = yaml.safe_load(out.read_text(encoding="utf-8"))
    assert isinstance(raw["confounders"], str) and raw["confounders"].startswith(CONFIRM)
    try:
        load_config(out)
    except ConfigError as exc:
        assert any(p.startswith("`confounders` still reads") for p in exc.problems)
    else:
        raise AssertionError("a skeleton with the confounders slot unfilled was accepted")


# -- defects 2, 3, 4: the screen the bundle shows ---------------------------

def test_the_bundle_shows_a_person_the_summary_and_the_next_command_and_leaves_nothing_beside_itself(effect_book, tmp_path, capsys):
    bundle = tmp_path / "build_pack.py"
    assert main(["bundle", str(effect_book / "config.yaml"), "-o", str(bundle)]) == 0
    capsys.readouterr()
    desk = tmp_path / "desk"
    desk.mkdir()
    shutil.copy(bundle, desk / bundle.name)
    env = dict(os.environ)
    run = lambda *args: subprocess.run([sys.executable, bundle.name, *args], cwd=desk,  # noqa: E731
                                       capture_output=True, text=True, timeout=600, env=env)
    r = run("--synth", "demo", "--loans", "2000")
    assert r.returncode == 0, r.stderr[-800:]
    # defect 2: the JSON status stays off the screen unless asked
    assert r.stdout.strip() == "" and r.stderr.startswith("wrote demo/loans.csv")
    # defect 3: the make-a-book step says what to type next, in this script's spelling
    assert "then: python build_pack.py --data demo/loans.csv --asof 2026-06-30 --config demo/config.yaml -o demo/pack.xlsx" in r.stderr
    # defect 4: nothing appeared beside the emailed file
    assert sorted(p.name for p in desk.iterdir()) == ["build_pack.py", "demo"]
    r = run("--validate", "demo/loans.csv", "--asof", "2026-06-30", "--config", "demo/config.yaml")
    assert r.returncode == 0 and "then: python build_pack.py --data demo/loans.csv --asof 2026-06-30 --config demo/config.yaml -o " in r.stderr
    r = run("--data", "demo/loans.csv", "--asof", "2026-06-30", "--config", "demo/config.yaml", "-o", "demo/pack.xlsx", "--json")
    assert r.returncode == 0, r.stderr[-800:]
    assert json.loads(r.stdout)["ok"] is True, "with --json the status is printed for a program to read"
    assert "then: open demo/pack.xlsx in Excel" in r.stderr
    # defect 7: the summary in plain words
    assert "events per coefficient" in r.stderr and "EPP" not in r.stderr and "no engine" not in r.stderr
    assert "formula check: runs when Excel opens the file" in r.stderr
    # defect 8: no run date given, and the pack says so rather than guess
    assert "run date: not given, so the pack says so" in r.stderr
    assert sorted(p.name for p in desk.iterdir()) == ["build_pack.py", "demo"]
    assert not any(p.name.startswith("analysis_pack_bundle_src") for p in desk.iterdir())


# -- defects 5 and 6: headings a person can read -------------------------------

def test_no_heading_is_cut_by_the_next_column(effect_pack):
    """Every header cell either fits its column or wraps into a row tall
    enough for it; every bold label in column A with a number beside it fits
    column A. Cut headings were the labels the cover pointed at (defect 5)."""
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    cut = []
    for ws in wb.worksheets:
        widths = {}
        merged = {str(r).split(":")[0] for r in ws.merged_cells.ranges}   # banners and notes span columns
        for row in ws.iter_rows():
            for c in row:
                if not isinstance(c.value, str) or not c.value or c.coordinate in merged:
                    continue
                letter = c.column_letter
                widths.setdefault(letter, ws.column_dimensions[letter].width or 8.43)
                is_header = c.fill is not None and c.fill.fgColor is not None and str(c.fill.fgColor.rgb).endswith(ks.INK)
                if is_header:
                    lines = ks.header_lines(c.value, widths[letter])
                    height = ws.row_dimensions[c.row].height or 15
                    if lines > 1 and (not c.alignment.wrap_text or height < 15 * lines):
                        cut.append((ws.title, c.coordinate, c.value))
                elif letter == "A" and c.font is not None and c.font.bold:
                    right = ws.cell(c.row, 2).value
                    if right not in (None, "") and len(c.value) > widths["A"] and not c.alignment.wrap_text:
                        cut.append((ws.title, c.coordinate, c.value))
    assert not cut, cut


def test_working_columns_are_marked_as_not_results(effect_pack):
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    grad = _texts(wb["3_Gradient"])
    assert "chart: bar up" in grad and "Bar up" not in grad
    assert any("Columns headed chart: feed the chart" in t for t in grad)
    strat = _texts(wb["4_Stratified"])
    assert "working: P" in strat and "chart: flagged bar up" in strat
    assert "P" not in strat and "Flagged bar up" not in strat
    assert any("neither is a result" in t for t in strat)
    # the quieter face is what says "not a result" at a glance
    ws = wb["3_Gradient"]
    hdr = next(c for row in ws.iter_rows() for c in row if c.value == "chart: bar up")
    assert hdr.font.italic and not hdr.font.bold


# -- defect 8: the run date is never invented ----------------------------------

def test_without_a_run_date_the_pack_says_so_and_still_builds_the_same_bytes(effect_pack):
    cfg, pop, table = effect_pack["cfg"], effect_pack["pop"], effect_pack["table"]
    a, _, _ = build_pack(cfg, pop, table, None)
    b, _, _ = build_pack(cfg, pop, table, None)
    assert a == b
    wb = load_workbook(io.BytesIO(a))
    cover = "\n".join(_texts(wb["Cover"]))
    assert "run date not given" in cover and "run 2026" not in cover
    for name in ("1_Capture", "6_Model"):
        assert any("run date not given" in t for t in _texts(wb[name])), name
    prov = {ws_row[0].value: ws_row[1].value for ws_row in wb["_provenance"].iter_rows(min_col=1, max_col=2) if ws_row[0].value}
    assert str(prov["Run date"]).startswith("not given")
    with zipfile.ZipFile(io.BytesIO(a)) as z:
        core = z.read("docProps/core.xml").decode()
    assert f"{pop.asof.isoformat()}T00:00:00Z" in core, "the file's timestamp is the as-of date, and the provenance says so"
    with_date, _, _ = build_pack(cfg, pop, table, date(2026, 9, 22))
    assert "run 2026-09-22" in "\n".join(_texts(load_workbook(io.BytesIO(with_date))["Cover"]))


# -- defect 9: the cover reads as English for a noun outcome --------------------

def test_the_cover_reads_as_english_when_the_outcome_is_a_noun(effect_pack):
    cfg = effect_pack["cfg"]
    cover = _texts(load_workbook(io.BytesIO(effect_pack["bytes"]))["Cover"])
    question = next(t for t in cover if t.startswith("Loans where"))
    assert f"is {cfg.outcome.label} more common among them" in question
    assert f"do they {cfg.outcome.label}" not in question
    # the line is a live formula: its prefix is a quoted string joined to the word cells
    line = next(t for t in cover if "Survives or collapses" in t)
    assert f'for {cfg.outcome.label}: "&"{cfg.confounders[0].name}' in line, line
    assert re.search(rf'for {cfg.outcome.label}:"', line) is None


# -- defect 10: the knob note and the outcome line -------------------------------

def test_the_knob_note_and_the_outcome_line_say_what_is_true(effect_pack):
    cfg = effect_pack["cfg"]
    conf = _texts(load_workbook(io.BytesIO(effect_pack["bytes"]))["_config"])
    assert not any("not in this version" in t for t in conf)
    assert any(t.startswith("Step 4's word: the share of the crude effect") for t in conf)
    assert f"{cfg.outcome.label}, from {cfg.outcome.date_field}" in conf
    assert not any(t.endswith("(event_date)") for t in conf)


def test_the_filled_skeleton_still_builds_with_confounders_listed(effect_book, tmp_path, capsys):
    """The walk filled the skeleton and got half an answer; a person who lists
    the confounders the way the skeleton's own comment shows gets step 4."""
    out = tmp_path / "q.yaml"
    assert main(["init", str(effect_book / "loans.csv"), "-o", str(out)]) == 0
    capsys.readouterr()
    raw = yaml.safe_load(out.read_text(encoding="utf-8"))
    ans = yaml.safe_load((effect_book / "config.yaml").read_text(encoding="utf-8"))
    raw["name"] = "filled"
    raw["population"]["loan_id"] = ans["population"]["loan_id"]
    raw["population"]["origination_date"] = ans["population"]["origination_date"]
    raw["fields"] = {col: {"known": "later" if col == ans["outcome"]["date_field"] else "at_origination"}
                     for col in raw["fields"] if col in ans["fields"]}
    raw["rule"]["field_a"], raw["rule"]["field_b"] = ans["rule"]["field_a"], ans["rule"]["field_b"]
    raw["outcome"] = {"label": ans["outcome"]["label"], "date_field": ans["outcome"]["date_field"]}
    raw["confounders"] = [{"name": "size_band", "field": ans["rule"]["field_b"], "edges": [100000, 250000, 1000000]}]
    raw["existing_control"] = "none"
    out.write_text(yaml.safe_dump(raw, sort_keys=False), encoding="utf-8")
    cfg = load_config(out)
    pop = build_population(cfg, read_table(effect_book / "loans.csv").rows, ASOF)
    blob, data, _ = build_pack(cfg, pop, read_table(effect_book / "loans.csv"), None)
    assert data.strata and all(st.word for st in data.strata)
    cover = "\n".join(_texts(load_workbook(io.BytesIO(blob))["Cover"]))
    assert 'for event: "&"size_band' in cover and "nothing to compare within" not in cover
