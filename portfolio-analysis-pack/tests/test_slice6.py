"""Slice 6 (issue #370): two logistic regressions and a tree in plain Python,
events per parameter, and the not-estimable path."""

from __future__ import annotations

import io
import json
import math
from datetime import date

import pytest
import yaml
from openpyxl import load_workbook

from analysis_pack import ladder, model, synth
from analysis_pack.cli import main
from analysis_pack.config import load_config
from analysis_pack.ingest import Table, read_table
from analysis_pack.population import add_months, build_population
from analysis_pack.workbook import build_pack

from conftest import ASOF, RUN


def _fits(pack):
    return pack["data"].models[0]


def test_the_planted_effect_is_recovered_within_its_interval(effect_pack):
    mr = _fits(effect_pack)
    planted = effect_pack["planted"]["true_marginal_flag_odds_ratio"]
    for fit in (mr.m1, mr.m2):
        assert fit.estimable, fit.reason
        f = fit.flag()
        assert f.lo <= planted <= f.hi, (fit.label, f.odds_ratio, f.lo, f.hi, planted)
        assert abs(f.odds_ratio - planted) / planted < 0.25, (fit.label, f.odds_ratio, planted)
    assert mr.m1.epp is not None and mr.m1.epp >= 10 and mr.m1.warning is None


def test_the_null_book_has_flag_intervals_containing_one(tmp_path_factory):
    d = tmp_path_factory.mktemp("null6")
    synth.generate(d, mode="null")
    cfg = load_config(d / "config.yaml")
    table = read_table(d / "loans.csv")
    pop = build_population(cfg, table.rows, ASOF)
    data = ladder.run(cfg, pop)
    for fit in (data.models[0].m1, data.models[0].m2):
        f = fit.flag()
        assert fit.estimable and f.lo <= 1.0 <= f.hi, (fit.label, f.odds_ratio, f.lo, f.hi)


def test_the_confounded_book_loses_the_flag_when_the_confounders_enter(tmp_path_factory):
    d = tmp_path_factory.mktemp("conf6")
    synth.generate(d, mode="confounded")
    cfg = load_config(d / "config.yaml")
    table = read_table(d / "loans.csv")
    pop = build_population(cfg, table.rows, ASOF)
    data = ladder.run(cfg, pop)
    mr = data.models[0]
    f1, f2 = mr.m1.flag(), mr.m2.flag()
    assert mr.m1.estimable and f1.lo > 1.0, (f1.odds_ratio, f1.lo, f1.hi)        # M1 sees an effect
    assert mr.m2.estimable and f2.lo <= 1.0 <= f2.hi, (f2.odds_ratio, f2.lo, f2.hi)   # M2 does not
    assert mr.tree.first_split == "size_band", mr.tree.first_split
    assert mr.tree.leaves and all(leaf.loans >= cfg.model["min_leaf_loans"] for leaf in mr.tree.leaves)


def test_a_perfectly_separating_term_prints_not_estimable_naming_it(effect_book, tmp_path):
    raw = yaml.safe_load((effect_book / "config.yaml").read_text())
    raw["controls"].append({"name": "leaky", "field": "category_2"})
    p = tmp_path / "sep.yaml"
    p.write_text(yaml.safe_dump(raw, sort_keys=False))
    cfg = load_config(p)
    orig = add_months(ASOF, -40)
    rows = []
    for i in range(600):
        event = i % 20 == 0
        rows.append({"loan_id": f"S{i:04d}", "origination_date": orig.isoformat(), "field_a": "10" if i % 2 else "2",
                     "field_b": "5", "amount": str(1000 + i), "category_1": "x",
                     "category_2": "bad" if event else "good",           # separates the outcome perfectly
                     "code_1": "111111", "outcome_date": add_months(orig, 6).isoformat() if event else "",
                     "flag_1": 0, "measure_a": "1", "measure_b": "2"})
    pop = build_population(cfg, rows, ASOF)
    data = ladder.run(cfg, pop)
    mr = data.models[0]
    assert not mr.m1.estimable and mr.m1.reason.startswith("not estimable")
    assert any("leaky" in t for t in mr.m1.implicated), mr.m1.implicated
    table = Table(path="x.csv", sha256="0" * 64, columns=list(rows[0].keys()), rows=rows, kind="csv")
    blob, _, _ = build_pack(cfg, pop, table, RUN)
    wb = load_workbook(io.BytesIO(blob))
    cover = [str(c.value) for row in wb["Cover"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(l.startswith("Model (step 6): unknown") for l in cover), cover
    tab = [str(c.value) for row in wb["6_Model"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("not estimable" in l and "leaky" in l for l in tab)


def test_thin_data_prints_the_events_per_parameter_warning(tmp_path_factory, capsys):
    d = tmp_path_factory.mktemp("thin")
    synth.generate(d, mode="effect", loans=2500)
    rc = main(["build", str(d / "config.yaml"), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
               "-o", str(d / "thin.xlsx")])
    out = capsys.readouterr()
    assert rc == 0
    status = json.loads(out.out)
    m1 = status["models"][0]["m1"]
    assert "epp" in m1 and m1["epp"] < 10 and m1["warning"] and "Thin" in m1["warning"]
    assert "THIN" in out.err
    wb = load_workbook(d / "thin.xlsx")
    tab = [str(c.value) for row in wb["6_Model"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any(l.startswith("Thin:") for l in tab)


def test_the_full_book_does_not_warn_and_the_build_prints_model_time(effect_book, capsys, tmp_path):
    rc = main(["build", str(effect_book / "config.yaml"), "--data", str(effect_book / "loans.csv"), "--asof", "2026-06-30",
               "-o", str(tmp_path / "full.xlsx")])
    out = capsys.readouterr()
    assert rc == 0
    status = json.loads(out.out)
    assert status["models"][0]["m1"]["warning"] is None
    assert status["models"][0]["seconds"] > 0
    assert "models " in out.err and "s · workbook" in out.err


def test_the_solver_agrees_with_a_closed_form_on_a_one_predictor_problem():
    """One binary predictor: the logistic coefficient is the log of the crude
    odds ratio from the 2x2, exactly."""
    x = [[1.0, 1.0]] * 40 + [[1.0, 0.0]] * 160
    y = [1] * 10 + [0] * 30 + [1] * 8 + [0] * 152
    fit = model.fit_logistic(x, y, ["flag"], "M1", 1.959964)
    assert fit.estimable
    expected = math.log((10 / 30) / (8 / 152))
    assert abs(fit.flag().coef - expected) < 1e-8
    se_expected = math.sqrt(1 / 10 + 1 / 30 + 1 / 8 + 1 / 152)
    assert abs(fit.flag().se - se_expected) < 1e-6


def test_model_numbers_are_values_and_the_check_tab_carries_none_of_them(effect_pack):
    assert not any(c.sheet == "6_Model" for c in effect_pack["checks"])
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    formulas = [c for row in wb["6_Model"].iter_rows() for c in row if c.data_type == "f"]
    assert formulas == []
