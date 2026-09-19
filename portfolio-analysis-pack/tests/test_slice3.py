"""Slice 3 (issue #366): step 1 capture by quarter, step 2's two prevalence
series that never share a cell."""

from __future__ import annotations

import csv
import io
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook

from analysis_pack.config import load_config
from analysis_pack.ingest import read_table
from analysis_pack.population import add_months, build_population
from analysis_pack.workbook import build_pack

from conftest import ASOF, RUN


def _formulas(wb, sheet):
    out = {}
    for row in wb[sheet].iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and c.data_type == "f":
                out[c.coordinate] = c.value
    return out


def test_capture_counts_blanks_by_quarter_and_the_2019_gap_shows(effect_recalc, effect_pack):
    data = effect_pack["data"]
    rows = {r.quarter: r for r in data.capture}
    assert "2019Q1" in rows
    q1 = rows["2019Q1"]
    # the generator blanks field_b on about 40% of 2019Q1 loans on top of the base 8%
    assert q1.b_blank_share > 0.35, q1.b_blank_share
    later = rows["2022Q3"]
    assert later.b_blank_share < 0.15
    # and the recalculated tab agrees with the twins
    shares = effect_recalc.column("1_Capture", "G")
    vals = [v for r, v in sorted(shares.items()) if isinstance(v, (int, float))]
    assert any(abs(v - q1.b_blank_share) < 1e-12 for v in vals)


def test_a_quarter_with_no_captured_field_shows_zero_capture_and_a_blank_flag_rate(effect_book, tmp_path):
    cfg = load_config(effect_book / "config.yaml")
    orig = add_months(ASOF, -40)
    rows = []
    for i in range(20):
        rows.append({"loan_id": f"Q{i}", "origination_date": orig.isoformat(), "field_a": "", "field_b": "5",
                     "amount": "1", "category_1": "x", "category_2": "y", "code_1": "1111", "outcome_date": "",
                     "flag_1": 0, "measure_a": "1", "measure_b": "2"})
    other = add_months(ASOF, -30)
    for i in range(20):
        rows.append({"loan_id": f"R{i}", "origination_date": other.isoformat(), "field_a": "10", "field_b": "5",
                     "amount": "1", "category_1": "x", "category_2": "y", "code_1": "1111", "outcome_date": "",
                     "flag_1": 0, "measure_a": "1", "measure_b": "2"})
    pop = build_population(cfg, rows, ASOF)
    from analysis_pack import ladder
    data = ladder.run(cfg, pop)
    empty = [r for r in data.prevalence if r.capture.events == 0][0]
    assert empty.capture_rate == 0.0 and empty.flag_rate is None
    full = [r for r in data.prevalence if r.capture.events == 20][0]
    assert full.capture_rate == 1.0 and full.flag_rate == 1.0
    # and in the workbook the blank flag rate is a blank cell, not a zero
    from analysis_pack.ingest import Table
    table = Table(path="x.csv", sha256="0" * 64, columns=list(rows[0].keys()), rows=rows, kind="csv")
    blob, _, _ = build_pack(cfg, pop, table, RUN)
    from recalc import Recalc
    rec = Recalc(blob)
    h = rec.column("2_Prevalence", "H")
    vals = [v for r, v in sorted(h.items()) if r > 1]
    assert "" in vals or None in vals
    assert 1.0 in vals


def test_capture_and_flag_series_never_share_a_formula(effect_pack):
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    fs = _formulas(wb, "2_Prevalence")
    # capture rate = C/B (both present over seasoned); flag rate = G over the flag row's own n on _cube
    for coord, f in fs.items():
        col = re.match(r"[A-Z]+", coord).group(0)
        if col in ("D", "E", "F"):
            assert re.search(r"\bG\d+\b", f) is None, (coord, f)
        if col in ("H", "I", "J"):
            assert re.search(r"\bB\d+\b", f) is None, (coord, f)
    cube = wb["_cube"]
    blocks = [row[0].value for row in cube.iter_rows(min_row=2, max_col=1) if row[0].value]
    assert any(b.endswith(".capture") for b in blocks) and any(b.endswith(".flag") for b in blocks)


def test_the_23_month_loan_is_in_its_quarters_unseasoned_count(effect_book):
    cfg = load_config(effect_book / "config.yaml")
    orig = add_months(ASOF, -23)
    rows = [{"loan_id": "B", "origination_date": orig.isoformat(), "field_a": "10", "field_b": "5",
             "amount": "1", "category_1": "x", "category_2": "y", "code_1": "1111", "outcome_date": "",
             "flag_1": 0, "measure_a": "1", "measure_b": "2"}]
    pop = build_population(cfg, rows, ASOF)
    from analysis_pack import ladder
    data = ladder.run(cfg, pop)
    assert data.capture[0].unseasoned == 1 and data.capture[0].loans == 1
    assert data.prevalence == []


def test_switching_the_method_moves_the_prevalence_intervals_and_the_checks_still_agree(effect_pack):
    """The check twins are computed at build for the build's method; the sheet
    must switch too, so a Clopper-Pearson build's twins are CP and agree."""
    import yaml
    from analysis_pack import config as C
    cfg = effect_pack["cfg"]
    cp_cfg = C.Config(**{**cfg.__dict__, "method": "Clopper-Pearson"})
    blob, data, checks = build_pack(cp_cfg, effect_pack["pop"], effect_pack["table"], RUN)
    from recalc import Recalc
    rec = Recalc(blob)
    verdicts = [v for r, v in rec.column("_check", "G").items() if r >= 2]
    assert verdicts and all(v == "OK" for v in verdicts)
    lo_w = effect_pack["data"].prevalence[5].capture_lo
    lo_cp = data.prevalence[5].capture_lo
    assert lo_w is not None and lo_cp is not None and lo_w != lo_cp


def test_every_check_row_still_agrees_with_the_two_new_tabs(effect_recalc, effect_pack):
    verdicts = [v for r, v in effect_recalc.column("_check", "G").items() if r >= 2]
    assert len(verdicts) == len(effect_pack["checks"]) and all(v == "OK" for v in verdicts)
    assert any(c.sheet == "1_Capture" for c in effect_pack["checks"])
    assert any(c.sheet == "2_Prevalence" for c in effect_pack["checks"])
