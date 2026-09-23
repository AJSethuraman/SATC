"""The form: designation at the desk in Excel, a dropdown beside each column.

The firm, 22 Sep 2026, offered a pop-up window or a form in Excel: "Actually
excel version is fine." `pack setup EXTRACT` run once writes `question.xlsx`
beside the extract; the person picks in Excel and saves; the same command
read back builds the pack. These tests fill the form the way a person would
(openpyxl in place of a hand) and check what the tool does with it.
"""

from __future__ import annotations

import io
import json
import os
import shutil
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import load_workbook

from analysis_pack import form as F
from analysis_pack.cli import main
from analysis_pack.config import load_config
from analysis_pack.ingest import inspect_columns, read_table
from analysis_pack.picker import Picker

from conftest import ASOF
from test_picker import EFFECT_ANSWERS, Script


def fill(path: Path, roles: dict[str, str], answers: dict[str, object], later: dict[str, str] = {},
         edges: dict[str, object] = {}) -> None:
    """What a person does in Excel: pick from the dropdowns, type the answers, save."""
    wb = load_workbook(path)
    ws, wa = wb[F.COLUMNS_TAB], wb[F.ANSWERS_TAB]
    r = F.HEADER_ROW + 1
    while ws.cell(r, 2).value:
        col = ws.cell(r, 2).value
        ws.cell(r, F.ROLE_COL).value = roles.get(col)
        ws.cell(r, F.LATER_COL).value = later.get(col)
        ws.cell(r, F.EDGES_COL).value = edges.get(col)
        r += 1
    questions = {q: k for k, q, _, _ in F.ANSWERS}
    r = F.HEADER_ROW + 1
    while wa.cell(r, 1).value:
        key = questions[wa.cell(r, 1).value]
        if key in answers:
            wa.cell(r, 2).value = answers[key]
        r += 1
    wb.save(path)


# the same designation the picker's EFFECT_ANSWERS make
ROLES = {"loan_id": "loan number", "origination_date": "origination date", "field_a": "rule top",
         "field_b": "rule bottom + group", "outcome_date": "went bad: date", "amount": "group", "category_2": "group"}
ANSWERS = {"label": "event", "asof": datetime(2026, 6, 30)}
EDGES = {"field_b": "100000, 200000, 400000, 800000, 1600000"}


def _desk(effect_book, tmp_path) -> Path:
    data = tmp_path / "extract.csv"
    shutil.copy(effect_book / "loans.csv", data)
    return data


def test_the_first_run_writes_the_form_and_the_next_builds_from_it(effect_book, tmp_path, capsys):
    data = _desk(effect_book, tmp_path)
    assert main(["setup", str(data)]) == 0
    err = capsys.readouterr().err
    form = tmp_path / "question.xlsx"
    assert form.exists() and not (tmp_path / "question.yaml").exists()
    assert "wrote" in err and "12 columns" in err and "open it in Excel" in err and "run the same command again" in err
    assert f"pack setup {data}" in err
    # run again before picking anything: told so, nothing built
    assert main(["setup", str(data)]) == 0
    err = capsys.readouterr().err
    assert "nothing picked yet" in err and not (tmp_path / "question.yaml").exists()

    fill(form, ROLES, ANSWERS, edges=EDGES)
    rc = main(["setup", str(data)])
    captured = capsys.readouterr()
    assert rc == 0, captured.err[-1500:]
    err = captured.err
    assert "read " in err and "rule              field_a / field_b, fires above 1, steps 0.5, 1, 2, 5" in err
    assert "went bad          outcome_date (the date it happened)" in err
    assert "field_b (edges 100,000, 200,000, 400,000, 800,000, 1,600,000)" in err
    assert "amount (edges" in err and "its quartiles from the data)" in err and "category_2 (level by level)" in err
    assert "as-of 2026-06-30; run date not given" in err
    assert "wrote" in err and "question.yaml from question.xlsx" in err
    assert "building" in err and "step 4 field_b (edges): survives" in err
    status = json.loads(captured.out)
    assert status["ok"] and status["seasoned"] == 29343
    assert sorted(p.name for p in tmp_path.iterdir()) == ["extract.csv", "field_a_vs_field_b.xlsx", "question.xlsx", "question.yaml"]

    # the file it wrote is the file the picker writes from the same answers
    table = read_table(data)
    s = Script(EFFECT_ANSWERS)
    p = Picker(table, inspect_columns(table), ask=s.ask, say=s.say)
    p.run()
    picked = yaml.safe_load(p.to_yaml())
    formed = yaml.safe_load((tmp_path / "question.yaml").read_text(encoding="utf-8"))
    assert formed == picked
    assert load_config(tmp_path / "question.yaml").name == "field_a_vs_field_b"


def test_every_problem_is_listed_at_once_with_its_cell(effect_book, tmp_path, capsys):
    data = _desk(effect_book, tmp_path)
    assert main(["setup", str(data)]) == 0
    capsys.readouterr()
    form = tmp_path / "question.xlsx"
    roles = dict(ROLES)
    del roles["loan_id"]                                    # 1 no loan number
    roles["category_1"] = "rule top"                        # 2 a text column where a number is needed
    roles["field_a"] = "origination date"                   # 3 origination date on two columns
    fill(form, roles, {"asof": ANSWERS["asof"], "buckets": "5, 1"},   # 4 no event word; 5 edges that fall
         later={"field_b": "yes"},                          # 6 a leak: the rule bottom known only later
         edges={"category_2": "1, 2"})                      # 7 band edges on a text column
    rc = main(["setup", str(data)])
    captured = capsys.readouterr()
    assert rc == 2 and not (tmp_path / "question.yaml").exists()
    err = captured.err
    assert "cannot be built from yet" in err
    assert 'no column is marked "loan number"' in err
    assert "Columns!F9: category_1 reads as text" in err and 'needs a number column' in err
    assert '"origination date" is on 2 columns (origination_date in F5, field_a in F6)' in err
    assert "Answers!B7: the word for the event is needed" in err
    assert "Answers!B6: the edges must rise" in err
    assert "Columns!G7: field_b is the rule bottom + group" in err and "leak" in err
    assert "Columns!H10: category_2 reads as text and is taken level by level" in err
    assert "fix these in Excel, save, and run the same command again" in err
    status = json.loads(captured.out)
    assert status["refused"] == "form" and len(status["problems"]) == 7


def test_the_form_reads_what_excel_stores(effect_book, tmp_path, capsys):
    """Excel turns a typed date into a date cell and a typed number into a
    number; a person may type the date as text; the form takes each."""
    data = _desk(effect_book, tmp_path)
    assert main(["setup", str(data)]) == 0
    capsys.readouterr()
    form = tmp_path / "question.xlsx"
    fill(form, ROLES, {"label": "bad", "asof": "2026-06-30", "run_date": datetime(2026, 9, 22), "months": 24.0,
                       "buckets": 2, "fires": "1"}, edges={"field_b": 250000, "amount": "100,000, 200,000; 400,000"})
    table = read_table(data)
    a = F.read_form(form, table, inspect_columns(table))
    assert a["asof"] == ASOF and a["run_date"].isoformat() == "2026-09-22"
    assert a["window_months"] == 24 and a["buckets"] == [2.0] and a["fires_value"] == 1.0
    by = {c["name"]: c for c in a["confounders"]}
    assert by["field_b"]["edges"] == [250000.0]
    assert by["amount"]["edges"] == [100000.0, 200000.0, 400000.0], "thousands commas with a space between are one number each"
    assert a["label"] == "bad"
    # run together they cannot be told from four numbers: refused, with the way out
    fill(form, ROLES, {"label": "bad", "asof": "2026-06-30"}, edges={"amount": "100,000,200,000"})
    try:
        F.read_form(form, table, inspect_columns(table))
        assert False, "run-together thousands were guessed at"
    except F.FormError as exc:
        assert any("needs a space after it" in p for p in exc.problems)


def test_a_form_written_for_another_extract_is_refused(effect_book, tmp_path, capsys):
    data = _desk(effect_book, tmp_path)
    assert main(["setup", str(data)]) == 0
    capsys.readouterr()
    other = tmp_path / "other.csv"
    lines = (effect_book / "loans.csv").read_text(encoding="utf-8").splitlines()
    other.write_text("\n".join([lines[0].replace("amount", "size")] + lines[1:]) + "\n", encoding="utf-8")
    (tmp_path / "question.xlsx").rename(tmp_path / "keep.xlsx")
    shutil.copy(tmp_path / "keep.xlsx", tmp_path / "question.xlsx")
    fill(tmp_path / "question.xlsx", ROLES, ANSWERS, edges=EDGES)
    rc = main(["setup", str(other)])
    err = capsys.readouterr().err
    assert rc == 2 and "written for a different extract" in err and "it lists amount which other.csv does not have" in err


def test_the_flag_and_measure_outcomes_through_the_form(effect_book, tmp_path):
    data = _desk(effect_book, tmp_path)
    table = read_table(data)
    report = inspect_columns(table)
    form = tmp_path / "question.xlsx"
    F.write_form(table, report, form)
    roles = {k: v for k, v in ROLES.items() if k != "outcome_date"}
    roles["flag_1"] = "went bad: flag"
    fill(form, roles, ANSWERS, edges=EDGES)
    a = F.read_form(form, table, report)
    q = F.to_yaml(a, table.columns)
    raw = yaml.safe_load(q)
    assert raw["outcome"] == {"label": "event", "field": "flag_1", "op": "==", "value": 1.0, "basis": "windowed_by_bank"}
    assert raw["fields"]["flag_1"] == {"known": "later"}
    (tmp_path / "flag.yaml").write_text(q, encoding="utf-8")
    assert load_config(tmp_path / "flag.yaml").outcome.form == "flag"

    F.write_form(table, report, form)
    roles = {k: v for k, v in ROLES.items() if k != "outcome_date"}
    roles.update({"measure_a": "went bad: measure top", "measure_b": "went bad: measure bottom"})
    fill(form, roles, {**ANSWERS, "measure_cut": 0.9}, edges=EDGES)
    a = F.read_form(form, table, report)
    raw = yaml.safe_load(F.to_yaml(a, table.columns))
    assert raw["outcome"]["measure"] == {"kind": "ratio", "field_a": "measure_a", "field_b": "measure_b"}
    assert raw["outcome"]["op"] == ">=" and raw["outcome"]["value"] == 0.9 and raw["outcome"]["basis"] == "snapshot_at_asof"
    (tmp_path / "measure.yaml").write_text(F.to_yaml(a, table.columns), encoding="utf-8")
    assert load_config(tmp_path / "measure.yaml").outcome.form == "snapshot"
    # a measure with neither a cut nor edges, or with both, is refused
    fill(form, roles, {**ANSWERS, "measure_cut": None, "measure_edges": None}, edges=EDGES)
    try:
        F.read_form(form, table, report)
        assert False, "a measure with no cut was accepted"
    except F.FormError as exc:
        assert any("one cut" in p for p in exc.problems)


def test_the_form_has_a_dropdown_beside_every_column_and_the_lists_out_of_sight(effect_book, tmp_path):
    data = _desk(effect_book, tmp_path)
    table = read_table(data)
    form = tmp_path / "question.xlsx"
    F.write_form(table, inspect_columns(table), form)
    wb = load_workbook(form)
    ws = wb[F.COLUMNS_TAB]
    n = len(table.columns)
    ranges = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
    assert ranges[f"F{F.HEADER_ROW + 1}:F{F.HEADER_ROW + n}"] == f"=lists!$A$1:$A${len(F.ROLES)}"
    assert ranges[f"G{F.HEADER_ROW + 1}:G{F.HEADER_ROW + n}"] == "=lists!$B$1:$B$2"
    lists = wb[F.LISTS_TAB]
    assert lists.sheet_state == "hidden"
    assert [lists.cell(i, 1).value for i in range(1, len(F.ROLES) + 1)] == list(F.ROLES)
    # every column of the extract has its row, with what it reads as and samples
    names = [ws.cell(F.HEADER_ROW + i, 2).value for i in range(1, n + 1)]
    assert names == table.columns
    assert ws.cell(F.HEADER_ROW + 9, 3).value == "date" and ws.cell(F.HEADER_ROW + 3, 3).value == "decimal"
    # the quartiles of a number column are shown beside its edges cell, from the data
    amount = names.index("amount") + 1
    q = ws.cell(F.HEADER_ROW + amount, F.QUARTILE_COL).value
    assert q and len(q.split(", ")) == 3 and "," not in q.replace(", ", ""), "written the way the edges cell takes them"
    wa = wb[F.ANSWERS_TAB]
    questions = [wa.cell(F.HEADER_ROW + i, 1).value for i in range(1, len(F.ANSWERS) + 1)]
    assert questions == [q for _, q, _, _ in F.ANSWERS]
    assert [str(dv.sqref) for dv in wa.data_validations.dataValidation] == ["B4", "B13"]


def test_the_bundle_writes_the_form_and_builds_from_it(effect_book, tmp_path, capsys):
    bundle = tmp_path / "build_pack.py"
    assert main(["bundle", str(effect_book / "config.yaml"), "-o", str(bundle)]) == 0
    capsys.readouterr()
    desk = tmp_path / "desk"
    desk.mkdir()
    shutil.copy(bundle, desk / bundle.name)
    shutil.copy(effect_book / "loans.csv", desk / "extract.csv")
    run = lambda: subprocess.run([sys.executable, bundle.name, "--setup", "extract.csv"], cwd=desk, capture_output=True,
                                 text=True, timeout=900, env=dict(os.environ))
    r = run()
    assert r.returncode == 0, r.stderr[-2000:]
    assert "wrote question.xlsx" in r.stderr and "python build_pack.py --setup extract.csv" in r.stderr
    assert sorted(p.name for p in desk.iterdir()) == ["build_pack.py", "extract.csv", "question.xlsx"]
    fill(desk / "question.xlsx", ROLES, ANSWERS, edges=EDGES)
    r = run()
    assert r.returncode == 0, r.stderr[-2000:]
    assert "read question.xlsx:" in r.stderr and "wrote question.yaml from question.xlsx" in r.stderr
    assert "then: open field_a_vs_field_b.xlsx in Excel" in r.stderr
    assert r.stdout.strip() == ""
    assert sorted(p.name for p in desk.iterdir()) == ["build_pack.py", "extract.csv", "field_a_vs_field_b.xlsx",
                                                      "question.xlsx", "question.yaml"]
