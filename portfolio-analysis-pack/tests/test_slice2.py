"""Slice 2 (issue #365): inspect, refusals on dirt and on ambiguous dates,
the filter, XLSX input, the three outcome forms and percentage bands, and
the range-check line that says 'not run' rather than passing."""

from __future__ import annotations

import csv
import io
import json
import shutil
from datetime import date, datetime
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook

from analysis_pack.cli import main
from analysis_pack.config import load_config
from analysis_pack.ingest import detect_date_format, inspect_columns, read_table
from analysis_pack.population import AmbiguousDates, build_population
from analysis_pack.workbook import build_pack

from conftest import ASOF, RUN


def _run(argv, capsys):
    rc = main(argv)
    out = capsys.readouterr()
    status = json.loads(out.out) if out.out.strip() else None
    return rc, status, out.err


def _book(effect_book: Path, tmp_path: Path) -> Path:
    d = tmp_path / "book"
    shutil.copytree(effect_book, d)
    return d


def _config_with(d: Path, mutate, name="edited.yaml") -> Path:
    raw = yaml.safe_load((d / "config.yaml").read_text())
    mutate(raw)
    p = d / name
    p.write_text(yaml.safe_dump(raw, sort_keys=False, allow_unicode=True))
    return p


def _rewrite_dates(csv_path: Path, fmt: str, only_day_le_12: bool = False, limit: int | None = None):
    rows = list(csv.DictReader(csv_path.read_text().splitlines()))
    out = []
    for r in rows:
        d = date.fromisoformat(r["origination_date"])
        if only_day_le_12 and d.day > 12:
            d = d.replace(day=d.day % 12 + 1)
        r["origination_date"] = d.strftime(fmt)
        if r["outcome_date"]:
            e = date.fromisoformat(r["outcome_date"])
            if only_day_le_12 and e.day > 12:
                e = e.replace(day=e.day % 12 + 1)
            r["outcome_date"] = e.strftime(fmt)
        out.append(r)
        if limit and len(out) >= limit:
            break
    with csv_path.open("w", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(rows[0].keys()))
        w.writeheader(); w.writerows(out)


# -- inspect -----------------------------------------------------------------

def test_inspect_lists_every_column_with_kind_blank_share_distinct_and_samples(effect_book, capsys):
    rc, status, err = _run(["inspect", str(effect_book / "loans.csv")], capsys)
    assert rc == 0
    cols = {c["column"]: c for c in status["columns"]}
    assert set(cols) == set(read_table(effect_book / "loans.csv").columns)
    assert cols["field_a"]["kind"] == "decimal" and 0.03 < cols["field_a"]["null_share"] < 0.08
    assert cols["category_1"]["kind"] == "text" and cols["category_1"]["distinct"] == 5
    assert cols["origination_date"]["kind"] == "date-like"
    assert cols["origination_date"]["dates"]["resolved"] == "%Y-%m-%d"
    assert len(cols["loan_id"]["samples"]) == 5
    assert "every value fits %Y-%m-%d" in err


# -- dates -------------------------------------------------------------------

def test_one_pattern_fitting_every_value_is_used_and_recorded_without_a_declared_format(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    _rewrite_dates(d / "loans.csv", "%d-%b-%Y", limit=3000)
    p = _config_with(d, lambda raw: raw["population"].pop("date_format"))
    rc, status, err = _run(["build", str(p), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 0, err
    assert status["date_formats"]["origination_date"] == "%d-%b-%Y"
    wb = load_workbook(d / "x.xlsx")
    prov = {r[0].value: r[1].value for r in wb["_provenance"].iter_rows(min_col=1, max_col=2) if r[0].value}
    assert prov["Date format: origination_date"].startswith("%d-%b-%Y, 3,000 of 3,000 parsed")


def test_two_patterns_fitting_every_value_is_refused_with_both_readings_and_the_line_to_add(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    _rewrite_dates(d / "loans.csv", "%m/%d/%Y", only_day_le_12=True, limit=2000)
    p = _config_with(d, lambda raw: raw["population"].pop("date_format"))
    rc, status, err = _run(["build", str(p), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 2 and status["refused"] == "dates"
    assert set(status["candidates"]) >= {"%m/%d/%Y", "%d/%m/%Y"}
    assert len(status["readings"]) >= 2 and status["readings"][0][1] != status["readings"][1][1]
    assert "date_format:" in err and "reads it as" in err
    assert not (d / "x.xlsx").exists()
    # declaring the pattern makes the same file build
    p2 = _config_with(d, lambda raw: raw["population"].update({"date_format": "%m/%d/%Y"}), name="declared.yaml")
    rc, status, err = _run(["build", str(p2), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "y.xlsx")], capsys)
    assert rc == 0 and status["date_formats"]["origination_date"] == "%m/%d/%Y"


def test_no_pattern_fitting_every_value_leaves_the_failures_as_unparseable_dirt(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    lines = (d / "loans.csv").read_text().splitlines()
    header = lines[0].split(","); i = header.index("origination_date")
    bad = lines[1].split(","); bad[i] = "sometime in 2021"
    (d / "loans.csv").write_text("\n".join([lines[0], ",".join(bad)] + lines[2:2001]) + "\n")
    p = _config_with(d, lambda raw: raw["population"].pop("date_format"))
    rc, status, err = _run(["build", str(p), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 2 and status["refused"] == "hygiene"
    assert status["counts"] == {"unparseable date": 1}


def test_typed_xlsx_date_cells_need_no_pattern(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    rows = list(csv.DictReader((d / "loans.csv").read_text().splitlines()))[:1500]
    wb = Workbook(); ws = wb.active; ws.title = "extract"
    cols = list(rows[0].keys()); ws.append(cols)
    for r in rows:
        vals = []
        for c in cols:
            v = r[c]
            if c in ("origination_date", "outcome_date"):
                vals.append(datetime.fromisoformat(v) if v else None)
            elif c in ("field_a", "field_b", "amount", "measure_a", "measure_b"):
                vals.append(float(v) if v else None)
            elif c == "flag_1":
                vals.append(int(v))
            else:
                vals.append(v)
        ws.append(vals)
    wb.save(d / "extract.xlsx")
    p = _config_with(d, lambda raw: raw["population"].pop("date_format"))
    rc, status, err = _run(["build", str(p), "--data", str(d / "extract.xlsx"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 0, err
    assert status["date_formats"] == {"origination_date": "typed date cells", "outcome_date": "typed date cells"}


# -- filter ------------------------------------------------------------------

def test_the_filter_reduces_the_population_and_is_recorded(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    p = _config_with(d, lambda raw: raw["population"].update({"filter": {"category_1": ["north", "south"]}}))
    rc, status, err = _run(["build", str(p), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 0
    wb = load_workbook(d / "x.xlsx")
    prov = {r[0].value: r[1].value for r in wb["_provenance"].iter_rows(min_col=1, max_col=2) if r[0].value}
    assert prov["Rows read"] == 40000 and 12000 < prov["Rows after filter"] < 20000
    notes_text = " ".join(str(c.value) for row in wb["_method"].iter_rows() for c in row if c.value)
    assert "keeping rows where category_1 in ['north', 'south']" in notes_text


# -- outcome forms -----------------------------------------------------------

def test_a_bank_windowed_flag_outcome_builds_and_says_the_pack_did_not_window_it(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    p = _config_with(d, lambda raw: raw.update({"outcome": {"label": "flagged", "field": "flag_1", "op": "==",
                                                             "value": 1, "basis": "windowed_by_bank"}}))
    rc, status, err = _run(["build", str(p), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 0, err
    assert status["outcomes"][0]["label"] == "flagged" and status["outcomes"][0]["events"] > 100
    wb = load_workbook(d / "x.xlsx")
    notes_text = " ".join(str(c.value) for row in wb["_method"].iter_rows() for c in row if c.value)
    assert "The bank applied the window before the extract; the pack did not enforce it" in notes_text


def test_a_snapshot_measure_with_one_cut_builds_with_the_months_on_book_sentence(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    p = _config_with(d, lambda raw: raw.update({"outcome": {"label": "drawn to the line",
                                                             "measure": {"kind": "ratio", "field_a": "measure_a", "field_b": "measure_b"},
                                                             "op": ">=", "value": 0.9, "basis": "snapshot_at_asof"}}))
    rc, status, err = _run(["build", str(p), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 0, err
    wb = load_workbook(d / "x.xlsx")
    notes_text = " ".join(str(c.value) for row in wb["_method"].iter_rows() for c in row if c.value)
    assert "Months on book vary across the base" in notes_text
    assert status["outcomes"][0]["label"] == "drawn to the line"


def test_a_measure_with_edges_gives_one_outcome_per_edge_and_a_band_table(effect_book, tmp_path, capsys):
    from recalc import Recalc
    d = _book(effect_book, tmp_path)
    p = _config_with(d, lambda raw: raw.update({"outcome": {"label": "drawn",
                                                             "measure": {"kind": "ratio", "field_a": "measure_a", "field_b": "measure_b"},
                                                             "edges": [0.5, 0.75, 0.9], "basis": "snapshot_at_asof"}}))
    rc, status, err = _run(["build", str(p), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 0, err
    labels = [o["label"] for o in status["outcomes"]]
    assert labels == ["drawn ≥ 0.5", "drawn ≥ 0.75", "drawn ≥ 0.9"]
    events = [o["events"] for o in status["outcomes"]]
    assert events[0] > events[1] > events[2] > 0
    rec = Recalc((d / "x.xlsx").read_bytes())
    g = rec.column("_check", "G")
    verdicts = [v for r, v in g.items() if r >= 2]
    assert verdicts and all(v == "OK" for v in verdicts)
    assert len(rec.find_text("Cover", "Gradient:")) == 3
    wb = load_workbook(d / "x.xlsx")
    heads = [str(c.value) for row in wb["3_Gradient"].iter_rows() for c in row if c.value and "(share)" in str(c.value)]
    assert heads == ["< 0.5 (share)", "0.5 – 0.75 (share)", "0.75 – 0.9 (share)", "≥ 0.9 (share)"]
    notes_text = " ".join(str(c.value) for row in wb["_method"].iter_rows() for c in row if c.value)
    assert "no single cut was chosen" in notes_text


def test_a_measure_outcome_needs_exactly_one_of_a_cut_or_edges(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    p = _config_with(d, lambda raw: raw.update({"outcome": {"label": "drawn",
                                                             "measure": {"kind": "ratio", "field_a": "measure_a", "field_b": "measure_b"},
                                                             "basis": "snapshot_at_asof"}}))
    rc, status, err = _run(["validate", str(p), "--data", str(d / "loans.csv")], capsys)
    assert rc == 2 and "edges: [0.5, 0.75, 0.9]" in err


# -- ranges ------------------------------------------------------------------

def test_a_plausible_range_refuses_values_outside_it_and_a_field_without_one_says_not_run(effect_book, tmp_path, capsys):
    d = _book(effect_book, tmp_path)
    p = _config_with(d, lambda raw: raw["fields"].update({"amount": {"known": "at_origination", "plausible": [1000, 200000]}}))
    rc, status, err = _run(["build", str(p), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "x.xlsx")], capsys)
    assert rc == 2 and status["refused"] == "hygiene"
    assert set(status["counts"]) == {"outside the plausible range"} and status["counts"]["outside the plausible range"] > 1000
    rc, status, err = _run(["build", str(d / "config.yaml"), "--data", str(d / "loans.csv"), "--asof", "2026-06-30",
                            "-o", str(d / "y.xlsx")], capsys)
    assert rc == 0
    wb = load_workbook(d / "y.xlsx")
    prov = {r[0].value: r[1].value for r in wb["_provenance"].iter_rows(min_col=1, max_col=2) if r[0].value}
    assert prov["amount"] == "range check: not run (no plausible range given)"
