"""Guards from day one: the package knows no domain, reads no clock, and
writes the newer functions with the prefix Excel needs."""

from __future__ import annotations

import io
import re
from pathlib import Path

import yaml
from openpyxl import load_workbook

PKG = Path(__file__).resolve().parents[1] / "src" / "analysis_pack"
EXAMPLES = Path(__file__).resolve().parents[1] / "configs" / "examples"

DOMAIN_WORDS = ["income", "sales", "naics", "charge", "chargeoff", "charge-off", "dti", "fico", "ltv",
                "utilization", "sector", "sba", "small business"]


def _package_sources() -> dict[str, str]:
    return {p.name: p.read_text(encoding="utf-8") for p in PKG.glob("*.py") if p.name != "keybank_style.py"}


def _example_names() -> set[str]:
    names: set[str] = set()
    for p in EXAMPLES.glob("*.yaml"):
        raw = yaml.safe_load(p.read_text(encoding="utf-8"))
        names.update(raw.get("fields", {}).keys())
        for c in raw.get("confounders", []) + raw.get("controls", []):
            if "field" in c:                 # a derived control (origination_year) is the tool's own name
                names.add(c["name"])
    return names


def test_no_domain_word_or_example_column_name_appears_in_package_code():
    names = _example_names()
    assert names, "no example config found to take names from"
    hits = []
    for fname, text in _package_sources().items():
        for w in DOMAIN_WORDS:                       # domain words: any case
            if re.search(r"(?<![A-Za-z0-9_])" + re.escape(w) + r"(?![A-Za-z0-9_])", text.lower()):
                hits.append((fname, w))
        for w in names:                              # column and band names: as written (STATE is not "state")
            if re.search(r"(?<![A-Za-z0-9_])" + re.escape(w) + r"(?![A-Za-z0-9_])", text):
                hits.append((fname, w))
    assert not hits, hits


def test_the_wording_file_carries_no_domain_word_either():
    text = (PKG / "wording.yaml").read_text(encoding="utf-8").lower()
    hits = [w for w in DOMAIN_WORDS if re.search(r"(?<![a-z0-9_])" + re.escape(w) + r"(?![a-z0-9_])", text)]
    assert not hits, hits


def test_the_package_never_reads_the_clock():
    hits = [f for f, t in _package_sources().items() if re.search(r"datetime\.now|date\.today|time\.time\(|utcnow", t)]
    # cli.py measures elapsed seconds with perf_counter, which is a stopwatch, not a clock, and never reaches the workbook
    assert hits == [], hits


def test_every_newer_function_in_a_written_formula_carries_the_xlfn_prefix(effect_pack):
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    bare = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("=") and c.data_type == "f":
                    for fn in ("NORM.S.INV", "BETA.INV"):
                        for m in re.finditer(re.escape(fn), c.value):
                            if c.value[max(0, m.start() - 6):m.start()] != "_xlfn.":
                                bare.append((ws.title, c.coordinate, fn))
    assert not bare, bare


def test_the_check_tab_formula_column_is_text_not_live(effect_pack):
    wb = load_workbook(io.BytesIO(effect_pack["bytes"]))
    ws = wb["_check"]
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        c = row[0]
        assert c.data_type == "s" and str(c.value).startswith("="), (c.coordinate, c.data_type)
