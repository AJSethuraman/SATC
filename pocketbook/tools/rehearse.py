"""Run PocketBook's workbook route headlessly on an extract, answering the workbook from a file of explicit answers.

    python tools/rehearse.py ANSWERS.yaml --extract EXTRACT.csv --memory MEMORY.yaml [--out RESULT.yaml]

The public-data rehearsal (docs/rehearsal-public-data-2026-09.md) drives the product the analyst uses: Set up writes
`<extract> - PocketBook.xlsx` beside the extract, the answers are written into Control and Columns the way a person
would (as tests/test_generic.py does), and Run fills the result tabs. Nothing is picked for the harness (OC-13): every
judgment the workbook asks is in ANSWERS.yaml, and a question the file doesn't answer stops the run before Run,
naming the cell, exactly as the launcher would.

ANSWERS.yaml:
    kind: bleed | new_variable
    step: scout | prespec              # new_variable only
    prespec: FILE                      # new_variable + prespec only
    bands: [COL, ...]                  # cut into bands (bleed)
    segments: [COL, ...]               # segment by (bleed)
    split: COL                         # optional
    test: [COL, ...]                   # new_variable: the inputs tested
    hold: [COL, ...]                   # new_variable: held fixed
    outcome: COL                       # new_variable
    few_values: 12                     # the two Set up limits, stated even when they are the usual ones
    many_values: 50
    meanings: {COL: CODE}              # a meaning code from settings.yaml (booked, category, amount, ...)
    outcome_is: {COL: VALUE}           # "Yes means", only for an outcome that isn't 0/1
    treat: {COL: Real | Missing}       # every odd value found in COL; "*" answers any column not named
    control: {KEY: "the option's label as Control shows it" or {own: VALUE}}
    checked_every_column: Yes

--memory is required so the rehearsal never writes the tool's memory in the home folder. Wall time is measured per
step; peak memory is the largest resident set seen for this process and, separately, for this process plus its
children (the shuffle test's workers), sampled every 0.2 s with psutil. Without psutil it says "not measured".
"""

from __future__ import annotations

import argparse
import os
import sys
import threading
import time
from pathlib import Path

import yaml


class Peak:
    """The largest resident memory seen for this process, and for it with its children, sampled on a thread."""

    def __init__(self, every: float = 0.2):
        try:
            import psutil
        except ImportError:
            self.ok, self.me, self.tree = False, None, None
            return
        self.ok, self.psutil, self.every = True, psutil, every
        self.proc = psutil.Process()
        self.me = self.tree = 0
        self._stop = threading.Event()
        self._t = threading.Thread(target=self._loop, daemon=True)
        self._t.start()

    def _loop(self):
        while not self._stop.is_set():
            try:
                me = self.proc.memory_info().rss
                kids = 0
                for c in self.proc.children(recursive=True):
                    try:
                        kids += c.memory_info().rss
                    except self.psutil.Error:
                        pass
                self.me, self.tree = max(self.me, me), max(self.tree, me + kids)
            except self.psutil.Error:
                pass
            self._stop.wait(self.every)

    def stop(self) -> dict:
        if not self.ok:
            return {"peak_memory": "not measured (psutil is not installed)"}
        self._stop.set()
        self._t.join()
        gb = lambda b: round(b / 2 ** 30, 2)
        return {"peak_rss_gb_main": gb(self.me), "peak_rss_gb_with_workers": gb(self.tree)}


def answer(book_path: Path, ans: dict) -> list[str]:
    """Write the answers into Control and Columns; return what could not be placed."""
    from openpyxl import load_workbook

    from pocketbook import book, control, meanings

    cat = meanings.catalog()
    wb = load_workbook(book_path)
    unplaced = []
    ws = wb[control.SHEET]
    for key, pick in (ans.get("control") or {}).items():
        r = control.row_of(ws, key)
        if r is None:
            unplaced.append(f"Control has no row {key}")
            continue
        if isinstance(pick, dict):
            ws.cell(row=r, column=control.OWN_COL).value = pick["own"]
        else:
            ws.cell(row=r, column=control.CHOOSE_COL).value = pick
    cols = wb["Columns"]
    rows = {r[book.C_NAME - 1].value: r for r in book.table_rows(cols) if r[book.C_NAME - 1].value}
    for col, code in (ans.get("meanings") or {}).items():
        if col not in rows:
            unplaced.append(f"Columns has no row for {col}")
        elif code not in cat:
            unplaced.append(f"{code} is not a meaning (settings.yaml lists them)")
        else:
            rows[col][book.C_MEANS - 1].value = cat[code].label
    for col, v in (ans.get("outcome_is") or {}).items():
        if col in rows:
            rows[col][book.C_IS - 1].value = v
        else:
            unplaced.append(f"Columns has no row for {col}")
    treat = ans.get("treat") or {}
    for r in book.table_rows(cols):
        key = r[book.C_QKEY - 1].value
        if isinstance(key, str) and key.count("|") == 2:
            col = key.split("|")[0]
            if col in treat or "*" in treat:
                r[book.C_TREAT - 1].value = treat.get(col, treat.get("*"))
    if ans.get("checked_every_column") in (True, "Yes"):
        cols[book.CONFIRM_CELL] = "Yes"
    wb.save(book_path)
    return unplaced


def odd_values(book_path: Path) -> list[str]:
    """Every odd-value question on Columns, as 'COL: question -> answer'."""
    from openpyxl import load_workbook

    from pocketbook import book

    out = []
    for r in book.table_rows(load_workbook(book_path, read_only=False)["Columns"]):
        key = r[book.C_QKEY - 1].value
        if isinstance(key, str) and key.count("|") == 2:
            out.append(f"{key} -> {r[book.C_TREAT - 1].value or '(blank: used as recorded)'}")
    return out


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(prog="rehearse")
    p.add_argument("answers", type=Path)
    p.add_argument("--extract", type=Path, required=True)
    p.add_argument("--memory", type=Path, required=True, help="the memory file; never the home folder's")
    p.add_argument("--out", type=Path, help="where to write the timings and the Run's lines (YAML)")
    a = p.parse_args(argv)
    if not a.extract.exists():
        print(f"REFUSED: {a.extract} isn't there. The public extracts live outside git; build one with "
              f"tools/public_extract.py first.", file=sys.stderr)
        return 2
    ans = yaml.safe_load(a.answers.read_text(encoding="utf-8"))
    kind = ans.get("kind")
    if kind not in ("bleed", "new_variable"):
        print("REFUSED: the answers must say kind: bleed or kind: new_variable", file=sys.stderr)
        return 2
    for k in ("few_values", "many_values"):
        if not isinstance(ans.get(k), int):
            print(f"REFUSED: {k} is asked in the launcher and is not defaulted here; state it", file=sys.stderr)
            return 2
    os.environ["POCKETBOOK_MEMORY"] = str(a.memory)       # belt and braces: every call is also given memory_path

    from pocketbook import book, choices as ch

    tup = lambda k: tuple(ans[k]) if ans.get(k) is not None else None
    if ans.get("prespec"):
        ans["prespec"] = str(Path(ans["prespec"]).resolve())
    c = ch.Choices(run_kind=kind, bands=tup("bands"), segments=tup("segments"), split=ans.get("split"),
                   outcome=ans.get("outcome"), test=tuple(ans.get("test") or ()), hold=tuple(ans.get("hold") or ()),
                   shortlist=ans.get("prespec"), few_values=ans["few_values"], many_values=ans["many_values"])
    result = {"extract": str(a.extract), "answers": str(a.answers), "memory": str(a.memory)}
    peak = Peak()
    t0 = time.perf_counter()
    out = book.set_up(a.extract, memory_path=a.memory, choices=c)
    t1 = time.perf_counter()
    result["set_up_seconds"] = round(t1 - t0, 1)
    result["set_up_lines"] = out.lines
    print(f"Set up: {t1 - t0:.1f} s")
    print("\n".join(f"  {x}" for x in out.lines))
    if not out.ok:
        result.update(peak.stop())
        _write(a.out, result)
        return 3
    if kind == "new_variable":
        from pocketbook import control
        step = ans.get("step")
        if step not in ("scout", "prespec"):
            print("REFUSED: a new_variable run must say step: scout or step: prespec", file=sys.stderr)
            return 2
        ans.setdefault("control", {})
        labels = {o.value: o.label for s in control.load_settings() if s.key == "new_variable_step" for o in s.options}
        ans["control"].setdefault("new_variable_step", labels[step])
    unplaced = answer(out.book, ans)
    if unplaced:
        print("Couldn't place these answers:\n" + "\n".join(f"  - {u}" for u in unplaced))
    result["odd_values"] = odd_values(out.book)
    _, left, _ = book.read_book(out.book, a.memory)
    result["answers_still_needed"] = left
    if left:
        print("Still needs an answer (nothing is picked for you):\n" + "\n".join(f"  - {x}" for x in left))
        result.update(peak.stop())
        _write(a.out, result)
        return 3
    t2 = time.perf_counter()
    ran = book.run(out.book, a.extract, memory_path=a.memory)
    t3 = time.perf_counter()
    result.update(peak.stop())
    result.update({"run_ok": ran.ok, "run_seconds": round(t3 - t2, 1), "run_lines": ran.lines,
                   "run_problems": list(ran.problems or []), "workbook": str(out.book)})
    print(f"Run: {t3 - t2:.1f} s, ok={ran.ok}")
    print("\n".join(f"  {x}" for x in ran.lines))
    print({k: v for k, v in result.items() if k.startswith("peak")})
    _write(a.out, result)
    return 0 if ran.ok else 1


def _write(path: Path | None, result: dict) -> None:
    if path:
        path.write_text(yaml.safe_dump(result, sort_keys=False, allow_unicode=True, width=120), encoding="utf-8")


if __name__ == "__main__":
    sys.exit(main())
