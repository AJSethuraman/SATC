"""Road 1: build the synthetic book, Run it, and read every Summary cell PocketBook shows, for every band column and
every value of the Filter by column, each view set in the dropdowns and calculated by LibreOffice.

    cd pocketbook && python3 docs/tie-out/2026-09-30-summary/build.py

Writes, beside this file: "Summary book.csv" (the loans), "Summary book - PocketBook.xlsx" (the workbook as the Run
left it) and shown.csv (band column, filter, row, heading, value), one line per cell read."""

import csv
import os
import shutil
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path[:0] = [str(ROOT / "src"), str(ROOT / "tests")]

from openpyxl import load_workbook  # noqa: E402

from pocketbook import book, choices as ch, perm, results, synth  # noqa: E402
from recalc import recalc  # noqa: E402
from test_book import _answer  # noqa: E402

N, SEED = 8000, 7
BANDS, SEGMENTS = ("FICO", "REV_DEBT"), ("CHANNEL", "ASSET_CLASS")


def main() -> None:
    work = Path(tempfile.mkdtemp(prefix="summary-tie-out-"))
    os.environ["POCKETBOOK_MEMORY"] = str(work / "memory.yaml")
    perm.SHUFFLES = 200                     # Summary tests nothing; the shuffles only feed other tabs
    x = synth.write_extract(work, n=N, seed=SEED)
    csv_out = HERE / "Summary book.csv"
    shutil.copy(x, csv_out)
    out = book.set_up(csv_out, choices=ch.Choices(run_kind=ch.BLEED, bands=BANDS, segments=SEGMENTS,
                                                  filter=ch.ORIG_YEAR, outcome="BAD_FLAG"))
    _answer(out.book)                       # the tests' answers: FICO's -9999 answered Missing on Columns
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    wb = load_workbook(out.book)
    ws = wb[results.SUMMARY]
    at = {}
    for row in ws.iter_rows(max_row=60):
        for c in row:
            if c.value in ("BAND COLUMN", f"ONLY LOANS WHERE {ch.ORIG_YEAR} IS"):
                at[c.value] = ws.cell(row=c.row + 1, column=c.column).coordinate
    years = [v for v in _options(wb, at[f"ONLY LOANS WHERE {ch.ORIG_YEAR} IS"]) if v != results.ALL_LOANS]
    lines = []
    for band in BANDS:
        for only in [results.ALL_LOANS] + years:
            view = work / f"{band}-{only}.xlsx"
            wb = load_workbook(out.book)
            wb[results.SUMMARY][at["BAND COLUMN"]] = band
            wb[results.SUMMARY][at[f"ONLY LOANS WHERE {ch.ORIG_YEAR} IS"]] = only
            wb.save(view)
            calc = recalc(view, work / f"rc-{band}-{only}")[results.SUMMARY]
            h = next(r for r in range(1, calc.max_row + 1) if calc.cell(row=r, column=3).value == "Loans")
            assert calc.cell(row=h, column=2).value == band, "the dropdown didn't switch the column"
            heads = []
            c = 3
            while calc.cell(row=h, column=c).value not in (None, ""):
                heads.append(calc.cell(row=h, column=c).value)
                c += 1
            r = h + 1
            while calc.cell(row=r, column=2).value not in (None, ""):
                label = calc.cell(row=r, column=2).value
                for j, head in enumerate(heads):
                    v = calc.cell(row=r, column=3 + j).value
                    lines.append([band, only, label, head, "" if v is None else repr(v)])
                r += 1
    with open(HERE / "shown.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["band column", "only loans where", "row", "heading", "shown"])
        w.writerows(lines)
    print(f"{len(lines):,} cells read from {len(BANDS) * (1 + len(years))} views")


def _options(wb, coord: str) -> list:
    from openpyxl.utils import range_boundaries
    ws = wb[results.SUMMARY]
    for dv in ws.data_validations.dataValidation:
        if coord in str(dv.sqref):
            sheet, rng = dv.formula1.lstrip("=").split("!")
            c0, r0, _, r1 = range_boundaries(rng.replace("$", ""))
            src = wb[sheet.strip("'")]
            return [src.cell(row=r, column=c0).value for r in range(r0, r1 + 1)]
    raise KeyError(coord)


if __name__ == "__main__":
    main()
