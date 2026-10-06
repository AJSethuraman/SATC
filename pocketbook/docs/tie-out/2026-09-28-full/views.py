"""Step 1 of the PocketBook road: every view of the workbook an analyst can pick, calculated.

    python3 views.py "Consumer book Q3 - PocketBook.xlsx" OUT

The result tabs show one thing at a time through dropdown cells: Pockets (MEASURE C18, POCKETS D18, SHOW E18), Grids
(GRID B13, MEASURE F13), Split (GRID B15, MEASURE B26) and Paid, cost, kept (GRID B16). This makes one copy of the
workbook per view -- 40 copies, since Grids has 8 grids x 5 measures and the other tabs' choices ride along in the
same copies -- sets the dropdowns the way an analyst would pick them, and has LibreOffice calculate every copy (the
repository's own helper, pocketbook/tests/recalc.py: "recalculate Excel files on load: always"). SHOW is set to All,
which lists every pocket the others list. The options are read from the workbook's own lists (_choices), not typed
here. Nothing here imports PocketBook.
"""
import json
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2] / "tests"))          # pocketbook/tests, for recalc.py
import recalc  # noqa: E402


def options(wb):
    ws = wb["_choices"]
    heads = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}

    def col(name):
        c = heads[name]
        return [ws.cell(r, c).value for r in range(2, ws.max_row + 1) if ws.cell(r, c).value not in (None, "")]
    return {"pockets_measure": col("Pockets: Measure"), "pockets_kind": col("Pockets: Pockets"),
            "grids_grid": col("Grids: Grid"), "grids_measure": col("Grids: Measure"),
            "split_grid": col("Split: Grid"), "split_measure": col("Split: Measure"),
            "pck_grid": col("Paid, cost, kept: Grid")}


def main():
    book, out = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
    (out / "set").mkdir(parents=True, exist_ok=True)
    wb = load_workbook(book)
    opt = options(wb)
    pockets = [(m, k) for k in opt["pockets_kind"] for m in opt["pockets_measure"]]
    grids = [(g, m) for g in opt["grids_grid"] for m in opt["grids_measure"]]
    split = [(g, m) for g in opt["split_grid"] for m in opt["split_measure"]]
    pck = opt["pck_grid"]
    n = max(len(pockets), len(grids), len(split), len(pck))
    plan = []
    for i in range(n):
        v = {"pockets": pockets[i % len(pockets)], "grids": grids[i % len(grids)], "split": split[i % len(split)],
             "pck": pck[i % len(pck)]}
        wb = load_workbook(book)
        wb["Pockets"]["C18"], wb["Pockets"]["D18"], wb["Pockets"]["E18"] = v["pockets"][0], v["pockets"][1], "All"
        wb["Grids"]["B13"], wb["Grids"]["F13"] = v["grids"]
        wb["Split"]["B15"], wb["Split"]["B26"] = v["split"]
        wb["Paid, cost, kept"]["B16"] = v["pck"]
        name = f"view-{i:02d}.xlsx"
        wb.save(out / "set" / name)
        plan.append({"file": name, **v})
    (out / "plan.json").write_text(json.dumps(plan, indent=1))
    # one LibreOffice call for every copy: the same headless conversion recalc.recalc_file makes, many files at once
    files = sorted(str(p) for p in (out / "set").glob("view-*.xlsx"))
    got = subprocess.run([recalc.SOFFICE, f"-env:UserInstallation={recalc._profile_dir().as_uri()}", "--headless",
                          "--calc", "--convert-to", "xlsx", "--outdir", str(out / "calculated"), *files],
                         capture_output=True, text=True, timeout=3000)
    done = sorted((out / "calculated").glob("view-*.xlsx"))
    print(f"{len(done)} of {len(files)} views calculated", got.returncode)
    shutil.copy(book, out / "as-written.xlsx")


if __name__ == "__main__":
    main()
