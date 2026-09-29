"""Step 2 of the PocketBook road: read every figure out of the calculated views of the Where-the-book-bleeds
workbook, cell by cell, with openpyxl. Nothing here imports PocketBook or works anything out: it reads what the
analyst would read, and names each figure by the labels printed beside it on the same tab.

    python3 read_bleed.py VIEWS_DIR  > (writes VIEWS_DIR/figures.json)

VIEWS_DIR is what views.py made. Tabs whose content does not depend on a dropdown (Start here, Control, Columns,
Look, Record) are read once, from the first view; Pockets, Grids, Split and Paid, cost, kept once per view of their
own. The chart on Look draws its bars from a hidden range the workbook keeps for it (_look), and those bar heights are
read there, since the chart itself holds no numbers.

Every cell on a visible result tab that holds a number, or text with a digit in it, is recorded: either as a named
figure (its key says which pocket, measure and column), or, when it is words that happen to hold a digit (a heading's
date, a cell name in an instruction), as `words` -- so the roster can account for every one.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

DIGIT = re.compile(r"\d")


def has_digit(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) or (isinstance(v, str) and DIGIT.search(v))


class Figures:
    def __init__(self):
        self.out = []
        self.seen = set()

    def add(self, tab, view, cell, key, value, kind="figure"):
        self.out.append({"tab": tab, "view": view, "cell": cell, "key": key, "value": value, "kind": kind})
        self.seen.add((tab, view, cell))

    def sweep(self, ws, tab, view, max_col=60, skip_cols=()):
        """Every digit-bearing cell the analyst can see on the tab (hidden rows and columns are the workbook's own
        workings, and are skipped) and not already named: recorded as words, for the roster's count."""
        hid_c, hid_r = hidden(ws)
        for row in ws.iter_rows(max_col=max_col):
            for c in row:
                if c.value is None or c.column in hid_c or c.row in hid_r or (tab, view, c.coordinate) in self.seen:
                    continue
                if c.value is not None and has_digit(c.value):
                    self.add(tab, view, c.coordinate, ["words"], c.value, kind="words")


def hidden(ws):
    """The column numbers and row numbers the sheet hides."""
    from openpyxl.utils import column_index_from_string
    cols, rows = set(), set()
    for k, d in ws.column_dimensions.items():
        if d.hidden:
            lo = d.min or column_index_from_string(k)
            hi = d.max or lo
            cols.update(range(lo, hi + 1))
    for k, d in ws.row_dimensions.items():
        if d.hidden:
            rows.add(k)
    return cols, rows


def find(ws, text, max_row=200, max_col=40, min_row=12):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col):
        for c in row:
            if c.value == text:
                return c.row, c.column
    return None


def header(ws, r, c0):
    """Labels along row r from column c0 until a blank."""
    out, c = [], c0
    while ws.cell(r, c).value not in (None, ""):
        out.append((c, str(ws.cell(r, c).value)))
        c += 1
    return out


# --------------------------------------------------------------------------- Pockets

def read_lines(F, ws, tab, view, row, cols, measure):
    """The lines in use, live from Control, shown above each result tab's rows."""
    for c, what in cols:
        v = ws[f"{c}{row}"].value
        if v is not None:
            F.add(tab, view, f"{c}{row}", ["line", what, measure if what == "material" else "*"], v)


def read_pockets(F, ws, view):
    measure, kind = ws["C18"].value, ws["D18"].value
    read_lines(F, ws, "Pockets", view, 15, (("D", "worse"), ("E", "better"), ("F", "sure"), ("H", "material")),
               measure)
    F.add("Pockets", view, "F18", ["pockets-count", kind, measure], ws["F18"].value)
    heads = {str(ws.cell(20, c).value): c for c in range(2, 20) if ws.cell(20, c).value}
    r = 21
    while ws.cell(r, 3).value not in (None, ""):
        band, seg = ws.cell(r, 3).value, ws.cell(r, 4).value
        half = ws.cell(r, heads["REV_DEBT half"]).value if "REV_DEBT half" in heads else None
        who = [kind, measure, band, seg, half]
        for name, c in heads.items():
            if name in ("Band", "Segment", "REV_DEBT half"):
                continue
            v = ws.cell(r, c).value
            if v is None:
                continue
            field = {"#": "rank", "Loans": "loans", "This pocket": "rate", "Worse?": "worse",
                     "p-value": "p", "Material?": "material", "Could have caught": "caught",
                     "Holds FICO fixed?": "holds"}.get(name)
            if field is None:
                field = "rest-rate" if name.startswith("Rest of") else "gap" if name.startswith(("×", "Gap")) \
                    else "dollars"
            F.add("Pockets", view, ws.cell(r, c).coordinate, ["pocket", *who, field], v)
        r += 1
    F.sweep(ws, "Pockets", view, max_col=16)


# --------------------------------------------------------------------------- Grids

def grid_block(F, ws, view, grid, measure, title_rc, block):
    r0, c0 = title_rc
    heads = header(ws, r0 + 1, c0 + 1)
    r = r0 + 2
    while ws.cell(r, c0).value not in (None, ""):
        rowlab = str(ws.cell(r, c0).value)
        for c, collab in heads:
            v = ws.cell(r, c).value
            if v is not None:
                F.add("Grids", view, ws.cell(r, c).coordinate, ["grid", grid, measure, block, rowlab, collab], v)
        if rowlab == "All":
            break
        r += 1


def read_grids(F, ws, view):
    grid, measure = ws["B13"].value, ws["F13"].value
    rate = find(ws, f"Rate · {measure}")
    grid_block(F, ws, view, grid, measure, rate, "rate")
    book = find(ws, "vs the book")
    grid_block(F, ws, view, grid, measure, book, "book")
    band = find(ws, "vs rest of band")
    grid_block(F, ws, view, grid, measure, band, "band")
    loans = (band[0], book[1])            # "Loans" sits beside "vs rest of band", under "vs the book"
    assert ws.cell(*loans).value == "Loans"
    grid_block(F, ws, view, grid, measure, loans, "loans")
    # how common each group is (two-way grids): pocket rows under a two-line header
    g = find(ws, "How common each group is: a count, not a test")
    if g:                                     # the same block under every measure: counted once, by its key
        hr = g[0] + 3
        if ws.cell(hr, 2).value not in (None, "") and ws.cell(hr, 4).value == "Loans":
            groups = [(c, str(ws.cell(hr, c).value)) for c in range(6, 16) if ws.cell(hr, c).value]
            r = hr + 2
            while ws.cell(r, 2).value not in (None, ""):
                band_, seg = str(ws.cell(r, 2).value), str(ws.cell(r, 3).value)
                cols = {4: "loans", 5: "booked"}
                for c, lab in groups:
                    cols[c], cols[c + 1] = f"{lab}|loans", f"{lab}|booked"
                for c, what in cols.items():
                    v = ws.cell(r, c).value
                    if v is not None:
                        F.add("Grids", view, ws.cell(r, c).coordinate, ["groups", grid, band_, seg, what], v)
                r += 1
    F.sweep(ws, "Grids", view, max_col=23)


# --------------------------------------------------------------------------- Split

def read_split(F, ws, view, first_of_grid):
    grid, measure = ws["B15"].value, ws["B26"].value
    read_lines(F, ws, "Split", view, 12, (("C", "worse"), ("D", "better"), ("E", "sure"), ("F", "material")),
               "Charge-offs")
    if first_of_grid:
        F.add("Split", view, "E15", ["split-chip", grid], ws["E15"].value)
        heads = {c: str(ws.cell(17, c).value) for c in range(3, 10)}
        for r in range(18, 23):
            m = ws.cell(r, 2).value
            for c, h in heads.items():
                v = ws.cell(r, c).value
                if v is not None:
                    F.add("Split", view, ws.cell(r, c).coordinate, ["split-sum", grid, m, h.split(" (")[0]], v)
        F.add("Split", view, "B23", ["split-steady", grid], ws["B23"].value)
    cols = header(ws, 29, 3)
    pcols = header(ws, 29, 9)
    r = 30
    while ws.cell(r, 2).value not in (None, ""):
        band = str(ws.cell(r, 2).value)
        for c, seg in cols:
            v = ws.cell(r, c).value
            if v is not None:
                F.add("Split", view, ws.cell(r, c).coordinate, ["split-pocket", grid, measure, band, seg, "gap"], v)
        for c, seg in pcols:
            v = ws.cell(r, c).value
            if v is not None:
                F.add("Split", view, ws.cell(r, c).coordinate, ["split-pocket", grid, measure, band, seg, "p"], v)
        r += 1
    if first_of_grid:
        F.sweep(ws, "Split", view, max_col=12)
    else:
        for r in range(28, 40):
            for c in range(2, 13):
                cell = ws.cell(r, c)
                if (("Split", view, cell.coordinate) not in F.seen and cell.value is not None
                        and has_digit(cell.value)):
                    F.add("Split", view, cell.coordinate, ["words"], cell.value, kind="words")


# --------------------------------------------------------------------------- Paid, cost, kept

def read_pck(F, ws, view):
    grid = ws["B16"].value
    read_lines(F, ws, "Paid, cost, kept", view, 13, (("C", "worse"), ("D", "better"), ("F", "sure"),
                                                      ("H", "material")), "Charge-offs")
    names = {4: "loans", 5: "earned-gap", 6: "earned-dollars", 7: "gco-multiple", 8: "gco-dollars",
             9: "kept-gap", 10: "kept-dollars", 11: "together"}
    r = 20
    while ws.cell(r, 2).value not in (None, ""):
        band, seg = str(ws.cell(r, 2).value), str(ws.cell(r, 3).value)
        for c, what in names.items():
            v = ws.cell(r, c).value
            if v is not None and v != "":
                F.add("Paid, cost, kept", view, ws.cell(r, c).coordinate, ["pck", grid, band, seg, what], v)
        r += 1
    F.sweep(ws, "Paid, cost, kept", view, max_col=12)


# --------------------------------------------------------------------------- the tabs read once

def read_once(F, wb, view):
    ws = wb["Start here"]
    for c in ("B17", "D17", "F17"):
        F.add("Start here", view, c, ["start-tile", ws.cell(16, ws[c].column).value], ws[c].value)
    r = 20
    while ws.cell(r, 2).value not in (None, "", "No pocket is worse and material."):
        for col, what in ((4, "loans"), (5, "multiple"), (6, "dollars")):
            F.add("Start here", view, ws.cell(r, col).coordinate,
                  ["start-top", r - 19, ws.cell(r, 2).value, ws.cell(r, 3).value, what], ws.cell(r, col).value)
        r += 1
    F.sweep(ws, "Start here", view, max_col=10)

    ws = wb["Control"]
    for r in range(15, 20):
        for c in "MNOP":
            F.add("Control", view, f"{c}{r}", ["ladder", r - 14, c], ws[f"{c}{r}"].value)
    for cell in ("E15", "F15", "I15", "E16", "F16", "I16", "E19", "F24", "I24"):
        if ws[cell].value is not None:
            F.add("Control", view, cell, ["control", cell], ws[cell].value)
    F.sweep(ws, "Control", view, max_col=16)

    ws = wb["Columns"]
    r = 13
    while ws.cell(r, 2).value not in (None, ""):
        col = ws.cell(r, 2).value
        for c, what in ((5, "why"), (6, "blank"), (7, "odd"), (12, "check")):
            v = ws.cell(r, c).value
            if v is not None and has_digit(v):
                F.add("Columns", view, ws.cell(r, c).coordinate, ["columns", col, what], v)
        r += 1
    F.sweep(ws, "Columns", view, max_col=16)

    ws = wb["Look"]
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, 2).value
        if lab in ("FICO", "ORIG_BAL", "REV_DEBT") and ws.cell(r + 1, 2).value == "Loans":
            col = lab
        if lab in ("Loans", "Blank", "Not a number", "Smallest", "Median", "Mean", "Largest",
                   "Loans with both values", "Dots shown", "Moves together (correlation)") or \
                (isinstance(lab, str) and (lab.startswith("At ") or lab == "Likely a code")):
            if isinstance(lab, str) and lab.startswith("REV_DEBT against"):
                continue
            for c in (3, 4):
                v = ws.cell(r, c).value
                if v is not None and has_digit(v):
                    F.add("Look", view, ws.cell(r, c).coordinate, ["look", col, lab, c], v)
        if isinstance(lab, str) and lab.startswith("REV_DEBT against"):
            col = lab
    F.sweep(ws, "Look", view, max_col=16, skip_cols=("C",) if False else ())
    # the chart bars: _look's bar loans (20 bars), the low and high end bars, and the code bar, per column
    lk = wb["_look"]
    for fcol, gcol, hcol in (("F", "G", "H"), ("S", "T", "U"), ("AF", "AG", "AH")):
        name = lk[f"{fcol}4"].value
        for r in range(4, 4 + int(lk[f"{fcol}13"].value)):
            F.add("Look", view, f"_look!{hcol}{r}",
                  ["look-bar", name, float(lk[f"{gcol}{r}"].value), float(lk[f"{fcol}16"].value),
                   r == 3 + int(lk[f"{fcol}13"].value)], lk[f"{hcol}{r}"].value)
        for r, what in ((18, "low end"), (19, "high end"), (10, "at the code")):
            if lk[f"{fcol}{r}"].value is not None and not (what == "at the code" and lk[f"{fcol}9"].value is None):
                F.add("Look", view, f"_look!{fcol}{r}", ["look-end", name, what], lk[f"{fcol}{r}"].value)

    ws = wb["Record"]
    for r in range(8, ws.max_row + 1):
        for c in (3, 6):
            v = ws.cell(r, c).value
            if v is not None and has_digit(v):
                F.add("Record", view, ws.cell(r, c).coordinate, ["record", ws.cell(r, c).coordinate,
                                                                   ws.cell(r, c - 1).value], v)
    F.sweep(ws, "Record", view, max_col=10)


def main():
    d = Path(sys.argv[1])
    plan = json.loads((d / "plan.json").read_text())
    F = Figures()
    seen_split_grids = set()
    seen = {"pockets": set(), "grids": set(), "pck": set(), "split": set()}
    for i, p in enumerate(plan):
        wb = load_workbook(d / "calculated" / p["file"], data_only=True)
        view = p["file"][:-5]
        if i == 0:
            read_once(F, wb, view)
            # for attributing differences only (compare.py, "sampling flips"): the p-value each pocket's verdict
            # used, from the workbook's own hidden record of every pocket (_pockets, columns A to U)
            pk = wb["_pockets"]
            head = [c.value for c in pk[3]]
            rows = [dict(zip(head, [c.value for c in r])) for r in pk.iter_rows(min_row=4, max_col=21)]
            (d / "wb-pockets.json").write_text(json.dumps([r for r in rows if r.get("Grid")], default=str))
        if tuple(p["pockets"]) not in seen["pockets"]:
            seen["pockets"].add(tuple(p["pockets"]))
            read_pockets(F, wb["Pockets"], view)
        if tuple(p["grids"]) not in seen["grids"]:
            seen["grids"].add(tuple(p["grids"]))
            read_grids(F, wb["Grids"], view)
        if tuple(p["split"]) not in seen["split"]:
            seen["split"].add(tuple(p["split"]))
            first = p["split"][0] not in seen_split_grids
            seen_split_grids.add(p["split"][0])
            read_split(F, wb["Split"], view, first)
        if p["pck"] not in seen["pck"]:
            seen["pck"].add(p["pck"])
            read_pck(F, wb["Paid, cost, kept"], view)
    (d / "figures.json").write_text(json.dumps(F.out, indent=0, default=str))
    kinds = {}
    for f in F.out:
        kinds[(f["tab"], f["kind"])] = kinds.get((f["tab"], f["kind"]), 0) + 1
    for k, v in sorted(kinds.items()):
        print(k, v)
    print("total", len(F.out))


if __name__ == "__main__":
    main()
