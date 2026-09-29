"""The PocketBook road for the second run kind: read every figure out of the calculated Test-new-variables workbook
(Start here, Control, Columns, Look, Scouting, New variables, Record), with openpyxl, and name each by the labels
printed beside it. Like read_bleed.py it imports nothing from PocketBook and works nothing out.

    python3 read_scout.py CALCULATED.xlsx OUT/figures-scout.json

This workbook has no dropdown views, so it is calculated once (views.py is not needed: recalc.py's recalc_file).
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from read_bleed import Figures, has_digit, hidden  # noqa: E402


def text_cells(F, ws, tab, cells, prefix):
    for c in cells:
        v = ws[c].value
        if v is not None and has_digit(v):
            F.add(tab, "one", c, [prefix, c], v)


def read_scouting(F, ws):
    tab = "Scouting"
    text_cells(F, ws, tab, ("C5", "C8", "C9", "C10", "D19", "G19", "H19"), "sc-text")
    r = 24
    heads = {c: ws.cell(23, c).value for c in range(3, 11)}
    while ws.cell(r, 2).value not in (None, ""):
        name = ws.cell(r, 2).value
        for c, h in heads.items():
            v = ws.cell(r, c).value
            if v is not None and has_digit(v) or h in ("Proposed?", "Why") and v:
                F.add(tab, "one", ws.cell(r, c).coordinate, ["sc-cand", name, h], v)
        r += 1
    for row in ws.iter_rows(min_row=r, max_row=ws.max_row):
        b = row[1]
        if b.value == "Column" and ws.cell(b.row, 4).value == "Importance":
            rr = b.row + 1
            while ws.cell(rr, 2).value not in (None, ""):
                F.add(tab, "one", f"D{rr}", ["sc-held", ws.cell(rr, 2).value], ws.cell(rr, 4).value)
                rr += 1
        if isinstance(b.value, str) and b.value.startswith("    bins: "):
            F.add(tab, "one", b.coordinate, ["sc-prespec", "bins", ws.cell(b.row - 1, 2).value.split(": ")[1]], b.value)
        if isinstance(b.value, str) and b.value.startswith(("holdout:", "development:")):
            F.add(tab, "one", b.coordinate, ["sc-prespec", b.value.split(":")[0]], b.value)
        m = re.match(r"(\S+): the forest's bad rate as \S+ moves", str(b.value or ""))
        if m:
            name, rr, i = m.group(1), b.row + 2, 0
            while ws.cell(rr, 2).value is not None and isinstance(ws.cell(rr, 2).value, (int, float)):
                F.add(tab, "one", f"B{rr}", ["sc-curve", name, i, "x"], ws.cell(rr, 2).value)
                F.add(tab, "one", f"D{rr}", ["sc-curve", name, i, "rate"], ws.cell(rr, 4).value)
                rr, i = rr + 1, i + 1
    F.sweep(ws, tab, "one", max_col=12)


def read_new_variables(F, ws):
    tab = "New variables"
    text_cells(F, ws, tab, ("C7", "C8", "C11", "C17", "C29", "C39", "F23", "L23", "N23"), "nv-text")
    rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=16))

    def table(r0, cols, key, name_col=2):
        """rows under a header at r0 until a blank in column C; a blank name repeats the one above."""
        heads = {}
        for c in cols:
            if ws.cell(r0, c).value:
                sec = next((ws.cell(r0 - 1, cc).value for cc in range(c, 3, -1) if ws.cell(r0 - 1, cc).value), "")
                heads[c] = (str(sec).split(" · ")[0] + ": " if sec else "") + str(ws.cell(r0, c).value)
        r, name = r0 + 1, None
        while ws.cell(r, 3).value not in (None, ""):
            name = ws.cell(r, name_col).value or name
            comp = ws.cell(r, 3).value
            for c, h in heads.items():
                v = ws.cell(r, c).value
                if v is not None and (has_digit(v) or isinstance(v, str) and v in ("Yes", "No") or
                                      h.endswith("In words")):
                    F.add(tab, "one", ws.cell(r, c).coordinate, [key, name, comp, h], v)
            r += 1
    for row in rows:
        b = row[1]
        if b.value == "Candidate" and ws.cell(b.row, 8).value == "Odds ratio, together":
            table(b.row, range(6, 15), "nv-joint")
        if b.value == "Candidate" and ws.cell(b.row, 4).value == "Gap":
            table(b.row, range(4, 15), "nv-main")
    cand, joint = None, False
    for row in rows:
        b = row[1]
        v = str(b.value or "")
        m = re.match(r"Does (\S+) matter\?", v)
        if m:
            cand = m.group(1)
        if v.startswith(("Found,", "Confirmed on the held-back loans,")) and ":" in v:
            block = v.split(":")[0]
            rr = b.row + 1
            while ws.cell(rr, 2).value and ws.cell(rr, 2).value != "What it found":
                test = ws.cell(rr, 2).value
                for c, h in ((6, "Statistic"), (7, "Degrees of freedom"), (8, "p-value"), (9, "Reading")):
                    vv = ws.cell(rr, c).value
                    if vv is not None:
                        F.add(tab, "one", ws.cell(rr, c).coordinate, ["nv-test", cand, block, test, h], vv)
                rr += 1
            if ws.cell(rr, 2).value == "What it found":            # the tests above, said in words
                F.add(tab, "one", f"C{rr}", ["nv-reading", cand, block], ws.cell(rr, 3).value)
        if v.startswith("Group of ") and ws.cell(b.row, 9).value == "Odds ratio":
            head = {c: ws.cell(b.row, c).value for c in range(6, 15) if ws.cell(b.row, c).value}
            rr = b.row + 1
            block = ws.cell(rr, 2).value.split(":")[0]           # Found / Confirmed on the held-back loans
            rr += 1
            while ws.cell(rr, 2).value and ws.cell(rr, 2).value != "What it found":
                grp = ws.cell(rr, 2).value.replace(" (reference)", "")
                seen = {}
                for c, h in head.items():
                    seen[h] = seen.get(h, 0) + 1
                    hh = h if seen[h] == 1 else h + " (held fixed)"
                    vv = ws.cell(rr, c).value
                    if vv is not None and has_digit(vv):
                        F.add(tab, "one", ws.cell(rr, c).coordinate, ["nv-group", cand, block, grp, hh], vv)
                rr += 1
        if v.startswith("Group of ") and ws.cell(b.row, 7).value == "Share of loans":
            head = {c: ws.cell(b.row, c).value for c in range(6, 14) if ws.cell(b.row, c).value}
            rr = b.row + 1
            while ws.cell(rr, 2).value and ws.cell(rr, 2).value != "What it found":
                grp = ws.cell(rr, 2).value.replace(" (reference)", "")
                for c, h in head.items():
                    vv = ws.cell(rr, c).value
                    if vv is not None:
                        F.add(tab, "one", ws.cell(rr, c).coordinate, ["nv-conc", cand, grp, h], vv)
                rr += 1
        if v.startswith("All ") and v.endswith("together, on the held-back loans"):
            joint = True
        if v == "What it found" and cand is None and joint:          # the together table's verdict in words
            F.add(tab, "one", f"C{b.row}", ["nv-joint-found"], ws.cell(b.row, 3).value)
            joint = False
        if v == "What it found" and cand is not None and has_digit(ws.cell(b.row, 3).value or ""):
            t = ws.cell(b.row, 3).value
            where = "Found" if t.startswith("On development") else "Confirmed"
            kind = "conc" if " holds " in t else "plain" if "with nothing held fixed" in t else "held"
            F.add(tab, "one", f"C{b.row}", ["nv-words", cand, where, kind], t)
    F.sweep(ws, tab, "one", max_col=16)


def read_start(F, ws):
    tab = "Start here"
    text_cells(F, ws, tab, ("C1", "B17", "D17", "F17"), "st-text")
    heads = {c: ws.cell(19, c).value for c in range(3, 9)}
    r = 20
    while ws.cell(r, 2).value not in (None, "") and ":" in str(ws.cell(r, 2).value):
        name = str(ws.cell(r, 2).value).replace(" (reference)", "")
        for c, h in heads.items():
            v = ws.cell(r, c).value
            if v is not None:
                F.add(tab, "one", ws.cell(r, c).coordinate, ["st-row", name, h], v)
        r += 1
    F.sweep(ws, tab, "one", max_col=10)


def read_rest(F, wb):
    ws = wb["Control"]
    text_cells(F, ws, "Control", ("E30", "F30", "I30", "E19", "E15", "F15", "I15"), "control")
    F.sweep(ws, "Control", "one", max_col=16)
    ws = wb["Columns"]
    r = 13
    while ws.cell(r, 2).value not in (None, ""):
        col = ws.cell(r, 2).value
        for c, what in ((5, "why"), (6, "blank"), (7, "odd"), (12, "check")):
            v = ws.cell(r, c).value
            if v is not None and has_digit(v):
                F.add("Columns", "one", ws.cell(r, c).coordinate, ["columns", col, what], v)
        r += 1
    F.sweep(ws, "Columns", "one", max_col=16)
    ws = wb["Look"]
    col = None
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, 2).value
        if isinstance(lab, str) and ws.cell(r + 1, 2).value == "Loans":
            col = lab
        if isinstance(lab, str) and lab.startswith(("REV_DEBT against", "Loans with both")):
            pass
        if lab in ("Loans", "Blank", "Not a number", "Smallest", "Median", "Mean", "Largest") or \
                (isinstance(lab, str) and (lab.startswith("At ") or lab == "Likely a code")):
            for c in (3, 4):
                v = ws.cell(r, c).value
                if v is not None and (has_digit(v) or lab == "Likely a code"):
                    F.add("Look", "one", ws.cell(r, c).coordinate, ["look", col, lab, c], v)
    F.sweep(ws, "Look", "one", max_col=16)
    ws = wb["Record"]
    for r in range(8, ws.max_row + 1):
        for c in (3, 6, 7):
            v = ws.cell(r, c).value
            if v is not None and has_digit(v):
                F.add("Record", "one", ws.cell(r, c).coordinate, ["record", ws.cell(r, c).coordinate], v)
    F.sweep(ws, "Record", "one", max_col=10)
    # the Look chart bars, from the chart's own data range on _look (as read_bleed.py does)
    lk = wb["_look"]
    c = 6
    while c < lk.max_column:
        name = lk.cell(4, c).value
        if lk.cell(4, c - 1).value == "column" and name:
            bars = lk.cell(13, c).value
            for r in range(4, 4 + int(bars)):
                F.add("Look", "one", f"_look!{lk.cell(r, c + 2).coordinate}",
                      ["look-bar", name, float(lk.cell(r, c + 1).value), float(lk.cell(16, c).value),
                       r == 3 + int(bars)], lk.cell(r, c + 2).value)
            for r, what in ((18, "low end"), (19, "high end"), (10, "at the code")):
                if lk.cell(r, c).value is not None and not (what == "at the code" and lk.cell(9, c).value is None):
                    F.add("Look", "one", f"_look!{lk.cell(r, c).coordinate}", ["look-end", name, what],
                          lk.cell(r, c).value)
        c += 1


def main():
    wb = load_workbook(sys.argv[1], data_only=True)
    F = Figures()
    read_start(F, wb["Start here"])
    read_scouting(F, wb["Scouting"])
    read_new_variables(F, wb["New variables"])
    read_rest(F, wb)
    Path(sys.argv[2]).write_text(json.dumps(F.out, indent=0, default=str))
    kinds = {}
    for f in F.out:
        kinds[(f["tab"], f["kind"])] = kinds.get((f["tab"], f["kind"]), 0) + 1
    for k, v in sorted(kinds.items()):
        print(k, v)
    print("total", len(F.out))


if __name__ == "__main__":
    main()
