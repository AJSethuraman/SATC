"""Road 1: what PocketBook shows. Makes the Kiosk book, runs PocketBook on it, and for every choice of Paid, cost,
kept's Grid dropdown has LibreOffice calculate the tab, then reads every gross cell (Booked, GCO, RANR, RANR rate,
Loans), whether each RANR and RANR rate cell is drawn red, the count line beside the dropdown, and the totals under
the table. Writes work/road1.csv.

    python3 docs/tie-out/2026-09-30-pck-gross/road1.py      (from pocketbook/)
"""
import csv
import os
import random
import re
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[2]
sys.path[:0] = [str(ROOT / "src"), str(ROOT / "tests")]
WORK = HERE / "work"
os.environ["POCKETBOOK_MEMORY"] = str(WORK / "memory.yaml")

from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import range_boundaries  # noqa: E402
from openpyxl.utils.cell import column_index_from_string  # noqa: E402

from pocketbook import book, choices as ch, house, results as rs, synth  # noqa: E402
from recalc import recalc  # noqa: E402
from test_book import _answer  # noqa: E402

BOOK_CSV = HERE / "Kiosk book.csv"


def make_file() -> Path:
    """The synthetic book of 8,000 loans, and 90 through a Kiosk that lose money outright: 4% of what they booked,
    and their charge-offs on top. Every 9th Kiosk loan charged off 30%; one Kiosk loan's RANR is left blank."""
    rows = synth.make_rows(8000, seed=30)
    rng = random.Random(930)
    for k in range(90):
        bal = round(rng.uniform(8000, 40000), 2)
        gco = round(bal * 0.3, 2) if k % 9 == 0 else 0.0
        rows.append({**rows[10], "LOAN_NBR": f"K{k:03d}", "FICO": rng.randint(560, 800), "CHANNEL": "Kiosk",
                     "ORIG_BAL": bal, "BAD_FLAG": 1 if gco else 0, "GCO_AMT": gco,
                     "ASSET_CLASS": rng.choice((1, 2, 3, 4)), "RANR_AMT": round(-0.04 * bal - gco, 2)})
    rows[-5]["RANR_AMT"] = ""                   # a blank RANR: left out of RANR and of the booked under it
    with BOOK_CSV.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return BOOK_CSV


RED = re.compile(r'^(?:AND\(\$[A-Z]+\d+<>"",)?AND\(ISNUMBER\(\$([A-Z]+)(\d+)\),\$[A-Z]+\d+<0\)\)?$')


def red(written, calc, r: int, c: int) -> bool:
    """Whether the red-text rule the tab carries holds on a cell, worked out from the calculated values."""
    for rng in written.conditional_formatting:
        for bounds in str(rng.sqref).split():
            c0, r0, c1, r1 = range_boundaries(bounds)
            if not (c0 <= c <= c1 and r0 <= r <= r1):
                continue
            for rule in rng.rules:
                font = rule.dxf.font if rule.dxf is not None else None
                m = RED.match(rule.formula[0])
                if m and font is not None and font.color.rgb[-6:] == house.CRIMSON:
                    v = calc.cell(row=int(m.group(2)) + r - r0, column=column_index_from_string(m.group(1))).value
                    return isinstance(v, (int, float)) and v < 0
    return False


def main() -> None:
    if WORK.exists():
        shutil.rmtree(WORK)
    WORK.mkdir()
    x = make_file()
    b = WORK / "Kiosk book - PocketBook.xlsx"
    out = book.set_up(x, b, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",),
                                               segments=("CHANNEL", "ASSET_CLASS")))
    _answer(out.book)
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    wb = load_workbook(b)
    ws = wb[rs.PCK]
    grid_cell = next(ws.cell(row=c.row + 1, column=c.column) for row in ws.iter_rows(max_row=60) for c in row
                     if c.value == "GRID")
    dv = next(d for d in ws.data_validations.dataValidation if grid_cell.coordinate in str(d.sqref))
    sheet, rng = dv.formula1.lstrip("=").split("!")
    c0, r0, _, r1 = range_boundaries(rng.replace("$", ""))
    grids = [wb[sheet.strip("'")].cell(row=r, column=c0).value for r in range(r0, r1 + 1)]
    lines = []
    for i, grid in enumerate(grids):
        pick = WORK / f"grid{i}.xlsx"
        shutil.copy(b, pick)
        w = load_workbook(pick)
        w[rs.PCK][grid_cell.coordinate] = grid
        w.save(pick)
        written = load_workbook(pick)[rs.PCK]
        calc = recalc(pick, WORK / f"rc{i}")[rs.PCK]
        head = next(r for r in range(1, calc.max_row + 1) if calc.cell(row=r, column=rs.C_TOG).value == "Together")
        lines.append((grid, "line", "", "", "count line", calc.cell(row=grid_cell.row, column=rs.C_LOANS).value))
        lines.append((grid, "line", "", "", "count line red",
                      bool(calc.cell(row=grid_cell.row, column=rs.C_H_UN).value)))
        for r in range(head + 1, calc.max_row + 1):
            band = calc.cell(row=r, column=rs.C_BAND).value
            if band in (None, ""):
                continue
            kind = "total" if band in (rs.PCK_LISTED, rs.PCK_UNLISTED, rs.PCK_BOOK) else "pocket"
            seg = calc.cell(row=r, column=rs.C_SEG).value if kind == "pocket" else ""
            for name, c in (("loans", rs.C_LOANS), ("booked", rs.C_BOOK), ("gco", rs.C_GCO), ("ranr", rs.C_RANR),
                            ("ranr_rate", rs.C_RATE)):
                lines.append((grid, kind, band, seg, name, calc.cell(row=r, column=c).value))
            for name, c in (("ranr red", rs.C_RANR), ("ranr_rate red", rs.C_RATE), ("booked red", rs.C_BOOK),
                            ("gco red", rs.C_GCO)):
                lines.append((grid, kind, band, seg, name, red(written, calc, r, c)))
    with (WORK / "road1.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["grid", "kind", "band", "segment", "figure", "value"])
        w.writerows(lines)
    print(f"{len(grids)} grids, {len(lines)} figures read -> {WORK / 'road1.csv'}")


if __name__ == "__main__":
    main()
