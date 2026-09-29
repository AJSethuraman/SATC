"""Did the tie-out skip anything? An independent check of the roster's coverage, written after the firm asked
"how do we know it checked every figure if they aren't in the pdf?" (29 Sep 2026).

    python3 completeness.py roster.csv RUN WORK/views-RUN/calculated [RUN FOLDER ...]

(29 Sep 2026: any number of runs, each named as roster.csv names it, with the folder of its calculated workbooks.)

It opens every calculated workbook the tie-out read (all 40 dropdown views of the bleed workbook, and the new-variable
workbook) with openpyxl alone -- none of the tie-out's own reading code -- lists every number on every visible tab
that an analyst can see (hidden rows and columns are the workbook's own helpers and are left out), and sorts each one:

  checked in this view                  its (tab, cell, view) is a row of roster.csv
  same figure, checked in another view  the same cell shows the same value in a view whose row is in roster.csv
                                        (a figure the dropdown doesn't change, e.g. the Loans block under every measure)
  NEVER CHECKED                         neither -- the count that has to be zero

Then the same for text inside the result tables (below each tab's frozen header), and it prints every distinct
text that was never checked, so a reader can see for themselves that none is a verdict.
"""
import collections
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


def roster_keys(path):
    have = set()
    for r in csv.DictReader(open(path, encoding="utf-8")):
        m = re.match(r"^(.*?)!([A-Z]+\d+)(?: \[(view-\d+)\])?$", r["cell"])
        if m:
            have.add((r["run"], m.group(1), m.group(2), m.group(3)))
    return have


def hidden(ws):
    cols = set()
    for k, d in ws.column_dimensions.items():          # a hidden range R:U is filed under R alone
        if d.hidden:
            first = d.min or range_boundaries(k + "1")[0]
            cols |= {get_column_letter(i) for i in range(first, (d.max or first) + 1)}
    return cols, {k for k, d in ws.row_dimensions.items() if d.hidden}


def survey(run, have, files):
    viewed = {t for r, t, c, v in have if r == run and v}
    cells = {}
    for f in files:
        view = f.stem if f.stem.startswith("view-") else None
        wb = load_workbook(f, data_only=True)
        for ws in wb.worksheets:
            if ws.sheet_state != "visible" or (ws.title not in viewed and view not in (None, "view-00")):
                continue
            hc, hr = hidden(ws)
            top = range_boundaries(f"{ws.freeze_panes or 'A1'}:{ws.freeze_panes or 'A1'}")[1]
            for row in ws.iter_rows():
                for c in row:
                    v = c.value
                    if v is None or c.column_letter in hc or c.row in hr:
                        continue
                    kind = "number" if isinstance(v, (int, float)) and not isinstance(v, bool) else "text"
                    if kind == "text" and (not str(v).strip() or c.row < top):
                        continue
                    cells[(ws.title, c.coordinate, view if ws.title in viewed else None)] = (kind, v)
    seen = {(t, c, k, round(v, 9) if k == "number" else v) for (t, c, w), (k, v) in cells.items()
            if (run, t, c, w) in have}
    tally, never = collections.Counter(), collections.Counter()
    for (t, c, w), (k, v) in cells.items():
        if (run, t, c, w) in have:
            what = "checked in this view"
        elif (t, c, k, round(v, 9) if k == "number" else v) in seen:
            what = "same figure, checked in another view"
        else:
            what = "NEVER CHECKED"
            if k == "text":
                never[v] += 1
        tally[(k, what)] += 1
    return tally, never


def main():
    have = roster_keys(sys.argv[1])
    for run, folder in zip(sys.argv[2::2], sys.argv[3::2]):
        tally, never = survey(run, have, sorted(Path(folder).glob("*.xlsx")))
        print(f"{run}:")
        for (k, what), n in sorted(tally.items()):
            print(f"  {k:6s} {what:38s} {n:6d}")
        if never:
            print(f"  {len(never)} distinct table texts never checked (labels, if none is a verdict):")
            for v, n in never.most_common():
                print(f"    {n:5d}  {str(v)[:80]!r}")


if __name__ == "__main__":
    main()
