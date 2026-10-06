"""Step 2 of the PocketBook road: read every figure out of the calculated views of a Where-the-book-bleeds workbook,
cell by cell, with openpyxl. Nothing here imports PocketBook or works anything out: it reads what the analyst would
read, and names each figure by the labels printed beside it on the same tab.

    python3 read_book.py VIEWS_DIR          (writes VIEWS_DIR/figures.json)

The full tie-out's read_bleed.py (28 Sep 2026), made to find its cells by their labels rather than by fixed rows --
a category split's tabs carry one more line of notes, so every row below moves -- and extended to what this build
added:

  - Grids: "What one cell says", the six lines under the blocks, for the Row and Column each view picked;
  - Split: a category split (each value against the rest of its pocket) as well as the halves, and the line "Do the
    values of ... differ at all?";
  - Paid, cost, kept: every dot of the chart, from the hidden sheet it is drawn from (_chart): its x and y, whether
    it is in the red series (H_RX/H_RY) or the green one (H_GX/H_GY), and its number (H_NUM); and the list
    "Numbered on the chart" under it;
  - Look: the row "Answered missing, left out", every column's block wherever it sits, and the scatters' dots
    (_dots), so a value answered missing can be looked for in them;
  - Record: each figure is named by the label it sits under and how many rows below it (the rows move).

Every other rule is read_bleed.py's: every visible cell holding a number or a digit is recorded, as a named figure or
as `words`, so the roster can account for every one.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

DIGIT = re.compile(r"\d")
STAMP = re.compile(r"^\d{4}-\d\d-\d\d \d\d:\d\d$")


def has_digit(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool) or (isinstance(v, str) and DIGIT.search(v))


class Figures:
    def __init__(self):
        self.out = []
        self.seen = set()

    def add(self, tab, view, cell, key, value, kind="figure"):
        self.out.append({"tab": tab, "view": view, "cell": cell, "key": key, "value": value, "kind": kind})
        self.seen.add((tab, view, cell))

    def sweep(self, ws, tab, view, max_col=60, rows=None):
        """Every digit-bearing cell the analyst can see on the tab and not already named: recorded as words."""
        hid_c, hid_r = hidden(ws)
        for row in ws.iter_rows(max_col=max_col):
            for c in row:
                if c.value is None or c.column in hid_c or c.row in hid_r or (tab, view, c.coordinate) in self.seen:
                    continue
                if rows is not None and c.row not in rows:
                    continue
                if has_digit(c.value):
                    self.add(tab, view, c.coordinate, ["words"], c.value, kind="words")


def hidden(ws):
    cols, rows = set(), set()
    for k, d in ws.column_dimensions.items():
        if d.hidden:
            lo = d.min or column_index_from_string(k)
            hi = d.max or lo
            cols.update(range(lo, hi + 1))
    for k, d in ws.row_dimensions.items():
        if d.hidden:
            rows.add(k)
    return cols, rows


def find(ws, text, max_row=400, max_col=40, min_row=12, col=None):
    for row in ws.iter_rows(min_row=min_row, max_row=max_row, max_col=max_col):
        for c in row:
            if c.value == text and (col is None or c.column == col):
                return c.row, c.column
    return None


def find_start(ws, prefix, max_row=400, col=2):
    for r in range(1, max_row + 1):
        v = ws.cell(r, col).value
        if isinstance(v, str) and v.startswith(prefix):
            return r
    return None


def below(ws, label):
    """The dropdown under its label; a label ending in * is matched by its start ("ONLY LOANS WHERE SYS_FLAG IS")."""
    for row in ws.iter_rows(min_row=8, max_row=40, max_col=30):
        for c in row:
            if c.value == label or (label.endswith("*") and str(c.value or "").startswith(label[:-1])):
                return ws.cell(c.row + 1, c.column)
    raise KeyError(label)


def header(ws, r, c0):
    out, c = [], c0
    while ws.cell(r, c).value not in (None, ""):
        out.append((c, str(ws.cell(r, c).value)))
        c += 1
    return out


def read_lines(F, ws, tab, view, measure):
    """The lines in use, live from Control, shown above each result tab's rows: the row under LINES IN USE NOW,
    named by the heading over each cell."""
    at = find(ws, "LINES IN USE NOW", max_row=30, max_col=4, min_row=8)
    r = at[0]
    names = {"Worse at": "worse", "Better at": "better", "How sure": "sure", "Material at": "material"}
    for c in range(at[1] + 1, at[1] + 12):
        what = names.get(ws.cell(r, c).value)
        v = ws.cell(r + 1, c).value
        if what and v is not None:
            F.add(tab, view, ws.cell(r + 1, c).coordinate, ["line", what, measure if what == "material" else "*"], v)


# --------------------------------------------------------------------------- Pockets

def read_pockets(F, ws, view):
    mc = below(ws, "MEASURE")
    measure, kind = mc.value, below(ws, "POCKETS").value
    read_lines(F, ws, "Pockets", view, measure)
    count = ws.cell(mc.row, mc.column + 3)
    F.add("Pockets", view, count.coordinate, ["pockets-count", kind, measure], count.value)
    hr = find(ws, "#", max_row=40, max_col=3)[0]
    heads = {str(ws.cell(hr, c).value): c for c in range(2, 20) if ws.cell(hr, c).value}
    half_col = next((c for h, c in heads.items() if h.endswith(" half") or c == 5), None)
    r = hr + 1
    while ws.cell(r, 3).value not in (None, ""):
        band, seg = ws.cell(r, 3).value, ws.cell(r, 4).value
        half = ws.cell(r, half_col).value if half_col else None
        who = [kind, measure, band, seg, half]
        for name, c in heads.items():
            if name in ("Band", "Segment") or c == half_col:
                continue
            v = ws.cell(r, c).value
            if v is None:
                continue
            field = {"#": "rank", "Loans": "loans", "This pocket": "rate", "Worse?": "worse",
                     "p-value": "p", "Material?": "material", "Could have caught": "caught",
                     "Holds FICO fixed?": "holds"}.get(name)
            if field is None:
                field = "rest-rate" if name.startswith("Rest of") else "gap" if name.startswith(("×", "Gap")) \
                    else "dollars"
            F.add("Pockets", view, ws.cell(r, c).coordinate, ["pocket", *who, field], v)
        r += 1
    F.sweep(ws, "Pockets", view, max_col=16)


# --------------------------------------------------------------------------- Grids

def grid_block(F, ws, view, grid, measure, title_rc, block, extra=None):
    """One block: its header row is the first row under the title with the band column's name in it (a split grid's
    header has a segment row over its parts, merged: the label is "segment · part", as the Row and Column lists
    name them). Returns {(row label, column label): (coordinate, value)}."""
    r0, c0 = title_rc
    hr = r0 + 1
    while ws.cell(hr, c0).value in (None, ""):
        hr += 1
    heads = header(ws, hr, c0 + 1)
    if hr > r0 + 1:                                      # the segment row, merged over its parts
        seg, out = "", []
        for c, lab in heads:
            s = ws.cell(hr - 1, c).value
            seg = str(s) if s not in (None, "") else seg
            out.append((c, f"{seg} · {lab}" if lab != "All" and seg else lab))
        heads = out
    got = {}
    r = hr + 1
    while ws.cell(r, c0).value not in (None, ""):
        rowlab = str(ws.cell(r, c0).value)
        for c, collab in heads:
            v = ws.cell(r, c).value
            got[(rowlab, collab)] = (ws.cell(r, c).coordinate, v)
            if v is not None:
                F.add("Grids", view, ws.cell(r, c).coordinate, ["grid", grid, measure, block, rowlab, collab], v)
        if rowlab == "All":
            break
        r += 1
    return got


def read_grids(F, ws, view, first, meta):
    """`meta`: the as-written workbook's _views rows, for each view's heat kind, bound and fewest loans (the hidden
    cells the colour rules read)."""
    grid, measure = below(ws, "GRID").value, below(ws, "MEASURE").value
    row_, col_ = below(ws, "ROW").value, below(ws, "COLUMN").value
    flt = below(ws, "ONLY LOANS WHERE*").value
    gkey = grid if flt in (None, "", "All loans") else f"{grid} | where {flt}"
    rate = find(ws, f"Rate · {measure}")
    grid_block(F, ws, view, gkey, measure, rate, "rate")
    book = next((r, c) for row in ws.iter_rows(min_row=12, max_row=40, max_col=40) for c_ in row
                for r, c in [(c_.row, c_.column)] if isinstance(c_.value, str) and c_.value.startswith("vs the book"))
    F.add("Grids", view, ws.cell(*book).coordinate, ["grid-head", gkey, measure], ws.cell(*book).value)
    bk_ = grid_block(F, ws, view, gkey, measure, book, "book")
    band = find(ws, "vs rest of band")
    bd_ = grid_block(F, ws, view, gkey, measure, band, "band")
    loans = (band[0], book[1])
    assert ws.cell(*loans).value == "Loans"
    ln_ = grid_block(F, ws, view, gkey, measure, loans, "loans")
    if gkey == grid:
        # build 44734da4: each band label with the loans its band holds (the Loans block's All column), so the
        # loan-file road can count the loans inside the label's own range
        for (rl, cl), (coord, v) in ln_.items():
            if cl == "All" and rl != "All" and isinstance(v, (int, float)):
                F.add("Grids", view, coord + " (band label)", ["band-label", grid.split(" x ")[0], rl], v)
    # grey (change 1 of 30 Sep): the colour rules' test, a number in vs the book or vs rest of band whose pocket has
    # fewer loans than the fewest loans the Run used, evaluated with the calculated values; and the bound the heat
    # scale is read against, from the view's own row on _views
    mkey, vkey = meta["keys"].get(measure), grid if gkey == grid else f"{grid}|where {flt}"
    m = meta["rows"].get(f"G|{vkey}|{mkey}|meta")
    if m is not None:
        kind, bound, _, bookfig, few = (m + [None] * 5)[:5]
        F.add("Grids", view, "_views (meta)", ["grid-bound", gkey, measure], bound)
        F.add("Grids", view, "_views (meta) fewest", ["grid-few", gkey, measure], few)
        F.add("Grids", view, "_views (meta) book", ["grid-bookfig", gkey, measure], bookfig)
        for blk, got in (("book", bk_), ("band", bd_)):
            for (rl, cl), (coord, v) in got.items():
                if isinstance(v, (int, float)) and not isinstance(v, bool):
                    n = ln_.get((rl, cl), (None, None))[1]
                    grey = isinstance(n, (int, float)) and isinstance(few, (int, float)) and n < few
                    F.add("Grids", view, coord + " (grey?)", ["grid-grey", gkey, measure, blk, rl, cl], grey)
    if mkey == "loan_size":
        # the median booked per loan: shown only in What one cell says, read from its own row on _views
        for (rl, cl), (coord, v) in ln_.items():
            pass
        for i, rl in enumerate(meta["rows"].get(f"G|{vkey}|rows", []) or [], start=1):
            med = meta["rows"].get(f"G|{vkey}|loan_size|median|{i}") or []
            cols = meta["rows"].get(f"G|{vkey}|cols") or []
            for cl, v in zip(cols, med):
                if v is not None and rl is not None:
                    F.add("Grids", view, f"_views (median {rl} {cl})", ["size-median", gkey, str(rl), str(cl)], v)
    at = find_start(ws, "What one cell says")
    F.add("Grids", view, f"B{at + 1}", ["panel", gkey, measure, row_, col_, "name"], ws.cell(at + 1, 2).value)
    for k in range(2, 7):
        lab = ws.cell(at + k, 2).value
        F.add("Grids", view, f"C{at + k}", ["panel", gkey, measure, row_, col_, lab], ws.cell(at + k, 3).value)
    g = find(ws, "How common each group is: a count, not a test")
    if g and first:                           # the same block under every measure and pick: read once per grid
        hr = g[0] + 3
        if ws.cell(hr, 2).value not in (None, "") and ws.cell(hr, 4).value == "Loans":
            groups = [(c, str(ws.cell(hr, c).value)) for c in range(6, 20) if ws.cell(hr, c).value]
            r = hr + 2
            while ws.cell(r, 2).value not in (None, ""):
                band_, seg = str(ws.cell(r, 2).value), str(ws.cell(r, 3).value)
                cols = {4: "loans", 5: "booked"}
                for c, lab in groups:
                    cols[c], cols[c + 1] = f"{lab}|loans", f"{lab}|booked"
                for c, what in cols.items():
                    v = ws.cell(r, c).value
                    if v is not None:
                        F.add("Grids", view, ws.cell(r, c).coordinate, ["groups", grid, band_, seg, what], v)
                r += 1
    F.sweep(ws, "Grids", view, max_col=23)


def as_written_meta(path):
    """The as-written workbook's _views rows and the Grids measure names' keys (_choices)."""
    wb = load_workbook(path, read_only=True)
    rows = {}
    for r in wb["_views"].iter_rows(values_only=True):
        if r and isinstance(r[0], str) and r[0].startswith("G|"):
            vals = list(r[1:])
            while vals and vals[-1] is None:
                vals.pop()
            rows[r[0]] = vals
    ch = wb["_choices"]
    head = [c.value for c in next(ch.iter_rows(min_row=1, max_row=1))]
    i, j = head.index("Grids: Measure"), head.index("Grids: Measure (1)")
    keys = {r[i]: r[j] for r in ch.iter_rows(min_row=2, values_only=True) if r[i]}
    return {"rows": rows, "keys": keys}


# --------------------------------------------------------------------------- Split

def read_split(F, ws, view, first_of_grid):
    gc = below(ws, "GRID")
    mc = below(ws, "MEASURE")
    grid, measure = gc.value, mc.value
    read_lines(F, ws, "Split", view, "Charge-offs")
    hr = find(ws, "Measure", max_row=40, max_col=3, col=2)[0]
    if first_of_grid:
        chip = ws.cell(gc.row, gc.column + 3)
        if chip.value not in (None, ""):
            F.add("Split", view, chip.coordinate, ["split-chip", grid], chip.value)
        heads = {c: str(ws.cell(hr, c).value) for c in range(3, 10)}
        r = hr + 1
        while ws.cell(r, 2).value not in (None, "") and not str(ws.cell(r, 2).value).startswith("Same in"):
            m = ws.cell(r, 2).value
            for c, h in heads.items():
                v = ws.cell(r, c).value
                if v is not None:
                    F.add("Split", view, ws.cell(r, c).coordinate, ["split-sum", grid, m, h.split(" (")[0]], v)
            r += 1
        F.add("Split", view, f"B{r}", ["split-steady", grid], ws.cell(r, 2).value)
        if str(ws.cell(r + 1, 2).value or "").startswith("Do the values"):
            F.add("Split", view, f"B{r + 1}", ["split-differ", grid], ws.cell(r + 1, 2).value)
    top = mc.row + 3                                   # the two grids' column labels
    cols = header(ws, top, 3)
    pcols = header(ws, top, 9)
    r = top + 1
    while ws.cell(r, 2).value not in (None, ""):
        band = str(ws.cell(r, 2).value)
        for c, seg in cols:
            v = ws.cell(r, c).value
            if v is not None:
                F.add("Split", view, ws.cell(r, c).coordinate, ["split-pocket", grid, measure, band, seg, "gap"], v)
        for c, seg in pcols:
            v = ws.cell(r, c).value
            if v is not None:
                F.add("Split", view, ws.cell(r, c).coordinate, ["split-pocket", grid, measure, band, seg, "p"], v)
        r += 1
    if first_of_grid:
        F.sweep(ws, "Split", view, max_col=12)
    else:
        F.sweep(ws, "Split", view, max_col=12, rows=set(range(mc.row, r + 2)))


# --------------------------------------------------------------------------- Paid, cost, kept

def read_pck(F, ws, cs, view):
    grid = below(ws, "GRID").value
    read_lines(F, ws, "Paid, cost, kept", view, "Charge-offs")
    names = {4: "loans", 5: "earned-gap", 6: "earned-dollars", 7: "gco-multiple", 8: "gco-dollars",
             9: "kept-gap", 10: "kept-dollars", 11: "together"}
    first = find(ws, "Band", max_row=40, max_col=3)[0] + 1
    r = first
    rows = []
    while ws.cell(r, 2).value not in (None, ""):
        band, seg = str(ws.cell(r, 2).value), str(ws.cell(r, 3).value)
        rows.append((band, seg))
        for c, what in names.items():
            v = ws.cell(r, c).value
            if v is not None and v != "":
                F.add("Paid, cost, kept", view, ws.cell(r, c).coordinate, ["pck", grid, band, seg, what], v)
        r += 1
    # the chart (change 3): _chart row k is the tab's k-th row. A: x, B: y, C: its name when numbered, D/E: that
    # point again, F/G: the red series (Net drain), H/I: the green (Strong or Priced for it), J: its number
    for k, (band, seg) in enumerate(rows, start=1):
        get = [cs.cell(k, c).value for c in range(1, 11)]
        num = lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else None  # noqa: E731
        who = ["pck-dot", grid, band, seg]
        F.add("Paid, cost, kept", view, f"_chart!A{k}", who + ["x"], num(get[0]))
        F.add("Paid, cost, kept", view, f"_chart!B{k}", who + ["y"], num(get[1]))
        red = num(get[5]) is not None and num(get[6]) is not None
        green = num(get[7]) is not None and num(get[8]) is not None
        colour = "red" if red and not green else "green" if green and not red else "none" if not red else "both"
        F.add("Paid, cost, kept", view, f"_chart!F{k}:I{k}", who + ["colour"], colour)
        if red:
            F.add("Paid, cost, kept", view, f"_chart!F{k}:G{k}", who + ["red-at"],
                  [get[5] == get[0], get[6] == get[1]])
        if green:
            F.add("Paid, cost, kept", view, f"_chart!H{k}:I{k}", who + ["green-at"],
                  [get[7] == get[0], get[8] == get[1]])
        F.add("Paid, cost, kept", view, f"_chart!J{k}", who + ["number"], num(get[9]))
    # the list under the chart: "Numbered on the chart", then "k  band / seg"
    at = find(ws, "Numbered on the chart", max_row=200, max_col=30)
    if at:
        for k in range(1, 9):
            c = ws.cell(at[0] + k, at[1])
            if c.value not in (None, ""):
                m = re.match(r"^(\d+)  (.*)$", str(c.value))
                n, name = int(m.group(1)), m.group(2)
                F.add("Paid, cost, kept", view, c.coordinate, ["pck-listed", grid, n], name)
    F.sweep(ws, "Paid, cost, kept", view, max_col=30)


# --------------------------------------------------------------------------- the tabs read once

def record_anchor(ws, r, c):
    """The label a Record cell sits under, and how many rows below it: the nearest label at or above it in the
    column to its left. A run's own time as a label reads as "(the run's time)"."""
    for k in range(0, 40):
        lab = ws.cell(r - k, c - 1).value
        if lab not in (None, ""):
            lab = "(the run's time)" if STAMP.match(str(lab)) else str(lab)
            return lab, k
    return None, 0


def read_look(F, wb, view):
    ws = wb["Look"]
    col = None
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, 2).value
        if isinstance(lab, str) and ws.cell(r + 1, 2).value == "Loans" and " against " not in lab:
            col = lab
        if isinstance(lab, str) and " against " in lab:
            col = lab
            continue
        if lab in ("Loans", "Blank", "Not a number", "Smallest", "Median", "Mean", "Largest",
                   "Answered missing, left out", "Loans with both values", "Dots shown",
                   "Moves together (correlation)") or \
                (isinstance(lab, str) and "th percentile (P" in lab) or \
                (isinstance(lab, str) and (lab.startswith("At ") or lab == "Likely a code")):
            for c in (3, 4):
                v = ws.cell(r, c).value
                if v is not None and (has_digit(v) or lab == "Likely a code"):
                    F.add("Look", view, ws.cell(r, c).coordinate, ["look", col, lab, c], v)
        elif lab in ("Bars", "From", "To") and col is not None:
            # the analyst's view of the chart, shown back (named by its label: Look's rows moved on 30 Sep 2026)
            v = ws.cell(r, 3).value
            if v is not None:
                F.add("Look", view, ws.cell(r, 3).coordinate, ["look-setting", col, lab], v)
    F.sweep(ws, "Look", view, max_col=16)
    lk = wb["_look"]
    g = 6
    while lk.cell(4, g).value not in (None, ""):
        fcol, gcol, hcol = (get_column_letter(g + i) for i in range(3))
        name = lk[f"{fcol}4"].value
        for r in range(4, 4 + int(lk[f"{fcol}13"].value)):
            F.add("Look", view, f"_look!{hcol}{r}",
                  ["look-bar", name, float(lk[f"{gcol}{r}"].value), float(lk[f"{fcol}16"].value),
                   r == 3 + int(lk[f"{fcol}13"].value)], lk[f"{hcol}{r}"].value)
        # 30 Sep 2026: `g` is the group's value column (G_VALUE, its 5th): the bar labels are 8 to its right
        # (G_BAR_LABEL), the grey lines' x 5 (G_EX), at rows 52, 54, .. (PCT_TOP)
        for r in range(4, 4 + int(lk[f"{fcol}13"].value)):                 # each bar's short label
            lab = lk.cell(r, g + 8).value
            F.add("Look", view, f"_look!{get_column_letter(g + 8)}{r}",
                  ["look-label", name, float(lk[f"{gcol}{r}"].value)], lab)
        for k in range(5):                                                  # the grey P10..P90 lines' x
            r = 52 + 2 * k
            x = lk.cell(r, g + 5).value
            F.add("Look", view, f"_look!{get_column_letter(g + 5)}{r}", ["look-pline", name, k],
                  x if isinstance(x, (int, float)) else None)
        for r, what in ((18, "low end"), (19, "high end"), (10, "at the code")):
            if lk[f"{fcol}{r}"].value is not None and not (what == "at the code" and lk[f"{fcol}9"].value is None):
                F.add("Look", view, f"_look!{fcol}{r}", ["look-end", name, what], lk[f"{fcol}{r}"].value)
        g += 13
    # the scatters' dots (_dots: row 3 names each column, the pairs from row 4)
    if "_dots" in wb.sheetnames:
        ds = wb["_dots"]
        c = 1
        while ds.cell(3, c).value not in (None, ""):
            xs, ys = ds.cell(3, c).value, ds.cell(3, c + 1).value
            pairs = []
            r = 4
            while ds.cell(r, c).value is not None:
                pairs.append([ds.cell(r, c).value, ds.cell(r, c + 1).value])
                r += 1
            F.add("Look", view, f"_dots!{get_column_letter(c)}4:{get_column_letter(c + 1)}{r - 1}",
                  ["look-dots", f"{ys} against {xs}"], pairs)
            c += 2


#: _pockets' Borderline columns (30 Sep 2026: live.P_SE_BOOK .. P_WORSE_SAID)
SE_BOOK = "Shuffle's standard error of the p-value vs book, after the allowance"
SE_BAND = "Shuffle's standard error of the p-value vs band, after the allowance"
BTXT = ("Borderline: the p-value that decides, when the flag turns on it and it is within 2 standard errors of the "
        "bar")


def seg_of(v):
    """A segment as Start here names it, without the borderline words after it (30 Sep 2026)."""
    return str(v).split(" · borderline (p ")[0] if v is not None else v


def read_border(F, rows, view):
    """Every pocket on _pockets, two-way and split: its borderline p-value as printed (blank when not borderline), the
    one Worse? prints, and the standard error of the p-value that decides. And the rule itself, applied to the
    workbook's own p and SE (docs/statistics.md B2a), which compare.py checks against what the workbook printed."""
    for i, r in enumerate(rows):
        who = ["pk-border", r["Grid"], str(r["Band"]), str(r["Segment"]), r["Rate (key)"]]
        at = f"_pockets!row {i + 4}"
        band = r.get("Judged against its band?") in (True, "TRUE", 1)
        se = r.get(SE_BAND) if band else r.get(SE_BOOK)
        F.add("Pockets", view, at + " (Borderline)", who + ["any"], r.get(BTXT) or "")
        F.add("Pockets", view, at + " (Borderline, for Worse?)", who + ["worse"], r.get("Borderline, for Worse?") or "")
        if isinstance(se, (int, float)):
            F.add("Pockets", view, at + " (SE)", who + ["se"], se)
        F.add("Pockets", view, at + " (the rule on its own p)", ["pk-rule", *who[1:]],
              [r.get(BTXT) or "", r.get("Borderline, for Worse?") or ""])
        F.out[-1]["inputs"] = {"p": r.get("p-value that decides"), "se": se, "flag": r.get("Flag"),
                                 "gap": r.get("Gap that decides"), "profit": r["Rate (key)"] in ("ranr_rate",
                                                                                                 "contribution_rate")}


def read_once(F, wb, view):
    ws = wb["Start here"]
    for c in ("B17", "D17", "F17"):
        F.add("Start here", view, c, ["start-tile", ws.cell(16, ws[c].column).value], ws[c].value)
    r = 20
    while ws.cell(r, 2).value not in (None, "", "No pocket is worse and material.") and r < 25:   # the five largest
        for col, what in ((4, "loans"), (5, "multiple"), (6, "dollars")):
            F.add("Start here", view, ws.cell(r, col).coordinate,
                  ["start-top", r - 19, ws.cell(r, 2).value, seg_of(ws.cell(r, 3).value), what],
                  ws.cell(r, col).value)
        F.add("Start here", view, ws.cell(r, 3).coordinate,
              ["start-top", r - 19, ws.cell(r, 2).value, seg_of(ws.cell(r, 3).value), "segment"], ws.cell(r, 3).value)
        r += 1
    F.sweep(ws, "Start here", view, max_col=10)

    ws = wb["Control"]
    for r in range(15, 20):
        for c in "MNOP":
            F.add("Control", view, f"{c}{r}", ["ladder", r - 14, c], ws[f"{c}{r}"].value)
    for cell in ("E15", "F15", "I15", "E16", "F16", "I16", "E19", "F24", "I24"):
        if ws[cell].value is not None:
            F.add("Control", view, cell, ["control", cell], ws[cell].value)
    F.sweep(ws, "Control", view, max_col=16)

    ws = wb["Columns"]
    _, hid_r = hidden(ws)
    r = 13
    while ws.cell(r, 2).value not in (None, ""):
        col = ws.cell(r, 2).value
        for c, what in ((5, "why"), (6, "blank"), (7, "odd"), (12, "check")):
            v = ws.cell(r, c).value
            if v is not None and has_digit(v) and r not in hid_r:
                F.add("Columns", view, ws.cell(r, c).coordinate, ["columns", col, what], v)
        r += 1
    F.sweep(ws, "Columns", view, max_col=16)

    read_look(F, wb, view)

    ws = wb["Record"]
    for r in range(8, ws.max_row + 1):
        for c in (3, 6):
            v = ws.cell(r, c).value
            if v is not None and has_digit(v):
                lab, k = record_anchor(ws, r, c)
                F.add("Record", view, ws.cell(r, c).coordinate,
                      ["record", "C" if c == 3 else "F", lab, k], v)
    F.sweep(ws, "Record", view, max_col=10)


def main():
    d = Path(sys.argv[1])
    plan = json.loads((d / "plan.json").read_text())
    F = Figures()
    seen = {"pockets": set(), "grids": set(), "pck": set(), "split": set()}
    meta = as_written_meta(d / "as-written.xlsx")
    split_grids, group_grids = set(), set()
    for i, p in enumerate(plan):
        wb = load_workbook(d / "calculated" / p["file"], data_only=True)
        view = p["file"][:-5]
        if i == 0:
            read_once(F, wb, view)
            pk = wb["_pockets"]
            head = [c.value for c in pk[3]]
            rows = [dict(zip(head, [c.value for c in r])) for r in pk.iter_rows(min_row=4, max_col=41)]
            rows = [r for r in rows if r.get("Grid")]
            (d / "wb-pockets.json").write_text(json.dumps(rows, default=str))
            read_border(F, rows, view)
        if tuple(p["pockets"]) not in seen["pockets"]:
            seen["pockets"].add(tuple(p["pockets"]))
            read_pockets(F, wb["Pockets"], view)
        g = (*p["grids"], p.get("filter"), *p.get("cell", ()))
        if g not in seen["grids"]:
            seen["grids"].add(g)
            read_grids(F, wb["Grids"], view, p["grids"][0] not in group_grids, meta)
            group_grids.add(p["grids"][0])
        if tuple(p["split"]) not in seen["split"]:
            seen["split"].add(tuple(p["split"]))
            first = p["split"][0] not in split_grids
            split_grids.add(p["split"][0])
            read_split(F, wb["Split"], view, first)
        if p["pck"] not in seen["pck"]:
            seen["pck"].add(p["pck"])
            read_pck(F, wb["Paid, cost, kept"], wb["_chart"], view)
    (d / "figures.json").write_text(json.dumps(F.out, indent=0, default=str))
    kinds = {}
    for f in F.out:
        kinds[(f["tab"], f["kind"])] = kinds.get((f["tab"], f["kind"]), 0) + 1
    for k, v in sorted(kinds.items()):
        print(k, v)
    print("total", len(F.out))


if __name__ == "__main__":
    main()
