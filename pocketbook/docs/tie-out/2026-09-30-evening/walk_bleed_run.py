"""PocketBook's side (it imports PocketBook freely): the walk of 27 Sep 2026 to its first Run, Where the book bleeds,
made again on this build. It is the full tie-out's walk_both_runs.py (28 Sep) with the window taken out: this
machine's Python has no Tk, so the launcher's Flow -- the object the window draws and every launcher test drives --
is driven directly, the same calls in the same order: pick the file, Set up, cut FICO and ORIG_BAL into bands,
segment by CHANNEL and ASSET_CLASS, split by REV_DEBT, pick the outcome and say yes (new on this build: nothing is
picked for the analyst, f9db3452), Next, Run (refused: answers needed), answer Control and Columns as the walk did,
Run. No screenshot is taken. The walk's scouting Run is not made here (the scouting book is make_scout_book.py's).

    python3 walk_bleed_run.py ../../../src OUT "/tmp/credit/Loan files"
    python3 walk_bleed_run.py ../../../src OUT "/tmp/credit/Flag files" SYS_FLAG "Flag book.csv"

    python3 walk_bleed_run.py ../../../src OUT "/tmp/credit/Bureau files" REV_DEBT "Bureau book.csv" FICO,SHORT_HIST \
        CHANNEL SHORT_HIST

The last three: the band columns, the segment columns, and the columns whose odd values Columns answers Missing
(FICO's -9999 is always answered Missing and RANR_AMT's negatives Real, as on the walk). With a split column and a
file name, the same walk splits by that column instead of REV_DEBT (a category: every
value against the rest of its pocket) -- how the category-split books of this tie-out were made.
"""
import os
import shutil
import sys
from pathlib import Path

SRC, OUT, FOLDER = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve(), Path(sys.argv[3])
SPLIT = sys.argv[4] if len(sys.argv) > 4 else "REV_DEBT"
NAME = sys.argv[5] if len(sys.argv) > 5 else "Consumer book Q3.csv"
BANDS = set(sys.argv[6].split(",")) if len(sys.argv) > 6 else {"FICO", "ORIG_BAL"}
SEGS = set(sys.argv[7].split(",")) if len(sys.argv) > 7 else {"CHANNEL", "ASSET_CLASS"}
MISSING = sys.argv[8].split(",") if len(sys.argv) > 8 else []
HOME = FOLDER.parent
os.environ["HOME"] = str(HOME)
os.environ["POCKETBOOK_MEMORY"] = str(HOME / f"memory-{Path(NAME).stem}.yaml")
os.environ["GIT_CEILING_DIRECTORIES"] = str(HOME)
sys.path.insert(0, str(SRC))

from openpyxl import load_workbook  # noqa: E402
from pocketbook import book as bk, control, deps, launcher  # noqa: E402


def answer(book: Path, **by_key):
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    for key, v in by_key.items():
        ws.cell(row=control.row_of(ws, key), column=control.CHOOSE_COL).value = v
    wb.save(book)


def columns(book: Path, odd: dict, confirm=True):
    wb = load_workbook(book)
    ws = wb["Columns"]
    for r in bk.table_rows(ws):
        key = r[bk.C_QKEY - 1].value
        if isinstance(key, str) and key.count("|") == 2 and key.split("|")[0] in odd:
            r[bk.C_TREAT - 1].value = odd[key.split("|")[0]]
    if confirm:
        ws[bk.CONFIRM_CELL] = "Yes"
    wb.save(book)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    x = FOLDER / NAME
    launcher.PREFS = HOME / "launcher.json"
    flow = launcher.Flow(gate=launcher.AddOns())
    flow.gate.optional = []
    flow.pick(str(x))
    flow.gate.missing, flow.gate.got = [], list(deps.NEEDED)
    flow.after_install(False)
    flow.set_up()
    assert flow.screen() == "L2", flow.message
    # the walk's cut: FICO and ORIG_BAL into bands, CHANNEL and ASSET_CLASS as segments, split by REV_DEBT
    flow.cut = BANDS | ({"REV_DEBT"} if SPLIT == "REV_DEBT" else set())
    flow.seg = SEGS
    flow.click(SPLIT, "c")
    # 30 Sep 2026, build 44734da4: Filter by, a pick of its own (the Grids' "Only loans where"); FILTER=ORIG_YEAR or
    # a category. The category books filter by the column they split by, which is what the filter read before
    if os.environ.get("FILTER"):
        flow.click(os.environ["FILTER"], "d")
        assert flow.filter == os.environ["FILTER"], (flow.filter, flow.message)
    flow.pick_outcome("BAD_FLAG")
    flow.answer_outcome(True)
    print("summary", *flow.summary(), flush=True)
    flow.next()
    assert flow.page == "answer", flow.message
    book = launcher.book_for(x)
    flow.run()
    print("needs", len(flow.needs), flow.needs_head(), flush=True)
    answer(book, worse_at="The smallest significant gap in a typical pocket (suggested)",
           better_at="The smallest significant gap in a typical pocket (suggested)",
           revenue_line="Each pocket's own test (suggested)", confidence="95% sure",
           materiality="1% of the book's total losses", compare_to="The rest of its band",
           min_loans="Enough for 5 expected losses (suggested)", min_events="10 losses")
    columns(book, {"FICO": "Missing", "RANR_AMT": "Real", **{c: "Missing" for c in MISSING}})
    flow.refresh()
    flow.run()
    print("page", flow.page, [n.says for n in flow.needs], flush=True)
    print("first", flow.first_lines(), flush=True)
    assert flow.page == "done"
    shutil.copy(book, OUT / book.name)
    print("kept", OUT / book.name, flush=True)


if __name__ == "__main__":
    main()
