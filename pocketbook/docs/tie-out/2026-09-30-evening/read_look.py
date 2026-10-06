"""PocketBook's road for change 4: Look, read out of the calculated Bureau book workbook (Treat as answered Missing
for FICO's -9999 and SHORT_HIST's negative bureau codes), with read_book.py's own Look reader. Also the survey this
change asks for: every number anywhere in the calculated workbook, every sheet hidden or not, at or below
-1,000,000 -- which should be nowhere but Columns' "Samples" cell.

    python3 read_look.py "WORK/calculated-bureau/Bureau book - PocketBook.xlsx" OUT.json

Nothing here imports PocketBook.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import read_book  # noqa: E402

book, out = Path(sys.argv[1]), Path(sys.argv[2])
wb = load_workbook(book, data_only=True)
F = read_book.Figures()
read_book.read_look(F, wb, "calculated")
# the survey: numbers <= -1,000,000 in any cell of any sheet, and text holding one (a sample of values, a sentence)
low = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, (int, float)) and not isinstance(v, bool) and v <= -1_000_000:
                low.append([ws.title, c.coordinate, v, ws.sheet_state])
            elif isinstance(v, str) and re.search(r"-\s?\d{1,3}(,\d{3}){2,}|-\d{7,}", v):
                low.append([ws.title, c.coordinate, v[:120], ws.sheet_state])
F.add("Look", "calculated", "(every sheet)", ["low-values"], low)
out.write_text(json.dumps(F.out, indent=0, default=str))
print(len(F.out), "figures;", len(low), "cells at or below -1,000,000:", low)
