"""Addendum, 30 Sep 2026: change 9's band-label check again, on the build that cuts dollar bands at whole dollars (the
firm: "Cut at whole dollars is fine"). The evening's 18 band-label DIFFERS were ORIG_BAL's two labels either side of
the $37,950.548 edge, on the six grid books and on the three re-run against 29 September (the same Q3, Flag and
Two-flag workbooks, read under 29 September's views).

For each of the six books: the workbook made again on the build given, by the same walk (walk_bleed_run.py, with
Filter by picked as run-it-all.sh picks it); its Grids tab turned to each ORIG_BAL grid and calculated by LibreOffice
(pocketbook/tests/recalc.py, the repository's own helper, as views.py uses it); every band label read with the loans
its band holds (the Loans block's All column, read_book.py's grid_block); and the loans whose value lies inside the
label's own range counted from the loan file, exactly as by_hand_book.py counts them on a column with cents: a label
names whole dollars, so "a - b" holds a <= value < b + 1. The same count for FICO, a whole-number column (a <= value
<= b), to show it did not move. No PocketBook code on the counting side; the workbooks are PocketBook's.

    python3 recheck_whole_dollars.py ../../../src OUT          # writes OUT/recheck.json and prints one line a label
"""
import csv
import json
import math
import os
import re
import shutil
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2] / "tests"))           # pocketbook/tests, for recalc.py
sys.path.insert(0, str(HERE))
import recalc  # noqa: E402
import read_book  # noqa: E402

SRC, OUT = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
Q3 = HERE.parent / "2026-09-28" / "Consumer book Q3.csv"
BOOKS = {                        # run-it-all.sh's books: file, split, Filter by
    "q3": (Q3, "REV_DEBT", ""),
    "flag": (HERE / "Flag book.csv", "SYS_FLAG", "SYS_FLAG"),
    "two": (HERE / "Two-flag book.csv", "SYS_FLAG", "SYS_FLAG"),
    "grey": (HERE / "Grey book.csv", "REV_DEBT", ""),
    "year": (HERE / "Year book.csv", "REV_DEBT", "ORIG_YEAR"),
    "yearsplit": (HERE / "Year-split book.csv", "ORIG_YEAR", ""),
}


class Figures:
    def add(self, *a):
        pass


def make(b):
    csv_, split, filt = BOOKS[b]
    home = OUT / b
    files = home / "files"
    files.mkdir(parents=True, exist_ok=True)
    shutil.copy(csv_, files / csv_.name)
    made = home / "made"
    if made.exists() and list(made.glob("*.xlsx")):
        return next(made.glob("*.xlsx")), csv_              # made already, on this build (a run picked up again)
    env = {**os.environ, "FILTER": filt, "PYTHONDONTWRITEBYTECODE": "1"}
    r = subprocess.run([sys.executable, "-B", str(HERE / "walk_bleed_run.py"), str(SRC), str(made), str(files), split,
                        csv_.name], capture_output=True, text=True, env=env, timeout=1500)
    if r.returncode:
        raise SystemExit(f"{b}: the walk failed\n{r.stdout[-2000:]}\n{r.stderr[-2000:]}")
    return next(made.glob("*.xlsx")), csv_


def grid_list(book):
    """The Grids tab's Grid dropdown's list, from the workbook's own _choices sheet (as views.py reads it)."""
    ws = load_workbook(book)["_choices"]
    heads = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    c = heads["Grids: Grid"]
    return [ws.cell(r, c).value for r in range(2, ws.max_row + 1) if ws.cell(r, c).value not in (None, "")]


def labels_on(book, grid, work):
    wb = load_workbook(book)
    ws = wb["Grids"]
    read_book.below(ws, "GRID").value = grid
    one = work / f"{re.sub(r'[^A-Za-z0-9]+', '_', grid)}.xlsx"
    wb.save(one)
    done = recalc.recalc_file(one, work / "calculated")
    ws = load_workbook(done, data_only=True)["Grids"]
    rate = read_book.find(ws, [c.value for row in ws.iter_rows(min_row=12, max_row=40, max_col=40) for c in row
                               if isinstance(c.value, str) and c.value.startswith("Rate · ")][0])
    book_ = next((c_.row, c_.column) for row in ws.iter_rows(min_row=12, max_row=40, max_col=40) for c_ in row
                 if isinstance(c_.value, str) and c_.value.startswith("vs the book"))
    band = read_book.find(ws, "vs rest of band")
    loans = (band[0], book_[1])
    assert ws.cell(*loans).value == "Loans", ws.cell(*loans).value
    got = read_book.grid_block(Figures(), ws, "recheck", grid, "", loans, "loans")
    assert rate is not None
    return {rl: v for (rl, cl), (_, v) in got.items() if cl == "All" and rl != "All" and isinstance(v, (int, float))}


def file_values(csv_, col):
    rows = list(csv.DictReader(open(csv_, newline="", encoding="utf-8")))
    out = []
    for r in rows:
        s = r[col].strip()
        if s == "" or (col == "FICO" and float(s) < -1000):      # blank; FICO's -9999 answered Missing
            continue
        out.append(float(s))
    return out


def label_range(lab):
    a, b = lab.split(" - ")
    return float(a.replace(",", "")), float(b.replace(",", ""))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    report = {}
    for b in BOOKS:
        book, csv_ = make(b)
        grids = [g for g in grid_list(book) if str(g).split(" x ")[0] in ("ORIG_BAL", "FICO")
                 and "/" not in str(g)]                             # the two-way grids; a split grid bands alike
        seen = {}
        for g in grids:
            col = g.split(" x ")[0]
            for lab, n in labels_on(book, g, OUT / b).items():
                if " - " not in lab:
                    continue                                        # (blank), (marked missing): bands of their own
                prev = seen.setdefault((col, lab), n)
                assert prev == n, (b, g, lab, prev, n)             # the same band on every grid of its column
        vals = {c: file_values(csv_, c) for c in ("ORIG_BAL", "FICO")}
        rows = []
        for (col, lab), n in sorted(seen.items(), key=lambda kv: (kv[0][0], label_range(kv[0][1])[0])):
            a, z = label_range(lab)
            cents = not all(v.is_integer() for v in vals[col])
            inside = sum(1 for v in vals[col] if (a <= v < z + 1 if cents else a <= v <= z))
            verdict = "TIED" if inside == n else "DIFFERS"
            rows.append({"column": col, "label": lab, "workbook": n, "loan file": inside, "verdict": verdict})
            print(f"{b:10} {col:9} {lab:20} workbook {n:>6,}  loan file {inside:>6,}  {verdict}", flush=True)
        edge = [v for v in vals["ORIG_BAL"] if v == 37950.99]
        report[b] = {"labels": rows, "loan of 37,950.99": len(edge),
                     "its band": next((r["label"] for r in rows if r["column"] == "ORIG_BAL"
                                       and label_range(r["label"])[0] <= math.floor(37950.99)
                                       <= label_range(r["label"])[1]), None)}
    json.dump(report, open(OUT / "recheck.json", "w"), indent=1)
    n = sum(1 for b in report.values() for r in b["labels"])
    bad = sum(1 for b in report.values() for r in b["labels"] if r["verdict"] != "TIED")
    print(f"{n} labels read, {n - bad} TIED, {bad} DIFFERS")


if __name__ == "__main__":
    main()
