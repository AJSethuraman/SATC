"""Photograph the analyst's workbook as LibreOffice Calc shows it, recalculated on load (as Excel would), one view
per shot. Run under xvfb-run:

    xvfb-run -a -s "-screen 0 1600x1300x24" python3 workbook-shots.py
    xvfb-run -a -s "-screen 0 2200x1300x24" python3 workbook-shots.py pockets-badloans grids-loans

(the two wide views need the wider screen). Each view is a copy of the workbook with the sheet, the top-left cell and (for
Pockets) the Measure dropdown set the way an analyst would set them; nothing else is touched."""
import os
import shutil
import signal
import subprocess
import sys
import time
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
BOOK = HERE / "Consumer book Q3 - PocketBook.xlsx"
PROFILE = Path(os.environ.get("TMPDIR", "/tmp")) / "tieout-lo-profile"
SHOTS = HERE / "raw"

XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema"
           xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop></item>
<item oor:path="/org.openoffice.Office.Common/Misc"><prop oor:name="ShowTipOfTheDay" oor:op="fuse"><value>false</value></prop></item>
<item oor:path="/org.openoffice.Setup/Office"><prop oor:name="FirstStartWizardCompleted" oor:op="fuse"><value>true</value></prop></item>
</oor:items>
"""

VIEWS = {
    # name: (sheet, top-left cell, selected cell, {cell: value set as the analyst would})
    "start-here": ("Start here", "A2", "F20", {}),
    "pockets-chargeoffs": ("Pockets", "A21", "J21", {("Pockets", "C18"): "Charge-offs"}),
    "pockets-badloans": ("Pockets", "A21", "L21", {("Pockets", "C18"): "Bad loans"}),
    "grids-loans": ("Grids", "A14", "Q35", {}),
    "grids-booked": ("Grids", "A38", "E63", {}),
    "record": ("Record", "A2", "C25", {}),
}


def main():
    (PROFILE / "user").mkdir(parents=True, exist_ok=True)
    (PROFILE / "user" / "registrymodifications.xcu").write_text(XCU, encoding="utf-8")
    SHOTS.mkdir(exist_ok=True)
    only = sys.argv[1:] or list(VIEWS)
    for name in only:
        sheet, top, sel, sets = VIEWS[name]
        d = Path(os.environ.get("TMPDIR", "/tmp")) / "tieout-views" / name
        d.mkdir(parents=True, exist_ok=True)
        f = d / BOOK.name
        shutil.copy(BOOK, f)
        wb = load_workbook(f)
        for (sh, cell), v in sets.items():
            wb[sh][cell].value = v
        ws = wb[sheet]
        wb.active = wb.sheetnames.index(sheet)
        for w in wb.worksheets:
            w.sheet_view.tabSelected = w.title == sheet
        if ws.sheet_view.pane is not None:
            ws.sheet_view.pane.topLeftCell = top      # the scrolling part under the frozen rows
        else:
            ws.sheet_view.topLeftCell = top
        ws.sheet_view.selection[0].activeCell = sel
        ws.sheet_view.selection[0].sqref = sel
        wb.save(f)
        p = subprocess.Popen(["soffice", f"-env:UserInstallation={PROFILE.as_uri()}", "--norestore", "--nologo",
                              "--calc", str(f)], start_new_session=True,
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        time.sleep(float(os.environ.get("WAIT", "14")))
        subprocess.run(["import", "-window", "root", str(SHOTS / f"{name}.png")], check=False)
        os.killpg(p.pid, signal.SIGTERM)
        try:
            p.wait(timeout=20)
        except subprocess.TimeoutExpired:
            os.killpg(p.pid, signal.SIGKILL)
        time.sleep(2)
        print("shot", name, flush=True)


if __name__ == "__main__":
    main()
