"""The PocketBook road: read the figures out of the workbook the analyst opens, after its formulas are calculated.

    python3 read-workbook.py "Consumer book Q3 - PocketBook.xlsx"

The workbook is recalculated by LibreOffice through the repository's own helper (pocketbook/tests/recalc.py: a
headless conversion under a profile set to "recalculate Excel files on load: always"), as Excel would on opening.
Then openpyxl reads the calculated values, cell by cell. Nothing here imports PocketBook or re-runs it.

Pockets shows one measure at a time (the MEASURE dropdown, Pockets!C18). The file opens on Bad loans; for
Charge-offs, a copy is made with C18 set to "Charge-offs" -- the one thing an analyst does to see that view --
and that copy is recalculated too.
"""
import shutil
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[2] / "tests"))          # pocketbook/tests
import recalc  # noqa: E402


def calculated(book: Path, measure: str | None = None):
    work = Path(tempfile.mkdtemp(prefix="tieout-"))
    copy = work / book.name
    shutil.copy(book, copy)
    if measure:
        wb = load_workbook(copy)
        wb["Pockets"]["C18"].value = measure
        wb.save(copy)
    return load_workbook(recalc.recalc_file(copy, work / "calculated"), data_only=True)


def show(wb, sheet, cells):
    ws = wb[sheet]
    for c in cells:
        print("  %s!%-4s %r" % (sheet, c, ws[c].value))


book = Path(sys.argv[1]).resolve()
wb = calculated(book)
print("As the workbook opens (Pockets on Bad loans)")
show(wb, "Start here", ["C1", "B20", "C20", "D20", "E20", "F20"])
show(wb, "Pockets", ["C18", "C21", "D21", "F21", "G21", "H21", "I21", "J21", "K21", "L21"])
show(wb, "Pockets", ["C37", "D37", "K37", "L37"])
show(wb, "Grids", ["B13", "E45"])
print("  Grids loans block, FICO x CHANNEL (rows 28-35: band; N Branch, O Broker, P Online, Q All)")
g = wb["Grids"]
for r in range(28, 36):
    print("    %-18s" % g.cell(r, 13).value, [g.cell(r, c).value for c in range(14, 18)])
print("  Grids booked dollars by pocket (rows 44-63: D loans, E booked dollars)")
for r in range(44, 64):
    print("    %-18s %-8s" % (g.cell(r, 2).value, g.cell(r, 3).value), g.cell(r, 4).value, g.cell(r, 5).value)
show(wb, "Record", ["C10", "C11", "C16", "C25", "C47", "C70"])

wb = calculated(book, "Charge-offs")
print("\nPockets with MEASURE set to Charge-offs")
show(wb, "Pockets", ["C18", "C21", "D21", "F21", "G21", "H21", "I21", "J21", "K21", "L21"])
