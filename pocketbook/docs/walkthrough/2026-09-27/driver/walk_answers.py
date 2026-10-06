"""The walk of 27 Sep 2026 again, scripted: the real window under xvfb, and the workbook at each stage.

    xvfb-run -a -s "-screen 0 1000x760x24" venv/bin/python walk_answers.py SRC OUT [bleed|all]

SRC is a pocketbook/src folder (the build to walk); OUT gets shots/*.png (the window), books/<stage>.xlsx (the
workbook as the analyst left it at each stage). The actions are the procedure's, step by step."""
import os
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

SRC, OUT = Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve()
WHAT = sys.argv[3] if len(sys.argv) > 3 else "all"
HOME = OUT / "home"
os.environ["HOME"] = str(HOME)
os.environ["POCKETBOOK_MEMORY"] = str(HOME / "memory.yaml")
os.environ["GIT_CEILING_DIRECTORIES"] = str(OUT)
sys.path.insert(0, str(SRC))

import tkinter as tk  # noqa: E402

from openpyxl import load_workbook  # noqa: E402
from pocketbook import control, launcher, synth  # noqa: E402

SHOTS, BOOKS = OUT / "shots", OUT / "books"


def settle(root, n=8):
    for _ in range(n):
        root.update()
        time.sleep(0.05)


def shot(root, name):
    settle(root)
    g = f"{root.winfo_width()}x{root.winfo_height()}+{root.winfo_rootx()}+{root.winfo_rooty()}"
    subprocess.run(["import", "-window", "root", "-crop", g, "+repage", str(SHOTS / f"{name}.png")], check=False)
    print("shot", name, flush=True)


def keep(book: Path, stage: str):
    shutil.copy(book, BOOKS / f"{stage}.xlsx")
    print("kept", stage, flush=True)


def answer(book: Path, **by_key):
    wb = load_workbook(book)
    ws = wb[control.SHEET]
    for key, v in by_key.items():
        ws.cell(row=control.row_of(ws, key), column=control.CHOOSE_COL).value = v
    wb.save(book)


def columns(book: Path, odd: dict, confirm=True):
    from pocketbook import book as bk
    wb = load_workbook(book)
    ws = wb["Columns"]
    for r in bk.table_rows(ws):
        key = r[bk.C_QKEY - 1].value
        if isinstance(key, str) and key.count("|") == 2 and key.split("|")[0] in odd:
            r[bk.C_TREAT - 1].value = odd[key.split("|")[0]]
    if confirm:
        ws[bk.CONFIRM_CELL] = "Yes"
    wb.save(book)


def main():
    for d in (SHOTS, BOOKS, HOME):
        d.mkdir(parents=True, exist_ok=True)
    x = OUT / "Consumer book Q3.csv"
    if not x.exists():
        p = synth.write_extract(OUT / "tmp", n=8000)
        Path(p).rename(x)
    launcher.PREFS = HOME / "launcher.json"
    root = tk.Tk()
    w = launcher.build(root)
    flow = w["flow"]
    render = w["render"]
    root.geometry("720x560+0+0")
    settle(root)
    w["extract"].set(str(x))
    flow.set_up()
    render()
    shot(root, "s06-choose-tests-bleed")
    # the walk's cut: FICO and ORIG_BAL into bands, CHANNEL and ASSET_CLASS as segments, split by REV_DEBT
    flow.cut, flow.seg = {"FICO", "ORIG_BAL", "REV_DEBT"}, {"CHANNEL", "ASSET_CLASS"}
    flow.click("REV_DEBT", "c")
    render()
    shot(root, "s07-split-rev-debt")
    flow.next()
    render()
    shot(root, "s08-answer-in-workbook")
    book = launcher.book_for(x)
    keep(book, "01-written")
    flow.run()
    render()
    shot(root, "s09-run-before-answering")
    print("needs", len(flow.needs), flow.needs_head(), flush=True)
    flow.page = "answer"
    answer(book, worse_at="The smallest significant gap in a typical pocket (suggested)",
           better_at="The smallest significant gap in a typical pocket (suggested)",
           revenue_line="Each pocket's own test (suggested)", confidence="95% sure",
           materiality="1% of the book's total losses", compare_to="The rest of its band",
           min_loans="Enough for 5 expected losses (suggested)", min_events="10 losses")
    columns(book, {"FICO": "Missing", "RANR_AMT": "Real"})
    keep(book, "02-answered")
    # Excel's owner file, while the workbook is open
    owner = book.with_name("~$" + book.name)
    owner.write_text("x")
    flow.refresh()
    render()
    shot(root, "s11-workbook-open")
    owner.unlink()
    flow.refresh()
    flow.run()
    render()
    shot(root, "s19-run-finished")
    print("first", flow.first_lines() if hasattr(flow, "first_lines") else None, flush=True)
    keep(book, "03-run1")
    # a Changes-now answer, live
    b4 = BOOKS / "04-live.xlsx"
    shutil.copy(book, b4)
    answer(b4, worse_at="2 times")
    # a Needs-a-Run answer, waiting
    answer(book, min_loans="100 loans")
    keep(book, "05-pending")
    flow.run()
    render()
    shot(root, "s20-run-again")
    keep(book, "06-run2")
    if WHAT == "bleed":
        root.destroy()
        return
    # ---- test new variables
    flow.go(2)
    flow.set_mode("new")
    flow.gate.optional = ["scikit-learn"]           # as on the walk: scikit-learn not installed yet
    render()
    shot(root, "s22-test-new-variables")
    flow.gate.optional, flow.gate.got = [], ["scikit-learn"]
    flow.after_install(optional=True)
    render()
    shot(root, "s23b-sklearn-installed")
    flow.test = ["REV_DEBT", "ASSET_CLASS", "CHANNEL", "ORIG_BAL"]
    flow.hold = ["FICO"]
    flow.outcome = "BAD_FLAG"
    render()
    shot(root, "s24-candidates-ticked")
    flow.next()
    render()
    shot(root, "s25-answer-new")
    keep(book, "07-new-written")
    flow.run()
    render()
    shot(root, "s26-scout-finished")
    keep(book, "08-scout")
    pre = book.with_name("Consumer book Q3 - pre-spec.yaml")
    text = pre.read_text(encoding="utf-8")
    m = re.search(r"bins: \[[^\]]*\]", text)
    print("bins", m.group(0) if m else None, flush=True)
    pre.write_text(text.replace(m.group(0), "bins: [10000, 16000]", 1), encoding="utf-8")
    flow.run()
    render()
    shot(root, "s27-edited-prespec-run")
    keep(book, "09-edited")
    root.destroy()


if __name__ == "__main__":
    import multiprocessing
    multiprocessing.freeze_support()
    main()
