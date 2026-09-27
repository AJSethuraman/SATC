"""What an analyst sees in a tab's cells and dropdowns (openpyxl; formulas as written, lists resolved).
    python3 dump.py BOOK.xlsx SHEET [first_row last_row]"""
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1])
ws = wb[sys.argv[2]]
a, b = (int(sys.argv[3]), int(sys.argv[4])) if len(sys.argv) > 4 else (1, ws.max_row)
dv = {}
for v in ws.data_validations.dataValidation:
    for rng in v.sqref.ranges:
        for row in ws.iter_rows(min_row=rng.min_row, max_row=rng.max_row, min_col=rng.min_col, max_col=rng.max_col):
            for c in row:
                f = v.formula1 or ""
                if f.startswith("=") or "!" in f or "$" in f:
                    ref = f.lstrip("=")
                    try:
                        sh, cells = ref.split("!")
                        sh = sh.strip("'")
                        vals = [x.value for r in wb[sh][cells.replace("$", "")] for x in r if x.value is not None]
                        f = " | ".join(map(str, vals))
                    except Exception:
                        pass
                dv[c.coordinate] = f
for row in ws.iter_rows(min_row=a, max_row=b):
    cells = [(c.coordinate, c.value) for c in row if c.value is not None]
    lists = [(c.coordinate, dv[c.coordinate]) for c in row if c.coordinate in dv]
    if cells or lists:
        print(row[0].row, cells)
        for k, v in lists:
            print("     LIST", k, ":", v)
