"""Type answers into the workbook, as the analyst would in Excel, then save.
    python3 edit.py BOOK.xlsx 'Control!C15=value' 'Columns!C3=Yes' 'odd:FICO=Missing' ...
A value of '' clears the cell. 'odd:COL=Real' answers Treat as on the odd-value row for that column."""
import sys
from openpyxl import load_workbook

wb = load_workbook(sys.argv[1])
for a in sys.argv[2:]:
    if a.startswith("odd:"):
        col, v = a[4:].split("=", 1)
        ws = wb["Columns"]
        head = next(r for r in ws.iter_rows() if r[1].value == "Column" or r[0].value == "Column")
        hrow = head[0].row
        names = {c.value: c.column for c in ws[hrow] if c.value}
        cname = next(k for k in names if str(k).startswith("Column"))
        treat = next(k for k in names if str(k).startswith("Treat as"))
        for r in ws.iter_rows(min_row=hrow + 1):
            if r[names[cname] - 1].value == col:
                r[names[treat] - 1].value = v
                print("Columns", r[names[treat] - 1].coordinate, "=", v)
                break
        continue
    ref, v = a.split("=", 1)
    sh, cell = ref.split("!")
    try:
        v = float(v) if v.replace(".", "", 1).isdigit() else v
    except ValueError:
        pass
    wb[sh][cell] = v if v != "" else None
    print(sh, cell, "=", repr(v))
wb.save(sys.argv[1])
