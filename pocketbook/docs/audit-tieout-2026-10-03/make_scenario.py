"""Builds the tie-out scenario with PocketBook itself: a synthetic extract, the main workbook and the audit workbook.

This is the production side of the tie-out, and the only file in this folder that imports pocketbook. tieout.py,
beside it, imports nothing from pocketbook.

    python make_scenario.py OUT_DIR

The extract is synth.write_extract(OUT_DIR, n=50000, seed=7), so the same file is produced every time (tieout.py
prints its SHA-256). One band column (FICO, cut by PocketBook into five bands of about equal loans: edges 653, 686,
713, 746) by one segment column (CHANNEL).
Control is answered as the test suite answers it, plus "Also write the audit workbook?" = Yes. Shuffles: PocketBook's
default, 10,000.
"""

import hashlib
import os
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path[:0] = [str(HERE.parents[1] / "src"), str(HERE.parents[1] / "tests")]

from openpyxl import load_workbook  # noqa: E402

from pocketbook import audit, book, choices as ch, control, perm, synth  # noqa: E402
from test_book import _answer  # noqa: E402

N, SEED = 50_000, 7


def main(out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    os.environ["POCKETBOOK_MEMORY"] = str(out / "memory.yaml")
    assert perm.SHUFFLES == 10_000
    src = synth.write_extract(out, n=N, seed=SEED)
    made = book.set_up(src, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                               outcome="BAD_FLAG"))
    _answer(made.book)                          # Control as the suite answers it; FICO's -9999 treated as missing
    wb = load_workbook(made.book)
    for r in wb[control.SHEET].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == audit.KEY:
            r[control.CHOOSE_COL - 1].value = "Yes"
    for r in book.table_rows(wb["Columns"]):     # every other odd value is real
        k = r[book.C_QKEY - 1].value
        if isinstance(k, str) and k.count("|") == 2 and not r[book.C_TREAT - 1].value:
            r[book.C_TREAT - 1].value = "Real"
    wb.save(made.book)
    ran = book.run(made.book)
    print("\n".join(ran.lines))
    assert ran.ok
    print("extract ", src, hashlib.sha256(src.read_bytes()).hexdigest())
    print("main    ", made.book)
    print("audit   ", audit.path_for(made.book))


if __name__ == "__main__":
    main(Path(sys.argv[1]))
