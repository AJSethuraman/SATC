"""The pictures in TIEOUT.html: the loan file's first rows with the excluded values ringed, and the audit workbook's
One pocket sheet as LibreOffice calculates it.

    python pictures.py RUN_DIR OUT_DIR

The One pocket picture is taken from the calculated copy compare.py leaves in RUN_DIR/compare/audit-default: every
sheet but One pocket is removed from a values-only copy, which LibreOffice prints to PDF and pdftoppm turns into a PNG.
"""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont

MONO = "/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf"
SANS_B = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
RED = (200, 16, 16)


def csv_picture(csv_path: Path, out: Path) -> None:
    """The header and first six records of the loan file, as text, with the five values PocketBook must exclude or
    set aside ringed in red."""
    lines = csv_path.read_text(encoding="utf-8").splitlines()[:7]
    cols = ["LOAN_NBR", "FICO", "CHANNEL", "ORIG_BAL", "BAD_FLAG", "GCO_AMT", "RANR_AMT"]
    rows = [r.split(",")[:7] for r in lines]
    font = ImageFont.truetype(MONO, 26)
    head = ImageFont.truetype(SANS_B, 24)
    widths = [max(len(r[i]) for r in rows) + 3 for i in range(7)]
    cw, ch = font.getbbox("M")[2], 44
    W = sum(widths) * cw + 60
    H = ch * len(rows) + 110
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    d.text((30, 18), f"{csv_path.name} (first 7 lines of 50,001), as the file holds them", fill=(20, 20, 20), font=head)
    y0 = 70
    marks = {(2, 1), (3, 1), (4, 5), (5, 3), (6, 4)}            # (line, column): -9999, blank, #N/A, blank, 2
    for li, r in enumerate(rows):
        x = 30
        y = y0 + li * ch
        if li == 0:
            d.rectangle([20, y - 4, W - 20, y + ch - 8], fill=(232, 236, 242))
        for ci, v in enumerate(r):
            d.text((x, y), v, fill=(20, 20, 20), font=font)
            if (li + 1, ci) in marks:
                w = max(len(v), 2) * cw
                d.ellipse([x - 16, y - 8, x + w + 16, y + ch - 6], outline=RED, width=5)
            x += widths[ci] * cw
    img.save(out)


def one_pocket_picture(run: Path, out: Path, work: Path) -> None:
    src = next((run / "compare" / "audit-default").glob("*.xlsx"))
    work.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(src, data_only=True)
    for name in list(wb.sheetnames):
        if name != "One pocket":
            del wb[name]
    ws = wb["One pocket"]
    ws.print_area = "B1:H56"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    p = work / "one-pocket.xlsx"
    wb.save(p)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(work), str(p)],
                   capture_output=True, timeout=300)
    subprocess.run(["pdftoppm", "-png", "-r", "220", "-singlefile", str(work / "one-pocket.pdf"),
                    str(work / "one-pocket")], check=True)
    img = Image.open(work / "one-pocket.png").convert("RGB")
    box = Image.eval(img, lambda v: 255 - v).getbbox()           # crop the page's white margin
    img = img.crop((box[0] - 10, box[1] - 10, box[2] + 10, box[3] + 10))
    # ring the Ties? column in red: the column the audit's own verdict sits in
    W, H = img.size
    d = ImageDraw.Draw(img)
    x0, x1 = int(W * 0.7175), int(W * 0.748)
    d.rounded_rectangle([x0, int(H * 0.16), x1, int(H * 0.955)], radius=18, outline=RED, width=6)
    img.save(out)


def sheet_picture(src: Path, sheet: str, area: str, out: Path, work: Path, dpi: int = 150) -> None:
    """One sheet of a calculated workbook, values only, as LibreOffice prints the area `area`: used by the re-run
    section (build_rerun.py)."""
    work.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(src, data_only=True)
    for name in list(wb.sheetnames):
        if name != sheet:
            del wb[name]
    ws = wb[sheet]
    ws.print_area = area
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    p = work / f"{out.stem}.xlsx"
    wb.save(p)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(work), str(p)],
                   capture_output=True, timeout=300)
    subprocess.run(["pdftoppm", "-png", "-r", str(dpi), "-singlefile", str(work / f"{out.stem}.pdf"),
                    str(work / out.stem)], check=True)
    img = Image.open(work / f"{out.stem}.png").convert("RGB")
    box = Image.eval(img, lambda v: 255 - v).getbbox()
    img.crop((box[0] - 10, box[1] - 10, box[2] + 10, box[3] + 10)).save(out)


if __name__ == "__main__":
    run, outd = Path(sys.argv[1]), Path(sys.argv[2])
    outd.mkdir(parents=True, exist_ok=True)
    csv_picture(run / "loans.csv", outd / "source-loans-csv.png")
    one_pocket_picture(run, outd / "audit-one-pocket.png", run / "pictures")
    print("pictures in", outd)
