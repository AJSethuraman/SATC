"""Significant at N%? says which side, its p-value has a column of its own, and Pockets sorts for investigation.

The firm, 17 Oct 2026, on significantly better pockets: direction labels, "Yes, worse (p = 1.2%)" and "Yes, better
(p = 0.08%)", Borderline and No as they are, and "i want it sorted in a sensical manner for investigation and
slicers that get added would need to adapt i guess". So the category and the p-value sit in two columns, each with
few distinct values to filter or slice on, the p-value a number; and the rows come in seven groups, worst first:
Yes, worse; Borderline losing more; No losing more; no test; No losing less; Borderline losing less; Yes, better,
each by its excess, so the pocket that loses least against the rest comes last."""

from __future__ import annotations

import pytest
from openpyxl import load_workbook

import tabs
from conftest import TEST_SHUFFLES
from pocketbook import book, choices as ch, engine, live, perm, results, synth
from test_book import _answer


@pytest.fixture(scope="module")
def built(tmp_path_factory):
    """The synthetic book with the line-size plant and a charge-off date, cut as the example book is, after a Run,
    and Pockets calculated on each measure the order test reads."""
    d = tmp_path_factory.mktemp("sigcols")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
        x = synth.write_extract(d, n=8000, chargeoff=True, big_lines=True)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO", "ORIG_BAL"),
                                                segments=("CHANNEL", "ASSET_CLASS"), outcome="BAD_FLAG"))
        assert out.ok, out.lines
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    tabs_ = {m: tabs.calculated(tabs.choose(out.book, d / f"{m[:3]}.xlsx", results.POCKETS, measure=m),
                                results.POCKETS) for m in ("GCOs ($)", "Bad loans", "RANR")}
    return {"b": out.book, "written": load_workbook(out.book)[results.POCKETS], "tabs": tabs_}


# --------------------------------------------------------------------------
# the direction


@pytest.mark.parametrize("word, gap, p, profit, said", [
    (engine.WORSE, 1.6, 0.012, False, "Yes, worse"),
    (engine.BETTER, 0.5, 0.0008, False, "Yes, better"),
    (engine.IN_LINE, 1.1, 0.0001, False, "Yes, worse"),            # the side is the gap's, whatever Worse? reads
    (engine.WORSE, -0.004, 0.001, True, "Yes, worse"),             # profit: short of the rest is worse
    (engine.BETTER, 0.004, 0.001, True, "Yes, better"),
    (engine.UNSURE_WORSE, 1.6, 0.3, False, "No"),
    (engine.UNSURE_BETTER, 0.7, 0.3, False, "No"),
])
def test_a_significant_pocket_says_which_side_it_is_on(word, gap, p, profit, said):
    assert live.sig_said(word, gap, p, False, profit, 0.05) == said
    assert live.sig_said(word, gap, p, True, profit, 0.05) == live.SIG_BORDER      # Borderline as it was


def test_the_tab_reads_yes_worse_and_yes_better_by_the_excess(built):
    """On the tab, Yes, worse sits on a pocket over the rest and Yes, better on one under it."""
    seen = set()
    for ws in built["tabs"].values():
        for x in tabs.pockets(ws):
            if x["chance"] == live.SIG_WORSE:
                assert x["excess"] > 0, x
            elif x["chance"] == live.SIG_BETTER:
                assert x["excess"] < 0, x
            seen.add(x["chance"])
    assert {live.SIG_WORSE, live.SIG_BETTER} <= seen, seen


# --------------------------------------------------------------------------
# the p-value column


def test_the_p_value_is_a_number_in_its_own_column_right_beside_the_category(built):
    assert results.K_PV == results.K_P + 1
    ws = built["tabs"]["GCOs ($)"]
    h = tabs.header_row(ws, results.K_NUM, "#")
    assert ws.cell(row=h, column=results.K_P).value == "Significant at 95%?"
    assert ws.cell(row=h, column=results.K_PV).value == live.P_HEAD == "p-value"
    rows = tabs.pockets(ws)
    with_p = [x for x in rows if x["chance"] in (live.SIG_WORSE, live.SIG_BETTER, live.SIG_BORDER, live.SIG_NO)]
    assert with_p and all(isinstance(x["p"], float) for x in with_p)
    assert all(x["p"] is None for x in rows if x not in with_p)
    # the category column holds only the category: a handful of values, none carrying a p-value
    assert len({x["chance"] for x in rows}) <= 7 and not any("(p " in str(x["chance"]) for x in rows)
    # shown as a percentage: the cell's own format, and the conditional formats for the rest of the precision
    w = built["written"]
    first = h + 1
    # the firm's precision as far as a number format's two conditions carry it, written out rather than read back
    assert w.cell(row=first, column=results.K_PV).number_format == '[<0.0001]"under 0.01%";[<0.01]0.00%;0.0%' == live.P_FMT
    # and the house format every other tab's p-value uses, so the two styles never sit side by side
    assert live.P_FMT == results.P_FMT
    rules = [(r.formula[0], r.dxf.numFmt.formatCode if r.dxf and r.dxf.numFmt else None)
             for rng in w.conditional_formatting if str(rng.sqref).startswith(results.col(results.K_PV))
             for r in rng.rules]
    want = live.p_rules(f"{results.col(results.K_PV)}{first}")
    line = f'${results.col(results.K_NUM)}{first}<>""'
    assert [(f, fmt) for f, fmt in rules if fmt] == [(f"AND({line},{f})", fmt) for f, fmt in want]


def test_excels_filter_arrows_cover_every_heading_the_p_value_included(built):
    w = built["written"]
    h = tabs.header_row(w, results.K_NUM, "#")
    ref = w.auto_filter.ref
    # to the hidden index columns, so a sort carries each row's pocket (tests/test_sorting_2026_10_17.py)
    assert ref and ref.startswith(f"{results.col(results.K_NUM)}{h}:{results.col(results.K_ROW)}"), ref
    lo, hi = (results.col(c) for c in (results.K_NUM, results.K_HOLDS))
    assert lo <= results.col(results.K_P) <= hi and results.col(results.K_PV) <= hi


# --------------------------------------------------------------------------
# the order


def _group(x) -> int:
    return live.sig_group(x["chance"], (x["excess"] or 0) > 0)


def test_pockets_come_in_the_seven_groups_worst_first_each_by_its_excess(built):
    """Every measure's rows run through the groups in order, and within a group by excess, largest first, so in the
    losing-less groups the pocket that loses least comes last; GCOs ($) on this book has all seven groups."""
    for measure, ws in built["tabs"].items():
        rows = tabs.pockets(ws)
        groups = [_group(x) for x in rows]
        assert groups == sorted(groups), (measure, groups)
        for g in set(groups):
            ex = [x["excess"] for x in rows if _group(x) == g and x["excess"] is not None]
            assert ex == sorted(ex, reverse=True), (measure, g, ex)
    gco = tabs.pockets(built["tabs"]["GCOs ($)"])
    assert {_group(x) for x in gco} == set(range(1, 8)), sorted({_group(x) for x in gco})
    assert gco[-1]["chance"] == live.SIG_BETTER and gco[0]["chance"] == live.SIG_WORSE


def test_the_groups_are_the_firms_order():
    assert [live.sig_group(w, m) for w, m in ((live.SIG_WORSE, True), (live.SIG_BORDER, True), (live.SIG_NO, True),
                                              (live.FEW_LOSSES, True), (live.NO_REST, False), (live.NOT_MORE, False),
                                              (results.MIXED + "FICO", True), (live.SIG_NO, False),
                                              (live.SIG_BORDER, False), (live.SIG_BETTER, False))] == \
        [1, 2, 3, 4, 4, 4, 4, 5, 6, 7]


def test_the_note_says_the_order_and_that_control_does_not_move_it(built):
    ws = built["tabs"]["GCOs ($)"]
    note = " ".join(str(c.value) for row in ws.iter_rows(max_row=20) for c in row if isinstance(c.value, str))
    assert "Sorted as of the last Run in seven groups: Yes, worse;" in note
    assert "the pocket that loses least against the rest comes last" in note and "but not the order" in note
