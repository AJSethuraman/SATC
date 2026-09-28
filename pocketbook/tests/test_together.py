"""Every candidate together (OC-51; docs/statistics.md B10): one logistic regression of the outcome on every
shortlisted input's groups at once, on the held-back loans, with the held-fixed columns' pockets as control dummies.

The firm, 27 Sep 2026: *"shouldn't it regress all of those identified variables if it actually deems them
important? import --> tree runs --> tree guesses on 2024 data if 2022-2023 are used to build branches --> regress
shortlist?"* and then *"yes build it out"*.

The books are test_shortlist's: the first synthetic book with income / sales (cliffs planted below 0.1 and above
2.0), UTIL (planted above 0.9) and TENURE (nothing planted), confirming a saved shortlist with FICO and CHANNEL held
fixed; and the same book with UTIL_COPY, a near copy of UTIL drawn from a stream of its own. Every figure the tests
compare with is worked out here from the extract: the design by hand, the fit by scikit-learn, the allowance by
Benjamini-Hochberg's own steps."""

import bisect
import csv
import math
import random
from datetime import date

import numpy as np
import pytest

from pocketbook import confirm_tab, joint, prevalence, synth
from test_shortlist import BASE, INCOME, UTIL, TENURE, _bh, _book, _run, _write_spec, three  # noqa: F401
from test_book_dates import _choose
from test_shortlist import NOT_CUT

HOLD = (date(2024, 1, 1), date(2024, 12, 31))
COPY = {"column": "UTIL_COPY", "bins": [0.9], "reference": 0}
UTIL_ONE = {"column": "UTIL", "bins": [0.9], "reference": 0}


# --------------------------------------------------------------------------
# By hand, from the extract


def _rows(x):
    with open(x, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _value(r, column):
    if column == "income_to_sales":
        return None if r["SALES"] in ("", "0") else float(r["INCOME"]) / float(r["SALES"])
    try:
        return float(r[column])
    except ValueError:
        return None


def _design(res, x, terms, bins):
    """The held-back loans as the joint model sees them, built here: each loan's pocket (FICO's band, cut at the
    run's own edges, a missing or blank score a pocket of its own, crossed with CHANNEL) and its group of each
    input; the pockets where no loan, or every one, went bad left out. Returns (X, y, how many pockets, the column of
    each (input, group))."""
    edges = prevalence.edges_of(res, "FICO")[0]
    loans = []
    for r in _rows(x):
        d = date.fromisoformat(r["ORIG_DATE"])
        if not HOLD[0] <= d <= HOLD[1] or r["BAD_FLAG"] not in ("0", "1"):
            continue
        vals = [_value(r, t) for t in terms]
        if any(v is None for v in vals):
            continue
        fico = r["FICO"]
        band = "blank" if fico == "" else "missing" if float(fico) < -1000 else bisect.bisect_right(edges,
                                                                                                    float(fico))
        loans.append(((band, r["CHANNEL"]), [bisect.bisect_right(bins[t], v) for t, v in zip(terms, vals)],
                      int(r["BAD_FLAG"])))
    per = {}
    for p, _, y in loans:
        per.setdefault(p, []).append(y)
    live = sorted((p for p, ys in per.items() if 0 < sum(ys) < len(ys)), key=repr)
    loans = [x_ for x_ in loans if x_[0] in live]
    cols = [(i, g) for i, t in enumerate(terms) for g in range(len(bins[t]) + 1) if g != REFS[t]]
    X = np.zeros((len(loans), len(live) + len(cols)))
    for n, (p, gs, _) in enumerate(loans):
        X[n, live.index(p)] = 1.0
        for c, (i, g) in enumerate(cols):
            X[n, len(live) + c] = float(gs[i] == g)
    return X, np.array([y for _, _, y in loans]), len(live), cols


REFS = {"income_to_sales": 2, "UTIL": 0, "TENURE": 1, "UTIL_COPY": 0}


def _loglik(model, X, y):
    p = model.predict_proba(X)[:, 1]
    return float(np.sum(y * np.log(p) + (1 - y) * np.log(1 - p)))


def _sk(X, y):
    """scikit-learn's logistic regression with no penalty: penalty=None, which scikit-learn 1.8 renamed C=inf."""
    import sklearn
    from sklearn.linear_model import LogisticRegression
    major, minor = (int(v) for v in sklearn.__version__.split(".")[:2])
    none = {"C": np.inf} if (major, minor) >= (1, 8) else {"penalty": None}
    return LogisticRegression(**none, fit_intercept=False, solver="newton-cg", tol=1e-12, max_iter=10_000).fit(X, y)


# --------------------------------------------------------------------------
# The three-input shortlist: income / sales, UTIL and TENURE


def test_the_joint_model_keeps_the_planted_inputs_and_tenure_adds_nothing(three):
    j = three["st"].joint
    assert j.problem is None and [t.column for t in j.terms] == ["income_to_sales", "UTIL", "TENURE"]
    by = {t.column: t for t in j.terms}
    bar = 0.05
    assert by["income_to_sales"].p_allowed < bar and by["UTIL"].p_allowed < bar, [t.p_allowed for t in j.terms]
    assert by["TENURE"].p_allowed >= bar
    # the planted sizes sit inside their ranges, net of the other candidates (synth: x2 below 0.1, x3 above 2.0,
    # x2.5 above 0.9 of the line)
    inc, util = by["income_to_sales"], by["UTIL"]
    for t, k, planted in ((inc, 0, 2.0), (inc, 5, 3.0), (util, 3, synth.UTIL_ODDS)):
        lo, hi = (math.exp(t.beta[k] + s * 1.96 * t.se[k]) for s in (-1, 1))
        assert lo < planted < hi, (t.column, k, lo, hi)
    # the allowance is Benjamini-Hochberg's own steps over the three likelihood ratio p-values
    assert [t.p_allowed for t in j.terms] == pytest.approx(_bh([t.p_lr for t in j.terms]), rel=1e-12)
    assert all(t.df == len(t.groups) - 1 for t in j.terms) and all(t.partners == [] for t in j.terms)
    # Record names the test once, with the others the run used
    import tabs
    assert tabs.record(three["b"])["Tests: all together"].startswith("One logistic regression on the held-back loans")


def test_the_joint_fit_is_scikit_learns_on_a_design_built_by_hand(three):
    """The pockets, the groups and the loans rebuilt from the extract; scikit-learn's LogisticRegression with no
    penalty fits them. Every coefficient, standard error, and every likelihood ratio (the model with and without each
    candidate) agrees."""
    res, j = three["res"], three["st"].joint
    terms = [t.column for t in j.terms]
    bins = {"income_to_sales": INCOME["bins"], "UTIL": UTIL["bins"], "TENURE": TENURE["bins"]}
    X, y, npk, cols = _design(res, three["x"], terms, bins)
    assert (len(y), int(y.sum()), npk) == (j.loans, j.bad, j.pockets)
    full = _sk(X, y)
    want = full.coef_[0][npk:]
    got = [j.terms[i].beta[g] for i, g in cols]
    assert np.allclose(got, want, rtol=0, atol=1e-5), np.max(np.abs(np.array(got) - want))
    # the standard errors: the inverse of the information at scikit-learn's answer
    p = full.predict_proba(X)[:, 1]
    cov = np.linalg.inv((X * (p * (1 - p))[:, None]).T @ X)
    se = np.sqrt(np.diag(cov))[npk:]
    assert np.allclose([j.terms[i].se[g] for i, g in cols], se, rtol=1e-5)
    ll = _loglik(full, X, y)
    assert math.isclose(j.loglik, ll, abs_tol=1e-6)
    for i, t in enumerate(j.terms):
        keep = list(range(npk)) + [npk + c for c, (ii, _) in enumerate(cols) if ii != i]
        red = _sk(X[:, keep], y)
        lr = 2 * (ll - _loglik(red, X[:, keep], y))
        assert math.isclose(t.lr, lr, abs_tol=1e-5), (t.column, t.lr, lr)


@pytest.fixture(scope="module")
def copy(tmp_path_factory):
    """The book with UTIL_COPY (UTIL and a little noise, rank correlation about 0.96) on the shortlist beside UTIL
    and income / sales, FICO and CHANNEL held fixed: 20,000 loans."""
    folder = tmp_path_factory.mktemp("copy")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("GIT_CEILING_DIRECTORIES", str(folder))
        mp.setenv("POCKETBOOK_MEMORY", str(folder / "memory.yaml"))
        real = synth.write_extract

        def with_copy(out, n=20000, seed=7, ratio=False, shortlist=False):
            x = real(out, n, seed, ratio, shortlist)
            rows = _rows(x)
            rng = random.Random("util-copy")
            for r in rows:
                r["UTIL_COPY"] = round(float(r["UTIL"]) + rng.gauss(0, 0.1), 4)
            with open(x, "w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=list(rows[0]))
                w.writeheader()
                w.writerows(rows)
            return x
        mp.setattr(synth, "write_extract", with_copy)
        x, b = _book(folder, 20000, {"inputs": [UTIL_ONE, COPY, INCOME]})
        mp.undo()
        _choose(b, drop=NOT_CUT + ("UTIL_COPY",), shortlist="prespec.yaml")
        mp.setenv("GIT_CEILING_DIRECTORIES", str(folder))
        mp.setenv("POCKETBOOK_MEMORY", str(folder / "memory.yaml"))
        ran, st, res = _run(b, mp)
    assert ran.ok, ran.lines
    return {"x": x, "b": b, "st": st, "res": res}


def test_a_near_copy_is_flagged_beside_both_and_adds_nothing_net_of_the_original(copy):
    st = copy["st"]
    j = st.joint
    by = {t.column: t for t in j.terms}
    util, dup = by["UTIL"], by["UTIL_COPY"]
    # flagged beside both, with the rank correlation worked out here on the development loans
    dev = [r for r in _rows(copy["x"]) if date(2022, 1, 1) <= date.fromisoformat(r["ORIG_DATE"]) <= date(2023, 12, 31)]
    from scipy.stats import spearmanr
    rho = spearmanr([float(r["UTIL"]) for r in dev], [float(r["UTIL_COPY"]) for r in dev]).statistic
    assert [n for n, _ in util.partners] == ["UTIL_COPY"] and [n for n, _ in dup.partners] == ["UTIL"]
    assert math.isclose(util.partners[0][1], rho, abs_tol=1e-9) and rho > 0.9
    # on its own the copy holds up on the held-back loans; together it adds nothing the original doesn't say
    alone = next(t for t in st.tests if t.column == "UTIL_COPY").holdout
    assert alone.fit.odds[1] > 1.5 and alone.allowed[1] < 0.05
    assert dup.p_allowed > 0.05 and dup.lr < 3.84 and abs(math.log(dup.odds[1])) < abs(math.log(alone.fit.odds[1]))
    assert util.p_allowed < 0.05 and by["income_to_sales"].p_allowed < 0.05
    # and the tab says so next to both, and once in words under the table
    from openpyxl import load_workbook
    ws = load_workbook(copy["b"])[confirm_tab.SHEET]
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=confirm_tab.N_CAND).value == "Candidate")
    with_ = {ws.cell(row=r, column=confirm_tab.J_CAND).value: ws.cell(row=r, column=confirm_tab.J_WITH).value
             for r in range(head + 1, head + 8) if ws.cell(row=r, column=confirm_tab.J_CAND).value}
    assert with_["UTIL"] == f"UTIL_COPY ({rho:+.2f})" and with_["UTIL_COPY"] == f"UTIL ({rho:+.2f})"
    lines = [ws.cell(row=r, column=confirm_tab.FIRST + 1).value for r in range(head, ws.max_row + 1)]
    said = [v for v in lines if isinstance(v, str) and v.startswith("UTIL and UTIL_COPY move together")]
    assert len(said) == 1, said


# --------------------------------------------------------------------------
# Where the likelihood has no finite answer: refused in words, never printed


def _terms(*groups):
    return [joint.Term(f"V{i}", tuple(f"g{k}" for k in range(n)), 0) for i, n in enumerate(groups)]


def _cells(spec):
    """{(pocket, group of V0, group of V1, ...): [loans, bad]} from (pocket, groups..., loans, bad) tuples."""
    return {tuple(s[:-2]): [s[-2], s[-1]] for s in spec}


BASE_CELLS = [("p1", 0, 0, 200, 20), ("p1", 1, 0, 150, 30), ("p1", 0, 1, 120, 18), ("p1", 1, 1, 100, 25),
              ("p2", 0, 0, 180, 10), ("p2", 1, 0, 160, 20), ("p2", 0, 1, 110, 9), ("p2", 1, 1, 90, 14)]


def test_a_group_with_no_bad_loan_has_no_odds_ratio_and_the_rest_is_fitted_as_if_it_werent_there():
    cells = _cells(BASE_CELLS + [("p1", 2, 0, 40, 0), ("p2", 2, 1, 30, 0)])
    j = joint.fit_cells(joint.Joint(_terms(3, 2), ()), cells)
    assert j.problem is None
    v0 = j.terms[0]
    assert v0.odds[2] is None and v0.not_estimable == {2: joint.NO_BAD}
    alone = joint.fit_cells(joint.Joint(_terms(3, 2), ()), _cells(BASE_CELLS))
    assert alone.terms[0].not_estimable == {2: joint.NO_LOANS}
    assert v0.odds[1] == pytest.approx(alone.terms[0].odds[1], rel=1e-9)
    assert j.terms[1].odds[1] == pytest.approx(alone.terms[1].odds[1], rel=1e-9)


def test_a_group_holding_the_same_loans_as_another_candidates_is_said_so():
    same = [(p, a, b, b, n, y) for p, a, b, n, y in BASE_CELLS]
    j = joint.fit_cells(joint.Joint(_terms(2, 2, 2), ()), _cells(same))
    assert j.problem is None and j.terms[2].not_estimable == {1: joint.ALIASED} and j.terms[2].df == 0
    assert j.terms[2].lr is None and j.terms[1].odds[1] is not None


def test_a_combination_with_no_bad_loan_refuses_the_model_in_words():
    """No group is at a limit on its own, but one direction sends the odds of a combination to nothing: the fit
    runs off, and the model refuses rather than print its odds ratios."""
    cells = _cells([("p", 0, 0, 100, 0), ("p", 1, 0, 100, 0), ("p", 2, 1, 50, 50), ("p", 2, 0, 100, 30),
                    ("p", 0, 1, 100, 30), ("p", 1, 1, 100, 30)])
    j = joint.fit_cells(joint.Joint(_terms(3, 2), ()), cells)
    assert j.problem is not None and j.problem.startswith("the regression didn't settle")
    assert all(t.odds == [] for t in j.terms)


def test_a_reference_with_no_bad_loan_leaves_that_candidate_out_in_words():
    cells = {k + ((k[1] + k[2]) % 2,): v for k, v in _cells(BASE_CELLS).items()}
    cells = {**cells, ("p1", 0, 0, 2): [60, 0], ("p2", 1, 1, 2): [40, 0]}
    terms = _terms(2, 2) + [joint.Term("V2", ("g0", "g1", "g2"), 2)]
    j = joint.fit_cells(joint.Joint(terms, ()), cells)
    assert j.problem is None
    assert j.terms[2].problem == "its reference group, g2, has no bad loan, so nothing can be compared with it"
    assert j.terms[0].odds[1] is not None and j.terms[1].odds[1] is not None


def test_pockets_too_thin_for_a_constant_each_are_refused():
    """Forty pockets holding four bad loans each: under joint.MIN_BAD_PER_POCKET, so one constant a pocket would bias
    the odds ratios (statistics.md B10), and the model says so instead."""
    cells = {(f"p{p}", a, b): [20, 1] for p in range(40) for a in (0, 1) for b in (0, 1)}
    j = joint.fit_cells(joint.Joint(_terms(2, 2), ()), cells)
    assert j.problem.startswith("the pockets are too thin for a constant each: 160 bad loans across 40 pockets")
    cells = {k: [20, 2] if k[1] else v for k, v in cells.items()}         # six bad a pocket: fitted
    assert joint.fit_cells(joint.Joint(_terms(2, 2), ()), cells).problem is None


def test_a_solver_that_breaks_down_refuses_the_same_way(monkeypatch):
    """Where numpy gives up solving instead of running off (CI's did, 28 Sep 2026, on the case above), the model
    refuses with the same words: the columns were checked independent first, so the cause is the same."""
    import numpy as np

    def breaks(*a, **k):
        raise np.linalg.LinAlgError("Singular matrix")
    monkeypatch.setattr(joint.kgroups, "logistic", breaks)
    cells = _cells([("p", 0, 0, 100, 10), ("p", 1, 0, 100, 20), ("p", 2, 1, 100, 30), ("p", 2, 0, 100, 30),
                    ("p", 0, 1, 100, 30), ("p", 1, 1, 100, 30)])
    j = joint.fit_cells(joint.Joint(_terms(3, 2), ()), cells)
    assert j.problem == joint.DIDNT_SETTLE and all(t.odds == [] for t in j.terms)
