"""Two asks from the firm, 30 Sep 2026 (BACKLOG §6d).

Grids: the four blocks stacked, one under another from the left column. Side by side, the right-hand pair started
after the widest grid's columns, so a 4-segment grid left a blank middle the width of the widest (the firm's photo).
Now: Rate, vs the book, vs rest of band, Loans, then the note, What one cell says and the Groups tables.

Summary: "the band column should also allow for categories because we can still view it that way, and the logic
should still make sense". The left-column dropdown lists the band columns, then the segment/category columns; a
category's values run down the side in natural order, then (blank) and (marked missing), then All.

Every figure checked against the loan file is worked out here from the CSV's text with plain arithmetic: nothing in
the checks is imported from pocketbook."""

import csv

import pytest
from openpyxl import load_workbook

from pocketbook import book, choices as ch, engine, results, synth
from test_book import _answer

import tabs
from recalc import need_soffice

FLAG_EVERY = 97                     # CHANNEL is blanked on every 97th loan, so a category has a (blank) row
SUMMARY_ONLY = f"Only loans where {ch.ORIG_YEAR} is"
PICKER = "band or category column"
GRID = "FICO x ASSET_CLASS"         # four segments: the grid the firm photographed


@pytest.fixture(scope="module")
def made(tmp_path_factory):
    from pocketbook import perm
    d = tmp_path_factory.mktemp("grids_summary")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 100)
        src = synth.write_extract(d / "src", n=3000)
        rows = list(csv.DictReader(open(src, encoding="utf-8")))
        for i, r in enumerate(rows):
            if i % FLAG_EVERY == 5:
                r["CHANNEL"] = ""
        x = d / "loans.csv"
        with open(x, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("ASSET_CLASS", "CHANNEL"),
                                                filter=ch.ORIG_YEAR, outcome="BAD_FLAG"))
        _answer(out.book)
        wb = load_workbook(out.book)
        for r in book.table_rows(wb["Columns"]):
            key = r[book.C_QKEY - 1].value
            if isinstance(key, str) and key.count("|") == 2 and not key.startswith("FICO"):
                r[book.C_TREAT - 1].value = "Real"
        wb.save(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return out.book, list(csv.DictReader(open(x, encoding="utf-8")))


# --------------------------------------------------------------------------
# the arithmetic, from the CSV's text alone


def _f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _figures(rs, whole):
    read = [r for r in rs if r["BAD_FLAG"] in ("0", "1")]
    bad = sum(1 for r in read if r["BAD_FLAG"] == "1")
    booked = sum(_f(r["ORIG_BAL"]) for r in rs if _f(r["ORIG_BAL"]) is not None)
    all_booked = sum(_f(r["ORIG_BAL"]) for r in whole if _f(r["ORIG_BAL"]) is not None)

    def rate(rows, col):
        pairs = [(_f(r[col]), _f(r["ORIG_BAL"])) for r in rows]
        pairs = [(a, b) for a, b in pairs if a is not None and b is not None]
        return sum(a for a, _ in pairs), sum(b for _, b in pairs)

    gco, gden = rate(rs, "GCO_AMT")
    bg, bden = rate(whole, "GCO_AMT")
    ranr, rden = rate(rs, "RANR_AMT")
    return {"Loans": len(rs), "Bad loans": bad, "Bad loans %": bad / len(read) if read else None,
            "Booked $": booked, "GCOs ($)": gco, "GCOs ÷ Booked": gco / gden if gden else None,
            "× book": (gco / gden) / (bg / bden) if gden else None, "RANR $": ranr,
            "RANR ÷ Booked": ranr / rden if rden else None, "_all_booked": all_booked}


def _summary(b, out, pick, only=None):
    """Summary with its dropdowns set, calculated: {row label: {heading: value}} in order, and the label heading."""
    picks = {PICKER: pick, **({SUMMARY_ONLY: only} if only else {})}
    ws = tabs.calculated(tabs.choose(b, out, results.SUMMARY, **picks), results.SUMMARY)
    h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
    heads, c = [], 3
    while ws.cell(row=h, column=c).value not in (None, ""):
        heads.append(ws.cell(row=h, column=c).value)
        c += 1
    rows, r = {}, h + 1
    while ws.cell(row=r, column=2).value not in (None, ""):
        rows[ws.cell(row=r, column=2).value] = {hd: ws.cell(row=r, column=3 + j).value for j, hd in enumerate(heads)}
        r += 1
    return rows, ws.cell(row=h, column=2).value


def _same(got, want):
    if want is None:
        return got in (None, "")
    return got == pytest.approx(want, rel=1e-9, abs=1e-9)


# --------------------------------------------------------------------------
# Grids


def _tops(ws):
    col_b = [(c.row, c.value) for c in ws["B"] if isinstance(c.value, str)]
    rate = next(r for r, v in col_b if v.startswith('="Rate · "'))
    return [rate, next(r for r, v in col_b if r > rate and v.startswith("=IF(ISNUMBER(")),
            next(r for r, v in col_b if r > rate and v == "vs rest of band"),
            next(r for r, v in col_b if r > rate and v == "Loans")]


def test_grids_stacks_its_four_blocks_one_under_another_from_the_left(made):
    b, _ = made
    ws = load_workbook(b)[results.GRIDS]
    tops = _tops(ws)
    assert tops == sorted(tops) and len({y - x for x, y in zip(tops, tops[1:])}) == 1
    # the first block's heading spans from B; nothing of a block starts to the right of it any more
    heads = {c.value for row in ws.iter_rows(min_row=tops[0], max_row=tops[-1], min_col=3, max_col=21) for c in row
             if isinstance(c.value, str) and (c.value in ("vs rest of band", "Loans") or c.value.startswith("=IF(ISN"))}
    assert not heads
    # the dropdowns stay where the bank's checklist sends the analyst
    assert (ws["B12"].value, ws["F12"].value) == ("GRID", "MEASURE")
    # the note under the blocks, then What one cell says, then the groups, all below Loans
    blank = next(c.row for c in ws["B"] if isinstance(c.value, str) and c.value.startswith("A blank: alone"))
    one = next(c.row for c in ws["B"] if c.value == "The colour")
    assert tops[-1] < blank < one


def test_a_four_segment_grid_reads_back_from_the_stacked_blocks_as_the_loan_file_counts_it(made, tmp_path):
    """The Loans block and What one cell says, after LibreOffice calculates them, against the CSV: each FICO band by
    asset class, counted here from the text (the -9999 code answered missing, a blank its own row)."""
    need_soffice()
    b, rows = made
    out = tabs.choose(b, tmp_path / "g.xlsx", results.GRIDS, grid=GRID)
    ws = tabs.calculated(out, results.GRIDS)
    loans = tabs.block(ws, "Loans")
    bands = [lab for lab, _ in loans if lab != "All" and " - " in str(lab)]
    bands = list(dict.fromkeys(bands))

    def band(v):
        if v in ("", None):
            return "(blank)"
        x = float(v)
        if x < -1000:
            return "(marked missing)"
        return next(lab for lab in bands if float(lab.split(" - ")[0]) <= x <= float(lab.split(" - ")[1]))

    count: dict = {}
    for r in rows:
        k = (band(r["FICO"]), r["ASSET_CLASS"])
        count[k] = count.get(k, 0) + 1
    assert {str(c) for _, c in loans} == {"1", "2", "3", "4", "All"}
    checked = 0
    for (bl, seg), n in loans.items():
        if bl == "All" or seg == "All":
            continue
        assert (n or 0) == count.get((bl, str(seg)), 0), (bl, seg)     # an empty pocket shows blank
        checked += 1
    assert checked >= 4 * len(bands)
    assert loans[("All", "All")] == len(rows)
    # the Rate block, bad loans (the default measure), pocket by pocket
    rate = tabs.block(ws, "Rate")
    for (bl, seg), v in rate.items():
        if bl in bands and seg != "All":
            mine = [r for r in rows if band(r["FICO"]) == bl and r["ASSET_CLASS"] == str(seg)
                    and r["BAD_FLAG"] in ("0", "1")]
            assert _same(v, sum(r["BAD_FLAG"] == "1" for r in mine) / len(mine)), (bl, seg)
    # What one cell says reads the stacked Loans block: its Loans line names the pocket's count
    picked_row, picked_col = bands[1], "3"
    ws = tabs.calculated(tabs.choose(b, tmp_path / "one.xlsx", results.GRIDS, grid=GRID, row=picked_row,
                                     column=picked_col), results.GRIDS)
    said = next(ws.cell(row=c.row, column=3).value for c in ws["B"] if c.value == "Loans"
                and c.row > _tops(ws.formulas)[-1] + 3)
    n = count[(picked_row, str(picked_col))]
    assert f"{n:,}" in str(said), (said, n)


# --------------------------------------------------------------------------
# Summary


def test_summary_lists_the_category_columns_after_the_band_columns(made):
    b, _ = made
    got = tabs.options(load_workbook(b), results.SUMMARY, PICKER)
    assert got[0] == "FICO" and sorted(got[1:]) == ["ASSET_CLASS", "CHANNEL"]


def test_summary_by_a_category_is_the_loans_worked_out_again(made, tmp_path):
    """CHANNEL down the side: Branch, Broker, Online, (blank), then All, every figure against the CSV; the shares
    add to 100% and the header names the column picked."""
    need_soffice()
    b, rows = made
    shown, head = _summary(b, tmp_path / "ch.xlsx", "CHANNEL")
    assert head == "CHANNEL"
    assert list(shown) == ["Branch", "Broker", "Online", "(blank)", "All"]
    by = {lab: [r for r in rows if (r["CHANNEL"] or "(blank)") == lab] for lab in shown if lab != "All"}
    by["All"] = rows
    assert len(by["(blank)"]) == len(range(5, len(rows), FLAG_EVERY))
    for lab, rs in by.items():
        want = _figures(rs, rows)
        for hd in ("Loans", "Bad loans", "Bad loans %", "Booked $", "GCOs ($)", "GCOs ÷ Booked", "× book", "RANR $",
                   "RANR ÷ Booked"):
            assert _same(shown[lab][hd], want[hd]), (lab, hd, shown[lab][hd], want[hd])
        assert _same(shown[lab]["% of loans"], len(rs) / len(rows))
        assert _same(shown[lab]["% of booked"], want["Booked $"] / want["_all_booked"])
    for hd in ("% of loans", "% of booked", "% of GCOs", "% of RANR"):
        assert sum(v[hd] for lab, v in shown.items() if lab != "All") == pytest.approx(1.0, rel=1e-9)
    assert shown["All"]["× book"] == pytest.approx(1.0)


def test_summary_by_a_category_on_one_year_keeps_the_whole_books_rows(made, tmp_path):
    """ASSET_CLASS on one origination year: 1 to 4 in natural order, then All; × book still against the whole
    book."""
    need_soffice()
    b, rows = made
    year = sorted({r["ORIG_DATE"][:4] for r in rows})[1]
    mine = [r for r in rows if r["ORIG_DATE"][:4] == year]
    shown, head = _summary(b, tmp_path / "ac.xlsx", "ASSET_CLASS", only=year)
    assert head == "ASSET_CLASS" and [str(x) for x in shown] == ["1", "2", "3", "4", "All"]
    for lab in list(shown):
        rs = mine if lab == "All" else [r for r in mine if r["ASSET_CLASS"] == str(lab)]
        want = _figures(rs, rows)
        for hd in ("Loans", "Bad loans", "GCOs ($)", "× book", "RANR ÷ Booked"):
            assert _same(shown[lab][hd], want[hd]), (lab, hd)
    assert shown["All"]["Loans"] == len(mine) < len(rows)


def test_the_engine_keeps_a_summary_for_every_category_and_every_filter_view(made, tmp_path):
    """The Run's summaries for a category column, keyed (column, value, value2): the whole book and each filter
    view, on _views for the tab; and, straight from the engine, a category's rows in natural order then its
    special rows then All, adding up to the book."""
    b, rows = made
    wb = load_workbook(b)
    views = {r[0] for r in wb[results.VIEWS].iter_rows(values_only=True) if isinstance(r[0], str)}
    for col in ("CHANNEL", "ASSET_CLASS"):
        assert f"S|{col}|1" in views
        for y in sorted({r["ORIG_DATE"][:4] for r in rows}):
            assert f"S|{col}|where {y}|1" in views, (col, y)
    from pocketbook import config as cfgmod
    from pocketbook.ingest import read_table
    cfg, data = synth.write(tmp_path / "cube", n=1500)
    res = engine.run(cfgmod.load(cfg), read_table(data))
    for d in res.config.dimensions:
        got = engine.summary_rows(res, d.name)
        labels = [lab for lab, _ in got]
        assert labels[-1] == engine.ALL and labels[:-1] == engine._order(labels[:-1])
        assert sum(v["loans"] for lab, v in got[:-1]) == got[-1][1]["loans"] == res.rows
