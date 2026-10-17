"""Significant at N%? (the firm, 17 Oct 2026): Pockets' p-value column in standard statistics wording.

The firm: "i don't like how often by chance or could be chance as much as luck at this point. or another sort of
standard term, i just think could be chance is too vague for a typical user". The layout they chose:

    Significant at 95%?
    Yes (p = 0.09%)
    Borderline (p = 4.8%)
    No (p = 23%)
    Too few losses to test
    Not losing more

It replaces How often by chance (5 Oct 2026), whose Yes pocket read "1 in X" and whose other readings were Could be
chance and Too close to call. The p-value is the same one "1 in X" was built from: the p-value that decides, after
the allowance for many tests. The heading follows How sure on Control live."""

from __future__ import annotations

import re
import shutil

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

import tabs
from conftest import TEST_SHUFFLES
from pocketbook import book, choices as ch, control, engine, live, perm, results, stats, synth
from test_book import _answer


@pytest.fixture(scope="module")
def sig_book(tmp_path_factory):
    """The synthetic book, FICO by CHANNEL, judged against the rest of its band: the workbook, the engine's result,
    every call to the allowance for many tests (its raw p-values and what it gave back), and the book calculated."""
    from recalc import calculated_book
    d = tmp_path_factory.mktemp("sig")
    got: dict = {"adjust": []}
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
        run, adjust = engine.run, engine.adjust

        def kept(*a, **k):
            got["res"] = run(*a, **k)
            return got["res"]

        def watched(ps, how):
            out = adjust(ps, how)
            got["adjust"].append((list(ps), list(out), how))
            return out
        mp.setattr(engine, "run", kept)
        mp.setattr(engine, "adjust", watched)
        out = book.set_up(synth.write_extract(d / "src", n=8000),
                          choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",)))
        assert out.ok, out.lines
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return {"b": out.book, "res": got["res"], "adjust": got["adjust"], "calc": calculated_book(out.book), "dir": d}


def _with_confidence(src, out, label: str | None = None, own: float | None = None):
    """A copy of the workbook with How sure on Control set to `label` ("99%"), or typed under Or your own."""
    shutil.copy(src, out)
    wb = load_workbook(out)
    for r in wb[control.SHEET].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == "confidence":
            if label:
                r[control.CHOOSE_COL - 1].value = label
            if own is not None:
                r[control.OWN_COL - 1].value = own
    wb.save(out)
    return out


def _head(ws) -> str:
    return ws.cell(row=tabs.header_row(ws, results.K_NUM, "#"), column=results.K_P).value


def _note_labels(ws) -> list:
    top = next(c.row for row in ws.iter_rows(max_row=40) for c in row if c.value == "How this tab works")
    return [ws.cell(row=r, column=results.K_BAND).value for r in range(top + 1, top + 12)]


# --------------------------------------------------------------------------
# the heading


def test_the_heading_follows_how_sure_on_control(sig_book, tmp_path):
    """Significant at 95%? at the Run's 95%, and Significant at 99%? once How sure reads 99%, with no Run between:
    the heading and the method note's label are live, as the tab's other live headings are."""
    ws = sig_book["calc"][results.POCKETS]
    assert _head(ws) == "Significant at 95%?"
    assert "Significant at 95%?" in _note_labels(ws)
    sure = _with_confidence(sig_book["b"], tmp_path / "99.xlsx", "99%")
    ws = tabs.calculated(sure, results.POCKETS)
    assert _head(ws) == "Significant at 99%?"
    assert "Significant at 99%?" in _note_labels(ws)
    # and the readings follow it: at 99% sure a Yes is a p-value under 1%
    yes = [x["p"] for x in tabs.pockets(ws) if str(x["chance"]).startswith("Yes, ")]
    assert yes and all(p < 0.01 for p in yes), yes


@pytest.mark.parametrize("own, sure, line", [(0.975, "97.5%", "2.5%"), (0.999, "99.9%", "0.1%")])
def test_a_typed_how_sure_reads_as_typed_in_the_heading_and_the_note(sig_book, tmp_path, own, sure, line):
    """How sure takes a typed value from 0.5 to 0.999. The heading and the note word it as _live does: 97.5%, never
    98%; and its line 2.5%, never 3%."""
    ws = tabs.calculated(_with_confidence(sig_book["b"], tmp_path / f"{own}.xlsx", own=own), results.POCKETS)
    assert _head(ws) == f"Significant at {sure}?"
    top = next(c.row for row in ws.iter_rows(max_row=40) for c in row if c.value == "How this tab works")
    note = {ws.cell(row=r, column=results.K_BAND).value: ws.cell(row=r, column=results.K_BAND + 1).value
            for r in range(top + 1, top + 12)}
    assert f"Yes, worse means its p-value is under {line} and" in note[f"Significant at {sure}?"]
    assert live.pct_words(own) == sure and live.pct_words(round(1 - own, 12)) == line


# --------------------------------------------------------------------------
# the three readings


@pytest.mark.parametrize("word, gap, p, border, said", [
    (engine.WORSE, 1.6, 0.0009, False, live.SIG_WORSE),
    (engine.WORSE, 1.6, 0.048, True, live.SIG_BORDER),
    (engine.UNSURE_WORSE, 1.6, 0.052, True, live.SIG_BORDER),
    (engine.UNSURE_WORSE, 1.6, 0.23, False, live.SIG_NO),
    # the column answers its own question (the review, 17 Oct 2026): under the Worse at line but significant is Yes
    (engine.IN_LINE, 1.15, 0.00001, False, live.SIG_WORSE),
    (engine.BETTER, 0.6, 0.002, False, live.SIG_BETTER),
])
def test_yes_borderline_and_no_read_by_the_p_value(word, gap, p, border, said):
    assert live.sig_said(word, gap, p, border, False, bar=0.05) == said


def test_the_tab_prints_only_the_new_readings_and_never_the_old(sig_book):
    """Every row on Pockets, for every measure, reads Yes, Borderline or No with its p-value beside it, or one of the reasons
    kept from How often by chance. Nothing a reader sees says 1 in X, Could be chance, Too close to call or How often
    by chance."""
    kept = {live.FEW_LOSSES, live.NO_REST, live.NOT_MORE}
    seen = set()
    for measure in [results.plain(m) for m in results.rates(sig_book["res"])]:
        ws = tabs.calculated(tabs.choose(sig_book["b"], sig_book["dir"] / f"m-{len(seen)}-{measure[:3]}.xlsx",
                                         results.POCKETS, measure=measure), results.POCKETS)
        for x in tabs.pockets(ws):
            said = str(x["chance"])
            seen.add(said)
            if said in kept:
                assert x["p"] is None, x
                continue
            assert said in (live.SIG_WORSE, live.SIG_BETTER, live.SIG_BORDER, live.SIG_NO), said
            assert isinstance(x["p"], float) and 0 <= x["p"] <= 1, x
            # Yes is a p-value under the bar and No one at or over it; Borderline may be either side
            if said.startswith("Yes, "):
                assert x["p"] < 0.05, x
            elif said == live.SIG_NO:
                assert x["p"] >= 0.05, x
            if x["worse"] == "Yes":
                assert said in (live.SIG_WORSE, live.SIG_BORDER), x
    assert {live.SIG_WORSE, live.SIG_NO} <= seen, seen
    # every pocket with a p-value is answered with it; only one without a p-value gets a reason
    pk = sig_book["calc"][live.POCKETS]
    for r in range(live.P_FIRST, pk.max_row + 1):
        p, said = pk.cell(row=r, column=live.P_P).value, pk.cell(row=r, column=live.P_SIG).value
        if pk.cell(row=r, column=live.P_MEASURE).value in (None, ""):
            continue
        if isinstance(p, float | int) and not isinstance(p, bool):
            assert said in (live.SIG_WORSE, live.SIG_BETTER, live.SIG_BORDER, live.SIG_NO), (r, p, said)
        else:
            assert said in kept, (r, said)
    wb = sig_book["calc"]
    text = " ".join(str(c.value) for ws in wb.worksheets if ws.sheet_state == "visible" for row in ws.iter_rows()
                    for c in row if isinstance(c.value, str))
    for gone in ("How often by chance", "Could be chance", "Too close to call", "Rarer than 1 in"):
        assert gone not in text, gone
    assert not re.search(r"\b1 in [\d,]+\b", " ".join(str(c.value) for row in wb[results.POCKETS].iter_rows()
                                                      for c in row if isinstance(c.value, str)))


# --------------------------------------------------------------------------
# the p format


@pytest.mark.parametrize("p, bar, said", [
    (0.0009, 0.05, "p = 0.09%"), (0.0037, 0.05, "p = 0.37%"), (0.048, 0.05, "p = 4.8%"), (0.23, 0.05, "p = 23%"),
    (0.0001, 0.05, "p = 0.01%"), (0.0000999, 0.05, "p < 0.01%"), (0.0, 0.05, "p < 0.01%"),
    # the tiers are by the p-value itself, as a number format's conditions are: under 1% two decimals, so 0.00995 is
    # 1.00%; under 10% one, so 0.0996 is 10.0%
    (0.00995, 0.05, "p = 1.00%"), (0.0996, 0.05, "p = 10.0%"), (1.0, 0.05, "p = 100%"),
    # a p-value that would round onto the bar it is under, or over, takes one more decimal
    (0.0498, 0.05, "p = 4.98%"), (0.0503, 0.05, "p = 5.03%"), (0.05, 0.05, "p = 5.0%"),
    (0.0099, 0.01, "p = 0.99%"), (0.00999, 0.01, "p = 0.999%"), (0.0101, 0.01, "p = 1.01%"),
    # LibreOffice rounds the binary double: 0.00145 x 100 is 0.14499...
    (0.00145, 0.05, "p = 0.14%"),
    # just under the line at 90%, 95% and 97.5% sure, and past two extra decimals it says which side
    (0.0999, 0.1, "p = 9.99%"), (0.09996, 0.1, "p = 9.996%"), (0.10004, 0.1, "p > 10%"),
    (0.0499996, 0.05, "p < 5%"), (0.02499, 0.025, "p = 2.499%"), (0.024996, 0.025, "p < 2.5%"),
    (0.000996, 0.001, "p = 0.0996%"),
])
def test_the_p_value_is_a_percentage_at_the_firms_precision(p, bar, said):
    assert live.p_words(p, bar) == said


def test_the_p_columns_formats_choose_as_python_does_in_a_spreadsheet(tmp_path):
    """The p-value column's conditional formats (live.p_rules), each condition calculated by LibreOffice, choose the
    format live.p_format chooses, the first rule that holds or else the cell's own (live.P_FMT), at five lines."""
    from recalc import recalc
    ps = [0.0009, 0.0037, 0.048, 0.23, 0.0001, 0.0000999, 0.0, 0.00995, 0.0996, 1.0, 0.0498, 0.0503, 0.05, 0.0099,
          0.012345, 0.000151, 0.4567, 0.00999, 0.0101, 0.00145, 0.0999, 0.09996, 0.10004, 0.0499996, 0.02499,
          0.024996, 0.000996, 0.001, 0.01, 0.025, 0.1]
    own = lambda p: '"< 0.01%"' if p < live.RAREST else "0.00%" if p < 0.01 else "0.0%"     # noqa: E731
    for bar in (0.05, 0.01, 0.1, 0.025, 0.001):
        wb = Workbook()
        ws = wb.active
        ws["Z1"] = bar
        wb.defined_names[live.BAR] = DefinedName(live.BAR, attr_text=f"'{ws.title}'!$Z$1")
        for i, p in enumerate(ps, start=1):
            ws.cell(row=i, column=1, value=p)
            for j, (f, _) in enumerate(live.p_rules(f"A{i}"), start=2):
                ws.cell(row=i, column=j, value="=" + f)
        wb.save(tmp_path / f"p-{bar}.xlsx")
        calc = recalc(tmp_path / f"p-{bar}.xlsx", tmp_path / f"out-{bar}").active
        for i, p in enumerate(ps, start=1):
            rules = live.p_rules(f"A{i}")
            hit = next((fmt for j, (_, fmt) in enumerate(rules, start=2) if calc.cell(row=i, column=j).value is True),
                       own(p))
            assert hit == live.p_format(p, bar), (p, bar, hit)


# --------------------------------------------------------------------------
# the same p-value "1 in X" was built from, tied by hand


def _bh(raw: list) -> list:
    """Benjamini-Hochberg by hand: each p-value x m / its rank, then the smallest from it upward, capped at 1."""
    tested = sorted((p, i) for i, p in enumerate(raw) if p is not None)
    m = len(tested)
    out: list = [None] * len(raw)
    running = 1.0
    for rank in range(m, 0, -1):
        p, i = tested[rank - 1]
        running = min(running, p * m / rank)
        out[i] = running
    return out


def test_a_yes_pocket_prints_the_p_value_that_decides_tied_by_hand(sig_book):
    """One Yes pocket on Bad loans, tied end to end: its p-value worked out here from its own counts (the pooled
    two-proportion z test against the rest of its band), the allowance worked out here over its grid's p-values,
    and the two side by side with the engine's, with _pockets' p-value that decides (what "1 in X" was built from)
    and with what Pockets prints."""
    res, calc = sig_book["res"], sig_book["calc"]
    assert res.config.benchmark.many_tests == "bh" and res.config.benchmark.compare_to == "peers"
    g = next(x for x in res.grids)
    # the Yes pocket with the largest p-value, so the figure printed is a percentage rather than "p < 0.01%"
    yes = [(c.rates["outcome_loans"].p_band, bl, dl, c, c.rates["outcome_loans"]) for (bl, dl), c in g.inner()
           if c.rates["outcome_loans"].flag == engine.WORSE and c.rates["outcome_loans"].by_band
           and c.rates["outcome_loans"].units >= res.config.benchmark.min_units]
    assert yes, "no Yes pocket on Bad loans"
    pick = max(yes, key=lambda x: x[0])[1:]
    bl, dl, c, s = pick
    band = g.cells[(bl, engine.ALL)].rates["outcome_loans"]
    raw = stats.two_prop_z(s.num, s.units, band.num - s.num, band.units - s.units)[1]
    same = lambda r: r is not None and abs(r - raw) <= 1e-12 * raw                 # noqa: E731
    call = next(x for x in sig_book["adjust"] if any(same(r) for r in x[0]) and x[1][[same(r) for r in x[0]].index(
        True)] == pytest.approx(s.p_band, rel=1e-12))
    i = [same(r) for r in call[0]].index(True)
    by_hand = _bh(call[0])[i]
    assert by_hand == pytest.approx(s.p_band, rel=1e-12) and by_hand >= raw
    # _pockets' p-value that decides, the cell the old "1 in X" read, is the same number
    pk = calc[live.POCKETS]
    row = next(r for r in range(live.P_FIRST, pk.max_row + 1)
               if (pk.cell(row=r, column=live.P_KIND).value, pk.cell(row=r, column=live.P_BAND).value,
                   pk.cell(row=r, column=live.P_SEG).value, pk.cell(row=r, column=live.P_MEASURE).value)
               == ("grids", bl, dl, "outcome_loans") and pk.cell(row=r, column=live.P_GRID).value.startswith("FICO"))
    assert pk.cell(row=row, column=live.P_P).value == pytest.approx(by_hand, rel=1e-12)
    want = live.SIG_WORSE
    assert pk.cell(row=row, column=live.P_SIG).value == want
    # and Pockets prints it on that pocket's row
    ws = tabs.calculated(tabs.choose(sig_book["b"], sig_book["dir"] / "bad.xlsx", results.POCKETS,
                                     measure="Bad loans"), results.POCKETS)
    names = {b.name: b.field for b in res.config.bands}
    rows = [x for x in tabs.pockets(ws) if x["band"] == f"{names[g.band]} {bl}" and str(dl) in str(x["seg"])]
    assert [x["chance"] for x in rows] == [want], rows
    assert rows[0]["p"] == pytest.approx(by_hand, rel=1e-12)          # the p-value column, beside it, the same number


@pytest.mark.parametrize("many, claims", [("none", False), ("bh", True), ("bonferroni", True)])
def test_the_note_claims_an_allowance_only_when_control_sets_one(sig_book, many, claims):
    """With No allowance on Control, the note says each p-value is the pocket's own; with one, that the raw p-value
    is no larger (never "smaller still": a pocket's allowed p-value can equal its raw one)."""
    import dataclasses
    res = sig_book["res"]
    b = dataclasses.replace(res.config.benchmark, many_tests=many)
    got = dict(results._pockets_note(dataclasses.replace(res, config=dataclasses.replace(res.config, benchmark=b))))
    note = got[results.SIG_HEAD][1:]                     # the live.text formula, words and all
    assert ("allows for testing many pockets at once" in note) is claims
    assert ("no allowance for testing many pockets at once" in note) is not claims
    assert "smaller still" not in note and ("raw p-value is no larger" in note) is claims


def test_the_reading_formula_matches_python_over_every_combination(tmp_path):
    """sig_words, calculated by LibreOffice, against sig_said over every engine word, gap, p-value and standard error,
    for a loss and for profit, including the combinations no Run here happens to produce (a pocket that loses more
    with no p-value)."""
    import itertools
    from recalc import recalc
    from pocketbook import stats
    combos = list(itertools.product(engine.READINGS + ("",), (None, 1.6, 1.1, 0.8, -0.002, 0.002),
                                    (None, 0.0004, 0.049, 0.3), (None, 0.001), (False, True)))
    wb = Workbook()
    ws = wb.active
    ws["Z1"] = 0.05
    wb.defined_names[live.BAR] = DefinedName(live.BAR, attr_text=f"'{ws.title}'!$Z$1")
    for i, (word, gap, p, se, profit) in enumerate(combos, start=1):
        for c, v in enumerate((word, gap if gap is not None else "", p if p is not None else "",
                               se if se is not None else ""), start=1):
            ws.cell(row=i, column=c, value=v)
        ws.cell(row=i, column=5, value="=" + live.sig_words(f"A{i}", f"B{i}", f"C{i}", f"D{i}", profit))
    wb.save(tmp_path / "sig.xlsx")
    calc = recalc(tmp_path / "sig.xlsx", tmp_path / "out").active
    for i, (word, gap, p, se, profit) in enumerate(combos, start=1):
        near = stats.borderline(p, se, 0.95)
        want = live.sig_said(word or None, gap, p, near, profit, 0.05)
        assert calc.cell(row=i, column=5).value == want, (word, gap, p, se, profit)
