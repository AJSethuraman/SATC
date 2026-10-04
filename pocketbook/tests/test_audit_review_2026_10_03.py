"""The audit workbook's review, 3 Oct 2026: four bugs found by reading audit.py and book.py, each proved here.

1. A pocket was found on _pocketbook by MATCH on its name. MATCH reads ~ * ? as wildcards: on the code before the
   fix, a segment "Br~anch" found no row (56 figures ✗), and "<B*" and 'On"line' each left a ✗ in part D of
   Shuffle test, which found another pocket's row. COUNTIFS also reads a leading < > = as an operator, and "=A" was
   written to _pocketbook and _lists as a formula. Every name is now found character for character, every
   criterion built from text matches it literally, and every name is written as text.
2. The Loans sheet's rows are written straight into its XML; a control character in a loan number made the sheet
   unreadable.
3. Any failure of the audit workbook other than "open in Excel" ended the Run after the main workbook was saved,
   with no "what ran.yaml" and none of the Run's lines.
4. More loans than an Excel sheet has rows wrote a Loans sheet Excel cannot open.

Every check of a figure reads the workbook as LibreOffice calculates it."""

import csv
import shutil
import xml.etree.ElementTree as ET
import zipfile

import pytest
from openpyxl import load_workbook

from pocketbook import audit, book, choices as ch, perm, synth
from test_book import _answer
from test_audit_2026_10_03 import _picks, _set_control, _table, _ties_in

from recalc import need_soffice, recalc

N = 1500
SHUFFLES = 400
TICK = audit.TICK

#: segment values Excel reads as something other than text: a wildcard (~ * ?), an operator at the start (< > =), a
#: quote, and a trailing space. "1?" read as a wildcard also matches "10"
CHANNEL = {"Branch": "Br~anch", "Online": 'On"line ', "Broker": "<B*"}
ASSET = {"1": "10", "2": "1?", "3": "=A", "4": ">500"}


def _build(d, name, *, edit=None, answer="Yes"):
    """A bleed Run on a synthetic book, its CSV changed by `edit` (a function of the rows) first."""
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / f"memory-{name}.yaml"))
        mp.setattr(perm, "SHUFFLES", SHUFFLES)
        src = synth.write_extract(d / name, n=N)
        if edit is not None:
            rows = list(csv.DictReader(open(src, encoding="utf-8", newline="")))
            edit(rows)
            with open(src, "w", encoding="utf-8", newline="") as f:
                w = csv.DictWriter(f, fieldnames=list(rows[0]))
                w.writeheader()
                w.writerows(rows)
        out = book.set_up(src, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO", "REV_DEBT"),
                                                  segments=("CHANNEL", "ASSET_CLASS"), outcome="BAD_FLAG"))
        _answer(out.book)
        _set_control(out.book, audit.KEY, answer)
        return src, out.book


def _odd(rows):
    for r in rows:
        r["CHANNEL"] = CHANNEL.get(r["CHANNEL"], r["CHANNEL"])
        r["ASSET_CLASS"] = ASSET.get(r["ASSET_CLASS"], r["ASSET_CLASS"])


def _pb(path) -> dict:
    """_pocketbook as written: (grid, band, seg) -> its figures."""
    ws = load_workbook(path, read_only=True)[audit.PB]
    heads = [c.value for c in ws[1]]
    return {(r[1], r[2], r[3]): dict(zip(heads, r)) for r in ws.iter_rows(min_row=2, values_only=True)}


@pytest.fixture(scope="module")
def odd(tmp_path_factory):
    need_soffice()
    d = tmp_path_factory.mktemp("audit-odd")
    src, b = _build(d, "odd", edit=_odd)
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(perm, "SHUFFLES", SHUFFLES)
        ran = book.run(b)
    assert ran.ok, ran.lines
    path = audit.path_for(b)
    yield {"audit": path, "calc": recalc(path, d / "rc"), "dir": d}
    shutil.rmtree(d, ignore_errors=True)


def _pick(odd, name, key):
    d = odd["dir"] / name
    d.mkdir(exist_ok=True)
    p = d / odd["audit"].name
    wb = load_workbook(odd["audit"])
    ws = wb[audit.ONE]
    at = _picks(ws)
    for cell, v in zip((at["GRID"], at["BAND"], at["SEGMENT"]), key):
        ws[cell] = v
        ws[cell].data_type = "s"                  # as a pick from the dropdown: "=A" is the value, not a formula
    wb.save(p)
    return recalc(p, d / "rc")


def test_a_segment_excel_reads_as_a_wildcard_or_an_operator_is_found_as_written(odd):
    """For every odd segment, a pocket of it picked on One pocket: every figure is that pocket's and ties, and part
    D of Shuffle test finds it in its family."""
    pockets = _pb(odd["audit"])
    wanted = set(CHANNEL.values()) | set(ASSET.values())
    segs = {k[2] for k in pockets}
    # the trailing space is the loan file's: PocketBook reads 'On"line ' as 'On"line'
    assert {s.rstrip() for s in wanted} == {s for s in segs if s not in ("(blank)", "(marked missing)")}
    checked = tested_n = 0
    for seg in sorted(segs - {"(blank)", "(marked missing)"}):
        # a tested pocket of this segment, the largest, in the grid whose segment column holds it
        key = max((k for k in pockets if k[2] == seg),
                  key=lambda k: (isinstance(pockets[k]["raw"], float), pockets[k]["loans"]))
        tested = isinstance(pockets[key]["raw"], float)
        calc = _pick(odd, f"seg-{checked}", key)
        ws = calc[audit.ONE]
        rows = _table(ws)
        bad = [(r["step"], r["ties"], r["excel"], r["pb"]) for r in rows
               if r["ties"] != TICK and not (r["ties"] == "–" and r["step"].startswith("Shuffles"))]
        assert bad == [], (seg, bad)
        got = {r["step"]: r["excel"] for r in rows}
        assert got["Loans"] == pockets[key]["loans"], seg
        assert got["GCOs"] == pytest.approx(pockets[key]["gco"], abs=1e-6), seg
        # part D lists the pocket's family and marks this pocket in it, every row tied
        sh = calc[audit.SHUFFLE]
        marks = [r for r in range(1, sh.max_row + 1) if sh.cell(row=r, column=11).value == "◀ selected pocket"]
        assert len(marks) == tested, seg
        if tested:
            assert sh.cell(row=marks[0], column=3).value == f"{key[1]}, {key[2]}"
        assert all(v == TICK for _, v in _ties_in(sh, 10)), seg
        checked += 1
        tested_n += tested
    assert checked == len(segs - {"(blank)", "(marked missing)"}) == 7
    assert tested_n >= 5


def test_odd_names_are_written_as_text_and_every_count_ties(odd):
    """"=A" is a segment, never a formula, wherever the audit writes it; Rows in and out and Bands tie."""
    wb = load_workbook(odd["audit"])
    pb, lists = wb[audit.PB], wb[audit.LISTS]
    for ws in (pb, lists):
        found = [c for row in ws.iter_rows() for c in row if c.value == "=A"]
        assert found and all(c.data_type == "s" for c in found), ws.title
    assert all(v == TICK for _, v in _ties_in(odd["calc"][audit.ROWS], 5))
    assert all(v == TICK for _, v in _ties_in(odd["calc"][audit.BANDS], 7))
    # the random pocket ties as it opens
    rows = _table(odd["calc"][audit.ONE])
    assert rows and all(r["ties"] == TICK for r in rows)


def test_a_criterion_from_text_matches_that_text_only():
    """crit escapes the wildcards and puts "=" first, so no character of the text is read as an operator."""
    assert audit.crit("Br~anch") == '"=Br~~anch"'
    assert audit.crit("1?") == '"=1~?"'
    assert audit.crit("<B*") == '"=<B~*"'
    assert audit.crit('On"line') == '"=On""line"'
    assert audit.crit(">500") == '"=>500"'


# --------------------------------------------------------------------------
# 2. a control character in the loans


def test_a_control_character_in_a_loan_number_leaves_the_loans_sheet_readable(tmp_path):
    need_soffice()

    # a loan well down the file, past the first values Set up's Columns tab shows
    def edit(rows):
        rows[700]["LOAN_NBR"] = "L\x015"
    _, b = _build(tmp_path, "ctrl", edit=edit)
    with pytest.MonkeyPatch.context() as mp:
        mp.setattr(perm, "SHUFFLES", SHUFFLES)
        ran = book.run(b)
    assert ran.ok, ran.lines
    path = audit.path_for(b)
    z = zipfile.ZipFile(path)
    from pocketbook.excel_lists import _sheet_files
    ET.fromstring(z.read(_sheet_files(z)[audit.LOANS]))          # parses: no character XML forbids
    calc = recalc(path, tmp_path / "rc")
    loans = calc[audit.LOANS]
    # read_table shows a forbidden character as its control picture (U+2401 for \x01), so two loan numbers differing
    # only by one stay two (the main workbook's fix of 4 Oct 2026); the audit carries the same text
    assert loans.cell(row=702, column=2).value == "L\u24015"
    assert len(_ties_in(calc[audit.ROWS], 5)) >= 10 and all(v == TICK for _, v in _ties_in(calc[audit.ROWS], 5))
    rows = _table(calc[audit.ONE])
    assert rows and all(r["ties"] == TICK for r in rows)


# --------------------------------------------------------------------------
# 3 and 4: the Run finishes whatever happens to the audit workbook


def _record(b):
    return b.with_name(f"{b.stem} - what ran.yaml")


def test_an_audit_workbook_that_fails_does_not_end_the_run(tmp_path, monkeypatch):
    _, b = _build(tmp_path, "boom")

    def boom(*a, **k):
        raise RuntimeError("the disk is full")
    monkeypatch.setattr(perm, "SHUFFLES", SHUFFLES)
    monkeypatch.setattr(audit, "write", boom)
    audit.path_for(b).write_bytes(b"an earlier Run's audit workbook")      # left beside the book by an earlier Run
    ran = book.run(b)
    assert ran.ok, ran.lines
    assert (f"Couldn't write {audit.path_for(b).name}: RuntimeError: the disk is full The one from an earlier Run "
            f"was removed, so it can't be mistaken for this Run's.") in ran.lines
    assert _record(b).exists()
    assert ran.lines[-1].startswith(f"Open {b.name}: start with")
    assert not audit.path_for(b).exists()


def test_more_loans_than_an_excel_sheet_holds_skips_the_audit_workbook(tmp_path, monkeypatch):
    assert audit.MOST_LOANS == 1_048_576 - 1               # an Excel sheet's rows, less the header
    _, b = _build(tmp_path, "rows")
    monkeypatch.setattr(perm, "SHUFFLES", SHUFFLES)
    monkeypatch.setattr(audit, "MOST_LOANS", N - 1)
    ran = book.run(b)
    assert ran.ok, ran.lines
    said = [x for x in ran.lines if audit.path_for(b).name in x]
    assert said == [f"Didn't write {audit.path_for(b).name}: the book has {N:,} loans, and its Loans sheet holds at "
                    f"most {N - 1:,}, one a row under the header."]
    assert not audit.path_for(b).exists() and _record(b).exists()
    # at the limit exactly, it is written
    monkeypatch.setattr(audit, "MOST_LOANS", N)
    ran = book.run(b)
    assert ran.ok and audit.path_for(b).exists()


def test_a_long_text_in_a_formula_is_split_into_pieces_excel_accepts():
    """Excel drops a formula holding a string constant longer than 255 characters when it repairs the file, and
    LibreOffice takes it, so the recalculated tests can't see it (the review of 4 Oct 2026). q() splits a long text
    into pieces of at most 255 joined with &, and the pieces put back together are the text."""
    import re as _re
    text = 'A "quoted" segment, ' * 40                          # 800 characters, quotes included
    lit = audit.q(text)
    pieces = _re.findall(r'"((?:[^"]|"")*)"', lit)
    assert len(pieces) == 4 and all(len(p.replace('""', '"')) <= 255 for p in pieces)
    assert "".join(p.replace('""', '"') for p in pieces) == text
    assert lit == "&".join('"' + p + '"' for p in pieces)
    assert audit.q("short") == '"short"' and audit.q("") == '""'


def test_the_audit_cleans_a_control_character_itself_if_one_reaches_it(tmp_path, monkeypatch):
    """The extract is cleaned as it is read (ingest.cleaned shows \\x01 as U+2401), so in a Run no control character
    reaches the audit. The audit still cleans its own Loans sheet, so its XML stays readable whatever it is handed:
    with the reader's cleaning switched off, the loan number "L\\x015" still leaves a sheet that parses."""
    from pocketbook import ingest
    monkeypatch.setattr(ingest, "cleaned", lambda v: v)

    def edit(rows):
        rows[700]["LOAN_NBR"] = "L\x015"
    _, b = _build(tmp_path, "ctrl2", edit=edit)
    monkeypatch.setattr(perm, "SHUFFLES", SHUFFLES)
    ran = book.run(b)
    assert ran.ok, ran.lines
    z = zipfile.ZipFile(audit.path_for(b))
    from pocketbook.excel_lists import _sheet_files
    ET.fromstring(z.read(_sheet_files(z)[audit.LOANS]))
