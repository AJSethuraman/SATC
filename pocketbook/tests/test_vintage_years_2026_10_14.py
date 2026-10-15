"""Summary's vintage chart on a book with many origination years (14 Oct 2026).

Found: every _views lookup reads `results.AREA`, columns B to ZZ (701 numbers a row). The vintage block kept each
line as one row holding every measure, measure after measure, year after year, so a row was 17 x years wide on a
book with a Charge-off date marked. From 42 years the last measures' later years fell past ZZ: the helper cells read
#REF!, the drawn point became #N/A and the table under the chart showed a blank, with nothing said. At 100 years
whole measures went.

Fixed: one _views row per line and year, every measure across, so a row is as wide as Summary and the years are
rows. The review of the first fix (a row per line and measure) found it multiplied the rows by the measures, toward
Excel's 1,048,576 at the bank's size; so _views is also kept under results.MOST_VIEWS_ROWS: past it the vintage chart
is one sentence, and if the rest alone would pass it the Run refuses and saves nothing.

The figures checked against the loan file are worked out here from the CSV's text with plain arithmetic, as
test_summary_vintage_2026_10_03 and test_co_months_2026_10_01 do: nothing in the checks is imported from pocketbook."""

import csv
import re
import shutil
import statistics
from collections import Counter
from datetime import date

import pytest
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from pocketbook import book, results, summary_chart, synth
from test_firm_answers_2026_09_29 import _summary_label
from test_summary_vintage_2026_10_03 import _build, _measures, _views

import tabs
from recalc import need_soffice

YEARS = 42                          # the first year count the flat rows broke at, with a Charge-off date (17 measures)
MED = "Median months to charge-off"


def _spread(rows, n):
    """The synthetic loans spread over `n` origination years, loan i in year (2026 - n) + i mod n; each charge-off
    date moved by the same number of years, so the months to charge-off stay as they were."""
    base = 2026 - n
    for i, r in enumerate(rows):
        old = date.fromisoformat(r["ORIG_DATE"])
        ny = base + i % n
        r["ORIG_DATE"] = date(ny, old.month, min(old.day, 28)).isoformat()
        if r.get("CO_DATE"):
            c = date.fromisoformat(r["CO_DATE"])
            r["CO_DATE"] = date(c.year + ny - old.year, c.month, min(c.day, 28)).isoformat()
    return rows


@pytest.fixture(scope="module")
def made(tmp_path_factory):
    """The synthetic book with a charge-off date, its loans made over 42 years, filtered by CHANNEL."""
    d = tmp_path_factory.mktemp("vintage_years")
    src = synth.write_extract(d / "src", n=3000, chargeoff=True)
    rows = _spread(list(csv.DictReader(open(src, encoding="utf-8"))), YEARS)
    return _build(d, rows, "years", "CHANNEL"), rows


def _area_width(area: str) -> int:
    """How many numbers a lookup through `area` ('_views'!$B:$ZZ) can read on a row."""
    first, last = re.search(r"!\$([A-Z]+):\$([A-Z]+)$", area).groups()
    return column_index_from_string(last) - column_index_from_string(first) + 1


# --------------------------------------------------------------------------
# no LibreOffice


def test_no_views_row_is_wider_than_the_lookups_read(made):
    """Every row on _views fits in the columns every lookup reads (B to ZZ, 701), on a book whose 17 Summary measures
    over 42 years would not fit on one row; and no key is on _views twice, since a lookup finds only the first."""
    b, _ = made
    wide = _area_width(results.AREA)
    assert wide == 701
    heads = _measures(b)
    assert heads[-1] == MED and len(heads) * YEARS > wide      # the book the flat rows broke on
    over, keys = {}, []
    for row in load_workbook(b)[results.VIEWS].iter_rows(values_only=True):
        keys.append(row[0])
        last = max((j for j, x in enumerate(row) if x is not None), default=0)     # column A is j = 0
        if last > wide:
            over[row[0]] = last
    assert over == {}, f"{len(over)} rows wider than {wide}, e.g. {next(iter(over.items()))}"
    assert None not in keys and [k for k, n in Counter(keys).items() if n > 1] == []


def test_the_vintage_rows_are_a_line_and_year_each_as_wide_as_summary():
    """The vintage block alone, no Run: 120 years, 17 measures and a column of 300 rows (a long category list). Each
    row is one line in one year, as wide as Summary, the rows' labels one a row; each holds what the Run worked out
    for that line and year, measure by measure; no key twice; and a line with nothing in a year has no row."""
    from types import SimpleNamespace
    years = [str(1900 + y) for y in range(120)]
    keys = [f"m{k}" for k in range(17)]
    labels = [f"row {i}" for i in range(300)]
    Y, K = len(years), len(keys)
    flat = lambda tag: [f"{tag}{k}/{y}" for k in range(K) for y in range(Y)]     # noqa: E731
    book_ = flat("b")
    for k in range(K):
        book_[k * Y + 5] = None                         # 1905: nothing on the whole book's line
    key = ("COL", None, None, None)
    v = summary_chart.Vintage(years=years, keys=keys, book={key: book_},
                              rows={key: [(flat(f"p{i}."), flat(f"r{i}.")) for i in range(len(labels))]})
    res = SimpleNamespace(vintage=v, summaries={key: SimpleNamespace(labels=labels + ["All"])})
    put = []
    n = summary_chart.vintage_views(res, SimpleNamespace(put=lambda k, vals: put.append((k, list(vals)))),
                                    ["COL"], ["Col"])
    got = dict(put)
    assert n == len(put) == len(got) == 1 + len(labels) + (1 + 2 * len(labels)) * Y - 1
    assert max(len(x) for x in got.values()) == K
    assert got["V|Col|rows"] == [300] and [got[f"V|Col|rows|{i}"][0] for i in range(1, 301)] == labels
    assert "V|Col|book|y1905" not in got
    for y in (0, 41, 119):
        assert got[f"V|Col|book|y{1900 + y}"] == [f"b{k}/{y}" for k in range(K)]
        assert got[f"V|Col|300|y{1900 + y}"] == [f"p299.{k}/{y}" for k in range(K)]
        assert got[f"V|Col|300|rest|y{1900 + y}"] == [f"r299.{k}/{y}" for k in range(K)]


def _run_again(b, out, tmp_path, monkeypatch, most):
    """The book run again with _views allowed `most` rows: (the Run's outcome, the copy)."""
    from pocketbook import perm
    shutil.copy(b, out)
    monkeypatch.setattr(results, "MOST_VIEWS_ROWS", most)
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / f"memory-{out.stem}.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    return book.run(out), out


def test_views_past_its_most_rows_drops_the_vintage_chart_then_refuses(tmp_path, monkeypatch):
    """_views is never written past results.MOST_VIEWS_ROWS (1,000,000; Excel opens 1,048,576). Tried on the
    synthetic book with the most set low. Room for every row but the vintage chart's: the Run completes, and the chart
    is one sentence, in the note and where it would be, its rows left off and every other row as it was. Room for
    fewer: the Run refuses in words and the workbook is not saved, so it keeps the last Run's results."""
    assert results.MOST_VIEWS_ROWS == 1_000_000
    src = synth.write_extract(tmp_path / "src", n=3000)
    b = _build(tmp_path, list(csv.DictReader(open(src, encoding="utf-8"))), "rows", "CHANNEL")
    keys = [r[0] for r in load_workbook(b)[results.VIEWS].iter_rows(values_only=True)]
    vint = [k for k in keys if k.startswith("V|")]
    rest = [k for k in keys if not k.startswith("V|")]
    assert len(vint) > 100 and len(rest) > 100

    ran, b2 = _run_again(b, tmp_path / "no chart.xlsx", tmp_path, monkeypatch, len(rest))
    assert ran.ok, ran.lines
    wb = load_workbook(b2)
    assert [r[0] for r in wb[results.VIEWS].iter_rows(values_only=True)] == rest
    ws = wb[results.SUMMARY]
    want = (f"No vintage chart: its points would need {len(vint):,} more rows on the hidden _views sheet, taking it "
            f"past the {len(rest):,} this workbook writes (Excel opens at most 1,048,576).")
    said = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)
            and c.value.startswith("No vintage chart")]
    assert said == [want] * 2                                  # the note, and the block
    assert ws._charts == [] and wb[results.COMPARE]._charts          # Compare, written after Summary, is still there

    before = b.read_bytes()
    ran, b3 = _run_again(b, tmp_path / "refused.xlsx", tmp_path, monkeypatch, len(rest) - 1)
    assert not ran.ok
    assert ran.lines == [f"Couldn't write the results: even without the vintage chart they need {len(rest):,} rows "
                         f"on the hidden _views sheet, and the most this workbook writes is {len(rest) - 1:,} (Excel "
                         f"opens at most 1,048,576). Nothing was written, so the workbook still has the last Run's "
                         f"results. Use fewer filters, fewer filter values, or fewer band and category columns."]
    assert b3.read_bytes() == before


def test_views_past_its_most_rows_are_counted_not_written(monkeypatch):
    """Rows past results.MOST_VIEWS_ROWS are counted and not written, so the refusal above is reached with every row
    counted even when the rows without the vintage chart's pass Excel's 1,048,576, where openpyxl stops with a
    ValueError (Codex on #441). Tried with the most set at 5 and 8 rows put."""
    from openpyxl import Workbook
    monkeypatch.setattr(results, "MOST_VIEWS_ROWS", 5)
    views = results.Views(Workbook())
    assert [views.put(f"k{i}", [i, i + 1]) for i in range(8)] == list(range(1, 9))
    assert views.r == 8
    assert [r for r in views.ws.iter_rows(values_only=True)] == [(f"k{i}", i, i + 1) for i in range(5)]


# --------------------------------------------------------------------------
# LibreOffice


def _ymd(s):
    m = re.match(r"^(\d{4})-(\d\d)-(\d\d)$", s or "")
    return tuple(int(x) for x in m.groups()) if m else None


def _months(r):
    """A loan's months to charge-off, or None: whole calendar months, the day ignored (test_co_months_2026_10_01)."""
    o, c = _ymd(r["ORIG_DATE"]), _ymd(r["CO_DATE"])
    if o is None or c is None or c < o:
        return None
    return (c[0] - o[0]) * 12 + (c[1] - o[1])


def _median(rs):
    got = [m for m in (_months(r) for r in rs) if m is not None]
    return statistics.median(got) if got else None


def test_every_year_of_the_last_measure_is_in_the_table_and_on_the_chart(made, tmp_path):
    """LibreOffice: Vintage measure Median months to charge-off, the last of the 17, on the 42-year book. Every year
    of all three lines shows the loan file's median and loans in the table under the chart; no helper cell reads an
    error but #N/A; and each line of the chart reads its year's median, or #N/A where the year has fewer loans than
    Fewest loans in a pocket or no charged-off loan."""
    need_soffice()
    b, rows = made
    ws = tabs.calculated(tabs.choose(b, tmp_path / "med.xlsx", results.SUMMARY, **{"Vintage measure": MED}),
                         results.SUMMARY)
    h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
    labels = [ws.cell(row=r, column=2).value for r in range(h + 1, ws.max_row + 1)]
    labels = labels[:labels.index("All")]               # Summary's FICO rows, as the tab shows them
    band = labels[0]                                    # the Vintage row picked when the Run writes the tab
    bands = [x for x in labels if not x.startswith("(")]
    errors = sorted({(c.coordinate, c.value) for row in ws.iter_rows() for c in row
                     if isinstance(c.value, str) and c.value.startswith("#") and c.value != "#N/A"})
    assert errors == []
    note = next(c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                and "in Fewest loans in a pocket on Control is left off its line" in c.value)
    few = int(re.search(r"fewer loans than the ([\d,]+) in Fewest", note).group(1).replace(",", ""))
    t = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "Year")
    assert ws.cell(row=t, column=3).value == band
    years = sorted({r["ORIG_DATE"][:4] for r in rows})
    assert len(years) == YEARS
    shown, numbers = {}, 0
    for k, y in enumerate(years):
        r = t + 1 + k
        assert str(ws.cell(row=r, column=2).value) == y
        yr = [x for x in rows if x["ORIG_DATE"][:4] == y]
        mine = [x for x in yr if _summary_label(x["FICO"], bands) == band]
        others = [x for x in yr if _summary_label(x["FICO"], bands) != band]
        for c, rs in ((3, mine), (5, others), (7, yr)):
            want = _median(rs)
            v = ws.cell(row=r, column=c).value
            if want is None:
                assert v in (None, ""), (y, c, v)
            else:
                assert v == pytest.approx(want), (y, c, v)
                numbers += 1
            assert ws.cell(row=r, column=c + 1).value == len(rs), (y, c)
            shown[(c, y)] = (want, len(rs))
    assert numbers >= 2 * YEARS                     # the rest of the book and the whole book, at least, every year
    series = ws.formulas._charts[0].series
    assert len(series) == 3
    drawn = {}
    for c, s in zip((3, 5, 7), series):
        first = s.yVal.numRef.f.split("!")[1].replace("$", "").split(":")[0]
        col_, r0 = re.match(r"([A-Z]+)(\d+)", first).groups()
        for k, y in enumerate(years):
            v = ws[f"{col_}{int(r0) + k}"].value
            want, loans = shown[(c, y)]
            if want is None or loans < few:
                assert v == "#N/A", (c, y, v)
            else:
                assert v == pytest.approx(want), (c, y, v)          # months: drawn as they are, not in per cent
                drawn[c] = drawn.get(c, 0) + 1
    assert drawn.get(7) == YEARS and sum(drawn.values()) > 1.5 * YEARS, drawn       # the whole book, every year
