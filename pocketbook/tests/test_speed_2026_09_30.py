"""Where a Run's time goes, and what the workbook costs Excel to open (the firm, 30 Sep 2026).

"578 s" for a Run on about 17,000 loans x 70 columns at the bank, against about 22 s here, and nothing said which
part. And: "in excel it seems to work quickly enough but it takes quite some time to open particularly in the last
stretch of loading".

- Each stage of a Run is timed by the wall clock: Record's "Where the time went", a line of the Run's ("Took ...: the
  three biggest"), the record file, and `Outcome.timings`; `progress(stage)` is told as each starts. No test here
  asserts a number of seconds: only that the table is there, in the Run's order, and adds up.
- The extract and the workbook are each read from disk once (a OneDrive folder and a virus scanner at the bank).
- The workbook: no OFFSET (Excel works it out again at every change), no whole column of the Run's hidden tables,
  and Look's line ends read one "tallest bar" cell; every number reads as it did, checked in a calculated copy.
"""

from __future__ import annotations

import csv
import random
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from pocketbook import book, bounds, choices as ch, excel_lists, ingest, look, perm, synth, timing
from test_book import _answer

STAGES_EVERY_RUN = ("Opening the workbook", "Reading the answers", "Reading the extract", "Reading each loan's values",
                    "Cutting the bands", "Building the grids", "The shuffle test", "Building each filter's grids",
                    "Writing Pockets", "Writing Grids", "Writing Look", "Writing Record", "Saving the workbook")
SECONDS = re.compile(r"^(\d+(?:\.\d)?) s|^(\d+) min (\d+) s")


def _file(d: Path, n: int = 1500) -> Path:
    src = synth.write_extract(d / "src", n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    rng = random.Random(930)
    for i, r in enumerate(rows):
        r["SYS_FLAG"] = "" if i % 97 == 0 else rng.choice("YN")
    out = d / "loans.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return out


CHOSEN = ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",), split="REV_DEBT", filter="SYS_FLAG",
                    outcome="BAD_FLAG")


@pytest.fixture(scope="module")
def ran(tmp_path_factory):
    """A Run of the kind the bank makes (a score cut into bands, a segment, a split, a filter), with the stages each
    written down as the launcher's progress hears them."""
    d = tmp_path_factory.mktemp("speed")
    heard_set_up, heard = [], []
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 200)
        x = _file(d)
        out = book.set_up(x, choices=CHOSEN, progress=heard_set_up.append)
        assert out.ok, out.lines
        _answer(out.book)
        got = book.run(out.book, progress=heard.append)
        assert got.ok, got.lines
    return out, got, heard_set_up, heard


def _seconds(words: str) -> float:
    m = SECONDS.match(words)
    assert m, words
    return float(m.group(1)) if m.group(1) else int(m.group(2)) * 60 + int(m.group(3))


def _time_rows(b):
    import tabs
    rows = [(str(label).strip(), v) for label, v, *_ in tabs.record_rows(b) if label is not None]
    at = next(i for i, (label, _) in enumerate(rows) if label == book.TIME_HEAD)
    stages = []
    for label, v in rows[at + 1:]:
        if label == book.TIME_AFTER:
            break
        stages.append((label, v))
    return rows[at][1], stages


# --------------------------------------------------------------------------
# Seconds per stage


def test_every_stage_is_timed_in_the_order_it_ran_and_progress_hears_each_start(ran):
    out, got, _, heard = ran
    names = [k for k, _ in got.timings]
    for stage in STAGES_EVERY_RUN:
        assert stage in names, (stage, names)
    assert names.index("Reading the extract") < names.index("The shuffle test") < names.index("Writing Pockets") \
        < names.index("Saving the workbook")
    assert all(s >= 0 for _, s in got.timings)
    # progress is told each stage as it starts: every stage timed was heard, first heard first
    assert list(dict.fromkeys(heard)) == names
    assert heard[0] == "Opening the workbook"


def test_record_shows_where_the_time_went_and_it_adds_up(ran):
    out, got, _, _ = ran
    head, stages = _time_rows(out.book)
    labels = [k for k, _ in stages]
    ran_names = [k for k, _ in got.timings]
    # every stage finished before Record was written, in the Run's order; Record and the save come after
    assert labels == ran_names[:ran_names.index("Writing Record")]
    shown = [_seconds(v) for _, v in stages]
    total = _seconds(head)
    # each row rounded to a tenth (or a second), so the rows add up to the head within the rounding
    assert abs(sum(shown) - total) <= 0.05 * len(shown) + 1.0, (head, stages)
    assert head.startswith(f"{timing.took(total)} before Record was written. The biggest: ")
    notes = dict(stages)
    assert "(loans.csv: csv, 1,500 rows x 11 columns)" in notes["Reading the extract"]
    assert "200 shuffles of 1,500 loans on 1 process" in notes["The shuffle test"]
    assert "2 grids: 1 banded column by 1 segment, each split" in notes["Building the grids"]


def test_the_runs_lines_say_how_long_it_took_and_names_the_three_biggest(ran):
    _, got, _, _ = ran
    last = got.lines[-2]                  # just above "Open ...: start with", which stays the last line
    total = sum(s for _, s in got.timings)
    assert last.startswith(f"Took {timing.took(total)}: "), last
    big = sorted(got.timings, key=lambda kv: -kv[1])[:3]
    assert last == f"Took {timing.took(total)}: " + ", ".join(
        f"{k[0].lower() + k[1:]} {timing.took(s)}" for k, s in big) + "."
    # and the record file beside the workbook carries every stage, the save included
    said = (got.book.with_name(f"{got.book.stem} - what ran.yaml")).read_text(encoding="utf-8")
    assert "# Took " in said and ": Saving the workbook\n" in said


def test_set_up_is_timed_and_heard_too(ran):
    out, _, heard, _ = ran
    names = [k for k, _ in out.timings]
    assert names[0] == "Reading the extract" and "Saving the workbook" in names
    assert list(dict.fromkeys(heard)) == names


def test_a_progress_that_breaks_never_stops_a_run(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    x = _file(tmp_path, n=600)

    def broken(stage):
        raise RuntimeError("the window was closed")
    out = book.set_up(x, choices=CHOSEN, progress=broken)
    assert out.ok, out.lines
    _answer(out.book)
    got = book.run(out.book, progress=broken)
    assert got.ok, got.lines


def test_worker_processes_that_will_not_start_are_said_on_record(tmp_path, monkeypatch):
    """At the bank the shuffles go to worker processes; if Windows won't start them the test is dealt in one, which
    is correct and slow. Record says which it was."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 600)
    monkeypatch.setattr(perm, "workers_for", lambda *a, **k: 2)

    def refused(*a, **k):
        raise OSError("no processes on this machine")
    monkeypatch.setattr(perm, "_in_pool", refused)
    x = _file(tmp_path, n=600)
    out = book.set_up(x, choices=CHOSEN)
    _answer(out.book)
    assert book.run(out.book).ok
    _, stages = _time_rows(out.book)
    said = dict(stages)["The shuffle test"]
    assert "on 1 process; the worker processes didn't start (OSError: no processes on this machine)" in said


def test_the_clock_adds_a_stage_marked_twice_and_says_times_plainly(monkeypatch):
    ticks = iter(range(100))                     # a clock that moves one second each time it is read
    monkeypatch.setattr(timing.time, "perf_counter", lambda: float(next(ticks)))
    c = timing.Clock()
    with timing.running(c):
        timing.mark("A")
        timing.mark("B")
        timing.mark("A")
    assert c.rows() == [("A", 2.0), ("B", 1.0)]
    assert timing.took(0.04) == "0.0 s" and timing.took(9.84) == "9.8 s" and timing.took(42.4) == "42 s"
    assert timing.took(578) == "9 min 38 s"
    timing.mark("nobody is timing")          # no clock running: nothing happens
    c2 = timing.Clock()
    c2.seconds = {"Reading the extract": 64.0, "The shuffle test": 422.0, "Writing Look": 30.0, "Saving": 2.0}
    assert timing.took_line(c2) == ("Took 8 min 38 s: the shuffle test 7 min 2 s, reading the extract 1 min 4 s, "
                                    "writing Look 30 s.")


# --------------------------------------------------------------------------
# Each file read once


def test_an_xlsx_extract_is_read_from_the_bytes_already_read(tmp_path, monkeypatch):
    import openpyxl
    x = _file(tmp_path, n=50)
    rows = list(csv.reader(open(x, encoding="utf-8")))
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    p = tmp_path / "loans.xlsx"
    wb.save(p)
    given = []
    real = openpyxl.load_workbook
    monkeypatch.setattr(openpyxl, "load_workbook", lambda src, **kw: given.append(src) or real(src, **kw))
    t = ingest.read_table(p)
    assert len(given) == 1 and not isinstance(given[0], (str, Path)), given
    assert t.columns == rows[0] and [["" if v is None else v for v in r.values()] for r in t.rows] == rows[1:]
    assert t.kind == "xlsx"


def test_the_workbook_is_read_from_disk_once_and_a_sheet_without_excels_dropdowns_is_not_parsed_twice(ran,
                                                                                                        monkeypatch):
    out, _, _, _ = ran
    opened = []
    real_read = Path.read_bytes
    monkeypatch.setattr(Path, "read_bytes", lambda self: opened.append(self) or real_read(self))
    parsed = []
    real = excel_lists._from_sheet
    monkeypatch.setattr(excel_lists, "_from_sheet", lambda root: parsed.append(root) or real(root))
    wb = excel_lists.load(out.book)
    assert opened == [out.book]
    assert parsed == []                     # openpyxl wrote it: no sheet carries Excel's extension block
    assert "Control" in wb.sheetnames


# --------------------------------------------------------------------------
# What the workbook costs Excel to open


def _formulas(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    yield ws.title, c.coordinate, c.value
        for rng in ws.conditional_formatting:
            for rule in rng.rules:
                for f in rule.formula or []:
                    yield ws.title, f"CF {rng.sqref}", f
        for dv in ws.data_validations.dataValidation:
            for f in (dv.formula1, dv.formula2):
                if f:
                    yield ws.title, f"DV {dv.sqref}", f


def test_no_offset_and_no_whole_column_of_the_runs_tables(ran):
    out, _, _, _ = ran
    wb = load_workbook(out.book)
    every = list(_formulas(wb))
    assert not [f for f in every if "OFFSET(" in f[2].upper()]
    whole = re.compile(r"'?(_views|_pockets)'?!\$[A-Z]{1,3}:\$[A-Z]{1,3}(?![\w$])|\$1048576")
    wide = [f for f in every if whole.search(f[2]) and f[0] not in bounds.SKIP]
    assert not wide, wide[:5]
    last = wb["_pockets"].max_row
    for name in ("pk_kind", "pk_dollars", "pk_flag"):
        assert wb.defined_names[name].attr_text.endswith(f"${last}"), wb.defined_names[name].attr_text
    # Look's line ends read the one tallest-bar cell, not a MAX over the chart's 120 slots each
    assert not [f for s, _, f in every if s == look.DATA and f.startswith("=IF(ISNA(") and "MAX(" in f]


def test_bounding_changes_no_number(ran, tmp_path, monkeypatch):
    """The same Run's workbook with its whole columns put back, and as saved: calculated by LibreOffice, every cell
    of every tab reads the same."""
    import math
    from recalc import recalc
    out, _, _, _ = ran
    wb = load_workbook(out.book)
    last = {t: wb[t].max_row for t in bounds.TABLES}
    unbound = tmp_path / "whole.xlsx"
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    v = c.value
                    for t, n in last.items():
                        v = re.sub(rf"('{t}'!\$[A-Z]+)\$1(:\$[A-Z]+)\${n}(?![\d])", r"\1\2", v)
                        v = re.sub(rf"('{t}'!\$[A-Z]+\$\d+:\$[A-Z]+)\${n}(?![\d])", r"\1$1048576", v)
                    c.value = v
    wb.save(unbound)
    assert bounds.bound(load_workbook(unbound)) > 0          # there was something to bound
    a, b = recalc(unbound, tmp_path / "a"), recalc(out.book, tmp_path / "b")
    seen = diff = 0
    for ws in a.worksheets:
        other = b[ws.title]
        for row in ws.iter_rows():
            for c in row:
                y = other[c.coordinate].value
                seen += 1
                if not (c.value == y or isinstance(c.value, float) and isinstance(y, float)
                        and math.isclose(c.value, y, rel_tol=1e-12, abs_tol=1e-12)):
                    diff += 1
    assert seen > 10_000 and diff == 0


def test_looks_tallest_bar_is_the_tallest_bar(ran, tmp_path):
    from recalc import recalc
    out, _, _, _ = ran
    hs = recalc(out.book, tmp_path / "rc")[look.DATA]
    g = 2
    while hs.cell(row=look.DATA_TOP - 1, column=g).value:
        slots = [hs.cell(row=look.DATA_TOP + k, column=g + look.G_SLOT).value for k in range(look.SLOTS)]
        tall = hs.cell(row=look.S_TALL, column=g + look.G_VALUE).value
        assert tall == max(v for v in slots if isinstance(v, (int, float)))
        tops = [hs.cell(row=look.PCT_TOP + 2 * k + 1, column=g + look.G_EY).value for k in range(5)]
        assert all(t == tall or t in (None, "#N/A") for t in tops) and tall in tops
        g += look.GROUP
    assert g > 2
