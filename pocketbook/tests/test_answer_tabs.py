"""The redesign, phase 2 (docs/redesign-2026-09-26/README.md, "Global rules" and sections 1 to 4): the tabs the
analyst fills in. Start here, Control, Columns and Look, each with the title band, one method note that folds
away, the tab colour, gridlines off and the input styles; Control in three blocks with Status and the
materiality panel; Columns with the odd values and the memory on each column's row; Look with the mean, live
bars and range, red edge lines from Columns and the likely code on a bar of its own. And a Run that loads the
workbook once and saves it once, and doesn't redraw Look.

The live parts are proved by calculating the workbook through LibreOffice (tests/recalc.py): an answer is changed
with openpyxl, the copy recalculated, and the cells the charts and tiles read compared with a count by hand."""

import csv
import re
import statistics
from pathlib import Path

import openpyxl
import pytest
from openpyxl import load_workbook
from openpyxl.chart import BarChart, ScatterChart

from conftest import TEST_SHUFFLES
from pocketbook import book, choices as ch, control, house, launcher, look, memory, perm, synth
from recalc import recalc
from test_book import _answer
from test_book_dates import _choose, _columns, _control

ANSWER_TABS = ("Start here", "Control", "Columns", "Look")
EDGES = "620; 680; 740"


@pytest.fixture(scope="module")
def ran(tmp_path_factory):
    """The synthetic book, answered, FICO cut at 620; 680; 740 and the pockets split by REV_DEBT, and Run."""
    d = tmp_path_factory.mktemp("answer-tabs")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
        x = synth.write_extract(d, n=4000)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, split="REV_DEBT"))
        assert out.ok, out.lines
        _answer(out.book)
        _columns(out.book, "FICO", C_EDGES=EDGES)
        got = book.run(out.book)
        assert got.ok, got.lines
    return x, out.book


def _key_rows(ws) -> dict:
    return {r[control.KEY_COL - 1].value: r[0].row for r in ws.iter_rows(min_row=control.FIRST_ROW)
            if r[control.KEY_COL - 1].value}


def _hex(color) -> str:
    return (color.rgb or "")[-6:] if color is not None and isinstance(color.rgb, str) else ""


def _calc(path: Path, tmp: Path, edit=None):
    """The workbook at `path` calculated by LibreOffice, after `edit(wb)` on a copy."""
    wb = load_workbook(path)
    if edit:
        edit(wb)
    tmp.mkdir(parents=True, exist_ok=True)
    p = tmp / path.name
    wb.save(p)
    return recalc(p, tmp / "calc")


# --------------------------------------------------------------------------
# Global rules


def test_the_four_tabs_you_fill_in_come_first_red_and_nothing_folded_is_left(ran):
    _, b = ran
    wb = load_workbook(b)
    assert wb.sheetnames[:4] == list(ANSWER_TABS)
    for gone in ("Odd values", "Learned", "Materiality"):
        assert gone not in wb.sheetnames
    for t in ANSWER_TABS:
        assert wb[t].sheet_properties.tabColor.rgb[-6:] == house.KEY_RED, t


def test_each_tab_opens_with_the_title_band_and_one_method_note_that_folds_away(ran):
    _, b = ran
    wb = load_workbook(b)
    for t in ANSWER_TABS:
        ws = wb[t]
        title = ws["B1"]
        assert title.value in (t, book.NAME) and _hex(title.fill.fgColor) == house.INK, t
        assert title.font.b and title.font.sz == 16 and title.font.name == "Arial", t
        assert title.border.bottom.style == "thick" and _hex(title.border.bottom.color) == house.KEY_RED, t
        assert ws.sheet_view.showGridLines is False, t
        notes = [c for row in ws.iter_rows() for c in row if c.value == "How this tab works"]
        assert len(notes) == 1, t                                   # T1: one method note, never beside rows
        top = notes[0].row
        grouped = [r for r in range(top, top + 10) if ws.row_dimensions[r].outline_level == 1]
        assert grouped and grouped[0] == top and grouped == list(range(top, top + len(grouped))), t
        assert _hex(notes[0].fill.fgColor) == house.CANVAS, t
        assert ws.freeze_panes, t


def test_result_rows_on_columns_are_one_line_the_same_height_numbers_centred(ran):
    _, b = ran
    ws = load_workbook(b)["Columns"]
    rows = list(book.table_rows(ws))
    assert rows
    assert {ws.row_dimensions[r[0].row].height for r in rows} == {18}
    for r in rows:
        assert not any(c.alignment.wrap_text for c in r[book.C_NAME - 1:book.C_DEFINE]), r[0].row
        assert r[book.C_BLANK - 1].alignment.horizontal == "center"
        assert r[book.C_NAME - 1].alignment.horizontal in (None, "left")


# --------------------------------------------------------------------------
# Control


def test_control_holds_changes_now_then_needs_a_run_then_the_launchers_choices(ran):
    _, b = ran
    ws = load_workbook(b)[control.SHEET]
    rows = _key_rows(ws)
    a, bb, c = rows[control.BLOCK_NOW], rows[control.BLOCK_RUN], rows["launcher|head"]
    assert a < bb < c
    assert [k for k, r in sorted(rows.items(), key=lambda t: t[1]) if a < r < bb] == list(control.NOW_KEYS)
    # the cutoff (OC-51) closes Block B: asked only when scouting, it moves no row a bleed run reads
    assert [k for k, r in sorted(rows.items(), key=lambda t: t[1]) if bb < r < c] == list(control.RUN_KEYS) + ["cutoff"]
    assert "run_kind" in rows and rows["run_kind"] > c
    # the bands: INK with a red rule, SLATE with a STONE rule, MIST
    assert _hex(ws.cell(row=a, column=2).fill.fgColor) == house.INK
    assert _hex(ws.cell(row=a, column=2).border.bottom.color) == house.KEY_RED
    assert _hex(ws.cell(row=bb, column=2).fill.fgColor) == house.SLATE
    assert _hex(ws.cell(row=bb, column=2).border.bottom.color) == house.STONE
    assert _hex(ws.cell(row=c, column=2).fill.fgColor) == house.MIST
    heads = [ws.cell(row=a + 1, column=k).value for k in (2, 3, 5, 6)]
    assert heads == ["Setting", "Your answer", "Comes to", "Last Run used"]
    assert ws.cell(row=bb + 1, column=control.STATUS_COL).value == "Status"
    assert ws.cell(row=bb + 1, column=3).value.endswith("↻")


def test_the_answer_cells_carry_the_three_input_styles_and_the_legend_shows_them(ran):
    _, b = ran
    ws = load_workbook(b)[control.SHEET]
    rows = _key_rows(ws)
    now = ws.cell(row=rows["worse_at"], column=control.CHOOSE_COL)
    run = ws.cell(row=rows["min_events"], column=control.CHOOSE_COL)
    assert _hex(now.fill.fgColor) == house.CANVAS and now.border.left.style == "thin" and now.font.b
    assert _hex(run.fill.fgColor) == house.PAPER and run.border.left.style == "dashed" and run.font.b
    assert _hex(run.border.left.color) == house.SLATE
    legend = rows["legend"]
    said = [ws.cell(row=legend, column=k).value for k in (2, 3, 6)]
    assert said == ["Changes now", "Needs a Run", "Still needs an answer"]
    assert _hex(ws.cell(row=legend, column=6).fill.fgColor) == house.ALERT_FG
    # the shade on a blank answer is ALERT_FG, by conditional format
    fills = {_hex(rule.dxf.fill.fgColor) for rng in ws.conditional_formatting for rule in rng.rules
             if rule.dxf is not None and rule.dxf.fill is not None and "=\"\"" in "".join(rule.formula)}
    assert house.ALERT_FG in fills


def test_control_keeps_the_values_worked_out_from_the_extract_beside_their_settings(ran):
    _, b = ran
    ws = load_workbook(b)[control.SHEET]
    rows = _key_rows(ws)
    for key in book.SUGGEST_KEYS:
        said = ws.cell(row=rows[key], column=book.SUGGEST_COL).value
        assert said.startswith("suggested: ") and said.endswith("from this extract at the last Run"), key
    assert ws.cell(row=rows[control.BLOCK_NOW] + 1, column=book.SUGGEST_COL).value == "Worked out from the loans"


def test_status_says_waiting_only_for_an_answer_that_differs_from_the_last_run(ran, tmp_path):
    _, b = ran
    same = _calc(b, tmp_path / "same")
    ws = same[control.SHEET]
    rows = _key_rows(load_workbook(b)[control.SHEET])
    assert [ws.cell(row=rows[k], column=control.STATUS_COL).value for k in control.RUN_KEYS] == \
        [house.SAME] * len(control.RUN_KEYS)
    assert all(ws.cell(row=rows[k], column=control.STATUS_COL).value in (None, "") for k in control.NOW_KEYS)
    start = same["Start here"]
    assert _tile(start, "Changes waiting for a Run") == 0
    assert not _banner(start)

    def change(wb):
        c = wb[control.SHEET].cell(row=rows["min_events"], column=control.CHOOSE_COL)
        c.value = "20 losses"

    moved = _calc(b, tmp_path / "moved", change)
    ws = moved[control.SHEET]
    got = {k: ws.cell(row=rows[k], column=control.STATUS_COL).value for k in control.RUN_KEYS}
    assert got.pop("min_events") == house.WAITING
    assert set(got.values()) == {house.SAME}
    start = moved["Start here"]
    assert _tile(start, "Changes waiting for a Run") == 1
    banner = _banner(start)
    assert banner.startswith("↻ 1 change is waiting for a Run. Fewest loans with a loss before a loss rate is "
                             "tested: 10 losses → 20 losses (Control C")
    assert banner.endswith("The result tabs still show the last Run. Save, close, and press Run in the launcher.")


def _tile(ws, label):
    for row in ws.iter_rows():
        for c in row:
            if c.value == label:
                return ws.cell(row=c.row + 1, column=c.column).value
    raise KeyError(label)


def _banner(ws) -> str:
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("↻ "):
                return c.value
    return ""


def _pockets(values_wb):
    """Every grid pocket's deciding charge-off dollars on _pockets, as calculated."""
    from pocketbook import live
    ws = values_wb[live.POCKETS]
    out = []
    for r in range(live.P_FIRST, ws.max_row + 1):
        if ws.cell(row=r, column=live.P_KIND).value == "grids" and \
                ws.cell(row=r, column=live.P_MEASURE).value == "gco_rate":
            v = ws.cell(row=r, column=live.P_DOLLARS).value
            out.append(v if isinstance(v, (int, float)) else None)
    return out


def test_the_materiality_panel_keeps_what_each_level_would_keep_and_marks_the_one_in_use(ran, tmp_path):
    _, b = ran
    got = _calc(b, tmp_path / "panel")
    ws, live_ws = got[control.SHEET], got["_live"]
    from pocketbook import live
    total = live_ws.cell(row=live.L_GCO_TOTAL, column=3).value
    dollars = [d for d in _pockets(got) if d is not None and d > 0]
    head = next(c for row in ws.iter_rows() for c in row if c.value == "What each materiality level keeps")
    levels = []
    for r in range(head.row + 2, head.row + 7):
        level, line, kept, share = (ws.cell(row=r, column=head.column + k).value for k in range(4))
        levels.append(level)
        pct = float(level.split("%")[0]) / 100
        assert line == pytest.approx(pct * total)
        assert kept == sum(1 for d in dollars if d >= pct * total)
        assert share == pytest.approx(sum(d for d in dollars if d >= pct * total) / sum(dollars))
    assert levels == ["0.5%", "1% ◂", "2%", "5%", "10%"]            # the answer in use: 1%

    def five(wb):
        c = wb[control.SHEET].cell(row=_key_rows(wb[control.SHEET])["materiality"], column=control.CHOOSE_COL)
        c.value = "5% of the book's total losses"

    moved = _calc(b, tmp_path / "panel5", five)
    ws = moved[control.SHEET]
    assert [ws.cell(row=r, column=head.column).value for r in range(head.row + 2, head.row + 7)] == \
        ["0.5%", "1%", "2%", "5% ◂", "10%"]


def test_a_new_variable_run_is_not_asked_the_profit_line_and_check_does_not_echo_it(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = synth.write_extract(tmp_path, n=1500)
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.NEW_VARIABLE))
    wb = load_workbook(out.book)
    ws = wb[control.SHEET]
    for r in ws.iter_rows(min_row=control.FIRST_ROW):
        k = r[control.KEY_COL - 1].value
        if k in {**_PICK, "revenue_line": None}:
            r[control.CHOOSE_COL - 1].value = _PICK.get(k)
    wb.save(out.book)
    used = control.read_control(out.book)
    assert "revenue_line" not in used
    # blank and unasked, it isn't shaded as needing an answer
    rule = next(r for rng in ws.conditional_formatting for r in rng.rules
                if f"$C{_key_rows(ws)['revenue_line']}" in "".join(r.formula))
    assert "Where the book bleeds" in "".join(rule.formula)
    assert "revenue_line" not in book._settings_of(_config_without_profit_line())


_PICK = {"min_loans": "30", "min_events": "10", "materiality": "1% of the book's total losses",
         "compare_to": "The rest of its band", "worse_at": "1.25 times", "better_at": "0.8 times",
         "confidence": "95%", "cutoff": "The month start nearest 70% of the loans (suggested)"}


def _config_without_profit_line():
    from conftest import cube
    return cube()                                   # a cube file with no profit line, as a new-variable run reads


# --------------------------------------------------------------------------
# Columns


def test_columns_names_each_column_once_with_its_odd_values_and_memory_on_its_row(ran):
    x, b = ran
    ws = load_workbook(b)["Columns"]
    assert ws[book.CONFIRM_CELL].value == "Yes" and ws["B3"].value == "Checked every column?"
    heads = [ws.cell(row=book.COL_HEAD, column=c).value for c in range(book.C_NAME, book.C_REMEMBERED + 1)]
    assert heads == ["Column", "Samples", "What it is ↻", "Why we think so", "Blank", "Odd values",
                     "Treat as ↻", "Band edges ↻", "Remembered"]
    with open(x, newline="", encoding="utf-8") as fh:
        header = next(csv.reader(fh))
    names = [r[book.C_NAME - 1].value for r in book.table_rows(ws) if r[book.C_NAME - 1].value]
    assert names == header
    fico = next(r for r in book.table_rows(ws) if r[book.C_NAME - 1].value == "FICO")
    assert fico[book.C_ODD - 1].value == "-9999 on 80 loans"          # every 50th loan of 4,000
    assert fico[book.C_TREAT - 1].value == "Missing" and fico[book.C_EDGES - 1].value == EDGES
    ranr = next(r for r in book.table_rows(ws) if r[book.C_NAME - 1].value == "RANR_AMT")
    assert ranr[book.C_ODD - 1].value.startswith("Negative on ") and ranr[book.C_TREAT - 1].value is None
    # the odd value is said once, beside its answer, not again under Check first
    assert "Treat as" not in str(fico[book.C_LOOK - 1].value or "")
    # the new-column block sits under the table
    keys = [r[book.C_QKEY - 1].value for r in ws.iter_rows(min_row=book.COL_FIRST) if len(r) >= book.C_QKEY]
    assert keys.index(book.TABLE_END) < keys.index("derived|head") < keys.index("derived|1")


def test_treat_as_on_columns_is_what_the_run_reads(ran):
    _, b = ran
    raw, problems, _ = book.read_book(b)
    assert not problems
    fico = [q for q in raw["questions"] if q["column"] == "FICO"]
    assert fico and all(q["answer"] == "missing" for q in fico) and fico[0]["rows"] == 80
    ranr = [q for q in raw["questions"] if q["column"] == "RANR_AMT"]
    assert ranr and all(q["answer"] is None for q in ranr)


def test_treat_as_takes_real_or_missing_and_a_typed_word_is_refused_by_cell(ran, tmp_path):
    _, b = ran
    copy = tmp_path / b.name
    wb = load_workbook(b)
    row = book.table_rows(wb["Columns"])
    r = next(r for r in row if r[book.C_NAME - 1].value == "FICO")[0].row
    wb["Columns"].cell(row=r, column=book.C_TREAT).value = "maybe"
    wb.save(copy)
    _, problems, _ = book.read_book(copy)
    assert f"Columns!{book._col(book.C_TREAT)}{r}: Treat as takes Real or Missing, or blank." in problems


def test_remembered_shows_on_each_row_and_forget_drops_it_at_the_next_run(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    x = synth.write_extract(tmp_path, n=2000)
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED))
    _answer(out.book)
    assert book.run(out.book).ok
    assert "FICO" in memory.load()["columns"]
    again = book.set_up(x)
    ws = load_workbook(again.book)["Columns"]
    fico = next(r for r in book.table_rows(ws) if r[book.C_NAME - 1].value == "FICO")
    assert fico[book.C_REMEMBERED - 1].value == "Yes · 1 Run" and fico[book.C_FORGET - 1].value == "No"
    _columns(again.book, "FICO", C_FORGET="Yes")
    ran = book.run(again.book)
    assert ran.ok, ran.lines
    assert "FICO" not in memory.load()["columns"]
    ws = load_workbook(again.book)["Columns"]
    assert ws[book.CONFIRM_CELL].value is None                        # waits for the analyst to confirm it again
    fico = next(r for r in book.table_rows(ws) if r[book.C_NAME - 1].value == "FICO")
    assert fico[book.C_REMEMBERED - 1].value == "No" and "Forgotten" in fico[book.C_LOOK - 1].value
    assert "Forgotten: FICO" in ws["D3"].value


# --------------------------------------------------------------------------
# Start here


def test_start_here_counts_what_is_left_and_what_the_last_run_found(ran, tmp_path):
    _, b = ran
    got = _calc(b, tmp_path / "start")
    ws = got["Start here"]
    assert _tile(ws, book.NEEDED) == 0
    assert _tile(ws, "Odd values to answer") == 1                     # RANR_AMT's negatives
    found = _tile(ws, "Pockets worse and material, charge-offs")
    k, of = found.split(" of ")
    dollars = [d for d in _pockets(got) if d is not None]
    assert int(of.replace(",", "")) == len(dollars)
    # the tie-out tile is gone: a Run that doesn't tie out stops, so it could only read fine (tenet T2)
    assert [c.coordinate for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "Tie-out" in c.value] \
        == []
    heads = [c.value for row in ws.iter_rows() for c in row if c.value == "Largest, worse and material"]
    assert heads
    groups = [c.value for row in ws.iter_rows() for c in row if c.value in ("You answer", "Results", "Record")]
    assert groups == ["You answer", "Results", "Record"]
    links = {c.hyperlink.location if c.hyperlink.location else c.hyperlink.target
             for row in load_workbook(b)["Start here"].iter_rows() for c in row if c.hyperlink}
    assert any("Control" in str(x) for x in links) and any("Look" in str(x) for x in links)


def test_a_blank_answer_and_an_odd_value_are_counted_on_start_here_as_they_are_left(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    out = book.set_up(synth.write_extract(tmp_path, n=1500), choices=ch.Choices(run_kind=ch.BLEED))
    got = _calc(out.book, tmp_path / "calc0")
    ws = got["Start here"]
    # the bleed's: the cutoff is asked only when scouting (OC-51)
    judged = [s for s in control.load_settings() if s.judgment and not s.in_launcher
              and control.asked(s, {"run_kind": ch.BLEED})]
    # the Run's refusal's own count: each blank answer on Control, and Checked every column
    assert _tile(ws, book.NEEDED) == len(judged) + 1 == len(book.read_book(out.book)[1])
    assert _tile(ws, "Odd values to answer") >= 2
    assert any(str(c.value).startswith("Nothing yet") for row in ws.iter_rows() for c in row)


# --------------------------------------------------------------------------
# Look


def _fico(x) -> list[float]:
    with open(x, newline="", encoding="utf-8") as fh:
        vals = [r["FICO"] for r in csv.DictReader(fh)]
    return [float(v) for v in vals if v not in ("",) and float(v) != -9999]


def _block_row(ws, name) -> int:
    return next(c.row for c in ws["B"] if c.value == name and _hex(c.fill.fgColor) == house.INK)


def test_look_shows_the_mean_beside_the_median_and_the_code_on_a_red_bar_of_its_own(ran):
    x, b = ran
    wb = load_workbook(b)
    ws = wb["Look"]
    r = _block_row(ws, "FICO")
    vals = _fico(x)
    lines = {ws.cell(row=r + k, column=2).value: ws.cell(row=r + k, column=3).value for k in range(1, 9)}
    assert lines["Median"] == statistics.median(vals)
    assert lines["Mean"] == pytest.approx(statistics.fmean(vals))
    assert lines["At -9999, likely a code"] == 80
    reds = [c for c in ws._charts if isinstance(c, BarChart) and c.anchor._from.row + 1 == r + 1
            and c.anchor._from.col + 1 == 6]
    assert len(reds) == 1
    fill = reds[0].series[0].graphicalProperties.solidFill
    rgb = fill if isinstance(fill, str) else fill.srgbClr
    assert (rgb if isinstance(rgb, str) else rgb.val) == house.KEY_RED
    # ORIG_BAL has no code, so no red bar
    r2 = _block_row(ws, "ORIG_BAL")
    assert not [c for c in ws._charts if isinstance(c, BarChart) and c.anchor._from.row + 1 == r2 + 1
                and c.anchor._from.col + 1 == 6]


def _feed(values_wb, name):
    """What the FICO chart draws, as calculated: the bars (each bar's count once), the two end bars, and the
    edge lines' positions."""
    hs = values_wb[look.DATA]
    g = next(c for c in range(2, hs.max_column + 1) if hs.cell(row=look.DATA_TOP - 1, column=c).value == name)
    bars = [hs.cell(row=look.DATA_TOP + j, column=g + look.G_BAR).value for j in range(max(look.BARS))]
    bars = [v for v in bars if isinstance(v, (int, float))]
    low = hs.cell(row=look.S_LOWTAIL, column=g + look.G_VALUE).value
    high = hs.cell(row=look.S_HIGHTAIL, column=g + look.G_VALUE).value
    slots = [hs.cell(row=look.DATA_TOP + k, column=g + look.G_SLOT).value for k in range(look.SLOTS)]
    edges = [hs.cell(row=look.DATA_TOP + k, column=g + look.G_EDGE).value for k in range(look.EDGE_LINES)]
    xs = [hs.cell(row=look.DATA_TOP + k, column=g + look.G_EX).value for k in range(2 * look.EDGE_LINES)]
    return bars, low, high, slots, [e for e in edges if isinstance(e, (int, float))], xs


def test_the_bars_regroup_live_when_bars_and_range_change(ran, tmp_path):
    x, b = ran
    vals = _fico(x)
    lo, hi = look.span(sorted(vals))
    got = _calc(b, tmp_path / "bars20")
    bars, low, high, slots, _, _ = _feed(got, "FICO")
    want, wlow, whigh = look.regroup(sorted(vals), 20, lo, hi, lo, hi)
    assert (bars, low, high) == (want, wlow, whigh) and sum(bars) + low + high == len(vals)
    # the chart's slots hold each bar's count, in order, with a gap after each
    per = look.MIDDLE // 20
    assert [slots[look.LOW + per * j] for j in range(20)] == want
    assert all(slots[look.LOW + per * j + per - 1] == 0 for j in range(20))

    ws = load_workbook(b)["Look"]
    r = _block_row(ws, "FICO")

    def narrower(wb):
        w = wb["Look"]
        w.cell(row=r + look.R_BARS, column=look.VALUE_COL).value = 10
        w.cell(row=r + look.R_FROM, column=look.VALUE_COL).value = 600
        w.cell(row=r + look.R_TO, column=look.VALUE_COL).value = 800

    got = _calc(b, tmp_path / "bars10", narrower)
    bars, low, high, slots, _, _ = _feed(got, "FICO")
    want, wlow, whigh = look.regroup(sorted(vals), 10, lo, hi, 600, 800)
    assert (bars, low, high) == (want, wlow, whigh) and len(bars) == 10
    assert sum(bars) + low + high == len(vals)
    assert slots[look.LOW - 1] == 0 and slots[1] == wlow          # the low end bar, then the gap before the bars


def test_the_red_lines_follow_the_band_edges_typed_on_columns(ran, tmp_path):
    x, b = ran
    vals = sorted(_fico(x))
    lo, hi = look.span(vals)
    _, _, _, _, edges, xs = _feed(_calc(b, tmp_path / "edges"), "FICO")
    assert edges == [620, 680, 740]
    for k, e in enumerate(edges):
        at = look.LOW + 0.5 + look.MIDDLE * (e - lo) / (hi - lo)
        assert xs[2 * k] == pytest.approx(at) and xs[2 * k + 1] == pytest.approx(at)
    assert all(not isinstance(v, (int, float)) for v in xs[2 * len(edges):])      # no fourth line

    def every(wb):
        for r in book.table_rows(wb["Columns"]):
            if r[book.C_NAME - 1].value == "FICO":
                r[book.C_EDGES - 1].value = "every 100"

    _, _, _, _, edges, _ = _feed(_calc(b, tmp_path / "every", every), "FICO")
    first = (min(vals) // 100) * 100 + 100
    assert edges == [first + 100 * k for k in range(len(edges))] and edges[-1] <= max(vals) < edges[-1] + 100


def test_look_is_not_redrawn_at_a_run_unless_the_split_moves(ran, tmp_path, monkeypatch):
    _, b = ran
    copy = tmp_path / b.name
    copy.write_bytes(b.read_bytes())
    before = load_workbook(copy)
    feed = [[c.value for c in row] for row in before[look.DATA].iter_rows()]
    charts = len(before["Look"]._charts)
    drawn = []
    real = look._scatters
    monkeypatch.setattr(look, "_scatters", lambda *a, **k: drawn.append(a[5]) or real(*a, **k))
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    assert book.run(copy, extract=b.with_name("loans.csv")).ok
    assert drawn == []
    after = load_workbook(copy)
    assert [[c.value for c in row] for row in after[look.DATA].iter_rows()] == feed
    assert len(after["Look"]._charts) == charts
    assert {type(c) for c in after["Look"]._charts} >= {BarChart, ScatterChart}
    _choose(copy, split=None)
    assert book.run(copy, extract=b.with_name("loans.csv")).ok
    assert drawn == [None]
    assert not [c for c in load_workbook(copy)["Look"]._charts if isinstance(c, ScatterChart)]


# --------------------------------------------------------------------------
# One load and one save


def test_a_run_loads_the_workbook_once_and_saves_it_once(ran, tmp_path, monkeypatch):
    _, b = ran
    copy = tmp_path / b.name
    copy.write_bytes(b.read_bytes())
    loads, saves = [], []
    real_load, real_save = openpyxl.load_workbook, openpyxl.Workbook.save

    def load(path, *a, **k):
        if Path(str(path)).name == copy.name:
            loads.append(path)
        return real_load(path, *a, **k)

    def save(self, path):
        if Path(str(path)).name == copy.name:
            saves.append(path)
        return real_save(self, path)

    from pocketbook import excel_lists
    for mod in (openpyxl, book, control, excel_lists):             # excel_lists: book opens it through there
        monkeypatch.setattr(mod, "load_workbook", load)
    monkeypatch.setattr(openpyxl.Workbook, "save", save)
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    assert book.run(copy, extract=b.with_name("loans.csv")).ok
    assert (len(loads), len(saves)) == (1, 1)
    # a refused Run too: its Log line and nothing else
    loads.clear(), saves.clear()
    _control(copy, min_events=None)
    loads.clear(), saves.clear()
    assert not book.run(copy, extract=b.with_name("loans.csv")).ok
    assert (len(loads), len(saves)) == (1, 1)


# --------------------------------------------------------------------------
# Words


BANNED = re.compile(r"\b(luck|cube)\b", re.IGNORECASE)


def test_nothing_the_analyst_reads_says_luck_or_cube(ran):
    _, b = ran
    wb = load_workbook(b)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    assert not BANNED.search(c.value), f"{ws.title}!{c.coordinate}: {c.value[:120]}"
    for s in control.load_settings():
        for o in s.options:
            assert not BANNED.search(o.label + " " + o.explains), o.label
    flow = launcher.Flow()
    words = list(launcher.START) + list(flow.heads()) + list(launcher.STEPS)
    flow.mode = "new"
    words += list(flow.heads())
    for w in words:
        assert not BANNED.search(w), w
    assert "Split by" in launcher.Flow().heads()
