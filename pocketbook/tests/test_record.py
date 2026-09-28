"""Record (the redesign, section 10; phase 4): Check and the Log on one grey tab, in a 2 x 3 grid of paired sections.

    This Run          | Settings
    Does it add up    | Tests used
    Left out          | Every Run, newest first

Every section has the ONYX band, a CANVAS row naming its columns, label columns of one width, rows of one height
unless their words need more, and a 2 px ONYX rule under it; a pair starts and ends on the same row. Every line
Check carried is on it, and so is every Log entry: the Log is kept on the hidden _log, so a refused Run shows on
Every Run at once. Settings has each setting's answer in use now beside the one the last Run used, shaded while
they differ.

The lines are checked against what the Run worked out (book._record_rows, which Check's writer became) and against
the engine's own numbers, never against the tab's own reading of itself."""

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

import tabs
from pocketbook import book, control, house, live, record, synth, choices as ch
from recalc import recalc
from test_book import _answer
from test_book_dates import _control


@pytest.fixture(scope="module")
def ran(tmp_path_factory):
    folder = tmp_path_factory.mktemp("record")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(folder / "memory.yaml"))
        x = synth.write_extract(folder, n=3000, ratio=True)
        b = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO", "ORIG_BAL"),
                                              segments=("CHANNEL",), split="REV_DEBT")).book
        _answer(b)
        seen = {}
        real = book._record_rows

        def spy(*a, **k):
            seen["rows"] = real(*a, **k)
            return seen["rows"]

        mp.setattr(book, "_record_rows", spy)
        first = book.run(b)
        second = book.run(b)
    assert first.ok and second.ok, (first.lines, second.lines)
    return {"x": x, "b": b, "rows": seen["rows"], "ran": second}


def _sections(ws) -> dict:
    """Each section's band row, header row and last row, by its title, read from the tab's cells."""
    titles = {v[0]: k for k, v in record.SECTIONS.items()}
    out = {}
    for r in range(1, ws.max_row + 1):
        for c in (record.L_LABEL, record.R_LABEL):
            v = ws.cell(row=r, column=c).value
            if v in titles and ws.cell(row=r, column=c).fill.fgColor.rgb[-6:] == house.ONYX:
                out[titles[v]] = {"band": r, "col": c}
    for kind, s in out.items():
        last_col = record.L_TEXT if s["col"] == record.L_LABEL else record.R_B
        r = s["band"] + 2
        while ws.cell(row=r, column=last_col).border.bottom.style != "medium":
            r += 1
            assert r < ws.max_row + 2, kind
        s["last"] = r
    return out


def test_record_is_one_grey_tab_where_check_and_the_log_were(ran):
    wb = load_workbook(ran["b"])
    assert record.SHEET in wb.sheetnames and not {"Check", "Log"} & set(wb.sheetnames)
    ws = wb[record.SHEET]
    assert ws.sheet_properties.tabColor.rgb[-6:] == house.STONE
    assert ws["B1"].value == "Record" and ws["B1"].fill.fgColor.rgb[-6:] == house.SLATE
    assert ws["B1"].border.bottom.color.rgb[-6:] == house.STONE and ws["B1"].border.bottom.style == "thick"
    assert wb[record.LOG].sheet_state == "hidden"
    assert ws.sheet_view.showGridLines is False


def test_the_six_sections_sit_in_three_pairs_that_start_and_end_on_the_same_row(ran):
    ws = load_workbook(ran["b"])[record.SHEET]
    got = _sections(ws)
    assert set(got) == set(record.SECTIONS)
    for left, right in record.PAIRS:
        assert got[left]["col"] == record.L_LABEL and got[right]["col"] == record.R_LABEL
        assert got[left]["band"] == got[right]["band"], (left, right)
        assert got[left]["last"] == got[right]["last"], (left, right)
        for kind in (left, right):
            band, head = got[kind]["band"], got[kind]["band"] + 1
            c = got[kind]["col"]
            assert ws.cell(row=head, column=c).fill.fgColor.rgb[-6:] == house.CANVAS
            assert ws.cell(row=head, column=c).value == record.SECTIONS[kind][1][0]
            assert ws.row_dimensions[band].height == 24
    # the pairs follow one another down the tab, in the spec's order
    tops = [got[left]["band"] for left, _ in record.PAIRS]
    assert tops == sorted(tops)
    # paired sections share their label column's width
    dims = ws.column_dimensions
    assert dims[get_column_letter(record.L_LABEL)].width == dims[get_column_letter(record.R_LABEL)].width


def test_every_line_check_carried_is_on_record_in_its_section(ran):
    """What the Run worked out for Check (book._record_rows) is on the tab, each line in the section it was sorted
    into, none dropped and none added."""
    want = [(kind, label, v) for kind, rows in ran["rows"].items() if kind != record.SETTINGS
            for label, v in rows]
    shown = [(kind, label, v) for label, v, _, kind, _ in tabs.record_rows(ran["b"]) if kind != record.SETTINGS]
    assert [x for x in want if x not in shown] == []
    assert [x for x in shown if x not in want] == []
    by = {label: kind for kind, label, _ in want}
    for label, kind in (("Extract", record.THIS), ("Band edges used: FICO", record.THIS),
                        ("Tie-out checks", record.ADDS), ("Pockets tested", record.ADDS),
                        ("Pocket budget", record.ADDS), ("Tests", record.TESTS), ("p-value", record.TESTS),
                        ("Decides each pocket", record.TESTS), ("Left out of Charge-offs", record.LEFT)):
        assert by.get(label) == kind, label


def test_every_run_shows_every_log_entry_newest_first(ran):
    wb = load_workbook(ran["b"])
    entries = record.entries(wb)
    assert len(entries) == 2 and entries[0][0] >= entries[1][0]
    assert tabs.runs(wb) == [line for _, lines in entries for line in lines]
    assert tabs.runs(wb)[0] == ran["ran"].lines[0]


def test_settings_has_each_setting_the_run_asked_now_and_at_the_last_run(ran, tmp_path):
    got = tabs.settings(recalc(ran["b"], tmp_path / "calc"))
    asked = [s.question for s in control.load_settings() if not s.in_launcher and control.asked(s, {"run_kind": "bleed"})]
    assert sorted(got) == sorted(asked)
    assert [k for k, (_, _, flag) in got.items() if flag != 0] == []
    # a change on Control shows at once beside what the last Run used, shaded: a Changes-now line and a
    # Needs-a-Run one alike
    wb = load_workbook(ran["b"])
    ws = wb[control.SHEET]
    ws.cell(row=control.row_of(ws, "worse_at"), column=control.CHOOSE_COL).value = "2 times"
    ws.cell(row=control.row_of(ws, "min_loans"), column=control.CHOOSE_COL).value = "300 loans"
    changed = tmp_path / "changed.xlsx"
    wb.save(changed)
    now = tabs.settings(recalc(changed, tmp_path / "calc2"))
    moved = sorted(k for k, (a, b, flag) in now.items() if flag == 1)
    q = {s.key: s.question for s in control.load_settings()}
    assert moved == sorted([q["worse_at"], q["min_loans"]])
    assert now[q["worse_at"]][0] == "worse at 2.00x" and now[q["min_loans"]][0] == "300 loans"
    assert now[q["worse_at"]][1] == got[q["worse_at"]][1]              # the last Run's words stay
    # the shading is a rule on the Settings pair, on the flag
    rec = load_workbook(changed)[record.SHEET]
    fills = [r.dxf.fill.fgColor.rgb[-6:] for cf in rec.conditional_formatting for r in cf.rules
             if f"${get_column_letter(record.FLAG)}" in r.formula[0]]
    assert fills == [house.ALERT_FG]


def test_a_refused_run_shows_on_every_run_at_once_and_leaves_the_rest_of_record(ran, tmp_path):
    import shutil
    b = tmp_path / ran["b"].name
    shutil.copy(ran["b"], b)
    shutil.copy(ran["x"], tmp_path / ran["x"].name)
    before = [x for x in tabs.record_rows(b)]
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = None
    wb.save(b)
    out = book.run(b)
    assert not out.ok
    lines = tabs.runs(b)
    assert lines[0].startswith("Couldn't run") and any("Columns!C3" in x for x in lines[:4])
    assert len(record.entries(load_workbook(b))) == 3
    assert tabs.record_rows(b) == before
    got = _sections(load_workbook(b)[record.SHEET])
    assert got[record.LEFT]["last"] == got[record.RUNS]["last"]


def test_an_older_workbooks_log_becomes_every_run_and_its_check_goes(ran, tmp_path):
    """A workbook written before Record has a Check and a Log tab. The Log's entries carry on under Every Run and
    count for the holdout; Check is written afresh as Record."""
    import shutil
    b = tmp_path / ran["b"].name
    shutil.copy(ran["b"], b)
    shutil.copy(ran["x"], tmp_path / ran["x"].name)
    wb = load_workbook(b)
    del wb[record.SHEET]
    wb[record.LOG].title = "Log"
    wb["Log"].sheet_state = "visible"
    wb.create_sheet("Check")["B4"] = "Extract"
    wb.save(b)
    old = record.entries(load_workbook(b))
    assert len(old) == 2
    assert book.run(b).ok
    wb = load_workbook(b)
    assert not {"Check", "Log"} & set(wb.sheetnames)
    assert record.entries(wb)[1:] == old
    assert tabs.runs(wb)[-len(old[-1][1]):] == old[-1][1]


def test_record_rows_are_one_line_high_unless_their_words_need_more(ran):
    ws = load_workbook(ran["b"])[record.SHEET]
    heights = {}
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(row=r, column=record.KEY).value or "").startswith("row|"):
            heights[r] = ws.row_dimensions[r].height
    short = [r for r in heights if all(len(str(ws.cell(row=r, column=c).value or "")) < 30
                                       for c in (record.L_LABEL, record.L_TEXT, record.R_LABEL, record.R_A, record.R_B))]
    assert short and [r for r in short if heights[r] != record.ROW_H] == []
    assert [r for r, h in heights.items() if h is None or h < record.ROW_H] == []


def test_a_long_line_goes_on_in_the_rows_under_it_and_reads_back_whole():
    """A sentence too long for one row is wrapped at its spaces into rows one line high; joined with one
    space, the rows give the words back exactly, double spaces and all, over thousands of made-up lines."""
    import random
    rng = random.Random(7)
    bad = []
    for _ in range(3000):
        s = "".join(rng.choice("ab .;:,  ") for _ in range(rng.randint(0, 400)))
        for w in (5, 30, 104):
            rows = record._chunks(s, w)
            if " ".join(rows) != s:
                bad.append((s, w))
            if any(len(x) > w and " " in x[1:w] for x in rows):       # longer only where no space let it be cut
                bad.append(("too long", s, w))
    assert bad == []


def test_record_keeps_a_long_value_on_rows_one_line_high(ran):
    ws = load_workbook(ran["b"])[record.SHEET]
    rows = tabs.record_rows(ran["b"])
    tests = next(v for label, v, *_ in rows if label == "Tests")
    assert len(tests) > 2 * record.CHARS["RAB"]                    # long enough to need several rows
    kept = [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=record.JOIN_R).value == record.SAME_LINE]
    # each piece fits its row; a row is taller only when the section beside it needs the room
    assert kept and [r for r in kept if len(ws.cell(row=r, column=record.R_A).value) > record.CHARS["RAB"]] == []
    beside = lambda r: max(record._lines(ws.cell(row=r, column=c).value, record.CHARS[k])       # noqa: E731
                           for c, k in ((record.L_LABEL, record.L_LABEL), (record.L_TEXT, record.L_TEXT),
                                        (record.R_LABEL, record.R_LABEL)))
    assert [r for r in kept if ws.row_dimensions[r].height != (record.ROW_H if beside(r) == 1 else
                                                               14 * beside(r) + 4)] == []
    assert tests == next(v for kind, label, v, _, _ in record.read(ws) if label == "Tests")


def test_no_analyst_facing_word_on_record_names_a_tab_that_is_gone(ran, tmp_path):
    ws = recalc(ran["b"], tmp_path / "calc")[record.SHEET]
    said = [str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    bad = [s for s in said if "Check tab" in s or "Log tab" in s or "Confirmatory test tab" in s or "on Check" in s]
    assert bad == []
    assert [s for s in said if s.startswith(("#N/A", "#VALUE", "#NAME", "#REF", "#DIV")) or "Err:" in s] == []
