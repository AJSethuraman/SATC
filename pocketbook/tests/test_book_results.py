"""The result tabs added on 25 Sep 2026 (median per pocket, the split, losses
against revenue, materiality) and the second walkthrough's defects, each held
by a test that goes red if it comes back."""

import math
import shutil
import pytest

from openpyxl import load_workbook

from recalc import recalc

from pocketbook import book, control, engine, house, live, meanings, results, synth
import tabs
from pocketbook.ingest import read_table
from test_book import _answer
from test_book_dates import _choose



def _tab(path, name):
    """A result tab as LibreOffice calculates it (tests/recalc.py): its readings, dollars and headings are
    formulas over Control's lines (OC-40). The tab as written, formulas and rules, rides along as .formulas."""
    ws = recalc(path)[name]
    ws.formulas = load_workbook(path)[name]
    return ws


def _row(ws, name):
    for r in book.table_rows(ws):
        if r[book.C_NAME - 1].value == name:
            return r[0].row
    raise KeyError(name)


def _at(path, name, col) -> str:
    """The cell on Columns for one extract column's row, as a refusal names it: "I14"."""
    return f"{book._col(col)}{_row(load_workbook(path)['Columns'], name)}"


def _set(path, name, col, value):
    wb = load_workbook(path)
    ws = wb["Columns"]
    ws.cell(row=_row(ws, name), column=col).value = value     # value=None in ws.cell() doesn't clear a cell
    wb.save(path)


def _ready(tmp_path, n=6000):
    out = book.set_up(synth.write_extract(tmp_path, n=n))
    _answer(out.book)
    return out.book


def test_median_per_pocket_is_shown_on_the_grids(tmp_path):
    """A column shown per pocket is one of the Grids tab's measures: its median in the Rate block."""
    b = _ready(tmp_path)
    _set(b, "REV_DEBT", book.C_SHOW, "median")
    assert book.run(b).ok
    assert "Median REV_DEBT per pocket" in tabs.options(load_workbook(b), results.GRIDS, "Measure")
    ws = tabs.calculated(tabs.choose(b, tmp_path / "median.xlsx", results.GRIDS, measure="Median REV_DEBT per pocket"),
                         results.GRIDS)
    rate = tabs.block(ws, "Rate · Median REV_DEBT per pocket")
    assert sum(1 for v in rate.values() if isinstance(v, (int, float)) and v > 1000) > 10
    assert all(v is None for v in tabs.block(ws, "vs the book").values())       # a median isn't compared


def test_median_of_a_category_is_refused_by_cell(tmp_path):
    b = _ready(tmp_path, n=2000)
    _set(b, "CHANNEL", book.C_SHOW, "average")
    ran = book.run(b)
    assert not ran.ok and any(f"Columns!{_at(b, 'CHANNEL', book.C_SHOW)}" in x for x in ran.lines)


def test_split_by_a_number_finds_the_planted_revolving_debt_effect(tmp_path):
    b = _ready(tmp_path, n=8000)
    _choose(b, split="REV_DEBT")
    raw, problems, _ = book.read_book(b)
    assert not problems and raw["split"] == {"field": "REV_DEBT", "how": "own_median"}
    assert "REV_DEBT" not in {x["field"] for x in raw["bands"]}      # it splits; it isn't also cut
    assert book.run(b).ok
    ws = tabs.calculated(b, results.SPLIT)
    # the method is said once, at the top (asked for on 25 Sep 2026), not under every grid
    labels = [ws.cell(row=r, column=2).value for r in range(3, 10)]
    assert labels == ["How this tab works", "What it does", "High vs low", "High vs low, all", "p-value",
                      "What it holds", "As of"]
    grids = tabs.options(load_workbook(b), results.SPLIT, "Grid")
    assert grids[0].startswith("FICO x")                     # grids that hold the score fixed come first
    assert tabs.dropdown(ws, "Grid").value == grids[0]
    assert ws.cell(row=tabs.dropdown(ws, "Grid").row, column=results.SPLIT_CHIP).value == "Holds FICO fixed"
    head = tabs.header_row(ws, 2, "Measure")
    assert ws.cell(row=head + 1, column=2).value == "Bad loans"
    ratio = ws.cell(row=head + 1, column=5).value
    assert 1.3 < ratio < 2.6                                   # planted: 1.8x the bad rate
    chips = [r[1] for r in load_workbook(b)[results.VIEWS].iter_rows(values_only=True) if str(r[0]).endswith("|chip")]
    assert "Doesn't hold FICO fixed" in chips and "Holds FICO fixed" in chips


def test_split_by_a_category_repeats_the_grid_once_per_value(tmp_path):
    """Each grid split by every value is a grid of its own on Grids, its segments each value of each segment."""
    b = _ready(tmp_path)
    _choose(b, drop=("ASSET_CLASS",), split="ASSET_CLASS")
    assert book.read_book(b)[0]["split"]["how"] == "each_value"
    assert book.run(b).ok
    grids = tabs.options(load_workbook(b), results.GRIDS, "Grid")
    split = [g for g in grids if g.endswith(" / ASSET_CLASS")]
    assert split and len(split) * 2 == len(grids)
    ws = tabs.calculated(tabs.choose(b, tmp_path / "g.xlsx", results.GRIDS, grid=split[0]), results.GRIDS)
    heads = set(c for _, c in tabs.block(ws, "vs the book"))
    assert {"Broker · 1", "Broker · 4"} <= heads


def test_paid_cost_kept_follows_the_lines_on_control(tmp_path):
    """Ruling OC-26; the third walk, defects 1 and 5. Each side reads by the Control lines, against the same
    comparison as its flag: charge-offs by their multiple, what they paid and what we kept by their gap in
    points (NEXT-GOAL 3.2 to 3.4). The shading shows the flag; a side that isn't significant is left plain."""
    untested = (engine.THIN, engine.FEW)
    b = _ready(tmp_path)
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=4):
        if r[6].value == "revenue_line":
            r[2].value = "0.5 points either way"
    wb.save(b)
    assert book.run(b).ok
    ws = tabs.calculated(b, results.PCK)
    rows = [x for x in tabs.pck(ws) if not set(x["flags"].values()) & set(untested)]
    assert rows
    for x in rows:
        g, r, c = x["flags"]["g"], x["flags"]["r"], x["flags"]["c"]
        # the lines on Control decide each side (the firm, 25 Sep 2026)
        want_g = ((engine.WORSE, engine.UNSURE_WORSE) if x["cost"] >= 1.25 else (engine.BETTER, engine.UNSURE_BETTER)
                  if x["cost"] <= 0.8 else (engine.IN_LINE,))
        assert g in want_g, x
        for gap, flag in ((x["kept"], r), (x["paid"], c)):
            assert flag in ((engine.WORSE, engine.UNSURE_WORSE) if gap <= -0.5 else (engine.BETTER, engine.UNSURE_BETTER)
                            if gap >= 0.5 else (engine.IN_LINE,)), x
        # shaded by the flag: pink worse and real, green better and real, plain otherwise
        for fill, flag in ((x["g_fill"], g), (x["r_fill"], r), (x["c_fill"], c)):
            assert fill == {engine.WORSE: house.ALERT_FG, engine.BETTER: house.POSITIVE_BG}.get(flag), x
    # what they paid us, less what they cost us, is what we kept (contribution = RANR + GCO, OC-35), pocket by
    # pocket on _pockets. Within a hair: the one loan with a GCO of "#N/A" is out of GCO and contribution only
    v = recalc(b, tmp_path / "all")
    rate = {}
    for r_ in v[live.POCKETS].iter_rows(min_row=live.P_FIRST, values_only=True):
        if r_[live.P_KIND - 1] == "grids":
            rate[(r_[live.P_GRID - 1], r_[live.P_BAND - 1], r_[live.P_SEG - 1], r_[live.P_MEASURE - 1])] = \
                r_[live.P_RATE - 1]
    pockets = {k[:3] for k in rate if None not in (rate[(*k[:3], m)] for m in ("contribution_rate", "gco_rate",
                                                                             "ranr_rate"))}
    assert pockets and all(rate[(*k, "contribution_rate")] - rate[(*k, "gco_rate")] == pytest.approx(
        rate[(*k, "ranr_rate")], abs=5e-4) for k in pockets)
    # the grid picked is FICO x CHANNEL; its first row is the planted pocket, priced like its band, losing far
    # more and keeping less
    assert tabs.dropdown(ws, "Grid").value == "FICO x CHANNEL"
    head = tabs.header_row(ws, results.C_TOG, "Together")
    assert tabs.heads(ws, head, results.C_BAND, results.C_TOG) == [
        "Band", "Segment", "Loans", "Booked", "GCOs", "RANR", "RANR ÷ Booked", "Avg line", "Line × book", "Gap pts",
        "Dollars", "Rest", "× rest of band", "Dollars", "Rest", "Gap pts", "Dollars", "Rest", "Together"]
    assert [ws.cell(row=head - 1, column=c).value for c in (results.C_PAID, results.C_COST, results.C_KEPT)] == [
        "RANR + GCOs · gap vs rest of band", "GCOs · × rest of band", "RANR · gap vs rest of band"]
    firsts = [int(str(x["band"]).split(" - ")[0]) for x in tabs.pck(ws) if str(x["band"])[0].isdigit()]
    assert int(str(tabs.pck(ws)[0]["band"]).split(" - ")[0]) == min(firsts)     # band order, the lowest first
    first = next(x for x in tabs.pck(ws) if int(str(x["band"]).split(" - ")[0]) == min(firsts)
                 and x["seg"] == "Broker")
    assert first["flags"]["g"] == engine.WORSE and first["flags"]["r"] == engine.WORSE
    assert first["together"] == "Net drain"
    assert len(tabs.options(load_workbook(b), results.PCK, "Grid")) == 6
    assert len(load_workbook(b)[results.PCK]._charts) == 1                 # one chart: the grid picked


def test_the_suggested_profit_line_is_each_pockets_own_test(tmp_path):
    b = _ready(tmp_path)
    assert book.run(b).ok
    v = recalc(b, tmp_path / "rc")
    check = tabs.record(v)
    assert check["Profit counts as more or less"] == "each pocket's own test: only a gap that is significant at 95%"
    reads = [r[live.P_SAID - 1] for r in v[live.POCKETS].iter_rows(min_row=live.P_FIRST, values_only=True)
             if r[live.P_MEASURE - 1] in ("ranr_rate", "contribution_rate") and r[live.P_SAID - 1]]
    # nothing past a line is left to mark: a gap its own test doesn't call significant says so in words
    assert reads and not any(x.endswith("(not significant)") for x in reads)
    assert all(x.endswith(", not significant") for x in reads if not x.startswith(("short of", "ahead of", "too few")))
    assert any(x.endswith(", not significant") for x in reads)
    assert check["How profit reads"].startswith("As the gap: how many points of booked dollars a pocket is short of "
                                                "or ahead of its band")


def test_a_copied_workbook_runs_the_extract_picked_not_the_old_path(tmp_path):
    """Second walk, defect 1."""
    b = _ready(tmp_path / "a", n=3000)
    other = tmp_path / "b"
    other.mkdir()
    copy = shutil.copy(b, other / b.name)
    x2 = synth.write_extract(other, n=2500, seed=11)
    ran = book.run(copy, x2)
    assert ran.ok and any("Ran on 2,500 loans from loans.csv" in x for x in ran.lines)


def test_a_moved_pair_finds_the_extract_beside_the_workbook(tmp_path):
    b = _ready(tmp_path / "a", n=3000)
    moved = tmp_path / "b"
    shutil.copytree(tmp_path / "a", moved)
    shutil.rmtree(tmp_path / "a")
    ran = book.run(moved / b.name)
    assert ran.ok and any("Ran on 3,000 loans" in x for x in ran.lines)


def test_edges_excel_read_as_one_number_are_refused(tmp_path):
    """Second walk, defect 2: 620,680,740 became 620680740."""
    b = _ready(tmp_path, n=2000)
    _set(b, "FICO", book.C_EDGES, 620680740)
    ran = book.run(b)
    assert not ran.ok and any("Excel dropped the commas" in x and f"Columns!{_at(b, 'FICO', book.C_EDGES)}" in x
                              for x in ran.lines)


def test_edges_with_semicolons_are_read(tmp_path):
    b = _ready(tmp_path, n=2000)
    _set(b, "FICO", book.C_EDGES, "620; 680; 740")
    assert book.read_book(b)[0]["bands"][0]["edges"] == [620, 680, 740]


def test_a_new_column_takes_the_yes_back(tmp_path):
    """Second walk, defect 4."""
    x = synth.write_extract(tmp_path, n=2000)
    out = book.set_up(x)
    _answer(out.book)
    text = x.read_text().splitlines()
    x.write_text("\n".join([text[0] + ",NEW_LINE"] + [t + ",1" for t in text[1:]]) + "\n")
    again = book.set_up(x)
    assert any("NEW_LINE" in line for line in again.lines)
    ws = load_workbook(out.book)["Columns"]
    assert ws[book.CONFIRM_CELL].value is None and "NEW_LINE" in ws["D3"].value


def test_set_up_again_keeps_the_last_results(tmp_path):
    """Second walk, defect 6."""
    x = synth.write_extract(tmp_path, n=3000)
    out = book.set_up(x)
    _answer(out.book)
    assert book.run(out.book).ok
    book.set_up(x)
    names = load_workbook(out.book).sheetnames
    assert {results.POCKETS, results.GRIDS, "Record", results.PCK} <= set(names)


def test_set_up_carries_the_answers_over_from_a_workbook_under_the_old_name(tmp_path):
    """The rename (26 Sep 2026, swept 27 Sep): a workbook made as "<extract> - Origination Cube.xlsx" is where
    Set up takes the answers from when the PocketBook one is not there yet. The old file is left as it was."""
    from test_book import PICK
    x = synth.write_extract(tmp_path, n=2000)
    first = book.set_up(x).book
    _answer(first)
    old = first.with_name(first.name.replace(" - PocketBook.xlsx", " - Origination Cube.xlsx"))
    first.rename(old)
    before = old.read_bytes()
    again = book.set_up(x)
    assert again.ok and again.book == first and first.exists()
    carried = book._answers(first)
    assert {k: carried["control"][k][0] for k in PICK} == dict(PICK)
    assert carried["confirmed"] == book._answers(old)["confirmed"] and carried["odd"] == book._answers(old)["odd"]
    assert old.read_bytes() == before


def test_a_workbook_open_in_excel_is_refused_before_anything_changes(tmp_path, monkeypatch):
    """Second walk, defect 8."""
    b = _ready(tmp_path, n=2000)
    before = b.read_bytes()
    monkeypatch.setattr(book, "_writable", lambda p: False)
    ran = book.run(b)
    assert not ran.ok and "open in Excel" in ran.lines[0]
    assert b.read_bytes() == before
    assert not b.with_name(f"{b.stem} - what ran.yaml").exists()


def test_start_here_says_when_it_last_ran(tmp_path):
    """Second walk, defect 9."""
    b = _ready(tmp_path, n=3000)
    assert book.run(b).ok
    ws = load_workbook(b)["Start here"]
    assert "last Run 20" in ws["C1"].value                     # the subtitle: when it last ran
    said = {c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)}
    assert any(s.startswith("=IFERROR(SUM(answers_needed),0)") for s in said)


def test_ranr_is_marked_more_is_better_and_its_gap_reads_or_less(tmp_path):
    """Second walk, defects 13 and 14: profit reads in points, a shortfall below zero, never a multiple
    (NEXT-GOAL 3.2), on Pockets and on the Grids' heat scale."""
    b = _ready(tmp_path)
    assert book.run(b).ok
    ws = tabs.calculated(tabs.choose(b, tmp_path / "k.xlsx", results.POCKETS, measure="RANR"),
                         results.POCKETS)
    rows = tabs.pockets(ws)
    assert rows and all(x["caught"] < 0 for x in rows if x["caught"] is not None)
    head = tabs.header_row(ws, results.K_NUM, "#")
    assert ws.cell(row=head, column=results.K_GAP).value == "Gap in pts"
    assert ws.cell(row=head, column=results.K_EX).value == "Dollars short of band"
    fmts = [r.dxf.numFmt.formatCode for rng in ws.formulas.conditional_formatting
            if str(rng.sqref).startswith(results.col(results.K_CAUGHT)) for r in rng.rules if r.dxf.numFmt]
    assert fmts and set(fmts) == {results.PTS_FMT}
    # Grids reads it as a gap in points where more is better: its heat runs the other way
    meta = {r[0]: r[1:3] for r in load_workbook(b)[results.VIEWS].iter_rows(values_only=True)}
    assert meta["G|FICO x CHANNEL|ranr_rate|meta"][0] == "pts" and meta["G|FICO x CHANNEL|gco_rate|meta"][0] == "x"


def test_an_edge_outside_the_columns_values_is_refused_by_cell(tmp_path):
    """Second walk, defect 2: one edge above every FICO made a single band."""
    b = _ready(tmp_path, n=2000)
    _set(b, "FICO", book.C_EDGES, "620; 680; 9000")
    ran = book.run(b)
    assert not ran.ok
    said = next(x for x in ran.lines if f"Columns!{_at(b, 'FICO', book.C_EDGES)}" in x)
    assert "9,000" in said or "9000" in said
    assert "-9,999" not in said and "-9999" not in said          # the answered missing code is not the range


def test_three_way_pockets_are_tested_and_ranked(tmp_path):
    """Ruling OC-27; the third walk, defect 3: a category split drew pictures only. On Pockets, the split view."""
    b = _ready(tmp_path)
    _choose(b, split="ASSET_CLASS")
    ran = book.run(b)
    assert ran.ok and any("Split by ASSET_CLASS" in x for x in ran.lines)
    assert "Split by ASSET_CLASS" in tabs.options(load_workbook(b), results.POCKETS, "Pockets")
    ws = tabs.calculated(tabs.choose(b, tmp_path / "s.xlsx", results.POCKETS, pockets="Split by ASSET_CLASS"),
                         results.POCKETS)
    head = tabs.header_row(ws, results.K_NUM, "#")
    assert ws.cell(row=head, column=results.K_HALF).value == "ASSET_CLASS"
    rows = tabs.pockets(ws)
    assert rows and {x["half"] for x in rows} <= {"1", "2", "3", "4"} and all(x["worse"] for x in rows)
    check = tabs.record(b)
    assert "ASSET_CLASS" in check["Split"] and "isn't cut on its own" in check["Split"]


def test_a_split_on_a_column_that_cant_split_is_refused_by_cell(tmp_path):
    """The third walk, defect 7: GCO split by itself read 93.83x."""
    b = _ready(tmp_path, n=2000)
    _choose(b, split="GCO_AMT")
    ran = book.run(b)
    split = f"Control!C{control.row_of(load_workbook(b)[control.SHEET], 'launcher|split')}"
    assert not ran.ok and any(x.strip(" -").startswith(split) and f"Columns!{_at(b, 'GCO_AMT', book.C_MEANS)}" in x and "can't split the pockets"
                              in x for x in ran.lines), ran.lines


def test_a_dollar_materiality_line_is_gco_and_profit_is_held_to_it(tmp_path):
    """The third walk, defect 8: a dollar line is GCO's, and the outcome rates
    don't borrow it. A profit shortfall is dollars too: since NEXT-GOAL 3.2 it is
    material at the same line (it was a share of |total RANR|, which collapsed
    when the book's profit was near zero)."""
    from test_book import PICK
    b = _ready(tmp_path, n=3000)
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=4):
        if r[6].value == "materiality":
            r[2].value, r[3].value = None, 100000
    wb.save(b)
    assert book.run(b).ok, PICK
    check = tabs.record(_tab(b, "Record"))
    assert check["Materiality line: GCOs ($)"] == "100,000 GCO_AMT dollars"
    assert check["Materiality line: RANR"] == (
        "a shortfall of 100,000 RANR_AMT dollars: the same dollar line as GCO (Control's materiality answer)")
    assert "GCO amount" in check["Materiality line: Bad loans"]
    assert check["Smallest excess loss worth reporting"] == "$100,000 of GCO"


def test_a_forget_holds_until_a_person_confirms_again(tmp_path):
    """The third walk, defect 9: the next Run re-learned the column silently."""
    from pocketbook import memory
    b = _ready(tmp_path, n=3000)
    assert book.run(b).ok
    _set(b, "CHANNEL", book.C_FORGET, book.FORGET_YES)       # Forget? on its Columns row, filled in by the Run
    assert book.run(b).ok
    ws = load_workbook(b)["Columns"]
    assert ws[book.CONFIRM_CELL].value is None and "CHANNEL" in ws["D3"].value
    ran = book.run(b)
    assert not ran.ok and any("Columns!C3" in x for x in ran.lines)
    assert "CHANNEL" not in memory.load()["columns"]


def test_the_workbook_is_refused_as_its_own_extract(tmp_path):
    """The third walk, defect 10."""
    b = _ready(tmp_path, n=500)
    out = book.set_up(b)
    assert not out.ok and "is the workbook, not the loan file" in out.lines[0]
    assert not b.with_name(f"{b.stem}{book.SUFFIX}").exists()


def test_a_renamed_column_says_to_press_set_up(tmp_path):
    b = _ready(tmp_path, n=1000)
    x = tmp_path / "loans.csv"
    text = x.read_text().splitlines()
    x.write_text("\n".join([text[0].replace("CHANNEL", "CHNL")] + text[1:]) + "\n")
    ran = book.run(b)
    said = " ".join(ran.lines)
    assert not ran.ok and "a segment" in said and "press Set up again" in said and "dimension" not in said


def test_the_heat_maps_leave_out_pockets_under_the_minimum(tmp_path):
    """The third walk, defect 4: a 1-loan row was among the strongest colours."""
    b = _ready(tmp_path, n=4000)
    assert book.run(b).ok
    ws = tabs.calculated(b, results.GRIDS)
    vs = tabs.block(ws, "vs the book")
    blank = [v for (bl, _), v in vs.items() if bl == "(blank)"]
    assert blank and all(v is None for v in blank)
    assert any(isinstance(v, float) for v in vs.values())


def test_band_width_every_20_cuts_and_is_remembered(tmp_path):
    """The firm, 25 Sep 2026: "20 point bands look very different", and yes to
    remembering a column's edges."""
    from pocketbook import memory
    b = _ready(tmp_path, n=6000)
    _set(b, "FICO", book.C_EDGES, "every 20")
    ran = book.run(b)
    assert ran.ok, ran.lines
    check = tabs.record(b)
    edges = check["Band edges used: FICO"]
    pts = [float(x) for x in edges.split("(")[0].replace(",", "").split(";")]
    assert all(p % 20 == 0 for p in pts) and len(pts) > 10
    assert memory.load()["columns"]["FICO"]["edges"] == "every 20"
    other = tmp_path / "next"
    x2 = synth.write_extract(other, n=2000, seed=3)
    again = book.set_up(x2)
    assert load_workbook(again.book)["Columns"][_at(again.book, "FICO", book.C_EDGES)].value == "every 20"


def test_a_band_width_too_narrow_is_refused(tmp_path):
    b = _ready(tmp_path, n=2000)
    _set(b, "FICO", book.C_EDGES, "every 1")
    ran = book.run(b)
    assert not ran.ok and any(f"Columns!{_at(b, 'FICO', book.C_EDGES)}" in x and "50 bands or fewer" in x
                              for x in ran.lines)


def test_suggested_answers_are_worked_out_from_the_book(tmp_path):
    """The firm, 25 Sep 2026: suggestions "where there's a calculation", never pre-chosen."""
    b = _ready(tmp_path, n=6000)
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=4):
        if r[6].value == "min_loans":
            r[2].value = "Enough for 5 expected losses (suggested)"
        if r[6].value in ("worse_at", "better_at"):
            r[2].value = "The smallest significant gap in a typical pocket (suggested)"
    wb.save(b)
    ran = book.run(b)
    assert ran.ok, ran.lines
    check = tabs.record(b)
    said = check["Worked out from this book"]
    loans = int(said.split("fewest loans ")[1].split(" ")[0].replace(",", ""))
    assert 60 < loans < 90                                     # 5 / 7.1% bad = about 71 (the firm, 25 Sep 2026)
    assert "worse at" in said and "better at" in said
    import yaml
    ran_with = yaml.safe_load(b.with_name(f"{b.stem} - what ran.yaml").read_text())["benchmark"]
    assert ran_with["min_units"] == loans and ran_with["worse_at"] > 1


def test_check_says_what_the_allowance_covers(tmp_path):
    b = _ready(tmp_path, n=2000)
    assert book.run(b).ok
    check = tabs.record(b)
    assert "each grid and measure on its own" in check["The allowance for many tests covers"]


def test_paid_cost_kept_dollars_agree_with_the_flag_and_untested_pockets_get_none(tmp_path):
    """The fourth walk, defects 2 and 3."""
    b = _ready(tmp_path)
    assert book.run(b).ok
    ws = tabs.calculated(b, results.PCK)
    rows = tabs.pck(ws)
    few = {engine.THIN, engine.FEW}
    untested = [x for x in rows if set(x["flags"].values()) & few]
    for x in rows:
        g, r, c = x["flags"]["g"], x["flags"]["r"], x["flags"]["c"]
        if g in (engine.WORSE, engine.UNSURE_WORSE):
            assert x["cost_d"] > 0, x
        if g in (engine.BETTER, engine.UNSURE_BETTER):
            assert x["cost_d"] < 0, x
        if r in (engine.BETTER, engine.UNSURE_BETTER):
            assert x["kept_d"] > 0, x
        if r in (engine.WORSE, engine.UNSURE_WORSE):
            assert x["kept_d"] < 0, x
        if c in (engine.BETTER, engine.UNSURE_BETTER):
            assert x["paid_d"] > 0, x
    assert any(x["flags"]["g"] == engine.WORSE for x in rows)
    # untested on any side: no colour on any side, and nothing read together
    assert untested and all(x["c_fill"] is None and x["g_fill"] is None and x["r_fill"] is None
                            and not x["together"] for x in untested)
    # and they sit where their band puts them, as every row does (the firm, 30 Sep 2026: band order)
    lows = [float(str(x["band"]).split(" - ")[0]) if str(x["band"])[0].isdigit() else math.inf for x in rows]
    assert lows == sorted(lows)
    note = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(3, 12)}
    assert "RANR already has the GCOs taken out, so this adds them back" in note["RANR + GCOs"]


def test_split_rows_say_what_their_grid_holds_fixed(tmp_path):
    """The fourth walk, defect 1: the Three-way tab carried no caveat at all. On Pockets' split view, each row
    says whether its grid holds the split's partner fixed; those that don't come last, grey, with no verdict
    colour (the firm, 25 Sep 2026, after the seventh walk)."""
    b = _ready(tmp_path)
    _choose(b, split="REV_DEBT")
    assert book.run(b).ok
    ws = tabs.calculated(tabs.choose(b, tmp_path / "s.xlsx", results.POCKETS, pockets="Split by REV_DEBT"),
                         results.POCKETS)
    head = tabs.header_row(ws, results.K_NUM, "#")
    assert ws.cell(row=head, column=results.K_HOLDS).value == "Holds FICO fixed?"
    rows = tabs.pockets(ws)
    assert rows[0]["band"].startswith("FICO") and rows[0]["holds"] == "Yes"
    assert any(x["band"].startswith("ORIG_BAL") and x["holds"] == "No: may be mostly FICO" for x in rows)
    held = [x["holds"] == "Yes" for x in rows]
    assert held == sorted(held, reverse=True)                              # held-fixed grids first
    # no verdict colour on those rows: their first rule greys them and the verdict rules come after it
    rules = [r.formula[0] for rng in ws.formulas.conditional_formatting
             if str(rng.sqref).startswith(results.col(results.K_WORSE)) for r in sorted(rng.rules,
                                                                                    key=lambda x: x.priority)]
    assert 'LEFT($' in rules[0] and '="No:"' in rules[0] and 'Yes"' in rules[1]
    split = tabs.calculated(b, results.SPLIT)
    chips = [r[1] for r in load_workbook(b)[results.VIEWS].iter_rows(values_only=True) if str(r[0]).endswith("|chip")]
    assert chips.count("Holds FICO fixed") == 2                              # once per FICO grid, not per measure
    assert tabs.dropdown(split, "Grid").value.startswith("FICO")


def test_the_luck_line_is_luck_alone_not_the_catch_rate(tmp_path):
    """The fourth walk, defect 4: at 80% caught it was the gap a pocket can find."""
    b = _ready(tmp_path)
    assert book.run(b).ok
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=4):
        if r[6].value == "worse_at":
            r[2].value = "The smallest significant gap in a typical pocket (suggested)"
    wb.save(b)
    assert book.run(b).ok
    check = tabs.record(b)
    worse = float(check["Worked out from this book"].split("worse at ")[1].split("x")[0])
    assert 1.1 < worse < 1.45                                  # luck alone; the catch-rate gap is bigger


def test_a_profit_line_of_95_points_is_refused_with_its_range(tmp_path):
    """The fourth walk, defect 5: the advice given must be allowed. The profit
    line is points now (NEXT-GOAL 3.2), so 95 is out of range, and the range
    is said."""
    from pocketbook import control
    b = _ready(tmp_path, n=1000)
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=4):
        if r[6].value == "revenue_line":
            r[2].value, r[3].value = None, 95
    wb.save(b)
    ran = book.run(b)
    said = next(x for x in ran.lines if "How far profit" in x)
    assert "a number of points, such as 0.3, from 0.01 to 20; got 95" in said and "type 0.95" not in said


def test_the_planted_pocket_is_a_net_drain(tmp_path):
    """The fourth walk's render had under 620 / Broker at RANR 1.17x its band on
    176 loans, which its own test called luck: the synthetic RANR was a flat draw
    that ignored losses. RANR is profit after losses now (NEXT-GOAL 3.5): the
    pocket is priced like its band, so what it pays is not significantly
    different, it loses far more, and it falls short of its band on profit.
    Together: a net drain."""
    b = _ready(tmp_path, n=8000)
    _set(b, "FICO", book.C_EDGES, "620; 680; 740")
    assert book.run(b).ok
    ws = tabs.calculated(b, results.PCK)
    x = next(y for y in tabs.pck(ws) if str(y["band"]).endswith(" - 619") and y["seg"] == "Broker")
    assert x["flags"]["c"] in (engine.IN_LINE, engine.UNSURE_WORSE, engine.UNSURE_BETTER)     # not significant
    assert (x["flags"]["g"], x["flags"]["r"], x["together"]) == (engine.WORSE, engine.WORSE, "Net drain")
    assert x["kept"] < 0 and x["cost"] > 1.25
    # pink on the two sides that are worse and real, none on what they paid
    assert x["g_fill"] == house.ALERT_FG and x["r_fill"] == house.ALERT_FG and x["c_fill"] is None


def test_boxes_follow_the_lines_exactly():
    """A side reads more or less only past its line; between the lines it's the same. (The book-level test can
    pass by chance when no pocket sits between 1.00x and a line, which let 'boxes by 1.00x again' slip past the
    mutation check in CI on 25 Sep 2026.)"""
    from conftest import cube
    bench = cube(benchmark={"min_units": 2, "min_events": 1, "worse_at": 1.25, "better_at": 0.8,
                            "confidence": 0.95, "power": 0.8, "compare_to": "topline", "many_tests": "none",
                            "materiality": "none", "shuffles": 200}).benchmark
    read = lambda x: engine.reading_of(x, 100, bench, 0, 0.001)             # noqa: E731
    assert (read(1.10), read(0.90), read(1.25), read(0.80)) == (engine.IN_LINE, engine.IN_LINE, engine.WORSE,
                                                                   engine.BETTER)
    # a side is more or less only when its flag is worse or better and real; the same when tested and neither
    assert results.side_of(engine.WORSE, "worse") == "more" and results.side_of(engine.WORSE, "better") == "less"
    assert results.side_of(engine.UNSURE_WORSE, "worse") == "same" and results.side_of(engine.FEW, "worse") is None
    # charge-offs and what we kept read together (NEXT-GOAL 3.4; the redesign's five, section 6)
    t = results.together_of
    assert t(engine.WORSE, engine.BETTER) == "Priced for it"
    assert t(engine.WORSE, engine.WORSE) == "Net drain"
    assert t(engine.BETTER, engine.BETTER) == "Strong"
    assert t(engine.BETTER, engine.WORSE) == "Safe but idle"
    assert t(engine.IN_LINE, engine.WORSE) == t(engine.UNSURE_WORSE, engine.WORSE) == "Earns less, not from losses"
    assert t(engine.WORSE, engine.IN_LINE) == t(engine.WORSE, engine.UNSURE_BETTER) == "Losing more, profit holding"
    assert [t(g, r) for g, r in ((engine.IN_LINE, engine.BETTER),
                                 (engine.FEW, engine.WORSE), (engine.WORSE, engine.FEW),
                                 (engine.IN_LINE, engine.IN_LINE))] == [""] * 4


def test_a_gap_that_is_not_significant_is_left_plain(tmp_path):
    """The firm, 25 Sep 2026, after the fifth walk: the lines decide, and a gap that could be chance is marked
    (its flag says not significant) and left plain."""
    b = _ready(tmp_path, n=8000)
    _set(b, "FICO", book.C_EDGES, "620; 680; 740")
    assert book.run(b).ok
    rows = tabs.pck_all(b, tmp_path / "grids")                  # every grid, each picked in turn
    marked = [x for x in rows if x["flags"]["g"] in (engine.UNSURE_WORSE, engine.UNSURE_BETTER)]
    assert marked and all(x["g_fill"] is None for x in marked)
    assert any(x["g_fill"] == house.ALERT_FG for x in rows)
    # and the other way round: a gap past a line on Control that isn't coloured is marked, either side
    past = [x for x in rows if x["flags"]["g"] not in (engine.FEW, engine.THIN) and (x["cost"] >= 1.25 or x["cost"] <= 0.8)]
    plain = [x for x in past if x["g_fill"] is None]
    assert any(x["cost"] >= 1.25 for x in plain)
    assert all(x["flags"]["g"] in (engine.UNSURE_WORSE, engine.UNSURE_BETTER) for x in plain)


def test_remembered_edges_only_fill_a_column_the_workbook_has_not_seen(tmp_path):
    """The fifth walk: another copy's edges came into this workbook, and clearing didn't undo it."""
    from pocketbook import memory
    b = _ready(tmp_path, n=3000)
    _set(b, "FICO", book.C_EDGES, "every 20")
    assert book.run(b).ok
    _set(b, "FICO", book.C_EDGES, None)                       # cleared on this workbook
    assert book.run(b).ok
    assert "edges" not in memory.load()["columns"]["FICO"]     # and forgotten with it
    book.set_up(tmp_path / "loans.csv")
    assert load_workbook(b)["Columns"][_at(b, "FICO", book.C_EDGES)].value is None
    # a remembered edge fills only a workbook that hasn't seen the column, and says so
    _set(b, "FICO", book.C_EDGES, "every 25")
    assert book.run(b).ok
    other = book.set_up(synth.write_extract(tmp_path / "q4", n=1000, seed=5))
    ws = load_workbook(other.book)["Columns"]
    assert ws[_at(other.book, "FICO", book.C_EDGES)].value == "every 25"
    assert "remembered from before" in ws[_at(other.book, "FICO", book.C_LOOK)].value


def test_control_shows_what_the_last_run_used(tmp_path):
    """The fifth walk: a suggested option showed no number anywhere on Control."""
    from pocketbook import control
    b = _ready(tmp_path, n=3000)
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == "min_loans":
            r[control.CHOOSE_COL - 1].value = "Enough for 5 expected losses (suggested)"
    wb.save(b)
    ran = book.run(b)
    assert ran.ok and any(x.startswith("Worked out from this book: fewest loans") for x in ran.lines)
    ws = load_workbook(b)["Control"]
    col = control.LAST_COL
    for block in (control.BLOCK_NOW, control.BLOCK_RUN):
        assert ws.cell(row=control.row_of(ws, block) + 1, column=col).value == "Last Run used"
    used = {r[control.KEY_COL - 1].value: ws.cell(row=r[0].row, column=col).value
            for r in ws.iter_rows(min_row=control.FIRST_ROW) if r[control.KEY_COL - 1].value}
    assert "worked out from this book" in used["min_loans"] and used["revenue_line"] == "each pocket's own test"
    book.set_up(tmp_path / "loans.csv")                        # and it stays through Set up again
    ws = load_workbook(b)["Control"]
    again = {r[control.KEY_COL - 1].value: ws.cell(row=r[0].row, column=col).value
             for r in ws.iter_rows(min_row=control.FIRST_ROW) if r[control.KEY_COL - 1].value}
    assert again["min_loans"] == used["min_loans"]


def test_another_workbooks_edges_stay_out_of_this_one(tmp_path):
    """The fifth walk's scenario: edges typed on another copy must not come into a
    workbook that already has the column, not even as a note."""
    a = _ready(tmp_path / "a", n=2000)
    assert book.run(a).ok                                    # FICO confirmed here with no edges
    b = _ready(tmp_path / "b", n=2000)
    _set(b, "FICO", book.C_EDGES, "every 25")
    assert book.run(b).ok                                    # remembered from the other copy
    book.set_up(tmp_path / "a" / "loans.csv")
    ws = load_workbook(a)["Columns"]
    assert ws[_at(a, "FICO", book.C_EDGES)].value is None
    assert "remembered from before" not in str(ws[_at(a, "FICO", book.C_LOOK)].value or "")


def test_a_suggestion_with_nothing_to_work_from_says_so(tmp_path):
    """The sixth walk, defect 3: a default was reported as "worked out from this book",
    and a book where nothing could be tested read "Nothing is worse"."""
    from pocketbook import control
    b = _ready(tmp_path, n=2000)
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        key = r[control.KEY_COL - 1].value
        if key == "min_loans":
            r[control.CHOOSE_COL - 1].value, r[control.OWN_COL - 1].value = None, 100000
        # since 25 Sep 2026 fewest loans only picks the test (exact below it), so nothing is untested
        # unless fewest losses says so
        if key == "min_events":
            r[control.CHOOSE_COL - 1].value, r[control.OWN_COL - 1].value = None, 100000
        if key in ("worse_at", "better_at"):
            r[control.CHOOSE_COL - 1].value = "The smallest significant gap in a typical pocket (suggested)"
    wb.save(b)
    ran = book.run(b)
    assert ran.ok
    said = " ".join(ran.lines)
    assert "Nothing in this book to work these out from" in said and "Worked out from this book" not in said
    for loss in ("Outcome, share of loans", "Outcome, share of booked dollars", "GCOs per booked dollar"):
        assert f"No pocket had enough losses to test {loss}" in said and f"Nothing is worse for {loss}" not in said
    ws = load_workbook(b)["Control"]
    used = [ws.cell(row=r, column=control.LAST_COL).value for r in range(control.FIRST_ROW, ws.max_row + 1)]
    assert any(isinstance(x, str) and "the usual value" in x for x in used)
    # Check names the fallback as Control does and never calls it worked out (the seventh walk, defect 6)
    check = tabs.record(b)
    fell = check["Suggested values"]
    # the worked-out wording never sits beside a fallback ("luck alone can make" until NEXT-GOAL 3.1)
    assert "can call significant" not in fell and "how much better (0.80x) and how much worse (1.25x)" in fell
    assert "_at" not in fell and check["Pockets tested"].startswith("none of")


def test_split_gaps_that_are_not_significant_are_bracketed(tmp_path):
    """The sixth walk, defect 9: a 2.33x at 61% was deep red."""
    b = _ready(tmp_path)
    _choose(b, split="REV_DEBT")
    assert book.run(b).ok
    ws = tabs.calculated(b, results.SPLIT)
    vals = list(tabs.block(ws, "Bad loans, high vs low").values())
    assert any(isinstance(v, str) and v.startswith("(") and v.endswith("×)") for v in vals)
    assert any(isinstance(v, float) for v in vals)


def test_material_pockets_too_small_to_test_are_pointed_out(tmp_path):
    """The firm, 25 Sep 2026: whether a small pocket matters is a materiality thing. On Pockets: Worse? reads Too
    few losses and Material? Yes, the two separate columns."""
    from pocketbook import control
    b = _ready(tmp_path, n=4000)
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        # too few losses is what leaves a pocket untested since 25 Sep 2026: below fewest loans the
        # share of loans gets the exact test and a dollar rate is shuffled, so 400 loans no longer did it
        if r[control.KEY_COL - 1].value == "min_events":
            r[control.CHOOSE_COL - 1].value, r[control.OWN_COL - 1].value = None, 60
    wb.save(b)
    ran = book.run(b)
    assert ran.ok and any("material but too small to test" in x for x in ran.lines)
    said = next(x for x in ran.lines if "material but too small to test" in x)
    assert "Worse? reads Too few losses and Material? Yes on Pockets" in said
    # the window counts pockets, once each, and the rows (the seventh walk, defect 7)
    v = recalc(b, tmp_path / "rc")
    key = {r[0].row: tuple(c.value for c in (r[live.P_GRID - 1], r[live.P_BAND - 1], r[live.P_SEG - 1]))
           for r in v[live.POCKETS].iter_rows(min_row=live.P_FIRST)}
    small = [r for r in v[results.LIST].iter_rows(min_row=2, values_only=True)
             if r[results.L_KKEY - 1] == "grids" and r[results.L_WORSE - 1] == live.TOO_FEW
             and r[results.L_MAT - 1] == live.YES]
    pockets = {key[r[results.L_ROW - 1]] for r in small}
    assert pockets and said.startswith(f"{len(pockets)} pocket")
    if len(small) != len(pockets):
        assert f"({len(small)} rows" in said
    # and a measure's view shows them
    ws = v[results.POCKETS]
    assert any(x["worse"] == live.TOO_FEW and x["material"] == live.YES for x in tabs.pockets(ws)) or any(
        r[results.L_MEAS - 1] != "Bad loans" for r in small)


def test_a_pocket_under_fewest_loans_is_tested_and_says_how(tmp_path):
    """Walk 6, defect 8, on the workbook: under fewest loans the share of loans gets the exact test (statistics.md
    B1) and shows its p-value, never read as too few to test. Which test ran is said once, in Pockets' method
    note (tenet T1), and kept pocket by pocket on _pockets."""
    from pocketbook import control
    b = _ready(tmp_path, n=4000)
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == "min_loans":
            r[control.CHOOSE_COL - 1].value, r[control.OWN_COL - 1].value = None, 230
    wb.save(b)
    assert book.run(b).ok
    v = recalc(b, tmp_path / "rc")
    ws = v[results.POCKETS]
    note = " ".join(str(ws.cell(row=r, column=results.K_BAND + 1).value) for r in range(3, 14))
    assert "the exact test for a pocket under 230 loans" in note
    assert "Test" not in tabs.heads(ws, tabs.header_row(ws, results.K_NUM, "#"))
    rows = [x for x in tabs.pockets(ws) if x["worse"] != live.TOO_FEW]
    small, big = [x for x in rows if x["loans"] < 230], [x for x in rows if x["loans"] >= 230]
    assert small and all(x["p"] is not None for x in small) and big
    tests = {(r[live.P_LOANS - 1] < 230): r[live.P_TEST_BOOK - 1]
             for r in v[live.POCKETS].iter_rows(min_row=live.P_FIRST, values_only=True)
             if r[live.P_MEASURE - 1] == "outcome_loans" and r[live.P_TEST_BOOK - 1]}
    assert tests == {True: "exact test", False: "z test"}
    assert not any(x["worse"] == "too few loans to test" for x in tabs.pockets(ws))


def test_a_real_loss_keeps_its_red_when_profit_is_not_significant(tmp_path):
    """The seventh walk, defect 1: under a fixed revenue line, under 620 / Broker (GCO 5.35x, a finding) lost its
    shading because its revenue side could be chance. Each side is coloured on its own."""
    b = _ready(tmp_path, n=8000)
    _set(b, "FICO", book.C_EDGES, "620; 680; 740")
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=4):
        if r[6].value == "revenue_line":
            r[2].value = "0.25 points either way"
    wb.save(b)
    assert book.run(b).ok
    rows = tabs.pck_all(b, tmp_path / "grids")                  # every grid, each picked in turn
    planted = next(x for x in rows if str(x["band"]).endswith(" - 619") and x["seg"] == "Broker")
    assert planted["flags"]["g"] == engine.WORSE and planted["g_fill"] == house.ALERT_FG
    marked = [x for x in rows if x["flags"]["r"] in (engine.UNSURE_WORSE, engine.UNSURE_BETTER)]
    assert marked and all(x["r_fill"] is None for x in marked)            # the fixed line marks some
    # read together, a real loss beside a kept gap that could be chance is "Losing more, profit holding" (C, the
    # firm, 27 Sep 2026); any other pair with that kept side reads nothing together
    assert all(x["together"] == ("Losing more, profit holding" if x["flags"]["g"] == engine.WORSE else None)
               for x in marked), [(x["flags"], x["together"]) for x in marked]
    assert any(x["g_fill"] == house.ALERT_FG for x in marked)            # and their loss side keeps its colour
    real_worse = [x for x in rows if x["flags"]["g"] == engine.WORSE]
    assert real_worse and all(x["g_fill"] == house.ALERT_FG for x in real_worse)
    # Control explains the suggested option as it works now: each pocket's own test (defect 2)
    opts = [r[4] for r in load_workbook(b)["_options"].iter_rows(min_row=2, values_only=True)
            if r[1] == "revenue_line" and r[6] == "Each pocket's own test (suggested)"]
    assert opts and "own test" in opts[0] and "typical size" not in opts[0]


def test_the_category_limits_on_control_apply_at_set_up(tmp_path):
    """Found checking the seventh walk's blank "Last Run used" rows: Set up
    always used the recommended 12 and 50, whatever Control said."""
    from pocketbook import control
    b = _ready(tmp_path)
    assert book.run(b).ok
    wb = load_workbook(b)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == "few_values":
            r[control.CHOOSE_COL - 1].value, r[control.OWN_COL - 1].value = None, 3
    wb.save(b)
    extract = next(tmp_path.rglob("*.csv"))
    assert book.set_up(extract, b).ok
    ws = load_workbook(b)["Control"]
    # applied at Set up: the limit shows in the launcher's block, where it is chosen since the redesign
    assert control.answer_of("few_values", *(ws.cell(row=control.row_of(ws, "few_values"), column=c).value
                                             for c in (control.CHOOSE_COL, control.OWN_COL))) == 3
    # at 3, a column of four numbers is an amount, not a category (Set up's guess; a confirmed meaning wins)
    table = read_table(extract)
    four = next(c for c in table.columns if c == "ASSET_CLASS")
    assert meanings.suggest(table, few_values=3)[four].means != "category"
    assert meanings.suggest(table)[four].means == "category"
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    assert book.run(b).ok
    used = {r[control.KEY_COL - 1].value: r[control.LAST_COL - 1].value
            for r in load_workbook(b)["Control"].iter_rows(min_row=control.FIRST_ROW)}
    assert used["band_count"] and used["band_cut"]


def test_the_category_limits_chosen_in_the_launcher_change_set_ups_guesses(tmp_path):
    """The limits moved to the launcher (redesign phase 1) and the test above checks only that they are
    shown: CI's planted bug "category limits ignored" went uncaught. What Set up guesses must follow
    them: at 3, a column of four values is an amount; at the usual 12, a category."""
    from pocketbook import choices as ch
    got = {}
    for few in (3, 12):
        made = book.set_up(synth.write_extract(tmp_path / str(few), n=2000), choices=ch.Choices(few_values=few))
        ws = load_workbook(made.book)["Columns"]
        got[few] = next(r[book.C_SUGG - 1].value for r in ws.iter_rows(min_row=book.COL_FIRST)
                        if r[book.C_NAME - 1].value == "ASSET_CLASS")
    assert got == {3: "amount", 12: "category"}


def test_revenue_reads_the_same_on_both_tabs(tmp_path):
    """The seventh walk, defect 3, and the firm's call (25 Sep 2026): the profit line on Control decides profit on
    Pockets too, so a pocket can't read "keeps less" on one tab and "in line" on the other. Both read the one
    flag on _pockets: Worse? on Pockets is RANR vs GCOs' RANR side, pocket by pocket."""
    for option in ("Each pocket's own test (suggested)", "0.25 points either way"):
        b = _ready(tmp_path / option[:4], n=8000)
        wb = load_workbook(b)
        for r in wb["Control"].iter_rows(min_row=4):
            if r[6].value == "revenue_line":
                r[2].value = option
        wb.save(b)
        assert book.run(b).ok
        k = tabs.choose(b, tmp_path / option[:4] / "k.xlsx", results.POCKETS, measure="RANR")
        v = recalc(k, tmp_path / option[:4] / "rc")
        pk = v[results.PCK]
        kept = {pk.cell(row=x["row"], column=results.C_H_RR).value: x["flags"]["r"] for x in tabs.pck(pk)}
        ws = v[results.POCKETS]
        seen = 0
        for x in tabs.pockets(ws):
            prow = ws.cell(row=x["row"], column=results.K_ROW).value
            if prow in kept:
                assert x["worse"] == live.worse_of(kept[prow]), (option, x, kept[prow])
                seen += 1
        assert seen, option


def test_words_after_a_number_start_clear_of_it(tmp_path):
    """The render of 26 Sep 2026 printed "10.9%worse" and "80.4loans". The redesign's rule 4: numbers, verdicts
    and short answers centred; labels left, with an indent (LibreOffice draws no indent on a cell left to the
    default)."""
    b = _ready(tmp_path, n=3000)
    assert book.run(b).ok
    wb = load_workbook(b)
    ws = wb[results.POCKETS]
    first = tabs.header_row(ws, results.K_NUM, "#") + 1
    for c in (results.K_BAND, results.K_SEG, results.K_HOLDS):
        got = ws.cell(row=first, column=c).alignment
        assert got.indent >= 1 and got.horizontal == "left", c
    for c in (results.K_LOANS, results.K_GAP, results.K_EX, results.K_WORSE, results.K_P, results.K_MAT):
        assert ws.cell(row=first, column=c).alignment.horizontal == "center", c
    pk = wb[results.PCK]
    first = tabs.header_row(pk, results.C_TOG, "Together") + 1
    got = pk.cell(row=first, column=results.C_TOG).alignment
    assert got.indent >= 1 and got.horizontal == "left"


