"""Two choices the firm made on 1 Oct 2026.

1. Start here's "Dollars above their share, in those" added every worse-and-material pocket on every grid, and every
   loan sits in every grid: the bank's workbook read $2,904,231,129 above share on a book whose GCOs were
   $37,767,925. The firm chose distinct loans: each loan counts once, in the pocket where its own dollars above share
   are largest. The RANR tile ("N short $X") had the same fault and gets the same rule.
2. Filter 1 starts on Origination year when the extract has a column marked Origination date.

The checking arithmetic below reads the loan file with the csv module and works each loan's dollars out itself; it
takes from the Run only what the Run decided (which pockets are worse and material, and against what), never a
total."""

from __future__ import annotations

import csv
import math
from bisect import bisect_right
from pathlib import Path

import pytest
from openpyxl import load_workbook

from pocketbook import book, choices as ch, control, engine, launcher, perm, synth
from test_book import _answer


# --------------------------------------------------------------------------
# helpers: a Run kept with its engine result, and the tiles read back


def _ran(d: Path, bands: tuple, segments: tuple, n: int = 6000):
    got = {}
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 400)
        run = engine.run

        def kept(*a, **k):
            got["res"] = run(*a, **k)
            return got["res"]
        mp.setattr(engine, "run", kept)
        x = synth.write_extract(d / "src", n=n)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=bands, segments=segments))
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return {"x": x, "b": out.book, "res": got["res"], "ran": ran}


@pytest.fixture(scope="module")
def many(tmp_path_factory):
    """Two band columns by two segment columns: four grids, every loan in each."""
    return _ran(tmp_path_factory.mktemp("many"), ("FICO", "ORIG_BAL"), ("CHANNEL", "ASSET_CLASS"))


@pytest.fixture(scope="module")
def one(tmp_path_factory):
    """One grid: every loan in exactly one pocket."""
    return _ran(tmp_path_factory.mktemp("one"), ("FICO",), ("CHANNEL",))


def _tile(ws, label: str):
    for row in ws.iter_rows():
        for c in row:
            if c.value == label:
                return ws.cell(row=c.row + 1, column=c.column).value
    raise KeyError(label)


def _calc(path: Path, out: Path):
    from recalc import recalc
    return recalc(path, out)


# --------------------------------------------------------------------------
# the independent count: plain Python over the loan file


def _num(raw):
    try:
        return float(raw)
    except (TypeError, ValueError):
        return None


def _band_index(raw, edges) -> object:
    """Which numeric band a loan's value falls in (0 is below the first edge), or what it is instead."""
    v = _num(raw)
    if raw in (None, ""):
        return "(blank)"
    if v is None:
        return "(not a number)"
    if v == -9999:                       # _answer treats FICO's -9,999 as missing
        return "(marked missing)"
    return bisect_right(edges, v)


def _pockets_of(res, measure: str):
    """What the Run decided, and nothing it added up: each worse-and-material pocket on each two-way grid as
    (grid, band field, band, segment field, segment, whether it is judged against its band), with the pocket's
    own top and bottom for checking that the loans found here are the pocket's."""
    fields = {b.name: b.field for b in res.config.bands}
    fields.update({d.name: d.field for d in res.config.dimensions})
    out = []
    for gi, g in enumerate(res.grids):
        for (bl, dl), c in g.inner():
            s = c.rates[measure]
            if s.flag == engine.WORSE and s.material and s.dollars and s.dollars > 0:
                out.append((gi, g, fields[g.band], bl, fields[g.dimension], dl, s.by_band, s.num, s.den))
    return out


def _independent(x: Path, res, measure: str) -> dict:
    """Each loan once, worked out from the loan file: every loan in at least one of the pockets counts its top less
    its bottom at the rate its pocket is compared with (the rest of the book, or the rest of its band), and where it
    is in several, the largest of those. A loss rate's top is GCO, profit's is RANR and its shortfall is turned."""
    m = next(m for m in res.measures if m.name == measure)
    worse = m.higher_is == "worse"
    with open(x, newline="", encoding="utf-8") as fh:
        loans = list(csv.DictReader(fh))
    yx = [(_num(r[m.value]), _num(r[m.per])) for r in loans]
    enters = [y is not None and x_ is not None for y, x_ in yx]
    book_y = math.fsum(y for (y, _), e in zip(yx, enters) if e)
    book_x = math.fsum(x_ for (_, x_), e in zip(yx, enters) if e)
    best: dict[int, float] = {}
    pockets = _pockets_of(res, measure)
    for gi, g, bfield, bl, dfield, dl, by_band, num, den in pockets:
        numeric = [lab for lab in g.band_labels if not lab.startswith("(")]
        edges = res.band_edges[g.band]
        assert len(numeric) == len(edges) + 1, (g.band, numeric, edges)
        want = numeric.index(bl) if bl in numeric else bl
        band_of = [_band_index(r[bfield], edges) for r in loans]
        inside = [i for i, r in enumerate(loans) if band_of[i] == want and r[dfield] == dl and enters[i]]
        # the loans found here are the pocket's: the same top and bottom as the Run's pocket
        assert math.fsum(yx[i][0] for i in inside) == pytest.approx(num, abs=1e-6)
        assert math.fsum(yx[i][1] for i in inside) == pytest.approx(den, abs=1e-6)
        if by_band:
            peers = [i for i in range(len(loans)) if band_of[i] == want and enters[i]]
            ry = math.fsum(yx[i][0] for i in peers) - num
            rx = math.fsum(yx[i][1] for i in peers) - den
        else:
            ry, rx = book_y - num, book_x - den
        r = ry / rx
        for i in inside:
            y, x_ = yx[i]
            e = y - r * x_ if worse else r * x_ - y
            best[i] = max(best.get(i, e), e)
    return {"dollars": math.fsum(best.values()), "loans": len(best), "pockets": len(pockets),
            "grids": len({p[0] for p in pockets}), "book_top": book_y,
            "top_of_counted": math.fsum(yx[i][0] for i in best)}


# --------------------------------------------------------------------------
# 1. each loan once


def test_each_loan_once_ties_to_the_loan_file_and_never_exceeds_the_books_gcos(many):
    res = many["res"]
    want = _independent(many["x"], res, "gco_rate")
    got = res.once["gco_rate"]
    assert want["pockets"] >= 2 and want["grids"] >= 2, want            # the book is one where it matters
    assert got.dollars == pytest.approx(want["dollars"], rel=1e-9)
    assert (got.loans, got.pockets, got.grids) == (want["loans"], want["pockets"], want["grids"])
    assert 0 < got.dollars <= want["top_of_counted"] <= want["book_top"]
    # and the old sum counted a loan once per pocket it sits in: more than each loan once on this book
    assert got.pocket_sum > got.dollars
    # the launcher's tile says the same number
    assert many["ran"].summary["dollars"] == pytest.approx(want["dollars"], rel=1e-9)


def test_the_ranr_shortfall_counts_each_loan_once_too(many):
    res = many["res"]
    want = _independent(many["x"], res, "ranr_rate")
    got = res.once["ranr_rate"]
    assert want["pockets"] >= 1, want
    assert got.dollars == pytest.approx(want["dollars"], rel=1e-9)
    assert (got.loans, got.pockets, got.grids) == (want["loans"], want["pockets"], want["grids"])


def test_one_grid_each_loan_once_is_the_pockets_sum(one):
    res = one["res"]
    got = res.once["gco_rate"]
    pockets = [c.rates["gco_rate"] for g in res.grids for _, c in g.inner()]
    old = math.fsum(s.dollars for s in pockets if s.flag == engine.WORSE and s.material and s.dollars > 0)
    assert got.pockets >= 1 and got.grids == 1
    assert got.dollars == pytest.approx(old, rel=1e-9) == pytest.approx(got.pocket_sum, rel=1e-9)
    assert got.dollars == pytest.approx(_independent(one["x"], res, "gco_rate")["dollars"], rel=1e-9)


def test_start_here_shows_each_loan_once_and_says_so(many, tmp_path):
    res = many["res"]
    got = res.once
    ws = _calc(many["b"], tmp_path / "calc")["Start here"]
    assert _tile(ws, "Dollars above share, each loan once") == pytest.approx(got["gco_rate"].dollars, abs=0.01)
    n, g = got["gco_rate"].pockets, got["gco_rate"].grids
    count = _tile(ws, "Pockets worse and material, GCOs")
    assert count.startswith(f"{n} of ") and f" · in {g} grids" in count, count
    short = _tile(ws, "Pockets short on RANR, each loan once")
    assert short == f"{got['ranr_rate'].pockets} short ${got['ranr_rate'].dollars:,.0f}", short
    note = " ".join(str(c.value) for row in ws.iter_rows(max_row=8) for c in row if c.value)
    assert "each loan once" in note and "once per grid" in note


def test_start_here_says_run_again_when_control_has_moved_the_pockets(many, tmp_path):
    """Materiality takes effect live: the pockets on Start here move with it, and a total of each loan once is the
    Run's, so it says to Run rather than show the Run's number beside other pockets."""
    b = tmp_path / "moved.xlsx"
    wb = load_workbook(many["b"])
    ws = wb[control.SHEET]
    ws.cell(row=control.row_of(ws, "materiality"), column=control.CHOOSE_COL).value = \
        "5% of the book's total losses"
    wb.save(b)
    start = _calc(b, tmp_path / "calc")["Start here"]
    count = _tile(start, "Pockets worse and material, GCOs")
    assert not count.startswith(f"{many['res'].once['gco_rate'].pockets} of "), count     # it did move
    assert " in " not in count.split(" · borderline")[0].split(" of ")[1]
    assert _tile(start, "Dollars above share, each loan once") == book.ONCE_STALE


def test_a_found_from_before_each_loan_once_reads_run_again(many, tmp_path):
    """A workbook whose last Run predates 1 Oct 2026 has no total of each loan once: it says to Run, never the old
    sum."""
    b = tmp_path / "old.xlsx"
    wb = load_workbook(many["b"])
    ws = wb[book.FOUND]
    for r in range(ws.max_row, 0, -1):
        if str(ws.cell(row=r, column=1).value or "").startswith("once"):
            ws.delete_rows(r)
    # Set up writes Start here again from _found, as it does on the bank's machine
    del wb["Start here"]
    ws = wb.create_sheet("Start here", 0)
    book._start_here(ws, wb, many["x"], 6000, 10)
    wb.save(b)
    start = _calc(b, tmp_path / "calc")["Start here"]
    assert _tile(start, "Dollars above share, each loan once") == book.ONCE_STALE


# --------------------------------------------------------------------------
# 2. Filter 1 starts on Origination year


def _flow(x) -> launcher.Flow:
    f = launcher.Flow(gate=launcher.AddOns())
    f.pick(str(x))
    f.set_up()
    assert f.screen() == "L2", f.message
    return f


def test_filter_1_starts_on_origination_year_and_can_be_cleared_or_changed(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = synth.write_extract(tmp_path / "src", n=1500)
    f = _flow(x)
    assert f.filter == ch.ORIG_YEAR and f.filter2 is None
    by = {r["name"]: r for r in f.rows()}
    assert by[ch.ORIG_YEAR]["d"] == {"on": True, "radio": True}
    assert by[ch.ORIG_YEAR]["e"] == {"on": False, "radio": True}           # Filter 2 is the analyst's
    assert f.choices().filter == ch.ORIG_YEAR and f.choices().filter2 is None
    f.click(ch.ORIG_YEAR, "d")                                            # cleared
    assert f.filter is None and f.choices().filter is None
    f.click("CHANNEL", "d")                                               # or changed
    assert f.filter == "CHANNEL"
    # cleared, written, and read back by the next Set up: the analyst's clearing holds
    f.click("CHANNEL", "d")
    f.pick_outcome("BAD_FLAG")
    f.answer_outcome(True)
    f.next()
    assert f.page == "answer", f.message
    assert _flow(x).filter is None


def test_no_origination_date_no_filter(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = synth.write_extract(tmp_path / "src", n=800)
    with open(x, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    bare = tmp_path / "bare" / "nodate.csv"
    bare.parent.mkdir()
    with open(bare, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=[c for c in rows[0] if not c.endswith("_DATE")])
        w.writeheader()
        w.writerows([{k: v for k, v in r.items() if not k.endswith("_DATE")} for r in rows])
    f = _flow(bare)
    assert f.filter is None and all(r["name"] != ch.ORIG_YEAR for r in f.rows())


def test_too_many_years_start_with_no_filter(tmp_path, monkeypatch):
    """A default Next would refuse is no default."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(ch, "FILTER_MOST_VALUES", 2)
    f = _flow(synth.write_extract(tmp_path / "src", n=800))
    assert f.filter is None


def test_the_window_draws_filter_1_on_origination_year(tmp_path, monkeypatch):
    """The same on screen. Needs a display (xvfb-run on Linux)."""
    from test_deps import _window
    from pocketbook import deps
    root = _window()
    try:
        monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
        monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
        monkeypatch.setattr(deps, "missing", lambda: [])
        monkeypatch.setattr(deps, "missing_optional", lambda: [])
        w = launcher.build(root)
        flow = w["flow"]
        x = synth.write_extract(tmp_path / "src", n=800)
        w["extract"].set(str(x))
        flow.pick(str(x))
        flow.set_up()
        w["render"]()
        root.update()
        boxes = w["boxes"]
        on = lambda name, which: len(boxes[(name, which)].find_all()) == 2       # noqa: E731  a dot in the ring
        assert on(ch.ORIG_YEAR, "d") and not on(ch.ORIG_YEAR, "e") and not on("CHANNEL", "d")
        flow.click(ch.ORIG_YEAR, "d")
        w["render"]()
        root.update()
        assert not w["boxes"][(ch.ORIG_YEAR, "d")].find_all()[1:]
    finally:
        root.destroy()
