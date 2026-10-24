"""American spelling (the firm, 14 Oct 2026, on the wording review's item S: "American spelling", everywhere, tab
and workbook names included): gray, color, colored, labeled, recognizes.

Only the words moved, and some of them are read back. A workbook written before the change carries the British
words in its cells, and one of those cells is looked up by its label: the number typed beside Summary's "Grey rows
under" is read before a Run rewrites the tab, and put back after. So a workbook the old code wrote must still Set up
and Run, keep its answers and that number, and come out in the new words. It is held here on a workbook turned back
into the old words cell by cell, as the code before the change wrote them (the pattern of
test_literal_names_2026_09_30, which held the "Paid, cost, kept" rename on a workbook the old code wrote)."""

from __future__ import annotations

import re
import shutil

import pytest
from openpyxl import load_workbook

from pocketbook import book, control, perm, results, summary_chart, synth
from test_book import _answer

#: the British forms the code wrote before 14 Oct 2026; none may come back
BRITISH = re.compile(r"\b(?:grey|greys|greyed|colour|colours|coloured|labelled|labelling|recognise|recognises|"
                     r"recognised)\b", re.I)
#: the change undone, word for word: what an older workbook holds where this one holds the American word
BACK = [("Gray", "Grey"), ("gray", "grey"), ("Colored", "Coloured"), ("colored", "coloured"), ("Color", "Colour"),
        ("color", "colour"), ("labeled", "labelled"), ("recognizes", "recognises"), ("recognized", "recognised")]


def _british(text: str) -> str:
    for new, old in BACK:
        text = re.sub(rf"\b{new}\b", old, text)
    return text


def _strings(wb):
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    yield ws, c


def _answers(wb) -> dict:
    """Control's answers, by key: the two answer cells of every setting row."""
    ws = wb[control.SHEET]
    return {r[control.KEY_COL - 1].value: (r[control.CHOOSE_COL - 1].value, r[control.OWN_COL - 1].value)
            for r in ws.iter_rows(min_row=control.FIRST_ROW) if r[control.KEY_COL - 1].value}


@pytest.fixture(scope="module")
def ran(tmp_path_factory):
    d = tmp_path_factory.mktemp("american")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 200)
        x = synth.write_extract(d, n=3000)
        b = book.set_up(x).book
        _answer(b)
        assert book.run(b).ok
    return d, x, b


def test_a_run_writes_no_british_spelling_on_any_sheet(ran):
    _, _, b = ran
    wb = load_workbook(b)
    left = [f"{ws.title}!{c.coordinate}: {c.value[:100]}" for ws, c in _strings(wb) if BRITISH.search(c.value)]
    assert left == []
    assert wb[results.SUMMARY]["B2"].value == summary_chart.GREY_LABEL == "Gray rows under"


def test_a_workbook_written_in_the_old_words_still_sets_up_and_runs_with_its_answers_kept(ran, tmp_path,
                                                                                            monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    d, x0, b0 = ran
    x, b = tmp_path / x0.name, tmp_path / b0.name
    shutil.copy(x0, x)
    shutil.copy(b0, b)
    wb = load_workbook(b)
    # the workbook as the code before the change wrote it: every word back in British spelling
    for ws, c in _strings(wb):
        c.value = _british(c.value)
    ws = wb[results.SUMMARY]
    assert ws["B2"].value == "Grey rows under"
    ws["C2"].value = 75                                     # the analyst's own number beside the old label
    wb[book.ABOUT]["B1"] = str(x.resolve())
    wb.save(b)
    old = load_workbook(b)
    assert [c.coordinate for _, c in _strings(old) if BRITISH.search(c.value)]      # it is the old words
    assert summary_chart.typed_grey(old) == 75                                      # read by the old label
    answers = _answers(old)
    # Set up again (the analyst's Next, with the new kit), then Run
    assert book.set_up(x, b).ok
    assert _answers(load_workbook(b)) == answers
    out = book.run(b, x)
    assert out.ok, out.lines
    wb = load_workbook(b)
    assert _answers(wb) == answers
    assert (wb[results.SUMMARY]["B2"].value, wb[results.SUMMARY]["C2"].value) == ("Gray rows under", 75)
    left = [f"{ws.title}!{c.coordinate}: {c.value[:100]}" for ws, c in _strings(wb) if BRITISH.search(c.value)]
    assert left == []


def test_the_grey_number_is_put_back_beside_either_label():
    """keep_grey finds the number's cell by its label, old or new."""
    from openpyxl import Workbook
    for label in ("Grey rows under", "Gray rows under"):
        wb = Workbook()
        ws = wb.active
        ws.title = results.SUMMARY
        ws.cell(row=summary_chart.GREY_ROW, column=2, value=label)
        ws.cell(row=summary_chart.GREY_ROW, column=3, value=40)
        assert summary_chart.typed_grey(wb) == 40
        summary_chart.keep_grey(wb, 60)
        assert ws.cell(row=summary_chart.GREY_ROW, column=3).value == 60
