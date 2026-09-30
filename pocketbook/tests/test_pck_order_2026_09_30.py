"""RANR vs GCOs: in band order, sortable in Excel, headed against the rest, with each side's Rest (the firm,
30 Sep 2026).

"it would be really nice if the bands could be sorted or at least make it easier to look at what's happening on
this tab. Hard to see these in a sensical order." The rows came out by GCO dollars, bands and segments jumbled.
Now: band ascending, the lowest first, (blank) last; then segment; and Excel's sort and filter arrows on the
headings, so any column can be sorted.

The headings said "gap vs book" and "× book", which read as the whole book; the numbers are against the rest of
the book, the pocket left out (or the rest of its band, when Control judges against the band). Checked by hand on
the firm's row, 720-739 Non-Customer/VLA: booked 3,613,815, GCOs 464,248, RANR -37,084, in a book of 234,472,291
booked, 10,816,762 GCOs and 10,439,188 RANR. GCOs ÷ Booked 12.85% against the rest's 4.48% is 2.86×; against the
whole book's 4.61% it would be 2.78×. Each side now shows the rest's rate beside its gap: RANR + GCOs 9.02%, GCOs
4.48%, RANR 4.54%.

The book here is built to those numbers. Every expected rest is worked out in this file from the loans alone,
with nothing from pocketbook in the arithmetic."""

from __future__ import annotations

import math
import re
import shutil
import tempfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.utils import range_boundaries

from conftest import cube, row, table
from recalc import values_of
from pocketbook import engine, results
import tabs

BANDS = [{"name": "score", "field": "SCORE", "edges": [720, 740]}]
VLA, BRANCH, BROKER = "Non-Customer/VLA", "Customer/Branch", "Broker"
#: the firm's row and the whole book, in cents
FIRM = {"booked": 361_381_500, "gco": 46_424_800, "ranr": -3_708_400}
BOOK = {"booked": 23_447_229_100, "gco": 1_081_676_200, "ranr": 1_043_918_800}


def _bench(compare_to: str) -> dict:
    return {"min_units": 30, "min_events": 5, "worse_at": 1.25, "better_at": 0.8, "confidence": 0.95, "power": 0.8,
            "compare_to": compare_to, "many_tests": "none", "materiality": 5_000, "shuffles": 300}


def _split(total: int, weights: list[int]) -> list[int]:
    """Whole cents shared by weight, the last taking what rounding leaves, so the parts add to the total exactly."""
    out = [total * w // sum(weights) for w in weights[:-1]]
    return out + [total - sum(out)]


def _loans() -> list[dict]:
    """The firm's pocket (730 is in 720 - 739; 100 loans, 20 charged off), and eight others plus a blank-score
    pocket making up the rest of the book to the firm's totals, each with its own loss rate."""
    rest = {k: BOOK[k] - FIRM[k] for k in BOOK}
    others = [(s, c) for s in (700, 730, 760) for c in (BROKER, BRANCH, VLA) if (s, c) != (730, VLA)] + [("", BROKER)]
    n = len(others)
    booked, gco = _split(rest["booked"], [1] * n), _split(rest["gco"], list(range(1, n + 1)))
    ranr = _split(rest["ranr"], [n + 1 - k for k in range(1, n + 1)])
    pockets = [((730, VLA), 100, 20, FIRM["booked"], FIRM["gco"], FIRM["ranr"])]
    pockets += [(o, 300, 15 + 3 * k, booked[k], gco[k], ranr[k]) for k, o in enumerate(others)]
    out, i = [], 0
    for (score, chan), loans, bad, b, g, r in pockets:
        bs, gs, rs_ = _split(b, [1] * loans), _split(g, [1] * bad), _split(r, [1] * loans)
        for j in range(loans):
            out.append(row(i, score, chan, bs[j] / 100, int(j < bad), gs[j] / 100 if j < bad else 0.0,
                           rs_[j] / 100))
            i += 1
    return out


LOANS = _loans()


def _band_of(score, labels):
    """A loan's band from the grid's labels ("720 - 739"), worked out here."""
    if score == "":
        return "(blank)"
    for lab in labels:
        m = re.match(r"^\s*(-?[\d.]+)\s*-\s*(-?[\d.]+)\s*$", str(lab))
        if m and float(m.group(1)) <= score <= float(m.group(2)):
            return lab
    raise AssertionError((score, labels))


def _sums(loans) -> dict:
    return {"booked": math.fsum(x["BAL"] for x in loans), "gco": math.fsum(x["GCO"] for x in loans),
            "ranr": math.fsum(x["RANR"] for x in loans)}


def _rates(s) -> dict:
    """Each side's rate: (RANR + GCOs) ÷ Booked, GCOs ÷ Booked, RANR ÷ Booked."""
    return {"paid_r": (s["ranr"] + s["gco"]) / s["booked"], "cost_r": s["gco"] / s["booked"],
            "kept_r": s["ranr"] / s["booked"]}


def _rest_of(band, seg, labels, by_band: bool) -> dict:
    """The rest's rates for one pocket: every loan but the pocket's, in the book or in its band; a pocket alone in
    its band is compared with the rest of the book."""
    mine = lambda x: _band_of(x["SCORE"], labels) == band                          # noqa: E731
    alone = len({x["CHAN"] for x in LOANS if mine(x)}) == 1
    rest = [x for x in LOANS if not (mine(x) and x["CHAN"] == seg) and (mine(x) or not by_band or alone)]
    return _rates(_sums(rest))


@pytest.fixture(scope="module")
def made():
    """The tab for each comparison: the workbook as written, and calculated."""
    out = {}
    d = Path(tempfile.mkdtemp(prefix="pck-order-"))
    try:
        for compare_to in ("topline", "peers"):
            res = engine.run(cube(bands=BANDS, benchmark=_bench(compare_to)), table(LOANS))
            wb = Workbook()
            results.write(wb, res, "now")
            v = values_of(wb, d / compare_to)
            ws = v[results.PCK]
            ws.formulas = wb[results.PCK]
            out[compare_to] = {"res": res, "wb": wb, "ws": ws, "chart": v[results.CHART]}
        yield out
    finally:
        shutil.rmtree(d, ignore_errors=True)


def _rows(ws) -> list[dict]:
    rows = tabs.pck(ws)
    for x in rows:
        for k, c in (("paid_r", results.C_PAID_R), ("cost_r", results.C_COST_R), ("kept_r", results.C_KEPT_R)):
            x[k] = ws.cell(row=x["row"], column=c).value
    return rows


def test_the_book_is_the_firms():
    """The book adds up to the firm's figures, and the firm's row to theirs, and the hand check holds: 2.86× is
    against the rest, not the whole book."""
    whole = _sums(LOANS)
    assert (round(whole["booked"]), round(whole["gco"]), round(whole["ranr"])) == (234_472_291, 10_816_762,
                                                                                  10_439_188)
    firm = _sums([x for x in LOANS if x["SCORE"] == 730 and x["CHAN"] == VLA])
    assert (round(firm["booked"]), round(firm["gco"]), round(firm["ranr"])) == (3_613_815, 464_248, -37_084)
    rest = _rates({k: whole[k] - firm[k] for k in whole})
    assert (round(rest["paid_r"], 4), round(rest["cost_r"], 4), round(rest["kept_r"], 4)) == (0.0902, 0.0448,
                                                                                              0.0454)
    assert round(firm["gco"] / firm["booked"] / rest["cost_r"], 2) == 2.86
    assert round(firm["gco"] / firm["booked"] / (whole["gco"] / whole["booked"]), 2) == 2.78


def test_the_firms_row_shows_the_rest_it_is_measured_against(made):
    ws = made["topline"]["ws"]
    got = {(x["band"], x["seg"]): x for x in _rows(ws)}
    x = got[("720 - 739", VLA)]
    assert (x["booked"], x["gco"], x["ranr"]) == pytest.approx((3_613_815, 464_248, -37_084), abs=0.005)
    assert (round(x["paid_r"], 4), round(x["cost_r"], 4), round(x["kept_r"], 4)) == (0.0902, 0.0448, 0.0454)
    assert round(x["cost"], 2) == 2.86
    assert ws.formulas.cell(row=x["row"], column=results.C_COST_R).number_format == results.RATE_FMT


@pytest.mark.parametrize("compare_to", ["topline", "peers"])
def test_every_rows_rest_is_the_rest_of_the_book_or_its_band(made, compare_to):
    ws, res = made[compare_to]["ws"], made[compare_to]["res"]
    labels = res.grids[0].band_labels
    rows = _rows(ws)
    assert len(rows) == 10
    for x in rows:
        want = _rest_of(x["band"], x["seg"], labels, compare_to == "peers")
        for k in ("paid_r", "cost_r", "kept_r"):
            assert x[k] == pytest.approx(want[k], abs=1e-12), (compare_to, x["band"], x["seg"], k)
    # the multiple is the pocket's own rate over the rest's
    for x in rows:
        assert x["cost"] == pytest.approx(x["gco"] / x["booked"] / x["cost_r"], rel=1e-9)


@pytest.mark.parametrize("compare_to, rest", [("topline", "book"), ("peers", "band")])
def test_the_headings_say_the_rest(made, compare_to, rest):
    ws = made[compare_to]["ws"]
    head = tabs.header_row(ws, results.C_TOG, "Together")
    assert [ws.cell(row=head - 1, column=c).value for c in (results.C_PAID, results.C_COST, results.C_KEPT)] == [
        f"RANR + GCOs · gap vs rest of {rest}", f"GCOs · × rest of {rest}", f"RANR · gap vs rest of {rest}"]
    assert tabs.heads(ws, head, results.C_PAID, results.C_TOG) == [
        "Gap pts", "Dollars", "Rest", f"× rest of {rest}", "Dollars", "Rest", "Gap pts", "Dollars", "Rest", "Together"]
    # each group's rule runs under its Rest too
    assert all(ws.formulas.cell(row=head - 1, column=c).border.bottom.style == "medium"
               for c in range(results.C_PAID, results.C_TOG + 1))


def _band_key(label):
    m = re.match(r"^\s*(-?[\d.]+)", str(label))
    return (0, float(m.group(1))) if m else (1, 0.0)


def test_rows_are_by_band_lowest_first_then_segment(made):
    ws, res = made["topline"]["ws"], made["topline"]["res"]
    got = [(x["band"], x["seg"]) for x in _rows(ws)]
    want = sorted(got, key=lambda k: (_band_key(k[0]), k[1].lower()))
    assert got == want
    assert got[0][0].startswith("7") and got[-1][0] == "(blank)"          # the lowest first, (blank) last
    assert got == [(x["band"], x["seg"]) for x in results.pck_rows(res, res.grids[0])]
    note = " ".join(str(c.value) for r in ws.iter_rows(max_row=30) for c in r if isinstance(c.value, str))
    assert "Rows are by band, lowest first, then by segment" in note
    assert "read together first" not in note and "GCO dollars" not in note


def test_the_arrows_cover_the_pockets_and_not_the_totals(made):
    ws = made["topline"]["ws"]
    f = ws.formulas
    head = tabs.header_row(ws, results.C_TOG, "Together")
    rows = _rows(ws)
    c0, r0, c1, r1 = range_boundaries(f.auto_filter.ref)
    assert (c0, r0) == (results.C_BAND, head)
    assert r1 == rows[-1]["row"] and c1 >= results.C_H_BR >= results.C_TOG      # the workings sort with each row
    totals = [r for r in range(1, ws.max_row + 1)
              if ws.cell(row=r, column=results.C_BAND).value in (results.PCK_LISTED, results.PCK_BOOK)]
    assert len(totals) == 2 and all(r > r1 for r in totals)
    # the chart sits right of every sorted column, and the hidden columns are all inside the arrows' range
    assert all(f.column_dimensions[results.col(c)].hidden for c in range(results.C_H, results.C_H_BR + 1))
    assert not f.column_dimensions[results.col(results.C_CH)].hidden
    (chart,) = f._charts
    assert range_boundaries(chart.anchor)[0] == results.C_CH


REF = re.compile(r"(?<![A-Za-z_'!])\$?([A-Z]{1,3})(\$?)(\d+)(?![\d(])")


def test_every_rule_on_the_rows_reads_its_own_row(made):
    """A sort moves rows; a rule that read another row would colour the wrong pocket."""
    f = made["topline"]["ws"].formulas
    rows = _rows(made["topline"]["ws"])
    first, last = rows[0]["row"], rows[-1]["row"]
    seen = 0
    for rng in f.conditional_formatting:
        for bounds in str(rng.sqref).split():
            _, r0, _, r1 = range_boundaries(bounds)
            if r0 != first:
                continue
            for rule in rng.rules:
                for _c, fixed, n in REF.findall(rule.formula[0]):
                    assert not fixed and int(n) == first, (bounds, rule.formula[0])
                seen += 1
    assert seen >= 8


def _sorted_by_excel(wb, rows, order):
    """The tab as Excel leaves it after a sort: each row's cells in the arrows' range moved to its new row, a formula's
    relative rows moving with it and its absolute ones staying, as a copy and paste does."""
    ws = wb[results.PCK]
    c0, r0, c1, r1 = range_boundaries(ws.auto_filter.ref)
    was = {r: [ws.cell(row=r, column=c).value for c in range(c0, c1 + 1)] for r in range(r0 + 1, r1 + 1)}
    for new, old in zip(range(r0 + 1, r1 + 1), order):
        for c, v in zip(range(c0, c1 + 1), was[old]):
            if isinstance(v, str) and v.startswith("="):
                v = Translator(v, origin=f"{results.col(c)}{old}").translate_formula(f"{results.col(c)}{new}")
            ws.cell(row=new, column=c, value=v)


def test_a_sort_in_excel_keeps_every_row_whole_its_colours_and_the_chart(made, tmp_path):
    ws = made["topline"]["ws"]
    made["topline"]["wb"].save(tmp_path / "book.xlsx")
    wb = load_workbook(tmp_path / "book.xlsx")                   # a copy: the fixture's stays as written
    rows = _rows(ws)
    before = {(x["band"], x["seg"]): x for x in rows}
    # sorted by GCOs' multiple, the worst first, as a person would with the arrow
    order = [x["row"] for x in sorted(rows, key=lambda x: -x["cost"])]
    _sorted_by_excel(wb, rows, order)
    v = values_of(wb, tmp_path / "rc")
    after_ws = v[results.PCK]
    after_ws.formulas = wb[results.PCK]
    after = _rows(after_ws)
    assert [(x["band"], x["seg"]) for x in after] == [(x["band"], x["seg"]) for x in
                                                      sorted(rows, key=lambda x: -x["cost"])]
    keys = ("loans", "booked", "gco", "ranr", "ranr_rate", "paid", "paid_d", "paid_r", "cost", "cost_d", "cost_r",
            "kept", "kept_d", "kept_r", "together_said", "c_fill", "g_fill", "r_fill", "flags")
    for x in after:
        b = before[(x["band"], x["seg"])]
        assert {k: x[k] for k in keys} == {k: b[k] for k in keys}, (x["band"], x["seg"])
    # the red line above and the totals under the table are untouched
    assert any("1 pocket lost money outright" in str(c.value) for r in after_ws.iter_rows(max_row=40) for c in r)
    # the chart numbers the first pockets read together in the new order, and names them under it
    named = [f"{x['band']} / {x['seg']}" for x in after if x["together"]][:results.LABELLED]
    hs = v[results.CHART]
    assert [hs.cell(row=j, column=results.H_LNAME).value for j in range(1, len(named) + 1)] == named
    assert named, "the book reads some pockets together"
    for j, name in enumerate(named, start=1):
        x = next(y for y in after if f"{y['band']} / {y['seg']}" == name)
        assert hs.cell(row=j, column=results.H_NX).value == pytest.approx(x["cost"])
        assert hs.cell(row=j, column=results.H_NY).value == pytest.approx(x["kept"])
        assert hs.cell(row=j, column=results.H_NUM).value == j
    listed = [str(c.value) for r in after_ws.iter_rows() for c in r
              if isinstance(c.value, str) and re.match(r"^\d  ", c.value)]
    assert listed == [f"{j}  {n}" for j, n in enumerate(named, start=1)]

