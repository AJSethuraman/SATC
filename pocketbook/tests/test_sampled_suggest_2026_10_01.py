"""Set up's suggestions from a sample of the grids (the firm, 1 Oct 2026).

The bank's book: 184,937 loans, 12 band columns x 14 segment columns, 168 grids, and Set up's "Working out the
suggestions" took over 3.5 minutes. The firm: "Will the quick estimates be as accurate? ... Test it and let's see".

Worse at and better at read one thing from the grids: how many loans each pocket holds (book.luck_gap). Fewest
loans reads the whole book's rate and no grid at all. So on a big book Set up builds a sample of the grids
(book.suggest_pairs): every segment column equally often, every band column within one of equally often. On nine
synthetic books of 17,000 and 185,000 loans (BACKLOG.md 6d) a sample of 28 matched every grid at the 0.01x shown;
Run still works each one out from every grid, and says so on Control where Set up's estimate was different.
"""

from __future__ import annotations

import csv
import math
import random
from collections import Counter

import pytest
from openpyxl import load_workbook

from pocketbook import book, choices as ch, config as cfgmod, control, engine
from conftest import cube, table

NB, ND = 8, 10                       # 80 grids: past the sample's 30, as the bank's 168 are past its 28


def _loans(n: int, nb: int, nd: int, seed: int) -> tuple[list[dict], list[str], list[str]]:
    """A book shaped like the bank's: number columns of several shapes, categories of 3 to 7 values in an uneven
    mix, the last one skewed (one value on about 88% of loans), and an outcome that leans on some of each."""
    rng = random.Random(seed)
    nums = [f"N{i:02d}" for i in range(nb)]
    cats = []
    for j in range(nd):
        k = rng.randint(3, 7)
        w = [0.88] + [0.12 / (k - 1)] * (k - 1) if j == nd - 1 else [1 / (i + 1) ** rng.uniform(0.3, 1.6)
                                                                       for i in range(k)]
        cats.append((f"S{j:02d}", [f"v{j}{chr(65 + i)}" for i in range(k)], w,
                     [math.exp(rng.gauss(0, 0.35)) for _ in range(k)]))
    slopes = [rng.choice([0, 0.3, -0.4, 0.6]) for _ in nums]
    rows = []
    for i in range(n):
        r = {"ID": f"L{i}"}
        lean = 0.0
        for c, s in zip(nums, slopes):
            z = rng.gauss(0, 1)
            lean += s * z
            r[c] = round(700 + 50 * z) if c[-1] in "02468" else round(math.exp(9 + 0.8 * z), 2)
        for c, vals, w, eff in cats:
            k = rng.choices(range(len(vals)), w)[0]
            r[c] = vals[k]
            lean += math.log(eff[k])
        bad = int(rng.random() < min(0.9, 0.04 * math.exp(0.35 * lean) / 1.6))
        r.update(BAL=round(rng.uniform(5000, 60000), 2), BAD=bad)
        r["GCO"] = round(r["BAL"] * 0.5, 2) if bad else 0.0
        r["RANR"] = round(r["BAL"] * 0.05 - r["GCO"], 2)
        rows.append(r)
    return rows, nums, [c for c, *_ in cats]


@pytest.fixture(scope="module")
def wide():
    """20,000 loans, 8 band columns x 10 segment columns, cut at Set up's draft bands (5 of equal loans), with
    the first pass's settings: no shuffle test, and the suggestions' own floors."""
    rows, nums, cats = _loans(20_000, NB, ND, seed=1001)
    cfg = cube(bands=[{"name": c, "field": c, "count": 5, "cut": "equal_loans"} for c in nums],
               dimensions=[{"name": c, "field": c} for c in cats], measures=[{"name": "loans", "mode": "count"}])
    cfg = cfgmod.Config(**{**cfg.__dict__, "benchmark": cfgmod.Benchmark(**{**cfg.benchmark.__dict__,
                                                                             "shuffles": 0})})
    return cfg, table(rows)


def _shown(values: dict) -> dict:
    return {k: book._said(k, v) for k, v in values.items()}


def test_the_sample_suggests_what_every_grid_does_at_the_rounding_shown(wide):
    cfg, t = wide
    pairs = book.suggest_pairs([b.name for b in cfg.bands], [d.name for d in cfg.dimensions])
    assert pairs is not None and len(pairs) == 30
    every, fb_every = book._suggest_values(engine.run(cfg, t), set(book.SUGGEST_KEYS))
    some, fb_some = book._suggest_values(engine.run(cfg, t, pairs=pairs), set(book.SUGGEST_KEYS))
    assert fb_every == fb_some == set()
    assert _shown(some) == _shown(every)
    assert some["min_loans"] == every["min_loans"]                 # the whole book's rate: never sampled


def test_the_sample_builds_only_its_grids_and_leaves_the_book_whole(wide):
    """The pocket sizes come from the sampled grids; every loan, every rate the book needs, is still read."""
    cfg, t = wide
    pairs = book.suggest_pairs([b.name for b in cfg.bands], [d.name for d in cfg.dimensions])
    res = engine.run(cfg, t, pairs=pairs)
    assert {(g.band, g.dimension) for g in res.grids} == pairs
    assert res.total.rows == len(t.rows)


def test_the_sample_takes_every_segment_column_equally_often():
    bands = [f"b{i}" for i in range(12)]
    segs = [f"s{i}" for i in range(14)]
    got = book.suggest_pairs(bands, segs)
    assert len(got) == 28                                          # the bank's 168 grids, and the size tested
    per_seg = Counter(d for _, d in got)
    per_band = Counter(b for b, _ in got)
    assert set(per_seg) == set(segs) and set(per_seg.values()) == {2}
    assert set(per_band) == set(bands) and max(per_band.values()) - min(per_band.values()) <= 1
    assert book.suggest_pairs(bands, segs) == got                  # the same on every Set up
    # more band columns than segment columns: still every segment equally often, every band at least once
    wider = book.suggest_pairs([f"b{i}" for i in range(20)], [f"s{i}" for i in range(6)])
    assert set(Counter(d for _, d in wider).values()) == {len(wider) // 6}
    assert len({b for b, _ in wider}) == 20


def test_a_small_book_reads_every_grid():
    assert book.suggest_pairs([f"b{i}" for i in range(4)], [f"s{i}" for i in range(7)]) is None    # 28 grids
    assert book.suggest_pairs(["b"], ["s"]) is None
    assert book.suggest_pairs([], ["s"]) is None


def _write(rows: list[dict], path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)


@pytest.fixture(scope="module")
def set_up_wide(tmp_path_factory):
    """Set up on 3,000 loans, 6 number columns x 6 categories (36 grids, sampled to 30), as the launcher chooses."""
    d = tmp_path_factory.mktemp("sampled")
    rows, nums, cats = _loans(3_000, 6, 6, seed=1002)
    for r in rows:
        r["LOAN_NBR"] = r.pop("ID")
        r["BAD_FLAG"] = r.pop("BAD")
        r["ORIG_BAL"], r["GCO_AMT"], r["RANR_AMT"] = r.pop("BAL"), r.pop("GCO"), r.pop("RANR")
    _write(rows, d / "loans.csv")
    built = []
    real = engine.run

    def spy(cfg, t, *a, **k):
        res = real(cfg, t, *a, **k)
        built.append(len(res.grids))
        return res
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(engine, "run", spy)
        out = book.set_up(d / "loans.csv", choices=ch.Choices(run_kind=ch.BLEED, bands=tuple(nums),
                                                              segments=tuple(cats), outcome="BAD_FLAG"))
    assert out.ok, out.lines
    return out, built


def _suggest_cells(wb) -> dict[str, str]:
    ws = wb[control.SHEET]
    return {r[control.KEY_COL - 1].value: ws.cell(row=r[0].row, column=book.SUGGEST_COL).value
            for r in ws.iter_rows(min_row=control.FIRST_ROW) if r[control.KEY_COL - 1].value in book.SUGGEST_KEYS}


def test_set_up_builds_only_the_sampled_grids_and_says_so_on_control(set_up_wide):
    out, built = set_up_wide
    assert built == [30]                                           # one first pass, 30 of the 36 grids
    wb = load_workbook(out.book)
    said = _suggest_cells(wb)
    for key in book.SAMPLED_KEYS:
        assert said[key].startswith("suggested: ") and "a quick estimate from 30 of its 36 grids" in said[key], said
    assert said["min_loans"].endswith(", from this extract")          # never sampled, so never called an estimate
    assert wb[book.ABOUT]["A5"].value == book.QUICK
    quick = book._quick_of(wb)
    assert quick[1] == (30, 36) and quick[0]["worse_at"] == out.summary["suggested"]["worse_at"]


def test_run_says_where_every_grid_disagrees_with_set_ups_estimate(set_up_wide):
    """What Run writes beside each suggestion, through the function it writes with: the same words as before where
    Set up's estimate agrees; where it doesn't, what the estimate was; and where the answer chosen is that estimate,
    that every grid says otherwise."""
    out, _ = set_up_wide
    wb = load_workbook(out.book)
    when = "from this extract at the last Run"
    every = {"min_loans": 70, "worse_at": 1.33, "better_at": 0.75}
    agree = ({"worse_at": 1.33, "better_at": 0.75}, (30, 36))
    book._suggestions(wb[control.SHEET], every, set(), when=when, quick=agree, used={"worse_at": 1.33})
    said = _suggest_cells(wb)
    assert said["worse_at"] == f"suggested: 1.33x, {when}" and said["better_at"] == f"suggested: 0.75x, {when}"

    differ = ({"worse_at": 1.32, "better_at": 0.76}, (30, 36))
    book._suggestions(wb[control.SHEET], every, set(), when=when, quick=differ,
                      used={"worse_at": 1.32, "better_at": "luck"})
    said = _suggest_cells(wb)
    assert said["worse_at"] == (f"suggested: 1.33x, {when}. The answer chosen, 1.32x, is Set up's quick estimate "
                                f"from 30 of 36 grids: every grid says 1.33x")
    assert said["better_at"] == f"suggested: 0.75x, {when}. Set up's quick estimate, from 30 of 36 grids, was 0.76x"
    assert said["min_loans"] == f"suggested: 70, {when}"


def test_a_run_on_the_estimate_says_every_grid_disagrees(set_up_wide, tmp_path):
    """Through a real Run: the answer typed on Control is Set up's estimate, and the estimate (written here, so the
    test does not wait for a book where the sample happens to miss) is not what every grid says."""
    import shutil
    from test_book import _answer
    out, _ = set_up_wide
    path = tmp_path / out.book.name
    shutil.copy(out.book, path)
    shutil.copy(out.book.parent / "loans.csv", tmp_path / "loans.csv")
    _answer(path, odd=False)
    wb = load_workbook(path)
    wb[book.ABOUT]["B5"] = "grids=30/36;worse_at=9.99;better_at=0.1"
    wb[book.ABOUT]["B1"] = str((tmp_path / "loans.csv").resolve())
    for r in wb[control.SHEET].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == "worse_at":
            r[control.OWN_COL - 1].value = 9.99
    wb.save(path)
    got = book.run(path)
    assert got.ok, got.lines
    said = _suggest_cells(load_workbook(path))["worse_at"]
    every = said.split()[1].rstrip(",")
    assert said == (f"suggested: {every}, from this extract at the last Run. The answer chosen, 9.99x, is Set up's "
                    f"quick estimate from 30 of 36 grids: every grid says {every}")
