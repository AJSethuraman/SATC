"""The bank-machine checklist (docs/BANK-MACHINE-CHECKLIST.md, Goal 3 item 3) says what to type and where to
click on a machine nobody here can reach. Every command, file and cell it names is held here, so the checklist
can't go stale while the code moves under it: the practice-book commands make the books whose fingerprints it
prints, the kit holds the files it lists and no bank data, the add-ons it installs are the ones PocketBook asks
for, and every cell it sends the analyst to still holds what it says."""

from __future__ import annotations

import hashlib
import html
import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from openpyxl import load_workbook

from pocketbook import book, choices as ch, control, deps, synth

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "docs" / "BANK-MACHINE-CHECKLIST.md"
TEXT = MD.read_text(encoding="utf-8")
sys.path.insert(0, str(ROOT / "tools"))
import bank_kit  # noqa: E402


def _blocks() -> list[str]:
    return re.findall(r"```\n(.*?)\n```", TEXT, re.S)


def test_the_practice_book_commands_make_the_books_the_checklist_fingerprints(tmp_path):
    """3.1 and 5.2: each `py -c` line, run as typed from the PocketBook folder, prints the path the checklist
    says, and certutil's MD5 of that file is the one the checklist prints."""
    try:
        (tmp_path / "src").symlink_to(ROOT / "src", target_is_directory=True)
    except OSError:       # Windows refuses a symlink without Developer Mode or admin (WinError 1314): copy instead
        shutil.copytree(ROOT / "src", tmp_path / "src", ignore=shutil.ignore_patterns("__pycache__"))
    commands = [line for b in _blocks() for line in b.splitlines() if line.startswith("py -c") and "synth" in line]
    assert len(commands) == 2
    for line in commands:
        code = re.fullmatch(r'py -c "(.*)"', line).group(1)
        assert "%" not in code                   # Command Prompt would read %...% as a variable
        out = subprocess.run([sys.executable, "-c", code], cwd=tmp_path, capture_output=True, text=True, check=True)
        printed = out.stdout.strip()
        windows = printed.replace("/", "\\")
        assert f"It prints `{windows}`" in TEXT, windows
        said = re.search(rf'certutil -hashfile "{re.escape(windows)}" MD5.{{0,400}}?`([0-9a-f]{{32}})`', TEXT, re.S)
        assert said, f"no MD5 printed for {windows}"
        assert hashlib.md5((tmp_path / printed).read_bytes()).hexdigest() == said.group(1)


def test_the_kit_holds_what_the_checklist_lists_and_no_bank_data(tmp_path):
    z = bank_kit.make_zip(tmp_path)
    names = zipfile.ZipFile(z).namelist()
    assert all(n.startswith("PocketBook/") for n in names)
    for listed in ("PocketBook.pyw", "Install add-ons.bat", "Install add-ons from this folder.bat", "VERSION.txt",
                   "README.md", "src/pocketbook/__init__.py", "src/pocketbook/settings.yaml",
                   "docs/walkthrough/2026-09-27/PROCEDURE-pocketbook-analyst.pdf", "docs/BANK-MACHINE-CHECKLIST.pdf"):
        assert f"PocketBook/{listed}" in names, listed
    # 2.2's dir listing names these
    for said in ("PocketBook.pyw", "Install add-ons.bat", "Install add-ons from this folder.bat", "VERSION.txt",
                 "README.md"):
        assert f"`{said}`" in TEXT
    # nothing an analyst could have left in the folder: no extract, workbook, pre-spec, memory or cache
    assert not [n for n in names if n.endswith((".csv", ".xlsx", ".xls", ".pyc", ".json")) or "__pycache__" in n
                or (n.endswith(".yaml") and n != "PocketBook/src/pocketbook/settings.yaml")]
    version = zipfile.ZipFile(z).read("PocketBook/VERSION.txt").decode()
    assert version.startswith("PocketBook ") and "\nCommit: " in version
    # Part 7 says the commit line is what matters
    assert 'type "%USERPROFILE%\\PocketBook\\VERSION.txt"' in TEXT


def test_the_add_ons_are_the_same_everywhere_the_checklist_installs_them():
    wanted = set(deps.NEEDED) | set(deps.OPTIONAL)
    kit = {re.split(r"[<>=]", a)[0] for a in bank_kit.ADD_ONS}
    assert kit == wanted
    bat = (ROOT / "Install add-ons from this folder.bat").read_text(encoding="utf-8")
    offline = [ln for b in _blocks() for ln in b.splitlines() if "--no-index" in ln]
    assert len(offline) == 1
    for line in [offline[0], *[ln for ln in bat.splitlines() if "pip install" in ln]]:
        assert "--user" in line and "--no-index" in line
        assert set(line.split("scikit-learn")[0].split()[-3:] + ["scikit-learn"]) == wanted, line
    assert '"%~dp0add-ons"' in bat and "%USERPROFILE%\\PocketBook\\add-ons" in offline[0]
    assert (ROOT / "Install add-ons from this folder.bat").read_bytes().count(b"\r\n") >= 5   # a .bat keeps CRLF


def test_the_html_is_built_from_the_markdown_as_it_stands():
    """The HTML and PDF are what the firm prints; a heading or command added to the Markdown and not built
    would be missing from them."""
    page = html.unescape((ROOT / "docs" / "BANK-MACHINE-CHECKLIST.html").read_text(encoding="utf-8"))
    for h in re.findall(r"^#{2,3} (.+)$", TEXT, re.M):
        assert re.sub(r"[*`]", "", h) in page, h
    for b in _blocks():
        assert b in page, b
    assert (ROOT / "docs" / "BANK-MACHINE-CHECKLIST.pdf").stat().st_size > 100_000
    for img in re.findall(r"\]\(([^)]+\.png)\)", TEXT):
        assert (ROOT / "docs" / img).exists(), img


#: every cell Part 4 sends the analyst to, and what it must hold there
PICK = {"min_loans": "Enough for 5 expected losses (suggested)", "min_events": "10 losses",
        "materiality": "1% of the book's total losses", "compare_to": "The rest of its band",
        "worse_at": "The smallest significant gap in a typical pocket (suggested)",
        "better_at": "The smallest significant gap in a typical pocket (suggested)", "confidence": "95% sure",
        "revenue_line": "Each pocket's own test (suggested)"}


def _lists(ws) -> dict[str, object]:
    out = {}
    for v in ws.data_validations.dataValidation:
        for cell in str(v.sqref).split():
            out[cell] = v
    return out


def _listed(wb, v) -> list:
    sheet, cells = v.formula1.lstrip("=").split("!")
    return [c.value for row in wb[sheet.strip("'")][cells.replace("$", "")] for c in row if c.value is not None]


def test_every_cell_the_checklist_names_holds_what_it_says(tmp_path):
    extract = synth.write_extract(tmp_path, n=3000)
    c = ch.Choices(run_kind=ch.BLEED, bands=("FICO", "ORIG_BAL"), segments=("CHANNEL", "ASSET_CLASS"),
                   split="REV_DEBT")
    out = book.set_up(extract, choices=c)
    wb = load_workbook(out.book)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value in PICK:
            r[control.CHOOSE_COL - 1].value = PICK[r[control.KEY_COL - 1].value]
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    for r in book.table_rows(wb["Columns"]):
        key = r[book.C_QKEY - 1].value
        if isinstance(key, str) and key.count("|") == 2:
            r[book.C_TREAT - 1].value = "Missing" if key.startswith("FICO") else "Real"
    wb.save(out.book)
    assert book.run(out.book, extract).ok
    wb = load_workbook(out.book)

    def said(cell_words: str):
        assert cell_words in TEXT, cell_words

    ctl, lists = wb["Control"], _lists(wb["Control"])
    said("Control C15")
    assert ctl["B15"].value == "How much worse than its comparison a pocket must be"
    assert _listed(wb, lists["C15"])[0] == PICK["worse_at"] and _listed(wb, lists["C15"])[-1] == "2 times"
    assert len(_listed(wb, lists["C15"])) == 4
    said("type `0.9` in D15")
    d15 = lists["D15"]
    assert (d15.type, d15.operator, d15.formula1, d15.errorTitle) == ("decimal", "between", "1.01", "Out of range")
    assert d15.error in TEXT
    said("Control C19")
    assert ctl["B19"].value == "Smallest excess loss worth reporting"
    said("**100 loans** in C24")
    assert ctl["B24"].value.startswith("Fewest loans in a pocket") and "100 loans" in _listed(wb, lists["C24"])
    said("Control H24")
    assert ctl["H23"].value == "Status" and "Waiting for a Run" in ctl["H24"].value

    for tab, cells, sizes in (("Pockets", ("C18", "D18", "E18"), (5, 2, 3)), ("Paid, cost, kept", ("B16",), (4,)),
                              ("Grids", ("B13", "F13"), (8, 5)), ("Split", ("B15", "B26"), (4, 5))):
        lv = _lists(wb[tab])
        for cell, n in zip(cells, sizes):
            assert len(_listed(wb, lv[cell])) == n, (tab, cell)
    said("Pockets C18, D18, E18")
    assert "Charge-offs" in _listed(wb, _lists(wb["Pockets"])["C18"])
    assert "Worse and material" in _listed(wb, _lists(wb["Pockets"])["E18"])
    assert "FICO x ASSET_CLASS" in _listed(wb, _lists(wb["Grids"])["B13"])
    for tab, cell in (("Pockets", "K15"), ("Paid, cost, kept", "M13"), ("Split", "I12")):
        said(f"{tab} {cell}")
        assert "waits for a Run" in str(wb[tab][cell].value), (tab, cell)

    look = wb["Look"]
    said("In C20 (**Bars**)")
    assert look["B10"].value == "FICO" and look["B20"].value == "Bars" and _lists(look)["C20"].formula1 == '"10,20,50"'
    assert (look["B21"].value, look["B22"].value) == ("From", "To")
    said("in I14 (FICO's band edges)")
    row = next(r[0].row for r in book.table_rows(wb["Columns"]) if r[book.C_NAME - 1].value == "FICO")
    assert f"{book._col(book.C_EDGES)}{row}" == "I14"


def _paste(tmp_path, text: str):
    """Paste the script as the firm will on the bank machine: into PocketBook.py, with Windows line endings."""
    import runpy

    f = tmp_path / "PocketBook.py"
    f.write_bytes(text.replace("\n", "\r\n").encode("utf-8"))
    return runpy.run_path(str(f), run_name="pasted")["unpack"]


def test_the_paste_script_writes_back_every_file_exactly(tmp_path):
    """27 Sep 2026, the firm: "i expect the file to be sent in script form ... i will just paste it into py's".
    PocketBook.py carries the window and all its code; pasted with Windows line endings and run, it writes a
    PocketBook folder holding each file byte for byte as it is here, and nothing else."""
    n = _paste(tmp_path, bank_kit.paste_text())()
    want = {rel: f.read_text(encoding="utf-8") for f, rel in bank_kit.paste_files()}
    assert n == len(want) and "src/pocketbook/launcher.py" in want and "PocketBook.pyw" in want
    got = {p.relative_to(tmp_path / "PocketBook").as_posix(): p.read_text(encoding="utf-8")
           for p in (tmp_path / "PocketBook").rglob("*") if p.is_file()}
    assert got == want


def test_a_paste_cut_short_is_refused_and_writes_nothing(tmp_path):
    """A paste that stops before the last line names how far it got and writes no folder: half a PocketBook
    would open and fail somewhere deep inside a Run."""
    text = bank_kit.paste_text()
    last = text.rindex("#=== FILE ")
    unpack = _paste(tmp_path, text[:last])
    try:
        unpack()
        raise AssertionError("a cut-short paste was unpacked")
    except SystemExit as e:
        assert "stops early" in str(e)
    assert not (tmp_path / "PocketBook").exists()


def test_a_paste_with_one_line_changed_is_refused(tmp_path):
    """One character different in one file (an editor's autocorrect, a lost indent) fails that file's SHA-256,
    and the refusal names the file."""
    text = bank_kit.paste_text()
    at = text.index("#|", text.index("#=== FILE src/pocketbook/stats.py"))
    unpack = _paste(tmp_path, text[:at] + "#| " + text[at + 2:])
    try:
        unpack()
        raise AssertionError("a changed paste was unpacked")
    except SystemExit as e:
        assert "src/pocketbook/stats.py" in str(e)
    assert not (tmp_path / "PocketBook").exists()


def test_the_checklist_says_how_many_files_the_paste_writes():
    """2.2 tells the analyst what a good paste prints; the count is the script's own."""
    said = re.search(r"`(\d+) files, every one checked\. Opening the window\.`", TEXT)
    assert said and int(said.group(1)) == len(bank_kit.paste_files())
