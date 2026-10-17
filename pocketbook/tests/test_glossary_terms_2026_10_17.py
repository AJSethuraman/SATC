"""The Glossary's custom terms (the firm, 17 Oct 2026).

The firm: "our custom terms - like bad / good loans, these i do want in the glossary along with our key terms as
fairly plain language and simple explanations to walk a user through what it means. keep my voice in mind". And on
the p-value: "luck is seriously only a term we will understand, make sure it's clear that luck is basically the
other side of confidence", which the Glossary now says through How sure and Significant at N%?.

The guard: every verdict word, reading, Show option, Measure option and Together reading the workbook can print has
a Glossary row whose term names it. Each word is read from the constant the tabs themselves print, so a reading
added later without a Glossary row fails here. Then every new entry renders in a built workbook, before and after
a Run, and the examples that use the book's own figures tie to the extract."""

from __future__ import annotations

import csv
import re
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook

from conftest import TEST_SHUFFLES
from pocketbook import book, choices as ch, engine, glossary, live, perm, results, synth
from test_book import _answer

#: the rows added on 17 Oct 2026, in the order the Glossary shows them
NEW = ["Judged against", "Good loan", "Bad loans / Bad loans % / Bad loans ÷ Loans",
       "Bad dollars / Bad dollars ÷ Booked", "Loan size", "Two-way / Split by",
       "Across the bottom / Lines by / Panels by",
       "Excess / Bad loans above share / Dollars above share / Dollars short of band / Dollars short of book",
       "Worse? (Yes / Not sure / Too few losses / No)", "Material? / Material at",
       "Show (All / Worse and material / Worse or not sure)", "How sure / Significant at N%?",
       "Significant at N%? (Yes / Borderline / No)",
       "Too few losses to test / Nothing to compare with / Not losing more",
       "Profit line (How far profit must move before it counts as more or less)",
       "Holds (column) fixed? / Mixed with (column)", "Could have caught", "Together", "Priced for it", "Net drain",
       "Strong", "Safe but idle", "Earns less, not from losses", "Losing more, profit holding",
       "Show on Line size (Pockets where something tripped / 2 or more signals / All pockets)",
       "Treat as (Real / Missing)", "Changes now / Needs a Run", "Comes to", "Worked out from the loans",
       "Last Run used / Status", "Lines in use now", "Gray rows under", "Found on / Held back", "Pre-spec",
       "Holds up? / Still holds?"]
#: the rows whose example uses the book's own figures after a Run, on the synthetic book (no split, no filter)
FROM_THE_BOOK = ["Judged against", "Good loan", "Bad loans / Bad loans % / Bad loans ÷ Loans",
                 "Bad dollars / Bad dollars ÷ Booked", "Loan size",
                 "Excess / Bad loans above share / Dollars above share / Dollars short of band / Dollars short of book",
                 "Worse? (Yes / Not sure / Too few losses / No)", "Material? / Material at",
                 "Show (All / Worse and material / Worse or not sure)", "How sure / Significant at N%?",
                 "Significant at N%? (Yes / Borderline / No)", "Profit line (How far profit must move before it "
                 "counts as more or less)", "Could have caught", "Together", "Priced for it", "Net drain", "Strong",
                 "Safe but idle", "Earns less, not from losses", "Losing more, profit holding",
                 "Treat as (Real / Missing)", "Comes to", "Lines in use now"]


def parts(term: str) -> list[str]:
    """A term's names: "Worse? (Yes / Not sure / No)" names Worse?, Yes, Not sure and No."""
    return [x.strip() for x in re.split(r" / |\(|\)", term) if x.strip()]


def named(word: str, terms: list[str]) -> list[str]:
    """The terms that name `word`: as one of their names, or with its own parenthetical dropped ("Months to
    charge-off (avg)"), or, for a word that is the start of a reading ("Mixed with FICO"), as that start."""
    bare = re.sub(r"\s*\([^)]*\)$", "", word).strip()
    return [t for t in terms if word.strip() in parts(t) or bare in parts(t)
            or (word.endswith(" ") and any(x.startswith(word.strip()) for x in parts(t)))]


def _sig_readings() -> list[str]:
    """Every reading Significant at N%? can print, from live.sig_said itself: every engine word, with and without a
    gap and a p-value, near the line or not, for a loss and for profit. A p-value's reading is its word."""
    out = set()
    for word in engine.READINGS + (None,):
        for gap in (None, 1.6, 1.1, 0.8, -0.002, 0.002):
            for p in (None, 0.001, 0.3):
                for near in (False, True):
                    for profit in (False, True):
                        out.add(live.sig_said(word, gap, p, near, profit, 0.05).split(" (p ")[0])
    return sorted(out)


def _readings() -> dict[str, list[str]]:
    """Everything the workbook can print as a verdict, a reading, a Show or Measure option or a Together reading,
    from the functions and constants the tabs print it with."""
    split_by = results.split_label(SimpleNamespace(config=SimpleNamespace(split=("X", "own_median"))))
    return {
        "Worse?": sorted({live.worse_of(w) for w in engine.READINGS} - {""}),
        "Significant at N%?": [live.SIG_HEAD.format("N%"), results.MIXED] + _sig_readings(),
        "Show": list(results.SHOW) + list(results.LINE_SHOW),
        "Measure": list(results.PLAIN.values()) + [results.SIZE_NAME, results.CO_MONTHS_NAME]
                   + [x[1] for x in results.LINE_MEASURES],
        "Together": ["Together"] + list(results.TOGETHER.values()),
        "Pockets": [results.TWO_WAY, split_by.removesuffix(" X")],
        "Line size": [s[0] for s in results.SIGNALS],
        "Treat as": list(book.TREAT),
    }


# --------------------------------------------------------------------------
# the guard


@pytest.mark.parametrize("family", list(_readings()))
def test_every_word_the_tabs_can_print_has_a_glossary_row(family):
    terms = [t for t, _, _ in glossary.rows(None)]
    missing = [w for w in _readings()[family] if not named(w, terms)]
    assert not missing, (family, missing)


def test_each_verdict_rows_term_names_all_its_readings():
    """Worse?'s four readings in Worse?'s own row, Significant at N%?'s three in its own, and Tripped by's reasons
    in what Tripped by means."""
    rows = {t: (m, e) for t, m, e in glossary.rows(None)}
    worse = next(t for t in rows if t.startswith("Worse? ("))
    assert set(_readings()["Worse?"]) <= set(parts(worse)), worse
    sig = next(t for t in rows if t.startswith(live.SIG_HEAD.format("N%") + " ("))
    assert {live.SIG_YES, live.SIG_BORDER, live.SIG_NO} <= set(parts(sig)), sig
    assert set(_sig_readings()) == set(live.SIG_READINGS)
    assert all(r in rows["Tripped by"][0] for r in results.TRIPPED_BY)
    # the pairing the firm asked for: How sure and the p-value in one row, 95% to 5% to 1 in 20
    pair = rows[f"How sure / {live.SIG_HEAD.format('N%')}"][0]
    assert "95% sure" in pair and "under 5%" in pair and "1 in 20" in pair
    # the condition the p-value rests on, and that the printed one is after the allowance (the review, 17 Oct 2026)
    assert "if the pocket were no different from the rest" in pair and "allowance" in pair


def test_the_guard_fails_on_a_reading_with_no_row(monkeypatch):
    """The control: a Together reading added to the code without a Glossary row is caught."""
    monkeypatch.setitem(results.TOGETHER, ("less", "same"), "Quietly better")
    terms = [t for t, _, _ in glossary.rows(None)]
    assert len(terms) == len(set(terms))
    assert not named("Quietly better", terms)
    assert "Quietly better" in _readings()["Together"]


def test_no_glossary_row_says_luck_or_the_old_chance_words():
    text = " ".join(" ".join(r) for r in glossary.rows(None)).lower()
    for gone in ("luck", "how often by chance", "could be chance", "too close to call", "rarer than 1 in"):
        assert gone not in text, gone


# --------------------------------------------------------------------------
# in a built workbook, before and after a Run


@pytest.fixture(scope="module")
def books(tmp_path_factory):
    d = tmp_path_factory.mktemp("glossary-terms")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
        x = synth.write_extract(d, n=3000)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED))
        assert out.ok, out.lines
        set_up = load_workbook(out.book)
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
        return {"extract": x, "set_up": set_up, "run": load_workbook(out.book)}


def _rows(wb) -> dict[str, tuple[str, str]]:
    ws = wb[glossary.SHEET]
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=glossary.TERM).value == "Term")
    return {ws.cell(row=r, column=glossary.TERM).value: (ws.cell(row=r, column=glossary.MEANS).value,
                                                        ws.cell(row=r, column=glossary.EXAMPLE).value)
            for r in range(head + 1, ws.max_row + 1) if ws.cell(row=r, column=glossary.TERM).value}


def test_every_new_entry_renders_before_and_after_a_run_in_the_glossarys_order(books):
    for when in ("set_up", "run"):
        got = _rows(books[when])
        assert [t for t in got if t in NEW] == NEW, when
        for t in NEW:
            means, example = got[t]
            assert means and example and means.endswith(".") and example.endswith("."), (when, t)


def test_after_a_run_the_examples_use_the_books_figures(books):
    before, after = _rows(books["set_up"]), _rows(books["run"])
    for t in FROM_THE_BOOK:
        assert after[t][1] != before[t][1], t
    # the rows the book has no figure for keep their made-up example, never one presented as the book's
    for t in set(NEW) - set(FROM_THE_BOOK):
        assert not after[t][1].startswith("This book"), t


def _num(v):
    try:
        return float(v.replace(",", ""))
    except (AttributeError, ValueError):
        return None


def test_good_bad_and_bad_dollars_examples_tie_to_the_extract(books):
    """Good loans, bad loans, Bad loans % and Bad dollars ÷ Booked, worked out from the CSV without PocketBook:
    BAD_FLAG 1 is bad and 0 is good, any other value is neither, and Bad dollars are those loans' ORIG_BAL."""
    good = bad = 0
    bad_usd = usd = 0.0
    with open(books["extract"], newline="", encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            flag, bal = r["BAD_FLAG"].strip(), _num(r["ORIG_BAL"])
            if flag not in ("0", "1"):
                continue
            good, bad = good + (flag == "0"), bad + (flag == "1")
            if bal is not None:
                usd += bal
                bad_usd += bal if flag == "1" else 0.0
    rows = _rows(books["run"])
    assert rows["Good loan"][1] == (f"{good:,} of this book's {good + bad:,} loans with an outcome are good, and "
                                    f"{bad:,} are bad.")
    assert f"so its Bad loans % is {bad / (good + bad):.2%}." in rows["Bad loans / Bad loans % / Bad loans ÷ Loans"][1]
    assert rows["Bad dollars / Bad dollars ÷ Booked"][1] == (
        f"This book's bad loans booked ${bad_usd:,.0f} of ${usd:,.0f}, so Bad dollars ÷ Booked is "
        f"{bad_usd / usd:.2%}.")


def test_the_significant_rows_name_pockets_with_the_readings_pockets_prints(books):
    """Significant at N%?'s example after a Run names this book's pockets, each with a reading in the tab's own
    format, a Yes first."""
    ex = _rows(books["run"])[f"{live.SIG_HEAD.format('N%')} ({live.SIG_YES} / {live.SIG_BORDER} / {live.SIG_NO})"][1]
    said = re.findall(r"(.+?) read ((?:Yes|Borderline|No) \(p [=<] [\d.]+%\)) on (GCOs \(\$\)|Bad loans)\.", ex)
    assert said, ex
    assert said[0][1].startswith("Yes")


def test_the_allowance_is_worded_as_control_set_it():
    """With no allowance for many tests, no row claims one; with Benjamini-Hochberg, the row says what it holds."""
    def text(mt):
        rows = {t: m for t, m, _ in glossary.rows({"many_tests": mt, "co_date": True})}
        return " ".join(m for t, m in rows.items() if t.startswith(("How sure", live.SIG_HEAD.format("N%"))))
    assert "no allowance" in text("none") and "after the allowance" not in text("none")
    assert "share of false finds near the line" in text("bh")
    assert "chance of any false find" in text("bonferroni")


@pytest.mark.parametrize("conf", [0.9, 0.95, 0.975, 0.999])
def test_the_confidence_examples_follow_how_sure(conf):
    """At any How sure, Borderline's example is under its own line, and the readings say the line as _live does."""
    rows = {t: (m, e) for t, m, e in glossary.rows({"confidence": conf, "co_date": True})}
    bar = round(1 - conf, 12)
    ex = rows["Borderline"][1]
    m = re.search(r"borderline \(p ([\d.]+)\) has a p-value under the ([\d.]+)% threshold", ex)
    assert m and float(m.group(1)) < bar and m.group(2) + "%" == live.pct_words(bar), ex


def test_made_up_examples_use_made_up_names_after_a_run(books):
    """After a Run, an example that is not the book's names no column of this book, only made-up ones such as
    Column A, and the note says which examples are the book's."""
    import csv as csv_
    with open(books["extract"], newline="", encoding="utf-8") as fh:
        columns = next(csv_.reader(fh))
    after = _rows(books["run"])
    for t in set(NEW) - set(FROM_THE_BOOK):
        ex = after[t][1]
        assert not [c for c in columns if re.search(rf"\b{re.escape(c)}\b", ex)], (t, ex)
        assert not re.search(r"\b(SCORE|SOURCE|Direct|UTILIZATION|BALANCE|INCOME_TO_SALES)\b", ex), (t, ex)
    note = [c.value for row in books["run"][glossary.SHEET].iter_rows(max_col=3) for c in row]
    assert any(isinstance(v, str) and v.startswith("Examples naming this book, the last Run or one of its pockets")
               and "made-up names such as Column A" in v for v in note)


def test_a_split_run_names_its_own_split_and_partner(tmp_path, monkeypatch):
    """Where the Run knows the names, the example uses them: the split column, the column it moves with, and a grid
    that does not hold that column fixed."""
    monkeypatch.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
    out = book.set_up(synth.write_extract(tmp_path, n=3000),
                      choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO", "ORIG_BAL"), segments=("CHANNEL",),
                                         split="REV_DEBT"))
    assert out.ok, out.lines
    _answer(out.book)
    assert book.run(out.book).ok
    rows = _rows(load_workbook(out.book))
    assert rows["Two-way / Split by"][1].startswith("This book is split by REV_DEBT: each pocket is halved at its "
                                                    "own median REV_DEBT")
    holds = rows["Holds (column) fixed? / Mixed with (column)"][1]
    assert re.fullmatch(r"REV_DEBT moves with FICO \(correlation [+-][\d.]+\), so a split pocket in ORIG_BAL x "
                        r"CHANNEL / REV_DEBT reads No: may be mostly FICO, and Mixed with FICO\.", holds), holds
