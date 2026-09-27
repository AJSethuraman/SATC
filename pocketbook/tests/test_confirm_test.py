"""The confirmatory test (capabilities 4b and 4e): the pre-spec's column in its own groups, against its reference
group, inside its pockets, on the development range and the holdout; and how concentrated the bad loans are on
the holdout.

The goal it finishes (the firm, 25 Sep 2026): "the cube should be ready to test a derived column (income /
sales) against a dated outcome, on a holdout, from a committed pre-spec." It ends when, on the dated synthetic
book with the example pre-spec, a run finds the cliffs on development and confirms them on the holdout, and
Check no longer says "deviates" on the reference group. The first tests here are those two sentences.

Every count is checked against the extract read by hand (csv), never against the cube's own numbers."""

import bisect
import csv
import math
import shutil
import subprocess
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from conftest import cube, table
from pocketbook import book, confirm_tab, confirmatory, control, engine, house, kgroups, prespec, stats, synth
from openpyxl.utils import get_column_letter
from recalc import recalc
from test_book import _answer
from test_book_dates import _check, _choose, _columns, _control

needs_git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed on this machine")
EXAMPLE = Path(__file__).resolve().parents[1] / "docs" / "prespec-example.yaml"
BINS = [0.1, 0.25, 0.5, 1.0, 2.0]
PLANTED = {0: 2.0, 5: 3.0}                  # synth.RATIO_ODDS: below 0.1 twice the odds, above 2.0 three times


def _git(repo, *args):
    subprocess.run(["git", "-c", "user.name=Test", "-c", "user.email=test@example.com", "-c", "commit.gpgsign=false",
                    *args], cwd=repo, check=True, capture_output=True)


def _set_up(folder: Path, n: int, spec: Path = EXAMPLE, cut_no=("ORIG_BAL", "REV_DEBT", "ASSET_CLASS", "INCOME",
                                                                 "SALES", "income_to_sales")):
    """The dated synthetic book, income / sales made on Control under the example pre-spec's name, testing from
    that pre-spec, committed beside the workbook."""
    x = synth.write_extract(folder, n=n, ratio=True)
    b = book.set_up(x).book
    _answer(b)
    _control(b, run_kind="Finding and testing a new variable", new_variable_step="Test from a pre-spec",
             **{"derived|1": ("income_to_sales", "INCOME", "SALES")})
    book.set_up(x)
    _choose(b, drop=cut_no, shortlist="prespec.yaml")        # the saved shortlist, picked in the launcher
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    shutil.copy(spec, folder / "prespec.yaml")
    _git(folder, "init", "-q")
    _git(folder, "add", "prespec.yaml")
    _git(folder, "commit", "-q", "-m", "pre-spec")
    return x, b


@pytest.fixture(scope="module")
def route(tmp_path_factory):
    """The goal's run: 20,000 loans (synth's default), the example pre-spec as committed in docs/."""
    folder = tmp_path_factory.mktemp("goal")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("GIT_CEILING_DIRECTORIES", str(folder))
        mp.setenv("POCKETBOOK_MEMORY", str(folder / "memory.yaml"))
        x, b = _set_up(folder, 20000)
        seen = {}
        real = confirmatory.state

        def spy(*a, **k):
            seen["st"] = real(*a, **k)
            return seen["st"]

        mp.setattr(confirmatory, "state", spy)
        ran = book.run(b)
    assert ran.ok, ran.lines
    return {"x": x, "b": b, "ran": ran, "st": seen["st"], "t": seen["st"].test}


def _rows(x):
    with open(x, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _by_hand(x, lo: str, hi: str):
    """Loans made lo to hi (both ends in) with a readable ratio and outcome, as (group, bad, gco or None)."""
    out = []
    for r in _rows(x):
        if not lo <= r["ORIG_DATE"] <= hi or r["SALES"] in ("", "0") or r["BAD_FLAG"] not in ("0", "1"):
            continue
        g = bisect.bisect_right(BINS, float(r["INCOME"]) / float(r["SALES"]))
        try:
            gco = float(r["GCO_AMT"])
        except ValueError:
            gco = None
        out.append((g, int(r["BAD_FLAG"]), gco))
    return out


# --------------------------------------------------------------------------
# The goal


def test_the_goal_the_cliffs_are_found_on_development_and_confirmed_on_the_holdout(route):
    t = route["t"]
    assert t.problem is None and t.method == kgroups.CONDITIONAL
    assert t.column == "income_to_sales" and t.groups[t.ref] == "0.25 - 0.49" and t.strata == ("FICO", "CHANNEL")
    for side in (t.development, t.holdout):
        a, f = side.association, side.fit
        assert a.p_general < 0.05 and f.p_block < 0.05 and f.settled
        for k, planted in PLANTED.items():
            # each cliff is found: worse than the reference, significant, and the planted odds inside its range
            lo, hi = f.interval(k, 1.96)
            assert f.odds[k] > 1 and f.p[k] < 0.05 and lo < planted < hi, (side.name, k, f.odds[k], lo, hi)
    # on development, where nothing was planted between the cliffs, nothing between them reads significant
    assert all(t.development.fit.p[k] > 0.05 for k in (1, 3, 4))


def test_the_goal_check_no_longer_says_deviates(route):
    chk = _check(route["b"])
    warns = chk.get("Warning", [])
    warns = warns if isinstance(warns, list) else [warns]
    assert not [w for w in warns if w.startswith(confirmatory.DEVIATES)]
    assert chk["Differs from the pre-spec"] == "nowhere: this run used what it says"
    assert route["st"].deviations == []
    assert any(line.startswith("Follows pre-spec prespec.yaml (commit ") for line in route["ran"].lines)
    assert route["ran"].lines[-1].endswith("start with New variables.")


# --------------------------------------------------------------------------
# What went in: counted by hand from the extract


def test_the_holdout_holds_only_its_own_loans_counted_by_hand(route):
    t = route["t"]
    for side, (lo, hi) in ((t.holdout, ("2024-01-01", "2024-12-31")), (t.development, ("2022-01-01", "2023-12-31"))):
        mine = _by_hand(route["x"], lo, hi)
        assert side.n == len(mine)
        assert side.loans == [sum(1 for g, _, _ in mine if g == k) for k in range(6)]
        assert side.bad == [sum(y for g, y, _ in mine if g == k) for k in range(6)]
    rows = _rows(route["x"])
    outside = sum(1 for r in rows if not ("2022-01-01" <= r["ORIG_DATE"] <= "2024-12-31"))
    assert t.left_out[confirmatory.OUTSIDE] == outside
    chk = _check(route["b"])["Confirmatory test"]
    assert chk.startswith(f"Development 2022-01-01 to 2023-12-31: {t.development.n:,} loans, "
                          f"{t.development.n_bad:,} bad; holdout 2024-01-01 to 2024-12-31: {t.holdout.n:,} loans")
    assert f"{outside:,} made outside both ranges" in chk


def test_the_pockets_are_the_pre_specs_strata_as_the_grids_cut_them(route):
    """B3 worked out again here from the extract: each loan's FICO band from the edges Check says were used,
    its channel, its group, and nothing else."""
    t = route["t"]
    edges = [float(e) for e in _check(route["b"])["Band edges used: FICO"].split("(")[0].split(";")]
    cells = {}
    for r in _rows(route["x"]):
        if not "2024-01-01" <= r["ORIG_DATE"] <= "2024-12-31" or r["SALES"] in ("", "0") or r["BAD_FLAG"] not in ("0", "1"):
            continue
        f = r["FICO"]
        band = "blank" if f == "" else "missing" if f == "-9999" else bisect.bisect_right(edges, float(f))
        c = cells.setdefault((band, r["CHANNEL"]), [[0] * 6, [0] * 6])
        g = bisect.bisect_right(BINS, float(r["INCOME"]) / float(r["SALES"]))
        c[0][g] += 1
        c[1][g] += int(r["BAD_FLAG"])
    mine = kgroups.association([kgroups.Pocket(*c) for c in cells.values()], range(1, 7))
    assert t.holdout.association.general == pytest.approx(mine.general, rel=1e-12)
    assert t.holdout.association.trend == pytest.approx(mine.trend, rel=1e-12)


# --------------------------------------------------------------------------
# 4e: concentration on the holdout


def test_4e_capture_is_on_the_holdout_only_counted_by_hand(route):
    t = route["t"]
    mine = _by_hand(route["x"], "2024-01-01", "2024-12-31")
    n, nbad = len(mine), sum(y for _, y, _ in mine)
    gco = math.fsum(g for _, _, g in mine if g is not None)
    got = t.concentration()
    ws = load_workbook(route["b"])[confirm_tab.SHEET]
    table_rows = {ws.cell(row=r, column=2).value: r for r in range(1, ws.max_row + 1)}
    for k, name in enumerate(t.groups):
        inside = [x for x in mine if x[0] == k]
        want_flag, want_cap = len(inside) / n, sum(y for _, y, _ in inside) / nbad
        want_gco = math.fsum(g for _, _, g in inside if g is not None) / gco
        want_lift = (sum(y for _, y, _ in inside) / len(inside)) / (nbad / n)
        c = got[k]
        assert (c.flag_rate, c.capture, c.gco_capture) == pytest.approx((want_flag, want_cap, want_gco), rel=1e-12)
        assert c.lift == pytest.approx(want_lift, rel=1e-12)
        # and the tab's own table says the same (its last block: rows named by the group, share columns)
        r = max(rr for label, rr in table_rows.items() if label and str(label).startswith(name))
        C = confirm_tab.CONC_FIRST
        assert ws.cell(row=r, column=C + 1).value == pytest.approx(want_flag, rel=1e-12)
        assert ws.cell(row=r, column=C + 3).value == pytest.approx(want_cap, rel=1e-12)
        assert ws.cell(row=r, column=C + 5).value == pytest.approx(want_gco, rel=1e-12)
        assert ws.cell(row=r, column=C + 7).value == pytest.approx(want_lift, rel=1e-12)


# --------------------------------------------------------------------------
# The tab, calculated


@pytest.fixture(scope="module")
def calculated(route, tmp_path_factory):
    return recalc(route["b"], tmp_path_factory.mktemp("calc"))[confirm_tab.SHEET]


def _texts(ws) -> list[str]:
    return [str(c.value) for row in ws.iter_rows() for c in row if isinstance(c.value, str)]


def _under(ws, head: str) -> dict:
    """The rows under the CANVAS row starting `head`, by their label in B, up to the next such row or blank."""
    top = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=2).value or "").startswith(head))
    out = {}
    for r in range(top + 1, ws.max_row + 1):
        v = ws.cell(row=r, column=2).value
        if v in (None, "", "What it found"):
            break
        out.setdefault(v, r)
        out.setdefault(str(v).removesuffix(" (reference)"), r)
    return out


def _found(ws, lead: str) -> str:
    return next(s for s in _texts(ws) if s.startswith(lead) and "pockets held fixed" in s)


def test_the_tab_says_what_it_found_in_one_plain_line_per_range(route, calculated):
    t = route["t"]
    texts = _texts(calculated)
    assert not [s for s in texts if s.startswith("#") or "Err:" in s]
    for side, lead in ((t.development, "On development,"), (t.holdout, "On the holdout,")):
        line = _found(calculated, lead)
        f = side.fit
        for k in PLANTED:
            lo, hi = f.interval(k, stats.z_for_confidence(0.95))
            assert (f"{t.groups[k]} goes bad {f.odds[k]:.2f} times as often as 0.25 - 0.49 ({lo:.2f}x to "
                    f"{hi:.2f}x)") in line, line
        for k in range(6):
            if k != t.ref and f.p[k] >= 0.05:
                assert t.groups[k] + " goes" not in line
    first = next(s for s in texts if s.startswith("On the holdout, the bad rate differs across the groups"))
    assert first.endswith("rises steadily as income_to_sales rises.")          # general and trend both at 95%
    assert "Choices made here that no ruling settles yet" in texts
    for term in ("Mantel-Haenszel", "conditional logistic regression", "Odds are bad loans divided by good ones"):
        assert any(term in s for s in texts), term


@needs_git
def test_the_readings_follow_the_confidence_level_on_control(route, tmp_path):
    """OC-40: at 99% sure, a holdout p-value between 1% and 5% stops reading significant and the ranges widen,
    with no Run."""
    t = route["t"]
    mid = [k for k in range(6) if k != t.ref and t.holdout.fit.p[k] is not None and 0.01 < t.holdout.fit.p[k] < 0.05]
    assert mid, "the book has no holdout group between 1% and 5% to show the change"
    assert 0.01 < t.holdout.association.p_trend < 0.05
    copy = tmp_path / "at-99.xlsx"
    wb = load_workbook(route["b"])
    ws = wb[control.SHEET]
    ws.cell(row=control.row_of(ws, "confidence"), column=control.CHOOSE_COL).value = "99% sure"
    wb.save(copy)
    calc = recalc(copy, tmp_path / "calc")[confirm_tab.SHEET]
    line = _found(calc, "On the holdout,")
    for k in mid:
        assert t.groups[k] + " goes" not in line
    f = t.holdout.fit
    lo, hi = f.interval(5, stats.z_for_confidence(0.99))
    assert f"2.00 - {t.groups[5].split(' - ')[1]} goes bad {f.odds[5]:.2f} times as often as 0.25 - 0.49 " \
           f"({lo:.2f}x to {hi:.2f}x)" in line
    texts = _texts(calc)
    assert any(s.startswith("On the holdout, the bad rate differs across the groups, but not in one direction")
               for s in texts)
    assert "Range (99% sure)" in texts
    # the tables follow too: the trend's reading on the holdout, and the range beside each odds ratio, both with
    # the pockets held fixed (the tests in full: each set of loans under a CANVAS row naming it)
    tests = _under(calc, "Confirmed on the held-back loans, FICO and CHANNEL held fixed: ")
    D = confirm_tab.DATA
    assert calc.cell(row=tests["A steady climb or fall (trend)"], column=D + 3).value == "not significant"
    assert calc.cell(row=tests["Any difference across the groups (general)"], column=D + 3).value == "significant"
    groups = _under(calc, "Confirmed on the held-back loans: ")
    assert calc.cell(row=groups[t.groups[5]], column=D + 8).value == f"{lo:.2f}x to {hi:.2f}x"


# --------------------------------------------------------------------------
# The New variables tab's table (the redesign, section 9): each candidate found, confirmed, and confirmed with the
# columns held fixed, its excess and whether it is material, read as the analyst sees it


def _table(ws) -> dict:
    """The table's rows, by what each compares ("0.02 - 0.09 vs 0.25 - 0.49"): {column: value}, keyed by the
    table's own column constants."""
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=confirm_tab.N_CAND).value == "Candidate")
    out = {}
    for r in range(head + 1, ws.max_row + 1):
        comp = ws.cell(row=r, column=confirm_tab.N_COMP).value
        if not comp:
            break
        out[comp] = {c: ws.cell(row=r, column=c).value for c in range(confirm_tab.N_CAND, confirm_tab.N_WORDS + 1)}
        out[comp]["row"] = r
    return out


def _one_pocket(x, lo: str, hi: str):
    """The loans made lo to hi, counted by hand from the extract, all in one pocket: nothing held fixed."""
    mine = _by_hand(x, lo, hi)
    return kgroups.Pocket([sum(1 for g, _, _ in mine if g == k) for k in range(6)],
                          [sum(y for g, y, _ in mine if g == k) for k in range(6)])


def _book_gco(x) -> float:
    """The book's GCO as GCO per booked dollar counts it: every loan with a readable GCO and booked amount."""
    total = 0.0
    for r in _rows(x):
        try:
            gco, bal = float(r["GCO_AMT"]), float(r["ORIG_BAL"])
        except ValueError:
            continue
        if bal > 0:
            total += gco
    return total


def test_without_the_held_fixed_columns_each_gap_is_the_test_on_one_pocket_counted_by_hand(route):
    """The firm's lean pre-spec (26 Sep 2026): each input reported with and without the columns held fixed. Without
    them, every loan of a range sits in one pocket: the same regression, on counts made here from the extract."""
    t = route["t"]
    for side, (lo, hi) in ((t.development_plain, ("2022-01-01", "2023-12-31")),
                           (t.holdout_plain, ("2024-01-01", "2024-12-31"))):
        pocket = _one_pocket(route["x"], lo, hi)
        want = kgroups.conditional_fit([pocket], t.ref, 6)
        assert len(side.pockets) == 1 and side.loans == [int(n) for n in pocket.loans]
        for k in range(6):
            if k != t.ref:
                assert side.fit.odds[k] == pytest.approx(want.odds[k], rel=1e-9)
                assert side.fit.p[k] == pytest.approx(want.p[k], rel=1e-9)
    # and it differs from the test with FICO and CHANNEL held fixed, which is the point of reporting both
    assert [round(o, 6) for o in t.holdout_plain.fit.odds if o] != [round(o, 6) for o in t.holdout.fit.odds if o]


def test_each_candidate_row_is_found_confirmed_and_confirmed_with_the_columns_held_fixed(route, calculated):
    t = route["t"]
    rows = _table(calculated)
    ref = t.groups[t.ref]
    assert sorted(rows) == sorted(f"{g} vs {ref}" for k, g in enumerate(t.groups) if k != t.ref)
    for k, g in enumerate(t.groups):
        if k == t.ref:
            continue
        x = rows[f"{g} vs {ref}"]
        assert x[confirm_tab.N_CAND] == t.column
        for (odds_col, p_col), side in (((confirm_tab.N_FG, confirm_tab.N_FP), t.development_plain),
                                        ((confirm_tab.N_CG, confirm_tab.N_CP), t.holdout_plain),
                                        ((confirm_tab.N_HG, confirm_tab.N_HP), t.holdout)):
            assert x[odds_col] == pytest.approx(side.fit.odds[k], rel=1e-9)
            # the p-value after the allowance for testing many at once (OC-49): here the one candidate's own groups
            assert x[p_col] == pytest.approx(side.allowed[k], rel=1e-9)
            assert side.allowed[k] >= side.fit.p[k] * (1 - 1e-12)


def _verdicts_by_hand(t, bar: float) -> dict:
    """Holds up? and Still holds?: significant on the held-back loans after the allowance for testing many at once
    (Benjamini-Hochberg across the groups, worked out here: statistics.md A2), on the side of 1 the found loans
    showed."""
    allowed = {}
    for side in (t.holdout_plain, t.holdout):
        ks = [k for k in range(len(t.groups)) if k != t.ref and side.fit.p[k] is not None]
        order = sorted(ks, key=lambda k: side.fit.p[k])
        m = len(order)
        scaled = [side.fit.p[k] * m / (j + 1) for j, k in enumerate(order)]
        for j, k in enumerate(order):
            allowed[(id(side), k)] = min(1.0, min(scaled[j:]))
    out = {}
    for k, g in enumerate(t.groups):
        if k == t.ref:
            continue
        found = t.development_plain.fit.odds[k]

        def holds(side):
            o, p = side.fit.odds[k], allowed.get((id(side), k))
            return "Yes" if o is not None and p is not None and p < bar and (o > 1) == (found > 1) else "No"
        out[f"{g} vs {t.groups[t.ref]}"] = (holds(t.holdout_plain), holds(t.holdout))
    return out


def test_holds_up_and_still_holds_are_worked_out_live_at_the_confidence_on_control(route, calculated, tmp_path):
    t = route["t"]
    want = _verdicts_by_hand(t, 0.05)
    got = {k: (x[confirm_tab.N_HOLDS], x[confirm_tab.N_STILL]) for k, x in _table(calculated).items()}
    assert got == want
    assert ("Yes", "Yes") in want.values()               # the planted cliffs hold up, FICO and CHANNEL held fixed
    words = {k: x[confirm_tab.N_WORDS] for k, x in _table(calculated).items()}
    said = {("Yes", "Yes"): "Holds up, and not just FICO and CHANNEL", ("Yes", "No"): "Holds up, but it was mostly "
            "FICO and CHANNEL", ("No", "Yes"): "Holds up only with FICO and CHANNEL held fixed",
            ("No", "No"): "Didn't hold up on held-back loans"}
    assert words == {k: said[v] for k, v in want.items()}
    # at 99% sure a verdict between the two bars turns, with no Run
    wb = load_workbook(route["b"])
    ws = wb[control.SHEET]
    ws.cell(row=control.row_of(ws, "confidence"), column=control.CHOOSE_COL).value = "99% sure"
    copy = tmp_path / "at-99.xlsx"
    wb.save(copy)
    calc = recalc(copy, tmp_path / "calc")[confirm_tab.SHEET]
    strict = _verdicts_by_hand(t, 0.01)
    assert strict != want
    assert {k: (x[confirm_tab.N_HOLDS], x[confirm_tab.N_STILL]) for k, x in _table(calc).items()} == strict
    # and at 90% a verdict between 5% and 10% turns the other way (since the allowance, OC-49, it is a Still holds?
    # on this book: the Holds up? p-value between the bars is on the other side of 1 from the found one)
    ws.cell(row=control.row_of(ws, "confidence"), column=control.CHOOSE_COL).value = "90% sure"
    copy = tmp_path / "at-90.xlsx"
    wb.save(copy)
    calc = recalc(copy, tmp_path / "calc90")[confirm_tab.SHEET]
    loose = _verdicts_by_hand(t, 0.10)
    assert [k for k in want if loose[k] != want[k]], "no verdict between 5% and 10% to show the change"
    assert {k: (x[confirm_tab.N_HOLDS], x[confirm_tab.N_STILL]) for k, x in _table(calc).items()} == loose


def test_excess_is_each_groups_charge_offs_above_its_share_scaled_to_the_book_and_material_follows_control(
        route, calculated, tmp_path):
    t = route["t"]
    mine = _by_hand(route["x"], "2024-01-01", "2024-12-31")
    hold_gco = math.fsum(g for _, _, g in mine if g is not None)
    scale = _book_gco(route["x"]) / hold_gco
    rows = _table(calculated)
    want = {}
    for k, g in enumerate(t.groups):
        if k == t.ref:
            continue
        inside = [m for m in mine if m[0] == k]
        excess = (math.fsum(x for _, _, x in inside if x is not None) - len(inside) / len(mine) * hold_gco) * scale
        key = f"{g} vs {t.groups[t.ref]}"
        assert rows[key][confirm_tab.N_EX] == pytest.approx(excess, rel=1e-9)
        want[key] = excess
    line = 0.01 * _book_gco(route["x"])                   # tests/test_book.py answers 1% of the book's losses
    assert {k: x[confirm_tab.N_MAT] for k, x in rows.items()} == \
        {k: "Yes" if e > 0 and e >= line else "No" for k, e in want.items()}
    assert "Yes" in {x[confirm_tab.N_MAT] for x in rows.values()} and "No" in {x[confirm_tab.N_MAT] for x in rows.values()}
    # a higher line, live
    wb = load_workbook(route["b"])
    ws = wb[control.SHEET]
    ws.cell(row=control.row_of(ws, "materiality"), column=control.CHOOSE_COL).value = "10% of the book's total losses"
    copy = tmp_path / "at-10.xlsx"
    wb.save(copy)
    calc = recalc(copy, tmp_path / "calc")[confirm_tab.SHEET]
    assert {k: x[confirm_tab.N_MAT] for k, x in _table(calc).items()} == \
        {k: "Yes" if e > 0 and e >= 10 * line else "No" for k, e in want.items()}


def test_a_saved_shortlist_hides_the_found_columns_and_record_names_the_file(route):
    wb = load_workbook(route["b"])
    ws = wb[confirm_tab.SHEET]
    hidden = {c for c in range(confirm_tab.N_CAND, confirm_tab.N_WORDS + 1)
              if ws.column_dimensions[get_column_letter(c)].hidden}
    assert hidden == {confirm_tab.N_FG, confirm_tab.N_FP}
    assert _check(route["b"])["Pre-spec"].endswith("prespec.yaml")
    # nothing the tiles or the tests in full show sits in a hidden column
    heads = [ws.cell(row=r, column=c).value for r in range(1, ws.max_row + 1) for c in sorted(hidden)
             if ws.cell(row=r, column=c).value not in (None, "")]
    assert [h for h in heads if h not in ("Gap", confirm_tab.P_ALLOWED, f"Found · {route['t'].development.n:,} loans")
            and not isinstance(h, (int, float))] == []


def test_the_chart_draws_found_confirmed_and_held_fixed_from_the_table_with_the_worse_line(route):
    ws = load_workbook(route["b"])[confirm_tab.SHEET]
    (chart,) = ws._charts
    rows = _table(ws)
    first, last = min(x["row"] for x in rows.values()), max(x["row"] for x in rows.values())
    col = lambda c: f"'{confirm_tab.SHEET}'!${get_column_letter(c)}${first}:${get_column_letter(c)}${last}"  # noqa
    bars = [s.val.numRef.f for s in chart.series]
    # the Found bars are left off with the Found columns (a saved shortlist)
    assert bars == [col(confirm_tab.N_CG), col(confirm_tab.N_HG)]
    assert [s.graphicalProperties.solidFill.srgbClr for s in chart.series] == [house.INK, house.KEY_RED]
    (line,) = chart._charts[1].series
    assert line.val.numRef.f == col(confirm_tab.LINE_COL) and line.graphicalProperties.line.dashStyle == "dash"
    assert {ws.cell(row=r, column=confirm_tab.LINE_COL).value for r in range(first, last + 1)} == {"=worse_at"}
    assert chart.visible_cells_only is False


def _dated(rows_spec):
    """rows_spec: (date, chan, ratio, bad) per loan."""
    return [{"ID": f"L{i}", "SCORE": 600, "CHAN": ch, "BAL": 100, "BAD": bad, "GCO": 50 * bad, "RANR": 3, "R": v,
             "ORIG": d} for i, (d, ch, v, bad) in enumerate(rows_spec)]


SPEC = {"prespec": 1, "written": "2026-09-26", "column": "R", "bins": [1.0, 2.0], "reference": 1,
        "strata": ["CHAN"], "confidence": 0.95, "holdout": {"from": "2024-01-01", "to": "2024-12-31"},
        "development": {"from": "2022-01-01", "to": "2023-12-31"}}


def _run(rows, **spec):
    ps = prespec.parse({**SPEC, **spec})
    res = engine.run(cube(origination_date="ORIG"), table(_dated(rows)))
    ps = prespec.named(ps, *confirmatory.column_range(res, ps.column))
    return res, ps, confirmatory.run_test(res, ps)


def _mixed_book():
    out = []
    for y, n in (("2021-06-01", 7), ("2022-03-01", 40), ("2023-11-30", 40), ("2024-01-01", 30), ("2024-12-31", 30),
                 ("2025-01-01", 9)):
        for i in range(n):
            out.append((y, "AB"[i % 2], (0.5, 1.5, 2.5)[i % 3], 1 if i % 5 == 0 or (i % 3 == 2 and i % 4 == 0) else 0))
    return out


def test_the_holdout_takes_no_development_loan_and_the_range_held_to_is_the_pre_specs():
    rows = _mixed_book()
    res, ps, t = _run(rows)
    assert t.holdout.n == 60 and t.development.n == 80
    assert t.left_out == {confirmatory.OUTSIDE: 16}
    assert t.holdout.range == ps.holdout and t.development.range == ps.development
    used = confirmatory.in_use(res, ps, t)
    assert used["holdout"] is ps.holdout
    # the extract runs from 2021 to 2025, and the test held itself to the holdout: no deviation (final check F5)
    assert prespec.deviations(ps, used, where="in this run") == []


def test_the_reference_group_is_the_pre_specs():
    rows = _mixed_book()
    for ref in (0, 1, 2):
        res, ps, t = _run(rows, reference=ref)
        assert t.ref == ref and confirmatory.in_use(res, ps, t)["reference"] == t.groups[ref]
        pockets = {}
        for d, ch, v, bad in rows:
            if "2024-01-01" <= d <= "2024-12-31":
                c = pockets.setdefault(ch, [[0] * 3, [0] * 3])
                g = bisect.bisect_right([1.0, 2.0], v)
                c[0][g] += 1
                c[1][g] += bad
        want = kgroups.conditional_fit([kgroups.Pocket(*c) for c in pockets.values()], ref)
        assert t.holdout.fit.odds[ref] == 1.0
        assert t.holdout.fit.beta == pytest.approx(want.beta, abs=1e-12)


def test_the_run_fits_conditionally_so_thin_pockets_are_not_biased():
    """Matched pairs through the whole route: every pocket (CHAN, the pre-spec's stratum) is one loan in the
    group and one in the reference, one of them bad. The conditional odds ratio is 60 / 30 = 2.0; a constant per
    pocket would give 4.0 (B5, "Breaks when")."""
    rows = []
    for i in range(90):
        worse = i < 60
        rows.append(("2024-05-01", f"P{i:03d}", 2.5, 1 if worse else 0))
        rows.append(("2024-05-01", f"P{i:03d}", 1.5, 0 if worse else 1))
    rows.append(("2023-05-01", "P000", 0.5, 1))           # development needs a loan; its own business
    res, ps, t = _run(rows, reference=1)
    assert t.holdout.fit.method == kgroups.CONDITIONAL and t.method == kgroups.CONDITIONAL
    assert t.holdout.fit.odds[2] == pytest.approx(2.0, rel=1e-9)
    assert len(t.holdout.pockets) == 90


def test_a_pocket_below_every_floor_still_counts_in_the_pooled_test():
    """The goal's rule: fewest loans (here 1,000) governs per-pocket readings only."""
    rows = _mixed_book() + [("2024-06-01", "Z", 0.5, 1), ("2024-06-01", "Z", 1.5, 0), ("2024-06-01", "Z", 2.5, 0)]
    ps = prespec.parse(SPEC)
    bench = {**cube().raw["benchmark"], "min_units": 1000, "min_events": 50}
    res = engine.run(cube(origination_date="ORIG", benchmark=bench), table(_dated(rows)))
    t = confirmatory.run_test(res, prespec.named(ps, *confirmatory.column_range(res, "R")))
    assert len(t.holdout.pockets) == 3 and t.holdout.association.pockets == 3


@needs_git
def test_a_stratum_the_run_does_not_cut_by_is_refused_by_the_pre_spec_cell(tmp_path, monkeypatch):
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    spec = tmp_path / "spec.yaml"
    spec.write_text(EXAMPLE.read_text(encoding="utf-8").replace("strata: [FICO, CHANNEL]",
                                                                "strata: [FICO, ASSET_CLASS]"), encoding="utf-8")
    x, b = _set_up(tmp_path / "book", 1500, spec)
    ran = book.run(b)
    assert not ran.ok
    cell = f"Control!C{control.read_prespec(load_workbook(b)[control.SHEET])[1].split('!C')[1]}"
    want = (f"{cell}: the pre-spec cuts the pockets by ASSET_CLASS, and the launcher doesn't cut by it. Tick it "
            f"under Choose tests (a band or a segment), then press Next. Or fix the pre-spec.")
    assert any(want in line for line in ran.lines), ran.lines
    _choose(b, segments=("CHANNEL", "ASSET_CLASS"))
    assert book.run(b).ok


@needs_git
def test_a_holdout_with_no_loans_in_the_extract_says_so_and_tests_nothing(tmp_path, monkeypatch):
    """A pre-spec whose holdout is still to come: development is tested, and the holdout's lines say there is
    nothing to test rather than "no difference shows"."""
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    spec = tmp_path / "spec.yaml"
    spec.write_text(EXAMPLE.read_text(encoding="utf-8").replace("holdout: {from: 2024-01-01, to: 2024-12-31}",
                                                                "holdout: {from: 2030-01-01, to: 2030-12-31}"),
                    encoding="utf-8")
    x, b = _set_up(tmp_path / "book", 1500, spec)
    ran = book.run(b)
    assert ran.ok, ran.lines
    ws = recalc(b, tmp_path / "calc")[confirm_tab.SHEET]
    said = _texts(ws)
    lines = [s for s in said if s.startswith("On the holdout, nothing to test: no loan in this test was made "
                                             "2030-01-01 to 2030-12-31.")]
    assert len(lines) == 3                                     # the tests, the odds ratios, the concentration
    assert any(s.startswith("On development,") for s in said)
    assert not [s for s in said if s.startswith("#") or "Err:" in s]
    chk = _check(b)
    assert chk["Holdout"] == "None of this extract's loans were made in the holdout (2030-01-01 to 2030-12-31)."
    assert chk["Differs from the pre-spec"] == "nowhere: this run used what it says"


def test_without_readable_dates_the_test_says_why_and_the_run_goes_on():
    rows = _mixed_book()
    ps = prespec.parse(SPEC)
    res = engine.run(cube(), table(_dated(rows)))                # no column marked Origination date
    t = confirmatory.run_test(res, ps)
    assert t.problem.startswith("the loans can't be split into development and holdout: no column is marked")
    assert confirmatory.in_use(res, ps, t)["reference"] != ps.reference     # it falls back to what the tabs used
