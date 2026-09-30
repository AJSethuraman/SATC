"""The firm, 30 Sep 2026, three things at once:

1. Columns shows every column. "for some reason in my testing it is hiding random rows from the columns tab which
   makes it hard to make sure it's right". The columns this Run doesn't use are greyed and say so under Check
   first; nothing is asked about them, as before.
2. The mouse wheel after a page change: `_tkinter.TclError: invalid command name ".!frame2.!frame3.!frame100.
   !frame3.!canvas"` (at the bank). The wheel was bound to Choose tests' own table and outlived it.
3. openpyxl's "Data Validation extension is not supported and will be removed" (and the Conditional Formatting
   one) on every read of a workbook Excel had saved; and "it seemed pocketbook hanging and it didn't before": a
   live progress line, the stage and the seconds, while Set up and Run work."""

from __future__ import annotations

import io
import re
import threading
import time
import warnings
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from pocketbook import book, choices as ch, control, excel_lists, house, ingest, launcher, memory, perm, synth
from test_book import _answer
from test_firm_answers_2026_09_29 import _bank_file

PICKED = ch.Choices(run_kind=ch.BLEED, bands=("ORIG_BAL",), segments=("CHANNEL",), outcome="BAD_FLAG")
CF_EXT = "{78C0D931-6437-407D-A8EE-F0AAD7539E65}"       # Excel's name for the conditional formats' extension


def _hex(color) -> str:
    return (color.rgb or "")[-6:] if color is not None and isinstance(color.rgb, str) else ""


def _rows(ws) -> dict:
    return {r[book.C_NAME - 1].value: r for r in book.table_rows(ws) if r[book.C_NAME - 1].value}


# ---- 1. Columns: every column in sight, the ones this Run skips greyed


def test_columns_shows_every_column_and_greys_the_ones_this_run_does_not_use(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = synth.write_extract(tmp_path, n=1500)
    out = book.set_up(x, choices=PICKED)
    ws = load_workbook(out.book)["Columns"]
    rows = _rows(ws)
    assert not [r[0].row for r in book.table_rows(ws) if ws.row_dimensions[r[0].row].hidden]   # nothing hidden
    grey = {n for n, r in rows.items() if str(r[book.C_LOOK - 1].value or "").startswith(book.NOT_USED)}
    assert {"FICO", "ASSET_CLASS", "REV_DEBT"} <= grey
    assert not grey & {"LOAN_NBR", "ORIG_BAL", "CHANNEL", "BAD_FLAG", "GCO_AMT", "RANR_AMT"}
    for name, r in rows.items():
        cells = r[book.C_NAME - 1:book.C_DEFINE]
        if name in grey:
            assert {_hex(c.fill.fgColor) for c in cells} == {house.CANVAS}, name
            assert {_hex(c.font.color) for c in cells} == {house.SLATE}, name
        else:
            assert house.CANVAS not in {_hex(c.fill.fgColor) for c in cells}, name
            assert _hex(r[book.C_NAME - 1].font.color) != house.SLATE, name
    # nothing asked about a grey column: FICO's -9999 has no Treat as row
    keys = [r[book.C_QKEY - 1].value for r in ws.iter_rows(min_row=book.COL_FIRST) if len(r) >= book.C_QKEY]
    assert not [k for k in keys if k and str(k).startswith("FICO|")]
    assert book.UNUSED_NOTE.format(len(grey)) in ws["D3"].value and "hidden" not in ws["D3"].value
    # "Not used this Run." is not something still to do: Check first's shading leaves a grey row alone
    look = book._col(book.C_LOOK)
    rules = [f for rng in ws.conditional_formatting if str(rng.sqref).startswith(look) for rule in rng.rules
             for f in rule.formula]
    assert rules and all(book.NOT_USED in f and "LEFT(" in f for f in rules)
    _answer(out.book)
    _, problems, _ = book.read_book(out.book)
    assert not problems, problems


def test_a_grey_column_picked_later_is_shown_plain_and_asked_about(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = synth.write_extract(tmp_path, n=1500)
    book.set_up(x, choices=PICKED)
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("ORIG_BAL", "FICO"), segments=("CHANNEL",),
                                            outcome="BAD_FLAG"))
    ws = load_workbook(out.book)["Columns"]
    fico = _rows(ws)["FICO"]
    assert not str(fico[book.C_LOOK - 1].value or "").startswith(book.NOT_USED)
    assert _hex(fico[book.C_NAME - 1].fill.fgColor) != house.CANVAS
    keys = [r[book.C_QKEY - 1].value for r in ws.iter_rows(min_row=book.COL_FIRST) if len(r) >= book.C_QKEY]
    assert [k for k in keys if k and str(k).startswith("FICO|")]


# ---- 3a. openpyxl's warnings about Excel's extension blocks, on every read of a workbook Excel saved


def _with_extensions(path, sheets) -> None:
    """Add, to each named sheet, the two extension blocks Excel writes (a dropdown and a conditional format kept
    in its 2010 block), the way a workbook comes back from Excel."""
    src = path.read_bytes()
    with zipfile.ZipFile(io.BytesIO(src)) as zin:
        parts = excel_lists._sheet_files(zin)
        wanted = {parts[s] for s in sheets}
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
            for item in zin.infolist():
                data = zin.read(item.filename)
                if item.filename in wanted:
                    block = (f'<extLst><ext uri="{excel_lists.DV_EXT}" xmlns:x14="{excel_lists.X14}">'
                             f'<x14:dataValidations count="1" xmlns:xm="{excel_lists.XM}"><x14:dataValidation '
                             f'type="list" allowBlank="1"><x14:formula1><xm:f>\'_options\'!$A$1:$A$2</xm:f>'
                             f'</x14:formula1><xm:sqref>Z99</xm:sqref></x14:dataValidation></x14:dataValidations>'
                             f'</ext><ext uri="{CF_EXT}" xmlns:x14="{excel_lists.X14}"><x14:conditionalFormattings/>'
                             f'</ext></extLst>')
                    xml = data.decode("utf-8")
                    assert "<extLst>" not in xml
                    data = xml.replace("</worksheet>", block + "</worksheet>").encode("utf-8")
                z.writestr(item, data)


def _said(fn) -> list[str]:
    with warnings.catch_warnings(record=True) as got:
        warnings.simplefilter("always")
        fn()
    return [str(w.message) for w in got if "extension is not supported" in str(w.message)]


def test_reading_a_workbook_excel_saved_says_nothing_about_its_extension_blocks(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = synth.write_extract(tmp_path, n=800)
    b = book.set_up(x, choices=PICKED).book
    _answer(b)
    _with_extensions(b, ["Control", "Columns", "Start here"])
    # openpyxl alone says both: what the firm saw
    assert {m.split(" extension")[0] for m in _said(lambda: load_workbook(b))} == {
        "Data Validation", "Conditional Formatting"}
    assert _said(lambda: book._answers(book._earlier(b))) == []
    assert _said(lambda: book.read_extract(x)) == []                 # reads the launcher's picks from Control
    assert _said(lambda: control.read_control(b)) == []
    assert _said(lambda: excel_lists.load(b)) == []                  # Set up and Run's own load
    # the extract, an .xlsx Excel saved
    xl = tmp_path / "extract.xlsx"
    wb = Workbook()
    wb.active.append(["LOAN_NBR", "BAL"])
    wb.active.append([1, 100])
    wb.active.title = "loans"
    wb.save(xl)
    _with_extensions(xl, ["loans"])
    assert _said(lambda: ingest.read_table(xl)) == []
    assert ingest.read_table(xl).columns == ["LOAN_NBR", "BAL"]
    # the memory's review sheet
    rv = tmp_path / "review.xlsx"
    wb = Workbook()
    wb.active.title = "Learned"
    wb.save(rv)
    _with_extensions(rv, ["Learned"])
    assert _said(lambda: memory.apply_review(rv, tmp_path / "memory.yaml")) == []
    # any other warning still comes through
    with pytest.warns(UserWarning, match="something else"):
        with warnings.catch_warnings():
            excel_lists.quiet(b)
            warnings.warn("something else", UserWarning)


# ---- 3b. conditional formats Excel may move into its extension block


def _cross_sheet(ws) -> list[str]:
    return [f for rng in ws.conditional_formatting for r in rng.rules for f in (r.formula or []) if "!" in f]


def test_no_tab_a_run_keeps_has_a_shading_rule_excel_would_move_out_of_reach(tmp_path, monkeypatch):
    """Excel keeps a conditional format in its 2010 extension block when its formula reads another sheet (as it
    does a dropdown whose list is on another sheet), and openpyxl drops that block. The tabs a Run keeps as they
    are (Control, Columns, Look) have no such rule, so Excel has nothing of theirs to move; the one tab that has
    them, Grids, is deleted and drawn again by every Run."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    x = synth.write_extract(tmp_path, n=1500)
    b = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO", "ORIG_BAL"), segments=("CHANNEL",),
                                          outcome="BAD_FLAG")).book
    _answer(b)
    assert book.run(b, x).ok
    wb = load_workbook(b)
    for t in ("Control", "Columns", "Look", "Start here"):
        assert _cross_sheet(wb[t]) == [], t
    before = _cross_sheet(wb["Grids"])
    assert before                                   # the one tab with such rules
    # Excel moves them: gone from where openpyxl reads them, into the extension block
    src = b.read_bytes()
    with zipfile.ZipFile(io.BytesIO(src)) as zin, zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED) as z:
        part = excel_lists._sheet_files(zin)["Grids"]
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == part:
                xml = re.sub(r"<conditionalFormatting [^>]*>(?:(?!</conditionalFormatting>).)*!(?:(?!"
                             r"</conditionalFormatting>).)*</conditionalFormatting>", "", data.decode("utf-8"),
                             flags=re.S)
                xml = xml.replace("</worksheet>", f'<extLst><ext uri="{CF_EXT}" xmlns:x14="{excel_lists.X14}">'
                                  f'<x14:conditionalFormattings/></ext></extLst></worksheet>')
                data = xml.encode("utf-8")
            z.writestr(item, data)
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        assert _cross_sheet(load_workbook(b)["Grids"]) == []
    assert book.run(b, x).ok
    assert sorted(_cross_sheet(load_workbook(b)["Grids"])) == sorted(before)


# ---- 3c. the progress line: the stage, and the seconds since the button was pressed


def test_the_progress_line_says_the_stage_and_the_time_taken():
    assert launcher.elapsed_words(0) == "0 s"
    assert launcher.elapsed_words(12.7) == "12 s"
    assert launcher.elapsed_words(60) == "1 min 0 s"
    assert launcher.elapsed_words(100) == "1 min 40 s"
    f = launcher.Flow(gate=launcher.AddOns())
    assert f.progress_line(now=5) == ""                          # nothing running, nothing said
    f.busy, f.began = "Running", 1000.0
    assert f.progress_line(now=1012) == "Running… 12 s"          # before the first stage is said
    f.progress("Running the shuffle test")
    assert f.progress_line(now=1100) == "Running the shuffle test… 1 min 40 s"


def test_set_up_and_run_say_each_stage_as_it_starts(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    x = synth.write_extract(tmp_path, n=1500)
    said: list[str] = []
    b = book.set_up(x, choices=PICKED, progress=said.append).book
    assert said == ["Reading the extract", "Looking at each column", "Drawing Look",
                    "Working out the suggested settings", "Saving the workbook"]
    _answer(b)
    said.clear()
    assert book.run(b, x, progress=said.append).ok
    at = [said.index(s) for s in ("Reading the workbook", "Reading the extract", "Cutting bands",
                                  "Running the shuffle test", "Writing the workbook", "Saving the workbook")]
    assert at == sorted(at), said
    assert book.run(b, x).ok                                      # and without anyone listening


# ---- on the window (needs a display: xvfb-run on Linux)


def _pump(root, until, seconds=30.0):
    end = time.monotonic() + seconds
    while time.monotonic() < end:
        root.update()
        if until():
            return True
        time.sleep(0.02)
    return False


def test_the_wheel_scrolls_the_table_on_screen_and_never_a_page_that_has_gone(tmp_path, monkeypatch):
    """The bank's traceback, and the progress line while Next writes the workbook: the window keeps drawing."""
    from test_deps import _window
    root = _window()
    errors: list[str] = []
    root.report_callback_exception = lambda kind, exc, tb: errors.append(f"{kind.__name__}: {exc}")
    try:
        from pocketbook import deps
        monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
        monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
        monkeypatch.setattr(deps, "missing", lambda: [])
        monkeypatch.setattr(deps, "missing_optional", lambda: [])
        w = launcher.build(root)
        root.geometry("900x520")                 # a small window, so the table scrolls
        root.deiconify()
        flow = w["flow"]
        x = _bank_file(tmp_path)
        w["extract"].set(str(x))
        flow.pick(str(x))
        flow.set_up()
        flow.set_mode("new")
        w["render"]()
        root.update()
        t = w["table"]
        assert t.yview()[0] == 0
        root.event_generate("<MouseWheel>", delta=-120)               # Windows' wheel, one notch down
        root.update()
        assert t.yview()[0] > 0
        down = t.yview()
        root.event_generate("<Button-4>")                              # X11's wheel, one notch up
        root.update()
        assert t.yview()[0] < down[0]
        # Next, held part way: the progress line says the stage and counts the seconds
        go, real = threading.Event(), book.set_up

        def held(*a, progress=book._quiet, **k):
            progress("Working out the suggested settings")
            go.wait(20)
            return real(*a, progress=progress, **k)
        monkeypatch.setattr(book, "set_up", held)
        flow.pick_outcome("EVER GCO")
        flow.answer_outcome(True)
        flow.click("FICO", "b")
        w["render"]()
        root.update()
        w["next"].invoke()
        assert _pump(root, lambda: "progress" in w and w["progress"].cget("text").startswith(
            "Working out the suggested settings… "))
        flow.began -= 100
        assert _pump(root, lambda: w["progress"].cget("text") == "Working out the suggested settings… 1 min 40 s"
                     or w["progress"].cget("text").startswith("Working out the suggested settings… 1 min 4"))
        root.event_generate("<MouseWheel>", delta=-120)               # the table still scrolls while it works
        go.set()
        assert _pump(root, lambda: not flow.busy, 120)
        root.update()
        assert flow.page == "answer", flow.message
        assert not t.winfo_exists() and "progress" not in w
        # the page that had the table is gone: the wheel does nothing, and says nothing
        for seq in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
            root.event_generate(seq, **({"delta": -120} if seq == "<MouseWheel>" else {}))
        root.update()
        assert errors == []
    finally:
        root.destroy()
