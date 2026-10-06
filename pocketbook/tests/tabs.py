"""Reading the result tabs as the analyst sees them (the redesign, phase 3: Pockets, RANR vs GCOs, Grids, Split).

Each tab shows what its dropdowns pick, by formulas, so a test reads it calculated (tests/recalc.py), and to see
another measure, grid or view it sets the dropdown on a copy and calculates again, as a person would. The fill a
conditional format gives a cell is worked out from the rule as written and the calculated values, for the one
shape of rule each tab writes; anything else fails rather than guess."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, range_boundaries

from pocketbook import results as rs
from recalc import recalc

POCKET_KEYS = {rs.K_NUM: "num", rs.K_BAND: "band", rs.K_SEG: "seg", rs.K_HALF: "half", rs.K_LOANS: "loans",
               rs.K_THIS: "this", rs.K_REST: "rest", rs.K_GAP: "gap", rs.K_EX: "excess", rs.K_WORSE: "worse",
               rs.K_P: "chance", rs.K_MAT: "material", rs.K_CAUGHT: "caught", rs.K_HOLDS: "holds"}
PCK_KEYS = {rs.C_BAND: "band", rs.C_SEG: "seg", rs.C_LOANS: "loans", rs.C_BOOK: "booked", rs.C_GCO: "gco",
            rs.C_RANR: "ranr", rs.C_RATE: "ranr_rate", rs.C_AVG: "avg_line", rs.C_AVGX: "avg_x",
            rs.C_PAID: "paid", rs.C_PAID_D: "paid_d",
            rs.C_COST: "cost", rs.C_COST_D: "cost_d", rs.C_KEPT: "kept", rs.C_KEPT_D: "kept_d",
            rs.C_TOG: "together"}


def blank(v):
    return None if v == "" else v


def word(v):
    """A verdict without its borderline words ("Yes · borderline (p 0.048)" is Yes): what every test of the
    verdict itself compares. The words are the Borderline tests' (test_firm_answers_2026_09_29), read from `*_said`."""
    return v.split(" · borderline (", 1)[0] if isinstance(v, str) else v


def calculated(path, name: str):
    """A tab as LibreOffice calculates it, the tab as written riding along as .formulas."""
    ws = recalc(path, Path(path).parent / f"rc-{Path(path).stem}")[name]
    ws.formulas = load_workbook(path)[name]
    return ws


def dropdown(ws, label: str):
    """The dropdown cell under its label ("MEASURE", "GRID", ...)."""
    for row in ws.iter_rows(max_row=80):
        for c in row:
            if c.value == label.upper():
                return ws.cell(row=c.row + 1, column=c.column)
    raise KeyError(label)


def choose(path, out, tab: str, **picks) -> Path:
    """A copy of the workbook with a tab's dropdowns set: choose(b, out, "Pockets", measure="GCOs ($)")."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(path, out)
    wb = load_workbook(out)
    for label, value in picks.items():
        dropdown(wb[tab], label).value = value
    wb.save(out)
    return out


def options(wb, tab: str, label: str) -> list:
    """What a dropdown offers: the range its data validation lists, read off _choices."""
    cell = dropdown(wb[tab], label)
    for dv in wb[tab].data_validations.dataValidation:
        if cell.coordinate in str(dv.sqref):
            sheet, rng = dv.formula1.lstrip("=").split("!")
            c0, r0, _, r1 = range_boundaries(rng.replace("$", ""))
            ws = wb[sheet.strip("'")]
            return [ws.cell(row=r, column=c0).value for r in range(r0, r1 + 1)]
    raise KeyError(label)


def header_row(ws, col: int, text: str, start: int = 1) -> int:
    for r in range(start, ws.max_row + 1):
        if ws.cell(row=r, column=col).value == text:
            return r
    raise KeyError(text)


def pockets(ws) -> list[dict]:
    """Pockets' rows as shown (calculated), by the column's key, each with its row."""
    h = header_row(ws, rs.K_NUM, "#")
    out = []
    for r in range(h + 1, ws.max_row + 1):
        if ws.cell(row=r, column=rs.K_NUM).value in (None, ""):
            break
        x = {k: blank(ws.cell(row=r, column=c).value) for c, k in POCKET_KEYS.items()}
        x["row"] = r
        x["worse_said"], x["worse"] = x["worse"], word(x["worse"])
        out.append(x)
    return out


def heads(ws, row: int, first: int = 2, last: int = 20) -> list:
    return [ws.cell(row=row, column=c).value for c in range(first, last + 1)]


RULE = re.compile(r'^AND\(\$([A-Z]+)(\d+)<>"",AND\(\$([A-Z]+)(\d+)=0,\$([A-Z]+)(\d+)="([^"]+)"\)\)$')


def pck(ws) -> list[dict]:
    """RANR vs GCOs' rows as shown (calculated), with the fill each pair's rule gives it: the rules are
    conditional formats over each side's flag, so they are worked out here from the calculated flags."""
    h = header_row(ws, rs.C_TOG, "Together")
    out = []
    for r in range(h + 1, ws.max_row + 1):
        if ws.cell(row=r, column=rs.C_BAND).value in (None, ""):
            break
        x = {k: blank(ws.cell(row=r, column=c).value) for c, k in PCK_KEYS.items()}
        x["row"] = r
        x["together_said"], x["together"] = x["together"], word(x["together"])
        for k, c in (("c_fill", rs.C_PAID), ("g_fill", rs.C_COST), ("r_fill", rs.C_KEPT)):
            x[k] = cf_fill(ws, r, c) if hasattr(ws, "formulas") else None      # the rules ride on .formulas
        x["flags"] = {k: blank(ws.cell(row=r, column=c).value) for k, c in (("c", rs.C_H_FC), ("g", rs.C_H_FG),
                                                                           ("r", rs.C_H_FR))}
        out.append(x)
    return out


def pck_all(path, out) -> list[dict]:
    """RANR vs GCOs' rows for every grid its dropdown offers, each grid picked in turn and calculated, as a
    person would page through them. Each row says its grid."""
    rows = []
    for i, grid in enumerate(options(load_workbook(path), rs.PCK, "Grid")):
        ws = calculated(choose(path, Path(out) / f"grid{i}.xlsx", rs.PCK, grid=grid), rs.PCK)
        rows += [{**x, "grid": grid} for x in pck(ws)]
    return rows


def cf_fill(ws, r: int, col: int) -> str | None:
    """The fill RANR vs GCOs' rules give a cell: the first rule that holds (see results.cf)."""
    for rng in ws.formulas.conditional_formatting:
        for bounds in str(rng.sqref).split():
            c0, r0, c1, r1 = range_boundaries(bounds)
            if not (c0 <= col <= c1 and r0 <= r <= r1):
                continue
            for rule in sorted(rng.rules, key=lambda x: x.priority):
                if rule.dxf is None or rule.dxf.fill is None:
                    continue
                m = RULE.match(rule.formula[0])
                assert m, rule.formula[0]
                lc, lr, uc, ur, fc, fr, word = m.groups()
                dr = r - r0
                val = lambda c, rr: ws.cell(row=int(rr) + dr, column=column_index_from_string(c)).value  # noqa
                if val(lc, lr) not in (None, "") and val(uc, ur) == 0 and val(fc, fr) == word:
                    return rule.dxf.fill.fgColor.rgb[-6:]
    return None


def _merged_value(ws, r: int, c: int):
    """A cell's value, or its merged range's first cell's when it sits inside a merge."""
    for rng in ws.merged_cells.ranges:
        if rng.min_row <= r <= rng.max_row and rng.min_col <= c <= rng.max_col:
            return ws.cell(row=rng.min_row, column=rng.min_col).value
    return ws.cell(row=r, column=c).value


def header_of(ws, r: int, c: int) -> tuple[int, dict]:
    """A block's column labels, the block's title at (r, c): the header's last row, and {label: column}. A split
    grid's header on Grids is two rows (the widths change, 29 Sep 2026: G4), the segment merged over its parts and
    then each part; its labels read back as the one label "<segment> · <part>", as before."""
    two = any(m.min_row == r + 1 and m.min_col > c for m in ws.merged_cells.ranges)
    head = r + 2 if two else r + 1
    labels, j = {}, c + 1
    while ws.cell(row=head, column=j).value not in (None, ""):
        low = ws.cell(row=head, column=j).value
        up = _merged_value(ws, r + 1, j) if two else None
        labels[f"{up} · {low}" if up not in (None, "") else low] = j
        j += 1
    return head, labels


def block(ws, title: str, start: int | None = None) -> dict:
    """A block on Grids or Split, by its title (the dark band): {(row label, column label): value}. It is looked
    for under the dropdowns, clear of the method note's labels."""
    if start is None:
        start = min((c.row for row in ws.iter_rows(max_row=80) for c in row if c.value in ("GRID", "MEASURE")),
                    default=1)
    for r in range(start, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.startswith(title):
                out = {}
                head, labels = header_of(ws, r, c)
                cols = {j: label for label, j in labels.items()}
                rr = head + 1
                while ws.cell(row=rr, column=c).value not in (None, ""):
                    for j, name in cols.items():
                        out[(ws.cell(row=rr, column=c).value, name)] = blank(ws.cell(row=rr, column=j).value)
                    rr += 1
                return out
    raise KeyError(title)


# --------------------------------------------------------------------------
# Record (the redesign, phase 4: Check and the Log on one tab)


def _record_ws(src):
    from pocketbook import record as rc
    if isinstance(src, (str, Path)):
        return load_workbook(src)[rc.SHEET]
    if hasattr(src, "sheetnames"):
        return src[rc.SHEET]
    return src


def record_rows(src) -> list[tuple]:
    """Record's rows as (label, value, last Run used, section, changed), every section but Every Run, in the tab's order; a
    list that runs over several rows (a value of several lines) comes back as one value, its lines joined by
    newlines. Settings' value is what is in use now; `last` what the last Run used."""
    from pocketbook import record as rc
    out: list[tuple] = []
    for kind, label, v, last, changed in rc.read(_record_ws(src)):
        if kind == rc.RUNS:
            continue
        if label is None and out and out[-1][3] == kind:
            k, prev, pl, pk, pc = out[-1]
            out[-1] = (k, f"{prev}\n{v}", pl, pk, pc)
        else:
            out.append((label, v, last, kind, changed))
    return out


def record(src) -> dict:
    """Record as Check read: {label: value}, a label on more than one row (Warning) giving a list. A setting's value
    is what the last Run used (Check's "What the last Run used" rows); `settings` has what is in use now."""
    from pocketbook import record as rc
    got: dict = {}
    for label, v, last, kind, _ in record_rows(src):
        got.setdefault(label, []).append(last if kind == rc.SETTINGS else v)
    return {k: v[0] if len(v) == 1 else v for k, v in got.items()}


def settings(src) -> dict:
    """Record's Settings: {setting: (in use now, last Run used, 1 while they differ)}."""
    from pocketbook import record as rc
    return {label: (v, last, changed) for label, v, last, kind, changed in record_rows(src) if kind == rc.SETTINGS}


def runs(src) -> list[str]:
    """Every line of every Run's entry, newest first, as Record's Every Run shows them (the Log tab's column B)."""
    from pocketbook import record as rc
    return [v for kind, _, v, _, _ in rc.read(_record_ws(src)) if kind == rc.RUNS and v]
