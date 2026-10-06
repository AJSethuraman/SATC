"""What the walk of 27 Sep 2026 found on the screens, each held here (docs/walkthrough/2026-09-27/).

The analyst's walk through the launcher's five steps, both run kinds, on the synthetic book. None of these was caught
by the 723 tests before it: each is a thing a screen said, not a number the engine got wrong."""

import re
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook

from pocketbook import book, deps, engine, launcher, results, scout, scout_tab, synth
from test_book import _answer
from test_launcher import _read
import tabs


@pytest.fixture(autouse=True)
def _nothing_opens(monkeypatch):
    monkeypatch.setattr(launcher, "open_file", lambda p: None)


# ---- the window


def test_a_run_that_stopped_is_not_called_an_answer_needed(tmp_path, monkeypatch):
    """Defect 2: a failed tie-out read "1 answer needed before Run", "1 left" on the rail, and "Something went wrong"
    under it: the one check the finished screen shows, failing, looked like a question the analyst had missed."""
    monkeypatch.setattr(launcher.Path, "home", lambda: tmp_path)
    f = _read(tmp_path)
    f.next()
    _answer(launcher.book_for(f.extract))

    def fails(grid, total, measures, n):
        raise engine.TieOutError(f"grid {grid.band} x {grid.dimension}: rows adds up to 7999 across the cells, "
                                 f"but the book says 8000")
    monkeypatch.setattr(engine, "tie_out", fails)
    f.run()
    assert f.screen() == "L3" and not f.answers
    assert f.needs_head()[0] == "Run stopped" and "answer" not in f.needs_head()[0]
    assert f.steps()[3].sub == "Run stopped"
    said = f.needs[0].says
    assert said.startswith("Run stopped: the grids didn't add up to the book, so nothing was written. grid "), said
    assert "Something went wrong" not in said
    # a refusal that is waiting for answers still says so
    monkeypatch.undo()
    monkeypatch.setattr(launcher.Path, "home", lambda: tmp_path)
    g = _read(tmp_path / "second")
    g.next()
    g.run()
    assert g.answers and g.needs_head()[0] == f"{len(g.needs)} answers needed before Run"


def test_the_answers_needed_run_top_to_bottom_as_the_workbook_shows_them(tmp_path):
    """Defect 5: the list read Control C15, C18, C19, C16, C17, C20, C24, C25, then Columns C3."""
    f = _read(tmp_path)
    f.next()
    f.run()
    tabs_seen = [n.sheet for n in f.needs if n.sheet]
    assert tabs_seen == sorted(tabs_seen, key=launcher.TAB_ORDER.index), tabs_seen
    rows = [int(re.sub(r"\D", "", n.cell)) for n in f.needs if n.sheet == "Control"]
    assert len(rows) > 3 and rows == sorted(rows), rows


def test_a_workbook_that_cannot_be_opened_from_here_says_so(tmp_path, monkeypatch):
    """Defect 3: with no program for .xlsx, Open at... raised inside the window, and the window said nothing."""
    f = _read(tmp_path)
    f.next()
    f.run()

    def cannot(path):
        raise FileNotFoundError(2, "No such file or directory", "xdg-open")
    monkeypatch.setattr(launcher, "open_file", cannot)
    b = launcher.book_for(f.extract)
    need = f.needs[0]
    said = f.open_at(need)
    assert said.startswith(f"Couldn't open {b.name} from here (No such file or directory). Open it yourself: it is in "
                           f"{b.parent}."), said
    assert said.endswith(f"Then go to {need.sheet} {need.cell}.")
    assert f.open_book().startswith(f"Couldn't open {b.name} from here")


def test_an_install_that_worked_is_good_news_not_a_refusal(tmp_path, monkeypatch):
    """Defect 4: "Installed numpy, openpyxl and PyYAML. Everything PocketBook needs is here." was drawn in the red of
    a refusal, on the extract page and again, after scikit-learn, on Choose tests."""
    monkeypatch.setattr(deps, "missing", lambda: [])
    monkeypatch.setattr(deps, "missing_optional", lambda: [])
    f = launcher.Flow(gate=launcher.AddOns())
    f.gate.got = ["numpy", "openpyxl", "PyYAML"]
    f.after_install(optional=False)
    assert f.note == ["Installed numpy, openpyxl and PyYAML. Everything PocketBook needs is here."] and f.message == []
    f.gate.got = ["scikit-learn"]
    f.after_install(optional=True)
    assert f.note[0].startswith("Installed scikit-learn.") and f.message == []
    # one that didn't install is still the note for IT, in the message
    f.gate.optional = ["scikit-learn"]
    f.after_install(optional=True)
    assert f.note == [] and f.message[0] == "Couldn't install scikit-learn from here."


def test_the_line_under_the_extract_counts_the_add_ons_missing():
    """Defect 9: "while the add-on installs", with three missing and nothing installing."""
    assert launcher.waiting_words(["numpy", "openpyxl", "PyYAML"]).startswith("Until they are in, ")
    assert launcher.waiting_words(["numpy"]).startswith("Until it is in, ")
    assert "installs" not in launcher.waiting_words(["numpy", "PyYAML"])


# ---- the workbook


def test_record_names_each_measure_as_the_result_tabs_do(tmp_path, monkeypatch):
    """Defect 6: Record's "Does it add up" said "Outcome, share of loans", "GCO per booked dollar" and "Profit after
    losses: RANR per booked dollar" for what Pockets, Grids and Split call Bad loans, Charge-offs and Kept after
    losses (GCOs ($) and RANR since 30 Sep 2026)."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = synth.write_extract(tmp_path, n=3000)
    b = book.set_up(x).book
    _answer(b)
    assert book.run(b).ok
    labels = [label for label, *_ in tabs.record_rows(b) if label]
    heads = ("Worse now: ", "Materiality line: ", "Left out of ", "Smallest gap a typical pocket could show: ")
    named = [lb[len(h):] for lb in labels for h in heads if lb.startswith(h)]
    named += [lb.split("gap: ", 1)[1] for lb in labels if lb.startswith("Loans needed for ")]
    plain = set(results.PLAIN.values())
    assert len(named) >= 10
    assert [n for n in named if n not in plain] == []


def test_the_scouting_curve_shows_a_dollar_column_as_whole_numbers():
    """Defect 8: REV_DEBT's curve on Scouting read 733.1264, 1584.205, 4405.724."""
    assert scout_tab.curve_format([(0.0, 0.05), (733.1264, 0.047), (29764.8957, 0.12)]) == "#,##0"
    assert scout_tab.curve_format([(0.41, 0.05), (0.93, 0.08)]) == "0.####"


# ---- the pre-spec, on the window and on Start here


def _found(b) -> dict:
    ws = load_workbook(b)[book.FOUND]
    return {r[0]: r[1] for r in ws.iter_rows(values_only=True) if r and r[0]}


def test_a_pre_spec_changed_or_just_written_never_reads_as_followed(tmp_path, monkeypatch):
    """Defect 1, the worst: the analyst edited the pre-spec after the held-back run and pressed Run again, and the
    window and Start here said "Follows the pre-spec: Yes", in green. Only Record's warnings said it had changed.
    And on the run whose own scouting wrote it, "Yes" could not have been anything else."""
    from test_scout import _book, _run
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x, b = _book(tmp_path, 6000)
    ran, _ = _run(b, monkeypatch)
    assert ran.ok, ran.lines
    tile = launcher.confirm_tiles(ran.summary)[-1]
    assert tile[:3] == ("Follows the pre-spec", "Written now", "by this Run's scouting") and tile[4] != "POSITIVE"
    assert _found(b)["follows"] == "Written now: by this Run's scouting"
    # the analyst edits it after that held-back run
    p = scout.prespec_for(b)
    text = p.read_text(encoding="utf-8")
    m = re.search(r"bins: \[([0-9.]+)", text)
    assert m, text
    p.write_text(text.replace(m.group(0), f"bins: [{float(m.group(1)) * 0.9:g}", 1), encoding="utf-8")
    ran, res = _run(b, monkeypatch)
    assert ran.ok, ran.lines
    assert res.prespec.changed_after is not None
    tile = launcher.confirm_tiles(ran.summary)[-1]
    assert tile[:3] == ("Follows the pre-spec", "Changed", "after a held-back run: see Record"), tile
    assert tile[4] == "CRIMSON"
    assert _found(b)["follows"] == "Changed: after a held-back run: see Record"
    ws = load_workbook(b)["New variables"]
    said = [c.value for row in ws.iter_rows(max_col=3) for c in row if isinstance(c.value, str)
            and (c.value.startswith("from the ") or c.value.startswith("changed after"))]
    assert said == ["changed after a held-back run"], said
    # a third Run on the same file: the change is the file as it stands now, followed
    ran, _ = _run(b, monkeypatch)
    assert ran.ok and launcher.confirm_tiles(ran.summary)[-1][1:3] == ("Yes", "differs nowhere")


def test_start_here_says_the_loans_tested_in_the_words_new_variables_uses(monkeypatch):
    """Defect 11: Start here's "Loans tested" tile read "5,604 on development, 2,395 on the h", cut off by the next
    tile, where New variables' own tile says "5,604 found · 2,395 held back"."""
    from openpyxl import Workbook
    from pocketbook import confirm_tab, confirmatory
    h = {"kind": "confirm", "column": "REV_DEBT", "candidates": ["REV_DEBT"], "reference": "11,000 - 14,999",
         "development": 5604, "holdout": 2395, "deviations": 0}
    monkeypatch.setattr(confirmatory, "headline", lambda res: h)
    monkeypatch.setattr(confirmatory, "groups_found", lambda res: [])
    ws = Workbook().active
    confirm_tab.write_found(ws, None)
    got = {r[0]: r[1] for r in ws.iter_rows(values_only=True)}
    assert got["loans"] == "5,604 found · 2,395 held back"


class _Line:
    """A label on the black bar: gone once the bar is, and Tk refuses to set a label that is gone."""

    def __init__(self, alive):
        self.alive, self.text = alive, None

    def winfo_exists(self):
        return self.alive

    def configure(self, text):
        if not self.alive:
            raise RuntimeError('invalid command name ".!frame.!frame3.!frame.!label2"')
        self.text = text


def test_the_bar_clock_leaves_a_bar_that_is_gone_alone(monkeypatch):
    """The second run's defect, without a display: the loop that waits for pip set the black bar's line, which an
    earlier install had taken away with the bar; that raised, the loop stopped, and the page never finished."""
    monkeypatch.setattr(deps, "missing", lambda: ["numpy"])
    monkeypatch.setattr(deps, "missing_optional", lambda: ["scikit-learn"])
    gate = launcher.AddOns()
    gate.start()
    gone = _Line(alive=False)
    launcher.banner_clock({"status": gone}, gate, optional=True)
    launcher.banner_clock({"status": gone}, gate, optional=False)
    launcher.banner_clock({}, gate, optional=False)
    live = _Line(alive=True)
    launcher.banner_clock({"status": live}, gate, optional=False)
    assert live.text.startswith("Installing numpy...")
    launcher.banner_clock({"status": _Line(alive=True)}, gate, optional=True)     # no bar for an optional one


def test_scikit_learn_installed_after_the_add_ons_in_the_same_window_finishes(monkeypatch, tmp_path):
    """Found on the walk's second run: the analyst installs the add-ons from the black bar, then scikit-learn from
    Choose tests, in the same window. The bar's status line was gone; the loop that waits for pip tried to update it,
    died, and the page said "Installing scikit-learn..." for ever. Needs a display (xvfb-run on Linux)."""
    from test_deps import _wait, _window, fake_pip, only_missing
    root = _window()
    try:
        monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
        monkeypatch.setattr(launcher.Path, "home", lambda: tmp_path)
        gone = only_missing(monkeypatch, "numpy", "sklearn")
        w = launcher.build(root)
        flow = w["flow"]
        root.update()
        assert flow.screen() == "L4" and flow.gate.optional == ["scikit-learn"]
        fake_pip(monkeypatch, 0, "Successfully installed numpy\n", then=lambda: gone.discard("numpy"))
        w["install"].invoke()
        _wait(root, flow.gate)
        root.update()
        assert not flow.gate.missing and not w["status"].winfo_exists()      # the bar, and its line, are gone
        x = synth.write_extract(tmp_path, n=800)
        w["extract"].set(str(x))
        flow.set_up()
        flow.set_mode("new")
        w["render"]()
        root.update()
        # pip takes a while, as it does: the loop that waits for it goes round at least twice
        fake_pip(monkeypatch, 0, "Successfully installed scikit-learn\n",
                 then=lambda: (time.sleep(1.2), gone.discard("sklearn")))
        w["install_optional"].invoke()
        _wait(root, flow.gate)
        root.update()
        assert flow.gate.optional == [] and flow.note[0].startswith("Installed scikit-learn."), (flow.note,
                                                                                                 flow.message)
    finally:
        root.destroy()


def test_a_finished_tile_grows_to_its_words(monkeypatch, tmp_path):
    """The second run: with the column named, the first tile's heading took three lines and "on the holdout, 95%
    sure" was cut off under the tile's fixed 96 px. Every tile is as tall as its words. Needs a display."""
    from test_deps import _window
    root = _window()
    try:
        monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
        w = launcher.build(root)
        flow = w["flow"]
        h = {"kind": "confirm", "column": "INCOME_TO_SALES_RATIO", "candidates": ["INCOME_TO_SALES_RATIO"],
             "reference": "0.41 - 0.93", "worse": 1, "groups": 2, "confidence": 0.95, "capture": 0.54,
             "deviations": 0, "changed": True}
        flow.finished, flow.page = book.Outcome(True, Path(), [], summary=h), "done"
        w["render"]()
        root.update()
        short = [(t.winfo_reqheight(), sum(c.winfo_reqheight() for c in t.winfo_children()) + 6) for t in w["tiles"]]
        assert len(short) == 3 and [x for x in short if x[0] < x[1]] == [], short
    finally:
        root.destroy()
