"""A control character in the extract (reproduced 3 Oct 2026).

Set up stopped with openpyxl's IllegalCharacterError ("L5, L0000001, L0000002 cannot be used in worksheets") when
the first loan number carried a control character (\\x01): Columns writes each column's first values, and a
worksheet cannot hold the characters XML 1.0 forbids. A bank's extract can carry them: a stray byte from the system
that wrote it. Every piece of the extract's text the workbook shows (values, labels, dropdown lists, column names)
has the same problem, and a label shown cleaned must still select its own loans.

Every figure checked here is worked out from the loan file with the csv module alone.
"""

import csv
import re
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from pocketbook import book, choices as ch, synth
from test_book import _answer
from test_compare_2026_09_30 import _bad, _close

#: what XML 1.0 forbids in a document, and openpyxl refuses in a cell
FORBIDDEN = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
DIRTY_LOAN, DIRTY_CHANNEL, DIRTY_COLUMN = "L\x015", "Bro\x02ker", "ASSET\x03CLASS"


def _file(d: Path, n: int = 3000, channel: bool = False, column: bool = False) -> Path:
    """The synthetic book with the first loan number carrying \\x01; `channel` writes Broker as Bro\\x02ker, and
    `column` names ASSET_CLASS ASSET\\x03CLASS."""
    src = synth.write_extract(d / "src", n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    rows[0]["LOAN_NBR"] = DIRTY_LOAN
    for r in rows:
        if channel and r["CHANNEL"] == "Broker":
            r["CHANNEL"] = DIRTY_CHANNEL
        if column:
            r[DIRTY_COLUMN] = r.pop("ASSET_CLASS")
    out = d / "loans.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return out


def _forbidden_in(path: Path) -> list[str]:
    """Every part of the saved file holding a character XML forbids: cells, dropdown lists, names, anything."""
    with zipfile.ZipFile(path) as z:
        return [n for n in z.namelist() if FORBIDDEN.search(z.read(n).decode("utf-8", "replace"))]


def _offered(x: Path) -> list[str]:
    """The column names as the launcher offers them, read from the extract by Set up's own reader."""
    read = book.read_extract(x)
    assert read.problem is None, read.problem
    return [c.name for c in read.columns]


def test_ingest_uses_openpyxls_own_list_of_forbidden_characters():
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from pocketbook import ingest
    assert ingest.ILLEGAL_CHARACTERS_RE.pattern == ILLEGAL_CHARACTERS_RE.pattern
    for c in map(chr, range(32)):
        assert bool(ingest.ILLEGAL_CHARACTERS_RE.search(ingest.cleaned(f"a{c}b"))) is False, repr(c)


def test_cleaning_keeps_values_apart_and_reads_as_the_reader_did():
    """Two values differing only by a control character stay two values, and nothing else changes."""
    from pocketbook import ingest
    assert ingest.cleaned("Broker") == "Broker" and ingest.cleaned(" A b ") == " A b "
    assert ingest.cleaned("Bro\x02ker") not in ("Broker", "Brokker", "Bro ker")
    assert ingest.cleaned("A\x01") != ingest.cleaned("A") != ingest.cleaned("A\x02")
    # what the reader already trimmed (a \x1f at either end is whitespace to Python) still reads the same
    assert ingest.cell_text(ingest.cleaned("A\x1f")) == ingest.cell_text("A\x1f") == "A"
    assert ingest.cleaned(5) == 5 and ingest.cleaned(None) is None


def test_set_up_survives_a_control_character_in_the_first_loan_number(tmp_path, monkeypatch):
    """The reproduction of 3 Oct 2026, exactly: Set up writes, and the workbook opens again."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = _file(tmp_path)
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO", "REV_DEBT"),
                                            segments=("CHANNEL", "ASSET_CLASS"), outcome="BAD_FLAG"))
    assert out.ok, out.lines
    wb = load_workbook(out.book)
    shown = [c.value for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(re.search(r"(^|\s)L[^\s,]5, L0000001", v) for v in shown), \
        "the loan number is shown, a mark where the control character was, among Columns' first values"
    assert not _forbidden_in(out.book)


@pytest.fixture(scope="module")
def dirty(tmp_path_factory):
    """Set up and Run over a loan file with \\x01 in the first loan number, \\x02 in a channel (a segment, and the
    Grids' filter) and \\x03 in a column's name (a segment)."""
    from pocketbook import perm
    from test_book import at
    d = tmp_path_factory.mktemp("illegal_chars")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 100)
        x = _file(d, channel=True, column=True)
        offered = _offered(x)
        asset = next(c for c in offered if c.startswith("ASSET"))
        up = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL", asset),
                                               filter="CHANNEL", outcome="BAD_FLAG"))
        assert up.ok, up.lines
        _answer(up.book)
        wb = load_workbook(up.book)
        wb["Columns"][at(wb, "FICO", book.C_EDGES)] = "620; 680; 740"
        wb.save(up.book)
        ran = book.run(up.book)
    return {"book": up.book, "csv": x, "ran": ran, "asset": asset, "offered": offered, "dir": d}


def _rows(x) -> list[dict]:
    return list(csv.DictReader(open(x, encoding="utf-8")))


def test_set_up_and_run_both_write_and_the_workbook_opens_again(dirty):
    assert dirty["ran"].ok, dirty["ran"].lines
    wb = load_workbook(dirty["book"])
    assert {"Columns", "Grids", "Summary"} <= set(wb.sheetnames)
    assert not _forbidden_in(dirty["book"])
    assert not any(FORBIDDEN.search(line) for line in dirty["ran"].lines)


def test_a_column_named_with_a_control_character_is_offered_and_listed(dirty):
    """The launcher offers it under a name a worksheet can hold, and Columns lists it under that same name."""
    asset = dirty["asset"]
    assert not FORBIDDEN.search(asset) and asset != "ASSETCLASS" and "ASSET" in asset and "CLASS" in asset
    assert len(dirty["offered"]) == len(set(dirty["offered"]))
    names = [r[book.C_NAME - 1].value for r in book.table_rows(load_workbook(dirty["book"])["Columns"])]
    assert asset in names


def test_the_dirty_channel_is_a_filter_value_of_its_own_and_its_loans_tie(dirty, tmp_path):
    """The Grids' filter lists Branch, Online and the dirty channel as three values; picking it on Summary shows its
    loans, and only its: the loans and their bad-loan rate are the loan file's."""
    import tabs
    rows = _rows(dirty["csv"])
    wb = load_workbook(dirty["book"])
    only = "Only loans where CHANNEL is"
    opts = tabs.options(wb, "Grids", only)
    assert opts[0] == "All loans" and "Branch" in opts and "Online" in opts and len(opts) == 4, opts
    label = next(o for o in opts[1:] if o not in ("Branch", "Online"))
    assert not FORBIDDEN.search(label) and label.startswith("Bro") and label.endswith("ker")
    assert tabs.options(wb, "Summary", only) == opts
    mine = [r for r in rows if r["CHANNEL"] == DIRTY_CHANNEL]
    assert mine
    out = tabs.choose(dirty["book"], tmp_path / "s.xlsx", "Summary",
                      **{"Band or category column": "FICO", only: label})
    ws = tabs.calculated(out, "Summary")
    h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
    heads = {ws.cell(row=h, column=c).value: c for c in range(3, ws.max_column + 1)}
    r = next(r for r in range(h + 1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "All")
    assert ws.cell(row=r, column=heads["Loans"]).value == len(mine)
    assert _close(ws.cell(row=r, column=heads["Bad loans %"]).value, _bad(mine))


def test_the_dirty_channel_and_column_are_segments_whose_loans_tie(dirty, tmp_path):
    """Summary by CHANNEL shows the dirty channel's loans under its cleaned label; Summary by the dirty-named
    column shows each asset class's loans: both segments, both tied to the loan file."""
    import tabs
    rows = _rows(dirty["csv"])
    wb = load_workbook(dirty["book"])
    cats = tabs.options(wb, "Summary", "Band or category column")
    assert "CHANNEL" in cats and dirty["asset"] in cats, cats
    for col, raw_col in (("CHANNEL", "CHANNEL"), (dirty["asset"], DIRTY_COLUMN)):
        out = tabs.choose(dirty["book"], tmp_path / f"c{len(col)}.xlsx", "Summary", **{"Band or category column": col})
        ws = tabs.calculated(out, "Summary")
        h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
        heads = {ws.cell(row=h, column=c).value: c for c in range(3, ws.max_column + 1)}
        shown, r = {}, h + 1
        while ws.cell(row=r, column=2).value not in (None, ""):
            shown[str(ws.cell(row=r, column=2).value)] = (ws.cell(row=r, column=heads["Loans"]).value,
                                                          ws.cell(row=r, column=heads["Bad loans %"]).value)
            r += 1
        values = sorted({x[raw_col] for x in rows})
        for v in values:
            lab = next(k for k in shown if k == v or (FORBIDDEN.search(v) and k.startswith(v[:3])
                                                     and k.endswith(v[-3:])))
            mine = [x for x in rows if x[raw_col] == v]
            assert shown[lab][0] == len(mine), (col, v, shown)
            assert _close(shown[lab][1], _bad(mine)), (col, v)


def test_a_control_character_that_python_counts_as_a_line_break_does_not_split_a_row(tmp_path):
    """\\x0b, \\x0c and \\x1c to \\x1e end a line to str.splitlines(), so a CSV read through it split one loan in two
    and shifted every field after it, with no error (the review of 4 Oct 2026). A CSV row ends only at its own line
    break: three loans read as three, each value in its own column, the character shown as its control picture."""
    from pocketbook import ingest
    p = tmp_path / "breaks.csv"
    p.write_text("ID,NOTE,FICO\nL1,a\x0bb,700\nL2,plain,720\x1c\nL3,x\x0cy\x1dz\x1e,690\n", encoding="utf-8")
    t = ingest.read_table(p)
    assert len(t.rows) == 3
    assert [r["ID"] for r in t.rows] == ["L1", "L2", "L3"]
    assert t.rows[0]["NOTE"] == "a␋b" and t.rows[0]["FICO"] == "700"
    assert t.rows[1]["FICO"] == "720"                    # \x1c at the end is whitespace, trimmed like any other
    assert t.rows[2]["NOTE"] == "x␌y␝z" and t.rows[2]["FICO"] == "690"
