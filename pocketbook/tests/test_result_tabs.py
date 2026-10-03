"""The result tabs as the firm's redesign draws them (docs/redesign-2026-09-26/README.md, sections 5 to 8; the
redesign's phase 3): Pockets, RANR vs GCOs, Grids and Split, in place of Where it bleeds, Three-way, Losses vs
revenue and Prevalence.

Each tab's structure is held here, and its live behaviour through LibreOffice (tests/recalc.py): the dropdowns
are changed on a copy, as the analyst would, and so are Control's live lines, and the rows shown, their verdicts,
the caption and the colours are held to the engine run again in Python. The order of the rows is as of the last
Run, and the tab says so."""

from __future__ import annotations

import dataclasses
import re

import pytest
from openpyxl import load_workbook

import tabs
from pocketbook import book, config as cfgmod, engine, house, live, perm, results, synth
from pocketbook.ingest import read_table
from recalc import SOFFICE, recalc
from test_book import _answer
from test_book_dates import _choose
from test_live import _set

pytestmark = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice (soffice) isn't installed, so the "
                                                         "workbook's formulas can't be calculated here")


@pytest.fixture(scope="module")
def ran(tmp_path_factory):
    """A whole Run on the synthetic book, split by revolving debt, and the engine's result for it."""
    d = tmp_path_factory.mktemp("results")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 2_000)
        extract = synth.write_extract(d, n=4000)
        b = book.set_up(extract).book
        _answer(b)
        _choose(b, split="REV_DEBT")
        assert book.run(b).ok
        raw, problems, _ = book.read_book(b)
        assert not problems
        cfg = cfgmod.parse(raw)
        table_ = read_table(extract)
        res = engine.run(cfg, table_)
    return {"book": b, "dir": d, "cfg": cfg, "table": table_, "res": res, "values": recalc(b, d / "rc")}


def _names(res):
    return book._names(res)


def _key(res, kind, g, bl, dl, m):
    return (kind, f"{_names(res)[g.band]} x {_names(res)[g.dimension]}", bl, dl, m)


def _pocket_of(values):
    """Each _pockets row's pocket: (kind, grid, band, segment, measure)."""
    ws = values[live.POCKETS]
    return {r[0].row: tuple(r[c - 1].value for c in (live.P_KIND, live.P_GRID, live.P_BAND, live.P_SEG,
                                                      live.P_MEASURE))
            for r in ws.iter_rows(min_row=live.P_FIRST) if r[live.P_KIND - 1].value is not None}


def _shown(values) -> list[tuple]:
    """The pockets Pockets shows now, in order."""
    ws = values[results.POCKETS]
    of = _pocket_of(values)
    return [of[ws.cell(row=x["row"], column=results.K_ROW).value] for x in tabs.pockets(ws)]


def _order(res, kind, mname) -> list[tuple]:
    """Pockets' order as of the Run, from the engine: results.candidates."""
    m = next(x for x in res.measures if x.name == mname)
    return [_key(res, kind, g, bl, dl, mname) for g, bl, dl, _ in results.candidates(res, kind, m)]


def _verdicts(res) -> dict:
    """Worse? and Material?, as the engine says them, by pocket."""
    out = {}
    for kind, grids in (("grids", res.grids), ("three-way", res.three_way)):
        for g in grids:
            for (bl, dl), c in g.inner():
                for m in res.measures:
                    if m.is_rate:
                        s = c.rates[m.name]
                        out[_key(res, kind, g, bl, dl, m.name)] = (
                            live.worse_of(s.flag), {True: live.YES, False: live.NO, None: ""}[s.material], s.dollars)
    return out


def _caption(values) -> str:
    ws = values[results.POCKETS]
    return tabs.dropdown(ws, "Measure").offset(column=results.K_LOANS - results.K_BAND).value


def _counts(caption: str) -> tuple[int, int, int]:
    got = re.fullmatch(r"(\d+) worse and material · (\d+) worse · (\d+) shown", caption)
    assert got, caption
    return tuple(int(x) for x in got.groups())


# --------------------------------------------------------------------------
# The tabs themselves


def test_the_old_result_tabs_are_gone_and_the_four_new_ones_sit_in_order(ran):
    wb = load_workbook(ran["book"])
    names = wb.sheetnames
    assert not set(results.OLD_TABS) & set(names)
    visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    assert visible == ["Start here", "Glossary", "Control", "Columns", "Look", results.POCKETS, results.PCK, results.GRIDS,
                       results.SUMMARY, results.SPLIT, "Record"]
    for t in results.TABS:
        assert wb[t].sheet_properties.tabColor.rgb.endswith(house.INK), t
        assert wb[t].sheet_view.showGridLines is False, t
    for t in (*results.HIDDEN, results.CHART):
        assert wb[t].sheet_state == "hidden", t


def test_an_older_workbook_loses_the_tabs_the_redesign_replaced(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    b = book.set_up(synth.write_extract(tmp_path, n=1500)).book
    _answer(b)
    wb = load_workbook(b)
    for t in results.OLD_TABS:
        wb.create_sheet(t)["B2"] = "left by an older Run"
    wb.save(b)
    assert book.run(b).ok
    names = load_workbook(b).sheetnames
    assert not set(results.OLD_TABS) & set(names) and results.POCKETS in names


def test_each_result_tab_opens_with_its_title_band_and_one_method_note_that_folds(ran):
    wb = load_workbook(ran["book"])
    for t in results.TABS:
        ws = wb[t]
        first = next(c for c in ws[1] if c.value == t)
        assert first.fill.fgColor.rgb.endswith(house.INK) and first.font.b and first.font.size == 16, t
        assert first.border.bottom.color.rgb.endswith(house.KEY_RED), t
        notes = [c for row in ws.iter_rows() for c in row if c.value == "How this tab works"]
        assert len(notes) == 1, t                                     # T1: one method note, said once
        r = notes[0].row
        assert ws.row_dimensions[r].outline_level == 1 and ws.row_dimensions[r + 1].outline_level == 1, t
        assert ws.freeze_panes, t


def test_pockets_offers_the_three_dropdowns_with_the_redesigns_options(ran):
    wb = load_workbook(ran["book"])
    assert tabs.options(wb, results.POCKETS, "Measure") == ["Bad loans", "Bad dollars", "GCOs ($)",
                                                            "RANR", "RANR + GCOs"]
    assert tabs.options(wb, results.POCKETS, "Pockets") == ["Two-way", "Split by REV_DEBT"]
    assert tabs.options(wb, results.POCKETS, "Show") == ["All", "Worse and material", "Worse or not sure"]
    for label in ("Measure", "Pockets", "Show"):
        cell = tabs.dropdown(wb[results.POCKETS], label)
        assert cell.fill.fgColor.rgb.endswith(house.CANVAS) and cell.font.b, label       # changes now


def test_pockets_has_worse_and_material_as_two_columns_and_no_test_column(ran):
    ws = ran["values"][results.POCKETS]
    head = tabs.heads(ws, tabs.header_row(ws, results.K_NUM, "#"), results.K_NUM, results.K_HOLDS)
    assert head[:4] == ["#", "Band", "Segment", None]                   # the half is blank on the two-way view
    assert head[4:13] == ["Loans", "This pocket", "Rest of band", "× rest of band", "Bad loans above share",
                          "Worse?", "p-value", "Material?", "Could have caught"]
    assert "Test" not in head and not any("luck" in str(h).lower() for h in head)


def test_pockets_ranks_the_worse_pockets_first_by_their_dollars(ran):
    res, v = ran["res"], ran["values"]
    shown = _shown(v)
    assert shown == _order(res, "grids", "outcome_loans")               # All, two-way, Bad loans: every candidate
    got = _verdicts(res)
    worse = [got[k][2] for k in shown if got[k][0] == live.YES]
    assert worse == sorted(worse, reverse=True) and len(worse) >= 1
    first_not = next(i for i, k in enumerate(shown) if got[k][0] != live.YES)
    assert all(got[k][0] != live.YES for k in shown[first_not:])        # the worse ones first
    said = " ".join(str(c.value) for row in v[results.POCKETS].iter_rows(max_row=25) for c in row if c.value)
    assert 'Order and "Could have caught" are from the last Run, 20' in said and "not the order" in said


@pytest.mark.parametrize("measure, kind, show", [("GCOs ($)", "Two-way", "Worse or not sure"),
                                                 ("RANR", "Split by REV_DEBT", "All"),
                                                 ("Bad dollars", "Split by REV_DEBT", "Worse and material")])
def test_the_dropdowns_pick_the_rows_and_the_caption_counts_them(ran, tmp_path, measure, kind, show):
    res = ran["res"]
    b = tabs.choose(ran["book"], tmp_path / "picked.xlsx", results.POCKETS, measure=measure, pockets=kind, show=show)
    v = recalc(b, tmp_path / "rc")
    mname = next(m.name for m in res.measures if m.is_rate and results.plain(m) == measure)
    k = "grids" if kind == "Two-way" else "three-way"
    got = _verdicts(res)
    keep = {"All": lambda w, mt: True, "Worse and material": lambda w, mt: w == live.YES and mt == live.YES,
            "Worse or not sure": lambda w, mt: w in (live.YES, live.NOT_SURE)}[show]
    want = [x for x in _order(res, k, mname) if keep(*got[x][:2])]
    assert _shown(v) == want and want
    rows = tabs.pockets(v[results.POCKETS])
    assert [(x["worse"], x["material"]) for x in rows] == [(got[x][0] or None, got[x][1] or None) for x in want]
    every = [x for x in got if x[0] == k and x[4] == mname]
    n_worse = sum(1 for x in every if got[x][0] == live.YES)
    n_both = sum(1 for x in every if got[x][0] == live.YES and got[x][1] == live.YES)
    assert _counts(_caption(v)) == (n_both, n_worse, len(want))
    if k == "three-way":                                          # the split view names each half and its grid
        assert all(x["half"] in ("high", "low", "none") for x in rows) and all(x["holds"] for x in rows)


def test_worse_and_material_leaves_out_a_worse_pocket_below_the_line(ran, tmp_path):
    """Worse? and Material? are judged apart: with the materiality line between the worse pockets' dollars, Show
    "Worse and material" keeps the worse pockets over the line and leaves out the ones under it."""
    res0 = ran["res"]
    worse = sorted(c.rates["gco_rate"].dollars for g in res0.three_way for _, c in g.inner()
                   if c.rates["gco_rate"].flag == engine.WORSE)
    assert len(worse) >= 2                                   # the split pockets: several worse on charge-offs
    line = (worse[0] + worse[-1]) / 2
    b = _set(ran["book"], tmp_path / "line.xlsx", {"materiality": (None, line)})
    b = tabs.choose(b, tmp_path / "picked.xlsx", results.POCKETS, measure="GCOs ($)", show="Worse and material",
                    pockets="Split by REV_DEBT")
    v = recalc(b, tmp_path / "rc")
    cfg = ran["cfg"]
    res = engine.run(dataclasses.replace(cfg, benchmark=dataclasses.replace(cfg.benchmark,
                                                                            materiality=("dollars", line))),
                     ran["table"])
    got = _verdicts(res)
    want = [x for x in _order(res0, "three-way", "gco_rate") if got[x][:2] == (live.YES, live.YES)]
    left_out = [x for x in _order(res0, "three-way", "gco_rate") if got[x][:2] == (live.YES, live.NO)]
    assert want and left_out
    assert _shown(v) == want


def test_the_rows_verdicts_caption_and_colours_follow_control_without_a_run(ran, tmp_path):
    """Worse at 2 times and judged against the rest of the book, changed on Control: the rows shown, their
    verdicts and the caption are the engine's run again with those lines, in the last Run's order."""
    res0 = ran["res"]
    b = _set(ran["book"], tmp_path / "lines.xlsx", {"worse_at": ("2 times", None),
                                                    "compare_to": ("The rest of the book", None)})
    b = tabs.choose(b, tmp_path / "picked.xlsx", results.POCKETS, measure="GCOs ($)", show="Worse or not sure")
    v = recalc(b, tmp_path / "rc")
    cfg = ran["cfg"]
    res = engine.run(dataclasses.replace(cfg, benchmark=dataclasses.replace(cfg.benchmark, worse_at=2.0,
                                                                            compare_to="topline")), ran["table"])
    got = _verdicts(res)
    want = [x for x in _order(res0, "grids", "gco_rate") if got[x][0] in (live.YES, live.NOT_SURE)]
    assert _shown(v) == want and want                              # the Run's order, the new verdicts
    before = _verdicts(res0)
    assert any(before[x][:2] != got[x][:2] for x in got)           # the change moved something
    rows = tabs.pockets(v[results.POCKETS])
    for x, k in zip(rows, want):
        assert (x["worse"], x["material"]) == (got[k][0] or None, got[k][1] or None), (x, k)
        assert x["excess"] == pytest.approx(got[k][2])             # the rest of the book's dollars (Option A)
    every = [x for x in got if x[0] == "grids" and x[4] == "gco_rate"]
    assert _counts(_caption(v))[:2] == (sum(1 for x in every if got[x][:2] == (live.YES, live.YES)),
                                        sum(1 for x in every if got[x][0] == live.YES))
    ws = v[results.POCKETS]
    head = tabs.header_row(ws, results.K_NUM, "#")
    assert ws.cell(row=head, column=results.K_REST).value == "Rest of book"
    tiles = {ws.cell(row=r, column=c).value: ws.cell(row=r + 1, column=c).value
             for r in range(3, head) for c in range(3, 12) if ws.cell(row=r, column=c).value in ("Worse at",
                                                                                                "Judged against")}
    assert tiles == {"Worse at": 2.0, "Judged against": "Rest of the book"}


def test_the_verdict_colours_are_the_redesigns(ran):
    ws = load_workbook(ran["book"])[results.POCKETS]
    by_col = {}
    for rng in ws.conditional_formatting:
        col = re.match(r"[A-Z]+", str(rng.sqref)).group(0)
        by_col.setdefault(col, []).extend(sorted(rng.rules, key=lambda r: r.priority))
    worse = by_col[results.col(results.K_WORSE)]
    yes = next(r for r in worse if r.formula[0].endswith(f'="{live.YES}")'))
    assert yes.dxf.fill.fgColor.rgb.endswith(house.ALERT_FG) and yes.dxf.font.color.rgb.endswith(house.CRIMSON)
    unsure = next(r for r in worse if r.formula[0].endswith(f'="{live.NOT_SURE}")'))
    assert unsure.dxf.fill.fgColor.rgb.endswith(house.CANVAS)
    mat = next(r for r in by_col[results.col(results.K_MAT)] if r.formula[0].endswith(f'="{live.YES}")'))
    assert mat.dxf.fill.fgColor.rgb.endswith(house.MIST) and mat.dxf.font.b


def test_judged_against_the_book_counts_points_and_dollars_against_the_rest_of_the_book(ran):
    """Option A (NEXT-GOAL item 5): a pocket's gap and its dollars both against every other loan in the book, so
    a pocket is worse only when it also loses dollars over that rest. OC-4's excess over the whole book stays,
    for the tie-out."""
    cfg = ran["cfg"]
    res = engine.run(dataclasses.replace(cfg, benchmark=dataclasses.replace(cfg.benchmark, compare_to="topline")),
                     ran["table"])
    seen = 0
    for g in res.grids:
        for _, c in g.inner():
            for m in res.measures:
                if not m.is_rate:
                    continue
                s, t = c.rates[m.name], res.total.rates[m.name]
                if s.den == t.den or s.vs_rest is None:
                    continue
                rest = (t.num - s.num) / (t.den - s.den)
                want = (s.num - rest * s.den) if m.higher_is == "worse" else (rest * s.den - s.num)
                assert s.dollars == pytest.approx(want) and s.dollars == s.excess_rest
                if s.flag == engine.WORSE:
                    assert s.dollars > 0
                seen += 1
        inner = [c.rates["gco_rate"].excess for _, c in g.inner()]
        assert sum(inner) == pytest.approx(0, abs=1e-3 * abs(res.total.rates["gco_rate"].num))
    assert seen > 100


def test_the_rest_a_pocket_is_read_against_follows_judged_against(ran):
    """Rest of band on Pockets is the rest the verdict used: the rest of the pocket's band when Control says
    so (the fixture's answer), not the rest of the book. Redesign phase 3 rewrote the only test that read it,
    and CI's planted bug "rest rate from the book" went uncaught."""
    ws = ran["values"][live.POCKETS]
    wrong, seen = [], 0
    for r in ws.iter_rows(min_row=live.P_FIRST):
        shown, band, whole = (r[c - 1].value for c in (live.P_REST, live.P_REST_BAND, live.P_REST_BOOK))
        if not all(isinstance(v, (int, float)) for v in (shown, band, whole)) or band == pytest.approx(whole):
            continue
        seen += 1
        if shown != pytest.approx(band):
            wrong.append((r[0].row, shown, band, whole))
    assert seen > 50 and not wrong, wrong[:5]


# --------------------------------------------------------------------------
# RANR vs GCOs


def test_paid_cost_kept_shows_one_grid_at_a_time_in_its_column_groups(ran, tmp_path):
    res = ran["res"]
    wb = load_workbook(ran["book"])
    grids = tabs.options(wb, results.PCK, "Grid")
    assert grids == [f"{_names(res)[g.band]} x {_names(res)[g.dimension]}" for g in res.grids]
    ws = ran["values"][results.PCK]
    head = tabs.header_row(ws, results.C_TOG, "Together")
    assert [ws.cell(row=head - 1, column=c).value for c in (results.C_PAID, results.C_COST, results.C_KEPT)] == [
        "RANR + GCOs · gap vs rest of band", "GCOs · × rest of band", "RANR · gap vs rest of band"]
    assert all(wb[results.PCK].cell(row=head - 1, column=c).border.bottom.style == "medium"
               for c in range(results.C_PAID, results.C_TOG + 1))
    # another grid picked: its own pockets, in its own last-Run order
    g = res.grids[1]
    v = recalc(tabs.choose(ran["book"], tmp_path / "g.xlsx", results.PCK, grid=grids[1]), tmp_path / "rc")
    rows = tabs.pck(v[results.PCK])
    assert [(x["band"], x["seg"]) for x in rows] == [(x["band"], x["seg"]) for x in results.pck_rows(res, g)]
    for x in rows:
        k = lambda m: _key(res, "grids", g, x["band"], x["seg"], m)          # noqa: E731
        s = {m: g.cells[(x["band"], x["seg"])].rates[m] for m in ("gco_rate", "ranr_rate", "contribution_rate")}
        assert x["flags"] == {"c": s["contribution_rate"].flag, "g": s["gco_rate"].flag, "r": s["ranr_rate"].flag}
        untested = any(y.flag in (engine.FEW, engine.THIN) for y in s.values())
        assert (x["together"] or "") == ("" if untested else results.together_of(s["gco_rate"].flag,
                                                                                 s["ranr_rate"].flag)), k("gco_rate")
        assert x["cost_d"] == pytest.approx(s["gco_rate"].dollars)


def test_the_together_formula_reads_all_five_pairs_and_nothing_untested(tmp_path):
    """Together as the tab writes it, calculated, against results.together_of for every pair of flags."""
    from openpyxl import Workbook
    from recalc import values_of
    flags = [engine.WORSE, engine.BETTER, engine.IN_LINE, engine.UNSURE_WORSE, engine.UNSURE_BETTER, engine.FEW]
    wb = Workbook()
    ws = wb.active
    cases = [(g, r, u) for g in flags for r in flags for u in (0, 1)]
    for i, (g, r, u) in enumerate(cases, start=1):
        ws.cell(row=i, column=1, value=g)
        ws.cell(row=i, column=2, value=r)
        ws.cell(row=i, column=3, value=u)
        ws.cell(row=i, column=4, value="=" + results.together_formula(f"$A{i}", f"$B{i}", f"$C{i}"))
    got = values_of(wb, tmp_path).active
    for i, (g, r, u) in enumerate(cases, start=1):
        want = "" if u else results.together_of(g, r)
        assert (got.cell(row=i, column=4).value or "") == want, (g, r, u)
    assert {results.together_of(g, r) for g in flags for r in flags} - {""} == set(results.TOGETHER.values())


def test_the_scatter_is_the_grid_picked_on_a_log_scale_with_lines_at_one_and_zero(ran):
    wb = load_workbook(ran["book"])
    (chart,) = wb[results.PCK]._charts
    assert chart.x_axis.scaling.logBase == 10
    assert chart.x_axis.scaling.min <= 0.1 and chart.x_axis.scaling.max >= 10
    hs = wb[results.CHART]
    lines = [s for s in chart.series if s.tx is not None and getattr(s.tx, "v", None) == "line"]
    assert len(lines) == 2
    xs = [[hs[c].value for c in ref.numRef.f.split("!")[1].replace("$", "").split(":")] for ref in
          (s.xVal for s in lines)]
    ys = [[hs[c].value for c in ref.numRef.f.split("!")[1].replace("$", "").split(":")] for ref in
          (s.yVal for s in lines)]
    assert xs[0] == [1.0, 1.0] and ys[1] == [0.0, 0.0]                    # dashed at 1x and at 0
    assert all(s.graphicalProperties.line.dashStyle == "dash" for s in lines)
    corners = {s.tx.v for s in chart.series if s.tx is not None and getattr(s.tx, "v", None)} - {"line", "Pockets"}
    assert corners == {"Priced for it", "Net drain", "Strong", "Safe but idle"}
    named = [s for s in chart.series if s.tx is not None and s.tx.strRef is not None]
    assert named and all(s.tx.strRef.f.startswith(f"'{results.CHART}'!") for s in named)



def test_the_scatter_colours_by_verdict_numbers_its_named_pockets_and_lists_them_under_it(ran):
    """At the bank, 29 Sep 2026: a Net drain drawn black read as nothing, and the names ran over the dots."""
    wb = load_workbook(ran["book"])
    ws, hs = wb[results.PCK], wb[results.CHART]
    (chart,) = ws._charts
    assert chart.x_axis.title is None and chart.y_axis.title is None
    red = hs.cell(row=1, column=results.H_RX).value
    green = hs.cell(row=1, column=results.H_GX).value
    assert f'"{results.BAD_TOGETHER[0]}"' in red and not any(f'"{g}"' in red for g in results.GOOD_TOGETHER)
    assert all(f'"{g}"' in green for g in results.GOOD_TOGETHER) and f'"{results.BAD_TOGETHER[0]}"' not in green
    fills = {getattr(s.marker.graphicalProperties.solidFill.srgbClr, "val", s.marker.graphicalProperties.solidFill.srgbClr)
             for s in chart.series
             if s.marker is not None and s.marker.graphicalProperties is not None
             and s.marker.graphicalProperties.solidFill is not None
             and s.marker.graphicalProperties.solidFill.srgbClr is not None}
    assert {results.house.KEY_RED, results.POSITIVE} <= fills
    named = [s for s in chart.series if s.tx is not None and s.tx.strRef is not None]
    assert named and all(s.tx.strRef.f.split("!")[1].startswith(f"${results.col(results.H_NUM)}$") for s in named)
    cells = [c.value for r in ws.iter_rows() for c in r if isinstance(c.value, str)]
    assert "Numbered on the chart" in cells
    listed = [v for v in cells if v.startswith(f"=IF('{results.CHART}'!${results.col(results.H_LNAME)}$")]
    assert len(listed) == results.LABELLED


# --------------------------------------------------------------------------
# Grids


def test_grids_fill_the_four_blocks_for_the_grid_and_measure_picked(ran, tmp_path):
    res = ran["res"]
    g = res.grids[1]
    name = f"{_names(res)[g.band]} x {_names(res)[g.dimension]}"
    wb = load_workbook(ran["book"])
    grids = tabs.options(wb, results.GRIDS, "Grid")
    assert name in grids and any(x.endswith(" / REV_DEBT") for x in grids)        # two-way, then split
    v = recalc(tabs.choose(ran["book"], tmp_path / "g.xlsx", results.GRIDS, grid=name, measure="GCOs ($)"),
               tmp_path / "rc")
    ws = v[results.GRIDS]
    rate, book_, band, loans = (tabs.block(ws, t) for t in ("Rate · GCOs ÷ Booked", "vs the book", "vs rest of band",
                                                            "Loans"))
    few = (engine.THIN, engine.FEW)
    for (bl, d), c in g.inner():
        s = c.rates["gco_rate"]
        assert rate[(bl, d)] == pytest.approx(s.rate)
        assert book_[(bl, d)] == (None if s.reading_topline in few else pytest.approx(s.vs_topline))
        assert band[(bl, d)] == (None if s.reading_band in few or s.vs_band is None else pytest.approx(s.vs_band))
        assert loans[(bl, d)] == c.rows
    assert loans[("All", "All")] == res.rows
    # a pocket with no loans is blank, not nought (the first grid, as the tab opens)
    g0 = res.grids[0]
    rate0 = tabs.block(ran["values"][results.GRIDS], "Rate · Bad loans ÷ Loans")
    empty = [(bl, d) for bl in g0.band_labels for d in g0.dim_labels if (bl, d) not in g0.cells]
    assert empty and all(rate0[k] is None for k in empty)
    assert all(v_ is None for (bl, d), v_ in band.items() if "All" in (bl, d))


def test_the_heat_scale_is_the_redesigns_and_the_loans_block_has_no_red_or_green(ran):
    ws = load_workbook(ran["book"])[results.GRIDS]
    loans_fill, heat_fill = set(), set()
    for rng in ws.conditional_formatting:
        fills = {r.dxf.fill.fgColor.rgb[-6:] for r in rng.rules if r.dxf is not None and r.dxf.fill is not None}
        formulas = " ".join(r.formula[0] for r in rng.rules)
        if '="All"' in formulas:                                            # a share of the grid's loans
            loans_fill |= fills
        elif "LOG(" in formulas:
            heat_fill |= fills
    assert heat_fill >= {house.HEAT_GOOD, house.HEAT_MID, house.HEAT_BAD, house.HEAT_BAD2}
    assert loans_fill and loans_fill <= {c for _, c in results.LOANS_STEPS}
    assert not loans_fill & {house.HEAT_GOOD, house.HEAT_BAD, house.HEAT_BAD2, house.ALERT_FG, house.POSITIVE_BG}


# --------------------------------------------------------------------------
# Split


def test_split_shows_the_summary_and_the_two_grids_for_what_is_picked(ran, tmp_path):
    res = ran["res"]
    wb = load_workbook(ran["book"])
    grids = tabs.options(wb, results.SPLIT, "Grid")
    g = next(x for x in res.grids if f"{_names(res)[x.band]} x {_names(res)[x.dimension]}" == grids[-1])
    b = tabs.choose(ran["book"], tmp_path / "s.xlsx", results.SPLIT, grid=grids[-1], measure="GCOs ($)")
    v = recalc(b, tmp_path / "rc")
    ws = v[results.SPLIT]
    chip = tabs.dropdown(ws, "Grid").offset(column=results.SPLIT_CHIP - 2).value
    assert chip == ("Holds FICO fixed" if book._holds_partner(res, g) else "Doesn't hold FICO fixed")
    head = tabs.header_row(ws, 2, "Measure")
    rows = {ws.cell(row=r, column=2).value: [ws.cell(row=r, column=c).value for c in range(3, 10)]
            for r in range(head + 1, head + 6)}
    p = g.split_pooled["gco_rate"]
    got = rows["GCOs ($)"]
    # a shuffled p-value near the bar prints as Borderline (29 Sep 2026): results.split_said
    pooled_p = results.split_said(p["ratio_p"], p.get("ratio_se"), res.config.benchmark.confidence)
    assert got[0] == p["pockets"] and got[2] == pytest.approx(p["ratio"])
    assert got[4] == (pooled_p if isinstance(pooled_p, str) else pytest.approx(pooled_p))
    assert got[3].endswith("×") and " to " in got[3]
    said = next(c.value for row in ws.iter_rows() for c in row if str(c.value).startswith("Same in every pocket? "))
    assert said.startswith("Same in every pocket? Bad loans: ")
    hl, pv = tabs.block(ws, "GCOs ($), high vs low"), tabs.block(ws, "p-value per pocket")
    for (bl, d), x in g.split_compare.items():
        if "gco_rate" not in x or x["gco_rate"][0] is None:
            continue
        gap, pval = x["gco_rate"][0], x["gco_rate"][1]
        assert (bl, str(d)) in pv, (sorted(pv)[:8], len(pv))
        want = results.split_said(pval, g.split_se.get((bl, d), {}).get("gco_rate"),
                                   res.config.benchmark.confidence)
        assert pv[(bl, str(d))] == (want if isinstance(want, str) else pytest.approx(pval))
        shown = hl[(bl, str(d))]
        if pval < 0.05:
            assert shown == pytest.approx(gap)
        else:
            assert shown == f"({gap:.2f}×)"


def test_the_split_brackets_follow_the_confidence_on_control(ran, tmp_path):
    b = _set(ran["book"], tmp_path / "c.xlsx", {"confidence": ("90% sure", None)})
    before = tabs.block(ran["values"][results.SPLIT], "Bad loans, high vs low")
    after = tabs.block(recalc(b, tmp_path / "rc")[results.SPLIT], "Bad loans, high vs low")
    brackets = lambda blk: sum(1 for v in blk.values() if isinstance(v, str))       # noqa: E731
    assert brackets(after) < brackets(before)            # a wider bar: fewer gaps left in brackets


# --------------------------------------------------------------------------
# Words


def test_nothing_on_the_result_tabs_or_in_their_dropdowns_says_luck_or_cube(ran):
    banned = re.compile(r"\b(luck|cube)\b", re.IGNORECASE)
    v = ran["values"]
    for t in (*results.TABS, results.CHOICES):
        for row in v[t].iter_rows():
            for c in row:
                assert not (isinstance(c.value, str) and banned.search(c.value)), (t, c.coordinate, c.value)
    wb = load_workbook(ran["book"])
    for chart in wb[results.PCK]._charts:
        titles = [a.title.tx.rich.p[0].r[0].t for a in (chart, chart.x_axis, chart.y_axis) if a.title is not None]
        assert not any(banned.search(x) for x in titles)
