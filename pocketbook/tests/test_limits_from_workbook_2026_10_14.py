"""The two column limits, when the extract already has a workbook (HANDOFF.md open item, 4 Oct 2026: "The launcher's
Next overwrites a category limit typed on Control").

Before: the launcher kept its own last choice of the two limits and Next wrote it over Control, so a value typed on
Control was replaced without a word. The firm, 14 Oct 2026, chose the workbook's: picking an extract loads the limits
Control holds into the launcher's two dropdowns, as bands, segments and filters already are, and Next writes back
what the dropdowns show. A value typed on Control survives; a change made in the launcher replaces it.
"""

from openpyxl import load_workbook

from pocketbook import book, control, launcher
from test_launcher import _flow, _graded


def _type_on_control(b, key, choose=None, own=None):
    """What an analyst does in Excel: pick an option in C, or type a number of their own in D."""
    wb = load_workbook(b)
    ws = wb[control.SHEET]
    r = control.row_of(ws, key)
    ws.cell(row=r, column=control.CHOOSE_COL).value = choose
    ws.cell(row=r, column=control.OWN_COL).value = own
    wb.save(b)


def _limits(b) -> dict:
    ws = load_workbook(b)[control.SHEET]
    return {k: control.answer_of(k, ws.cell(row=control.row_of(ws, k), column=control.CHOOSE_COL).value,
                                 ws.cell(row=control.row_of(ws, k), column=control.OWN_COL).value)
            for k in ("few_values", "many_values")}


def _written(tmp_path, few=6, many=50):
    """An extract whose workbook the launcher wrote with these limits."""
    x = _graded(tmp_path)
    f = _flow(x, few=few, many=many)
    f.set_up()
    f.next()
    assert f.page == "answer", f.message
    b = launcher.book_for(x)
    assert _limits(b) == {"few_values": few, "many_values": many}
    return x, b


def test_a_limit_picked_on_control_survives_next(tmp_path):
    """The reproduction: 24 picked on Control, the launcher reopened with its own last choice (6), Next pressed."""
    x, b = _written(tmp_path)
    _type_on_control(b, "few_values", choose="24 values")
    f = _flow(x, few=6)                                   # the launcher's saved preference is still 6
    assert (f.few, f.many) == (24, 50)
    f.set_up()
    f.next()
    assert _limits(b) == {"few_values": 24, "many_values": 50}


def test_a_limit_typed_off_the_list_on_control_survives_next(tmp_path):
    """A number of the analyst's own in D (60 is not one of the options), not only an option picked in C."""
    x, b = _written(tmp_path)
    _type_on_control(b, "many_values", own=60)
    f = _flow(x)
    assert (f.few, f.many) == (6, 60)
    f.set_up()
    f.next()
    assert _limits(b) == {"few_values": 6, "many_values": 60}


def test_a_limit_changed_in_the_launcher_after_reopening_is_written(tmp_path):
    x, b = _written(tmp_path)
    _type_on_control(b, "few_values", choose="24 values")
    f = _flow(x)
    assert f.few == 24 and f.limits_said
    f.choose_limit("few", 6)
    assert f.few == 6 and not f.limits_said               # the line no longer describes the dropdowns
    f.set_up()
    f.next()
    assert _limits(b) == {"few_values": 6, "many_values": 50}


def test_the_launcher_says_where_the_limits_came_from(tmp_path):
    x, b = _written(tmp_path)
    f = _flow(x)
    assert f.limits_said == launcher.LIMITS_FROM_BOOK
    (tmp_path / "other").mkdir()
    f.pick(str(_graded(tmp_path / "other")))              # an extract with no workbook beside it
    assert f.limits_said == "" and (f.few, f.many) == (6, 50)


def test_an_extract_with_no_workbook_keeps_the_launchers_limits(tmp_path):
    x = _graded(tmp_path)
    f = _flow(x, few=24, many=25)
    assert (f.few, f.many, f.limits_said) == (24, 25, "")


def test_the_limits_load_when_the_launcher_opens_on_the_last_extract(tmp_path):
    """The window starts a Flow with the extract from its saved preferences, not through pick."""
    x, b = _written(tmp_path)
    _type_on_control(b, "few_values", choose="24 values")
    f = launcher.Flow(gate=launcher.AddOns(), extract=str(x), few=6)
    assert f.few == 24 and f.limits_said == launcher.LIMITS_FROM_BOOK


def test_a_blank_limit_on_control_leaves_the_launchers(tmp_path):
    x, b = _written(tmp_path, few=6, many=25)
    _type_on_control(b, "few_values")                     # both cells cleared
    f = _flow(x, few=12)
    assert (f.few, f.many) == (12, 25)


def test_an_unreadable_workbook_leaves_the_launchers_limits(tmp_path):
    x = _graded(tmp_path)
    launcher.book_for(x).write_bytes(b"not a workbook")
    assert book.limits_in(x) == {}
    f = _flow(x, few=24)
    assert (f.few, f.limits_said) == (24, "")


def test_read_limits_reads_a_read_only_worksheet(tmp_path):
    """book.limits_in opens the workbook read-only, whose rows stop at their last filled cell."""
    _, b = _written(tmp_path, few=6, many=25)
    _type_on_control(b, "many_values", own=60)
    wb = load_workbook(b, read_only=True)
    try:
        assert control.read_limits(wb[control.SHEET]) == {"few_values": 6, "many_values": 60}
    finally:
        wb.close()
