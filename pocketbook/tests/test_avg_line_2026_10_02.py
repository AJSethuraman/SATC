"""Avg line and × book (the firm, 2 Oct 2026: "I want to start including and using booked dollar averages so more
easily demonstrate how line assignments look in pockets and I think it adds to the story if our basis is commitments
would be stronger elsewhere"). Booked is the committed line, so booked dollars per loan is the average line.

Placed where the firm chose: Summary, right after Booked $, on every row and every filter view; and RANR vs GCOs, at
the end of the gross block. Avg line is a row's booked dollars over its loans with a booked amount (a loan with none is
left out of both, as Grids' Loan size does); × book is that over the whole book's Avg line, filtered or not. Blank,
never 0 and never an error, where a row has no such loan. Shaded as Grids' Loan size is, one neutral hue, never red
or green.

Every figure checked against the loan file is worked out here from the CSV's text with plain arithmetic: nothing in
the checks is imported from pocketbook. The checks of _views and of the formulas as written need no LibreOffice and
no display; the two that read the tabs calculated need LibreOffice and skip without it."""

import csv

import pytest
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from pocketbook import book, choices as ch, results, synth
from test_book import _answer

import tabs
from recalc import need_soffice

PICKER = "band or category column"
ONLY = f"Only loans where {ch.ORIG_YEAR} is"
GRID = "FICO x CHANNEL"
LONE = "Kiosk"                 # a channel with loans in one origination year only: empty in every other year's view
BLANK_EVERY = 211              # every 211th loan's booked amount is blanked: left out of Avg line's top and bottom


def _write(rows, path):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return path


@pytest.fixture(scope="module")
def made(tmp_path_factory):
    """The synthetic book with ORIG_BAL blank on every 211th loan (under 1%, so Set up still reads it as the booked
    amount), and a Kiosk channel of 25 loans all made in the first origination year. Bands FICO, segments CHANNEL,
    filtered by origination year."""
    from pocketbook import perm
    d = tmp_path_factory.mktemp("avg_line")
    src = synth.write_extract(d / "src", n=3000)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    for i, r in enumerate(rows):
        if i % BLANK_EVERY == 5:
            r["ORIG_BAL"] = ""
    first = min(r["ORIG_DATE"][:4] for r in rows if r["ORIG_DATE"])
    kiosk = [r for r in rows if r["ORIG_DATE"][:4] == first and r["ORIG_BAL"]][:25]
    for r in kiosk:
        r["CHANNEL"] = LONE
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 100)
        x = _write(rows, d / "loans.csv")
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                                filter=ch.ORIG_YEAR, outcome="BAD_FLAG"))
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return out.book, rows


# --------------------------------------------------------------------------
# the arithmetic, from the CSV's text alone


def _bal(r):
    try:
        return float(r["ORIG_BAL"])
    except ValueError:
        return None


def _avg(rs):
    got = [b for b in (_bal(r) for r in rs) if b is not None]
    return sum(got) / len(got) if got else None


def _band(fico, labels) -> str:
    """A loan's FICO band from the tab's own labels ("496 - 653", both ends in the band): -9999 was answered
    missing, a blank is its own row."""
    if fico in ("", None):
        return "(blank)"
    v = float(fico)
    if v < -1000:
        return "(marked missing)"
    for lab in labels:
        lo, _, hi = lab.partition(" - ")
        if hi and float(lo) <= v <= float(hi):
            return lab
    raise KeyError(fico)


def _same(got, want):
    if want is None:
        return got in (None, "")
    return got == pytest.approx(want, rel=1e-9)


def _views(b, prefix):
    ws = load_workbook(b)[results.VIEWS]
    return {r[0]: list(r[1:]) for r in ws.iter_rows(values_only=True) if isinstance(r[0], str)
            and r[0].startswith(prefix)}


def _heads(b):
    """Summary's column headings as written, from Loans on."""
    ws = load_workbook(b)[results.SUMMARY]
    h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
    heads, c = [], 3
    while ws.cell(row=h, column=c).value not in (None, ""):
        heads.append(ws.cell(row=h, column=c).value)
        c += 1
    return heads


# --------------------------------------------------------------------------


def test_the_fixture_has_blank_lines_and_a_channel_missing_from_most_years(made):
    """The checks below lean on both: a blank booked amount makes Avg line's bottom differ from Loans, and the Kiosk
    is an empty row in every later year's view."""
    _, rows = made
    assert sum(1 for r in rows if _bal(r) is None) >= 10
    years = {r["ORIG_DATE"][:4] for r in rows if r["CHANNEL"] == LONE}
    assert len(years) == 1 and len({r["ORIG_DATE"][:4] for r in rows if r["ORIG_DATE"]}) >= 3


def test_summary_puts_avg_line_and_x_book_right_after_booked(made):
    b, _ = made
    heads = _heads(b)
    i = heads.index("Booked $")
    assert heads[i + 1:i + 3] == ["Avg line", "× book"]
    assert heads[i + 3] == "% of booked"


def test_every_summary_view_ties_avg_line_and_x_book_to_the_loan_file(made):
    """No LibreOffice: every Summary view the Run wrote, by CHANNEL and by FICO band, for All loans and each
    origination year. Avg line: the row's booked dollars over its loans with a booked amount; × book: that over the
    whole book's, in a filtered view too. A row with no loans is blank."""
    b, rows = made
    views = _views(b, "S|")
    heads = _heads(b)
    a = 1 + heads.index("Avg line")
    book_avg = _avg(rows)
    checked, filtered, empty = 0, 0, 0
    fico_labels = sorted({v[0] for k, v in views.items() if k.startswith("S|FICO|") and " - " in str(v[0])})
    assert fico_labels
    for k, v in views.items():
        _, col, *rest = k.split("|")
        if col not in ("CHANNEL", "FICO"):
            continue
        year = rest[0].removeprefix("where ") if len(rest) == 2 else None
        mine = [r for r in rows if year is None or r["ORIG_DATE"][:4] == year]
        lab = v[0]
        if col == "CHANNEL":
            rs = mine if lab == "All" else [r for r in mine if (r["CHANNEL"] or "(blank)") == lab]
        else:
            rs = mine if lab == "All" else [r for r in mine if _band(r["FICO"], fico_labels) == lab]
        want = _avg(rs)
        got_avg, got_x = (v[a] if len(v) > a else None), (v[a + 1] if len(v) > a + 1 else None)
        assert _same(got_avg, want), (k, got_avg, want)
        assert _same(got_x, want / book_avg if want is not None else None), (k, got_x)
        checked += 1
        filtered += year is not None
        empty += not rs
    assert checked > 40 and filtered > 30
    assert empty >= 2                       # the Kiosk in the later years: blank, not 0 and not an error


def test_x_book_on_the_all_row_is_one_unfiltered_and_the_years_own_average_filtered(made):
    b, rows = made
    views = _views(b, "S|CHANNEL|")
    a = 1 + _heads(b).index("Avg line")
    alls = {k: v for k, v in views.items() if v[0] == "All"}
    assert sum(k.count("|") == 2 for k in alls) == 1 and len(alls) >= 4
    for k, v in alls.items():
        if k.count("|") == 2:
            assert v[a + 1] == pytest.approx(1.0)
        else:
            year = k.split("|")[2].removeprefix("where ")
            assert v[a + 1] == pytest.approx(_avg([r for r in rows if r["ORIG_DATE"][:4] == year]) / _avg(rows))


def test_ranr_vs_gcos_every_pockets_avg_line_and_x_book_tie_to_the_loan_file(made):
    """No LibreOffice: every pocket of FICO x CHANNEL as the Run put it on _views, and the totals under the table."""
    b, rows = made
    views = _views(b, f"C|{GRID}|")
    pockets = {k: v for k, v in views.items() if k.split("|")[-1].isdigit() and "total" not in k}
    labels = sorted({str(v[0]) for v in pockets.values() if " - " in str(v[0])})
    book_avg = _avg(rows)
    assert len(pockets) > 10
    for k, v in pockets.items():
        band, seg = str(v[0]), str(v[1])
        rs = [r for r in rows if _band(r["FICO"], labels) == band and r["CHANNEL"] == seg]
        assert len(rs) == v[2], k
        want = _avg(rs)
        assert _same(v[11], want), (k, v[11], want)
        assert _same(v[12], want / book_avg), (k, v[12])
    totals = {v[0]: v for k, v in views.items() if "|total|" in k}
    whole = totals[results.PCK_BOOK]
    assert whole[6] == pytest.approx(book_avg, rel=1e-9) and whole[7] == pytest.approx(1.0)
    listed = [r for r in rows if (_band(r["FICO"], labels), r["CHANNEL"]) in
              {(str(v[0]), str(v[1])) for v in pockets.values()}]
    assert totals[results.PCK_LISTED][6] == pytest.approx(_avg(listed), rel=1e-9)


def test_ranr_vs_gcos_shows_them_in_the_gross_block_from_its_own_row(made):
    """As written: the headings, under the gross block's group heading; each cell picks its own row's _views entry,
    the one Avg line and × book sit at, so a sort moves them with the row; and the shading is Loan size's neutral
    hue on the row's own × book, never red or green."""
    b, _ = made
    ws = load_workbook(b)[results.PCK]
    head = tabs.header_row(ws, results.C_TOG, "Together")
    assert [ws.cell(row=head, column=c).value for c in (results.C_RATE, results.C_AVG, results.C_AVGX,
                                                         results.C_PAID)] == ["RANR ÷ Booked", "Avg line", "× book",
                                                                              "Gap pts"]
    group = [str(m) for m in ws.merged_cells.ranges if m.min_row == head - 1 and m.min_col == results.C_BOOK]
    assert group == [f"{results.col(results.C_BOOK)}{head - 1}:{results.col(results.C_AVGX)}{head - 1}"]
    for rr in (head + 1, head + 5):
        v = f"${results.col(results.C_H_ROW)}{rr}"
        for c, j in ((results.C_AVG, 12), (results.C_AVGX, 13)):
            f = ws.cell(row=rr, column=c).value
            assert f.startswith(f'=IF({v}="",""') and f.count(f",{v},{j})") == 2, (c, f)
    rng = f"{results.col(results.C_AVG)}{head + 1}"
    rules = [r for cr in ws.conditional_formatting if str(cr.sqref).startswith(rng) for r in cr.rules]
    fills = {r.dxf.fill.fgColor.rgb[-6:] for r in rules if r.dxf.fill is not None}
    assert fills == {c for _, c in results.SIZE_STEPS}
    assert all(f"${results.col(results.C_AVGX)}{head + 1}>=" in r.formula[0] for r in rules if r.dxf.fill is not None)
    assert all(r.dxf.font is None for r in rules)                 # no red or green text either
    lo, hi = ws.auto_filter.ref.split(":")
    assert column_index_from_string(lo.rstrip("0123456789")) <= results.C_AVG
    assert column_index_from_string(hi.rstrip("0123456789")) >= results.C_AVGX


# --------------------------------------------------------------------------
# as the analyst sees them: LibreOffice calculates the tab


def _summary(b, out, pick, only=None):
    picks = {PICKER: pick, **({ONLY: only} if only else {})}
    ws = tabs.calculated(tabs.choose(b, out, results.SUMMARY, **picks), results.SUMMARY)
    h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
    heads, c = [], 3
    while ws.cell(row=h, column=c).value not in (None, ""):
        heads.append(ws.cell(row=h, column=c).value)
        c += 1
    a = 3 + heads.index("Avg line")
    rows, r = {}, h + 1
    while ws.cell(row=r, column=2).value not in (None, ""):
        rows[ws.cell(row=r, column=2).value] = (ws.cell(row=r, column=a).value, ws.cell(row=r, column=a + 1).value)
        r += 1
    return rows


def test_summary_calculated_on_one_year_ties_to_the_loan_file(made, tmp_path):
    need_soffice()
    b, rows = made
    year = sorted({r["ORIG_DATE"][:4] for r in rows if r["ORIG_DATE"]})[1]
    mine = [r for r in rows if r["ORIG_DATE"][:4] == year]
    shown = _summary(b, tmp_path / "yr.xlsx", "CHANNEL", only=year)
    assert LONE in shown and shown[LONE] == (None, None)     # no loans in this year: blank, not 0 or #DIV/0!
    for lab, (avg, x) in shown.items():
        want = _avg(mine if lab == "All" else [r for r in mine if r["CHANNEL"] == lab])
        assert _same(avg, want), (lab, avg, want)
        assert _same(x, want / _avg(rows) if want is not None else None), (lab, x)


def test_ranr_vs_gcos_calculated_ties_to_the_loan_file(made, tmp_path):
    need_soffice()
    b, rows = made
    ws = tabs.calculated(tabs.choose(b, tmp_path / "pck.xlsx", results.PCK, grid=GRID), results.PCK)
    shown = tabs.pck(ws)
    labels = sorted({str(x["band"]) for x in shown if " - " in str(x["band"])})
    assert len(shown) > 10
    for x in shown:
        want = _avg([r for r in rows if _band(r["FICO"], labels) == str(x["band"]) and r["CHANNEL"] == x["seg"]])
        assert _same(x["avg_line"], want), x
        assert _same(x["avg_x"], want / _avg(rows)), x
