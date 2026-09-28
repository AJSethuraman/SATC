"""A test of a new variable builds no bleed analysis (OC-42; the firm, 26 Sep 2026, on BACKLOG 6d: "Yes. Seems
obvious I think. They have entirely different outputs generally").

A Run from a pre-spec runs the confirmatory test and what it needs: the pre-spec's strata cut into the bands the
Columns tab gives them, the development and holdout ranges, and 4e's concentration; then Check and the Log. It
builds no pocket grid, no three-way or split grid, and no shuffle test, and it writes none of the bleed's tabs. A
workbook that held them from an earlier bleed Run has them taken off, and Check says so in one line, so no tab
shows an old Run's pockets beside this Run's test.

Start here's "What the last Run found" and the launcher's last step read the confirmation, not nought pockets of
nought. A bleed Run is unchanged: it still builds its grids and runs one shuffle test for all of them.

Everything counted here is counted from the test's own groups by hand, never from the tiles' numbers."""

import pytest
from openpyxl import load_workbook

from pocketbook import book, confirm_tab, confirmatory, control, engine, launcher, perm, results
import tabs
from recalc import recalc
from test_book_dates import _check, _control
from test_confirmatory import _held, _log, _ready, _spec, needs_git

BLEED, NEW = "Where the book bleeds", "Finding and testing a new variable"
FROM_SPEC = "Test from a pre-spec"
BLEED_TABS = set(results.TABS)          # Pockets, Paid cost kept, Grids, Split (the redesign, phase 3)


def _counting(mp):
    """How many grids the engine builds and how many shuffle tests it runs, from here on."""
    seen = {"grids": 0, "shuffles": 0}
    build, shuffle = engine._build_grid, perm.run

    def built(*a, **k):
        seen["grids"] += 1
        return build(*a, **k)

    def shuffled(*a, **k):
        seen["shuffles"] += 1
        return shuffle(*a, **k)

    mp.setattr(engine, "_build_grid", built)
    mp.setattr(perm, "run", shuffled)
    return seen


@pytest.fixture(scope="module")
def runs(tmp_path_factory):
    """One workbook: a bleed Run first, so it holds the bleed's tabs, then a test of a new variable from a
    committed pre-spec on the same extract."""
    folder = tmp_path_factory.mktemp("new-variable")
    out = {}
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(folder / "memory.yaml"))
        x, b = _ready(folder, mp, n=20000)
        _spec(x.parent)
        _control(b, run_kind=BLEED)
        with pytest.MonkeyPatch.context() as spy:
            out["bleed_counts"] = _counting(spy)
            out["bleed"] = book.run(b)
        out["bleed_tabs"] = set(load_workbook(b).sheetnames)
        _control(b, run_kind=NEW, new_variable_step=FROM_SPEC)
        _held(b, "prespec.yaml")
        seen = {}
        real = confirmatory.state

        def state(*a, **k):
            seen["st"] = real(*a, **k)
            return seen["st"]

        with pytest.MonkeyPatch.context() as spy:
            out["counts"] = _counting(spy)
            spy.setattr(confirmatory, "state", state)
            out["ran"] = book.run(b)
    out.update(x=x, b=b, st=seen["st"], t=seen["st"].test)
    assert out["bleed"].ok, out["bleed"].lines
    assert out["ran"].ok, out["ran"].lines
    return out


def _worse_by_hand(t, confidence: float = 0.95) -> list[int]:
    """The groups going bad significantly more often than the reference on the holdout."""
    f = t.holdout.fit
    return [k for k in range(len(t.groups)) if k != t.ref and f.odds[k] is not None and f.odds[k] > 1
            and f.p[k] is not None and f.p[k] < round(1 - confidence, 12)]


# --------------------------------------------------------------------------
# What a test of a new variable builds, and doesn't


@needs_git
def test_a_new_variable_run_builds_no_grid_and_runs_no_shuffle_test(runs):
    assert runs["counts"] == {"grids": 0, "shuffles": 0}
    # the counter can see them: the bleed Run on the same workbook built its grids and ran one shuffle test
    assert runs["bleed_counts"]["grids"] > 0 and runs["bleed_counts"]["shuffles"] == 1


@needs_git
def test_a_new_variable_run_writes_no_bleed_tab_and_takes_off_the_ones_a_bleed_run_left(runs):
    assert BLEED_TABS <= runs["bleed_tabs"]
    tabs = set(load_workbook(runs["b"]).sheetnames)
    assert not tabs & BLEED_TABS, tabs & BLEED_TABS
    assert not tabs & {*results.HIDDEN, results.CHART}                  # nor the hidden sheets they read
    assert {confirm_tab.SHEET, "Record", "_log", "Start here"} <= tabs
    assert not tabs & {"Check", "Log", "Confirmatory test"}                # the tabs Record and New variables replaced


@needs_git
def test_check_says_so_in_one_line_and_says_nothing_of_pockets_or_tie_outs(runs):
    chk = _check(runs["b"])
    assert chk[book.NO_BLEED[0]] == book.NO_BLEED[1]
    for label in ("Tie-out checks", "Pockets tested", "Pocket budget", "Pockets in all", "Families of tests",
                  "Tests", "Decides each pocket", "Split"):
        assert label not in chk, label
    assert not [k for k in chk if str(k).startswith(("Worse now: ", "Pockets: ", "Loans needed for "))]
    # what the confirmatory test rests on is still there: the strata's band edges, and the pre-spec's lines
    assert "Band edges used: FICO" in chk and chk["Differs from the pre-spec"] == "nowhere: this run used what it says"
    assert "tie-out" not in _log(runs["b"])[0] and "tie-out" not in runs["ran"].lines[0]
    assert runs["ran"].lines[-1].endswith("start with New variables.")


# --------------------------------------------------------------------------
# What the last Run found, read from the confirmation


@needs_git
def test_the_launchers_last_step_reads_the_confirmation(runs):
    t, h = runs["t"], runs["ran"].summary
    worse = _worse_by_hand(t)
    assert worse, "the synthetic book plants two cliffs; the test should find at least one"
    assert h["kind"] == "confirm" and h["reference"] == t.groups[t.ref] and h["column"] == t.column
    assert (h["worse"], h["groups"]) == (len(worse), len(t.groups) - 1)
    assert h["capture"] == pytest.approx(sum(t.holdout.bad[k] for k in worse) / sum(t.holdout.bad), rel=1e-12)
    assert (h["development"], h["holdout"], h["deviations"]) == (t.development.n, t.holdout.n, 0)
    tiles = launcher.confirm_tiles(h)
    assert tiles[0][:2] == (f"{t.column} groups worse than {t.groups[t.ref]}",
                            f"{len(worse)} of {len(t.groups) - 1}")
    assert tiles[1][1] == f"{h['capture']:.0%}" and tiles[2][1] == "Yes"
    assert "Pockets worse and material" not in [x[0] for x in tiles]


def test_the_launcher_says_why_when_the_test_could_not_run():
    tiles = launcher.confirm_tiles({"kind": "confirm", "problem": "no column is marked Origination date"})
    assert tiles == (("Confirmatory test", "Not run", "no column is marked Origination date", "KEY_RED", "INK"),)


@needs_git
def test_start_here_shows_each_group_on_the_holdout_and_follows_the_confidence(runs, tmp_path):
    t = runs["t"]
    worse = _worse_by_hand(t)
    ws = recalc(runs["b"], tmp_path / "calc")["Start here"]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    at = {v: k for k, v in cells.items() if isinstance(v, str)}
    label = f"Groups worse than {t.groups[t.ref]}, on the holdout"
    r, c = at[label]
    assert cells[(r + 1, c)] == f"{len(worse)} of {len(t.groups) - 1}"
    share = cells[at["Their share of the holdout's bad loans"][0] + 1, at["Their share of the holdout's bad loans"][1]]
    assert share == pytest.approx(sum(t.holdout.bad[k] for k in worse) / sum(t.holdout.bad), rel=1e-9)
    assert cells[at["Follows the pre-spec"][0] + 1, at["Follows the pre-spec"][1]] == "Yes"
    # one row per group, the reference marked, each worse group saying so
    head = at[f"Group of {t.column}"][0]
    rows = {cells.get((head + 1 + k, 2)): head + 1 + k for k in range(len(t.groups))}
    assert f"{t.groups[t.ref]} (reference)" in rows
    for k, name in enumerate(t.groups):
        rr = rows[name + (" (reference)" if k == t.ref else "")]
        assert cells[(rr, 3)] == t.holdout.loans[k]
        if k != t.ref:
            assert cells[(rr, 7)] == ("Yes, worse" if k in worse else cells[(rr, 7)])
            assert (cells[(rr, 7)] == "Yes, worse") == (k in worse)
    assert "Nothing yet" not in " ".join(v for v in cells.values() if isinstance(v, str))

    # at 99% sure, a group whose p-value sits between 1% and 5% stops counting (OC-40: the lines are live)
    wb = load_workbook(runs["b"])
    from pocketbook import control
    cws = wb[control.SHEET]
    cws.cell(row=control.row_of(cws, "confidence"), column=control.CHOOSE_COL).value = "99%"
    changed = tmp_path / "changed.xlsx"
    wb.save(changed)
    ws = recalc(changed, tmp_path / "calc99")["Start here"]
    got = ws.cell(row=r + 1, column=c).value
    assert len(_worse_by_hand(t, 0.99)) != len(worse)       # this book has a group between the two bars
    assert got == f"{len(_worse_by_hand(t, 0.99))} of {len(t.groups) - 1}"


# --------------------------------------------------------------------------
# The engine, directly


def test_the_engine_builds_no_grid_for_a_new_variable_and_the_same_grids_as_before_for_the_bleed(book):
    from conftest import cube, table
    tb = table(book)
    bled = engine.run(cube(), tb)
    tested = engine.run(cube(run_kind="new_variable"), tb)
    assert bled.bleed and bled.grids and bled.tie_outs > 0
    assert not tested.bleed and tested.grids == [] and tested.three_way == [] and tested.tie_outs == 0
    # what the confirmatory test reads is the same either way: the bands' edges, the loans, the book's totals
    assert tested.band_edges == bled.band_edges and tested.rows == bled.rows
    assert {k: (v.num, v.den) for k, v in tested.total.rates.items()} == \
        {k: (v.num, v.den) for k, v in bled.total.rates.items()}


# --------------------------------------------------------------------------
# Control asks a new variable only what it uses (the redesign, phase 4)

#: what a test of a new variable reads on Control: the worse line (the New variables chart), materiality,
#: confidence, and how the held-fixed columns are banded
#: and, since the shortlist (OC-49), the allowance for many tests: its groups are many tests at once
USED = ("worse_at", "materiality", "confidence", "band_count", "band_cut", "many_tests")
UNUSED = ("min_loans", "min_events", "better_at", "compare_to", "power", "revenue_line")


def _dev_by_hand(x) -> tuple[list[int], list[int]]:
    """Development loans (2022-2023) per group of INCOME / SALES, and their bad loans, counted from the extract."""
    import bisect
    import csv as csv_
    from test_confirm_test import BINS
    loans, bad = [0] * 6, [0] * 6
    with open(x, newline="", encoding="utf-8") as fh:
        for r in csv_.DictReader(fh):
            if not "2022-01-01" <= r["ORIG_DATE"] <= "2023-12-31" or r["SALES"] in ("", "0") or \
                    r["BAD_FLAG"] not in ("0", "1"):
                continue
            g = bisect.bisect_right(BINS, float(r["INCOME"]) / float(r["SALES"]))
            loans[g] += 1
            bad[g] += int(r["BAD_FLAG"])
    return loans, bad


def _copied(runs, folder):
    """The fixture's workbook, extract and pre-spec, copied into a folder of the test's own."""
    import shutil
    folder.mkdir(exist_ok=True)
    for f in (runs["b"], runs["x"], runs["x"].parent / "prespec.yaml"):
        shutil.copy(f, folder / f.name)
    return folder / runs["b"].name, folder / runs["x"].name


@needs_git
def test_a_new_variable_run_is_asked_only_what_it_uses(runs, tmp_path):
    b, _ = _copied(runs, tmp_path / "copy")
    ws = load_workbook(b)[control.SHEET]
    hidden = {k for k in USED + UNUSED if ws.row_dimensions[control.row_of(ws, k)].hidden}
    assert hidden == set(UNUSED)
    # left blank, none of them is asked for: the Run goes ahead, and Record lists only what the run used
    wb = load_workbook(b)
    ws = wb[control.SHEET]
    for k in UNUSED:
        r = control.row_of(ws, k)
        ws.cell(row=r, column=control.CHOOSE_COL).value = None
        if ws.cell(row=r, column=control.OWN_COL).value != "n/a":
            ws.cell(row=r, column=control.OWN_COL).value = None
    wb.save(b)
    ran = book.run(b)
    assert ran.ok, ran.lines
    q = {s.key: s.question for s in control.load_settings()}
    assert sorted(tabs.settings(b)) == sorted(q[k] for k in USED)
    # a bleed workbook still asks every one of them
    bl = load_workbook(runs["b"])
    cws = bl[control.SHEET]
    cws.cell(row=control.row_of(cws, "run_kind"), column=control.CHOOSE_COL).value = BLEED
    control.fold_launcher_rows(cws)
    assert [k for k in USED + UNUSED if cws.row_dimensions[control.row_of(cws, k)].hidden] == []


@needs_git
def test_the_worse_line_is_suggested_from_the_confirmations_own_groups(runs):
    """The smallest odds ratio a group of typical size could call significant against the reference, on the loans
    the groups were found on: worked out here from the extract's counts, at 95% sure."""
    import math
    import statistics
    loans, bad = _dev_by_hand(runs["x"])
    t = runs["t"]
    p = sum(bad) / sum(loans)
    pq = p * (1 - p)
    want = round(statistics.median(math.exp(1.959963984540054 * math.sqrt(1 / (n * pq) + 1 / (loans[t.ref] * pq)))
                                   for k, n in enumerate(loans) if k != t.ref), 2)
    assert book.test_gap(t, 0.95) == want
    ws = load_workbook(runs["b"])[control.SHEET]
    said = ws.cell(row=control.row_of(ws, "worse_at"), column=book.SUGGEST_COL).value
    assert said == f"suggested: {want:.2f}x, from this extract at the last Run"


@needs_git
def test_a_new_variable_set_up_builds_no_grid_and_suggests_the_worse_line(runs, tmp_path, monkeypatch):
    """Found 27 Sep 2026: a pre-spec Set up worked its suggestions out from a bleed run of every grid, which a test
    of a new variable never shows; at 17,000 x 80 with every column cut that pass held 2.4 GB."""
    b, x = _copied(runs, tmp_path / "again")
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(tmp_path))
    with pytest.MonkeyPatch.context() as spy:
        seen = _counting(spy)
        out = book.set_up(x)
    assert out.ok and seen == {"grids": 0, "shuffles": 0}
    ws = load_workbook(b)[control.SHEET]
    said = ws.cell(row=control.row_of(ws, "worse_at"), column=book.SUGGEST_COL).value
    assert said == f"suggested: {book.test_gap(runs['t'], 0.95):.2f}x, from this extract"
