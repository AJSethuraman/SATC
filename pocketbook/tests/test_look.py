"""The Look tab (fix 3.8, the audit's item f: band edges were chosen blind; the redesign's section 4).

Set up shows each number column's spread and its likely code beside a live column chart; a number column
splitting the pockets adds a scatter of it against each band column (docs/statistics.md A9). Every number is
checked against a count made by hand from the CSV, not against look.py's own arithmetic. The live bars, the
edge lines and the one-load Run are in tests/test_answer_tabs.py."""

import csv

import pytest
import statistics

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, ScatterChart

from pocketbook import book, choices as ch, house, look, synth
from pocketbook.ingest import read_table
from test_book_results import _ready
from test_book_dates import _choose

#: the number columns Look draws: those that can be cut into bands. GCO_AMT and RANR_AMT are outcomes (H, the firm,
#: 27 Sep 2026)
NUMBER_COLUMNS = ["FICO", "ORIG_BAL", "REV_DEBT"]


def _csv(path, col) -> list[str]:
    with open(path, newline="", encoding="utf-8") as f:
        return [row[col].strip() for row in csv.DictReader(f)]


def _numbers(raw: list[str]) -> list[float]:
    out = []
    for v in raw:
        try:
            out.append(float(v))
        except ValueError:
            pass
    return out


def _blocks(ws) -> dict[str, dict]:
    """Each block by column: its row and its labelled lines as (value, share)."""
    out: dict[str, dict] = {}
    for c in ws["B"]:
        if c.value in NUMBER_COLUMNS and (c.fill.fgColor.rgb or "")[-6:] == house.INK and c.row >= look.FIRST:
            r = c.row
            lines = {ws.cell(row=r + k, column=2).value: (ws.cell(row=r + k, column=3).value,
                                                          ws.cell(row=r + k, column=4).value)
                     for k in range(1, look.R_MAX + 1)}
            fmt = {ws.cell(row=r + k, column=2).value: ws.cell(row=r + k, column=3).number_format
                   for k in range(1, look.R_MAX + 1)}
            out[c.value] = {"row": r, "lines": lines, "fmt": fmt}
    return out


def _title(chart) -> str:
    return chart.title.tx.rich.p[0].r[0].t


def _main_charts(ws):
    """The column charts beside the blocks (anchored at H), not the red code bars (at F)."""
    return [c for c in ws._charts if isinstance(c, BarChart) and c.anchor._from.col + 1 == 8]


def _dots(ds, header: str) -> list[tuple]:
    for col in range(1, ds.max_column + 1):
        if ds.cell(row=3, column=col).value == header:
            return [(ds.cell(row=r, column=col).value, ds.cell(row=r, column=col + 1).value)
                    for r in range(4, ds.max_row + 1) if ds.cell(row=r, column=col).value is not None]
    raise KeyError(header)


def test_set_up_writes_a_block_and_a_chart_for_every_number_column(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=3000))
    wb = load_workbook(out.book)
    assert wb.sheetnames.index("Look") == wb.sheetnames.index("Columns") + 1
    assert wb["_look"].sheet_state == "hidden"
    ws = wb["Look"]
    assert ws["B1"].value == "Look" and "before you choose its band edges" in ws["C1"].value
    blocks = _blocks(ws)
    # the loan number, the channel, the 0/1 flag and the four asset classes have no edges to choose
    assert list(blocks) == NUMBER_COLUMNS
    assert not [c for c in ws._charts if isinstance(c, ScatterChart)]                   # no split chosen
    assert sorted(c.anchor._from.row for c in _main_charts(ws)) == sorted(b["row"] for b in blocks.values())
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any(t.startswith("None: nothing splits the pockets.") for t in text)


def test_the_numbers_match_a_count_by_hand(tmp_path):
    x = synth.write_extract(tmp_path, n=3000)
    wb = load_workbook(book.set_up(x).book)
    blocks = _blocks(wb["Look"])

    fico = _csv(x, "FICO")
    blank = sum(1 for v in fico if v == "")
    nums = _numbers(fico)
    rest = [v for v in nums if v != -9999]
    f = blocks["FICO"]["lines"]
    assert f["Loans"][0] == 3000
    assert f["Blank"] == (blank, blank / 3000) and blank == 1
    assert f["At -9,999, likely a code"] == (60, 60 / 3000)               # every 50th loan
    assert f["Smallest"][0] == min(rest) and f["Largest"][0] == max(rest)
    assert f["Median"][0] == statistics.median(rest)
    assert f["Mean"][0] == statistics.fmean(rest)                        # the firm: the mean beside the median
    assert min(rest) > 300                                                # the code is out of the spread
    assert blocks["FICO"]["fmt"]["Median"] == "#,##0"                     # a score reads as a whole number
    # the slices hold every loan in the spread, and none of the code
    hs = wb["_look"]
    g = next(c for c in range(2, hs.max_column + 1) if hs.cell(row=look.DATA_TOP - 1, column=c).value == "FICO")
    counts = [hs.cell(row=look.DATA_TOP + i, column=g + look.G_COUNT).value for i in range(look.SLICES)]
    below = hs.cell(row=look.S_BELOW, column=g + look.G_VALUE).value
    above = hs.cell(row=look.S_ABOVE, column=g + look.G_VALUE).value
    assert sum(counts) + below + above == len(rest)

    bal = _csv(x, "ORIG_BAL")
    got = blocks["ORIG_BAL"]["lines"]
    vals = _numbers(bal)
    assert got["Blank"] == (1, 1 / 3000) and got["Likely a code"] == ("none found", None)
    assert (got["Smallest"][0], got["Median"][0], got["Largest"][0]) == \
        (min(vals), statistics.median(vals), max(vals))
    assert blocks["ORIG_BAL"]["fmt"]["Median"] == "#,##0"                 # dollars too
    assert look.shape_of(read_table(x), "GCO_AMT").text == 1              # the "#N/A", counted as not a number


def test_the_slices_sit_on_round_numbers_around_the_1st_and_99th_percentile():
    vals = sorted([500.0] + [600.0 + i % 250 for i in range(5000)] + [990.0])
    lo, hi = look.span(vals)
    assert (lo, hi) == (600.0, 850.0)
    counts, below, above = look.slices(vals, lo, hi)
    assert len(counts) == look.SLICES and below == 1 and above == 1 and sum(counts) == 5000
    bars, low, high = look.regroup(vals, 10, lo, hi, lo, hi)
    assert (low, high) == (1, 1) and sum(bars) == 5000 and len(bars) == 10


def test_a_run_split_by_revolving_debt_adds_a_scatter_against_each_band_column(tmp_path):
    b = _ready(tmp_path, n=3000)
    _choose(b, split="REV_DEBT")
    raw, _, _ = book.read_book(b)
    bands = [x["field"] for x in raw["bands"]]
    assert bands == ["FICO", "ORIG_BAL"]
    assert book.run(b).ok
    wb = load_workbook(b)
    assert wb.sheetnames.index("Look") == wb.sheetnames.index("Columns") + 1        # the Run kept its place
    ws, ds = wb["Look"], wb["_dots"]
    assert len(_main_charts(ws)) == len(NUMBER_COLUMNS)                              # the blocks' charts kept
    scatters = [c for c in ws._charts if isinstance(c, ScatterChart)]
    assert sorted(_title(c) for c in scatters) == [f"REV_DEBT against {x}" for x in bands]
    text = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)]
    assert any("random sample of 2,000" in t for t in text)                 # it says it's a sample
    assert "Moves together (correlation)" in text
    x = b.with_name("loans.csv")
    rev = _csv(x, "REV_DEBT")
    for band in bands:
        # every dot is a real loan's (band, REV_DEBT), at most 2,000 of them, and never the -9999 code
        dots = _dots(ds, band)
        loans = set(zip(_numbers_or_none(_csv(x, band)), _numbers_or_none(rev)))
        assert len(dots) == 2000 and all(d in loans for d in dots)
        assert all(d[0] != -9999 for d in dots)


def test_the_split_chosen_in_the_launcher_draws_its_scatters_at_set_up(tmp_path):
    out = book.set_up(synth.write_extract(tmp_path, n=2000), choices=ch.Choices(run_kind=ch.BLEED, split="REV_DEBT"))
    ws = load_workbook(out.book)["Look"]
    assert sorted(_title(c) for c in ws._charts if isinstance(c, ScatterChart)) == \
        ["REV_DEBT against FICO", "REV_DEBT against ORIG_BAL"]


def _numbers_or_none(raw):
    out = []
    for v in raw:
        try:
            out.append(float(v))
        except ValueError:
            out.append(None)
    return out


def test_the_sample_is_the_same_every_time(tmp_path):
    table = read_table(synth.write_extract(tmp_path, n=3000))
    got = []
    for _ in range(2):
        wb = Workbook()
        look.write_look(wb, table, NUMBER_COLUMNS, "REV_DEBT", ["FICO"])
        got.append(_dots(wb["_dots"], "FICO"))
    assert got[0] == got[1] and len(got[0]) == look.SAMPLE


def test_set_up_again_rebuilds_the_tab_without_a_copy(tmp_path):
    x = synth.write_extract(tmp_path, n=2000)
    book.set_up(x)
    out = book.set_up(x)
    wb = load_workbook(out.book)
    assert [n for n in wb.sheetnames if "look" in n.lower()] == ["Look", "_look"]
    assert len(_main_charts(wb["Look"])) == len(NUMBER_COLUMNS)
    assert wb.sheetnames.index("Look") == wb.sheetnames.index("Columns") + 1


@pytest.mark.parametrize("values", [[0.0, 0.05, 0.24, 0.5, 1.1999] * 50, [300, 512, 640, 849] * 50,
                                    [0.001, 0.07, 0.29] * 50, [0.3, 0.7, 2.1] * 50,
                                    [0.0, 1e-14, 2e-14] * 30])       # below any fixed rounding (Codex on #404)
def test_a_value_on_a_slice_edge_is_counted_in_that_slice(values):
    """The full tie-out of 28 Sep 2026: UTIL's chart counted every loan at exactly 0.24 a bar low (543 against the
    538 in the file), because the top came out 1.2000000000000002 and 0.24 / 0.006 is 39.99999999. The span's
    edges are the round numbers themselves, and a value on any slice's lower edge is counted in that slice."""
    lo, hi = look.span(sorted(values))
    assert hi > lo and lo == float(repr(lo)) and len(repr(hi)) <= 8     # round numbers, not 1.2000000000000002
    w = (hi - lo) / look.SLICES
    for k in range(look.SLICES):
        edge = float(f"{lo + k * w:.12g}")
        counts, below, above = look.slices([edge], lo, hi)
        assert counts.index(1) == k, (lo, hi, k, edge)


# ---- at the bank, 29 Sep 2026, evening. The firm: "Shouldn't we have SD markings on the look tab? Our FICO seems
# fairly distributed but other stuff is not"; offered SD or percentiles, they chose percentiles. And the labels
# under the bars: Excel wrapped 24,000 as "24,0/00"; it should read 24k.


def _by_hand(values: list[float], p: int) -> float:
    """The p-th percentile as Excel's PERCENTILE.INC works it out, written out here rather than taken from
    look.py or statistics: the value at position p/100 x (n - 1) of the sorted values, read between its two
    neighbours."""
    s = sorted(values)
    at = p / 100 * (len(s) - 1)
    k = int(at)
    return s[k] if k + 1 == len(s) else s[k] + (at - k) * (s[k + 1] - s[k])


def _group(hs, name) -> int:
    return next(c for c in range(2, hs.max_column + 1) if hs.cell(row=look.DATA_TOP - 1, column=c).value == name)


def _pct_rows(ws, r) -> dict[str, float]:
    return {ws.cell(row=r + look.R_PCT + k, column=2).value: ws.cell(row=r + look.R_PCT + k, column=3).value
            for k in range(len(look.PERCENTILES))}


def test_look_percentiles_match_a_count_by_hand_from_the_csv(tmp_path):
    x = synth.write_extract(tmp_path, n=3000)
    wb = load_workbook(book.set_up(x).book)
    ws, hs, blocks = wb["Look"], wb[look.DATA], _blocks(wb["Look"])
    fico = [v for v in _numbers(_csv(x, "FICO")) if v != -9999]      # the blank and the code left out
    bal = _numbers(_csv(x, "ORIG_BAL"))
    for name, vals in (("FICO", fico), ("ORIG_BAL", bal)):
        want = [_by_hand(vals, p) for p in (10, 25, 50, 75, 90)]
        r = blocks[name]["row"]
        got = _pct_rows(ws, r)
        assert list(got) == ["10th percentile (P10)", "25th percentile (P25)", "50th percentile (P50), the median",
                             "75th percentile (P75)", "90th percentile (P90)"]
        assert list(got.values()) == pytest.approx(want, abs=1e-9)
        assert got["50th percentile (P50), the median"] == pytest.approx(statistics.median(vals))
        assert ws.cell(row=r + look.R_PCT, column=3).number_format == blocks[name]["fmt"]["Median"]
        # the same five on _look, where the grey lines are placed from
        g = _group(hs, name)
        stored = [hs.cell(row=look.S_PCT + k, column=g + look.G_VALUE).value for k in range(5)]
        assert stored == pytest.approx(want, abs=1e-9)
        assert not ws.cell(row=r + look.BLOCK - 1, column=2).value                # a gap row under the list
    # numpy's default, where it is installed, agrees
    np = pytest.importorskip("numpy")
    assert [_by_hand(fico, p) for p in (10, 50, 90)] == pytest.approx(list(np.percentile(fico, [10, 50, 90])))


def test_look_percentiles_leave_out_what_columns_answered_missing(tmp_path):
    """A value answered missing on Columns is no loan on the chart, so it moves no percentile."""
    from pocketbook import config as cfgmod
    rows = [{"ID": str(i), "HIST": str(-99000901 - i % 4 if i % 10 == 0 else (i * 37) % 436)} for i in range(400)]
    p = tmp_path / "hist.csv"
    with open(p, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["ID", "HIST"])
        w.writeheader()
        w.writerows(rows)
    table = read_table(p)
    kept = [float(r["HIST"]) for r in rows if float(r["HIST"]) >= 0]
    wb = Workbook()
    look.write_look(wb, table, ["HIST"], rules={"HIST": cfgmod.MissingRule(below=0.0)})
    got = _pct_rows(wb["Look"], look.FIRST)
    assert list(got.values()) == pytest.approx([_by_hand(kept, q) for q in (10, 25, 50, 75, 90)])
    assert min(got.values()) >= 0
    # unanswered, the codes are used as recorded and pull the 10th down
    wb = Workbook()
    look.write_look(wb, table, ["HIST"])
    assert _pct_rows(wb["Look"], look.FIRST)["10th percentile (P10)"] < 0


def _look_book(tmp_path, n=3000, edit=None):
    """Look alone, on a bare workbook, calculated by LibreOffice after `edit(ws)` on the Look tab; and the
    workbook as written."""
    from recalc import values_of
    table = read_table(synth.write_extract(tmp_path, n=n))
    wb = Workbook()
    look.write_look(wb, table, NUMBER_COLUMNS)
    if edit:
        edit(wb["Look"])
    return table, values_of(wb, tmp_path / "calc"), wb


def _set(column, bars=None, frm=None, to=None):
    def edit(ws):
        r = look.FIRST + look.BLOCK * NUMBER_COLUMNS.index(column)
        for row, v in ((look.R_BARS, bars), (look.R_FROM, frm), (look.R_TO, to)):
            if v is not None:
                ws.cell(row=r + row, column=look.VALUE_COL).value = v
    return edit


def _grey(hs, name):
    """The grey lines as calculated: each one's x and y at its two points (None where no line), and the tallest
    bar."""
    g = _group(hs, name)
    num = lambda v: v if isinstance(v, (int, float)) else None     # noqa: E731
    xs = [num(hs.cell(row=look.PCT_TOP + k, column=g + look.G_EX).value) for k in range(2 * len(look.PERCENTILES))]
    ys = [num(hs.cell(row=look.PCT_TOP + k, column=g + look.G_EY).value) for k in range(2 * len(look.PERCENTILES))]
    return xs, ys, max(hs.cell(row=look.DATA_TOP + k, column=g + look.G_SLOT).value for k in range(look.SLOTS))


def _written_chart(wb, column):
    """A block's column chart, in a workbook not yet saved (its anchor still a cell name)."""
    at = f"{look.CHART_AT}{look.FIRST + look.BLOCK * NUMBER_COLUMNS.index(column) + 1}"
    return next(c for c in wb["Look"]._charts if isinstance(c, BarChart) and c.anchor == at)


def _rgb(fill) -> str:
    rgb = fill if isinstance(fill, str) else fill.srgbClr
    return rgb if isinstance(rgb, str) else rgb.val


def _at(p, frm, to):
    """Where the red lines put a value (tests/test_answer_tabs.py holds them to it): worked out here."""
    return look.LOW + 0.5 + look.MIDDLE * (p - frm) / (to - frm)


def test_look_percentile_lines_sit_where_the_edge_lines_would_and_follow_from_and_to(tmp_path):
    table, got, wb = _look_book(tmp_path / "a")
    fico = look.shape_of(table, "FICO")
    lo, hi = look.span(fico.values)
    xs, ys, tallest = _grey(got[look.DATA], "FICO")
    for k, p in enumerate(fico.pcts):
        assert lo <= p <= hi
        assert xs[2 * k] == pytest.approx(_at(p, lo, hi)) and xs[2 * k + 1] == pytest.approx(_at(p, lo, hi))
        assert (ys[2 * k], ys[2 * k + 1]) == (0, tallest)              # from the floor to the tallest bar
    # the chart draws them: five thin grey lines named P10 to P90, each named at its top point only, before the red
    main = _written_chart(wb, "FICO")
    grey = main._charts[1].series[:len(look.PERCENTILES)]
    assert [s.tx.v for s in grey] == ["P10", "P25", "P50", "P75", "P90"]
    for s in grey:
        assert _rgb(s.graphicalProperties.line.solidFill) == house.SLATE and s.graphicalProperties.line.prstDash in (None, "solid")
        assert [(d.idx, d.showSerName) for d in s.dLbls.dLbl] == [(1, True)] and not s.dLbls.showSerName
    red = main._charts[1].series[len(look.PERCENTILES):]
    assert len(red) == look.EDGE_LINES and all(s.graphicalProperties.line.prstDash == "dash" for s in red)
    # a narrower range: the 10th and the 90th fall outside it and aren't drawn; the rest move to the new scale
    frm, to = fico.pcts[0] + 5, fico.pcts[4] - 5
    _, got, _ = _look_book(tmp_path / "b", edit=_set("FICO", bars=10, frm=frm, to=to))
    xs, ys, _ = _grey(got[look.DATA], "FICO")
    assert xs[0:2] == [None, None] and ys[0:2] == [None, None]
    assert xs[8:10] == [None, None] and ys[8:10] == [None, None]
    for k in (1, 2, 3):
        assert xs[2 * k] == pytest.approx(_at(fico.pcts[k], frm, to)) and ys[2 * k] == 0


def test_look_labels_under_the_bars_are_short(tmp_path):
    """The labels as LibreOffice calculates them: FICO keeps its own (550), dollars read 12k, a step of 500
    keeps its half (20.5k), millions read 1.2M, and a negative keeps its sign."""
    def labels(got, name):
        g = _group(got[look.DATA], name)
        return [v for v in (got[look.DATA].cell(row=look.DATA_TOP + k, column=g + look.G_SLOT_LABEL).value
                            for k in range(look.SLOTS)) if v not in (None, "")]

    table, got, wb = _look_book(tmp_path / "a")
    lo, hi = look.span(look.shape_of(table, "FICO").values)
    assert hi < 1000 and labels(got, "FICO") == [f"{lo + k * (hi - lo) / 5:,.0f}" for k in range(5)]
    blo, bhi = look.span(look.shape_of(table, "ORIG_BAL").values)
    starts = [blo + k * (bhi - blo) / 5 for k in range(5)]
    assert bhi >= 10000 and labels(got, "ORIG_BAL") == [f"{v / 1000:g}k" if v else "0" for v in starts]
    assert "24k" in labels(got, "ORIG_BAL")                                # the bank's 24,000
    # the axis sets them flat and never wraps them (Excel's "Wrap text in shape" off), in the file as saved
    import re
    import zipfile
    with zipfile.ZipFile(tmp_path / "a" / "calc" / "book.xlsx") as z:
        charts = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/charts/chart")]
    cat_axes = [m for c in charts for m in re.findall(r"<(?:c:)?catAx>.*?</(?:c:)?catAx>", c, re.S)
                if f'tickMarkSkip val="{look.MIDDLE}"' in m]                      # not the red code bars'
    assert len(cat_axes) == len(NUMBER_COLUMNS)
    assert all(re.search(r'<a:bodyPr[^>]*rot="0"[^>]*wrap="none"', a) for a in cat_axes)

    def several(ws):
        _set("ORIG_BAL", bars=10, frm=20000, to=22500)(ws)
        _set("REV_DEBT", bars=10, frm=0, to=3000000)(ws)
        _set("FICO", bars=10, frm=-30000, to=0)(ws)
    _, got, _ = _look_book(tmp_path / "b", edit=several)
    assert labels(got, "ORIG_BAL") == ["20k", "20.5k", "21k", "21.5k", "22k"]
    assert labels(got, "REV_DEBT") == ["0", "600k", "1.2M", "1.8M", "2.4M"]
    assert labels(got, "FICO") == ["-30k", "-24k", "-18k", "-12k", "-6k"]


def test_look_axis_formats_read_short_in_libreoffice(tmp_path):
    """The scatters' axis formats, as a spreadsheet shows them: when every tick is a whole thousand it reads 24k,
    a millions tick 1.5M; a score keeps its own format."""
    from recalc import values_of
    cases = [(700, 0, 50, "700"), (1500, 0, 500, "1,500"), (0, 0, 1000, "0"), (24000, 0, 1000, "24k"), (1500000, 0, 500000, "1.5M"),
             (200000, 0, 100000, "200k"), (-24000, -30000, 2000, "-24k"), (-3000000, -4000000, 1000000, "-3M")]
    wb = Workbook()
    ws = wb.active
    for i, (v, lo, unit, _) in enumerate(cases, start=1):
        f = look.axis_format(lo, unit, "#,##0").replace('"', '""')
        ws.cell(row=i, column=1, value=f'=TEXT({v},"{f}")')
    got = values_of(wb, tmp_path)
    assert [got.active.cell(row=i, column=1).value for i in range(1, len(cases) + 1)] == [c[3] for c in cases]
