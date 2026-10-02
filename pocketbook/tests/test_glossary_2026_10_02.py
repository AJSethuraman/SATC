"""The Glossary tab (the firm, 2 Oct 2026, after explaining the grids' figures to their boss: "Maybe a nice
glossary of terms in the workbook should be there"; they keep "points" as the unit).

Written by every Set up and every Run, right after Start here, one term a row: the term, what it means in one or
two plain sentences, and an example in this book's own figures from the last Run (made up before the first Run).
The examples are tied here to the extract itself, worked out from the CSV without PocketBook."""

from __future__ import annotations

import csv
import re

import pytest
from openpyxl import load_workbook

from conftest import TEST_SHUFFLES
from pocketbook import book, choices as ch, glossary, house, perm, synth
from test_book import _answer

#: every term, in the firm's order (2 Oct 2026)
TERMS = ["Booked / Booked $", "GCOs ($)", "RANR", "RANR + GCOs", "Rate (GCOs ÷ Booked, RANR ÷ Booked)",
         "Points (pts)", "× book / × rest of book", "Rest of book / rest of band", "Bad loan", "Band", "Segment",
         "Pocket", "Grid", "Filter / Only loans where", "Origination year", "Worse and material", "Borderline",
         "Shuffle test / p-value", "Fewest loans in a pocket", "Worse at / Better at",
         "Dollars above their share (each loan once)", "Lifetime-to-date", "Months to charge-off", "Odd values",
         "(marked missing) / (blank)"]
#: the contract-desk words website/copy.spec.py refuses on every page (CLAUDE.md, client-facing copy)
CONTRACT_WORDS = ["governs", "governed by", "constitutes", "in accordance with", "pursuant", "herein", "thereof",
                  "aforementioned", "accompanies", "at our discretion", "in the event that", "utilize", "commence",
                  "deemed", "whereupon", "notwithstanding", "shall be", "hereby", "aforesaid"]
MOST_WORDS = 28


@pytest.fixture(scope="module")
def books(tmp_path_factory):
    """The synthetic book: the workbook after Set up, after a Run, and after Set up again."""
    d = tmp_path_factory.mktemp("glossary")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
        x = synth.write_extract(d, n=3000)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED))
        assert out.ok, out.lines
        set_up = load_workbook(out.book)
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
        after_run = load_workbook(out.book)
        again = book.set_up(x)
        assert again.ok, again.lines
        return {"extract": x, "set_up": set_up, "run": after_run, "again": load_workbook(again.book)}


def _rows(wb) -> dict[str, tuple[str, str]]:
    ws = wb[glossary.SHEET]
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=glossary.TERM).value == "Term")
    return {ws.cell(row=r, column=glossary.TERM).value: (ws.cell(row=r, column=glossary.MEANS).value,
                                                        ws.cell(row=r, column=glossary.EXAMPLE).value)
            for r in range(head + 1, ws.max_row + 1) if ws.cell(row=r, column=glossary.TERM).value}


def _num(v: str) -> float | None:
    try:
        return float(v.replace(",", ""))
    except (AttributeError, ValueError):
        return None


def _from_csv(path) -> dict:
    """The whole book's figures, worked out from the extract alone: a loan enters a rate only when both its
    amounts read as numbers, as PocketBook counts them (a blank or text amount is left out, never a zero)."""
    booked = gco = gco_den = ranr = ranr_den = 0.0
    with open(path, newline="", encoding="utf-8") as f:
        for r in csv.DictReader(f):
            bal, g, k = _num(r["ORIG_BAL"]), _num(r["GCO_AMT"]), _num(r["RANR_AMT"])
            if bal is None:
                continue
            booked += bal
            if g is not None:
                gco, gco_den = gco + g, gco_den + bal
            if k is not None:
                ranr, ranr_den = ranr + k, ranr_den + bal
    return {"booked": booked, "gco_per100": gco / gco_den * 100, "ranr_per100": ranr / ranr_den * 100,
            "ranr": ranr}


# --------------------------------------------------------------------------


def test_the_glossary_sits_right_after_start_here_after_set_up_a_run_and_set_up_again(books):
    for when in ("set_up", "run", "again"):
        wb = books[when]
        assert wb.sheetnames[:2] == ["Start here", glossary.SHEET], when
        ws = wb[glossary.SHEET]
        assert ws.sheet_state == "visible" and ws.sheet_properties.tabColor.rgb[-6:] == house.TAB_RECORD, when
        assert ws["B1"].value == glossary.SHEET and ws["B3"].value == "How this tab works", when


def test_every_glossary_term_is_there_once_in_the_firms_order(books):
    for when in ("set_up", "run"):
        got = list(_rows(books[when]))
        assert got == TERMS, when
        for term, (means, example) in _rows(books[when]).items():
            assert means and example, (when, term)


def test_glossary_examples_are_made_up_before_the_first_run_and_the_note_says_so(books):
    note = [c.value for row in books["set_up"][glossary.SHEET].iter_rows(max_col=3) for c in row]
    assert glossary.MADE_UP in note
    ranr = _rows(books["set_up"])["RANR"][1]
    assert "$4.45 per $100 booked" in ranr and not ranr.startswith("This book")


def test_glossary_ranr_and_gco_examples_tie_to_the_extract(books):
    """The firm's example, "the book keeps $4.45 per $100 booked", in this book's own figures: the whole book's
    RANR over its booked dollars x 100, worked out from the CSV."""
    want = _from_csv(books["extract"])
    rows = _rows(books["run"])
    assert f"The book keeps ${want['ranr_per100']:.2f} per $100 booked." in rows["RANR"][1], rows["RANR"]
    assert f"This book's RANR came to ${want['ranr']:,.0f}." in rows["RANR"][1]
    assert f"${want['gco_per100']:.2f} per $100 booked." in rows["GCOs ($)"][1], rows["GCOs ($)"]
    assert f"is {want['gco_per100']:.2f}%: ${want['gco_per100']:.2f} lost per $100 booked." in \
        rows["Rate (GCOs ÷ Booked, RANR ÷ Booked)"][1]
    assert f"This book booked ${want['booked']:,.0f} over" in rows["Booked / Booked $"][1]
    note = [c.value for row in books["run"][glossary.SHEET].iter_rows(max_col=3) for c in row]
    assert any(isinstance(v, str) and v.startswith("This book's own figures, for the whole book, from the last Run")
               for v in note)


def test_glossary_points_and_multiples_add_up_as_printed(books):
    """The Points example's gap is its two rates' difference, and × book is the pocket's rate over the book's."""
    rows = _rows(books["run"])
    pts = rows["Points (pts)"][1]
    m = re.fullmatch(r".+ keeps \$(-?[\d.]+) per \$100 booked; the rest of the book keeps \$(-?[\d.]+)\. "
                     r"That is ([+-][\d.]+) pts\.", pts)
    assert m, pts
    a, b, gap = (float(x) for x in m.groups())
    assert abs((a - b) - gap) <= 0.011
    x = rows["× book / × rest of book"][1]
    m = re.fullmatch(r".+ loses \$([\d.]+) per \$100 booked: ([\d.]+)× book and ([\d.]+)× rest of book\.", x)
    assert m, x
    book_rate = _from_csv(books["extract"])["gco_per100"]
    assert abs(float(m.group(1)) / book_rate - float(m.group(2))) < 0.01


def test_set_up_again_keeps_the_last_runs_figures_in_the_glossary(books):
    assert _rows(books["again"]) == _rows(books["run"])


def test_start_here_links_to_the_glossary(books):
    for when in ("set_up", "run"):
        ws = books[when]["Start here"]
        links = [c for row in ws.iter_rows() for c in row if c.hyperlink]
        to = f"'{glossary.SHEET}'!A1"
        hit = [c for c in links if (c.hyperlink.location or c.hyperlink.target or "").lstrip("#") == to]
        assert len(hit) == 1, when
        assert str(hit[0].value).startswith(f"{glossary.SHEET}: "), hit[0].value


def test_each_glossary_sentence_is_short_and_plain(books):
    """CLAUDE.md's client-facing rules 3 and 5, as website/copy.spec.py runs them: no contract-desk word, and no
    sentence past about 28 words, in what each term means and in its example, before and after a Run."""
    for when in ("set_up", "run"):
        for term, cells in _rows(books[when]).items():
            for text in cells:
                low = text.lower()
                assert not [w for w in CONTRACT_WORDS if w in low], (term, text)
                for s in re.split(r"(?<=[.?!])\s+(?=[A-Z0-9($-])", text):
                    assert len(s.split()) <= MOST_WORDS, (term, s)
