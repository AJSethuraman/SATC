"""The firm's own terms for the three dollar measures, 30 Sep 2026:

    "let's rename this list of stuff for a couple things and be more literal - Charge-offs = GCOs ($), kept after
    losses = RANR, earned before losses = RANR + GCOs"

and on Grids, asking for the simpler ratio under each rate: "Like if I'm viewing charge-offs i can view the
COs/booked and same with the RANR". The Rate block already IS GCOs over booked, so its heading says so, and follows
the Measure dropdown. Then: "I want to use the terms I gave you out of the box so it can be understood by
insiders", so the tab "Paid, cost, kept" is "RANR vs GCOs".

Only the words moved. The keys (gco_rate, ranr_rate, contribution_rate) are what a workbook, a pre-spec and memory
hold, so a workbook set up and run before the change still Runs; that is held here on one the old code wrote."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

import tabs
from pocketbook import book, perm, results, synth
from recalc import SOFFICE
from test_book import _answer
from test_book_dates import _choose

pytestmark = pytest.mark.skipif(SOFFICE is None, reason="LibreOffice (soffice) isn't installed, so the "
                                                         "workbook's formulas can't be calculated here")

ROOT = Path(__file__).resolve().parents[1]
#: a workbook the code before the change wrote and ran (the evening tie-out, 30 Sep 2026), and its extract
OLD = ROOT / "docs" / "tie-out" / "2026-09-30-evening"
OLD_BOOK, OLD_EXTRACT = OLD / "Bureau book - PocketBook.xlsx", OLD / "Bureau book.csv"

#: the three measures' old names, and the stand-ins for them the tabs used; none may name a measure any more
OLD_NAMES = ("Charge-offs", "Kept after losses", "Earned before losses", "Paid us", "Cost us", "Paid, cost, kept")
#: the same, in any case, read inside a sentence ("pockets have charge-offs above their share")
OLD_WORDS = re.compile(r"charge-off|charged off|charges off|kept after losses|earned before losses|paid us|cost us"
                       r"|paid, cost, kept|profit after losses|contribution before losses|Kept ·|keeping (less|more)",
                       re.I)
NEW = {"gco_rate": "GCOs ($)", "ranr_rate": "RANR", "contribution_rate": "RANR + GCOs"}
#: what the Rate block's heading says for each Measure option, literally: what is divided by what
HEADS = {"Bad loans": "Rate · Bad loans ÷ Loans", "Bad dollars": "Rate · Bad dollars ÷ Booked",
         "GCOs ($)": "Rate · GCOs ÷ Booked", "RANR": "Rate · RANR ÷ Booked",
         "RANR + GCOs": "Rate · (RANR + GCOs) ÷ Booked", "Loan size": "Rate · Average booked per loan"}
#: sheets the analyst never reads that still hold words, each with why its words may stay: the meanings'
#: definitions ("gross charge-off dollars" says what GCO is, and is no measure's name)
ALLOWED = {"_meanings"}


@pytest.fixture(scope="module")
def ran(tmp_path_factory):
    d = tmp_path_factory.mktemp("literal")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 500)
        extract = synth.write_extract(d, n=3000)
        b = book.set_up(extract).book
        _answer(b)
        _choose(b, split="REV_DEBT")
        assert book.run(b).ok
    return b


def _strings(wb, sheets=None):
    """Every string the workbook holds, (sheet, cell, text), on the sheets named or every sheet."""
    for ws in wb.worksheets:
        if sheets is not None and ws.title not in sheets:
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    yield ws.title, c.coordinate, c.value


def old_names_in(wb) -> list[str]:
    """Where a built workbook still names a measure or the tab the old way: every sheet but ALLOWED, the hidden
    ones included, since _choices is what the dropdowns list and _found what Start here's tiles say."""
    return [f"{s}!{c}: {v[:120]}" for s, c, v in _strings(wb) if s not in ALLOWED and OLD_WORDS.search(v)]


def test_no_tab_names_a_measure_or_the_tab_the_old_way(ran):
    wb = load_workbook(ran)
    assert results.OLD_PCK not in wb.sheetnames and results.PCK in wb.sheetnames
    assert old_names_in(wb) == []
    # and the exceptions are what they say they are: definitions, never a measure's name
    kept = [v for s, _, v in _strings(wb, ALLOWED) if OLD_WORDS.search(v)]
    assert all(not any(n in v for n in OLD_NAMES) for v in kept), kept


def test_every_list_offers_the_new_names_and_every_tab_uses_them(ran):
    wb = load_workbook(ran)
    five = ["Bad loans", "Bad dollars", *NEW.values()]
    assert tabs.options(wb, results.POCKETS, "Measure") == five
    assert tabs.options(wb, results.GRIDS, "Measure") == five + ["Loan size"]
    assert tabs.options(wb, results.SPLIT, "Measure") == five
    said = {s: " ".join(v for t, _, v in _strings(wb) if t == s) for s in wb.sheetnames}
    for name in NEW.values():
        assert f"Worse now: {name}" in said["Record"] and f"Materiality line: {name}" in said["Record"], name
    for name in ("RANR + GCOs", "GCOs", "RANR"):              # RANR vs GCOs heads its sides with them
        assert name in said[results.PCK]
    assert "GCOs ÷ Booked" in said[results.SUMMARY] and "RANR ÷ Booked" in said[results.SUMMARY]
    ws = wb[results.PCK]
    head = tabs.header_row(ws, results.C_TOG, "Together")
    groups = [ws.cell(row=head - 1, column=c).value for c in (results.C_PAID, results.C_COST, results.C_KEPT)]
    assert [g.split(" · ")[0].lstrip('="') for g in groups] == ["RANR + GCOs", "GCOs", "RANR"]
    # the tab list on Start here links the new name
    start = wb["Start here"]
    assert any(results.PCK in f"{c.hyperlink.location or ''}{c.hyperlink.target or ''}"
               for row in start.iter_rows() for c in row if c.hyperlink is not None)


def _rate_head(ws) -> str | None:
    return next((c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                 and c.value.startswith("Rate · ")), None)


def test_the_rate_heading_says_what_is_divided_by_what_and_follows_the_measure(ran, tmp_path):
    """Picked on a copy and calculated, as the analyst would: the heading changes with the Measure, and the vs the
    book heading still carries the book's own figure."""
    wb = load_workbook(ran)
    for pick in tabs.options(wb, results.GRIDS, "Measure"):
        ws = tabs.calculated(tabs.choose(ran, tmp_path / f"{len(pick)}-{pick[:3]}.xlsx", results.GRIDS,
                                         measure=pick), results.GRIDS)
        assert _rate_head(ws) == HEADS[pick], pick
        heads = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                 and c.value.startswith("vs the book (")]
        assert len(heads) == 1 and re.fullmatch(r"vs the book \(book: (\$[\d,]+|\d+\.\d\d%)\)", heads[0]), \
            (pick, heads)


def test_the_one_cell_panel_names_the_measure_in_the_new_words(ran, tmp_path):
    for pick, word in (("GCOs ($)", ": GCOs were "), ("RANR", ": RANR was "), ("RANR + GCOs", ": RANR + GCOs came to ")):
        ws = tabs.calculated(tabs.choose(ran, tmp_path / f"one-{len(pick)}.xlsx", results.GRIDS, measure=pick),
                             results.GRIDS)
        top = next(c.row for row in ws.iter_rows() for c in row if c.value == "What one cell says")
        said = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(top + 1, top + 7)}
        assert word in said["Rate"] and said["Rate"].endswith(" booked dollars."), said["Rate"]
        assert not OLD_WORDS.search(" ".join(str(v) for v in said.values())), said


def test_a_workbook_the_old_code_wrote_still_runs_and_comes_out_in_the_new_words(tmp_path, monkeypatch):
    """Found in the tie-out folder, written and run by the code before the change: its Measure dropdowns set to the
    old names, its _found and Record in them, and a Paid, cost, kept tab. A Run reads its keys, not its words, so it
    runs; and it takes the old tab off rather than leave two."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 500)
    b, x = tmp_path / OLD_BOOK.name, tmp_path / OLD_EXTRACT.name
    shutil.copy(OLD_BOOK, b)
    shutil.copy(OLD_EXTRACT, x)
    wb = load_workbook(b)
    assert results.OLD_PCK in wb.sheetnames and results.PCK not in wb.sheetnames       # it is the old shape
    assert old_names_in(wb)
    # the analyst had picked the old names on the tabs, as the dropdowns offered them then
    tabs.dropdown(wb[results.GRIDS], "Measure").value = "Charge-offs"
    tabs.dropdown(wb[results.POCKETS], "Measure").value = "Kept after losses"
    wb.save(b)
    # Set up again first (a new extract, the same workbook): Start here is written from _found, which still says
    # "charge-offs", and reads in the new words at once
    assert book.set_up(x, b).book == b
    start = load_workbook(b)["Start here"]
    tiles = [v for _, _, v in _strings(start.parent, {"Start here"})]
    assert "Pockets worse and material, GCOs" in tiles and not any("charge-offs" in v for v in tiles)
    out = book.run(b, x)
    assert out.ok, out.lines
    wb = load_workbook(b)
    assert results.OLD_PCK not in wb.sheetnames and results.PCK in wb.sheetnames
    assert old_names_in(wb) == []
    assert tabs.dropdown(wb[results.GRIDS], "Measure").value in tabs.options(wb, results.GRIDS, "Measure")
