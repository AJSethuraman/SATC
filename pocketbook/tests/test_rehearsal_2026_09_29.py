"""Defects found running PocketBook on public loan data (docs/rehearsal-public-data-2026-09.md), each held here on a
tiny synthetic book: no public row is read.

1. A category limit the launcher's list doesn't offer (Choices(many_values=60): State has 51 values and a blank, one
   over the usual 50) was written into Control's pick cell, where only a listed option is read. The Run then
   refused "60 is not an option" on a row the analyst can't edit (it is the launcher's), after Set up had already
   used 60; and reading the block back (Set up again, the launcher) quietly took the usual 50 instead.
"""

from openpyxl import load_workbook

from pocketbook import book, choices as ch, control, synth
from test_book import PICK


def _limits_book(tmp_path, **limits):
    extract = synth.write_extract(tmp_path, n=2000)
    c = ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",), **limits)
    out = book.set_up(extract, choices=c)
    assert out.ok, out.lines
    return out.book


def _answer(b):
    wb = load_workbook(b)
    for r in wb[control.SHEET].iter_rows(min_row=control.FIRST_ROW):
        k = r[control.KEY_COL - 1].value
        if k in PICK:
            r[control.CHOOSE_COL - 1].value = PICK[k]
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)


def test_a_limit_off_the_launchers_list_is_read_back_as_chosen(tmp_path):
    b = _limits_book(tmp_path, few_values=3, many_values=60)
    got, _ = control.read_choices(load_workbook(b)[control.SHEET])
    assert (got.few_values, got.many_values) == (3, 60)


def test_a_limit_off_the_launchers_list_runs(tmp_path):
    b = _limits_book(tmp_path, many_values=60)
    _answer(b)
    _, problems, _ = book.read_book(b)
    assert not [p for p in problems if "not an option" in p], problems
    ran = book.run(b)
    assert ran.ok, ran.lines


def test_a_listed_limit_is_still_picked_from_the_list(tmp_path):
    b = _limits_book(tmp_path, few_values=24, many_values=100)
    ws = load_workbook(b)[control.SHEET]
    row = {k: control.row_of(ws, k) for k in ("few_values", "many_values")}
    assert ws.cell(row=row["few_values"], column=control.CHOOSE_COL).value == "24 values"
    assert ws.cell(row=row["many_values"], column=control.CHOOSE_COL).value == "100 values"
    assert ws.cell(row=row["many_values"], column=control.OWN_COL).value is None


# --------------------------------------------------------------------------
# 2. The confirmatory test's conditional likelihood took minutes per pocket at SBA scale: kgroups.pocket_terms
#    multiplied every group's polynomial in full with np.convolve, though almost every coefficient of a tilted
#    binomial is exactly 0.0 (it underflows) and only the coefficient at m is ever read. One evaluation of a
#    pocket of 217,800 loans took 191 s; the pre-registered test on the FOIA file was stopped after 10 minutes
#    inside its first fit. Now each polynomial carries only its nonzero coefficients, and a coefficient that is
#    read once is one dot product: the same sums, without the zeros.

import math
import time

import numpy as np

from pocketbook import kgroups


def _old_pocket_terms(p, b):
    """kgroups.pocket_terms as it was before 29 Sep 2026, kept here as the reference the new one must equal."""
    K = len(p.loans)
    R = p.loans.astype(int)
    m = int(round(p.m))
    lam = kgroups._tilt(p.loans, b, m)
    mul = lambda a, c: np.convolve(a, c)[: m + 1]

    def prod(ps):
        out = np.ones(1)
        for q in ps:
            out = mul(out, q)
        return out
    coef = lambda poly: float(poly[m]) if m < len(poly) else 0.0
    f, g, h = [], [], []
    for k in range(K):
        top = min(R[k], m)
        j = np.arange(top + 1, dtype=float)
        x = b[k] + lam
        log1p = math.log1p(math.exp(-abs(x))) + max(x, 0.0)
        pmf = np.exp(kgroups._log_choose(R[k], top) + j * x - R[k] * log1p)
        f.append(pmf), g.append(j * pmf), h.append(j * j * pmf)
    B = coef(prod(f))
    logD = math.log(B) + sum(float(R[k]) * (math.log1p(math.exp(-abs(b[k] + lam))) + max(b[k] + lam, 0.0))
                             for k in range(K)) - m * lam
    ll = float((p.bad * b).sum()) - logD
    others = [prod([f[l] for l in range(K) if l != k]) for k in range(K)]
    E = np.array([coef(mul(g[k], others[k])) / B for k in range(K)])
    cov = np.zeros((K, K))
    for k in range(K):
        cov[k, k] = coef(mul(h[k], others[k])) / B - E[k] ** 2
        for l in range(k + 1, K):
            rest = prod([f[q] for q in range(K) if q not in (k, l)])
            cov[k, l] = cov[l, k] = coef(mul(mul(g[k], g[l]), rest)) / B - E[k] * E[l]
    return ll, p.bad - E, -cov


def _pockets():
    yield kgroups.Pocket([4000, 4100, 3900, 4000, 3800], [800, 820, 700, 560, 460])
    yield kgroups.Pocket([12, 3000, 7, 2500], [1, 400, 7, 300])              # a group with every loan bad
    yield kgroups.Pocket([50, 60, 0, 40], [0, 10, 0, 5])                     # an empty group, one with none bad
    yield kgroups.Pocket([9000, 200], [1500, 190])


def test_the_conditional_likelihood_is_unchanged(tmp_path):
    for p in _pockets():
        for b in (np.zeros(len(p.loans)), np.linspace(-0.4, 0.5, len(p.loans))):
            new, old = kgroups.pocket_terms(p, b), _old_pocket_terms(p, b)
            assert math.isclose(new[0], old[0], rel_tol=1e-12, abs_tol=1e-9), (new[0], old[0])
            assert np.allclose(new[1], old[1], rtol=1e-10, atol=1e-8), (new[1], old[1])
            assert np.allclose(new[2], old[2], rtol=1e-9, atol=1e-7), (new[2], old[2])


def test_a_pocket_the_size_of_an_sba_stratum_takes_seconds_not_minutes():
    loans = np.array([44000, 44000, 44000, 44000, 41800])
    p = kgroups.Pocket(loans, np.round(loans * np.array([.2, .2, .18, .14, .12])))
    t = time.perf_counter()
    kgroups.pocket_terms(p, np.zeros(5))
    took = time.perf_counter() - t
    assert took < 20, f"{took:.1f} s for one evaluation (it was 191 s)"
