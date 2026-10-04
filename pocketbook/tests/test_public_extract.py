"""The public-data rehearsal's converter (tools/public_extract.py) and harness (tools/rehearse.py), on tiny
made-up rows only: no public file, no real loan, is read here (the raw files are outside git, by licence and size).

Each derived column is held to the rule its manifest states, at the edges of that rule; the rows left out are
counted by reason and add up; the judgments (--ranr, LendingClub's term and window) refuse rather than default;
names, addresses and free text never reach the extract; and PocketBook reads what the converter writes.
"""

import csv
import importlib.util
import io
import sys
import zipfile
from pathlib import Path

import pytest
import yaml

from pocketbook import ingest

TOOLS = Path(__file__).resolve().parents[1] / "tools"


def _tool(name):
    spec = importlib.util.spec_from_file_location(name, TOOLS / f"{name}.py")
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod                  # dataclasses look their module up by name
    spec.loader.exec_module(mod)
    return mod


px = _tool("public_extract")

FOIA_HEAD = ["AsOfDate", "Program", "LocationID", "BorrName", "BorrStreet", "BorrCity", "BorrState", "BorrZip",
             "BankName", "BankFDICNumber", "BankNCUANumber", "BankStreet", "BankCity", "BankState", "BankZip",
             "GrossApproval", "SBAGuaranteedApproval", "ApprovalDate", "ApprovalFY", "FirstDisbursementDate",
             "ProcessingMethod", "InitialInterestRate", "FixedorVariableInterestInd", "TermInMonths", "NaicsCode",
             "NaicsDescription", "FranchiseCode", "FranchiseName", "ProjectCounty", "ProjectState",
             "SBADistrictOffice", "CongressionalDistrict", "BusinessType", "BusinessAge", "LoanStatus",
             "PaidInFullDate", "ChargeOffDate", "GrossChargeOffAmount", "RevolverStatus", "JobsSupported",
             "CollateralInd", "SoldSecMrktInd"]


def _foia_row(status, gco="0.0", term="120", disb="2005-01-15", naics="531110", **kw):
    r = {h: "" for h in FOIA_HEAD}
    r.update({"AsOfDate": "2026-06-30", "Program": " 7A", "BorrName": "ZEBRA WIDGETS LLC", "BorrStreet": "1 Elm St",
              "BorrCity": "Springfield", "BorrState": "OH", "BorrZip": "45501", "BankName": "FIRST MADEUP BANK",
              "GrossApproval": "150000.0", "ApprovalDate": "2004-12-01", "ApprovalFY": "2005",
              "FirstDisbursementDate": disb, "ProcessingMethod": "SBA Express Program", "TermInMonths": term,
              "NaicsCode": naics, "BusinessType": "CORPORATION", "BusinessAge": "Existing, 5 or more years",
              "LoanStatus": status, "GrossChargeOffAmount": gco, "RevolverStatus": "N", "JobsSupported": "4",
              "CollateralInd": "Y", "SoldSecMrktInd": ""})
    r.update(kw)
    return r


def _write_csv(path, head, rows):
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=head)
        w.writeheader()
        w.writerows(rows)
    return path


def _read(path):
    with path.open(newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def _foia(tmp_path, rows):
    raw = _write_csv(tmp_path / "FOIA_7a_FY2000_FY2009_asof_260630.csv", FOIA_HEAD, rows)
    out = tmp_path / "x" / "foia.csv"
    m = px.convert("sba-foia", raw, out, ranr="neg-gco")
    return m, _read(out), out


def test_foia_keeps_known_outcomes_and_counts_the_rest_by_reason(tmp_path):
    rows = [_foia_row("P I F"), _foia_row("CHGOFF", gco="81234.5"), _foia_row("CANCLD"), _foia_row("EXEMPT"),
            _foia_row("CHGOFF", gco="2000.0"), _foia_row("PIF")]      # "PIF" is the dictionary's word, not the data's
    m, got, out = _foia(tmp_path, rows)
    assert m["rows_read"] == 6 and m["rows_kept"] == 3
    assert m["rows_left_out_by_reason"] == {"LoanStatus CANCLD: no known outcome": 1,
                                            "LoanStatus EXEMPT: no known outcome": 1,
                                            "LoanStatus PIF: no known outcome": 1}
    assert [r["CHGOFF_FLAG"] for r in got] == ["0", "1", "1"]
    assert m["rows_kept_by_status"] == {"LoanStatus CHGOFF": 2, "LoanStatus P I F": 1}
    # the key is the file and the data row, so the rows left out leave gaps: row 5 is the third kept
    assert [r["ROW_KEY"] for r in got] == ["FY0009-0000001", "FY0009-0000002", "FY0009-0000005"]
    assert [r["GrossChargeOffAmount"] for r in got] == ["0.0", "81234.5", "2000.0"]
    assert [r["RANR_NEG_GCO"] for r in got] == ["0", "-81234.50", "-2000"]
    assert m["raw_sha256"] == px.sha256_of(tmp_path / "FOIA_7a_FY2000_FY2009_asof_260630.csv")
    assert m["extract_sha256"] == px.sha256_of(out)


def test_foia_writes_no_name_address_or_bank(tmp_path):
    _, _, out = _foia(tmp_path, [_foia_row("P I F"), _foia_row("CHGOFF", gco="10.0")])
    text = out.read_text(encoding="utf-8")
    for leak in ("ZEBRA", "Elm St", "Springfield", "45501", "FIRST MADEUP"):
        assert leak not in text, leak
    assert "BorrName" not in text.splitlines()[0]


def test_real_estate_is_240_months_or_more(tmp_path):
    _, got, _ = _foia(tmp_path, [_foia_row("P I F", term="239"), _foia_row("P I F", term="240"),
                                 _foia_row("P I F", term="300"), _foia_row("P I F", term="")])
    assert [r["REAL_ESTATE"] for r in got] == ["N", "Y", "Y", ""]


def test_recession_is_the_papers_sas_to_the_day(tmp_path):
    """Disbursed + term x 30 days, from 1 Dec 2007 to 30 Jun 2009 inclusive (JSE 26(1) footnote 6)."""
    def disb_for(end):                     # a 12-month loan: 360 days
        return (end - px.timedelta(days=360)).isoformat()
    ends = [px.date(2007, 11, 30), px.date(2007, 12, 1), px.date(2009, 6, 30), px.date(2009, 7, 1)]
    _, got, _ = _foia(tmp_path, [_foia_row("P I F", term="12", disb=disb_for(e)) for e in ends]
                      + [_foia_row("P I F", disb="")])
    assert [r["RECESSION"] for r in got] == ["N", "Y", "Y", "N", ""]


def test_naics2_is_the_first_two_digits_and_blank_stays_blank(tmp_path):
    m, got, _ = _foia(tmp_path, [_foia_row("P I F", naics="531110"), _foia_row("P I F", naics="484121"),
                                 _foia_row("P I F", naics="")])
    assert [r["NAICS2"] for r in got] == ["53", "48", ""]
    assert m["derived_values_left_blank"] == {"NAICS2: NaicsCode blank": 1}


def test_every_column_says_what_it_is_and_why(tmp_path):
    m, got, _ = _foia(tmp_path, [_foia_row("P I F")])
    names = [c["name"] for c in m["columns"]]
    assert names == list(got[0].keys())                          # the manifest describes exactly what was written
    for c in m["columns"]:
        assert c["status"] in (px.NATIVE, px.CONSTRUCTED, px.APPROX) and c["rule"] and c["why"], c
    status = {c["name"]: c["status"] for c in m["columns"]}
    assert status["ROW_KEY"] == px.CONSTRUCTED and status["GrossChargeOffAmount"] == px.NATIVE
    assert status["RANR_NEG_GCO"] == status["REAL_ESTATE"] == status["RECESSION"] == px.APPROX
    meant = {c["read_as"] for c in m["columns"]}
    assert {"key", "booked", "outcome", "gco", "ranr", "origination_date"} <= meant


def test_ranr_is_a_judgment_and_is_refused_until_named(tmp_path):
    raw = _write_csv(tmp_path / "f.csv", FOIA_HEAD, [_foia_row("P I F")])
    out = tmp_path / "e.csv"
    with pytest.raises(px.Refused, match="--ranr"):
        px.convert("sba-foia", raw, out)
    with pytest.raises(px.Refused, match="isn't built"):
        px.convert("sba-foia", raw, out, ranr="interest-proxy")
    assert not out.exists()
    assert px.main(["sba-foia", str(raw), "--out", str(out)]) == 2 and not out.exists()


def test_a_missing_raw_file_is_refused_in_words_and_writes_nothing(tmp_path, capsys):
    out = tmp_path / "e.csv"
    assert px.main(["sba-foia", str(tmp_path / "gone.csv"), "--out", str(out), "--ranr", "neg-gco"]) == 2
    err = capsys.readouterr().err
    assert "REFUSED" in err and "isn't there" in err
    assert not out.exists() and not list(tmp_path.iterdir())


def test_the_wrong_file_is_refused_naming_what_it_lacks(tmp_path):
    raw = _write_csv(tmp_path / "f.csv", ["id", "loan_status"], [{"id": "1", "loan_status": "Fully Paid"}])
    out = tmp_path / "e.csv"
    with pytest.raises(px.Refused, match="isn't a sba-foia file: it has no Program"):
        px.convert("sba-foia", raw, out, ranr="neg-gco")
    assert not out.exists() and not (tmp_path / "e.csv.part").exists()


# --------------------------------------------------------------------------
# SBAnational, read out of its zip

NAT_HEAD = ["LoanNr_ChkDgt", "Name", "City", "State", "Zip", "Bank", "BankState", "NAICS", "ApprovalDate",
            "ApprovalFY", "Term", "NoEmp", "NewExist", "CreateJob", "RetainedJob", "FranchiseCode", "UrbanRural",
            "RevLineCr", "LowDoc", "ChgOffDate", "DisbursementDate", "DisbursementGross", "BalanceGross",
            "MIS_Status", "ChgOffPrinGr", "GrAppv", "SBA_Appv"]


def _nat_row(key, status, appr, fy, disb, gco="$0.00 ", term="84", naics="451120"):
    return {"LoanNr_ChkDgt": key, "Name": "ACME PRETEND CO", "City": "NOWHERE", "State": "FL", "Zip": "33101",
            "Bank": "MADEUP BANK", "BankState": "FL", "NAICS": naics, "ApprovalDate": appr, "ApprovalFY": fy,
            "Term": term, "NoEmp": "4", "NewExist": "2", "CreateJob": "0", "RetainedJob": "0",
            "FranchiseCode": "1", "UrbanRural": "1", "RevLineCr": "N", "LowDoc": "N", "ChgOffDate": "",
            "DisbursementDate": disb, "DisbursementGross": "$60,000.00 ", "BalanceGross": "$0.00 ",
            "MIS_Status": status, "ChgOffPrinGr": gco, "GrAppv": "$60,000.00 ", "SBA_Appv": "$48,000.00 "}


def _nat_zip(tmp_path, rows):
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=NAT_HEAD, lineterminator="\n")
    w.writeheader()
    w.writerows(rows)
    z = tmp_path / "should-this-loan.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.writestr("SBAnational.csv", buf.getvalue())
        zf.writestr("paper.pdf", b"%PDF")
    return z


def test_national_settles_each_two_digit_year_from_the_fiscal_year(tmp_path):
    rows = [_nat_row("1", "P I F", "28-Feb-97", "1997", "31-Mar-97"),
            _nat_row("2", "CHGOFF", "15-Oct-05", "2006", "1-Nov-05", gco="$12,345.60 "),   # October starts FY2006
            _nat_row("3", "P I F", "29-Feb-00", "2000", "1-Mar-00"),      # a leap day only 2000 has
            _nat_row("4", "P I F", "3-Mar-98", "1976A", "1-Apr-98"),      # a fiscal year that isn't a year
            _nat_row("5", "P I F", "3-Mar-98", "2004", "1-Apr-98"),       # a fiscal year neither century fits
            _nat_row("6", "", "3-Mar-98", "1998", "1-Apr-98")]            # no outcome
    out = tmp_path / "nat.csv"
    m = px.convert("sba-national", _nat_zip(tmp_path, rows), out, ranr="neg-gco")
    got = _read(out)
    assert [r["APPROVAL_DATE"] for r in got] == ["1997-02-28", "2005-10-15", "2000-02-29", "", ""]
    assert m["derived_values_left_blank"]["APPROVAL_DATE: ApprovalFY not a plain year"] == 1
    assert m["derived_values_left_blank"]["APPROVAL_DATE: neither century falls in ApprovalFY"] == 1
    assert m["rows_left_out_by_reason"] == {"MIS_Status (blank): no known outcome": 1}
    assert [r["CHGOFF_FLAG"] for r in got] == ["0", "1", "0", "0", "0"]
    assert got[1]["RANR_NEG_GCO"] == "-12345.60" and got[1]["ChgOffPrinGr"] == "$12,345.60 "
    # passed through as written; PocketBook reads it
    assert got[0]["DisbursementGross"] == "$60,000.00 " and ingest.parse_number(got[0]["DisbursementGross"]) == 60000
    text = out.read_text(encoding="utf-8")
    assert "ACME" not in text and "MADEUP" not in text and "33101" not in text


def test_national_naics_zero_is_no_industry(tmp_path):
    out = tmp_path / "nat.csv"
    rows = [_nat_row("1", "P I F", "3-Mar-98", "1998", "1-Apr-98", naics="0"),
            _nat_row("2", "P I F", "3-Mar-98", "1998", "1-Apr-98", naics="621111")]
    px.convert("sba-national", _nat_zip(tmp_path, rows), out, ranr="neg-gco")
    assert [r["NAICS2"] for r in _read(out)] == ["", "62"]


def test_national_recession_reads_the_disbursement_in_the_century_nearest_approval(tmp_path):
    # approved Nov 2006, disbursed "1-Jan-07", 12 months (360 days): ends 27 Dec 2007, inside the window
    rows = [_nat_row("1", "P I F", "20-Nov-06", "2007", "1-Jan-07", term="12")]
    out = tmp_path / "nat.csv"
    px.convert("sba-national", _nat_zip(tmp_path, rows), out, ranr="neg-gco")
    assert _read(out)[0]["RECESSION"] == "Y"


# --------------------------------------------------------------------------
# LendingClub

def _lc_head():
    return list(px.LC_NEEDS) + ["url", "desc", "emp_title", "title", "last_fico_range_high", "recoveries"]


def _lc_row(i, status, issued="Jun-2010", term=" 36 months", funded="10000", pymnt="11500.5", prncp="10000",
            fee="0"):
    r = {h: "" for h in _lc_head()}
    r.update({"id": str(i), "loan_status": status, "issue_d": issued, "funded_amnt": funded, "total_pymnt": pymnt,
              "total_rec_prncp": prncp, "collection_recovery_fee": fee, "term": term, "grade": "B",
              "sub_grade": "B3", "emp_title": "Secret Agent Person", "desc": "my private story", "url": "https://x",
              "title": "Debt", "last_fico_range_high": "700", "recoveries": "12", "fico_range_low": "700",
              "purpose": "credit_card", "home_ownership": "RENT"})
    return r


def test_lendingclub_builds_gco_and_ranr_from_cash_flows(tmp_path):
    rows = [_lc_row(1, "Fully Paid"),
            _lc_row(2, "Charged Off", pymnt="4200.25", prncp="3500", fee="50.5"),
            _lc_row(3, "Current"),
            _lc_row(4, "Fully Paid", issued="Jan-2012"),
            _lc_row(5, "Fully Paid", term=" 60 months"),
            _lc_row(6, "Does not meet the credit policy. Status:Charged Off", pymnt="1000", prncp="800"),
            {**{h: "" for h in _lc_head()}, "id": "Total amount funded in policy code 1: 1234"}]
    raw = _write_csv(tmp_path / "lc.csv", _lc_head(), rows)
    out = tmp_path / "lc-out.csv"
    m = px.convert("lendingclub", raw, out, term=36, issued="2008-01:2011-12")
    got = _read(out)
    assert [r["id"] for r in got] == ["1", "2", "6"]
    assert [r["BAD"] for r in got] == ["0", "1", "1"]
    assert m["rows_kept_by_status"] == {"loan_status Charged Off": 1, "loan_status Fully Paid": 1,
                                        "loan_status Does not meet the credit policy. Status:Charged Off": 1}
    assert [r["GCO_APPROX"] for r in got] == ["0.00", "6500.00", "9200.00"]
    # total_pymnt - funded_amnt - collection_recovery_fee
    assert [r["RANR_APPROX"] for r in got] == ["1500.50", "-5850.25", "-9000.00"]
    assert got[0]["ISSUE_DATE"] == "2010-06-01"
    assert m["rows_left_out_by_reason"] == {"issued outside the window": 1, "loan_status Current: no outcome yet": 1,
                                            "not a loan row (the file's summary lines)": 1, "term not 36 months": 1}
    head = out.read_text(encoding="utf-8").splitlines()[0].split(",")
    for never in ("url", "desc", "emp_title", "title", "last_fico_range_high", "recoveries", "total_pymnt"):
        assert never not in head, never
    assert "Secret Agent" not in out.read_text(encoding="utf-8")


def test_lendingclub_own_download_skips_its_notes_line_and_names_a_column_it_lacks(tmp_path):
    # LendingClub's own LoanStats3a.csv: a "Notes offered by Prospectus" line above the header, and no fico_range_low
    # (the revived rehearsal, 3 Oct 2026). Without --absent it is refused naming the column; with it, the column is
    # left out of the extract and the manifest says so; the notes line is skipped and recorded, never the header
    head = [h for h in _lc_head() if h != "fico_range_low"]
    rows = [{k: v for k, v in {**_lc_row(1, "Fully Paid"), "int_rate": " 12.25%", "revol_util": "45.5%"}.items()
             if k in head},
            {k: v for k, v in {**_lc_row(2, "Charged Off", pymnt="4200.25", prncp="3500"), "int_rate": "7.5",
                               "revol_util": ""}.items() if k in head}]
    raw = _write_csv(tmp_path / "LoanStats3a.csv", head, rows)
    raw.write_text('Notes offered by Prospectus (https://www.lendingclub.com/info/prospectus.action)\n'
                   + raw.read_text(encoding="utf-8"), encoding="utf-8")
    out = tmp_path / "lc-out.csv"
    with pytest.raises(px.Refused, match="it has no fico_range_low$"):
        px.convert("lendingclub", raw, out, term=36, issued="2008-01:2011-12")
    m = px.convert("lendingclub", raw, out, term=36, issued="2008-01:2011-12", absent=("fico_range_low",))
    got = _read(out)
    assert [r["id"] for r in got] == ["1", "2"] and [r["BAD"] for r in got] == ["0", "1"]
    assert "fico_range_low" not in got[0] and "grade" in got[0]
    assert m["choices"]["absent"] == ["fico_range_low"] and "fico_range_low" in m["absent_from_raw"]
    assert m["preamble_skipped"][0].startswith("Notes offered by Prospectus")
    assert "fico_range_low" not in [c["name"] for c in m["columns"]]
    # its rates are written "12.25%", which PocketBook reads as not a number: the sign comes off, counted
    assert [(r["int_rate"], r["revol_util"]) for r in got] == [("12.25", "45.5"), ("7.5", "")]
    assert m["percent_sign_removed"]["values"] == {"int_rate": 1, "revol_util": 1}
    # a file that has the column refuses --absent for it; the key and the cash flows can never be absent
    full = _write_csv(tmp_path / "lc.csv", _lc_head(), [_lc_row(1, "Fully Paid")])
    with pytest.raises(px.Refused, match="has it"):
        px.convert("lendingclub", full, tmp_path / "o.csv", term=36, issued="2008-01:2011-12",
                   absent=("fico_range_low",))
    with pytest.raises(px.Refused, match="never optional"):
        px.convert("lendingclub", full, tmp_path / "o.csv", term=36, issued="2008-01:2011-12", absent=("total_pymnt",))
    m2 = px.convert("lendingclub", full, tmp_path / "o.csv", term=36, issued="2008-01:2011-12")
    assert "preamble_skipped" not in m2 and "absent_from_raw" not in m2 and "ROW_KEY" not in _read(tmp_path / "o.csv")[0]


def test_lendingclub_with_every_id_blank_is_keyed_by_its_row_only_when_named(tmp_path):
    # the copy of LoanStats3a.csv reachable on 3 Oct 2026 has id and member_id blank on every row, a text line between
    # the policy sections and two totals at the end. Without --row-key nothing is kept and the reason says why
    rows = [{**_lc_row(1, "Fully Paid"), "id": ""},
            {**{h: "" for h in _lc_head()}, "id": "Loans that do not meet the credit policy"},
            {**_lc_row(2, "Does not meet the credit policy. Status:Charged Off", pymnt="1000", prncp="800"), "id": ""},
            {**{h: "" for h in _lc_head()}, "id": "Total amount funded in policy code 1: 1234"}]
    raw = _write_csv(tmp_path / "lc.csv", _lc_head(), rows)
    m = px.convert("lendingclub", raw, tmp_path / "a.csv", term=36, issued="2008-01:2011-12")
    assert m["rows_kept"] == 0
    assert m["rows_left_out_by_reason"] == {"id blank: no key (if every id is blank, name --row-key)": 2,
                                            "not a loan row (the file's summary lines)": 2}
    m = px.convert("lendingclub", raw, tmp_path / "b.csv", term=36, issued="2008-01:2011-12", row_key=True)
    got = _read(tmp_path / "b.csv")
    assert [r["ROW_KEY"] for r in got] == ["LC-0000001", "LC-0000003"] and "id" not in got[0]
    assert [r["BAD"] for r in got] == ["0", "1"]
    assert m["rows_left_out_by_reason"] == {"not a loan row (the file's summary lines)": 2}
    key = next(c for c in m["columns"] if c["read_as"] == "key")
    assert key["name"] == "ROW_KEY" and key["status"] == px.CONSTRUCTED
    with pytest.raises(px.Refused, match="lendingclub source only"):
        px.convert("sba-foia", raw, tmp_path / "c.csv", ranr="neg-gco", row_key=True)


def test_lendingclub_term_and_window_are_refused_until_named(tmp_path):
    raw = _write_csv(tmp_path / "lc.csv", _lc_head(), [_lc_row(1, "Fully Paid")])
    for kw in ({}, {"term": 36}, {"issued": "2008-01:2011-12"}):
        with pytest.raises(px.Refused, match="not defaulted"):
            px.convert("lendingclub", raw, tmp_path / "o.csv", **kw)
    with pytest.raises(px.Refused, match="runs backwards"):
        px.convert("lendingclub", raw, tmp_path / "o.csv", term=36, issued="2011-12:2008-01")
    assert not (tmp_path / "o.csv").exists()


# --------------------------------------------------------------------------
# PocketBook reads what the converter writes, and the harness refuses without an extract


def test_pocketbook_reads_the_extract_as_the_manifest_says(tmp_path):
    rows = [_foia_row("P I F" if i % 4 else "CHGOFF", gco="0.0" if i % 4 else f"{1000 + i}.0",
                      naics=f"{11 + i % 5}0000") for i in range(40)]
    _, _, out = _foia(tmp_path, rows)
    t = ingest.read_table(out)
    det = ingest.detect_date_format("ApprovalDate", [r["ApprovalDate"] for r in t.rows])
    assert det.resolved == "%Y-%m-%d"
    for col in ("GrossApproval", "GrossChargeOffAmount", "RANR_NEG_GCO", "TermInMonths"):
        assert all(isinstance(ingest.parse_number(r[col]), float) for r in t.rows), col
    assert sum(float(r["RANR_NEG_GCO"]) for r in t.rows) == -sum(float(r["GrossChargeOffAmount"]) for r in t.rows)


def test_the_harness_refuses_cleanly_without_its_extract(tmp_path, capsys):
    rh = _tool("rehearse")
    ans = tmp_path / "a.yaml"
    ans.write_text(yaml.safe_dump({"kind": "bleed", "few_values": 12, "many_values": 50}), encoding="utf-8")
    assert rh.main([str(ans), "--extract", str(tmp_path / "gone.csv"), "--memory", str(tmp_path / "m.yaml")]) == 2
    assert "isn't there" in capsys.readouterr().err
    assert not (tmp_path / "m.yaml").exists()


def test_the_harness_will_not_default_the_launcher_limits(tmp_path, capsys):
    rh = _tool("rehearse")
    ex = _write_csv(tmp_path / "e.csv", ["A"], [{"A": "1"}])
    ans = tmp_path / "a.yaml"
    ans.write_text(yaml.safe_dump({"kind": "bleed"}), encoding="utf-8")
    assert rh.main([str(ans), "--extract", str(ex), "--memory", str(tmp_path / "m.yaml")]) == 2
    assert "few_values" in capsys.readouterr().err
    assert not list(tmp_path.glob("*PocketBook*"))


def test_the_workbook_check_can_fail(tmp_path):
    """The rehearsal's workbook check reads clean on a clean workbook, and finds a planted #REF!, a formula
    pointing at a tab that isn't there, and a dropdown listing from one."""
    from openpyxl import Workbook
    from openpyxl.worksheet.datavalidation import DataValidation
    rh = _tool("rehearse")
    wb = Workbook()
    ws = wb.active
    ws.title = "Pockets"
    wb.create_sheet("_live")["A1"] = 1
    ws["A1"] = "='_live'!A1*2"
    clean = tmp_path / "clean.xlsx"
    wb.save(clean)
    got = rh.check_workbook(clean)
    assert got["error_token_count"] == 0 and got["sheets_referenced_but_missing"] == []
    assert got["formulas_per_tab"] == {"Pockets": 1, "_live": 0}
    ws["A2"] = "=#REF!+1"
    ws["A3"] = "=Gone!B2"
    dv = DataValidation(type="list", formula1="='Old tab'!$A$1:$A$3")
    ws.add_data_validation(dv)
    dv.add("B1")
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    got = rh.check_workbook(broken)
    assert got["error_token_count"] == 1 and "Pockets!A2" in got["error_tokens"][0]
    assert got["sheets_referenced_but_missing"] == ["Gone", "Old tab"]


# --------------------------------------------------------------------------
# Found in review, 29 Sep 2026: on both SBA files the recorded term depends on the outcome (paid loans' terms are
# mostly whole years, charged-off loans' mostly not), and columns recorded after booking went in unlabelled. The
# manifest now counts the one and labels the other.

def test_the_manifest_counts_whole_year_terms_by_outcome(tmp_path):
    rows = ([_foia_row("P I F", term=t) for t in ("84", "120", "60", "85")]
            + [_foia_row("CHGOFF", gco="100.0", term=t) for t in ("47", "53", "84")]
            + [_foia_row("CHGOFF", gco="100.0", term="")])                  # no term: not counted
    m, _, _ = _foia(tmp_path, rows)
    tc = m["term_check"]["by_outcome"]
    assert tc == {"outcome 0": {"loans_with_a_term": 4, "whole_years": 3, "share_whole_years": 0.75},
                  "outcome 1": {"loans_with_a_term": 3, "whole_years": 1, "share_whole_years": 0.3333}}
    assert "circular" in m["term_check"]["why"]


def test_national_counts_its_term_the_same_way(tmp_path):
    rows = [_nat_row("1", "P I F", "28-Feb-97", "1997", "31-Mar-97", term="84"),
            _nat_row("2", "CHGOFF", "28-Feb-97", "1997", "31-Mar-97", gco="$5.00 ", term="41")]
    out = tmp_path / "n.csv"
    m = px.convert("sba-national", _nat_zip(tmp_path, rows), out, ranr="neg-gco")
    assert m["term_check"]["by_outcome"]["outcome 0"]["share_whole_years"] == 1.0
    assert m["term_check"]["by_outcome"]["outcome 1"]["share_whole_years"] == 0.0


def test_columns_not_known_at_booking_are_labelled(tmp_path):
    m, _, _ = _foia(tmp_path, [_foia_row("P I F")])
    after = {c["name"]: c.get("after_booking", "") for c in m["columns"]}
    assert "term_check" in after["TermInMonths"] and "secondary market" in after["SoldSecMrktInd"]
    assert after["REAL_ESTATE"] and after["RECESSION"]                      # built from the term
    assert not after["RevolverStatus"] and not after["GrossApproval"]      # known when the loan is made
    out = tmp_path / "n.csv"
    n = px.convert("sba-national", _nat_zip(tmp_path, [_nat_row("1", "P I F", "28-Feb-97", "1997", "31-Mar-97")]),
                   out, ranr="neg-gco")
    nat = {c["name"]: c.get("after_booking", "") for c in n["columns"]}
    assert "Disbursed" in nat["DisbursementGross"] and "term_check" in nat["Term"] and nat["RECESSION"]


def test_lendingclub_has_no_term_check(tmp_path):
    """Its term is one value, chosen on the command line: nothing to count."""
    raw = _write_csv(tmp_path / "lc.csv", _lc_head(), [_lc_row(1, "Fully Paid"), _lc_row(2, "Charged Off")])
    m = px.convert("lendingclub", raw, tmp_path / "o.csv", term=36, issued="2008-01:2011-12")
    assert "term_check" not in m and not any("after_booking" in c for c in m["columns"])


def test_the_sba_ranr_stand_in_says_its_revenue_is_invented(tmp_path):
    m, _, _ = _foia(tmp_path, [_foia_row("P I F")])
    why = {c["name"]: c["why"] for c in m["columns"]}["RANR_NEG_GCO"]
    assert "invented" in why and "NOT evidence" in why and "InitialInterestRate" in why
    assert "invents no revenue" not in why


# --------------------------------------------------------------------------
# tools/rehearsal_effects.py: a pocket alone in its band is judged against the book, and its count says so (found in
# review, 29 Sep 2026: on Term bands, "N worse 3" came from bands where N was the whole band)

def _effects_book(path):
    from openpyxl import Workbook
    wb = Workbook()
    v = wb.active
    v.title = "_views"
    for row in (["G|t x s|cols", "N", "Y", "All"], ["G|t x s|rows", "short", "long", "All"],
                ["G|t x s|outcome_loans|rate|1", 0.5, 0.2, 0.4], ["G|t x s|loans|1", 100, 50, 150],
                ["G|t x s|outcome_loans|rate|2", 0.1, None, 0.1], ["G|t x s|loans|2", 200, 0, 200],
                ["G|t x s|outcome_loans|rate|3", 0.2, 0.2, 0.2], ["G|t x s|loans|3", 300, 50, 350]):
        v.append(row)
    p = wb.create_sheet("_pockets")
    head = ["Grids or three-way", "Grid", "Rate (key)", "Segment", "Too few losses to test (fewest losses is a Run setting)",
            "Alone in its band", "Vs rest of band", "p-value vs band, after the allowance",
            "Vs rest of book (a multiple; profit: the gap, 0.01 = 1 point)", "p-value vs book, after the allowance"]
    p.append(head)
    p.append(["grids", "t x s", "outcome_loans", "N", None, False, 2.5, 0.001, 2.5, 0.001])   # worse vs its band
    p.append(["grids", "t x s", "outcome_loans", "Y", None, False, 0.4, 0.001, 1.0, 0.9])     # better vs its band
    p.append(["grids", "t x s", "outcome_loans", "N", None, True, None, None, 0.5, 0.001])    # alone: vs the book
    wb.save(path)
    return path


def test_effects_say_which_pockets_were_judged_against_the_book(tmp_path):
    fx = _tool("rehearsal_effects")
    got = fx.read(_effects_book(tmp_path / "b.xlsx"), "outcome_loans", 1.25, 0.8, 0.95, "band")["t x s"]["pockets"]
    assert got["N"]["worse"] == 1 and got["N"]["better"] == 1 and got["N"]["better_vs_book"] == 1
    assert got["N"]["worse_vs_book"] == 0 and got["Y"]["better"] == 1 and got["Y"]["better_vs_book"] == 0
    assert fx.say(got["N"]) == ("pockets worse 1, better 1 (1 alone in its band, so against the book), "
                                "neither/untested 0")


# --------------------------------------------------------------------------
# The scripts behind the report's answer key and timing split are in git (review, 29 Sep 2026: they lived only in a
# session's scratch folder). Both refuse in words without the data folder, which can be purged.

def test_the_answer_key_and_the_timing_split_refuse_without_their_data(tmp_path, capsys):
    key, timing = _tool("rehearsal_answer_key"), _tool("rehearsal_timing")
    assert key.main([str(tmp_path)]) == 2 and "isn't there" in capsys.readouterr().err
    assert timing.main([str(tmp_path), "gone"]) == 2 and "isn't there" in capsys.readouterr().err


def test_the_answer_key_reads_rates_and_months_as_the_report_quotes_them():
    key = _tool("rehearsal_answer_key")
    assert key.show("grade", {"A": [200, 10], "B": [100, 20], "": [4, 1]}) == \
        "grade (blank)=25.00%(n=4)  B=20.00%(n=100)  A=5.00%(n=200)"
    assert key.show("t", {"N": [10, 1], "Y": [10, 5]}, ("Y", "N", "(blank)")) == "t Y=50.00%(n=10)  N=10.00%(n=10)"
    assert key.months("2005-01-15", "2007-03-02") == 26
