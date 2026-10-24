"""RANR vs GCOs: Avg line and Line × book sit beside Booked, the way RANR ÷ Booked sits beside RANR.

The firm, 17 Oct 2026: "i want avg line and booked metrics to be closer together the same way RANR and RANR/booked
is, like avg line and line x book and such belong in the same relative place". The gross block was Booked, GCOs,
RANR, RANR ÷ Booked, Avg line, Line × book; it is now Booked, Avg line, Line × book, GCOs, RANR, RANR ÷ Booked, one
group heading over all six. Summary already reads Booked $, Avg line, Line × book, % of booked, the shape its GCO and
RANR columns have, so it does not move. Checked on the workbook as written, with no LibreOffice and no display."""

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from pocketbook import results
from test_avg_line_2026_10_02 import made  # noqa: F401  (the module's synthetic book)

import tabs

GROSS_ORDER = ["Booked", "Avg line", "Line × book", "GCOs", "RANR", "RANR ÷ Booked"]


def test_avg_line_and_line_x_book_sit_beside_booked(made):  # noqa: F811
    b, _ = made
    ws = load_workbook(b)[results.PCK]
    head = tabs.header_row(ws, results.C_TOG, "Together")
    assert tabs.heads(ws, head, results.C_BOOK, results.C_RATE) == GROSS_ORDER
    assert ws.cell(row=head, column=results.C_RATE + 1).value == "Gap pts"     # the sides start after RANR ÷ Booked
    group = [str(m) for m in ws.merged_cells.ranges if m.min_row == head - 1 and m.min_col == results.C_BOOK]
    assert group == [f"{results.col(results.C_BOOK)}{head - 1}:{results.col(results.C_RATE)}{head - 1}"]
    assert ws.cell(row=head - 1, column=results.C_BOOK).value == "Gross · this pocket alone"


def test_each_gross_cell_sits_in_one_conditional_format_range(made):  # noqa: F811
    """LibreOffice applies one conditional format to a cell (results.cf), so a cell in two ranges is shaded by
    whichever comes first: the Avg line columns, now inside Band to GCOs, must not also sit in the plain range."""
    b, _ = made
    ws = load_workbook(b)[results.PCK]
    head = tabs.header_row(ws, results.C_TOG, "Together")
    row = head + 1
    for c in range(results.C_BAND, results.C_RATE + 1):
        hits = 0
        for cr in ws.conditional_formatting:
            for part in str(cr.sqref).split():
                c0, r0, c1, r1 = range_boundaries(part)
                hits += c0 <= c <= c1 and r0 <= row <= r1
        assert hits == 1, (results.col(c), hits)
