"""The firm's answers to the walk-through's design calls, A to K (docs/walkthrough/2026-09-27/WALKTHROUGH-DEFECTS.md,
"Left for the firm"; BACKLOG.md §6d, 27 Sep 2026), and tenet T2: a check figure is shown only where it can fail.

Each answer is held by one test here, named for what the analyst now sees. The two that need a display skip
without one, as the walk's own window tests do."""

import shutil
from pathlib import Path
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from conftest import TEST_SHUFFLES
from pocketbook import book, choices as ch, confirm_tab, confirmatory, control, engine, house, launcher, live, look
from pocketbook import perm, results, synth
from recalc import recalc, values_of
from test_book import _answer
from test_launcher import _read
import tabs

CUT = dict(run_kind=ch.BLEED, bands=("FICO", "ORIG_BAL"), segments=("CHANNEL", "ASSET_CLASS"))


@pytest.fixture(autouse=True)
def _nothing_opens(monkeypatch):
    monkeypatch.setattr(launcher, "open_file", lambda p: None)


@pytest.fixture(scope="module")
def bled(tmp_path_factory):
    """The walk's cut: FICO and ORIG_BAL into bands, CHANNEL and ASSET_CLASS as segments, answered and Run."""
    d = tmp_path_factory.mktemp("firm-answers")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", TEST_SHUFFLES)
        x = synth.write_extract(d, n=6000)
        out = book.set_up(x, choices=ch.Choices(**CUT))
        assert out.ok, out.lines
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return {"x": x, "b": out.book, "ran": ran, "dir": d}


def _copy(bled, to: Path) -> tuple[Path, Path]:
    """The extract and its workbook, copied, so a test can change them without touching the others'."""
    to.mkdir(parents=True, exist_ok=True)
    x = Path(shutil.copy(bled["x"], to / bled["x"].name))
    b = Path(shutil.copy(bled["b"], to / bled["b"].name))
    return x, b


def _edit(b: Path, to: Path, **answers) -> Path:
    """A copy of the workbook with Control's answers changed, as an analyst types them."""
    wb = load_workbook(b)
    ws = wb[control.SHEET]
    for key, v in answers.items():
        ws.cell(row=control.row_of(ws, key), column=control.CHOOSE_COL).value = v
    to.mkdir(parents=True, exist_ok=True)
    p = to / b.name
    wb.save(p)
    return p


def _tile(ws, label):
    for row in ws.iter_rows():
        for c in row:
            if c.value == label:
                return ws.cell(row=c.row + 1, column=c.column).value
    raise KeyError(label)


def _banner(ws) -> str:
    return next((c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str)
                 and c.value.startswith("↻ ")), "")


# ---- A


def test_a_choice_changed_in_the_launcher_waits_for_a_run_on_start_here_control_and_the_banner(bled, tmp_path,
                                                                                               monkeypatch):
    """A: after Next changed what is cut, Start here said 0 changes waiting, under result tabs from the Run before."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    x, b = _copy(bled, tmp_path / "book")
    calm = recalc(b, tmp_path / "calm")
    assert _tile(calm["Start here"], "Changes waiting for a Run") == 0
    # the launcher's Next, with ORIG_BAL no longer cut into bands
    assert book.set_up(x, choices=ch.Choices(**{**CUT, "bands": ("FICO",)})).ok
    got = recalc(b, tmp_path / "moved")
    ws = got[control.SHEET]
    status = {ws.cell(row=r, column=control.KEY_COL).value: ws.cell(row=r, column=control.STATUS_COL).value
              for r in control.launcher_rows_of(load_workbook(b)[control.SHEET])}
    assert status.pop("launcher|bands") == house.WAITING
    assert [k for k, v in status.items() if v != house.SAME] == []
    start = got["Start here"]
    assert _tile(start, "Changes waiting for a Run") == 1
    assert "Cut into bands: FICO, ORIG_BAL → FICO (launcher)" in _banner(start), _banner(start)
    # and the Run that uses it clears it
    assert book.run(b).ok
    again = recalc(b, tmp_path / "again")
    assert _tile(again["Start here"], "Changes waiting for a Run") == 0 and not _banner(again["Start here"])
    # a different kind of run is one change, not one for every row it changes with it (the walk's case)
    assert book.set_up(x, choices=ch.Choices(run_kind=ch.NEW_VARIABLE, bands=("FICO",), segments=(),
                                             outcome="BAD_FLAG", test=("REV_DEBT",), hold=("FICO",))).ok
    new = recalc(b, tmp_path / "new")["Start here"]
    assert _tile(new, "Changes waiting for a Run") == 1
    assert "What are you running?: Where the book bleeds → Finding and testing a new variable (launcher)" in \
        _banner(new), _banner(new)


# ---- B, and tenet T2


def test_no_screen_shows_a_tie_out_figure_that_could_only_read_fine(bled):
    """B: "702 / 702" on the finished screen and Start here could never read anything else: a Run that doesn't tie out
    stops and writes nothing. Record keeps the check, as the number of checks."""
    h = bled["ran"].summary
    tiles = launcher.finished_tiles(h)
    assert [t[0] for t in tiles] == ["Pockets worse and material", "Charge-offs above their share"]
    wb = load_workbook(bled["b"])
    said = [c.coordinate for row in wb["Start here"].iter_rows() for c in row
            if isinstance(c.value, str) and ("tie-out check" in c.value.lower() or c.value.endswith(" agree"))]
    assert said == []
    assert [x for x in bled["ran"].lines if "tie-out" in x] == []          # nor in the Run's lines, or the Log's
    got = tabs.record(bled["b"])["Tie-out checks"]
    assert got == f"{h['tie_outs']:,}: every grid adds up to the book" and " of " not in got


# ---- C


def test_charge_offs_worse_and_real_with_kept_about_the_same_reads_losing_more_profit_holding(tmp_path):
    """C: a pocket shaded worse on charge-offs whose kept gap could be chance had a blank Together."""
    for kept in (engine.IN_LINE, engine.UNSURE_WORSE, engine.UNSURE_BETTER):
        assert results.together_of(engine.WORSE, kept) == "Losing more, profit holding"
        assert results.together_of(engine.IN_LINE, kept) == ""
    assert results.together_of(engine.IN_LINE, engine.WORSE) == "Earns less, not from losses"     # kept as it was
    wb = Workbook()
    ws = wb.active
    for i, (g, k) in enumerate(((engine.WORSE, engine.IN_LINE), (engine.WORSE, engine.UNSURE_BETTER),
                                (engine.BETTER, engine.IN_LINE)), start=1):
        ws.cell(row=i, column=1, value=g)
        ws.cell(row=i, column=2, value=k)
        ws.cell(row=i, column=3, value=0)
        ws.cell(row=i, column=4, value="=" + results.together_formula(f"$A{i}", f"$B{i}", f"$C{i}"))
    got = values_of(wb, tmp_path).active
    assert [got.cell(row=i, column=4).value or "" for i in (1, 2, 3)] == ["Losing more, profit holding"] * 2 + [""]


# ---- D


def _listed(values_wb) -> list[tuple]:
    ws = values_wb["Start here"]
    head = next(c.row for row in ws.iter_rows() for c in row if c.value == "Largest, worse and material")
    rows = [tuple(ws.cell(row=head + k, column=c).value for c in (2, 3, 6)) for k in range(1, book.TOP_ROWS + 1)]
    return [r for r in rows if r[0]]


def _worse_and_material_now(values_wb) -> list[tuple]:
    """Every pocket _found lists, in its order, that _pockets reads worse and material now: worked out here from
    _pockets' own cells, not from _found's running count."""
    pk = values_wb[live.POCKETS]
    out = []
    for r in values_wb[book.FOUND].iter_rows(values_only=True):
        if r[0] == "top":
            prow = r[book.TOP_PROW - 1]
            if pk.cell(row=prow, column=live.P_WORSE).value == live.YES and \
                    pk.cell(row=prow, column=live.P_MAT).value == live.YES:
                out.append((r[1], r[2]))
    return out


def test_start_here_lists_only_the_pockets_worse_and_material_now_largest_first(bled, tmp_path):
    """D: after worse at went to 2 times, two of the four rows under "Largest, worse and material" read Worse? No.
    The rows are picked by formula from the verdicts now, in the last Run's order of dollars."""
    before = recalc(bled["b"], tmp_path / "before")
    listed = _listed(before)
    want = _worse_and_material_now(before)[:book.TOP_ROWS]
    assert listed and [r[:2] for r in listed] == want
    dollars = [r[2] for r in listed]
    assert dollars == sorted(dollars, reverse=True)                      # largest first, as of the last Run
    after = recalc(_edit(bled["b"], tmp_path / "twice", worse_at="2 times"), tmp_path / "after")
    now = _listed(after)
    assert [r[:2] for r in now] == _worse_and_material_now(after)[:book.TOP_ROWS]
    gone = [r[:2] for r in listed if r[:2] not in [x[:2] for x in now]]
    assert gone, "worse at 2 times dropped none of the listed pockets: the test proves nothing on this book"
    # no Worse? or Material? column: every row listed is both, so the column could only say Yes (tenet T2)
    heads = [c.value for row in load_workbook(bled["b"])["Start here"].iter_rows() for c in row
             if c.value in ("Worse?", "Material?")]
    assert heads == []


# ---- E


def test_the_launcher_start_here_and_the_refusal_give_one_count_of_answers_needed(tmp_path, monkeypatch):
    """E: after Next the window said "7 columns to look at first", Start here "Columns to confirm: 10", and the Run's
    refusal listed 9 answers. Each now says the refusal's count."""
    monkeypatch.setattr(launcher.Path, "home", lambda: tmp_path)
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    f = _read(tmp_path)
    f.next()
    b = launcher.book_for(f.extract)
    n = f.written.summary["needed"]
    assert n > 1
    said = [x for x in f.written.lines if "needed before Run" in x]
    assert said == [f"Next: {n} answers needed before Run, on Control and Columns. Fill in the shaded cells, save, "
                    f"close, and press Run."], f.written.lines
    assert [x for x in f.written.lines if "to look at first" in x] == []
    assert _tile(recalc(b, tmp_path / "calc")["Start here"], book.NEEDED) == n
    f.run()
    assert len(f.needs) == n and f.needs_head()[0] == f"{n} answers needed before Run"


# ---- F


def test_the_launcher_says_what_each_term_means_where_it_first_uses_it(tmp_path, monkeypatch):
    """F: GCO dollars, RANR dollars, worse at 1.34x, grids with five measures each, and scouting were first read on
    the window, and explained only on a tab. A short plain line under each, with no term of art in it."""
    monkeypatch.setattr(launcher.Path, "home", lambda: tmp_path)
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    f = _read(tmp_path)
    f.pick_outcome("BAD_FLAG")                       # 29 Sep 2026: nothing is picked for the analyst
    f.answer_outcome(True)
    got = f.meanings()
    assert [x.split(":")[0] for x in got] == ["GCO dollars", "RANR dollars", "A grid", "Five measures"], got
    f.gate.optional = ["scikit-learn"]
    f.set_mode("new")
    assert "(scouting)" in f.summary()[1] and [x.split(":")[0] for x in f.meanings()][-1] == "Scouting"
    f.gate.optional = []
    f.set_mode("bleed")
    f.next()
    worse = [x for x in f.meanings() if x.startswith("worse at ")]
    assert len(worse) == 1 and f.suggested_line() and worse[0].split(":")[0] in f.suggested_line()
    every = got + f.meanings() + launcher.plain_words("scouting")
    jargon = ("GCO", "RANR", "pocket", "multiple", "p-value", "significan", "topline", "holdout", "pre-spec", "×")
    long_ = [x for x in every if len(x.split()) > 15]
    loose = [x for x in every if any(j in x.split(": ", 1)[1] for j in jargon)]
    assert long_ == [] and loose == []


# ---- G


def test_comes_to_shows_the_worked_out_value_before_the_first_run_never_the_options_words(tmp_path, monkeypatch):
    """G: before the first Run, Comes to read "The smallest significant gap in a typical pocket (suggested)",
    spilling over Or your own, where after the Run it read 1.34x. Set up already knows the value."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    out = book.set_up(synth.write_extract(tmp_path, n=3000), choices=ch.Choices(**CUT))
    sugg = out.summary["suggested"]
    assert "worse_at" in sugg
    wb = load_workbook(out.book)
    ws = wb[control.SHEET]
    labels = {o.label: o for s in control.load_settings() if s.key in ("worse_at", "better_at") for o in s.options}
    suggestion = next(lab for lab, o in labels.items() if o.value == "luck")
    edited = _edit(out.book, tmp_path / "answered", worse_at=suggestion, better_at="0.8 times",
                   materiality="1% of the book's total losses")
    got = recalc(edited, tmp_path / "calc")[control.SHEET]
    comes = {k: got.cell(row=control.row_of(ws, k), column=control.COMES_COL).value
             for k in ("worse_at", "better_at", "materiality")}
    assert comes["materiality"] in (None, "") and \
        (comes["worse_at"], comes["better_at"]) == (f"{sugg['worse_at']:.2f}×", "0.80×"), comes
    words = {o.label for s in control.load_settings() for o in s.options} | {o.shown for s in control.load_settings()
                                                                              for o in s.options}
    shown = [c.value for row in got.iter_rows(min_col=control.COMES_COL, max_col=control.COMES_COL) for c in row
             if c.value in words]
    assert shown == []


# ---- H


def test_look_draws_only_the_columns_that_can_be_cut_into_bands(tmp_path, monkeypatch):
    """H: GCO_AMT and RANR_AMT each had a block and a chart on Look saying "no edges on Columns yet". An outcome is
    never cut into bands, so it has no edges to choose: the firm, "Not sure why we even have the info?"."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    out = book.set_up(synth.write_extract(tmp_path, n=2000), choices=ch.Choices(**CUT))
    wb = load_workbook(out.book)
    drawn = [c for c in wb[look.DATA][1][1:] if c.value]
    assert [c.value for c in drawn] == ["FICO", "ORIG_BAL", "REV_DEBT"]
    named = [c.coordinate for row in wb[look.LOOK].iter_rows() for c in row
             if isinstance(c.value, str) and any(o in c.value for o in ("GCO_AMT", "RANR_AMT", "BAD_FLAG"))]
    assert named == []


# ---- I


def test_start_heres_group_table_says_the_odds_hold_the_columns_fixed(monkeypatch):
    """I: after the scouting Run, Start here's 0 to 10,999 read 1.24x, New variables' "Confirmed, FICO held fixed", while
    its header said only "x 11,000 - 14,999's odds"."""
    h = {"kind": "confirm", "column": "REV_DEBT", "candidates": ["REV_DEBT"], "reference": "11,000 - 14,999",
         "development": 5604, "holdout": 2395, "deviations": 0}
    monkeypatch.setattr(confirmatory, "headline", lambda res: h)
    monkeypatch.setattr(confirmatory, "groups_found", lambda res: [])

    def header(strata) -> list:
        t = SimpleNamespace(problem=None, holdout=object(), strata=strata, held=bool(strata))
        wb = Workbook()
        found = wb.create_sheet(book.FOUND)
        confirm_tab.write_found(found, SimpleNamespace(prespec=SimpleNamespace(tests=[t])))
        value = {r[0]: r[1] for r in found.iter_rows(values_only=True)}.get
        ws = wb.active
        confirm_tab.found_block(ws, wb, 1, value)
        return [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("×")]

    assert header(("FICO",)) == ["× 11,000 - 14,999's odds, FICO held fixed"]
    assert header(("FICO", "CHANNEL")) == ["× 11,000 - 14,999's odds, FICO and CHANNEL held fixed"]
    assert header(()) == ["× 11,000 - 14,999's odds"]


# ---- J


def test_the_finished_screen_shows_the_runs_first_two_lines(bled, tmp_path):
    """J: the launcher built "Open ... start with Pockets" and, after an edit, "Changed pre-spec: ...", and showed
    neither: only Record did."""
    ran = bled["ran"]
    f = launcher.Flow(gate=launcher.AddOns())
    f.finished = ran
    first = f.first_lines()
    assert first[0] == f"Open {bled['b'].name}: start with Pockets." and first == ran.summary["first"]
    small = [x for x in ran.lines if " material but too small to test" in x]
    assert first[1] == (small[0] if small else ran.lines[0]) and first[1] in ran.lines
    changed = f"{confirmatory.CHANGED} held.yaml changed after the held-back run of 2026-09-27: fingerprint a then, b now."
    lines = ["Ran on 8,000 loans from loans.csv.", "Follows pre-spec held.yaml.", changed,
             "Open loans - PocketBook.xlsx: start with New variables."]
    assert book.first_lines(lines) == [lines[-1], changed]
    split = "Split by REV_DEBT: each pocket halved at its own median."
    assert book.first_lines([lines[0], split, lines[-1]]) == [lines[-1], split]


def test_the_finished_window_draws_two_tiles_and_the_runs_first_two_lines(monkeypatch, tmp_path):
    """B and J on the window itself. Needs a display (xvfb-run on Linux)."""
    from test_deps import _window
    root = _window()
    try:
        monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
        w = launcher.build(root)
        flow = w["flow"]
        first = ["Open loans - PocketBook.xlsx: start with Pockets.", "2 pockets are material but too small to test."]
        h = {"measure": "GCO per booked dollar", "gco": True, "worse": 2, "pockets": 116, "dollars": 614831.0,
             "tie_outs": 162, "open": [], "first": first}
        flow.finished, flow.page = book.Outcome(True, Path(), [], summary=h), "done"
        w["render"]()
        root.update()
        heads = [t.winfo_children()[1].cget("text") for t in w["tiles"]]
        assert heads == ["Pockets worse and material", "Charge-offs above their share"]
        assert w["first"].cget("text") == "\n".join(first) and w["first"].winfo_manager() == "pack"
    finally:
        root.destroy()


# ---- Choose tests, in the order the analyst works down it (the firm, 27 Sep 2026: "it just isn't necessary to
# have anything there, really. i would prefer that screens are ordered more sensibly- this one seems all over the
# place")

ORDER = ["FICO", "ORIG_BAL", "REV_DEBT",          # cut into bands, or split by
         "CHANNEL", "ASSET_CLASS",                # segment by
         "BAD_FLAG", "GCO_AMT", "RANR_AMT",       # what is measured
         "LOAN_NBR", "ORIG_DATE"]                 # the key and the date, greyed


def test_choose_tests_rows_run_numbers_categories_outcomes_then_key_and_date(tmp_path):
    """The extract reads LOAN_NBR, FICO, CHANNEL, ORIG_BAL, BAD_FLAG, ...: the table no longer does. The same order
    for both run kinds, so the toggle never reshuffles it; an outcome row carries no box and no words."""
    f = _read(tmp_path)
    assert [c.name for c in f.read.columns] != ORDER
    for mode in ("bleed", "new", "bleed"):
        f.set_mode(mode)
        rows = f.rows()
        assert [r["name"] for r in rows] == ORDER, mode
        assert [r["group"] for r in rows] == [0, 0, 0, 1, 1, 2, 2, 2, 3, 3]
        assert [r["name"] for r in rows if r["grey"]] == ["LOAN_NBR", "ORIG_DATE"]
        assert not any("every" in r for r in rows)
    for r in f.rows()[5:]:
        assert r["a"] is r["b"] is r["c"] is None


def test_choose_tests_rows_keep_the_extracts_order_within_a_group(tmp_path):
    """Grouped, not sorted: the extract written back to front reads each group back to front too."""
    x = synth.write_extract(tmp_path / "a", n=1500)
    lines = [ln.split(",") for ln in x.read_text(encoding="utf-8").splitlines()]
    back = tmp_path / "b" / "loans.csv"
    back.parent.mkdir()
    back.write_text("\n".join(",".join(reversed(ln)) for ln in lines) + "\n", encoding="utf-8")
    f = launcher.Flow(gate=launcher.AddOns())
    f.pick(str(back))
    f.set_up()
    assert f.screen() == "L2", f.message
    assert [r["name"] for r in f.rows()] == ["REV_DEBT", "ORIG_BAL", "FICO", "ASSET_CLASS", "CHANNEL",
                                            "RANR_AMT", "GCO_AMT", "BAD_FLAG", "ORIG_DATE", "LOAN_NBR"]


def test_choose_tests_window_draws_the_rows_in_order_with_a_quiet_gap_and_no_every_measure(monkeypatch, tmp_path):
    """On the window: the rows top to bottom as Flow.rows says, a gap with no words where one group ends, and no
    "every measure" anywhere. Needs a display (xvfb-run on Linux)."""
    from test_deps import _window
    root = _window()
    try:
        monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
        w = launcher.build(root)
        root.geometry("1180x628")                # 29 Sep 2026: the window a 1366 x 768 laptop opens it at
        root.deiconify()                         # the table takes the window's height, so it has to be shown
        flow = w["flow"]
        flow.pick(str(synth.write_extract(tmp_path, n=1500)))
        flow.set_up()
        for mode in ("bleed", "new"):
            flow.set_mode(mode)
            w["render"]()
            root.update()
            ys = [w[f"row_{n}"].winfo_y() for n in ORDER]
            assert ys == sorted(ys) and len(set(ys)) == len(ys), mode
            assert len(w["gaps"]) == 3 and all(not g.winfo_children() for g in w["gaps"])
            last = w["row_ORIG_DATE"]                    # ten columns and three gaps fit without scrolling
            assert last.winfo_y() + last.winfo_reqheight() + 1 <= w["table"].winfo_height(), mode
            said, todo = [], [root]
            while todo:
                x = todo.pop()
                todo += x.winfo_children()
                try:
                    said.append(str(x.cget("text")))
                except Exception:
                    pass
            assert said and not [s for s in said if "every measure" in s], mode
            assert all(g.winfo_class() == "Frame" for g in w["gaps"])            # frames: nothing written
    finally:
        root.destroy()


def test_choose_tests_gaps_fall_where_a_group_starts_and_the_table_fits_unscrolled(tmp_path):
    """Headless, so CI checks it too: a gap before the first category, the first outcome and the key, nowhere
    else; and ten columns with their three gaps fit in the room each run kind gives the table, so no row hides
    below it."""
    flow = launcher.Flow()
    flow.pick(str(synth.write_extract(tmp_path, n=1500)))
    flow.set_up()
    for mode in ("bleed", "new"):
        flow.set_mode(mode)
        rows = flow.rows()
        assert [r["name"] for r in rows if r["gap_before"]] == ["CHANNEL", "BAD_FLAG", "LOAN_NBR"], mode
        # 29 Sep 2026: the table takes the window's height, no longer a fixed 250 or 280 px; on a laptop's
        # 1366 x 768 screen the window opens 628 px tall, room for these rows with the rest of the screen
        assert launcher.table_height(rows) <= launcher.window_size(1366, 768)[1] - 330, mode
