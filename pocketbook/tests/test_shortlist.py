"""The lean pre-spec (Goal 2 item 3): an outcome plus a shortlist of inputs, each with its bins and reference group,
the columns held fixed, and the allowance for testing the whole shortlist at once.

The firm, 26 Sep 2026: "well it cannot be one column, but a shortlist whatever. one column makes no sense - it can't
be used in a tree", and "we might want to test a set once with and once without" (FICO): each input is reported
with and without the held-fixed columns.

The book is the first synthetic book with two more columns (synth.add_shortlist): UTIL, whose odds of going bad
are planted at 2.5 times above 0.9 of the line, and TENURE, with nothing planted, which must not hold up. The
income / sales cliffs are planted as before. Every count is made here from the extract by hand, and every allowed
p-value is worked out here by Benjamini-Hochberg's own steps (statistics.md A2), never read back from the cube."""

import bisect
import csv
import math
import shutil
import subprocess
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from pocketbook import book, confirm_tab, confirmatory, control, kgroups, launcher, prespec, synth
from recalc import recalc
from test_book import _answer
from test_book_dates import _check, _choose, _control
from test_confirm_test import BINS, PLANTED

needs_git = pytest.mark.skipif(shutil.which("git") is None, reason="git is not installed on this machine")
DOCS = Path(__file__).resolve().parents[1] / "docs"
EXAMPLE = DOCS / "prespec-example.yaml"
SHORTLIST_EXAMPLE = DOCS / "prespec-shortlist-example.yaml"
UTIL_BINS, TENURE_BINS = [0.3, 0.6, 0.9], [2.0, 5.0, 10.0]
INCOME = {"column": "income_to_sales", "bins": BINS, "reference": "0.25 - 0.49"}
UTIL = {"column": "UTIL", "bins": UTIL_BINS, "reference": 0}
TENURE = {"column": "TENURE", "bins": TENURE_BINS, "reference": 1}
BASE = {"prespec": 1, "written": "2026-09-27", "outcome": "BAD_FLAG", "strata": ["FICO", "CHANNEL"],
        "confidence": 0.95, "holdout": {"from": "2024-01-01", "to": "2024-12-31"},
        "development": {"from": "2022-01-01", "to": "2023-12-31"}}
HOLD, DEV = ("2024-01-01", "2024-12-31"), ("2022-01-01", "2023-12-31")
NOT_CUT = ("ORIG_BAL", "REV_DEBT", "ASSET_CLASS", "INCOME", "SALES", "income_to_sales", "UTIL", "TENURE")


def _git(repo, *args):
    subprocess.run(["git", "-c", "user.name=Test", "-c", "user.email=test@example.com", "-c", "commit.gpgsign=false",
                    *args], cwd=repo, check=True, capture_output=True)


def _write_spec(folder: Path, name: str = "prespec.yaml", **raw) -> Path:
    f = folder / name
    f.write_text(yaml.safe_dump({**BASE, **raw}, sort_keys=False), encoding="utf-8")
    return f


def _book(folder: Path, n: int, spec: dict | None = None, text: str | None = None):
    """The first synthetic book with UTIL and TENURE, income / sales made on Columns, confirming the shortlist
    `spec` (or the file `text`), committed beside the workbook."""
    x = synth.write_extract(folder, n=n, ratio=True, shortlist=True)
    b = book.set_up(x).book
    _answer(b)
    _control(b, run_kind="Finding and testing a new variable", new_variable_step="Test from a pre-spec",
             **{"derived|1": ("income_to_sales", "INCOME", "SALES")})
    book.set_up(x)
    _choose(b, drop=NOT_CUT, shortlist="prespec.yaml")
    wb = load_workbook(b)
    wb["Columns"][book.CONFIRM_CELL] = "Yes"
    wb.save(b)
    if text is not None:
        (folder / "prespec.yaml").write_text(text, encoding="utf-8")
    else:
        _write_spec(folder, **spec)
    _git(folder, "init", "-q")
    _git(folder, "add", "prespec.yaml")
    _git(folder, "commit", "-q", "-m", "pre-spec")
    return x, b


def _run(b, mp):
    seen = {}
    real = confirmatory.state

    def spy(*a, **k):
        seen["st"] = real(*a, **k)
        seen["res"] = a[2]
        return seen["st"]
    mp.setattr(confirmatory, "state", spy)
    ran = book.run(b)
    return ran, seen.get("st"), seen.get("res")


@pytest.fixture(scope="module")
def three(tmp_path_factory):
    """Three inputs, FICO and CHANNEL held fixed: 20,000 loans, as the goal's own run."""
    folder = tmp_path_factory.mktemp("three")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("GIT_CEILING_DIRECTORIES", str(folder))
        mp.setenv("POCKETBOOK_MEMORY", str(folder / "memory.yaml"))
        x, b = _book(folder, 20000, {"inputs": [INCOME, UTIL, TENURE]})
        ran, st, res = _run(b, mp)
    assert ran.ok, ran.lines
    return {"x": x, "b": b, "ran": ran, "st": st, "res": res, "tests": {t.column: t for t in st.tests},
            "folder": folder}


@pytest.fixture(scope="module")
def three_calc(three, tmp_path_factory):
    return recalc(three["b"], tmp_path_factory.mktemp("calc3"))


@pytest.fixture(scope="module")
def two(tmp_path_factory):
    """Two inputs, nothing held fixed (`strata: []`)."""
    folder = tmp_path_factory.mktemp("two")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("GIT_CEILING_DIRECTORIES", str(folder))
        mp.setenv("POCKETBOOK_MEMORY", str(folder / "memory.yaml"))
        x, b = _book(folder, 20000, {"inputs": [UTIL, INCOME], "strata": []})
        ran, st, res = _run(b, mp)
    assert ran.ok, ran.lines
    return {"x": x, "b": b, "ran": ran, "st": st, "tests": {t.column: t for t in st.tests}}


# --------------------------------------------------------------------------
# By hand, from the extract


def _rows(x):
    with open(x, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def _value(r, column):
    if column == "income_to_sales":
        return None if r["SALES"] in ("", "0") else float(r["INCOME"]) / float(r["SALES"])
    return float(r[column])


def _counts(x, column, bins, rng):
    """Loans and bad loans per group of `column`, made in `rng` (both ends in), counted from the extract."""
    k = len(bins) + 1
    loans, bad = [0] * k, [0] * k
    for r in _rows(x):
        v = _value(r, column)
        if not rng[0] <= r["ORIG_DATE"] <= rng[1] or v is None or r["BAD_FLAG"] not in ("0", "1"):
            continue
        g = bisect.bisect_right(bins, v)
        loans[g] += 1
        bad[g] += int(r["BAD_FLAG"])
    return loans, bad


def _bh(ps: list[float]) -> list[float]:
    """Benjamini-Hochberg's adjusted p-values, by the steps statistics.md A2 writes out: sort ascending, scale the
    i-th smallest by m / i, take the smallest of those from each one up, cap at 1."""
    m = len(ps)
    order = sorted(range(m), key=lambda i: ps[i])
    scaled = [ps[order[j]] * m / (j + 1) for j in range(m)]
    out = [0.0] * m
    for j in range(m):
        out[order[j]] = min(1.0, min(scaled[j:]))
    return out


def _family(tests, side_name):
    """Every group of every candidate against its own reference on one set of loans: (column, group, raw p)."""
    out = []
    for t in tests:
        side = getattr(t, side_name)
        for k in range(len(t.groups)):
            if k != t.ref and side.fit.p[k] is not None:
                out.append((t.column, k, side.fit.p[k]))
    return out


def _table(ws) -> dict:
    """New variables' table as the analyst reads it: {(candidate, comparison): {column: value, "row": r}}."""
    # each candidate on its own: the table whose Found column reads Gap (OC-51: all together sits above it)
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=confirm_tab.N_CAND).value == "Candidate"
                and ws.cell(row=r, column=confirm_tab.N_FG).value == "Gap")
    out = {}
    for r in range(head + 1, ws.max_row + 1):
        comp = ws.cell(row=r, column=confirm_tab.N_COMP).value
        if not comp:
            break
        cand = ws.cell(row=r, column=confirm_tab.N_CAND).value
        out[(cand, comp)] = {c: ws.cell(row=r, column=c).value for c in range(confirm_tab.N_CAND,
                                                                             confirm_tab.N_WORDS + 1)}
        out[(cand, comp)]["row"] = r
    return out


# --------------------------------------------------------------------------
# The file


def test_a_shortlist_reads_each_input_with_its_bins_and_reference():
    ps = prespec.parse({**BASE, "inputs": [INCOME, UTIL, TENURE]})
    assert ps.shortlist and ps.outcome == "BAD_FLAG" and ps.columns == ("income_to_sales", "UTIL", "TENURE")
    got = [(i.column, i.bins, i.reference, i.reference_index) for i in ps.inputs]
    assert got == [("income_to_sales", tuple(BINS), "0.25 - 0.49", 2), ("UTIL", (0.3, 0.6, 0.9), "up to 0.2", 0),
                   ("TENURE", (2.0, 5.0, 10.0), "2 - 4", 1)]
    named = prespec.named(ps, ranges={"UTIL": (0.0, 1.2), "TENURE": (0.0, 61.3)})
    assert [i.groups[0] for i in named.inputs] == ["up to 0.09", "0.0 - 0.2", "0 - 1"]
    # a label read in whole units drops the fraction (30 Sep 2026, "Cut at whole dollars is fine"): 61.3 reads 61
    assert named.inputs[1].reference == "0.0 - 0.2" and named.inputs[2].groups[-1] == "10 - 61"


def test_the_one_column_pre_spec_still_reads_as_a_shortlist_of_one():
    old = prespec.load(EXAMPLE)
    assert not old.shortlist and old.outcome is None and len(old.inputs) == 1
    one = prespec.parse({**BASE, "strata": list(old.strata), "written": "2026-09-25",
                         "inputs": [{"column": old.column, "bins": list(old.bins), "reference": old.reference}]})
    assert one.inputs == old.inputs and (one.strata, one.confidence, one.holdout, one.development) == \
        (old.strata, old.confidence, old.holdout, old.development)


@pytest.mark.parametrize("key", prespec.SHORTLIST_KEYS)
def test_each_missing_shortlist_line_is_refused_with_the_line_to_add(key):
    raw = {**BASE, "inputs": [INCOME]}
    raw.pop(key)
    with pytest.raises(prespec.PreSpecError) as exc:
        prespec.parse(raw)
    said = [p for p in exc.value.problems if p.startswith(f"missing line `{key}:`")]
    assert said and "\n" + key + ":" in said[0], exc.value.problems


@pytest.mark.parametrize("inputs,words", [
    ([], "`inputs:` must list one or more inputs"),
    ([{"column": "UTIL", "bins": [0.3]}], "missing line `inputs[0].reference:`"),
    ([UTIL, {**UTIL, "bins": [0.5]}], "`inputs:` lists UTIL more than once"),
    ([{**UTIL, "edges": [1]}], "unknown line `inputs[0].edges:`"),
    ([{**UTIL, "bins": [0.9, 0.3]}], "`inputs[0].bins:` must rise"),
    ([{**UTIL, "reference": "middle"}], "`inputs[0].reference: 'middle'` is not one of the groups"),
    ([{**UTIL, "reference": "[CONFIRM: which group?]"}], "`inputs[0].reference` still reads"),
    ([UTIL, {**TENURE, "column": "FICO"}], "`strata:` includes `FICO`, the column being tested"),
    ([{**UTIL, "column": "BAD_FLAG"}], "`inputs:` tests `BAD_FLAG`, the outcome itself"),
])
def test_a_shortlist_is_refused_line_by_line(inputs, words):
    with pytest.raises(prespec.PreSpecError) as exc:
        prespec.parse({**BASE, "inputs": inputs})
    assert [p for p in exc.value.problems if words in p], exc.value.problems


def test_a_shortlist_with_a_top_level_column_as_well_is_refused():
    with pytest.raises(prespec.PreSpecError) as exc:
        prespec.parse({**BASE, "inputs": [UTIL], "column": "UTIL"})
    assert exc.value.problems[0].startswith("`column:` is written at the top as well as under `inputs:`")


def test_strata_are_suggested_and_left_blank_in_what_pocketbook_writes():
    """OC-13, the firm's answer on the 26 Sep docket: the line PocketBook offers for missing strata suggests them and
    leaves them unanswered, so the file refuses until the analyst says; so does the committed shortlist example."""
    raw = {**BASE, "inputs": [UTIL]}
    raw.pop("strata")
    with pytest.raises(prespec.PreSpecError) as exc:
        prespec.parse(raw)
    offered = next(p for p in exc.value.problems if p.startswith("missing line `strata:`")).split("Add:\n")[1]
    with pytest.raises(prespec.PreSpecError) as exc:
        prespec.parse({**yaml.safe_load(offered), **{k: v for k, v in raw.items()}})
    assert exc.value.problems[0].startswith("`strata` still reads '[CONFIRM: the columns to hold fixed, such as")
    with pytest.raises(prespec.PreSpecError) as exc:
        prespec.load(SHORTLIST_EXAMPLE)
    assert len(exc.value.problems) == 1 and exc.value.problems[0].startswith("`strata` still reads")
    answered = yaml.safe_load(SHORTLIST_EXAMPLE.read_text(encoding="utf-8"))
    answered["strata"] = ["FICO", "CHANNEL"]
    assert prespec.parse(answered).columns == ("income_to_sales", "UTIL", "TENURE")


def test_each_input_that_differs_is_said_by_its_column():
    ps = prespec.parse({**BASE, "inputs": [INCOME, UTIL]})
    same = {"outcome": "BAD_FLAG", "strata": ["FICO", "CHANNEL"], "confidence": 0.95,
            "holdout": {"from": "2024-01-01", "to": "2024-12-31"},
            "inputs": [{"column": "UTIL", "bins": UTIL_BINS, "reference": 0},
                       {"column": "income_to_sales", "edges": BINS, "reference": 2}]}
    assert prespec.deviations(ps, same) == []
    moved = {**same, "outcome": "OTHER_FLAG",
             "inputs": [{"column": "UTIL", "bins": [0.3, 0.6], "reference": 0}, {"column": "TENURE", "bins": [2],
                                                                                 "reference": 0}]}
    assert prespec.deviations(ps, moved, where="in this run") == [
        "The outcome is `OTHER_FLAG` in this run; the pre-spec says `BAD_FLAG`.",
        "`income_to_sales` isn't tested in this run; the pre-spec lists it.",
        "For UTIL: the bins are 0.3, 0.6 in this run; the pre-spec says 0.3, 0.6, 0.9.",
        "`TENURE` is tested in this run; the pre-spec doesn't list it."]


# --------------------------------------------------------------------------
# Three inputs on the first book: the planted ones hold up, the empty one doesn't


@needs_git
def test_three_inputs_the_planted_hold_up_with_and_without_the_held_fixed_columns_and_the_empty_one_does_not(three):
    tests = three["tests"]
    assert list(tests) == ["income_to_sales", "UTIL", "TENURE"] and all(t.problem is None for t in tests.values())
    bar = 0.05
    for side_name in ("holdout", "holdout_plain"):
        inc = getattr(tests["income_to_sales"], side_name)
        for k, planted in PLANTED.items():
            lo, hi = inc.fit.interval(k, 1.96)
            assert inc.fit.odds[k] > 1 and inc.allowed[k] < bar and lo < planted < hi, (side_name, k)
        util = getattr(tests["UTIL"], side_name)
        lo, hi = util.fit.interval(3, 1.96)
        assert util.fit.odds[3] > 1 and util.allowed[3] < bar and lo < synth.UTIL_ODDS < hi, (side_name, util.fit.odds)
        assert all(util.allowed[k] >= bar for k in (1, 2)), (side_name, util.allowed)
        tenure = getattr(tests["TENURE"], side_name)
        assert all(p is None or p >= bar for p in tenure.allowed), (side_name, tenure.allowed)
    # the counts behind them, made here from the extract
    for t in tests.values():
        for side, rng in ((t.holdout, HOLD), (t.development, DEV)):
            assert (side.loans, side.bad) == _counts(three["x"], t.column, list(t.bins), rng)


@needs_git
def test_the_allowance_is_benjamini_hochberg_across_every_candidates_groups(three):
    tests = list(three["tests"].values())
    for side_name in confirmatory.SETS:
        fam = _family(tests, side_name)
        assert len(fam) == 5 + 3 + 3                           # every group but the references, all three inputs
        want = _bh([p for _, _, p in fam])
        got = [getattr(three["tests"][c], side_name).allowed[k] for c, k, _ in fam]
        wrong = [(c, k, g, w) for (c, k, _), g, w in zip(fam, got, want) if not math.isclose(g, w, rel_tol=1e-12)]
        assert wrong == []


@needs_git
def test_new_variables_shows_one_block_per_candidate_with_the_allowed_p_value(three, three_calc):
    ws = three_calc[confirm_tab.SHEET]
    rows = _table(ws)
    tests = three["tests"]
    assert [c for c, _ in rows] == [t.column for t in tests.values() for k in range(len(t.groups)) if k != t.ref]
    head = ws.cell(row=min(x["row"] for x in rows.values()) - 1, column=confirm_tab.N_CP).value
    assert head == confirm_tab.P_ALLOWED
    wrong = []
    for t in tests.values():
        ref = t.groups[t.ref]
        for k, g in enumerate(t.groups):
            if k == t.ref:
                continue
            x = rows[(t.column, f"{g} vs {ref}")]
            for (o, p), side in (((confirm_tab.N_CG, confirm_tab.N_CP), t.holdout_plain),
                                 ((confirm_tab.N_HG, confirm_tab.N_HP), t.holdout),
                                 ((confirm_tab.N_FG, confirm_tab.N_FP), t.development_plain)):
                if not (math.isclose(x[o], side.fit.odds[k], rel_tol=1e-9)
                        and math.isclose(x[p], side.allowed[k], rel_tol=1e-9)):
                    wrong.append((t.column, g, o, x[o], x[p]))
            holds = "Yes" if side_ok(t.holdout_plain, t, k) else "No"
            still = "Yes" if side_ok(t.holdout, t, k) else "No"
            if (x[confirm_tab.N_HOLDS], x[confirm_tab.N_STILL]) != (holds, still):
                wrong.append((t.column, g, x[confirm_tab.N_HOLDS], x[confirm_tab.N_STILL], holds, still))
    assert wrong == []
    tenure = [x for (c, _), x in rows.items() if c == "TENURE"]
    assert {x[confirm_tab.N_HOLDS] for x in tenure} == {"No"} and {x[confirm_tab.N_STILL] for x in tenure} == {"No"}


def side_ok(side, t, k, bar=0.05):
    """Holds up? by hand: significant after the allowance, on the side of 1 the found loans showed."""
    o, p, found = side.fit.odds[k], side.allowed[k], t.development_plain.fit.odds[k]
    return o is not None and p is not None and p < bar and (o > 1) == (found > 1)


@needs_git
def test_the_tests_in_full_keep_every_raw_p_value(three):
    """The raw p-values stay, candidate by candidate: each group's on the held-back loans with the columns held
    fixed, under the tests in full."""
    ws = load_workbook(three["b"])[confirm_tab.SHEET]
    heads = [r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=2).value or "").startswith(
        "The tests in full: ")]
    assert [ws.cell(row=r, column=2).value for r in heads] == [f"The tests in full: {c}" for c in three["tests"]]
    missing = []
    for i, t in enumerate(three["tests"].values()):
        end = heads[i + 1] if i + 1 < len(heads) else ws.max_row + 1
        seen = [ws.cell(row=r, column=c).value for r in range(heads[i], end) for c in range(2, 15)]
        nums = [v for v in seen if isinstance(v, float)]
        for k in range(len(t.groups)):
            p = t.holdout.fit.p[k]
            if k != t.ref and p is not None and not any(math.isclose(v, p, rel_tol=1e-12) for v in nums):
                missing.append((t.column, k))
    assert missing == []


@needs_git
def test_the_method_note_names_the_allowance_once(three):
    ws = load_workbook(three["b"])[confirm_tab.SHEET]
    note = [str(ws.cell(row=r, column=3).value) for r in range(4, 40) if ws.cell(row=r, column=2).value]
    naming = [s for s in note if "Benjamini-Hochberg" in s]
    assert len(naming) == 1 and "across the 3 candidates' 11 groups" in naming[0]
    assert "Allowing for testing many pockets at once" in naming[0]


@needs_git
def test_one_chart_per_candidate_read_from_its_own_block(three):
    ws = load_workbook(three["b"])[confirm_tab.SHEET]
    rows = _table(ws)
    assert len(ws._charts) == 3
    for chart, t in zip(ws._charts, three["tests"].values()):
        mine = [x["row"] for (c, _), x in rows.items() if c == t.column]
        col = lambda c: (f"'{confirm_tab.SHEET}'!${get_column_letter(c)}${min(mine)}:"          # noqa: E731
                         f"${get_column_letter(c)}${max(mine)}")
        assert [s.val.numRef.f for s in chart.series] == [col(confirm_tab.N_CG), col(confirm_tab.N_HG)]
        assert chart.title.tx.rich.p[0].r[0].t.startswith(f"{t.column}: the gap against")


@needs_git
def test_the_tiles_count_the_candidates_holding_up(three, three_calc):
    ws = three_calc[confirm_tab.SHEET]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows(max_row=80) for c in row if c.value is not None}
    at = {v: k for k, v in cells.items() if isinstance(v, str)}
    r, c = at["Candidates · live"]
    assert cells[(r + 1, c)] == "3 · 2 hold up"
    h = confirmatory.headline(three["res"])
    assert (h["holding"], len(h["candidates"]), h["capture"]) == (2, 3, None)
    tiles = launcher.confirm_tiles(h)
    assert tiles[1][:2] == ("Candidates with a group worse", "2 of 3")


@needs_git
def test_the_worse_line_is_suggested_from_every_candidates_groups(three):
    """The median over every candidate's groups of the smallest odds ratio a group that size could call
    significant against its own reference, on the development loans, worked out here from the extract's counts."""
    import statistics
    gaps = []
    for t in three["tests"].values():
        loans, bad = _counts(three["x"], t.column, list(t.bins), DEV)
        p = sum(bad) / sum(loans)
        pq = p * (1 - p)
        gaps += [math.exp(1.959963984540054 * math.sqrt(1 / (n * pq) + 1 / (loans[t.ref] * pq)))
                 for k, n in enumerate(loans) if k != t.ref]
    want = round(statistics.median(gaps), 2)
    assert book.test_gap(list(three["tests"].values()), 0.95) == want
    assert want != book.test_gap(three["tests"]["income_to_sales"], 0.95)
    ws = load_workbook(three["b"])[control.SHEET]
    assert ws.cell(row=control.row_of(ws, "worse_at"), column=book.SUGGEST_COL).value == \
        f"suggested: {want:.2f}x, from this extract at the last Run"


@needs_git
def test_start_here_lists_every_candidates_groups(three, three_calc):
    ws = three_calc["Start here"]
    cells = {(c.row, c.column): c.value for row in ws.iter_rows() for c in row if c.value is not None}
    at = {v: k for k, v in cells.items() if isinstance(v, str)}
    r, c = at["Candidates with a group worse"]
    assert cells[(r + 1, c)] == "2 of 3"
    head = at["Candidate: group"][0]
    names = [cells.get((head + i, 2)) for i in range(1, 1 + sum(len(t.groups) for t in three["tests"].values()))]
    assert names[0] == f"income_to_sales: {three['tests']['income_to_sales'].groups[0]}"
    assert "UTIL: " + three["tests"]["UTIL"].groups[0] + " (reference)" in names
    worse = [n for i, n in enumerate(names) if cells.get((head + 1 + i, 7)) == "Yes, worse"]
    assert {n.split(": ")[0] for n in worse} == {"income_to_sales", "UTIL"}


@needs_git
def test_record_says_each_input_and_the_allowance(three):
    chk = _check(three["b"])
    says = chk["What the pre-spec says"].splitlines()
    assert "outcome: BAD_FLAG" in says and "inputs: 3" in says and "2. column: UTIL" in says
    assert "3. reference: 2 - 4" in says
    assert chk["Differs from the pre-spec"] == "nowhere: this run used what it says"
    tests = chk["Confirmatory test"]
    assert [s.split(":")[0] for s in tests] == ["income_to_sales", "UTIL", "TENURE"]
    assert chk["Tests: the allowance for many at once"].startswith("Benjamini-Hochberg, as Control's")


@needs_git
def test_the_allowance_turns_a_borderline_verdict(three, tmp_path, monkeypatch):
    """A group significant on its own and not once the shortlist's other ten comparisons are allowed for: Still
    holds? reads No with the allowance and Yes with Control's "No allowance", after a Run (the setting needs one).
    Read at the first confidence on Control where this book has such a group."""
    found = [(bar, label, t, k) for bar, label in ((0.05, "95% sure"), (0.10, "90% sure"))
             for t in three["tests"].values() for k in range(len(t.groups))
             if k != t.ref and t.holdout.fit.p[k] is not None and t.holdout.fit.p[k] < bar <= t.holdout.allowed[k]
             and (t.holdout.fit.odds[k] > 1) == (t.development_plain.fit.odds[k] > 1)]
    assert found, "no group on this book sits between its own p-value and the allowed one"
    bar, label, t, k = found[0]
    key = (t.column, f"{t.groups[k]} vs {t.groups[t.ref]}")

    def at(b, name):
        copy = tmp_path / f"{name}.xlsx"
        shutil.copy(b, copy)
        _control(copy, confidence=label)
        return _table(recalc(copy, tmp_path / name)[confirm_tab.SHEET])
    before = at(three["b"], "before")
    assert before[key][confirm_tab.N_STILL] == "No"
    assert before[key][confirm_tab.N_HP] == pytest.approx(t.holdout.allowed[k], rel=1e-9)
    folder = tmp_path / "none"
    shutil.copytree(three["folder"], folder)
    b = folder / three["b"].name
    _control(b, many_tests="No allowance")
    monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(folder))
    ran, st, res = _run(b, monkeypatch)
    assert ran.ok, ran.lines
    after = at(b, "after")
    assert after[key][confirm_tab.N_STILL] == "Yes"
    assert after[key][confirm_tab.N_HP] == pytest.approx(t.holdout.fit.p[k], rel=1e-9)
    ws = load_workbook(b)[confirm_tab.SHEET]
    assert ws.cell(row=min(x["row"] for x in after.values()) - 1, column=confirm_tab.N_HP).value == "p-value"
    assert _check(b)["Tests: the allowance for many at once"].startswith("None: each p-value")


# --------------------------------------------------------------------------
# Two inputs, nothing held fixed


@needs_git
def test_two_inputs_with_nothing_held_fixed(two, tmp_path):
    tests = two["tests"]
    assert list(tests) == ["UTIL", "income_to_sales"]
    for t in tests.values():
        assert not t.held and t.strata == ()
        assert (t.holdout.loans, t.holdout.bad) == _counts(two["x"], t.column, list(t.bins), HOLD)
    fam = _family(list(tests.values()), "holdout")
    assert len(fam) == 3 + 5
    want = dict(zip([(c, k) for c, k, _ in fam], _bh([p for _, _, p in fam])))
    assert tests["UTIL"].holdout.allowed[3] == pytest.approx(want[("UTIL", 3)], rel=1e-12)
    assert want[("UTIL", 3)] < 0.05
    rows = _table(recalc(two["b"], tmp_path / "calc")[confirm_tab.SHEET])
    util = rows[("UTIL", f"{tests['UTIL'].groups[3]} vs {tests['UTIL'].groups[0]}")]
    assert util[confirm_tab.N_HOLDS] == "Yes" and util[confirm_tab.N_WORDS] == "Holds up on held-back loans"
    assert _check(two["b"])["Differs from the pre-spec"] == "nowhere: this run used what it says"


# --------------------------------------------------------------------------
# The one-column pre-spec, figure for figure


@needs_git
def test_the_one_column_pre_spec_is_confirmed_figure_for_figure_as_a_shortlist_of_one(tmp_path, monkeypatch):
    """The same book, once with the committed one-column example and once with the same input written as a
    shortlist of one: every number on New variables is the same, and so is each raw statistic phase 4 showed."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    got = {}
    for name, text in (("old", EXAMPLE.read_text(encoding="utf-8")),
                       ("list", yaml.safe_dump({**BASE, "written": "2026-09-25", "outcome": "BAD_FLAG",
                                                "inputs": [INCOME]}, sort_keys=False))):
        folder = tmp_path / name
        monkeypatch.setenv("GIT_CEILING_DIRECTORIES", str(folder))
        x, b = _book(folder, 6000, text=text)
        ran, st, res = _run(b, monkeypatch)
        assert ran.ok, ran.lines
        ws = load_workbook(b)[confirm_tab.SHEET]
        got[name] = ([ws.cell(row=r, column=c).value for r in range(1, ws.max_row + 1) for c in range(2, 27)
                      if isinstance(ws.cell(row=r, column=c).value, (int, float))], st.test)
    assert got["old"][0] == got["list"][0] and len(got["old"][0]) > 100
    t = got["old"][1]
    # the raw statistics are phase 4's own: the regression and the tests on counts made here from the extract
    want = kgroups.conditional_fit([kgroups.Pocket(*_counts(tmp_path / "old" / "loans.csv", "income_to_sales", BINS,
                                                            HOLD))], t.ref, 6)
    assert [round(p, 12) for p in t.holdout_plain.fit.p if p is not None] == \
        [round(p, 12) for p in want.p if p is not None]
    # with Control's "No allowance", every p-value the table shows is phase 4's raw one
    confirmatory.allow([t], "none")
    assert [getattr(t, s).allowed for s in confirmatory.SETS] == \
        [[None if k == t.ref else getattr(t, s).fit.p[k] for k in range(6)] for s in confirmatory.SETS]


# --------------------------------------------------------------------------
# The launcher fills Test it and Hold fixed from a saved shortlist


def test_the_launcher_fills_test_it_and_hold_fixed_from_a_shortlist(tmp_path):
    f = launcher.Flow(gate=launcher.AddOns())
    f.pick(str(synth.write_extract(tmp_path, n=1500, ratio=True, shortlist=True)))
    f.set_up()
    assert f.screen() == "L2", f.message
    f.set_mode("new")
    f.outcome = None
    spec = _write_spec(tmp_path, inputs=[{**UTIL}, {**TENURE}, {**INCOME}], strata=["FICO", "CHANNEL"])
    f.pick_shortlist(str(spec))
    assert f.test == ["UTIL", "TENURE", "income_to_sales"] and f.hold == ["FICO", "CHANNEL"]
    assert f.asking == "BAD_FLAG" and f.outcome is None       # the file's outcome, asked about, not picked
    f.answer_outcome(True)
    assert f.outcome == "BAD_FLAG"
    rows = {r["name"]: r for r in f.rows()}
    assert [n for n, r in rows.items() if r["b"] and r["b"]["on"]] == ["UTIL", "TENURE"]   # the ratio isn't made yet
    assert all(r["locked"] for r in rows.values())
    ok, said = f.summary()
    assert ok and said == ("the saved shortlist prespec.yaml: 3 inputs (UTIL, TENURE, income_to_sales) against "
                           "BAD_FLAG, with FICO and CHANNEL held fixed, confirmed on the loans it held back.")
    got = f.choices()
    assert got.test == ("UTIL", "TENURE", "income_to_sales") and got.hold == ("FICO", "CHANNEL")
    assert got.split is None and got.bands == ("FICO",) and got.segments == ("CHANNEL",)
    f.next()
    wb = load_workbook(launcher.book_for(f.extract))
    ws = wb[control.SHEET]
    said = {r[control.KEY_COL - 1].value: r[control.CHOOSE_COL - 1].value for r in ws.iter_rows(min_row=control.FIRST_ROW)}
    assert said["launcher|test"] == "UTIL, TENURE, income_to_sales" and said["launcher|hold"] == "FICO, CHANNEL"
    assert said[control.PRESPEC_KEY] == str(spec)


# --------------------------------------------------------------------------
# Control on a new-variable run (phase 4's loose ends)


@needs_git
def test_control_hides_the_levels_panel_and_names_only_worse_at_on_a_new_variable_run(three, tmp_path):
    ws = load_workbook(three["b"])[control.SHEET]
    panel = [get_column_letter(c) for c in range(control.PANEL_COL, control.PANEL_COL + 4)]
    assert [ws.column_dimensions[c].hidden for c in panel] == [True] * 4
    note = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(4, 12)}
    assert note["Worked out"] == control.METHOD_NEW["Worked out"]
    assert "Fewest loans" not in note["Worked out"] and "better at" not in note["Worked out"]
    assert "pockets" not in note["Materiality levels"].split(".")[0]
    # the allowance is asked on this run now: the shortlist's groups are many tests at once
    assert not ws.row_dimensions[control.row_of(ws, "many_tests")].hidden
    # and a bleed run shows the panel and says what it works out
    wb = load_workbook(three["b"])
    cws = wb[control.SHEET]
    cws.cell(row=control.row_of(cws, "run_kind"), column=control.CHOOSE_COL).value = "Where the book bleeds"
    control.fold_launcher_rows(cws)
    assert [cws.column_dimensions[c].hidden for c in panel] == [False] * 4
    note = {cws.cell(row=r, column=2).value: cws.cell(row=r, column=3).value for r in range(4, 12)}
    assert note["Worked out"] == dict(control.METHOD)["Worked out"]
