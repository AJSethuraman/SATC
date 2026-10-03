"""The audit workbook (the firm, 3 Oct 2026).

"We will need to make this auditable... a demo output mode that would take a file and show the calculations it makes
on one set of things and prove out each one so someone could take their current population run it and then
independently understand the calculation and its steps." Then: "we definitely want to be able to demonstrate and
explain what the formulas are and how to do them by hand." Agreed: a separate workbook, written by a Run when
Control's "Also write the audit workbook?" is Yes, proving every figure for one pocket picked by a live dropdown (the
top flagged pocket to start with), recomputed by Excel from the raw loans.

Every check reads the audit workbook as LibreOffice calculates it. What is checked against the loan file is worked out
here from the CSV's text with the csv module and plain arithmetic, and the shuffles from perm alone."""

import csv
import hashlib
import shutil

import pytest
from openpyxl import load_workbook

from pocketbook import audit, book, choices as ch, control, engine, perm, synth
from test_book import _answer

from recalc import need_soffice, recalc

N = 3000
SHUFFLES = 2000
TICK = audit.TICK


def _set_control(path, key, value):
    wb = load_workbook(path)
    for r in wb[control.SHEET].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == key:
            r[control.CHOOSE_COL - 1].value = value
    for r in book.table_rows(wb["Columns"]):           # every odd value but FICO's is real
        key_ = r[book.C_QKEY - 1].value
        if isinstance(key_, str) and key_.count("|") == 2 and not r[book.C_TREAT - 1].value:
            r[book.C_TREAT - 1].value = "Real"
    wb.save(path)


def _build(d, name, audit_answer):
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / f"memory-{name}.yaml"))
        mp.setattr(perm, "SHUFFLES", SHUFFLES)
        src = synth.write_extract(d / name, n=N)
        out = book.set_up(src, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO", "REV_DEBT"),
                                                  segments=("CHANNEL", "ASSET_CLASS"), outcome="BAD_FLAG"))
        _answer(out.book)
        _set_control(out.book, audit.KEY, audit_answer)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return src, out.book, ran


@pytest.fixture(scope="module")
def made(tmp_path_factory):
    need_soffice()
    d = tmp_path_factory.mktemp("audit")
    src, b, ran = _build(d, "on", "Yes")
    path = audit.path_for(b)
    calc = recalc(path, d / "rc")
    yield {"src": src, "book": b, "ran": ran, "audit": path, "calc": calc, "dir": d}
    shutil.rmtree(d, ignore_errors=True)


def _picked(made, name, grid, band, seg):
    """The audit workbook with One pocket's dropdowns set, as LibreOffice calculates it."""
    d = made["dir"] / name
    d.mkdir(exist_ok=True)
    p = d / made["audit"].name
    wb = load_workbook(made["audit"])
    ws = wb[audit.ONE]
    at = _picks(ws)
    ws[at["GRID"]], ws[at["BAND"]], ws[at["SEGMENT"]] = grid, band, seg
    wb.save(p)
    return recalc(p, d / "rc")


def _picks(ws) -> dict:
    out = {}
    for row in ws.iter_rows(max_row=20, max_col=6):
        for c in row:
            if c.value in ("GRID", "BAND", "SEGMENT"):
                out[c.value] = f"{c.column_letter}{c.row + 1}"
    return out


def _table(ws) -> list[dict]:
    """One pocket's rows: step, Excel's figure, PocketBook's, Ties?."""
    head = next(r for r in range(1, 40) if ws.cell(row=r, column=2).value == "Step")
    out = []
    for r in range(head + 1, ws.max_row + 1):
        step, ties = ws.cell(row=r, column=2).value, ws.cell(row=r, column=7).value
        if ties is None:
            continue
        out.append({"step": step, "excel": ws.cell(row=r, column=5).value, "pb": ws.cell(row=r, column=6).value,
                    "ties": ties, "written": ws.cell(row=r, column=4).value, "row": r})
    return out


def _engine_pockets(made):
    """PocketBook's own figures, read from the _pocketbook sheet the audit stores: (grid, band, seg) -> row."""
    ws = load_workbook(made["audit"], read_only=True)[audit.PB]       # as written: the figures at full precision
    heads = [c.value for c in ws[1]]
    return {(r[1], r[2], r[3]): dict(zip(heads, r)) for r in ws.iter_rows(min_row=2, values_only=True)}


# --------------------------------------------------------------------------
# One pocket


EXPECTED_STEPS = {"Loans", "Bad loans", "Bad loan rate", "Booked, loans with a GCO", "GCOs", "GCO rate",
                  "Booked, loans with a RANR", "RANR", "RANR rate", "GCO rate, rest of the book", "× book",
                  "× rest of the book", "Gap in points", "Dollars above share, against the book",
                  "RANR gap in points", "GCO rate, rest of its band", "× rest of its band",
                  "Dollars above share, against its band", "Dollars above share, as shown", "Avg line",
                  "Line × book", "p-value", "p-value after the allowance for many tests"}


def test_every_figure_of_the_default_pocket_ties(made):
    ws = made["calc"][audit.ONE]
    rows = _table(ws)
    assert EXPECTED_STEPS <= {r["step"] for r in rows}
    assert [r for r in rows if r["ties"] != TICK] == []
    # the default is the top flagged pocket: worse and material on GCOs, largest dollars above its share
    at = _picks(ws)
    picked = (ws[at["GRID"]].value, ws[at["BAND"]].value, ws[at["SEGMENT"]].value)
    pockets = _engine_pockets(made)
    flagged = [(v["dollars"], k) for k, v in pockets.items() if v["worse_material"] is True]
    assert flagged and picked == max(flagged)[1]
    # a figure written out reads its own numbers: "$765,455 ÷ $6,779,970 = 11.290%"
    rate = next(r for r in rows if r["step"] == "GCO rate")
    assert rate["written"].endswith(f"= {rate['excel'] * 100:.3f}%") and " ÷ " in rate["written"]
    assert all(isinstance(r["excel"], (int, float)) for r in rows), "every figure of the default pocket is shown"


def _f(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def test_the_loans_sheet_carries_the_extract_so_the_pocket_adds_up_from_the_csv(made):
    """The default pocket's loans, booked and GCOs worked out from the CSV's text, with the band's edges as the Bands
    sheet states them and FICO's -9999 left out as the Run was told."""
    ws = made["calc"][audit.ONE]
    at = _picks(ws)
    grid, band, seg = ws[at["GRID"]].value, ws[at["BAND"]].value, ws[at["SEGMENT"]].value
    bcol, scol = grid.split(" x ")
    bands = made["calc"][audit.BANDS]
    lo = hi = None
    for r in range(1, bands.max_row + 1):
        if bands.cell(row=r, column=2).value == band:
            lo, hi = bands.cell(row=r, column=3).value, bands.cell(row=r, column=4).value
            break
    assert lo is not None
    rows = list(csv.DictReader(open(made["src"], encoding="utf-8")))
    inside = [r for r in rows if r[scol] == seg and _f(r[bcol]) is not None and _f(r[bcol]) != -9999
              and lo <= _f(r[bcol]) < hi]
    pairs = [(_f(r["GCO_AMT"]), _f(r["ORIG_BAL"])) for r in inside]
    pairs = [(a, b) for a, b in pairs if a is not None and b is not None]
    got = {r["step"]: r["excel"] for r in _table(ws)}
    assert got["Loans"] == len(inside)
    assert got["GCOs"] == pytest.approx(sum(a for a, _ in pairs), abs=1e-6)
    assert got["Booked, loans with a GCO"] == pytest.approx(sum(b for _, b in pairs), abs=1e-6)
    assert got["Bad loans"] == sum(1 for r in inside if r["BAD_FLAG"] == "1")


@pytest.mark.parametrize("which", ["other grid", "small pocket", "no number band"])
def test_every_figure_ties_for_a_pocket_picked_with_the_dropdowns(made, which):
    pockets = _engine_pockets(made)
    ws0 = made["calc"][audit.ONE]
    at = _picks(ws0)
    default = (ws0[at["GRID"]].value, ws0[at["BAND"]].value, ws0[at["SEGMENT"]].value)
    if which == "other grid":
        key = next(k for k in pockets if k[0] == "REV_DEBT x ASSET_CLASS" and pockets[k]["raw"] != audit.NONE)
    elif which == "small pocket":
        key = min((k for k in pockets if k != default), key=lambda k: pockets[k]["loans"])
    else:
        key = next(k for k in pockets if k[1] == engine.MISSING_RULE_LABEL)
    ws = _picked(made, which.replace(" ", "-"), *key)[audit.ONE]
    rows = _table(ws)
    assert {r["step"] for r in rows} >= EXPECTED_STEPS
    # every figure ties; the shuffles' count is listed for the default pocket only, so it reads "–" there
    assert [r for r in rows if r["ties"] != TICK and not (r["ties"] == "–" and r["step"].startswith("Shuffles"))] == []
    got = {r["step"]: r["excel"] for r in rows}
    assert got["Loans"] == pockets[key]["loans"]
    assert ws.cell(row=rows[0]["row"], column=5).value == pockets[key]["loans"]


def test_the_dropdowns_list_by_offset_never_index_to_index(made):
    """Excel removed an INDEX:INDEX list from Grids as unreadable (1 Oct 2026): every list here is a range or OFFSET."""
    ws = load_workbook(made["audit"])[audit.ONE]
    forms = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert len(forms) == 3
    assert all("INDEX" not in f for f in forms) and sum("OFFSET(" in f for f in forms) == 2


# --------------------------------------------------------------------------
# Rows in and out, Bands, Run stamp


def _ties_in(ws, col: int) -> list:
    return [(r, ws.cell(row=r, column=col).value) for r in range(1, ws.max_row + 1)
            if ws.cell(row=r, column=col).value in (audit.TICK, audit.CROSS)]


def test_rows_in_and_out_tie_to_the_run(made):
    ws = made["calc"][audit.ROWS]
    got = _ties_in(ws, 5)
    assert len(got) >= 10 and all(v == TICK for _, v in got)
    rows = list(csv.DictReader(open(made["src"], encoding="utf-8")))
    assert ws.cell(row=got[0][0], column=3).value == len(rows)
    # the synthetic book's dirt, each counted where it is left out: '#N/A' as a GCO, a blank balance, a flag of 2
    words = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r, _ in got}
    assert words["Less: GCO_AMT not a number"] == 1
    assert words["Less: ORIG_BAL blank"] == 1
    assert words["Less: BAD_FLAG not 0 or 1"] == 1


def test_every_loans_band_by_formula_is_the_engines_band(made):
    calc = made["calc"]
    ws = calc[audit.LOANS]
    heads = [c.value for c in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == N
    for field in ("FICO", "REV_DEBT"):
        mine, theirs = heads.index(f"{field} band"), heads.index(f"{field} band, PocketBook's")
        assert all(r[mine] == r[theirs] for r in rows), field
        assert all(isinstance(r[mine], str) and r[mine] for r in rows)
    # PocketBook's column is the engine's: the default pocket's loans are exactly the loans in its band and segment
    assert all(v == TICK for _, v in _ties_in(calc[audit.BANDS], 7))
    written = load_workbook(made["audit"])[audit.LOANS]
    assert str(written.cell(row=5, column=heads.index("FICO band") + 1).value).startswith("=IF(ISNUMBER(")


def test_the_run_stamp_fingerprint_is_the_files_sha256(made):
    ws = made["calc"][audit.STAMP]
    said = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(1, ws.max_row + 1)}
    assert said["Its SHA-256 fingerprint"] == hashlib.sha256(made["src"].read_bytes()).hexdigest()
    assert said["Input file"] == made["src"].name and said["Rows read"] == N
    assert said["Shuffles per test"] == SHUFFLES
    assert said["Shuffle seed"] == str(perm.seed_of(engine.SHUFFLE_SEED_NAME))


# --------------------------------------------------------------------------
# The shuffle test


def _cells(ws, label: str) -> tuple:
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == label:
            return ws.cell(row=r, column=4).value, ws.cell(row=r, column=5).value, ws.cell(row=r, column=6).value
    raise KeyError(label)


def test_the_shuffled_gaps_countif_reproduces_the_p_value_exactly(made):
    ws = made["calc"][audit.SHUFFLE]
    pockets = _engine_pockets(made)
    one = made["calc"][audit.ONE]
    at = _picks(one)
    key = (one[at["GRID"]].value, one[at["BAND"]].value, one[at["SEGMENT"]].value)
    hits, hits_pb, tie = _cells(ws, "Shuffles that count")
    assert hits == hits_pb == pockets[key]["hits"] and tie == TICK
    # the count is a whole number, so it reproduces exactly; the p-value is (count + 1) / (shuffles + 1), which
    # LibreOffice saves to 15 significant digits, so it is read back to that. The sheet's own Ties? compares the
    # numbers as Excel holds them
    rows = [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "p-value"
            and ws.cell(row=r, column=4).value is not None]
    p, tie = ws.cell(row=rows[-1], column=4).value, ws.cell(row=rows[-1], column=6).value
    assert tie == TICK
    # openpyxl writes a number to 16 significant digits, so PocketBook's stored p-value is read back to that
    assert (hits + 1) / (SHUFFLES + 1) == pytest.approx(pockets[key]["raw"], rel=1e-15, abs=0)
    assert p == pytest.approx(pockets[key]["raw"], rel=1e-14, abs=0)
    # the list holds every shuffle, and counting it by hand gives the same
    head = next(r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=2).value).startswith("Every shuffled"))
    gaps = [ws.cell(row=r, column=3).value for r in range(head + 2, head + 2 + SHUFFLES)]
    assert len([g for g in gaps if g is not None]) == SHUFFLES
    line = _cells(ws, "The line")[0]
    assert sum(1 for g in gaps if isinstance(g, str) or abs(g) >= line) == hits
    # and the 10-loan example ties shuffle by shuffle to perm's own draws
    assert all(v == TICK for _, v in _ties_in(ws, 11)) and len(_ties_in(ws, 11)) == audit.EXAMPLE_SHUFFLES
    assert all(v == TICK for _, v in _ties_in(ws, 6))


def test_the_example_is_perms_own_shuffle(made):
    """The worked example's p-value is perm.pocket_vs_rest's, and its loans dealt the label are the first of each
    shuffled order: what the Run does, on ten loans."""
    ex = audit.EXAMPLE
    y = [float(g) for _, g in ex]
    x = [float(b) for b, _ in ex]
    k = audit.EXAMPLE_IN
    ans, _ = perm.pocket_vs_rest(y[:k], x[:k], y[k:], x[k:], shuffles=audit.EXAMPLE_SHUFFLES, seed=audit.EXAMPLE_SEED)
    ws = made["calc"][audit.SHUFFLE]
    rows = [r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "p-value"
            and ws.cell(row=r, column=4).value is not None]
    assert ws.cell(row=rows[0], column=4).value == pytest.approx(ans.p, abs=1e-15)
    assert 0 < ans.hits < audit.EXAMPLE_SHUFFLES          # some shuffles count and some don't: worth following


def test_the_run_says_it_wrote_the_audit_workbook(made):
    assert any(made["audit"].name in x for x in made["ran"].lines)


def test_audit_off_writes_no_audit_workbook(tmp_path):
    for answer in (None, "No"):
        _, b, ran = _build(tmp_path, f"off-{answer}", answer)
        assert not audit.path_for(b).exists()
        assert not any("audit" in x for x in ran.lines)
