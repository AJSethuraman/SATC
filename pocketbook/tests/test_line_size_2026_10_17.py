"""Line size (the firm, 17 Oct 2026: "we give assignments that are too high to these people"), in a view of its own
"so it isn't all in one spot and too much for a user", shown when an indicator trips.

The firm's decisions: literal names; each line editable on Control, starting where they said (1.25×, 1.25×, 80%,
0.90×), marked suggested; "Worse" for tripping on a switch, Bad loans / GCOs / Either, starting on Either; the four
ratios that point in a risk direction counted as Signals; and "let's just look at an example after".

The figures: Line on bad loans ÷ line on good loans, Share of line lost on bad loans, Months to charge-off weighted
by line, and Line-weighted ÷ straight, beside Avg line and Line × book. Checked twice: by hand on a book of a few
loans whose every number is written below, and on the synthetic book with the line-size plant (synth.add_big_lines)
by plain arithmetic on the CSV's text, nothing imported from pocketbook for the arithmetic. The tab's Show
dropdown, its signals, Tripped by and the switch are formulas, so they are read after a LibreOffice recalculation."""

from __future__ import annotations

import csv
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from pocketbook import book, choices as ch, control, engine, glossary, live, results, synth
from test_book import _answer

import tabs
from recalc import need_soffice

TAB = results.LINE
GRID = "FICO x CHANNEL"
COLS = {results.Y_BAND: "band", results.Y_SEG: "seg", results.Y_LOANS: "loans", results.Y_AVG: "avg",
        results.Y_AVGX: "avg_x", results.Y_BADX: "bad_vs_good", results.Y_LOST: "lost_share",
        results.Y_MO: "co_straight", results.Y_MOW: "co_weighted", results.Y_MOX: "co_ratio",
        results.Y_SIG: "signals", results.Y_TRIP: "tripped"}
#: the starting lines, as the firm set them
START = {"line_big": 1.25, "line_bad_bigger": 1.25, "line_lost": 0.8, "line_sooner": 0.9}
SIGNAL_OF = {"line_big": ("Big line", "avg_x"), "line_bad_bigger": ("Bad got bigger lines", "bad_vs_good"),
             "line_lost": ("Lost the whole line", "lost_share"), "line_sooner": ("Big lines fail sooner", "co_ratio")}
SHOW = ("Pockets where something tripped", "2 or more signals", "All pockets")


def _build(d: Path, name: str, co: bool = True, n: int = 4000) -> Path:
    from pocketbook import perm
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / f"memory-{name}.yaml"))
        mp.setattr(perm, "SHUFFLES", 100)
        x = synth.write_extract(d / name, n=n, chargeoff=co, big_lines=True)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                                outcome="BAD_FLAG"))
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return out.book


@pytest.fixture(scope="module")
def made(tmp_path_factory):
    """The synthetic book with the line-size plant, with and without a charge-off date column."""
    d = tmp_path_factory.mktemp("line_size")
    # 8,000 loans: enough for a pocket worse on GCOs ($) alone, so each choice of the switch reads differently
    return {"co": _build(d, "co", n=8000), "noco": _build(d, "noco", co=False), "d": d}


# --------------------------------------------------------------------------
# reading the tab


def _rows(ws) -> list[dict]:
    h = tabs.header_row(ws, results.Y_BAND, "Band")
    out = []
    for r in range(h + 1, ws.max_row + 1):
        if ws.cell(row=r, column=results.Y_BAND).value in (None, ""):
            break
        out.append({**{k: tabs.blank(ws.cell(row=r, column=c).value) for c, k in COLS.items()}, "row": r})
    return out


def _shown(path, out, **picks) -> list[dict]:
    """The tab's rows with its dropdowns set, calculated."""
    return _rows(tabs.calculated(tabs.choose(path, out, TAB, **picks), TAB))


def _control(path, out, own: dict | None = None, pick: dict | None = None) -> Path:
    """A copy with Control's line-size answers changed: a number typed under Or your own, or an option picked."""
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(path, out)
    wb = load_workbook(out)
    ws = wb[control.SHEET]
    for key, v in (own or {}).items():
        ws.cell(row=control.row_of(ws, key), column=control.OWN_COL).value = v
    for key, v in (pick or {}).items():
        ws.cell(row=control.row_of(ws, key), column=control.CHOOSE_COL).value = v
    wb.save(out)
    return out


def _views(path) -> dict:
    """Line size's pockets as the Run put them on _views: {(band, segment): figures by name}, and their order."""
    ws = load_workbook(path)[results.VIEWS]
    got, order = {}, []
    for r in ws.iter_rows(values_only=True):
        if isinstance(r[0], str) and r[0].startswith(f"L|{GRID}|"):
            x = dict(zip(results.LINE_VIEW, r[1:]))
            got[(str(x["band"]), str(x["seg"]))] = x
            order.append((int(r[0].rsplit("|", 1)[1]), (str(x["band"]), str(x["seg"]))))
    return got, [k for _, k in sorted(order)]


def _worse(path, d) -> dict:
    """Worse? on Bad loans and on GCOs ($) for every pocket of the grid, as _pockets calculates it now."""
    ws = tabs.calculated(path, live.POCKETS) if not hasattr(path, "worksheets") else path[live.POCKETS]
    out = {}
    for r in range(live.P_FIRST, ws.max_row + 1):
        if ws.cell(row=r, column=live.P_KIND).value != "grids" or ws.cell(row=r, column=live.P_GRID).value != GRID:
            continue
        key = (str(ws.cell(row=r, column=live.P_BAND).value), str(ws.cell(row=r, column=live.P_SEG).value))
        m = ws.cell(row=r, column=live.P_MEASURE).value
        if m in ("outcome_loans", "gco_rate"):
            out.setdefault(key, {})[m] = ws.cell(row=r, column=live.P_WORSE).value == live.YES
    return out


def _signals(x: dict, lines: dict, co: bool = True) -> list[str]:
    """The signals a pocket trips, worked out here: each ratio against its line, ≥ but ≤ for the months."""
    out = []
    for key, (name, fig) in SIGNAL_OF.items():
        if key == "line_sooner" and not co:
            continue
        v = x.get(fig)
        if isinstance(v, (int, float)) and (v <= lines[key] if key == "line_sooner" else v >= lines[key]):
            out.append(name)
    return out


def _tripped(worse: dict, signals: list[str], on: str) -> list[str]:
    out = []
    if worse.get("outcome_loans") and on in ("either", "bad"):
        out.append("Worse on bad loans")
    if worse.get("gco_rate") and on in ("either", "gco"):
        out.append("Worse on GCOs")
    if "Bad got bigger lines" in signals:
        out.append("Bad got bigger lines")
    return out


def _said(signals: list[str], n: int = 4) -> str:
    return f"{len(signals)} of {n}" + (": " + ", ".join(signals) if signals else "")


# --------------------------------------------------------------------------
# the arithmetic, by hand on a few loans


def _small():
    """One pocket a person can check: score under 650, channel A. Six loans:

        L0  line 1,000  bad   GCO   900  charged off 3 months after it was made
        L1  line 3,000  bad   GCO 1,500  charged off after 1 month
        L2  line 2,000  good
        L3  no line     bad   GCO   500  charged off after 6 months: out of every line figure, the straight months
                                         too (they are over the same loans as the weighted), but in Grids' own
        L4  line 4,000  good
        L5  line 1,000  bad   GCO "#N/A": in the bad loans' line, out of the share lost (no GCO amount)

    Line on bad ÷ good: (1,000 + 3,000 + 1,000) / 3 over (2,000 + 4,000) / 2 = 1,666.67 / 3,000 = 0.5556.
    Share of line lost: (900 + 1,500) / (1,000 + 3,000) = 0.6000.
    Weighted months: (3 × 1,000 + 1 × 3,000) / 4,000 = 1.5000; straight, over the same two loans: (3 + 1) / 2 = 2.0.
    Line-weighted ÷ straight: 1.5 / 2.0 = 0.7500. Grids' and Summary's own average, every charged-off loan:
    (3 + 1 + 6) / 3 = 3.3333.
    Channel B, at score 700, is there so the grid has a second pocket: one good and one bad loan, each line 500."""
    from conftest import cube, table
    made = "2024-01-15"
    rows = [("L0", 600, "A", 1000, 1, 900, "2024-04-10"), ("L1", 600, "A", 3000, 1, 1500, "2024-02-20"),
            ("L2", 600, "A", 2000, 0, 0, ""), ("L3", 600, "A", "", 1, 500, "2024-07-01"),
            ("L4", 600, "A", 4000, 0, 0, ""), ("L5", 600, "A", 1000, 1, "#N/A", ""),
            ("L6", 700, "B", 500, 0, 0, ""), ("L7", 700, "B", 500, 1, 250, "2024-03-01")]
    data = [{"ID": i, "SCORE": s, "CHAN": c, "BAL": b, "BAD": bad, "GCO": g, "RANR": 10, "OPEN_DT": made,
             "CO_DT": co} for i, s, c, b, bad, g, co in rows]
    cfg = cube(key=None, booked=None, outcome=None, gco=None, ranr=None, columns_confirmed=True,
               columns={"ID": "key", "BAL": "booked", "BAD": "outcome", "GCO": "gco", "RANR": "ranr",
                        "SCORE": "fico", "CHAN": "category", "OPEN_DT": "origination_date",
                        "CO_DT": "chargeoff_date"})
    return engine.run(cfg, table(data))


WANT_A = {"bad_vs_good": (5000 / 3) / 3000, "lost_share": 2400 / 4000, "co_weighted": 6000 / 4000,
          "co_straight": 4 / 2, "co_ratio": (6000 / 4000) / (4 / 2)}


def test_line_size_figures_by_hand_on_a_small_book():
    res = _small()
    g = res.grids[0]
    key = next(k for k, _ in g.inner() if k[1] == "A")
    got = engine.line_vs(g.lines, key)
    for k, want in WANT_A.items():
        assert got[k] == pytest.approx(want, rel=1e-12), k
    assert got["co_ratio"] == pytest.approx(0.75, rel=1e-12)
    # Grids' own months average stays every charged-off loan's, L3 included
    assert g.cells[key].medians[engine.CO_MONTHS].mean == pytest.approx(10 / 3, rel=1e-12)
    # Summary's row for channel A is the same pocket, and carries the four figures
    rows = dict(engine.summary_rows(res, "chan"))
    assert rows["A"]["line_bad_x"] == pytest.approx(WANT_A["bad_vs_good"], rel=1e-12)
    assert rows["A"]["line_lost"] == pytest.approx(WANT_A["lost_share"], rel=1e-12)
    assert rows["A"]["co_line"] == pytest.approx(WANT_A["co_weighted"], rel=1e-12)
    assert rows["A"]["co_line_x"] == pytest.approx(WANT_A["co_ratio"], rel=1e-12)
    assert rows["A"]["co_avg"] == pytest.approx(10 / 3, rel=1e-12)          # Summary's own, unchanged
    # All is every loan, never an average of the rows: bad lines 1,000 + 3,000 + 1,000 + 500 over 4, good 2,000 +
    # 4,000 + 500 over 3
    assert rows["All"]["line_bad_x"] == pytest.approx((5500 / 4) / (6500 / 3), rel=1e-12)
    assert rows["All"]["line_lost"] == pytest.approx((2400 + 250) / (4000 + 500), rel=1e-12)


def test_line_size_figures_are_blank_where_there_is_nothing_to_divide_by():
    """A pocket with no good loan has no Line on bad ÷ good; one with no bad loan has no share lost; no charge-off
    date, no months."""
    x = engine.LineSize(bad_loans=2, bad_booked=100.0)
    assert x.bad_vs_good is None and x.lost_share is None and x.weighted_months is None
    y = engine.LineSize(good_loans=2, good_booked=100.0)
    assert y.bad_vs_good is None and y.lost_share is None
    assert engine.line_vs({}, ("a", "b")) == {"bad_vs_good": None, "lost_share": None, "co_straight": None,
                                                    "co_weighted": None, "co_ratio": None}


# --------------------------------------------------------------------------
# the figures on the synthetic book, from the CSV's text alone


def _csv(path: Path) -> list[dict]:
    src = Path(str(path).replace(" - PocketBook.xlsx", ".csv"))
    return list(csv.DictReader(open(src, encoding="utf-8")))


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _months(r):
    if not r.get("CO_DATE") or not r.get("ORIG_DATE"):
        return None
    o, c = [tuple(int(p) for p in d.split("-")) for d in (r["ORIG_DATE"], r["CO_DATE"])]
    return (c[0] - o[0]) * 12 + (c[1] - o[1]) if c >= o else None


def _band(fico: str, labels) -> str:
    if fico in ("", None):
        return "(blank)"
    v = float(fico)
    if v < -1000:
        return "(marked missing)"
    for lab in labels:
        lo, _, hi = lab.partition(" - ")
        if hi and float(lo) <= v <= float(hi):
            return lab
    raise KeyError(fico)


def _by_hand(rows: list[dict]) -> dict:
    bad = [_num(r["ORIG_BAL"]) for r in rows if r["BAD_FLAG"] == "1" and _num(r["ORIG_BAL"]) is not None]
    good = [_num(r["ORIG_BAL"]) for r in rows if r["BAD_FLAG"] == "0" and _num(r["ORIG_BAL"]) is not None]
    lost = [(_num(r["GCO_AMT"]), _num(r["ORIG_BAL"])) for r in rows if r["BAD_FLAG"] == "1"
            and _num(r["ORIG_BAL"]) is not None and _num(r["GCO_AMT"]) is not None]
    wt = [(_months(r), _num(r["ORIG_BAL"])) for r in rows if _months(r) is not None
          and _num(r["ORIG_BAL"]) is not None]
    mo = [m for m, _ in wt]                 # the straight average over the same loans as the weighted
    out = {"bad_vs_good": (sum(bad) / len(bad)) / (sum(good) / len(good)) if bad and good else None,
           "lost_share": sum(g for g, _ in lost) / sum(b for _, b in lost) if lost else None,
           "co_straight": sum(mo) / len(mo) if mo else None,
           "co_weighted": sum(m * b for m, b in wt) / sum(b for _, b in wt) if wt else None}
    out["co_ratio"] = out["co_weighted"] / out["co_straight"] if out["co_weighted"] is not None \
        and out["co_straight"] else None
    return out


def test_line_size_figures_tie_to_the_loan_file_and_sort_highest_first(made, tmp_path):
    need_soffice()
    b = made["co"]
    shown = _shown(b, tmp_path / "all.xlsx", show=SHOW[2])
    stored, order = _views(b)
    assert [(str(x["band"]), str(x["seg"])) for x in shown] == order and len(shown) == len(stored)
    labels = [x["band"] for x in shown if " - " in str(x["band"])]
    loans = _csv(b)
    for x in shown:
        mine = [r for r in loans if _band(r["FICO"], labels) == x["band"] and r["CHANNEL"] == x["seg"]]
        want = _by_hand(mine)
        assert x["loans"] == len(mine)
        for k, v in want.items():
            if v is None:
                assert x[k] is None, (x["band"], x["seg"], k)
            else:
                assert x[k] == pytest.approx(v, rel=1e-9), (x["band"], x["seg"], k)
    # sorted by Line on bad loans ÷ line on good loans, highest first, a pocket without one last
    ratios = [x["bad_vs_good"] for x in shown]
    have = [v for v in ratios if v is not None]
    assert have == sorted(have, reverse=True) and ratios[:len(have)] == have
    # the plant: Branch, scores 620 to 679, its bad loans on 1.8 times the line (the bands are cut from the book, so
    # a band holds loans either side of the plant's scores)
    top = shown[0]
    assert top["seg"] == synth.BIG_LINES[0] and top["bad_vs_good"] >= START["line_bad_bigger"]


def _view_rows(path) -> dict:
    """Every _views row by its key, the key left off."""
    return {r[0]: list(r[1:]) for r in load_workbook(path)[results.VIEWS].iter_rows(values_only=True)
            if isinstance(r[0], str)}


def test_line_size_grids_measures_tie_to_the_loan_file_inner_and_margins(made):
    """Grids' two line-size Measures, as the Run put them on _views (G|<grid>|line_bad_x|rate|i): an inner pocket,
    its band's All, its segment's All and the grid's All, against the CSV by hand. The margins are every loan of the
    band, segment or grid, never an average of the cells."""
    b = made["co"]
    v = _view_rows(b)
    rows_, cols_ = v[f"G|{GRID}|rows"], v[f"G|{GRID}|cols"]
    rows_ = [x for x in rows_ if x is not None]
    cols_ = [x for x in cols_ if x is not None]
    labels = [x for x in rows_ if " - " in str(x)]
    loans = _csv(b)
    band, seg = labels[1], cols_[0]
    for key, fig in (("line_bad_x", "bad_vs_good"), ("line_lost", "lost_share")):
        for bl, d in ((band, seg), (band, "All"), ("All", seg), ("All", "All")):
            i, j = rows_.index(bl) + 1, cols_.index(d)
            mine = [r for r in loans if (bl == "All" or _band(r["FICO"], labels) == bl)
                    and (d == "All" or r["CHANNEL"] == d)]
            got = v[f"G|{GRID}|{key}|rate|{i}"][j]
            assert got == pytest.approx(_by_hand(mine)[fig], rel=1e-9), (key, bl, d)


def test_line_size_a_blank_switch_reads_as_either(made, tmp_path):
    """Control's Tripped by being worse on cleared: the tab reads it as Either, its starting answer, rather than
    every Tripped by, Shown now and count going to #N/A and the tab going blank."""
    need_soffice()
    b = made["co"]
    blank = _control(b, tmp_path / "blank.xlsx", pick={"line_worse_on": None})
    ws = load_workbook(blank)[control.SHEET]
    assert ws.cell(row=control.row_of(ws, "line_worse_on"), column=control.CHOOSE_COL).value is None
    rows = _shown(blank, tmp_path / "blank-all.xlsx", show=SHOW[2])
    either = _shown(b, tmp_path / "either-all.xlsx", show=SHOW[2])
    assert rows and [(x["band"], x["seg"], x["tripped"]) for x in rows] == \
        [(x["band"], x["seg"], x["tripped"]) for x in either]
    calc = tabs.calculated(tabs.choose(blank, tmp_path / "blank-cap.xlsx", TAB, show=SHOW[0]), TAB)
    cap = [c.value for row in calc.iter_rows() for c in row if isinstance(c.value, str) and " tripped · " in c.value]
    assert cap and cap[0].endswith(f"{sum(1 for x in either if x['tripped'])} shown"), cap
    tile = [c.value for row in calc.iter_rows(max_row=20) for c in row if c.value == "Either"]
    assert tile                                         # the lines in use say what is used


# --------------------------------------------------------------------------
# the tab


def test_line_size_sits_after_ranr_vs_gcos_and_leads_with_its_band_and_note(made):
    wb = load_workbook(made["co"])
    visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    assert visible[visible.index(results.PCK) + 1] == TAB
    ws = wb[TAB]
    first = ws.cell(row=1, column=results.Y_BAND)
    assert first.value == TAB and first.fill.fgColor.rgb.endswith(results.house.INK) and first.font.size == 16
    assert ws.cell(row=3, column=results.Y_BAND).value == "How this tab works"
    assert ws.row_dimensions[4].outlineLevel == 1                       # the note folds, as every result tab's
    assert wb[results.LINES].sheet_state == "hidden"
    head = tabs.header_row(ws, results.Y_BAND, "Band")
    assert tabs.heads(ws, head, results.Y_BAND, results.Y_TRIP) == list(results.LINE_HEADS)
    assert tabs.options(wb, TAB, "Grid")[0] == GRID and tabs.options(wb, TAB, "Show") == list(SHOW)
    assert tabs.dropdown(ws, "Show").value == SHOW[0]                   # tripped, by default
    # Start here lists it, a link, among the results
    assert any(c.value and str(c.value).startswith(f"{TAB}: ") and c.hyperlink
               for row in wb["Start here"].iter_rows() for c in row)


def test_line_size_show_picks_the_rows(made, tmp_path):
    """Pockets where something tripped (worse on the measures the switch picks, or Bad got bigger lines), 2 or more
    signals, and every pocket: each the rows of All pockets that qualify, in the same order."""
    need_soffice()
    b = made["co"]
    every = _shown(b, tmp_path / "all.xlsx", show=SHOW[2])
    worse = _worse(b, tmp_path)
    stored, _ = _views(b)
    for x in every:
        sig = _signals(stored[(str(x["band"]), str(x["seg"]))], START)
        assert x["signals"] == _said(sig), (x["band"], x["seg"])
        assert (x["tripped"] or "") == ", ".join(_tripped(worse[(str(x["band"]), str(x["seg"]))], sig, "either"))
    key = lambda xs: [(x["band"], x["seg"]) for x in xs]                # noqa: E731
    tripped = [x for x in every if x["tripped"]]
    two = [x for x in every if int(str(x["signals"]).split(" of ")[0]) >= 2]
    assert tripped and two and len(two) < len(every) and len(tripped) < len(every)
    assert key(_shown(b, tmp_path / "tripped.xlsx", show=SHOW[0])) == key(tripped)
    assert key(_shown(b, tmp_path / "two.xlsx", show=SHOW[1])) == key(two)
    # the caption counts what the dropdowns select
    ws = tabs.calculated(tabs.choose(b, tmp_path / "cap.xlsx", TAB, show=SHOW[1]), TAB)
    cap = next(c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and " shown" in c.value
               and " tripped · " in c.value)
    assert cap == f"{len(tripped)} tripped · {len(two)} with 2 or more signals · {len(two)} shown"


def test_line_size_each_signal_trips_at_its_line(made, tmp_path):
    """Each line typed on Control at exactly one pocket's ratio trips that pocket (≥, and ≤ for the months), and a
    hair past it does not. The ratios are the Run's own, read off _views, so the comparison is exact."""
    need_soffice()
    b = made["co"]
    stored, _ = _views(b)
    target = {}
    for key, (name, fig) in SIGNAL_OF.items():
        pocket = next(k for k, x in stored.items() if isinstance(x[fig], float))
        target[key] = (pocket, stored[pocket][fig])
    at = _shown(_control(b, tmp_path / "at.xlsx", own={k: v for k, (_, v) in target.items()}),
                tmp_path / "at-all.xlsx", show=SHOW[2])
    nudge = {k: v * (1 - 1e-9) if k == "line_sooner" else v * (1 + 1e-9) for k, (_, v) in target.items()}
    past = _shown(_control(b, tmp_path / "past.xlsx", own=nudge), tmp_path / "past-all.xlsx", show=SHOW[2])
    lines_at = {k: v for k, (_, v) in target.items()}
    for key, (pocket, _) in target.items():
        name = SIGNAL_OF[key][0]
        x = next(r for r in at if (str(r["band"]), str(r["seg"])) == pocket)
        y = next(r for r in past if (str(r["band"]), str(r["seg"])) == pocket)
        assert name in x["signals"], (key, x["signals"])
        assert name not in y["signals"], (key, y["signals"])
    # and every other pocket reads by the same rule, on the Run's own figures (a calculated copy's are rounded)
    for got, lines in ((at, lines_at), (past, nudge)):
        for x in got:
            assert x["signals"] == _said(_signals(stored[(str(x["band"]), str(x["seg"]))], lines)), (x["band"],
                                                                                                   x["seg"])


def test_line_size_switch_changes_tripped_by(made, tmp_path):
    """Tripped by being worse on: Bad loans counts Worse? on Bad loans only, GCOs on GCOs ($) only, Either both; Bad
    got bigger lines counts whatever the switch says."""
    need_soffice()
    b = made["co"]
    worse = _worse(b, tmp_path)
    stored, _ = _views(b)
    seen = set()
    for label, on in (("Bad loans", "bad"), ("GCOs", "gco"), ("Either (suggested)", "either")):
        rows = _shown(_control(b, tmp_path / f"{on}.xlsx", pick={"line_worse_on": label}), tmp_path / f"{on}-all.xlsx",
                      show=SHOW[2])
        said = []
        for x in rows:
            k = (str(x["band"]), str(x["seg"]))
            want = _tripped(worse[k], _signals(stored[k], START), on)
            assert (x["tripped"] or "") == ", ".join(want), (on, x["band"], x["seg"])
            said.append(x["tripped"])
        seen.add(tuple(said))
    assert len(seen) == 3                               # each choice reads differently on this book


def test_line_size_without_a_charge_off_date_counts_out_of_3(made, tmp_path):
    need_soffice()
    b = made["noco"]
    wb = load_workbook(b)
    ws = wb[TAB]
    for c in results.Y_MONTHS:
        assert ws.column_dimensions[results.col(c)].hidden, c
    rows = _shown(b, tmp_path / "all.xlsx", show=SHOW[2])
    assert rows and all(str(x["signals"]).split(": ")[0].endswith(" of 3") for x in rows)
    assert not any("Big lines fail sooner" in str(x["signals"]) for x in rows)
    stored, _ = _views(b)
    for x in rows:
        assert x["signals"] == _said(_signals(stored[(str(x["band"]), str(x["seg"]))], START, co=False), 3)
    note = " ".join(str(c.value) for row in ws.iter_rows(max_row=20) for c in row if c.value)
    assert "no column marked Charge-off date" in note and "out of 3" in note
    # Summary carries the two figures that need no dates, and not the months
    heads = _summary_heads(b)
    assert {results.LINE_HEADS[5], results.LINE_HEADS[6]} <= set(heads)
    assert not {results.LINE_HEADS[8], results.LINE_HEADS[9]} & set(heads)


def _summary_heads(path) -> list:
    ws = load_workbook(path)[results.SUMMARY]
    h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
    heads, c = [], 3
    while ws.cell(row=h, column=c).value not in (None, ""):
        heads.append(ws.cell(row=h, column=c).value)
        c += 1
    return heads


def test_line_size_shading_is_red_past_each_line_and_one_range_a_cell(made):
    """Each ratio cell is red (the workbook's worse fill and crimson) when it passes its line, by a rule that reads the
    line by name; nothing else is judged; every cell sits in one conditional-format range (results.cf)."""
    ws = load_workbook(made["co"])[TAB]
    first = tabs.header_row(ws, results.Y_BAND, "Band") + 1
    rules = {}
    for cr in ws.conditional_formatting:
        for part in str(cr.sqref).split():
            c0, r0, c1, r1 = range_boundaries(part)
            if r0 <= first <= r1:
                for c in range(c0, c1 + 1):
                    rules.setdefault(c, []).append(cr)
    for c in range(results.Y_BAND, results.Y_TRIP + 1):
        assert len(rules.get(c, [])) == 1, (results.col(c), len(rules.get(c, [])))
    lines = {results.Y_AVGX: ("line_big", ">="), results.Y_BADX: ("line_bad_bigger", ">="),
             results.Y_LOST: ("line_lost", ">="), results.Y_MOX: ("line_sooner", "<=")}
    for c in range(results.Y_BAND, results.Y_TRIP + 1):
        filled = [r for r in rules[c][0].rules if r.dxf is not None and r.dxf.fill is not None]
        if c not in lines:
            assert not filled, results.col(c)
            continue
        name, how = lines[c]
        assert len(filled) == 1, results.col(c)
        assert f"{how}{name}" in filled[0].formula[0] and filled[0].dxf.fill.fgColor.rgb.endswith(results.ALERT)
        assert filled[0].dxf.font.color.rgb.endswith(results.CRIMSON)


def test_line_size_lines_sit_on_control_suggested_in_a_block_of_their_own(made):
    need_soffice()
    wb = load_workbook(made["co"])
    ws = wb[control.SHEET]
    head = control.row_of(ws, control.BLOCK_LINES)
    last_launcher = max(control.launcher_rows_of(ws))
    assert head > last_launcher                          # under the launcher's block: no row above moved
    settings = {s.key: s for s in control.load_settings()}
    for key in control.LINE_KEYS:
        r = control.row_of(ws, key)
        s = settings[key]
        assert head < r and s.takes_effect == "live" and s.added
        picked = ws.cell(row=r, column=control.CHOOSE_COL).value
        assert picked == s.start().shown and picked.endswith("(suggested)"), key
        if key in START:
            assert s.start().value == START[key]
    assert settings["line_worse_on"].start().value == "either"
    # Record keeps them: in use now and what the Run used, in _live's words
    got = tabs.settings(tabs.calculated(made["co"], "Record").parent)
    assert got[settings["line_big"].question][:2] == ("1.25×", "1.25×")
    assert got[settings["line_lost"].question][:2] == ("80%", "80%")
    assert got[settings["line_worse_on"].question][:2] == ("Either", "Either")


def test_line_size_an_older_workbook_without_the_line_rows_still_runs(made, tmp_path, monkeypatch):
    """A workbook set up before 17 Oct 2026 has no line-size rows on Control. Run reads the starting lines, writes the
    tab and records them; the next Set up writes the rows, on their starting lines."""
    need_soffice()
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(made["d"] / "memory-co.yaml"))
    from pocketbook import perm
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    b = tmp_path / "old" / Path(made["co"]).name
    b.parent.mkdir()
    shutil.copy(made["co"], b)
    shutil.copy(str(made["co"]).replace(" - PocketBook.xlsx", ".csv"), b.parent / "loans.csv")
    wb = load_workbook(b)
    ws = wb[control.SHEET]
    top = control.row_of(ws, control.BLOCK_LINES)
    ws.delete_rows(top, ws.max_row - top + 1)
    wb.save(b)
    assert control.row_of(load_workbook(b)[control.SHEET], "line_big") is None
    ran = book.run(b, extract=b.parent / "loans.csv")
    assert ran.ok, ran.lines
    rows = _shown(b, tmp_path / "old-all.xlsx", show=SHOW[2])
    stored, _ = _views(b)
    assert rows and all(x["signals"] == _said(_signals(stored[(str(x["band"]), str(x["seg"]))], START))
                        for x in rows)
    got = tabs.settings(tabs.calculated(b, "Record").parent)
    big = next(s for s in control.load_settings() if s.key == "line_big")
    assert got[big.question][:2] == ("1.25 times", "1.25 times")
    again = book.set_up(b.parent / "loans.csv", b)
    assert again.ok, again.lines
    ws = load_workbook(b)[control.SHEET]
    assert ws.cell(row=control.row_of(ws, "line_big"), column=control.CHOOSE_COL).value == big.start().shown


def test_line_size_summary_and_grids_carry_the_figures(made):
    wb = load_workbook(made["co"])
    heads = _summary_heads(made["co"])
    i = heads.index("Booked $")
    assert heads[i:i + 6] == ["Booked $", "Avg line", "Line × book", "% of booked", results.LINE_HEADS[5],
                              results.LINE_HEADS[6]]
    j = heads.index("Median months to charge-off")
    assert heads[j + 1:j + 3] == [results.LINE_HEADS[8], results.LINE_HEADS[9]]
    # Grids: shown and not compared, like Months to charge-off (avg)
    opts = tabs.options(wb, results.GRIDS, "Measure")
    k = opts.index(results.SIZE_NAME)
    assert opts[k + 1:k + 3] == [results.LINE_HEADS[5], results.LINE_HEADS[6]]
    views = {r[0] for r in wb[results.VIEWS].iter_rows(values_only=True) if isinstance(r[0], str)}
    assert f"G|{GRID}|line_bad_x|rate|1" in views and f"G|{GRID}|line_lost|meta" in views
    assert f"G|{GRID}|line_bad_x|book|1" not in views                  # never compared


def test_line_size_glossary_has_a_line_per_term(made):
    wb = load_workbook(made["co"])
    terms = [r[0] for r in wb[glossary.SHEET].iter_rows(min_col=glossary.TERM, max_col=glossary.TERM,
                                                       values_only=True) if r[0]]
    for t in (results.LINE_HEADS[5], results.LINE_HEADS[6], results.LINE_HEADS[8], results.LINE_HEADS[9], "Signals",
              "Big line", "Bad got bigger lines", "Lost the whole line", "Big lines fail sooner", "Tripped by"):
        assert terms.count(t) == 1, t
    without = [r[0] for r in load_workbook(made["noco"])[glossary.SHEET].iter_rows(
        min_col=glossary.TERM, max_col=glossary.TERM, values_only=True) if r[0]]
    for t in (results.LINE_HEADS[8], results.LINE_HEADS[9], "Big lines fail sooner"):
        assert t not in without
