"""What is learned can be seen and pruned. The firm: "we also need an intuitive
way to go and prune rules that shouldn't have been added as we learn." """

import re

from openpyxl import load_workbook

from pocketbook import cli, config as cfgmod, memory, profile, synth
from pocketbook.ingest import read_table

ANSWERS = dict(min_units=30, min_events=10, worse_at=1.25, better_at=0.8, confidence=0.95, compare_to='peers', materiality='"1% of losses"')


def _confirmed_file(tmp_path, n=2000):
    _, data = synth.write(tmp_path, n=n)
    path, _, _ = profile.write_cube_file(read_table(data), tmp_path / "cube.yaml")
    s = path.read_text().replace("columns_confirmed: no", "columns_confirmed: yes")
    for k, v in ANSWERS.items():
        s = re.sub(rf'{k}: "\[CONFIRM[^"]*"', f"{k}: {v}", s)
    s = s.replace("{column: FICO, pattern: repeated_value, value: -9999, rows: 40, answer: }",
                  "{column: FICO, pattern: repeated_value, value: -9999, rows: 40, answer: missing}")
    path.write_text(s)
    return path, data


def test_a_confirmed_run_is_remembered_and_an_unconfirmed_one_is_not(tmp_path, capsys):
    path, data = _confirmed_file(tmp_path)
    assert cli.main(["validate", str(path), "--data", str(data)]) == 0
    mem = memory.load()
    assert mem["columns"]["FICO"]["means"] == "fico" and mem["columns"]["GCO_AMT"]["means"] == "gco"
    assert any(v["answer"] == "missing" for v in mem["answers"].values())
    path.write_text(path.read_text().replace("columns_confirmed: yes", "columns_confirmed: no"))
    before = memory.load()
    assert cli.main(["validate", str(path), "--data", str(data)]) == 2
    assert memory.load() == before


def test_the_next_init_uses_what_was_learned(tmp_path):
    path, data = _confirmed_file(tmp_path)
    cli.main(["validate", str(path), "--data", str(data)])
    again, _, _ = profile.write_cube_file(read_table(data), tmp_path / "again.yaml")
    text = again.read_text()
    assert "REMEMBERED - you confirmed this as a FICO score" in text
    assert "answer: missing}   # REMEMBERED - you answered missing" in text


def test_nothing_from_the_data_is_stored(tmp_path):
    path, data = _confirmed_file(tmp_path)
    cli.main(["validate", str(path), "--data", str(data)])
    stored = memory.default_path().read_text()
    for r in read_table(data).rows[:50]:
        assert r["LOAN_NBR"] not in stored


def test_forget_by_name(tmp_path, capsys):
    path, data = _confirmed_file(tmp_path)
    cli.main(["validate", str(path), "--data", str(data)])
    assert cli.main(["memory", "--forget", "FICO"]) == 0
    mem = memory.load()
    assert "FICO" not in mem["columns"] and not any(v["column"] == "FICO" for v in mem["answers"].values())
    assert "GCO_AMT" in mem["columns"]


def test_prune_in_excel(tmp_path):
    path, data = _confirmed_file(tmp_path)
    cli.main(["validate", str(path), "--data", str(data)])
    review = tmp_path / "learned.xlsx"
    assert cli.main(["memory", "--out", str(review)]) == 0
    wb = load_workbook(review)
    ws = wb["Learned"]
    for row in ws.iter_rows(min_row=4):
        if row[2].value == "CHANNEL":
            row[0].value = memory.FORGET
    wb.save(review)
    assert cli.main(["memory", "--read", str(review)]) == 0
    mem = memory.load()
    assert "CHANNEL" not in mem["columns"] and "FICO" in mem["columns"]


# ---- the rename (Goal 3 item 1, 27 Sep 2026): a machine that ran the Origination Cube keeps what it learned

def _old_home(tmp_path, monkeypatch):
    """A home directory as the Origination Cube left it: its memory and launcher choices under the old folder."""
    home = tmp_path / "home"
    old = home / ".origination-cube"
    old.mkdir(parents=True)
    memory.save({"columns": {"FICO": {"means": "fico", "first": "2026-09-25", "last": "2026-09-25", "times": 1}},
                 "answers": {}}, old / "memory.yaml")
    (old / "launcher.json").write_text('{"extract": "C:/loans/sep.csv", "few": 6, "many": 25}', encoding="utf-8")
    monkeypatch.setattr(memory.Path, "home", lambda: home)
    monkeypatch.delenv("POCKETBOOK_MEMORY", raising=False)
    monkeypatch.delenv("CUBE_MEMORY", raising=False)
    return home


def test_memory_under_the_old_folder_is_still_read_and_the_next_save_writes_the_new_one(tmp_path, monkeypatch):
    home = _old_home(tmp_path, monkeypatch)
    old_text = (home / ".origination-cube" / "memory.yaml").read_text()
    assert memory.read_path() == home / ".origination-cube" / "memory.yaml"
    mem = memory.load()
    assert mem["columns"]["FICO"]["means"] == "fico"
    mem["columns"]["DTI"] = {"means": "dti", "first": "2026-09-27", "last": "2026-09-27", "times": 1}
    assert memory.save(mem) == home / ".pocketbook" / "memory.yaml"
    assert memory.read_path() == home / ".pocketbook" / "memory.yaml"          # the new file now wins
    assert set(memory.load()["columns"]) == {"FICO", "DTI"}
    assert (home / ".origination-cube" / "memory.yaml").read_text() == old_text  # the old file is left alone


def test_the_old_environment_variable_is_read_and_the_new_one_wins(tmp_path, monkeypatch):
    _old_home(tmp_path, monkeypatch)
    monkeypatch.setenv("CUBE_MEMORY", str(tmp_path / "old-env.yaml"))
    assert memory.default_path() == memory.read_path() == tmp_path / "old-env.yaml"
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "new-env.yaml"))
    assert memory.default_path() == memory.read_path() == tmp_path / "new-env.yaml"


def test_the_launchers_last_choices_under_the_old_folder_are_still_read(tmp_path, monkeypatch):
    from pocketbook import launcher
    home = _old_home(tmp_path, monkeypatch)
    monkeypatch.setattr(launcher, "PREFS", None)
    assert launcher._prefs() == {"extract": "C:/loans/sep.csv", "few": 6, "many": 25}
    launcher._save_prefs({"extract": "C:/loans/oct.csv", "few": 6, "many": 25})
    assert (home / ".pocketbook" / "launcher.json").exists()
    assert launcher._prefs()["extract"] == "C:/loans/oct.csv"
    assert "sep.csv" in (home / ".origination-cube" / "launcher.json").read_text()
