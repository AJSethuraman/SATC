"""Three filters, with a size limit (the firm, 1 Oct 2026).

The firm: *"I thought we discussed two filters plus date"*. Chosen: Origination year plus two more filters, with a
size limit. Filter 3 ("and <column> is") sits beside Filter 2 as Filter 2 sits beside Filter 1: each value alone,
each pair and each triple, all holding together. Every grid is built again for each view, so the three are limited
together: (n1 + 1) x (n2 + 1) x (n3 + 1) views, All loans counted for each, at most 150 (6 x 5 x 5).

Every figure the tie-outs check is worked out again here from the loan file with the csv module alone: they read
the calculated workbook with openpyxl, and the arithmetic imports nothing from pocketbook.
"""

import csv
import random
import re
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import range_boundaries

from pocketbook import book, choices as ch, launcher, synth
from test_book import _answer
from test_compare_2026_09_30 import _bad, _close, _fico_of, _fewest, _series, _tables

YEARS = ("2022", "2023", "2024")
FLAGS = ("N", "Y", "(blank)")
REGIONS = ("East", "West")
ONLY_YEAR = f"Only loans where {ch.ORIG_YEAR} is"
AND_FLAG = "and SYS_FLAG is"
AND_REGION = "and REGION is"


def _file(d: Path, n: int, extra: dict | None = None) -> Path:
    """The synthetic book with an origination date (2022 to 2024; every 151st loan has none), SYS_FLAG (Y or N; every
    97th loan blank) and REGION (East or West). `extra` adds columns, each {name: its values}, drawn at random."""
    from datetime import date, timedelta
    src = synth.write_extract(d / "src", n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    rng = random.Random(1001)
    for i, r in enumerate(rows):
        made = date(2022, 1, 1) + timedelta(days=rng.randint(0, 1094))
        r["ORIG_DATE"] = "" if i % 151 == 5 else made.isoformat()
        r["SYS_FLAG"] = "" if i % 97 == 0 else rng.choice("YN")
        r["REGION"] = rng.choice(REGIONS)
        for name, vals in (extra or {}).items():
            r[name] = vals[i % len(vals)] if i < len(vals) else rng.choice(vals)    # every value present
    out = d / "loans.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return out


def _set_up(d: Path, x: Path, **filters):
    from test_book import at
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                            outcome="BAD_FLAG", **filters))
    _answer(out.book)
    wb = load_workbook(out.book)
    wb["Columns"][at(wb, "FICO", book.C_EDGES)] = "620; 680; 740"     # whole scores: labels say the loans
    wb.save(out.book)
    return out.book


@pytest.fixture(scope="module")
def three(tmp_path_factory):
    """Filter 1 ORIG_YEAR, Filter 2 SYS_FLAG, Filter 3 REGION: 5 x 4 x 3 = 60 views of the one grid."""
    from pocketbook import perm
    d = tmp_path_factory.mktemp("three_filters")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 100)
        x = _file(d, 5000)
        b = _set_up(d, x, filter=ch.ORIG_YEAR, filter2="SYS_FLAG", filter3="REGION")
        ran = book.run(b)
        assert ran.ok, ran.lines
    return {"book": b, "csv": x, "ran": ran, "dir": d}


# ---- worked out from the loan file, with nothing from pocketbook


def _rows(x) -> list[dict]:
    return list(csv.DictReader(open(x, encoding="utf-8")))


def _year(r) -> str:
    return r["ORIG_DATE"][:4] if r["ORIG_DATE"] else "(no date)"


def _flag(r) -> str:
    return r["SYS_FLAG"] or "(blank)"


def _keep(rows, year, flag, region, all_="All loans"):
    return [r for r in rows if year in (all_, _year(r)) and flag in (all_, _flag(r))
            and region in (all_, r["REGION"])]


# ---- Grids and Summary: every number of a triple filter is the loan file's


@pytest.mark.parametrize("year,flag,region,where", [
    ("2023", "Y", "East", ", only loans where ORIG_YEAR is 2023 and SYS_FLAG is Y and REGION is East"),
    ("All loans", "All loans", "West", ", only loans where REGION is West"),
    ("2024", "All loans", "East", ", only loans where ORIG_YEAR is 2024 and REGION is East"),
])
def test_three_filters_grids_every_cell_from_the_loan_file(three, tmp_path, year, flag, region, where):
    """Grids with all three dropdowns: the loans with every value picked, cell by cell, and vs the book still
    against the whole book; one cell's words name each filter picked."""
    import tabs
    from test_firm_answers_2026_09_29 import _bad_rate, _cells, _grids
    b, rows = three["book"], _rows(three["csv"])
    wb = load_workbook(b)
    assert tabs.options(wb, "Grids", AND_REGION) == ["All loans", *REGIONS]
    assert tabs.options(wb, "Grids", AND_FLAG) == ["All loans", *FLAGS]
    ws, blocks, said = _grids(b, tmp_path / "g.xlsx", **{ONLY_YEAR: year, AND_FLAG: flag, AND_REGION: region})
    mine = _keep(rows, year, flag, region)
    assert mine
    rate, bk, loans = blocks["Rate"], blocks["vs the book"], blocks["Loans"]
    cells = _cells(mine, loans)
    assert {k for k, n in loans.items() if n} == set(cells)
    book_rate = _bad_rate(rows)
    for k, got in cells.items():
        assert loans[k] == len(got) and rate[k] == pytest.approx(_bad_rate(got)), k
        if isinstance(bk[k], (int, float)):
            assert bk[k] == pytest.approx(_bad_rate(got) / book_rate), k
    assert said["name"].endswith(where), said["name"]


def test_three_filters_grids_pickers_side_by_side_grid_and_measure_where_they_were(three):
    """Filter 3's dropdown is the next one right of Filter 2's, and Row and Column move right of it: no two
    dropdowns' cells overlap, and Grid and Measure keep their columns (B and F, which the bank's checklist names)."""
    import tabs
    ws = load_workbook(three["book"])["Grids"]
    at = {k: tabs.dropdown(ws, k) for k in ("Grid", "Measure", ONLY_YEAR, AND_FLAG, AND_REGION, "Row", "Column")}
    assert len({c.row for c in at.values()}) == 1
    assert (at["Grid"].column_letter, at["Measure"].column_letter) == ("B", "F")
    cols = [at[k].column for k in (ONLY_YEAR, AND_FLAG, AND_REGION, "Row", "Column")]
    assert cols == sorted(cols) and cols[2] - cols[1] == cols[1] - cols[0] == 4
    spans = [range_boundaries(str(m)) for m in ws.merged_cells.ranges if m.min_row == m.max_row == at["Grid"].row]
    for i, a in enumerate(spans):
        for z in spans[i + 1:]:
            assert a[2] < z[0] or z[2] < a[0], (a, z)


def test_three_filters_summary_is_the_loans_with_all_three(three, tmp_path):
    import tabs
    b, rows = three["book"], _rows(three["csv"])
    _fico = _fico_of(rows)
    assert tabs.options(load_workbook(b), "Summary", AND_REGION) == ["All loans", *REGIONS]
    for year, flag, region in (("2024", "N", "West"), ("All loans", "Y", "East")):
        out = tabs.choose(b, tmp_path / f"s{year[:2]}{flag}{region}.xlsx", "Summary",
                          **{"Band or category column": "FICO", ONLY_YEAR: year, AND_FLAG: flag, AND_REGION: region})
        ws = tabs.calculated(out, "Summary")
        h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
        heads = {ws.cell(row=h, column=c).value: c for c in range(3, ws.max_column + 1)}
        shown, r = {}, h + 1
        while ws.cell(row=r, column=2).value not in (None, ""):
            shown[ws.cell(row=r, column=2).value] = (ws.cell(row=r, column=heads["Loans"]).value,
                                                     ws.cell(row=r, column=heads["Bad loans %"]).value)
            r += 1
        mine = _keep(rows, year, flag, region)
        assert shown["All"][0] == len(mine) and _close(shown["All"][1], _bad(mine))
        bands = [lab for lab in shown if " - " in lab]
        assert len(bands) == 4
        for lab in bands:
            pick = [r for r in mine if _fico(r) == lab]
            assert shown[lab][0] == len(pick), (year, flag, region, lab)
            if pick:
                assert _close(shown[lab][1], _bad(pick)), (year, flag, region, lab)


def test_three_filters_named_in_the_runs_lines_and_record(three):
    import tabs
    rows = _rows(three["csv"])
    counts = {g: sum(1 for r in rows if r["REGION"] == g) for g in REGIONS}
    listed = ", ".join(f"{g} ({n:,} loans)" for g, n in counts.items())
    said = f" And by REGION: {listed}; the three together keep the loans with all three."
    assert any(said in line for line in three["ran"].lines), three["ran"].lines
    assert said in tabs.record(three["book"])["Grids filter"]
    ws = load_workbook(three["book"])["Control"]
    from pocketbook import control
    shown = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=control.CHOOSE_COL).value
             for r in range(control.FIRST_ROW, ws.max_row + 1)}
    assert (shown["Filter the Grids by"], shown["And filter them by"], shown["And then by"]) == \
        (ch.ORIG_YEAR, "SYS_FLAG", "REGION")


# ---- Compare: the third filter as Panels by, and as Lines by


def test_compare_third_filter_as_panels_every_point_is_the_loan_files(three, tmp_path):
    """FICO across, lines by Filter 1 (the years), one panel per REGION (Filter 3); SYS_FLAG, in neither, is All
    loans. Each point is that year's loans of that region in that band; thin points are left off the chart."""
    import tabs
    from recalc import recalc_file
    b, rows = three["book"], _rows(three["csv"])
    _fico = _fico_of(rows)
    assert tabs.options(load_workbook(b), "Compare", "Panels by") == \
        ["None", f"Filter 1: {ch.ORIG_YEAR}", "Filter 2: SYS_FLAG", "Filter 3: REGION"]
    calc = recalc_file(tabs.choose(b, tmp_path / "cmp.xlsx", "Compare",
                                   **{"Across the bottom": "FICO", "Measure": "Bad loans",
                                      "Lines by": f"Filter 1: {ch.ORIG_YEAR}", "Panels by": "Filter 3: REGION"}),
                       tmp_path / "rc")
    few = _fewest(calc)
    tables = _tables(calc)
    assert list(tables) == [f"REGION is {g}" for g in REGIONS]
    drawn = _series(tmp_path / "cmp.xlsx", calc)
    for p, region in enumerate(REGIONS):
        shown = tables[f"REGION is {region}"]
        assert len(shown) == 4
        for i, lab in enumerate(shown):
            for k, y in enumerate(YEARS + ("(no date)",)):
                pick = [r for r in rows if _year(r) == y and r["REGION"] == region and _fico(r) == lab]
                rate, n = shown[lab][y]
                if not pick:                                     # no loan: the view is empty, the point blank
                    assert n in (None, "", 0) and drawn[p][1 + k][i] is None, (region, lab, y)
                    continue
                assert n == len(pick) and _close(rate, _bad(pick)), (region, lab, y)
                assert (drawn[p][1 + k][i] is None) == (len(pick) < few), (region, lab, y)
                if len(pick) >= few:
                    assert _close(drawn[p][1 + k][i], _bad(pick))
            whole = [r for r in rows if _fico(r) == lab]
            assert shown[lab]["Whole book"][1] == len(whole)


def test_compare_third_filter_as_lines_panels_by_filter_2(three, tmp_path):
    """Lines by REGION (Filter 3), a panel per SYS_FLAG value (Filter 2); the years, in neither, are All loans."""
    import tabs
    from recalc import recalc_file
    b, rows = three["book"], _rows(three["csv"])
    _fico = _fico_of(rows)
    calc = recalc_file(tabs.choose(b, tmp_path / "cmp.xlsx", "Compare",
                                   **{"Across the bottom": "FICO", "Measure": "Bad loans",
                                      "Lines by": "Filter 3: REGION", "Panels by": "Filter 2: SYS_FLAG"}),
                       tmp_path / "rc")
    tables = _tables(calc)
    assert list(tables) == [f"SYS_FLAG is {f}" for f in FLAGS]
    for flag in FLAGS:
        shown = tables[f"SYS_FLAG is {flag}"]
        for lab in shown:
            for g in REGIONS:
                pick = [r for r in rows if _flag(r) == flag and r["REGION"] == g and _fico(r) == lab]
                rate, n = shown[lab][g]
                if not pick:
                    assert n in (None, "", 0), (flag, lab, g)
                    continue
                assert n == len(pick) and _close(rate, _bad(pick)), (flag, lab, g)


def test_compare_a_filter_across_lines_by_the_third(three, tmp_path):
    """Origination year across the bottom, lines by REGION: one chart, each point the year's loans of the region."""
    import tabs
    from recalc import recalc_file
    b, rows = three["book"], _rows(three["csv"])
    calc = recalc_file(tabs.choose(b, tmp_path / "cmp.xlsx", "Compare",
                                   **{"Across the bottom": "Origination year", "Lines by": "Filter 3: REGION"}),
                       tmp_path / "rc")
    tables = _tables(calc)
    assert list(tables) == ["All loans"] and list(tables["All loans"]) == list(YEARS)
    for y in YEARS:
        for g in REGIONS:
            pick = [r for r in rows if _year(r) == y and r["REGION"] == g]
            rate, n = tables["All loans"][y][g]
            assert n == len(pick) and _close(rate, _bad(pick)), (y, g)
    ws = load_workbook(calc, data_only=True)["Compare"]
    words = [c.value for row in ws.iter_rows(max_row=40) for c in row if isinstance(c.value, str)]
    assert "Origination year is across the bottom, so the lines are by REGION and one chart is drawn." in words


# ---- the size limit: 150 views, refused past it by the launcher and the Run alike


def test_the_limit_is_the_product_of_the_three_all_loans_counted():
    assert ch.FILTER_MOST_VIEWS3 == 150 and ch.FILTER_MOST_VIEWS == 49
    assert ch.too_many_views("A", 5, "B", 4, "C", 4) is None                      # 6 x 5 x 5 = 150
    for n1 in range(9):
        for n2 in range(9):
            for n3 in range(9):
                said = ch.too_many_views("A", n1, "B", n2, "C", n3)
                assert (said is None) == ((n1 + 1) * (n2 + 1) * (n3 + 1) <= 150), (n1, n2, n3)
    said = ch.too_many_views("A", 5, "B", 4, "C", 5)
    assert said == ("A (5 values), B (4 values) and C (5 values) together make 6 x 5 x 6 = 180 views of every "
                    "grid, counting All loans in each. Three filters can make 150 at most. Drop a filter, or pick a "
                    "column with fewer values.")
    assert ch.views_of(5, 4, 4) == 150 and ch.views_of() == 1


FIVE = ("a1", "a2", "a3", "a4", "a5")
FOUR = ("b1", "b2", "b3", "b4")


@pytest.fixture(scope="module")
def sized(tmp_path_factory):
    """FA with 5 values, FB with 4, FC4 with 4 and FC5 with 5: 6 x 5 x 5 = 150 views, or 6 x 5 x 6 = 180."""
    d = tmp_path_factory.mktemp("sized")
    return _file(d, 1500, {"FA": FIVE, "FB": FOUR, "FC4": FOUR, "FC5": FIVE})


def _flow(x) -> launcher.Flow:
    f = launcher.Flow(gate=launcher.AddOns())
    f.gate.optional = []
    f.pick(str(x))
    f.set_up()
    assert f.screen() == "L2", f.message
    f.pick_outcome("BAD_FLAG")
    f.answer_outcome(True)
    return f


def _pick(f, *names):
    if f.filter:
        f.click(f.filter, "d")
    for name, which in zip(names, "def"):
        f.click(name, which)


def test_the_launcher_refuses_180_views_before_next_and_takes_150(sized, tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    f = _flow(sized)
    _pick(f, "FA", "FB", "FC5")
    ok, said = f.summary()
    assert not ok and said == ch.too_many_views("FA", 5, "FB", 4, "FC5", 5) and "6 x 5 x 6 = 180" in said
    assert f.states()["next"] == "disabled"
    f.click("FC4", "f")                                                       # a column with fewer values
    ok, said = f.summary()
    g = int(re.search(r"= ([\d,]+) grids?, five measures", said).group(1).replace(",", ""))
    assert ok and said.endswith("Grids can show only the loans of one FA, one FB, one FC4, or any of them "
                                "together. Each grid is built for 6 × 5 × 5 = 150 views, All loans counted: "
                                f"{g:,} grid{'s' * (g != 1)} × 150 = {g * 150:,} to build."), said
    assert f.states()["next"] != "disabled"


def test_the_run_refuses_180_views_and_takes_150(sized, tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    d = tmp_path / "over"
    d.mkdir()
    shutil.copy(sized, d / "loans.csv")
    b = _set_up(d, d / "loans.csv", filter="FA", filter2="FB", filter3="FC5")
    ran = book.run(b)
    assert not ran.ok and ran.lines == [f"Couldn't run: {ch.too_many_views('FA', 5, 'FB', 4, 'FC5', 5)}"]
    from pocketbook import perm
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    d = tmp_path / "at"
    d.mkdir()
    shutil.copy(sized, d / "loans.csv")
    b = _set_up(d, d / "loans.csv", filter="FA", filter2="FB", filter3="FC4")
    ran = book.run(b)
    assert ran.ok, ran.lines
    views = load_workbook(b)["_views"]
    keys = {r[0] for r in views.iter_rows(values_only=True) if isinstance(r[0], str) and r[0].startswith("G|")
            and r[0].endswith("|cols")}
    # one grid, and every view of it with a loan in it: from the loan file, each loan is in the views of every value
    # it has, alone, in pairs and all three, and in the whole book
    seen = set()
    for r in _rows(sized):
        for a in (None, r["FA"]):
            for z in (None, r["FB"]):
                for c in (None, r["FC4"]):
                    seen.add((a, z, c))
    assert len(keys) == len(seen) > 140


# ---- the launcher's Filter 3 picker


def test_filter_3_in_the_launcher_needs_filter_2_another_column_and_is_written_to_control(tmp_path, monkeypatch):
    from pocketbook import control
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = _file(tmp_path, 1500)
    f = _flow(x)
    by = {r["name"]: r for r in f.rows()}
    assert by["REGION"]["f"] == {"on": False, "radio": True} and by[ch.ORIG_YEAR]["f"] is not None
    assert by["FICO"]["f"] is None                                   # a number column filters nothing
    assert f.filter == ch.ORIG_YEAR and f.filter3 is None             # Filter 1's default stays; 3 is the analyst's
    f.click("REGION", "f")                                           # Filter 3 without Filter 2 is the second
    got = f.choices()
    assert (got.filter, got.filter2, got.filter3) == (ch.ORIG_YEAR, "REGION", None)
    f.click("SYS_FLAG", "e")
    got = f.choices()
    assert (got.filter, got.filter2, got.filter3) == (ch.ORIG_YEAR, "SYS_FLAG", "REGION")
    f.click("SYS_FLAG", "f")                                         # Filter 3 = Filter 2: refused
    assert f.summary() == (False, ch.same_filter_twice("SYS_FLAG", 2, 3)) and f.states()["next"] == "disabled"
    assert "Filter 2 and Filter 3 are both SYS_FLAG." in f.summary()[1]
    f.click(ch.ORIG_YEAR, "f")                                       # Filter 3 = Filter 1: refused
    assert f.summary() == (False, ch.same_filter_twice(ch.ORIG_YEAR, 1, 3))
    f.click("REGION", "f")
    ok, said = f.summary()
    assert ok and ("Grids can show only the loans of one ORIG_YEAR, one SYS_FLAG, one REGION, or any of them "
                   "together. Each grid is built for 5 × 4 × 3 = 60 views") in said, said
    f.next()
    assert f.page == "answer", f.message
    ws = load_workbook(f.book())[control.SHEET]
    shown = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=control.CHOOSE_COL).value
             for r in range(control.FIRST_ROW, ws.max_row + 1)}
    assert shown["And filter them by"] == "SYS_FLAG" and shown["And then by"] == "REGION"
    again = _flow(x)
    assert (again.filter, again.filter2, again.filter3) == (ch.ORIG_YEAR, "SYS_FLAG", "REGION")


def test_the_run_refuses_one_column_as_filters_1_and_3(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = _file(tmp_path, 1500)
    b = _set_up(tmp_path, x, filter="REGION", filter2="SYS_FLAG", filter3="REGION")
    ran = book.run(b)
    assert not ran.ok and ran.lines == [f"Couldn't run: {ch.same_filter_twice('REGION', 1, 3)}"]
    assert ch.same_filter_twice("REGION", 1, 3) == ("Filter 1 and Filter 3 are both REGION. The third filter narrows "
                                                   "the other two, so it must be another column. Pick a different "
                                                   "one, or none.")


def test_the_window_draws_a_filter_3_column(tmp_path, monkeypatch):
    """The same on screen: a Filter 3 heading and a ring on every category row. Needs a display (xvfb-run on
    Linux)."""
    from test_deps import _window
    from pocketbook import deps
    root = _window()
    try:
        monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
        monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
        monkeypatch.setattr(deps, "missing", lambda: [])
        monkeypatch.setattr(deps, "missing_optional", lambda: [])
        w = launcher.build(root)
        flow = w["flow"]
        x = _file(tmp_path, 800)
        w["extract"].set(str(x))
        flow.pick(str(x))
        flow.set_up()
        w["render"]()
        root.update()
        boxes = w["boxes"]
        on = lambda name, which: len(w["boxes"][(name, which)].find_all()) == 2  # noqa: E731  a dot in the ring
        assert ("REGION", "f") in boxes and ("FICO", "f") in boxes and not on("REGION", "f")
        assert not boxes[("FICO", "f")].winfo_ismapped()                         # no ring on a number column
        flow.click("SYS_FLAG", "e")
        flow.click("REGION", "f")
        w["render"]()
        root.update()
        assert on("REGION", "f") and on("SYS_FLAG", "e") and on(ch.ORIG_YEAR, "d")
    finally:
        root.destroy()
