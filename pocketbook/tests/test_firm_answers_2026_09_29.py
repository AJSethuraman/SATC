"""The firm at the bank machine, 29 Sep 2026 (BACKLOG §6d, "Held for the firm's final decision", items 2 to 6).

- "I have no idea why it automatically decided this random column was an outcome ... it should be a pop-up to say
  explicitly this is going to be what our outcome is." It had picked "% orig commitments" over "EVER GCO".
- "yes no or 01 ... that can be part of the hygiene process, but I need to know how it's accepting stuff"
- "every time I click the button screen kind of blinks scroll all the way up" and a way to "select everything or
  not by the press of a button".
- "there's a lot of white space to use, and it should really use it".
- 12 and 50 "read like one scale".
"""

import csv
import math
import random
import re
import statistics

import pytest
from openpyxl import load_workbook

from pocketbook import book, choices as ch, control, launcher, meanings, synth
from test_book import _answer

import yaml


def _bank_file(tmp_path, n=1500):
    """The synthetic book as the bank's file had it: the outcome named EVER GCO, and a 0/1 percent column whose
    name, run together, holds both `co` and `gco`."""
    src = synth.write_extract(tmp_path / "src", n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    rng = random.Random(29)
    out = tmp_path / "bank.csv"
    names = [("EVER GCO" if c == "BAD_FLAG" else c) for c in rows[0]] + ["% orig commitments"]
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(names)
        for r in rows:
            w.writerow(list(r.values()) + [rng.choice([0, 1])])
    return out


def _flow(x) -> launcher.Flow:
    f = launcher.Flow(gate=launcher.AddOns())
    f.gate.optional = []                             # scouting's add-on, there or not: these tests aren't about it
    f.pick(str(x))
    f.set_up()
    assert f.screen() == "L2", f.message
    return f


# ---- the outcome: names read as words, never a tie broken by the alphabet, never picked for the analyst


def test_a_column_name_is_read_as_words_not_as_letters_run_together():
    cat = meanings.catalog()
    hint = lambda c, m: meanings.hint_in(c, cat[m].hints)  # noqa: E731
    assert meanings.words("% orig commitments") == ["orig", "commitments"]
    assert meanings.words("EVER GCO") == ["ever", "gco"] and meanings.words("LoanNumber") == ["loan", "number"]
    assert hint("% orig commitments", "outcome") is None and hint("% orig commitments", "gco") is None
    assert hint("EVER GCO", "outcome") and hint("EVER_30DPD", "outcome") == "ever"
    # what the old reading found still reads the same
    assert hint("BAD_FLAG", "outcome") == "bad" and hint("GCO_AMT", "gco") == "gco"
    assert hint("ORIG_BAL", "booked") == "bal" and hint("ORIGBAL", "booked") == "bal"
    assert hint("LOAN_NBR", "key") == "loan" and hint("CHARGEOFF_FLAG", "outcome") == "chargeoff"


def test_the_bank_file_reads_ever_gco_as_the_yes_no_column_and_the_percent_column_as_yes_no_too(tmp_path):
    got = {c.name: c for c in book.read_extract(_bank_file(tmp_path)).columns}
    assert got["EVER GCO"].kind == "out" and got["% orig commitments"].kind != "out"
    assert got["% orig commitments"].yes is not None                  # still offered: the analyst decides
    assert got["EVER GCO"].yes + got["EVER GCO"].no + got["EVER GCO"].other == 1500
    assert got["ORIG_BAL"].yes is None                                # a column of amounts is never offered


def test_two_names_that_fit_the_outcome_equally_leave_it_for_the_analyst(tmp_path):
    x = tmp_path / "two.csv"
    rng = random.Random(1)
    with open(x, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["LOAN_NBR", "BAD_A", "BAD_B", "ORIG_BAL"])
        for i in range(400):
            w.writerow([1000 + i, int(rng.random() < 0.1), int(rng.random() < 0.2), rng.randint(1000, 9000)])
    table = book.read_table(x)
    sugg = meanings.suggest(table)
    assert not [c for c, s in sugg.items() if s.means == "outcome"]


def test_the_outcome_is_never_suggested_from_its_values_alone(tmp_path):
    x = tmp_path / "plain.csv"
    rng = random.Random(2)
    with open(x, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["LOAN_NBR", "WENT_SOUTH", "ORIG_BAL"])
        for i in range(400):
            w.writerow([1000 + i, int(rng.random() < 0.1), rng.randint(1000, 9000)])
    sugg = meanings.suggest(book.read_table(x))
    assert sugg["WENT_SOUTH"].means != "outcome"


def test_nothing_is_picked_as_the_outcome_and_a_pick_is_asked_about_in_counts(tmp_path):
    f = _flow(_bank_file(tmp_path))
    for mode in ("bleed", "new"):
        f.set_mode(mode)
        assert f.outcome is None and f.summary() == (False, launcher.NO_OUTCOME)
    assert [c.name for c in f.outcome_choices()] == ["EVER GCO", "% orig commitments"]
    f.pick_outcome("% orig commitments")
    q = f.outcome_question()
    assert q.startswith("Use % orig commitments as the outcome?") and "1 means the loan went bad" in q
    f.answer_outcome(False)                                            # no: nothing changes
    assert f.outcome is None and f.asking is None
    f.click("CHANNEL", "b")
    f.pick_outcome("EVER GCO")
    c = next(c for c in f.outcome_choices() if c.name == "EVER GCO")
    assert f"{c.yes:,} loans" in f.outcome_question() and f"{c.no:,}." in f.outcome_question()
    f.answer_outcome(True)
    assert f.outcome == "EVER GCO"
    rows = {r["name"]: r for r in f.rows()}
    assert rows["EVER GCO"]["what"].startswith("The outcome · 1 on ")
    assert rows["% orig commitments"]["what"].startswith("Yes/no · 1 on ")
    assert rows["EVER GCO"]["b"] is None and rows["EVER GCO"]["c"] is None    # the outcome isn't tested or held
    f.pick_outcome("ORIG_BAL")                                         # not a yes/no column: not offered
    assert f.asking is None


def test_all_and_none_tick_every_input_or_none_and_all_leaves_hold_fixed_alone(tmp_path):
    f = _flow(_bank_file(tmp_path))
    f.set_mode("new")
    f.click("FICO", "c")
    f.test_every(True)
    offered = [r["name"] for r in f.rows() if r["b"] is not None]
    assert "FICO" in offered and "FICO" not in f.test and f.hold == ["FICO"]
    assert f.test == [n for n in offered if n != "FICO"]
    f.test_every(False)
    assert f.test == [] and f.hold == ["FICO"]
    f.set_mode("bleed")
    f.test_every(True)                                                 # only for a new variable
    assert f.test == []


def test_the_picked_outcome_goes_on_columns_and_the_run_reads_it(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    f = _flow(_bank_file(tmp_path))
    f.pick_outcome("EVER GCO")
    f.answer_outcome(True)
    assert "EVER GCO" not in f.choices().segments and "EVER GCO" not in f.choices().bands
    f.next()
    b = launcher.book_for(f.extract)
    ws = load_workbook(b)["Columns"]
    means = {r[book.C_NAME - 1].value: r[book.C_MEANS - 1].value for r in book.table_rows(ws)}
    assert means["EVER GCO"] == meanings.catalog()["outcome"].label
    assert [c for c, m in means.items() if m == meanings.catalog()["outcome"].label] == ["EVER GCO"]
    _answer(b)
    _, problems, _ = book.read_book(b)
    assert not [p for p in problems if "Outcome" in p], problems
    # picked again the other way: it moves, and only one column is ever the outcome
    f.go(2)
    f.pick_outcome("% orig commitments")
    f.answer_outcome(True)
    assert "% orig commitments" not in f.choices().segments          # the outcome never cuts the pockets
    f.next()
    means = {r[book.C_NAME - 1].value: r[book.C_MEANS - 1].value
             for r in book.table_rows(load_workbook(b)["Columns"])}
    assert [c for c, m in means.items() if m == meanings.catalog()["outcome"].label] == ["% orig commitments"]


# ---- the window: its size, and the two settings' words


def test_the_window_opens_at_the_screen_it_is_on():
    assert launcher.window_size(1366, 768) == (1180, 628)
    assert launcher.window_size(1920, 1080) == (1180, 860)
    assert launcher.window_size(800, 600) == (720, 560)


def test_the_two_column_limits_say_which_columns_they_are_for():
    got = {s.key: s.question for s in control.load_settings() if s.key in ("few_values", "many_values")}
    assert got["few_values"].startswith("Number columns:") and "cut into bands" in got["few_values"]
    assert got["many_values"].startswith("Text columns:") and "never cut into bands" in got["many_values"]
    assert yaml.safe_load(open(control.__file__.replace("control.py", "settings.yaml"), encoding="utf-8"))


def test_a_click_keeps_the_table_where_it_was_scrolled_and_the_outcome_is_asked(monkeypatch, tmp_path):
    """On the window (needs a display: xvfb-run on Linux): a tick changes its box in place, so a table scrolled
    half way down stays there; All ticks every input; picking the outcome asks first, in counts."""
    from test_deps import _window
    root = _window()                                 # skips where there is no Tk or no display
    from tkinter import messagebox
    try:
        from pocketbook import deps
        monkeypatch.setattr(launcher, "PREFS", tmp_path / "launcher.json")
        monkeypatch.setattr(deps, "missing", lambda: [])
        monkeypatch.setattr(deps, "missing_optional", lambda: [])     # scouting's add-on, there or not
        asked = []
        monkeypatch.setattr(messagebox, "askyesno", lambda title, text, parent=None: asked.append(text) or True)
        w = launcher.build(root)
        root.geometry("900x520")                 # a small window, so the table scrolls
        root.deiconify()                         # shown: the table only has a height on a window that is
        flow = w["flow"]
        flow.pick(str(_bank_file(tmp_path)))
        flow.set_up()
        flow.set_mode("new")
        w["render"]()
        root.update()
        t = w["table"]
        t.yview_moveto(0.5)
        root.update()
        before = t.yview()
        assert before[0] > 0
        w["box_FICO_b"].event_generate("<Button-1>")
        root.update()
        assert flow.test == ["FICO"] and t.yview() == before
        w["test_all"].event_generate("<Button-1>")
        root.update()
        assert len(flow.test) > 3 and t.yview() == before
        w["box_EVER GCO_a"].event_generate("<Button-1>")
        root.update()
        assert flow.outcome == "EVER GCO" and asked and asked[0].startswith("Use EVER GCO as the outcome?")
        assert w["outcome"].get() == "EVER GCO" and w["next"].cget("state") == "normal"
    finally:
        root.destroy()


# ---- the same day, later: All · None for the bleed, and Columns shows only what this Run uses


def test_all_and_none_for_the_bleeds_bands_and_segments_leave_the_split_alone(tmp_path):
    f = _flow(_bank_file(tmp_path))
    f.click("REV_DEBT", "c")                                            # split by REV_DEBT
    f.pick_every("a", False)
    f.pick_every("b", False)
    assert f.cut == set() and f.seg == set() and f.split == "REV_DEBT"
    f.pick_every("a", True)
    nums = {r["name"] for r in f.rows() if r["a"] is not None}
    assert f.cut == nums - {"REV_DEBT"} and f.split == "REV_DEBT"
    f.pick_every("b", True)
    assert f.seg == {r["name"] for r in f.rows() if r["b"] is not None}
    f.pick_every("c", False)                                            # the split has no All or None
    assert f.split == "REV_DEBT"


def test_columns_hides_what_the_launcher_didnt_pick_and_asks_nothing_about_it(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = synth.write_extract(tmp_path, n=1500)
    picked = ch.Choices(run_kind=ch.BLEED, bands=("ORIG_BAL",), segments=("CHANNEL",), outcome="BAD_FLAG")
    out = book.set_up(x, choices=picked)
    ws = load_workbook(out.book)["Columns"]
    rows = {r[book.C_NAME - 1].value: r for r in book.table_rows(ws) if r[book.C_NAME - 1].value}
    hidden = {n for n, r in rows.items() if ws.row_dimensions[r[0].row].hidden}
    assert "FICO" in hidden and "ASSET_CLASS" in hidden and "REV_DEBT" in hidden
    assert not hidden & {"LOAN_NBR", "ORIG_BAL", "CHANNEL", "BAD_FLAG", "GCO_AMT", "RANR_AMT"}
    keys = [r[book.C_QKEY - 1].value for r in ws.iter_rows(min_row=book.COL_FIRST) if len(r) >= book.C_QKEY]
    assert not [k for k in keys if k and str(k).startswith("FICO|")]  # FICO's -9999 isn't asked about
    assert f"{len(hidden)} not picked in the launcher are hidden" in ws["D3"].value
    _answer(out.book)
    _, problems, _ = book.read_book(out.book)
    assert not problems, problems
    # nothing picked (Set up without the launcher): every column shows, and FICO's code is asked about again
    ws = load_workbook(book.set_up(x).book)["Columns"]
    assert not any(ws.row_dimensions[r[0].row].hidden for r in book.table_rows(ws))
    keys = [r[book.C_QKEY - 1].value for r in ws.iter_rows(min_row=book.COL_FIRST) if len(r) >= book.C_QKEY]
    assert [k for k in keys if k and str(k).startswith("FICO|")]


def test_look_charts_have_no_axis_title_for_excel_to_draw_over_the_numbers(tmp_path, monkeypatch):
    """At the bank, in Excel: "The axis title is out of place" -- "Loans" sat on top of 800 and 1,000."""
    import zipfile
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    out = book.set_up(synth.write_extract(tmp_path, n=1500), choices=ch.Choices(
        run_kind=ch.BLEED, bands=("FICO", "ORIG_BAL"), segments=("CHANNEL",), outcome="BAD_FLAG"))
    with zipfile.ZipFile(out.book) as z:
        charts = [z.read(n).decode("utf-8") for n in z.namelist() if n.startswith("xl/charts/chart")]
    assert charts and not [c for c in charts if "<a:t>Loans</a:t>" in c]


def _as_excel_saves_it(path):
    """Rewrite a workbook the way Excel saves it: every dropdown whose list sits on another sheet moves
    out of the sheet's dropdowns and into Excel's 2010 extension block (the x14 dataValidations)."""
    import re
    import zipfile
    from pocketbook import excel_lists as xl
    src = path.read_bytes()
    moved = 0
    with zipfile.ZipFile(__import__("io").BytesIO(src)) as zin, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.startswith("xl/worksheets/sheet"):
                xml = data.decode("utf-8")
                ext = []
                def pull(m):
                    body = m.group(0)
                    f = re.search(r"<formula1>(.*?)</formula1>", body).group(1)
                    if "!" not in f:
                        return body
                    attrs = re.search(r"<dataValidation ([^>]*)>", body).group(1)
                    sq = re.search(r'sqref="([^"]*)"', attrs).group(1)
                    attrs = re.sub(r'\s*sqref="[^"]*"', "", attrs)
                    ext.append(f'<x14:dataValidation {attrs}><x14:formula1><xm:f>{f.lstrip("=")}</xm:f>'
                               f'</x14:formula1><xm:sqref>{sq}</xm:sqref></x14:dataValidation>')
                    return ""
                xml = re.sub(r"<dataValidation [^>]*>.*?</dataValidation>", pull, xml, flags=re.S)
                if ext:
                    moved += len(ext)
                    xml = re.sub(r"<dataValidations[^>]*>\s*</dataValidations>", "", xml)
                    xml = re.sub(r'(<dataValidations count=")\d+', lambda m: m.group(1) + str(
                        len(re.findall(r"<dataValidation ", xml))), xml)
                    block = (f'<extLst><ext uri="{xl.DV_EXT}" xmlns:x14="{xl.X14}"><x14:dataValidations '
                             f'count="{len(ext)}" xmlns:xm="{xl.XM}">{"".join(ext)}</x14:dataValidations></ext></extLst>')
                    xml = xml.replace("</worksheet>", block + "</worksheet>")
                data = xml.encode("utf-8")
            z.writestr(item, data)
    return moved


def _lists(path):
    wb = load_workbook(path)
    return sorted((ws.title, str(d.sqref), (d.formula1 or "").lstrip("="))
                  for ws in wb for d in ws.data_validations.dataValidation)


def test_dropdowns_excel_saved_on_another_sheet_survive_the_next_run(tmp_path, monkeypatch):
    """At the bank: "I have no drop downs in most places now", after changing an answer on Control. Excel keeps
    a dropdown whose list is on another sheet in its extension block; openpyxl dropped that block, so the next
    Run saved the workbook without them. Only Yes/No and the other typed-in lists were left."""
    import warnings
    from pocketbook import excel_lists, perm
    from test_book import _answer
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 300)
    b = book.set_up(synth.write_extract(tmp_path, n=1500)).book
    _answer(b)
    assert book.run(b).ok
    before = _lists(b)
    assert _as_excel_saves_it(b) > 10
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        assert len(_lists(b)) < len(before)                    # openpyxl alone loses them: the bug
    assert sum(len(v) for v in excel_lists.extended(b).values()) > 10
    assert book.run(b).ok
    after = _lists(b)
    for sheet in ("Control", "Columns"):                       # kept in place by a Run, never rebuilt
        assert [x for x in after if x[0] == sheet] == [x for x in before if x[0] == sheet]
    assert {x[0] for x in after} == {x[0] for x in before}


def test_opening_the_workbook_at_a_cell_keeps_the_dropdowns_excel_saved(tmp_path, monkeypatch):
    """The launcher's Open the workbook button saves it once to pick the cell: that save lost them too."""
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    b = book.set_up(synth.write_extract(tmp_path, n=800)).book
    before = _lists(b)
    _as_excel_saves_it(b)
    assert book.open_at(b, "Control", "C15")
    assert _lists(b) == before


# ---- later the same day, at the bank: a category splits the pockets too
# The firm: "I kind of figured I'd be able to see a view with system flag and origination FICO and asset segment
# somehow" ... "I'd rather make it able to happen within the config or launcher. Like I know it can't break down too
# far but can we not make something work?"


def _flag_file(tmp_path, n=1500, seed=29):
    """The synthetic book with two categories added after every other value is drawn, so no planted number
    moves: SYS_FLAG (Y or N, a blank on every 97th loan) and REGION (seven values)."""
    src = synth.write_extract(tmp_path / "src", n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    rng = random.Random(seed)
    out = tmp_path / "flagged.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(list(rows[0]) + ["SYS_FLAG", "REGION"])
        for i, r in enumerate(rows):
            flag = "" if i % 97 == 0 else rng.choice("YN")
            w.writerow(list(r.values()) + [flag, rng.choice(["R1", "R2", "R3", "R4", "R5", "R6", "R7"])])
    return out


def test_category_split_launcher_offers_split_by_on_a_category_and_never_segments_by_it_too(tmp_path):
    f = _flow(_flag_file(tmp_path))
    rows = {r["name"]: r for r in f.rows()}
    assert rows["SYS_FLAG"]["c"] == {"on": False, "radio": True}          # a category can split
    assert rows["CHANNEL"]["c"] is not None and rows["FICO"]["c"] is not None
    f.click("SYS_FLAG", "c")
    rows = {r["name"]: r for r in f.rows()}
    assert f.split == "SYS_FLAG" and "SYS_FLAG" not in f.seg               # unticked from Segment by
    assert rows["SYS_FLAG"]["c"]["on"] and not rows["SYS_FLAG"]["b"]["on"]
    f.pick_outcome("BAD_FLAG")
    f.answer_outcome(True)
    got = f.choices()
    assert got.split == "SYS_FLAG" and "SYS_FLAG" not in got.segments and "CHANNEL" in got.segments
    ok, said = f.summary()
    assert ok and "split by SYS_FLAG" in said
    f.pick_every("b", True)                                                 # All leaves the split alone
    assert "SYS_FLAG" not in f.seg and f.split == "SYS_FLAG"
    f.click("SYS_FLAG", "b")                                                # ticked as a segment: no longer splits
    assert f.split is None and "SYS_FLAG" in f.seg
    f.click("SYS_FLAG", "c")
    f.click("REV_DEBT", "c")                                                # still one split column, or none
    assert f.split == "REV_DEBT" and "SYS_FLAG" not in f.seg


def test_category_split_too_many_values_is_refused_in_words_in_the_launcher_and_at_the_run(tmp_path, monkeypatch):
    assert ch.too_many_values("X", ch.SPLIT_MOST_VALUES) is None           # six values split
    said = ch.too_many_values("REGION", 7)
    assert said == ("REGION has 7 values. A category can split the pockets by 6 values at most: with more, each "
                    "pocket's parts are too thin to read. Split by a column with fewer values, or by none.")
    f = _flow(_flag_file(tmp_path))
    f.pick_outcome("BAD_FLAG")
    f.answer_outcome(True)
    f.click("REGION", "c")
    assert f.summary() == (False, said) and f.states()["next"] == "disabled"
    # the Run refuses it too, in the same words, whatever wrote the workbook
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    out = book.set_up(f.extract, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                                    split="REGION", outcome="BAD_FLAG"))
    _answer(out.book)
    ran = book.run(out.book)
    assert not ran.ok and ran.lines == [f"Couldn't run: {said}"]


def _category_run(tmp_path, field, many_tests="bh", n=12000):
    from pocketbook import config as cfgmod, engine
    from pocketbook.ingest import read_table
    cfg, _ = synth.write(tmp_path / "cube", n=n)
    raw = cfgmod.load(cfg).raw
    raw["split"] = {"field": field, "how": "each_value"}
    raw["benchmark"]["many_tests"] = many_tests
    return engine.run(cfgmod.parse(raw), read_table(_flag_file(tmp_path, n=n)))


def test_category_split_engine_sets_each_value_against_the_rest_of_its_pocket(tmp_path):
    """Asset class 4 goes bad 1.4 times as often as the others (synth.py): each value against the rest of its
    pocket finds it, pocket by pocket and pooled, and every figure is worked out again here from the parts."""
    from pocketbook import engine, stats
    res = _category_run(tmp_path, "ASSET_CLASS")
    g = res.grids[0]                                                         # FICO x CHANNEL
    assert g.split_parts == ["1", "2", "3", "4"] and not g.split_compare     # no halves
    four = g.part_pooled["4"]["outcome_loans"]
    assert 1.2 < four["ratio"] < 1.8 and four["odds_p"] < 0.01 and four["pockets"] >= 8
    assert g.part_pooled["1"]["outcome_loans"]["ratio"] < 1
    # every value is one comparison: the value's rate over the rest of its pocket's, and the p-value of the z test,
    # after one allowance for many tests across every value and pocket of the grid
    raw, want = {}, {}
    for v in g.split_parts:
        for (b, d) in g.part_tested[v]["outcome_loans"]:
            me = g.split_cells[(b, d, v)].rates["outcome_loans"]
            rest = [g.split_cells[(b, d, x)].rates["outcome_loans"] for x in g.split_parts
                    if x != v and (b, d, x) in g.split_cells]
            num, den, units = sum(r.num for r in rest), sum(r.den for r in rest), sum(r.units for r in rest)
            assert g.part_compare[v][(b, d)]["outcome_loans"][0] == pytest.approx(stats.multiple(me.rate, num / den))
            raw[(v, b, d)] = stats.two_prop_z(me.num, me.units, num, units)[1]
    for k, p in zip(raw, engine.adjust(list(raw.values()), "bh")):
        want[k] = p
    got = {(v, b, d): x["outcome_loans"][1] for v in g.split_parts for (b, d), x in g.part_compare[v].items()
           if x.get("outcome_loans", (None, None))[1] is not None}
    assert got.keys() == want.keys() and all(got[k] == pytest.approx(want[k]) for k in want)
    # a dollar rate's parts are shuffled within their pocket, like the halves
    gco = [x["gco_rate"][1] for x in g.part_compare["4"].values() if x.get("gco_rate", (None,))[0] is not None]
    assert gco and all(p is not None for p in gco)
    # whether the values differ at all: B3 on three degrees of freedom, and the planted class shows
    b3 = g.split_general["outcome_loans"]
    assert b3["df"] == 3 and b3["p"] < 0.001
    assert "gco_rate" not in g.split_general                                 # no k-group test for a dollar rate
    # the families Record counts: every value's pockets in one family per grid and rate
    from pocketbook import checks
    fam = [f for f in checks.families(res) if f[0] == "split values" and f[2] == "outcome_loans"
           and f[1] == f"{g.band} x {g.dimension}"]
    assert fam and fam[0][4] == len(want)


def test_category_split_engine_with_two_values_is_the_cmh_test_for_the_differ_at_all(tmp_path):
    """With two values, B3 is Cochran-Mantel-Haenszel's chi-square (docs/statistics.md B3: "With two groups it
    collapses to A6"), worked out here from every pocket's four counts; and Y against the rest is Y against N."""
    from pocketbook import stats
    g = _category_run(tmp_path, "SYS_FLAG", many_tests="none", n=8000).grids[0]
    assert g.split_parts == ["N", "Y", "(blank)"]                          # a blank is a value of its own
    assert g.split_general["outcome_loans"]["df"] == 2
    t = _two_valued(tmp_path).grids[0]
    assert t.split_parts == ["N", "Y"]
    strata = []
    for (b, d), _ in t.inner():
        cells = [t.split_cells.get((b, d, v)) for v in ("Y", "N")]
        s = [c.rates["outcome_loans"] if c else None for c in cells]
        strata.append(tuple(x for r in s for x in ((r.events, r.units - r.events) if r else (0, 0))))
    chi, p = stats.cmh(strata)
    assert t.split_general["outcome_loans"]["df"] == 1
    assert t.split_general["outcome_loans"]["q"] == pytest.approx(chi, rel=1e-9)
    assert t.split_general["outcome_loans"]["p"] == pytest.approx(p, rel=1e-9)
    for (b, d), x in t.part_compare["Y"].items():                           # Y against the rest is Y against N
        if x["outcome_loans"][0] is not None:
            y, n_ = (t.split_cells[(b, d, v)].rates["outcome_loans"] for v in ("Y", "N"))
            assert x["outcome_loans"][0] == pytest.approx(stats.multiple(y.rate, n_.rate))
            assert x["outcome_loans"][1] == pytest.approx(stats.two_prop_z(y.num, y.units, n_.num, n_.units)[1])


def _two_valued(tmp_path, n=8000):
    """SYS_FLAG with its blanks read as N: two values."""
    from pocketbook import config as cfgmod, engine
    from pocketbook.ingest import read_table
    src = _flag_file(tmp_path / "two", n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    out = tmp_path / "two.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        for r in rows:
            w.writerow({**r, "SYS_FLAG": r["SYS_FLAG"] or "N"})
    cfg, _ = synth.write(tmp_path / "cube2", n=n)
    raw = cfgmod.load(cfg).raw
    raw["split"] = {"field": "SYS_FLAG", "how": "each_value"}
    raw["benchmark"]["many_tests"] = "none"
    return engine.run(cfgmod.parse(raw), read_table(out))


def test_category_split_workbook_carries_the_flag_through_every_tab(tmp_path, monkeypatch):
    """FICO x asset segment by the system flag, as the firm asked: Grids, Pockets, Split, Record and Start here."""
    from pocketbook import results
    import tabs
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = _flag_file(tmp_path, n=6000)
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL", "ASSET_CLASS"),
                                            split="SYS_FLAG", outcome="BAD_FLAG"))
    _answer(out.book)
    assert book.read_book(out.book)[0]["split"] == {"field": "SYS_FLAG", "how": "each_value"}
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    b = out.book
    wb = load_workbook(b)
    assert "FICO x ASSET_CLASS / SYS_FLAG" in tabs.options(wb, results.GRIDS, "Grid")
    assert tabs.options(wb, results.POCKETS, "Pockets") == ["Two-way", "Split by SYS_FLAG"]
    split = tabs.options(wb, results.SPLIT, "Grid")
    assert "FICO x ASSET_CLASS · SYS_FLAG Y vs rest" in split and "FICO x CHANNEL · SYS_FLAG N vs rest" in split
    ws = tabs.calculated(tabs.choose(b, tmp_path / "g.xlsx", results.GRIDS, grid="FICO x ASSET_CLASS / SYS_FLAG"),
                         results.GRIDS)
    heads = set(c for _, c in tabs.block(ws, "vs the book"))
    assert {"ASSET_CLASS 4 · Y", "ASSET_CLASS 4 · N"} <= heads or {"4 · Y", "4 · N"} <= heads
    ws = tabs.calculated(tabs.choose(b, tmp_path / "s.xlsx", results.SPLIT, grid="FICO x ASSET_CLASS · SYS_FLAG Y vs "
                                                                                 "rest"), results.SPLIT)
    text = [str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    assert "Worse than the rest in" in text and "Value vs rest, all" in text
    assert "Bad loans, value vs rest" in text
    differ = next(t for t in text if t.startswith("Do the values of SYS_FLAG differ at all? Bad loans: "))
    assert "degrees of freedom" in differ and "GCOs ($)" in differ and "not tested: dollar rate" in differ
    assert not [t for t in text if "high half" in t.lower() or "High vs low" in t]
    rec = {k: v for k, v, *_ in tabs.record_rows(b)}
    assert "each value set against the rest of its pocket" in rec["Split"]
    assert "for the split's values" in rec["Families of tests"]
    assert "K-group Mantel-Haenszel" in rec["Tests"]
    start = [str(v) for row in load_workbook(b)["Start here"].iter_rows(values_only=True) for v in row if v]
    assert "Split: each pocket split" in start


# ---- Look reads the Treat as answers (at the bank, 29 Sep 2026: "This median call seems to ignore that I said the
# -99... values that represent things missing from the bureau are treating as missing. Look into this and see if
# this leaks elsewhere"). Look was drawn at Set up, before any answer, and a Run never drew its blocks again.

def _bureau_file(tmp_path, n=1500):
    """The synthetic book with SHORT_HIST: months of credit history, and where the bureau has none one of four
    codes from -99,000,901, none on 1% of loans, so no one of them is "likely a code" and Columns asks about the
    negatives, as it did at the bank."""
    src = synth.write_extract(tmp_path, n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    rng = random.Random(3)
    for i, r in enumerate(rows):
        r["SHORT_HIST"] = -99000901 - (i // 40) % 4 if i % 40 == 0 else rng.randint(0, 435)
    out = tmp_path / "bureau.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return out


def _look_block(path, column):
    from pocketbook import look
    wb = load_workbook(path)
    i = look.drawn_columns(wb).index(column)
    ws, r = wb[look.LOOK], look.FIRST + look.BLOCK * i
    return {str(ws.cell(row=r + k, column=look.STATS_COL).value): ws.cell(row=r + k, column=look.VALUE_COL).value
            for k in range(1, look.R_TREAT)}


def test_look_leaves_out_what_columns_answered_missing_once_run(tmp_path, monkeypatch):
    from pocketbook import perm
    from test_book import _answer, treat_odd
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    b = book.set_up(_bureau_file(tmp_path), choices=ch.Choices(
        run_kind=ch.BLEED, bands=("FICO", "SHORT_HIST"), segments=("CHANNEL",), outcome="BAD_FLAG")).book
    before = _look_block(b, "SHORT_HIST")
    assert before["Smallest"] == -99000904 and before["Mean"] < 0            # unanswered: used as recorded
    wb = load_workbook(b)
    assert treat_odd(wb, "SHORT_HIST", "Missing")
    wb.save(b)
    _answer(b)                                                               # FICO's -9999 answered Missing too
    assert book.run(b).ok
    after = _look_block(b, "SHORT_HIST")
    assert after["Smallest"] >= 0 and after["Largest"] <= 435 and after["Mean"] > 0
    assert after["Answered missing, left out"] == sum(1 for i in range(1500) if i % 40 == 0)
    fico = _look_block(b, "FICO")
    assert fico["Smallest"] > 0 and fico["Answered missing, left out"] > 0
    wb = load_workbook(b)
    note = [c.value for row in wb["Look"].iter_rows() for c in row if isinstance(c.value, str)
            and "negative values" in c.value]
    assert note                                                              # it says what was answered


def test_look_keeps_a_bars_from_or_to_the_analyst_typed_when_it_is_drawn_again(tmp_path, monkeypatch):
    from pocketbook import look, perm
    from test_book import _answer, treat_odd
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 200)
    b = book.set_up(_bureau_file(tmp_path), choices=ch.Choices(
        run_kind=ch.BLEED, bands=("FICO", "SHORT_HIST"), segments=("CHANNEL",), outcome="BAD_FLAG")).book
    wb = load_workbook(b)
    i = look.drawn_columns(wb).index("SHORT_HIST")
    r = look.FIRST + look.BLOCK * i
    wb[look.LOOK].cell(row=r + look.R_BARS, column=look.VALUE_COL).value = 50
    wb[look.LOOK].cell(row=r + look.R_TO, column=look.VALUE_COL).value = 250
    treat_odd(wb, "SHORT_HIST", "Missing")
    wb.save(b)
    _answer(b)
    assert book.run(b).ok
    got = _look_block(b, "SHORT_HIST")
    assert got["Bars"] == 50 and got["To"] == 250


def test_looks_dots_leave_out_what_columns_answered_missing(tmp_path):
    """The scatters read each loan's value the same way: an answered missing is no dot, as it is no loan in a band."""
    from pocketbook import config as cfgmod, look
    from pocketbook.ingest import read_table
    t = read_table(_bureau_file(tmp_path, n=400))
    got = look._readable(t, "SHORT_HIST", None, cfgmod.MissingRule(below=0.0))
    assert [v for v in got if v is not None] and min(v for v in got if v is not None) >= 0
    assert sum(v is None for v in got) == sum(1 for i in range(400) if i % 40 == 0)


# ---- later still, at the bank: one cell of a grid read out in words
# The firm: "It would be useful to be able to maybe select a particular line and say I want this as an example and
# it fills in the band saying what versus book means what versus band means and what loans means and it changes
# obviously depending on the grid and measure you're looking at, but the measures should be constant from run to
# run the grid may change, but like that's really just the band."


@pytest.fixture(scope="module")
def one_cell_book(tmp_path_factory):
    """A Run on the synthetic book, shuffled only a little: these tests read words, not p-values."""
    from pocketbook import perm
    d = tmp_path_factory.mktemp("one_cell")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 200)
        b = book.set_up(synth.write_extract(d, n=4000)).book
        _answer(b)
        ran = book.run(b)
        assert ran.ok, ran.lines
    return b


def _said(ws) -> dict:
    """What one cell says, as calculated: the cell's name under "name", then each line by its label."""
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=2).value == "What one cell says":
            out = {"name": ws.cell(row=r + 1, column=2).value}
            for k in range(2, 7):
                out[ws.cell(row=r + k, column=2).value] = ws.cell(row=r + k, column=3).value
            return out
    raise KeyError("What one cell says")


def _grids(b, out, **picks):
    """Grids with its dropdowns set, calculated: the sheet, its four blocks by title, and what one cell says."""
    from pocketbook import results
    import tabs
    ws = tabs.calculated(tabs.choose(b, out, results.GRIDS, **picks), results.GRIDS)
    blocks = {t: tabs.block(ws, t) for t in ("Rate", "vs the book", "vs rest of band", "Loans")}
    return ws, blocks, _said(ws)


def _pockets(loans: dict) -> dict:
    """The inner pockets with loans in them, {(row, column): loans}."""
    return {k: v for k, v in loans.items() if "All" not in k and isinstance(v, (int, float)) and v > 0}


def _shade(v, kind: str, bound: float) -> str:
    """The heat scale's colour for a gap, worked out here from results.HEAT_STEPS' limits and not its formula."""
    import math
    if not isinstance(v, (int, float)):
        return "blank"
    t = -v / bound if kind == "pts" else (math.log2(v) if v > 0 else 0)
    for lim, word in ((1.0, "deep red"), (0.585, "red"), (0.263, "light red")):
        if t >= lim:
            return word
    for lim, word in ((-0.585, "green"), (-0.263, "light green")):
        if t <= lim:
            return word
    return "pale"


def _fewest(b) -> int:
    """The fewest loans in a pocket the last Run used (Control's answer, as Record keeps it)."""
    from pocketbook import config as cfgmod
    return cfgmod.parse(book.read_book(b)[0]).benchmark.min_units


def _min_losses(b) -> int:
    from pocketbook import config as cfgmod
    return cfgmod.parse(book.read_book(b)[0]).benchmark.min_events


def test_one_cell_reads_the_blocks_own_numbers_in_words_for_a_multiple_and_a_gap_in_points(one_cell_book, tmp_path):
    b = one_cell_book
    for measure, kind in (("GCOs ($)", "x"), ("RANR", "pts")):
        _, blocks, _ = _grids(b, tmp_path / f"{kind}-0.xlsx", measure=measure)
        rate, bk, bd, loans = (blocks[t] for t in ("Rate", "vs the book", "vs rest of band", "Loans"))
        both = [k for k in _pockets(loans) if isinstance(bk[k], (int, float)) and isinstance(bd[k], (int, float))]
        bl, d = both[-1]                                          # not the pocket the tab opens on
        assert (bl, d) != both[0]
        ws, blocks, said = _grids(b, tmp_path / f"{kind}.xlsx", measure=measure, row=bl, column=d)
        assert blocks["Rate"] == rate                             # picking a cell changes nothing above it
        n, v, w = loans[(bl, d)], bk[(bl, d)], bd[(bl, d)]
        r = f"{rate[(bl, d)] * 100:.2f}%"
        others = [c for (x, c) in _pockets(loans) if x == bl and c != d]
        named = f" (the {', '.join(others)} loans)" if 1 <= len(others) <= 3 else ""
        assert said["name"] == f"{bl} · {d}, {measure}"
        if kind == "x":
            assert said["Rate"] == f"These {n:,} loans: GCOs were {r} of their booked dollars."
            assert said["vs the book"] == f"{v:.2f}× the GCO rate of the whole book."
            assert said["vs rest of band"] == f"{w:.2f}× the GCO rate of the other loans in {bl}{named}."
        else:
            side = lambda g: "less" if g < 0 else "more"                                 # noqa: E731
            assert said["Rate"] == f"These {n:,} loans: RANR was {r} of their booked dollars."
            assert said["vs the book"] == (f"RANR was {abs(v):.2f} points {side(v)} of booked dollars than for the "
                                           "whole book.")
            assert said["vs rest of band"] == (f"RANR was {abs(w):.2f} points {side(w)} of booked dollars than for "
                                               f"the other loans in {bl}{named}.")
        assert said["Loans"] == f"{n:,} loans; {bl} has {loans[(bl, 'All')]:,} in all."
        # the scale's bound leaves out the cells under the fewest loans, which are grey (the firm, 29 Sep 2026)
        few = _fewest(b)
        bound = max([abs(x) for blk in (bk, bd) for k, x in blk.items() if isinstance(x, (int, float))
                     and isinstance(loans.get(k), (int, float)) and loans[k] >= few] + [0.01])
        if n < few:
            assert said["The colour"] == f"Grey: only {n:,} loans, fewer than the {few:,} set on Control, so not coloured."
        else:
            assert said["The colour"].startswith(f"vs the book is {_shade(v, kind, bound)}, vs rest of band "
                                                 f"{_shade(w, kind, bound)}.")


def test_one_cell_says_why_a_blank_is_blank_alone_in_its_band_or_too_few_losses(one_cell_book, tmp_path):
    from pocketbook import results
    b = one_cell_book
    few = f"Blank: fewer losses than the minimum ({_min_losses(b)} losses), so not compared."
    _, blocks, _ = _grids(b, tmp_path / "blank-0.xlsx", measure="GCOs ($)")
    bk, bd, loans = blocks["vs the book"], blocks["vs rest of band"], blocks["Loans"]
    pockets = _pockets(loans)
    mates = lambda k: sum(1 for j in pockets if j[0] == k[0])                                  # noqa: E731
    alone = next(k for k in pockets if mates(k) == 1)
    thin = next(k for k in pockets if mates(k) > 1 and bk[k] is None and bd[k] is None)
    assert bd[alone] is None
    for measure in ("GCOs ($)", "RANR"):
        _, _, said = _grids(b, tmp_path / f"alone-{measure[0]}.xlsx", measure=measure, row=alone[0], column=alone[1])
        assert said["vs rest of band"] == f"Blank: alone in its band. Nothing else in {alone[0]} to compare with."
    _, _, said = _grids(b, tmp_path / "alone.xlsx", measure="GCOs ($)", row=alone[0], column=alone[1])
    if loans[alone] == 1:                                         # one loan reads as one, not "1 loans"
        assert said["Rate"].startswith("This one loan: GCOs were ") and said["Rate"].endswith(" of its booked "
                                                                                            "dollars.")
        assert said["Loans"].startswith("1 loan; ")
    if bk[alone] is None:
        assert said["vs the book"] == few                        # only one reason a comparison with the book is blank
    _, _, said = _grids(b, tmp_path / "few.xlsx", measure="GCOs ($)", row=thin[0], column=thin[1])
    assert said["vs the book"] == few and said["vs rest of band"] == few
    assert "alone" not in said["The colour"] and said["The colour"].startswith("vs the book is blank, vs rest of band "
                                                                               "blank.")
    ws = load_workbook(b)[results.GRIDS]
    notes = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("A blank")]
    assert notes == ["A blank: alone in its band, or fewer losses than the minimum, so not compared."]


def _listed(ws, label: str) -> list:
    """A Row or Column dropdown's list as it stands (calculated): its OFFSET over the hidden labels, worked out."""
    import re
    import tabs
    cell = tabs.dropdown(ws.formulas, label)
    dv = next(v for v in ws.formulas.data_validations.dataValidation if cell.coordinate in str(v.sqref))
    m = re.fullmatch(r"=?OFFSET\(\$([A-Z]+)\$(\d+),0,0,MAX\(1,\$([A-Z]+)\$(\d+)\),1\)", dv.formula1)
    assert m, dv.formula1
    c, r, nc_, nr_ = m.groups()
    n = max(1, int(ws[f"{nc_}{nr_}"].value or 0))
    return [ws[f"{c}{int(r) + k}"].value for k in range(n)]


def test_one_cell_row_and_column_lists_follow_the_grid_picked(one_cell_book, tmp_path):
    from pocketbook import results
    import tabs
    b = one_cell_book
    grids = tabs.options(load_workbook(b), results.GRIDS, "Grid")
    ws, blocks, said = _grids(b, tmp_path / "g0.xlsx")
    rows0 = [bl for bl, d in blocks["Loans"] if d == "All" and bl != "All"]
    cols0 = [d for bl, d in blocks["Loans"] if bl == "All" and d != "All"]
    assert _listed(ws, "Row") == rows0 and _listed(ws, "Column") == cols0
    assert said["name"] == f"{tabs.dropdown(ws, 'Row').value} · {tabs.dropdown(ws, 'Column').value}, Bad loans"
    other = next(g for g in grids if " / " not in g and g.split(" x ")[0] != grids[0].split(" x ")[0]
                 and g.split(" x ")[1] != grids[0].split(" x ")[1])
    ws, blocks, said = _grids(b, tmp_path / "g1.xlsx", grid=other)
    rows1 = [bl for bl, d in blocks["Loans"] if d == "All" and bl != "All"]
    cols1 = [d for bl, d in blocks["Loans"] if bl == "All" and d != "All"]
    assert rows1 != rows0 and cols1 != cols0
    assert _listed(ws, "Row") == rows1 and _listed(ws, "Column") == cols1
    # the Row and Column left from the other grid aren't in this one: it asks, and reads out nothing
    assert said["name"] == results.SAY_PICK and not any(said[k] for k in said if k != "name")
    k = next(iter(_pockets(blocks["Loans"])))
    _, blocks, said = _grids(b, tmp_path / "g2.xlsx", grid=other, row=k[0], column=k[1])
    assert said["name"] == f"{k[0]} · {k[1]}, Bad loans"
    assert said["Loans"].startswith(f"{blocks['Loans'][k]:,} loans; {k[0]} has ")


def test_segments_read_their_numbers_as_numbers():
    """At the bank, 29 Sep 2026, Grids by Loan Amount Bucket ($5k): "$5k-<$10k" sat after "$40k+", as text sorts."""
    from pocketbook import engine
    got = engine._order(["$40k+", "$5k–<$10k", "$0k–<$5k", "$10k–<$15k", engine.BLANK_LABEL, "Tier 10", "Tier 2"])
    assert got == ["$0k–<$5k", "$5k–<$10k", "$10k–<$15k", "$40k+", "Tier 2", "Tier 10", engine.BLANK_LABEL]


# ---- that evening, by pop-up: four changes to Grids, each answered yes
# Raised at the bank from a photo of Grids: 620-659 · $20k-<$25k, 3 loans, Kept after losses -50.38 pts against the
# book, was the deepest red on the grid and paled every real gap. The firm on filtering: "it would be nice to be able
# to filter by that category which would probably solve a lot of ... having multiway views", and, decided the same
# day: "we keep things compared to the whole book that's just kind of the point". On loan size: "we tend to give
# these loan amounts to these FICO scores within this category". Yes to all four; a number column as a segment was
# answered "Later".

KIOSK = "Kiosk"
ONLY = "Only loans where SYS_FLAG is"


def _grids_file(tmp_path, n=6000):
    """The flagged book, and three loans through a channel of their own, in one score band, each keeping -50% of
    what it booked: a 3-loan pocket far deeper than any real one, as on the bank's photo."""
    src = _flag_file(tmp_path, n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    for k, fico in enumerate((641, 645, 649)):
        bal = 20000.0 + 1000 * k
        rows.append({**rows[10], "LOAN_NBR": f"K{k}", "FICO": fico, "CHANNEL": KIOSK, "ORIG_BAL": bal,
                     "BAD_FLAG": 1, "GCO_AMT": bal * 0.5, "RANR_AMT": -bal * 0.5, "SYS_FLAG": "Y"})
    out = tmp_path / "grids.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return out


@pytest.fixture(scope="module")
def grids_book(tmp_path_factory):
    """FICO x CHANNEL and FICO x ASSET_CLASS, split by SYS_FLAG, fewest loans 30: the book and its loan file."""
    from pocketbook import perm
    d = tmp_path_factory.mktemp("grids")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 200)
        x = _grids_file(d)
        # since 30 Sep 2026 the filter is its own pick (Filter by), not the split: this book picks both
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL", "ASSET_CLASS"),
                                                split="SYS_FLAG", filter="SYS_FLAG", outcome="BAD_FLAG"))
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return out.book, x


def _loans_file(x) -> list[dict]:
    return list(csv.DictReader(open(x, encoding="utf-8")))


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _band(fico, labels) -> str:
    """A loan's FICO band, read from the grid's own labels ("496 - 653", both ends in the band), worked out here and
    not by the engine: the -9999 code was answered missing, and a blank is its own row."""
    if fico in ("", None):
        return "(blank)"
    v = float(fico)
    if v < -1000:
        return "(marked missing)"
    for lab in labels:
        lo, _, hi = lab.partition(" - ")
        if hi and float(lo) <= v <= float(hi):
            return lab
    raise KeyError(fico)


def _cells(rows, loans_block) -> dict:
    """{(band, channel): [loans]}, margins included, over `rows`, for the cells the Loans block has."""
    labels = [bl for bl, d in loans_block if d == "All" and bl not in ("All", "(blank)", "(marked missing)")]
    out: dict = {}
    for r in rows:
        bl, d = _band(r["FICO"], labels), r["CHANNEL"]
        for k in {(bl, d), (bl, "All"), ("All", d), ("All", "All")}:
            out.setdefault(k, []).append(r)
    return out


def _bad_rate(rows):
    read = [int(r["BAD_FLAG"]) for r in rows if r["BAD_FLAG"] in ("0", "1")]
    return sum(read) / len(read) if read else None


def _size(rows):
    bals = [_num(r["ORIG_BAL"]) for r in rows if _num(r["ORIG_BAL"]) is not None]
    return (sum(bals) / len(bals), statistics.median(bals)) if bals else (None, None)


def _views_rows(b) -> dict:
    from pocketbook import results
    return {r[0]: list(r[1:]) for r in load_workbook(b)[results.VIEWS].iter_rows(values_only=True) if r[0]}


def _below(ws) -> int:
    """The row the blocks start under: the dropdowns', clear of the method note's labels."""
    import tabs
    return tabs.dropdown(ws, "Grid").row


def _head(ws, title: str) -> str:
    return next(c.value for row in ws.iter_rows(min_row=_below(ws)) for c in row if isinstance(c.value, str)
                and c.value.startswith(title))


def _inner(formula: str) -> str:
    """A rule's own formula, without the row's divider every rule on a block carries (results.cf)."""
    m = re.fullmatch(r'AND\(\$[A-Z]+\d+<>"",(.*)\)', formula)
    return m.group(1) if m else formula


GREY = re.compile(r"^AND\(ISNUMBER\(([A-Z]+)(\d+)\),ISNUMBER\(([A-Z]+)(\d+)\),ISNUMBER\((\$[A-Z]+\$\d+)\),"
                  r"([A-Z]+)(\d+)<(\$[A-Z]+\$\d+)\)$")


def _rules_at(ws, coord: str) -> list:
    """The conditional formats over one cell of the sheet as written, in the order they are tried."""
    from openpyxl.utils import range_boundaries
    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
    c, r = coordinate_from_string(coord)
    c = column_index_from_string(c)
    out = []
    for rng in ws.formulas.conditional_formatting:
        for bounds in str(rng.sqref).split():
            c0, r0, c1, r1 = range_boundaries(bounds)
            if c0 <= c <= c1 and r0 <= r <= r1:
                out += [(x, c - c0, r - r0) for x in sorted(rng.rules, key=lambda x: x.priority)]
    return out


def _is_grey(ws, coord: str) -> bool:
    """Whether a cell's grey rule holds (its second: the first is the same with a gap in points' number format),
    worked out from the calculated sheet by moving its relative references as a spreadsheet does."""
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.cell import column_index_from_string
    from pocketbook import house
    rule, dc, dr = _rules_at(ws, coord)[1]
    m = GREY.match(_inner(rule.formula[0]))
    assert m, rule.formula[0]
    assert rule.dxf.fill is None and rule.dxf.font.color.rgb[-6:] == house.DISABLED_TEXT
    at = lambda col, row: ws[f"{get_column_letter(column_index_from_string(col) + dc)}{int(row) + dr}"].value  # noqa
    loans, few = at(m.group(3), m.group(4)), ws[m.group(5).replace("$", "")].value
    return isinstance(at(m.group(1), m.group(2)), (int, float)) and isinstance(loans, (int, float)) and loans < few


def _coord(ws, title: str, key) -> str:
    """Where a block's cell is on the sheet: the block by its title, the cell by its row and column labels."""
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith(title) and c.row > _below(ws):
                import tabs
                head, cols = tabs.header_of(ws, c.row, c.column)
                rr = head + 1
                while ws.cell(row=rr, column=c.column).value != key[0]:
                    rr += 1
                return ws.cell(row=rr, column=cols[key[1]]).coordinate
    raise KeyError(title)


def test_grids_grey_a_cell_under_the_fewest_loans_and_leave_it_out_of_the_scale(grids_book, tmp_path):
    b, _ = grids_book
    few = _fewest(b)
    ws, blocks, _ = _grids(b, tmp_path / "k0.xlsx", measure="RANR")
    bk, bd, loans = blocks["vs the book"], blocks["vs rest of band"], blocks["Loans"]
    kiosk = next(k for k, v in loans.items() if k[1] == KIOSK and k[0] != "All" and v)
    assert loans[kiosk] == 3 and bk[kiosk] < -50                     # the photo: 3 loans, the deepest gap
    real = [abs(v) for blk in (bk, bd) for k, v in blk.items() if isinstance(v, (int, float)) and loans[k] >= few]
    assert abs(bk[kiosk]) > max(real)
    # the colour scale's bound is the largest gap among the cells with enough loans
    meta = _views_rows(b)["G|FICO x CHANNEL|ranr_rate|meta"]
    assert meta[1] == pytest.approx(max(real)) and meta[4] == few
    # every cell under the fewest loans is grey, with no colour; every other number is coloured as before
    for title, blk in (("vs the book", bk), ("vs rest of band", bd)):
        for k, v in blk.items():
            if isinstance(v, (int, float)):
                assert _is_grey(ws, _coord(ws, title, k)) == (loans[k] < few), (title, k)
        rules = _rules_at(ws, _coord(ws, title, kiosk))
        grey = _inner(rules[1][0].formula[0])
        assert [r.dxf.fill for r, *_ in rules[:2]] == [None, None]
        coloured = [_inner(r.formula[0]) for r, *_ in rules[2:] if r.dxf.fill is not None]
        assert coloured and all(f.startswith(f"AND(NOT({grey})") for f in coloured)
    # what one cell says, and the note
    _, _, said = _grids(b, tmp_path / "k1.xlsx", measure="RANR", row=kiosk[0], column=KIOSK)
    assert said["The colour"] == f"Grey: only 3 loans, fewer than the {few} set on Control, so not coloured."
    note = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(3, 16)}
    assert note["Colour"].endswith(f"Grey: fewer loans than the {few} in Fewest loans in a pocket on Control, so "
                                   f"not coloured, and left out of the largest gap.")


def test_grids_fewest_loans_is_the_number_the_run_used_and_no_filter_by_offers_all_loans(tmp_path, monkeypatch):
    """Fewest loans left at its suggestion: the grey line is the number the Run worked out, not the old 30. With no
    Filter by picked in the launcher, nothing can be filtered, and the tab says where to pick one."""
    import math
    from pocketbook import control, perm, results
    import tabs
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    x = synth.write_extract(tmp_path / "src", n=4000)
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                            split="REV_DEBT", outcome="BAD_FLAG"))
    _answer(out.book)
    wb = load_workbook(out.book)
    for r in wb["Control"].iter_rows(min_row=control.FIRST_ROW):
        if r[control.KEY_COL - 1].value == "min_loans":
            r[control.CHOOSE_COL - 1].value = "Enough for 5 expected losses (suggested)"
    wb.save(out.book)
    assert book.run(out.book).ok
    want = max(2, math.ceil(5 / _bad_rate(_loans_file(x))))
    assert want != 30 and _views_rows(out.book)["G|fewest"][0] == want
    wb = load_workbook(out.book)
    assert tabs.options(wb, results.GRIDS, "Only loans where") == [results.ALL_LOANS]
    ws = wb[results.GRIDS]
    at = tabs.dropdown(ws, "Only loans where")
    assert ws.cell(row=at.row + 1, column=at.column).value == "Pick a Filter by in the launcher."
    ws, blocks, _ = _grids(out.book, tmp_path / "f0.xlsx")
    thin = next(k for k, v in _pockets(blocks["Loans"]).items() if v < want)
    _, _, said = _grids(out.book, tmp_path / "f1.xlsx", row=thin[0], column=thin[1])
    n = blocks["Loans"][thin]
    assert said["The colour"] == (f"Grey: only {n} loan{'' if n == 1 else 's'}, fewer than the {want} set on Control, "
                                  f"so not coloured.")
    assert f"Grey: fewer loans than the {want} in Fewest loans" in \
        next(ws.cell(row=r, column=3).value for r in range(3, 16) if ws.cell(row=r, column=2).value == "Colour")


def test_grids_the_books_own_figure_heads_vs_the_book(grids_book, tmp_path):
    """The book's rate, worked out here from the loan file, in the heading, following the Measure picked; the same
    with a value filtered, since the comparison stays the whole book."""
    b, x = grids_book
    rows = _loans_file(x)
    both = [(_num(r["RANR_AMT"]), _num(r["ORIG_BAL"])) for r in rows]
    kept = sum(a for a, c in both if a is not None and c is not None) / sum(c for a, c in both
                                                                         if a is not None and c is not None)
    want = {"Bad loans": f"{_bad_rate(rows) * 100:.2f}%", "RANR": f"{kept * 100:.2f}%",
            "Loan size": f"${_size(rows)[0]:,.0f}"}
    for k, (measure, fig) in enumerate(want.items()):
        for only in ("All loans", "Y"):
            ws, _, _ = _grids(b, tmp_path / f"h{k}{only[0]}.xlsx", measure=measure, **{ONLY: only})
            assert _head(ws, "vs the book") == f"vs the book (book: {fig})", (measure, only)
            assert _head(ws, "vs rest of band") == "vs rest of band"


def test_grids_only_loans_where_shows_one_values_grid_against_the_whole_book(grids_book, tmp_path):
    """Y's loans only: every count, rate and gap worked out again from the loan file. vs the book is against the
    whole book's rate; vs rest of band against the rest of the band among Y's loans."""
    from pocketbook import results
    import tabs
    b, x = grids_book
    rows = _loans_file(x)
    assert tabs.options(load_workbook(b), results.GRIDS, ONLY) == [results.ALL_LOANS, "N", "Y", "(blank)"]
    ws, blocks, said = _grids(b, tmp_path / "y.xlsx", **{ONLY: "Y"})
    rate, bk, bd, loans = (blocks[t] for t in ("Rate", "vs the book", "vs rest of band", "Loans"))
    cells = _cells([r for r in rows if r["SYS_FLAG"] == "Y"], loans)
    book_rate = _bad_rate(rows)
    min_events = _min_losses(b)
    assert {k for k, v in loans.items() if v} == set(cells)
    shown = 0
    for k, got in cells.items():
        assert loans[k] == len(got)
        assert rate[k] == pytest.approx(_bad_rate(got))
        bad = sum(1 for r in got if r["BAD_FLAG"] == "1")
        if bad < min_events:
            assert bk[k] is None and bd[k] is None
            continue
        assert bk[k] == pytest.approx(_bad_rate(got) / book_rate)                  # the whole book's rate
        if "All" not in k:
            rest = [r for kk, v in cells.items() if kk[0] == k[0] and "All" not in kk and kk != k for r in v]
            if rest and _bad_rate(rest):
                assert bd[k] == pytest.approx(_bad_rate(got) / _bad_rate(rest))
                shown += 1
    assert shown >= 3
    assert said["name"].endswith(", Bad loans, only loans where SYS_FLAG is Y")
    # its own heat scale, from its own cells with enough loans
    few = _fewest(b)
    v = _views_rows(b)
    key = "G|FICO x CHANNEL|where Y|ranr_rate"
    labels = v["G|FICO x CHANNEL|where Y|rows"]
    got = [abs(g) for i in range(1, len([x for x in labels if x]) + 1) for what in ("book", "band")
           for g, n in zip(v[f"{key}|{what}|{i}"], v[f"G|FICO x CHANNEL|where Y|loans|{i}"])
           if isinstance(g, (int, float)) and isinstance(n, int) and n >= few]
    assert v[f"{key}|meta"][1] == pytest.approx(max(got))
    assert v[f"{key}|meta"][1] != v["G|FICO x CHANNEL|ranr_rate|meta"][1]
    # the 3-loan pocket is grey here too, by its count among Y's loans
    kiosk = next(k for k in cells if k[1] == KIOSK and k[0] != "All")
    _, _, said = _grids(b, tmp_path / "y1.xlsx", measure="RANR", row=kiosk[0], column=KIOSK,
                        **{ONLY: "Y"})
    assert said["The colour"].startswith("Grey: only 3 loans, ")


def test_grids_loan_size_is_booked_per_loan_described_never_red_or_green(grids_book, tmp_path):
    from pocketbook import results
    import tabs
    b, x = grids_book
    rows = _loans_file(x)
    wb = load_workbook(b)
    assert results.SIZE_NAME in tabs.options(wb, results.GRIDS, "Measure")
    assert results.SIZE_NAME not in tabs.options(wb, results.POCKETS, "Measure")          # never tested
    assert results.SIZE_NAME not in tabs.options(wb, results.SPLIT, "Measure")
    book_avg = _size(rows)[0]
    for only, pick in (("All loans", rows), ("Y", [r for r in rows if r["SYS_FLAG"] == "Y"])):
        ws, blocks, _ = _grids(b, tmp_path / f"s{only[0]}.xlsx", measure=results.SIZE_NAME, **{ONLY: only})
        rate, bk, bd, loans = (blocks[t] for t in ("Rate", "vs the book", "vs rest of band", "Loans"))
        cells = _cells(pick, loans)
        for k, got in cells.items():
            avg = _size(got)[0]
            assert rate[k] == pytest.approx(avg) and bk[k] == pytest.approx(avg / book_avg)
            rest = [r for kk, v in cells.items() if kk[0] == k[0] and "All" not in kk and kk != k for r in v]
            assert bd[k] == (None if "All" in k or not rest else pytest.approx(avg / _size(rest)[0])), k
    # one cell in words, with the median, and no red or green
    k = next(k for k, v in _cells(rows, loans).items() if "All" not in k and len(v) > 100
             and all(_num(r["ORIG_BAL"]) is not None for r in v))
    got = _cells(rows, loans)[k]
    ws, blocks, said = _grids(b, tmp_path / "s1.xlsx", measure=results.SIZE_NAME, row=k[0], column=k[1])
    avg, med = _size(got)
    assert said["Rate"] == f"These {len(got):,} loans averaged ${avg:,.0f} booked, median ${med:,.0f}."
    assert said["vs the book"] == f"{avg / book_avg:.2f}× the average loan of the whole book."
    assert said["The colour"] == results.SAY_SIZE
    fills = {r.dxf.fill.fgColor.rgb[-6:] for r, *_ in _rules_at(ws, _coord(ws, "vs the book", k))
             if r.dxf.fill is not None and '="size"' in r.formula[0]}
    assert fills == {c for _, c in results.SIZE_STEPS}
    heat = [r.formula[0] for r, *_ in _rules_at(ws, _coord(ws, "vs the book", k))
            if r.dxf.fill is not None and "LOG(" in r.formula[0]]
    assert heat and all('<>"size"' in f for f in heat)


def test_grids_loan_size_is_not_offered_without_a_booked_amount(tmp_path):
    import dataclasses
    from openpyxl import Workbook
    from pocketbook import config as cfgmod, engine, results
    from pocketbook.ingest import read_table
    cfg, _ = synth.write(tmp_path / "cube", n=3000)
    raw = cfgmod.load(cfg).raw
    raw["split"] = {"field": "SYS_FLAG", "how": "each_value"}
    c = cfgmod.parse(raw)
    table = read_table(_flag_file(tmp_path, n=3000))
    res = engine.run(c, table)
    assert res.book_size is not None and res.grids[0].sizes
    bare = engine.run(dataclasses.replace(c, booked=""), table)
    assert bare.book_size is None and not bare.grids[0].sizes
    for r, offered in ((res, True), (bare, False)):
        wb = Workbook()
        results.write_grids(wb, r, results.Choices(wb), results.Views(wb))
        ch_ = wb[results.CHOICES]
        heads = {ch_.cell(row=1, column=j).value: j for j in range(1, ch_.max_column + 1)}
        opts = [ch_.cell(row=i, column=heads["Grids: Measure"]).value for i in range(2, ch_.max_row + 1)]
        assert (results.SIZE_NAME in opts) == offered


# ---- Borderline: a verdict whose shuffled p-value sits near the bar
# BACKLOG §6d item 7: "A verdict on a shuffled p-value near 5% can fall either way with another seed ... Recommended:
# flag them." The firm, 29 Sep 2026, by pop-up: "I don't like 'could fall either way' but flag it somehow", and chose
# "Borderline": "Net drain · borderline (p 0.048)". The tie-out of the same day found a shuffled p-value 3.5 standard
# errors from where a million shuffles put it. The rule (docs/statistics.md B2a): the p-value that decides the verdict,
# after the allowance, came from shuffling and sits within 2 of its own standard errors of the bar, either side; the
# standard error is sqrt(p (1 - p) / shuffles), times what the allowance multiplied the p-value by. A z test's or an
# exact test's p-value is the same on every run, so it is never borderline.

UNDER, OVER, FAR, NEAR_Z, TINY = 0.048, 0.052, 0.3, 0.049, 0.001
BORDER_SAID = ("Borderline: the test's p-value is within the shuffle's own margin of the 5% bar, so another run "
               "could read it the other way.")


def test_borderline_is_two_of_the_shuffles_own_standard_errors_either_side_of_the_bar():
    import math
    from pocketbook import stats
    assert stats.BORDERLINE_SE == 2
    assert stats.shuffle_se(0.05, 10_000) == pytest.approx(math.sqrt(0.05 * 0.95 / 10_000))       # 0.00218
    assert stats.shuffle_se(0.05, None) is None and stats.shuffle_se(None, 2_000) is None
    se = lambda p: stats.shuffle_se(p, 2_000)                                                 # noqa: E731
    # both sides of the bar, a pass and a fail alike
    assert stats.borderline(UNDER, se(UNDER), 0.95) and stats.borderline(OVER, se(OVER), 0.95)
    # the edge: 1.5 standard errors in, 2.5 out, either side
    s05 = se(0.05)
    for k, want in ((1.5, True), (2.5, False)):
        for p in (0.05 - k * s05, 0.05 + k * s05):
            assert stats.borderline(p, se(p), 0.95) is want, (k, p)
    # far from the bar, or with no standard error (a z or exact test), never
    assert not stats.borderline(FAR, se(FAR), 0.95) and not stats.borderline(TINY, se(TINY), 0.95)
    assert not stats.borderline(NEAR_Z, None, 0.95)
    # at 10,000 shuffles the margin is narrower: 0.048 is 0.9 standard errors from 5%, 0.045 is 2.4
    assert stats.borderline(0.048, stats.shuffle_se(0.048, 10_000), 0.95)
    assert not stats.borderline(0.045, stats.shuffle_se(0.045, 10_000), 0.95)
    # the bar follows the confidence: 0.098 is borderline at 90%, not at 95%
    assert stats.borderline(0.098, se(0.098), 0.90) and not stats.borderline(0.098, se(0.098), 0.95)
    assert stats.borderline_words(UNDER, 0.95) == "borderline (p 0.048)"
    assert stats.p_text(0.0496, 0.95) == "0.0496"          # three decimals would print it on the bar, 0.050
    # docs/statistics.md B2a's worked example
    se10 = lambda p: stats.shuffle_se(p, 10_000)                                              # noqa: E731
    assert round((0.05 - 0.048) / se10(0.048), 1) == 0.9 and round((0.05 - 0.045) / se10(0.045), 1) == 2.4
    assert round(se10(0.0048) * 20 / 2, 4) == 0.0069


def test_borderline_standard_error_follows_the_allowance_that_set_the_p_value():
    """The tie-out of 29 Sep 2026: one pocket's raw p set 21 others' through Benjamini-Hochberg. The standard error
    is the one of the raw p-value that sets each adjusted one, times the same m / j."""
    from pocketbook import engine
    a, b, c = 0.001, 0.002, 0.003
    # every adjusted p is 0.03, set by the third (0.03 x 3 / 3): so is every standard error, at x 1
    assert engine.adjust([0.01, 0.02, 0.03], "bh") == pytest.approx([0.03] * 3)
    assert engine.adjust_se([0.01, 0.02, 0.03], [a, b, c], "bh") == pytest.approx([c, c, c])
    # 0.01 sets its own (0.01 x 2 / 1 = 0.02): its error x 2; 0.5 sets its own at x 1
    assert engine.adjust_se([0.01, 0.5], [a, b], "bh") == pytest.approx([2 * a, b])
    assert engine.adjust_se([0.01, 0.2], [a, b], "bonferroni") == pytest.approx([2 * a, 2 * b])
    # capped at 1, nowhere near the bar: no standard error, so never borderline
    assert engine.adjust_se([0.01, 0.6], [a, b], "bonferroni")[1] is None
    assert engine.adjust_se([0.7, 0.8], [a, b], "bh") == pytest.approx([b, b])       # both set by 0.8, x 1
    assert engine.adjust_se([0.9, 1.0], [a, b], "bh") == [None, None]                # capped at 1
    assert engine.adjust_se([0.01, None, 0.02], [a, None, b], "none") == [a, None, b]
    # a p-value set by one that wasn't shuffled has none
    assert engine.adjust_se([0.01, 0.02], [a, None], "bh") == [None, None]


def _plant(real):
    """engine._shuffle_tests, then every tested pocket's p-value in two grids planted (both comparisons), so each
    grid's family is equal p-values and the allowance leaves them as they are: FICO x CHANNEL's charge-offs just
    under the bar and what it kept far under it, with Bad loans' z and exact tests at 0.049; FICO x ASSET_CLASS's
    charge-offs just over the bar and what it kept far over it. The split's halves of FICO x CHANNEL: charge-offs
    just under, and the pooled figure just over."""

    def planted(config, measures, per_row, n, built, halved):
        real(config, measures, per_row, n, built, halved)
        two = {g.dimension.lower(): g for g, _, _ in halved}
        want = [(two["channel"], {"gco_rate": UNDER, "ranr_rate": TINY, "outcome_loans": NEAR_Z}),
                (two["asset_class"], {"gco_rate": OVER, "ranr_rate": FAR})]
        for g, ps in want:
            for _, c in g.inner():
                for m, p in ps.items():
                    s = c.rates[m]
                    if s.p_book is not None:
                        s.p_book = p
                    if s.p_band is not None:
                        s.p_band = p
        g = two["channel"]
        for got in g.split_compare.values():
            x = got.get("gco_rate")
            if x is not None and x[1] is not None:
                got["gco_rate"] = (x[0], UNDER, x[2], x[3])
        if g.split_pooled.get("gco_rate", {}).get("ratio_p") is not None:
            g.split_pooled["gco_rate"]["ratio_p"] = OVER
    return planted


@pytest.fixture(scope="module")
def border_book(tmp_path_factory):
    """The synthetic book, FICO by CHANNEL and by ASSET_CLASS, split in halves by REV_DEBT, with p-values planted
    either side of the bar (_plant): the book, the engine's result for it, the Run's outcome, and the book
    calculated."""
    from pocketbook import engine, perm
    from recalc import calculated_book
    d = tmp_path_factory.mktemp("border")
    got = {}
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 2_000)
        mp.setattr(engine, "_shuffle_tests", _plant(engine._shuffle_tests))
        run = engine.run

        def kept(*a, **k):
            got["res"] = run(*a, **k)
            return got["res"]
        mp.setattr(engine, "run", kept)
        out = book.set_up(synth.write_extract(d / "src", n=8000),
                          choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL", "ASSET_CLASS"),
                                             split="REV_DEBT"))
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return {"b": out.book, "res": got["res"], "ran": ran, "calc": calculated_book(out.book)}


def _two(res):
    """The two two-way grids, FICO x CHANNEL and FICO x ASSET_CLASS."""
    by = {g.dimension.lower(): g for g in res.grids}
    return by["channel"], by["asset_class"]


def _flagged(m: str, grid) -> list:
    return [c.rates[m].borderline for _, c in grid.inner() if c.rates[m].borderline]


def test_borderline_engine_flags_a_shuffled_verdict_just_either_side_of_the_bar_and_nothing_else(border_book):
    from pocketbook import engine
    res = border_book["res"]
    ch_, ac = _two(res)
    for grid, p, words in ((ch_, UNDER, "borderline (p 0.048)"), (ac, OVER, "borderline (p 0.052)")):
        turned = 0
        for _, c in grid.inner():
            s = c.rates["gco_rate"]
            if s.flag in (engine.WORSE, engine.BETTER, engine.UNSURE_WORSE, engine.UNSURE_BETTER):
                assert s.borderline == words, (grid.dimension, s.flag)
                assert (s.p_band if s.by_band else s.p_book) == pytest.approx(p)
                assert (s.se_band if s.by_band else s.se_book) == pytest.approx((p * (1 - p) / 2_000) ** 0.5)
                turned += 1
            else:
                assert s.borderline is None, (grid.dimension, s.flag)     # in line inside the loss line: not p's call
            assert s.worse_borderline == (words if s.flag in (engine.WORSE, engine.UNSURE_WORSE) else None)
        assert turned >= 3, grid.dimension
    # a pass just under the bar reads worse, a fail just over reads not significant: both flagged
    assert any(c.rates["gco_rate"].flag == engine.WORSE for _, c in ch_.inner())
    assert any(c.rates["gco_rate"].flag == engine.UNSURE_WORSE for _, c in ac.inner())
    # a z or exact test at 0.049 is never borderline; nor is a shuffled p-value far from the bar either side
    zs = [c.rates["outcome_loans"] for _, c in ch_.inner() if c.rates["outcome_loans"].p_book is not None]
    assert zs and all(s.p_book == pytest.approx(NEAR_Z) and s.test in (engine.Z_TEST, engine.EXACT_TEST)
                      and s.se_book is None and s.borderline is None for s in zs)
    assert any(s.flag == engine.WORSE for s in zs)                     # a real verdict, still not flagged
    assert _flagged("ranr_rate", ch_) == [] and _flagged("ranr_rate", ac) == []
    assert any(c.rates["ranr_rate"].flag == engine.WORSE for _, c in ch_.inner())


def _pk_rows(res):
    """(grid, band, segment, measure, RateStat) in _pockets' own order (live._write_pockets)."""
    rates = [m for m in res.measures if m.is_rate]
    for grids in (res.grids, res.three_way):
        for g in grids:
            for (bl, dl), c in g.inner():
                for m in rates:
                    yield g, bl, dl, m, c.rates[m.name]


def _note(ws) -> dict:
    """A tab's method note, {label: words}, from its calculated cells."""
    out = {}
    top = next(c.row for row in ws.iter_rows(max_row=40) for c in row if c.value == "How this tab works")
    col = next(c.column for c in ws[top] if c.value == "How this tab works")
    for r in range(top + 1, top + 20):
        a, b = ws.cell(row=r, column=col).value, ws.cell(row=r, column=col + 1).value
        if a is None:
            break
        out[a] = b
    return out


def test_borderline_pockets_worse_says_it_beside_the_word_and_keeps_the_words_colour(border_book, tmp_path):
    import tabs
    from pocketbook import live, results
    res, calc = border_book["res"], border_book["calc"]
    # every pocket on _pockets, as the tabs read it, against the engine's own flag
    pk = calc[live.POCKETS]
    n = 0
    for r, (g, bl, dl, m, s) in enumerate(_pk_rows(res), start=live.P_FIRST):
        assert (pk.cell(row=r, column=live.P_BAND).value, pk.cell(row=r, column=live.P_MEASURE).value) == (bl, m.name)
        btxt = pk.cell(row=r, column=live.P_BTXT).value or None
        assert btxt == (s.borderline[len("borderline (p "):-1] if s.borderline else None), (g.dimension, bl, dl,
                                                                                               m.name)
        worse = pk.cell(row=r, column=live.P_WORSE).value
        assert pk.cell(row=r, column=live.P_WORSE_SAID).value == live.flagged(worse, s.worse_borderline)
        n += bool(btxt)
    assert n >= 6
    # the Pockets tab: charge-offs, a pass and a fail, each flagged; Bad loans at 0.049 by the z test, never
    ws = tabs.calculated(tabs.choose(border_book["b"], tmp_path / "gco.xlsx", results.POCKETS,
                                     measure="GCOs ($)"), results.POCKETS)
    said = {x["worse_said"] for x in tabs.pockets(ws)}
    assert "Yes · borderline (p 0.048)" in said and "Not sure · borderline (p 0.052)" in said, said
    ws = tabs.calculated(tabs.choose(border_book["b"], tmp_path / "bad.xlsx", results.POCKETS,
                                     measure="Bad loans"), results.POCKETS)
    rows = tabs.pockets(ws)
    assert any(x["worse"] == "Yes" for x in rows) and not any("borderline" in str(x["worse_said"]) for x in rows)
    # its colour is the word's, borderline or not: every rule on Worse? compares the word without the flag
    written = load_workbook(border_book["b"])[results.POCKETS]
    col = results.col(results.K_WORSE)
    rules = [r.formula[0] for rng in written.conditional_formatting for r in rng.rules
             if str(rng.sqref).startswith(col) and f'="{live.YES}"' in r.formula[0]]
    assert rules and all(f'IFERROR(LEFT(${col}' in f and 'FIND(" · "' in f for f in rules), rules
    # the method note says what it means, in the firm's plain words and under 25, at the end of its p-value item
    # so that no row of the tab moves
    note = _note(calc[results.POCKETS])
    assert note["p-value"].endswith(" " + BORDER_SAID)
    assert len(BORDER_SAID.split()) <= 25 + 1


def test_borderline_paid_cost_kept_together_says_it_and_its_colour_and_chart_stay_the_words(border_book, tmp_path):
    import tabs
    from pocketbook import results
    res = border_book["res"]
    ch_, _ = _two(res)
    path = tabs.choose(border_book["b"], tmp_path / "pck.xlsx", results.PCK, grid="FICO x CHANNEL")
    ws = tabs.calculated(path, results.PCK)
    rows = tabs.pck(ws)
    by = {(bl, dl): c for (bl, dl), c in ch_.inner()}
    said = set()
    for x in rows:
        c = by[(x["band"], x["seg"])]
        want = results.together_said(x["together"], c.rates["gco_rate"].borderline, c.rates["ranr_rate"].borderline)
        assert x["together_said"] == want, x
        said.add(x["together_said"])
    assert "Net drain · borderline (p 0.048)" in said, said
    # red and green follow the word: the Net drain rows' dots are the red ones
    from recalc import recalc
    chart = recalc(path, tmp_path / "pck-chart")[results.CHART]
    drains = [k for k, x in enumerate(rows, start=1) if x["together"] == "Net drain"]
    assert drains and all(isinstance(chart.cell(row=k, column=results.H_RX).value, (int, float)) for k in drains)
    tog = [r.formula[0] for rng in ws.formulas.conditional_formatting for r in rng.rules
           if str(rng.sqref).startswith(results.col(results.C_TOG))]
    named = [f for f in tog if "Net drain" in f or "Strong" in f]
    assert named and all(results.col(results.C_H_TOG) in f for f in named)
    note = _note(ws)
    assert note["Together"].endswith(f" {BORDER_SAID} Together gives the p-value of each side that is, GCOs "
                                     f"first.")


def _summary_p(ws, measure: str):
    h = next(c.row for row in ws.iter_rows() for c in row if c.value == "High half worse in")
    for r in range(h + 1, h + 8):
        if ws.cell(row=r, column=2).value == measure:
            return ws.cell(row=r, column=7).value
    raise KeyError(measure)


def test_borderline_split_says_it_in_place_of_the_p_value_and_keeps_it_bold(border_book, tmp_path):
    import tabs
    from pocketbook import results
    res = border_book["res"]
    ch_, _ = _two(res)
    ws = tabs.calculated(tabs.choose(border_book["b"], tmp_path / "split.xlsx", results.SPLIT,
                                     grid="FICO x CHANNEL", measure="GCOs ($)"), results.SPLIT)
    ps = tabs.block(ws, "p-value per pocket")
    flagged = 0
    for (bl, dl), c in ch_.inner():
        x = ch_.split_compare.get((bl, dl), {}).get("gco_rate")
        se = ch_.split_se.get((bl, dl), {}).get("gco_rate")
        want = results.split_said(x[1] if x else None, se, 0.95)
        got = ps[(bl, dl)]
        assert got == (pytest.approx(want) if isinstance(want, float) else want), (bl, dl, got, want)
        flagged += want == "borderline (p 0.048)"
    assert flagged >= 3
    assert _summary_p(ws, "GCOs ($)") == "borderline (p 0.052)"
    assert isinstance(_summary_p(ws, "Bad loans"), float)                    # the z test's: never
    # the bold rule reads the number behind the words, one range per column
    rules = [r.formula[0] for rng in ws.formulas.conditional_formatting for r in rng.rules
             if "ISTEXT" in r.formula[0] and "significance_bar" in r.formula[0]]
    assert len(rules) >= len(ch_.dim_labels)
    assert _note(ws)["p-value"].endswith(" " + BORDER_SAID)


def test_borderline_start_here_record_and_the_launcher_count_and_name_them(border_book):
    import tabs
    from pocketbook import engine, launcher, live
    res, calc, ran = border_book["res"], border_book["calc"], border_book["ran"]
    names = book._names(res)
    gco = [(g, bl, dl, c.rates["gco_rate"]) for g in res.grids for (bl, dl), c in g.inner()]
    worse = [x for x in gco if x[3].flag == engine.WORSE and x[3].material is not False and (x[3].dollars or 0) > 0]
    border = [x for x in worse if x[3].worse_borderline]
    assert border and len(border) < len(gco)
    # the launcher's headline
    assert ran.summary["borderline"] == len(border)
    tiles = launcher.finished_tiles(ran.summary)
    assert tiles[0][2].endswith(f" · {len(border):,} borderline")
    # Start here: the tile, and each of the five largest that is borderline says so beside its segment
    ws = calc["Start here"]
    tile = next(ws.cell(row=c.row + 1, column=c.column).value for row in ws.iter_rows() for c in row
                if c.value == "Pockets worse and material, GCOs")
    assert tile.endswith(f" · {len(border)} borderline"), tile
    head = next(c.row for row in ws.iter_rows() for c in row if c.value == "Largest, worse and material")
    listed = [(ws.cell(row=head + k, column=2).value, ws.cell(row=head + k, column=3).value)
              for k in range(1, book.TOP_ROWS + 1)]
    want = {(f"{names[g.band]} {bl}", str(dl)): s.worse_borderline for g, bl, dl, s in worse}
    assert listed[0][0] and any("borderline (p 0.048)" in str(seg) for _, seg in listed), listed
    for band, seg in (x for x in listed if x[0]):
        assert seg == live.flagged(tabs.word(seg), want[(band, tabs.word(seg))]), (band, seg)
    # Record: the rule, and the count now; the Run's line names the worst with its flag
    rec = tabs.record(calc)
    assert "within 2 of its own standard errors of the 5% bar" in rec["Borderline"]
    assert "never borderline" in rec["Borderline"]
    n = len(gco)
    assert rec["Borderline now: GCOs ($)"] == (f"{sum(1 for x in gco if x[3].borderline)} of {n:,} pockets on "
                                                  f"the grids have a borderline verdict "
                                                  f"({sum(1 for x in gco if x[3].worse_borderline)} on Worse?).")
    assert not any(k.startswith("Borderline now: Bad loans") for k in rec)          # never shuffled
    top = max(worse, key=lambda x: x[3].dollars)
    title = next(m.title for m in res.measures if m.name == "gco_rate")
    line = f"Worst for {title}: {names[top[0].band]} {top[1]} / {names[top[0].dimension]} {top[2]}"
    assert top[3].worse_borderline and f"{line} · {top[3].worse_borderline}." in ran.lines, ran.lines


# ---- Filter by, apart from Split by (30 Sep 2026)
# The firm found the Grids' filter worked only off Split by: "Wait only works on split by? Isn't that for like above and
# below median". Offered a separate Filter by (any category of six values or fewer, or the origination year), they
# answered "Yes hoping to have this by morning". Their use: 2022 to 2024 originations, flipped year by year to show
# the pockets hold across vintages. Every rule decided for a filtered view stays: vs the book is the whole book ("we
# keep things compared to the whole book that's just kind of the point"), vs rest of band is within the filtered
# loans, grey and the heat scale go by the view's own cells.

ONLY_YEAR = "Only loans where ORIG_YEAR is"
KIOSK_2024_RANR = -100000.0
YEARS = ("2022", "2023", "2024", "(no date)")


def _vintage_file(tmp_path, n=4000, first=2022, years=3):
    """The synthetic book, its loans made from `first` over `years` years; every 151st has no origination date (a
    blank), so it belongs to no year. A second date column, FIRST_PAY_DATE, sits beside it 400 days later, so a year
    read from the wrong column lands in the wrong year. Three Kiosk loans at FICO 700: two booked $1,000,000 in 2022
    keeping nothing, one booked $1,000 in 2024 losing $100,000, so the 2024 view shows a figure far wider than any
    whole-book cell."""
    from datetime import date, timedelta
    src = synth.write_extract(tmp_path / "src", n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    rng = random.Random(30)
    span = (date(first + years, 1, 1) - date(first, 1, 1)).days - 1
    for i, r in enumerate(rows):
        made = date(first, 1, 1) + timedelta(days=rng.randint(0, span))
        r["ORIG_DATE"] = "" if i % 151 == 5 else made.isoformat()
        r["FIRST_PAY_DATE"] = (made + timedelta(days=400)).isoformat()
    for k, (bal, ranr, made) in enumerate(((1e6, 0.0, "2022-03-01"), (1e6, 0.0, "2022-06-01"),
                                           (1000.0, KIOSK_2024_RANR, "2024-05-01"))):
        rows.append({**rows[10], "LOAN_NBR": f"K{k}", "FICO": 700, "CHANNEL": KIOSK, "ORIG_BAL": bal, "BAD_FLAG": 0,
                     "GCO_AMT": 0, "RANR_AMT": ranr, "ORIG_DATE": made, "FIRST_PAY_DATE": "2025-01-01"})
    out = tmp_path / "vintages.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return out


def _year(r) -> str:
    """A loan's origination year, worked out here from the loan file: its ORIG_DATE's first four characters."""
    return r["ORIG_DATE"][:4] if r["ORIG_DATE"] else "(no date)"


def test_filter_by_launcher_offers_every_category_and_the_origination_year(tmp_path, monkeypatch):
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    x = _vintage_file(tmp_path, n=1500)
    loans = _loans_file(x)
    f = _flow(x)
    assert f.heads() == ("Cut into bands", "Segment by", "Split by", "Filter 1", "Filter 2")
    rows = f.rows()
    by = {r["name"]: r for r in rows}
    for name in ("CHANNEL", "ASSET_CLASS"):                                  # every category can filter
        assert by[name]["d"] == {"on": False, "radio": True}
    for name in ("FICO", "ORIG_BAL", "REV_DEBT", "BAD_FLAG", "LOAN_NBR", "ORIG_DATE"):
        assert by[name]["d"] is None, name
    # ORIG_YEAR sits with the categories: it splits and filters, never segments
    none = sum(1 for r in loans if not r["ORIG_DATE"])
    yr = by[ch.ORIG_YEAR]
    assert yr["what"] == f"Origination year, from ORIG_DATE · 3 values · {none} with no date"
    assert yr["b"] is None and yr["c"] == {"on": False, "radio": True} and yr["d"] == {"on": False, "radio": True}
    names = [r["name"] for r in rows]
    assert names.index("ASSET_CLASS") < names.index(ch.ORIG_YEAR) < names.index("BAD_FLAG")
    # Filter by is its own pick: the split and the segments stay as they are
    f.click(ch.ORIG_YEAR, "d")
    f.click("REV_DEBT", "c")
    assert f.filter == ch.ORIG_YEAR and f.split == "REV_DEBT" and {"CHANNEL", "ASSET_CLASS"} <= f.seg
    f.click("CHANNEL", "d")                                                  # one column filters, or none
    assert f.filter == "CHANNEL" and "CHANNEL" in f.seg and f.split == "REV_DEBT"
    f.click("CHANNEL", "d")
    assert f.filter is None
    f.click(ch.ORIG_YEAR, "d")
    f.pick_outcome("BAD_FLAG")
    f.answer_outcome(True)
    got = f.choices()
    assert got.filter == ch.ORIG_YEAR and got.split == "REV_DEBT"
    ok, said = f.summary()
    assert ok and said.endswith("adds 4 more. Grids can show only the loans of one ORIG_YEAR."), said
    # written to Control beside the split, and read back by the next Set up
    f.next()
    assert f.page == "answer", f.message
    ws = load_workbook(f.book())[control.SHEET]
    shown = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=control.CHOOSE_COL).value
             for r in range(control.FIRST_ROW, ws.max_row + 1)}
    assert shown["Split every pocket by"] == "REV_DEBT" and shown["Filter the Grids by"] == ch.ORIG_YEAR
    again = _flow(x)
    assert again.filter == ch.ORIG_YEAR and again.split == "REV_DEBT"
    # a test of a new variable has no Filter by; an extract with no origination date has no ORIG_YEAR row
    f.set_mode("new")
    assert all(r["d"] is None for r in f.rows()) and f.choices().filter is None
    bare = tmp_path / "bare" / "nodate.csv"
    bare.parent.mkdir()
    with open(bare, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=[c for c in loans[0] if not c.endswith("_DATE")])
        w.writeheader()
        w.writerows([{k: v for k, v in r.items() if not k.endswith("_DATE")} for r in loans])
    assert ch.ORIG_YEAR not in [r["name"] for r in _flow(bare).rows()]


def test_filter_by_refuses_more_than_six_values_in_the_launcher_and_at_the_run(tmp_path, monkeypatch):
    assert ch.too_many_to_filter("X", ch.FILTER_MOST_VALUES) is None           # six values filter
    said = ch.too_many_to_filter("REGION", 7)
    assert said == ("REGION has 7 values. The Grids can be filtered by a column of 6 values at most: with more, each "
                    "value's loans are too few to fill a grid. Filter by a column with fewer values, or by none.")
    f = _flow(_flag_file(tmp_path))
    f.pick_outcome("BAD_FLAG")
    f.answer_outcome(True)
    f.click("REGION", "d")
    assert f.summary() == (False, said) and f.states()["next"] == "disabled"
    f.click("SYS_FLAG", "d")                                                 # Y, N and a blank: two values
    assert f.summary()[0]
    # nine years of originations are too many too; the loans with no date aren't counted
    many = _flow(_vintage_file(tmp_path / "nine", n=1500, first=2016, years=9))
    many.pick_outcome("BAD_FLAG")
    many.answer_outcome(True)
    many.click(ch.ORIG_YEAR, "d")
    assert many.summary() == (False, ch.too_many_to_filter(ch.ORIG_YEAR, 9))
    # the Run refuses it too, in the same words, whatever wrote the workbook
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    out = book.set_up(f.extract, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                                    filter="REGION", outcome="BAD_FLAG"))
    _answer(out.book)
    ran = book.run(out.book)
    assert not ran.ok and ran.lines == [f"Couldn't run: {said}"]


@pytest.fixture(scope="module")
def vintage_book(tmp_path_factory):
    """FICO x CHANNEL on 2022 to 2024 originations, split in halves by REV_DEBT (a number) and filtered by ORIG_YEAR."""
    from pocketbook import perm
    d = tmp_path_factory.mktemp("vintage")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 200)
        x = _vintage_file(d)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                                split="REV_DEBT", filter=ch.ORIG_YEAR, outcome="BAD_FLAG"))
        _answer(out.book)
        # edges typed at whole scores, so each band's label says exactly which loans it holds (an equal-loans edge
        # such as 654.2 puts FICO 654 in a band labelled "... - 653")
        from test_book import at
        wb = load_workbook(out.book)
        wb["Columns"][at(wb, "FICO", book.C_EDGES)] = "620; 680; 740"
        wb.save(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return out.book, x, ran


def test_filter_by_year_under_a_number_split_every_cell_from_the_loan_file(vintage_book, tmp_path):
    """Filter by works whatever Split by is doing: here REV_DEBT halves every pocket, and ORIG_YEAR filters the Grids.
    For every year, and the loans with no date, every count, rate and comparison worked out again from the loan
    file: vs the book against the whole book's rate, vs rest of band against the rest of the band among that
    year's loans; and each view's heat scale from its own cells."""
    from pocketbook import results
    import tabs
    b, x, ran = vintage_book
    rows = _loans_file(x)
    wb = load_workbook(b)
    assert tabs.options(wb, results.GRIDS, ONLY_YEAR) == [results.ALL_LOANS, *YEARS]
    assert tabs.options(wb, results.POCKETS, "Pockets") == ["Two-way", "Split by REV_DEBT"]    # the split is its own
    book_rate = _bad_rate(rows)
    few, min_events = _fewest(b), _min_losses(b)
    v = _views_rows(b)
    bounds = []
    for year in YEARS:
        ws, blocks, said = _grids(b, tmp_path / f"y{year[:3]}.xlsx", **{ONLY_YEAR: year})
        rate, bk, bd, loans = (blocks[t] for t in ("Rate", "vs the book", "vs rest of band", "Loans"))
        cells = _cells([r for r in rows if _year(r) == year], loans)
        assert {k for k, n in loans.items() if n} == set(cells), year
        shown = 0
        for k, got in cells.items():
            assert loans[k] == len(got), (year, k)
            assert rate[k] == pytest.approx(_bad_rate(got)), (year, k)
            bad = sum(1 for r in got if r["BAD_FLAG"] == "1")
            if bad < min_events:
                if "All" not in k:                                          # a pocket; a total isn't held back
                    assert bk[k] is None and bd[k] is None, (year, k)
                continue
            assert bk[k] == pytest.approx(_bad_rate(got) / book_rate), (year, k)            # the whole book
            if "All" not in k:
                rest = [r for kk, vv in cells.items() if kk[0] == k[0] and "All" not in kk and kk != k for r in vv]
                if rest and _bad_rate(rest):
                    assert bd[k] == pytest.approx(_bad_rate(got) / _bad_rate(rest)), (year, k)
                    shown += 1
        assert shown >= (3 if year != "(no date)" else 0), year
        assert said["name"].endswith(f", Bad loans, only loans where ORIG_YEAR is {year}"), said["name"]
        assert _head(ws, "vs the book") == f"vs the book (book: {book_rate * 100:.2f}%)"      # the whole book's
        # its own heat scale: the largest gap in points among its own cells with enough loans
        key = f"G|FICO x CHANNEL|where {year}"
        labels = v[f"{key}|rows"]
        gaps = [abs(g) for i in range(1, len([y for y in labels if y]) + 1) for what in ("book", "band")
                for g, n in zip(v[f"{key}|ranr_rate|{what}|{i}"], v[f"{key}|loans|{i}"])
                if isinstance(g, (int, float)) and isinstance(n, int) and n >= few]
        if gaps:
            assert v[f"{key}|ranr_rate|meta"][1] == pytest.approx(max(gaps)), year
            bounds.append(v[f"{key}|ranr_rate|meta"][1])
    assert len(set(bounds)) == len(bounds) >= 3                             # each year its own scale
    # no note under the dropdown: a Filter by was picked; the method note says what ORIG_YEAR is
    ws = wb[results.GRIDS]
    at = tabs.dropdown(ws, ONLY_YEAR)
    assert ws.cell(row=at.row + 1, column=at.column).value != results.SAY_NO_FILTER
    note = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(3, 16)}
    assert note["Only loans where"].endswith("ORIG_YEAR is the year in ORIG_DATE; (no date) holds the loans without "
                                             "a readable date.")
    # Record and the Run name the filter where they name the split
    counts = {y: sum(1 for r in rows if _year(r) == y) for y in YEARS}
    listed = ", ".join(f"{y} ({n:,} loans)" for y, n in counts.items())
    rec = tabs.record(b)
    assert rec["Grids filter"].startswith(f"ORIG_YEAR (the year in ORIG_DATE): {listed}. ")
    assert rec["Split"].startswith("REV_DEBT, each pocket halved")
    assert f"Grids filter by ORIG_YEAR (the year in ORIG_DATE): {listed}. Pick one in Grids' Only loans where." \
        in ran.lines


def test_filter_by_grids_widths_fit_the_filtered_values(vintage_book):
    """The one data width fits every value any view shows, filtered ones too: the 2024 view's Kiosk pocket keeps
    -10,000% of what it booked, far wider than any whole-book figure."""
    import dataclasses
    import math
    from openpyxl import Workbook
    from pocketbook import config as cfgmod, engine, results
    from pocketbook.ingest import read_table
    b, x, _ = vintage_book
    cfg = cfgmod.parse(book.read_book(b)[0])
    table = read_table(x)
    widths = {}
    for filt in (ch.ORIG_YEAR, None):
        res = engine.run(dataclasses.replace(cfg, filter_by=filt), table)
        _, _, _, _, fit = results.grid_views(res, results.Views(Workbook()))
        widths[filt] = results.grid_widths(fit)[0]
    kiosk = f"{KIOSK_2024_RANR / 1000 * 100:.2f}%"                          # the 2024 view's Kept after losses
    assert widths[ch.ORIG_YEAR] >= min(results.DATA_CAP, len(kiosk) + 2)
    assert widths[ch.ORIG_YEAR] > widths[None]                                # so the filtered views set it
    got = load_workbook(b)[results.GRIDS].column_dimensions["C"].width
    assert math.isclose(got, widths[ch.ORIG_YEAR], abs_tol=0.01)


def test_filter_by_origination_year_splits_too_with_each_year_against_the_rest(tmp_path, monkeypatch):
    """ORIG_YEAR as Split by: each year set against the rest of its pocket on the Split tab, and whether the years
    differ at all, the consistency test across vintages. One definition of the year serves both picks. With no
    Filter by, the Grids offer All loans only, though the split is a category."""
    import dataclasses
    from pocketbook import config as cfgmod, engine, perm, results
    from pocketbook.ingest import read_table
    import tabs
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    x = _vintage_file(tmp_path, n=6000)
    rows = _loans_file(x)
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",), segments=("CHANNEL",),
                                            split=ch.ORIG_YEAR, outcome="BAD_FLAG"))
    _answer(out.book)
    raw = book.read_book(out.book)[0]
    assert raw["split"] == {"field": ch.ORIG_YEAR, "how": "each_value"} and "filter_by" not in raw
    ran = book.run(out.book)
    assert ran.ok, ran.lines
    wb = load_workbook(out.book)
    split = tabs.options(wb, results.SPLIT, "Grid")
    assert [s for s in split if s.startswith("FICO x CHANNEL · ")] == [
        f"FICO x CHANNEL · ORIG_YEAR {y} vs rest" for y in YEARS]
    assert tabs.options(wb, results.GRIDS, "Only loans where") == [results.ALL_LOANS]
    at = tabs.dropdown(wb[results.GRIDS], "Only loans where")
    assert wb[results.GRIDS].cell(row=at.row + 1, column=at.column).value == results.SAY_NO_FILTER
    ws = tabs.calculated(tabs.choose(out.book, tmp_path / "s.xlsx", results.SPLIT,
                                     grid="FICO x CHANNEL · ORIG_YEAR 2023 vs rest"), results.SPLIT)
    text = [str(v) for row in ws.iter_rows(values_only=True) for v in row if v is not None]
    differ = next(t for t in text if t.startswith("Do the values of ORIG_YEAR differ at all? Bad loans: "))
    assert "on 3 degrees of freedom" in differ
    assert any(t.startswith("Same in every pocket? Bad loans: ") for t in text)
    assert tabs.record(out.book)["Split"].startswith("ORIG_YEAR (the year in ORIG_DATE), each pocket split by each")
    # the engine's parts are the years of the loan file: every loan in one, a loan with no date in none of the years
    res = engine.run(cfgmod.parse(raw), read_table(x))
    g = res.grids[0]
    assert g.split_parts == list(YEARS)
    assert g.split_general["outcome_loans"]["df"] == 3
    for y in YEARS:
        n = sum(c.rows for (b_, d_, p), c in g.split_cells.items() if p == y and engine.ALL not in (b_, d_))
        assert n == sum(1 for r in rows if _year(r) == y), y
    assert [r[ch.ORIG_YEAR] for r in res.table.rows] == [_year(r) for r in rows]
    assert engine.origination_years(read_table(x), "ORIG_DATE") == [_year(r) for r in rows]
    # no column marked Origination date: refused in words, never guessed
    with pytest.raises(engine.DataRefused, match="no column in this extract is marked so"):
        engine.run(dataclasses.replace(cfgmod.parse(raw), origination_date=None), read_table(x))


def test_a_band_cut_between_whole_numbers_is_labelled_by_the_values_it_holds():
    """Found building Filter by, 30 Sep 2026: FICO cut at an equal-loan point of 654.2 read "496 - 653" and held
    654. A band is labelled from the first shown value it holds to the last."""
    from pocketbook import engine
    assert engine.band_labels((654.2, 700.0), 496, 850, whole=True) == ["496 - 654", "655 - 699", "700 - 850"]
    # dollars with cents read to the nearest dollar, as before: 26,803.10 is 26,803
    assert engine.band_labels((26803.1, 38548.6), 5000, 90000) == ["5,000 - 26,802", "26,803 - 38,548", "38,549 - 90,000"]
    assert engine.band_labels((620, 680, 740), 500, 850) == ["500 - 619", "620 - 679", "680 - 739", "740 - 850"]
    for edge, value in ((654.2, 654), (654.2, 655), (700.0, 699), (700.0, 700)):
        label = engine.band_labels((edge,), 496, 850, whole=True)[0 if value < edge else 1]
        lo, hi = (float(x.replace(",", "")) for x in label.split(" - "))
        assert lo <= value <= hi, (edge, value, label)


# ---- Cut at whole dollars (30 Sep 2026)
# BACKLOG §6d, "Open, for the firm: a dollar band's label when its edge has cents": the evening tie-out found ORIG_BAL
# cut at an equal-loan point of $37,950.548, its bands labelled "26,324 - 37,950" and "37,951 - 49,151", and a loan of
# $37,950.99 in the second, whose label starts above it. The firm: "Cut at whole dollars is fine". The rule
# (engine.whole_cut): an edge PocketBook cuts on a column whose labels read in whole units (every edge 100 or more
# either way) and whose values carry cents is raised to the next whole number. A whole-unit label reads a value with
# its cents dropped ($37,950.99 reads 37,950), so every loan's reading lies inside its own band's label and no other.
# Typed edges are the analyst's and stay as typed; a column of whole numbers (FICO) and a ratio are left as they were.

def _reading_in(label: str, v: float) -> bool:
    import math
    lo, hi = (float(x.replace(",", "")) for x in label.split(" - "))
    return lo <= math.floor(v) <= hi


def _dollar_file(tmp_path, n, bal=None):
    """The synthetic book (ORIG_BAL in dollars and cents), each loan's ORIG_BAL replaced by `bal(i, as_read)`."""
    cfg, data = synth.write(tmp_path, n=n)
    if bal is not None:
        rows = list(csv.DictReader(open(data, encoding="utf-8")))
        for i, r in enumerate(rows):
            r["ORIG_BAL"] = bal(i, r["ORIG_BAL"])
        with open(data, "w", newline="", encoding="utf-8") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
    return cfg, data


def _dollar_run(tmp_path, band, n=6000, bal=None):
    import copy
    from pocketbook import config as cfgmod, engine
    from pocketbook.ingest import read_table
    cfg, data = _dollar_file(tmp_path, n, bal)
    raw = copy.deepcopy(cfgmod.load(cfg).raw)
    raw["bands"] = [band]
    tbl = read_table(data)
    return engine.run(cfgmod.parse(raw), tbl), tbl


def test_whole_dollars_the_firms_loan_of_37950_99_sits_under_the_label_that_names_it():
    from pocketbook import engine
    vals = [26324.50, 30001.25, 37950.10, 37950.99, 44000.00, 49151.40]
    raw_edge = (37950.10 + 37950.99) / 2                                     # 37,950.545: the equal-loan point
    edges = engine.cut_edges(vals, 2, "equal_loans")
    assert edges == (37951.0,) and raw_edge < 37951
    labels = engine.band_labels(edges, min(vals), max(vals), whole=engine.all_whole(vals))
    assert labels == ["26,324 - 37,950", "37,951 - 49,151"]
    assert engine.band_of(37950.99, edges, labels) == "26,324 - 37,950"
    for v in vals:                                                            # inside its own label, no other
        assert [lab for lab in labels if _reading_in(lab, v)] == [engine.band_of(v, edges, labels)], v
    # the edge as the build before cut it: 37,950.99 sat in the band whose label starts at 37,951
    old = engine.band_labels((raw_edge,), min(vals), max(vals))
    assert old == labels and engine.band_of(37950.99, (raw_edge,), old) == "37,951 - 49,151"
    assert not _reading_in(old[1], 37950.99)


def test_whole_dollars_every_loan_reads_inside_its_bands_label_on_a_run(tmp_path):
    import math
    from pocketbook import engine
    res, tbl = _dollar_run(tmp_path, {"name": "bal", "field": "ORIG_BAL", "count": 5, "cut": "equal_loans"})
    vals = [float(r["ORIG_BAL"]) for r in tbl.rows if str(r.get("ORIG_BAL") or "").strip()]
    assert not engine.all_whole(vals)                                          # the column carries cents
    raw = [engine._quantile(sorted(vals), i / 5) for i in range(1, 5)]
    edges = res.band_edges["bal"]
    assert edges == tuple(float(math.ceil(e)) for e in raw) and any(not e.is_integer() for e in raw)
    g = res.grids[0]
    ranged = [lab for lab in g.band_labels if " - " in lab]                     # not "(blank)"
    assert len(ranged) == 5
    held = {lab: g.cell(lab, engine.ALL).rows for lab in ranged}
    inside = {lab: sum(1 for v in vals if _reading_in(lab, v)) for lab in ranged}
    assert inside == held and sum(held.values()) == len(vals)
    for v in vals:
        assert sum(1 for lab in ranged if _reading_in(lab, v)) == 1, v
    assert all(f"{e:,.0f} - " in " ".join(ranged) for e in edges)              # each band starts at its edge
    # and a loan planted a cent under each raised edge, where the build before cut, reads inside its label
    plant = {i * 997: math.ceil(e) - 0.01 for i, e in enumerate(raw, 1)}
    res, tbl = _dollar_run(tmp_path / "plant", {"name": "bal", "field": "ORIG_BAL", "count": 5, "cut": "equal_loans"},
                           bal=lambda i, was: plant.get(i, was))
    g = res.grids[0]
    got = [float(r["ORIG_BAL"]) for r in tbl.rows if str(r.get("ORIG_BAL") or "").strip()]
    assert all(v in got for v in plant.values())
    ranged = [lab for lab in g.band_labels if " - " in lab]
    held = {lab: g.cell(lab, engine.ALL).rows for lab in ranged}
    assert held == {lab: sum(1 for v in got if _reading_in(lab, v)) for lab in ranged}


def test_whole_dollars_typed_edges_with_cents_stay_exactly_as_typed(tmp_path):
    typed = [26324.2, 37950.548, 49151.3]
    res, _ = _dollar_run(tmp_path, {"name": "bal", "field": "ORIG_BAL", "edges": typed})
    assert res.band_edges["bal"] == tuple(typed)


def test_whole_dollars_leave_a_whole_number_column_and_a_ratio_as_they_were():
    from pocketbook import engine
    fico = [float(x) for x in range(600, 852)]                                  # whole scores, cut at 650.2 ...
    raw = tuple(engine._quantile(sorted(fico), i / 5) for i in range(1, 5))
    assert any(not e.is_integer() for e in raw) and engine.cut_edges(fico, 5, "equal_loans") == raw
    ratio = [0.01 + (i * 0.7919) % 0.98 for i in range(997)]                   # under 100: read to its decimals
    raw = tuple(engine._quantile(sorted(ratio), i / 5) for i in range(1, 5))
    assert engine.cut_edges(ratio, 5, "equal_loans") == raw
    small = [5.5 + (i * 7.31) % 90 for i in range(997)]                         # dollars under 100: not whole units
    assert engine.cut_edges(small, 4, "equal_loans") == tuple(engine._quantile(sorted(small), i / 4) for i in (1, 2, 3))


def test_whole_dollars_raising_an_edge_never_leaves_two_the_same_or_an_empty_band():
    from pocketbook import engine
    vals = [100.2, 100.4, 100.6, 100.8, 500.5]
    assert engine.whole_cut((100.3, 100.5, 100.7), vals) == (101.0,)            # three raised onto one: kept once
    assert engine.whole_cut((100.3, 500.1), vals) == (101.0,)                   # past the largest value: no band
    assert engine.whole_cut((100.3, 250.5), vals) == (101.0, 251.0)
    got = engine.whole_cut((150.2, 150.9, 151.0, 300.4), [100.5, 160.1, 220.7, 400.3])
    assert got == (151.0, 301.0)


def test_whole_dollars_a_run_whose_edges_meet_says_it_got_fewer_bands(tmp_path):
    """Three equal-loan points between $100 and $101 are all raised to 101: two edges, and the Run's own warning.
    The column has five values; few_values is set under that so the cut, not one band per value (few_values), is
    what is tested here."""
    from pocketbook import engine
    res, tbl = _dollar_run(tmp_path, {"name": "bal", "field": "ORIG_BAL", "count": 5, "cut": "equal_loans",
                                      "few_values": 4},
                           n=500, bal=lambda i, was: (100.1, 100.2, 100.6, 100.7, 900.5)[i % 5])
    vals = sorted(float(r["ORIG_BAL"]) for r in tbl.rows)
    assert len({engine._quantile(vals, k / 5) for k in range(1, 5)}) == 4        # four edges before raising
    assert res.band_edges["bal"] == (101.0, 261.0)                              # three of them raised onto 101
    assert any("band bal: asked for 5 bands, got 3" in w for w in res.warnings), res.warnings


def test_whole_dollars_a_column_inside_one_dollar_keeps_its_cut_rather_than_losing_every_band():
    """Every loan between $100 and $101 reads 100: raising would leave no edge at all, and the Run would refuse the
    column as having nothing to cut. The cut is kept as it was."""
    from pocketbook import engine
    vals = [100.1, 100.2, 100.3, 100.6, 100.7, 100.8]
    assert engine.cut_edges(vals, 2, "equal_loans") == (engine._quantile(vals, 0.5),)


def test_whole_dollars_scoutings_suggested_bins_on_a_dollar_column_are_whole():
    import numpy as np
    from pocketbook import scout
    v = np.array([37000.25 + i * 0.37 for i in range(2000)] + [float("nan")] * 5)
    assert scout.bins_at([0], [(37300.55, 1.0)], [37300.2, 37300.9], v, [0.1, 0.3]) == [37301.0]
    ratio = np.array([0.1 + i * 0.0004 for i in range(2000)])
    assert scout.bins_at([0], [(0.4555, 1.0)], [0.455, 0.456], ratio, [0.1, 0.3]) == [
        scout.edge([(0.4555, 1.0)], 0.455, 0.456)] == [0.456]                   # a ratio: as scouting rounds it


# ---- RANR vs GCOs: gross booked, GCO and RANR (30 Sep 2026)
# The firm: "On the paid cost kept tab I would like to work on gross GCO gross booked and gross RANR as well so we can
# also see if pockets are straight negative on returns". Each pocket's own booked dollars, charge-offs and RANR, and
# RANR per booked dollar, compared with nothing; a pocket whose RANR is below zero in red and counted above the table;
# totals under it that add up to the whole book. Every figure is checked here against the loan file, read without the
# engine, for every grid the dropdown offers.

KIOSK_LOSS = 0.04          # the planted pocket keeps -4% of what it booked, before its charge-offs


def _gross_file(tmp_path, n=6000):
    """The synthetic book, and 60 loans through a Kiosk spread over every score band, each losing money outright."""
    src = synth.write_extract(tmp_path / "src", n=n)
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    for k in range(60):
        bal = 10000.0 + 250 * k
        gco = round(bal * 0.3, 2) if k % 5 == 0 else 0.0
        rows.append({**rows[10], "LOAN_NBR": f"K{k:03d}", "FICO": 560 + (k * 7) % 240, "CHANNEL": KIOSK,
                     "ORIG_BAL": bal, "BAD_FLAG": 1 if gco else 0, "GCO_AMT": gco,
                     "RANR_AMT": round(-KIOSK_LOSS * bal - gco, 2)})
    out = tmp_path / "gross.csv"
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    return out


@pytest.fixture(scope="module")
def gross_book(tmp_path_factory):
    """FICO x CHANNEL and FICO x ASSET_CLASS over the Kiosk book, each grid picked in turn and calculated."""
    import tabs
    from pocketbook import perm, results
    d = tmp_path_factory.mktemp("gross")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 200)
        x = _gross_file(d)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO",),
                                                segments=("CHANNEL", "ASSET_CLASS")))
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    grids = tabs.options(load_workbook(out.book), results.PCK, "Grid")
    calc = {g: tabs.calculated(tabs.choose(out.book, d / f"grid{i}.xlsx", results.PCK, grid=g), results.PCK)
            for i, g in enumerate(grids)}
    return {"b": out.book, "x": x, "calc": calc}


def _gross_sums(rows) -> dict:
    """Loans, booked, GCO and RANR over loans, read from the file as the Run reads a dollar rate: a loan counts
    towards one when its amount and its booked amount both read as numbers ("#N/A" and a blank don't). Booked is
    under RANR, so RANR rate is RANR over it."""
    def s(top):
        return math.fsum(_num(r[top]) for r in rows if _num(r[top]) is not None and _num(r["ORIG_BAL"]) is not None)
    booked = math.fsum(_num(r["ORIG_BAL"]) for r in rows if _num(r["ORIG_BAL"]) is not None
                       and _num(r["RANR_AMT"]) is not None)
    ranr = s("RANR_AMT")
    return {"loans": len(rows), "booked": booked, "gco": s("GCO_AMT"), "ranr": ranr,
            "ranr_rate": ranr / booked if booked else None}


def _gross_by_pocket(x, grid: str, labels) -> dict:
    """{(band, segment): sums} for one grid, and ("All", "All") for the book, from the loan file alone."""
    seg = grid.split(" x ")[1]
    cells: dict = {}
    for r in _loans_file(x):
        for k in ((_band(r["FICO"], labels), r[seg]), ("All", "All")):
            cells.setdefault(k, []).append(r)
    return {k: _gross_sums(v) for k, v in cells.items()}


def _pck_labels(rows) -> set:
    return {str(x["band"]) for x in rows if " - " in str(x["band"])}


def _pck_totals(ws) -> dict:
    """The totals under RANR vs GCOs' table, {label: {row, loans, booked, gco, ranr, ranr_rate}}."""
    from pocketbook import results as rs
    out = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=rs.C_BAND).value
        if label in (rs.PCK_LISTED, rs.PCK_UNLISTED, rs.PCK_BOOK):
            out[label] = {"row": r, "loans": ws.cell(row=r, column=rs.C_LOANS).value,
                          **{k: ws.cell(row=r, column=c).value
                             for k, c in zip(("booked", "gco", "ranr", "ranr_rate"), rs.GROSS)}}
    return out


def test_pck_gross_every_pockets_booked_gco_ranr_and_rate_are_the_loan_files_for_every_grid(gross_book):
    import tabs
    from pocketbook import results
    assert set(gross_book["calc"]) == {"FICO x CHANNEL", "FICO x ASSET_CLASS"}
    for grid, ws in gross_book["calc"].items():
        rows = tabs.pck(ws)
        head = tabs.header_row(ws, results.C_TOG, "Together")
        assert tabs.heads(ws, head, results.C_LOANS, results.C_RATE) == ["Loans", "Booked", "GCOs", "RANR",
                                                                         "RANR ÷ Booked"]
        assert ws.cell(row=head - 1, column=results.C_BOOK).value == "Gross · this pocket alone"
        want = _gross_by_pocket(gross_book["x"], grid, _pck_labels(rows))
        assert len(rows) > 10, grid
        for x in rows:
            w = want[(str(x["band"]), str(x["seg"]))]
            assert x["loans"] == w["loans"], (grid, x)
            for k in ("booked", "gco", "ranr"):
                assert x[k] == pytest.approx(w[k], abs=0.005), (grid, x["band"], x["seg"], k)
            assert x["ranr_rate"] == pytest.approx(w["ranr_rate"], abs=1e-12), (grid, x)


def test_pck_gross_totals_add_up_to_the_whole_book_on_every_grid(gross_book):
    import tabs
    from pocketbook import results
    for grid, ws in gross_book["calc"].items():
        rows, tot = tabs.pck(ws), _pck_totals(ws)
        whole, want = tot[results.PCK_BOOK], _gross_by_pocket(gross_book["x"], grid, _pck_labels(rows))[("All", "All")]
        for k in ("loans", "booked", "gco", "ranr"):
            assert whole[k] == pytest.approx(want[k], abs=0.01), (grid, k)
            listed = tot[results.PCK_LISTED][k]
            assert listed == pytest.approx(math.fsum(x[k] for x in rows), abs=0.01), (grid, k)
            unlisted = tot.get(results.PCK_UNLISTED, {}).get(k) or 0
            assert listed + unlisted == pytest.approx(whole[k], abs=0.01), (grid, k)
        assert whole["ranr_rate"] == pytest.approx(want["ranr_rate"], abs=1e-12)
        assert tot[results.PCK_LISTED]["row"] > rows[-1]["row"] + 1        # under the table, a row clear of it


NEG_RULE = re.compile(r'^(?:AND\(\$[A-Z]+\d+<>"",)?AND\(ISNUMBER\(\$([A-Z]+)(\d+)\),\$[A-Z]+\d+<0\)\)?$')


def _red(ws, r: int, c: int) -> bool:
    """Whether a red-text rule holds on a cell: the RANR it reads below zero, its reference moved to the row as a
    spreadsheet moves it."""
    from openpyxl.utils import range_boundaries
    from openpyxl.utils.cell import column_index_from_string
    from pocketbook import house
    for rng in ws.formulas.conditional_formatting:
        for bounds in str(rng.sqref).split():
            c0, r0, c1, r1 = range_boundaries(bounds)
            if not (c0 <= c <= c1 and r0 <= r <= r1):
                continue
            for rule in sorted(rng.rules, key=lambda x: x.priority):
                font = rule.dxf.font if rule.dxf is not None else None
                if font is None or font.color is None or font.color.rgb[-6:] != house.CRIMSON:
                    continue
                m = NEG_RULE.match(rule.formula[0])
                assert m, rule.formula[0]
                v = ws.cell(row=int(m.group(2)) + r - r0, column=column_index_from_string(m.group(1))).value
                if isinstance(v, (int, float)) and v < 0:
                    return True
    return False


def test_pck_gross_a_pocket_that_lost_money_outright_is_red_and_counted_above_the_table(gross_book):
    import tabs
    from pocketbook import house, results
    kiosks = 0
    for grid, ws in gross_book["calc"].items():
        rows = tabs.pck(ws)
        want = _gross_by_pocket(gross_book["x"], grid, _pck_labels(rows))
        neg = [x for x in rows if want[(str(x["band"]), str(x["seg"]))]["ranr"] < 0]
        # the Kiosk's 60 loans sit in one asset class, among thousands that make money: that grid has none
        assert bool(neg) is (grid == "FICO x CHANNEL"), grid
        kiosks += sum(x["seg"] == KIOSK for x in neg)
        for x in rows:
            is_neg = x in neg
            assert _red(ws, x["row"], results.C_RANR) is is_neg, (grid, x)
            assert _red(ws, x["row"], results.C_RATE) is is_neg, (grid, x)
            assert not _red(ws, x["row"], results.C_BOOK) and not _red(ws, x["row"], results.C_GCO)
        # the count line beside the Grid dropdown, from the loan file: how many, and how much between them
        lost = -math.fsum(want[(str(x["band"]), str(x["seg"]))]["ranr"] for x in neg)
        s = tabs.dropdown(ws, "Grid").row
        n = len(neg)
        assert ws.cell(row=s, column=results.C_LOANS).value == (
            f"{n} pocket{'s' if n != 1 else ''} lost money outright, totalling ${lost:,.0f}: RANR below zero, in red."
            if n else "No pocket listed lost money outright: every one's RANR is zero or more.")
        rules = [r for rng in ws.formulas.conditional_formatting
                 if str(rng.sqref) == f"{results.col(results.C_LOANS)}{s}" for r in rng.rules]
        assert rules and rules[0].dxf.font.color.rgb[-6:] == house.CRIMSON
        assert ws.cell(row=s, column=results.C_H_UN).value == n                # the rule's cell: red above 0
    assert kiosks >= 5                          # the planted Kiosk pockets, one in each score band
    note = _note(gross_book["calc"]["FICO x CHANNEL"])
    assert "Negative RANR: the pocket lost money outright, before comparing it with anyone." in note["RANR"]


def test_pck_gross_a_pocket_not_listed_gets_its_own_total_and_the_three_still_add_up():
    """A pocket with nothing to compare it with is not listed; its dollars go to Not listed, never lost from the
    book. Neither synthetic grid has one, so the rows are stated here."""
    from pocketbook import engine, results

    def cell(n, booked, gco, ranr):
        return engine.Cell(rows=n, rates={"ranr_rate": engine.RateStat(num=ranr, den=booked),
                                          "gco_rate": engine.RateStat(num=gco, den=booked)})

    class G:
        cells = {("a", "x"): cell(3, 300.0, 10.0, 20.0), ("a", "y"): cell(2, 200.0, 50.0, -30.0),
                 (engine.ALL, engine.ALL): cell(5, 500.0, 60.0, -10.0)}

        def inner(self):
            return [(k, c) for k, c in self.cells.items() if engine.ALL not in k]
    got = results.pck_totals(G(), [{"band": "a", "seg": "x"}])
    assert got == [(results.PCK_LISTED, 3, 300.0, 10.0, 20.0, 20.0 / 300), (results.PCK_UNLISTED, 2, 200.0, 50.0,
                                                                            -30.0, -0.15),
                   (results.PCK_BOOK, 5, 500.0, 60.0, -10.0, -0.02)]
    assert [t[0] for t in results.pck_totals(G(), [{"band": "a", "seg": "x"}, {"band": "a", "seg": "y"}])] == [
        results.PCK_LISTED, results.PCK_BOOK]


def test_pck_gross_no_negative_pocket_says_so_and_one_says_it_in_the_singular():
    from pocketbook import results
    assert results.negative_line([{"ranr": 5.0}, {"ranr": 0.0}]) == (
        "No pocket listed lost money outright: every one's RANR is zero or more.")
    assert results.negative_line([{"ranr": -1234.4}, {"ranr": 3.0}]) == (
        "1 pocket lost money outright, totalling $1,234: RANR below zero, in red.")


def test_pck_gross_dollars_fit_their_columns_and_turn_to_thousands_past_the_cap(gross_book):
    from pocketbook import results
    ws = load_workbook(gross_book["b"])[results.PCK]
    rows = [v for k, v in _views_rows(gross_book["b"]).items() if k.startswith("C|") and k.split("|")[-1].isdigit()]
    for j, c in enumerate(results.GROSS):
        vals = [v[7 + j] for v in rows if len(v) > 7 + j and v[7 + j] is not None]
        assert vals, c
        longest = max(len(f"{v:.2%}") if j == 3 else len(f"${abs(v):,.0f}") + (v < 0) for v in vals)
        assert ws.column_dimensions[results.col(c)].width >= longest + 2, c
    assert results.gross_widths([(1.9e8, 8e6, -2.4e7, 0.12)])[0] is False
    k, w = results.gross_widths([(2.5e9, 8e7, -2.4e8, -0.12)])
    assert k is True and results.gross_shown(-2.4e8, True) == "-$240,000k"
    assert w[results.C_BOOK] >= len("$2,500,000k") + 2



# ---- 30 Sep 2026: a Summary tab, one band column and the book's plain figures across
# The firm: "I want to add some easy high value views as well. Like a few matrices where it lists out a chosen band on
# the left and shows real calculated metrics... Maybe I want to see bands of FICO on the left and straight up unit
# counts, loan amounts, % of units, % of loan amounts, charged off dollars, ratio, percentage of units. Same with RANR.
# They'd be across the top." The ratio: "Charged off / booked". Bad loans columns: "Yes do this".

SUMMARY_ONLY = f"Only loans where {ch.ORIG_YEAR} is"
SUMMARY_HEADS = ["Loans", "% of loans", "Bad loans", "Bad loans %", "Booked $", "% of booked", "GCOs ($)",
                 "GCOs ÷ Booked", "× book", "% of GCOs", "RANR $", "RANR ÷ Booked", "% of RANR"]
SHARES = ("% of loans", "% of booked", "% of GCOs", "% of RANR")
SPECIAL = ("(blank)", "(not a number)", "(marked missing)")


@pytest.fixture(scope="module")
def summary_book(tmp_path_factory):
    """FICO and REV_DEBT (dollars) cut into bands, by CHANNEL, filtered by ORIG_YEAR: the book and its loan file."""
    from pocketbook import perm
    d = tmp_path_factory.mktemp("summary")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 100)
        x = synth.write_extract(d / "src", n=4000)
        out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=("FICO", "REV_DEBT"), segments=("CHANNEL",),
                                                filter=ch.ORIG_YEAR, outcome="BAD_FLAG"))
        _answer(out.book)
        ran = book.run(out.book)
        assert ran.ok, ran.lines
    return out.book, x


def _summary_tab(b, out, band=None, only=None):
    """Summary with its dropdowns set, calculated: the sheet, the headings across, {row label: {heading: value}} in
    the order shown, and the label column's heading."""
    from pocketbook import results
    import tabs
    picks = {k: v for k, v in (("band or category column", band), (SUMMARY_ONLY, only)) if v}
    ws = tabs.calculated(tabs.choose(b, out, results.SUMMARY, **picks), results.SUMMARY)
    h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
    heads = []
    c = 3
    while ws.cell(row=h, column=c).value not in (None, ""):
        heads.append(ws.cell(row=h, column=c).value)
        c += 1
    rows = {}
    r = h + 1
    while ws.cell(row=r, column=2).value not in (None, ""):
        rows[ws.cell(row=r, column=2).value] = {hd: ws.cell(row=r, column=3 + j).value for j, hd in enumerate(heads)}
        r += 1
    return ws, heads, rows, ws.cell(row=h, column=2).value


def _summary_label(v, labels) -> str:
    """A loan's row on Summary, worked out here from the tab's own band labels ("496 - 653", read in whole units)
    and the extract's text, not by the engine: a blank is (blank), text is (not a number), a FICO under -1000 was
    answered missing."""
    if v in ("", None):
        return "(blank)"
    try:
        x = float(v)
    except ValueError:
        return "(not a number)"
    if x < -1000:
        return "(marked missing)"
    got = [lab for lab in labels if _reading_in(lab, x)]
    assert len(got) == 1, (v, got)
    return got[0]


def _road2(rows, field, labels, whole_rows, label_of=None) -> dict:
    """Every Summary figure for `rows`, from the CSV's text alone: the bands, the special rows, then All. `label_of`
    reads a loan's row from its text when the bands aren't ranges (one band per value)."""
    label_of = label_of or _summary_label
    def f(v):
        try:
            return float(v)
        except (TypeError, ValueError):
            return None

    def figures(rs):
        read = [r for r in rs if r["BAD_FLAG"] in ("0", "1")]
        bad = sum(1 for r in read if r["BAD_FLAG"] == "1")
        booked = sum(f(r["ORIG_BAL"]) for r in rs if f(r["ORIG_BAL"]) is not None)
        g = [(f(r["GCO_AMT"]), f(r["ORIG_BAL"])) for r in rs]
        g = [(a, c) for a, c in g if a is not None and c is not None]
        k = [(f(r["RANR_AMT"]), f(r["ORIG_BAL"])) for r in rs]
        k = [(a, c) for a, c in k if a is not None and c is not None]
        gco, gden = sum(a for a, _ in g), sum(c for _, c in g)
        ranr, rden = sum(a for a, _ in k), sum(c for _, c in k)
        return {"Loans": len(rs), "Bad loans": bad, "Bad loans %": bad / len(read) if read else None,
                "Booked $": booked, "GCOs ($)": gco, "GCOs ÷ Booked": gco / gden if gden else None,
                "RANR $": ranr, "RANR ÷ Booked": ranr / rden if rden else None}

    by: dict = {}
    for r in rows:
        by.setdefault(label_of(r[field], labels), []).append(r)
    order = list(labels) + [s for s in SPECIAL if s in by]
    out = {lab: figures(by.get(lab, [])) for lab in order}
    out["All"] = figures(rows)
    book_rate = figures(whole_rows)["GCOs ÷ Booked"]
    top = out["All"]
    for x in out.values():
        x["% of loans"] = x["Loans"] / top["Loans"]
        x["% of booked"] = x["Booked $"] / top["Booked $"]
        x["% of GCOs"] = x["GCOs ($)"] / top["GCOs ($)"]
        x["% of RANR"] = x["RANR $"] / top["RANR $"]
        x["× book"] = x["GCOs ÷ Booked"] / book_rate if x["GCOs ÷ Booked"] is not None else None
    return out


def _same(got, want) -> bool:
    if want is None:
        return got in (None, "")
    return got == pytest.approx(want, rel=1e-9, abs=1e-9)


def _check_summary(shown, want) -> None:
    assert list(shown) == list(want)                                 # bands in order, the special rows, then All
    for lab, cols in want.items():
        for hd, v in cols.items():
            assert _same(shown[lab][hd], v), (lab, hd, shown[lab][hd], v)


def test_summary_every_cell_of_fico_and_a_dollar_band_column_is_the_loans_worked_out_again(summary_book, tmp_path):
    """FICO, then REV_DEBT picked in the dropdown: every cell against the loan file, the shares adding to 100% over
    the bands and the special rows, and All the book's own totals."""
    from pocketbook import config as cfgmod, engine, house
    from pocketbook.ingest import read_table
    b, x = summary_book
    rows = _loans_file(x)
    seen = {}
    for field in ("FICO", "REV_DEBT"):
        ws, heads, shown, head = _summary_tab(b, tmp_path / f"{field}.xlsx", band=field)
        assert head == field and heads == SUMMARY_HEADS                  # the dropdown switched the column
        labels = [lab for lab in shown if lab != "All" and lab not in SPECIAL]
        assert len(labels) == 5 and all(" - " in lab for lab in labels)
        _check_summary(shown, _road2(rows, field, labels, rows))
        for hd in SHARES:
            assert sum(v[hd] for lab, v in shown.items() if lab != "All") == pytest.approx(1.0, rel=1e-9)
        top = shown["All"]
        assert top["Loans"] == len(rows) and top["× book"] == pytest.approx(1.0)
        assert all(top[hd] == pytest.approx(1.0) for hd in SHARES)
        seen[field] = shown
    assert "(marked missing)" in seen["FICO"] and "(blank)" in seen["FICO"]
    assert not set(SPECIAL) & set(seen["REV_DEBT"])
    assert seen["FICO"]["All"] == pytest.approx(seen["REV_DEBT"]["All"])       # one book, however it is cut
    # the book's totals, as the engine keeps them for every other tab
    res = engine.run(cfgmod.parse(book.read_book(b)[0]), read_table(x))
    t = res.total.rates
    top = seen["FICO"]["All"]
    assert top["Bad loans"] == t["outcome_loans"].num and top["Bad loans %"] == pytest.approx(t["outcome_loans"].rate)
    assert top["GCOs ($)"] == pytest.approx(t["gco_rate"].num)
    assert top["RANR $"] == pytest.approx(t["ranr_rate"].num)
    assert top["Booked $"] == pytest.approx(res.book_size.booked)
    # light neutral shading at most: the All row, CANVAS; nothing red or green
    fills = {r.dxf.fill.fgColor.rgb[-6:] for rng in ws.formulas.conditional_formatting for r in rng.rules
             if r.dxf is not None and r.dxf.fill is not None}
    assert fills == {house.CANVAS}


def test_summary_only_loans_where_a_year_is_that_years_loans_against_the_whole_book(summary_book, tmp_path):
    from pocketbook import results
    import tabs
    b, x = summary_book
    rows = _loans_file(x)
    years = sorted({r["ORIG_DATE"][:4] for r in rows})
    assert tabs.options(load_workbook(b), results.SUMMARY, SUMMARY_ONLY) == [results.ALL_LOANS] + years
    _, _, whole, _ = _summary_tab(b, tmp_path / "all.xlsx", band="FICO")
    labels = [lab for lab in whole if lab != "All" and lab not in SPECIAL]
    for year in (years[0], years[len(years) // 2]):
        mine = [r for r in rows if r["ORIG_DATE"][:4] == year]
        assert 0 < len(mine) < len(rows)
        _, heads, shown, head = _summary_tab(b, tmp_path / f"y{year}.xlsx", band="FICO", only=year)
        assert head == "FICO" and heads == SUMMARY_HEADS
        want = _road2(mine, "FICO", labels, rows)
        # the whole book's rows stay put: a special row this year has no loans in shows none
        for s in SPECIAL:
            if s in whole and s not in want:
                assert shown[s]["Loans"] == 0 and shown[s]["Bad loans %"] in (None, "")
                del shown[s]
        _check_summary(shown, want)
        assert shown["All"]["Loans"] == len(mine) and shown["All"]["× book"] != pytest.approx(1.0)
        for hd in SHARES:
            assert sum(v[hd] for lab, v in shown.items() if lab != "All") == pytest.approx(1.0, rel=1e-9)


def test_summary_leaves_off_the_columns_a_run_has_no_source_for_and_says_so(tmp_path):
    import dataclasses
    from openpyxl import Workbook
    from pocketbook import config as cfgmod, engine, results
    from pocketbook.ingest import read_table
    import tabs
    cfg, data = synth.write(tmp_path / "cube", n=2000)
    c = cfgmod.load(cfg)
    table = read_table(data)
    bare = dataclasses.replace(c, booked="", measures=tuple(m for m in c.measures
                                                          if m.name not in ("gco_rate", "contribution_rate")))
    for conf, want in ((c, SUMMARY_HEADS), (bare, ["Loans", "% of loans", "Bad loans", "Bad loans %", "RANR $",
                                                   "RANR ÷ Booked", "% of RANR"])):
        res = engine.run(conf, table)
        wb = Workbook()
        results.write_summary(wb, res, results.Choices(wb), results.Views(wb))
        ws = wb[results.SUMMARY]
        h = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=3).value == "Loans")
        assert [ws.cell(row=h, column=j).value for j in range(3, 3 + len(want) + 1)] == want + [None]
        note = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=3).value for r in range(3, h)}
        if conf is bare:
            assert note["Not shown"] == ("This Run has no booked amount and no GCO dollars, so those columns "
                                         "are left off.")
            assert "GCOs ($)" not in note and "Booked $" not in note
        else:
            assert "Not shown" not in note and "Only loans where" not in note     # no Filter by, no second dropdown
            with pytest.raises(KeyError):
                tabs.dropdown(ws, "Only loans where")


# ---- A number column too few-valued to cut: one band per value (30 Sep 2026, at the bank)
# The firm: "So it refuses to run some stuff because it cannot band. Which makes sense for the examples so far - they
# are things like major derogs which do not include too many numbers." And to the fix: "Yes that's fine". A column
# like Major Derogatories (0 to 8, most loans at 0) marked Amount or number, cut into equal-loan bands, had every cut
# fall on the zeros, and the Run refused it: "the extract has no column ... (a band: no readable numbers to cut)".
# Now: with Control's few values (12) or fewer, each value is its own band, named by the value; with more, it is cut
# as far as it can be; one value is refused in words naming the two fixes. Typed Band edges always win.

DEROGS = "Major Derogatories"


def _derog(i: int, rng) -> int:
    return 0 if rng.random() < 0.85 else rng.randint(1, 8)


def _with_derogs(src, out, value=_derog, seed=30):
    """The loan file at `src` with a Major Derogatories column added: by default 85% zeros, the rest 1 to 8."""
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    rng = random.Random(seed)
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(list(rows[0]) + [DEROGS])
        for i, r in enumerate(rows):
            w.writerow(list(r.values()) + [value(i, rng)])
    return out


def _derog_run(tmp_path, band: dict, value=_derog, n=3000):
    """The engine on the synthetic cube cut by one band on Major Derogatories: the result and the loan file's rows."""
    import copy
    from pocketbook import config as cfgmod, engine
    from pocketbook.ingest import read_table
    cfg, data = synth.write(tmp_path, n=n)
    x = _with_derogs(data, tmp_path / "derogs.csv", value)
    raw = copy.deepcopy(cfgmod.load(cfg).raw)
    raw["bands"] = [{"name": "derogs", "field": DEROGS, **band}]
    return engine.run(cfgmod.parse(raw), read_table(x)), _loans_file(x)


EQUAL = {"count": 5, "cut": "equal_loans"}


def _value_of(v, labels) -> str:
    """A loan's row on a tab cut one band per value, from its text alone: the whole number itself."""
    if v in ("", None):
        return "(blank)"
    return str(int(float(v)))


def _derogs_workbook(d, value=_derog, n=3000, filt=ch.ORIG_YEAR, few=12):
    """Set up on the loan file with Major Derogatories, marked Amount or number on Columns, and answered. `few`: the
    launcher's "Number columns: this many values or fewer is a category"."""
    from test_book import at
    x = _with_derogs(synth.write_extract(d / "src", n=n), d / "derogs.csv", value)
    out = book.set_up(x, choices=ch.Choices(run_kind=ch.BLEED, bands=(DEROGS,), segments=("CHANNEL",), filter=filt,
                                            outcome="BAD_FLAG", few_values=few))
    wb = load_workbook(out.book)
    wb["Columns"][at(wb, DEROGS, book.C_MEANS)] = "Amount or number"
    wb.save(out.book)
    _answer(out.book)
    return out.book, x


@pytest.fixture(scope="module")
def derogs_book(tmp_path_factory):
    """The bank's shape: Major Derogatories marked Amount or number and cut into bands, by CHANNEL, filtered by
    ORIG_YEAR. The book, its loan file, and the Run."""
    from pocketbook import perm
    d = tmp_path_factory.mktemp("derogs")
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("POCKETBOOK_MEMORY", str(d / "memory.yaml"))
        mp.setattr(perm, "SHUFFLES", 100)
        b, x = _derogs_workbook(d)
        ran = book.run(b)
    return b, x, ran


def test_few_values_the_bank_column_was_refused_before_and_runs_now_one_band_per_value(tmp_path):
    from pocketbook import engine
    res, rows = _derog_run(tmp_path, EQUAL)
    vals = sorted(float(r[DEROGS]) for r in rows)
    # the reproduction: every equal-loan cut falls on the zeros, which is what the Run refused
    assert {engine._quantile(vals, k / 5) for k in range(1, 5)} == {0.0}
    assert engine.cut_edges(vals, 5, "equal_loans") == ()
    assert res.band_edges["derogs"] == (1.0, 2.0, 3.0, 4.0, 5.0, 6.0, 7.0, 8.0)
    assert res.value_bands["derogs"] == tuple(float(v) for v in range(9))
    g = res.grids[0]
    assert g.band_labels == [str(v) for v in range(9)]                     # "0", "1", ... never "0.0" or "0 - 0"
    held = {lab: g.cell(lab, engine.ALL).rows for lab in g.band_labels}
    assert held == {str(v): sum(1 for r in rows if int(r[DEROGS]) == v) for v in range(9)}
    assert f"{DEROGS}: too few values to cut into equal bands, so each value is its own band" in res.warnings
    assert not [w for w in res.warnings if "asked for" in w]


def test_few_values_the_run_says_so_and_columns_suggests_category_while_the_answer_stays(derogs_book):
    import tabs
    from test_book import at
    b, x, ran = derogs_book
    said = f"{DEROGS}: too few values to cut into equal bands, so each value is its own band."
    assert ran.ok, ran.lines
    assert said in ran.lines                                              # the Run's own lines
    rec = tabs.record(b)
    warn = rec["Warning"] if isinstance(rec["Warning"], list) else [rec["Warning"]]
    assert said[:-1] in warn                                              # and Record
    wb = load_workbook(b)
    why = wb["Columns"][at(wb, DEROGS, book.C_WHY)].value
    assert why.endswith(" Few values (0 to 8): Category may read better.") and why.count("Few values") == 1
    assert wb["Columns"][at(wb, DEROGS, book.C_MEANS)].value == "Amount or number"   # a suggestion, never a change
    assert "Few values" not in str(wb["Columns"][at(wb, "FICO", book.C_WHY)].value)
    # the setting the launcher chose rides with the band: Control's 12
    raw = book.read_book(b)[0]
    assert [x.get("few_values") for x in raw["bands"]] == [12]


def test_few_values_grids_every_value_band_is_the_loan_files(derogs_book, tmp_path):
    from pocketbook import results
    import tabs
    b, x, _ = derogs_book
    rows = _loans_file(x)
    grid = next(g for g in tabs.options(load_workbook(b), results.GRIDS, "Grid") if DEROGS in g)
    for only in (None, "2023"):
        picks = {"grid": grid, **({ONLY_YEAR: only} if only else {})}
        ws, blocks, _ = _grids(b, tmp_path / f"g{only}.xlsx", **picks)
        loans, rate = blocks["Loans"], blocks["Rate"]
        mine = [r for r in rows if only is None or _year(r) == only]
        assert 0 < len(mine) < len(rows) or only is None
        want: dict = {}
        for r in mine:
            v = str(int(r[DEROGS]))
            for k in {(v, r["CHANNEL"]), (v, "All"), ("All", r["CHANNEL"]), ("All", "All")}:
                want.setdefault(k, []).append(r)
        assert [k[0] for k in loans if k[1] == "All"] == [str(v) for v in range(9)] + ["All"]
        assert {k: n for k, n in loans.items() if n} == {k: len(v) for k, v in want.items()}, only
        for k, got in want.items():
            assert rate[k] == pytest.approx(_bad_rate(got)), (only, k)


def test_few_values_summary_every_value_band_is_the_loan_files(derogs_book, tmp_path):
    b, x, _ = derogs_book
    rows = _loans_file(x)
    for only in (None, "2023"):
        ws, heads, shown, head = _summary_tab(b, tmp_path / f"s{only}.xlsx", band=DEROGS, only=only)
        assert head == DEROGS and heads == SUMMARY_HEADS
        labels = [lab for lab in shown if lab != "All"]
        assert labels == [str(v) for v in range(9)]
        mine = [r for r in rows if only is None or _year(r) == only]
        _check_summary(shown, _road2(mine, DEROGS, labels, rows, label_of=_value_of))


def test_few_values_typed_edges_always_win(tmp_path):
    res, rows = _derog_run(tmp_path, {"edges": [1, 2, 5]})
    assert res.band_edges["derogs"] == (1.0, 2.0, 5.0) and not res.value_bands
    assert res.grids[0].band_labels == ["0 - 0", "1 - 1", "2 - 4", "5 - 8"]
    assert not [w for w in res.warnings if "its own band" in w]


def test_few_values_typed_edges_on_columns_win_and_take_the_suggestion_off(derogs_book, tmp_path, monkeypatch):
    import shutil
    from pocketbook import perm
    from test_book import at
    b, x, _ = derogs_book
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    mine = tmp_path / b.name
    shutil.copy(b, mine)
    wb = load_workbook(mine)
    wb["Columns"][at(wb, DEROGS, book.C_EDGES)] = "1; 2; 5"
    wb.save(mine)
    ran = book.run(mine)
    assert ran.ok, ran.lines
    assert not [ln for ln in ran.lines if "its own band" in ln]
    wb = load_workbook(mine)
    assert "Few values" not in str(wb["Columns"][at(wb, DEROGS, book.C_WHY)].value)


def test_few_values_a_single_value_is_refused_in_words_that_name_both_fixes(tmp_path):
    from pocketbook import engine
    with pytest.raises(engine.DataRefused) as got:
        _derog_run(tmp_path, EQUAL, value=lambda i, rng: 0)
    assert str(got.value) == (f"`{DEROGS}` reads 0 on every loan, so there is nothing to cut into bands. On Columns, "
                              f"set What it is to Category, or type Band edges like 1; 2; 5")


def test_few_values_a_single_value_column_in_the_workbook_says_the_two_fixes(tmp_path, monkeypatch):
    from pocketbook import perm
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    b, _ = _derogs_workbook(tmp_path, value=lambda i, rng: 3, n=1500, filt=None)
    ran = book.run(b)
    assert not ran.ok
    assert ran.lines == [f'Couldn\'t run: "{DEROGS}" reads 3 on every loan, so there is nothing to cut into bands. '
                         f'On Columns, set What it is to Category, or type Band edges like 1; 2; 5']


def test_few_values_more_values_than_the_setting_cut_what_they_can(tmp_path):
    """Twenty values, heavy at zero: more than 12, so no value bands. The equal-loan cuts that survive are kept and
    the Run says it got fewer bands; when none survive, the zeros against the rest."""
    from pocketbook import engine
    res, rows = _derog_run(tmp_path, EQUAL, value=lambda i, rng: 0 if rng.random() < 0.7 else rng.randint(1, 20))
    vals = sorted(float(r[DEROGS]) for r in rows)
    assert len(set(vals)) > engine.FEW_VALUES and not res.value_bands
    edges = res.band_edges["derogs"]
    assert edges == engine.cut_edges(vals, 5, "equal_loans") and 1 <= len(edges) < 4
    assert (f"band derogs: asked for 5 bands, got {len(edges) + 1} (`{DEROGS}` has too many repeated values to cut "
            f"finer)") in res.warnings
    res, rows = _derog_run(tmp_path / "b", EQUAL, value=lambda i, rng: 0 if rng.random() < 0.9 else rng.randint(1, 20))
    assert engine.cut_edges([float(r[DEROGS]) for r in rows], 5, "equal_loans") == ()
    assert res.band_edges["derogs"] == (1.0,) and not res.value_bands
    assert res.grids[0].band_labels == ["0 - 0", "1 - 20"]
    assert "band derogs: asked for 5 bands, got 2 (`Major Derogatories` has too many repeated values to cut finer)" \
        in res.warnings


def test_few_values_the_setting_is_the_threshold(tmp_path):
    """Nine values: one band each at 9 or more; at 6 (the setting, not 12) cut as far as they go."""
    res, _ = _derog_run(tmp_path, {**EQUAL, "few_values": 9})
    assert len(res.value_bands["derogs"]) == 9
    res, _ = _derog_run(tmp_path / "six", {**EQUAL, "few_values": 6})
    assert not res.value_bands and res.band_edges["derogs"] == (1.0,)


def test_few_values_the_launchers_setting_is_the_one_the_run_uses(tmp_path, monkeypatch):
    """The launcher's 6: nine values are more than that, so the Run cuts what it can and says it got fewer bands."""
    from pocketbook import perm
    monkeypatch.setattr(perm, "SHUFFLES", 100)
    monkeypatch.setenv("POCKETBOOK_MEMORY", str(tmp_path / "memory.yaml"))
    b, _ = _derogs_workbook(tmp_path, n=1500, filt=None, few=6)
    assert [x.get("few_values") for x in book.read_book(b)[0]["bands"]] == [6]
    ran = book.run(b)
    assert ran.ok, ran.lines
    assert not [ln for ln in ran.lines if "its own band" in ln]
    import tabs
    rec = tabs.record(b)
    warn = rec["Warning"] if isinstance(rec["Warning"], list) else [rec["Warning"]]
    assert any("asked for 5 bands, got 2" in w for w in warn), warn


def test_few_values_labels_are_the_values_themselves():
    from pocketbook import engine
    assert engine.labels_for((1.0, 2.0), [0.0, 1.0, 2.0], (0.0, 1.0, 2.0)) == ["0", "1", "2"]
    assert engine.labels_for((1.5,), [0.5, 1.5], (0.5, 1.5)) == ["0.5", "1.5"]
    assert engine.labels_for((2000.0,), [1000.0, 2000.0], (1000.0, 2000.0)) == ["1,000", "2,000"]
    # a column cut into ranges reads as before, whole-number and whole-dollar rules included
    assert engine.labels_for((620.0, 680.0), [500.0, 850.0]) == ["500 - 619", "620 - 679", "680 - 850"]
    assert engine.labels_for((654.2,), [496.0, 850.0]) == ["496 - 654", "655 - 850"]


def test_few_values_prevalence_counts_value_bands_as_the_grids_cut_them(tmp_path):
    from pocketbook import prevalence
    res, rows = _derog_run(tmp_path, EQUAL)
    assert prevalence._labels_by_band(res, res.table.rows)["derogs"] == [str(int(r[DEROGS])) for r in rows]
    # a subset keeps the same names whatever values it happens to hold (its smallest isn't 0)
    some = [r for r in res.table.rows if int(r[DEROGS]) >= 3]
    assert prevalence._labels_by_band(res, some)["derogs"] == [str(int(r[DEROGS])) for r in some]
    got = prevalence.count(res, res.grids[0], prevalence.Grouping("ASSET_CLASS", "values", "t"))
    assert got is not None                                                 # tied out to the grid's own pockets
    order, out, shown = got
    assert {k[0] for k in out} == set(res.grids[0].band_labels)
    assert prevalence._bands_of([(float(r[DEROGS]), None) for r in rows], res.band_edges["derogs"],
                                res.value_bands["derogs"])[1] == [str(v) for v in range(9)]
