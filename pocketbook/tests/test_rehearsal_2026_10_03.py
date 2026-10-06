"""Defects found when the public-data rehearsal was revived on current PocketBook (docs/rehearsal-public-data-2026-09.md,
section 7), each held here on a tiny synthetic book: no public row is read.

1. A band column that is there but holds no value PocketBook reads as a number was refused as missing. LendingClub's
   own LoanStats3a.csv writes revol_util with a % sign ("45.5%"); marked Amount and cut into bands, the Run said "the extract has
   no column "revol_util" (a band: no readable numbers to cut). Its columns are: ..., revol_util, ... If a column
   was renamed or dropped, press Set up again." The column was listed in the same sentence, and Set up again would
   not have helped. Now it says none of the values reads as a number, how many of each kind, and the two fixes.
"""

import copy
import csv

import pytest

from pocketbook import config as cfgmod, engine, synth
from pocketbook.ingest import read_table

UTIL = "Revolving Util"


def _with(src, out, value):
    rows = list(csv.DictReader(open(src, encoding="utf-8")))
    with open(out, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(list(rows[0]) + [UTIL])
        for i, r in enumerate(rows):
            w.writerow(list(r.values()) + [value(i)])
    return out


def _run(tmp_path, value, n=600):
    cfg, data = synth.write(tmp_path, n=n)
    x = _with(data, tmp_path / "util.csv", value)
    raw = copy.deepcopy(cfgmod.load(cfg).raw)
    raw["bands"] = [{"name": "util", "field": UTIL, "count": 5, "cut": "equal_loans"}]
    return engine.run(cfgmod.parse(raw), read_table(x))


def test_a_band_column_of_percent_text_is_refused_as_unreadable_not_as_missing(tmp_path):
    with pytest.raises(engine.DataRefused) as got:
        _run(tmp_path, lambda i: "" if i % 10 == 0 else f"{(i * 7) % 100}.5%")
    msg = str(got.value)
    assert not isinstance(got.value, engine.ColumnsMissing)
    assert "has no column" not in msg and "Set up again" not in msg
    assert msg.startswith(f"`{UTIL}` is cut into bands, but none of its 600 values reads as a number: ")
    assert "540 aren't numbers (a value with a % sign is not read as one), 60 are blank" in msg
    assert msg.endswith("Save it as plain numbers, or on Columns set What it is to Category")
    # a NothingToCut, so book.run and the CLI say it in words, without a traceback (both catch NothingToCut)
    assert isinstance(got.value, engine.NothingToCut)


def test_the_same_column_as_plain_numbers_is_cut(tmp_path):
    res = _run(tmp_path, lambda i: "" if i % 10 == 0 else f"{(i * 7) % 100}.5")
    assert res.grids


# 2. A p-value too small for a spreadsheet written into a formula. Scouting on SBAnational (4 Oct 2026) confirmed a
#    group on the held-back loans at p = 2.5e-315, a subnormal double. New variables writes each group's p-value into
#    its "What it found" formula as a literal; LibreOffice calculated that line to #VALUE!, so the tab's summary of
#    what held up on the holdout read as an error. Now a subnormal literal is written 0.0.

def test_a_subnormal_number_is_written_into_a_formula_as_zero():
    from pocketbook import live
    assert live.num(2.525564887e-315) == "0.0" and live.num(-3e-320) == "0.0"
    assert live.num(live.SMALLEST_NORMAL) == repr(live.SMALLEST_NORMAL)
    assert live.num(0.0) == "0.0" and live.num(0.25) == "0.25" and live.num(1e-300) == "1e-300"


def test_new_variables_line_with_a_subnormal_p_value_calculates(tmp_path):
    # the formula shape confirm_tab writes for "What it found", calculated through LibreOffice
    from openpyxl import Workbook
    from pocketbook import live
    from recalc import need_soffice, values_of
    need_soffice()
    wb = Workbook()
    ws = wb.active
    p = live.num(2.525564887e-315)
    ws["A1"] = f'=IF(AND(ISNUMBER({p}),{p}<0.05),"significant","not")'
    ws["A2"] = f'=IF(AND(ISNUMBER({live.num(0.3)}),{live.num(0.3)}<0.05),"significant","not")'
    got = values_of(wb, tmp_path)[ws.title]
    assert got["A1"].value == "significant" and got["A2"].value == "not"
