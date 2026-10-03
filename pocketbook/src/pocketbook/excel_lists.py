"""Open a workbook Excel has saved without losing its dropdowns.

Excel saves a dropdown whose list sits on another sheet (every answer on
Control, Meaning on Columns, the pickers on Grids, Pockets and RANR vs
GCOs) in its 2010 extension block rather than beside the others. openpyxl
reads that block, warns, and drops it, so the next save wrote the workbook
back with none of those dropdowns. Found at the bank, 29 Sep 2026: "I have
no drop downs in most places now", after an answer on Control was changed,
saved in Excel, and the workbook was run again.

`load` reads those dropdowns out of the file itself and puts them back on
the sheets openpyxl opened, so whatever is saved next still has them.
"""

from __future__ import annotations

import io
import posixpath
import re
import warnings
from contextlib import contextmanager
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
X14 = "http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"
XM = "http://schemas.microsoft.com/office/excel/2006/main"
DV_EXT = "{CCE6A557-97BC-4B89-ADB6-D9C93CAAB3DF}"      # Excel's name for the dropdowns' extension

# the attributes a dropdown carries, as Excel writes them and as openpyxl takes them
FLAGS = {"allowBlank": "allow_blank", "showErrorMessage": "showErrorMessage",
         "showInputMessage": "showInputMessage", "showDropDown": "showDropDown"}
TEXTS = ("type", "operator", "errorStyle", "errorTitle", "error", "promptTitle", "prompt")


#: what openpyxl prints for each of Excel's extension blocks it can't keep: said on every read of a workbook Excel
#: has saved, to a window nobody can act on (the firm, 30 Sep 2026). The dropdowns are put back by `load`; the
#: conditional formats PocketBook draws are all drawn again by the Set up or Run that saves the workbook
QUIET = ("Data Validation extension is not supported", "Conditional Formatting extension is not supported")


@contextmanager
def hushed():
    """Without openpyxl's two warnings about Excel's extension blocks, and nothing else. A read-only workbook reads
    each sheet as it is walked, so its warnings come then, not at the load: the walk goes inside this too."""
    with warnings.catch_warnings():
        for said in QUIET:
            warnings.filterwarnings("ignore", message=said)
        yield


def quiet(path, **kw):
    """load_workbook, without openpyxl's two warnings about Excel's extension blocks."""
    with hushed():
        return load_workbook(path, **kw)


def load(path, **kw):
    """load_workbook, with the dropdowns Excel kept in its extension block put back. The file is read once and
    both passes read those bytes (30 Sep 2026: the workbook sits in a OneDrive folder at the bank, where every
    open of the file goes through the sync client and the virus scanner)."""
    if isinstance(path, (str, Path)):
        path = Path(path).read_bytes()
    wb = quiet(io.BytesIO(path) if isinstance(path, bytes) else path, **kw)
    if kw.get("read_only"):
        return wb
    for sheet, dvs in extended(path).items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        held = {str(d.sqref) for d in ws.data_validations.dataValidation}
        for dv in dvs:
            if str(dv.sqref) not in held:
                ws.add_data_validation(dv)
    return wb


def extended(path) -> dict[str, list[DataValidation]]:
    """Every dropdown in Excel's extension block, by sheet name. Empty for a
    workbook only openpyxl has written, or anything that is not a workbook. `path` may be the file's bytes."""
    out: dict[str, list[DataValidation]] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(path) if isinstance(path, bytes) else Path(path)) as z:
            files = _sheet_files(z)
            for name, part in files.items():
                if part not in z.namelist():
                    continue
                xml = z.read(part)
                if not _HAS_EXT(xml):
                    continue            # no extension block: nothing to put back, and the sheet isn't parsed twice
                dvs = _from_sheet(ET.fromstring(xml))
                if dvs:
                    out[name] = dvs
    except (zipfile.BadZipFile, KeyError, ET.ParseError, OSError):
        return {}
    return out


#: whether a sheet's XML names the dropdowns' extension at all (its uri, any case), before it is parsed
_HAS_EXT = re.compile(re.escape(DV_EXT).encode(), re.IGNORECASE).search


def _sheet_files(z: zipfile.ZipFile) -> dict[str, str]:
    """Sheet name -> its part in the zip, read from the workbook and its relationships."""
    wb = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels.iter(f"{{{PKG}}}Relationship")}
    out = {}
    for s in wb.iter(f"{{{MAIN}}}sheet"):
        t = target.get(s.get(f"{{{REL}}}id"), "")
        out[s.get("name")] = t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join("xl", t))
    return out


def _from_sheet(root: ET.Element) -> list[DataValidation]:
    out = []
    for ext in root.iter(f"{{{MAIN}}}ext"):
        if (ext.get("uri") or "").upper() != DV_EXT:
            continue
        for d in ext.iter(f"{{{X14}}}dataValidation"):
            kw = {py: d.get(xl) in ("1", "true") for xl, py in FLAGS.items() if d.get(xl) is not None}
            kw.update({k: d.get(k) for k in TEXTS if d.get(k) is not None})
            for n in (1, 2):
                f = d.find(f"{{{X14}}}formula{n}/{{{XM}}}f")
                if f is not None and f.text:
                    kw[f"formula{n}"] = f.text
            sq = d.find(f"{{{XM}}}sqref")
            if sq is None or not sq.text:
                continue
            kw.setdefault("type", "list")
            out.append(DataValidation(sqref=sq.text, **kw))
    return out
