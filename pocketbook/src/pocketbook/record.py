"""Record (the redesign, section 10): Check and Log on one grey tab, for the reviewer.

Six sections in a 2 x 3 grid of pairs, read across:

    This Run               | Settings
    Does it add up         | Tests used
    Left out               | Every Run, newest first

Every section has the same ONYX header band, a CANVAS row naming its columns, a label column of the same width on
both sides, rows of equal height, and a 2 px ONYX rule under it. A pair's two sections start on the same row and
end on the same row, the shorter one carried down in white space, so no colour block is ragged.

Every line Check and the Log carried is here. Check's lines are sorted into the four sections by what they are
(`section_of`); Settings has one row per Control setting the run asked: the answer in use now (a formula over
Control, so a change shows at once), what the last Run used, and the pair shaded when they differ. The Log's
entries are kept on the hidden `_log` sheet, one line a row as the Log tab held them, so a refused Run can add an
entry without a Run's other lines; Every Run shows them, newest first. A workbook written before this tab has a
Check and a Log: the Log becomes `_log` as it is, and Check goes (a Run writes it afresh anyway).
"""

from __future__ import annotations

import math
import re
from datetime import datetime

from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import house

SHEET = "Record"
LOG = "_log"
OLD_CHECK, OLD_LOG = "Check", "Log"
LOG_FIRST = 4             # the newest line on _log (and on the Log tab it replaces)
LOG_NOTE = ("Every Run and every refusal, newest first. Each entry is what that Run used: a line changed on "
            "Control afterwards shows on the result tabs, not here.")
NOTHING_YET = "Nothing yet: no Run has finished."
NOTHING_HERE = "Nothing for this Run."

THIS, SETTINGS, ADDS, TESTS, LEFT, RUNS = "this", "settings", "adds", "tests", "left", "runs"
SECTIONS = {THIS: ("This Run", ("Item", "Detail")),
            SETTINGS: ("Settings", ("Setting", "In use now", "Last Run used")),
            ADDS: ("Does it add up", ("Check", "Result")),
            TESTS: ("Tests used", ("Where", "Test")),
            LEFT: ("Left out", ("Where", "What and how many")),
            RUNS: ("Every Run, newest first", ("When", "What happened"))}
PAIRS = ((THIS, SETTINGS), (ADDS, TESTS), (LEFT, RUNS))

# the columns: a label and its words on the left, a label and two columns on the right, the key hidden in K
L_LABEL, L_TEXT, R_LABEL, R_A, R_B, KEY, FLAG, JOIN_L, JOIN_R = 2, 3, 5, 6, 7, 11, 12, 13, 14
WIDTHS = {1: 2, L_LABEL: 33, L_TEXT: 92, 4: 3, R_LABEL: 33, R_A: 44, R_B: 44, 8: 2, KEY: 12, FLAG: 6, JOIN_L: 4,
          JOIN_R: 4}
ROW_H = 18                # a row's height in points, one line
CHARS = {L_TEXT: 112, R_A: 52, R_B: 52, "RAB": 104, L_LABEL: 38, R_LABEL: 38}   # about how many fit on a line
#: how a row goes on from the one above it (JOIN_L / JOIN_R, hidden): the next line of a list, or the rest of a
#: sentence too long for one row, so a section's rows stay one line high and read back whole
NEXT_LINE, SAME_LINE = "nl", "sp"
#: a line that is a warning, a changed pre-spec or a refusal reads in CRIMSON, every row of it
WARNINGS = ("Warning", "Changed pre-spec", "Couldn't run")

METHOD = [
    ("What it holds", "What the last Run was, whether it adds up, which tests it used, what it left out, and every "
                      "Run and refusal. Nothing here is a result: the result tabs hold those."),
    ("Settings", "Each Control setting: the answer in use now, and what the last Run used. A setting changed since "
                 "the Run is shaded; the result tabs follow a Changes-now setting at once, and a Needs-a-Run one at "
                 "the next Run."),
    ("Every Run", "Newest first. Each entry is what that Run used, and a refusal says what it was waiting for."),
]

# --------------------------------------------------------------------------
# Where each of Check's lines goes

_THIS = ("Extract", "Loans run", "What was run", "Record of this run", "Origination dates", "Band edges used",
         "New column", "What ", "Split", "How closely ", "Pre-spec", "What the pre-spec says", "Holdout",
         "Confirmatory test", "Runs that touched the holdout", "This pre-spec's held-back runs",
         "Differs from the pre-spec", "Scouting")
_TESTS = ("Tests", "p-value", "Standard error", "The allowance for many tests covers", "Contribution before losses",
          "Decides each pocket", "Profit counts as more or less", "How profit reads", "Reading a single red",
          "Families of tests")
_LEFT = ("Left out of ", "Alone in its band")


def section_of(label: str, before: str | None = None) -> str:
    """The section a Check line belongs in, by its label. A warning goes with the line before it (a pre-spec's
    warning with the pre-spec, a budget's with the budget); `before` is that line's section."""
    if label == "Warning":
        return before or LEFT
    if label.startswith(_LEFT):
        return LEFT
    if label.startswith(_TESTS):
        return TESTS
    if label.startswith(_THIS):
        return THIS
    return ADDS


# --------------------------------------------------------------------------
# The Log, kept on _log


def log_sheet(wb):
    """_log, made on first use. A workbook written before Record has a Log tab: it becomes _log as it is."""
    if LOG in wb.sheetnames:
        return wb[LOG]
    if OLD_LOG in wb.sheetnames:
        ws = wb[OLD_LOG]
        ws.title = LOG
    else:
        ws = wb.create_sheet(LOG)
        ws["A1"] = "Log"
        ws["A2"] = LOG_NOTE
    ws.sheet_state = "hidden"
    return ws


def add(wb, lines: list[str], stamp: str | None = None) -> str:
    """One entry, newest first: the time on its first line, one line a row. Returns the time written."""
    ws = log_sheet(wb)
    if ws["A1"].value != "Log":
        if ws.max_row > 1 or ws["A1"].value:
            ws.insert_rows(1, amount=LOG_FIRST - 1)
        ws["A1"] = "Log"
    ws["A2"] = LOG_NOTE
    stamp = stamp or datetime.now().strftime("%Y-%m-%d %H:%M")
    ws.insert_rows(LOG_FIRST, amount=len(lines) + 1)
    for i, line in enumerate(lines):
        ws.cell(row=LOG_FIRST + i, column=1, value=stamp if i == 0 else None)
        ws.cell(row=LOG_FIRST + i, column=2, value=line)
    return stamp


def entries(wb) -> list[tuple[str | None, list[str]]]:
    """Every entry in the workbook's log, newest first: (when, its lines). Reads _log, or an older workbook's Log
    tab, and works on a workbook opened read-only."""
    name = LOG if LOG in wb.sheetnames else OLD_LOG if OLD_LOG in wb.sheetnames else None
    if name is None:
        return []
    out: list[tuple[str | None, list[str]]] = []
    for row in wb[name].iter_rows(min_row=LOG_FIRST, max_col=2, values_only=True):
        when, line = (tuple(row) + (None, None))[:2]
        if when not in (None, ""):
            out.append((str(when), []))
        if isinstance(line, str) and line:
            if not out:
                out.append((None, []))
            out[-1][1].append(line)
    return out


# --------------------------------------------------------------------------
# Writing the tab


def _fill(hex_: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_, bgColor=hex_)


def _lines(v, chars: int) -> int:
    """About how many lines a value takes in a column `chars` wide: a formula by the words it joins."""
    if v is None:
        return 1
    s = str(v)
    if s.startswith("="):
        s = "".join(re.findall(r'"((?:[^"]|"")*)"', s)) + " " * 6 * s.count("TEXT(")
    return max(1, math.ceil(len(s) / chars))


def _band(ws, r: int, first: int, last: int, title: str, heads: tuple, spans: list[tuple[int, int]]) -> None:
    """A section's header band (ONYX, 32 px) and the CANVAS row naming its columns."""
    for c in range(first, last + 1):
        ws.cell(row=r, column=c).fill = _fill(house.ONYX)
    h = ws.cell(row=r, column=first, value=title)
    h.font = Font(name="Arial", bold=True, size=10, color=house.PAPER)
    h.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[r].height = 24
    for (a, b), text in zip(spans, heads):
        for c in range(a, b + 1):
            x = ws.cell(row=r + 1, column=c)
            x.fill = _fill(house.CANVAS)
            x.border = Border(bottom=Side(style="thin", color=house.MIST))
        if b > a:
            ws.merge_cells(start_row=r + 1, start_column=a, end_row=r + 1, end_column=b)
        x = ws.cell(row=r + 1, column=a, value=text)
        x.font = Font(name="Arial", bold=True, size=9, color=house.SLATE)
        x.alignment = Alignment(horizontal="left" if a in (L_LABEL, L_TEXT, R_LABEL) or len(heads) == 2
                                else "center", vertical="center", indent=1)
    ws.row_dimensions[r + 1].height = ROW_H


def _chunks(text: str, width: int) -> list[str]:
    """Words too long for one row, wrapped into rows the way a cell wraps them: as many words as fit, then the next
    row. Each cut takes out exactly one space, so the rows joined with one space give the words back as they were.
    A word longer than a row stays whole, and that row wraps."""
    if len(text) <= width:
        return [text]
    words = text.split(" ")
    out, cur = [], words[0]
    for w in words[1:]:
        if len(cur) + 1 + len(w) <= width:
            cur += " " + w
        else:
            out.append(cur)
            cur = w
    out.append(cur)
    return out


def _rows_of(kind: str, rows: list, left: bool) -> list:
    """A section's rows as the grid draws them: (label, a, b, flag, join). A value of several lines is one row a
    line, and a sentence too long for one row goes on in the rows under it, its label on the first, so rows stay
    one line high; `join` says how a row goes on from the one above (NEXT_LINE, SAME_LINE). A formula can't be cut:
    it wraps."""
    width = CHARS[L_TEXT] if left else CHARS["RAB"]
    out = []
    for x in rows:
        label, a = x[0], x[1]
        b = x[2] if len(x) > 2 else None
        flag = x[3] if len(x) > 3 else None
        if kind == SETTINGS or not isinstance(a, str) or a.startswith("="):
            out.append((label, a, b, flag, None))
            continue
        first = True
        warn = "warn" if label == "Warning" or a.startswith(WARNINGS) else None      # every row of it, red
        for i, line in enumerate(a.split("\n")):
            for j, piece in enumerate(_chunks(line, width)):
                join = None if first else (SAME_LINE if j else NEXT_LINE)
                out.append((label if first else None, piece, None, warn, join))
                first = False
    return out


def _put(ws, r: int, kind: str, row: tuple, left: bool) -> int:
    """One row of a section; returns the lines it needs."""
    label, a, b, flag, join = row
    ws.cell(row=r, column=JOIN_L if left else JOIN_R, value=join)
    lab = ws.cell(row=r, column=L_LABEL if left else R_LABEL, value=label)
    lab.font = Font(name="Calibri", size=10, color=house.SLATE)
    lab.alignment = Alignment(horizontal="left", vertical="center", indent=1, wrap_text=True)
    need = _lines(label, CHARS[L_LABEL])
    if left:
        cells = [(L_TEXT, a, CHARS[L_TEXT])]
    elif kind == SETTINGS:
        cells = [(R_A, a, CHARS[R_A]), (R_B, b, CHARS[R_B])]
    else:
        ws.merge_cells(start_row=r, start_column=R_A, end_row=r, end_column=R_B)
        cells = [(R_A, a, CHARS["RAB"])]
    for c, v, chars in cells:
        x = ws.cell(row=r, column=c, value=v)
        x.font = Font(name="Calibri", size=10, color=house.INK_TEXT, bold=kind == SETTINGS and c == R_A)
        x.alignment = Alignment(horizontal="center" if kind == SETTINGS else "left", vertical="center", indent=1,
                                wrap_text=True)
        if kind != SETTINGS and flag == "warn":
            x.font = Font(name="Calibri", size=10, color=house.CRIMSON, bold=True)
        need = max(need, _lines(v, chars))
    if kind == SETTINGS and flag:
        ws.cell(row=r, column=FLAG, value=flag)
    for c in ((L_LABEL, L_TEXT) if left else (R_LABEL, R_A, R_B)):
        ws.cell(row=r, column=c).border = Border(bottom=Side(style="thin", color=house.ROW_RULE))
    return need


def _rule(ws, r: int, first: int, last: int) -> None:
    """The 2 px ONYX rule under a section."""
    for c in range(first, last + 1):
        x = ws.cell(row=r, column=c)
        x.border = Border(bottom=Side(style="medium", color=house.ONYX))


def _pair(ws, top: int, left: str, right: str, rows: dict) -> int:
    """Two sections side by side from row `top`, ending on the same row. Returns the row after them."""
    ws.cell(row=top, column=KEY, value=f"pair|{left}|{right}")
    _band(ws, top, L_LABEL, L_TEXT, SECTIONS[left][0], SECTIONS[left][1], [(L_LABEL, L_LABEL), (L_TEXT, L_TEXT)])
    spans = [(R_LABEL, R_LABEL), (R_A, R_A), (R_B, R_B)] if right == SETTINGS else [(R_LABEL, R_LABEL), (R_A, R_B)]
    _band(ws, top, R_LABEL, R_B, SECTIONS[right][0], SECTIONS[right][1], spans)
    empty = NOTHING_YET if rows.get("none") else NOTHING_HERE
    a = _rows_of(left, rows.get(left, []), True) or [(None, empty, None, None, None)]
    b = _rows_of(right, rows.get(right, []), False) or [(None, empty, None, None, None)]
    r = top + 2
    for i in range(max(len(a), len(b))):
        need = 1
        if i < len(a):
            need = max(need, _put(ws, r, left, a[i], True))
        if i < len(b):
            need = max(need, _put(ws, r, right, b[i], False))
        elif right != SETTINGS:
            ws.merge_cells(start_row=r, start_column=R_A, end_row=r, end_column=R_B)
        ws.cell(row=r, column=KEY, value=f"row|{left}|{right}")
        ws.row_dimensions[r].height = ROW_H if need == 1 else 14 * need + 4
        r += 1
    _rule(ws, r - 1, L_LABEL, L_TEXT)
    _rule(ws, r - 1, R_LABEL, R_B)
    if right == SETTINGS and b:
        # a setting changed since the last Run: the pair shaded ALERT_FG (FLAG is 1 while they differ)
        ws.conditional_formatting.add(f"{_col(R_A)}{top + 2}:{_col(R_B)}{r - 1}", FormulaRule(
            formula=[f"${_col(FLAG)}{top + 2}=1"], fill=_fill(house.ALERT_FG),
            font=Font(bold=True, color=house.CRIMSON)))
    return r


def _col(n: int) -> str:
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def log_rows(wb) -> list[tuple[str | None, str]]:
    """Every Run's rows: its time on its first line."""
    return [(when if i == 0 else None, line) for when, lines in entries(wb) for i, line in enumerate(lines)]


def write(wb, rows: dict[str, list] | None) -> None:
    """Record, afresh. `rows`: each section's rows, (label, words) or, for Settings, (setting, in use now, last Run
    used, changed) where `changed` is a formula that is 1 while the two differ. None: a workbook with no Run yet
    (a refusal before the first Run), whose sections say so."""
    at = wb.sheetnames.index(SHEET) if SHEET in wb.sheetnames else None
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    for old in (OLD_CHECK,):
        if old in wb.sheetnames:
            del wb[old]
    log_sheet(wb)
    ws = wb.create_sheet(SHEET, at) if at is not None else wb.create_sheet(SHEET)
    for c, w in WIDTHS.items():
        ws.column_dimensions[_col(c)].width = w
    house.title_band(ws, SHEET, "What ran, whether it adds up, what was left out, and every Run. For the reviewer.",
                     L_LABEL, R_B, tab=house.TAB_RECORD, fill_hex=house.SLATE, rule=house.STONE,
                     sub_color=house.MIST)
    r = house.method_note(ws, 3, L_LABEL, R_B, METHOD)
    rows = dict(rows) if rows is not None else {"none": True}
    rows[RUNS] = log_rows(wb)
    for left, right in PAIRS:
        r = _pair(ws, r, left, right, rows) + 1
    for c in (KEY, FLAG, JOIN_L, JOIN_R):
        ws.column_dimensions[_col(c)].hidden = True
    ws.freeze_panes = "A2"
    ws.print_area = f"B1:{_col(R_B)}{r}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def refresh_runs(wb) -> None:
    """Every Run drawn again after the log gains an entry, the rest of Record left as the last Run wrote it. With
    no Record yet, one is written whose other sections say no Run has finished."""
    if SHEET not in wb.sheetnames:
        return write(wb, None)
    ws = wb[SHEET]
    top = next((r for r in range(1, ws.max_row + 1) if str(ws.cell(row=r, column=KEY).value or "")
                == f"pair|{LEFT}|{RUNS}"), None)
    if top is None:
        return write(wb, None)
    none = any(ws.cell(row=r, column=L_TEXT).value == NOTHING_YET for r in range(top + 2, ws.max_row + 1)
               if ws.cell(row=r, column=KEY).value == f"row|{LEFT}|{RUNS}")
    left = [(label, v) for kind, label, v, _, _ in read(ws) if kind == LEFT]
    for rng in [m for m in ws.merged_cells.ranges if m.min_row >= top]:
        ws.unmerge_cells(str(rng))
    for r in range(top, ws.max_row + 1):
        for c in range(1, JOIN_R + 1):
            x = ws.cell(row=r, column=c)
            x.value, x.border, x.fill, x.font = None, Border(), PatternFill(), Font()
        ws.row_dimensions[r].height = None
    end = _pair(ws, top, LEFT, RUNS, {LEFT: left, RUNS: log_rows(wb), "none": none}) + 1
    ws.print_area = f"B1:{_col(R_B)}{end}"


# --------------------------------------------------------------------------
# Reading it back (the launcher's tests and the Run's own checks read it this way)


def read(ws) -> list[tuple[str, str | None, object, object, object]]:
    """Every row of every section as (section, label, value, last, changed): Settings' value is what is in use now,
    `last` what the last Run used and `changed` 1 while they differ (the shading follows it); a line carried over
    from the row above (a list of several lines, or a sentence too long for one row) is joined back to it, so each
    value reads whole; a label of None is a row with no label of its own (Every Run's later lines)."""
    got: dict[str, list] = {k: [] for k in SECTIONS}
    for r in range(1, ws.max_row + 1):
        key = str(ws.cell(row=r, column=KEY).value or "")
        if not key.startswith("row|"):
            continue
        _, left, right = key.split("|")
        for kind, lab, a, b, j in ((left, L_LABEL, L_TEXT, None, JOIN_L),
                                   (right, R_LABEL, R_A, R_B if right == SETTINGS else None, JOIN_R)):
            label, v = ws.cell(row=r, column=lab).value, ws.cell(row=r, column=a).value
            if label is None and v in (None, NOTHING_YET, NOTHING_HERE):
                continue
            join = ws.cell(row=r, column=j).value
            if join in (NEXT_LINE, SAME_LINE) and got[kind]:
                k, lab_, prev, last, changed = got[kind][-1]
                got[kind][-1] = (k, lab_, f"{prev}{chr(10) if join == NEXT_LINE else ' '}{v}", last, changed)
                continue
            got[kind].append((kind, label, v, ws.cell(row=r, column=b).value if b else None,
                              ws.cell(row=r, column=FLAG).value if b else None))
    return [x for left, right in PAIRS for k in (left, right) for x in got[k]]
