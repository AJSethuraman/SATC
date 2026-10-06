"""The New variables tab (the redesign, section 9; capabilities 4b and 4e): does each candidate still tell good loans
from bad on loans it was never found on, and once the held-fixed columns are held fixed? It replaces the
Confirmatory test tab and keeps every statistic that tab showed.

It leads (OC-51; the firm, 27 Sep 2026: "tree guesses on 2024 data if 2022-2023 are used to build branches -->
regress shortlist?") with the tree on loans it never saw, one line, when this Run scouted; then every candidate
together in one regression on the held-back loans (joint.py): each group's odds ratio net of the other candidates,
and what each candidate adds net of them; then each candidate on its own, as below.

A test of a new variable today confirms a saved shortlist, the pre-spec: its column cut into groups, each group
compared with the reference group. Each row of the table is one candidate comparison, a group against the reference:

    Found                    on the development loans, where the groups came from, nothing held fixed
    Confirmed · held back    the same comparison on the holdout, nothing held fixed; Holds up?
    Confirmed, X held fixed  the same again inside the pockets the held-fixed columns make; Still holds?
    Excess, Material?        the group's losses on the holdout above its share, scaled to the whole book, against
                             Control's materiality line
    In words                 the two verdicts together

Each input is reported with and without the columns held fixed (the firm's lean pre-spec, 26 Sep 2026). A saved
shortlist was found on another run, so its Found columns are hidden, and Record names the file (the spec).

Under the chart, the tests in full: B3 and B4 (does the column matter at all, and in one direction), B5 (each
group's odds ratio with its range) on every set of loans with and without the columns held fixed, and 4e's
concentration on the holdout. Each has its finding in one plain line.

Live (OC-40): Holds up?, Still holds?, Material?, In words, every "significant", every range and the chart's worse
line follow Control (confidence, materiality, worse at). The odds ratios, statistics and p-values are the Run's;
none depends on those settings. No SORT, FILTER or LET.
"""

from __future__ import annotations

import math

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import house, joint, live

SHEET = "New variables"
OLD_SHEET = "Confirmatory test"                            # the tab this replaces; a Run takes it off
INK, SLATE, PAPER = house.INK_TEXT, house.SLATE, house.PAPER
WORSE_FILL, BETTER_FILL = house.ALERT_FG, house.POSITIVE_BG
FIRST, LAST = 2, 14                                        # columns B to N
#: the table's columns
(N_CAND, N_COMP, N_FG, N_FP, N_CG, N_CP, N_HOLDS, N_HG, N_HP, N_STILL, N_EX, N_MAT, N_WORDS) = range(2, 15)
LINE_COL = 24                                              # X, hidden: the chart's worse line, =worse_at a row
HELPER = 26                                                # Z, hidden: the pieces the plain lines join
Z = "NORMSINV(1-(1-confidence)/2)"                         # the live z, as the Split tab writes it
P_FMT = '[<0.0001]"under 0.01%";[<0.01]0.00%;0.0%'         # book.P_FMT
X_FMT = '0.00"×"'
LIFT_TAIL = " times the holdout's bad rate"
YES, NO = "Yes", "No"
WIDTHS = {1: 2, N_CAND: 20, N_COMP: 30, N_FG: 10, N_FP: 11, N_CG: 16, N_CP: 12, N_HOLDS: 16, N_HG: 16, N_HP: 12,
          N_STILL: 16, N_EX: 14, N_MAT: 12, N_WORDS: 42}
#: the tests in full keep their numbers in F to N, never in Found's two columns, which a saved shortlist hides
DATA = N_CG
TEXT_WIDE = 190                                            # about how many characters fit across C to N


def _c(n: int) -> str:
    return live.col(n)


def _q(s: str) -> str:
    return live.q(s)


def _cell(ws, r: int, c: int, v=None, *, bold=False, color=INK, h="center", fmt=None, size=10):
    x = ws.cell(row=r, column=c, value=v)
    x.font = Font(name="Calibri", bold=bold, size=size, color=color)
    x.alignment = Alignment(horizontal=h, vertical="center")
    if fmt:
        x.number_format = fmt
    return x


def _rule(ws, r: int, first: int = FIRST, last: int = LAST) -> None:
    for c in range(first, last + 1):
        ws.cell(row=r, column=c).border = Border(bottom=Side(style="thin", color=house.ROW_RULE))
    ws.row_dimensions[r].height = 18


def _line(ws, r: int, label: str | None, text, chars: int | None = None) -> int:
    """A finding in words: a label in B and the words across C to N. Prose, so it may wrap. Returns the next row."""
    if label is not None:
        _cell(ws, r, FIRST, label, bold=True, h="left").alignment = Alignment(horizontal="left", vertical="top")
    ws.merge_cells(start_row=r, start_column=FIRST + 1, end_row=r, end_column=LAST)
    c = ws.cell(row=r, column=FIRST + 1, value=text)
    c.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    c.font = Font(name="Calibri", size=10, color=INK)
    n = chars if chars is not None else len(str(text))
    ws.row_dimensions[r].height = 15.0 * max(1, math.ceil(n / TEXT_WIDE)) + 4
    return r + 1


def _sig_word(ref: str) -> str:
    return f'IF({ref}="","",IF({live.sig(ref)},"significant","not significant"))'


def _held_words(strata) -> str:
    return " and ".join(strata)


def _num(v) -> float | None:
    return None if v is None or (isinstance(v, float) and not math.isfinite(v)) else v


class _Views:
    """The four sets of numbers a candidate is read on, by the words the tab uses for them."""

    def __init__(self, t):
        held = _held_words(t.strata)
        self.held = held
        dev, hold = t.development, t.holdout
        if t.held:
            self.all = [(t.development_plain, "Found, nothing held fixed", "On development, with nothing held fixed,"),
                        (dev, f"Found, {held} held fixed", "On development,"),
                        (t.holdout_plain, "Confirmed on the held-back loans, nothing held fixed",
                         "On the holdout, with nothing held fixed,"),
                        (hold, f"Confirmed on the held-back loans, {held} held fixed", "On the holdout,")]
        else:
            self.all = [(dev, "Found, nothing held fixed", "On development,"),
                        (hold, "Confirmed on the held-back loans, nothing held fixed", "On the holdout,")]


def _loans_said(tests, side_name: str) -> str:
    """How many loans and bad loans a set holds, once when every candidate has the same, else candidate by
    candidate (a loan with no readable value of one candidate is left out of that candidate's test only)."""
    sides = [getattr(t, side_name) for t in tests]
    if len({(s.n, s.n_bad) for s in sides}) == 1:
        s = sides[0]
        rate = s.n_bad / s.n if s.n else None
        return f"{s.n:,} loans, {s.n_bad:,} bad" + (f" ({rate:.2%})" if rate is not None else "")
    return "; ".join(f"{t.column} {s.n:,} loans, {s.n_bad:,} bad" for t, s in zip(tests, sides))


def _method(res, st, tests, excesses, stamp: str) -> list[tuple[str, object]]:
    """The tab's one method note (tenet T1): every test and choice behind the rows, once, in plain words. A shortlist
    of one reads as the one-column pre-spec's did; several candidates are named together."""
    from . import confirmatory
    t = tests[0]
    one = len(tests) == 1
    lead = _method_lead(res, st)
    Ks = sorted({len(x.groups) for x in tests})
    K = Ks[0] if len(Ks) == 1 else None
    ref = t.groups[t.ref] if one else "its reference group"
    held = _held_words(t.strata)
    dev, hold = t.development, t.holdout
    if one:
        out = lead + [("The column", f"{t.column}, in {K} groups: {'; '.join(t.groups)}. The groups are the pre-spec's."),
               ("Compared with", f"{ref}, the pre-spec's reference group. Every other group is compared with it.")]
    else:
        where = "the shortlist scouting wrote" if getattr(st, "scouted", False) else "the saved shortlist"
        out = lead + [("The candidates", f"{len(tests)} inputs on {where}, each in its own groups, the "
                                  f"pre-spec's: " + "; ".join(f"{x.column} in {len(x.groups)} groups "
                                                             f"({', '.join(x.groups)})" for x in tests) + "."),
               ("Compared with", "Each candidate's own reference group, the pre-spec's: "
                                 + "; ".join(f"{x.column} against {x.groups[x.ref]}" for x in tests)
                                 + ". Every other group of a candidate is compared with its reference.")]
    said = []
    for side_name, side, what in (("development", dev, "found on: where the groups were chosen"),
                                  ("holdout", hold, "held back: the test that counts")):
        said.append(f"Loans made {side.range.text()}: {_loans_said(tests, side_name)}; {what}.")
    def left_words(left):
        return "; ".join(f"{v:,} {k}" if k.startswith("made") else f"{v:,} with {k}" for k, v in left.items())
    if one or len({tuple(x.left_out.items()) for x in tests}) == 1:
        if t.left_out:
            said.append(f"Left out of {'this test' if one else 'each test'}: {left_words(t.left_out)}. Record counts "
                        f"every loan.")
    else:
        said.append(" ".join(f"Left out of {x.column}'s test: {left_words(x.left_out)}." for x in tests
                             if x.left_out) + " Record counts every loan.")
    out.append(("Found, then confirmed", " ".join(said) + " The comparison is made once on each; only the "
                                                          "held-back result counts."))
    out.append(("Gap", f"The odds ratio against {ref}. Odds are bad loans divided by good ones, and an odds ratio is "
                       f"a group's odds divided by {ref}'s: 2.00× means twice the odds of going bad. While few loans "
                       f"go bad it is close to how many times as often they do. Worked out by conditional logistic "
                       f"regression, which compares the groups inside each pocket and takes each pocket's own count "
                       f"of bad loans as given, instead of estimating a rate for every pocket, which biases the "
                       f"answer when pockets are thin."))
    how = confirmatory.allowance(res)
    after = "" if how == "none" else ", after the allowance for testing many at once (below)"
    out.append(("p-value", live.text("The chance of a gap at least this big if the group were no different from "
                                     f"{ref}{after}. Below ", ('TEXT(significance_bar,"0%")',), " is significant, at ",
                                     ('TEXT(confidence,"0%")',), " sure. Two-sided: a gap either way counts.")))
    out.append(("Many at once", confirmatory.allowance_words(res, tests)))
    out.append(("Holds up?", "Yes when the held-back loans show the gap on the same side of 1 as the found loans "
                             "did, and significant. Only the held-back result counts: the found one is where the "
                             "idea came from."))
    what_col = t.column if one else "the candidate"
    if t.held:
        out.append((f"{held} held fixed",
                    f"The held-back comparison made again with every loan compared only with loans in its own "
                    f"pocket, cut by {held} at the band edges on Record. A difference in the {held} mix then can't "
                    f"pass for a difference in {what_col}. Still holds? says whether the gap stays; if it goes, "
                    f"the candidate was mostly telling you {held}. Every pocket counts, however small: fewest "
                    f"loans and fewest losses decide only whether one pocket can be read on its own, and a "
                    f"pocket where no loan, or every loan, went bad says nothing and adds nothing."))
    else:
        out.append(("Nothing held fixed", "The pre-spec holds no column fixed, so the whole book is one pocket and "
                                          "there is one confirmation."))
    ex = next((e for e in excesses if e is not None), None)
    if ex is not None:
        what = "GCOs" if ex.unit == "dollars" else "bad loans"
        out.append(("Material?", f"Excess is a group's {what} on the held-back loans above its share of them (its "
                                 f"share of the held-back loans), times the whole book's {what} over the held-back "
                                 f"loans', so it sits on the same scale as Control's materiality line. Material? is "
                                 f"Yes when it reaches the line."))
    out.append(("In words", f"The two verdicts read together: holds up, and not just {held or 'the columns held fixed'}; "
                            f"holds up, but it was mostly {held or 'those columns'}; holds up only with them held "
                            f"fixed; or didn't hold up on held-back loans." if t.held else
                "Holds up on held-back loans, or didn't."))
    k_words = (f"the Mantel-Haenszel test for {K} groups" if K else "the Mantel-Haenszel test for its groups")
    df_words = (f"{K - 1} degrees of freedom" if K else "one fewer degree of freedom than its groups")
    out.append(("The tests in full", f"Under the chart{'' if one else ', one block per candidate'}. Does "
                                     f"{t.column if one else 'each candidate'} matter at all: the general test "
                                     f"({k_words}: a standard way to add pockets up without "
                                     f"mixing their loans), read against a chi-square on {df_words}"
                                     f"{', one fewer than the groups' if K else ''}; the trend test, which gives the "
                                     f"groups the scores 1 to {K if K else 'their number'} and asks whether the bad "
                                     f"rate climbs or falls steadily with them, on 1 "
                                     f"degree of freedom; and the regression's block test (a likelihood ratio test, "
                                     f"on {df_words}), whether the groups together add anything. A "
                                     f"difference with no trend is a U or a hump, which a straight-line test would "
                                     f"call nothing. Then each group's odds ratio with its range and its raw p-value, "
                                     f"before any allowance, and how much of the "
                                     f"held-back loans' losses sit in each group (on the holdout only, and on the "
                                     f"book as a whole, not pocket by pocket: the lift, the group's bad rate over "
                                     f"the holdout's, is the finding, never the share of all losses)."))
    gaps = [s.mh_gap for x in tests for s in (x.development, x.holdout) if s.mh_gap is not None]
    out.append(("The range", live.text("Where the true odds ratio sits, ", ('TEXT(confidence,"0%")',), " sure: the "
                                       "log odds ratio plus and minus the multiple of its standard error that the "
                                       "confidence level sets. It is each group's own, before the allowance."
                                       + (f" As a check, each odds ratio held fixed was worked out a second way, pair "
                                          f"by pair with Mantel-Haenszel: the largest gap between the two is "
                                          f"{max(gaps):.1%}." if gaps else ""))))
    if st is not None and st.scouted:
        out.append(("Written by scouting", f"This Run found the candidates on its development loans (the Scouting "
                                           f"tab), wrote the pre-spec {st.name} from them, and then confirmed it on "
                                           f"the loans held back. Found is those development loans, where the "
                                           f"groups were chosen. Record names the file."))
    elif st is not None:
        out.append(("Written shortlist", f"This run confirms the saved shortlist {st.name}, the pre-spec. Its Found "
                                         f"columns are hidden (unhide columns D and E to see them) and Record names "
                                         f"the file. The tests in full show both sets of loans."))
    for x in getattr(st, "tests", None) or ():
        if x.problem:
            out.append(("Not run", f"{x.column} couldn't be tested: {x.problem}. Record says why."))
    out.append(("Choices made here that no ruling settles yet", " ".join((
        f"The trend test scores the groups 1 to {K}, evenly spaced, rather than by where their numbers sit." if K else
        "The trend test scores each candidate's groups 1, 2, 3 and on, evenly spaced, rather than by where their "
        "numbers sit.",
        "Conditional logistic regression is used in every pocket, whatever its size: the unconditional fit the scope "
        "allows for pockets over a few thousand loans isn't used, as the conditional one is exact and quick at any "
        "size.",
        "A loan with no readable value in a column the pockets are cut by sits in a pocket of its own, as on the "
        "grids.",
        "A group with no bad loan, or only bad loans, in the pockets that say something has no odds ratio; the "
        "others are fitted as if its loans weren't there, which is where the regression heads anyway.",
        "Excess is scaled to the whole book by its losses over the held-back loans' losses.",
        "The allowance treats each set of loans as one family, every candidate's groups together, as a grid's "
        "pockets are one family per rate.",
        *(["All together gives each pocket its own constant rather than conditioning the pockets out, and refuses "
           f"when they average under {joint.MIN_BAD_PER_POCKET} bad loans each (statistics.md B10)."]
          if _joint_shown(st) else [])))))
    out.append(("As of", f"The numbers are the last Run's, {stamp}. Holds up?, Still holds?, Material?, In words, "
                         f"every range and the chart's worse line follow Control."))
    return out


def _joint_shown(st) -> bool:
    """Whether the tab shows every candidate together: two candidates or more ran (with one, together is alone)."""
    j = getattr(st, "joint", None)
    return j is not None and len([t for t in getattr(st, "tests", None) or () if t.problem is None]) > 1


def _method_lead(res, st) -> list[tuple[str, object]]:
    """The method note's first items (OC-51): the tree on loans it never saw, and every candidate together."""
    from . import scout
    out = []
    sc = getattr(res, "scout", None)
    if sc is not None and getattr(sc, "oot", None) is not None:
        held = " and ".join(sc.hold)
        out.append(("The tree, unseen", "The random forest from the Scouting tab, grown on every loan made before the "
                                        f"cutoff ({sc.cutoff.isoformat()}), scores the loans made on or after it, "
                                        "which it never saw. AUC is the chance a random bad loan scores above a "
                                        "random good one: 0.5 is a coin flip, 1 is perfect. Built is the forest's "
                                        "AUC on the development loans, each scored by a forest grown on the others "
                                        "(the Scouting tab); unseen is on the held-back loans. Unseen well under "
                                        "built means the tree learnt those years, not the book."
                                        + (f" Said twice: from the candidates alone, and with {held} in the forest "
                                           f"too." if sc.hold else "")))
    if _joint_shown(st):
        j = st.joint
        held = " and ".join(j.strata)
        pockets = (f"one constant for each pocket {held} make, so the pockets are held fixed as they are below"
                   if j.strata else "one constant for the whole book, as nothing is held fixed")
        out.append(("All together", "One logistic regression on the held-back loans with every candidate's groups in "
                                    f"it at once, each against its own reference, and {pockets}. Each odds ratio is "
                                    "then net of every other candidate. Adds is a likelihood ratio test of taking "
                                    "the candidate out, all its groups at once: its test statistic, on as many "
                                    "degrees of freedom as it has groups less one, and its p-value"
                                    + ("" if confirmatory_allowance(res) == "none" else
                                       ", after the allowance for testing the shortlist at once") + ". Adds? is Yes "
                                    "when that p-value is under the bar. Moves with names candidates whose values "
                                    f"rise and fall together (a rank correlation of {scout.CORRELATED:.1f} or more, "
                                    "on the development loans): what one adds, the other may already say. The "
                                    "loans are the held-back ones with a value of every candidate, in pockets where "
                                    "some went bad and some didn't."))
    return out


def confirmatory_allowance(res) -> str:
    from . import confirmatory
    return confirmatory.allowance(res)


def _unseen(ws, res, r: int) -> int:
    """The tree on loans it never saw, in one line (OC-51); nothing when this Run didn't scout."""
    from . import scout
    sc = getattr(res, "scout", None)
    o = getattr(sc, "oot", None)
    if o is None:
        return r
    house.section(ws, r, FIRST, LAST, "The tree on loans it never saw")
    said = scout.auc_words(sc) or f"The tree couldn't be checked on the held-back loans: {o.problem}."
    r = _line(ws, r + 1, "What it found", said)
    return r + 1


JOINT_HEADS = ("Candidate", "Compared", None, None, "Loans", "Bad rate", "Odds ratio, together", None, "p-value",
               "Adds: statistic", None, "Adds?", "Moves with")
(J_CAND, J_COMP, J_LOANS, J_RATE, J_ODDS, J_RANGE, J_P, J_STAT, J_PADD, J_ADDS, J_WITH) = (
    N_CAND, N_COMP, N_CG, N_CP, N_HOLDS, N_HG, N_HP, N_STILL, N_EX, N_MAT, N_WORDS)


def _together(ws, res, st, r: int, helper_rows: list) -> tuple[int, int | None]:
    """Every candidate in one regression on the held-back loans (OC-51, joint.py): a row per group against its
    reference, net of the other candidates, and per candidate what it adds net of them. Returns (the next row, the
    table's header row or None). Nothing with one candidate: together it is the same as on its own, below."""
    if not _joint_shown(st):
        return r, None
    j = st.joint
    house.section(ws, r, FIRST, LAST, f"All {len([t for t in j.terms])} candidates together, on the held-back loans")
    r += 1
    if j.problem:
        r = _line(ws, r, "Not fitted", f"The candidates couldn't be fitted together: {j.problem}.")
        return r + 1, None
    raw = j.allowance == "none"
    rng_head = live.text("Range (", ('TEXT(confidence,"0%")',), " sure)")
    heads = list(JOINT_HEADS)
    heads[J_RANGE - FIRST] = rng_head
    heads[J_PADD - FIRST] = "Adds: p-value" if raw else "Adds: p, allowed"
    head = r
    house.header(ws, head, FIRST, heads, centre_from=2)
    ws.merge_cells(start_row=head, start_column=J_COMP, end_row=head, end_column=J_COMP + 2)
    _wrap(ws, head)
    r += 1
    first_row = r
    adds_cells = []
    for i, t in enumerate(j.terms):
        if t.problem:
            continue
        ref = t.groups[t.ref]
        top = r
        for k, name in enumerate(t.groups):
            if k == t.ref:
                continue
            ws.merge_cells(start_row=r, start_column=J_COMP, end_row=r, end_column=J_COMP + 2)
            _cell(ws, r, J_CAND, t.column if r == top else None, bold=r == top, h="left")
            _cell(ws, r, J_COMP, f"{name} vs {ref}", color=SLATE)
            _cell(ws, r, J_LOANS, t.loans[k], fmt="#,##0")
            _cell(ws, r, J_RATE, t.bad[k] / t.loans[k] if t.loans[k] else None, fmt="0.00%")
            if t.odds[k] is None:
                _cell(ws, r, J_ODDS, "none")
            else:
                _cell(ws, r, J_ODDS, round(t.odds[k], 10), bold=True, fmt=X_FMT)
                b, se = live.num(t.beta[k]), live.num(t.se[k])
                _cell(ws, r, J_RANGE, f'=TEXT(EXP({b}-{Z}*{se}),"0.00")&"x to "&TEXT(EXP({b}+{Z}*{se}),"0.00")&"x"')
                _cell(ws, r, J_P, t.p[k], fmt=P_FMT)
            if r == top:
                if t.lr is not None:
                    _cell(ws, r, J_STAT, f"{t.lr:.2f} on {t.df}")
                    padd = t.p_lr if raw else t.p_allowed
                    _cell(ws, r, J_PADD, padd, fmt=P_FMT, bold=True)
                    P = f"${_c(J_PADD)}{r}"
                    _cell(ws, r, J_ADDS, f'=IF(NOT(ISNUMBER({P})),"",IF({live.sig(P)},"{YES}","{NO}"))', bold=True)
                    adds_cells.append((t.column, P))
                _cell(ws, r, J_WITH, ", ".join(f"{n} ({rho:+.2f})" for n, rho in t.partners), h="left")
            _rule(ws, r)
            if i and r == top:
                for c in range(FIRST, LAST + 1):
                    ws.cell(row=r, column=c).border = Border(top=Side(style="medium", color=house.INK),
                                                             bottom=Side(style="thin", color=house.ROW_RULE))
            r += 1
    last_row = r - 1
    if last_row >= first_row:
        rng = f"{_c(J_ADDS)}{first_row}:{_c(J_ADDS)}{last_row}"
        ws.conditional_formatting.add(rng, FormulaRule(
            formula=[f'{_c(J_ADDS)}{first_row}="{YES}"'], font=Font(bold=True, color=house.POSITIVE),
            fill=PatternFill("solid", fgColor=house.POSITIVE_BG, bgColor=house.POSITIVE_BG)))
        ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{_c(J_ADDS)}{first_row}="{NO}"'],
                                                       font=Font(bold=True, color=SLATE)))
    r += 1
    # what it found, live: which candidates add something net of the others
    yes = "&".join(f'IF({live.sig(P)},{_q("; " + name)},"")' for name, P in adds_cells) or '""'
    no = "&".join(f'IF(AND(ISNUMBER({P}),NOT({live.sig(P)})),{_q("; " + name)},"")' for name, P in adds_cells) \
        or '""'
    hy, hn = _helper(ws, r, yes), _helper(ws, r + 1, no)
    none_ = _q("No candidate adds anything significant once the others are in. ")
    adds, adds_not = _q("Adds something the others don't: "), _q("Adds nothing significant once the others are in: ")
    f = (f'=IF({hy}="",{none_},{adds}&MID({hy},3,2000)&". ")&'
         f'IF({hn}="","",{adds_not}&MID({hn},3,2000)&".")')
    r = _line(ws, r, "What it found", f, chars=60 + 25 * len(adds_cells))
    helper_rows.append(r - 1)
    r += 1
    # the words the table can't hold: pairs that move together, groups with no odds ratio, candidates left out
    said = set()
    for t in j.fitted:
        for n, rho in t.partners:
            other = n.removesuffix(" (held fixed)")
            key = tuple(sorted((t.column, other)))
            if key in said:
                continue
            said.add(key)
            if n.endswith(" (held fixed)"):
                r = _line(ws, r, None, f"{t.column} moves with {other} ({rho:+.2f}), which is held fixed: what it "
                                       f"adds here is what it says beyond {other}.")
            else:
                r = _line(ws, r, None, f"{t.column} and {other} move together ({rho:+.2f}): what one adds, the other "
                                       f"mostly says already, so each can read as adding little while the pair "
                                       f"matters.")
        for k, why in t.not_estimable.items():
            r = _line(ws, r, None, f"{t.column}, {t.groups[k]}: no odds ratio together, because {why}.")
    for t in j.terms:
        if t.problem:
            r = _line(ws, r, None, f"{t.column} is left out of the candidates together: {t.problem}.")
    if j.left_out:
        r = _line(ws, r, None, "Held-back loans left out of the candidates together: "
                               + "; ".join(f"{v:,} with {k}" for k, v in j.left_out.items()) + ".")
    return r + 1, head


def write(wb, res, stamp: str = "") -> None:
    st = getattr(res, "prespec", None)
    tests = list(getattr(st, "tests", None) or ()) if st is not None else []
    if not tests:
        return
    from . import confirmatory, results
    for old in (SHEET, OLD_SHEET):
        if old in wb.sheetnames:
            del wb[old]
    ws = wb.create_sheet(SHEET)
    live.ensure(wb, res)
    for c, w in WIDTHS.items():
        ws.column_dimensions[_c(c)].width = w
    ran = [t for t in tests if not t.problem]
    first = ran[0] if ran else tests[0]
    held_w = _held_words(first.strata) or "the columns held fixed"
    verb = "are" if len(first.strata) != 1 else "is"
    what = first.column if len(tests) == 1 else f"each of the {len(tests)} candidates"
    house.title_band(ws, SHEET, f"Does {what} still tell good loans from bad on loans it was never found on, "
                                f"and once {held_w} {verb} held fixed?", FIRST, LAST, tab=house.TAB_RESULT)
    if not ran:
        r = house.method_note(ws, 3, FIRST, LAST, [("Not run", f"The confirmatory test couldn't be run: "
                                                              f"{first.problem}. Record says why.")])
        _finish(ws, r)
        return
    excesses = [confirmatory.excess(res, t) for t in ran]
    r = house.method_note(ws, 3, FIRST, LAST, _method(res, st, ran, excesses, stamp or "as of the last Run"))

    # ---------------------------------------------------------------- the table's rows, worked out first: the
    # Candidates tile counts the candidates holding up over them
    t0 = ran[0]
    held = _held_words(t0.strata)
    unit_dollars = any(e is not None and e.unit == "dollars" for e in excesses)
    raw_p = confirmatory.allowance(res) == "none"
    p_head = "p-value" if raw_p else P_ALLOWED

    # ---------------------------------------------------------------- what was tested: the tiles
    outcome = next((m.flag for m in res.measures if m.name == "outcome_loans"), "the outcome")
    _cell(ws, r, FIRST, "WHAT WAS TESTED", bold=True, color=SLATE, h="left", size=8)
    # a pre-spec changed after a held-back run says so here, in red, in place of where it came from: appended, the
    # words ran past the cell and "changed" was the part cut off (walk of 27 Sep 2026)
    changed = getattr(st, "changed_after", None) is not None
    _cell(ws, r + 1, FIRST, "changed after a held-back run" if changed else
          "from the shortlist scouting wrote" if getattr(st, "scouted", False) else "from the saved shortlist",
          color=house.CRIMSON if changed else SLATE, h="left", size=8)
    tile_row = r
    r += 3
    n = f'COUNTIF(Status,"{house.WAITING}")'
    _cell(ws, r, FIRST, f"The numbers are from the last Run, {stamp or 'as of the last Run'}. The verdicts, the "
                        f"materiality line and the chart's worse line follow Control.", color=SLATE, h="left", size=9)
    _cell(ws, r + 1, FIRST, f'=IFERROR(IF({n}=0,"","↻ "&{n}&IF({n}=1," Control change waits for a Run.",'
                            f'" Control changes wait for a Run.")),"")', bold=True, color=house.CRIMSON, h="left",
          size=9)
    r += 3

    # ---------------------------------------------------------------- OC-51: the tree unseen, then all together
    helper_rows: list[int] = []
    r = _unseen(ws, res, r)
    r, joint_head = _together(ws, res, st, r, helper_rows)
    if joint_head is not None or getattr(getattr(res, "scout", None), "oot", None) is not None:
        house.section(ws, r, FIRST, LAST, "Each candidate on its own" if len(ran) > 1 else
                      f"{t0.column} on the held-back loans")
        r += 2

    # ---------------------------------------------------------------- the table: one block of rows per candidate
    dev_n = t0.development.n if len({t.development.n for t in ran}) == 1 else None
    hold_n = t0.holdout.n if len({t.holdout.n for t in ran}) == 1 else None
    groups = [(N_FG, N_FP, f"Found · {dev_n:,} loans" if dev_n is not None else "Found", house.STONE, SLATE),
              (N_CG, N_HOLDS, f"Confirmed · {hold_n:,} held back" if hold_n is not None else
               "Confirmed on the held-back loans", house.INK, INK),
              (N_HG, N_STILL, f"Confirmed, {held} held fixed" if t0.held else "Nothing held fixed: the pre-spec holds "
                                                                             "no column fixed", house.INK, INK)]
    for a, b, text, rule, colour in groups:
        ws.merge_cells(start_row=r, start_column=a, end_row=r, end_column=b)
        _cell(ws, r, a, text, bold=True, color=colour, size=9)
        for c in (a, b):
            ws.cell(row=r, column=c).border = Border(bottom=Side(style="medium", color=rule))
    head = r + 1
    house.header(ws, head, FIRST, ["Candidate", "Compared", "Gap", p_head, "Gap", p_head, "Holds up?", "Gap",
                                   p_head, "Still holds?", "Excess $" if unit_dollars else "Excess bad loans",
                                   "Material?", "In words"], centre_from=2)
    ws.cell(row=head, column=N_WORDS).alignment = Alignment(horizontal="left", vertical="center")
    first_row = head + 1
    r = first_row
    blocks = []                                                # (test, first row, last row) per candidate
    for i, (t, excess) in enumerate(zip(ran, excesses)):
        start = r
        r = _rows(ws, t, excess, r, unit_dollars, raw_p, top=i > 0)
        blocks.append((t, start, r - 1))
    last_row = r - 1
    rng = lambda c: f"{_c(c)}{first_row}:{_c(c)}{last_row}"                 # noqa: E731
    for c in (N_HOLDS, N_STILL):
        ws.conditional_formatting.add(rng(c), FormulaRule(
            formula=[f'{_c(c)}{first_row}="{YES}"'], font=Font(bold=True, color=house.POSITIVE),
            fill=PatternFill("solid", fgColor=house.POSITIVE_BG, bgColor=house.POSITIVE_BG)))
        ws.conditional_formatting.add(rng(c), FormulaRule(formula=[f'{_c(c)}{first_row}="{NO}"'],
                                                          font=Font(bold=True, color=SLATE)))
    ws.conditional_formatting.add(rng(N_MAT), FormulaRule(
        formula=[f'{_c(N_MAT)}{first_row}="{YES}"'], font=Font(bold=True),
        fill=PatternFill("solid", fgColor=house.MIST, bgColor=house.MIST)))
    # a saved shortlist: found on another run (the spec); one scouting wrote was found on this Run's own loans
    found_hidden = st is not None and not st.scouted
    if found_hidden:
        for c in (N_FG, N_FP):
            ws.column_dimensions[_c(c)].hidden = True
    ws.freeze_panes = f"A{(joint_head or head) + 1}"

    # the tiles, now the rows exist: Candidates counts those holding up, live (tiles sit only over columns that are
    # never hidden: Found's two are, for a saved shortlist)
    if len(ran) == 1:
        cand = t0.column
    else:
        C, H = (f"${_c(c)}${first_row}:${_c(c)}${last_row}" for c in (N_CAND, N_HOLDS))
        up = "+".join(f'(COUNTIFS({C},{_q(t.column)},{H},"{YES}")>0)' for t in ran)
        cand = f'=IFERROR("{len(ran)} · "&({up})&" hold up","{len(ran)}")'
    loans = (f"{dev_n:,} found · {hold_n:,} held back" if dev_n is not None and hold_n is not None else
             f"up to {max(t.development.n for t in ran):,} found · {max(t.holdout.n for t in ran):,} held back")
    tiles = [("Outcome", outcome, N_COMP, N_COMP),
             ("Candidates" if len(ran) == 1 else "Candidates · live", cand, N_CG, N_HOLDS),
             ("Held fixed", " · ".join(t0.strata) or "Nothing", N_HG, N_STILL),
             ("Material at · live", results.material_at(res), N_EX, N_MAT),
             ("Loans", loans, N_WORDS, N_WORDS)]
    for label, value, a, b in tiles:
        house.tile(ws, tile_row, a, b, label, value, top=house.KEY_RED if label.startswith("Material") else house.INK)
    r += 1

    # ---------------------------------------------------------------- the chart, one per candidate
    for t, a, b in blocks:
        r = _chart(ws, t, a, b, r, found_hidden, many=len(ran) > 1)

    # ---------------------------------------------------------------- the tests in full, candidate by candidate
    for t in ran:
        views = _Views(t)
        house.section(ws, r, FIRST, LAST, "The tests in full" + (f": {t.column}" if len(ran) > 1 else ""))
        r += 2
        r = _matters(ws, t, views, r)
        r = _odds(ws, t, views, r, helper_rows)
        r = _concentration(ws, t, r, helper_rows)
    _finish(ws, r, helper_rows)


P_ALLOWED = "p, allowed"


def _rows(ws, t, excess, r: int, unit_dollars: bool, raw_p: bool, top: bool) -> int:
    """One candidate's block: a row per group against its reference, found, confirmed, and confirmed with the columns
    held fixed, each p-value after the allowance (the raw ones are in the tests in full). Returns the next row."""
    ref = t.groups[t.ref]
    held = _held_words(t.strata)

    def fit(side, k):
        if side is None:
            return None, None
        p = side.fit.p[k] if raw_p or not side.allowed else side.allowed[k]
        return (None if side.fit.odds[k] is None else side.fit.odds[k]), p
    key = live.q("gco_rate" if unit_dollars else "outcome_loans")
    line = f'IF(materiality_kind="none",0,INDEX(line_values,MATCH({key},line_keys,0)))'
    first = r
    for k, name in enumerate(t.groups):
        if k == t.ref:
            continue
        fg, fp = fit(t.development_plain, k)
        cg, cp = fit(t.holdout_plain, k)
        hg, hp = fit(t.holdout, k) if t.held else (None, None)
        vals = {N_CAND: t.column, N_COMP: f"{name} vs {ref}", N_FG: _num(fg), N_FP: fp, N_CG: _num(cg), N_CP: cp,
                N_HG: _num(hg), N_HP: hp,
                N_EX: None if excess is None or excess.per_group[k] is None else round(excess.per_group[k], 6)}
        for c, v in vals.items():
            fmt = X_FMT if c in (N_FG, N_CG, N_HG) else P_FMT if c in (N_FP, N_CP, N_HP) else \
                ('"$"#,##0;-"$"#,##0' if unit_dollars else "#,##0.0") if c == N_EX else None
            _cell(ws, r, c, v, bold=c in (N_CG, N_HG) or (c == N_CAND and r == first),
                  h="left" if c in (N_CAND,) else "center", fmt=fmt,
                  color=SLATE if c in (N_FG, N_FP, N_COMP) or (c == N_CAND and r != first) else INK)
        F, G, C_, P, H, J = (f"${_c(c)}{r}" for c in (N_FG, N_FP, N_CG, N_CP, N_HG, N_HP))

        def same(x):
            return f"IF(ISNUMBER({F}),({F}>1)=({x}>1),{x}>1)"
        _cell(ws, r, N_HOLDS, f'=IF(NOT(ISNUMBER({C_})),"",IF(AND({live.sig(P)},{same(C_)}),"{YES}","{NO}"))', bold=True)
        _cell(ws, r, N_STILL, f'=IF(NOT(ISNUMBER({H})),"",IF(AND({live.sig(J)},{same(H)}),"{YES}","{NO}"))',
              bold=True)
        L = f"${_c(N_EX)}{r}"
        _cell(ws, r, N_MAT, f'=IFERROR(IF(NOT(ISNUMBER({L})),"",IF({line}="","",IF(AND({L}>0,{L}>={line}),'
                            f'"{YES}","{NO}"))),"")')
        hu, sh = f"${_c(N_HOLDS)}{r}", f"${_c(N_STILL)}{r}"
        if t.held:
            words = (f'=IF({hu}="","",IF({hu}="{YES}",IF({sh}="{YES}",{_q(f"Holds up, and not just {held}")},'
                     f'{_q(f"Holds up, but it was mostly {held}")}),IF({sh}="{YES}",'
                     f'{_q(f"Holds up only with {held} held fixed")},"Didn\'t hold up on held-back loans")))')
        else:
            words = f'=IF({hu}="","",IF({hu}="{YES}","Holds up on held-back loans","Didn\'t hold up on held-back loans"))'
        _cell(ws, r, N_WORDS, words, h="left")
        ws.cell(row=r, column=LINE_COL, value="=worse_at")
        _rule(ws, r)
        if top and r == first:                                 # a heavier rule between one candidate and the next
            for c in range(FIRST, LAST + 1):
                ws.cell(row=r, column=c).border = Border(top=Side(style="medium", color=house.INK),
                                                         bottom=Side(style="thin", color=house.ROW_RULE))
        r += 1
    return r


def _chart(ws, t, first_row: int, last_row: int, r: int, found_hidden: bool, many: bool = False) -> int:
    """One candidate's bars (found STONE, confirmed INK, held fixed KEY_RED) and a dashed red line at the worse line,
    read from its block of the table, so it follows the table."""
    ch = BarChart()
    ch.type = "col"
    ch.grouping = "clustered"
    ref = t.groups[t.ref]
    ch.title = (f"{t.column}: " if many else "") + f"the gap against {ref} (dashed: the worse line)"
    ch.y_axis.title = "Odds ratio"
    ch.y_axis.number_format = '0.00"×"'
    ch.y_axis.majorGridlines = None
    cats = Reference(ws, min_col=N_COMP, min_row=first_row, max_row=last_row)
    series = [] if found_hidden else [(N_FG, "Found", house.STONE)]
    series += [(N_CG, "Confirmed", house.INK)]
    if t.held:
        series += [(N_HG, f"Confirmed, {_held_words(t.strata)} held fixed", house.KEY_RED)]
    for c, name, colour in series:
        ch.add_data(Reference(ws, min_col=c, min_row=first_row, max_row=last_row), titles_from_data=False)
        s = ch.series[-1]
        s.tx = SeriesLabel(v=name)
        s.graphicalProperties.solidFill = colour
        s.graphicalProperties.line.solidFill = colour
    ch.set_categories(cats)
    line = LineChart()
    line.add_data(Reference(ws, min_col=LINE_COL, min_row=first_row, max_row=last_row), titles_from_data=False)
    s = line.series[0]
    s.tx = SeriesLabel(v="Worse at")
    s.graphicalProperties.line.solidFill = house.KEY_RED
    s.graphicalProperties.line.dashStyle = "dash"
    s.graphicalProperties.line.width = 19050
    s.marker.symbol = "none"
    s.smooth = False
    line.set_categories(cats)
    ch += line
    ch.legend.position = "b"
    ch.height, ch.width = (8.5, 26) if not many else (7.0, 26)
    ch.visible_cells_only = False                              # the worse line sits in a hidden column
    ws.add_chart(ch, f"{_c(FIRST)}{r}")
    return r + (19 if not many else 16)


def _matters(ws, t, views, r: int) -> int:
    """B3, B4 and the block test on every set of loans: does the column matter at all?"""
    K = len(t.groups)
    _cell(ws, r, FIRST, f"Does {t.column} matter?", bold=True, h="left", size=11)
    r += 1
    house.header(ws, r, FIRST, ["Test", None, None, None, "Statistic", "Degrees of freedom", "p-value", "Reading",
                                None], centre_from=2)
    ws.merge_cells(start_row=r, start_column=FIRST, end_row=r, end_column=FIRST + 1)
    _wrap(ws, r)
    r += 1
    tests = (("Any difference across the groups (general)", lambda s: (s.association.general, s.association.df,
                                                                       s.association.p_general)),
             ("A steady climb or fall (trend)", lambda s: (s.association.trend, 1 if s.association.trend is not None
                                                           else None, s.association.p_trend)),
             ("Any difference, from the regression (block test)", lambda s: (s.fit.block, s.fit.df, s.fit.p_block)))
    for side, where, lead in views.all:
        house.sub_header(ws, r, FIRST, [f"{where}: loans made {side.range.text()}"])
        for c in range(FIRST + 1, LAST + 1):
            ws.cell(row=r, column=c).fill = house.fill(house.CANVAS)
        # N1: the caption runs across the table in one merged cell, rather than widening the first column
        ws.merge_cells(start_row=r, start_column=FIRST, end_row=r, end_column=LAST)
        r += 1
        cells = []
        for label, get in tests:
            ws.merge_cells(start_row=r, start_column=FIRST, end_row=r, end_column=FIRST + 1)
            _cell(ws, r, FIRST, label, h="left")
            stat, df, p = get(side)
            _cell(ws, r, DATA, None if stat is None else round(stat, 6), fmt="0.00")
            _cell(ws, r, DATA + 1, df, fmt="0")
            _cell(ws, r, DATA + 2, p, fmt=P_FMT)
            ref_p = f"{_c(DATA + 2)}{r}"
            ws.merge_cells(start_row=r, start_column=DATA + 3, end_row=r, end_column=DATA + 4)
            _cell(ws, r, DATA + 3, f"={_sig_word(ref_p)}")
            cells.append(ref_p)
            _rule(ws, r, FIRST, DATA + 4)
            r += 1
        a = side.association
        g, tr = cells[0], cells[1]
        way = "rises" if a.direction > 0 else "falls"
        f = (f'=IF(AND({live.sig(g)},{live.sig(tr)}),{_q(f"{lead} the bad rate differs across the groups, and {way} steadily as {t.column} rises.")},'
             f'IF({live.sig(g)},{_q(f"{lead} the bad rate differs across the groups, but not in one direction: higher in some groups, lower in others.")},'
             f'IF({live.sig(tr)},{_q(f"{lead} the groups taken all at once do not differ significantly, but the bad rate {way} steadily as {t.column} rises (the trend test).")},'
             f'{_q(f"{lead} no difference across the groups shows at ")}&TEXT(confidence,"0%")&{_q(" sure. That is not proof there is none.")})))')
        r = _line(ws, r, "What it found", _empty(side, lead) or f, chars=150)
        r += 1
    return r


def _odds(ws, t, views, r: int, helper_rows: list) -> int:
    """B5: each group's odds ratio against the reference, with its range and p-value, on each set of loans, with
    nothing held fixed and with the columns held fixed side by side."""
    ref = t.groups[t.ref]
    K = len(t.groups)
    _cell(ws, r, FIRST, f"How much more often does each group go bad than {ref}?", bold=True, h="left", size=11)
    r += 1
    rng_head = live.text("Range (", ('TEXT(confidence,"0%")',), " sure)")
    held = _held_words(t.strata)
    heads = [f"Group of {t.column}", None, None, None, "Loans", "Bad loans", "Bad rate", "Odds ratio", "p-value",
             rng_head]
    if t.held:
        heads += [f"Odds ratio, {held} held fixed", "p-value", rng_head]
    ranges = ((t.development_plain, t.development, "Found", "On development,"),
              (t.holdout_plain, t.holdout, "Confirmed on the held-back loans", "On the holdout,"))
    for plain, side, where, lead in ranges:
        house.header(ws, r, FIRST, heads, centre_from=2)
        ws.merge_cells(start_row=r, start_column=FIRST, end_row=r, end_column=FIRST + 1)
        _wrap(ws, r)
        r += 1
        house.sub_header(ws, r, FIRST, [f"{where}: loans made {side.range.text()}"])
        for c in range(FIRST + 1, LAST + 1):
            ws.cell(row=r, column=c).fill = house.fill(house.CANVAS)
        r += 1
        first_group = r
        cells: dict = {}
        fits = [(plain.fit, DATA + 3, "plain"), (side.fit, DATA + 6, "held")] if t.held else \
            [(side.fit, DATA + 3, "held")]
        for k, name in enumerate(t.groups):
            ws.merge_cells(start_row=r, start_column=FIRST, end_row=r, end_column=FIRST + 1)
            _cell(ws, r, FIRST, name + (" (reference)" if k == t.ref else ""), h="left")
            _cell(ws, r, DATA, side.loans[k], fmt="#,##0")
            _cell(ws, r, DATA + 1, side.bad[k], fmt="#,##0")
            _cell(ws, r, DATA + 2, side.bad[k] / side.loans[k] if side.loans[k] else None, fmt="0.00%")
            for f, c, which in fits:
                if k == t.ref:
                    _cell(ws, r, c, "1.00")
                elif f.odds[k] is None:
                    _cell(ws, r, c, "none")
                else:
                    _cell(ws, r, c, round(f.odds[k], 10), fmt=X_FMT)
                    b, se = live.num(f.beta[k]), live.num(f.se[k])
                    _cell(ws, r, c + 1, f.p[k], fmt=P_FMT)
                    _cell(ws, r, c + 2, (f'=TEXT(EXP({b}-{Z}*{se}),"0.00")&"x to "&'
                                               f'TEXT(EXP({b}+{Z}*{se}),"0.00")&"x"'))
                    cells[(which, k)] = (f"{_c(c)}{r}", f"{_c(c + 1)}{r}", b, se)
            _rule(ws, r)
            r += 1
        for f, c, which in fits:
            # a significant odds ratio shaded, red above 1 and green below, against the one bar (OC-40)
            o, p = f"${_c(c)}{first_group}", f"${_c(c + 1)}{first_group}"
            for fill, cmp_ in ((WORSE_FILL, ">1"), (BETTER_FILL, "<1")):
                ws.conditional_formatting.add(f"{_c(c)}{first_group}:{_c(c + 2)}{r - 1}", FormulaRule(
                    formula=[f"AND(ISNUMBER({o}),ISNUMBER({p}),{p}<{live.BAR},{o}{cmp_})"],
                    fill=PatternFill("solid", fgColor=fill, bgColor=fill)))
        said_lines = [("plain", f"{lead[:-1]}, with nothing held fixed,"), ("held", lead)] if t.held else \
            [("held", lead)]
        for which, lead_ in said_lines:
            pieces = []
            for k, name in enumerate(t.groups):
                got = cells.get((which, k))
                if got is None:
                    continue
                o, p, b, se = got
                rng = f'" ("&TEXT(EXP({b}-{Z}*{se}),"0.00")&"x to "&TEXT(EXP({b}+{Z}*{se}),"0.00")&"x)"'
                said = (f'IF({o}>=1,{_q(f"; {name} goes bad ")}&TEXT({o},"0.00")&{_q(f" times as often as {ref}")},'
                        f'{_q(f"; {name} goes bad less often than {ref}, ")}&TEXT({o},"0.00")&" times as often")')
                pieces.append(f'IF({live.sig(p)},{said}&{rng},"")')
            h = _helper(ws, r, "&".join(pieces) if pieces else '""')
            tail = ", with the pockets held fixed." if which == "held" else "."
            f = (f'=IF({h}="",{_q(f"{lead_} no group differs from {ref} at ")}&TEXT(confidence,"0%")&'
                 f'{_q(" sure" + tail)},{_q(lead_ + " ")}&MID({h},3,2000)&{_q(tail)})')
            src = plain if which == "plain" else side
            r = _line(ws, r, "What it found", _empty(src, lead_) or f, chars=60 + 80 * (K - 1))
            helper_rows.append(r - 1)
        named = [(plain.fit, "nothing held fixed"), (side.fit, f"{held} held fixed")] if t.held else \
            [(side.fit, "the pockets held fixed")]
        for fit_, name_ in named:
            for k, why in fit_.not_estimable.items():
                r = _line(ws, r, None, f"{where}, {name_}, {t.groups[k]}: no odds ratio, because {why}.")
            if not fit_.settled:
                r = _line(ws, r, None, f"{where}, {name_}: the regression didn't settle on an answer, so its odds "
                                       f"ratios are not reliable.")
        r += 1
    return r


def _concentration(ws, t, r: int, helper_rows: list) -> int:
    """4e: how much of the holdout's losses sit in each group, on the holdout only."""
    K = len(t.groups)
    ref = t.groups[t.ref]
    hold = t.holdout
    _cell(ws, r, FIRST, "How much of the held-back loans' losses sit in each group?", bold=True, h="left", size=11)
    r += 1
    # the dollar columns only when the run has GCO: a test of a new variable needs none (Goal 2 item 2), and then
    # the table says what it does hold and nothing about dollars (the firm: "don't note what it does not include,
    # just note what it does")
    cols = [("Loans", "#,##0"), ("Share of loans", "0.0%"), ("Bad loans", "#,##0"), ("Share of bad loans", "0.0%")]
    if t.dollars:
        cols += [("GCO", "$#,##0"), ("Share of GCO", "0.0%")]
    cols += [("Bad rate", "0.00%"), ("Times the holdout's bad rate", '0.00"x"')]
    heads, fmts = [h for h, _ in cols], tuple(f for _, f in cols)
    at = {h: CONC_FIRST + j for j, h in enumerate(heads)}          # each column by its heading
    house.header(ws, r, FIRST, [f"Group of {t.column}", None, None, None] + heads, centre_from=2)
    ws.merge_cells(start_row=r, start_column=FIRST, end_row=r, end_column=FIRST + 1)
    _wrap(ws, r)
    r += 1
    conc = t.concentration()

    def keep(loans, flag, bad, capture, gco, gco_capture, rate, lift):
        return (loans, flag, bad, capture) + ((gco, gco_capture) if t.dollars else ()) + (rate, lift)

    lift_col = at["Times the holdout's bad rate"]
    crow = {}
    for k, name in enumerate(t.groups):
        x = conc[k]
        ws.merge_cells(start_row=r, start_column=FIRST, end_row=r, end_column=FIRST + 1)
        _cell(ws, r, FIRST, name + (" (reference)" if k == t.ref else ""), h="left")
        for j, (v, fm) in enumerate(zip(keep(x.loans, x.flag_rate, x.bad, x.capture, x.gco, x.gco_capture,
                                             x.bad_rate, x.lift), fmts)):
            _cell(ws, r, CONC_FIRST + j, v, fmt=fm)
        crow[k] = r
        _rule(ws, r)
        r += 1
    total_gco = math.fsum(g for g in hold.gco_of if g is not None)
    vals = keep(hold.n, 1.0 if hold.n else None, hold.n_bad, 1.0 if hold.n_bad else None, total_gco,
                1.0 if total_gco else None, hold.n_bad / hold.n if hold.n else None, 1.0 if hold.n else None)
    ws.merge_cells(start_row=r, start_column=FIRST, end_row=r, end_column=FIRST + 1)
    _cell(ws, r, FIRST, "All held-back loans in the test", bold=True, h="left")
    for j, (v, fm) in enumerate(zip(vals, fmts)):
        _cell(ws, r, CONC_FIRST + j, v, bold=True, fmt=fm)
    r += 2
    pieces = []
    for k, name in enumerate(t.groups):
        f = t.holdout.fit
        if k == t.ref or f.odds[k] is None:
            continue
        rr = crow[k]
        o, p = f"{live.num(f.odds[k])}", live.num(f.p[k]) if f.p[k] is not None else '""'
        share = lambda h: f'TEXT({_c(at[h])}{rr},"0.0%")'          # noqa: E731
        held = (f'{share("Share of loans")}&{_q(" of the loans, ")}&{share("Share of bad loans")}&'
                + (f'{_q(" of the bad loans and ")}&{share("Share of GCO")}&{_q(" of the GCO, at ")}' if t.dollars
                   else f'{_q(" of the bad loans, at ")}'))
        pieces.append(f'IF(AND({live.sig(p)},{o}>1),{_q(f"; {name} holds ")}&{held}&'
                      f'TEXT({_c(lift_col)}{rr},"0.00")&{_q(LIFT_TAIL)},"")')
    h = _helper(ws, r, "&".join(pieces) if pieces else '""')
    f = (f'=IF({h}="",{_q(f"On the holdout, no group goes bad significantly more often than {ref}, so none is singled out; every group is in the table.")},'
         f'{_q("On the holdout, ")}&MID({h},3,2000)&".")')
    r = _line(ws, r, "What it found", _empty(hold, "On the holdout,") or f, chars=40 + 125 * (K - 1))
    helper_rows.append(r - 1)
    if t.gco_unread:
        r = _line(ws, r, None, f"{t.gco_unread:,} holdout loans have no readable GCO and are left out of the GCO "
                               f"figures only.")
    return r + 1


CONC_FIRST = DATA                                          # the concentration table's first number column


def _wrap(ws, r: int) -> None:
    """A header row of the tests in full on two lines: their headings are longer than the table's columns."""
    for c in range(FIRST, LAST + 1):
        x = ws.cell(row=r, column=c)
        x.alignment = Alignment(horizontal=x.alignment.horizontal, vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30


def _empty(side, lead: str) -> str | None:
    """What a range's plain line says when nothing on it can be tested; None when something can."""
    if side.association.pockets:
        return None
    if not side.n:
        return f"{lead} nothing to test: no loan in this test was made {side.range.text()}."
    return f"{lead} nothing to test: no pocket holds both a bad loan and a good one."


def _helper(ws, r: int, expr: str) -> str:
    c = ws.cell(row=r, column=HELPER, value=f"={expr}")
    c.alignment = Alignment(wrap_text=False)
    return f"${_c(HELPER)}${r}"


def _finish(ws, last_row: int, helper_rows=()) -> None:
    for c in (LINE_COL, HELPER):
        ws.column_dimensions[_c(c)].hidden = True
    ws.print_area = f"B1:{_c(LAST)}{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# --------------------------------------------------------------------------
# Start here's "What the last Run found" for a test of a new variable (OC-42): it builds no pocket, so the block
# reads the confirmation. The Run's numbers go on the hidden _found sheet, one row per group; the tiles count them
# against significance_bar, so "significant" follows the confidence on Control as the tab's own words do.

FOUND_KIND = "confirm"
G_NAME, G_LOANS, G_BAD, G_RATE, G_ODDS, G_P, G_CAPTURE, G_REF = range(2, 10)     # _found's columns B to I
G_CAND = 10                                                                      # J: the candidate (Goal 2 item 3)


def write_found(ws, res) -> None:
    """What this test found, on _found (ws), under the stamp: one row per group of every candidate, its p-value after
    the allowance for testing many at once."""
    from . import confirmatory
    h = confirmatory.headline(res)
    ws.append(["kind", FOUND_KIND])
    if h.get("problem"):
        ws.append(["problem", h["problem"]])
        return
    ws.append(["column", h["column"]])
    ws.append(["candidates", ", ".join(h["candidates"])])
    ws.append(["reference", h["reference"]])
    # the New variables tab's own words; "5,604 on development, 2,395 on the holdout" ran into the next tile on
    # Start here (walk of 27 Sep 2026)
    ws.append(["loans", f"{h['development']:,} found · {h['holdout']:,} held back"])
    said, under, _ = confirmatory.follows_words(h)
    ws.append(["follows", said if said == "Yes" else f"{said}: {under}"])
    # the odds on the holdout are worked out with the pre-spec's columns held fixed: Start here's header says so (I)
    t0 = next(iter(confirmatory._ran(getattr(res, "prespec", None))), None)
    ws.append(["held", _held_words(t0.strata) if t0 is not None and t0.held else ""])
    for g in confirmatory.groups_found(res):
        ws.append(["group", g["group"], g["loans"], g["bad"], g["bad_rate"], g["odds"], g["p"], g["capture"],
                   "yes" if g["ref"] else None, g["candidate"]])


def found_block(ws, wb, r: int, value_of) -> int:
    """Four tiles and one row per group, on the holdout. value_of(key) reads _found. Returns the next free row. With
    several candidates, the rows run candidate by candidate, and the share tile, which would add up groups of
    different candidates that hold the same loans, counts the candidates holding up instead."""
    from . import house
    from .book import FOUND
    note = Font(name="Calibri", size=10, color=SLATE)
    problem = value_of("problem")
    if problem:
        ws.cell(row=r + 1, column=2, value=f"The confirmatory test couldn't be run: {problem}. Record says why.").font \
            = note
        return r + 3
    ref = value_of("reference")
    cands = [c for c in str(value_of("candidates") or value_of("column") or "").split(", ") if c]
    many = len(cands) > 1
    F = lambda c: f"'{FOUND}'!${_c(c)}:${_c(c)}"                           # noqa: E731
    worse = f'COUNTIFS({F(1)},"group",{F(G_ODDS)},">1",{F(G_P)},"<"&{live.BAR})'
    groups = sum(1 for row in wb[FOUND].iter_rows(min_row=1, max_col=1, values_only=True) if row[0] == "group")
    if not many:
        tiles = [(f"Groups worse than {ref}, on the holdout", f'=IFERROR({worse}&" of {groups - 1}","")', None),
                 ("Their share of the holdout's bad loans",
                  f'=IFERROR(SUMIFS({F(G_CAPTURE)},{F(1)},"group",{F(G_ODDS)},">1",{F(G_P)},"<"&{live.BAR}),"")',
                  "0%")]
    else:
        up = "+".join(f'(COUNTIFS({F(1)},"group",{F(G_CAND)},{live.q(c)},{F(G_ODDS)},">1",{F(G_P)},"<"&{live.BAR})>0)'
                      for c in cands)
        tiles = [("Groups worse than their reference, on the holdout",
                  f'=IFERROR({worse}&" of {groups - len(cands)}","")', None),
                 ("Candidates with a group worse", f'=IFERROR(({up})&" of {len(cands)}","")', None)]
    tiles += [("Loans tested", value_of("loans"), None), ("Follows the pre-spec", value_of("follows"), None)]
    for i, (label, f, fmt) in enumerate(tiles):
        house.tile(ws, r + 1, 2 + 2 * i, 3 + 2 * i, label, f, fmt)
    t = r + 4
    held = value_of("held")
    odds = ("× its reference's odds" if many else f"× {ref}'s odds") + (f", {held} held fixed" if held else "")
    house.header(ws, t, 2, ["Candidate: group" if many else f"Group of {value_of('column')}", "Loans", "Bad rate",
                            odds, "p-value", "Significant?",
                            "Share of bad loans", None], centre_from=2)
    if held:
        # the longer odds heading wraps in its column rather than running under its neighbours
        c = ws.cell(row=t, column=5)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[t].height = 40
    rows = [row for row in wb[FOUND].iter_rows(min_row=1) if row[0].value == "group"]
    for i, row in enumerate(rows, start=1):
        rr, src = t + i, row[0].row
        P = lambda c: f"'{FOUND}'!${_c(c)}${src}"                          # noqa: E731
        ref_row = row[G_REF - 1].value == "yes"
        name = (f'={P(G_CAND)}&": "&{P(G_NAME)}' if many else f"={P(G_NAME)}") + ('&" (reference)"' if ref_row else "")
        vals = [name, f"={P(G_LOANS)}",
                f'=IF({P(G_RATE)}="","",{P(G_RATE)})', f'=IF({P(G_ODDS)}="","none",{P(G_ODDS)})',
                f'=IF({P(G_P)}="","",{P(G_P)})',
                None if ref_row else f'=IF({live.sig(P(G_P))},IF({P(G_ODDS)}>1,"Yes, worse","Yes, better"),"No")',
                f"={P(G_CAPTURE)}"]
        for j, v in enumerate(vals):
            c = ws.cell(row=rr, column=2 + j, value=v)
            c.font = Font(name="Calibri", size=10, color=house.INK_TEXT)
            c.alignment = Alignment(horizontal="left" if j == 0 else "center", vertical="center")
        for j, fmt in ((3, "#,##0"), (4, "0.00%"), (5, '0.00"×"'), (6, P_FMT), (8, "0%")):
            ws.cell(row=rr, column=j).number_format = fmt
        ws.row_dimensions[rr].height = 18
    last = t + max(1, len(rows))
    ws.conditional_formatting.add(f"G{t + 1}:G{last}", FormulaRule(
        formula=[f'G{t + 1}="Yes, worse"'], font=Font(bold=True, color=house.CRIMSON),
        fill=PatternFill("solid", fgColor=house.ALERT_FG, bgColor=house.ALERT_FG)))
    link = ws.cell(row=last + 1, column=2, value="Found beside confirmed, with and without the columns held fixed, "
                                                 "and how it's worked out: the New variables tab.")
    link.hyperlink = f"#'{SHEET}'!A1"
    link.font = Font(name="Calibri", size=10, color=house.KEY_RED, underline="single")
    return last + 3
