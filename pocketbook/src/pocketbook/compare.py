"""The Compare tab: the filters' values as lines across one band column's bands (or across the years), in small
charts side by side on one scale.

The firm, 30 Sep 2026: *"can we make it so they can be visually compared in a graph? Like if we used origination
date as a filter it would essentially be vintage years"*, with lines (one per value) across FICO bands; then *"if we
are proving things exist across categories it should not be vintage analysis only so let's make sure that is the
case and how would we show that say vintage analysis mixed with like underwriter/system approved?"* A line chart
was chosen over bars.

What is across the bottom: any band column the Run cut, or either filter picked in the launcher (Origination year
across the bottom is the vintage view). The lines are by either filter (Filter 1 or Filter 2), one line per value;
the panels by the other, or none. With a filter across the bottom, the lines are by the other filter and there is one
chart. A third filter (the firm, 1 Oct 2026: "I thought we discussed two filters plus date") joins both pickers:
Lines by and Panels by each offer all three, and the filter picked in neither is All loans. Every number is worked out in the Run (engine.summaries, the same cells Summary shows) and put on _views; the
tab's formulas only pick them, as every result tab's do, so the dropdowns redraw the charts live.

How the charts are built, so they work alike in Excel and LibreOffice:
- Each chart's series read a hidden helper block on this tab whose cells change with the dropdowns (as Look's live
  charts do). A point whose cell has fewer loans than the Run's Fewest loans in a pocket is #N/A there, so it is
  not drawn.
- One scale for every panel, live. A chart's axis bounds can't be set by formula, and bounds fixed at write time
  could not follow the Measure dropdown (a bad-loan rate near 10% and RANR near -2% can't share fixed bounds). So
  every panel carries one more, invisible series: the smallest and the largest value plotted in ANY panel for what
  is picked now. Each chart scales itself to take in those two points, the same two in every panel, so every panel
  lands on the same axis, whatever is picked. The trade-off: the bounds are the programs' own rounding of that
  range, not numbers we choose.
"""

from __future__ import annotations

from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Font

from . import engine, house, live
from .choices import ORIG_YEAR, ORIG_YEAR_LABEL
from .results import (AND, AND3, CANVAS, COMPARE, WHERE, Choices, Views, YEAR_SAID, _cell, _fit, _hide, _widths, cf,
                      col, dropdown, fewest, match, pick, plain, rates, view_key, xk)

SHEET = COMPARE
#: the rows that aren't bands on the x axis: loans the column couldn't place, and loans with no date
SPECIAL = {engine.BLANK_LABEL, engine.NOT_NUMBER_LABEL, engine.MISSING_RULE_LABEL, engine.NO_DATE}
NONE = "None"
BOOK_LINE = "Whole book"
#: one hue per line, in a fixed order, never cycled (the reference palette, validated for colour-blind readers); the
#: whole book is a dashed grey line of its own
HUES = ("2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948")
PANEL_COLS = 6              # columns a panel's chart spans
CHART_ROWS = 17             # rows it covers
DATA_W = 11                 # the width of every column but the first
ACROSS = "Across the bottom"
SAY_THIN = "A point with fewer loans than the {few:,} in Fewest loans in a pocket on Control is left off its line."
SAY_SAME = "Panels by is the filter the lines are by, so one chart is drawn. Pick the other filter, or None."
SAY_SAME3 = "Panels by is the filter the lines are by, so one chart is drawn. Pick another filter, or None."
SAY_ACROSS = "{x} is across the bottom, so the lines are by {other} and one chart is drawn."
SAY_READ = "If the lines have the same shape in every panel, the pattern holds across both categories."
#: the helper block's three columns for each line: its y, the same to scale by, its x
Y_COL, CLEAN, X_COL = 0, 1, 2
#: how far under the smallest value the labels across sit, as a share of the values' range
LABEL_DROP = 0.12
#: where across the scale's two invisible points sit: on the axis's left end, clear of every label
SCALE_X = 0.5
SAY_ASOF = "As of the last Run: the numbers are the Run's, and Fewest loans in a pocket is the one it used."


def band_options(res) -> tuple[list[str], list[str]]:
    """Every band column the Run cut (engine name) and how the dropdown names it: its column, as Summary does."""
    from . import book as bk
    names = bk._names(res)
    bands = [b.name for b in res.config.bands if (b.name, None, None, None) in res.summaries]
    shown = [names[b] for b in bands]
    return bands, [x if shown.count(x) == 1 else f"{x} ({b})" for x, b in zip(shown, bands)]


def filters(res) -> list[tuple[str, list[str]]]:
    """The filters the Run used, in order, each (column, its values): Filter 1, then 2 and 3 when picked."""
    out = []
    for name, vals in ((res.config.filter_by, res.filter_values), (res.config.filter_by2, res.filter_values2),
                       (res.config.filter_by3, getattr(res, "filter_values3", []))):
        if not (name and vals):
            break
        out.append((name, list(vals)))
    return out


def filter_options(res) -> list[tuple[int, str]]:
    """The filters that can go across the bottom, as (1, 2 or 3, how the dropdown names it): every one, when the
    launcher picked two or three and the Run cut a band column; Origination year by its name. None with one filter:
    across the bottom and as the lines too, it would be one line."""
    fs = filters(res)
    if len(fs) < 2 or not band_options(res)[0]:
        return []
    taken = set(band_options(res)[1])
    out = []
    for k, (name, _) in enumerate(fs, start=1):
        shown = ORIG_YEAR_LABEL if name == ORIG_YEAR else name
        out.append((k, shown if shown not in taken else f"{shown} (filter)"))
    return out


def band_labels(res, band: str) -> list[str]:
    """The bands of one column in order: the x axis. The loans it couldn't place and All are left off."""
    return [x for x in res.summaries[(band, None, None, None)].labels if x != engine.ALL and x not in SPECIAL]


def filter_labels(res, k: int) -> list[str]:
    """Filter k's values across the bottom, in order: (blank) and (no date) are left off, as a band column's are."""
    return [v for v in filters(res)[k - 1][1] if v not in SPECIAL]


def other(k: int) -> int:
    """The filter the lines are by when filter k is across the bottom and the lines were picked by it too."""
    return 2 if k == 1 else 1


def points(res, band: str, measure: str, v1: str | None, v2: str | None,
           v3: str | None = None) -> list[tuple[float | None, int]]:
    """Each band's (rate, loans) on the loans with Filter 1 = v1, Filter 2 = v2 and Filter 3 = v3 (None: All loans),
    as _views holds them: (None, 0) where the view has no loans."""
    s = res.summaries.get((band, v1, v2, v3))
    out = []
    for lab in band_labels(res, band):
        c = s.cells.get(lab) if s is not None else None
        out.append((c.rates[measure].rate, c.rows) if c is not None and c.rows else (None, 0))
    return out


def filter_points(res, k: int, measure: str, w: str | None, j: int | None = None) -> list[tuple[float | None, int]]:
    """Filter k across the bottom: each of its values' (rate, loans) on the loans with filter j (the other one when
    there are two) = w (None: All loans), any third filter All loans. Read off any band column's Summary, whose All
    row is every loan of the view."""
    b0 = band_options(res)[0][0]
    j = other(k) if j is None else j
    out = []
    for v in filter_labels(res, k):
        at = {k: v, j: w}
        s = res.summaries.get((b0, at.get(1), at.get(2), at.get(3)))
        c = s.cells[engine.ALL] if s is not None else None
        out.append((c.rates[measure].rate, c.rows) if c is not None and c.rows else (None, 0))
    return out


def compare_views(res, views: Views) -> dict:
    """Every Compare view on _views: "C|<across>|labels", then for every filtered view (and the whole book) the
    loans at each point ("C|<across>|loans<view>") and each rate ("C|<across>|<measure><view>"). With a filter
    across the bottom, the views are the other filter's values, keyed as the lines by it are."""
    bands, shown = band_options(res)
    ms = rates(res)
    most = 2

    def put(opt, key, got):
        views.put(f"C|{opt}|loans{key}", [n for _, n in got[ms[0].name]] if ms else [])
        for m in ms:
            views.put(f"C|{opt}|{m.name}{key}", [r for r, _ in got[m.name]])

    for b, opt in zip(bands, shown):
        labels = band_labels(res, b)
        most = max(most, len(labels))
        views.put(f"C|{opt}|labels", labels)
        for (bb, v1, v2, v3) in [k for k in res.summaries if k[0] == b]:
            put(opt, view_key(v1, v2, v3), {m.name: points(res, b, m.name, v1, v2, v3) for m in ms})
    fs = filters(res)
    for k, opt in filter_options(res):
        labels = filter_labels(res, k)
        most = max(most, len(labels))
        views.put(f"C|{opt}|labels", labels)
        put(opt, "", {m.name: filter_points(res, k, m.name, None) for m in ms})      # the whole book
        for j in range(1, len(fs) + 1):
            if j == k:
                continue
            for w in fs[j - 1][1]:                                  # the lines by another filter, one per value
                at = {j: w}
                put(opt, view_key(at.get(1), at.get(2), at.get(3)),
                    {m.name: filter_points(res, k, m.name, w, j) for m in ms})
    return {"options": shown + [o for _, o in filter_options(res)], "most": most, "measures": ms}


def write_compare(wb, res, choices: Choices, views: Views) -> None:
    """Compare (after Summary): pickers for what is across the bottom, the measure, what the lines are by and what
    the panels are by; one small chart per panel, side by side on one scale; and under them every point's rate and
    loans. Written only when the launcher picked a Filter by and the Run cut a band column."""
    fs = filters(res)
    if not (fs and band_options(res)[0]):
        return
    names = [n for n, _ in fs]
    sf, sf2 = names[0], (names[1] if len(fs) > 1 else None)
    got = compare_views(res, views)
    ms, MB = got["measures"], got["most"]
    xf = filter_options(res)
    vals = [v for _, v in fs]
    L = max(len(v) for v in vals)                           # line slots
    P = max(len(v) for v in vals) if sf2 else 1            # panel slots
    few = fewest(res)
    ws = wb.create_sheet(SHEET)
    left = 2
    table_w = 1 + 2 * (L + 1)
    last = max(left + P * PANEL_COLS - 1, left + table_w - 1, left + 18)
    S0 = last + 3                                           # hidden: the picks worked out, then the helper block
    HB = S0 + 8
    # the point's label, its x and y (the labels drawn in the chart), the scale's two points' x and y, then the lines
    CATC, LY, PHY, HS = HB, HB + 1, HB + 2, HB + 3
    n_series = P * (L + 1)

    def scol(p: int, j: int, kind: int = Y_COL) -> int:
        """The helper column of panel p's line j (0 the whole book): its y (the rate drawn, or #N/A), the same as a
        number or blank (to scale by), or its x (the point's place across, or #N/A)."""
        return HS + kind * n_series + (p - 1) * (L + 1) + j

    def pcol(p: int, what: int) -> int:
        """Panel p's own x and y for the labels across (0, 2) and for the scale's two points (1, 3): #N/A in a panel
        not drawn, both of them, since LibreOffice draws a point with a y and no x at the first place across."""
        return HS + 3 * n_series + what * P + p - 1
    texts = [str(x) for b in band_options(res)[0] for x in band_labels(res, b)]
    texts += [str(x) for k, _ in xf for x in filter_labels(res, k)]
    lw = house.fit(texts + [BOOK_LINE, "All loans"], floor=12, cap=28, pad=3)
    _widths(ws, {1: 2, left: lw, **{c: DATA_W for c in range(left + 1, S0)}})
    house.title_band(ws, SHEET, "The filters' values as lines across the bands, side by side on one scale.", left,
                     last, tab=house.TAB_RESULT)
    each = "either" if len(fs) == 2 else "any"
    note = [
        ("What it is", "Pick what goes across the bottom and a measure. Each line is one value of the filter picked "
                       "in Lines by. The dashed grey line is the whole book."),
        ("Across", "A band column's bands" + (f", or {each} filter's values: Origination year there is the vintage "
                                              "view" if any(n == ORIG_YEAR_LABEL for _, n in xf) else
                                              f", or {each} filter's values" if xf else "") + "."),
        ("Panels", ("Panels by draws one small chart for each value of the other filter, side by side. Every panel "
                    "has the same scale, so higher in one panel is higher in all.") if len(fs) == 2 else
                   ("Panels by draws one small chart for each value of another filter, side by side; the filter in "
                    "neither Lines by nor Panels by is All loans. Every panel has the same scale, so higher in one "
                    "panel is higher in all.") if sf2 else
                   "Pick a second filter in the launcher (Filter 2) to draw one small chart for each of its values."),
        ("Thin points", (SAY_THIN.format(few=few) if few is not None else "Every point is drawn.")
                        + " The table under the charts still shows it, in grey. A panel with nothing to show for "
                          "what is picked stays empty, on the same scale."),
        ("Reading it", SAY_READ + " Nothing here is tested. " + SAY_ASOF),
    ]
    if ORIG_YEAR in names:
        note.append(("Years", YEAR_SAID.format(res.config.origination_date).strip()))
    r = house.method_note(ws, 3, left, last, note)
    across_opt = got["options"]
    lines_opts = [f"Filter {k}: {n}" for k, n in enumerate(names, start=1)]
    opt1, opt2 = lines_opts[0], (lines_opts[1] if sf2 else None)
    panel_opts = [NONE] + lines_opts
    b_rng, _ = choices.add("Compare: Across the bottom", across_opt)
    m_rng, (k_rng,) = choices.add("Compare: Measure", [plain(m) for m in ms], [m.name for m in ms])
    l_rng, _ = choices.add("Compare: Lines by", lines_opts)
    p_rng, _ = choices.add("Compare: Panels by", panel_opts)
    y_rng, _ = choices.add("Compare: Whole book line", ["Yes", "No"])
    s = r + 1
    B = dropdown(ws, s, left, ACROSS, b_rng, across_opt[0])
    M = dropdown(ws, s, left + 3, "Measure", m_rng, plain(ms[0]) if ms else "")
    LB_ = dropdown(ws, s, left + 6, "Lines by", l_rng, opt1)
    PB_ = dropdown(ws, s, left + 9, "Panels by", p_rng, opt2 or NONE)
    Y = dropdown(ws, s, left + 12, "Whole book line", y_rng, "Yes")
    for c0 in (left, left + 3, left + 6, left + 9, left + 12):
        ws.merge_cells(start_row=s, start_column=c0, end_row=s, end_column=c0 + 1)
    # the picks, worked out once in hidden cells
    hr = s
    at = lambda c, k: f"${col(c)}${hr + k}"                   # noqa: E731
    LB, PB, BAND, MKEY, FEW, SHOW, LABR, XF = (at(S0, k) for k in range(8))
    X = [at(S0 + 6, k) for k in range(3)]                   # the filters' names across the bottom, if offered
    for k, name in xf:
        ws[X[k - 1].replace("$", "")] = name

    def which(cell: str, start: int, none: str) -> str:
        """The filter (start, start + 1, ...) whose option a dropdown cell reads, else `none`."""
        out = none
        for k in range(len(lines_opts), start - 1, -1):
            out = f"IF({cell}={live.q(lines_opts[k - 1])},{k},{out})"
        return out
    xfs = "0"
    for k in range(len(xf), 0, -1):
        xfs = f'IF(AND({X[k - 1]}<>"",{B}={X[k - 1]}),{k},{xfs})'
    LBP = which(LB_, 2, "1")                               # the filter Lines by names
    f = {XF: xfs,
         # a filter across the bottom: the lines are by the one picked, or another when it is the one across, and
         # there are no panels
         LB: f'IF({XF}=0,{LBP},IF({LBP}<>{XF},{LBP},IF({XF}=1,2,1)))',
         PB: f'IF({XF}>0,0,IF({which(PB_, 1, "0")}={LB},0,{which(PB_, 1, "0")}))',
         BAND: f"{B}&\"\"", MKEY: f'IFERROR(INDEX({k_rng},MATCH({M},{m_rng},0)),"")',
         SHOW: f'{Y}="Yes"', LABR: match(xk("C|", (BAND,), "|labels"))}
    for cell, x in f.items():
        ws[cell.replace("$", "")] = f"={x}"
    ws[FEW.replace("$", "")] = few if few is not None else 0
    VC = (S0 + 4, S0 + 5, S0 + 7)                           # each filter's values, down a hidden column
    for i in range(max(L, 1)):
        for k, vs in enumerate(vals):
            ws.cell(row=hr + i, column=VC[k], value=str(vs[i]) if i < len(vs) else None)
    V = lambda k, i: f"${col(VC[k - 1])}${hr + i - 1}"       # noqa: E731

    def by(sel: str, cells: list[str]) -> str:
        """The cell of filter `sel` (1, 2 or 3) from one per filter."""
        out = cells[-1]
        for k in range(len(cells) - 1, 0, -1):
            out = f"IF({sel}={k},{cells[k - 1]},{out})"
        return out
    LV = lambda j: f"${col(S0 + 1)}${hr + j - 1}"            # noqa: E731
    PV = lambda p: f"${col(S0 + 2)}${hr + p - 1}"            # noqa: E731
    USED = lambda p: f"${col(S0 + 3)}${hr + p - 1}"          # noqa: E731
    for j in range(1, L + 1):
        ws[LV(j).replace("$", "")] = "=" + by(LB, [f'{V(k, j)}&""' for k in range(1, len(fs) + 1)])
    for p in range(1, P + 1):
        ws[PV(p).replace("$", "")] = (f'=IF({PB}=0,"",' + by(PB, [f'{V(k, p)}&""' for k in range(1, len(fs) + 1)])
                                      + ")")
        ws[USED(p).replace("$", "")] = f'=IF({PB}=0,{p}=1,{PV(p)}<>"")'
    # the helper block: two key rows (the rate's _views row, the loans'), the series' names, then one row per point
    R_RATE, R_LOANS, R_TITLE, DR = hr, hr + 1, hr + 2, hr + 3
    rows_ = range(DR, DR + MB)
    for i, rr in enumerate(rows_, start=1):
        ws.cell(row=rr, column=CATC, value=f"={pick(LABR, i)}")
    for p in range(1, P + 1):
        for j in range(0, L + 1):
            c = scol(p, j)
            if j == 0:                                  # the whole book, in every panel drawn
                on = f"AND({USED(p)},{SHOW})"
                sfx, title = "", live.q(BOOK_LINE)
            else:
                on = f'AND({USED(p)},{LV(j)}<>"")'
                # each filter's value in the view's key: the line's, the panel's, or All loans (blank)
                sfx = ""
                for k, part in zip(range(1, len(fs) + 1), (WHERE, AND, AND3)):
                    vk = f'IF({LB}={k},{LV(j)},IF({PB}={k},{PV(p)},""))'
                    sfx += f'&IF({vk}="","",{live.q(part.format(""))}&{vk})'
                title = LV(j)
            rk = f'"C|"&{BAND}&"|"&{MKEY}{sfx}'
            lk = f'"C|"&{BAND}&"|loans"{sfx}'
            ws.cell(row=R_RATE, column=c, value=f'=IF({on},{match(rk)},"")')
            ws.cell(row=R_LOANS, column=c, value=f'=IF({on},{match(lk)},"")')
            ws.cell(row=R_TITLE, column=c, value=f'=IF({col(c)}{R_RATE}="","",{title})')
            RR, RL = f"${col(c)}${R_RATE}", f"${col(c)}${R_LOANS}"
            for i, rr in enumerate(rows_, start=1):
                rate, loans = pick(RR, i), pick(RL, i)
                # left off the line: no rate, or fewer loans than the Run's fewest (the Grids' grey rule)
                ws.cell(row=rr, column=c, value=f"=IF(OR({rate}=\"\",NOT(ISNUMBER({loans}))),NA(),"
                                                f"IF({loans}<{FEW},NA(),{rate}))")
                ws.cell(row=rr, column=scol(p, j, CLEAN), value=f'=IF(ISNUMBER({col(c)}{rr}),{col(c)}{rr},"")')
                ws.cell(row=rr, column=scol(p, j, X_COL), value=f"=IF(ISNUMBER({col(c)}{rr}),{i},NA())")
    clean = f"${col(scol(1, 0, CLEAN))}${DR}:${col(scol(P, L, CLEAN))}${DR + MB - 1}"
    # one scale for every panel: the smallest and largest value drawn anywhere, as an invisible series in each; and
    # the labels across, a little under the smallest, drawn in the chart as each label series' name
    lo, hi = f"MIN({clean})", f"MAX({clean})"
    span = f"IF({hi}>{lo},{hi}-{lo},MAX(ABS({hi}),0.01))"
    for i, rr in enumerate(rows_, start=1):
        lab = f"${col(CATC)}${rr}"
        # at zero when nothing drawn is below it, so the rates' axis starts at zero; else under the smallest
        ws.cell(row=rr, column=LY, value=f"=IF(COUNT({clean})=0,0,IF({lo}>=0,0,{lo}-{LABEL_DROP}*{span}))")
        for p in range(1, P + 1):
            ws.cell(row=rr, column=pcol(p, 0), value=f'=IF(AND({USED(p)},{lab}<>""),{i},NA())')
            # the scale's points in every panel, drawn or not, so an empty panel sits on the same axis
            ws.cell(row=rr, column=pcol(p, 1), value=f"=IF(COUNT({clean})>0,{SCALE_X},NA())"
                    if i <= 2 else "=NA()")
            for x, y, what in ((pcol(p, 0), LY, 2), (pcol(p, 1), PHY, 3)):
                ws.cell(row=rr, column=pcol(p, what), value=f"=IF(ISNUMBER({col(x)}{rr}),{col(y)}{rr},NA())")
        # the bottom one where the labels sit (zero, or under the smallest), so a panel with no labels drawn still
        # reaches as low as one with them
        ws.cell(row=rr, column=PHY, value=f"=IF(COUNT({clean})=0,NA(),{col(LY)}{rr})" if i == 1 else
                f"=IF(COUNT({clean})=0,NA(),{hi})" if i == 2 else "=NA()")
    status = s + 1
    lb_name = by(LB, [live.q(n) for n in names])
    said = f'IF(AND({PB_}<>"{NONE}",{PB}=0),{live.q(SAY_SAME if len(fs) == 2 else SAY_SAME3)},"")'
    for k, n in reversed(xf):
        head, tail = SAY_ACROSS.split("{other}")
        head = head.format(x=n)
        said = f'IF({XF}={k},{live.q(head)}&{lb_name}&{live.q(tail)},{said})'
    _cell(ws, status, left, f"={said}", size=9, color=house.CRIMSON, h="left")
    key = status + 1
    # the key, in cells so it names only the lines drawn now: the whole book dashed grey, then each line in its hue
    _cell(ws, key, left, f'=IF({SHOW},"- - {BOOK_LINE}","")', bold=True, size=9, color=house.SLATE, h="left",
          name="Arial")
    for j in range(1, L + 1):
        c0 = left + 1 + 2 * (j - 1)
        _cell(ws, key, c0, f'=IF({LV(j)}="","","● "&{LV(j)})', bold=True, size=9, color=HUES[(j - 1) % len(HUES)],
              h="left", name="Arial")
        ws.merge_cells(start_row=key, start_column=c0, end_row=key, end_column=c0 + 1)
    top = key + 2
    pb_name = by(PB, [live.q(n) for n in names])
    for p in range(1, P + 1):
        c0 = left + (p - 1) * PANEL_COLS
        _cell(ws, top, c0, f'=IF({USED(p)},IF({PB}=0,"All loans",{pb_name}&" is "&{PV(p)}),"")', bold=True,
              h="left", name="Arial")
        ws.merge_cells(start_row=top, start_column=c0, end_row=top, end_column=c0 + PANEL_COLS - 1)
        ws.add_chart(_chart(ws, p, L, MB, DR, (CATC, pcol(p, 0), pcol(p, 2), pcol(p, 1), pcol(p, 3)), scol, R_TITLE), f"{col(c0)}{top + 1}")
    _cell(ws, top + CHART_ROWS + 1, left, f'="Across the bottom: "&{B}&" · "&{M}&" · lines by "&{lb_name}&". "&'
          + live.q(SAY_READ), size=9, color=house.SLATE, h="left")
    r = top + CHART_ROWS + 3
    for p in range(1, P + 1):
        r = _table(ws, r, left, p, L, MB, USED(p), PV(p), PB, pb_name, B, FEW, CATC, DR, scol, R_TITLE, R_RATE,
                   R_LOANS) + 1
    _hide(ws, S0, pcol(P, 3))
    ws.freeze_panes = f"A{s + 1}"
    _fit(ws)


def _chart(ws, p: int, L: int, MB: int, DR: int, cols: tuple, scol, R_TITLE: int) -> ScatterChart:
    """Panel p: its lines, the whole book dashed, the labels across, and the invisible series that gives every panel
    one scale. A scatter with lines, not a line chart: a point at #N/A is left off in Excel and LibreOffice alike,
    where LibreOffice draws a line chart's #N/A at zero; so each point has an x, its place across, and the labels
    across are drawn in the chart, each as a one-point series named by its label cell, so they follow the dropdown."""
    CATC, LX, LY, PHX, PHY = cols
    ch = ScatterChart()
    ch.scatterStyle = "lineMarker"
    ref = lambda c, r0=DR, r1=DR + MB - 1: Reference(ws, min_col=c, min_row=r0, max_row=r1)     # noqa: E731
    for j in range(0, L + 1):
        ser = Series(ref(scol(p, j)), ref(scol(p, j, X_COL)))
        ser.tx = SeriesLabel(strRef=StrRef(f"'{SHEET}'!${col(scol(p, j))}${R_TITLE}"))
        ser.smooth = False
        ln = ser.graphicalProperties.line
        if j == 0:
            ln.solidFill, ln.prstDash, ln.width = house.SLATE, "dash", 19050
            ser.marker.symbol = "none"
        else:
            hue = HUES[(j - 1) % len(HUES)]
            ln.solidFill, ln.width = hue, 22225
            ser.marker.symbol = "circle"
            ser.marker.size = 6
            ser.marker.graphicalProperties = GraphicalProperties(solidFill=hue)
            ser.marker.graphicalProperties.line.solidFill = house.PAPER
        ch.series.append(ser)
    for i in range(MB):
        lab = Series(ref(LY, DR + i, DR + i), ref(LX, DR + i, DR + i))
        lab.tx = SeriesLabel(strRef=StrRef(f"'{SHEET}'!${col(CATC)}${DR + i}"))
        lab.graphicalProperties.line.noFill = True
        lab.marker.symbol = "none"
        lab.smooth = False
        lab.dLbls = DataLabelList(showSerName=True, showVal=False, showCatName=False, showLegendKey=False,
                                  showPercent=False, dLblPos="b")
        lab.dLbls.txPr = RichText(p=[Paragraph(pPr=ParagraphProperties(defRPr=CharacterProperties(
            sz=800, solidFill=house.SLATE)), endParaRPr=CharacterProperties(sz=800))])
        ch.series.append(lab)
    ph = Series(ref(PHY), ref(PHX))
    ph.graphicalProperties.line.noFill = True
    ph.marker.symbol = "none"
    ph.smooth = False
    ch.series.append(ph)
    ch.legend = None                                # the key is in cells above the charts
    ch.display_blanks = "gap"
    ch.visible_cells_only = False                   # the helper block sits in hidden columns
    ch.y_axis.number_format = "0.0%"
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=house.ROW_RULE))
    ch.x_axis.majorGridlines = None
    ch.x_axis.scaling.min = 0.5
    ch.x_axis.delete = True                         # the places across are numbers; the labels say them
    ch.y_axis.delete = False
    ch.width, ch.height = PANEL_COLS * 2.05, 8.2
    return ch


def _table(ws, r: int, left: int, p: int, L: int, MB: int, USED: str, PV: str, PB: str, pb_name: str, B: str,
           FEW: str, CATC: int, DR: int, scol, R_TITLE: int, R_RATE: int, R_LOANS: int) -> int:
    """Panel p's points: what is across the bottom down the side; for the whole book and each line, the rate and the
    loans. A point left off its line is grey here, as a thin pocket is on Grids. Returns the row under it."""
    _cell(ws, r, left, f'=IF({USED},IF({PB}=0,"All loans",{pb_name}&" is "&{PV}),"")', bold=True, h="left",
          name="Arial", color=house.PAPER)
    wide = f"{col(left)}{r}:{col(left + 2 * (L + 1))}{r}"
    cf(ws, wide, [(USED, house.INK, None, None)])                  # a panel not drawn has no table to head
    cf(ws, f"{col(left)}{r + 1}:{col(left + 2 * (L + 1))}{r + 1}", [(USED, CANVAS, None, None)])
    _cell(ws, r + 1, left, f"=IF({USED},{B},\"\")", bold=True, size=9, color=house.SLATE, h="left", indent=1)
    for j in range(0, L + 1):
        c = left + 1 + 2 * j
        t = f"${col(scol(p, j))}${R_TITLE}"
        _cell(ws, r, c, f"={t}", bold=True, size=9, color=house.PAPER, name="Arial")
        ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + 1)
        _cell(ws, r + 1, c, f'=IF({t}="","","Rate")', bold=True, size=9, color=house.SLATE)
        _cell(ws, r + 1, c + 1, f'=IF({t}="","","Loans")', bold=True, size=9, color=house.SLATE)
    for i in range(1, MB + 1):
        rr = r + 1 + i
        _cell(ws, rr, left, f'=IF({USED},${col(CATC)}${DR + i - 1},"")', h="left", indent=1)
        for j in range(0, L + 1):
            c = left + 1 + 2 * j
            RR, RL = f"${col(scol(p, j))}${R_RATE}", f"${col(scol(p, j))}${R_LOANS}"
            _cell(ws, rr, c, f"={pick(RR, i)}", fmt="0.00%")
            _cell(ws, rr, c + 1, f"={pick(RL, i)}", fmt="#,##0")
            grey = f"AND(ISNUMBER(${col(c + 1)}{rr}),${col(c + 1)}{rr}<{FEW})"     # both cells read the loans
            cf(ws, f"{col(c)}{rr}:{col(c + 1)}{rr}", [(grey, None, Font(color=house.DISABLED_TEXT), None)])
        ws.row_dimensions[rr].height = 16
    return r + 2 + MB
