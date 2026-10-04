"""The audit workbook: one pocket's every figure, worked out again by Excel from the loans, beside PocketBook's own.

Its words follow pocketbook/VOICE.md (the firm, 3 Oct 2026: "it does not sound like something a human would type").

The firm, 3 Oct 2026: "We will need to make this auditable... a demo output mode that would take a file and show the
calculations it makes on one set of things and prove out each one so someone could take their current population run
it and then independently understand the calculation and its steps." Then: "we definitely want to be able to
demonstrate and explain what the formulas are and how to do them by hand." Agreed: a SEPARATE workbook, so the main
one doesn't get slower, written by a Run when Control's "Also write the audit workbook?" is Yes; it proves every
figure for ONE pocket, picked by a live dropdown, recomputed by Excel from the raw loans. It opens on a pocket
selected at random from the tested pockets, seeded from the input file's SHA-256 (the firm, 3 Oct 2026: "tying out
one thing that should prove everything if you picked randomly"; chosen "Random, seed stamped"), until then the top
flagged pocket.

`<book stem> - audit.xlsx`, beside the main workbook:

    Start here       what the file is and the order to read it
    Run stamp        the input file, its SHA-256, its rows, every setting the Run used, the version, the time, the seed
    Rows in and out  rows read, less each kind left out of each rate, = the loans in it; each line tied to the Run
    Bands            each band column's edges as typed and as used, and the range each band covers
    Loans            one row per loan: only the columns the Run used, each band worked out by a formula from the raw
                     value and the edges, beside PocketBook's band; and two live columns, In this pocket and In this
                     band, that follow the picks
    One pocket       a Grid, Band and Segment dropdown; one row per figure: Step | Definition | Calculation (live
                     text) | Excel's figure (COUNTIFS / SUMIFS on Loans) | PocketBook's figure | Ties? | By hand
    Shuffle test     a 10-loan worked example; every shuffled gap of the random pocket, whose COUNTIF gives its
                     p-value; a two-proportion z-test beside it as a textbook cross-check; the allowance for many
                     tests worked out from every tested pocket in the picked pocket's grid and comparison
    _pocketbook      hidden: PocketBook's figures for EVERY pocket, so any pocket picked is compared
    _lists           hidden: the dropdowns' lists, and each family of tests in rank order

Fast enough at a bank's size: the Loans sheet's rows are written straight into the file's XML (openpyxl spends half a
minute on 185,000 rows of cells), the only formulas over every loan are its bands and the two In this ... columns, and
One pocket holds about forty SUMIFS and COUNTIFS. The shuffles are dealt again for the random pocket only, from the
Run's own seed, with perm.run: listing 10,000 shuffles for every pocket would be millions of rows.
"""

from __future__ import annotations

import io
import os
import re
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

from . import __version__, engine, house, perm, stats, timing
from .engine import ALL, MISSING_RULE_LABEL, REASON_LABEL, classify_number, classify_text
from .ingest import cell_text, is_blank

START, STAMP, ROWS, ONE, SHUFFLE, BANDS, LOANS, PB, LISTS = (
    "Start here", "Run stamp", "Rows in and out", "One pocket", "Shuffle test", "Bands", "Loans", "_pocketbook",
    "_lists")
SHEETS = (START, STAMP, ROWS, BANDS, LOANS, ONE, SHUFFLE, PB, LISTS)
#: Control's question (settings.yaml)
KEY = "audit_book"
TICK, CROSS = "✓", "✗"
#: "is a number", as a SUMIFS / COUNTIFS criterion: a loan's cell holds its number, or the reason it has none in words
NUM = '">-1E+307"'
#: a figure PocketBook has none for (a pocket alone in its band has no rest of its band), on _pocketbook
NONE = "none"
#: how close two figures must be to tie: a billionth of the figure, for the rounding in the last digits of a sum
TOL = "1E-9"
#: why a loan has no number in a column, as the Loans sheet says it (engine.REASON_LABEL, and the outcome's own)
NOT_01 = "(not 0 or 1)"
REASON_WORDS = {"blank": "blank", "not a number": "not a number", "missing by rule": "marked missing",
                "flag not 0 or 1": "not 0 or 1"}

USD, USD0, PCT, MULT, PTS, PV, INT = ('$#,##0.00;-$#,##0.00', '$#,##0;-$#,##0', '0.000%', '0.000"×"',
                                      '+0.000;-0.000;0.000', '0.00000', '#,##0')
LOW_FMT = '[<-1E+300]"lowest";General'
HIGH_FMT = '[>1E+300]"and up";General'
LOW, HIGH = -1e307, 1e307
#: the most loans the Loans sheet holds: an Excel sheet's 1,048,576 rows, less the header
MOST_LOANS = 1_048_575


def path_for(book: str | Path) -> Path:
    book = Path(book)
    return book.with_name(f"{book.stem} - audit.xlsx")


# --------------------------------------------------------------------------
# A sheet laid out cell by cell, then streamed: the workbook is written write-only, row by row


_SIDE = Side(style="medium", color=house.KEY_RED)
STYLES = {
    "title": dict(font=Font(name="Arial", bold=True, size=16, color=house.PAPER), fill=house.fill(house.INK),
                  alignment=Alignment(vertical="center"), border=Border(bottom=Side(style="thick", color=house.KEY_RED))),
    "sub": dict(font=Font(name="Arial", size=10, color=house.STONE), fill=house.fill(house.INK),
                alignment=Alignment(vertical="center"), border=Border(bottom=Side(style="thick", color=house.KEY_RED))),
    "sec": dict(font=Font(name="Arial", bold=True, size=10, color=house.PAPER), fill=house.fill(house.INK),
                alignment=Alignment(vertical="center"), border=Border(bottom=_SIDE)),
    "head": dict(font=Font(name="Arial", bold=True, size=9, color=house.PAPER), fill=house.fill(house.SLATE),
                 alignment=Alignment(horizontal="left", vertical="center", wrap_text=True)),
    "label": dict(font=Font(name="Arial", bold=True, size=9, color=house.SLATE)),
    "notel": dict(font=Font(name="Calibri", bold=True, size=10, color=house.INK_TEXT), fill=house.fill(house.CANVAS),
                  alignment=Alignment(vertical="top", wrap_text=True)),
    "note": dict(font=Font(name="Calibri", size=10, color=house.INK_TEXT), fill=house.fill(house.CANVAS),
                 alignment=Alignment(vertical="top", wrap_text=True)),
    "text": dict(font=Font(name="Calibri", size=10, color=house.INK_TEXT),
                 alignment=Alignment(vertical="top", wrap_text=True)),
    "bold": dict(font=Font(name="Calibri", bold=True, size=10, color=house.INK_TEXT),
                 alignment=Alignment(vertical="top", wrap_text=True)),
    "num": dict(font=Font(name="Calibri", size=10, color=house.INK_TEXT), alignment=Alignment(vertical="top")),
    "numc": dict(font=Font(name="Calibri", size=10, color=house.INK_TEXT),
                 alignment=Alignment(horizontal="center", vertical="top")),
    "textc": dict(font=Font(name="Calibri", size=10, color=house.INK_TEXT),
                  alignment=Alignment(horizontal="center", vertical="top")),
    "numb": dict(font=Font(name="Calibri", bold=True, size=10, color=house.INK_TEXT),
                 alignment=Alignment(vertical="top")),
    "grey": dict(font=Font(name="Calibri", size=9, color=house.SLATE), alignment=Alignment(vertical="top",
                                                                                          wrap_text=True)),
    "tie": dict(font=Font(name="Calibri", bold=True, size=11, color=house.INK_TEXT),
                alignment=Alignment(horizontal="center", vertical="top")),
    "pick": dict(font=Font(name="Calibri", bold=True, size=10, color=house.INK_TEXT), fill=house.fill(house.CANVAS),
                 border=Border(*(Side(style="thin", color=house.INK_TEXT),) * 4),
                 alignment=Alignment(horizontal="left", vertical="center", indent=1)),
}


class Canvas:
    """One sheet's cells, kept until the sheet is written in row order."""

    def __init__(self, title: str):
        self.title = title
        self.cells: dict[tuple[int, int], tuple] = {}
        self.widths: dict[int, float] = {}
        self.heights: dict[int, float] = {}
        self.hidden_cols: set[int] = set()
        self.merges: list[str] = []
        self.dvs: list[DataValidation] = []
        self.rules: list[tuple[str, object]] = []
        self.hidden = False
        self.freeze: str | None = None
        self.tab = house.TAB_RESULT

    def put(self, r: int, c: int, v, st: str | None = "text", fmt: str | None = None) -> str:
        self.cells[(r, c)] = (v, st, fmt)
        return f"${get_column_letter(c)}${r}"

    def title_band(self, title: str, sub: str, last: int) -> None:
        for c in range(2, last + 1):
            self.put(1, c, None, "sub")
        self.put(1, 2, title, "title")
        self.put(1, 3, sub, "sub")
        self.heights[1] = 30

    def section(self, r: int, text: str, last: int, first: int = 2) -> None:
        for c in range(first, last + 1):
            self.put(r, c, None, "sec")
        self.put(r, first, text, "sec")
        self.heights[r] = 20

    def header(self, r: int, heads: list[str], first: int = 2) -> None:
        for i, h in enumerate(heads):
            self.put(r, first + i, h, "head")

    def note(self, r: int, items: list[tuple[str, str]], first: int, last: int, chars: int) -> int:
        """The sheet's note: a label and its words a row, on CANVAS. Returns the row under it."""
        for label, words in items:
            for c in range(first, last + 1):
                self.put(r, c, None, "note")
            self.put(r, first, label, "notel")
            self.put(r, first + 1, words, "note")
            if last > first + 1:
                self.merges.append(f"{get_column_letter(first + 1)}{r}:{get_column_letter(last)}{r}")
            self.heights[r] = 14 * house.lines_at(words, chars) + 3
            r += 1
        return r + 1

    def emit(self, wb) -> None:
        ws = wb.create_sheet(self.title)
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.tabColor = self.tab
        if self.hidden:
            ws.sheet_state = "hidden"
        for c, w in self.widths.items():
            ws.column_dimensions[get_column_letter(c)].width = w
        for c in self.hidden_cols:
            ws.column_dimensions[get_column_letter(c)].hidden = True
        for r, h in self.heights.items():
            ws.row_dimensions[r].height = h
        if self.freeze:
            ws.freeze_panes = self.freeze
        last = max((r for r, _ in self.cells), default=0)
        by_row: dict[int, dict[int, tuple]] = {}
        for (r, c), x in self.cells.items():
            by_row.setdefault(r, {})[c] = x
        for r in range(1, last + 1):
            row = by_row.get(r, {})
            out = []
            for c in range(1, max(row, default=0) + 1):
                if c not in row:
                    out.append(None)
                    continue
                v, st, fmt = row[c]
                cell = WriteOnlyCell(ws, value=v)
                if isinstance(v, Lit):
                    cell.data_type = "s"
                for k, val in STYLES.get(st, {}).items():
                    setattr(cell, k, val)
                if fmt:
                    cell.number_format = fmt
                out.append(cell)
            ws.append(out)
        for m in self.merges:
            ws.merged_cells.add(m)
        for dv in self.dvs:
            ws.data_validations.append(dv)
        for rng, rule in self.rules:
            ws.conditional_formatting.add(rng, rule)


def _tie_rules(cv: Canvas, rng: str) -> None:
    """✗ in red on pink, ✓ in green."""
    first = rng.split(":")[0].replace("$", "")
    cv.rules.append((rng, FormulaRule(formula=[f'{first}="{CROSS}"'], font=Font(bold=True, color=house.CRIMSON),
                                      fill=PatternFill("solid", fgColor=house.ALERT_FG, bgColor=house.ALERT_FG))))
    cv.rules.append((rng, FormulaRule(formula=[f'LEFT({first},1)="{TICK}"'],
                                      font=Font(bold=True, color=house.POSITIVE))))


def ties(e: str, f: str, scale: str | None = None) -> str:
    """Excel's figure against PocketBook's: ✓ when both are numbers within a billionth, or neither is a number.
    `scale`: the figure a difference of two totals is a billionth of (its own size can be 0)."""
    return (f'=IF(AND(ISNUMBER({e}),ISNUMBER({f})),IF(ABS({e}-{f})<={TOL}*MAX(1,ABS({scale or f})),"{TICK}",'
            f'"{CROSS}"),'
            f'IF(AND(NOT(ISNUMBER({e})),NOT(ISNUMBER({f}))),"{TICK}","{CROSS}"))')


def clean(s) -> str:
    """Text as a cell can hold it: the control characters XML forbids removed (a loan number "L\\x015" written as it
    is made the Loans sheet unreadable, the review of 3 Oct 2026)."""
    return ILLEGAL_CHARACTERS_RE.sub("", str(s))


class Lit(str):
    """Text from the loan file, written as text: a segment "=A" is a value, not a formula (Canvas.emit)."""


def lit(s) -> Lit:
    return Lit(clean(s))


#: Excel refuses a string constant longer than this inside a formula (LibreOffice takes it, so tests alone can't tell)
MOST_IN_LITERAL = 255


def q(s) -> str:
    """Text as a formula's string literal: in pieces of at most 255 characters joined with &, since Excel drops a
    formula holding a longer one when it repairs the file (the review of 4 Oct 2026)."""
    t = clean(s)
    parts = [t[i:i + MOST_IN_LITERAL] for i in range(0, len(t), MOST_IN_LITERAL)] or [""]
    return "&".join('"' + p.replace('"', '""') + '"' for p in parts)


def crit(s) -> str:
    """A COUNTIF criterion that matches the text `s` exactly as written: "=" first, so a leading <, > or = is not
    read as an operator, and ~ * ? escaped with ~, so none is read as a wildcard (the review of 3 Oct 2026)."""
    return q("=" + clean(s).replace("~", "~~").replace("*", "~*").replace("?", "~?"))


def find(value: str, rng: str, less: int = 0) -> str:
    """The row of `rng` holding `value`, character for character (EXACT), less `less`; #N/A when none does. MATCH
    is not used: it reads ~ * ? in the value as wildcards and ignores case: a segment "Br~anch" found no row, so every
    figure read ✗, and "<B*" in part D found another pocket's row (the review of 3 Oct 2026)."""
    return f"LOOKUP(2,1/EXACT({rng},{value}),ROW({rng})-{less})"


def text_cell(ws, v) -> WriteOnlyCell:
    """A write-only cell holding `v` as text, never as a formula."""
    c = WriteOnlyCell(ws, value=clean(v))
    c.data_type = "s"
    return c


# --------------------------------------------------------------------------
# What the Run worked out, per pocket


def _rate(num: float, den: float) -> float | None:
    return num / den if den else None


def _neg(v: float | None) -> float | None:
    return -v if v is not None else None


def grid_name(res, g) -> str:
    names = _names(res)
    return f"{names[g.band]} x {names[g.dimension]}"


def _names(res) -> dict[str, str]:
    out = {b.name: b.field for b in res.config.bands}
    out.update({d.name: d.field for d in res.config.dimensions})
    return out


#: the fixed name the audit's random pick is seeded from, with the input file's SHA-256 (the firm, 3 Oct 2026: "tying
#: out one thing that should prove everything if you picked randomly"; chosen "Random, seed stamped")
PICK_SEED_NAME = "the audit workbook's pocket, selected at random"


@dataclass(frozen=True)
class Pick:
    """The pocket the audit workbook opens on: drawn uniformly from `population` tested pockets with `seed`."""
    grid: object
    key: tuple
    seed: int
    population: int
    index: int


def tested_pockets(res) -> list[tuple[str, object, tuple]]:
    """The population the audit's pocket is drawn from: every pocket of the Run's grids with a GCO p-value on the
    comparison that decides it (the rest of its band, or the rest of the book), as (its key on _pocketbook, the
    grid, (band, segment)), in the order of that key as text, character by character."""
    out = []
    for g in res.grids:
        name = grid_name(res, g)
        for k, c in g.inner():
            s = c.rates["gco_rate"]
            if (s.p_band if s.by_band else s.p_book) is not None:
                out.append((f"{name}|{k[0]}|{k[1]}", g, k))
    return sorted(out, key=lambda t: t[0])


def pick_seed(sha256: str) -> int:
    """The pick's seed: perm.seed_of over the fixed name and the input file's SHA-256, so the same file always
    opens on the same pocket."""
    return perm.seed_of(PICK_SEED_NAME, sha256)


def pick_index(seed: int, n: int) -> int:
    """Which of n pockets, numbered from 0: the seed modulo n (uniform to within n / 2**64)."""
    return seed % n


def random_pocket(res, sha256: str) -> Pick:
    """The audit workbook's pocket, selected at random from tested_pockets. With none tested, the first pocket of
    the first grid, and a population of 0."""
    seed = pick_seed(sha256)
    pop = tested_pockets(res)
    if not pop:
        g = res.grids[0]
        return Pick(g, next(k for k, _ in g.inner()), seed, 0, 0)
    index = pick_index(seed, len(pop))
    _, g, k = pop[index]
    return Pick(g, k, seed, len(pop), index)


def _figures(res, g, k, c) -> dict:
    """Every figure One pocket shows, as the Run worked it out. None where the Run has none."""
    t = res.total.rates
    o, gc, rn = c.rates["outcome_loans"], c.rates["gco_rate"], c.rates["ranr_rate"]
    to, tg, tr = t["outcome_loans"], t["gco_rate"], t["ranr_rate"]
    shuffles = res.config.benchmark.shuffles
    band = g.cells[(k[0], ALL)].rates["gco_rate"]
    rbd, rbn = band.den - gc.den, band.num - gc.num
    size = (getattr(g, "sizes", {}) or {}).get(k)
    book = res.book_size
    rest = _rate(tg.num - gc.num, tg.den - gc.den)
    raw_book = (gc.hits_book + 1) / (shuffles + 1) if gc.p_book is not None and gc.hits_book is not None else None
    raw_band = (gc.hits_band + 1) / (shuffles + 1) if gc.p_band is not None and gc.hits_band is not None else None
    z, zp = stats.two_prop_z(o.num, o.units, to.num - o.num, to.units - o.units)
    avg = size.average if size is not None else None
    # RANR and RANR + GCOs against the rest of the book and of the band, as RANR vs GCOs shows them: a gap in
    # points, and the dollars above (+) or short of (-) the rest's rate. The Run keeps the shortfall, so its sign is
    # turned here, as the tab turns it
    cn, tc = c.rates["contribution_rate"], t["contribution_rate"]
    bandr, bandc = g.cells[(k[0], ALL)].rates["ranr_rate"], g.cells[(k[0], ALL)].rates["contribution_rate"]
    whole = book.booked if book is not None else None
    return {
        "ctb_bk": cn.den, "ctb": cn.num, "ctb_rate": cn.rate,
        "book_all_bk": whole, "book_nogco_bk": whole - tg.den if whole is not None else None,
        "ranr_usd_rest": _neg(rn.excess_rest),
        "book_ctb_bk": tc.den, "book_ctb": tc.num, "ctb_rest_rate": _rate(tc.num - cn.num, tc.den - cn.den),
        "ctb_gap": cn.vs_rest * 100 if cn.vs_rest is not None else None, "ctb_usd_rest": _neg(cn.excess_rest),
        "band_ranr_bk": bandr.den - rn.den, "band_ranr": bandr.num - rn.num,
        "band_ranr_rate": _rate(bandr.num - rn.num, bandr.den - rn.den),
        "ranr_gap_band": rn.vs_band * 100 if rn.vs_band is not None else None,
        "ranr_usd_band": _neg(rn.excess_band),
        "band_ctb_bk": bandc.den - cn.den, "band_ctb": bandc.num - cn.num,
        "band_ctb_rate": _rate(bandc.num - cn.num, bandc.den - cn.den),
        "ctb_gap_band": cn.vs_band * 100 if cn.vs_band is not None else None,
        "ctb_usd_band": _neg(cn.excess_band),
        "loans": c.rows, "bad": o.num, "bad_n": o.units, "bad_rate": o.rate,
        "gco_bk": gc.den, "gco": gc.num, "gco_rate": gc.rate,
        "ranr_bk": rn.den, "ranr": rn.num, "ranr_rate": rn.rate,
        "book_gco_bk": tg.den, "book_gco": tg.num, "book_rate": tg.rate,
        "rest_gco_bk": tg.den - gc.den, "rest_gco": tg.num - gc.num, "rest_rate": rest,
        "x_book": gc.vs_topline, "x_rest": gc.vs_rest,
        "gco_gap": (gc.rate - rest) * 100 if gc.rate is not None and rest is not None else None,
        "excess_rest": gc.excess_rest,
        "band_gco_bk": rbd, "band_gco": rbn, "band_rate": _rate(rbn, rbd), "x_band": gc.vs_band,
        "excess_band": gc.excess_band, "by_band": bool(gc.by_band), "dollars": gc.dollars,
        "book_ranr_bk": tr.den, "book_ranr": tr.num, "ranr_rest_bk": tr.den - rn.den, "ranr_rest": tr.num - rn.num,
        "ranr_rest_rate": _rate(tr.num - rn.num, tr.den - rn.den),
        "ranr_gap": rn.vs_rest * 100 if rn.vs_rest is not None else None,
        "avg_n": size.loans if size is not None else 0, "avg_bk": size.booked if size is not None else 0.0,
        "avg": avg, "book_avg": book.average if book is not None else None,
        "avg_x": engine.index_of(avg, book.average if book is not None else None),
        "events": gc.events, "hits_book": gc.hits_book if raw_book is not None else None,
        "hits_band": gc.hits_band if raw_band is not None else None,
        "raw_book": raw_book, "raw_band": raw_band, "adj_book": gc.p_book, "adj_band": gc.p_band,
        "hits": (gc.hits_band if gc.by_band else gc.hits_book) if (raw_band if gc.by_band else raw_book) is not None
        else None,
        "raw": raw_band if gc.by_band else raw_book, "p": gc.p_band if gc.by_band else gc.p_book,
        "flag": gc.flag, "material": gc.material, "worse_material": engine.worse_and_material(gc),
        "bad_tot": to.num, "bad_n_tot": to.units, "z": z, "z_p": zp,
    }


#: _pocketbook's columns, in order: the key, the grid, band and segment, then every figure
PB_KEYS = ["key", "grid", "band", "seg", "loans", "bad", "bad_n", "bad_rate", "gco_bk", "gco", "gco_rate", "ranr_bk",
           "ranr", "ranr_rate", "book_gco_bk", "book_gco", "book_rate", "rest_gco_bk", "rest_gco", "rest_rate",
           "x_book", "x_rest", "gco_gap", "excess_rest", "band_gco_bk", "band_gco", "band_rate", "x_band",
           "excess_band", "by_band", "dollars", "book_ranr_bk", "book_ranr", "ranr_rest_bk", "ranr_rest",
           "ranr_rest_rate", "ranr_gap", "avg_n", "avg_bk", "avg", "book_avg", "avg_x", "events", "hits_book",
           "hits_band", "raw_book", "raw_band", "adj_book", "adj_band", "hits", "raw", "p", "flag", "material",
           "worse_material", "bad_tot", "bad_n_tot", "z", "z_p",
           # RANR and RANR + GCOs against the rest of the book and of the band (RANR vs GCOs' gaps and dollars)
           "ctb_bk", "ctb", "ctb_rate", "book_all_bk", "book_nogco_bk", "ranr_usd_rest", "book_ctb_bk", "book_ctb",
           "ctb_rest_rate", "ctb_gap", "ctb_usd_rest", "band_ranr_bk", "band_ranr", "band_ranr_rate", "ranr_gap_band",
           "ranr_usd_band", "band_ctb_bk", "band_ctb", "band_ctb_rate", "ctb_gap_band", "ctb_usd_band",
           # formulas: the allowance for many tests, per grid and comparison (Benjamini-Hochberg's step-up)
           "m_book", "rank_book", "q_book", "m_band", "rank_band", "q_band",
           # the grid's number (its row on _lists less 1): a COUNTIFS matches a number exactly, while a grid's name
           # can hold ~ * ? or begin with < > =, which COUNTIFS reads as wildcards and operators
           "gridno"]
PBC = {k: i + 1 for i, k in enumerate(PB_KEYS)}


#: the hidden sheets, as a formula names them
PBQ, LISTSQ = f"'{PB}'", f"'{LISTS}'"


def pb_col(key: str) -> str:
    return get_column_letter(PBC[key])


# --------------------------------------------------------------------------
# The loans


def _loan_values(res):
    """Every column the Loans sheet shows, as the Run read it: each value a number, or the reason it has none."""
    cfg, rules = res.config, res.config.missing
    rows = res.table.rows
    ms = {m.name: m for m in res.measures}
    out = {}

    def numbers(col):
        got = []
        for r in rows:
            v, why = classify_number(r.get(col), rules.get(col))
            got.append(v if why is None else REASON_LABEL[why])
        return got

    out["booked"] = numbers(ms["gco_rate"].per)
    out["gco"] = numbers(ms["gco_rate"].value)
    out["ranr"] = numbers(ms["ranr_rate"].value)
    om = ms["outcome_loans"]
    bad = []
    for r in rows:
        raw = r.get(om.flag)
        if om.flag_is is not None:
            if is_blank(raw):
                bad.append(REASON_LABEL["blank"])
            elif classify_text(raw, rules.get(om.flag)) == MISSING_RULE_LABEL:
                bad.append(REASON_LABEL["missing by rule"])
            else:
                bad.append(1 if cell_text(raw) == cell_text(om.flag_is) else 0)
            continue
        v, why = classify_number(raw, rules.get(om.flag))
        if why is None and v not in (0.0, 1.0):
            bad.append(NOT_01)
        else:
            bad.append(int(v) if why is None else REASON_LABEL[why])
    out["bad"] = bad
    out["key"] = [cell_text(r.get(cfg.key)) if not is_blank(r.get(cfg.key)) else "" for r in rows]
    for b in cfg.bands:
        out[("band", b.name)] = numbers(b.field)
    out["date"] = ([cell_text(r.get(cfg.origination_date)) if not is_blank(r.get(cfg.origination_date)) else ""
                    for r in rows] if cfg.origination_date and cfg.origination_date in res.table.columns else None)
    return out


class Layout:
    """Where everything sits, worked out before anything is written: the sheets point at each other."""

    def __init__(self, res, values):
        cfg = res.config
        self.n = res.rows
        self.last = self.n + 1
        ms = {m.name: m for m in res.measures}
        cols = [("row", "Row in the extract"), ("key", f"{cfg.key} (loan)"),
                ("booked", f"{ms['gco_rate'].per} (booked)"), ("gco", f"{ms['gco_rate'].value} (GCO)"),
                ("ranr", f"{ms['ranr_rate'].value} (RANR)"), ("bad", f"{ms['outcome_loans'].flag} (bad)")]
        for b in cfg.bands:
            cols += [(("raw", b.name), b.field), (("band", b.name), f"{b.field} band"),
                     (("pb", b.name), f"{b.field} band, PocketBook's")]
        for d in cfg.dimensions:
            cols.append((("seg", d.name), d.field))
        if values["date"] is not None:
            cols.append(("date", f"{cfg.origination_date} (origination date)"))
        cols += [("in_pocket", "In this pocket"), ("in_band", "In this band")]
        self.cols = cols
        self.at = {k: i + 1 for i, (k, _) in enumerate(cols)}
        self.width = len(cols)

    def letter(self, key) -> str:
        return get_column_letter(self.at[key])

    def rng(self, key) -> str:
        c = self.letter(key)
        return f"{LOANS}!${c}$2:${c}${self.last}"


# --------------------------------------------------------------------------
# The book


def write(res, book: str | Path, src: str | Path, sha256: str, settings: list[tuple[str, str]] | None = None,
          when: datetime | None = None) -> Path:
    """Write the audit workbook for a bleed Run beside `book`; returns its path. `settings`: Control's answers as
    (question, what the Run used) in words."""
    out = path_for(book)
    when = when or datetime.now()
    timing.mark("Audit workbook: reading the loans")
    values = _loan_values(res)
    lay = Layout(res, values)
    pockets = [(g, k, c) for g in res.grids for k, c in g.inner()]
    pick = random_pocket(res, sha256)
    dg, dk = pick.grid, pick.key
    names = {id(g): grid_name(res, g) for g in res.grids}
    number = {id(g): i for i, g in enumerate(res.grids, start=1)}
    figs = [(names[id(g)], k, {**_figures(res, g, k, c), "gridno": number[id(g)]}) for g, k, c in pockets]
    # each pocket's row on _pocketbook, by its key: part D finds a pocket by its row, never by its name
    pb_row = {f"{gname}|{k[0]}|{k[1]}": i for i, (gname, k, _) in enumerate(figs, start=2)}

    wb = Workbook(write_only=True)
    bands_cv, band_ranges = _bands(res, lay, values)
    one_cv, one = _one_pocket(res, lay, names[id(dg)], dk, figs, pick)
    timing.mark("Audit workbook: dealing the random pocket's shuffles again")
    families = _families(res, names, number)
    shuffle_cv = _shuffle(res, dg, dk, names[id(dg)], one, figs, families)
    timing.mark("Audit workbook: writing the sheets")
    for cv in (_start(res, book), _stamp(res, src, sha256, settings or [], when, pick, names[id(dg)]),
               _rows(res, lay, values), bands_cv):
        cv.emit(wb)
    template = _loans_head(wb, lay)
    one_cv.emit(wb)
    shuffle_cv.emit(wb)
    _pocketbook(len(figs)).emit_rows(wb, figs)
    _lists(res, lay, families, pb_row).emit(wb)
    for name, ref in [("Loan_" + k, lay.rng(key)) for k, key in (
            ("Row", "row"), ("Booked", "booked"), ("GCO", "gco"), ("RANR", "ranr"), ("Bad", "bad"),
            ("InPocket", "in_pocket"), ("InBand", "in_band"))] + [
            ("Pick_Grid", f"'{ONE}'!{one['grid']}"), ("Pick_GridNo", f"'{ONE}'!$K$1"),
            ("Pick_Band", f"'{ONE}'!{one['band']}"),
            ("Pick_Seg", f"'{ONE}'!{one['seg']}"), ("Pick_BandCol", f"'{ONE}'!{one['bandcol']}"),
            ("Pick_SegCol", f"'{ONE}'!{one['segcol']}")]:
        wb.defined_names.add(DefinedName(name, attr_text=ref))
    buf = io.BytesIO()
    wb.save(buf)
    timing.mark("Audit workbook: writing the loans")
    data = _with_loans(buf.getvalue(), lay, values, res, band_ranges, template)
    part = out.with_name(out.name + ".part")   # written whole, then put in place: a failed write leaves no half file
    part.write_bytes(data)
    try:
        os.replace(part, out)
    finally:
        part.unlink(missing_ok=True)
    return out


# --------------------------------------------------------------------------
# Start here, Run stamp


#: Start here's heading over the list of sheets
ORDER = "Suggested order of review:"


def _start(res, book) -> Canvas:
    cv = Canvas(START)
    cv.tab = house.TAB_YOU
    cv.widths = {1: 2, 2: 14, 3: 96}               # B holds the title, "Start here", as well as the numbers
    cv.title_band("Start here", "The purpose of this workbook and the order in which to review it.", 3)
    lines = [
        ("", f"This workbook recalculates PocketBook's figures for a single pocket, step by step. It was produced by "
             f"the same Run as {Path(book).name}."),
        ("", "Each figure is recalculated in Excel from the loan-level records on the Loans sheet and shown beside "
             "the figure PocketBook reported. The Ties? column indicates whether the two agree."),
        ("", "The workbook opens on a pocket selected at random from the pockets the Run tested; Run stamp gives the "
             "seed and how it was drawn. To review a different pocket, select it at the top of the One pocket "
             "sheet; every figure updates accordingly."),
        ("", "This workbook is independent of the main workbook, so it can be edited freely without affecting it."),
        (None, None),
        ("", ORDER),
        ("1", "Run stamp — identifies the input file and its SHA-256 fingerprint, and lists every setting the Run "
              "used."),
        ("2", "Rows in and out — reconciles the rows read to the loans included in each rate, with the reason for "
              "each exclusion."),
        ("3", "Bands — shows where each band begins and ends."),
        ("4", "Loans — lists every loan, one row each, with only the columns the Run used. Each loan's band is "
              "assigned by formula from the edges on the Bands sheet."),
        ("5", "One pocket — shows every figure for the selected pocket: its definition, the calculation with the "
              "pocket's own numbers, and how to reproduce it by hand."),
        ("6", "Shuffle test — explains how the p-value is derived, with a small example that can be followed on "
              "paper, and works out the allowance for many tests from every tested pocket's p-value."),
        (None, None),
        ("", "To verify a figure independently, filter the Loans sheet to the selected pocket (In this pocket = 1) "
             "and select the relevant column; Excel's status bar shows its count and sum."),
        ("", "Excel calculates the figures when the file is opened. If the Ties? column is blank, press F9 to "
             "recalculate."),
    ]
    r = 3
    for num, text in lines:
        if num is not None:
            cv.put(r, 2, num or None, "bold")
            cv.put(r, 3, text, "bold" if text == ORDER else "text")
        r += 1
    return cv


#: how PocketBook placed a band column's edges, as the Run stamp says it (config.Band.cut)
CUT_WORDS = {"equal_loans": "each with roughly the same number of loans",
             "round": "each with roughly the same number of loans, edges rounded"}


def _settings_lines(res) -> list[tuple[str, str]]:
    cfg, b = res.config, res.config.benchmark
    ms = {m.name: m for m in res.measures}
    om = ms["outcome_loans"]
    out = [("Loan number", cfg.key), ("Booked amount", ms["gco_rate"].per),
           ("Outcome (bad loan)", om.flag + (f" = {om.flag_is}" if om.flag_is is not None else " = 1")),
           ("GCO dollars", ms["gco_rate"].value), ("RANR dollars", ms["ranr_rate"].value)]
    for x in cfg.bands:
        edges = res.band_edges.get(x.name, ())
        how = ("entered on Columns" if x.edges else
               "one band per distinct value" if x.name in (res.value_bands or {}) else
               f"{x.count} bands set by PocketBook, " + CUT_WORDS.get(x.cut, x.cut.replace('_', ' ')))
        out.append((f"Band column {x.field}", f"Edges {'; '.join(_g(e) for e in edges)} ({how})"))
    for d in cfg.dimensions:
        out.append(("Segment column", d.field))
    if cfg.split:
        out.append(("Split by", f"{cfg.split[0]} ({cfg.split[1].replace('_', ' ')})"))
    for i, f in enumerate((cfg.filter_by, cfg.filter_by2, cfg.filter_by3), start=1):
        if f:
            out.append((f"Filter {i}", f))
    if cfg.origination_date:
        out.append(("Origination date", cfg.origination_date))
    for col, rule in sorted(cfg.missing.items()):
        said = [f"below {_g(rule.below)}" if rule.below is not None else "",
                f"above {_g(rule.above)}" if rule.above is not None else "",
                f"at or below {_g(rule.at_or_below)}" if rule.at_or_below is not None else "",
                ("the values " + ", ".join(_g(v) if isinstance(v, (int, float)) else str(v) for v in rule.values))
                if rule.values else ""]
        out.append((f"Treated as missing in {col}", "; ".join(x for x in said if x)))
    if b is not None:
        out += [("Fewest loans in a pocket", f"{b.min_units:,}"), ("Fewest losses to test", f"{b.min_events:,}"),
                ("Worse at", f"{b.worse_at:g}×"), ("Better at", f"{b.better_at:g}×"),
                ("Confidence", f"{b.confidence:g}"), ("Power", f"{b.power:g}"),
                ("Judged against", "the rest of its band" if b.compare_to == "peers" else "the rest of the book"),
                ("Allowance for many tests", {"bh": "Benjamini-Hochberg", "bonferroni": "Bonferroni",
                                              "none": "none"}[b.many_tests]),
                ("Materiality", {"none": "none", "share": f"{b.materiality[1] * 100:g}% of the book's GCOs"}.get(
                    b.materiality[0], f"${b.materiality[1]:,.0f} of GCOs")),
                ("Profit line", str(b.revenue_line) if b.revenue_line is not None else "each pocket's own test"),
                ("Shuffles", f"{b.shuffles:,}")]
    return out


def _full(v) -> str:
    """An edge at full precision: 655, or 0.35000000000000003 as Python holds it."""
    v = float(v)
    return f"{v:.0f}" if v.is_integer() else repr(v)


def _g(v) -> str:
    v = float(v)
    return f"{v:,.0f}" if v.is_integer() else f"{v:,.6f}".rstrip("0").rstrip(".")


def _fit_col(cv: Canvas, col: int, chars: int) -> None:
    """Every row whose words in `col` wrap, tall enough for them (a row whose height is set already keeps it)."""
    for (r, c), (v, st, _) in list(cv.cells.items()):
        if c == col and isinstance(v, str) and not v.startswith("=") and r not in cv.heights and r > 1:
            n = house.lines_at(v, chars)
            if n > 1:
                cv.heights[r] = 12.5 * n + 3


def _fit_row(cv: Canvas, r: int, v, chars: int = 96) -> None:
    """Run stamp's wrapped words, each row tall enough for them."""
    if isinstance(v, str) and house.lines_at(v, chars) > 1:
        cv.heights[r] = 14 * house.lines_at(v, chars) + 3


def _stamp(res, src, sha256: str, settings, when: datetime, pick: Pick, dgrid: str) -> Canvas:
    cv = Canvas(STAMP)
    cv.tab = house.TAB_RECORD
    cv.widths = {1: 2, 2: 44, 3: 80}
    cv.title_band("Run stamp", "The input file, the software version and the settings used for this Run.", 3)
    seed = perm.seed_of(engine.SHUFFLE_SEED_NAME)
    r = 3
    cv.section(r, "The Run", 3)
    rows = [("Input file", Path(src).name), ("Its SHA-256 fingerprint", sha256),
            ("Rows read", res.rows), ("PocketBook version", __version__),
            ("Run at", when.strftime("%Y-%m-%d %H:%M:%S")),
            ("Shuffles per test", res.config.benchmark.shuffles),
            ("Shuffle seed", str(seed)),
            ("Where the seed comes from", f'The seed is derived by hashing a fixed name, "{engine.SHUFFLE_SEED_NAME}", '
                                          f'with SHA-256, so the same extract always produces the same shuffles.')]
    r += 1
    for label, v in rows:
        cv.put(r, 2, label, "bold")
        cv.put(r, 3, lit(v) if isinstance(v, str) else v, "num" if isinstance(v, int) else "text",
               INT if isinstance(v, int) else None)
        _fit_row(cv, r, v)
        r += 1
    r += 1
    cv.section(r, "The pocket this workbook opens on", 3)
    r += 1
    order = (f"The {pick.population:,} tested pockets are listed in order of their grid, band and segment names, "
             f"compared as text character by character, and numbered from 0. The pocket selected is number seed "
             f"modulo {pick.population:,}, which is {pick.index:,}." if pick.population else
             "No pocket was tested, so there was nothing to draw from.")
    for label, v in (("Pocket selected", pick_words(pick, dgrid)),
                     ("Population drawn from", f"{pick.population:,} pockets: every pocket of the Run's grids with a "
                                               f"GCO p-value on the comparison that decides it (the rest of its "
                                               f"band, or the rest of the book). A pocket with fewer loans with a "
                                               f"loss than the Run's minimum has no p-value."),
                     ("Pick seed", str(pick.seed)),
                     ("Where the pick seed comes from", f'The seed is derived with SHA-256 from a fixed name, '
                                                        f'"{PICK_SEED_NAME}", and the input file\'s SHA-256 '
                                                        f'fingerprint, so the same file always opens on the same '
                                                        f'pocket. {order}')):
        cv.put(r, 2, label, "bold")
        cv.put(r, 3, lit(v))
        _fit_row(cv, r, v)
        r += 1
    r += 1
    cv.section(r, "Settings used by the Run", 3)
    r += 1
    for label, v in _settings_lines(res):
        cv.put(r, 2, label, "bold")
        cv.put(r, 3, lit(v))
        _fit_row(cv, r, v)
        r += 1
    if settings:
        r += 1
        cv.section(r, "Answers on Control, as read by the Run", 3)
        r += 1
        for label, v in settings:
            cv.put(r, 2, label, "bold")
            cv.put(r, 3, lit(v))
            _fit_row(cv, r, v)
            r += 1
    return cv


#: on the Run stamp, the cells other sheets read
STAMP_SHUFFLES = f"'{STAMP}'!$C$9"
STAMP_SHA = f"'{STAMP}'!$C$5"


# --------------------------------------------------------------------------
# Rows in and out


def _rows(res, lay: Layout, values) -> Canvas:
    cv = Canvas(ROWS)
    cv.widths = {1: 2, 2: 58, 3: 16, 4: 16, 5: 8, 6: 60}
    cv.title_band("Rows in and out", "A reconciliation of the rows read to the loans included in each rate.", 6)
    r = cv.note(3, [("How it works", "Every loan in the extract is included in the Run. A loan with no numeric value "
                                     "in a column is excluded only from the rates that use that column. Excel's count "
                                     "is taken from the Loans sheet and shown beside PocketBook's count.")],
                2, 6, 120)
    cv.header(r, ["", "Excel's count", "PocketBook's count", "Ties?", "By hand"])
    first = r + 1
    r += 1
    cv.put(r, 2, f"Rows read from {Path(res.source).name}", "bold")
    cv.put(r, 3, "=COUNT(Loan_Row)", "numb", INT)
    cv.put(r, 4, res.rows, "numb", INT)
    cv.put(r, 5, ties(f"C{r}", f"D{r}"), "tie")
    cv.put(r, 6, "The number of data rows on the Loans sheet, excluding the header.", "grey")
    read_row = r
    r += 2
    ms = {m.name: m for m in res.measures}
    col_of = {ms["gco_rate"].per: "booked", ms["gco_rate"].value: "gco", ms["ranr_rate"].value: "ranr",
              ms["outcome_loans"].flag: "bad"}
    blocks = [("outcome_loans", "Bad loans ÷ Loans", "bad", None),
              ("gco_rate", "GCOs ÷ Booked", "gco", "booked"),
              ("ranr_rate", "RANR ÷ Booked", "ranr", "booked")]
    for mname, title, top, per in blocks:
        cv.section(r, title, 6)
        r += 1
        start = r
        lo = res.left_out.get(mname, {})
        for (column, why), k in sorted(lo.items(), key=lambda t: (str(t[0][0]), str(t[0][1]))):
            ck = col_of.get(column, top)
            label = NOT_01 if why == "flag not 0 or 1" else REASON_LABEL.get(why, f"({why})")
            cv.put(r, 2, f"Less: {column} {REASON_WORDS.get(why, why)}")
            if ck == top or per is None:
                f = f'=COUNTIF({lay.rng(ck)},{crit(label)})'
            else:
                f = f'=COUNTIFS({lay.rng(ck)},{crit(label)},{lay.rng(top)},{NUM})'
            cv.put(r, 3, f, "num", INT)
            cv.put(r, 4, k, "num", INT)
            cv.put(r, 5, ties(f"C{r}", f"D{r}"), "tie")
            # the count is of loans excluded for this reason alone: a loan with no number in the rate's other column
            # too is counted there (the tie-out, 3 Oct 2026: every By hand step filters what its formula counts)
            also = "" if ck == top or per is None else (
                f" and exclude text entries from {ms[mname].value}")
            cv.put(r, 6, f'Filter {column} on the Loans sheet to "{label}"{also}.', "grey")
            r += 1
        cv.put(r, 2, "Total excluded")
        both = f"COUNTIFS({lay.rng(top)},{NUM}" + (f",{lay.rng(per)},{NUM})" if per else ")")
        cv.put(r, 3, f"=C{read_row}-{both}", "num", INT)
        cv.put(r, 4, sum(lo.values()), "num", INT)
        cv.put(r, 5, ties(f"C{r}", f"D{r}"), "tie")
        cv.put(r, 6, "Rows read less the loans with a numeric value in every column the rate uses. "
                     + ("The lines above sum to this total." if lo else "No loans were excluded."), "grey")
        r += 1
        cv.put(r, 2, f"Loans included in {title}", "bold")
        cv.put(r, 3, "=" + both, "numb", INT)
        cv.put(r, 4, res.total.rates[mname].units, "numb", INT)
        cv.put(r, 5, ties(f"C{r}", f"D{r}"), "tie")
        cols = [ms[mname].flag if per is None else ms[mname].value] + ([ms[mname].per] if per else [])
        cv.put(r, 6, "On the Loans sheet, exclude text entries from " + " and ".join(cols)
               + "; the count of the remaining rows is this figure.", "grey")
        r += 2
        del start
    cv.section(r, "Band and segment columns: no loans are excluded", 6)
    r += 1
    cv.put(r, 2, "A loan with no numeric value in a band column, or no value in a segment column, is assigned to a "
                 "separate band or segment, such as (blank), and remains in every rate.", "grey")
    cv.merges.append(f"B{r}:F{r}")
    cv.heights[r] = 28
    r += 1
    labels = res.row_labels
    for b in res.config.bands:
        got = {}
        for v in labels[b.name]:
            if v in REASON_LABEL.values():
                got[v] = got.get(v, 0) + 1
        for v, k in sorted(got.items()):
            cv.put(r, 2, lit(f"{b.field} {v}"))
            cv.put(r, 3, f'=COUNTIF({lay.rng(("band", b.name))},{crit(v)})', "num", INT)
            cv.put(r, 4, k, "num", INT)
            cv.put(r, 5, ties(f"C{r}", f"D{r}"), "tie")
            cv.put(r, 6, f'Filter {b.field} band on the Loans sheet to "{v}".', "grey")
            r += 1
    for d in res.config.dimensions:
        got = {}
        for v in labels[d.name]:
            if v in REASON_LABEL.values():
                got[v] = got.get(v, 0) + 1
        for v, k in sorted(got.items()):
            cv.put(r, 2, lit(f"{d.field} {v}"))
            cv.put(r, 3, f'=COUNTIF({lay.rng(("seg", d.name))},{crit(v)})', "num", INT)
            cv.put(r, 4, k, "num", INT)
            cv.put(r, 5, ties(f"C{r}", f"D{r}"), "tie")
            cv.put(r, 6, f'Filter {d.field} on the Loans sheet to "{v}".', "grey")
            r += 1
    _tie_rules(cv, f"E{first}:E{r}")
    _fit_col(cv, 6, 92)
    return cv


# --------------------------------------------------------------------------
# Bands


def _bands(res, lay: Layout, values) -> tuple[Canvas, dict]:
    cv = Canvas(BANDS)
    cv.widths = {1: 2, 2: 26, 3: 18, 4: 18, 5: 16, 6: 16, 7: 8, 8: 50}
    cv.title_band("Bands", "The range of values each band covers.", 8)
    r = cv.note(3, [("How it works", "A loan falls in the band where its value is at least the From value and below "
                                     "the Up to value. On the Loans sheet, each loan's band is assigned by formula "
                                     "from its value and the From column on this sheet, and shown beside the band "
                                     "PocketBook assigned.")],
                2, 8, 120)
    ranges = {}
    first_tie = r
    for b in res.config.bands:
        edges = res.band_edges[b.name]
        each = (res.value_bands or {}).get(b.name)
        seen = [v for v in values[("band", b.name)] if not isinstance(v, str)]
        labels = engine.labels_for(edges, seen, each)
        cv.section(r, lit(b.field), 8)
        r += 1
        typed = "; ".join(_g(e) for e in b.edges) if b.edges else "None entered; PocketBook set the edges."
        how = ("Entered on Columns and used as entered." if b.edges else
               "One band per distinct value, because the column has too few values to divide into equal bands."
               if each else
               f"{b.count} bands set by PocketBook, each with roughly the same number of loans"
               + (", with the edges rounded." if b.cut == "round" else "."))
        for label, v in (("Edges as entered", typed), ("Edges used", "; ".join(_full(e) for e in edges)),
                         ("Method", how)):
            cv.put(r, 2, label, "bold")
            cv.put(r, 3, v)
            cv.merges.append(f"C{r}:H{r}")
            r += 1
        cv.header(r, ["Band", "From (at least)", "Up to (below)", "Loans (Excel)", "Loans (PocketBook)",
                      "Ties?", "By hand"])
        r += 1
        top = r
        lows = [LOW] + list(edges)
        highs = list(edges) + [HIGH]
        counts = {}
        for v in res.row_labels[b.name]:
            counts[v] = counts.get(v, 0) + 1
        col = lay.rng(("band", b.name))
        raw = lay.letter(("raw", b.name))
        for lab, lo, hi in zip(labels, lows, highs):
            cv.put(r, 2, lit(lab), "bold")
            cv.put(r, 3, lo, "num", LOW_FMT)
            cv.put(r, 4, hi, "num", HIGH_FMT)
            cv.put(r, 5, f"=COUNTIF({col},{crit(lab)})", "num", INT)
            cv.put(r, 6, counts.get(lab, 0), "num", INT)
            cv.put(r, 7, ties(f"E{r}", f"F{r}"), "tie")
            words = ("Filter " + b.field + " on the Loans sheet to values "
                     + (f"below {_g(hi)}" if lo == LOW else
                        f"of at least {_g(lo)}" + ("" if hi == HIGH else f" and below {_g(hi)}")))
            cv.put(r, 8, words + ", excluding text entries.", "grey")
            r += 1
        ranges[b.name] = (f"{BANDS}!$B${top}:$B${r - 1}", f"{BANDS}!$C${top}:$C${r - 1}")
        for lab in [x for x in REASON_LABEL.values() if counts.get(x)]:
            cv.put(r, 2, lit(lab), "bold")
            cv.put(r, 3, "no numeric value", "grey")
            cv.put(r, 5, f"=COUNTIF({col},{crit(lab)})", "num", INT)
            cv.put(r, 6, counts[lab], "num", INT)
            cv.put(r, 7, ties(f"E{r}", f"F{r}"), "tie")
            cv.put(r, 8, f'Filter {b.field} on the Loans sheet to "{lab}".', "grey")
            r += 1
        mine, theirs = lay.letter(("band", b.name)), lay.letter(("pb", b.name))
        cv.put(r, 2, "Loans where the formula band differs from PocketBook's band", "bold")
        cv.merges.append(f"B{r}:D{r}")
        cv.put(r, 5, f"=SUMPRODUCT(--({LOANS}!${mine}$2:${mine}${lay.last}<>{LOANS}!${theirs}$2:${theirs}${lay.last}))",
               "numb", INT)
        cv.put(r, 6, 0, "numb", INT)
        cv.put(r, 7, ties(f"E{r}", f"F{r}"), "tie")
        cv.put(r, 8, f"On the Loans sheet, compare the {b.field} band column with the {b.field} band, PocketBook's "
                     f"column.", "grey")
        r += 2
        del raw
    _tie_rules(cv, f"G{first_tie}:G{r}")
    _fit_col(cv, 8, 70)
    return cv, ranges


# --------------------------------------------------------------------------
# One pocket


#: how a By hand step opens, by the loans it starts from
SCOPE_WORDS = {"pocket": "Filter the Loans sheet to the selected pocket (In this pocket = 1)",
               "band": "Filter the Loans sheet to the selected band (In this band = 1)",
               "book": "With no filter on the Loans sheet"}


def _and(words) -> str:
    words = list(words)
    return words[0] if len(words) == 1 else ", ".join(words[:-1]) + " and " + words[-1]


def by_hand(scope: str, exclude: tuple, shows: str, then: str = "") -> str:
    """A By hand step over the Loans sheet: the loans it starts from, the columns whose text entries it excludes
    (exactly the columns the formula requires a number in), and what the status bar shows."""
    lead = SCOPE_WORDS[scope]
    if exclude:
        lead += (", " if scope == "book" else " and ") + "exclude text entries from the " + _and(exclude) + (
            " column" if len(exclude) == 1 else " columns")
    return f"{lead}; the status bar shows the {shows}." + (f" {then}" if then else "")


def _steps(res) -> list[tuple]:
    """One pocket's rows: (id or a section's title, step, in words, written out, Excel's formula, PocketBook's key,
    format, by hand). {id} in a formula is that step's Excel cell; {id|fmt} in the written-out line is its value as
    text."""
    b = res.config.benchmark
    peers = b.compare_to == "peers"
    IP, IB = "Loan_InPocket,1", "Loan_InBand,1"
    ref = "" if peers else " (for reference only; the Run judged pockets against the book)"
    alone = " (shown on RANR vs GCOs only for a pocket alone in its band)" if peers else ""
    return [
        "This pocket",
        ("loans", "Loans", "Records from the population in the selected band and segment.", "{loans|#,##0} loans",
         f"=COUNTIFS({IP})", "loans", INT, by_hand("pocket", (), "Count of the Row in the extract column")),
        ("bad", "Bad loans", "Loans in the pocket with an outcome of 1 (bad).",
         "{bad|#,##0} of {bad_n|#,##0} loans with an outcome",
         f"=SUMIFS(Loan_Bad,{IP})", "bad", INT, by_hand("pocket", (), "Sum of the bad column")),
        ("bad_n", "Loans with an outcome", "Loans in the pocket with an outcome of 0 or 1. Loans with any other value "
         "are excluded from the rate.",
         "{bad_n|#,##0} loans", f"=COUNTIFS({IP},Loan_Bad,{NUM})", "bad_n", INT,
         by_hand("pocket", (), "Numerical Count of the bad column")),
        ("bad_rate", "Bad loan rate", "Bad loans as a share of loans with an outcome.",
         "{bad|#,##0} ÷ {bad_n|#,##0} = {bad_rate|0.000%}", '=IFERROR({bad}/{bad_n},"")', "bad_rate", PCT,
         "Divide bad loans by loans with an outcome (the two rows above)."),
        ("gco_bk", "Booked, loans with a GCO", "Booked dollars for loans in the pocket that have both a booked amount "
         "and a GCO amount. This is the denominator of the GCO rate.", "{gco_bk|$#,##0.00}",
         f"=SUMIFS(Loan_Booked,{IP},Loan_GCO,{NUM})", "gco_bk", USD,
         by_hand("pocket", ("GCO",), "Sum of the booked column")),
        ("gco", "GCOs", "GCO dollars for the same loans.", "{gco|$#,##0.00}",
         f"=SUMIFS(Loan_GCO,{IP},Loan_Booked,{NUM})", "gco", USD,
         by_hand("pocket", ("booked",), "Sum of the GCO column")),
        ("gco_rate", "GCO rate", "GCOs as a share of booked dollars.",
         "{gco|$#,##0} ÷ {gco_bk|$#,##0} = {gco_rate|0.000%}",
         '=IFERROR({gco}/{gco_bk},"")', "gco_rate", PCT, "Divide GCOs by booked dollars (the two rows above)."),
        ("ranr_bk", "Booked, loans with a RANR", "Booked dollars for loans in the pocket that have both a booked "
         "amount and a RANR amount. This is the Booked figure on the RANR vs GCOs tab.", "{ranr_bk|$#,##0.00}",
         f"=SUMIFS(Loan_Booked,{IP},Loan_RANR,{NUM})", "ranr_bk", USD,
         by_hand("pocket", ("RANR",), "Sum of the booked column")),
        ("ranr", "RANR", "RANR dollars for the same loans.",
         "{ranr|$#,##0.00}", f"=SUMIFS(Loan_RANR,{IP},Loan_Booked,{NUM})", "ranr", USD,
         by_hand("pocket", ("booked",), "Sum of the RANR column")),
        ("ranr_rate", "RANR rate", "RANR as a share of booked dollars.",
         "{ranr|$#,##0} ÷ {ranr_bk|$#,##0} = {ranr_rate|0.000%}",
         '=IFERROR({ranr}/{ranr_bk},"")', "ranr_rate", PCT, "Divide RANR by booked dollars (the two rows above)."),
        ("ctb_bk", "Booked, loans with a GCO and a RANR", "Booked dollars for loans in the pocket that have a booked "
         "amount, a GCO amount and a RANR amount. This is the denominator of the RANR + GCOs rate.",
         "{ctb_bk|$#,##0.00}", f"=SUMIFS(Loan_Booked,{IP},Loan_GCO,{NUM},Loan_RANR,{NUM})", "ctb_bk", USD,
         by_hand("pocket", ("GCO", "RANR"), "Sum of the booked column")),
        ("ctb", "RANR + GCOs", "RANR dollars plus GCO dollars for the same loans.", "{ctb|$#,##0.00}",
         f"=SUMIFS(Loan_RANR,{IP},Loan_Booked,{NUM},Loan_GCO,{NUM})"
         f"+SUMIFS(Loan_GCO,{IP},Loan_Booked,{NUM},Loan_RANR,{NUM})", "ctb", USD,
         by_hand("pocket", ("booked", "GCO", "RANR"), "Sum of the RANR column, and then the Sum of the GCO column",
                 "Add the two.")),
        ("ctb_rate", "RANR + GCOs rate", "RANR + GCOs as a share of booked dollars. RANR already has the GCOs "
         "deducted, so this adds them back.", "{ctb|$#,##0} ÷ {ctb_bk|$#,##0} = {ctb_rate|0.000%}",
         '=IFERROR({ctb}/{ctb_bk},"")', "ctb_rate", PCT,
         "Divide RANR + GCOs by booked dollars (the two rows above)."),
        "The whole book",
        ("book_all_bk", "Booked, whole book, every loan with a booked amount", "Booked dollars for every loan in the "
         "book with a booked amount. This is the numerator of Avg line, whole book.", "{book_all_bk|$#,##0.00}",
         "=SUM(Loan_Booked)", "book_all_bk", USD, by_hand("book", (), "Sum of the booked column")),
        ("book_gco_bk", "Booked, whole book, loans with a GCO", "Booked dollars for every loan in the book with both "
         "a booked amount and a GCO amount. This is the denominator of the book's GCO rate.",
         "{book_gco_bk|$#,##0.00}", f"=SUMIFS(Loan_Booked,Loan_GCO,{NUM})", "book_gco_bk", USD,
         by_hand("book", ("GCO",), "Sum of the booked column")),
        ("book_nogco_bk", "Booked, whole book, loans with no GCO amount", "Booked dollars for loans with a booked "
         "amount but no numeric GCO amount. This reconciles the two booked totals above.",
         "{book_all_bk|$#,##0.00} − {book_gco_bk|$#,##0.00} = {book_nogco_bk|$#,##0.00}",
         "={book_all_bk}-{book_gco_bk}", "book_nogco_bk", USD,
         "Subtract the booked dollars for loans with a GCO from the booked dollars for every loan (the two rows "
         "above)."),
        ("book_gco", "GCOs, whole book", "GCO dollars for every loan in the book with both a booked amount and a GCO "
         "amount.", "{book_gco|$#,##0.00}", f"=SUMIFS(Loan_GCO,Loan_Booked,{NUM})", "book_gco", USD,
         by_hand("book", ("booked",), "Sum of the GCO column")),
        ("book_rate", "GCO rate, whole book", "The book's GCOs as a share of its booked dollars for loans with a GCO.",
         "{book_gco|$#,##0} ÷ {book_gco_bk|$#,##0} = {book_rate|0.000%}", '=IFERROR({book_gco}/{book_gco_bk},"")',
         "book_rate", PCT, "Divide the book's GCOs by its booked dollars for loans with a GCO."),
        ("book_ranr_bk", "Booked, whole book, loans with a RANR", "Booked dollars for every loan in the book with "
         "both a booked amount and a RANR amount. This is the Booked figure in the Whole book row on the RANR vs "
         "GCOs tab.", "{book_ranr_bk|$#,##0.00}", f"=SUMIFS(Loan_Booked,Loan_RANR,{NUM})", "book_ranr_bk", USD,
         by_hand("book", ("RANR",), "Sum of the booked column")),
        ("book_ranr", "RANR, whole book", "RANR dollars for the same loans.", "{book_ranr|$#,##0.00}",
         f"=SUMIFS(Loan_RANR,Loan_Booked,{NUM})", "book_ranr", USD,
         by_hand("book", ("booked",), "Sum of the RANR column")),
        ("book_ctb_bk", "Booked, whole book, loans with a GCO and a RANR", "Booked dollars for every loan in the book "
         "with a booked amount, a GCO amount and a RANR amount.", "{book_ctb_bk|$#,##0.00}",
         f"=SUMIFS(Loan_Booked,Loan_GCO,{NUM},Loan_RANR,{NUM})", "book_ctb_bk", USD,
         by_hand("book", ("GCO", "RANR"), "Sum of the booked column")),
        ("book_ctb", "RANR + GCOs, whole book", "RANR dollars plus GCO dollars for the same loans.",
         "{book_ctb|$#,##0.00}",
         f"=SUMIFS(Loan_RANR,Loan_Booked,{NUM},Loan_GCO,{NUM})+SUMIFS(Loan_GCO,Loan_Booked,{NUM},Loan_RANR,{NUM})",
         "book_ctb", USD,
         by_hand("book", ("booked", "GCO", "RANR"), "Sum of the RANR column, and then the Sum of the GCO column",
                 "Add the two.")),
        "GCOs against the rest of the book",
        ("rest_gco_bk", "Booked, rest of the book, loans with a GCO", "Booked dollars for loans with a GCO in the "
         "whole book, less the pocket's.",
         "{book_gco_bk|$#,##0} − {gco_bk|$#,##0} = {rest_gco_bk|$#,##0}", "={book_gco_bk}-{gco_bk}", "rest_gco_bk",
         USD, "Subtract the pocket's booked dollars for loans with a GCO from the book's."),
        ("rest_gco", "GCOs, rest of the book", "GCOs for the whole book, less the pocket's.",
         "{book_gco|$#,##0} − {gco|$#,##0} = {rest_gco|$#,##0}", "={book_gco}-{gco}", "rest_gco", USD,
         "Subtract the pocket's GCOs from the book's."),
        ("rest_rate", "GCO rate, rest of the book", "GCOs for the rest of the book as a share of its booked dollars.",
         "{rest_gco|$#,##0} ÷ {rest_gco_bk|$#,##0} = {rest_rate|0.000%}", '=IFERROR({rest_gco}/{rest_gco_bk},"")',
         "rest_rate", PCT, "Divide the rest of the book's GCOs by its booked dollars (the two rows above)."),
        ("x_book", "× book", "The pocket's GCO rate as a multiple of the whole book's GCO rate.",
         "{gco_rate|0.000%} ÷ {book_rate|0.000%} = {x_book|0.000}×", '=IFERROR({gco_rate}/{book_rate},"")',
         "x_book", MULT, "Divide the pocket's GCO rate by the book's."),
        ("x_rest", "× rest of the book", "The pocket's GCO rate as a multiple of the rest of the book's GCO rate.",
         "{gco_rate|0.000%} ÷ {rest_rate|0.000%} = {x_rest|0.000}×", '=IFERROR({gco_rate}/{rest_rate},"")',
         "x_rest", MULT, "Divide the pocket's GCO rate by the rest of the book's."),
        ("gco_gap", "GCO gap in points, against the book", "The pocket's GCO rate less the rest of the book's, in "
         "percentage points (1 point = 1%).",
         "({gco_rate|0.000%} − {rest_rate|0.000%}) × 100 = {gco_gap|+0.000;-0.000} points",
         '=IFERROR(({gco_rate}-{rest_rate})*100,"")', "gco_gap", PTS,
         "Subtract the rest of the book's GCO rate from the pocket's, then multiply by 100."),
        ("excess_rest", "Dollars above share, against the book",
         "The pocket's GCOs less the GCOs it would have incurred at the rest of the book's rate.",
         "{gco|$#,##0} − {rest_rate|0.000%} × {gco_bk|$#,##0} = {excess_rest|$#,##0}",
         '=IFERROR({gco}-{rest_rate}*{gco_bk},"")', "excess_rest", USD,
         "Multiply the rest of the book's GCO rate by the pocket's booked dollars, and subtract the result from "
         "the pocket's GCOs."),
        "RANR and RANR + GCOs against the rest of the book" + alone,
        ("ranr_rest_rate", "RANR rate, rest of the book", "RANR for the rest of the book as a share of its booked "
         "dollars, for loans with a booked amount and a RANR amount.",
         "({book_ranr|$#,##0} − {ranr|$#,##0}) ÷ ({book_ranr_bk|$#,##0} − {ranr_bk|$#,##0}) = "
         "{ranr_rest_rate|0.000%}",
         '=IFERROR(({book_ranr}-{ranr})/({book_ranr_bk}-{ranr_bk}),"")', "ranr_rest_rate", PCT,
         "Subtract the pocket's RANR from the book's, and the pocket's booked dollars for loans with a RANR from "
         "the book's, then divide the first result by the second."),
        ("ranr_gap", "RANR gap in points, against the book", "The pocket's RANR rate less the rest of the book's, in "
         "percentage points. This is RANR's Gap pts "
         "on RANR vs GCOs for a pocket judged against the book.",
         "({ranr_rate|0.000%} − {ranr_rest_rate|0.000%}) × 100 = {ranr_gap|+0.000;-0.000} points",
         '=IFERROR(({ranr_rate}-{ranr_rest_rate})*100,"")', "ranr_gap", PTS,
         "Subtract the rest of the book's RANR rate from the pocket's, then multiply by 100."),
        ("ranr_usd_rest", "RANR dollars, against the book", "The pocket's RANR less the RANR it would have earned at "
         "the rest of the book's rate. This is RANR's Dollars on RANR vs GCOs for a pocket judged against the book.",
         "{ranr|$#,##0} − {ranr_rest_rate|0.000%} × {ranr_bk|$#,##0} = {ranr_usd_rest|$#,##0}",
         '=IFERROR({ranr}-{ranr_rest_rate}*{ranr_bk},"")', "ranr_usd_rest", USD,
         "Multiply the rest of the book's RANR rate by the pocket's booked dollars for loans with a RANR, and "
         "subtract the result from the pocket's RANR."),
        ("ctb_rest_rate", "RANR + GCOs rate, rest of the book", "RANR + GCOs for the rest of the book as a share of "
         "its booked dollars, for loans with a booked amount, a GCO amount and a RANR amount.",
         "({book_ctb|$#,##0} − {ctb|$#,##0}) ÷ ({book_ctb_bk|$#,##0} − {ctb_bk|$#,##0}) = {ctb_rest_rate|0.000%}",
         '=IFERROR(({book_ctb}-{ctb})/({book_ctb_bk}-{ctb_bk}),"")', "ctb_rest_rate", PCT,
         "Subtract the pocket's RANR + GCOs from the book's, and the pocket's booked dollars for loans with a GCO "
         "and a RANR from the book's, then divide the first result by the second."),
        ("ctb_gap", "RANR + GCOs gap in points, against the book", "The pocket's RANR + GCOs rate less the rest of "
         "the book's, in percentage points. This is RANR + GCOs' Gap pts on RANR vs GCOs for a pocket judged "
         "against the book.",
         "({ctb_rate|0.000%} − {ctb_rest_rate|0.000%}) × 100 = {ctb_gap|+0.000;-0.000} points",
         '=IFERROR(({ctb_rate}-{ctb_rest_rate})*100,"")', "ctb_gap", PTS,
         "Subtract the rest of the book's RANR + GCOs rate from the pocket's, then multiply by 100."),
        ("ctb_usd_rest", "RANR + GCOs dollars, against the book", "The pocket's RANR + GCOs less what it would have "
         "earned at the rest of the book's rate. This is RANR + GCOs' Dollars on RANR vs GCOs for a pocket judged "
         "against the book.",
         "{ctb|$#,##0} − {ctb_rest_rate|0.000%} × {ctb_bk|$#,##0} = {ctb_usd_rest|$#,##0}",
         '=IFERROR({ctb}-{ctb_rest_rate}*{ctb_bk},"")', "ctb_usd_rest", USD,
         "Multiply the rest of the book's RANR + GCOs rate by the pocket's booked dollars for loans with a GCO and "
         "a RANR, and subtract the result from the pocket's RANR + GCOs."),
        "GCOs against the rest of its band" + ref,
        ("band_gco_bk", "Booked, rest of its band, loans with a GCO", "Booked dollars for loans with both a booked "
         "amount and a GCO amount in the same band, across every segment except this one.",
         "{band_gco_bk|$#,##0.00}", f"=SUMIFS(Loan_Booked,{IB},Loan_GCO,{NUM})-{{gco_bk}}", "band_gco_bk", USD,
         by_hand("band", ("GCO",), "Sum of the booked column", "Subtract the pocket's booked dollars for loans with "
                 "a GCO.")),
        ("band_gco", "GCOs, rest of its band", "GCO dollars for the same loans.", "{band_gco|$#,##0.00}",
         f"=SUMIFS(Loan_GCO,{IB},Loan_Booked,{NUM})-{{gco}}", "band_gco", USD,
         by_hand("band", ("booked",), "Sum of the GCO column", "Subtract the pocket's GCOs.")),
        ("band_rate", "GCO rate, rest of its band", "GCOs for the rest of the band as a share of its booked dollars.",
         "{band_gco|$#,##0} ÷ {band_gco_bk|$#,##0} = {band_rate|0.000%}", '=IFERROR({band_gco}/{band_gco_bk},"")',
         "band_rate", PCT, "Divide the two rows above. The cell is blank when the pocket is the only one in its "
         "band."),
        ("x_band", "× rest of its band", "The pocket's GCO rate as a multiple of the rest of its band's GCO rate.",
         "{gco_rate|0.000%} ÷ {band_rate|0.000%} = {x_band|0.000}×", '=IFERROR({gco_rate}/{band_rate},"")',
         "x_band", MULT, "Divide the pocket's GCO rate by the rest of its band's."),
        ("excess_band", "Dollars above share, against its band",
         "The pocket's GCOs less the GCOs it would have incurred at the rest of its band's rate.",
         "{gco|$#,##0} − {band_rate|0.000%} × {gco_bk|$#,##0} = {excess_band|$#,##0}",
         '=IFERROR({gco}-{band_rate}*{gco_bk},"")', "excess_band", USD,
         "Multiply the rest of the band's GCO rate by the pocket's booked dollars, and subtract the result from "
         "the pocket's GCOs."),
        ("dollars", "Dollars above share, as shown",
         "The figure PocketBook reports and ranks pockets by. It is measured against the rest of the band, unless "
         "the Run judged pockets against the book or the pocket is the only one in its band.", "{dollars|$#,##0}",
         f'=IFERROR(IF(INDEX({PBQ}!${pb_col("by_band")}:${pb_col("by_band")},$K$7),{{excess_band}},{{excess_rest}}),"")',
         "dollars", USD, "Equal to one of the two dollars-above-share rows above."),
        "RANR and RANR + GCOs against the rest of its band" + ref,
        ("band_ranr_bk", "Booked, rest of its band, loans with a RANR", "Booked dollars for loans with both a booked "
         "amount and a RANR amount in the same band, across every segment except this one.",
         "{band_ranr_bk|$#,##0.00}", f"=SUMIFS(Loan_Booked,{IB},Loan_RANR,{NUM})-{{ranr_bk}}", "band_ranr_bk", USD,
         by_hand("band", ("RANR",), "Sum of the booked column", "Subtract the pocket's booked dollars for loans with "
                 "a RANR.")),
        ("band_ranr", "RANR, rest of its band", "RANR dollars for the same loans.", "{band_ranr|$#,##0.00}",
         f"=SUMIFS(Loan_RANR,{IB},Loan_Booked,{NUM})-{{ranr}}", "band_ranr", USD,
         by_hand("band", ("booked",), "Sum of the RANR column", "Subtract the pocket's RANR.")),
        ("band_ranr_rate", "RANR rate, rest of its band", "RANR for the rest of the band as a share of its booked "
         "dollars.", "{band_ranr|$#,##0} ÷ {band_ranr_bk|$#,##0} = {band_ranr_rate|0.000%}",
         '=IFERROR({band_ranr}/{band_ranr_bk},"")', "band_ranr_rate", PCT,
         "Divide the two rows above. The cell is blank when the pocket is the only one in its band."),
        ("ranr_gap_band", "RANR gap in points, against its band", "The pocket's RANR rate less the rest of its "
         "band's, in percentage points. This is RANR's Gap pts on RANR vs GCOs for a pocket judged against its band.",
         "({ranr_rate|0.000%} − {band_ranr_rate|0.000%}) × 100 = {ranr_gap_band|+0.000;-0.000} points",
         '=IFERROR(({ranr_rate}-{band_ranr_rate})*100,"")', "ranr_gap_band", PTS,
         "Subtract the rest of the band's RANR rate from the pocket's, then multiply by 100."),
        ("ranr_usd_band", "RANR dollars, against its band", "The pocket's RANR less the RANR it would have earned at "
         "the rest of its band's rate. This is RANR's Dollars on RANR vs GCOs for a pocket judged against its band.",
         "{ranr|$#,##0} − {band_ranr_rate|0.000%} × {ranr_bk|$#,##0} = {ranr_usd_band|$#,##0}",
         '=IFERROR({ranr}-{band_ranr_rate}*{ranr_bk},"")', "ranr_usd_band", USD,
         "Multiply the rest of the band's RANR rate by the pocket's booked dollars for loans with a RANR, and "
         "subtract the result from the pocket's RANR."),
        ("band_ctb_bk", "Booked, rest of its band, loans with a GCO and a RANR", "Booked dollars for loans with a "
         "booked amount, a GCO amount and a RANR amount in the same band, across every segment except this one.",
         "{band_ctb_bk|$#,##0.00}", f"=SUMIFS(Loan_Booked,{IB},Loan_GCO,{NUM},Loan_RANR,{NUM})-{{ctb_bk}}",
         "band_ctb_bk", USD,
         by_hand("band", ("GCO", "RANR"), "Sum of the booked column", "Subtract the pocket's booked dollars for "
                 "loans with a GCO and a RANR.")),
        ("band_ctb", "RANR + GCOs, rest of its band", "RANR dollars plus GCO dollars for the same loans.",
         "{band_ctb|$#,##0.00}",
         f"=SUMIFS(Loan_RANR,{IB},Loan_Booked,{NUM},Loan_GCO,{NUM})"
         f"+SUMIFS(Loan_GCO,{IB},Loan_Booked,{NUM},Loan_RANR,{NUM})-{{ctb}}", "band_ctb", USD,
         by_hand("band", ("booked", "GCO", "RANR"), "Sum of the RANR column, and then the Sum of the GCO column",
                 "Add the two, and subtract the pocket's RANR + GCOs.")),
        ("band_ctb_rate", "RANR + GCOs rate, rest of its band", "RANR + GCOs for the rest of the band as a share of "
         "its booked dollars.", "{band_ctb|$#,##0} ÷ {band_ctb_bk|$#,##0} = {band_ctb_rate|0.000%}",
         '=IFERROR({band_ctb}/{band_ctb_bk},"")', "band_ctb_rate", PCT,
         "Divide the two rows above. The cell is blank when the pocket is the only one in its band."),
        ("ctb_gap_band", "RANR + GCOs gap in points, against its band", "The pocket's RANR + GCOs rate less the rest "
         "of its band's, in percentage points. This is RANR + GCOs' Gap pts on RANR vs GCOs for a pocket judged "
         "against its band.",
         "({ctb_rate|0.000%} − {band_ctb_rate|0.000%}) × 100 = {ctb_gap_band|+0.000;-0.000} points",
         '=IFERROR(({ctb_rate}-{band_ctb_rate})*100,"")', "ctb_gap_band", PTS,
         "Subtract the rest of the band's RANR + GCOs rate from the pocket's, then multiply by 100."),
        ("ctb_usd_band", "RANR + GCOs dollars, against its band", "The pocket's RANR + GCOs less what it would have "
         "earned at the rest of its band's rate. This is RANR + GCOs' Dollars on RANR vs GCOs for a pocket judged "
         "against its band.",
         "{ctb|$#,##0} − {band_ctb_rate|0.000%} × {ctb_bk|$#,##0} = {ctb_usd_band|$#,##0}",
         '=IFERROR({ctb}-{band_ctb_rate}*{ctb_bk},"")', "ctb_usd_band", USD,
         "Multiply the rest of the band's RANR + GCOs rate by the pocket's booked dollars for loans with a GCO and "
         "a RANR, and subtract the result from the pocket's RANR + GCOs."),
        "Avg line",
        ("avg_n", "Loans with a booked amount", "Loans in the pocket with a booked amount. This is the denominator of "
         "Avg line.", "{avg_n|#,##0} loans", f"=COUNTIFS({IP},Loan_Booked,{NUM})", "avg_n", INT,
         by_hand("pocket", (), "Numerical Count of the booked column")),
        ("avg_bk", "Booked, every loan with a booked amount", "Booked dollars for every loan in the pocket with a "
         "booked amount, including loans without a GCO or RANR amount.",
         "{avg_bk|$#,##0.00}", f"=SUMIFS(Loan_Booked,{IP})", "avg_bk", USD,
         by_hand("pocket", (), "Sum of the booked column")),
        ("avg", "Avg line", "The pocket's average committed line: booked dollars per loan with a booked amount.",
         "{avg_bk|$#,##0} ÷ {avg_n|#,##0} = {avg|$#,##0.00}", '=IFERROR({avg_bk}/{avg_n},"")', "avg", USD,
         "Divide the booked dollars by the number of loans with a booked amount (the two rows above)."),
        ("book_avg", "Avg line, whole book", "The same calculation across the whole book: booked dollars per loan "
         "with a booked amount.", "{book_avg|$#,##0.00}",
         "=SUM(Loan_Booked)/COUNT(Loan_Booked)", "book_avg", USD,
         by_hand("book", (), "Average of the booked column")),
        ("avg_x", "Line × book", "The pocket's Avg line as a multiple of the whole book's.",
         "{avg|$#,##0} ÷ {book_avg|$#,##0} = {avg_x|0.000}×", '=IFERROR({avg}/{book_avg},"")', "avg_x", MULT,
         "Divide the pocket's Avg line by the book's."),
        "The shuffle test (GCO rate)",
        ("events", "Loans with a loss", "Loans in the GCO rate with a GCO amount other than 0. A pocket with fewer "
         f"such loans than the Run's minimum ({b.min_events}) is not tested.", "{events|#,##0} loans",
         f'=COUNTIFS({IP},Loan_GCO,">0",Loan_Booked,{NUM})+COUNTIFS({IP},Loan_GCO,"<0",Loan_Booked,{NUM})', "events",
         INT, f"{SCOPE_WORDS['pocket']}, exclude text entries from the GCO and booked columns, and exclude 0 from the "
              "GCO column; the status bar shows the Numerical Count of the GCO column."),
        ("hits", "Shuffles with a gap at least as large, in either direction",
         "The number of the Run's shuffles that produced a gap at least as far from 0 as the pocket's actual gap. "
         "The shuffles are listed on the Shuffle test sheet for the pocket selected at random only.",
         "{hits|#,##0} of " + f"{b.shuffles:,}",
         f"=IF($K$8,'{SHUFFLE}'!$M$2,\"Listed for the pocket selected at random only\")", "hits", INT,
         "On the Shuffle test sheet, count the gaps in the list that are at or beyond the line in either "
         "direction."),
        ("raw", "p-value", "(Shuffles with a gap at least as large + 1) ÷ (total shuffles + 1). Adding 1 counts the "
         "actual assignment as one of the possible outcomes, so the p-value is never 0.",
         "({hits|#,##0} + 1) ÷ " + f"({b.shuffles:,} + 1) = " + "{raw|0.00000}",
         f'=IF({{events}}<{b.min_events},"",IF(ISNUMBER(IF($K$8,{{hits}},F{{hits_row}})),'
         f'(IF($K$8,{{hits}},F{{hits_row}})+1)/({STAMP_SHUFFLES}+1),""))', "raw", PV,
         "Add 1 to both counts, then divide."),
        ("p", "p-value after the allowance for many tests", _allowance_words(b.many_tests), "{p|0.00000}",
         _adjusted(b.many_tests), "p", PV, _allowance_hand(b.many_tests)),
    ]


def _allowance_words(how: str) -> str:
    return {"bh": "The Benjamini-Hochberg adjustment, applied across the pockets tested in this grid on the same "
                  "comparison. It is the smallest value of p × tests ÷ rank among pockets with this p-value or "
                  "higher, capped at 1.",
            "bonferroni": "The Bonferroni adjustment, which multiplies p by the number of pockets tested in this "
                          "grid, capped at 1.",
            "none": "No adjustment was applied; this is the unadjusted p-value."}[how]


def _allowance_hand(how: str) -> str:
    return {"bh": "Use part D of the Shuffle test sheet. It lists every pocket in this grid tested on the same "
                  "comparison with its unadjusted p-value, ranked from smallest to largest, and calculates "
                  "p × the number tested ÷ rank for each. This pocket's adjusted p-value is the smallest of those "
                  "values from its own rank to the bottom of the table, capped at 1.",
            "bonferroni": "Multiply the p-value by the number of pockets tested in the grid on the same comparison, "
                          "capped at 1. Part D of the Shuffle test sheet lists those pockets.",
            "none": "No calculation is needed."}[how]


def _adjusted(how: str) -> str:
    """The p-value after the allowance, from Excel's own p-value and PocketBook's p-values for the rest of the grid
    on _pocketbook, on the same comparison (the band's when this pocket is judged against its band). Pockets are
    ranked by their shuffles' count, a whole number, never by a p-value turned into text for a criterion: the same
    count gives the same p-value, and a count compares exactly."""
    raw, hits = "{raw}", "IF($K$8,{hits},F{hits_row})"
    band = f"INDEX({PBQ}!${pb_col('by_band')}:${pb_col('by_band')},$K$7)"

    def rng(key: str) -> str:
        return f"{PBQ}!${pb_col(key)}$2:${pb_col(key)}${{n}}"

    def one(side: str) -> str:
        if how == "none":
            return raw
        m = f"COUNTIFS({rng('gridno')},Pick_GridNo,{rng('hits_' + side)},\">=0\")"
        if how == "bonferroni":
            return f"MIN(1,{raw}*{m})"
        # Benjamini-Hochberg: the smallest p x m / rank from this pocket's own rank up. Pockets with the same count
        # share the highest rank among them, so this pocket's own term is p x m / (pockets at or under its count)
        rank = f"COUNTIFS({rng('gridno')},Pick_GridNo,{rng('hits_' + side)},\"<=\"&{hits})"
        above = f"COUNTIFS({rng('gridno')},Pick_GridNo,{rng('hits_' + side)},\">\"&{hits})"
        mins = f"_xlfn.MINIFS({rng('q_' + side)},{rng('gridno')},Pick_GridNo,{rng('hits_' + side)},\">\"&{hits})"
        return f"MIN(1,{raw}*{m}/{rank},IF({above}>0,{mins},1))"
    return f'=IFERROR(IF(NOT(ISNUMBER({raw})),"",IF({band},{one("band")},{one("book")})),"")'


def _written(template: str, cell: dict) -> str:
    """{id|fmt} parts as TEXT of that step's Excel cell; the rest as it is."""
    parts = []
    for lit, ref in re.findall(r"([^{]*)(\{[^}]*\})?", template):
        if lit:
            parts.append(q(lit))
        if ref:
            k, _, fmt = ref[1:-1].partition("|")
            parts.append(f'IF(ISNUMBER({cell[k]}),TEXT({cell[k]},{q(fmt)}),"none")')
    return '=IFERROR(' + "&".join(parts) + ',"")'


def pick_words(pick: Pick, dgrid: str) -> str:
    """The pick, as Run stamp and One pocket say it."""
    where = f"{dgrid}: {pick.key[0]}, {pick.key[1]}"
    if not pick.population:
        return f"{where}. No pocket was tested, so the workbook opens on the first pocket of the first grid."
    return (f"{where}, selected at random from the {pick.population:,} tested pockets (every pocket with a GCO "
            f"p-value, across the Run's grids); seed {pick.seed}.")


def _row_lines(text: str, chars: int) -> int:
    return house.lines_at(re.sub(r"\{[^}]*\}", "x" * 14, text), chars) if text else 1


def _one_pocket(res, lay: Layout, dgrid: str, dk, figs, pick: Pick) -> tuple[Canvas, dict]:
    cv = Canvas(ONE)
    cv.tab = house.TAB_YOU
    cv.widths = {1: 2, 2: 34, 3: 46, 4: 52, 5: 18, 6: 18, 7: 8, 8: 58, 10: 26, 11: 20}
    cv.hidden_cols = {10, 11}
    cv.title_band("One pocket", "Every figure for the selected pocket, recalculated from the loan-level data.", 8)
    r = cv.note(3, [
        ("Selection", f"The workbook opens on {pick_words(pick, dgrid)} Select a grid, then a band and a segment, to "
                      "review another pocket. Every figure below updates, as does the In this pocket column on the "
                      "Loans sheet."),
        ("Columns", "Excel's figure is a formula over the Loans sheet, and PocketBook's figure is the value the Run "
                    "produced. Ties? shows ✓ when the two agree to within one billionth of the figure. Calculation "
                    "shows the same arithmetic with this pocket's numbers."),
        ("Notation", 'In a formula, the criterion ">-1E+307" means "is a number". Where a loan has no numeric value in '
                     "a column, the Loans sheet shows the reason in text, and the loan is excluded from the rates "
                     "that use that column."),
    ], 2, 8, 150)
    n = len(figs) + 1
    lists = f"{LISTSQ}!"
    cv.put(r, 2, "GRID", "label")
    cv.put(r, 3, "BAND", "label")
    cv.put(r, 4, "SEGMENT", "label")
    r += 1
    grid = cv.put(r, 2, lit(dgrid), "pick")
    band = cv.put(r, 3, lit(dk[0]), "pick")
    seg = cv.put(r, 4, lit(dk[1]), "pick")
    cv.heights[r] = 20
    ng = len(res.grids)
    for cell, f in ((grid, f"={lists}$A$2:$A${ng + 1}"),
                    (band, f"=OFFSET({lists}$I$2,0,2*($K$1-1),MAX(1,$K$2),1)"),
                    (seg, f"=OFFSET({lists}$J$2,0,2*($K$1-1),MAX(1,$K$3),1)")):
        dv = DataValidation(type="list", formula1=f, allow_blank=False, showErrorMessage=True,
                            errorTitle="Invalid selection", error="Select a value from the list.")
        dv.add(cell.replace("$", ""))
        cv.dvs.append(dv)
    # the helpers, hidden in J:K: rows 1 to 9. Every name is found character for character (find), and the pocket
    # selected at random is recognised with EXACT: a band or segment can hold ~ * ? or begin with < > =
    helpers = [
        ("Grid's number", "=" + find(grid, f"{lists}$A$2:$A${ng + 1}", 1)),
        ("Its bands", f"=INDEX({lists}$D$2:$D${ng + 1},$K$1)"),
        ("Its segments", f"=INDEX({lists}$E$2:$E${ng + 1},$K$1)"),
        ("Band column on Loans", f"=INDEX({lists}$F$2:$F${ng + 1},$K$1)"),
        ("Segment column on Loans", f"=INDEX({lists}$G$2:$G${ng + 1},$K$1)"),
        ("Pocket's key", f'={grid}&"|"&{band}&"|"&{seg}'),
        ("Its row on _pocketbook", f'=IFERROR({find("$K$6", f"{PBQ}!$A$2:$A${n}")},"")'),
        ("The pocket selected at random?", f'=AND(EXACT({grid}&"",{q(dgrid)}),EXACT({band}&"",{q(dk[0])}),'
                                           f'EXACT({seg}&"",{q(dk[1])}))'),
        ("Selected at random", lit(f"{dgrid}: {dk[0]}, {dk[1]}")),
    ]
    for i, (label, f) in enumerate(helpers, start=1):
        cv.put(i, 10, label, "grey")
        cv.put(i, 11, f, "grey")
    r += 1
    reading = (f'=IF($K$7="","No pocket in this grid matches the selection.","At the Run, on GCOs, this pocket '
               f'was "&IF(INDEX({PBQ}!${pb_col("worse_material")}:${pb_col("worse_material")},$K$7),'
               f'"worse and material",'
               f'"not both worse and material")&", judged against the rest of "&'
               f'IF(INDEX({PBQ}!${pb_col("by_band")}:${pb_col("by_band")},$K$7),"its band","the book")&".")')
    cv.put(r, 2, reading, "bold")
    cv.merges.append(f"B{r}:H{r}")
    r += 1
    cv.put(r, 2, f'=IF($K$8,"This is the pocket selected at random; its shuffles are listed on the Shuffle test '
                 f'sheet.",'
                 + q(f"The pocket selected at random is {dgrid}: {dk[0]}, {dk[1]}. The Shuffle test sheet lists the "
                     f"shuffles for that pocket only.") + ")", "grey")
    cv.merges.append(f"B{r}:H{r}")
    r += 2
    cv.header(r, ["Step", "Definition", "Calculation", "Excel's figure", "PocketBook's figure", "Ties?", "By hand"])
    first = r + 1
    r += 1
    steps = _steps(res)
    rows = {}
    rr = r
    for s in steps:
        if isinstance(s, str):
            rr += 1
            continue
        rows[s[0]] = rr
        rr += 1
    cell = {k: f"$E${v}" for k, v in rows.items()}
    hits_cell = None
    for s in steps:
        if isinstance(s, str):
            cv.section(r, s, 8)
            r += 1
            continue
        key, step, words, written, formula, pbk, fmt, hand = s
        f = formula.replace("{n}", str(n))
        for k, v in cell.items():
            f = f.replace("{" + k + "}", v)
        f = f.replace("F{hits_row}", f"$F${rows['hits']}")
        cv.put(r, 2, step, "bold")
        cv.put(r, 3, words)
        cv.put(r, 4, _written(written, {**cell, key: f"$E${r}"}))
        cv.put(r, 5, f, "numb", fmt)
        cv.put(r, 6, f'=IF($K$7="","",INDEX({PBQ}!${pb_col(pbk)}:${pb_col(pbk)},$K$7))', "num", fmt)
        # a difference of two book totals is tied to a billionth of the larger total, not of itself (it can be 0)
        cv.put(r, 7, ties(f"E{r}", f"F{r}", cell["book_all_bk"] if key == "book_nogco_bk" else None)
               if key != "hits" else
               f'=IF($K$8,{ties(f"E{r}", f"F{r}")[1:]},"–")', "tie")
        cv.put(r, 8, hand, "grey")
        # every wrapped cell of the row fits, the step's own label too (the tie-out, 3 Oct 2026: "Booked, every loan
        # with a booked amount" wrapped over the row under it in LibreOffice)
        cv.heights[r] = 13 * max(_row_lines(step, 32), _row_lines(words, 52), _row_lines(written, 56),
                                 _row_lines(hand, 66)) + 4
        if key == "hits":
            hits_cell = (r, 5)
        r += 1
    _tie_rules(cv, f"G{first}:G{r}")
    r += 1
    cv.put(r, 2, "Only the pocket selected at random has its shuffles listed. Any other pocket's shuffles can be "
                 "regenerated from the Run's seed (on Run stamp); the same extract always produces the same "
                 "shuffles.", "grey")
    cv.merges.append(f"B{r}:H{r}")
    return cv, {"grid": grid, "band": band, "seg": seg, "bandcol": "$K$4", "segcol": "$K$5", "rows": rows,
                "hits": hits_cell, "default": "$K$8"}


# --------------------------------------------------------------------------
# The shuffle test


def _draws(res, g, k, by_band: bool):
    """The random pocket's shuffles, dealt again exactly as the Run dealt them: the same rows, the same seed, the
    same grouping (the band's loans only, when judged against its band), the same slice of each shuffled order.
    Returns (every shuffled gap in shuffle order, NaN where none could be worked out; the real gap; the line a
    shuffled gap must reach, |gap| less perm.TIE's allowance; how many reached it)."""
    np = perm.numpy()
    vals = res.per_row["gco_rate"]
    column = perm.Column("gco_rate", [v[0] if v is not None else None for v in vals],
                         [v[1] if v is not None else None for v in vals])
    keys = list(zip(res.row_labels[g.band], res.row_labels[g.dimension]))
    ids = {kk: i for i, (kk, _) in enumerate(g.inner())}
    group = None
    if by_band:
        codes = {lab: i for i, lab in enumerate(sorted({b for b, _ in keys}))}
        group = [codes[b] for b, _ in keys]
    st = perm.RestGap(keep=True)
    s = perm.Structure("the audit's pocket", group, [[ids[x] for x in keys]], {(0, "gco_rate"): st})
    perm.run(res.rows, [column], [s], res.config.benchmark.shuffles, perm.seed_of(engine.SHUFFLE_SEED_NAME))
    at = list(int(x) for x in st.pockets).index(ids[k])
    draws = np.concatenate([d[:, at] for d in st.draws]) if st.draws else np.zeros(0)
    got = st.answers.get(ids[k])
    return draws, float(st.g[at]), float(st.thr[at]), (got.hits if got else None)


#: the 10-loan example: booked and GCO dollars, the first four in the pocket
EXAMPLE = [(20000, 3000), (15000, 0), (25000, 4000), (10000, 0), (30000, 0), (12000, 1500), (18000, 0),
           (22000, 0), (16000, 2000), (14000, 0)]
EXAMPLE_IN = 4
EXAMPLE_SHUFFLES = 20
EXAMPLE_SEED = 20261003


def _shuffle(res, dg, dk, dgrid: str, one: dict, figs, families: dict) -> Canvas:
    cv = Canvas(SHUFFLE)
    cv.widths = {1: 2, 2: 30, 3: 30, 4: 16, 5: 16, 6: 16, 7: 16, 8: 16, 9: 16, 10: 8, 11: 30}
    cv.title_band("Shuffle test", "How a pocket's p-value is derived.", 11)
    b = res.config.benchmark
    r = cv.note(3, [
        ("The test", "The test reassigns the pocket's label to loans at random many times, and counts the "
                     "shuffles that produce a gap at least as large as the actual one."),
        ("What is shuffled", "The assignment of loans to the pocket. The pocket keeps the same number of loans in "
                             "every shuffle. When pockets are judged against the book, every loan in the GCO rate is "
                             "reshuffled. When they are judged against their band, only the loans in the pocket's "
                             "band are reshuffled, so the pocket is never compared with another band."),
        ("The gap", "The pocket's GCO rate (GCOs ÷ booked) less the rate for the rest. This is the same gap shown on "
                    "the One pocket sheet, before it is expressed as a multiple."),
        ("Two-sided", "A shuffle counts when its gap is at least as far from 0 as the actual gap, in either "
                      "direction. A small tolerance, one billionth of the rates' magnitude, ensures that an exact tie "
                      "is not lost to rounding."),
        ("p-value", "(Shuffles that count + 1) ÷ (total shuffles + 1). Adding 1 counts the actual assignment as one "
                    "of the possible outcomes, so the p-value is never 0."),
        ("Adjustment", "Each grid tests many pockets at once, so each p-value is adjusted for multiple tests. Part "
                       "D works out the adjustment, and the last row of the One pocket sheet shows the result."),
        ("The seed", "The seed is fixed (see Run stamp), so the same extract always produces the same shuffles."),
    ], 2, 11, 150)
    # (a) the worked example
    cv.section(r, "A. Worked example: ten loans", 11)
    r += 1
    cv.put(r, 2, f"Loans 1 to {EXAMPLE_IN} make up the pocket. Each shuffle assigns the pocket's label to "
                 f"{EXAMPLE_IN} loans chosen at random. Every number below is a formula that can be checked.", "grey")
    cv.merges.append(f"B{r}:K{r}")
    r += 1
    cv.header(r, ["Loan", "In the pocket?", "Booked", "GCO"])
    r += 1
    top = r
    for i, (bk, gco) in enumerate(EXAMPLE, start=1):
        cv.put(r, 2, i, "numc")
        cv.put(r, 3, "Yes" if i <= EXAMPLE_IN else "No", "textc")
        cv.put(r, 4, bk, "num", USD0)
        cv.put(r, 5, gco, "num", USD0)
        r += 1
    bot = r - 1
    BK, GC = f"$D${top}:$D${bot}", f"$E${top}:$E${bot}"
    np = perm.numpy()
    y = [float(x[1]) for x in EXAMPLE]
    x = [float(x[0]) for x in EXAMPLE]
    ans, st = perm.pocket_vs_rest(y[:EXAMPLE_IN], x[:EXAMPLE_IN], y[EXAMPLE_IN:], x[EXAMPLE_IN:],
                                  shuffles=EXAMPLE_SHUFFLES, seed=EXAMPLE_SEED, keep=True)
    pb_gaps = np.concatenate([d[:, 0] for d in st.draws])
    r += 1
    cv.put(r, 2, "Actual gap", "bold")
    real = (f'=SUMIFS({GC},$C${top}:$C${bot},"Yes")/SUMIFS({BK},$C${top}:$C${bot},"Yes")'
            f'-SUMIFS({GC},$C${top}:$C${bot},"No")/SUMIFS({BK},$C${top}:$C${bot},"No")')
    REAL = cv.put(r, 4, real, "numb", PCT)
    cv.put(r, 5, ans.gap, "num", PCT)
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    cv.put(r, 7, "The pocket's GCOs ÷ booked, less the same ratio for the other six loans. Excel's figure is on "
                 "the left and PocketBook's on the right.", "grey")
    cv.merges.append(f"G{r}:K{r}")
    real_row = r
    r += 2
    cv.header(r, ["Shuffle", "Loans assigned the label", "Pocket GCOs", "Pocket booked", "Pocket rate",
                  "Rest rate", "Gap", "As large?", "PocketBook's gap", "Ties?"])
    r += 1
    s_top = r
    for i in range(EXAMPLE_SHUFFLES):
        order = perm.order_of(np, len(EXAMPLE), EXAMPLE_SEED, i)
        dealt = sorted(int(j) + 1 for j in order[:EXAMPLE_IN])
        cv.put(r, 2, i + 1, "numc")
        cv.put(r, 3, ", ".join(str(j) for j in dealt), "textc")
        cv.put(r, 4, "=" + "+".join(f"INDEX({GC},{j})" for j in dealt), "num", USD0)
        cv.put(r, 5, "=" + "+".join(f"INDEX({BK},{j})" for j in dealt), "num", USD0)
        cv.put(r, 6, f"=D{r}/E{r}", "num", PCT)
        cv.put(r, 7, f"=(SUM({GC})-D{r})/(SUM({BK})-E{r})", "num", PCT)
        cv.put(r, 8, f"=F{r}-G{r}", "numb", PCT)
        cv.put(r, 9, f'=IF(ABS(H{r})>=ABS({REAL})*(1-1E-9),"Yes","No")', "textc")
        cv.put(r, 10, float(pb_gaps[i]), "num", PCT)
        cv.put(r, 11, ties(f"H{r}", f"J{r}"), "tie")
        r += 1
    s_bot = r - 1
    cv.widths[11] = 8
    r += 1
    cv.put(r, 2, "Count", "bold")
    cv.put(r, 3, "Shuffles with a gap at least as large, in either direction")
    cv.heights[r] = 28
    cv.put(r, 4, f'=COUNTIF($I${s_top}:$I${s_bot},"Yes")', "numb", INT)
    cv.put(r, 5, ans.hits, "num", INT)
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    cnt = r
    r += 1
    cv.put(r, 2, "p-value", "bold")
    cv.put(r, 3, f"(count + 1) ÷ ({EXAMPLE_SHUFFLES} + 1)")
    cv.put(r, 4, f"=(D{cnt}+1)/({EXAMPLE_SHUFFLES}+1)", "numb", PV)
    cv.put(r, 5, ans.p, "num", PV)
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    r += 1
    cv.put(r, 2, f"With only {EXAMPLE_SHUFFLES} shuffles, the smallest possible p-value is 1 ÷ "
                 f"{EXAMPLE_SHUFFLES + 1}. The Run uses {b.shuffles:,} shuffles.", "grey")
    cv.merges.append(f"B{r}:K{r}")
    ex_end = r
    del real_row
    r += 2
    # (b) the pocket selected at random
    cv.section(r, f"B. The pocket selected at random: {dgrid}, {dk[0]}, {dk[1]}", 11)
    r += 1
    c = dg.cells[dk]
    s = c.rates["gco_rate"]
    by_band = bool(s.by_band)
    tested = (s.p_band if by_band else s.p_book) is not None
    cv.hidden_cols = {13, 14}
    cv.put(1, 13, None, None)
    if not tested:
        cv.put(r, 2, "This pocket was not tested (too few losses, or no comparison group), so it has no p-value "
                     "and no shuffles are listed.", "text")
        cv.merges.append(f"B{r}:K{r}")
        cv.put(2, 13, "none", None)
        _tie_rules(cv, f"F{top}:F{ex_end}")
        _tie_rules(cv, f"K{s_top}:K{s_bot}")
        return cv
    draws, g, thr, hits = _draws(res, dg, dk, by_band)
    want = s.hits_band if by_band else s.hits_book
    cv.put(r, 2, "Judged against", "bold")
    cv.put(r, 3, "The rest of its band; only the loans in its band are reshuffled." if by_band else
           "The rest of the book; every loan in the GCO rate is reshuffled.")
    cv.merges.append(f"C{r}:K{r}")
    r += 1
    cv.header(r, ["", "", "Excel", "PocketBook", "Ties?"])
    r += 1
    first_tie = r
    cv.put(r, 2, "Actual gap", "bold")
    cv.put(r, 3, "The pocket's GCO rate less the rate for the rest")
    cv.heights[r] = 28
    gcell = f"'{ONE}'!$E${one['rows']['gco_rate']}"
    restc = f"'{ONE}'!$E${one['rows']['band_rate' if by_band else 'rest_rate']}"
    cv.put(r, 4, f"=IF('{ONE}'!{one['default']},{gcell}-{restc},\"\")", "numb", '0.000000%')
    cv.put(r, 5, g, "num", '0.000000%')
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    r += 1
    cv.put(r, 2, "The line", "bold")
    cv.put(r, 3, "The absolute value of the actual gap, less the small tolerance")
    cv.heights[r] = 28
    cv.put(r, 4, thr, "numb", '0.000000000%')
    line_at = r
    r += 1
    # the count reads the list of gaps under parts B to D, so its formulas are written once the list's place is known
    hits_row = r
    cv.put(r, 2, "Shuffles that count", "bold")
    cv.put(r, 3, "Shuffles with a gap at or beyond the line in either direction. A shuffle that leaves the pocket "
                 "with no dollars also counts.")
    cv.heights[r] = 42
    cv.put(r, 5, want, "num", INT)
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    cv.put(2, 13, f"=D{r}", None)              # One pocket reads the count here
    r += 1
    cv.put(r, 2, "Shuffles", "bold")
    cv.put(r, 5, b.shuffles, "num", INT)
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    r += 1
    cv.put(r, 2, "p-value", "bold")
    cv.put(r, 3, "(count + 1) ÷ (shuffles + 1)")
    cv.put(r, 4, f"=(D{hits_row}+1)/(D{hits_row + 1}+1)", "numb", PV)
    cv.put(r, 5, (want + 1) / (b.shuffles + 1) if want is not None else None, "num", PV)
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    r += 1
    cv.put(r, 2, "Regenerated here", "bold")
    cv.put(r, 3, f"This workbook regenerated the {b.shuffles:,} shuffles from the seed and counted {hits:,}; the "
                 f"Run counted {want:,}.", "grey")
    cv.merges.append(f"C{r}:K{r}")
    r += 1
    cv.put(r, 2, "Other pockets", "bold")
    cv.put(r, 3, "Only this pocket's shuffles are listed, because listing every pocket's would take millions of "
                 "rows. Any other pocket's shuffles can be regenerated from the Run's seed; the same extract always "
                 "produces the same shuffles. Part D lists every tested pocket's count and p-value as the Run "
                 "found them.", "grey")
    cv.merges.append(f"C{r}:K{r}")
    cv.heights[r] = 30
    r += 2
    # (c) the textbook cross-check, live for the pocket picked
    cv.section(r, "C. Cross-check: the two-proportion z-test on bad loans", 11)
    r += 1
    cv.put(r, 2, "This is not the test PocketBook applies to GCOs, because it counts loans rather than dollars. It "
                 "is the z-test PocketBook uses for the bad loan rate when a pocket meets the minimum number of "
                 "loans, and it follows the selection on One pocket.", "grey")
    cv.merges.append(f"B{r}:K{r}")
    cv.heights[r] = 30
    r += 1
    cv.header(r, ["", "", "Excel", "PocketBook", "Ties?"])
    r += 1
    OR = one["rows"]
    pick = f"'{ONE}'!$K$7"

    def pbv(k):
        return f'=IF({pick}="","",INDEX({PBQ}!${pb_col(k)}:${pb_col(k)},{pick}))'
    zs = [("x1", "Bad loans, pocket", f"='{ONE}'!$E${OR['bad']}", None),
          ("n1", "Loans with an outcome, pocket", f"='{ONE}'!$E${OR['bad_n']}", None),
          ("x2", "Bad loans, rest of the book", "=SUM(Loan_Bad)-D{x1}", None),
          ("n2", "Loans with an outcome, rest of the book", "=COUNT(Loan_Bad)-D{n1}", None),
          ("pp", "Pooled bad loan rate", "=(D{x1}+D{x2})/(D{n1}+D{n2})", None),
          ("se", "Standard error", "=SQRT(D{pp}*(1-D{pp})*(1/D{n1}+1/D{n2}))", None),
          ("z", "z", '=IFERROR((D{x1}/D{n1}-D{x2}/D{n2})/D{se},"")', "z"),
          ("zp", "Two-sided p-value", '=IFERROR(2*_xlfn.NORM.S.DIST(-ABS(D{z}),TRUE),"")', "z_p")]
    at = {}
    for k, *_ in zs:
        at[k] = r
        r += 1
    r = at["x1"]
    for k, label, f, pk in zs:
        for kk, v in at.items():
            f = f.replace("{" + kk + "}", str(v))
        cv.put(r, 2, label, "bold")
        if house.lines_at(label, 28) > 1:
            cv.heights[r] = 28
        fmt = INT if k in ("x1", "n1", "x2", "n2") else "0.000E+00" if k == "zp" else "0.000000"
        cv.put(r, 4, f, "numb", fmt)
        if pk:
            cv.put(r, 5, pbv(pk), "num", fmt)
            cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
        r += 1
    r += 1
    _tie_rules(cv, f"F{first_tie}:F{r}")
    _tie_rules(cv, f"F{top}:F{ex_end}")
    _tie_rules(cv, f"K{s_top}:K{s_bot}")
    # (d) the allowance for many tests, live for the pocket picked
    r = _many_tests(cv, r, res, one, families)
    # every shuffled gap, under the rest
    r += 1
    cv.section(r, f"Every shuffled gap for the pocket selected at random ({len(draws):,} shuffles)", 11)
    cv.header(r + 1, ["Shuffle", "Gap"])
    D0 = r + 2
    D1 = D0 + len(draws) - 1
    DR = f"$C${D0}:$C${max(D1, D0)}"
    for i, v in enumerate(draws.tolist()):
        cv.put(D0 + i, 2, i + 1, "numc")
        cv.put(D0 + i, 3, v if v == v else "no dollars in the shuffled pocket", "num", '0.000000%')
    cv.put(hits_row, 4, f'=COUNTIF({DR},">="&$D${line_at})+COUNTIF({DR},"<="&-$D${line_at})'
                        f'+COUNTIF({DR},"no dollars*")', "numb", INT)
    cv.put(hits_row + 1, 4, f"=COUNT({DR})+COUNTIF({DR},\"no dollars*\")", "numb", INT)
    return cv


#: part D's rule, written out, by the Run's allowance for many tests
RULE_WORDS = {
    "bh": ("Benjamini-Hochberg. The family is every pocket in the selected pocket's grid tested on the same "
           "comparison (the rest of its band, or the rest of the book). The table lists the family by shuffle count, "
           "smallest first. Rank is the number of pockets in the family with a count at or below this one, so "
           "pockets with the same count share the highest rank among them. Each pocket's term is p × tests ÷ rank. "
           "Its adjusted p-value is the smallest term from its own row to the bottom of the table, capped at 1. The "
           "counts and unadjusted p-values are the Run's; Excel calculates the rest."),
    "bonferroni": ("Bonferroni. The family is every pocket in the selected pocket's grid tested on the same "
                   "comparison (the rest of its band, or the rest of the book). Each adjusted p-value is p × tests, "
                   "capped at 1. The counts and unadjusted p-values are the Run's; Excel calculates the rest."),
    "none": ("The Run applied no allowance for many tests, so each adjusted p-value equals the unadjusted one. The "
             "family is listed for reference."),
}


def _many_tests(cv: Canvas, r: int, res, one: dict, families: dict) -> int:
    """Part D: every pocket tested in the selected pocket's grid on the same comparison, the Run's count and p-value
    for each, ranked, and the allowance for many tests worked out from them in Excel, tied to the Run's adjusted
    p-values and to One pocket's. Returns the row under it."""
    how = res.config.benchmark.many_tests
    cv.section(r, "D. The allowance for many tests", 11)
    r += 1
    cv.put(r, 2, RULE_WORDS[how], "grey")
    cv.merges.append(f"B{r}:K{r}")
    cv.heights[r] = 14 * house.lines_at(RULE_WORDS[how], 150) + 3
    r += 1
    most = max((len(v) for v in families.values()), default=0)
    if not most:
        cv.put(r, 2, "No pocket was tested, so there is no family to adjust.", "text")
        return r + 2
    pickrow = f"'{ONE}'!$K$7"
    side = "$N$3"
    cv.put(3, 14, f'=IF({pickrow}="","",IF(INDEX({PBQ}!${pb_col("by_band")}:${pb_col("by_band")},{pickrow}),'
                  f'"band","book"))', None)
    cv.put(r, 2, f'=IF({side}="","No pocket matches the selection on One pocket.","Family: the pockets in "&Pick_Grid'
                 f'&" tested against the rest of "&IF({side}="band","their band","the book")&".")', "bold")
    cv.merges.append(f"B{r}:K{r}")
    r += 1
    keys, vals, sizes, rows = fam_cols(len(res.grids))
    L = f"{LISTSQ}!${keys}$2:${keys}${1 + sum(len(v) for v in families.values())}"
    V = L.replace(f"${keys}$", f"${vals}$")
    S = L.replace(f"${keys}$", f"${sizes}$")
    R = L.replace(f"${keys}$", f"${rows}$")
    # a family is found by the grid's number, never its name; a pocket by its row on _pocketbook
    fam = f'Pick_GridNo&"|"&{side}&"|'
    cv.header(r, ["#", "Pocket (band, segment)", "Shuffles that count", "p-value", "Rank", "p × tests ÷ rank",
                  "Adjusted p-value, Excel", "Adjusted p-value, PocketBook", "Ties?"])
    cv.heights[r] = 28
    r += 1
    top, bot = r, r + most - 1
    D, F, H = (f"${x}${top}:${x}${bot}" for x in "DFH")
    m_at = bot + 3

    def side_of(key: str, row: int) -> str:
        return (f'IF({side}="band",INDEX({PBQ}!${pb_col(key + "_band")}:${pb_col(key + "_band")},$N{row}),'
                f'INDEX({PBQ}!${pb_col(key + "_book")}:${pb_col(key + "_book")},$N{row}))')
    for i in range(1, most + 1):
        on = f'$M{r}=""'
        cv.put(r, 13, f'=IFERROR(INDEX({V},MATCH({fam}{i}",{L},0)),"")', None)
        cv.put(r, 14, f'=IF({on},"",INDEX({R},MATCH({fam}{i}",{L},0)))', None)
        cv.put(r, 2, f'=IF({on},"",{i})', "numc")
        cv.put(r, 3, f'=IF({on},"",INDEX({PBQ}!${pb_col("band")}:${pb_col("band")},$N{r})&", "&'
                     f'INDEX({PBQ}!${pb_col("seg")}:${pb_col("seg")},$N{r}))', "text")
        cv.put(r, 4, f'=IF({on},"",{side_of("hits", r)})', "num", INT)
        cv.put(r, 5, f'=IF({on},"",{side_of("raw", r)})', "num", PV)
        cv.put(r, 6, f'=IF({on},"",COUNTIF({D},"<="&D{r}))', "num", INT)
        cv.put(r, 7, f'=IF({on},"",E{r}*$D${m_at}/F{r})', "num", PV)
        adj = {"bh": f"MIN(1,G{r}" + (f",H{r + 1})" if i < most else ")"), "bonferroni": f"MIN(1,E{r}*$D${m_at})",
               "none": f"E{r}"}[how]
        cv.put(r, 8, f'=IF({on},"",{adj})', "numb", PV)
        cv.put(r, 9, f'=IF({on},"",{side_of("adj", r)})', "num", PV)
        cv.put(r, 10, f'=IF({on},"",{ties(f"H{r}", f"I{r}")[1:]})', "tie")
        cv.put(r, 11, f"=IF(AND(NOT({on}),$N{r}={pickrow}),\"◀ selected pocket\",\"\")", "numb")
        r += 1
    _tie_rules(cv, f"J{top}:J{bot}")
    r += 1
    cv.header(r, ["", "", "Excel", "PocketBook", "Ties?"])
    r += 1
    cv.put(r, 2, "Pockets tested on this comparison", "bold")
    cv.put(r, 3, "The number of tests the allowance covers")
    cv.heights[r] = 28
    cv.put(r, 4, f"=COUNT({D})", "numb", INT)
    cv.put(r, 5, f'=IFERROR(INDEX({S},MATCH({fam}1",{L},0)),0)', "num", INT)
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    r += 1
    cv.put(r, 2, "The selected pocket's adjusted p-value", "bold")
    cv.put(r, 3, "Excel's figure from this table, against the Run's. One pocket's last row shows the same figure.")
    cv.heights[r] = 54
    cv.put(r, 4, f"=IFERROR(INDEX({H},MATCH({pickrow},$N${top}:$N${bot},0)),\"\")", "numb", PV)
    cv.put(r, 5, f'=IF({pickrow}="","",INDEX({PBQ}!${pb_col("p")}:${pb_col("p")},{pickrow}))', "num", PV)
    cv.put(r, 6, ties(f"D{r}", f"E{r}"), "tie")
    _tie_rules(cv, f"F{m_at}:F{r}")
    del F, one
    return r + 2


# --------------------------------------------------------------------------
# _pocketbook and _lists


class _pocketbook:
    """PocketBook's figures for every pocket, one row each (written straight, without a Canvas)."""

    def __init__(self, n: int):
        self.n = n

    def emit_rows(self, wb, figs) -> None:
        ws = wb.create_sheet(PB)
        ws.sheet_state = "hidden"
        ws.append(PB_KEYS)
        last = len(figs) + 1
        G = f"${pb_col('gridno')}$2:${pb_col('gridno')}${last}"
        for i, (gname, k, f) in enumerate(figs, start=2):
            row = []
            for key in PB_KEYS:
                if key == "key":
                    v = text_cell(ws, f"{gname}|{k[0]}|{k[1]}")
                elif key == "grid":
                    v = text_cell(ws, gname)
                elif key == "band":
                    v = text_cell(ws, k[0])
                elif key == "seg":
                    v = text_cell(ws, k[1])
                elif key.startswith(("m_", "rank_", "q_")):
                    side = key.split("_")[1]
                    raw, hits = f"{pb_col('raw_' + side)}{i}", f"{pb_col('hits_' + side)}{i}"
                    R = f"${pb_col('hits_' + side)}$2:${pb_col('hits_' + side)}${last}"
                    if key.startswith("m_"):
                        v = f'=COUNTIFS({G},${pb_col("gridno")}{i},{R},">=0")'
                    elif key.startswith("rank_"):
                        v = f'=IF(ISNUMBER({raw}),COUNTIFS({G},${pb_col("gridno")}{i},{R},"<="&{hits}),"{NONE}")'
                    else:
                        v = (f'=IF(ISNUMBER({raw}),{raw}*{pb_col("m_" + side)}{i}/{pb_col("rank_" + side)}{i},'
                             f'"{NONE}")')
                else:
                    v = f.get(key)
                    if v is None:
                        v = NONE
                row.append(v)
            ws.append(row)


def fam_cols(ng: int) -> tuple[str, str, str, str]:
    """On _lists, after every grid's band and segment lists: each family's pockets in rank order, as
    "grid's number|side|position" -> the pocket's key on _pocketbook, the family's size, and the pocket's row on
    _pocketbook."""
    first = 2 * ng + 10
    return tuple(get_column_letter(first + i) for i in range(4))


def _families(res, names: dict, number: dict) -> dict[tuple[int, str], list[str]]:
    """Each grid's tested pockets on each comparison (the family the allowance for many tests covers), by the Run's
    shuffle count, smallest first, then by key: (grid's number, "book" or "band") -> the pockets' keys on
    _pocketbook."""
    out = {}
    for g in res.grids:
        for side in ("book", "band"):
            got = []
            for k, c in g.inner():
                s = c.rates["gco_rate"]
                p, hits = (s.p_band, s.hits_band) if side == "band" else (s.p_book, s.hits_book)
                if p is not None and hits is not None:
                    got.append((hits, f"{names[id(g)]}|{k[0]}|{k[1]}"))
            if got:
                out[(number[id(g)], side)] = [key for _, key in sorted(got)]
    return out


def _lists(res, lay: Layout, families: dict, pb_row: dict) -> Canvas:
    cv = Canvas(LISTS)
    cv.hidden = True
    for i, h in enumerate(["Grid", "Band column", "Segment column", "Bands", "Segments", "Band column on Loans",
                           "Segment column on Loans"], start=1):
        cv.put(1, i, h, None)
    for j, g in enumerate(res.grids):
        r = j + 2
        bl = [b for b in g.band_labels if b != ALL]
        dl = [d for d in g.dim_labels if d != ALL]
        cv.put(r, 1, lit(grid_name(res, g)), None)
        cv.put(r, 2, lit(_names(res)[g.band]), None)
        cv.put(r, 3, lit(_names(res)[g.dimension]), None)
        cv.put(r, 4, len(bl), None)
        cv.put(r, 5, len(dl), None)
        cv.put(r, 6, lay.at[("band", g.band)], None)
        cv.put(r, 7, lay.at[("seg", g.dimension)], None)
        for i, v in enumerate(bl, start=2):
            cv.put(i, 9 + 2 * j, lit(v), None)
        for i, v in enumerate(dl, start=2):
            cv.put(i, 10 + 2 * j, lit(v), None)
    keys, *_ = fam_cols(len(res.grids))
    at = column_index_from_string(keys)
    for i, h in enumerate(("Family and position", "Pocket", "Family size", "Row on _pocketbook")):
        cv.put(1, at + i, h, None)
    r = 2
    for (gno, side), pockets in families.items():
        for i, key in enumerate(pockets, start=1):
            cv.put(r, at, f"{gno}|{side}|{i}", None)
            cv.put(r, at + 1, lit(key), None)
            cv.put(r, at + 2, len(pockets), None)
            cv.put(r, at + 3, pb_row[key], None)
            r += 1
    return cv


# --------------------------------------------------------------------------
# Loans: the header and one styled row through openpyxl, every other row straight into the sheet's XML


def _loans_head(wb, lay: Layout) -> None:
    ws = wb.create_sheet(LOANS)
    ws.freeze_panes = "C2"
    widths = {"row": 10, "key": 14, "in_pocket": 10, "in_band": 10}
    for i, (k, h) in enumerate(lay.cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(k, max(12, min(26, len(h) + 2)))
    ws.auto_filter.ref = f"A1:{get_column_letter(lay.width)}{lay.last}"
    head = []
    for _, h in lay.cols:
        c = text_cell(ws, h)                   # a column named "=A" is a name, not a formula
        for k, v in STYLES["head"].items():
            setattr(c, k, v)
        head.append(c)
    ws.append(head)
    fmts = {"booked": USD, "gco": USD, "ranr": USD}
    row = []
    for k, _ in lay.cols:
        c = WriteOnlyCell(ws, value=None)
        c.font = STYLES["num"]["font"]
        c.number_format = fmts.get(k, "General")
        row.append(c)
    ws.append(row)


def _with_loans(data: bytes, lay: Layout, values, res, band_ranges, _template) -> bytes:
    """The saved workbook with the Loans sheet's rows written into its XML: each value a number or words, each band
    and In this ... cell a formula. Row 2, which openpyxl wrote, carries each column's style; every row takes it."""
    zin = zipfile.ZipFile(io.BytesIO(data))
    part = _sheet_part(zin, LOANS)
    xml = zin.read(part).decode("utf-8")
    styles = dict(re.findall(r'<c r="([A-Z]+)2" s="(\d+)"', xml))
    row2 = re.search(r'<row r="2"[^>]*>.*?</row>|<row r="2"[^>]*/>', xml, re.S)
    letters = [get_column_letter(i) for i in range(1, lay.width + 1)]
    lastc = letters[-1]
    cols = []
    for k, _ in lay.cols:
        if k == "row":
            cols.append(("v", list(range(2, lay.n + 2))))
        elif k in ("key", "booked", "gco", "ranr", "bad", "date"):
            cols.append(("v", values[k]))
        elif k[0] == "raw":
            cols.append(("v", values[("band", k[1])]))
        elif k[0] == "pb":
            cols.append(("v", res.row_labels[k[1]]))
        elif k[0] == "seg":
            cols.append(("v", res.row_labels[k[1]]))
        elif k[0] == "band":
            labels, lows = band_ranges[k[1]]
            raw = lay.letter(("raw", k[1]))
            cols.append(("f", f"IF(ISNUMBER({raw}{{r}}),INDEX({labels},MATCH({raw}{{r}},{lows},1)),{raw}{{r}})"))
        # EXACT, character for character: "=" ignores case, so a segment "a" would take segment "A"'s loans too
        elif k == "in_band":
            cols.append(("f", f'IF(EXACT(INDEX($A{{r}}:${lastc}{{r}},Pick_BandCol)&"",Pick_Band&""),1,0)'))
        elif k == "in_pocket":
            cols.append(("f", f'IF(AND(EXACT(INDEX($A{{r}}:${lastc}{{r}},Pick_BandCol)&"",Pick_Band&""),'
                              f'EXACT(INDEX($A{{r}}:${lastc}{{r}},Pick_SegCol)&"",Pick_Seg&"")),1,0)'))
    cols = [(kind, escape(src) if kind == "f" else src) for kind, src in cols]
    sattr = [f' s="{styles[L]}"' if L in styles else "" for L in letters]
    out = []
    for i in range(lay.n):
        r = i + 2
        cells = []
        for j, (kind, src) in enumerate(cols):
            ref = f"{letters[j]}{r}"
            if kind == "f":
                cells.append(f'<c r="{ref}"{sattr[j]}><f>{src.replace("{r}", str(r))}</f></c>')
                continue
            v = src[i]
            if v is None or v == "":
                continue
            if isinstance(v, str):
                # the control characters XML forbids are removed, or the sheet does not open (a loan number
                # "L\x015", the review of 3 Oct 2026); a space at either end is kept as written
                v = clean(v)
                keep = ' xml:space="preserve"' if v != v.strip() else ""
                cells.append(f'<c r="{ref}"{sattr[j]} t="inlineStr"><is><t{keep}>{escape(v)}</t></is></c>')
            else:
                cells.append(f'<c r="{ref}"{sattr[j]}><v>{repr(float(v)) if isinstance(v, float) else v}</v></c>')
        out.append(f'<row r="{r}">' + "".join(cells) + "</row>")
    body = "".join(out)
    if row2:
        xml = xml[:row2.start()] + body + xml[row2.end():]
    else:
        xml = xml.replace("</sheetData>", body + "</sheetData>", 1)
    xml = re.sub(r'<dimension ref="[^"]*"', f'<dimension ref="A1:{lastc}{lay.last}"', xml, count=1)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, xml.encode("utf-8") if item.filename == part else zin.read(item.filename))
    return buf.getvalue()


def _sheet_part(z: zipfile.ZipFile, name: str) -> str:
    from .excel_lists import _sheet_files
    return _sheet_files(z)[name]
