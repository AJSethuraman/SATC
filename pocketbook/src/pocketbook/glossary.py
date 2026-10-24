"""Glossary: every term the tabs use, one a row, with an example in this book's own figures.

The firm, 2 Oct 2026, after explaining the grids' figures to their boss: "Maybe a nice glossary of terms in the
workbook should be there". Written by every Set up and every Run, right after Start here. Its examples use the
whole book's figures from the last Run (`figures`, kept on _found so Set up can write them again); before the first
Run they are made up, and the note says so. Where a tab already words a definition, the words here are that tab's.

Nothing here is a formula: the figures are the Run's, so a line changed on Control since does not move them.

Its words follow pocketbook/VOICE.md; the test keeps the contract-word list and the 28-word sentence cap here.
"""

from __future__ import annotations

import json

from . import house

SHEET = "Glossary"
#: _found's key for the figures, as JSON: what Set up writes the examples from after a Run
FOUND_KEY = "glossary"
TERM, MEANS, EXAMPLE = 2, 3, 4
HEADS = ("Term", "What it means", "Example")
WIDTHS = {1: 2, TERM: 24, MEANS: 62, EXAMPLE: 62}
#: the note's line, before and after the first Run
MADE_UP = "The examples are illustrative until the first Run. Each Run replaces them with this book's own figures."
FROM_RUN = ("Examples naming this book, the last Run or one of its pockets use this book's figures from the last Run "
            "({stamp}). The others are illustrations, with made-up names such as Column A.")
#: the generic book before the first Run (the firm's own example: "the book keeps $4.45 per $100 booked")
GENERIC = {"booked": 40_000_000, "gco": 880_000, "ranr": 1_780_000, "contrib": 2_660_000, "loans": 3_000,
           "bad": 312, "bad_den": 3_000}


def _usd(v: float) -> str:
    return f"${v:,.0f}" if v >= 0 else f"-${-v:,.0f}"


def per100(rate: float) -> str:
    """A rate as dollars per $100 booked: 0.0445 is "$4.45"."""
    v = rate * 100
    return f"${v:.2f}" if v >= 0 else f"-${-v:.2f}"


def _keeps(rate: float) -> str:
    """"keeps $4.45", or "loses $0.50" when RANR is below nought."""
    return f"keeps {per100(rate)}" if rate >= 0 else f"loses {per100(-rate)}"


def _rest(total, c) -> float | None:
    den = total.den - c.den
    return (total.num - c.num) / den if den else None


def figures(res, names: dict[str, str]) -> dict:
    """The whole book's figures from a Run, and the pockets the examples name: plain numbers and words only, so
    they keep as JSON on _found. A figure the Run has no measure for is left out, and its example is the made-up
    one."""
    from . import engine
    t = res.total.rates
    out: dict = {"loans": res.rows}
    if getattr(res, "book_size", None) is not None and res.book_size.booked:
        out["booked"], out["booked_loans"] = res.book_size.booked, res.book_size.loans
    elif "ranr_rate" in t:
        out["booked"] = t["ranr_rate"].den
    for key, m in (("gco", "gco_rate"), ("ranr", "ranr_rate"), ("contrib", "contribution_rate")):
        if m in t and t[m].rate is not None:
            out[key], out[f"{key}_rate"] = t[m].num, t[m].rate
    if "outcome_loans" in t and t["outcome_loans"].den:
        out["bad"], out["bad_den"] = int(t["outcome_loans"].num), int(t["outcome_loans"].den)
    if "outcome_booked" in t and t["outcome_booked"].den:
        out["bad_usd"], out["bad_usd_den"] = t["outcome_booked"].num, t["outcome_booked"].den
    b = res.config.benchmark
    if b is not None:
        out.update(min_units=b.min_units, worse_at=b.worse_at, better_at=b.better_at, confidence=b.confidence,
                   shuffles=b.shuffles, min_events=b.min_events, power=b.power,
                   judged=JUDGED.get(b.compare_to), profit=_profit_of(b))
        line = (getattr(res, "materiality_line", None) or {}).get("gco_rate")
        if line:
            out["material_line"] = line
            kind, v = b.materiality
            if kind == "share":
                out["material_share"] = v
    if b is not None:
        out["many_tests"] = b.many_tests
    if res.config.split:
        out["split"] = {"col": res.config.split[0], "median": res.config.split[1] == "own_median"}
        from .book import _holds_partner, _partner
        pt = _partner(res) if getattr(res, "split_moves_with", None) else None
        loose = next((g for g in getattr(res, "three_way", []) if _holds_partner(res, g) is False), None) if pt \
            else None
        if pt and loose is not None:
            out["split"].update(partner=pt[0], r=pt[1], loose=f"{names.get(loose.band, loose.band)} x "
                                                            f"{names.get(loose.dimension, loose.dimension)}")
    grids = [g for g in getattr(res, "grids", []) if g.band in names and g.dimension in names]
    if grids:
        g = grids[0]
        pockets = list(g.inner())
        out["grid"] = {"name": f"{names[g.band]} x {names[g.dimension]}", "pockets": len(pockets),
                       "bands": len({b_ for (b_, _), _ in pockets}), "segments": len({d for (_, d), _ in pockets})}
        n = sum(1 for g_ in grids for _ in g_.inner())
        if "gco_rate" in t:
            out["worse_material"] = {"n": sum(1 for g_ in grids for _, c in g_.inner()
                                              if engine.worse_and_material(c.rates["gco_rate"])), "of": n}

    def label(g, b_, d):
        seg = f"{names[g.dimension]} {d}" if str(d).replace(".", "").isdigit() else str(d)
        return {"band": f"{names[g.band]} {b_}", "band_col": names[g.band], "seg": seg,
                "seg_col": names[g.dimension], "seg_value": str(d)}

    # the pocket furthest above the rest of the book on GCOs, in dollars, and the one furthest short on RANR
    for key, m, sign in (("gco_pocket", "gco_rate", 1), ("ranr_pocket", "ranr_rate", -1)):
        if m not in t or t[m].rate is None:
            continue
        best = None
        for g in grids:
            for (b_, d), c in g.inner():
                s = c.rates[m]
                rest = _rest(t[m], s)
                if s.rate is None or rest is None or not s.units:
                    continue
                over = sign * (s.num - rest * s.den)
                if best is None or over > best[0]:
                    best = (over, g, b_, d, s, rest, c.rows, c)
        if best and best[0] > 0:
            over, g, b_, d, s, rest, loans, c = best
            out[key] = {**label(g, b_, d), "loans": loans, "rate": s.rate, "rest": rest, "over": over}
            # Could have caught, on Bad loans, for the same pocket
            bad = c.rates.get("outcome_loans")
            if key == "gco_pocket" and bad is not None and bad.smallest_gap is not None:
                out[key]["caught"] = bad.smallest_gap
    once = (getattr(res, "once", {}) or {}).get("gco_rate")
    if once is not None and once.pockets:
        out["once"] = {"dollars": once.dollars, "loans": once.loans, "pockets": once.pockets}
    fv = getattr(res, "filter_values", None)
    if fv and res.config.filter_by:
        out["filter"] = {"col": res.config.filter_by, "value": str(fv[0]), "picked": True}
    elif "gco_pocket" in out:
        p = out["gco_pocket"]
        out["filter"] = {"col": p["seg_col"], "value": p["seg_value"], "picked": False}
    # the band column with the most loans marked missing, and how many: the (marked missing) row's example
    gone = [(g.cells[(engine.MISSING_RULE_LABEL, engine.ALL)].rows, names[g.band]) for g in grids
            if (engine.MISSING_RULE_LABEL, engine.ALL) in g.cells]
    if gone:
        n, col = max(gone, key=lambda x: x[0])
        out["missing"] = {"col": col, "loans": n}
    dates = getattr(res, "dates", None)
    if dates is not None and getattr(dates.first, "year", None) and getattr(dates.last, "year", None):
        out["years"] = [dates.first.year, dates.last.year]
    out["co_date"] = bool(getattr(res.config, "chargeoff_date", None))
    out.update(_verdicts(res, grids, label))
    # the line-size figures, the whole book's (the firm, 17 Oct 2026)
    x = getattr(res, "book_lines", None)
    if x is not None and x.bad_loans and x.good_loans:
        out["line"] = {"bad": x.bad_booked / x.bad_loans, "good": x.good_booked / x.good_loans,
                       "x": x.bad_vs_good}
        if x.lost_share is not None:
            out["line"].update(lost=x.lost_share, lost_gco=x.lost_gco, lost_booked=x.lost_booked)
    return out


#: What a pocket is judged against, as Control words it
JUDGED = {"peers": "The rest of its band", "topline": "The rest of the book"}


def _profit_of(b) -> str | float:
    """The profit line the Run used: "test" (each pocket's own test), "materiality", or points (0.25)."""
    rl = getattr(b, "revenue_line", None)
    return "test" if rl in (None, "luck") else rl


def _verdicts(res, grids, label) -> dict:
    """Significant at N%? and Together as the last Run read them on the two-way grids: a pocket for each of Yes,
    Borderline and No (GCOs ($) first, then Bad loans), and each Together reading's count and first pocket."""
    from . import engine, live, results, stats
    out: dict = {}
    b = res.config.benchmark
    if b is None:
        return out
    bar = round(1 - b.confidence, 12)
    sig: dict = {}
    for m in ("gco_rate", "outcome_loans"):
        for g in grids:
            for (b_, d), c in g.inner():
                s = c.rates.get(m)
                if s is None:
                    continue
                gap = s.vs_band if s.by_band else s.vs_rest
                p = s.p_band if s.by_band else s.p_book
                se = s.se_band if s.by_band else s.se_book
                near = p is not None and se is not None and stats.borderline(p, se, b.confidence)
                word = live.sig_said(s.flag, gap, p, near, False, bar)
                # Yes, worse is shown on a pocket that is also worse, so Worse?'s example can name the same one
                if p is not None and word not in sig and (word != live.SIG_WORSE or s.flag == engine.WORSE):
                    sig[word] = {**label(g, b_, d), "said": word, "measure": results.PLAIN[m],
                                 "gap": gap, "p": p, "p_said": live.p_words(p, bar)}
    if sig:
        out["sig"] = sig
    sides = {k for k, *_ in results.SIDES}
    if all(any(k == m.name for m in res.measures) for k in sides):
        tog: dict = {}
        n = 0
        for g in grids:
            for row in results.pck_rows(res, g):
                n += 1
                if row["together"]:
                    t = tog.setdefault(row["together"], {"n": 0, **label(g, row["band"], row["seg"])})
                    t["n"] += 1
        out["together"] = {"pockets": n, "readings": tog}
    return out


def _plain(x: float) -> str:
    """A number as a plain decimal, never in scientific notation: 1/10001 is 0.00009999, not 9.999e-05."""
    from decimal import Decimal
    return f"{Decimal(f'{x:.6g}'):f}"


def _pocket(p: dict) -> str:
    return f"{p['band']} / {p['seg']}"


def rows(f: dict | None) -> list[tuple[str, str, str]]:
    """Every term, in the firm's order: (term, what it means, its example). `f` is `figures`' dict, or None before
    the first Run."""
    from . import engine, live, stats
    from .choices import NO_DATE
    run = f is not None
    f = f or {}
    g = {**GENERIC, **f}
    gco_r = g.get("gco_rate", GENERIC["gco"] / GENERIC["booked"])
    ranr_r = g.get("ranr_rate", GENERIC["ranr"] / GENERIC["booked"])
    contrib_r = g.get("contrib_rate", GENERIC["contrib"] / GENERIC["booked"])
    real = lambda *keys: run and all(k in f for k in keys)          # noqa: E731
    out = []

    out.append(("Booked / Booked $",
                "The dollar amount committed at origination. For cards, this is the credit line.",
                f"This book booked {_usd(f['booked'])} over {f.get('booked_loans', f['loans']):,} loans."
                if real("booked") else
                "A $10,000 loan is booked at $10,000, and a card with a $5,000 line is booked at $5,000."))
    # the firm, 2 Oct 2026: "booked dollar averages so more easily demonstrate how line assignments look in pockets"
    bl = f.get("booked_loans") if run else None
    avg = f["booked"] / bl if bl and real("booked") else None
    out.append(("Avg line / Line × book",
                "Avg line is the average committed line per loan: booked dollars divided by the number of loans "
                "with a booked amount. Line × book is a pocket's average line divided by the whole book's.",
                f"This book's average line is {_usd(avg)}, so a pocket averaging {_usd(avg * 1.24)} shows 1.24× on "
                f"Line × book." if avg else
                "A pocket with an average line of $6,200, in a book averaging $5,000, shows 1.24× on Line × book."))
    out.append(("GCOs ($)",
                "Gross charge-off (GCO) dollars: the amount charged off on each loan before any recovery, from the "
                "column marked GCO dollars on Columns. Each tab totals them for the loans in a row or pocket.",
                f"This book's GCOs totaled {_usd(f['gco'])}, or {per100(gco_r)} per $100 booked." if real("gco") else
                f"GCOs of {_usd(GENERIC['gco'])} on {_usd(GENERIC['booked'])} booked equal {per100(gco_r)} per $100 "
                f"booked."))
    out.append(("RANR",
                "Risk-adjusted net revenue: interest income and fees, less cost of funds and losses.",
                f"This book's RANR came to {_usd(f['ranr'])}. The book {_keeps(ranr_r)} per $100 booked."
                if real("ranr") else
                f"RANR of {_usd(GENERIC['ranr'])} on {_usd(GENERIC['booked'])} booked means the book "
                f"{_keeps(ranr_r)} per $100 booked."))
    out.append(("RANR + GCOs",
                "RANR with GCOs added back, which shows what the loans earned before losses.",
                f"This book earned {per100(contrib_r)} per $100 booked before losses and {_keeps(ranr_r)} after "
                f"them." if real("contrib", "ranr") else
                f"A book that earns {per100(contrib_r)} per $100 booked before losses and loses {per100(gco_r)} "
                f"{_keeps(ranr_r)}."))
    out.append(("Rate (GCOs ÷ Booked, RANR ÷ Booked)",
                "Dollars per $100 booked: the dollar figure divided by booked dollars. A rate of 2.20% equals $2.20 "
                "per $100 booked.",
                f"This book's GCOs ÷ Booked is {gco_r * 100:.2f}%, or {per100(gco_r)} lost per $100 booked."
                if real("gco") else
                f"GCOs ÷ Booked of {gco_r * 100:.2f}% means {per100(gco_r)} lost per $100 booked."))
    p = f.get("ranr_pocket") if run else None
    out.append(("Points (pts)",
                "The difference between two rates, in percentage points. A pocket at -3.00 pts keeps $3 less per "
                "$100 booked than its comparison.",
                f"{_pocket(p)} {_keeps(p['rate'])} per $100 booked and the rest of the book {_keeps(p['rest'])}. "
                f"The difference is {(p['rate'] - p['rest']) * 100:+.2f} pts." if p else
                "A pocket that keeps $1.45 per $100 booked, against $4.45 for the rest of the book, is at -3.00 pts."))
    p = f.get("gco_pocket") if run else None
    out.append(("× book / × rest of book",
                "A pocket's rate as a multiple of its comparison's: × book against the whole book, × rest of book "
                "against all other loans. 2.00× means twice the GCOs per booked dollar.",
                f"{_pocket(p)} loses {per100(p['rate'])} per $100 booked, which is {p['rate'] / gco_r:.2f}× book and "
                f"{p['rate'] / p['rest']:.2f}× rest of book." if p and real("gco_rate") and p["rest"] else
                "A pocket that loses $4.40 per $100 booked, in a book that loses $2.20, is at 2.00× book."))
    out.append(("Rest of book / rest of band",
                "The group a pocket is compared with: either every other loan in the book or every other loan in "
                "its band. The Judged against setting on Control determines which.",
                f"For {_pocket(p)}, the rest of band is every other {p['band']} loan. The rest of book is every "
                f"other loan in the book." if p else
                "For Column A 620 - 679 / B1, the rest of band is every other Column A 620 - 679 loan."))
    jd = f.get("judged") if run else None
    out.append(("Judged against",
                "The Control setting that picks each pocket's comparison: The rest of its band or The rest of the book. "
                "That comparison decides the gap, the excess and both verdicts. A pocket alone in its band is judged "
                "against the rest of the book.",
                f"At the last Run, Judged against was {jd}. A change on Control updates every verdict on the result "
                f"tabs immediately." if jd else
                "With The rest of its band, Column A 620 - 679 / B1 is compared with every other Column A 620 - 679 "
                "loan."))
    out.append(("Bad loan",
                "A loan flagged as bad in the outcome column, according to the Yes means setting on Columns.",
                f"{f['bad']:,} of this book's {f['bad_den']:,} loans went bad, a rate of "
                f"{f['bad'] / f['bad_den']:.1%}." if real("bad", "bad_den") else
                f"{GENERIC['bad']:,} of {GENERIC['bad_den']:,} loans went bad, a rate of "
                f"{GENERIC['bad'] / GENERIC['bad_den']:.1%}."))
    out.append(("Good loan",
                "A loan whose outcome is not bad: 0 in the outcome column or, with Yes means set on Columns, any value "
                "other than the one named. A blank outcome makes a loan neither good nor bad.",
                f"{f['bad_den'] - f['bad']:,} of this book's {f['bad_den']:,} loans with an outcome are good, and "
                f"{f['bad']:,} are bad." if real("bad", "bad_den") else
                f"Of {GENERIC['bad_den']:,} loans with an outcome, {GENERIC['bad']:,} went bad and "
                f"{GENERIC['bad_den'] - GENERIC['bad']:,} are good."))
    out.append(("Bad loans / Bad loans % / Bad loans ÷ Loans",
                "Bad loans is the number of bad loans in a row or pocket. Bad loans % divides it by the loans with an "
                "outcome of yes or no, and Grids heads the same rate Bad loans ÷ Loans.",
                f"This book has {f['bad']:,} bad loans of {f['bad_den']:,}, so its Bad loans % is "
                f"{f['bad'] / f['bad_den']:.2%}." if real("bad", "bad_den") else
                f"A pocket with {GENERIC['bad']:,} bad loans of {GENERIC['bad_den']:,} has a Bad loans % of "
                f"{GENERIC['bad'] / GENERIC['bad_den']:.2%}."))
    out.append(("Bad dollars / Bad dollars ÷ Booked",
                "Bad dollars are the booked dollars of the loans that went bad. Bad dollars ÷ Booked divides them by "
                "the booked dollars of every loan with an outcome.",
                f"This book's bad loans booked {_usd(f['bad_usd'])} of {_usd(f['bad_usd_den'])}, so Bad dollars ÷ "
                f"Booked is {f['bad_usd'] / f['bad_usd_den']:.2%}." if real("bad_usd", "bad_usd_den") else
                "Bad loans that booked $3,600,000, among loans with an outcome that booked $40,000,000, give Bad "
                "dollars ÷ Booked of 9.00%."))
    out.append(("Loan size",
                "A Grids measure: the average booked dollars per loan in each pocket, with the median. The average is "
                "also shown as a multiple of the whole book's and of the rest of the band's. It is never tested.",
                f"This book's loans average {_usd(avg)} booked, so a pocket averaging {_usd(avg * 1.5)} shows 1.50× "
                f"against the book." if avg else
                "A pocket averaging $7,500 booked per loan, in a book averaging $5,000, shows 1.50× against the "
                "book."))
    out.append(("Band",
                "A range of values in a numeric column. Band edges on Columns set where each band starts; a column "
                "with few distinct values gets one band per value.",
                f"{p['band']} is one band of {p['band_col']}." if p else "Column A 620 - 679 is one band of Column A."))
    out.append(("Segment",
                "A single value of a category column, such as one channel.",
                f"{p['seg']} is one segment of {p['seg_col']}." if p else "B1 is one segment of Column B."))
    out.append(("Pocket",
                "The loans that fall in both one band of a numeric column and one segment of a category column.",
                f"{_pocket(p)} contains the {p['loans']:,} loans in {p['band']} that are also in {p['seg']}." if p
                else "Column A 620 - 679 / B1 contains the loans in Column A 620 - 679 that are also in B1."))
    sp = f.get("split") if run else None
    out.append(("Two-way / Split by",
                "The Pockets dropdown on Pockets. Two-way lists each band crossed with each segment. Split by lists "
                "each of those pockets cut again by the split column chosen in the launcher, each part tested on its "
                "own.",
                (f"This book is split by {sp['col']}: each pocket is halved at its own median {sp['col']}, and the "
                 f"Split tab compares the high half with the low." if sp["median"] else
                 f"This book is split by {sp['col']}: each pocket is cut by each value of {sp['col']}, and each part "
                 f"is tested like any other pocket.") if sp else
                "Split by Column C halves Column A 620 - 679 / B1 at its own median Column C, and each half "
                "is tested like any other pocket."))
    gr = f.get("grid") if run else None
    out.append(("Grid",
                "Every band of one column crossed with every segment of another, giving one pocket for each "
                "combination that has loans.",
                f"{gr['name']} has {gr['bands']} bands by {gr['segments']} segments, for {gr['pockets']} pockets."
                if gr else "Column A x Column B has 5 bands by 3 segments, for 15 pockets."))
    fl = f.get("filter") if run else None
    out.append(("Filter / Only loans where",
                "Restricts Grids and Summary to the loans with one value of the Filter by column. Up to three "
                "filters can be selected; only loans matching every selected value are then shown. × book is still "
                "measured against the whole book.",
                (f"With Only loans where {fl['col']} is {fl['value']}, the grid shows those loans alone."
                 if fl["picked"] else
                 f"Choose {fl['col']} as Filter by in the launcher, then select {fl['value']} to see the grid for "
                 f"those loans alone.") if fl else
                "Choose Column B as Filter by in the launcher, then select B1 to see the grid for B1 loans "
                "alone."))
    yrs = f.get("years") if run else None
    out.append(("Origination year",
                f"The year of the loan's origination date. Loans without a readable date are grouped under "
                f"{NO_DATE}.",
                f"This book's loans were originated from {yrs[0]} to {yrs[1]}. Select a year under Only loans where "
                f"to view that year alone." if yrs else
                "Select 2023 under Only loans where to view only the loans originated in 2023."))
    # the firm, 3 Oct 2026: "Could the summary tab have vintage graphs as well?"; chosen: "Pocket vs rest vs book"
    out.append(("Vintage",
                "The loans originated in a given year. Summary's vintage chart plots one row against the rest of "
                "the book and the whole book, with a point for each year.",
                f"Selecting a row on Summary draws three lines, with one point per year from {yrs[0]} to {yrs[1]}."
                if yrs else
                "Selecting Column A 620 - 679 on Summary draws its line, the rest of the book's and the whole book's, "
                "from 2021 to 2024."))
    out.append(("Across the bottom / Lines by / Panels by",
                "The Compare tab's dropdowns. Across the bottom sets the horizontal axis: a band column's bands, or a "
                "filter's values. Lines by draws one line per value of one filter, and Panels by one chart per value "
                "of the other.",
                (f"With {gr['name'].split(' x ')[0]} across the bottom and Lines by {fl['col']}, each value of "
                 f"{fl['col']} is one line across the {gr['name'].split(' x ')[0]} bands.")
                if fl and fl["picked"] and gr else
                "With Column A across the bottom and Lines by Column B, each Column B value is one line across the "
                "Column A bands."))
    p = f.get("gco_pocket") if run else None
    out.append(("Excess / Bad loans above share / Dollars above share / Dollars short of band / Dollars short of book",
                "The pocket's losses above what it would have lost at the rest's rate, in bad loans for Bad loans and "
                "in dollars for the dollar measures. For RANR and RANR + GCOs, it is the dollars by which the pocket "
                "fell short of the rest.",
                f"{_pocket(p)} lost {_usd(p['over'])} more in GCOs than the rest of the book's rate gives on its "
                f"booked dollars." if p and "over" in p else
                "A pocket that booked $500,000 at a GCO rate of 4.40%, against 2.20% for the rest, has an excess of "
                "$11,000."))
    wm = f.get("worse_material") if run else None
    out.append(("Worse and material",
                "A pocket is worse when its rate is at least Worse at times its comparison's, with a p-value below "
                "the threshold on Control. It is material when its dollars above share reach the materiality "
                "threshold on Control.",
                f"At the last Run, {wm['n']:,} of {wm['of']:,} pockets were worse and material on GCOs." if wm else
                "If 4 of 81 pockets are worse and material, those 4 meet both tests and are listed on Start here."))
    conf = g.get("confidence", 0.95)
    bar = round(1 - conf, 12)
    sig = (f.get("sig") or {}) if run else {}
    w = f.get("worse_at") if run else None
    yes = sig.get(live.SIG_WORSE)
    mt = f.get("many_tests") if run else None
    allowance = ("Control sets no allowance for testing many pockets at once, so the p-value printed is the pocket's "
                 "own." if mt == "none" else
                 "The p-value printed is after the allowance for testing many pockets at once, which holds the share "
                 "of false finds near the line." if mt == "bh" else
                 "The p-value printed is after the allowance for testing many pockets at once, which holds the "
                 "chance of any false find to the line." if mt == "bonferroni" else
                 "The p-value printed is after any allowance for testing many pockets at once set on Control.")
    out.append((f"Worse? ({live.YES} / {live.NOT_SURE} / {live.TOO_FEW} / {live.NO})",
                "Yes means the pocket's rate is at least Worse at times the rest's, and the gap is significant at How "
                "sure. Not sure means the gap passes Worse at but is not significant. Too few losses means fewer "
                "losses than the fewest set on Control, so it is not tested. No covers every other pocket, such as "
                "one under the Worse at line or one not losing more.",
                f"{yes['band']} / {yes['seg']} is {yes['gap']:.2f}× the rest on {yes['measure']}, with "
                f"{yes['p_said']}, so its Worse? reads Yes." if yes and yes.get("gap") else
                "At Worse at 1.25×, a pocket at 1.60× the rest with p = 0.37% reads Yes; at 1.60× with p = 23% it "
                "reads Not sure."))
    ml = f.get("material_line") if run else None
    out.append(("Material? / Material at",
                "Material? reads Yes when the pocket's excess reaches the materiality line on Control, and No when it "
                "does not. Material at, on the tiles, shows that line in the measure's own unit. It is judged apart "
                "from Worse?.",
                f"At the last Run, Material at was {_usd(ml)} on GCOs ($), so a pocket reads Yes once its GCOs above "
                f"share reach {_usd(ml)}." if ml else
                "With Material at $8,800, a pocket $9,500 above share reads Yes, and one $4,000 above reads No."))
    from .results import LINE_SHOW, SHOW
    out.append((f"Show ({' / '.join(SHOW)})",
                f"The Show dropdown on Pockets. {SHOW[0]} lists every pocket losing more than its share. {SHOW[1]} "
                f"lists those whose Worse? and Material? both read Yes, and {SHOW[2]} those whose Worse? reads Yes or "
                f"Not sure.",
                f"At the last Run, {wm['n']:,} of {wm['of']:,} pockets on GCOs ($) were listed under {SHOW[1]}."
                if wm else
                f"A pocket whose Worse? reads Yes and Material? reads No is listed under {SHOW[2]}, not under "
                f"{SHOW[1]}."))
    # the firm, 17 Oct 2026: luck "is basically the other side of confidence", then Significant at N%?
    out.append((f"How sure / {live.SIG_HEAD.format('N%')}",
                f"How sure, on Control, is the confidence level, and {live.SIG_HEAD.format('N%')} on Pockets reads "
                f"each pocket against it. At 95% sure, a gap is significant when its p-value is under 5%. A gap that "
                f"large would then arise by chance less than 1 in 20 times if the pocket were no different from the "
                f"rest. {allowance}",
                f"At the last Run, How sure was {live.pct_words(conf)}, so the line was a p-value of "
                f"{live.pct_words(bar)}, 1 in {1 / bar:,.0f}. {yes['band']} / {yes['seg']} read {yes['said']}, with "
                f"{yes['p_said']}, on {yes['measure']}." if yes else
                f"At 95% sure the line is a p-value of 5%, 1 in 20, so a pocket that loses more at p = 0.37% reads "
                f"{live.SIG_WORSE}."))
    better, no, near = sig.get(live.SIG_BETTER), sig.get(live.SIG_NO), sig.get(live.SIG_BORDER)
    out.append((f"{live.SIG_HEAD.format('N%')} ({live.SIG_WORSE} / {live.SIG_BETTER} / {live.SIG_BORDER} / "
                f"{live.SIG_NO})",
                f"{live.SIG_WORSE} means the pocket's p-value is under the line How sure sets and it loses more than "
                f"the rest; {live.SIG_BETTER}, that it loses less. {live.SIG_NO} means the p-value is at or above the "
                f"line. {live.SIG_BORDER} means a shuffle test's p-value is within {stats.BORDERLINE_SE:g} standard "
                f"errors of that line, on either side.",
                " ".join(f"{x['band']} / {x['seg']} read {x['said']}, with {x['p_said']}, on {x['measure']}."
                         for x in (yes, better, near, no) if x)
                if run and (yes or better or near or no) else
                f"At 95% sure, a pocket losing more at p = 0.37% reads {live.SIG_WORSE}, and one losing less at p = "
                f"0.08% reads {live.SIG_BETTER}. One at p = 23% reads {live.SIG_NO}."))
    pv = yes or better or near or no
    out.append((f"{live.P_HEAD} (the column beside {live.SIG_HEAD.format('N%')})",
                f"The p-value that decides the reading beside it, shown as a percentage. It is a number, so the "
                f"column can be sorted and filtered on. Below 0.01% it shows under 0.01%, and one that would show "
                f"as the line takes more decimals. {allowance}",
                f"Beside {pv['band']} / {pv['seg']}'s {pv['said']} on {pv['measure']}, the column holds "
                f"{_plain(pv['p'])}, shown as {pv['p_said'].split(' ', 1)[1].lstrip('= ')}." if run and pv else
                f"A p-value of 0.0037 shows as {live.p_words(0.0037, 0.05)[4:]}, and at 95% sure one of 0.0498 shows "
                f"as {live.p_words(0.0498, 0.05)[4:]} rather than as the 5% line."))
    out.append(("Borderline",
                f"The shuffle test's p-value is within {stats.BORDERLINE_SE:g} standard errors of the threshold on "
                f"Control, on either side. Pockets reads {live.SIG_BORDER} in {live.SIG_HEAD.format('N%')}; RANR vs "
                f"GCOs, Split and Start here show the flag beside the verdict.",
                f"A Together cell reading Net drain · {stats.borderline_words(bar * 0.9, conf)} has a p-value under "
                f"the {live.pct_words(bar)} threshold and within {stats.BORDERLINE_SE:g} standard errors of it."))
    me = f.get("min_events") if run else None
    out.append((f"{live.FEW_LOSSES} / {live.NO_REST} / {live.NOT_MORE}",
                f"Readings in {live.SIG_HEAD.format('N%')} for a pocket with no p-value. {live.FEW_LOSSES} means "
                f"fewer losses than the fewest set on Control. {live.NO_REST} means the rest has no losses. "
                f"{live.NOT_MORE} means no test was run and the pocket loses no more than the rest.",
                f"At the last Run the fewest losses was {me:,}, so a pocket with {me - 1:,} bad loans reads "
                f"{live.FEW_LOSSES}." if isinstance(me, int) and me > 1 else
                f"With the fewest losses at 10, a pocket with 9 bad loans reads {live.FEW_LOSSES}."))
    sh = int(g.get("shuffles") or 10_000)
    out.append(("Shuffle test / p-value",
                "The shuffle test reassigns loans at random many times and counts the shuffles with a gap at least as "
                "large as the actual one. The p-value is (that count + 1) ÷ (shuffles + 1).",
                f"The loans are shuffled {sh:,} times. A p-value of 0.02 means that about {round(sh * 0.02):,} of "
                f"those shuffles produced a gap at least as large."))
    mu = f.get("min_units") if run else None
    out.append(("Fewest loans in a pocket",
                "A pocket with fewer loans than this is tested with the exact test rather than the usual one. Its "
                "cells on Grids are gray rather than colored.",
                f"At the last Run the minimum was {mu:,.0f}, so any pocket with fewer loans is gray on Grids."
                if isinstance(mu, (int, float)) else
                "With the minimum set at 30, a pocket of 20 loans is gray on Grids."))
    w = f.get("worse_at") if run else None
    b = f.get("better_at") if run else None
    out.append(("Worse at / Better at",
                "Worse at is the multiple of its comparison's rate a pocket must reach to be read as worse. Better "
                "at is the multiple it must fall to in order to be read as better. Both are set on Control.",
                f"At the last Run, Worse at was {w:.2f}× and Better at was {b:.2f}×. Against the book's "
                f"{per100(gco_r)} per $100 booked, a pocket needed {per100(gco_r * w)} to be read as worse."
                if w and b and real("gco_rate") else
                f"With Worse at set to 1.25× and a comparison rate of {per100(gco_r)} per $100 booked, a pocket needs "
                f"{per100(gco_r * 1.25)} to be read as worse."))
    pl = f.get("profit") if run else None
    out.append(("Profit line (How far profit must move before it counts as more or less)",
                "How far RANR or RANR + GCOs must fall short of, or run ahead of, the rest to count as less or more. "
                "With each pocket's own test, a gap counts when it is significant. With a line in points or dollars, "
                "a gap past it counts only when it is also significant; otherwise it is labeled not significant.",
                (f"At the last Run the profit line was each pocket's own test, so a RANR gap counted only when "
                 f"significant at {live.sure_words(conf)}." if pl == "test" else
                 "At the last Run the profit line was the materiality line, so a significant RANR shortfall counted "
                 "when it reached that dollar line." if pl == "materiality" else
                 f"At the last Run the profit line was {float(pl):.2f} points, so a pocket {float(pl) + 0.1:.2f} pts "
                 f"short of the rest, and significant, counted as less.") if pl is not None else
                "With 0.25 points either way, a pocket at -0.30 pts against the rest, and significant, counts as "
                "less."))
    from .results import MIXED
    out.append((f"Holds (column) fixed? / {MIXED}(column)",
                f"These apply to split pockets only. Holds (column) fixed? reads No when the grid does not hold fixed "
                f"the column the split moves with, so the gap may reflect that column. "
                f"{live.SIG_HEAD.format('N%')} then reads {MIXED}(column).",
                f"{sp['col']} moves with {sp['partner']} (correlation {sp['r']:+.2f}), so a split pocket in "
                f"{sp['loose']} reads No: may be mostly {sp['partner']}, and {MIXED}{sp['partner']}."
                if sp and sp.get("loose") else
                f"Column C moves with Column A, so a split pocket in Column D x Column B reads No: may be mostly "
                f"Column A, and {MIXED}Column A."))
    p = f.get("gco_pocket") if run else None
    pw = g.get("power", 0.8)
    out.append(("Could have caught",
                f"The smallest gap a pocket of its size would detect {live.pct_words(pw)} of the time, at the "
                f"confidence the last Run used. It is a multiple of the rest's rate, or points for RANR and RANR + GCOs.",
                f"{_pocket(p)}, with {p['loans']:,} loans, could have caught a gap of {p['caught']:.2f}× on Bad "
                f"loans." if p and "caught" in p else
                f"A pocket of 400 loans that could have caught 1.45× on Bad loans would detect a gap that large "
                f"{live.pct_words(pw)} of the time."))
    tg = f.get("together") if run else None
    out.append(("Together",
                "RANR vs GCOs' reading of a pocket's GCOs and RANR at once, against the rest. A side counts as more "
                "or less only when it passes Worse at or Better at (for RANR, the profit line) and is significant. "
                "A side that does not is about the same.",
                f"At the last Run, {sum(x['n'] for x in tg['readings'].values()):,} of {tg['pockets']:,} pockets on "
                f"RANR vs GCOs had a Together reading." if tg else
                "A pocket with significantly more GCOs and significantly less RANR than the rest reads Net drain."))
    from .results import TOGETHER
    said = {("more", "more"): ("more GCOs and more RANR than the rest, both significant",
                               "at 1.60× the rest on GCOs and +0.80 pts on RANR, both significant"),
            ("more", "less"): ("more GCOs and less RANR than the rest, both significant",
                               "at 1.60× the rest on GCOs and -0.80 pts on RANR, both significant"),
            ("less", "more"): ("fewer GCOs and more RANR than the rest, both significant",
                               "at 0.60× the rest on GCOs and +0.80 pts on RANR, both significant"),
            ("less", "less"): ("fewer GCOs and less RANR than the rest, both significant",
                               "at 0.60× the rest on GCOs and -0.80 pts on RANR, both significant"),
            ("same", "less"): ("less RANR than the rest, significant, while its GCOs count as about the same",
                               "at 1.05× the rest on GCOs and -0.80 pts on RANR, significant"),
            ("more", "same"): ("more GCOs than the rest, significant, while its RANR counts as about the same",
                               "at 1.60× the rest on GCOs, significant, and +0.05 pts on RANR, not significant")}
    for key, (means, gen) in said.items():
        word = TOGETHER[key]
        got = (tg or {}).get("readings", {}).get(word)
        out.append((word, f"The pocket has {means}.",
                    (f"At the last Run, one pocket read {word}: {got['band']} / {got['seg']}." if got and got["n"] == 1
                     else f"At the last Run, {got['n']:,} pockets read {word}, among them {got['band']} / "
                          f"{got['seg']}." if got else f"At the last Run, no pocket read {word}.")
                    if tg else f"A pocket {gen}, reads {word}."))
    o = f.get("once") if run else None
    out.append(("Dollars above their share (each loan once)",
                "The losses pockets incurred above their comparison's rate. Because every loan appears in every "
                "grid, each loan is counted once, in the pocket where it is furthest above its share.",
                f"At the last Run, {_usd(o['dollars'])} of GCOs were above share, across {o['loans']:,} loans in "
                f"{o['pockets']:,} pockets that were worse and material." if o else
                "A loan in two worse pockets, $900 above share in one and $400 in the other, is counted at $900."))
    out.append(("Lifetime-to-date",
                "RANR and GCOs are totals from each loan's origination date to the date of the extract. They are "
                "not annualized.",
                f"A loan originated in {yrs[0]} carries every RANR and GCO dollar recorded from {yrs[0]} to the "
                f"extract date." if yrs
                else "A loan originated in 2021 carries every RANR and GCO dollar recorded from 2021 to the extract "
                     "date."))
    if not run or f.get("co_date", True):   # a book with no charge-off date column never shows the measure
        out.append(("Months to charge-off",
                    "For a loan that charged off, the number of calendar months from its origination date to its "
                    "charge-off date, ignoring the day of the month.",
                    "A loan originated on 15 Jan 2023 and charged off on 3 Mar 2024 counts 14 months. 31 Jan to 1 "
                    "Feb counts as 1 month."))
    # the line-size terms (the firm, 17 Oct 2026), one a row, the Line size tab's literal names
    out += line_rows(f.get("line") if run else None, not run or f.get("co_date", True))
    out.append(("Odd values",
                "Values that may be codes rather than genuine numbers, such as one value far more frequent than any "
                "other, or negative values in a mostly positive column.",
                f"A value of -9999 in Column A on many loans is shown as an odd value. Missing under Treat as on "
                f"Columns places "
                f"those loans in {engine.MISSING_RULE_LABEL}."))
    from .book import TREAT
    mi = f.get("missing") if run else None
    out.append((f"Treat as ({' / '.join(TREAT)})",
                f"The answer on Columns to each odd value. {TREAT[0]} uses the value as recorded. {TREAT[1]} treats it "
                f"as blank, so its loans are placed in {engine.MISSING_RULE_LABEL} and still counted.",
                f"At the last Run, {mi['loans']:,} loans in {mi['col']} were answered {TREAT[1]}." if mi else
                f"Answering {TREAT[1]} for a value of -9999 places those loans in Column A {engine.MISSING_RULE_LABEL}."))
    mi = f.get("missing") if run else None
    out.append((f"{engine.MISSING_RULE_LABEL} / {engine.BLANK_LABEL}",
                f"Where a loan is placed when it cannot be assigned to a band. {engine.BLANK_LABEL} means the cell "
                f"was empty, and {engine.MISSING_RULE_LABEL} means its value was set to Missing on Columns.",
                f"{mi['loans']:,} loans are in {mi['col']} {engine.MISSING_RULE_LABEL} because their value was set "
                f"to Missing. They are still counted." if mi else
                f"A value of -9999 in Column A set to Missing is placed in Column A {engine.MISSING_RULE_LABEL} rather "
                f"than a band, "
                f"and is still counted."))
    out.append(("Changes now / Needs a Run",
                "The two blocks on Control. An answer under Changes now updates every reading, dollar figure, color "
                "and verdict immediately, through formulas. An answer under Needs a Run takes effect at the next Run.",
                f"Changing How sure from 95% to 99% renames {live.SIG_HEAD.format('95%')} to "
                f"{live.SIG_HEAD.format('99%')} immediately. Changing How many bands takes effect at the next Run."))
    ms = f.get("material_share") if run else None
    out.append(("Comes to",
                "The value an answer on Control produces now, beside the answer. A number typed under Or your own "
                "takes precedence over the pick.",
                f"At the last Run, {ms * 100:g}% of the book's total losses came to {_usd(ml)}." if ms and ml else
                "With 1% of the book's total losses picked, Comes to shows that share of the book's GCOs in dollars."))
    out.append(("Worked out from the loans",
                "The value Set up and each Run calculate from this extract for fewest loans, Worse at and Better at, "
                "shown beside the setting. It is never entered as your answer.",
                "A Worse at answer of 1.25 times can sit beside a worked-out value of 1.34x, and the answer stays "
                "1.25 times."))
    out.append(("Last Run used / Status",
                "Last Run used shows each answer on Control as the last Run used it. Status, in the Needs a Run block, "
                "says whether the answer now differs from it.",
                "Status reads Same as last Run until an answer such as How many bands is changed, and the change waits "
                "for the next Run."))
    jd = f.get("judged") if run else None
    out.append(("Lines in use now",
                "The tiles at the top of Pockets, RANR vs GCOs, Line size and Split. They show Worse at, Better at, "
                "How sure, Material at and Judged against, read live from Control.",
                f"At the last Run they read Worse at {w:.2f}×, Better at {f['better_at']:.2f}×, How sure "
                f"{live.pct_words(conf)} "
                f"and Judged against {jd}." if w and f.get("better_at") and jd else
                "The tiles read Worse at 1.25×, Better at 0.80×, How sure 95% and Judged against The rest of its "
                "band."))
    from .summary_chart import GREY_DEFAULT, GREY_LABEL
    out.append((GREY_LABEL,
                f"The number at the top of Summary. A row with fewer loans than it is shown in light gray, and typing "
                f"another number updates the shading immediately. Nothing else reads it.",
                f"With {GREY_LABEL} at {GREY_DEFAULT}, a row of {GREY_DEFAULT - 10} loans is shown in gray."))
    out.append(("Found on / Held back",
                "In a test of a new variable, the loans are cut at the date set on Control. Loans made before it are "
                "where the candidates and their groups are found; only the result on the held-back loans counts.",
                "With the date at 1 Jan 2024, loans made before it find the candidates and loans made on or after "
                "it are held back."))
    out.append(("Pre-spec",
                "The shortlist file a test of a new variable confirms: each candidate column, its groups and its "
                "reference group, written before the held-back loans are read. Record names the file.",
                "A pre-spec naming Column E in 6 groups, with 0.25 - 0.49 as the reference, is confirmed on the "
                "held-back loans."))
    out.append(("Holds up? / Still holds?",
                "Holds up? reads Yes when the gap on the held-back loans is significant and on the same side of 1 as "
                "on the found loans. Still holds? says whether that gap remains with each loan compared only within "
                "its own pocket.",
                "A group at 1.80× on the found loans and 1.60× on the held-back loans, significant, reads Yes under "
                "Holds up?."))
    return out


def _start(key: str) -> str:
    """A line-size line's starting value on Control, as the tab shows it: "1.25×", "80%"."""
    from . import control, live
    s = next(x for x in control.load_settings() if x.key == key)
    return live.line_words(key, s.start().value)


def line_rows(ln: dict | None, co: bool) -> list[tuple[str, str, str]]:
    """The Line size tab's terms: its figures, its signals and Tripped by. `ln` is the whole book's figures from the
    last Run, or None; `co`, whether the book has a charge-off date, without which three of them are never shown."""
    from .results import LINE_SHOW, SIGNALS, TRIPPED_BY
    out = [("Line on bad loans ÷ line on good loans",
            "The average line of the loans that went bad divided by the average line of those that did not. Only "
            "loans with a booked amount count.",
            f"This book's bad loans averaged a line of {_usd(ln['bad'])} and its good loans {_usd(ln['good'])}, "
            f"which is {ln['x']:.2f}×." if ln and ln.get("x") is not None else
            "Bad loans averaging a $6,250 line, against $5,000 for good loans, show 1.25×."),
           ("Share of line lost on bad loans",
            "GCOs on the bad loans divided by their booked dollars, over the bad loans with both amounts.",
            f"This book's bad loans charged off {_usd(ln['lost_gco'])} of {_usd(ln['lost_booked'])} booked, "
            f"{ln['lost']:.1%} of the line." if ln and "lost" in ln else
            "Bad loans that booked $1,000,000 and charged off $800,000 lost 80.0% of the line.")]
    if co:
        out += [("Months to charge-off weighted by line",
                 "The average months to charge-off with each loan counted in proportion to its booked amount: "
                 "Σ(months × booked) ÷ Σ booked.",
                 "A $10,000 loan charged off at 6 months and a $30,000 loan at 18 months average 15 months weighted "
                 "by line, and 12 months straight."),
                ("Line-weighted ÷ straight",
                 "Months to charge-off weighted by line divided by the straight average over the same loans.",
                 "15 months weighted by line against 12 months straight is 1.25×.")]
    n = 4 if co else 3
    out.append(("Signals",
                f"The line-size ratios a pocket passes, each against its line on Control, counted out of {n} and "
                f"named on the Line size tab.",
                f"2 of {n}: {SIGNALS[0][0]}, {SIGNALS[2][0]} means two of the ratios passed their lines."))
    out.append((SIGNALS[0][0], f"Line × book at or above its line on Control, which starts at {_start('line_big')}.",
                f"A pocket averaging $6,500 in a book averaging $5,000 is at 1.30× and trips {SIGNALS[0][0]} at "
                f"1.25×."))
    out.append((SIGNALS[1][0], f"Line on bad loans ÷ line on good loans at or above its line on Control, which "
                               f"starts at {_start('line_bad_bigger')}. It also trips the pocket.",
                f"Bad loans averaging $7,000 against $5,000 for good loans are at 1.40× and trip {SIGNALS[1][0]} at "
                f"1.25×."))
    out.append((SIGNALS[2][0], f"Share of line lost on bad loans at or above its line on Control, which starts at "
                               f"{_start('line_lost')}.",
                f"Bad loans that lost 85.0% of their line trip {SIGNALS[2][0]} at 80%."))
    if co:
        out.append((SIGNALS[3][0], f"Line-weighted ÷ straight at or below its line on Control, which starts at "
                                   f"{_start('line_sooner')}.",
                    f"14 months weighted by line against 16 months straight is 0.88× and trips {SIGNALS[3][0]} at "
                    f"0.90×."))
    out.append(("Tripped by",
                f"Why a pocket is tripped on Line size. {TRIPPED_BY[0]} and {TRIPPED_BY[1]} mean its Worse? reads Yes "
                f"on Bad loans or on GCOs ($), as Tripped by being worse on (Control) selects. {TRIPPED_BY[2]} means "
                f"that signal passed its line.",
                f"A pocket whose Worse? reads Yes on GCOs ($), with the switch on Either, reads {TRIPPED_BY[1]}."))
    out.append((f"Show on Line size ({' / '.join(LINE_SHOW)})",
                f"The Show dropdown on Line size. {LINE_SHOW[0]} lists the pockets with any Tripped by reason. "
                f"{LINE_SHOW[1]} lists those passing at least two of the {n} signals, and {LINE_SHOW[2]} every "
                f"pocket.",
                f"With {LINE_SHOW[1]}, a pocket whose Signals read 1 of {n} is left off the list."))
    return out


def stored(wb, found: str) -> dict | None:
    """The figures the last Run kept on _found, or None before the first Run."""
    if found not in wb.sheetnames:
        return None
    for a, b in wb[found].iter_rows(min_row=1, max_col=2, values_only=True):
        if a == FOUND_KEY and b:
            try:
                return json.loads(b)
            except ValueError:
                return None
    return None


def write(wb, f: dict | None, stamp: str | None = None) -> None:
    """The Glossary tab, afresh, right after Start here."""
    from openpyxl.styles import Alignment, Border, Font, Side
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    at = wb.sheetnames.index("Start here") + 1 if "Start here" in wb.sheetnames else 0
    ws = wb.create_sheet(SHEET, at)
    for c, w in WIDTHS.items():
        ws.column_dimensions[house._letter(c)].width = w
    house.title_band(ws, SHEET, "Definitions of the terms used on the tabs, each with an example.", TERM, EXAMPLE,
                     tab=house.TAB_RECORD)
    r = house.method_note(ws, 3, TERM, EXAMPLE, [
        ("Examples", MADE_UP if f is None else FROM_RUN.format(stamp=stamp) if stamp else
         FROM_RUN.replace(" ({stamp})", "")),
        ("Not live", "A setting changed on Control after a Run updates the result tabs, but not these examples."),
    ])
    house.header(ws, r, TERM, list(HEADS), centre_from=len(HEADS))
    head = r
    rule = Border(bottom=Side(style="thin", color=house.ROW_RULE))
    chars = {c: int(WIDTHS[c] * 1.15) for c in (TERM, MEANS, EXAMPLE)}      # as house.method_note reckons a line
    for term, means, example in rows(f):
        r += 1
        for col, v in ((TERM, term), (MEANS, means), (EXAMPLE, example)):
            c = ws.cell(row=r, column=col, value=v)
            c.font = Font(name="Calibri", size=10, bold=col == TERM, color=house.INK_TEXT)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = rule
        lines = max(house.lines_at(v, chars[col]) for col, v in ((TERM, term), (MEANS, means), (EXAMPLE, example)))
        ws.row_dimensions[r].height = 14 * lines + 4
    ws.freeze_panes = f"A{head + 1}"
    ws.print_title_rows = f"{head}:{head}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
