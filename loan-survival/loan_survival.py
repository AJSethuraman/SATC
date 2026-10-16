#!/usr/bin/env python3
"""
loan_survival.py - One-run loan performance testing with time handled properly.

What one run does, in order:
  1. BAD EVENT CHECK  Is 90+ days past due a good early stand-in for charge-off?
                      (what happened within N months of first 90+, lead time, charge-offs that skipped 90+)
  2. EARLY READ       From fully seasoned vintages: what share of eventual bads show up by month 6, 12, 18...
                      Tells you how much an early comparison of young cohorts can be trusted.
  3. TESTS            Any number of comparisons, each compared at the SAME month on book:
                        group_by = year          origination-year vintages
                        split_date = 2024-03-01  before/after a policy change (comma list allowed)
                        column + cutoffs         e.g. fico below vs at-or-above 640, 660, 680, 700
                        column (no cutoffs)      each value of a flag or category
                      Each can be limited to an origination window (orig_from / orig_to), e.g. seasoned loans only.
                        model = a, b per -50, c  several factors at once (Cox regression), alone vs together
  4. ECONOMICS        If a net revenue column is mapped: total, per loan, per $1,000 of amount, by group.

Usage:
    python loan_survival.py loans.xlsx                          # defaults: bad-event check + vintage test
    python loan_survival.py loans.xlsx --write-settings         # writes a starter settings file to edit
    python loan_survival.py loans.xlsx --settings tests.ini     # run every test in the settings file

Input: one row per loan. Map your column names in the [columns] section of the settings file.
    Required: orig_date, chargeoff_date (any 'went bad' date works here)
    Optional: loan_id, first_90_date, payoff_date, dpd_asof, net_revenue, amount, plus any test columns.

Requires pandas, numpy, openpyxl. Optional: matplotlib (charts), scipy (p-values; built-in fallback).
"""
from __future__ import annotations

import argparse
import configparser
import hashlib
import io
import math
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

VERSION = "2.1"


class UserError(Exception):
    """Something wrong with the file or settings, worded for the person running it (CLI prints it, GUI shows it)."""

Z95 = 1.959963984540054

STARTER = """\
; loan_survival.py settings. Lines starting with ; are comments.
; Every section named [test: ...] runs as its own comparison.

[data]
; as_of = 2026-09-30         ; leave commented to use the latest date in the file

[columns]
; left = what the script calls it, right = YOUR column name. Delete lines you don't have.
loan_id        = loan_id
orig_date      = orig_date
chargeoff_date = chargeoff_date
first_90_date  = first_90_date
payoff_date    = payoff_date
dpd_asof       = dpd_asof
net_revenue    = net_revenue
amount         = amount

[settings]
bad_event          = auto      ; auto | chargeoff | earliest  (earliest = first of 90+ or charge-off)
checkpoints        = 6,12,18,24,36
min_at_risk        = 30        ; don't report a month once fewer loans than this remain
dq_window          = 12        ; months after first 90+ to watch for charge-off
proxy_min_co_share = 0.60      ; auto uses 90+ only if at least this share of 90+ loans charge off
proxy_min_lead     = 3         ; ...and 90+ arrives at least this many months earlier (median)
early_read_horizon = 36        ; month treated as 'eventual' in the early-read table

[test: vintage]
group_by = year

; [test: before vs after policy change]
; split_date = 2024-03-01

; [test: FICO cutoffs on seasoned loans]
; column    = fico
; cutoffs   = 640,660,680,700
; orig_from = 2021-01-01
; orig_to   = 2022-12-31

; [test: what drives losses]
; model   = fico per -50, income_gt_sales, new_business   ; 'per -50' = hazard ratio per 50 points LOWER
; orig_to = 2022-12-31

; [test: income above sales]
; column  = income_gt_sales
; orig_to = 2022-12-31
"""


# ============================================================================ stats

def _gammaincc(a: float, x: float) -> float:
    """Regularized upper incomplete gamma Q(a, x) (Numerical Recipes)."""
    if x <= 0:
        return 1.0
    gln = math.lgamma(a)
    if x < a + 1:
        ap, s, d = a, 1.0 / a, 1.0 / a
        for _ in range(1000):
            ap += 1
            d *= x / ap
            s += d
            if abs(d) < abs(s) * 1e-15:
                break
        return max(0.0, 1.0 - s * math.exp(-x + a * math.log(x) - gln))
    b = x + 1 - a
    c, d = 1e300, 1.0 / b
    h = d
    for i in range(1, 1000):
        an = -i * (i - a)
        b += 2
        d = an * d + b
        d = 1e-300 if abs(d) < 1e-300 else d
        c = b + an / c
        c = 1e-300 if abs(c) < 1e-300 else c
        d = 1.0 / d
        delta = d * c
        h *= delta
        if abs(delta - 1) < 1e-15:
            break
    return math.exp(-x + a * math.log(x) - gln) * h


def chi2_sf(x: float, df: int) -> float:
    try:
        from scipy.stats import chi2
        return float(chi2.sf(x, df))
    except ImportError:
        return _gammaincc(df / 2.0, x / 2.0)


def months_between(start: pd.Series, end: pd.Series) -> pd.Series:
    """Whole calendar months (opened 3/15, observed to 4/14 = 0 months)."""
    m = (end.dt.year - start.dt.year) * 12 + (end.dt.month - start.dt.month)
    return m - (end.dt.day < start.dt.day).astype(int)


def kaplan_meier(time: np.ndarray, event: np.ndarray) -> pd.DataFrame:
    """Monthly Kaplan-Meier table, months 0..max. Events at a month are processed before censoring."""
    rows, surv, gw = [], 1.0, 0.0
    for t in range(int(time.max()) + 1 if len(time) else 0):
        n = int((time >= t).sum())
        d = int(((time == t) & (event == 1)).sum())
        c = int(((time == t) & (event == 0)).sum())
        if n > 0 and d > 0:
            surv *= 1 - d / n
            gw = gw + d / (n * (n - d)) if n > d else math.inf
        lo = hi = np.nan
        if 0 < surv < 1 and math.isfinite(gw):
            se = math.sqrt(gw) / abs(math.log(surv))
            lo, hi = surv ** math.exp(Z95 * se), surv ** math.exp(-Z95 * se)
        elif surv == 1:
            lo = hi = 1.0
        rows.append(dict(month_on_book=t, at_risk=n, bads=d, censored=c,
                         cum_bad_rate=1 - surv, cum_bad_lo95=1 - hi, cum_bad_hi95=1 - lo))
    return pd.DataFrame(rows)


def cum_bad_at(km: pd.DataFrame, month: int, min_at_risk: int) -> float:
    hit = km[km["month_on_book"] == month]
    if hit.empty or hit["at_risk"].iloc[0] < min_at_risk:
        return np.nan
    return float(hit["cum_bad_rate"].iloc[0])


def logrank(time: np.ndarray, event: np.ndarray, group: np.ndarray) -> tuple[float, int, float]:
    labels = list(pd.unique(group))
    k = len(labels)
    if k < 2 or event.sum() == 0:
        return np.nan, 0, np.nan
    O, E, V = np.zeros(k), np.zeros(k), np.zeros((k, k))
    for t in np.unique(time[event == 1]):
        at_risk = time >= t
        n = at_risk.sum()
        d = ((time == t) & (event == 1)).sum()
        ng = np.array([(at_risk & (group == g)).sum() for g in labels], float)
        dg = np.array([((time == t) & (event == 1) & (group == g)).sum() for g in labels], float)
        O += dg
        E += d * ng / n
        if n > 1:
            p = ng / n
            V += d * (n - d) / (n - 1) * (np.diag(p) - np.outer(p, p))
    diff = (O - E)[:-1]
    stat = float(diff @ np.linalg.pinv(V[:-1, :-1]) @ diff)
    return stat, k - 1, chi2_sf(stat, k - 1)


def cox_fit(time: np.ndarray, event: np.ndarray, X: np.ndarray, iters: int = 50):
    """Cox proportional hazards, Efron ties, Newton-Raphson. Returns (beta, se, pvalues)."""
    n, k = X.shape
    beta = np.zeros(k)
    times = np.unique(time[event == 1])
    for _ in range(iters):
        eta = X @ beta
        eta -= eta.max()
        w = np.exp(eta)
        grad, hess = np.zeros(k), np.zeros((k, k))
        for t in times:
            R = time >= t
            D = R & (time == t) & (event == 1)
            d = int(D.sum())
            wR, wD = w[R], w[D]
            XR, XD = X[R], X[D]
            s0R, s1R, s2R = wR.sum(), wR @ XR, (XR * wR[:, None]).T @ XR
            s0D, s1D, s2D = wD.sum(), wD @ XD, (XD * wD[:, None]).T @ XD
            grad += XD.sum(axis=0)
            for l in range(d):
                f = l / d
                s0, s1, s2 = s0R - f * s0D, s1R - f * s1D, s2R - f * s2D
                m = s1 / s0
                grad -= m
                hess -= s2 / s0 - np.outer(m, m)
        step = np.linalg.solve(hess, grad)
        beta -= step
        if np.max(np.abs(step)) < 1e-9:
            break
    se = np.sqrt(np.diag(np.linalg.inv(-hess)))
    p = np.array([chi2_sf((b / s) ** 2, 1) for b, s in zip(beta, se)])
    return beta, se, p


# ============================================================================ settings

class Settings:
    def __init__(self, text: str):
        cp = configparser.ConfigParser(inline_comment_prefixes=(";", "#"), interpolation=None)  # % in a column name is literal
        cp.optionxform = str
        cp.read_string(text)
        self.text = text
        self.as_of = cp.get("data", "as_of", fallback="").strip() or None
        self.cols = {k: v.strip() for k, v in cp["columns"].items()} if cp.has_section("columns") else {}
        s = cp["settings"] if cp.has_section("settings") else {}
        g = lambda k, d: (s.get(k, d) if s else d)
        self.bad_event = str(g("bad_event", "auto")).strip().lower()
        self.checkpoints = [int(x) for x in str(g("checkpoints", "6,12,18,24,36")).split(",") if x.strip()]
        self.min_at_risk = int(g("min_at_risk", 30))
        self.dq_window = int(g("dq_window", 12))
        self.proxy_min_co_share = float(g("proxy_min_co_share", 0.60))
        self.proxy_min_lead = float(g("proxy_min_lead", 3))
        self.horizon = int(g("early_read_horizon", 36))
        self.tests = []
        for sec in cp.sections():
            if sec.lower().startswith("test:"):
                self.tests.append((sec.split(":", 1)[1].strip(), {k: v.strip() for k, v in cp[sec].items()}))
        if self.bad_event not in ("auto", "chargeoff", "earliest"):
            raise UserError(f"bad_event must be auto, chargeoff or earliest (got '{self.bad_event}')")

    def col(self, key: str) -> str | None:
        return self.cols.get(key) or None


# ============================================================================ data

def load(path: Path) -> pd.DataFrame:
    return pd.read_excel(path) if path.suffix.lower() in (".xlsx", ".xlsm", ".xls") else pd.read_csv(path)


def prepare(raw: pd.DataFrame, st: Settings):
    """Parse dates, apply exclusions. Returns cleaned frame (original columns + parsed dates) and exclusions."""
    df = raw.copy()
    excluded = []

    def drop(mask, reason):
        nonlocal df
        mask = mask.fillna(False) if hasattr(mask, "fillna") else mask
        if mask.any():
            x = df[mask].copy()
            x.insert(0, "exclusion_reason", reason)
            excluded.append(x)
            df = df[~mask]

    for key in ("orig_date", "chargeoff_date"):
        if not st.col(key):
            raise UserError(f"[columns] must map '{key}'.")
    missing = [v for k, v in st.cols.items() if v not in df.columns]
    if missing:
        raise UserError(f"Mapped column(s) not found in file: {missing}\nColumns in file: {list(df.columns)}")

    for key in ("orig_date", "chargeoff_date", "first_90_date", "payoff_date"):
        c = st.col(key)
        if c:
            parsed = pd.to_datetime(df[c], errors="coerce")
            bad = df[c].notna() & (df[c].astype(str).str.strip() != "") & parsed.isna()
            df[c] = parsed
            drop(bad, f"unparseable date in {c}")
    if st.col("loan_id"):
        drop(df[st.col("loan_id")].duplicated(keep="first"), "duplicate loan id")
    o = st.col("orig_date")
    drop(df[o].isna(), "missing origination date")
    for key in ("chargeoff_date", "first_90_date", "payoff_date"):
        c = st.col(key)
        if c:
            drop(df[c] < df[o], f"{c} before origination")
    excl = pd.concat(excluded) if excluded else pd.DataFrame(columns=["exclusion_reason"])
    return df, excl


def add_time_event(df: pd.DataFrame, st: Settings, as_of: pd.Timestamp, bad_cols: list[str]) -> pd.DataFrame:
    df = df[df[st.col("orig_date")] <= as_of].copy()
    bad = df[bad_cols].min(axis=1)
    bad = bad.where(bad <= as_of)
    pif = df[st.col("payoff_date")] if st.col("payoff_date") else pd.Series(pd.NaT, index=df.index)
    obs_end = pif.where(pif < as_of, as_of).fillna(as_of)
    is_bad = bad.notna() & (bad <= obs_end)
    end = bad.where(is_bad, obs_end)
    df["_time"] = months_between(df[st.col("orig_date")], end).astype(int)
    df["_event"] = is_bad.astype(int)
    return df


# ============================================================================ 1. bad event check

def bad_event_check(df: pd.DataFrame, st: Settings, as_of: pd.Timestamp):
    f90c, coc, pifc, dpdc = st.col("first_90_date"), st.col("chargeoff_date"), st.col("payoff_date"), st.col("dpd_asof")
    if not f90c:
        return None, None, False, "No first_90_date column mapped; using charge-off as the bad event."
    f90, co = df[f90c], df[coc]
    pif = df[pifc] if pifc else pd.Series(pd.NaT, index=df.index)
    co_any = co.notna() & (co <= as_of)
    co_total, co_skip = int(co_any.sum()), int((co_any & (f90.isna() | (f90 > co))).sum())

    pop = df[f90.notna() & (f90 <= as_of - pd.DateOffset(months=st.dq_window))]
    pop = pop[~(pop[coc] < pop[f90c]).fillna(False)]
    end = pop[f90c] + pd.DateOffset(months=st.dq_window)
    co_in = pop[coc].notna() & (pop[coc] <= end)
    p_pif = pif.loc[pop.index]
    pif_in = ~co_in & p_pif.notna() & (p_pif <= end)
    other = ~co_in & ~pif_in
    out = pd.Series("", index=pop.index, dtype=object)
    out[co_in], out[pif_in] = "charged off", "paid off"
    later = other & pop[coc].notna()
    out[later] = "charged off after the window"
    rest = other & ~later
    if dpdc:
        dpd = pd.to_numeric(pop[dpdc], errors="coerce")
        out[rest & (dpd < 30)] = "cured (current at as-of)"
        out[rest & (dpd >= 30)] = "still delinquent at as-of"
        out[rest & dpd.isna()] = "not charged off (status unknown)"
    else:
        out[rest] = "not charged off (cured or still delinquent)"
    n = len(pop)
    tbl = out.value_counts().rename_axis(f"outcome within {st.dq_window} months of first 90+").reset_index(name="loans")
    tbl["share"] = tbl["loans"] / n if n else np.nan

    lead = months_between(pop.loc[co_in, f90c], pop.loc[co_in, coc])
    co_share = co_in.sum() / n if n else np.nan
    med = float(lead.median()) if len(lead) else np.nan
    skip_share = co_skip / co_total if co_total else np.nan
    use_90 = bool(n) and co_share >= st.proxy_min_co_share and med >= st.proxy_min_lead
    if not n:
        verdict = "Not enough loans reached 90+ early enough to judge. Using charge-off."
    elif use_90:
        verdict = (f"90+ is a good early stand-in: {co_share:.0%} of 90+ loans charged off within "
                   f"{st.dq_window} months, a median {med:.1f} months after first 90+.")
    elif co_share >= st.proxy_min_co_share:
        verdict = f"90+ predicts charge-off ({co_share:.0%}) but arrives only {med:.1f} months earlier (median). Charge-off is simpler."
    else:
        verdict = f"Only {co_share:.0%} of 90+ loans charged off within {st.dq_window} months; too many cure. Use charge-off."
    if co_total:
        verdict += f" {skip_share:.0%} of charge-offs never hit 90+ first."
    measures = pd.DataFrame([
        ("loans that first hit 90+ with enough time to resolve", n),
        (f"share charged off within {st.dq_window} months", co_share),
        ("median months from first 90+ to charge-off", med),
        ("25th / 75th percentile months", f"{lead.quantile(.25):.0f} / {lead.quantile(.75):.0f}" if len(lead) else ""),
        ("charge-offs that never hit 90+ first", co_skip),
        ("as a share of all charge-offs", skip_share),
        ("rule: use 90+ if share charged off is at least", st.proxy_min_co_share),
        ("rule: ...and median lead (months) is at least", st.proxy_min_lead),
        ("VERDICT", verdict),
    ], columns=["measure", "value"])
    return measures, tbl, use_90, verdict


# ============================================================================ 2. early read

def early_read(df: pd.DataFrame, st: Settings, year_col: pd.Series) -> pd.DataFrame:
    H, rows = st.horizon, []
    pooled = []
    for y in sorted(year_col.unique()):
        sub = df[year_col == y]
        km = kaplan_meier(sub["_time"].to_numpy(), sub["_event"].to_numpy())
        final = cum_bad_at(km, H, st.min_at_risk)
        if not final or np.isnan(final):
            continue
        pooled.append(sub)
        row = {"vintage": str(y), "loans": len(sub), f"cum_bad_m{H}": final}
        for m in [c for c in st.checkpoints if c < H]:
            row[f"share_of_m{H}_bads_seen_by_m{m}"] = cum_bad_at(km, m, st.min_at_risk) / final
        rows.append(row)
    if pooled:
        p = pd.concat(pooled)
        km = kaplan_meier(p["_time"].to_numpy(), p["_event"].to_numpy())
        final = cum_bad_at(km, H, st.min_at_risk)
        row = {"vintage": "all seasoned vintages", "loans": len(p), f"cum_bad_m{H}": final}
        for m in [c for c in st.checkpoints if c < H]:
            row[f"share_of_m{H}_bads_seen_by_m{m}"] = cum_bad_at(km, m, st.min_at_risk) / final
        rows.append(row)
    return pd.DataFrame(rows)


# ============================================================================ 3/4. tests + economics

def group_series(df: pd.DataFrame, spec: dict, st: Settings) -> list[tuple[str, pd.Series]]:
    """Returns a list of (comparison label, group labels) for one test definition."""
    o = st.col("orig_date")
    if spec.get("group_by", "").lower() == "year":
        return [("origination year", df[o].dt.year.astype(str))]
    if "split_date" in spec:
        cuts = sorted(pd.to_datetime([x.strip() for x in spec["split_date"].split(",")]))
        edges = [pd.Timestamp.min] + list(cuts) + [pd.Timestamp.max]
        names = []
        for lo, hi in zip(edges[:-1], edges[1:]):
            if lo == pd.Timestamp.min:
                names.append(f"before {hi:%Y-%m-%d}")
            elif hi == pd.Timestamp.max:
                names.append(f"{lo:%Y-%m-%d} on")
            else:
                names.append(f"{lo:%Y-%m-%d} to {hi - pd.Timedelta(days=1):%Y-%m-%d}")
        return [("split " + spec["split_date"], pd.cut(df[o], bins=edges, labels=names, right=False).astype(str))]
    col = spec.get("column")
    if not col:
        raise UserError(f"Test needs group_by = year, split_date, or column: {spec}")
    if col not in df.columns:
        raise UserError(f"Test column '{col}' not found in file.")
    if "cutoffs" in spec:
        x = pd.to_numeric(df[col], errors="coerce")
        out = []
        for c in [float(v) for v in spec["cutoffs"].split(",")]:
            lab = f"{c:g}"
            g = pd.Series(np.where(x < c, f"{col} below {lab}", f"{col} {lab}+"), index=df.index)
            g[x.isna()] = f"{col} missing"
            out.append((f"{col} cutoff {lab}", g))
        return out
    g = df[col].astype(str).where(df[col].notna(), "(blank)")
    if g.nunique() > 15:
        raise UserError(f"Column '{col}' has {g.nunique()} values. Use cutoffs = ... for numeric columns.")
    return [(f"{col}", col + " = " + g)]


def run_tests(df: pd.DataFrame, st: Settings, tests: list):
    rows, curves, charts = [], [], []
    o = st.col("orig_date")
    nr, amt = st.col("net_revenue"), st.col("amount")
    for name, spec in tests:
        pop = df
        if spec.get("orig_from"):
            pop = pop[pop[o] >= pd.Timestamp(spec["orig_from"])]
        if spec.get("orig_to"):
            pop = pop[pop[o] <= pd.Timestamp(spec["orig_to"])]
        window = f"{spec.get('orig_from', 'start')} to {spec.get('orig_to', 'as-of')}"
        for comp, g in group_series(pop, spec, st):
            t, e, gv = pop["_time"].to_numpy(), pop["_event"].to_numpy(), g.to_numpy()
            stat, dof, p = logrank(t, e, gv)
            chart_groups = {}
            for lab in sorted(pd.unique(gv)):
                m = gv == lab
                km = kaplan_meier(t[m], e[m])
                row = {"test": name, "comparison": comp, "orig_window": window, "group": lab,
                       "loans": int(m.sum()), "pct_of_tested": m.mean(), "bads_to_date": int(e[m].sum())}
                for cp in st.checkpoints:
                    row[f"cum_bad_m{cp}"] = cum_bad_at(km, cp, st.min_at_risk)
                row["logrank_p"] = p
                if nr:
                    rev = pd.to_numeric(pop.loc[m, nr], errors="coerce")
                    row["net_revenue_total"] = rev.sum()
                    row["net_revenue_per_loan"] = rev.mean()
                    if amt:
                        a = pd.to_numeric(pop.loc[m, amt], errors="coerce")
                        both = rev.notna() & a.notna()  # same loans top and bottom
                        row["net_revenue_per_1000"] = rev[both].sum() / a[both].sum() * 1000 if a[both].sum() else np.nan
                rows.append(row)
                km.insert(0, "group", lab)
                km.insert(0, "comparison", comp)
                km.insert(0, "test", name)
                curves.append(km)
                chart_groups[lab] = km
            charts.append((f"{name} | {comp}", chart_groups))
    res = pd.DataFrame(rows)
    if len(res):
        res["reading"] = np.where(res["logrank_p"] < 0.05, "groups differ (p < 0.05)",
                                  "difference could be chance (p >= 0.05)")
    return res, (pd.concat(curves, ignore_index=True) if curves else pd.DataFrame()), charts


def model_terms(df: pd.DataFrame, spec: str):
    """'fico per -50, income_gt_sales' -> design columns. 'per k' rescales: HR is per k units (negative = per k lower)."""
    cols, labels, bad = {}, [], []
    for term in [t.strip() for t in spec.split(",") if t.strip()]:
        parts = term.split()
        col = parts[0]
        if col not in df.columns:
            raise UserError(f"Model variable '{col}' not found in file.")
        scale = float(parts[2]) if len(parts) == 3 and parts[1].lower() == "per" else 1.0
        v = df[col]
        if not pd.api.types.is_numeric_dtype(v):
            v = v.astype(str).str.strip().str.upper().map({"Y": 1, "YES": 1, "TRUE": 1, "1": 1,
                                                          "N": 0, "NO": 0, "FALSE": 0, "0": 0})
        v = pd.to_numeric(v, errors="coerce") / scale
        if scale == 1:
            lab = col
        else:
            lab = f"{col} (per {abs(scale):g} {'lower' if scale < 0 else 'higher'})"
        cols[lab] = v
        labels.append(lab)
    return pd.DataFrame(cols, index=df.index), labels


def run_models(df: pd.DataFrame, st: Settings, tests: list) -> pd.DataFrame:
    o, rows = st.col("orig_date"), []
    for name, spec in tests:
        pop = df
        if spec.get("orig_from"):
            pop = pop[pop[o] >= pd.Timestamp(spec["orig_from"])]
        if spec.get("orig_to"):
            pop = pop[pop[o] <= pd.Timestamp(spec["orig_to"])]
        X, labels = model_terms(pop, spec["model"])
        ok = X.notna().all(axis=1)
        dropped = int((~ok).sum())
        if ok.sum() == 0 or pop.loc[ok, "_event"].sum() == 0:
            raise UserError(f"Model test '{name}': no usable rows (check the variables are numeric or Y/N, and that the window has bads).")
        X, t, e = X[ok].to_numpy(float), pop.loc[ok, "_time"].to_numpy(), pop.loc[ok, "_event"].to_numpy()
        alone = {lab: cox_fit(t, e, X[:, [i]]) for i, lab in enumerate(labels)}
        together = cox_fit(t, e, X) if len(labels) > 1 else None
        for i, lab in enumerate(labels):
            b1, s1, p1 = (x[0] for x in alone[lab])
            row = {"test": name, "variable": lab, "loans": int(ok.sum()), "rows_dropped_missing": dropped,
                   "bads": int(e.sum()), "HR_alone": math.exp(b1),
                   "HR_alone_low95": math.exp(b1 - Z95 * s1), "HR_alone_high95": math.exp(b1 + Z95 * s1),
                   "p_alone": p1}
            if together is not None:
                b2, s2, p2 = together[0][i], together[1][i], together[2][i]
                row.update({"HR_together": math.exp(b2), "HR_together_low95": math.exp(b2 - Z95 * s2),
                            "HR_together_high95": math.exp(b2 + Z95 * s2), "p_together": p2})
                lo, hi = row["HR_together_low95"], row["HR_together_high95"]
                if p2 < 0.05:
                    reading = "adds information beyond the others"
                    if lo < 1.1 and hi > 0.91 and (lo > 1 or hi < 1):
                        reading += " (but range is close to 1.0; treat as weak)"
                elif p1 < 0.05:
                    reading = "matters alone, but mostly explained by the others"
                else:
                    reading = "no clear effect"
            else:
                reading = "matters" if p1 < 0.05 else "no clear effect"
            row["reading"] = reading
            rows.append(row)
    return pd.DataFrame(rows)


def draw_charts(charts, st: Settings, folder: Path) -> list[Path]:
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except ImportError:
        return []
    paths = []
    for i, (title, groups) in enumerate(charts):
        if len(groups) > 8:
            continue
        fig, ax = plt.subplots(figsize=(8, 4.2), dpi=110)
        for lab, km in groups.items():
            km = km[km["at_risk"] >= st.min_at_risk]
            ax.step(km["month_on_book"], km["cum_bad_rate"] * 100, where="post", label=lab, linewidth=1.8)
        ax.set_title(title, fontsize=10)
        ax.set_xlabel("Months on book")
        ax.set_ylabel("Cumulative bad rate (%)")
        ax.grid(alpha=0.3)
        ax.legend(frameon=False, fontsize=8)
        for s in ("top", "right"):
            ax.spines[s].set_visible(False)
        fig.tight_layout()
        p = folder / f"_chart_{i}.png"
        fig.savefig(p)
        plt.close(fig)
        paths.append(p)
    return paths


# ============================================================================ output

README = [
    ("What this is", "Loan performance compared at the same month on book, so young loans aren't flattered."),
    ("BadEvent", "Whether 90+ DPD is a good early stand-in for charge-off, and which bad event this run used."),
    ("EarlyRead", "From fully seasoned vintages: share of eventual bads already seen by each month. High share at month 12 = a 12-month comparison of young cohorts is meaningful."),
    ("Tests", "One row per group per comparison. cum_bad_mN = share of the group that went bad within N months on book. Blank = too few loans still being watched at that month."),
    ("logrank_p", "Chance the curves differ this much if the groups truly performed the same. Below 0.05 is the usual bar."),
    ("net_revenue_*", "Only meaningful on seasoned loans (set orig_to). Compare per_1000 across groups, not totals, when group sizes differ."),
    ("Models", "Several factors at once. HR = how much the factor multiplies the monthly chance of going bad (1.0 = no effect, 2.0 = double). 'Alone' vs 'together': if a factor's HR shrinks toward 1.0 once the others are in, it was mostly standing in for them."),
    ("Curves", "Month-by-month detail behind every group, with 95% ranges."),
    ("RunLog", "Input file hash, settings used (verbatim), exclusions, versions."),
]


def write_xlsx(out: Path, sheets: dict, chart_pngs: list[Path]):
    from openpyxl.drawing.image import Image
    from openpyxl.styles import Font, Alignment
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, d in sheets.items():
            if d is None:
                continue
            (d if len(d.columns) else pd.DataFrame({"note": ["none"]})).to_excel(xw, sheet_name=name, index=False)
        if chart_pngs:
            ws = xw.book.create_sheet("Charts")
            for i, p in enumerate(chart_pngs):
                ws.add_image(Image(str(p)), f"A{1 + i * 24}")
        for ws in xw.book.worksheets:
            if ws.title == "Charts":
                continue
            ws.freeze_panes = "A2" if ws.title in ("ReadMe", "RunLog", "BadEvent") else "E2"
            for c in ws[1]:
                c.font = Font(bold=True)
            for col in ws.columns:
                head = str(col[0].value or "")
                w = max([len(head)] + [len(str(c.value)) for c in col[1:80] if c.value is not None]) + 2
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w, 9), 80)
                fmt = None
                if head.startswith("HR_"):
                    fmt = "0.00"
                elif head.startswith("p_"):
                    fmt = "0.0000"
                elif head.startswith(("cum_bad", "share", "pct_")) or head == "share":
                    fmt = "0.0%"
                elif head == "logrank_p":
                    fmt = "0.0000"
                elif head.startswith("net_revenue"):
                    fmt = "#,##0;[Red](#,##0)" if head.endswith(("total", "loan")) else "#,##0.00;[Red](#,##0.00)"
                if fmt:
                    for c in col[1:]:
                        c.number_format = fmt
            if ws.title in ("BadEvent", "ReadMe"):
                for r in ws.iter_rows(min_row=2):
                    for c in r:
                        c.alignment = Alignment(wrap_text=True, vertical="top")
                    if ws.title == "BadEvent" and isinstance(r[1].value, float) and "share" in str(r[0].value):
                        r[1].number_format = "0.0%"


def run(input_path: Path, settings_text: str | None = None, out_path: Path | None = None,
        settings_label: str | None = None) -> dict:
    """One full run. The CLI and the GUI both come through here, so the RunLog is the same either way.

    settings_text None = built-in defaults (column mappings kept only where the file has them).
    Returns a summary dict: out, as_of, rows_used, rows_excluded, bad_event, verdict, tests (DataFrame).
    Raises UserError for anything the person running it can fix.
    """
    input_path = Path(input_path)
    if out_path and Path(out_path).resolve() == input_path.resolve():
        raise UserError("The workbook would be saved over the loan file itself (same file). Pick a different name for the output.")
    st_text = STARTER if settings_text is None else settings_text
    st = Settings(st_text)
    raw = load(input_path)
    if settings_text is None:
        st.cols = {k: v for k, v in st.cols.items() if v in raw.columns}

    df, excluded = prepare(raw, st)
    date_cols = [st.col(k) for k in ("orig_date", "chargeoff_date", "first_90_date", "payoff_date") if st.col(k)]
    as_of = pd.Timestamp(st.as_of) if st.as_of else max(df[c].max() for c in date_cols)
    late = df[st.col("orig_date")] > as_of
    if late.any():
        x = df[late].copy()
        x.insert(0, "exclusion_reason", "originated after as-of")
        excluded = pd.concat([excluded, x])
        df = df[~late]

    # 1. bad event
    measures, outcomes, use_90, verdict = bad_event_check(df, st, as_of)
    if st.bad_event == "earliest" or (st.bad_event == "auto" and use_90):
        if not st.col("first_90_date"):
            raise UserError("bad_event = earliest needs first_90_date mapped.")
        bad_cols, bad_desc = [st.col("first_90_date"), st.col("chargeoff_date")], "earliest of first 90+ or charge-off"
    else:
        bad_cols, bad_desc = [st.col("chargeoff_date")], "charge-off"
    bad_desc += f" (setting: {st.bad_event})"
    if measures is None:
        measures = pd.DataFrame([("VERDICT", verdict)], columns=["measure", "value"])
    measures = pd.concat([measures, pd.DataFrame([("BAD EVENT USED IN THIS RUN", bad_desc)], columns=["measure", "value"])])

    loans = add_time_event(df, st, as_of, bad_cols)

    # 2. early read
    er = early_read(loans, st, loans[st.col("orig_date")].dt.year)
    if er.empty:
        er = pd.DataFrame({"note": [f"No vintage has {st.min_at_risk}+ loans still observed at month {st.horizon}."]})

    # 3/4. tests
    tests = st.tests or [("vintage", {"group_by": "year"})]
    res, curves, charts = run_tests(loans, st, [x for x in tests if "model" not in x[1]])
    models = run_models(loans, st, [x for x in tests if "model" in x[1]])

    out = Path(out_path) if out_path else input_path.with_name(input_path.stem + "_survival.xlsx")
    pngs = draw_charts(charts, st, out.parent)

    import openpyxl
    versions = {"pandas": pd.__version__, "numpy": np.__version__, "openpyxl": openpyxl.__version__}
    try:
        import scipy
        versions["scipy"] = scipy.__version__
    except ImportError:
        versions["scipy"] = "not installed (built-in p-values)"
    reasons = excluded["exclusion_reason"].value_counts().to_dict() if len(excluded) else {}
    log = [("script", f"loan_survival.py v{VERSION}"), ("run_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
           ("input_file", str(input_path.resolve())), ("input_sha256", hashlib.sha256(input_path.read_bytes()).hexdigest()),
           ("settings_file", settings_label or "(built-in defaults)"),
           ("as_of", f"{as_of:%Y-%m-%d}"), ("rows_read", len(raw)), ("rows_used", len(loans)),
           ("rows_excluded", len(excluded)), *[(f"excluded: {k}", v) for k, v in reasons.items()],
           ("bad event used", bad_desc), ("time unit", "whole calendar months on book; censored at payoff or as-of"),
           ("method", "Kaplan-Meier, 95% log(-log) Greenwood ranges, log-rank test"),
           *[(f"version: {k}", v) for k, v in versions.items()],
           ("---- settings used (verbatim) ----", "")]
    log += [("", line) for line in st_text.splitlines()]
    sheets = {"ReadMe": pd.DataFrame(README, columns=["sheet / column", "meaning"]),
              "BadEvent": measures, "BadEventOutcomes": outcomes, "EarlyRead": er, "Tests": res,
              "Models": models if len(models) else None, "Curves": curves, "Excluded": excluded, "RunLog": pd.DataFrame(log, columns=["item", "value"])}
    write_xlsx(out, sheets, pngs)
    for p in pngs:
        p.unlink(missing_ok=True)
    return dict(out=out, as_of=as_of, rows_read=len(raw), rows_used=len(loans), rows_excluded=len(excluded),
                bad_event=bad_desc, verdict=verdict, tests=res)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", type=Path)
    ap.add_argument("--settings", type=Path, help="settings .ini file (default: built-in defaults)")
    ap.add_argument("--write-settings", action="store_true", help="write a starter settings file next to the input and stop")
    ap.add_argument("--out", type=Path)
    a = ap.parse_args(argv)

    if a.write_settings:
        p = a.input.with_name(a.input.stem + "_settings.ini")
        if p.exists():
            sys.exit(f"{p} already exists; not overwriting.")
        p.write_text(STARTER)
        print(f"Wrote {p}. Edit [columns] and the [test: ...] sections, then run with --settings {p.name}")
        return

    try:
        r = run(a.input, a.settings.read_text() if a.settings else None, a.out,
                str(a.settings.resolve()) if a.settings else None)
    except UserError as e:
        sys.exit(str(e))
    res = r["tests"]
    print(f"As-of {r['as_of']:%Y-%m-%d} | rows used {r['rows_used']:,} | excluded {r['rows_excluded']:,}")
    print(f"Bad event: {r['bad_event']}\n  {r['verdict']}")
    show = [c for c in res.columns if c in ("comparison", "group", "loans") or c.startswith("cum_bad") or c == "logrank_p" or c == "net_revenue_per_1000"]
    with pd.option_context("display.width", 200, "display.max_columns", 30):
        print(res[show].to_string(index=False, float_format=lambda v: f"{v:.3f}"))
    print(f"\nWrote {r['out']}")


if __name__ == "__main__":
    main()
