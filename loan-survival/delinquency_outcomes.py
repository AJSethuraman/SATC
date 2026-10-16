#!/usr/bin/env python3
"""
delinquency_outcomes.py - Is 90+ days past due a good early stand-in for charge-off?

Two modes, chosen by the columns in the file:

LOAN mode (one row per loan) - "what happened after a loan first hit 90+?"
    loan_id          [required]
    first_90_date    date the loan first hit 90+ DPD; blank if never   [required]
    chargeoff_date   blank if not charged off                           [required]
    payoff_date      blank if not paid off                              [optional]
    dpd_asof         days past due at the as-of date                    [optional; splits cured vs still delinquent]

SNAPSHOT mode (one row per loan per month-end) - month-to-month roll-rate matrix
    loan_id, month_end, dpd                                             [required]
    status           optional; values listed in --co-values / --pif-values mark charge-off / paid off

Usage:
    python delinquency_outcomes.py loans.xlsx --as-of 2026-09-30
    python delinquency_outcomes.py snapshots.csv --mode snapshot --status-col status --co-values CO,CHGOFF
    python delinquency_outcomes.py loans.csv --window 12 --proxy-min-co-share 0.6 --proxy-min-lead 3

Output: <input>_delinquency.xlsx (Outcomes / Verdict / Excluded / RunLog, or RollRates / RollCounts / RunLog).

The verdict thresholds are YOUR policy choices, set by flags and printed in the RunLog.
The defaults are starting points, not standards.
"""
from __future__ import annotations

import argparse
import hashlib
import sys
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd

VERSION = "1.0"


def load(p: Path) -> pd.DataFrame:
    return pd.read_excel(p) if p.suffix.lower() in (".xlsx", ".xlsm", ".xls") else pd.read_csv(p)


def months_between(start: pd.Series, end: pd.Series) -> pd.Series:
    m = (end.dt.year - start.dt.year) * 12 + (end.dt.month - start.dt.month)
    return m - (end.dt.day < start.dt.day).astype(int)


def need(df, cols):
    missing = [c for c in cols if c and c not in df.columns]
    if missing:
        sys.exit(f"Missing column(s): {missing}. Columns found: {list(df.columns)}")


# ----------------------------------------------------------------------------- loan mode

def loan_mode(df: pd.DataFrame, a, as_of: pd.Timestamp):
    need(df, [a.id_col, a.first90_col, a.co_col])
    has_pif = a.pif_col in df.columns
    has_dpd = a.dpd_col in df.columns
    excluded, rows = [], []
    for c in [a.first90_col, a.co_col] + ([a.pif_col] if has_pif else []):
        parsed = pd.to_datetime(df[c], errors="coerce")
        # a blank date means "never happened"; a date we cannot read must not be mistaken for that
        unreadable = df[c].notna() & (df[c].astype(str).str.strip() != "") & parsed.isna()
        if unreadable.any():
            x = df[unreadable].copy()
            x.insert(0, "exclusion_reason", f"unparseable date in {c}")
            excluded.append(x)
            df, parsed = df[~unreadable], parsed[~unreadable]
        df[c] = parsed
    f90, co = df[a.first90_col], df[a.co_col]
    pif = df[a.pif_col] if has_pif else pd.Series(pd.NaT, index=df.index)

    co_any = co.notna() & (co <= as_of)
    co_without_90 = int((co_any & (f90.isna() | (f90 > co))).sum())  # no 90+ date, or one recorded after the charge-off
    co_total = int(co_any.sum())

    pop = df[f90.notna() & (f90 <= as_of)].copy()
    cutoff = as_of - pd.DateOffset(months=a.window)
    too_recent = pop[a.first90_col] > cutoff
    x = pop[too_recent].copy()
    x.insert(0, "exclusion_reason", f"first 90+ less than {a.window} months before as-of (not enough time to resolve)")
    excluded.append(x)
    pop = pop[~too_recent]
    bad_order = pop[a.co_col].notna() & (pop[a.co_col] < pop[a.first90_col])
    x = pop[bad_order].copy()
    x.insert(0, "exclusion_reason", "charge-off date before first 90+ date")
    excluded.append(x)
    pop = pop[~bad_order]

    end = pop[a.first90_col] + pd.DateOffset(months=a.window)
    p_co, p_pif = pop[a.co_col], (pop[a.pif_col] if has_pif else pd.Series(pd.NaT, index=pop.index))
    co_in = p_co.notna() & (p_co <= end)
    pif_in = ~co_in & p_pif.notna() & (p_pif <= end)
    other = ~co_in & ~pif_in
    outcome = pd.Series("", index=pop.index)
    outcome[co_in] = "charged off"
    outcome[pif_in] = "paid off"
    if has_dpd:
        dpd = pd.to_numeric(pop[a.dpd_col], errors="coerce")
        outcome[other & (dpd < 30)] = "cured (current at as-of)"
        outcome[other & (dpd >= 30)] = "still delinquent at as-of"
        outcome[other & dpd.isna()] = "not charged off (status unknown)"
        if (other & (p_co > end)).any():
            note = "charged off, but after the window"
            outcome[other & p_co.notna() & (p_co > end)] = note
    else:
        outcome[other & p_co.notna()] = "charged off, but after the window"
        outcome[other & p_co.isna()] = "not charged off (cured or still delinquent)"

    n = len(pop)
    counts = outcome.value_counts()
    tbl = pd.DataFrame({"outcome_within_window": counts.index, "loans": counts.values})
    tbl["share"] = tbl["loans"] / n if n else np.nan

    lead = months_between(pop.loc[co_in, a.first90_col], pop.loc[co_in, a.co_col])
    co_share = co_in.sum() / n if n else np.nan
    med_lead = float(lead.median()) if len(lead) else np.nan
    skip_share = co_without_90 / co_total if co_total else np.nan

    good_proxy = (co_share >= a.proxy_min_co_share) and (med_lead >= a.proxy_min_lead)
    if not n:
        verdict = "No loans hit 90+ early enough to evaluate."
    elif good_proxy:
        verdict = ("USE 90+ as the bad event: most 90+ loans charge off, and 90+ arrives months earlier.")
    elif co_share >= a.proxy_min_co_share:
        verdict = ("90+ predicts charge-off well, but the lead time is short. Either event works; charge-off is simpler.")
    else:
        verdict = (f"USE CHARGE-OFF (or '90+ and not cured within N months'): fewer than "
                   f"{a.proxy_min_co_share:.0%} of 90+ loans charge off within {a.window} months.")
    if co_total and skip_share >= a.skip_warn:
        verdict += (f" WARNING: {skip_share:.0%} of charge-offs never hit 90+ first, so 90+ alone misses them. "
                    f"Use 'earliest of 90+ or charge-off' as the bad event.")

    verdict_tbl = pd.DataFrame([
        ("loans that first hit 90+ (with enough time to resolve)", n),
        (f"share charged off within {a.window} months", co_share),
        ("median months from first 90+ to charge-off", med_lead),
        ("25th / 75th percentile months to charge-off",
         f"{lead.quantile(.25):.0f} / {lead.quantile(.75):.0f}" if len(lead) else ""),
        ("charge-offs (all time) that never hit 90+ first", co_without_90),
        ("as a share of all charge-offs", skip_share),
        ("rule: min share charged off to use 90+", a.proxy_min_co_share),
        ("rule: min median lead (months) to use 90+", a.proxy_min_lead),
        ("VERDICT", verdict),
    ], columns=["measure", "value"])

    lead_dist = lead.value_counts().sort_index().rename_axis("months_first90_to_chargeoff").reset_index(name="loans")
    excl = pd.concat(excluded) if excluded else pd.DataFrame()
    print(tbl.to_string(index=False, float_format=lambda v: f"{v:.1%}"))
    print("\n" + verdict)
    return {"Verdict": verdict_tbl, "Outcomes": tbl, "LeadTime": lead_dist, "Excluded": excl}, n, len(excl)


# ----------------------------------------------------------------------------- snapshot mode

BUCKETS = ["Current", "30-59", "60-89", "90-119", "120+"]


def bucket(dpd: pd.Series) -> pd.Series:
    b = pd.cut(dpd, [-np.inf, 29, 59, 89, 119, np.inf], labels=BUCKETS)
    return b.astype(str)


def snapshot_mode(df: pd.DataFrame, a, as_of: pd.Timestamp):
    need(df, [a.id_col, a.month_col, a.dpd_col])
    df[a.month_col] = pd.to_datetime(df[a.month_col], errors="coerce")
    df = df[df[a.month_col] <= as_of].copy()
    df["state"] = bucket(pd.to_numeric(df[a.dpd_col], errors="coerce"))
    # keep a month with no usable DPD in the chain as "Unknown", so its neighbours are not mistaken for a loan leaving the file
    df.loc[~df["state"].isin(BUCKETS), "state"] = "Unknown"
    if a.status_col and a.status_col in df.columns:
        s = df[a.status_col].astype(str).str.strip().str.upper()
        df.loc[s.isin(a.co_values), "state"] = "Charged off"
        df.loc[s.isin(a.pif_values), "state"] = "Paid off"
    df["ym"] = df[a.month_col].dt.year * 12 + df[a.month_col].dt.month
    df = df.sort_values([a.id_col, "ym"]).drop_duplicates([a.id_col, "ym"], keep="last")

    nxt = df.groupby(a.id_col)[["ym", "state"]].shift(-1)
    df["next_state"] = np.where(nxt["ym"] == df["ym"] + 1, nxt["state"], None)
    last_month = df["ym"].max()
    gone = df["next_state"].isna() & (df["ym"] < last_month) & ~df["state"].isin(["Charged off", "Paid off"])
    df.loc[gone, "next_state"] = "Left file (no status)"
    pairs = df[df["next_state"].notna() & ~df["state"].isin(["Charged off", "Paid off"])]

    order_from = BUCKETS
    order_to = BUCKETS + ["Charged off", "Paid off", "Left file (no status)"]
    # a blank or non-numeric DPD has no bucket; leave those transitions out of the count AND the matrix, and list them
    usable = pairs["state"].isin(order_from) & pairs["next_state"].isin(order_to)
    dropped = pairs[~usable].copy()
    pairs = pairs[usable]
    counts = pd.crosstab(pairs["state"], pairs["next_state"]).reindex(index=order_from, columns=order_to, fill_value=0)
    counts = counts.loc[:, (counts.sum() > 0) | counts.columns.isin(order_to[:-1])]
    rates = counts.div(counts.sum(axis=1).replace(0, np.nan), axis=0)
    counts.index.name = rates.index.name = "from (this month)"
    print("Month-to-month roll rates (row = this month, column = next month):")
    print(rates.to_string(float_format=lambda v: f"{v:.1%}"))
    sheets = {"RollRates": rates.reset_index(), "RollCounts": counts.reset_index()}
    if len(dropped):
        dropped.insert(0, "exclusion_reason", "DPD blank or not a number in this month or the next")
        sheets["Excluded"] = dropped.drop(columns=["ym", "next_state"], errors="ignore")
        print(f"\n{len(dropped):,} transition(s) left out because a month's DPD was blank or not a number (see Excluded).")
    return sheets, len(pairs), len(dropped)


# ----------------------------------------------------------------------------- main

def write(out: Path, sheets: dict, runlog: pd.DataFrame):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        for name, d in {**sheets, "RunLog": runlog}.items():
            (d if len(d.columns) else pd.DataFrame({"note": ["none"]})).to_excel(xw, sheet_name=name, index=False)
        for ws in xw.book.worksheets:
            ws.freeze_panes = "B2"
            for c in ws[1]:
                c.font = Font(bold=True)
            for col in ws.columns:
                head = str(col[0].value or "")
                w = max([len(head)] + [len(str(c.value)) for c in col[1:60] if c.value is not None]) + 2
                ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(w, 10), 70)
                if ws.title == "RollRates" and col[0].column > 1 or head == "share":
                    for c in col[1:]:
                        c.number_format = "0.0%"
        if "Verdict" in xw.book.sheetnames:
            ws = xw.book["Verdict"]
            for r in range(2, ws.max_row + 1):
                v = ws.cell(r, 2).value
                if isinstance(v, float) and 0 <= v <= 1 and "share" in str(ws.cell(r, 1).value):
                    ws.cell(r, 2).number_format = "0.0%"


def main(argv=None):
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("input", type=Path)
    p.add_argument("--mode", choices=["auto", "loan", "snapshot"], default="auto")
    p.add_argument("--as-of", help="YYYY-MM-DD (default: latest date in the data)")
    p.add_argument("--id-col", default="loan_id")
    p.add_argument("--first90-col", default="first_90_date")
    p.add_argument("--co-col", default="chargeoff_date")
    p.add_argument("--pif-col", default="payoff_date")
    p.add_argument("--dpd-col", default="dpd_asof", help="loan mode: DPD at as-of; snapshot mode: DPD each month")
    p.add_argument("--month-col", default="month_end")
    p.add_argument("--status-col", default="")
    p.add_argument("--co-values", default="CO,CHGOFF,CHARGED OFF")
    p.add_argument("--pif-values", default="PIF,PAID,PAID OFF")
    p.add_argument("--window", type=int, default=12, help="months after first 90+ to watch for charge-off")
    p.add_argument("--proxy-min-co-share", type=float, default=0.60)
    p.add_argument("--proxy-min-lead", type=float, default=3)
    p.add_argument("--skip-warn", type=float, default=0.20, help="warn if this share of charge-offs never hit 90+")
    p.add_argument("--out", type=Path)
    a = p.parse_args(argv)
    a.co_values = [v.strip().upper() for v in a.co_values.split(",")]
    a.pif_values = [v.strip().upper() for v in a.pif_values.split(",")]

    raw = load(a.input)
    mode = a.mode
    if mode == "auto":
        mode = "snapshot" if a.month_col in raw.columns else "loan"
    if mode == "snapshot" and a.dpd_col == "dpd_asof" and "dpd" in raw.columns:
        a.dpd_col = "dpd"
    date_cols = [c for c in [a.first90_col, a.co_col, a.pif_col, a.month_col] if c in raw.columns]
    as_of = pd.Timestamp(a.as_of) if a.as_of else max(pd.to_datetime(raw[c], errors="coerce").max() for c in date_cols)

    sheets, used, excl = (snapshot_mode if mode == "snapshot" else loan_mode)(raw.copy(), a, as_of)

    out = a.out or a.input.with_name(a.input.stem + "_delinquency.xlsx")
    log = [("script", f"delinquency_outcomes.py v{VERSION}"), ("run_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
           ("input_file", str(a.input.resolve())), ("input_sha256", hashlib.sha256(a.input.read_bytes()).hexdigest()),
           ("mode", mode), ("as_of", f"{as_of:%Y-%m-%d}"), ("rows_read", len(raw)),
           ("rows_used" if mode == "loan" else "month-to-month transitions used", used), ("rows_excluded", excl)]
    if mode == "loan":
        log += [("window_months", a.window), ("rule: min share charged off", a.proxy_min_co_share),
                ("rule: min median lead months", a.proxy_min_lead), ("rule: skip warning share", a.skip_warn)]
    else:
        log += [("buckets (DPD)", "Current 0-29, 30-59, 60-89, 90-119, 120+"),
                ("status col", a.status_col or "(none)"), ("charge-off values", ", ".join(a.co_values)),
                ("paid-off values", ", ".join(a.pif_values))]
    write(out, sheets, pd.DataFrame(log, columns=["item", "value"]))
    print(f"\nWrote {out}")


if __name__ == "__main__":
    main()
