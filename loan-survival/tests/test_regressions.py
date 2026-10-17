"""Regression tests for the Codex review of PR 445. Each one was red before its fix."""
import sys
from argparse import Namespace
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
import delinquency_outcomes as do  # noqa: E402
import loan_survival as ls  # noqa: E402

AS_OF = pd.Timestamp("2026-09-30")


def _args(**kw):
    base = dict(id_col="loan_id", first90_col="first_90_date", co_col="chargeoff_date", pif_col="payoff_date",
                dpd_col="dpd_asof", month_col="month_end", status_col="", co_values=["CO"], pif_values=["PIF"],
                window=12, proxy_min_co_share=0.6, proxy_min_lead=3, skip_warn=0.2)
    base.update(kw)
    return Namespace(**base)


def _measure(sheets, name):
    return sheets["Verdict"].set_index("measure").loc[name, "value"]


def test_charge_off_with_90_after_it_counts_as_skipping_90():
    df = pd.DataFrame({
        "loan_id": ["A", "B", "C"],
        "first_90_date": ["2024-01-31", "2025-06-30", None],
        "chargeoff_date": ["2024-05-31", "2025-03-31", "2024-02-29"],  # B: 90+ date recorded after the charge-off
        "payoff_date": [None, None, None], "dpd_asof": [0, 0, 0]})
    sheets, _, _ = do.loan_mode(df, _args(), AS_OF)
    assert _measure(sheets, "charge-offs (all time) that never hit 90+ first") == 2
    assert abs(_measure(sheets, "as a share of all charge-offs") - 2 / 3) < 1e-12


def test_malformed_loan_dates_are_excluded_not_read_as_blank():
    df = pd.DataFrame({
        "loan_id": ["A", "B", "C"],
        "first_90_date": ["2024-01-31", "2024-01-31", "2024-01-31"],
        "chargeoff_date": ["2024-05-31", "not a date", None],
        "payoff_date": [None, None, None], "dpd_asof": [0, 0, 0]})
    sheets, used, excl = do.loan_mode(df, _args(), AS_OF)
    assert used == 2 and excl == 1
    ex = sheets["Excluded"]
    assert list(ex["loan_id"]) == ["B"] and "chargeoff_date" in ex["exclusion_reason"].iloc[0]


def test_economics_per_1000_uses_rows_with_both_values():
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv")
    yr = pd.to_datetime(raw.orig_date).dt.year
    i_no_amt = raw.index[yr == 2022][0]
    i_no_rev = raw.index[yr == 2022][1]
    raw.loc[i_no_amt, "amount"] = np.nan   # has revenue, no amount: would inflate the ratio
    raw.loc[i_no_rev, "net_revenue"] = np.nan  # has amount, no revenue: would dilute it
    both = raw.net_revenue.notna() & raw.amount.notna()
    ref = raw[both & (yr == 2022)]
    expected = ref.net_revenue.sum() / ref.amount.sum() * 1000
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "b.csv"
        raw.to_csv(p, index=False)
        r = ls.run(p, (ROOT / "sample" / "settings_example.ini").read_text(), Path(tmp) / "o.xlsx")
    t = r["tests"]
    got = t[(t.comparison == "origination year") & (t.group.astype(str) == "2022")]["net_revenue_per_1000"].iloc[0]
    assert abs(got - expected) < 1e-9, (got, expected)


def test_snapshot_transitions_used_matches_the_matrix_denominator():
    rows = []
    for loan, dpds in (("L1", [0, None, 0]), ("L2", [0, 0, 0])):
        for m, d in zip(pd.date_range("2024-01-31", periods=3, freq="ME"), dpds):
            rows.append(dict(loan_id=loan, month_end=m, dpd=d))
    sheets, used, dropped = do.snapshot_mode(pd.DataFrame(rows), _args(dpd_col="dpd"), pd.Timestamp("2024-03-31"))
    counts = sheets["RollCounts"].drop(columns="from (this month)")
    assert used == int(counts.to_numpy().sum()) == 2
    assert dropped == 2 and "Excluded" in sheets


# ---- second Codex review of PR 445 (reviewed 23a0e90)

def test_snapshot_unreadable_month_is_listed_and_does_not_fake_a_loan_leaving():
    rows = []
    for loan, months in (("L1", ["2024-01-31", "garbage", "2024-03-31"]), ("L2", ["2024-01-31", "2024-02-29", "2024-03-31"])):
        for m in months:
            rows.append(dict(loan_id=loan, month_end=m, dpd=0))
    sheets, used, dropped = do.snapshot_mode(pd.DataFrame(rows), _args(dpd_col="dpd"), pd.Timestamp("2024-03-31"))
    counts = sheets["RollCounts"].drop(columns="from (this month)")
    assert "Left file (no status)" not in counts.columns or int(counts["Left file (no status)"].sum()) == 0
    assert used == int(counts.to_numpy().sum()) == 2          # only L2's two transitions
    assert set(sheets["Excluded"]["loan_id"]) == {"L1"} and dropped == 3  # all three of L1's rows, with the reason


def test_column_names_with_a_percent_sign_survive_the_settings_file():
    text = "[columns]\norig_date = Open %\nchargeoff_date = CO date\n\n[test: util]\ncolumn = Utilization %\ncutoffs = 50\n"
    st = ls.Settings(text)
    assert st.cols["orig_date"] == "Open %" and st.tests[0][1]["column"] == "Utilization %"


def test_output_path_cannot_be_the_input_file(tmp_path):
    import pytest
    src = tmp_path / "loans.xlsx"
    pd.read_csv(ROOT / "sample" / "fake_book.csv").to_excel(src, index=False)
    before = src.read_bytes()
    with pytest.raises(ls.UserError, match="same file"):
        ls.run(src, None, src)
    assert src.read_bytes() == before


# ---- third Codex review of PR 445 (reviewed 17d95f0)

def test_blank_loan_ids_are_not_duplicates():
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv")
    raw.loc[:2, "loan_id"] = np.nan
    st = ls.Settings((ROOT / "sample" / "settings_example.ini").read_text())
    df, excluded = ls.prepare(raw, st)
    assert len(df) == len(raw) and len(excluded) == 0


def test_header_with_a_space_hash_survives_the_settings_file():
    import loan_survival_gui as g
    state = g.new_state()
    state["cols"] = {"orig_date": "Open Date", "chargeoff_date": "CO Date", "loan_id": "Loan #"}
    state["tests"] = [{"name": "t", "spec": {"column": "Rate ;fixed", "cutoffs": "5"}}]
    text = g.build_settings_text(state)
    assert ls.Settings(text).cols["loan_id"] == "Loan #" and ls.Settings(text).tests[0][1]["column"] == "Rate ;fixed"
    assert g.parse_settings_text(text)["cols"]["loan_id"] == "Loan #"


def test_inline_comment_after_a_number_is_refused_in_words():
    import pytest
    with pytest.raises(ls.UserError, match="own line"):
        ls.Settings("[settings]\nmin_at_risk = 30   ; stop showing a month below this\n")


def test_run_does_not_touch_files_already_in_the_output_folder(tmp_path):
    (tmp_path / "_chart_0.png").write_bytes(b"not mine")
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv")
    src = tmp_path / "book.csv"
    raw.to_csv(src, index=False)
    ls.run(src, (ROOT / "sample" / "settings_example.ini").read_text(), tmp_path / "out.xlsx")
    assert (tmp_path / "_chart_0.png").read_bytes() == b"not mine"


def test_model_column_with_numbers_stored_as_text_keeps_the_numbers():
    df = pd.DataFrame({"fico": pd.Series(["650", "700", "Y"], dtype=object)})
    X, _ = ls.model_terms(df, "fico per -50")
    assert X.iloc[0, 0] == 650 / -50 and X.iloc[1, 0] == 700 / -50 and X.iloc[2, 0] == 1 / -50


def test_xls_without_its_reader_says_what_to_do(monkeypatch, tmp_path):
    import pytest

    def no_xlrd(*a, **k):
        raise ImportError("Missing optional dependency 'xlrd'")
    monkeypatch.setattr(pd, "read_excel", no_xlrd)
    f = tmp_path / "old.xls"
    f.write_bytes(b"x")
    with pytest.raises(ls.UserError, match="xlsx"):
        ls.load(f)


# ---- fourth Codex review of PR 445 (reviewed 0f83681)

def test_snapshot_rows_with_no_loan_id_are_listed_not_dropped():
    rows = [dict(loan_id=l, month_end=m, dpd=0) for l in ("L1", None)
            for m in pd.date_range("2024-01-31", periods=3, freq="ME")]
    sheets, used, dropped = do.snapshot_mode(pd.DataFrame(rows), _args(dpd_col="dpd"), pd.Timestamp("2024-03-31"))
    assert used == 2 and dropped == 3
    assert "no loan id" in sheets["Excluded"]["exclusion_reason"].iloc[0].lower()


def _two_loans_with_a_charge_off_after_as_of():
    cols = "[columns]\nloan_id = loan_id\norig_date = orig_date\nchargeoff_date = chargeoff_date\nfirst_90_date = first_90_date\ndpd_asof = dpd_asof\n"
    st = ls.Settings(cols)
    df = pd.DataFrame({
        "loan_id": ["A", "B", "C"], "orig_date": pd.to_datetime(["2022-01-01"] * 3),
        "first_90_date": pd.to_datetime(["2023-01-31"] * 3),
        # A: charged off in the window. B: charged off AFTER as-of (not observable yet). C: after the window but before as-of
        "chargeoff_date": pd.to_datetime(["2023-03-31", "2025-01-31", "2024-06-30"]), "dpd_asof": [0, 0, 0]})
    return st, df, pd.Timestamp("2024-12-31")


def test_charge_off_after_as_of_is_not_reported_as_an_event():
    st, df, as_of = _two_loans_with_a_charge_off_after_as_of()
    _, tbl, _, _ = ls.bad_event_check(df, st, as_of)
    got = dict(zip(tbl.iloc[:, 0], tbl["loans"]))
    assert got.get("charged off after the window", 0) == 1       # C only; B had not happened by as-of
    assert got.get("cured (current at as-of)", 0) == 1            # B read as not charged off


def test_standalone_charge_off_after_as_of_is_not_reported_as_an_event():
    st, df, as_of = _two_loans_with_a_charge_off_after_as_of()
    df = df.rename(columns={"first_90_date": "first_90_date"}).assign(payoff_date=pd.NaT)
    sheets, _, _ = do.loan_mode(df.copy(), _args(), as_of)
    got = dict(zip(sheets["Outcomes"].iloc[:, 0], sheets["Outcomes"]["loans"]))
    assert got.get("charged off, but after the window", 0) == 1 and got.get("cured (current at as-of)", 0) == 1


def test_model_column_names_may_contain_spaces():
    import loan_survival_gui as g
    df = pd.DataFrame({"Credit Score": [650.0, 700.0], "Has Flag": ["Y", "N"]})
    X, labels = ls.model_terms(df, "Credit Score per -50, Has Flag")
    assert X.iloc[0, 0] == 650 / -50 and X.iloc[1, 1] == 0
    assert g.split_model("Credit Score per -50, Has Flag") == [("Credit Score", "50", "lower"), ("Has Flag", "", "higher")]
    assert g.model_term("Credit Score", "50", "lower") == "Credit Score per -50"


def test_all_missing_revenue_total_is_blank_not_zero():
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv")
    raw.loc[pd.to_datetime(raw.orig_date).dt.year == 2025, "net_revenue"] = np.nan
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "b.csv"
        raw.to_csv(p, index=False)
        r = ls.run(p, (ROOT / "sample" / "settings_example.ini").read_text(), Path(tmp) / "o.xlsx")
    t = r["tests"]
    g25 = t[(t.comparison == "origination year") & (t.group.astype(str) == "2025")].iloc[0]
    assert pd.isna(g25["net_revenue_total"]) and pd.isna(g25["net_revenue_per_1000"])


# ---- fifth Codex review of PR 445 (reviewed 9e93bf1)

def test_model_columns_offered_by_the_gui_match_what_the_engine_can_use():
    import loan_survival_gui as g
    df = pd.DataFrame({"Score": pd.Series(["650", "700", "N/A", "720", "680"] * 4, dtype=object),
                       "Revenue per 100": [1.5, 2.5, 3.5, 4.5, 5.5] * 4, "Plain": [1.5, 2.5, 3.5, 4.5, 5.5] * 4,
                       "Words": ["a", "b", "c", "d", "e"] * 4})
    offered = g.model_columns(df)
    assert "Score" in offered and "Plain" in offered
    assert "Revenue per 100" not in offered and "Words" not in offered


def test_dates_in_mixed_formats_are_all_read():
    s = pd.Series(["2024-01-31", "02/28/2024", "March 5, 2024", None])
    import loan_survival_gui as g
    assert g.date_problems(s) == (0, 3, 1)
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv", dtype=str)
    raw.loc[0, "orig_date"] = "02/28/2022"
    st = ls.Settings((ROOT / "sample" / "settings_example.ini").read_text())
    df, excluded = ls.prepare(raw, st)
    assert len(excluded) == 0 and df.loc[0, "orig_date"] == pd.Timestamp("2022-02-28")


def test_the_same_factor_twice_is_refused_and_found_by_full_name():
    import pytest
    import loan_survival_gui as g
    df = pd.DataFrame({"fico": [1.0, 2.0]})
    with pytest.raises(ls.UserError, match="twice"):
        ls.model_terms(df, "fico per -50, fico")
    assert g.model_has_column(["Credit Score per -50"], "Credit Score") and not g.model_has_column(["Credit Score per -50"], "Credit")


def test_a_comparison_that_could_not_be_tested_says_so():
    import tempfile
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv")
    text = "[columns]\nloan_id=loan_id\norig_date=orig_date\nchargeoff_date=chargeoff_date\n\n[test: one year only]\ngroup_by = year\norig_from = 2023-01-01\norig_to = 2023-12-31\n"
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "b.csv"
        raw.to_csv(p, index=False)
        r = ls.run(p, text, Path(tmp) / "o.xlsx")
    t = r["tests"]
    assert len(t) == 1 and pd.isna(t["logrank_p"].iloc[0]) and "not testable" in t["reading"].iloc[0]
    assert "never observed over the same months" not in t["reading"].iloc[0]
    assert "no bad event happened while more than one group was still being followed" in t["reading"].iloc[0]


def test_excluded_sheet_shows_the_unreadable_text_not_a_blank():
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv", dtype=str)
    raw.loc[3, "orig_date"] = "TBD"
    st = ls.Settings((ROOT / "sample" / "settings_example.ini").read_text())
    _, excluded = ls.prepare(raw, st)
    assert list(excluded["orig_date"]) == ["TBD"] and "unparseable" in excluded["exclusion_reason"].iloc[0]


# ---- sixth Codex review of PR 445 (reviewed 7dbf378)

def test_standalone_loan_mode_leaves_out_repeated_loan_ids():
    df = pd.DataFrame({
        "loan_id": ["A", "A", "B"], "first_90_date": ["2024-01-31"] * 3,
        "chargeoff_date": ["2024-05-31", "2024-05-31", None], "payoff_date": [None] * 3, "dpd_asof": [0, 0, 0]})
    sheets, used, excl = do.loan_mode(df, _args(), AS_OF)
    assert used == 2 and excl == 1
    assert "duplicate loan id" in sheets["Excluded"]["exclusion_reason"].iloc[0]


def test_cox_fit_that_cannot_settle_says_so_instead_of_reporting_a_number():
    import pytest
    rng = np.random.default_rng(3)
    t = rng.integers(1, 30, 400)
    e = (rng.random(400) < 0.4).astype(int)
    with pytest.raises(ls.NotEstimable):
        ls.cox_fit(t, e, e[:, None].astype(float))            # a factor that perfectly separates bad from not bad


def test_model_sheet_marks_an_unestimable_factor_not_a_real_looking_ratio():
    rng = np.random.default_rng(3)
    n = 400
    df = pd.DataFrame({"orig_date": pd.Timestamp("2022-01-01"), "_time": rng.integers(1, 30, n),
                       "_event": (rng.random(n) < 0.4).astype(int), "noise": rng.normal(size=n)})
    df["sep"] = df["_event"]
    st = ls.Settings("[columns]\norig_date = orig_date\nchargeoff_date = x\n")
    out = ls.run_models(df, st, [("m", {"model": "noise, sep"})])
    row = out.set_index("variable")
    assert row.loc["sep", "reading"].startswith("not estimable") and pd.isna(row.loc["sep", "HR_alone"])
    assert not pd.isna(row.loc["noise", "HR_alone"])


def _payoff_before_90_frame():
    st = ls.Settings("[columns]\nloan_id = loan_id\norig_date = orig_date\nchargeoff_date = chargeoff_date\n"
                     "first_90_date = first_90_date\npayoff_date = payoff_date\ndpd_asof = dpd_asof\n")
    df = pd.DataFrame({
        "loan_id": ["A", "B"], "orig_date": pd.to_datetime(["2022-01-01"] * 2),
        "first_90_date": pd.to_datetime(["2023-01-31"] * 2), "chargeoff_date": pd.to_datetime(["2023-03-31", None]),
        "payoff_date": pd.to_datetime([None, "2022-12-31"]),       # B was 'paid off' before it was ever 90+: inconsistent
        "dpd_asof": [0, 0]})
    return st, df


def test_payoff_before_first_90_is_left_out_of_the_outcome_audit():
    st, df = _payoff_before_90_frame()
    measures, tbl, _, _ = ls.bad_event_check(df, st, pd.Timestamp("2026-09-30"))
    assert int(measures.set_index("measure").loc["loans that first hit 90+ with enough time to resolve", "value"]) == 1
    assert "paid off" not in set(tbl.iloc[:, 0])


def test_standalone_payoff_before_first_90_is_excluded():
    st, df = _payoff_before_90_frame()
    sheets, used, excl = do.loan_mode(df.copy(), _args(), pd.Timestamp("2026-09-30"))
    assert used == 1 and "payoff date before first 90+" in sheets["Excluded"]["exclusion_reason"].iloc[0]


# ---- seventh Codex review of PR 445 (reviewed 2bfc980)

def test_text_from_the_source_that_looks_like_a_formula_stays_text_in_the_workbook(tmp_path):
    import openpyxl
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv", dtype=str)
    raw.loc[3, "orig_date"] = "=TODAY()"
    src = tmp_path / "book.csv"
    raw.to_csv(src, index=False)
    out = tmp_path / "o.xlsx"
    ls.run(src, (ROOT / "sample" / "settings_example.ini").read_text(), out)
    ws = openpyxl.load_workbook(out)["Excluded"]
    cells = [c for row in ws.iter_rows() for c in row if c.value == "=TODAY()"]
    assert cells and all(c.data_type == "s" for c in cells)
    for sheet in openpyxl.load_workbook(out).worksheets:
        assert not [c for row in sheet.iter_rows() for c in row if c.data_type == "f"], sheet.title


# ---- eighth Codex review of PR 445 (reviewed 2956769)

def test_two_snapshot_rows_for_one_loan_and_month_are_refused_not_picked_between():
    def run(order):
        rows = [dict(loan_id="L1", month_end="2024-01-31", dpd=0), dict(loan_id="L1", month_end="2024-01-31", dpd=130),
                dict(loan_id="L1", month_end="2024-02-29", dpd=0), dict(loan_id="L2", month_end="2024-01-31", dpd=0),
                dict(loan_id="L2", month_end="2024-02-29", dpd=0)]
        df = pd.DataFrame([rows[i] for i in order])
        return do.snapshot_mode(df, _args(dpd_col="dpd"), pd.Timestamp("2024-02-29"))
    a, used_a, ex_a = run([0, 1, 2, 3, 4])
    b, used_b, ex_b = run([1, 0, 2, 3, 4])
    assert a["RollCounts"].equals(b["RollCounts"]) and used_a == used_b == 1     # same answer whichever row came first
    assert set(a["Excluded"]["loan_id"]) == {"L1"} and "same month" in a["Excluded"]["exclusion_reason"].iloc[0]


def _charge_off_after_payoff_frame():
    st = ls.Settings("[columns]\nloan_id = loan_id\norig_date = orig_date\nchargeoff_date = chargeoff_date\n"
                     "first_90_date = first_90_date\npayoff_date = payoff_date\ndpd_asof = dpd_asof\n")
    df = pd.DataFrame({
        "loan_id": ["A", "B"], "orig_date": pd.to_datetime(["2022-01-01"] * 2),
        "first_90_date": pd.to_datetime(["2023-01-31"] * 2),
        "chargeoff_date": pd.to_datetime(["2023-03-31", "2023-06-30"]),
        "payoff_date": pd.to_datetime([None, "2023-04-30"]),       # B: paid off in April, 'charged off' in June
        "dpd_asof": [0, 0]})
    return st, df


def test_charge_off_dated_after_payoff_is_left_out_of_the_outcome_audit():
    st, df = _charge_off_after_payoff_frame()
    measures, tbl, _, _ = ls.bad_event_check(df, st, pd.Timestamp("2026-09-30"))
    assert int(measures.set_index("measure").loc["loans that first hit 90+ with enough time to resolve", "value"]) == 1
    assert dict(zip(tbl.iloc[:, 0], tbl["loans"])) == {"charged off": 1}


def test_standalone_charge_off_dated_after_payoff_is_excluded():
    st, df = _charge_off_after_payoff_frame()
    sheets, used, excl = do.loan_mode(df.copy(), _args(), pd.Timestamp("2026-09-30"))
    assert used == 1 and "charge-off date after payoff" in sheets["Excluded"]["exclusion_reason"].iloc[0]


# ---- ninth Codex review of PR 445 (reviewed 254e21a)

def _skip_totals_frame():
    st = ls.Settings("[columns]\nloan_id = loan_id\norig_date = orig_date\nchargeoff_date = chargeoff_date\n"
                     "first_90_date = first_90_date\npayoff_date = payoff_date\ndpd_asof = dpd_asof\n")
    df = pd.DataFrame({
        "loan_id": ["A", "B"], "orig_date": pd.to_datetime(["2022-01-01"] * 2),
        "first_90_date": pd.to_datetime(["2023-01-31", None]),
        "chargeoff_date": pd.to_datetime(["2023-03-31", "2023-06-30"]),
        "payoff_date": pd.to_datetime([None, "2023-04-30"]),       # B: no 90+ date, and a charge-off dated after its payoff
        "dpd_asof": [0, 0]})
    return st, df


def test_charge_off_after_payoff_does_not_count_toward_the_skipped_90_share():
    st, df = _skip_totals_frame()
    measures, _, _, _ = ls.bad_event_check(df, st, pd.Timestamp("2026-09-30"))
    m = measures.set_index("measure")["value"]
    assert float(m["as a share of all charge-offs"]) == 0.0 and int(m["charge-offs that never hit 90+ first"]) == 0


def test_standalone_charge_off_after_payoff_does_not_count_toward_the_skipped_90_share():
    st, df = _skip_totals_frame()
    sheets, used, excl = do.loan_mode(df.copy(), _args(), pd.Timestamp("2026-09-30"))
    m = sheets["Verdict"].set_index("measure")["value"]
    assert float(m["as a share of all charge-offs"]) == 0.0 and excl == 1


def test_a_vintage_with_no_bad_loans_stays_in_the_early_read():
    st = ls.Settings("[columns]\norig_date = orig_date\nchargeoff_date = x\n")
    n = 100
    df = pd.DataFrame({"orig_date": pd.to_datetime(["2021-06-01"] * n + ["2022-06-01"] * n),
                       "_time": [40] * n + [40] * n, "_event": [0] * n + [1] * 10 + [0] * 90})
    er = ls.early_read(df, st, df["orig_date"].dt.year)
    assert list(er["vintage"]) == ["2021", "2022", "all seasoned vintages"]
    assert er.loc[0, "cum_bad_m36"] == 0.0 and er["loans"].iloc[-1] == 2 * n
    assert er.filter(like="share_of").iloc[0].isna().all()           # timing ratio is undefined when there are no bads
    all_good = ls.early_read(df.assign(_event=0), st, df["orig_date"].dt.year)
    assert len(all_good) == 3


def test_snapshot_last_month_survives_a_loan_set_aside_for_duplicate_rows():
    rows = [dict(loan_id="L1", month_end=m, dpd=0) for m in ("2024-01-31", "2024-02-29")]
    rows += [dict(loan_id="Z", month_end=m, dpd=0) for m in ("2024-01-31", "2024-02-29", "2024-03-31", "2024-03-31")]
    rows += [dict(loan_id="L2", month_end=m, dpd=0) for m in ("2024-01-31", "2024-02-29")]
    # March is only in loan Z, which is set aside for its duplicate row. The extract still runs to March, so L1 and L2 stopped early.
    sheets, _, _ = do.snapshot_mode(pd.DataFrame(rows), _args(dpd_col="dpd"), pd.Timestamp("2024-03-31"))
    counts = sheets["RollCounts"].set_index("from (this month)")
    assert counts.loc["Current", "Left file (no status)"] == 2


# ---- tenth Codex review of PR 445 (reviewed 03e685b)

def test_snapshot_last_month_counts_a_row_that_has_no_loan_id():
    rows = [dict(loan_id="L1", month_end=m, dpd=0) for m in ("2024-01-31", "2024-02-29")]
    rows += [dict(loan_id="L2", month_end=m, dpd=0) for m in ("2024-01-31", "2024-02-29")]
    rows += [dict(loan_id=None, month_end="2024-03-31", dpd=0)]       # the only March row has no id
    sheets, _, _ = do.snapshot_mode(pd.DataFrame(rows), _args(dpd_col="dpd"), pd.Timestamp("2024-03-31"))
    counts = sheets["RollCounts"].set_index("from (this month)")
    assert counts.loc["Current", "Left file (no status)"] == 2        # L1 and L2 stopped in February; the extract runs to March


def test_a_failed_write_leaves_the_existing_workbook_untouched(tmp_path, monkeypatch):
    import pytest
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv")
    src, out = tmp_path / "book.csv", tmp_path / "report.xlsx"
    raw.to_csv(src, index=False)
    out.write_bytes(b"last month's report")

    real = ls.write_xlsx

    def dies_halfway(path, sheets, pngs):
        real(path, {k: v for k, v in list(sheets.items())[:2]}, pngs)
        raise RuntimeError("closed mid-write")
    monkeypatch.setattr(ls, "write_xlsx", dies_halfway)
    with pytest.raises(RuntimeError):
        ls.run(src, (ROOT / "sample" / "settings_example.ini").read_text(), out)
    assert out.read_bytes() == b"last month's report"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["book.csv", "report.xlsx"]   # no half-written leftovers


# ---- eleventh Codex review of PR 445 (reviewed ffbb58c)

def test_one_column_cannot_fill_two_roles():
    import pytest
    import loan_survival_gui as g
    raw = pd.read_csv(ROOT / "sample" / "fake_book.csv")
    st = ls.Settings("[columns]\norig_date = orig_date\nchargeoff_date = orig_date\n")
    with pytest.raises(ls.UserError, match="both"):
        ls.prepare(raw, st)
    state = g.new_state()
    state["cols"] = {"orig_date": "Open", "chargeoff_date": "Open"}
    assert any("both" in p for p in g.validate_state(state))


def test_a_missing_month_in_the_middle_is_not_an_exit():
    rows = [dict(loan_id="L1", month_end=m, dpd=0) for m in ("2024-01-31", "2024-03-31")]            # no February
    rows += [dict(loan_id="L2", month_end=m, dpd=0) for m in ("2024-01-31", "2024-02-29", "2024-03-31")]
    sheets, used, excl = do.snapshot_mode(pd.DataFrame(rows), _args(dpd_col="dpd"), pd.Timestamp("2024-03-31"))
    counts = sheets["RollCounts"].set_index("from (this month)")
    assert "Left file (no status)" not in counts.columns or counts["Left file (no status)"].sum() == 0
    assert used == 2
    ex = sheets["Excluded"]
    assert list(ex["loan_id"]) == ["L1"] and "missing" in ex["exclusion_reason"].iloc[0] and excl == 1


def test_runlog_records_the_mappings_that_were_really_used(tmp_path):
    src = tmp_path / "small.csv"
    pd.read_csv(ROOT / "sample" / "fake_book.csv")[["loan_id", "orig_date", "chargeoff_date"]].to_csv(src, index=False)
    out = tmp_path / "o.xlsx"
    ls.run(src, None, out)
    log = pd.read_excel(out, sheet_name="RunLog").fillna("")
    text = "\n".join(log["value"].astype(str))
    start = text.index("[data]")
    logged = ls.Settings(text[start:])
    assert set(logged.cols) == {"loan_id", "orig_date", "chargeoff_date"}   # not first_90_date, net_revenue and the rest
    ls.run(src, text[start:], tmp_path / "again.xlsx")                        # and the logged text can be run again as it is


# ---- twelfth Codex review of PR 445 (reviewed 15be01f)

def test_two_model_factors_with_the_same_output_name_are_refused():
    import pytest
    df = pd.DataFrame({"score": [1.0, 2.0], "score (per 2 higher)": [3.0, 4.0]})
    with pytest.raises(ls.UserError, match="same name"):
        ls.model_terms(df, "score per 2, score (per 2 higher)")


def test_standalone_xls_without_its_reader_says_what_to_do(monkeypatch, tmp_path):
    import pytest

    def no_xlrd(*a, **k):
        raise ImportError("Missing optional dependency 'xlrd'")
    monkeypatch.setattr(pd, "read_excel", no_xlrd)
    f = tmp_path / "old.xls"
    f.write_bytes(b"x")
    with pytest.raises(SystemExit, match="xlsx"):
        do.load(f)


def test_records_after_a_terminal_status_set_the_loan_aside():
    rows = [dict(loan_id="L1", month_end=m, dpd=d, status=s) for m, d, s in
            (("2024-01-31", 0, "OPEN"), ("2024-02-29", 0, "CO"), ("2024-03-31", 0, "OPEN"), ("2024-04-30", 0, "OPEN"))]
    rows += [dict(loan_id="L2", month_end=m, dpd=0, status="OPEN") for m in ("2024-01-31", "2024-02-29", "2024-03-31", "2024-04-30")]
    sheets, used, excl = do.snapshot_mode(pd.DataFrame(rows), _args(dpd_col="dpd", status_col="status"), pd.Timestamp("2024-04-30"))
    assert used == 3                                           # L2's three transitions only
    ex = sheets["Excluded"]
    assert set(ex["loan_id"]) == {"L1"} and "after" in ex["exclusion_reason"].iloc[0] and excl == 4


# ---- thirteenth Codex review of PR 445 (reviewed adc4d5c)

def test_log_rank_with_no_overlap_between_groups_is_untestable_not_p_of_one():
    t = np.r_[np.full(30, 1), np.full(30, 10)]
    e = np.r_[np.zeros(30, int), np.ones(30, int)]          # A is censored at month 1; B's events come at month 10
    g = np.r_[np.full(30, "A"), np.full(30, "B")]
    stat, dof, p = ls.logrank(t, e, g)
    assert np.isnan(p) and dof == 0


def test_log_rank_degrees_of_freedom_follow_the_rank_of_the_covariance():
    rng = np.random.default_rng(5)
    n = 200
    t = np.r_[rng.integers(2, 30, n), rng.integers(2, 30, n), np.full(n, 1)]   # C is only ever observed at month 1
    e = np.r_[(rng.random(n) < .5).astype(int), (rng.random(n) < .5).astype(int), np.zeros(n, int)]
    g = np.r_[np.full(n, "A"), np.full(n, "B"), np.full(n, "C")]
    stat, dof, p = ls.logrank(t, e, g)
    assert dof == 1 and not np.isnan(p)                      # three labels, but only one testable difference


def test_standalone_loan_mode_leaves_out_rows_with_no_loan_id():
    df = pd.DataFrame({
        "loan_id": [None, "B", "C"], "first_90_date": ["2024-01-31"] * 3,
        "chargeoff_date": ["2024-05-31", None, None], "payoff_date": [None] * 3, "dpd_asof": [0, 0, 0]})
    sheets, used, excl = do.loan_mode(df, _args(), AS_OF)
    assert used == 2 and excl == 1 and "no loan id" in sheets["Excluded"]["exclusion_reason"].iloc[0]


# ---- a column that is one column against another (the firm, 17 Oct 2026: "income > sales ... could be anything > anything")

_BOOK = ROOT / "sample" / "fake_book.csv"
_COLS = "[columns]\nloan_id=loan_id\norig_date=orig_date\nchargeoff_date=chargeoff_date\n"


def _run_book(raw, extra, name="flag"):
    import tempfile
    with tempfile.TemporaryDirectory() as tmp:
        p = Path(tmp) / "b.csv"
        raw.to_csv(p, index=False)
        return ls.run(p, _COLS + extra, Path(tmp) / "o.xlsx"), pd.read_excel(Path(tmp) / "o.xlsx", sheet_name=None)


def test_a_flag_built_from_two_columns_is_compared_like_a_column():
    raw = pd.read_csv(_BOOK)
    raw["amt_ok"] = np.where(raw["amount"] >= 50000, "Yes", "No")
    r, _ = _run_book(raw, "\n[test: derived]\ncolumn = amount >= 50000\n")
    r2, _ = _run_book(raw, "\n[test: explicit]\ncolumn = amt_ok\n")
    a = r["tests"].set_index("group")["loans"].to_dict()
    b = r2["tests"].set_index("group")["loans"].to_dict()
    assert sorted(a.values()) == sorted(b.values()) and len(a) == 2
    assert all(g.startswith("amount >= 50000 = ") for g in a)


def test_a_flag_can_compare_two_columns():
    raw = pd.read_csv(_BOOK)
    r, _ = _run_book(raw, "\n[test: two]\ncolumn = net_revenue > amount\n")
    want = int((raw["net_revenue"] > raw["amount"]).sum())
    got = {g.split(" = ")[1]: n for g, n in r["tests"].set_index("group")["loans"].items()}
    assert got.get("Yes", 0) == want


def test_a_flag_in_a_model_matches_the_same_flag_as_a_column():
    raw = pd.read_csv(_BOOK)
    raw["f700"] = np.where(raw["fico"] >= 700, "Y", "N")
    _, s1 = _run_book(raw, "\n[test: m]\nmodel = fico >= 700\n")
    _, s2 = _run_book(raw, "\n[test: m]\nmodel = f700\n")
    a, b = s1["Models"].iloc[0], s2["Models"].iloc[0]
    assert a["variable"] == "fico >= 700" and abs(a["HR_alone"] - b["HR_alone"]) < 1e-12 and a["loans"] == b["loans"]


def test_a_real_column_with_the_flags_own_name_wins():
    raw = pd.read_csv(_BOOK)
    raw["amount > fico"] = np.where(raw.index % 2 == 0, "left", "right")      # nothing like amount > fico
    r, _ = _run_book(raw, "\n[test: t]\ncolumn = amount > fico\n")
    assert set(g.split(" = ")[1] for g in r["tests"]["group"]) == {"left", "right"}


def test_a_blank_on_either_side_leaves_the_loan_in_blank_not_in_no():
    raw = pd.read_csv(_BOOK)
    raw.loc[raw.index[:5], "fico"] = np.nan
    r, _ = _run_book(raw, "\n[test: t]\ncolumn = fico >= 700\n")
    got = {g.split(" = ")[1]: n for g, n in r["tests"].set_index("group")["loans"].items()}
    assert got.get("(blank)") == 5


def test_a_flag_on_a_missing_or_text_column_is_refused_in_words():
    raw = pd.read_csv(_BOOK)
    for bad in ("nope > 5", "income_gt_sales > 5"):
        try:
            _run_book(raw, f"\n[test: t]\ncolumn = {bad}\n")
            assert False, bad
        except ls.UserError as e:
            assert "not found" in str(e), str(e)
